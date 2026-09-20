"""core/files.py: 保存（分割読み sha256）と Excel の事前チェック。"""
import hashlib
import io
import os
import re
import time
import zipfile
from pathlib import Path

import pytest
from flask import Flask
from openpyxl import Workbook
from werkzeug.datastructures import FileStorage

from core import files
from core.files import UploadError, precheck_excel, remove_upload, save_upload, upload_path


@pytest.fixture
def core_app(tmp_path):
    app = Flask(__name__)
    app.config.update(TESTING=True, UPLOAD_DIR=tmp_path / "uploads")
    return app


def _storage(data: bytes, name: str) -> FileStorage:
    return FileStorage(stream=io.BytesIO(data), filename=name)


def _xlsx(path):
    wb = Workbook()
    wb.active["A1"] = "管理No"
    wb.save(path)
    return path


def test_save_upload_hash_and_remove(core_app, monkeypatch):
    monkeypatch.setattr(files, "CHUNK_SIZE", 7)  # 分割読みを確かめる
    data = "管理No,設備\nTR-1,CMP-101\n".encode("cp932") * 50
    with core_app.app_context():
        stored = save_upload(_storage(data, "C:\\fakepath\\故障履歴.csv"), "tables", {".csv", "xlsx"}, 10_000)
        assert stored.file_name == "故障履歴.csv"
        assert stored.stored_path.startswith("tables/") and stored.stored_path.endswith(".csv")
        assert stored.file_hash == hashlib.sha256(data).hexdigest()
        assert stored.size == len(data)
        path = upload_path(stored.stored_path)
        assert path.read_bytes() == data
        remove_upload(stored.stored_path)
        assert not path.exists()
        remove_upload(stored.stored_path)  # 2回目も例外にしない


def test_save_upload_rejects(core_app):
    with core_app.app_context():
        with pytest.raises(UploadError, match="保存し直して"):
            save_upload(_storage(b"x", "old.xls"), "forms", {".xlsx"}, 100)
        with pytest.raises(UploadError, match="大きすぎます"):
            save_upload(_storage(b"x" * 101, "big.csv"), "tables", {".csv"}, 100)
        with pytest.raises(UploadError, match="空です"):
            save_upload(_storage(b"", "empty.csv"), "tables", {".csv"}, 100)
        with pytest.raises(UploadError, match="選択"):
            save_upload(_storage(b"x", ""), "tables", {".csv"}, 100)
        # 失敗したファイルは残さない
        folder = core_app.config["UPLOAD_DIR"] / "tables"
        assert not folder.exists() or not any(folder.iterdir())
        with pytest.raises(UploadError):
            upload_path("../outside.txt")


def test_precheck_accepts_normal_xlsx(tmp_path):
    precheck_excel(_xlsx(tmp_path / "ok.xlsx"))


def test_precheck_ole_header(tmp_path):
    path = tmp_path / "locked.xlsx"
    path.write_bytes(bytes.fromhex("D0CF11E0A1B11AE1") + b"\x00" * 504)
    with pytest.raises(UploadError, match="パスワード付き"):
        precheck_excel(path)


def test_precheck_not_zip(tmp_path):
    path = tmp_path / "text.xlsx"
    path.write_text("管理No,設備", encoding="utf-8")
    with pytest.raises(UploadError, match="読み込めません"):
        precheck_excel(path)


def test_precheck_xlsb(tmp_path):
    path = tmp_path / "book.xlsx"
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("xl/workbook.bin", b"\x83\x01\x00")
    with pytest.raises(UploadError, match="xlsb"):
        precheck_excel(path)


def test_precheck_strict(tmp_path):
    path = tmp_path / "strict.xlsx"
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("xl/workbook.xml", '<?xml version="1.0"?><workbook '
                    'xmlns="http://purl.oclc.org/ooxml/spreadsheetml/main"><sheets/></workbook>')
    with pytest.raises(UploadError, match="Strict"):
        precheck_excel(path)


def test_precheck_zip_bomb_ratio(tmp_path):
    path = tmp_path / "bomb.xlsx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("xl/workbook.xml", "<workbook/>")
        zf.writestr("xl/worksheets/sheet1.xml", b"\x00" * (files.ZIP_RATIO_MIN_BYTES + 1024))
    assert path.stat().st_size < 1024 * 1024
    with pytest.raises(UploadError, match="圧縮率"):
        precheck_excel(path)


def test_precheck_zip_limits(tmp_path, monkeypatch):
    path = tmp_path / "large.xlsx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("xl/workbook.xml", "<workbook/>")
        zf.writestr("xl/worksheets/sheet1.xml", b"a" * 5000)
        zf.writestr("xl/worksheets/sheet2.xml", b"b" * 5000)
    monkeypatch.setattr(files, "ZIP_MAX_PART", 4000)
    with pytest.raises(UploadError, match="1つの部品"):
        precheck_excel(path)
    monkeypatch.setattr(files, "ZIP_MAX_PART", 6000)
    monkeypatch.setattr(files, "ZIP_MAX_TOTAL", 8000)
    with pytest.raises(UploadError, match="合計"):
        precheck_excel(path)


def test_remove_upload_empties_a_locked_file_and_reports_it(core_app, monkeypatch):
    """他のプロセスに掴まれていて消せないときは、例外を出さずに中身を空にして False を返す。

    「ダウンロードしたらこのPCから消えます」（design.md 3.3）と言い切っているので、ファイル名が残っても
    元の Excel/CSV の中身は残さない（空になったファイルは次の起動時に片付く）。
    """
    with core_app.app_context():
        stored = save_upload(_storage("社外秘の中身".encode("utf-8") * 10, "掴まれたブック.xlsx"),
                             "documents", {".xlsx"}, 10_000)
        path = upload_path(stored.stored_path)

        def locked(self, missing_ok=False):
            raise PermissionError(32, "別のプロセスが使用中です")

        monkeypatch.setattr(Path, "unlink", locked)
        monkeypatch.setattr(files, "REMOVE_RETRY_WAIT", 0)
        assert remove_upload(stored.stored_path) is False   # 例外を出さず、消せなかったことを返す
        assert path.exists() and path.read_bytes() == b""   # 中身は残さない

        monkeypatch.undo()
        assert remove_upload(stored.stored_path) is True
        assert not path.exists()


def test_remove_orphan_uploads(core_app):
    """DB から参照されていない取り込み済みファイルだけを片付ける。"""
    with core_app.app_context():
        used = save_upload(_storage(b"a" * 100, "使用中.xlsx"), "documents", {".xlsx"}, 10_000)
        orphan = save_upload(_storage(b"b" * 100, "残骸.xlsx"), "tables", {".xlsx"}, 10_000)
        fresh = save_upload(_storage(b"c" * 100, "保存したばかり.xlsx"), "documents", {".xlsx"}, 10_000)
        other = upload_path("documents/メモ.txt")
        other.write_text("save_upload が作った名前ではないファイル", encoding="utf-8")
        old = time.time() - 3600
        for stored in (used, orphan):
            os.utime(upload_path(stored.stored_path), (old, old))

        removed = files.remove_orphan_uploads(core_app.config["UPLOAD_DIR"], [used.stored_path])

        assert removed == 1
        assert upload_path(used.stored_path).exists()
        assert not upload_path(orphan.stored_path).exists()
        assert other.exists()
        # できたばかりのファイルは、別に起動しているアプリが DB の行を作る前かもしれないので消さない
        assert upload_path(fresh.stored_path).exists()


# ---- 結合セルの面積（openpyxl で開く前に止める） ------------------------------------------------

def _xlsx_with_merge(path, ref: str, *, pad: int = 0):
    """ふつうの xlsx を作り、シートの XML に <mergeCell ref="..."> を書き足す（pad: その前に入れる空白の量）。"""
    src = _xlsx(path.with_name("src_" + path.name))
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                merge = f'{" " * pad}<mergeCells count="1"><mergeCell ref="{ref}"/></mergeCells>'.encode()
                data = data.replace(b"</sheetData>", b"</sheetData>" + merge, 1)
                if b"<sheetData/>" in data:
                    data = data.replace(b"<sheetData/>", b"<sheetData/>" + merge, 1)
            zout.writestr(info, data)
    return path


def test_precheck_refuses_a_whole_sheet_merge_before_openpyxl(tmp_path):
    import time
    path = _xlsx_with_merge(tmp_path / "merge_full.xlsx", "A1:XFD1048576")
    started = time.monotonic()
    with pytest.raises(UploadError, match="結合セルの範囲が大きすぎます"):
        precheck_excel(path)
    assert time.monotonic() - started < 2


def test_precheck_accepts_whole_row_and_whole_column_merges(tmp_path):
    precheck_excel(_xlsx_with_merge(tmp_path / "row.xlsx", "A1:XFD1"))
    precheck_excel(_xlsx_with_merge(tmp_path / "col.xlsx", "A1:A1048576"))


def test_precheck_for_forms_refuses_many_whole_row_merges_quickly(tmp_path):
    """帳票は画面を開くたびに通常モードで開き直す。行全体の結合 120 個（約200万セル）は1回に約16秒かかるので断る。"""
    many = _xlsx_with_merge(tmp_path / "rows120.xlsx", "A1:XFD1")
    many.write_bytes(_rewrite_merges(many, [f"A{r}:XFD{r}" for r in range(1, 121)]))
    precheck_excel(many)   # 一覧表（読み取り専用で開き、結合を展開しない）はこれまでどおり通す
    started = time.monotonic()
    with pytest.raises(UploadError, match="結合セルの範囲が大きすぎます"):
        precheck_excel(many, max_merged=files.FORM_MAX_MERGED_CELLS)
    assert time.monotonic() - started < 2
    with pytest.raises(UploadError, match="結合セル"):   # 列全体の結合（罫線付きなら1分以上かかる）も
        precheck_excel(_xlsx_with_merge(tmp_path / "col.xlsx", "A1:A1048576"), max_merged=files.FORM_MAX_MERGED_CELLS)
    # 行全体の結合が数個の帳票は通す
    few = _xlsx_with_merge(tmp_path / "rows12.xlsx", "A1:XFD1")
    few.write_bytes(_rewrite_merges(few, [f"A{r}:XFD{r}" for r in range(1, 13)]))
    precheck_excel(few, max_merged=files.FORM_MAX_MERGED_CELLS)


def _rewrite_merges(path, refs):
    """_xlsx_with_merge で作ったブックの結合範囲を refs に置き換えた中身を返す。"""
    out = io.BytesIO()
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                merges = "".join(f'<mergeCell ref="{ref}"/>' for ref in refs)
                data = re.sub(rb"<mergeCells\b.*?</mergeCells>",
                              f'<mergeCells count="{len(refs)}">{merges}</mergeCells>'.encode(), data, flags=re.S)
            zout.writestr(info, data)
    return out.getvalue()


def test_precheck_finds_a_merge_split_across_read_chunks(tmp_path, monkeypatch):
    monkeypatch.setattr(files, "CHUNK_SIZE", 64)   # タグがチャンクの境目で切れるようにする
    for pad in range(0, 64, 7):
        path = _xlsx_with_merge(tmp_path / f"split{pad}.xlsx", "C1:Z200000", pad=pad)
        with pytest.raises(UploadError, match="結合セル"):
            precheck_excel(path)


def test_precheck_counts_each_merge_once(tmp_path, monkeypatch):
    monkeypatch.setattr(files, "CHUNK_SIZE", 64)
    monkeypatch.setattr(files, "MAX_MERGED_CELLS", 100)
    precheck_excel(_xlsx_with_merge(tmp_path / "exact.xlsx", "A1:J10"))   # ちょうど100セルは通す


def test_uploads_with_a_whole_sheet_merge_are_refused_in_japanese(app, client, tmp_path):
    path = _xlsx_with_merge(tmp_path / "merge_full.xlsx", "A1:XFD1048576")
    for url in ("/forms/upload", "/tables/upload"):
        res = client.post(url, data={"file": (io.BytesIO(path.read_bytes()), "結合.xlsx")},
                          content_type="multipart/form-data", follow_redirects=True)
        assert "結合セルの範囲が大きすぎます" in res.get_data(as_text=True), url
    assert [p for p in Path(app.config["UPLOAD_DIR"]).rglob("*") if p.is_file()] == []


def test_precheck_reports_a_damaged_compressed_part(tmp_path):
    """圧縮データが壊れている（zlib.error）ときも UploadError（500 にしない）。"""
    path = _xlsx(tmp_path / "ok.xlsx")
    data = bytearray(path.read_bytes())
    with zipfile.ZipFile(path) as zf:
        info = zf.getinfo("xl/workbook.xml")
    start = info.header_offset + 30 + len(info.filename.encode()) + len(info.extra)
    for i in range(start, start + min(info.compress_size, 64)):
        data[i] ^= 0xFF
    broken = tmp_path / "broken.xlsx"
    broken.write_bytes(bytes(data))
    with pytest.raises(UploadError, match="壊れている"):
        precheck_excel(broken)


def test_sheetless_workbook_is_refused_for_both_flows(app, client, tmp_path):
    """シートの無いブックは帳票でも一覧表でも受け付けず、アップロードしたファイルも残さない。"""
    import re

    path = _xlsx(tmp_path / "s.xlsx")
    out = io.BytesIO()
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(out, "w") as z:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/workbook.xml":
                data = re.sub(rb"<sheets>.*?</sheets>", b"<sheets/>", data, flags=re.S)
            z.writestr(item, data)
    sheetless = tmp_path / "sheetless.xlsx"
    sheetless.write_bytes(out.getvalue())
    with pytest.raises(UploadError, match="シートがないブック"):
        precheck_excel(sheetless)
    precheck_excel(path)   # ふつうのブックは通る
    for url in ("/forms/upload", "/tables/upload"):
        res = client.post(url, data={"file": (io.BytesIO(out.getvalue()), "シートなし.xlsx")},
                          content_type="multipart/form-data", follow_redirects=True)
        assert "シートがないブックです" in res.get_data(as_text=True), url
    assert [p for p in Path(app.config["UPLOAD_DIR"]).rglob("*") if p.is_file()] == []


def _xlsx_with_cells(path, count: int):
    """ふつうの xlsx のシートに、書式だけのセル <c/> を count 個書き足す（圧縮すると小さい）。"""
    src = _xlsx(path.with_name("src_" + path.name))
    rows = b'<row r="9999">' + b'<c s="0"/>' * count + b"</row>"
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                data, n = re.subn(rb"<sheetData\s*/>", b"<sheetData>" + rows + b"</sheetData>", data, 1)
                if not n:
                    data = data.replace(b"</sheetData>", rows + b"</sheetData>", 1)
            zout.writestr(info, data)
    return path


def test_precheck_refuses_too_many_cells_in_a_small_file(tmp_path):
    """圧縮すると小さいが展開すると大量のセルがあるブックは、openpyxl で開く前に断る。"""
    path = _xlsx_with_cells(tmp_path / "many.xlsx", files.MAX_CELLS + 1)
    assert path.stat().st_size < 200 * 1024
    with pytest.raises(UploadError, match="セル数が上限"):
        precheck_excel(path)
    precheck_excel(path, max_cells=files.MAX_CELLS + 100)   # 上限を上げれば通る


def test_precheck_counts_cells_split_across_read_chunks_once(tmp_path, monkeypatch):
    monkeypatch.setattr(files, "CHUNK_SIZE", 64)
    path = _xlsx_with_cells(tmp_path / "exact.xlsx", 500)
    with zipfile.ZipFile(path) as zf:
        existing = len(re.findall(rb"<(?:\w+:)?c(?=[\s/>])", zf.read("xl/worksheets/sheet1.xml")))
    precheck_excel(path, max_cells=existing)          # ちょうど上限は通す（二重に数えない）
    with pytest.raises(UploadError, match="セル数が上限"):
        precheck_excel(path, max_cells=existing - 1)  # 1つでも多ければ断る（数え落とさない）


def _xlsx_with_sheet_xml(path, edit, *, rename: str | None = None):
    """ふつうの xlsx のシートの XML を edit(bytes) で書き換える（rename: シートの部品の名前を変える）。"""
    src = _xlsx(path.with_name("src_" + path.name))
    sheet = "xl/worksheets/sheet1.xml"
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            name = info.filename
            if name == sheet:
                data = edit(data)
                name = rename or name
            elif rename and name in ("xl/_rels/workbook.xml.rels", "[Content_Types].xml"):
                data = data.replace(b"worksheets/sheet1.xml", rename.removeprefix("xl/").encode())
            zout.writestr(name, data)
    return path


def _add_merge(ref_attr: bytes):
    def edit(data: bytes) -> bytes:
        merge = b"<mergeCells count=\"1\"><mergeCell " + ref_attr + b"/></mergeCells>"
        if b"<sheetData/>" in data:
            return data.replace(b"<sheetData/>", b"<sheetData/>" + merge, 1)
        return data.replace(b"</sheetData>", b"</sheetData>" + merge, 1)
    return edit


def _add_cells(row: bytes):
    """<row> を1つシートに書き足す（sheetData が空要素でも中身があっても）。"""
    def edit(data: bytes) -> bytes:
        data, n = re.subn(rb"<sheetData\s*/>", b"<sheetData>" + row + b"</sheetData>", data, count=1)
        return data if n else data.replace(b"</sheetData>", row + b"</sheetData>", 1)
    return edit


def test_precheck_reads_a_sheet_part_whatever_its_name(tmp_path):
    """シートの部品の名前は workbook.xml.rels で決まる（.xml で終わらなくても openpyxl は開く）。"""
    import openpyxl

    small = _xlsx_with_sheet_xml(tmp_path / "small.xlsx", _add_merge(b'ref="C3:E5"'),
                                 rename="xl/worksheets/sheet1.dat")
    precheck_excel(small)
    assert [str(r) for r in openpyxl.load_workbook(small).active.merged_cells.ranges] == ["C3:E5"]
    big = _xlsx_with_sheet_xml(tmp_path / "big.xlsx", _add_merge(b'ref="C3:XFD1048576"'),
                               rename="xl/worksheets/sheet1.dat")
    with pytest.raises(UploadError, match="結合セルの範囲が大きすぎます"):
        precheck_excel(big)
    cells = _xlsx_with_sheet_xml(tmp_path / "cells.xlsx", _add_cells(b'<row r="9">' + b'<c s="0"/>' * 50 + b"</row>"),
                                 rename="xl/worksheets/sheet1.dat")
    with pytest.raises(UploadError, match="セル数が上限"):
        precheck_excel(cells, max_cells=40)


def test_precheck_reads_merge_refs_the_way_an_xml_parser_does(tmp_path):
    """= の前後の空白・文字参照（&#88; = X）も、XML として正しい書き方なので openpyxl は読む。数え落とさない。"""
    import openpyxl

    for i, attr in enumerate((b'ref = "C3:XFD1048576"', b'ref="C3:&#88;FD1048576"', b"ref\n=\n'C3:XFD1048576'")):
        path = _xlsx_with_sheet_xml(tmp_path / f"m{i}.xlsx", _add_merge(attr))
        with pytest.raises(UploadError, match="結合セルの範囲が大きすぎます"):
            precheck_excel(path)
    small = _xlsx_with_sheet_xml(tmp_path / "small.xlsx", _add_merge(b'ref = "C3:&#69;5"'))
    precheck_excel(small)
    assert [str(r) for r in openpyxl.load_workbook(small).active.merged_cells.ranges] == ["C3:E5"]


def test_precheck_counts_cells_under_any_namespace_prefix(tmp_path):
    """<x.y:c> のような接頭辞でも、SpreadsheetML の名前空間の <c> ならセルとして数える。"""
    main = b"http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    row = b'<row r="9" xmlns:x.y-z="' + main + b'">' + b'<x.y-z:c s="0"/>' * 50 + b"</row>"
    path = _xlsx_with_sheet_xml(tmp_path / "prefixed.xlsx", _add_cells(row))
    with pytest.raises(UploadError, match="セル数が上限"):
        precheck_excel(path, max_cells=40)
    precheck_excel(path, max_cells=100)


def test_save_upload_and_precheck_messages_carry_no_file_name_or_class_name(app, tmp_path):
    """flash に入るメッセージにはファイル名も例外の種類名も入れない（design.md 3.3）。"""
    from werkzeug.datastructures import FileStorage

    with app.app_context():
        for data, allowed in ((b"x", {".csv"}), (b"", {".xlsx"})):
            with pytest.raises(UploadError) as err:
                save_upload(FileStorage(io.BytesIO(data), "秘密の名前.xlsx"), "documents", allowed, 1024)
            assert "秘密の名前" not in str(err.value)
    broken = tmp_path / "broken.xlsx"
    broken.write_bytes(b"PK\x03\x04" + b"\x00" * 64)
    with pytest.raises(UploadError) as err:
        precheck_excel(broken)
    assert "BadZipFile" not in str(err.value) and "壊れている" in str(err.value)


def test_precheck_refuses_entity_definitions(tmp_path):
    """DTD で実体を定義して展開させる細工（正しいブックには無い）は、数える前に断る。"""
    def edit(data: bytes) -> bytes:
        return b'<!DOCTYPE worksheet [<!ENTITY a "aaaaaaaaaa">]>' + data.split(b"?>", 1)[-1]

    with pytest.raises(UploadError, match="不正なファイル"):
        precheck_excel(_xlsx_with_sheet_xml(tmp_path / "entity.xlsx", edit))


_XDR = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
_DRAWING_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing"


def _xlsx_sharing_one_drawing(path, sheets: int, anchors: int, absolute_target: bool = False):
    """多数のシートが同じ1つの描画部品（大量の図形）を参照するブック（圧縮後は小さい）。"""
    wb = Workbook()
    wb.active.title = "S1"
    for i in range(2, sheets + 1):
        wb.create_sheet(f"S{i}")
    wb.save(path)
    anchor = ('<xdr:absoluteAnchor><xdr:pos x="0" y="0"/><xdr:ext cx="1" cy="1"/>'
              '<xdr:clientData/></xdr:absoluteAnchor>')
    drawing = f'<?xml version="1.0"?><xdr:wsDr xmlns:xdr="{_XDR}">{anchor * anchors}</xdr:wsDr>'
    target = "/xl/drawings/drawing1.xml" if absolute_target else "../drawings/drawing1.xml"
    rels = ('<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            f'<Relationship Id="rIdD" Type="{_DRAWING_REL}" Target="{target}"/></Relationships>')
    src = path.with_suffix(".src.xlsx")
    path.rename(src)
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", item.filename):
                text = data.decode("utf-8").replace(
                    "<worksheet ", '<worksheet xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" ', 1)
                data = text.replace("</worksheet>", '<drawing r:id="rIdD"/></worksheet>').encode("utf-8")
            zout.writestr(item, data)
        for i in range(1, sheets + 1):
            zout.writestr(f"xl/worksheets/_rels/sheet{i}.xml.rels", rels)
        zout.writestr("xl/drawings/drawing1.xml", drawing)
    src.unlink()
    return path


def test_precheck_refuses_many_sheets_sharing_one_large_drawing_quickly(tmp_path):
    """図形の多い描画を多数のシートが共有するブックは、openpyxl が開く前に断る（開くたびに数十秒かかるため）。"""
    path = _xlsx_sharing_one_drawing(tmp_path / "shared.xlsx", sheets=20, anchors=40_000)
    assert path.stat().st_size < 200_000
    started = time.monotonic()
    with pytest.raises(UploadError, match="図形や画像の数が多すぎる"):
        precheck_excel(path)
    assert time.monotonic() - started < 2


def test_precheck_multiplies_drawing_anchors_by_referencing_sheets(tmp_path, monkeypatch):
    monkeypatch.setattr(files, "MAX_DRAWING_ANCHORS", 100)
    precheck_excel(_xlsx_sharing_one_drawing(tmp_path / "ok.xlsx", sheets=4, anchors=25))       # 100 は通す
    with pytest.raises(UploadError, match="図形や画像"):
        precheck_excel(_xlsx_sharing_one_drawing(tmp_path / "ng.xlsx", sheets=5, anchors=25))   # 125
    with pytest.raises(UploadError, match="図形や画像"):                                          # 絶対パスの Target も数える
        precheck_excel(_xlsx_sharing_one_drawing(tmp_path / "abs.xlsx", sheets=5, anchors=25, absolute_target=True))


def test_precheck_limits_drawing_bytes_times_referencing_sheets(tmp_path, monkeypatch):
    monkeypatch.setattr(files, "MAX_DRAWING_BYTES", 10_000)
    with pytest.raises(UploadError, match="図形や画像"):
        precheck_excel(_xlsx_sharing_one_drawing(tmp_path / "big.xlsx", sheets=10, anchors=20))


def test_precheck_accepts_a_workbook_with_a_few_images(tmp_path):
    from openpyxl.drawing.image import Image as XLImage
    from PIL import Image as PILImage

    png = tmp_path / "p.png"
    PILImage.new("RGB", (4, 4), "red").save(png)
    wb = Workbook()
    for i in range(3):
        ws = wb.active if i == 0 else wb.create_sheet(f"S{i}")
        ws["A1"] = "管理No"
        ws.add_image(XLImage(str(png)), "C3")
        ws.add_image(XLImage(str(png)), "E5")
    path = tmp_path / "images.xlsx"
    wb.save(path)
    precheck_excel(path)


_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_WORKSHEET_REL = f"{_R_NS}/worksheet"
_COMMENTS_REL = f"{_R_NS}/comments"
_PKG_RELS = "http://schemas.openxmlformats.org/package/2006/relationships"


def _rels(entries: str) -> str:
    return f'<?xml version="1.0"?><Relationships xmlns="{_PKG_RELS}">{entries}</Relationships>'


def _xlsx_sharing_one_sheet_part(path, sheets: int, merge_ref: str):
    """複数の <sheet> が同じ1つのシートの部品を指すブック（Excel は作らないが、手で書けば openpyxl は開く）。

    openpyxl は <sheet> の数だけその部品を読み直すので、結合セルもセルもその回数だけ作られる。
    """
    sheet_xml = (f'<?xml version="1.0"?><worksheet xmlns="{_MAIN_NS}"><sheetData/>'
                 f'<mergeCells count="1"><mergeCell ref="{merge_ref}"/></mergeCells></worksheet>')
    tabs = "".join(f'<sheet name="S{i}" sheetId="{i}" r:id="rId{i}"/>' for i in range(1, sheets + 1))
    links = "".join(f'<Relationship Id="rId{i}" Type="{_WORKSHEET_REL}" Target="worksheets/sheet1.xml"/>'
                    for i in range(1, sheets + 1))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("_rels/.rels", _rels(""))
        zf.writestr("xl/workbook.xml", f'<?xml version="1.0"?><workbook xmlns="{_MAIN_NS}" xmlns:r="{_R_NS}">'
                                       f"<sheets>{tabs}</sheets></workbook>")
        zf.writestr("xl/_rels/workbook.xml.rels", _rels(links))
        zf.writestr("xl/worksheets/sheet1.xml", sheet_xml)
    return path


def test_precheck_multiplies_one_sheet_part_by_the_sheets_that_share_it(tmp_path):
    """同じシートの部品を多数の <sheet> が指すブックは、部品1つ分だけ数えると上限をすり抜ける（1シートあたり約4秒）。"""
    one = _xlsx_sharing_one_sheet_part(tmp_path / "one.xlsx", sheets=1, merge_ref="A1:XFD1")   # 16,384 セル
    precheck_excel(one, max_merged=files.FORM_MAX_MERGED_CELLS)          # 1シート分は通す
    many = _xlsx_sharing_one_sheet_part(tmp_path / "many.xlsx", sheets=20, merge_ref="A1:XFD1")
    assert many.stat().st_size < 10_000
    started = time.monotonic()
    with pytest.raises(UploadError, match="結合セルの範囲が大きすぎます"):
        precheck_excel(many, max_merged=files.FORM_MAX_MERGED_CELLS)     # 20シート分＝約33万セル
    assert time.monotonic() - started < 2


def test_precheck_multiplies_one_comments_part_by_the_sheets_that_share_it(tmp_path):
    """1つのコメントの部品を多数のシートが参照すると、コメントの範囲のセルもシートの数だけ作られる。"""
    comments = (f'<?xml version="1.0"?><comments xmlns="{_MAIN_NS}"><commentList>'
                f'<comment ref="A1:D2500" authorId="0"/></commentList></comments>')   # 10,000 セル
    sheets = 6
    tabs = "".join(f'<sheet name="S{i}" sheetId="{i}" r:id="rId{i}"/>' for i in range(1, sheets + 1))
    links = "".join(f'<Relationship Id="rId{i}" Type="{_WORKSHEET_REL}" Target="worksheets/sheet{i}.xml"/>'
                    for i in range(1, sheets + 1))
    path = tmp_path / "comments.xlsx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("_rels/.rels", _rels(""))
        zf.writestr("xl/workbook.xml", f'<?xml version="1.0"?><workbook xmlns="{_MAIN_NS}" xmlns:r="{_R_NS}">'
                                       f"<sheets>{tabs}</sheets></workbook>")
        zf.writestr("xl/_rels/workbook.xml.rels", _rels(links))
        for i in range(1, sheets + 1):
            zf.writestr(f"xl/worksheets/sheet{i}.xml", f'<?xml version="1.0"?><worksheet xmlns="{_MAIN_NS}">'
                                                       "<sheetData/></worksheet>")
            zf.writestr(f"xl/worksheets/_rels/sheet{i}.xml.rels",
                        _rels(f'<Relationship Id="rIdC" Type="{_COMMENTS_REL}" Target="../comments1.xml"/>'))
        zf.writestr("xl/comments1.xml", comments)
    with pytest.raises(UploadError, match="コメントの範囲が大きすぎます"):
        precheck_excel(path)   # 6シート分＝60,000 セル（MAX_LINKED_CELLS 50,000 超え）


def test_precheck_refuses_a_small_file_that_expands_to_hundreds_of_megabytes(tmp_path, monkeypatch):
    """展開後が大きくなりすぎるブックは断る（部品ごとの圧縮率の確認は、小さい部品を見ないためすり抜ける）。"""
    monkeypatch.setattr(files, "ZIP_RATIO_MIN_BYTES", 200_000)   # 本物は 16MB。同じ形を小さく作る
    path = tmp_path / "many_small_parts.xlsx"
    part = b"<si/>" * 20_000                                     # 100,000 バイト（1部品ずつは上の値未満）
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("xl/workbook.xml", f'<?xml version="1.0"?><workbook xmlns="{_MAIN_NS}"><sheets/></workbook>')
        for i in range(30):
            zf.writestr(f"xl/pad{i}.xml", b"<sst>" + part + b"</sst>")
    assert path.stat().st_size < 20_000                          # 20KB が 3MB に展開される（150倍）
    with pytest.raises(UploadError, match="圧縮率"):
        precheck_excel(path)
