"""画面の経路: 列の対応づけの保存、帳票の種類の見本と状態、
帳票の元ファイルのダウンロード、Excel のシートの選択。"""
import io

import openpyxl

from models import database as db
from tables import store
from tests.tables_helpers import (columns_payload, editor_body, panel, panel_html, save_columns, save_layout,
                                  save_source, upload_bytes, upload_csv, wait_import_job)


# ---- 列の対応づけ（取り込みごとに決める。設定だけを開く画面は無い） ------------------------------------

def test_saving_the_columns_decides_everything_but_use_and_role_from_the_values(app, client):
    """画面が送るのは「使う・役割」だけ。キー・型・単位・出し方は見出しと値から決める。"""
    import_id = upload_csv(client, "a.csv")
    assert save_source(client, import_id, encoding="cp932", delimiter=",").status_code == 200
    assert save_layout(client, import_id).status_code == 200

    body = editor_body(client, import_id)
    assert set(body["columns"][0]) == {"index", "header", "use", "role"}
    assert body["name"] == "a"   # 表の名前の初期値はファイル名
    body["name"] = "トラブル対応一覧"
    assert save_columns(client, import_id, body).status_code == 200
    wait_import_job(app, import_id)

    with app.app_context():
        spec = store.get_import(import_id)["spec"]
    assert spec.name == "トラブル対応一覧"
    assert spec.column("occurred_at").type in ("date", "datetime")   # 型は値から
    assert spec.column("downtime").unit == "分"                # 単位は見出しから
    assert spec.column("symptom").md in ("body", "attribute")  # md での扱いも自動
    assert spec.markdown["group_by"] == "month"               # まとめ方は常に月
    assert spec.markdown["file_prefix"] == "" and spec.file_prefix == "トラブル対応一覧"


def test_saving_the_columns_again_replaces_the_imports_spec(app, client):
    """保存し直しても版は作らず、その取り込みの設定を上書きする。"""
    import_id = upload_csv(client, "a.csv")
    save_source(client, import_id, encoding="cp932", delimiter=",")
    save_layout(client, import_id)
    assert save_columns(client, import_id, columns_payload("1回目")).status_code == 200
    wait_import_job(app, import_id)
    with app.app_context():
        first = store.get_import(import_id)
    assert save_columns(client, import_id, columns_payload("2回目")).status_code == 200
    wait_import_job(app, import_id)
    with app.app_context():
        second = store.get_import(import_id)
    assert first["spec"].name == "1回目" and second["spec"].name == "2回目"
    assert second["spec_hash"] != first["spec_hash"]
    assert second["template_version_id"] == first["template_version_id"] == import_id


def test_the_table_name_is_required(app, client):
    import_id = upload_csv(client, "a.csv")
    save_source(client, import_id, encoding="cp932", delimiter=",")
    save_layout(client, import_id)
    res = save_columns(client, import_id, columns_payload(""))
    assert res.status_code == 400 and res.get_json()["error"] == "表の名前を入力してください"


def test_there_is_no_saved_settings_endpoint(client):
    """取り込み設定の経路（一覧・削除）は残していない。"""
    assert client.post("/tables/templates/1/delete").status_code == 404


# ---- 帳票の種類: 状態・削除、帳票の JSON と元ファイル ----------------------------------------------

def _make_form_type(app, client, sample_dir) -> int:
    from tests.test_forms_flow import add_field, create_type

    pattern_id = create_type(client, sample_dir / "standard.xlsx", "設備修理報告書")
    add_field(client, pattern_id, "修理報告書", "A3", "B3")
    add_field(client, pattern_id, "修理報告書", "A4", "B4")
    return pattern_id


def _upload_files(tmp_path) -> set[str]:
    return {p.name for p in (tmp_path / "uploads").rglob("*") if p.is_file()}


def test_form_type_status_and_delete(app, client, sample_dir, tmp_path):
    """使用開始・停止・削除。置いた Excel はどこにも残らない（2026-09-21 の利用者の指示）。"""
    from tests.test_forms_flow import book_part

    pattern_id = _make_form_type(app, client, sample_dir)
    other_id = _make_form_type(app, client, sample_dir)

    # 種類を作っても、セルをクリックしても、Excel はサーバーに残らない
    assert _upload_files(tmp_path) == set()
    with app.app_context():
        tables = {row[0] for row in db.get_db().execute(
            "SELECT name FROM sqlite_master WHERE type = 'table'")}
    assert "pattern_samples" not in tables

    res = client.post(f"/form-types/{pattern_id}/status", json={"status": "active"})
    assert res.status_code == 200 and "使用を開始しました" in res.get_json()["message"]
    assert "見本" not in res.get_json()["message"]
    with app.app_context():
        assert db.load_pattern(pattern_id).status == "active"
    assert client.post(f"/form-types/{pattern_id}/status", json={"status": "zzz"}).status_code == 400
    assert client.post("/form-types/9999/status", json={"status": "active"}).status_code == 404

    # Excel を置き直すとシートが出る（種類は増えない）。Excel でないファイルは断る
    res = client.post(f"/form-types/{pattern_id}/panel",
                      data={"book": book_part(sample_dir / "shifted.xlsx")},
                      content_type="multipart/form-data")
    assert res.status_code == 200 and 'data-cell="A1"' in res.get_json()["html"]
    assert _upload_files(tmp_path) == set()
    res = client.post(f"/form-types/{pattern_id}/panel",
                      data={"book": (io.BytesIO(b"not excel"), "x.xlsx")},
                      content_type="multipart/form-data")
    assert res.status_code == 400 and "Excelファイル" in res.get_json()["error"]
    assert _upload_files(tmp_path) == set()
    with app.app_context():
        assert len(db.list_patterns()) == 2      # 置き直しで種類は増えない

    res = client.post(f"/form-types/{pattern_id}/status", json={"status": "inactive"})
    assert "使用を停止しました" in res.get_json()["message"]
    with app.app_context():
        assert db.load_pattern(pattern_id).status == "inactive"

    res = client.post(f"/form-types/{pattern_id}/delete")
    assert "を削除しました" in res.get_json()["message"]
    with app.app_context():
        assert db.load_pattern(pattern_id) is None
        assert db.load_pattern(other_id) is not None
    assert client.post(f"/form-types/{pattern_id}/delete").status_code == 404


def test_the_original_file_can_be_downloaded_without_losing_the_form(app, client, sample_dir):
    """元のファイルのダウンロードでは何も消さない（消えるのは Markdown を渡したときだけ）。"""
    from tests.test_forms_flow import activate, read_form, upload_forms

    pattern_id = _make_form_type(app, client, sample_dir)
    activate(client, pattern_id)
    doc_id, = upload_forms(client, sample_dir / "standard.xlsx")
    read_form(client, doc_id, pattern_id, ["修理報告書"])

    res = client.get(f"/forms/{doc_id}/original")
    assert res.status_code == 200 and res.data == (sample_dir / "standard.xlsx").read_bytes()
    assert "standard.xlsx" in res.headers["Content-Disposition"]
    with app.app_context():
        assert db.get_document(doc_id) is not None
    assert client.get("/forms/9999/original").status_code == 404
    # 確定していない帳票の Markdown は渡さない
    assert client.get(f"/forms/{doc_id}/download.md").status_code == 404


# ---- 一覧表: Excel のシートの選択 ---------------------------------------------------------------

def test_save_source_uses_the_chosen_excel_sheet(app, client, tmp_path):
    book = openpyxl.Workbook()
    memo = book.active
    memo.title = "メモ"
    memo["A1"] = "このブックの説明"
    sheet = book.create_sheet("一覧")
    sheet.append(["管理No", "発生日", "現象"])
    for i in range(1, 6):
        sheet.append([f"TR-{i:03d}", f"2026-08-{i:02d}", f"アラーム停止{i}"])
    path = tmp_path / "二つのシート.xlsx"
    book.save(path)

    import_id = upload_bytes(client, path.read_bytes(), path.name)
    page = panel_html(client, import_id, "source")
    assert "メモ" in page and "一覧" in page

    res = save_source(client, import_id, sheet="一覧")
    assert res.status_code == 200 and res.get_json()["next"] == "layout"
    with app.app_context():
        assert store.get_import(import_id)["source"]["sheet"] == "一覧"
    page = panel_html(client, import_id, "layout")
    assert "TR-001" in page and "このブックの説明" not in page
    detect = client.post(f"/tables/imports/{import_id}/layout/detect", json={"header_rows": [1]}).get_json()
    assert detect["data_end"] == 6
    res = save_layout(client, import_id)
    assert res.status_code == 200 and res.get_json()["next"] == "columns"
    assert "発生日" in panel_html(client, import_id, "columns")

    # シートを選び直すと、前のシートで決めた見出し行・範囲は使わない
    save_source(client, import_id, sheet="メモ")
    with app.app_context():
        src = store.get_import(import_id)["source"]
    assert src["sheet"] == "メモ" and "header_rows" not in src
    assert "このブックの説明" in panel_html(client, import_id, "layout")


def test_the_section_of_a_clicked_cell_is_saved_with_the_field(app, client, tmp_path):
    """回答欄の中のセルをクリックした項目は「探す区画」を持つ（消えると回答側の欄を読めない）。"""
    from openpyxl.styles import Font, PatternFill

    from tests.test_forms_flow import add_field, create_type

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "報告書"
    for coord, text in (("A1", "■ 発行側"), ("A4", "▼ 回答欄")):
        ws[coord] = text
        ws[coord].font = Font(bold=True)
        ws[coord].fill = PatternFill("solid", fgColor="DDDDDD")
    for coord, text in (("A2", "処置内容"), ("A5", "処置内容")):
        ws[coord] = text
        ws[coord].fill = PatternFill("solid", fgColor="DDDDDD")
    ws["B2"], ws["B5"] = "発行側の処置", "回答側の処置"
    path = tmp_path / "回答欄.xlsx"
    wb.save(path)

    pattern_id = create_type(client, path, "工程異常連絡書")
    add_field(client, pattern_id, "報告書", "A5", "B5")
    with app.app_context():
        field = db.load_pattern(pattern_id).fields[0]
    assert field.section == "回答"
    # 同じ帳票の Excel を置き直しても区画は残り、読み取りテストは回答側の値を出す
    # （シートには両方の値が写っているので、Markdown の中身で確かめる）
    from tests.test_forms_flow import panel_html

    panel = panel_html(client, pattern_id, book=path)
    markdown = panel[panel.index("md-preview"):]
    assert "## 処置内容\n回答側の処置" in markdown and "発行側の処置" not in markdown
    with app.app_context():
        assert db.load_pattern(pattern_id).fields[0].section == "回答"
