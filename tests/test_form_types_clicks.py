"""帳票の種類をセルのクリックで作る: クリック→項目の組み立てと、その画面の動き。"""
import io

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from excel.workbook import load_workbook_info
from models import database as db
from pattern.clicks import click_field, table_cells


def _book(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "報告書"
    fill = PatternFill("solid", fgColor="FFD9E1F2")
    for coord, value in {"A1": "報告番号", "B1": "R-001", "A2": "設備番号", "B2": "EQ-001",
                         "A3": "作業時間(h)", "B3": 2.5, "A4": "所見", "A5": "異音あり",
                         "C1": "ライン：L6", "A7": "■ 交換部品",
                         "A8": "品番", "B8": "品名", "C8": "数量",
                         "A9": "PW-1", "B9": "ベアリング", "C9": 2,
                         "A10": "PW-2", "B10": "モータ", "C10": 1}.items():
        ws[coord] = value
    for coord in ("A1", "A2", "A3", "A4", "A7", "A8", "B8", "C8"):
        ws[coord].fill = fill
    wb.save(path)
    return path


@pytest.fixture()
def grid(tmp_path):
    return load_workbook_info(_book(tmp_path / "click.xlsx")).grids["報告書"]


def test_click_label_then_value_makes_a_field(grid):
    row, error = click_field(grid, "A1", "B1")
    assert error == "" and row["field_name"] == "report_id" and row["display_name"] == "報告番号"
    # 探す見出しはクリックした語＋辞書の同じ意味の語。値の位置・セル番地も控える
    assert row["candidates"].splitlines()[0] == "報告番号" and len(row["candidates"].splitlines()) > 1
    assert row["direction"] == "right" and row["label_cell"] == "A1" and row["cell"] == "B1"
    assert row["sheet_name"] == "報告書" and row["examples"] == ["R-001"]


def test_click_derives_type_and_unit_from_the_value(grid):
    row, _ = click_field(grid, "A3", "B3")
    assert row["data_type"] == "number" and row["unit"] == "時間" and row["display_name"] == "作業時間"
    below, _ = click_field(grid, "A4", "A5")
    assert below["direction"] == "below" and below["data_type"] == "string"


def test_clicking_the_same_cell_twice(grid):
    inline, _ = click_field(grid, "C1", "C1")   # 「ライン：L6」はセル内の見出しと値
    assert inline["direction"] == "same_cell" and inline["display_name"] == "ライン"
    only, _ = click_field(grid, "A5", "A5")     # 見出しのない値だけの項目
    assert only["candidates"] == "" and only["cell"] == "A5" and "A5" in only["display_name"]


def test_clicking_a_table_header_makes_one_table_field(grid):
    assert {"A7", "A8", "B8", "C8"} <= table_cells(grid)
    row, error = click_field(grid, "A8")
    assert error == "" and row["data_type"] == "table"
    assert row["table_columns"].splitlines() == ["品番", "品名", "数量"]


def test_a_table_filled_in_once_in_the_sample_is_still_a_table(tmp_path):
    """見本での記入が1行だけの表も、列見出しのすぐ上の見出しを指せば明細表になる。"""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "報告書"
    for coord, value in {"A1": "設備修理報告書", "F1": "承認", "G1": "確認",   # 帳票のタイトルと押印欄
                         "F2": "高橋", "G2": "中村", "F3": "7/5", "G3": "7/5",
                         "A5": "使用部品", "A6": "品番", "B6": "品名", "C6": "数量",
                         "A7": "PM26-1778", "B7": "センサ", "C7": 1}.items():
        ws[coord] = value
    fill = PatternFill("solid", fgColor="FFD9E1F2")
    for coord in ("A5", "A6", "B6", "C6", "F1", "G1"):
        ws[coord].fill = fill
    ws.merge_cells("A5:C5")   # 表の見出しは表の幅いっぱいに結合されている
    wb.save(tmp_path / "one_row.xlsx")
    grid = load_workbook_info(tmp_path / "one_row.xlsx").grids["報告書"]

    row, error = click_field(grid, "A5")
    assert error == "" and row["data_type"] == "table" and row["display_name"] == "使用部品"
    assert row["table_columns"].splitlines() == ["品番", "品名", "数量"]
    assert "A5" in table_cells(grid)
    # 帳票のタイトルは、下の押印欄から離れているので表の見出しにしない
    assert "A1" not in table_cells(grid)
    assert click_field(grid, "A1", "B1")[0]["data_type"] != "table"


def test_clicking_an_empty_cell_as_a_label_is_refused(grid):
    row, error = click_field(grid, "F20")
    assert row is None and "見出し" in error
    assert click_field(grid, "zz")[0] is None


def test_label_without_a_value_can_be_registered(grid):
    row, error = click_field(grid, "A2", "")
    assert error == "" and row["cell"] == "" and row["candidates"].startswith("設備番号")


# ---- 画面 ---------------------------------------------------------------------------

def _create(client, path, name="点検報告書"):
    res = client.post("/form-types/new",
                      data={"name": name, "samples": (io.BytesIO(path.read_bytes()), path.name)},
                      content_type="multipart/form-data")
    assert res.status_code == 200, res.get_data(as_text=True)
    return res.get_json()["pattern_id"]


def _panel(client, pattern_id, sample=None) -> str:
    url = f"/form-types/{pattern_id}/panel" + (f"?sample={sample}" if sample else "")
    return client.get(url).get_json()["html"]


def test_build_panel_adds_and_deletes_fields_by_clicking(app, client, tmp_path):
    path = _book(tmp_path / "click.xlsx")
    pattern_id = _create(client, path)
    page = client.get("/form-types/").get_data(as_text=True)
    assert "点検報告書" in page and "form_types.js" in page
    panel = _panel(client, pattern_id)
    assert 'data-cell="A1"' in panel and "まだ項目がありません" in panel
    # クリックを送るための仕掛け（form_types.js が使う）と、明細表の印
    assert 'id="cellBuilder"' in panel and "data-click-hint" in panel
    assert 'data-table-head="1"' in panel

    with app.app_context():
        sample_id = db.list_samples(pattern_id)[0]["id"]
    data = {"sample": sample_id, "sheet": "報告書", "label_cell": "A1", "value_cell": "B1"}
    body = client.post(f"/form-types/{pattern_id}/fields", data=data).get_json()
    assert "「報告番号」を項目にしました" in body["message"] and "R-001" in body["html"]
    assert "点検報告書" in body["list_html"]
    with app.app_context():
        pattern = db.load_pattern(pattern_id)
    assert [f.field_name for f in pattern.fields] == ["report_id"]
    assert pattern.title_fields == ["report_id"]          # タイトル項目は自動で決める
    assert [s.sheet_name for s in pattern.sheets] == ["報告書"]
    assert pattern.fields[0].cell == "B1" and pattern.fields[0].sheet_name == "報告書"

    # 同じセルをもう一度クリックしても増えない
    body = client.post(f"/form-types/{pattern_id}/fields", data=data).get_json()
    assert body["message"] == "そのセルはもう項目になっています"

    # 明細表は列見出しを1回クリックするだけ
    client.post(f"/form-types/{pattern_id}/fields",
                data={"sample": sample_id, "sheet": "報告書", "label_cell": "A8", "value_cell": ""})
    with app.app_context():
        table = [f for f in db.load_pattern(pattern_id).fields if f.data_type == "table"]
    assert len(table) == 1 and table[0].table_columns == ["品番", "品名", "数量"]

    body = client.post(f"/form-types/{pattern_id}/fields/report_id/delete").get_json()
    assert body["message"] == "項目を削除しました"
    with app.app_context():
        assert [f.field_name for f in db.load_pattern(pattern_id).fields] == ["parts"]
    assert client.post(f"/form-types/{pattern_id}/fields/zzz/delete").status_code == 404

    # 使用開始を押すまでは作成中のまま
    with app.app_context():
        assert db.load_pattern(pattern_id).status == "draft"
    assert client.post(f"/form-types/{pattern_id}/status", json={"status": "active"}).get_json()["status"] == "active"


def test_a_type_without_fields_cannot_be_used(app, client, tmp_path):
    pattern_id = _create(client, _book(tmp_path / "click.xlsx"))
    res = client.post(f"/form-types/{pattern_id}/status", json={"status": "active"})
    assert res.status_code == 400 and "読み取る項目がありません" in res.get_json()["error"]
    with app.app_context():
        assert db.load_pattern(pattern_id).status == "draft"


def test_the_name_comes_from_the_file_name_and_can_be_changed(app, client, tmp_path):
    path = _book(tmp_path / "設備点検表.xlsx")
    res = client.post("/form-types/new",
                      data={"name": "", "samples": (io.BytesIO(path.read_bytes()), path.name)},
                      content_type="multipart/form-data")
    pattern_id = res.get_json()["pattern_id"]
    with app.app_context():
        assert db.load_pattern(pattern_id).name == "設備点検表"
    assert client.post(f"/form-types/{pattern_id}/name", json={"name": "点検表"}).get_json()["ok"] is True
    with app.app_context():
        assert db.load_pattern(pattern_id).name == "点検表"
    res = client.post(f"/form-types/{pattern_id}/name", json={"name": " "})
    assert res.status_code == 400 and "名前を入れてください" in res.get_json()["error"]


def test_a_field_found_by_its_cell_when_the_label_is_missing(app, client, tmp_path):
    """見出しの無い「値だけ」の項目は、見本でクリックしたセルの番地から読む。"""
    path = _book(tmp_path / "click.xlsx")
    pattern_id = _create(client, path)
    with app.app_context():
        sample_id = db.list_samples(pattern_id)[0]["id"]
    client.post(f"/form-types/{pattern_id}/fields",
                data={"sample": sample_id, "sheet": "報告書", "label_cell": "A5", "value_cell": "A5"})
    assert "異音あり" in _panel(client, pattern_id)     # 読み取りテストは同じ欄に出る


# ---- 版によって書き方が違う帳票 -------------------------------------------------------

def _versions(tmp_path):
    """見本（v1: 見出しの右に値）と、書き方の違う同じ帳票（v2: セル内に「設備番号：…」）。"""
    from openpyxl import Workbook
    v1 = Workbook()
    ws = v1.active
    ws.title = "報告書"
    for coord, value in {"A1": "報告番号", "B1": "R-001",
                         "A2": "使用設備", "B2": "CMP-108　STI-CMP 8号機",
                         "A3": "所見", "B3": "異音あり"}.items():
        ws[coord] = value
    v1.save(tmp_path / "v1.xlsx")

    v2 = Workbook()
    ws2 = v2.active
    ws2.title = "報告書"
    # 同じ番地には別の欄の見出しが来る。設備は1つのセルに「設備番号：…」と書かれている
    for coord, value in {"A1": "報告番号", "B1": "R-002",
                         "A2": "所見", "B2": "発生日時",
                         "A4": "設備番号　：IMP-603", "A5": "所見", "B5": "振動あり"}.items():
        ws2[coord] = value
    v2.save(tmp_path / "v2.xlsx")
    return tmp_path / "v1.xlsx", tmp_path / "v2.xlsx"


def test_one_cell_holding_an_equipment_number_and_name_becomes_two_fields(tmp_path):
    """「使用設備：CMP-108　STI-CMP 8号機」を1回クリックすると、設備番号と設備名の2項目になる。"""
    from pattern.clicks import split_rows

    v1, _ = _versions(tmp_path)
    grid = load_workbook_info(v1).grids["報告書"]
    row, error = click_field(grid, "A2", "B2")
    assert error == ""
    parts = split_rows(row)
    assert [p["field_name"] for p in parts] == ["equipment_id", "equipment_name"]
    # 文字は消さず、同じセルの読む場所を分けるだけ。書き方の違う見出しは全部入れておく
    assert [p["examples"] for p in parts] == [["CMP-108"], ["STI-CMP 8号機"]]
    assert all(p["cell"] == "B2" for p in parts)
    assert "使用設備" in parts[0]["candidates"] and "設備番号" in parts[0]["candidates"]
    # 番号と名前に分けられない欄はそのまま1項目
    plain, _ = click_field(grid, "A1", "B1")
    assert split_rows(plain) == [plain]


def test_a_differently_written_version_is_read_by_the_label_not_by_the_cell_address(tmp_path):
    """見本と書き方が違う帳票でも、セル内の見出しから読む。別の欄の見出しは値にしない。"""
    from excel.extractor import extract_document
    from pattern.clicks import split_rows
    from pattern.forms import rows_to_pattern

    v1, v2 = _versions(tmp_path)
    grid = load_workbook_info(v1).grids["報告書"]
    rows = []
    for label, value in (("A2", "B2"), ("A3", "B3")):
        row, _ = click_field(grid, label, value, {r["field_name"] for r in rows})
        rows += split_rows(row, {r["field_name"] for r in rows})
    pattern = rows_to_pattern(1, {"name": "報告書", "version": "v1", "description": "",
                                  "image_processing": "none", "title_fields": [], "md_options": {}},
                              [{"use": True, "sheet_name": "報告書", "required": False}], rows)
    fields = {f["field_name"]: f for f in extract_document(load_workbook_info(v2), pattern, ["報告書"])["fields"]}
    # 見本では B2 だった番地に、この版では別の欄の見出し「発生日時」が来ている。それは値にしない
    assert fields["equipment_id"]["value"] == "IMP-603"
    assert fields["equipment_name"]["value"] in (None, "")
    assert fields["field_1"]["value"] == "振動あり"


def test_the_clicked_label_wins_over_a_dictionary_synonym_on_the_same_sheet(tmp_path):
    """「報告者」と「担当者」が並ぶ帳票で、クリックした方の値を読む。

    辞書は「担当者」の言い換えとして「報告者」「記入者」…も探す。その言い換えがシートの先に
    出てくると、クリックした欄ではなく言い換えの欄の値が読まれてしまっていた。
    """
    from excel.extractor import extract_document
    from pattern.forms import rows_to_pattern

    wb = Workbook()
    ws = wb.active
    ws.title = "報告書"
    fill = PatternFill("solid", fgColor="FFD9E1F2")
    for coord, value in {"A1": "報告者", "B1": "長谷川 聡", "C1": "担当者", "D1": "清水 彩花"}.items():
        ws[coord] = value
    for coord in ("A1", "C1"):
        ws[coord].fill = fill
    path = tmp_path / "二人.xlsx"
    wb.save(path)

    info = load_workbook_info(path)
    grid = info.grids["報告書"]
    meta = {"name": "報告書", "version": "v1", "description": "", "image_processing": "none",
            "title_fields": [], "md_options": {}}
    sheets = [{"use": True, "sheet_name": "報告書", "required": False}]
    for label_cell, value_cell, expected in (("C1", "D1", "清水 彩花"), ("A1", "B1", "長谷川 聡")):
        row, error = click_field(grid, label_cell, value_cell)
        assert error == ""
        # 辞書がもう一方の見出しも探す候補に入れている（この状況でクリックした方を読めること）
        assert "報告者" in row["candidates"] and "担当者" in row["candidates"]
        pattern = rows_to_pattern(1, meta, sheets, [row])
        field = extract_document(info, pattern, ["報告書"])["fields"][0]
        assert field["value"] == expected, (label_cell, field)
