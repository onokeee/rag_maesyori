from __future__ import annotations

import csv
import hashlib
import io
import json
import random
import re
import shutil
import subprocess
import tempfile
import time
import zipfile
from collections import Counter
from datetime import datetime, time as dtime, timedelta
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.datetime import CALENDAR_MAC_1904

from app import core
from app import tables
from app import views
from app.core import estimate_tokens, md_filename
from app.tables import (
    CsvSource,
    sniff_csv,
    classify_rows,
    guess_layout,
    is_month_label,
    list_kind,
    sample_data_rows,
    split_header_unit,
    STANDARD_COLUMNS,
    lookup_header,
    ExcelSource,
    suggest_columns,
    UploadError,
    open_source,
    value_kind,
    Issue,
    has_blocking,
    run_checks,
    record_block,
    render_all,
    split_entity_code,
    ConvertContext,
    convert_cell,
    find_year_context,
    read_records,
    ColumnSpec,
    resolve_columns,
    spec_from_dict,
    spec_from_suggestions,
    spec_hash,
    spec_to_dict,
    validate_spec,
    MAX_COLUMNS,
    parse_number_text,
    _Names,
    record_title,
)
from tests.conftest import (
    COLUMNS,
    CSV_TEXT,
    columns_payload,
    csv_source,
    panel,
    panel_html,
    preview_panel,
    save_columns,
    save_layout,
    save_source,
    upload,
    upload_csv,
    wait_import_job,
    editor_body,
    spec_import,
    imported,
    confirmed,
    upload_bytes,
)
from tests.test_ai import _read_import, ai_app_endpoints_ai, ai_client, fake_endpoints_ai


def _tables_js() -> str:
    """static/app.js のうち、表の取り込みの部分（「// ==== tables: 」の見出しから次の見出しまで）。"""
    js = (Path(__file__).resolve().parents[1] / "app" / "static" / "app.js").read_text(encoding="utf-8")
    start = js.index("// ==== tables: ")
    nxt = js.find("\n// ==== ", start + 1)
    return js[start:nxt if nxt >= 0 else len(js)]



# ====================================================================================================
# 元 tests/test_tables_read.py
# 一覧表の読み取り（表ソース・見出し帯・行分類・列の対応づけ候補）のテスト。
# ====================================================================================================

SAMPLES = Path(__file__).resolve().parents[1] / "samples" / "tables"


# ---- CSV ----

def _write(path: Path, text: str, encoding: str) -> Path:
    path.write_bytes(text.encode(encoding))
    return path


def _system_csv(n: int = 3) -> str:
    lines = [
        "出力日時,2026/09/14 18:05:33",
        "出力者,M10544,中村 浩二",
        '抽出条件,期間=2023/04/01〜2026/08/31,"記録区分=1:設備故障,2:チョコ停",部門=21000',
        "",
        "管理番号,設備コード,発生日時,現象,処置内容,停止時間(分)",
    ]
    for i in range(1, n + 1):
        lines.append(f'TR-2026-{i:05d},CMP-10{i},2026/08/0{i} 10:00:00,髙橋さん連絡①,"1. 点検\n2. 交換",{i * 10}')
    lines.append(f"合計件数,{n}")
    return "\r\n".join(lines) + "\r\n"


def test_sniff_cp932_preamble_multiline_and_trailer(tmp_path):
    path = _write(tmp_path / "故障履歴.csv", _system_csv(), "cp932")
    sniff = sniff_csv(path)
    assert sniff.encoding == "cp932" and not sniff.bom
    assert sniff.delimiter == ","
    assert sniff.header_row == 5 and sniff.preamble_rows == 4
    assert sniff.trailer_rows == 1
    assert sniff.warnings == []

    src = open_source(path, "故障履歴.csv")
    rows = list(src.rows("故障履歴.csv"))
    assert [r.index for r in rows] == list(range(1, 10))  # 引用符内の改行は1レコード
    assert rows[5].cells[4].text == "1. 点検\n2. 交換"
    assert rows[5].cells[3].text == "髙橋さん連絡①"
    assert rows[3].is_blank and rows[3].hidden is None

    layout = guess_layout(src, "故障履歴.csv")
    assert layout.table_kind == "list"
    assert layout.header_rows == [5] and layout.data_start == 6
    assert layout.headers == ["管理番号", "設備コード", "発生日時", "現象", "処置内容", "停止時間(分)"]
    assert layout.data_end == 9  # 末尾の件数行は合計行として範囲に含め、分類で除く
    assert layout.counts["data"] == 3 and layout.counts["subtotal"] == 1
    kinds = {rc.index: rc.kind for rc in layout.row_classes}
    assert kinds[1] == "title" and kinds[4] == "blank" and kinds[5] == "header" and kinds[9] == "subtotal"


def test_sniff_utf8_bom_embedded_newlines(tmp_path):
    text = '台帳No,起票日,不具合内容\n' + "".join(
        f'CA-2026-{i:04d},R8.4.{i},"■発生：R8.4.{i}\n■設備：CMP-101, ""A""系"\n' for i in range(1, 8)
    )
    path = tmp_path / "台帳.csv"
    path.write_bytes(b"\xef\xbb\xbf" + text.encode("utf-8"))
    sniff = sniff_csv(path)
    assert (sniff.encoding, sniff.bom, sniff.delimiter, sniff.header_row) == ("utf-8-sig", True, ",", 1)
    src = CsvSource(path)
    rows = list(src.rows())
    assert len(rows) == 8
    assert rows[0].cells[0].text == "台帳No"  # BOM は見出しに残らない
    assert rows[1].cells[2].text == '■発生：R8.4.1\n■設備：CMP-101, "A"系'
    layout = guess_layout(src, src.file_name)
    assert layout.header_rows == [1] and layout.counts["data"] == 7


def test_sniff_prefers_consistent_tab_over_comma(tmp_path):
    lines = ["管理No\t発生日\t処置"]
    for i in range(20):
        commas = "," * (i % 4)
        lines.append(f"A-{i:03d}\t2026/08/{i % 28 + 1:02d}\t交換{commas}、確認,完了")
    path = _write(tmp_path / "data.txt", "\n".join(lines) + "\n", "utf-8")
    sniff = sniff_csv(path)
    assert sniff.delimiter == "\t" and sniff.encoding == "utf-8"
    assert sniff.confidence == 1.0


def test_sniff_shift_jis_2004_and_utf16(tmp_path):
    body = "設備名,担当\n" + "".join(f"森鷗外{i}号機,剝離{i}\n" for i in range(5))
    sj = _write(tmp_path / "sj.csv", body, "shift_jis_2004")
    assert sniff_csv(sj).encoding == "shift_jis_2004"
    rows = list(CsvSource(sj).rows())
    assert rows[1].cells[0].text == "森鷗外0号機"

    u16 = tmp_path / "u16.csv"
    u16.write_bytes(body.encode("utf-16"))
    sniff = sniff_csv(u16)
    assert sniff.encoding == "utf-16" and sniff.bom
    assert list(CsvSource(u16).rows())[2].cells[1].text == "剝離1"


def test_wrong_encoding_choice_is_a_message_not_a_500(tmp_path):
    """画面の選択肢から文字コードを選び直せるように、読めない指定は UploadError にする。

    UTF-8 のファイルに UTF-16 を選ぶと、デコーダは UnicodeDecodeError ではなく UnicodeError（BOM が無い）を出す。
    """
    body = "設備名,担当\n" + "".join(f"CMP-10{i},田中\n" for i in range(5))
    path = _write(tmp_path / "utf8.csv", body, "utf-8")
    with pytest.raises(UploadError, match="文字コード"):
        list(CsvSource(path, options={"encoding": "utf-16", "delimiter": ","}).rows())
    with pytest.raises(UploadError, match="文字コード utf-8-x は使えません"):
        list(CsvSource(path, options={"encoding": "utf-8-x", "delimiter": ","}).rows())
    with pytest.raises(UploadError, match="区切り文字は1文字"):
        CsvSource(path, options={"encoding": "utf-8", "delimiter": "ABC"})
    # 正しい指定はそのまま読める
    assert list(CsvSource(path, options={"encoding": "utf-8", "delimiter": ","}).rows())[1].cells[0].text == "CMP-100"


def test_ascii_head_then_cp932_is_detected_by_full_decode(tmp_path):
    head = "code,name,memo\n" + "".join(f"C{i:06d},item{i},ok\n" for i in range(60000))  # 1MB超のASCII
    path = tmp_path / "big.csv"
    path.write_bytes(head.encode("ascii") + "C999999,設備,異常\n".encode("cp932"))
    assert path.stat().st_size > 1024 * 1024
    assert sniff_csv(path).encoding == "cp932"


def test_mixed_encoding_warns_strict_raises_and_replace_marks_rows(tmp_path):
    good = "管理No,現象\n" + "".join(f"T-{i},停止{i}\n" for i in range(5))
    path = tmp_path / "mixed.csv"
    path.write_bytes(good.encode("cp932") + "T-9,".encode("cp932") + b"\x82\xff\n" + "T-10,復旧\n".encode("cp932"))
    sniff = sniff_csv(path)
    assert sniff.encoding == "cp932"
    assert sniff.decode_error_line == 7
    assert "読めない文字" in sniff.warnings[0] and "〓" in sniff.warnings[0]

    with pytest.raises(UploadError, match="読めない文字"):
        list(CsvSource(path, options={"encoding": "cp932", "delimiter": ","}).rows())

    src = CsvSource(path, options={"encoding": "cp932", "delimiter": ",", "errors": "replace"})
    rows = list(src.rows())
    assert "〓" in rows[6].cells[1].text
    assert src.replaced_rows == [7]
    assert rows[7].cells[1].text == "復旧"


def test_csv_nul_removal_and_excel_wrapped_values(tmp_path):
    text = '設備コード,部品番号,数量\n="00123",="0045",1\nAB\x00C,"=""0099""",2\n'
    path = _write(tmp_path / "codes.csv", text, "utf-8")
    rows = list(CsvSource(path).rows())
    assert rows[1].cells[0].text == "00123" and rows[1].cells[1].text == "0045"
    assert rows[2].cells[0].text == "ABC"
    assert rows[2].cells[1].text == "0099"
    assert value_kind(rows[1].cells[0].value, rows[1].cells[0].text) == "code"  # 先頭ゼロは数値にしない
    info = CsvSource(path).sheets()[0]
    assert (info.name, info.max_row, info.max_col) == ("codes.csv", 3, 3)


def test_csv_start_and_limit(tmp_path):
    path = _write(tmp_path / "s.csv", "a,b\n" + "".join(f"{i},{i}\n" for i in range(10)), "utf-8")
    rows = list(CsvSource(path).rows(None, start=3, limit=2))
    assert [r.index for r in rows] == [3, 4]


def test_csv_two_row_header_without_merge(tmp_path):
    lines = ["不具合,,,対策,", "台帳No,起票日,不具合内容,暫定対策,期限"]
    lines += [f"CA-{i},2026/04/0{i},停止{i},交換,2026/05/0{i}" for i in range(1, 8)]
    path = _write(tmp_path / "two.csv", "\n".join(lines) + "\n", "utf-8")
    src = CsvSource(path)
    layout = guess_layout(src, src.file_name)
    assert layout.header_rows == [1, 2]
    assert layout.headers == ["不具合_台帳No", "不具合_起票日", "不具合_不具合内容", "対策_暫定対策", "対策_期限"]


def test_open_source_rejects_unknown_extension(tmp_path):
    p = tmp_path / "a.xls"
    p.write_bytes(b"x")
    with pytest.raises(UploadError, match="対応していない形式"):
        open_source(p, "a.xls")


# ---- Excel ----

def _list_workbook(path: Path) -> tuple[Path, int]:
    wb = Workbook()
    ws = wb.active
    ws.title = "故障履歴"
    ws["A1"] = "2026年8月 故障履歴一覧"
    ws["A1"].font = Font(bold=True)
    ws.merge_cells("A1:F1")
    ws["A2"] = "作成：保全1係"
    headers = ["管理No", "発生日", "設備番号", "設備名", "現象", "停止時間(分)"]
    for c, h in enumerate(headers, start=1):
        ws.cell(4, c, h).font = Font(bold=True)
    r = 5
    for i in range(1, 13):
        ws.cell(r, 1, f"TR-2026-{i:05d}")
        ws.cell(r, 2, datetime(2026, 8, i))
        ws.cell(r, 3, f"CMP-10{i % 3}")
        ws.cell(r, 4, "CMP研磨装置")
        ws.cell(r, 5, f"スラリー流量低下 {i}")
        ws.cell(r, 6, i * 5)
        r += 1
    # 継続行（キー列が空で文章だけ）
    ws.cell(r, 5, "（追記）フィルター交換後に再発なし")
    r += 1
    # 非表示の行と取り消し線の行
    for label, hide, strike in (("hidden", True, False), ("strike", False, True)):
        ws.cell(r, 1, f"TR-2026-9{r:04d}")
        ws.cell(r, 2, datetime(2026, 8, 20))
        ws.cell(r, 3, "CVD-201")
        ws.cell(r, 5, label)
        ws.cell(r, 6, 1)
        if hide:
            ws.row_dimensions[r].hidden = True
        if strike:
            for c in range(1, 7):
                ws.cell(r, c).font = Font(strike=True)
        r += 1
    ws.cell(r, 1, "合計").font = Font(bold=True)
    ws.cell(r, 6, 999)
    total_row = r
    ws.cell(r + 2, 1, "※停止時間は生産停止から復旧確認までの時間")
    wb.save(path)
    return path, total_row


def test_excel_list_rows_classes_and_end(tmp_path):
    path, total_row = _list_workbook(tmp_path / "list.xlsx")
    src = open_source(path, "list.xlsx")
    assert src.kind == "excel"
    info = src.sheets()[0]
    assert info.name == "故障履歴" and info.max_col == 6 and not info.date1904

    first = next(src.rows("故障履歴", 1, 1))
    assert first.cells[0].bold and first.cells[0].merged_anchor == (1, 1)
    assert first.cells[3].merged_anchor == (1, 1) and first.cells[3].text == ""

    layout = guess_layout(src, "故障履歴")
    assert layout.table_kind == "list"
    assert layout.header_rows == [4] and layout.data_start == 5
    assert layout.headers == ["管理No", "発生日", "設備番号", "設備名", "現象", "停止時間(分)"]
    assert layout.data_end == total_row
    kinds = {rc.index: (rc.kind, rc.reason) for rc in layout.row_classes}
    assert kinds[1][0] == "title" and kinds[4][0] == "header"
    assert kinds[17] == ("continuation", "キー列が空で文章だけの行")
    assert kinds[18] == ("excluded", "非表示の行")
    assert kinds[19] == ("excluded", "取り消し線の行")
    assert kinds[total_row] == ("subtotal", "合計")
    assert kinds[total_row + 2][0] == "note"
    assert layout.counts == {"data": 12, "continuation": 1, "excluded": 2, "subtotal": 1}
    assert any("除外した行が2行" in w for w in layout.warnings)

    classified = list(classify_rows(src, "故障履歴", layout))
    assert classified[0][0].index == 5 and classified[-1][1].kind == "subtotal"
    assert classified[0][0].cells[1].value == datetime(2026, 8, 1)
    assert classified[0][0].cells[1].text == "2026-08-01"
    assert len(sample_data_rows(src, "故障履歴", layout, n=5)) == 5
    assert list_kind(src, "故障履歴") == "list"


def test_excel_anchor_and_manual_header_row(tmp_path):
    path, _ = _list_workbook(tmp_path / "list.xlsx")
    src = ExcelSource(path)
    by_anchor = guess_layout(src, "故障履歴", anchors=["発生日", "設備番号", "現象"])
    assert by_anchor.header_rows == [4]
    manual = guess_layout(src, "故障履歴", header_row=4, data_end=10)
    assert manual.data_end == 10 and manual.counts["data"] == 6


def _crosstab_workbook(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "FY2025"
    ws["A1"] = "FY2025 月別・設備別 停止時間集計表"
    ws.merge_cells("A1:N1")
    ws["A2"] = "作成日"
    ws["B2"] = "R8.4.8"
    ws["A4"] = "集計期間：2025/4/1～2026/3/31（単位：分）"
    for col, text in ((1, "ライン"), (2, "設備番号"), (3, "設備名"), (16, "合計")):
        ws.cell(6, col, text).font = Font(bold=True)
        ws.merge_cells(start_row=6, start_column=col, end_row=7, end_column=col)
    ws.cell(6, 4, "2025年")
    ws.merge_cells("D6:L6")
    ws.cell(6, 13, "2026年")
    ws.merge_cells("M6:O6")
    months = [4, 5, 6, 7, 8, 9, 10, 11, 12, 1, 2, 3]
    for i, m in enumerate(months):
        ws.cell(7, 4 + i, f"{m}月").font = Font(bold=True)
    r = 8
    for line, equipments in (("L1", ["CMP-101", "CMP-102", "膜厚計-1"]), ("L2", ["CVD-201", "CVD-202"])):
        start = r
        for eq in equipments:
            ws.cell(r, 2, eq)
            ws.cell(r, 3, f"{eq} 号機")
            for i in range(12):
                ws.cell(r, 4 + i, 100 + i if eq != "CVD-202" or i < 6 else "－")
            ws.cell(r, 16, 1234)
            r += 1
        ws.cell(start, 1, line)
        ws.merge_cells(start_row=start, start_column=1, end_row=r - 1, end_column=1)
        ws.cell(r, 1, f"{line}小計").font = Font(bold=True)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        for i in range(13):
            ws.cell(r, 4 + i, 999)
        r += 1
    ws.cell(r, 1, "総合計")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    ws.cell(r + 2, 1, "※注記")
    ws.cell(r + 3, 1, "※1 前年比は年度合計どうしの比")
    wb.save(path)
    return path


def test_excel_crosstab_two_row_merged_header(tmp_path):
    src = ExcelSource(_crosstab_workbook(tmp_path / "ct.xlsx"))
    layout = guess_layout(src, "FY2025")
    assert layout.table_kind == "crosstab"
    assert layout.header_rows == [6, 7] and layout.data_start == 8
    assert layout.headers[:5] == ["ライン", "設備番号", "設備名", "2025年_4月", "2025年_5月"]
    assert layout.headers[12:16] == ["2026年_1月", "2026年_2月", "2026年_3月", "合計"]
    assert layout.header_levels[0][4] == "2025年" and layout.header_levels[1][4] == "5月"
    assert layout.data_end == 15  # 総合計の行まで
    kinds = {rc.index: (rc.kind, rc.reason) for rc in layout.row_classes}
    assert kinds[11] == ("subtotal", "小計") and kinds[14] == ("subtotal", "小計")
    assert kinds[15] == ("subtotal", "合計")
    assert kinds[10][0] == "data"  # 「膜厚計-1」は小計ではない
    assert kinds[9][0] == "data"  # 縦結合の中（ライン列が空）でも継続行にしない
    assert layout.counts == {"data": 5, "subtotal": 3}
    assert not any("下にも表" in w for w in layout.warnings)


def test_excel_multi_block_reads_first_block_and_warns(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "2025年3月"
    ws["A1"] = "2025年3月 部品交換・保全作業記録"
    r = 3
    for block in ("■ L1ライン", "■ L2ライン"):
        ws.cell(r, 1, block)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        r += 1
        for c, h in enumerate(["日付", "設備ID", "部品名", "数量", "金額(円)"], start=1):
            ws.cell(r, c, h)
        r += 1
        for i in range(6):
            ws.cell(r, 1, datetime(2025, 3, i + 1))
            ws.cell(r, 2, "CMP-101" if i % 2 else "〃")
            ws.cell(r, 3, "POUフィルタ")
            ws.cell(r, 4, "2個")
            ws.cell(r, 5, 37000)
            r += 1
        ws.cell(r, 3, "小計（6件）")
        ws.cell(r, 5, 222000)
        r += 1
        ws.cell(r, 1, "※金額は税抜")
        r += 2
    wb.save(tmp_path / "blocks.xlsx")
    src = ExcelSource(tmp_path / "blocks.xlsx")
    layout = guess_layout(src, "2025年3月")
    assert layout.header_rows == [4] and layout.data_end == 11
    assert layout.counts == {"data": 6, "subtotal": 1}
    assert any("下にも表らしい行" in w for w in layout.warnings)
    kinds = {rc.index: rc.kind for rc in layout.row_classes}
    assert kinds[3] == "title" and kinds[12] == "note" and kinds[14] == "title"


def test_excel_left_right_tables_and_repeated_header(tmp_path):
    wb = Workbook()
    ws = wb.active
    for c, h in enumerate(["日付", "設備ID", "部品名", None, "日付", "設備ID", "部品名"], start=1):
        if h:
            ws.cell(1, c, h)
    for r in range(2, 10):
        ws.cell(r, 1, datetime(2025, 3, r))
        ws.cell(r, 2, f"CMP-10{r}")
        ws.cell(r, 3, "パッド")
        ws.cell(r, 5, datetime(2025, 3, r))
        ws.cell(r, 6, f"CVD-20{r}")
        ws.cell(r, 7, "Oリング")
    ws.cell(10, 1, "日付")
    ws.cell(10, 2, "設備ID")
    ws.cell(10, 3, "部品名")
    ws.cell(11, 1, datetime(2025, 4, 1))
    ws.cell(11, 2, "ETC-301")
    ws.cell(11, 3, "リング")
    wb.save(tmp_path / "lr.xlsx")
    layout = guess_layout(ExcelSource(tmp_path / "lr.xlsx"), "Sheet")
    assert layout.headers == ["日付", "設備ID", "部品名"]
    assert layout.data_end == 9
    assert any("右側" in w for w in layout.warnings)
    assert any("見出しが再び出た" in w for w in layout.warnings)


def test_excel_cell_limit_hidden_columns_and_date1904(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "data"
    for r in range(1, 31):
        for c in range(1, 5):
            ws.cell(r, c, "見出し" if r == 1 else r * c)
    ws.column_dimensions["B"].hidden = True
    wb.epoch = CALENDAR_MAC_1904
    wb.save(tmp_path / "w.xlsx")

    with pytest.raises(UploadError, match="セル数が上限"):
        ExcelSource(tmp_path / "w.xlsx", options={"max_cells": 50})
    src = ExcelSource(tmp_path / "w.xlsx")
    info = src.sheets()[0]
    assert info.date1904 and info.hidden_columns == [2]
    assert (info.max_row, info.max_col) == (30, 4)


def test_excel_invalid_file_is_upload_error(tmp_path):
    p = tmp_path / "broken.xlsx"
    p.write_bytes(b"not a zip")
    with pytest.raises(UploadError, match="Excelファイルとして開けません"):
        ExcelSource(p)


def test_form_like_sheet_is_not_list(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "設備修理報告書"
    rows = [("報告番号", "R2026-00123"), ("発生日", datetime(2026, 9, 14)), ("設備", "CMP-101"),
            ("故障内容", "搬送アーム停止"), ("原因", "エンコーダ不良"), ("処置", "交換"), ("作業時間", 2.5)]
    for i, (label, value) in enumerate(rows, start=3):
        ws.cell(i, 1, label)
        ws.cell(i, 3, value)
    wb.save(tmp_path / "form.xlsx")
    src = ExcelSource(tmp_path / "form.xlsx")
    layout = guess_layout(src, "Sheet")
    assert layout.table_kind in ("form_like", "unknown")
    assert not list_kind(src, "Sheet")


# ---- 見出し・辞書・対応づけ ----

@pytest.mark.parametrize("header, expected", [
    ("停止時間(分)", ("停止時間", "分")),
    ("単価（円）", ("単価", "円")),
    ("作業工数(H)", ("作業工数", "H")),
    ("数量 [個]", ("数量", "個")),
    ("真因(なぜなぜ要約)", ("真因(なぜなぜ要約)", "")),
    ("作業内容（詳細）", ("作業内容（詳細）", "")),
    ("管理No", ("管理No", "")),
    ("(分)", ("(分)", "")),
])
def test_split_header_unit(header, expected):
    assert split_header_unit(header) == expected


def test_month_labels():
    for text in ("4月", "2025年_4月", "2026/04", "R8.4", "2026-04-01", "１２月"):
        assert is_month_label(text), text
    for text in ("合計", "前年比", "設備名", "2026年度"):
        assert not is_month_label(text), text


def test_dictionary_lookup():
    std, how = lookup_header("対応内容")
    assert (std.key, std.role, std.type, how) == ("response_log", "log", "text", "dictionary")
    assert lookup_header("停止時間(分)")[0].key == "downtime"
    assert lookup_header("交換部品_品番")[0].key == "part_no"
    assert lookup_header("管理　NO")[0].key == "record_no"
    std, how = lookup_header("作業内容（詳細）")
    assert (std.key, how) == ("action", "similar")
    assert lookup_header("2025年_4月") is None
    by_key = {c.key: c for c in STANDARD_COLUMNS}
    for key in ("record_no", "occurred_at", "equipment_id", "equipment_name", "line", "process", "failure_category",
                "severity", "symptom", "cause", "action", "response_log", "downtime", "work_hours", "cost", "status",
                "worker", "part_name", "quantity", "unit_price"):
        assert key in by_key, key
    assert by_key["worker"].md != "omit" and by_key["action"].log_candidate


def test_suggest_columns_types_rates_and_log():
    headers = ["管理No", "発生日", "停止時間(分)", "対応内容", "区分", "メモ欄", "担当者", "担当"]
    rows = []
    for i in range(20):
        rows.append([
            f"TR-{i:03d}",
            datetime(2026, 8, i + 1) if i < 19 else "8月末",
            str(i * 10) if i % 5 else "－",
            f"8/{i + 1} 10:00 連絡あり\n8/{i + 1} 11:00 交換完了（田中）",
            ["機械", "電気"][i % 2],
            None,
            "田中",
            "佐藤",
        ])
    result = {s.header: s for s in suggest_columns(headers, rows)}
    assert result["管理No"].key == "record_no" and result["管理No"].matched_by == "dictionary"
    assert result["発生日"].type == "datetime" and result["発生日"].type_error_rate == 0.05
    assert result["停止時間(分)"].unit == "分" and result["停止時間(分)"].display == "停止時間"
    assert result["停止時間(分)"].blank_rate == 0.2 and result["停止時間(分)"].type_error_rate == 0.0
    assert result["対応内容"].role == "log" and result["対応内容"].log and result["対応内容"].md == "body"
    assert result["区分"].key == "failure_category" and result["区分"].inferred_type == "enum"
    assert result["メモ欄"].blank_rate == 1.0 and result["メモ欄"].md == "omit"
    assert result["担当者"].key == "worker" and result["担当者"].md != "omit"
    assert result["担当"].key == "worker_2"  # 同じ標準キーは後の列に番号を付ける
    assert result["管理No"].examples == ["TR-000", "TR-001", "TR-002"]


def test_value_kind():
    assert value_kind(None, "") == "blank"
    assert value_kind(3.5, "3.5") == "number"
    assert value_kind(datetime(2026, 1, 1), "2026-01-01") == "date"
    assert value_kind(datetime(2026, 1, 1, 9), "2026-01-01 09:00") == "datetime"
    assert value_kind(dtime(9, 30), "09:30") == "time"
    assert value_kind("1,234", "1,234") == "number"
    assert value_kind("△120", "△120") == "number"
    assert value_kind("20260803", "20260803") == "date"
    assert value_kind("R5.4.12", "R5.4.12") == "date"
    assert value_kind("令和5年4月12日", "令和5年4月12日") == "date"
    assert value_kind("00123", "00123") == "code"
    assert value_kind("CMP-101", "CMP-101") == "code"
    assert value_kind("#N/A", "#N/A") == "error"
    assert value_kind("設備名", "設備名") == "string"
    assert value_kind("a\nb", "a\nb") == "text"


# ---- samples/tables の実ファイル ----

def _sample(name: str) -> Path:
    path = SAMPLES / name
    if not path.exists():
        pytest.skip(f"サンプルがありません: {name}")
    return path


@pytest.mark.samples
def test_sample_t1_clean_list():
    path = _sample("T1_トラブル対応一覧_2023-2026.xlsx")
    src = open_source(path, path.name)
    names = [s.name for s in src.sheets()]
    assert names[:3] == ["トラブル一覧", "設備マスタ", "集計メモ"]
    layout = guess_layout(src, "トラブル一覧")
    assert layout.table_kind == "list"
    assert layout.header_rows == [1] and (layout.data_start, layout.data_end) == (2, 8096)
    assert layout.counts == {"data": 8095}
    assert layout.headers[:4] == ["管理No", "発生日", "発生時刻", "設備番号"]
    assert "停止時間(分)" in layout.headers
    cols = {s.header: s for s in suggest_columns(layout.headers, sample_data_rows(src, "トラブル一覧", layout))}
    assert cols["停止時間(分)"].key == "downtime" and cols["停止時間(分)"].unit == "分"
    assert cols["処置内容"].key == "action" and cols["担当者"].md != "omit"
    assert list_kind(src, "トラブル一覧") == "list"


@pytest.mark.samples
def test_sample_t2_system_csv():
    path = _sample("T2_設備故障履歴_システム出力.csv")
    sniff = sniff_csv(path)
    assert (sniff.encoding, sniff.delimiter, sniff.preamble_rows, sniff.header_row, sniff.trailer_rows) == (
        "cp932", ",", 4, 5, 1)
    src = open_source(path, path.name)
    layout = guess_layout(src, path.name)
    assert layout.table_kind == "list" and layout.header_rows == [5]
    assert len(layout.headers) == 45 and layout.headers[0] == "管理番号" and layout.headers[-1] == "更新者ID"
    assert layout.counts == {"data": 33095, "subtotal": 1}
    assert layout.row_classes[-1].kind == "subtotal" and layout.row_classes[-1].index == layout.data_end


@pytest.mark.samples
def test_sample_t2_code_table():
    path = _sample("T2_コード表.csv")
    src = open_source(path, path.name)
    layout = guess_layout(src, path.name)
    assert layout.header_rows == [1]
    assert layout.headers == ["コード種別", "コード", "名称", "関連コード", "補足", "使用件数"]


@pytest.mark.samples
def test_sample_t3_crosstab():
    path = _sample("T3_月別停止時間集計.xlsx")
    src = open_source(path, path.name)
    expected = {"FY2023": ([6, 7], 73, 11), "FY2024": ([6, 7], 80, 12), "FY2025": ([6, 7], 80, 12),
                "FY2026": ([6, 7], 80, 12), "全期間": ([5, 6], 79, 12)}
    for sheet, (header_rows, end, subtotals) in expected.items():
        layout = guess_layout(src, sheet)
        assert layout.table_kind == "crosstab", sheet
        assert layout.header_rows == header_rows, sheet
        assert layout.data_end == end, sheet
        assert layout.counts.get("subtotal") == subtotals, sheet
        assert layout.headers[:3] == ["ライン", "設備番号", "設備名"], sheet
        assert layout.counts["data"] in (55, 61), sheet
    fy2025 = guess_layout(src, "FY2025")
    assert fy2025.headers[4] == "2025年_4月" and fy2025.headers[13] == "2026年_1月"
    counts_sheet = guess_layout(src, "件数")
    assert counts_sheet.data_end == 80 and any("下にも表らしい行" in w for w in counts_sheet.warnings)


@pytest.mark.samples
def test_sample_t4_multi_block():
    path = _sample("T4_保全作業記録_部品交換.xlsx")
    src = open_source(path, path.name)
    assert guess_layout(src, "記入要領").table_kind in ("unknown", "form_like")
    march = guess_layout(src, "2025年3月")
    assert march.table_kind == "list" and march.header_rows == [7]
    assert march.headers[:3] == ["実施日", "装置No", "装置名"]
    assert any("右側" in w for w in march.warnings)
    assert any("下にも表らしい行" in w for w in march.warnings)
    assert any("非表示の列" in w for w in march.warnings)
    aug = guess_layout(src, "2026年8月")
    assert aug.header_rows == [5, 6]
    assert "交換部品_品番" in aug.headers
    for sheet in [s.name for s in src.sheets()][1:]:
        layout = guess_layout(src, sheet)
        assert layout.table_kind == "list", sheet
        assert layout.counts.get("data", 0) >= 30, sheet
        assert layout.counts.get("subtotal") == 1, sheet


@pytest.mark.samples
def test_sample_t5_csv_and_two_row_header_xlsx():
    csv_path = _sample("T5_是正処置管理台帳.csv")
    sniff = sniff_csv(csv_path)
    assert (sniff.encoding, sniff.bom, sniff.header_row) == ("utf-8-sig", True, 1)
    src = open_source(csv_path, csv_path.name)
    layout = guess_layout(src, csv_path.name)
    assert layout.counts == {"data": 5043} and layout.headers[0] == "台帳No"

    xlsx = _sample("T5_是正処置管理台帳_2026上期.xlsx")
    xsrc = open_source(xlsx, xlsx.name)
    xl = guess_layout(xsrc, "2026上期")
    assert xl.header_rows == [1, 2] and xl.data_start == 3
    assert xl.headers[0] == "不具合_台帳No" and xl.headers[7] == "対策_暫定対策"
    assert xl.counts == {"data": 728}
    cols = suggest_columns(xl.headers, sample_data_rows(xsrc, "2026上期", xl))
    assert cols[0].key == "record_no" and cols[0].display == "台帳No" and cols[0].matched_by == "dictionary"

# ---- Excel: XML を順に読む実装と、通常モードで開いたときの一致 ----

_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml"


def _raw_workbook(path: Path) -> Path:
    """Excel が書くような形（共有文字列・書式・結合・リンク・コメント・テーブル）と、まれな形を手で組んだブック。"""
    import zipfile

    sheet1 = f"""<worksheet xmlns="{_MAIN}" xmlns:r="{_REL}"><dimension ref="A1:K12"/>
<sheetViews><sheetView workbookViewId="0"><pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/></sheetView></sheetViews>
<cols><col min="2" max="3" width="10" hidden="1" customWidth="1"/></cols>
<sheetData>
<row r="1"><c r="A1" s="1" t="s"><v>0</v></c><c r="B1" s="1" t="s"><v>1</v></c><c r="C1" s="1" t="s"><v>3</v></c>
<c r="D1" t="inlineStr"><is><r><t>リッチ</t></r><r><rPr><b/></rPr><t xml:space="preserve"> 文字</t></r><rPh sb="0" eb="1"><t>ヨミ</t></rPh></is></c><c r="E1" t="s"><v>6</v></c></row>
<row r="2" spans="1:5"><c r="A2" t="s"><v>2</v></c><c r="B2" s="2"><v>46237</v></c><c r="C2" t="str"><f>A1</f><v></v></c><c r="D2" t="b"><v>0</v></c><c r="E2" t="e"><v>#N/A</v></c></row>
<row r="4" hidden="1"><c r="A4"><v>1.5</v></c><c s="3"><v>2</v></c><c t="inlineStr"><is><t/></is></c><c t="inlineStr"/></row>
<row r="5" hidden="0" ht="20" customHeight="1"><c r="A5" s="4"><v>46237.5</v></c><c r="F5" s="3"/></row>
<row r="7"><c r="A7" t="s"><v>3</v></c><c r="B7"><v>99</v></c></row>
<row r="9"><c r="A9" t="d"><v>2026-08-03T10:00:00</v></c><c r="B9" t="s"><v>7</v></c></row>
<row r="10" hidden="true"><c r="G10" t="s"><v>5</v></c></row>
</sheetData>
<autoFilter ref="A1:E5"/>
<mergeCells count="1"><mergeCell ref="A7:C8"/></mergeCells>
<hyperlinks><hyperlink ref="K9" r:id="rId1"/><hyperlink ref="B12" location="Sheet2!A1"/><hyperlink ref="A2" location="x!A1"/><hyperlink ref="B8" location="y!A1"/></hyperlinks>
<tableParts count="1"><tablePart r:id="rId3"/></tableParts>
</worksheet>"""
    sheet2 = f"""<x:worksheet xmlns:x="{_MAIN}"><x:sheetData>
<x:row r="3"><x:c r="A3" t="inlineStr"><x:is><x:t>三</x:t></x:is></x:c></x:row>
<x:row r="2"><x:c r="A2"><x:v>2</x:v></x:c><x:c r="B2"><x:v>5</x:v></x:c></x:row>
<x:row r="3"><x:c r="B3"><x:v>7</x:v></x:c><x:c r="A3"><x:v>8</x:v></x:c></x:row>
<x:row r="5"><x:c r="A6" t="inlineStr"><x:is><x:t>行がずれたセル</x:t></x:is></x:c></x:row>
</x:sheetData></x:worksheet>"""
    sheet3 = f"""<worksheet xmlns="{_MAIN}"><sheetData/><mergeCells><mergeCell ref="B2:C3"/></mergeCells></worksheet>"""
    strings = ["管理No", "発生日", "", "設備", "リンク", "  ", "備考\r\n2行目", "=\"001\""]
    shared = f'<sst xmlns="{_MAIN}" count="{len(strings)}" uniqueCount="{len(strings)}">' + "".join(
        f'<si><t xml:space="preserve">{s}</t></si>' for s in strings) + "</sst>"
    styles = f"""<styleSheet xmlns="{_MAIN}"><numFmts count="1"><numFmt numFmtId="164" formatCode="yyyy/mm/dd hh:mm"/></numFmts>
<fonts count="3"><font><sz val="11"/><name val="Calibri"/></font><font><b/><sz val="11"/><name val="Calibri"/></font><font><strike/><sz val="11"/><name val="Calibri"/></font></fonts>
<fills count="3"><fill><patternFill patternType="none"/></fill><fill><patternFill patternType="gray125"/></fill><fill><patternFill patternType="solid"><fgColor rgb="FFFFFF00"/></patternFill></fill></fills>
<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>
<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>
<cellXfs count="5"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/><xf numFmtId="0" fontId="1" fillId="0" borderId="0" xfId="0" applyFont="1"/><xf numFmtId="14" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/><xf numFmtId="0" fontId="2" fillId="2" borderId="0" xfId="0" applyFont="1" applyFill="1"/><xf numFmtId="164" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/></cellXfs>
</styleSheet>"""
    files = {
        "[Content_Types].xml": f"""<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/xl/workbook.xml" ContentType="{_CT}.sheet.main+xml"/>
<Override PartName="/xl/worksheets/sheet1.xml" ContentType="{_CT}.worksheet+xml"/>
<Override PartName="/xl/worksheets/sheet2.xml" ContentType="{_CT}.worksheet+xml"/>
<Override PartName="/xl/worksheets/sheet3.xml" ContentType="{_CT}.worksheet+xml"/>
<Override PartName="/xl/styles.xml" ContentType="{_CT}.styles+xml"/>
<Override PartName="/xl/sharedStrings.xml" ContentType="{_CT}.sharedStrings+xml"/>
<Override PartName="/xl/comments1.xml" ContentType="{_CT}.comments+xml"/>
<Override PartName="/xl/tables/table1.xml" ContentType="{_CT}.table+xml"/></Types>""",
        "_rels/.rels": f'<Relationships xmlns="{_PKG_REL}"><Relationship Id="rId1" Type="{_REL}/officeDocument" Target="xl/workbook.xml"/></Relationships>',
        "xl/workbook.xml": f"""<workbook xmlns="{_MAIN}" xmlns:r="{_REL}"><sheets>
<sheet name="一覧" sheetId="1" r:id="rId1"/><sheet name="並び" sheetId="2" r:id="rId2"/><sheet name="空" sheetId="3" state="hidden" r:id="rId3"/></sheets></workbook>""",
        "xl/_rels/workbook.xml.rels": f"""<Relationships xmlns="{_PKG_REL}">
<Relationship Id="rId1" Type="{_REL}/worksheet" Target="worksheets/sheet1.xml"/><Relationship Id="rId2" Type="{_REL}/worksheet" Target="worksheets/sheet2.xml"/>
<Relationship Id="rId3" Type="{_REL}/worksheet" Target="worksheets/sheet3.xml"/><Relationship Id="rId4" Type="{_REL}/styles" Target="styles.xml"/>
<Relationship Id="rId5" Type="{_REL}/sharedStrings" Target="sharedStrings.xml"/></Relationships>""",
        "xl/worksheets/sheet1.xml": sheet1,
        "xl/worksheets/sheet2.xml": sheet2,
        "xl/worksheets/sheet3.xml": sheet3,
        "xl/worksheets/_rels/sheet1.xml.rels": f"""<Relationships xmlns="{_PKG_REL}">
<Relationship Id="rId1" Type="{_REL}/hyperlink" Target="https://example.com/a" TargetMode="External"/>
<Relationship Id="rId2" Type="{_REL}/comments" Target="../comments1.xml"/>
<Relationship Id="rId3" Type="{_REL}/table" Target="../tables/table1.xml"/></Relationships>""",
        "xl/comments1.xml": f"""<comments xmlns="{_MAIN}"><authors><author>a</author></authors><commentList>
<comment ref="C3" authorId="0"><text><t>メモ</t></text></comment></commentList></comments>""",
        "xl/tables/table1.xml": f"""<table xmlns="{_MAIN}" id="1" name="表1" displayName="表1" ref="A1:E5"><autoFilter ref="A1:E5"/>
<tableColumns count="5"><tableColumn id="1" name="管理No"/><tableColumn id="2" name="発生日"/><tableColumn id="3" name="設備"/>
<tableColumn id="4" name="列4"/><tableColumn id="5" name="列5"/></tableColumns></table>""",
        "xl/styles.xml": styles,
        "xl/sharedStrings.xml": shared,
    }
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, text in files.items():
            zf.writestr(name, '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n' + text)
    return path


def _normal_mode(path: Path):
    """通常モード（load_workbook）で開いたときの SheetInfo と行（以前の ExcelSource と同じ手順）。"""
    import warnings

    from openpyxl import load_workbook
    from openpyxl.utils import column_index_from_string

    from app.tables import CellInfo, SheetInfo, SourceRow, cell_text

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(path, data_only=True)
    infos, rows = [], {}
    for ws in wb.worksheets:
        max_row = max_col = 0
        for (r, c), cell in ws._cells.items():
            if cell.value is not None and cell.value != "":
                max_row, max_col = max(max_row, r), max(max_col, c)
        anchors = {}
        for rng in ws.merged_cells.ranges:
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    anchors[(r, c)] = (rng.min_row, rng.min_col)
        hidden = set()
        for key, dim in ws.column_dimensions.items():
            if dim.hidden:
                lo = dim.min or column_index_from_string(key)
                hidden.update(range(lo, (dim.max or lo) + 1))
        infos.append(SheetInfo(ws.title, ws.sheet_state != "visible", max_row, max_col, [t.ref for t in ws.tables.values()],
                               wb.epoch == CALENDAR_MAC_1904, sorted(hidden), ws.freeze_panes, ws.auto_filter.ref or None))
        out = []
        for r in range(1, max_row + 1):
            cells = []
            for c in range(1, max_col + 1):
                xl = ws._cells.get((r, c))
                if xl is None:
                    cells.append(CellInfo(None, "", merged_anchor=anchors.get((r, c))))
                    continue
                cells.append(CellInfo(xl.value, cell_text(xl.value), xl.number_format, bool(xl.font.b), bool(xl.font.strike),
                                      xl.fill.fill_type not in (None, "none"), anchors.get((r, c))))
            dim = ws.row_dimensions.get(r)
            out.append(SourceRow(r, cells, bool(dim is not None and dim.hidden)))
        rows[ws.title] = out
    return infos, rows


def test_excel_streaming_reader_matches_normal_mode(tmp_path):
    path = _raw_workbook(tmp_path / "raw.xlsx")
    infos, expected = _normal_mode(path)
    src = ExcelSource(path)
    assert src.sheets() == infos
    for info in infos:
        assert list(src.rows(info.name)) == expected[info.name], info.name
        assert list(src.rows(info.name, 3, 4)) == expected[info.name][2:6], info.name
    # 手で組んだ形がねらいどおり読めていること
    first = {r.index: r for r in expected["一覧"]}
    assert infos[0].max_col == 11 and infos[0].max_row == 12  # リンク先の値が入った K9・B12 まで
    assert first[9].cells[10].value == "https://example.com/a" and first[12].cells[1].value == "Sheet2!A1"
    assert first[2].cells[0].value == "" and first[1].cells[3].value == "リッチ 文字"
    assert first[7].cells[1].value is None and first[7].cells[1].merged_anchor == (7, 1)  # 結合範囲の値は消える
    assert first[3].cells[2].number_format == "General" and first[3].cells[1].number_format is None  # コメントのセル
    assert first[4].hidden and not first[5].hidden and first[10].hidden
    assert infos[0].table_ranges == ["A1:E5"] and infos[0].hidden_columns == [2, 3] and infos[0].freeze_panes == "A2"
    assert [r.cells[0].value for r in expected["並び"]][:3] == [None, 2, 8] and infos[2].hidden
    # 一度数えたシートの大きさを渡すと数え直さずに同じ結果
    again = ExcelSource(path, options={"sheet_stats": src.sheet_stats()})
    assert again.sheets() == infos and list(again.rows("一覧")) == expected["一覧"]


@pytest.mark.samples
def test_sample_excel_streaming_reader_matches_normal_mode():
    for name in ("T5_是正処置管理台帳_2026上期.xlsx", "T3_月別停止時間集計.xlsx"):
        path = _sample(name)
        infos, expected = _normal_mode(path)
        src = ExcelSource(path)
        assert src.sheets() == infos
        for info in infos:
            assert list(src.rows(info.name)) == expected[info.name], (name, info.name)


# ---- 行の分類: 集計の語・注記・縦結合のキー（R1） ----

def _auto_records(source, sheet, layout):
    from app.tables import read_records
    from app.tables import spec_from_suggestions

    suggestions = suggest_columns(layout.headers, sample_data_rows(source, sheet, layout))
    spec = spec_from_suggestions("テスト", layout, suggestions, {})
    records, _issues, stats = read_records(source, {"sheet": sheet}, layout, spec)
    return records, stats


def test_values_ending_in_kei_are_not_subtotals(tmp_path):
    """「設計」「会計」の部署や、文の中の「合計」は小計・合計の行にしない（語全体が集計の語のときだけ）。"""
    depts = ["設計", "製造", "品証", "会計"]
    lines = ["部署,日付,工数(h)"] + [f"{depts[i % 4]},2026/04/{i % 28 + 1:02d},{i % 7 + 1}" for i in range(40)]
    path = _write(tmp_path / "dept.csv", "\r\n".join(lines) + "\r\n", "cp932")
    src = CsvSource(path, "dept.csv")
    layout = guess_layout(src, "dept.csv")
    assert layout.data_end == 41 and layout.counts == {"data": 40}
    records, stats = _auto_records(src, "dept.csv", layout)
    assert len(records) == 40 and not stats.excluded

    # 日付などのキー列が空の「ブロック計」の行は、これまでどおり小計
    path = _write(tmp_path / "dept2.csv", "\r\n".join(lines + ["ブロック計,,160"]) + "\r\n", "cp932")
    src = CsvSource(path, "dept2.csv")
    layout = guess_layout(src, "dept2.csv")
    assert layout.counts == {"data": 40, "subtotal": 1}

    lines = ["管理No,発生日,設備,内容"]
    for i in range(1, 101):
        text = "合計カウンタ不良" if i == 40 else "搬送停止"
        lines.append(f"T{i:03d},2026/04/{i % 28 + 1:02d},CMP-101,{text}")
    path = _write(tmp_path / "total.csv", "\r\n".join(lines) + "\r\n", "cp932")
    src = CsvSource(path, "total.csv")
    layout = guess_layout(src, "total.csv")
    assert layout.data_end == 101 and layout.counts == {"data": 100}

    # 本物の集計の行はこれまでどおり: 「A部署 計」「4月計」「合計件数」
    lines = ["部署,日付,工数(h)", "設計,2026/04/01,3", "A部署 計,,3", "会計,2026/04/02,2", "4月計,,5", "合計件数,2,"]
    path = _write(tmp_path / "agg.csv", "\r\n".join(lines) + "\r\n", "cp932")
    src = CsvSource(path, "agg.csv")
    kinds = {rc.index: (rc.kind, rc.reason) for rc in guess_layout(src, "agg.csv").row_classes}
    assert kinds[2][0] == "data" and kinds[4][0] == "data"
    assert kinds[3] == ("subtotal", "小計") and kinds[5] == ("subtotal", "小計") and kinds[6] == ("subtotal", "合計")


def _log_book(path: Path, extra: dict[int, str]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "一覧"
    for c, h in enumerate(["管理No", "発生日", "設備", "現象", "対応内容"], start=1):
        ws.cell(1, c, h)
    r = 2
    for i in range(1, 41):
        ws.cell(r, 1, f"T{i:03d}")
        ws.cell(r, 2, datetime(2026, 4, i % 28 + 1))
        ws.cell(r, 3, "CMP-101")
        ws.cell(r, 4, "停止")
        ws.cell(r, 5, "点検した")
        r += 1
        if i in extra:
            ws.cell(r, 5, extra[i])
            r += 1
    wb.save(path)
    return path


def test_mark_rows_in_text_column_are_continuations(tmp_path):
    """文章の列だけに「●…」「※…」とある行は注記・表題ではなく前の記録の続き（表を途中で終わらせない）。"""
    path = _log_book(tmp_path / "cont.xlsx", {10: "●再発防止としてセンサ追加", 20: "※部品は後日交換予定"})
    src = ExcelSource(path)
    layout = guess_layout(src, "一覧")
    assert layout.data_end == 43
    records, stats = _auto_records(src, "一覧", layout)
    assert len(records) == 40 and not stats.excluded
    texts = {r.values.get("record_no") or r.key: "".join(str(v) for v in r.values.values()) for r in records}
    assert any("●再発防止としてセンサ追加" in t for t in texts.values())
    assert any("※部品は後日交換予定" in t for t in texts.values())

    # 範囲を手で決めても同じ（注記として捨てない）
    manual = guess_layout(src, "一覧", header_rows=[1], data_end=43)
    records, stats = _auto_records(src, "一覧", manual)
    assert len(records) == 40 and "注記" not in stats.excluded
    assert any("※部品は後日交換予定" in "".join(str(v) for v in r.values.values()) for r in records)

    # 左端の列の「※」はこれまでどおり注記（表の終わり）
    wb = Workbook()
    ws = wb.active
    ws.append(["管理No", "発生日", "設備", "現象", "対応内容"])
    for i in range(1, 11):
        ws.append([f"T{i:03d}", datetime(2026, 4, i), "CMP-101", "停止", "点検した"])
    ws.append(["※金額は税抜"])
    wb.save(tmp_path / "note.xlsx")
    layout = guess_layout(ExcelSource(tmp_path / "note.xlsx"), ws.title)
    assert layout.data_end == 11 and {rc.index: rc.kind for rc in layout.row_classes}[12] == "note"


def test_vertically_merged_keys_make_one_record(tmp_path):
    """管理No・日付・設備を縦に結合して対応内容を複数行に書いた表は、1件の記録にまとめる。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "一覧"
    for c, h in enumerate(["管理No", "発生日", "設備", "対応内容"], start=1):
        ws.cell(1, c, h)
    r = 2
    for t in range(10):
        ws.cell(r, 1, f"T{t:03d}")
        ws.cell(r, 2, datetime(2026, 4, t + 1))
        ws.cell(r, 3, "CMP-101")
        for k in range(3):
            ws.cell(r + k, 4, f"{k + 1}行目の対応")
        for c in (1, 2, 3):
            ws.merge_cells(start_row=r, start_column=c, end_row=r + 2, end_column=c)
        r += 3
    wb.save(tmp_path / "merged.xlsx")
    src = ExcelSource(tmp_path / "merged.xlsx")
    layout = guess_layout(src, "一覧")
    assert layout.data_end == 31
    records, stats = _auto_records(src, "一覧", layout)
    assert len(records) == 10 and stats.continuation_merged == 20
    assert all("#" not in r.key for r in records)
    joined = "".join(str(v) for v in records[0].values.values())
    assert "1行目の対応" in joined and "3行目の対応" in joined


# ---- CSV の壊れ方（R1） ----

def test_unclosed_quote_is_reported_as_an_error(tmp_path):
    """" の閉じ忘れで以降の行が1つの値になったら、黙って件数を減らさず、確定を止めるエラーにする。"""
    from app.tables import has_blocking, run_checks

    lines = ["管理No,設備名,現象", "T0001,搬送機,停止", 'T0002,"研磨機,異音']
    lines += [f"T{i:04d},搬送機,停止" for i in range(3, 5003)]
    path = _write(tmp_path / "q.csv", "\r\n".join(lines) + "\r\n", "cp932")
    src = open_source(path, "q.csv")
    rows = list(src.rows())
    assert len(rows) == 3 and src.unclosed_quote_row == 3

    layout = guess_layout(src, "q.csv")
    records, stats = _auto_records(src, "q.csv", layout)
    assert stats.unclosed_quote_row == 3
    from app.tables import spec_from_suggestions

    spec = spec_from_suggestions("テスト", layout, suggest_columns(layout.headers, sample_data_rows(src, "q.csv", layout)), {})
    issues = run_checks(records, spec, stats)
    assert has_blocking(issues) and any(i.code == "unclosed_quote" and "3行目" in i.message for i in issues)

    # 正しく閉じた複数行の値は問題にしない
    ok = _write(tmp_path / "ok.csv", '管理No,内容\r\nT1,"1行目\r\n2行目"\r\nT2,停止\r\nT3,"最後\r\nの行"\r\n', "cp932")
    src = open_source(ok, "ok.csv")
    assert len(list(src.rows())) == 4 and src.unclosed_quote_row is None


def test_csv_field_limit_error_becomes_upload_error(tmp_path):
    """1つの値が大きすぎて csv モジュールが読めないときは、500 ではなく案内付きの UploadError にする。"""
    import csv

    lines = ["管理No,設備名,現象", 'T0002,"研磨機,異音'] + [f"T{i:04d},搬送機,停止" for i in range(3, 200)]
    path = _write(tmp_path / "big.csv", "\r\n".join(lines) + "\r\n", "cp932")
    src = open_source(path, "big.csv")
    old = csv.field_size_limit(1000)
    try:
        with pytest.raises(UploadError, match="閉じていない"):
            list(src.rows())
    finally:
        csv.field_size_limit(old)


def test_excel_or_binary_named_csv_is_rejected(tmp_path):
    """拡張子が .csv でも中身が Excel・バイナリなら、文字コードの問題ではなくその旨を出す。"""
    wb = Workbook()
    wb.active.append(["管理No", "内容"])
    wb.save(tmp_path / "x.csv")
    with pytest.raises(UploadError, match="中身はExcelファイル"):
        open_source(tmp_path / "x.csv", "x.csv")
    (tmp_path / "b.txt").write_bytes(bytes(range(0, 32)) * 64)
    with pytest.raises(UploadError, match="テキストのCSVではありません"):
        open_source(tmp_path / "b.txt", "b.txt")
    # BOMなしの UTF-16 はバイナリ扱いしない
    u16 = tmp_path / "u16.csv"
    u16.write_bytes("管理No,内容\r\nT1,停止\r\nT2,異音\r\n".encode("utf-16-le"))
    assert open_source(u16, "u16.csv").encoding == "utf-16-le"


def test_damaged_sheet_part_is_an_upload_error_not_zero_cells(tmp_path):
    """シートの圧縮データが壊れていたら、セル数の目安を 0 にせず「壊れている」と伝える。"""
    import zipfile as _zip

    from app.tables import estimate_cell_count

    wb = Workbook()
    for r in range(1, 200):
        wb.active.append([f"管理No{r}", r, "対応内容の長い文章" * 5])
    path = tmp_path / "ok.xlsx"
    wb.save(path)
    data = bytearray(path.read_bytes())
    with _zip.ZipFile(path) as zf:
        info = zf.getinfo("xl/worksheets/sheet1.xml")
    start = info.header_offset + 30 + len(info.filename.encode()) + len(info.extra)
    for i in range(start, start + min(info.compress_size, 64)):
        data[i] ^= 0xFF
    broken = tmp_path / "broken.xlsx"
    broken.write_bytes(bytes(data))
    assert estimate_cell_count(path) > 0
    with pytest.raises(UploadError, match="壊れている"):
        estimate_cell_count(broken)


# ====================================================================================================
# 元 tests/test_tables_pipe.py
# 一覧表パイプライン（spec / normalize / checks / records / markdown / outputs / pipeline / store）のテスト。
# ====================================================================================================

SAMPLES_tables_pipe = Path(__file__).resolve().parent.parent / "samples" / "tables"


# ---- 合成データ ---------------------------------------------------------------------------

LIST_HEADERS = ["管理No", "発生日", "設備番号", "設備名", "故障区分", "現象", "処置", "停止時間(h)", "費用(円)", "担当者", "対応内容"]


def make_list_book(path: Path, include_tr004: bool = True, tr001_hours=1.5) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "2026年8月"
    ws["A1"] = "トラブル対応一覧 2026年度"
    ws.merge_cells("A1:K1")
    for i, h in enumerate(LIST_HEADERS, start=1):
        ws.cell(row=3, column=i, value=h).font = Font(bold=True)
    rows = [
        (4, ["TR-001", datetime(2026, 8, 3), "ｃｍｐ－101", "CMP研磨装置1号機", "機械", "スラリー流量低下", "フィルター交換",
             tr001_hours, "1,234", "田中", "8/3 10:00 田中：ライン停止の連絡あり。\n8/3 10:20 原点復帰で復旧。"]),
        (5, ["TR-002", "R8.8.10", "CMP-101", None, "電気", "#1 アラーム", "- 再起動", "90分", "▲500", "佐藤", None]),
        (6, [None, None, None, None, None, None, "続き：ケーブル交換", None, None, None, None]),
        (7, ["TR-003", "8/15", "CVD-201", "CVD 1号機", "電気", "真空度低下", "Oリング交換", 2, "300-", "高橋", None]),
        (8, ["TR-004", 20260820, "CVD-201", None, "機械", "+停止(セル先頭が記号)", "@cmd", 0.5, 100, "鈴木", None]),
        (9, ["TR-005", datetime(2026, 8, 21), "CVD-201", "CVD 1号機", "機械", "非表示の行", "x", 9, 0, "", None]),
        (10, ["TR-006", datetime(2026, 8, 22), "CVD-201", "CVD 1号機", "機械", "取り消し線の行", "x", 9, 0, "", None]),
        (11, ["TR-007", datetime(2026, 7, 30), "ETC-301", "エッチャー1号機", "その他", "7月の記録", "確認", 1, 0, "", None]),
    ]
    if not include_tr004:
        rows = [r for r in rows if r[1][0] != "TR-004"]
    hours = sum(1.5 if r[1][0] == "TR-002" else r[1][7] for r in rows
                if r[1][0] not in (None, "TR-005", "TR-006"))
    for row_no, values in rows:
        for col, v in enumerate(values, start=1):
            if v is not None:
                ws.cell(row=row_no, column=col, value=v)
    ws.merge_cells("D4:D5")  # 設備名の縦結合
    ws.row_dimensions[9].hidden = True
    for col in range(1, 12):
        ws.cell(row=10, column=col).font = Font(strike=True)
    ws.cell(row=12, column=1, value="小計")
    ws.cell(row=12, column=8, value=hours)
    ws.cell(row=13, column=1, value="合計")
    ws.cell(row=13, column=8, value=hours)
    ws.cell(row=15, column=1, value="※停止時間は時間で記入")
    wb.save(path)
    return path


def list_spec_dict(**markdown) -> dict:
    md = {"file_prefix": "トラブル対応一覧", "group_by": "month", "max_records_per_file": 300}
    md.update(markdown)
    return {
        "name": "トラブル対応一覧",
        "description": "保全課のトラブル対応記録",
        "columns": [
            {"key": "record_no", "display": "管理No", "headers": ["管理No"], "type": "code", "role": "key", "required": True},
            {"key": "occurred_at", "display": "発生日", "headers": ["発生日"], "type": "date", "role": "date", "required": True},
            {"key": "equipment_id", "display": "設備番号", "headers": ["設備番号"], "type": "code", "role": "entity",
             "normalize": ["nfkc", "upper"]},
            {"key": "equipment_name", "display": "設備名", "headers": ["設備名"], "type": "string", "role": "entity_label",
             "fill_down_blank": True},
            {"key": "failure_category", "display": "故障区分", "headers": ["故障区分"], "type": "enum", "role": "category",
             "allowed": ["機械", "電気"]},
            {"key": "symptom", "display": "現象", "headers": ["現象"], "type": "text", "role": "text", "md": "body"},
            {"key": "action", "display": "処置", "headers": ["処置"], "type": "text", "role": "text", "md": "body"},
            {"key": "downtime", "display": "停止時間", "headers": ["停止時間"], "type": "number", "role": "measure", "unit": "分"},
            {"key": "cost", "display": "費用", "headers": ["費用"], "type": "number", "role": "measure", "unit": "円"},
            {"key": "worker", "display": "担当者", "headers": ["担当者"], "type": "string", "role": "person"},
            {"key": "response_log", "display": "対応内容", "headers": ["対応内容"], "type": "text", "role": "log", "md": "body"},
        ],
        "record": {"key": ["record_no"], "fallback_key": ["occurred_at", "equipment_id", "symptom:20"]},
        "period": {"date_column": "occurred_at"},
        "log_stage": {"column": "response_log", "mask": []},
        "markdown": md,
    }


def read_list(path: Path, spec):
    source = open_source(path, path.name)
    layout = guess_layout(source, "2026年8月")
    return read_records(source, {"sheet": "2026年8月"}, layout, spec)


def august(records) -> list[dict]:
    return [r.to_dict() for r in records if str(r.values.get("occurred_at", "")).startswith("2026-08")]


def _no_blank_inside_records(text: str) -> bool:
    lines = text.split("\n")
    for i, line in enumerate(lines[:-1]):
        if line == "":
            nxt = lines[i + 1]
            prev = lines[i - 1] if i else ""
            if not (nxt.startswith("## ") or prev.startswith("# ")):
                return False
    return True


# ---- 値の変換 -------------------------------------------------------------------------------

@pytest.mark.parametrize("type_, value, text, expected, flag", [
    ("date", None, "令和8年8月3日", "2026-08-03", None),
    ("date", None, "R8.8.3", "2026-08-03", None),
    ("date", None, "H31.4.30", "2019-04-30", None),
    ("date", None, "昭和64年1月7日", "1989-01-07", None),
    ("date", None, "令和元年5月1日", "2019-05-01", None),
    ("date", None, "20260803", "2026-08-03", None),
    ("date", 20260803, "20260803", "2026-08-03", None),
    ("date", None, "2026/8/3(月)", "2026-08-03", None),
    ("date", None, "８/５", "2026-08-05", "year_inferred"),
    ("date", None, "2/5", "2027-02-05", "year_inferred"),
    ("date", 46237.0, "46237", "2026-08-03", None),
    ("datetime", None, "2026/08/03 14:20:00", "2026-08-03 14:20", None),
    ("datetime", None, "20260803 142000", "2026-08-03 14:20", None),
    ("datetime", None, "2026年8月3日 14時20分", "2026-08-03 14:20", None),
    ("datetime", datetime(2026, 8, 3, 0, 0), "", "2026-08-03", None),
    ("time", None, "142000", "14:20", None),
    ("time", None, "1420", "14:20", None),
    ("time", 0.5, "0.5", "12:00", None),
    ("time", timedelta(hours=1, minutes=5), "1:05", "01:05", None),
    ("time", dtime(9, 5), "09:05", "09:05", None),
])
def test_convert_dates_and_times(type_, value, text, expected, flag):
    col = ColumnSpec("c", "列", type=type_)
    cctx = ConvertContext(fiscal_year=2026, fiscal_start=4)
    out, error, got_flag = convert_cell(value, text, col, cctx)
    assert error is None
    assert out == expected
    assert got_flag == flag


@pytest.mark.parametrize("value, text, unit, header_unit, fmt, expected", [
    (None, "1,234", "", "", None, 1234),
    (None, "１，２３４", "", "", None, 1234),
    (None, "△500", "", "", None, -500),
    (None, "▲1,000", "円", "", None, -1000),
    (None, "300-", "", "", None, -300),
    (None, "98.5%", "%", "", None, 98.5),
    (0.985, "0.985", "%", "", "0.0%", 98.5),
    (None, "1.5h", "分", "", None, 90),
    (None, "2時間", "分", "", None, 120),
    (1.5, "1.5", "分", "h", None, 90),
    (None, "90分", "分", "h", None, 90),
    (timedelta(hours=2, minutes=30), "2:30", "分", "", None, 150),
    (timedelta(minutes=90), "1:30", "時間", "", None, 1.5),
    (None, "¥12,000", "円", "", None, 12000),
    (None, "1.23E+3", "", "", None, 1230),
    (3.0000000001, "3", "", "", None, 3),
])
def test_convert_numbers(value, text, unit, header_unit, fmt, expected):
    col = ColumnSpec("c", "列", type="number", unit=unit)
    out, error, _flag = convert_cell(value, text, col, ConvertContext(), fmt, header_unit)
    assert error is None
    assert out == expected


def test_convert_blank_na_errors_codes_and_maps():
    cctx = ConvertContext(na_tokens=frozenset({"-", "N/A"}))
    num = ColumnSpec("n", "数", type="number", unit="分")
    assert convert_cell(None, "－", num, cctx) == (None, None, None)
    assert convert_cell("#N/A", "#N/A", num, cctx) == (None, None, "error_value")
    out, error, _ = convert_cell(None, "約3", num, cctx)
    assert error and out == "約3"
    out, error, _ = convert_cell(None, "3件", num, cctx)
    assert error  # 分に換算できない単位
    date_col = ColumnSpec("d", "日付", type="date")
    out, error, flag = convert_cell(None, "8/5", date_col, ConvertContext())
    assert flag == "year_missing" and error
    code = ColumnSpec("k", "コード", type="code", normalize=["nfkc", "upper"])
    assert convert_cell(123, "123", code, cctx, "00000")[0] == "00123"
    assert convert_cell(None, "ｃｍｐ－１０１", code, cctx)[0] == "CMP-101"
    enum = ColumnSpec("e", "区分", type="enum", value_map={"電気系": "電気"})
    assert convert_cell(None, "電気系", enum, cctx)[0] == "電気"
    text = ColumnSpec("t", "文章", type="text")
    assert convert_cell(None, "  ＡＢＣ　 1\r\n  2行目  ", text, cctx)[0] == "ABC 1\n2行目"


def test_year_context():
    assert find_year_context([("タイトル行", "2026年度 故障一覧")])["fiscal_year"] == 2026
    assert find_year_context([("シート名", "FY25")])["fiscal_year"] == 2025
    assert find_year_context([("x", "令和8年度")])["fiscal_year"] == 2026
    ym = find_year_context([("ファイル名", "故障履歴_2026-08")])
    assert (ym["year"], ym["month"], ym["source"]) == (2026, 8, "ファイル名")
    assert find_year_context([("ファイル名", "T1_トラブル対応一覧_2023-2026")])["year"] is None
    assert find_year_context([("a", ""), ("b", "202608")])["month"] == 8


# ---- spec ------------------------------------------------------------------------------------

def test_spec_roundtrip_validate_and_resolve():
    spec = spec_from_dict(list_spec_dict())
    assert validate_spec(spec) == []
    again = spec_from_dict(json.loads(json.dumps(spec_to_dict(spec), ensure_ascii=False)))
    assert spec_hash(again) == spec_hash(spec)
    assert spec.date_key == "occurred_at" and spec.file_prefix == "トラブル対応一覧"

    bad = spec_from_dict({"name": "", "columns": [
        {"key": "1bad", "display": "x", "type": "weird", "role": "person"},
        {"key": "a", "display": "", "type": "number", "md": "nope"},
        {"key": "a", "display": "dup"},
    ], "record": {"key": ["missing"]}, "period": {"date_column": "a"},
        "markdown": {"group_by": "entity_month"},
        "log_stage": {"column": "zzz"}})
    errors = validate_spec(bad)
    joined = "\n".join(errors)
    for fragment in ("表の名前", "1bad", "重複", "表示名", "型「weird」", "mdでの扱い", "記録キーの列「missing」",
                     "日付", "対象×月", "zzz"):
        assert fragment in joined, fragment

    res = resolve_columns(spec, ["管理No", "発生日", "設備番号", "設備名", "故障区分", "現象", "処置内容", "停止時間(h)",
                                 "費用(円)", "担当者", "新しい列"])
    assert res.positions["downtime"] == 7 and res.positions["cost"] == 8
    assert "action" not in res.positions and "処置" in res.missing
    assert res.unused_headers == ["処置内容", "新しい列"]
    assert res.missing_required == []


def test_spec_from_suggestions(tmp_path):
    path = make_list_book(tmp_path / "list.xlsx")
    source = open_source(path, path.name)
    layout = guess_layout(source, "2026年8月")
    suggestions = suggest_columns(layout.headers, sample_data_rows(source, "2026年8月", layout))
    spec = spec_from_suggestions("トラブル", layout, suggestions, {"description": "説明", "group_by": "entity_month"})
    assert validate_spec(spec) == []
    keys = [c.key for c in spec.columns]
    assert keys[:4] == ["record_no", "occurred_at", "equipment_id", "equipment_name"]
    assert spec.record["key"] == ["record_no"]
    assert spec.record["fallback_key"][0] == "occurred_at"
    assert spec.period == {"date_column": "occurred_at"}
    assert spec.log_stage is not None and spec.log_stage.column == "response_log"
    assert spec.markdown["group_by"] == "entity_month"


# ---- 一覧: 読み込み → チェック → md → zip ---------------------------------------------------------

def test_list_read_normalize_and_checks(tmp_path):
    path = make_list_book(tmp_path / "list.xlsx")
    spec = spec_from_dict(list_spec_dict())
    records, row_issues, stats = read_list(path, spec)
    by_key = {r.key: r for r in records}
    assert list(by_key) == ["TR-001", "TR-002", "TR-003", "TR-004", "TR-007"]
    r1, r2, r3, r4 = (by_key[k].values for k in ("TR-001", "TR-002", "TR-003", "TR-004"))
    assert r1["equipment_id"] == "CMP-101" and by_key["TR-001"].originals["equipment_id"] == "ｃｍｐ－101"
    assert r1["downtime"] == 90 and r1["cost"] == 1234 and r1["occurred_at"] == "2026-08-03"
    assert r2["occurred_at"] == "2026-08-10" and r2["equipment_name"] == "CMP研磨装置1号機"  # 縦結合を埋める
    assert r2["action"] == "- 再起動\n続き:ケーブル交換"  # 継続行を連結（NFKC）
    assert r2["downtime"] == 90 and r2["cost"] == -500
    assert r3["occurred_at"] == "2026-08-15" and "年を" in by_key["TR-003"].warnings[0]
    assert r3["cost"] == -300 and r3["equipment_name"] == "CVD 1号機"
    assert r4["occurred_at"] == "2026-08-20" and r4["equipment_name"] == "CVD 1号機"  # 空欄＝上と同じ
    assert stats.filled_down == {"equipment_name": 1}
    assert stats.continuation_merged == 1
    assert stats.excluded == {"非表示の行": 1, "取り消し線の行": 1, "小計": 1, "合計": 1}
    assert stats.year_inferred == 1 and stats.year_context["fiscal_year"] == 2026
    assert stats.records == 5 and stats.months == {"2026-07": 1, "2026-08": 4}
    assert len(stats.reconcile) == 2 and all(rc["expected"] == rc["actual"] == 390 for rc in stats.reconcile)
    assert row_issues == []

    issues = run_checks(records, spec, stats)
    codes = [i.code for i in issues]
    assert "year_inferred" in codes and "not_allowed" in codes
    assert codes.count("excluded_rows") == 2
    assert not has_blocking(issues)

    # 見出しが足りない・型エラーが多い → 確定を止める
    broken = spec_from_dict(list_spec_dict())
    broken.column("symptom").required = True
    broken.column("symptom").headers = ["不具合内容"]
    broken.column("symptom").display = "不具合"
    broken.column("failure_category").type = "number"
    _recs, _row_issues, st2 = read_list(path, broken)
    issues2 = run_checks(_recs, broken, st2)
    assert {"required_missing", "type_error_rate"} <= {i.code for i in issues2 if i.level == "error"}
    assert len(_row_issues) == 5 and _row_issues[0].column == "故障区分"


def test_list_markdown_and_determinism(tmp_path):
    path = make_list_book(tmp_path / "list.xlsx")
    spec = spec_from_dict(list_spec_dict())
    records, _issues, _stats = read_list(path, spec)
    aug = august(records)
    ai = {
        "TR-001": {"status": "ok", "result": {
            "entries": [{"id": "e1", "segs": ["s1"], "t": ["連絡", "初動"]}],
            "incident": {"root_cause": {"v": "フィルター目詰まり", "certainty": "確定"},
                         "permanent_actions": [{"v": "フィルター交換"}],
                         "parts": [{"name": "フィルター", "model": "FL-10", "qty_q": "×1"}]}}},
        "TR-002": {"status": "flagged", "result": {"incident": {"root_cause": {"v": "出してはいけない"}}}},
    }
    files = render_all(spec, aug, ai)
    # 作るのは RAG に入れる記録ファイルだけ（集計・データセット説明は作らない）
    assert [f.name for f in files] == ["トラブル対応一覧_2026-08.md"]
    rec = files[0].text
    assert rec.startswith("# トラブル対応一覧 2026年8月の記録\n\n- データ種別: トラブル対応一覧（1行＝1件）の記録\n"
                          "- 対象期間: 2026-08-01〜2026-08-31\n- このファイルの記録: 4件（2026年8月の全件）\n\n## ")
    block1 = rec.split("\n\n")[2]
    assert block1.split("\n")[0] == "## 【TR-001】CMP研磨装置1号機（CMP-101）スラリー流量低下｜2026-08-03"
    for line in ("- 管理No: TR-001", "- 発生日: 2026-08-03（2026年8月）", "- 設備: CMP研磨装置1号機（CMP-101）",
                 "- 停止時間: 90分", "- 費用: 1,234円", "- 対応の要点（AI抽出）:", "  原因: フィルター目詰まり（確定）",
                 "  恒久処置: フィルター交換", "  使用部品: フィルター FL-10 ×1", "- 対応の時系列:",
                 "- 出典: list.xlsx（管理No TR-001）"):
        assert line in block1.split("\n"), line
    timeline = [ln for ln in block1.split("\n") if ln.startswith("  1. ")]
    assert timeline and timeline[0].startswith("  1. 2026-08-03 10:00［連絡・初動］") and "CMP研磨装置1号機" in timeline[0]
    assert "  2. 2026-08-03 10:20" in block1
    assert "- 担当者: 田中" in rec and "田中｜" in rec  # 人名の列も出す
    assert "出してはいけない" not in rec
    assert "- 処置:\n  \\- 再起動\n  続き:ケーブル交換" in rec
    assert "- 現象: +停止(セル先頭が記号)" in rec
    assert "|" not in rec.replace("｜", "")  # パイプ表を使わない
    assert _no_blank_inside_records(rec)
    assert "\r" not in rec and rec.endswith("\n") and not rec.endswith("\n\n")


    # 決定的: 入力の順番を変えても同じバイト列
    shuffled = list(aug)
    random.Random(3).shuffle(shuffled)
    again = render_all(spec, shuffled, ai)
    assert [hashlib.sha256(f.data).hexdigest() for f in again] == [hashlib.sha256(f.data).hexdigest() for f in files]

    # 記録ファイルは件数では分けない（月ごとに1ファイル）。設備×月の設定も同じ
    assert [f.name for f in files if f.kind == "records"] == ["トラブル対応一覧_2026-08.md"]
    files3 = render_all(spec_from_dict(list_spec_dict(group_by="entity_month")), aug, {})
    ent = next(f for f in files3 if f.name == "トラブル対応一覧_CMP-101_2026-08.md")
    assert "- このファイルの記録: 2件（CMP-101の2026年8月の全件）" in ent.text

    # 7月と8月の両方がある取り込みは月ごとのファイルになる
    all_names = [f.name for f in render_all(spec, [r.to_dict() for r in records], {})]
    assert "トラブル対応一覧_2026-07.md" in all_names and "トラブル対応一覧_2026-08.md" in all_names


def test_zip_csv_and_formula_guard():
    assert tables.guard_formula("=HYPERLINK(\"x\")") == "'=HYPERLINK(\"x\")"
    for s in ("+1", "-2", "@SUM", "\tx", "\rx"):
        assert tables.guard_formula(s).startswith("'")
    assert tables.guard_formula("通常の文字") == "通常の文字"
    assert tables.guard_formula(-5) == "-5" and tables.guard_formula(None) == ""
    data = tables.issues_csv([Issue("warning", "x", "=1+1", row=3, column="-列")])
    assert data.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    assert rows[1] == ["警告", "x", "3", "'-列", "'=1+1"]
    # zip は RAG に入れる md だけ（フォルダ分けなし・管理用CSVなし）。並びを変えても同じバイト列
    zipped = tables.build_zip([("b.md", b"B\n"), ("a.md", b"A\n")])
    assert zipped == tables.build_zip([("a.md", b"A\n"), ("b.md", b"B\n")])
    with zipfile.ZipFile(io.BytesIO(zipped)) as zf:
        assert zf.namelist() == ["a.md", "b.md"]


# ---- CSV ---------------------------------------------------------------------------------------

def test_system_csv_quirks(tmp_path):
    text = ("出力日時,2026/09/01 10:00:00\n抽出条件,期間=2026/08\n\n"
            "管理番号,発生日,発生時刻,設備コード,停止時間(分),費用(円),現象\n"
            '="00123",20260803,142000,cmp-101,95,"1,234",スラリー流量低下\n'
            '="00124",20260804,090500,CMP-101,30-,500,"複数行の\n現象"\n'
            "合計件数,2\n")
    path = tmp_path / "システム出力_2026-08.csv"
    path.write_bytes(text.encode("cp932"))
    spec = spec_from_dict({
        "name": "故障", "columns": [
            {"key": "record_no", "display": "管理番号", "headers": ["管理番号"], "type": "code", "role": "key"},
            {"key": "occurred_at", "display": "発生日", "headers": ["発生日"], "type": "date", "role": "date"},
            {"key": "occurred_time", "display": "発生時刻", "headers": ["発生時刻"], "type": "time"},
            {"key": "equipment_id", "display": "設備コード", "headers": ["設備コード"], "type": "code", "role": "entity",
             "normalize": ["nfkc", "upper"]},
            {"key": "downtime", "display": "停止時間", "headers": ["停止時間"], "type": "number", "role": "measure", "unit": "分"},
            {"key": "cost", "display": "費用", "headers": ["費用"], "type": "number", "role": "measure", "unit": "円"},
            {"key": "symptom", "display": "現象", "headers": ["現象"], "type": "text", "role": "text"},
        ],
        "record": {"key": ["record_no"]},
    })
    source = open_source(path, path.name)
    layout = guess_layout(source, path.name)
    records, row_issues, stats = read_records(source, {}, layout, spec)
    assert [r.values for r in records] == [
        {"record_no": "00123", "occurred_at": "2026-08-03", "occurred_time": "14:20", "equipment_id": "CMP-101",
         "downtime": 95, "cost": 1234, "symptom": "スラリー流量低下"},
        {"record_no": "00124", "occurred_at": "2026-08-04", "occurred_time": "09:05", "equipment_id": "CMP-101",
         "downtime": -30, "cost": 500, "symptom": "複数行の\n現象"},
    ]
    assert records[1].source == {"file": path.name, "sheet": "", "row": 6}
    assert stats.excluded == {"合計": 1} and row_issues == []
    files = render_all(spec, [r.to_dict() for r in records], {})
    rec = next(f for f in files if f.name == "故障_2026-08.md").text
    assert "- 発生日: 2026-08-03 14:20（2026年8月）" in rec and "発生時刻" not in rec
    assert "## 【00124】CMP-101 複数行の｜2026-08-04" in rec
    assert "- 現象:\n  複数行の\n  現象\n" in rec


# ---- DB・ジョブを通した流れ ------------------------------------------------------------------------

class FakeCtx:
    def __init__(self):
        self.progress_calls = []

    def progress(self, **kw):
        self.progress_calls.append(kw)

    def check_cancel(self):
        pass


def _new_import(app, spec_dict=None, name="list.xlsx"):
    """取り込みを1件つくり、その取り込みの設定（spec）を持たせる。戻り値: (spec, import_id)"""
    upload_dir = Path(app.config["UPLOAD_DIR"]) / "tables"
    upload_dir.mkdir(parents=True, exist_ok=True)
    path = make_list_book(upload_dir / name)
    spec = spec_from_dict(spec_dict or list_spec_dict())
    file_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    import_id = tables.create_import(name, file_hash, f"tables/{name}", {"sheet": "2026年8月"})
    tables.save_spec(import_id, spec)
    return spec, import_id


def test_pipeline_read_render_download(app):
    with app.app_context():
        spec, import_id = _new_import(app)
        ctx = FakeCtx()
        result = tables.run_read(ctx, import_id)
        assert result["records"] == 5 and result["error"] == 0
        assert any(c.get("phase") == "読み込み" for c in ctx.progress_calls)
        imp = tables.get_import(import_id)
        assert imp["status"] == "preview" and imp["stats"]["records"] == 5
        assert Path(imp["rows_path"]).exists() and Path(imp["issues_path"]).exists()
        assert len(tables.load_rows(import_id)) == 5 and len(tables.load_rows(import_id, 1, 2)) == 2

        listing = tables.preview_files(import_id, imp, tables.spec_for_import(imp))
        assert "トラブル対応一覧_2026-07.md" in [f["name"] for f in listing]
        preview_dir = tables.import_files(import_id)["preview"]
        assert tables.md_text(preview_dir, "トラブル対応一覧_2026-08.md").startswith("# トラブル対応一覧 2026年8月")
        assert tables.md_text(preview_dir, "../x.md") is None

        done = tables.run_render(FakeCtx(), import_id)
        imp = tables.get_import(import_id)
        assert imp["status"] == "confirmed" and imp["confirmed_at"] and done["records"] == 5
        assert len(tables.md_paths(import_id)) == done["files"]
        data = tables.build_download(import_id)
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            names = zf.namelist()
            # 渡すのは RAG に入れる md だけ（フォルダも管理用CSVも無い）
            assert names == [p.name for p in tables.md_paths(import_id)] and len(names) == done["files"]
            assert all(n.endswith(".md") and "/" not in n for n in names)

        # 設定は取り込みの行にあり、保存し直せば上書きされる（版は作らない）
        spec.description = "v2"
        tables.save_spec(import_id, spec)
        again = tables.get_import(import_id)
        assert again["spec"].description == "v2" and again["spec_hash"] != imp["spec_hash"]


def test_pipeline_blocking_and_broken_file(app):
    with app.app_context():
        spec_dict = list_spec_dict()
        spec_dict["columns"][5].update(required=True, headers=["不具合内容"], display="不具合内容")
        spec, import_id = _new_import(app, spec_dict)
        tables.run_read(FakeCtx(), import_id)
        assert tables.get_import(import_id)["stats"]["issue_counts"]["error"] == 1
        assert has_blocking(tables.load_issues(import_id))

        bad_path = Path(app.config["UPLOAD_DIR"]) / "tables" / "broken.xlsx"
        bad_path.write_bytes(b"not a zip")
        bad = tables.create_import("broken.xlsx", "x", "tables/broken.xlsx", {})
        tables.save_spec(bad, spec)
        with pytest.raises(Exception):
            tables.run_read(FakeCtx(), bad)
        failed = tables.get_import(bad)
        assert failed["status"] == "failed" and "Excel" in failed["stats"]["error"]


def test_read_job_runs_in_worker(app):
    from app import core

    with app.app_context():
        _spec, import_id = _new_import(app)
        job_id = tables.start_read_job(import_id)
        job = core.wait_job(job_id, timeout=60)
        assert job["status"] == "done", job
        assert job["result"]["records"] == 5
        imp = tables.get_import(import_id)
        assert imp["status"] == "preview" and imp["job_id"] == job_id


# ---- 取り込みの控え（source_cache）と、作り直さない工夫 ----------------------------------------------

def _must_not_open(_stats):
    raise AssertionError("元のファイルを開き直した")


def test_scan_memo_gives_same_layout(tmp_path):
    path = make_list_book(tmp_path / "list.xlsx")
    source = open_source(path, path.name)
    memo: dict = {}
    auto = guess_layout(source, "2026年8月", scan_memo=memo)
    assert auto == guess_layout(source, "2026年8月") and len(memo) == 1
    manual = guess_layout(source, "2026年8月", header_rows=auto.header_rows, scan_memo=memo)
    assert manual == guess_layout(source, "2026年8月", header_rows=auto.header_rows) and len(memo) == 1
    ended = guess_layout(source, "2026年8月", header_rows=auto.header_rows, data_end=8, scan_memo=memo)
    assert ended == guess_layout(source, "2026年8月", header_rows=auto.header_rows, data_end=8) and len(memo) == 2


def test_import_source_cache_reuses_results_without_reopening(tmp_path):
    from app.tables import list_kind
    from app.tables import CACHE_FILE, ImportSource

    path = make_list_book(tmp_path / "list.xlsx")
    opened = []

    def opener(stats):
        opened.append(stats)
        return open_source(path, path.name, {"sheet_stats": stats} if stats else None)

    direct = open_source(path, path.name)
    sheet = "2026年8月"
    first = ImportSource(tmp_path / "imp", "key1", "excel", path.name, opener)
    auto = first.layout(sheet)
    assert auto == guess_layout(direct, sheet)
    assert first.layout(sheet, max_scan_rows=200) == guess_layout(direct, sheet, max_scan_rows=200)
    assert first.list_kind(sheet) == list_kind(direct, sheet)
    samples = first.sample_rows(sheet, auto, 200)
    assert samples == sample_data_rows(direct, sheet, auto, 200)
    assert list(first.rows(sheet, 1, 60)) == list(direct.rows(sheet, 1, 60)) and first.sheets() == direct.sheets()
    assert len(opened) == 1 and (tmp_path / "imp" / CACHE_FILE).exists()

    # 同じ鍵なら元のファイルを開かずに同じ結果（自動判定と同じ見出し行を指定し直しても全行をなめ直さない）
    again = ImportSource(tmp_path / "imp", "key1", "excel", path.name, _must_not_open)
    assert again.layout(sheet) == auto and again.sample_rows(sheet, auto, 200) == samples
    assert list(again.rows(sheet, 5, 3)) == list(direct.rows(sheet, 5, 3)) and again.sheets() == direct.sheets()
    assert again.list_kind(sheet) == list_kind(direct, sheet)
    assert again.layout(sheet, header_rows=auto.header_rows) == guess_layout(direct, sheet, header_rows=auto.header_rows)

    # 鍵（ファイル・読み込み設定）が変われば作り直す。Excel を開き直すときはシートの大きさの控えを渡す
    changed = ImportSource(tmp_path / "imp", "key2", "excel", path.name, opener)
    assert changed.layout(sheet) == auto and len(opened) == 2 and opened[1] is None
    reopened = ImportSource(tmp_path / "imp", "key2", "excel", path.name, opener)
    assert list(reopened.rows(sheet)) == list(direct.rows(sheet)) and opened[2] == direct.sheet_stats()


def test_import_source_does_not_recreate_a_deleted_folder(tmp_path):
    """取り込みを消したあとに、開いたままの画面の処理が控えを書いてもフォルダごと復活させない。

    控えには元の表の先頭行・列の見本の行がそのまま入るので、復活すると「ダウンロードしたら消える」
    （design.md 3.3）が破れる。
    """
    import shutil

    from app.tables import CACHE_FILE, ImportSource

    path = make_list_book(tmp_path / "list.xlsx")
    directory = tmp_path / "imports" / "1"
    source = ImportSource(directory, "key1", "excel", path.name, lambda stats: open_source(path, path.name))
    source.sheets()
    assert (directory / CACHE_FILE).exists()

    shutil.rmtree(directory)                 # ダウンロード（purge_table_import）でフォルダごと消えた
    source.layout("2026年8月")                # 別のタブの画面処理が続きを控えに書こうとする
    assert not directory.exists()            # 復活しない


def test_rows_page_and_render_reuses_preview(app, monkeypatch):
    from app import tables

    with app.app_context():
        _spec, import_id = _new_import(app)
        tables.run_read(FakeCtx(), import_id)
        rows = tables.load_rows(import_id)
        assert tables.load_rows_page(import_id, 1, 2) == (rows[1:3], 5)
        assert tables.load_rows_page(import_id, 100, 10) == ([], 5)

        # 2回目の読み込みは、控えた表の形を使う
        monkeypatch.setattr(tables, "guess_layout", lambda *a, **k: pytest.fail("表の形を推定し直した"))
        tables.run_read(FakeCtx(), import_id)
        monkeypatch.undo()
        assert tables.load_rows(import_id) == rows

        imp = tables.get_import(import_id)
        spec = tables.spec_for_import(imp)
        listing = tables.preview_files(import_id, imp, spec)
        preview = tables.import_files(import_id)["preview"]
        monkeypatch.setattr(tables, "render_files", lambda *a, **k: pytest.fail("Markdown を作り直した"))
        done = tables.run_render(FakeCtx(), import_id)
        monkeypatch.undo()
        md = {p.name: p.read_bytes() for p in tables.md_paths(import_id)}
        assert done["files"] == len(listing) == len(md)
        assert md == {f["name"]: (preview / f["name"]).read_bytes() for f in listing}
        assert md == {f.name: f.data for f in tables.render_files(import_id, tables.get_import(import_id), spec)}

        # 読み込み直したあとはプレビューが消えるので、確定のときに作る
        tables.run_read(FakeCtx(), import_id)
        calls = []
        original = tables.render_files
        monkeypatch.setattr(tables, "render_files", lambda *a, **k: calls.append(1) or original(*a, **k))
        tables.run_render(FakeCtx(), import_id)
        assert calls == [1] and {p.name: p.read_bytes() for p in tables.md_paths(import_id)} == md


# ---- samples（実ファイル） ------------------------------------------------------------------------

@pytest.mark.samples
def test_sample_t1_end_to_end():
    path = SAMPLES_tables_pipe / "T1_トラブル対応一覧_2023-2026.xlsx"
    if not path.exists():
        pytest.skip(f"{path.name} がありません")
    source = open_source(path, path.name)
    layout = guess_layout(source, "トラブル一覧")
    spec = spec_from_suggestions("トラブル対応一覧", layout,
                                 suggest_columns(layout.headers, sample_data_rows(source, "トラブル一覧", layout)))
    assert validate_spec(spec) == []
    records, row_issues, stats = read_records(source, {"sheet": "トラブル一覧"}, layout, spec)
    assert stats.records == 8095 and len({r.key for r in records}) == 8095
    assert stats.type_errors == {} and row_issues == []
    assert Counter(str(r.values.get("occurred_at"))[:4] for r in records) == {"2023": 1534, "2024": 2195, "2025": 2553, "2026": 1813}
    assert not has_blocking(run_checks(records, spec, stats))
    files = render_all(spec, [r.to_dict() for r in records], {})
    heads = [ln for f in files if f.kind == "records" for ln in f.text.split("\n") if ln.startswith("## ")]
    assert sum(1 for h in heads if "（続き" not in h) == 8095   # 大きい記録は「（続きn/m）」に分かれる
    first = next(f for f in files if f.name == "トラブル対応一覧_2023-04.md")
    assert "- 発生日: 2023-04-01 02:21（2023年4月）" in first.text and "- 担当者: " in first.text
    assert _no_blank_inside_records(first.text)


# ---- LightRAG オフライン評価の反映（時系列の重複・記録の上限・コード列・設備の値の揺れ） ----

def _rec(key: str, **values) -> dict:
    return {"key": key, "values": values, "source": {"file": "x.xlsx", "row": 4}}


def test_timeline_keeps_every_sentence_even_when_other_columns_repeat_them():
    """時系列は原文のまま出す（他の列と同じ文でも省かない。入力の内容を削らない）。"""
    spec = spec_from_dict(list_spec_dict())
    log = ("8/3 10:00 田中：スラリー流量低下アラームで研磨停止。保全へ連絡\n"
           "8/3 11:00 田中：フィルター交換、流量の再校正を実施\n"
           "8/3 12:00 田中：復旧を確認。流量の再校正を実施\n")
    rec = _rec("TR-001", record_no="TR-001", occurred_at="2026-08-03", equipment_id="CMP-101",
               symptom="スラリー流量低下アラームで研磨停止",
               action="フィルター交換、流量の再校正を実施", response_log=log)
    lines = record_block(rec, spec)
    timeline = [ln.strip() for ln in lines[lines.index("- 対応の時系列:") + 1:]]
    assert timeline[0].endswith(": スラリー流量低下アラームで研磨停止。保全へ連絡")
    assert timeline[1].endswith(": フィルター交換、流量の再校正を実施")
    assert timeline[2].endswith(": 復旧を確認。流量の再校正を実施")
    assert "- 現象: スラリー流量低下アラームで研磨停止" in lines and "- 処置: フィルター交換、流量の再校正を実施" in lines


def test_long_record_is_split_into_continuation_blocks_without_losing_text():
    """上限を超える記録は「（続きn/m）」に分ける。文は1つも消さず、どの部分にも管理No・設備・日付を書く。"""
    from app.tables import RECORD_TOKEN_BUDGET, record_blocks

    assert RECORD_TOKEN_BUDGET == 400
    log = "".join(f"8/3 {9 + i // 6:02d}:{(i % 6) * 10:02d} 田中：{'対応の記録です。' * 12}\n" for i in range(40))
    rec = _rec("TR-003", record_no="TR-003", occurred_at="2026-08-03", equipment_id="CMP-101", response_log=log)
    blocks = record_blocks(rec, spec_from_dict(list_spec_dict()))
    assert len(blocks) > 1
    for n, block in enumerate(blocks, start=1):
        tail = f"（続き{n}/{len(blocks)}）" if n > 1 else f"（1/{len(blocks)}）"
        assert block[0].startswith("## ") and block[0].endswith(tail)
        for line in ("- 管理No: TR-003", "- 発生日: 2026-08-03（2026年8月）", "- 設備番号: CMP-101"):
            assert line in block, (n, line)
        assert estimate_tokens("\n".join(block)) <= RECORD_TOKEN_BUDGET
    body = "\n".join(ln for block in blocks for ln in block)
    assert body.count("対応の記録です。" * 12) == 40      # 40件の記入がすべて残っている
    assert "管理用の正規化CSVに収録" not in body


def test_code_and_person_columns_are_written_by_default():
    """コードの列も人名の列も既定で出す（入力の内容を勝手に落とさない）。空欄だけの列は出さない。"""
    headers = ["管理番号", "状態コード", "再発フラグ", "発見区分コード", "担当者", "現象", "備考"]
    rows = [[f"MS-{i:04d}", "9", "0", "H01", "田中", f"アラーム{i}が出た", ""] for i in range(30)]
    by_header = {s.header: s for s in suggest_columns(headers, rows)}
    assert [h for h, s in by_header.items() if s.md == "omit"] == ["備考"]
    assert by_header["備考"].omit_reason == "blank"
    assert by_header["担当者"].role == "person" and by_header["担当者"].md != "omit"


def test_entity_code_and_name_in_one_column_are_split_before_grouping():
    """設備名の列がない台帳で「ETC-302(OXIDEエッチャ 2号機)」と「ETC-302」が別設備に割れないようにする。"""
    assert split_entity_code("ETC-302(OXIDEエッチャ 2号機)") == ("ETC-302", "OXIDEエッチャ 2号機")
    assert split_entity_code("CVD-203 W-CVD 3号機") == ("CVD-203", "W-CVD 3号機")
    assert split_entity_code("CMP-101") == ("CMP-101", "")
    assert split_entity_code("CMP-101 / CMP-102") == ("CMP-101 / CMP-102", "")   # 名前の側も番号だけ
    assert split_entity_code("超純水製造装置 2系") == ("超純水製造装置 2系", "")  # 番号らしくない

    d = list_spec_dict(group_by="entity_month")
    d["columns"] = [c for c in d["columns"] if c["key"] != "equipment_name"]
    spec = spec_from_dict(d)
    records = [_rec("A-1", record_no="A-1", occurred_at="2026-08-03", equipment_id="ETC-302(OXIDEエッチャ 2号機)"),
               _rec("A-2", record_no="A-2", occurred_at="2026-08-10", equipment_id="ETC-302")]
    files = render_all(spec, records, {})
    names = [f.name for f in files]
    assert "トラブル対応一覧_ETC-302_2026-08.md" in names      # 1つのファイルにまとまる
    assert not any("OXIDE" in n for n in names)
    rec_file = next(f for f in files if f.name == "トラブル対応一覧_ETC-302_2026-08.md")
    assert "- このファイルの記録: 2件" in rec_file.text
    assert "- 設備: OXIDEエッチャ 2号機（ETC-302）" in rec_file.text


def test_entity_name_before_code_is_also_split():
    """「Oxideエッチャ 2号機（ETC-302）」のように番号が後ろにある並びも、番号にそろえる（ファイルが二重にならない）。"""
    assert split_entity_code("Oxideエッチャ 2号機（ETC-302）") == ("ETC-302", "Oxideエッチャ 2号機")
    assert split_entity_code("ARFスキャナ 2号機(LIT-402)") == ("LIT-402", "ARFスキャナ 2号機")
    assert split_entity_code("超純水製造装置（2系）") == ("超純水製造装置（2系）", "")   # 番号らしくない
    assert split_entity_code("CMP-101（CMP-102）") == ("CMP-101（CMP-102）", "")      # 名前の側も番号だけ

    d = list_spec_dict(group_by="entity_month")
    d["columns"] = [c for c in d["columns"] if c["key"] != "equipment_name"]
    spec = spec_from_dict(d)
    records = [_rec("A-1", record_no="A-1", occurred_at="2026-08-03", equipment_id="Oxideエッチャ 2号機（ETC-302）"),
               _rec("A-2", record_no="A-2", occurred_at="2026-08-10", equipment_id="ETC-302")]
    names = [f.name for f in render_all(spec, records, {})]
    assert "トラブル対応一覧_ETC-302_2026-08.md" in names
    assert not any("Oxide" in n for n in names)


def test_code_column_upper_keeps_the_equipment_name_as_written():
    """code 列の upper は番号だけに効かせる（同じセルの設備名まで大文字にすると帳票側の表記と割れる）。"""
    col = ColumnSpec(key="equipment_id", display="設備番号", type="code", role="entity", normalize=["nfkc", "upper"])
    cctx = ConvertContext(na_tokens=set())
    for raw, expected in (("etc-302", "ETC-302"),
                          ("Oxideエッチャ 2号機（ETC-302）", "Oxideエッチャ 2号機(ETC-302)"),  # 括弧は NFKC で半角
                          ("etc-302 Oxideエッチャ 2号機", "ETC-302 Oxideエッチャ 2号機"),
                          ("Oxideエッチャ 2号機", "Oxideエッチャ 2号機")):
        assert convert_cell(raw, raw, col, cctx)[0] == expected


def test_columns_not_in_the_records_stay_out_of_every_file():
    """記録に出していない列（意味の分からないコード値）は、どのファイルにも出さない。"""
    spec = spec_from_dict(list_spec_dict())
    spec.column("failure_category").md = "omit"
    records = [_rec("A-1", record_no="A-1", occurred_at="2026-08-03", equipment_id="CMP-101",
                    failure_category="F01", downtime=30),
               _rec("A-2", record_no="A-2", occurred_at="2026-08-04", equipment_id="CMP-101",
                    failure_category="F01", downtime=60)]
    text = "\n".join(f.text for f in render_all(spec, records, {}))
    assert "F01" not in text and "故障区分" not in text


def test_only_columns_set_to_not_output_are_left_out_of_the_records():
    """「出さない」にした列だけを出さない（人名の列は出す）。"""
    spec = spec_from_dict(list_spec_dict())
    spec.column("failure_category").md = "omit"
    records = [_rec("A-1", record_no="A-1", occurred_at="2026-08-03", equipment_id="CMP-101", worker="田中")]
    rec = next(f for f in render_all(spec, records, {}) if f.kind == "records").text
    assert "- 担当者: 田中" in rec and "故障区分" not in rec


def test_record_title_is_cut_at_a_readable_place():
    """見出しの切り詰めは、閉じない括弧を残さず、切ったことが分かるようにする。"""
    from app.tables import TITLE_TEXT_CHARS, record_title

    spec = spec_from_dict(list_spec_dict())
    long = "ドライポンプが過負荷停止(A-3501)、チャンバー圧力上昇(同型機CVD-205は正常なので単体の不具合)"
    values = {"record_no": "TR-1", "occurred_at": "2026-08-03", "equipment_id": "CMP-101", "symptom": long}
    title = record_title(values, spec)
    assert title.endswith("…｜2026-08-03")
    body = title.split("】")[1].split("｜")[0]
    assert body.count("(") == body.count(")") and len(body) <= TITLE_TEXT_CHARS + 1
    # 40字に収まる現象はそのまま（「…」を付けない）
    short = dict(values, symptom="ドライポンプが過負荷停止")
    assert record_title(short, spec).endswith("ドライポンプが過負荷停止｜2026-08-03")


def test_record_title_drops_a_date_only_preamble():
    """本文の1行目が「【発生】R05.04.01 11:45(休日)」なら、見出しの日付と同じものを繰り返さない。"""
    from app.tables import record_title

    spec = spec_from_dict(list_spec_dict())
    text = "【発生】R05.04.01 11:45(休日)\n【設備】ETC-305 Polyエッチャ 5号機\n【内容】PM4のVppが範囲外"
    values = {"record_no": "CA-1", "occurred_at": "2023-04-03", "equipment_id": "ETC-305", "symptom": text}
    # 日付は繰り返さず、設備の行（見出しの設備と同じ）も飛ばして、次の行の中身を使う（R5-MD-4）
    assert record_title(values, spec) == "【CA-1】ETC-305 PM4のVppが範囲外｜2023-04-03"
    # 日付のあとに中身が続くときは、日付だけ落として中身を見出しに使う
    values2 = dict(values, symptom="R5.4.2 16:58(休日)、ビーム電流が低下し停止")
    assert record_title(values2, spec) == "【CA-1】ETC-305 ビーム電流が低下し停止｜2023-04-03"
    # 札も日付も無い現象はそのまま
    values3 = dict(values, symptom="スラリー流量低下")
    assert record_title(values3, spec) == "【CA-1】ETC-305 スラリー流量低下｜2023-04-03"


def test_record_body_shows_the_split_entity_like_the_heading():
    """設備名の列がない台帳でも、本文の設備の行を見出し・集計と同じ書き方にそろえる。"""
    d = list_spec_dict()
    d["columns"] = [c for c in d["columns"] if c["key"] != "equipment_name"]
    spec = spec_from_dict(d)
    rec = _rec("A-1", record_no="A-1", occurred_at="2026-08-03", equipment_id="ETC-302(Oxideエッチャ 2号機)")
    lines = record_block(rec, spec)
    assert "- 設備番号: Oxideエッチャ 2号機（ETC-302）" in lines
    assert lines[0].startswith("## 【A-1】Oxideエッチャ 2号機（ETC-302）")
    # 番号だけの記録は今までどおり（同じ列の名前のまま）
    rec2 = _rec("A-2", record_no="A-2", occurred_at="2026-08-04", equipment_id="ETC-302")
    assert "- 設備番号: ETC-302" in record_block(rec2, spec)


# ---- ユーザーが決めた出力の変更（ヒント既定オン・時系列の重複削除の選択・丸数字） ----

def test_filenames_have_no_lightrag_hint_and_retired_settings_are_ignored():
    """ファイル名にヒント（.[...]）は付けない。保存済みの設定に残っていた項目は読み飛ばす。"""
    from app.tables import TableSpec

    assert "lightrag_hint" not in TableSpec("新しい設定").markdown
    saved = list_spec_dict(lightrag_hint=True, dedupe_timeline=False, omit_person=True, max_records_per_file=3)
    md = spec_from_dict(saved).markdown
    assert not ({"lightrag_hint", "dedupe_timeline", "omit_person", "max_records_per_file"} & set(md))
    assert validate_spec(spec_from_dict(saved)) == []

    files = render_all(spec_from_dict(list_spec_dict()),
                       [_rec("TR-1", record_no="TR-1", occurred_at="2026-08-03", equipment_id="CMP-101")], {})
    for f in files:
        assert ".[" not in f.name and "]" not in f.name and f.name.endswith(".md")


@pytest.mark.parametrize("type_, text, expected", [
    ("text", "①分解 ②清掃 ③組立", "①分解 ②清掃 ③組立"),          # 本文は丸数字のまま
    ("text", "Ⓐ系統の㋐弁を閉", "Ⓐ系統の㋐弁を閉"),
    ("string", "①機械", "①機械"),
    ("text", "㈱山田製作所へ連絡", "(株)山田製作所へ連絡"),          # 区切りが残る表記は NFKC のまま
    ("text", "ＡＢＣ　１２３", "ABC 123"),                          # 全角の英数字・空白は今までどおり
    ("code", "①-２", "1-2"),                                        # コードは突き合わせに使うので NFKC のみ
])
def test_enclosed_characters_are_kept_in_the_text_that_goes_into_markdown(type_, text, expected):
    col = ColumnSpec("c", "列", type=type_)
    out, error, flag = convert_cell(None, text, col, ConvertContext(na_tokens=frozenset()))
    assert (out, error, flag) == (expected, None, None)


def test_enclosed_characters_do_not_change_dates_numbers_or_na_tokens():
    """比較・解析に使う正規化は今までどおり NFKC（丸数字を残すのは md に出す文章だけ）。"""
    cctx = ConvertContext(fiscal_year=2026, fiscal_start=4, na_tokens=frozenset(["-", "該当なし"]))
    assert convert_cell(None, "２０２６/８/３", ColumnSpec("d", "発生日", type="date"), cctx)[0] == "2026-08-03"
    assert convert_cell(None, "１，２３４", ColumnSpec("n", "費用", type="number"), cctx)[0] == 1234
    assert convert_cell(None, "－", ColumnSpec("t", "現象", type="text"), cctx)[0] is None       # NA 判定
    assert convert_cell(None, "①", ColumnSpec("t", "現象", type="text"), cctx)[0] == "①"


def test_record_markdown_and_filename_for_a_circled_number(tmp_path):
    """md の本文は丸数字のまま、ファイル名は今までどおり NFKC（OS・LightRAG 側で揺れないように）。"""
    spec = spec_from_dict(list_spec_dict())
    rec = _rec("TR-①", record_no="TR-001", occurred_at="2026-08-03", equipment_id="CMP-101",
               equipment_name="CMP研磨装置1号機", action="①フィルター交換 ②流量再校正")
    assert "- 処置: ①フィルター交換 ②流量再校正" in record_block(rec, spec)
    assert md_filename(["トラブル対応一覧", "①"]) == "トラブル対応一覧_1.md"


# ---- R1: 番号の分け方・平均の分母・9999/12/31 ----

def test_split_entity_code_prefers_trailing_code_and_drops_qualifiers():
    """「Fab1 OHTシステム（OHT-801）」の番号は OHT-801。「IMP-602（推定）」の「推定」は設備名にしない。"""
    assert split_entity_code("Fab1 OHTシステム（OHT-801）") == ("OHT-801", "Fab1 OHTシステム")
    assert split_entity_code("Fab2 自動倉庫（ストッカ）（STK-831）") == ("STK-831", "Fab2 自動倉庫（ストッカ）")
    assert split_entity_code("IMP-602（推定）") == ("IMP-602", "")
    assert split_entity_code("CVD-203 W-CVD 3号機") == ("CVD-203", "W-CVD 3号機")
    assert split_entity_code("ETC-302(OXIDEエッチャ 2号機)") == ("ETC-302", "OXIDEエッチャ 2号機")

    d = list_spec_dict(group_by="entity_month")
    d["columns"] = [c for c in d["columns"] if c["key"] != "equipment_name"]
    spec = spec_from_dict(d)
    records = [_rec("A-1", record_no="A-1", occurred_at="2026-08-03", equipment_id="IMP-602（推定）"),
               _rec("A-2", record_no="A-2", occurred_at="2026-08-10", equipment_id="IMP-602")]
    names = [f.name for f in render_all(spec, records, {})]
    assert "トラブル対応一覧_IMP-602_2026-08.md" in names
    assert not any("推定" in n for n in names)


def test_far_future_date_does_not_break_the_records(): 
    """「期限なし」の 9999/12/31 があっても、月末日の計算で落ちずに md を作れる。"""
    from app.tables import month_last_day

    assert month_last_day("9999-12") == "9999-12-31" and month_last_day("2024-02") == "2024-02-29"
    spec = spec_from_dict(list_spec_dict())
    records = [_rec("A-1", record_no="A-1", occurred_at="2026-08-03", equipment_id="CMP-101", downtime=30),
               _rec("A-2", record_no="A-2", occurred_at="9999-12-31", equipment_id="CMP-101", downtime=10)]
    names = [f.name for f in render_all(spec, records, {})]
    assert names == ["トラブル対応一覧_2026-08.md", "トラブル対応一覧_9999-12.md"]


def test_the_file_list_counts_records_not_continuation_parts():
    """「内容の確認」のファイル一覧の「記録」は、記録の件数（「（続きn/m）」は1件の続きなので数えない）。"""
    from app.tables import MdFile
    from app.tables import _record_count

    text = ("# 見出し\n\n## 【A-1】長い記録（1/3）\n- a\n\n## 【A-1】長い記録（続き2/3）\n- b\n\n"
            "## 【A-1】長い記録（続き3/3）\n- c\n\n## 【A-2】短い記録\n- d\n")
    assert _record_count(MdFile("記録.md", text, "records")) == 2


# ====================================================================================================
# 元 tests/test_tables_flow.py
# 一覧表の1画面（/tables）の通し確認。
#
# 画面は1枚で、段（読み取り方→表の範囲→列の対応づけ→AI整形→内容の確認→確定してダウンロード）を
# fetch で出し入れする。テストも画面と同じ JSON のやりとりで進める。
# ====================================================================================================

def test_csv_import_flow(app, client, monkeypatch):
    import_id = upload_csv(client, "トラブル一覧.csv")

    source = panel(client, import_id, "source")
    # 取り込み設定は保存しないので、読み取り方の段に設定の選択は出ない
    assert "new_template_name" not in source["html"] and "取り込み設定" not in source["html"]
    assert csv_source(client, import_id)["next"] == "layout"

    layout = panel(client, import_id, "layout")
    assert "row-header" in layout["html"]
    detect = client.post(f"/tables/imports/{import_id}/layout/detect", json={"header_rows": [1]}).get_json()
    assert detect["data_end"] == 13 and detect["table_kind"] == "list"
    res = save_layout(client, import_id)
    assert res.status_code == 200 and res.get_json()["next"] == "columns"

    columns = panel_html(client, import_id, "columns")
    # 表の名前はファイル名から入れておく（設定は残らないので、この取り込みだけの名前）
    assert "対応内容" in columns and 'data-setting="name" value="トラブル一覧"' in columns
    # 2回目からは取り込みの控えで画面を作る（CSV を開き直さない）
    monkeypatch.setattr(tables, "open_import_source", lambda *a, **k: pytest.fail("CSV を開き直した"))
    assert "row-header" in panel_html(client, import_id, "layout")
    assert "対応内容" in panel_html(client, import_id, "columns")
    assert panel(client, import_id, "source")["html"]
    monkeypatch.undo()

    res = save_columns(client, import_id, columns_payload("トラブル対応一覧", ai_role="log"))
    assert res.status_code == 200, res.get_json()
    assert res.get_json()["next"] == "ai"
    imp = wait_import_job(app, import_id)
    assert imp["status"] == "preview" and imp["stats"]["records"] == 12

    assert client.get(f"/api/jobs/{imp['job_id']}").get_json()["status"] == "done"
    assert "分割プレビュー" in panel_html(client, import_id, "ai")
    split = client.post(f"/tables/imports/{import_id}/ai/split-preview", json={"row_key": "TR-001"}).get_json()
    assert len(split["segments"]) == 2

    preview = preview_panel(app, client, import_id)
    assert "確定してMarkdownを作成" in preview["html"] and "トラブル対応一覧_2026-08.md" in preview["html"]
    text = client.get(f"/tables/imports/{import_id}/preview/file",
                      query_string={"name": "トラブル対応一覧_2026-08.md"}).get_json()["text"]
    assert "TR-001" in text and "対応の時系列" in text

    res = client.post(f"/tables/imports/{import_id}/confirm")
    assert res.status_code == 200 and res.get_json()["next"] == "done"
    imp = wait_import_job(app, import_id)
    assert imp["status"] == "confirmed"
    assert "Markdownをまとめてダウンロード" in panel_html(client, import_id, "done")

    # zip は RAG に入れる md だけ（フォルダ分けも管理用CSVも無い）
    zdata = client.get(f"/tables/imports/{import_id}/download.zip").data
    names = zipfile.ZipFile(io.BytesIO(zdata)).namelist()
    assert names == ["トラブル対応一覧_2026-08.md"]
    with app.app_context():
        assert tables.get_import(import_id) is None      # ダウンロードしたら残さない（設定も一緒に消える）

    # 2回目も同じ順で進む（設定は残らないので、列の対応づけは毎回決める）
    second = upload_csv(client, "トラブル一覧2.csv")
    save_source(client, second, encoding="cp932", delimiter=",")
    res = save_layout(client, second)
    assert res.status_code == 200 and res.get_json()["next"] == "columns"
    res = save_columns(client, second, columns_payload("トラブル対応一覧2"))
    assert res.status_code == 200, res.get_json()
    assert wait_import_job(app, second)["status"] == "preview"


# ---- 読み込み中の中止（段に出る［中止］） ------------------------------------------------

def _pending_import(app, status="reading"):
    """読み込み中に見える取り込みを1件つくる（ジョブは待機中のまま）。"""
    from app import database
    with app.app_context():
        import_id = tables.create_import("大きい表.csv", "h" * 64, "uploads/大きい表.csv", {"encoding": "utf-8"})
        conn = database.get_db()
        ts = database.now()
        cur = conn.execute("""INSERT INTO jobs (kind, ref_type, ref_id, status, params_json, progress_json,
                              created_at, updated_at) VALUES ('table_read', 'table_import', ?, 'queued', '{}', '{}', ?, ?)""",
                           (import_id, ts, ts))
        conn.commit()
        tables.update_import(import_id, status=status, job_id=cur.lastrowid)
        return import_id, cur.lastrowid


def test_panels_show_the_running_read_job_with_a_cancel_url(app, client):
    import_id, _job_id = _pending_import(app)
    data = panel(client, import_id, "preview")
    assert data["reading"] is True
    assert data["job"]["kind"] == "table_read"
    assert data["job"]["cancel_url"] == f"/tables/imports/{import_id}/cancel"


def test_cancel_read_job_puts_the_import_back(app, client):
    import_id, job_id = _pending_import(app)
    res = client.post(f"/tables/imports/{import_id}/cancel")
    assert res.status_code == 200 and res.get_json()["ok"] is True
    with app.app_context():
        assert core.get_job(job_id)["status"] == "cancelled"
        # 待機中のまま中止したときは取り込みの状態も戻す（読み込み中のままにしない）
        assert tables.get_import(import_id)["status"] == "uploaded"
    # 待ち画面ではなく、読み込みからやり直せる段が出る
    assert panel(client, import_id, "preview").get("reading") is None


def test_cancel_render_job_returns_to_preview(app, client):
    from app import database

    import_id, job_id = _pending_import(app, status="confirming")
    with app.app_context():
        conn = database.get_db()
        conn.execute("UPDATE jobs SET kind = 'table_render' WHERE id = ?", (job_id,))
        conn.commit()
    assert client.post(f"/tables/imports/{import_id}/cancel").status_code == 200
    with app.app_context():
        assert core.get_job(job_id)["status"] == "cancelled"
        assert tables.get_import(import_id)["status"] == "preview"


def test_cancel_without_running_job_is_reported(app, client):
    with app.app_context():
        import_id = tables.create_import("表.csv", "h" * 64, "uploads/表.csv", {})
    res = client.post(f"/tables/imports/{import_id}/cancel")
    assert res.status_code == 400 and "中止できる処理がありません" in res.get_json()["error"]


# ---- 表の範囲の段の表示 ----------------------------------------------------------------

def test_layout_panel_without_headers(app, client):
    import_id = upload_csv(client, "1列.csv", "あ\r\nい\r\n", encoding="utf-8")
    csv_source(client, import_id, encoding="utf-8")
    data = panel(client, import_id, "layout")
    # 見出し行が見つからないときに「1〜0行目」のような存在しない行番号を出さない
    assert "データの行が見つかりません" in data["html"] and "〜0行目" not in data["html"]
    assert data["note"] == ""
    # 見出し行を指定すれば進める
    res = save_layout(client, import_id, header_rows="1")
    assert res.status_code == 200 and res.get_json()["next"] == "columns"


# ---- 取り込みの削除・段の表示 ----------------------------------------------------------

def test_delete_import_removes_the_file_and_the_read_rows(app, client):
    """間違えて取り込んだファイルを消せる（元のファイルも読み込んだ内容も残さない）。"""
    import_id = upload_csv(client, "間違い.csv")
    assert "data-import-delete" in panel_html(client, import_id, "layout")
    with app.app_context():
        base = tables.import_dir(import_id)
        base.mkdir(parents=True, exist_ok=True)
        (base / "rows.jsonl.gz").write_bytes(b"x")
        stored = tables.get_import(import_id)["stored_path"]
    res = client.post(f"/tables/imports/{import_id}/delete")
    assert res.status_code == 200 and res.get_json()["ok"] is True
    with app.app_context():
        from app.core import upload_path
        assert tables.get_import(import_id) is None
        assert not tables.import_dir(import_id).exists()
        assert not upload_path(stored).exists()
        assert tables.list_imports(limit=10) == []


def test_delete_is_refused_while_a_job_is_running(app, client):
    import_id, _job_id = _pending_import(app)
    res = client.post(f"/tables/imports/{import_id}/delete")
    assert res.status_code == 400 and "処理中の取り込みは削除できません" in res.get_json()["error"]
    with app.app_context():
        assert tables.get_import(import_id) is not None


def test_source_rejects_an_encoding_that_is_not_offered(app, client):
    """画面にない文字コードを保存させない（保存できると読み込みで落ちて段が開けなくなる）。"""
    import_id = upload_csv(client, "文字コード.csv")
    res = save_source(client, import_id, encoding="rot13", delimiter=",")
    assert res.status_code == 400 and "この画面にない文字コード" in res.get_json()["error"]
    with app.app_context():
        assert tables.get_import(import_id)["source"].get("encoding") != "rot13"
    # 画面にある文字コードは保存できる
    assert csv_source(client, import_id, encoding="utf-16")["next"] == "layout"
    # 中身と合わない文字コードでも、500 ではなく段を出して選び直せる（読めなければ理由を1行出す）
    assert panel(client, import_id, "source")["html"]
    layout = panel(client, import_id, "layout")
    assert layout["html"] or layout.get("locked")


# 「値は「日付」らしい」の注意書き（英語キーのまま出していないかを見ていたテスト）は、「知らせ」の列ごと
# 2026-09-21 にやめた（利用者の指示）。出なくなったことは tests/test_tables_column_words.py で見る。


def test_preview_panel_says_it_is_already_confirmed(app, client):
    """確定済みの取り込みを開き直したときに、黙って作り直せるボタンだけを出さない。"""
    import_id = upload_csv(client, "確定済み.csv")
    csv_source(client, import_id)
    save_layout(client, import_id)
    assert save_columns(client, import_id, columns_payload("確定済み")).status_code == 200
    wait_import_job(app, import_id)
    preview_panel(app, client, import_id)
    client.post(f"/tables/imports/{import_id}/confirm")
    assert wait_import_job(app, import_id)["status"] == "confirmed"
    data = preview_panel(app, client, import_id)
    assert data["confirmed"] is True
    assert "この取り込みは確定済みです" in data["html"]
    assert "確定し直してMarkdownを作り直す" in data["html"] and "data-confirm" in data["html"]


def test_source_panel_marks_a_crosstab_sheet(app, client):
    """クロス集計のシートに「一覧表らしい」と出さない（範囲の段で初めて止められると分かりにくい）。"""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "FY2023"
    ws.append(["設備"] + [f"2023-{m:02d}" for m in range(4, 13)] + [f"2024-{m:02d}" for m in range(1, 4)])
    for i in range(1, 16):
        ws.append([f"EQ-{i:02d}"] + [i * m for m in range(1, 13)])
    buf = io.BytesIO()
    wb.save(buf)
    res = upload(client, buf.getvalue(), "月別集計.xlsx")
    import_id = res.get_json()["import_id"]
    html = panel_html(client, import_id, "source")
    assert "クロス集計らしい（対応していません）" in html and "・一覧表らしい" not in html
    # 範囲の段の判定は今までどおり（この範囲では進めない）
    res = save_layout(client, import_id)
    assert res.status_code == 400 and "クロス集計" in res.get_json()["error"]


# ---- 処理中の変更・削除を止める -----------------------------------------------------------

def _queued_job(app, import_id, kind):
    from app import database
    with app.app_context():
        conn = database.get_db()
        ts = database.now()
        cur = conn.execute("""INSERT INTO jobs (kind, ref_type, ref_id, status, params_json, progress_json,
                              created_at, updated_at) VALUES (?, 'table_import', ?, 'queued', '{}', '{}', ?, ?)""",
                           (kind, import_id, ts, ts))
        conn.commit()
        return cur.lastrowid


def test_source_is_not_changed_while_reading(app, client):
    """読み込み中に読み取り方を保存しても、文字コード・区切り文字を書き換えない
    （古い設定の読み込み結果で確定させない）。"""
    import_id, _job_id = _pending_import(app)
    res = save_source(client, import_id, encoding="utf-8", delimiter=";")
    assert res.status_code == 409 and "処理中は変更できません" in res.get_json()["error"]
    assert save_layout(client, import_id, header_rows="2").status_code == 409
    assert save_columns(client, import_id, {"name": "x", "columns": []}).status_code == 409
    with app.app_context():
        imp = tables.get_import(import_id)
        assert imp["status"] == "reading" and imp["source"] == {"encoding": "utf-8"}


def test_utf16_without_bom_can_pass_the_source_panel(app, client):
    """BOMなしの UTF-16 と判定された CSV も、読み取り方をそのまま保存して先へ進める。"""
    import_id = upload_csv(client, "u16.csv", encoding="utf-16-le")
    with app.app_context():
        assert tables.get_import(import_id)["source"]["encoding"] == "utf-16-le"
    assert 'value="utf-16-le"' in panel_html(client, import_id, "source")
    assert csv_source(client, import_id, encoding="utf-16-le")["next"] == "layout"
    assert panel(client, import_id, "layout").get("locked") is None


def test_delete_and_confirm_are_refused_while_ai_formatting_runs(app, client, monkeypatch):
    """AI整形の実行中に削除すると、動いている呼び出しが消したあとの DB に結果を書き戻すので、止める。"""
    from types import SimpleNamespace

    from app import views as tables_view

    import_id = upload_csv(client, "AI中.csv")
    with app.app_context():
        tables.update_import(import_id, status="preview")
    _queued_job(app, import_id, "ai_format")
    res = client.post(f"/tables/imports/{import_id}/delete")
    assert res.status_code == 400 and "AI整形の実行中は削除できません" in res.get_json()["error"]
    with app.app_context():
        assert tables.get_import(import_id) is not None

    monkeypatch.setattr(tables_view, "_spec_for", lambda imp: SimpleNamespace(log_stage=object()))
    monkeypatch.setattr(tables, "start_render_job", lambda *a, **k: pytest.fail("AI整形の実行中に確定した"))
    res = client.post(f"/tables/imports/{import_id}/confirm")
    assert res.status_code == 400 and "AI整形の実行中は確定できません" in res.get_json()["error"]


def test_trial_run_needs_the_external_confirmation(app, client, monkeypatch):
    """試し実行も、外部の AI に送るときは画面の確認のチェックがなければ送らない（サーバー側で確かめる）。"""
    from types import SimpleNamespace

    from app import aiproc
    from app import llm
    from app import views as tables_view

    import_id = upload_csv(client, "試し.csv")
    monkeypatch.setattr(tables_view, "_spec_for", lambda imp: SimpleNamespace(log_stage=object()))
    monkeypatch.setattr(llm, "job_client_settings", lambda *a, **k: {"chat_url": "https://api.example.com/v1/chat"})
    monkeypatch.setattr(aiproc, "trial_row", lambda *a, **k: pytest.fail("確認なしで外部に送った"))
    monkeypatch.setattr(aiproc, "load_rows_for_ai", lambda *a, **k: pytest.fail("確認なしで外部に送った"))
    res = client.post(f"/tables/imports/{import_id}/ai/trial", json={"row_key": "TR-001"})
    assert res.status_code == 400 and "外部のAIサービス" in res.get_json()["error"]


def test_workbook_without_sheets_is_rejected(app, client, tmp_path):
    """シートのないブックは取り込みにせず、選び直しの案内を出す（読み取り方の段が 500 にならない）。"""
    import re

    from openpyxl import Workbook

    wb = Workbook()
    wb.save(tmp_path / "s.xlsx")
    src = zipfile.ZipFile(tmp_path / "s.xlsx")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/workbook.xml":
                data = re.sub(rb"<sheets>.*?</sheets>", b"<sheets/>", data, flags=re.S)
            z.writestr(item, data)
    res = upload(client, out.getvalue(), "シートなし.xlsx")
    assert res.status_code == 400 and "シート" in res.get_json()["error"]
    with app.app_context():
        assert tables.list_imports(limit=10) == []


@pytest.mark.parametrize("url, target", [("/tables/upload", "app.views.open_source"),
                                         ("/forms/upload", "app.views.precheck_excel")])
def test_unexpected_error_during_upload_leaves_no_file(app, client, monkeypatch, tmp_path, url, target):
    """読み込みの途中で思わぬエラーが出ても、アップロードしたファイルは残さない（design.md 3.3）。"""
    from pathlib import Path

    from openpyxl import Workbook

    def boom(*args, **kwargs):
        raise RuntimeError("想定外")

    monkeypatch.setattr(target, boom)
    wb = Workbook()
    wb.active["A1"] = "管理No"
    wb.save(tmp_path / "a.xlsx")
    with pytest.raises(RuntimeError):
        client.post(url, data={"file": (io.BytesIO((tmp_path / "a.xlsx").read_bytes()), "a.xlsx")},
                    content_type="multipart/form-data")
    assert [p for p in Path(app.config["UPLOAD_DIR"]).rglob("*") if p.is_file()] == []


# ====================================================================================================
# 元 tests/test_tables_columns_summary.py
# 「列の対応づけ」の段は、決めることが無ければ要約1行だけにする。
#
# 利用者の問い（2026-09-20）「列の対応付けを行う意味は？」への答え。この段で決められるのは
# 「出す／出さない」と四つの役割（識別番号・日付・対象・経過の記録）だけで、ふつうの一覧表なら
# どちらも見出しと値から決まっている。決まっているときは22行の表を出さず、要約1行と［変更する］にする。
# 決まっていないとき（識別番号が無い・読み取れない値がある など）は、今までどおり表を開く。
# ====================================================================================================

# 識別番号になる見出しが無い表（「作業メモ」は辞書に無い）
NO_KEY_CSV = "作業メモ,発生日,設備番号,対応内容\r\n" + "".join(
    f'メモ{i},2026-08-{i:02d},EQ-0{i % 3 + 1},"8/{i} 10:00 田中: 確認。\n8/{i} 11:00 佐藤: 復旧。"\r\n'
    for i in range(1, 13))


def _ready(app, client, text: str, name: str = "一覧.csv") -> int:
    """読み取り方と範囲まで決めて、「列の対応づけ」の段を開ける取り込みを作る。"""
    import_id = upload_csv(client, name, text)
    assert csv_source(client, import_id)["next"] == "layout"
    assert save_layout(client, import_id).status_code == 200
    return import_id


def _summary(html: str) -> str:
    found = re.search(r'<p class="col-summary">(.*?)</p>', html, re.S)
    return found.group(1).strip() if found else ""


def _todo(html: str) -> list[str]:
    if "data-columns-todo" not in html:
        return []
    block = html.split("data-columns-todo", 1)[1].split("</ul>", 1)[0]
    return [t.strip() for t in re.findall(r"<li>(.*?)</li>", block, re.S)]


# ---- 決まっているとき: 要約1行 -----------------------------------------------------------------

def test_the_summary_replaces_the_table_when_nothing_needs_deciding(app, client):
    import_id = _ready(app, client, CSV_TEXT)
    html = panel_html(client, import_id, "columns")

    # 四つの役割を名指しし、出す列と出さない列の数も書く
    assert _summary(html) == ("管理No＝識別番号、発生日＝日付、設備番号＝対象、対応内容＝経過の記録として読み取ります。"
                              "8列のうち8列を Markdown に出します（出さない列はありません）。")
    assert _todo(html) == []
    # 表は隠すだけで DOM に残す（［変更する］で開ける・保存で送る中身は変わらない）
    assert "data-columns-table hidden" in html and html.count("data-col ") == 8
    assert "変更する" in html and "この対応づけで読み込む" in html


def test_the_summary_says_so_when_there_is_no_equipment_column(app, client):
    """対象の列が無い表では、あるふりをせず「ありません」と書く（SPEC_CSV の設備名は設備番号ではない）。"""
    import_id = spec_import(app, client)
    summary = _summary(panel_html(client, import_id, "columns"))
    assert summary.startswith("管理No＝識別番号、発生日＝日付、対応内容＝経過の記録として読み取ります。")
    assert "対象の列はありません。" in summary
    assert "8列のうち8列を Markdown に出します（出さない列はありません）。" in summary


# ---- 決まっていないとき: 今までどおり表 ----------------------------------------------------------

def test_the_table_opens_when_the_record_number_is_missing(app, client):
    import_id = _ready(app, client, NO_KEY_CSV, "メモ.csv")
    html = panel_html(client, import_id, "columns")

    assert _summary(html) == ""
    assert _todo(html) == ["識別番号の列が決まっていません。1つ選んでください"]
    assert "data-columns-table hidden" not in html and html.count("data-col ") == 4


def test_the_table_opens_when_a_column_has_values_that_cannot_be_read(app, client):
    text = CSV_TEXT.replace(",60,田中", ",不明,田中")   # 停止時間（数値）に読み取れない値を混ぜる
    import_id = _ready(app, client, text)
    html = panel_html(client, import_id, "columns")

    assert _summary(html) == ""
    assert _todo(html) == ["列「停止時間(分)」に読み取れない値があります（8.3%）。出すかどうか決めてください"]


# ---- ［変更する］は表を出すだけ（送る中身は変わらない） ---------------------------------------------

@pytest.mark.skipif(shutil.which("node") is None, reason="node がない")
def test_the_change_button_only_unhides_the_table():
    js = _tables_js()
    start = js.index("function openColumnsTable(")
    end = js.index("function onEditorChange(")
    script = js[start:end] + """
const summary = { hidden: false }, table = { hidden: true };
const editor = { querySelector: (s) => (s === "[data-columns-summary]" ? summary : table) };
openColumnsTable(editor);
openColumnsTable(null);
process.stdout.write(JSON.stringify({ summary: summary.hidden, table: table.hidden }));
"""
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "check.mjs"
        path.write_text(script, encoding="utf-8")
        out = subprocess.run(["node", str(path)], capture_output=True, check=True, timeout=30, encoding="utf-8")
    assert json.loads(out.stdout) == {"summary": True, "table": False}


def test_the_spec_saved_through_the_summary_is_the_spec_saved_through_the_table(app, client):
    """要約のまま保存しても、表を開いて何も変えずに保存しても、取り込み設定は1文字も変わらない。"""
    # 要約の段から集めた中身（表は隠れているが DOM にあるので、画面と同じものが集まる）
    summary_id = _ready(app, client, CSV_TEXT, "要約.csv")
    assert "data-columns-table hidden" in panel_html(client, summary_id, "columns")
    body = editor_body(client, summary_id)
    assert save_columns(client, summary_id, body).status_code == 200
    wait_import_job(app, summary_id)

    # 表から集めたのと同じ中身（tests.tables_helpers.columns_payload が作る、画面の表そのままの送り方）
    table_id = _ready(app, client, CSV_TEXT, "表.csv")
    assert save_columns(client, table_id, columns_payload("要約", ai_role="log")).status_code == 200
    wait_import_job(app, table_id)

    with app.app_context():
        assert body["name"] == "要約"
        assert tables.get_import(summary_id)["spec_json"] == tables.get_import(table_id)["spec_json"]


# ====================================================================================================
# 元 tests/test_tables_column_words.py
# 「列の対応づけ」の言葉づかい（利用者の指示 2026-09-21）。
#
# 1. 表の5列目「知らせ」をやめる。ほとんどの行で空のうえ、書いていたことは表の上の行と重なっていた。
#    はじめから「使わない」にしている理由だけは、見出しの下に小さく残す（無いと、チェックの外れた列を
#    入れ直してよいのか分からない）。「値は「日付」らしい」は出さない（日付の役割は日付として読める列に
#    しか出ないので、読んでも直しようがない）。
# 2. 役割の名前を、設備の記録以外の表にも当てはまる言い方にする（設備 → 対象、AI整形の対象（追記ログ）
#    → 経過の記録）。中で使う名前（key/date/entity/log/attribute）は変えないので、保存した取り込み設定は
#    そのまま読める。
# 3. AI整形の段（⑤）からも「追記ログ」という言葉をなくす。
# ====================================================================================================

# 空欄だけの列（予備）と、型が合っていない列（状態コード＝数字だがコードの列）がある表
MIXED_CSV = "管理No,発生日,設備番号,予備,状態コード,対応内容,停止時間\r\n" + "".join(
    f'TR-{i:03d},2026-08-{i % 28 + 1:02d},EQ-0{i % 3 + 1},,{i % 3},'
    f'"8/{i % 28 + 1} 10:00 田中: 確認した。",{i * 5}\r\n' for i in range(1, 29))


def _columns_html(client, text: str = MIXED_CSV, name: str = "言葉.csv") -> str:
    """読み取り方と範囲まで決めて、「列の対応づけ」の段の HTML を返す。"""
    import_id = upload_csv(client, name, text)
    assert csv_source(client, import_id)["next"] == "layout"
    assert save_layout(client, import_id).status_code == 200
    return panel_html(client, import_id, "columns")


def _row(html: str, header: str) -> str:
    """その見出しの行（<tr>…</tr>）だけを取り出す。"""
    found = re.search(r'<tr data-col [^>]*data-header="' + re.escape(header) + r'".*?</tr>', html, re.S)
    assert found, f"列「{header}」の行が無い"
    return found.group(0)


def _head_cell(row_html: str) -> str:
    """その行の「見出し」のセル（2つ目の <td>）。"""
    cells = re.findall(r"<td[^>]*>.*?</td>", row_html, re.S)
    assert len(cells) == 4, cells   # 使う・見出し・役割・値の例の4つだけ
    return cells[1]


# ---- 1. 「知らせ」の列をやめる -----------------------------------------------------------------

def test_the_column_table_has_four_columns_and_no_notice_column(client):
    """表の見出しは「使う／見出し／役割／値の例」の4つだけ。"""
    html = _columns_html(client)
    head = re.search(r"<thead>.*?</thead>", html, re.S).group(0)
    assert re.findall(r"<th>(.*?)</th>", head, re.S) == ["使う", "見出し（そのまま項目名になります）", "役割", "値の例"]
    assert "知らせ" not in html


def test_the_reason_a_column_starts_unchecked_is_under_its_name(client):
    """チェックの外れている列は、その理由を見出しの下に小さく出す（無いと入れ直してよいか分からない）。"""
    html = _columns_html(client)
    row = _row(html, "予備")
    assert "is-unused" in row and "checked" not in row
    cell = _head_cell(row)
    assert "<strong>予備</strong>" in cell
    assert "空欄だけなので、はじめから使わない設定にしています" in cell
    assert 'class="muted small"' in cell   # 見出しより小さく、薄い字で出す

    # ふつうの列（出す列）には何も足さない
    assert _head_cell(_row(html, "対応内容")) == "<td><strong>対応内容</strong></td>"


def test_the_note_says_which_of_the_two_reasons_it_is():
    """理由は2つ（空欄だけ／記録に要らない管理用の列）。出す列には何も書かない。"""
    from app.views import _unused_note

    assert _unused_note(SimpleNamespace(md="omit", omit_reason="blank")) == \
        "空欄だけなので、はじめから使わない設定にしています"
    assert _unused_note(SimpleNamespace(md="omit", omit_reason="dictionary")) == \
        "記録に不要な管理用の列らしいので、はじめから使わない設定にしています"
    assert _unused_note(SimpleNamespace(md="attribute", omit_reason="")) == ""


def test_the_table_no_longer_guesses_the_type_of_a_column(client):
    """「値は「日付」らしい」は出さない（tests/test_tables_flow.py から移したテスト）。

    日付の役割は日付として読める列にしか出ないので、読んでも人には直せない行き止まりだった。
    型エラー・空欄の割合は、表の上の行（data-columns-todo）に出す。
    """
    html = _columns_html(client)
    assert not re.findall(r"値は「([^」]+)」らしい", html)
    assert "型エラー" not in html and "空欄 " not in html


def test_a_column_that_cannot_be_read_is_still_reported_above_the_table(client):
    """知らせの列をやめても、読み取れない値があることは表の上の行で分かる。"""
    html = _columns_html(client, CSV_TEXT.replace(",60,田中", ",不明,田中"), "読めない.csv")
    assert "data-columns-todo" in html
    todo = html.split("data-columns-todo", 1)[1].split("</ul>", 1)[0]
    assert "列「停止時間(分)」に読み取れない値があります（8.3%）。出すかどうか決めてください" in todo


# ---- 2. 役割の名前 ---------------------------------------------------------------------------

def test_the_role_choices_use_words_that_fit_any_table(client):
    """プルダウンの役割は、設備の記録だけでなくどんな表にも当てはまる言い方にする。"""
    html = _columns_html(client)
    # 日付として読める列（発生日）だけが五つとも選べる（ほかの列に「日付」は出さない）
    select = re.search(r'<select data-field="role".*?</select>', _row(html, "発生日"), re.S).group(0)
    labels = re.findall(r"<option [^>]*>(.*?)</option>", select, re.S)
    assert labels == ["識別番号", "日付", "対象（設備・製品・顧客など）",
                      "経過の記録（1つのセルに日付ごとに書き足した列）", "その他"]
    # 画面のどこにも古い言い方を出さない（「設備」は値の例・見出しには出てよい）
    assert ">設備</option>" not in html and "追記ログ" not in html and "AI整形の対象" not in html


def test_the_screen_explains_what_the_two_new_roles_are_for(client):
    """名前を変えただけでは分からないので、何に使う役割なのかを表と一緒に出す。"""
    html = _columns_html(client)
    block = html.split("data-role-help", 1)[1].split("</ul>", 1)[0]
    assert "その記録が何についてのものかを表す列" in block and "長い記録を分けたときも各かたまりに書きます" in block
    assert "例: 対応内容、対応履歴、経過、対応メモ" in block
    assert "日付ごとに切り分けて時系列で出します（AIを使わなくても出ます）" in block
    # 説明は表と一緒に出し入れする（要約1行のときは表ごと隠れる）
    assert html.index("data-columns-table") < html.index("data-role-help")


def test_the_summary_and_the_todo_lines_use_the_new_words(app, client):
    """要約1行と、表の上の「決めてください」の行も新しい言い方にそろえる。"""
    html = _columns_html(client, CSV_TEXT, "要約.csv")
    summary = re.search(r'<p class="col-summary">(.*?)</p>', html, re.S).group(1)
    assert "設備番号＝対象" in summary and "対応内容＝経過の記録" in summary
    assert "設備" not in summary.replace("設備番号", "") and "追記ログ" not in summary

    # 対象・経過の記録の列が無い表では、その名前で「ありません」と書く
    plain = _columns_html(client, "管理No,発生日,数量\r\n" + "".join(
        f"TR-{i:03d},2026-08-{i % 28 + 1:02d},{i}\r\n" for i in range(1, 29)), "数量.csv")
    line = re.search(r'<p class="col-summary">(.*?)</p>', plain, re.S).group(1)
    assert "対象の列はありません。経過の記録の列はありません。" in line


def test_the_role_keys_in_the_saved_spec_do_not_change(app, client):
    """画面の名前を変えても、取り込み設定に入る役割の名前（entity / log）は同じ。

    保存した取り込み設定（spec_json）が読めなくなると、開き直したときに選び直しになる。
    """
    import_id = upload_csv(client, "役割.csv", CSV_TEXT)
    csv_source(client, import_id)
    save_layout(client, import_id)
    body = editor_body(client, import_id)
    roles = {col["header"]: col["role"] for col in body["columns"]}
    assert roles["設備番号"] == "entity" and roles["対応内容"] == "log"   # 画面が送るのは中の名前
    assert save_columns(client, import_id, body).status_code == 200
    wait_import_job(app, import_id)
    with app.app_context():
        spec = tables.get_import(import_id)["spec_json"]
    assert '"role": "entity"' in spec and '"role": "log"' in spec


def test_two_log_columns_are_refused_with_the_new_words(app, client):
    """経過の記録は1列だけ。断るときも新しい言い方にする。"""
    import_id = upload_csv(client, "2列.csv", CSV_TEXT)
    csv_source(client, import_id)
    save_layout(client, import_id)
    body = editor_body(client, import_id)
    for col in body["columns"]:
        if col["header"] in ("現象", "対応内容"):
            col["role"] = "log"
    res = save_columns(client, import_id, body)
    assert res.status_code == 400
    assert res.get_json()["error"] == "経過の記録の列は1つだけにしてください"


# ---- 3. AI整形の段（⑤） ----------------------------------------------------------------------

def test_the_ai_step_says_what_it_does_without_the_old_word(app, client):
    """⑤の1行目は「経過の記録」の列をどうするかを書く（「追記ログ」とは言わない）。"""
    import_id = imported(app, client, "ai.csv", "AI", ai_role="log")
    data = panel(client, import_id, "ai")
    html = data["html"]
    assert "追記ログ" not in html and "AI整形の対象" not in html
    assert "④で「経過の記録」にした列「対応内容」" in html
    assert "ルールで日付・記入者ごとに分けて時系列にします" in html
    assert "AI なしで Markdown に出ます" in html
    assert "「対応の要点」" in html and "記録の種別" in html
    assert data["note"] == "経過の記録の列: 対応内容"


def test_the_ai_step_is_locked_with_the_new_words(app, client):
    """経過の記録の列が無い取り込みでは、その名前で「使いません」と出す。"""
    import_id = imported(app, client, "なし.csv", "なし")
    assert panel(client, import_id, "ai")["locked"] == "経過の記録の列がないので、この取り込みでは使いません"


def test_no_step_of_the_table_import_says_log_append(app, client):
    """表の取り込みのどの段にも「追記ログ」「AI整形の対象」を出さない。"""
    import_id = imported(app, client, "通し.csv", "通し", ai_role="log")
    pages = [client.get("/tables").get_data(as_text=True)]
    for step in ("source", "layout", "columns", "ai", "done"):
        data = panel(client, import_id, step)
        pages.append(data["html"] + (data.get("locked") or ""))
    for text in pages:
        assert "追記ログ" not in text and "AI整形の対象" not in text


# ====================================================================================================
# 元 tests/test_tables_fixes2.py
# 一覧表の不具合修正（2巡目）の確認。
# ====================================================================================================

def _auto_tables_fixes2(src, sheet):
    layout = guess_layout(src, sheet)
    spec = spec_from_suggestions("テスト", layout, suggest_columns(layout.headers, sample_data_rows(src, sheet, layout)), {})
    records, _issues, stats = read_records(src, {"sheet": sheet}, layout, spec)
    return layout, spec, records, stats


# ---- 読み込み（CSV） ----------------------------------------------------------------------------

def test_cp932_file_with_one_bad_byte_stays_cp932(tmp_path):
    """CP932 の拡張文字（髙・﨑）があるファイルに読めないバイトが1つあっても、Shift_JIS-2004 にして字を化けさせない。"""
    rows = "".join(f"TR-00{i},2026-08-01,髙砂ポンプ{i},㈱山﨑製作所で停止\r\n" for i in range(5))
    p = tmp_path / "a.csv"
    p.write_bytes(("管理No,発生日,設備名,内容\r\n" + rows).encode("cp932") + b"TR-999,2026-08-02,\x86\x9f,x\r\n")
    sniff = sniff_csv(p)
    assert sniff.encoding == "cp932" and sniff.decode_error_line == 7 and sniff.warnings
    src = CsvSource(p, "a.csv", {"encoding": "cp932", "delimiter": ",", "errors": "replace"})
    texts = [c.text for r in src.rows() for c in r.cells]
    assert "髙砂ポンプ0" in texts and "㈱山﨑製作所で停止" in texts


def test_long_but_closed_quoted_value_is_a_warning_not_an_error(tmp_path):
    """" が正しく閉じた長い値（600行の貼り付けログ）は確定を止めない。閉じ忘れの疑いとして警告だけ出す。"""
    log = "\n".join(f"{i}: 確認" for i in range(600))
    text = ("管理No,発生日,対応内容\r\nTR-001,2026-08-01,a\r\nTR-002,2026-08-02,b\r\nTR-003,2026-08-03,c\r\n"
            f'TR-004,2026-08-04,"{log}"\r\nTR-005,2026-08-05,e\r\n')
    p = tmp_path / "long.csv"
    p.write_bytes(text.encode("utf-8"))
    src = open_source(p, "long.csv")
    assert len(list(src.rows())) == 6 and src.unclosed_quote_row is None and src.long_record_row == 5
    layout, spec, records, stats = _auto_tables_fixes2(src, "long.csv")
    assert len(records) == 5 and stats.long_record_row == 5
    issues = run_checks(records, spec, stats)
    assert not has_blocking(issues)
    assert any(i.code == "long_record" and i.level == "warning" and "5行目" in i.message for i in issues)


def test_csv_wider_than_the_column_limit_is_refused(tmp_path):
    p = tmp_path / "wide.csv"
    header = ",".join(f"列{i}" for i in range(MAX_COLUMNS + 1))
    p.write_text(header + "\r\n" + ",".join("1" for _ in range(MAX_COLUMNS + 1)) + "\r\n", encoding="utf-8")
    with pytest.raises(UploadError, match="列数が上限"):
        open_source(p, "wide.csv")


# ---- 表の範囲・見出し ----------------------------------------------------------------------------

def test_row_with_its_own_no_and_date_is_not_a_total_row(tmp_path):
    """点検項目が「稼働時間累計」でも、自分の点検No・点検日がある行はデータ（合計行で表を切らない）。"""
    lines = ["点検No,点検日,点検項目,測定値,判定"]
    for i in range(40):
        item = "稼働時間累計" if i % 4 == 3 else f"振動{i}"
        lines.append(f"P-{i:03d},2026/04/{i % 28 + 1:02d},{item},{i * 10},{'' if i % 4 == 3 else 'OK'}")
    p = tmp_path / "check.csv"
    p.write_bytes(("\r\n".join(lines) + "\r\n").encode("utf-8"))
    src = open_source(p, "check.csv")
    layout = guess_layout(src, "check.csv")
    assert layout.counts.get("data") == 40 and not layout.counts.get("subtotal") and layout.warnings == []
    # 番号・日付のない本当の合計行は、これまでどおり合計
    p2 = tmp_path / "total.csv"
    p2.write_bytes(("\r\n".join(lines[:11] + [",,合計,450,"]) + "\r\n").encode("utf-8"))
    src2 = open_source(p2, "total.csv")
    assert guess_layout(src2, "total.csv").counts.get("data") == 10


def test_unlabeled_last_column_is_kept(tmp_path):
    """右端の見出しが空欄でも値のある列は捨てない（「列6」として取り込み、Markdown にも出る）。"""
    text = "管理No,発生日,設備名,現象,停止時間(分),\r\n" + "".join(
        f"T{i},2026/04/{i + 1:02d},ポンプ,停止,{i},ベアリング交換済み{i}\r\n" for i in range(20))
    p = tmp_path / "blank.csv"
    p.write_bytes(text.encode("utf-8"))
    src = open_source(p, "blank.csv")
    layout, spec, records, _stats = _auto_tables_fixes2(src, "blank.csv")
    assert layout.headers[-1] == "列6"
    col = next(c for c in spec.columns if "列6" in (c.headers or [c.display]))
    assert records[3].values.get(col.key) == "ベアリング交換済み3"


def test_repeated_headers_keep_their_name_as_display():
    """同じ見出し「備考」が3つあっても、表示名が「2」「3」にならない。"""
    from app.tables import _dedupe
    from app.tables import _display_name

    headers = _dedupe(["管理No", "発生日", "設備名", "備考", "備考", "備考"])
    assert headers[3:] == ["備考", "備考(2)", "備考(3)"]
    assert [_display_name(h) for h in headers[3:]] == ["備考", "備考(2)", "備考(3)"]
    sugg = suggest_columns(headers, [["T1", "2026/04/01", "ポンプ", "a", "b", "c"]])
    displays = [s["display"] if isinstance(s, dict) else s.display for s in (sugg.values() if isinstance(sugg, dict) else sugg)]
    assert "2" not in displays and "3" not in displays


# ---- 数値の変換 --------------------------------------------------------------------------------

def test_excel_number_under_an_unconvertible_unit_is_reported():
    from app.tables import _convert_number
    from app.tables import ColumnSpec

    col = ColumnSpec(key="downtime", display="停止時間", type="number", role="measure", unit="分")
    assert _convert_number(7200, "7200", col, None, "s") == (120, None, None)   # 秒→分
    assert _convert_number(3, "3", col, None, "時") == (180, None, None)        # 時→分
    value, error, _ = _convert_number(5, "5", col, None, "kg")
    assert value == 5 and error and "換算できません" in error


def test_huge_exponent_is_a_cell_issue_not_a_crash(tmp_path):
    assert parse_number_text("1e309") == (None, "")
    assert parse_number_text("9" * 400) == (None, "")
    assert parse_number_text("1e308") == (1e308, "")
    text = "管理No,発生日,停止時間(分)\r\n" + "".join(f"T{i},2026/04/{i + 1:02d},{i}\r\n" for i in range(10)) + "T99,2026/04/20,1e999\r\n"
    p = tmp_path / "big.csv"
    p.write_bytes(text.encode("utf-8"))
    src = open_source(p, "big.csv")
    layout = guess_layout(src, "big.csv")
    spec = spec_from_suggestions("テスト", layout, suggest_columns(layout.headers, sample_data_rows(src, "big.csv", layout)), {})
    num = next(c for c in spec.columns if c.display.startswith("停止時間"))
    num.type, num.role = "number", "measure"
    records, issues, _stats = read_records(src, {"sheet": "big.csv"}, layout, spec)
    assert len(records) == 11
    assert any("1e999" in i.message and "数値に変換できません" in i.message for i in issues)


def test_overflowing_sum_renders_without_error():
    from app.tables import render_all
    from app.tables import fmt_number

    assert fmt_number(float("inf")) == "inf"
    spec = spec_from_dict(list_spec_dict(group_by="entity_month"))
    records = [_rec("A-1", record_no="A-1", occurred_at="2026-08-03", equipment_id="CMP-101", cost=1e308),
               _rec("A-2", record_no="A-2", occurred_at="2026-08-10", equipment_id="CMP-101", cost=1e308)]
    assert render_all(spec, records, {})


def test_placeholder_equipment_is_not_an_equipment():
    """対象設備が「調査中」だけの記録は、設備別のファイル分けに入れない（本文には原文のまま出す）。"""
    from app.tables import render_all
    from app.tables import entity_value

    d = list_spec_dict(group_by="entity_month")
    d["columns"] = [c for c in d["columns"] if c["key"] != "equipment_name"]
    spec = spec_from_dict(d)
    assert entity_value({"equipment_id": "調査中"}, spec) == ("", "")
    assert entity_value({"equipment_id": "CMP-101"}, spec) == ("CMP-101", "")
    records = [_rec("A-1", record_no="A-1", occurred_at="2026-08-03", equipment_id="調査中", symptom="停止"),
               _rec("A-2", record_no="A-2", occurred_at="2026-08-10", equipment_id="CMP-101", symptom="異音")]
    files = render_all(spec, records, {})
    assert not any("調査中" in f.name for f in files)
    assert "- 設備番号: 調査中" in "".join(f.text for f in files if f.kind == "records")


# ---- 画面（取り込み・削除・ダウンロード） -----------------------------------------------------------------

def _confirmed_import(app, client) -> int:
    return confirmed(app, client, "トラブル一覧.csv", "トラブル対応一覧")


def test_the_import_carries_its_own_spec_and_takes_it_along_when_downloaded(app, client):
    """取り込み設定は保存しない。設定は取り込みの行が持ち、ダウンロードすると一緒に消える。"""
    import_id = _confirmed_import(app, client)
    with app.app_context():
        imp = tables.get_import(import_id)
        assert imp["spec"] is not None and imp["spec"].name == "トラブル対応一覧"
        assert imp["spec_hash"] and imp["template_version_id"] == import_id
    res = client.get(f"/tables/imports/{import_id}/download.zip")
    assert res.status_code == 200 and res.mimetype == "application/zip"
    with app.app_context():
        assert tables.get_import(import_id) is None


def test_second_download_after_purge_is_404_not_500(app, client, monkeypatch):
    import_id = _confirmed_import(app, client)

    def gone(*args, **kwargs):
        raise FileNotFoundError("purged")

    monkeypatch.setattr(views, "build_download", gone)
    assert client.get(f"/tables/imports/{import_id}/download.zip").status_code == 404


def test_delete_is_refused_while_the_preview_draft_is_being_made(app, client):

    import_id = upload_csv(client, "間違い.csv")
    _queued_job(app, import_id, "table_preview")
    res = client.post(f"/tables/imports/{import_id}/delete")
    assert res.status_code == 400 and "Markdownの下書きを作っている間は削除できません" in res.get_json()["error"]
    with app.app_context():
        assert tables.get_import(import_id) is not None


def test_failed_preview_draft_can_be_made_again(app, client, monkeypatch):
    import_id = upload_csv(client, "t.csv")
    csv_source(client, import_id)
    save_layout(client, import_id)
    save_columns(client, import_id, columns_payload("再作成テスト"))
    wait_import_job(app, import_id)

    real = tables.render_files
    calls = {"n": 0}

    def flaky(*args, **kwargs):
        calls["n"] += 1
        if calls["n"] == 1:
            raise OSError("一時的に使えない")
        return real(*args, **kwargs)

    monkeypatch.setattr(tables, "render_files", flaky)

    def wait_preview():
        with app.app_context():
            core.wait_job(core.latest_job("table_import", import_id, kind="table_preview")["id"], timeout=60)

    panel(client, import_id, "preview")
    wait_preview()
    data = panel(client, import_id, "preview")
    assert data["failed"] is True and "もう一度作る" in data["html"]
    # 下書きの失敗を「表を読み込めませんでした」と取り違えない（読み込みは終わっている）
    assert "表を読み込めませんでした" not in data["html"]
    res = client.post(f"/tables/imports/{import_id}/preview")
    assert res.status_code == 200 and res.get_json()["building"] is True
    wait_preview()
    assert "確定してMarkdownを作成" in panel(client, import_id, "preview")["html"]


def test_import_status_is_written_before_the_job_can_finish(app, monkeypatch):
    """ジョブがすぐ終わって書いた状態を、あとから reading / confirming で上書きしない。"""
    with app.app_context():
        import_id = tables.create_import("x.csv", "h" * 64, "tables/x.csv", {"encoding": "utf-8"})

        def instant(kind, ref_type, ref_id, fn, params):
            # ワーカーが先に終わった場合と同じ: 呼び出し元が状態を書く前にジョブが状態を書く
            tables.update_import(ref_id, status="preview" if kind == "table_read" else "confirmed")
            return 12345

        monkeypatch.setattr(tables, "start_job", instant)
        assert tables.start_read_job(import_id) == 12345
        imp = tables.get_import(import_id)
        assert imp["status"] == "preview" and imp["job_id"] == 12345
        tables.start_render_job(import_id)
        assert tables.get_import(import_id)["status"] == "confirmed"


# ---- 読み込み（Excel） --------------------------------------------------------------------------

def _small_book(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append(["管理No", "発生日", "設備名", "現象", "停止時間"])
    for i in range(10):
        ws.append([f"T{i}", f"2026/04/{i + 1:02d}", "ポンプ", "停止", i + 1])
    return wb, ws


def test_format_only_far_cell_does_not_block_upload(app, client, tmp_path):
    from app.tables import ExcelSource

    wb, ws = _small_book(tmp_path)
    ws["XFD1048576"].fill = PatternFill("solid", fgColor="FFFF00")
    path = tmp_path / "far.xlsx"
    wb.save(path)
    ExcelSource(path).close()
    res = upload(client, path.read_bytes(), "far.xlsx")
    assert res.status_code == 200 and res.get_json()["import_id"]


def test_workbook_with_unreadable_values_is_refused_at_upload(app, client, tmp_path):
    wb, _ws = _small_book(tmp_path)
    wb.save(tmp_path / "ok.xlsx")
    src = zipfile.ZipFile(tmp_path / "ok.xlsx")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb"<v>3</v>", b"<v>NaN</v>", data, count=1)
            z.writestr(item, data)
    res = upload(client, out.getvalue(), "nan.xlsx")
    assert res.status_code == 400 and res.get_json()["error"]
    with app.app_context():
        assert tables.list_imports(limit=10) == []
    assert [p for p in Path(app.config["UPLOAD_DIR"]).rglob("*") if p.is_file()] == []


def test_retention_note_only_says_the_data_is_deleted():
    """渡すのは zip だけになったので、案内も「消える」ことだけにする（CSVの順番の話は残さない）。"""
    from app.views import TABLES_DELETE_ON_DOWNLOAD_NOTE as DELETE_ON_DOWNLOAD_NOTE

    assert "より先" not in DELETE_ON_DOWNLOAD_NOTE and "CSV" not in DELETE_ON_DOWNLOAD_NOTE
    assert "管理用" not in DELETE_ON_DOWNLOAD_NOTE and "サーバーから消えます" in DELETE_ON_DOWNLOAD_NOTE


def test_confirmed_import_can_be_deleted_without_downloading(app, client):
    """ダウンロードせずに消す道を、確定したあとの段にも残す。"""
    import_id = _confirmed_import(app, client)
    assert "data-import-delete" in panel(client, import_id, "done")["html"]
    res = client.post(f"/tables/imports/{import_id}/delete")
    assert res.status_code == 200 and res.get_json()["ok"] is True
    with app.app_context():
        assert tables.get_import(import_id) is None


def test_preview_headers_show_the_unit(app, client):
    import_id = _confirmed_import(app, client)
    assert "<th>停止時間（分）</th>" in preview_panel(app, client, import_id)["html"]


def test_a_panel_of_a_missing_import_is_404(client):
    assert client.get("/tables/imports/999/panel/done").status_code == 404
    assert client.get("/tables/imports/999/panel/unknown").status_code == 404


# ====================================================================================================
# 元 tests/test_tables_fixes3.py
# 一覧表の不具合修正（3巡目）の確認。
# ====================================================================================================

def _csv(tmp_path, name: str, lines: list[str], encoding: str = "utf-8"):
    p = tmp_path / name
    p.write_bytes(("\r\n".join(lines) + "\r\n").encode(encoding))
    return open_source(p, name)


def _auto_tables_fixes3(src, sheet, tweak=None):
    layout = guess_layout(src, sheet)
    suggestions = suggest_columns(layout.headers, sample_data_rows(src, sheet, layout))
    if tweak:
        tweak(suggestions)
    spec = spec_from_suggestions("テスト", layout, suggestions, {})
    records, issues, stats = read_records(src, {"sheet": sheet}, layout, spec)
    return layout, spec, records, stats


# ---- T3-1: 日付を書き省いた行 ------------------------------------------------------------------

def _date_once_lines():
    lines = ["日付,設備,現象,処置,停止時間(分)"]
    for i in range(30):
        date = f"2026/08/{i // 3 + 1:02d}" if i % 3 == 0 else ""
        lines.append(f"{date},設備{i + 1},ポンプ停止{i + 1},部品を交換した{i + 1},{(i + 1) * 5}")
    return lines


@pytest.mark.parametrize("fill_down", [False, True])
def test_row_with_blank_date_but_own_numbers_is_its_own_record(tmp_path, fill_down):
    """日付を1日1回だけ書いた表で、日付が空の行も数値（停止時間）を持つなら別の記録。前の記録に黙って混ぜない。"""
    src = _csv(tmp_path, "once.csv", _date_once_lines())

    def tweak(suggestions):
        if fill_down:
            for s in suggestions:
                if s.header == "日付":
                    s.fill_down_blank = True

    _layout, _spec, records, stats = _auto_tables_fixes3(src, "once.csv", tweak)
    assert len(records) == 30 and stats.continuation_merged == 0
    downtime = [v for r in records for k, v in r.values.items() if isinstance(v, (int, float)) and k != "record_no"]
    assert sum(downtime) == sum((i + 1) * 5 for i in range(30))
    if fill_down:
        assert all(any(str(v).startswith("2026-08") for v in r.values.values()) for r in records)


def test_text_only_row_with_blank_keys_is_still_a_continuation(tmp_path):
    lines = ["管理No,発生日,現象,停止時間(分)"]
    for i in range(10):
        lines.append(f"TR-{i:03d},2026/08/{i + 1:02d},停止{i},{i + 1}")
        if i == 4:
            lines.append(",,追記: 後日点検")
    src = _csv(tmp_path, "cont.csv", lines)
    _layout, _spec, records, stats = _auto_tables_fixes3(src, "cont.csv")
    assert len(records) == 10 and stats.continuation_merged == 1
    assert any("後日点検" in str(v) for v in records[4].values.values())


# ---- T3-2: 秒の小数・時差つきの日時 ------------------------------------------------------------

@pytest.mark.parametrize("tail", [".000", "+09:00", "Z"])
def test_datetime_with_fraction_or_offset_is_converted(tmp_path, tail):
    """SQL Server 等の「2026-08-02 10:01:00.000」、Webの「…T10:01:00+09:00」「…Z」も日時として読む（確定を止めない）。"""
    sep = " " if tail == ".000" else "T"
    lines = ["管理No,発生日時,現象"] + [f"TR-{i:03d},2026-08-{i + 1:02d}{sep}10:{i:02d}:00{tail},停止{i}" for i in range(20)]
    src = _csv(tmp_path, "sqlexp.csv", lines)
    _layout, spec, records, stats = _auto_tables_fixes3(src, "sqlexp.csv")
    col = next(c for c in spec.columns if "発生日時" in c.headers)
    assert col.type == "datetime"
    assert not stats.type_errors
    assert records[3].values[col.key] == "2026-08-04 10:03"
    assert not has_blocking(run_checks(records, spec, stats))


# ---- T3-3: 空行の後の書き添え ------------------------------------------------------------------

def _footer_lines(footer: str):
    lines = ["管理No,発生日,設備,現象"] + [f"TR-{i:03d},2026/08/{i + 1:02d},設備{i},停止{i + 1}" for i in range(20)]
    return lines + ["", footer]


@pytest.mark.parametrize("footer", ["以上", "作成者：保全課 田中", ",,,以上"])
def test_footer_after_blank_row_is_not_a_record(tmp_path, footer):
    """表の下の「以上」「作成者：…」は記録にも、最後の記録の文章にもしない。"""
    src = _csv(tmp_path, "foot.csv", _footer_lines(footer))
    layout, _spec, records, stats = _auto_tables_fixes3(src, "foot.csv")
    assert layout.data_end == 21
    assert len(records) == 20 and stats.continuation_merged == 0
    assert all("以上" not in str(v) and "田中" not in str(v) for r in records for v in r.values.values())
    assert any(rc.index == 23 and rc.kind == "note" for rc in layout.row_classes)


def test_text_only_row_after_blank_row_is_not_merged_silently(tmp_path):
    """空行を挟んだ文章だけの行は、前の記録に黙って混ぜない（1件として出し、警告で気づけるようにする）。"""
    lines = ["管理No,発生日,設備,現象"] + [f"TR-{i:03d},2026/08/{i + 1:02d},設備{i},停止{i + 1}" for i in range(5)]
    lines += ["", ",,,別件の停止", "TR-100,2026/08/20,設備9,停止9"]
    src = _csv(tmp_path, "gap.csv", lines)
    _layout, _spec, records, stats = _auto_tables_fixes3(src, "gap.csv")
    assert stats.continuation_merged == 0 and len(records) == 7
    assert "別件" not in str(records[4].values)


# ---- T3-4: 結合した2段見出しを書き出したCSV ----------------------------------------------------

def test_csv_from_merged_two_row_header_reads_both_rows_as_header(tmp_path):
    lines = ["管理No,発生,,設備,,停止時間(分)", ",日付,時刻,番号,名称,"]
    lines += [f"TR-{i:03d},2026/08/{i + 1:02d},10:{i:02d},M-{i:02d},プレス{i},{i + 5}" for i in range(20)]
    src = _csv(tmp_path, "two_nopre.csv", lines)
    layout, _spec, records, stats = _auto_tables_fixes3(src, "two_nopre.csv")
    assert layout.header_rows == [1, 2]
    assert layout.headers[:5] == ["管理No", "発生_日付", "発生_時刻", "設備_番号", "設備_名称"]
    assert len(records) == 20 and not stats.type_errors


def test_csv_single_header_row_is_not_joined_with_first_data_row(tmp_path):
    lines = ["管理No,発生日,設備,現象"] + [f"TR-{i:03d},2026/08/{i + 1:02d},設備{i},停止{i + 1}" for i in range(20)]
    layout = guess_layout(_csv(tmp_path, "one.csv", lines), "one.csv")
    assert layout.header_rows == [1]


# ---- T3-5: CP932 のファイルに UTF-8 の行が混ざる -----------------------------------------------

def test_utf8_lines_appended_to_cp932_csv_are_warned(tmp_path):
    from app.tables import sniff_csv

    head = "管理No,発生日,現象\r\n" + "".join(f"TR-{i:03d},2026/08/{i % 28 + 1:02d},ポンプ停止{i}\r\n" for i in range(20))
    tail = "".join(f"TR-{i:03d},2026/08/{i % 28 + 1:02d},ポンプ停止{i}\r\n" for i in range(20, 28))
    p = tmp_path / "mixed.csv"
    p.write_bytes(head.encode("cp932") + tail.encode("utf-8"))
    sniff = sniff_csv(p)
    assert sniff.encoding == "cp932"
    assert any("UTF-8 の行が混ざっています" in w and "22・23" in w for w in sniff.warnings)


def test_plain_cp932_csv_has_no_utf8_warning(tmp_path):
    from app.tables import sniff_csv

    text = "管理No,発生日,現象\r\n" + "".join(f"TR-{i:03d},2026/08/{i % 28 + 1:02d},ポンプ停止ｱｲｳ髙{i}\r\n" for i in range(30))
    p = tmp_path / "sjis.csv"
    p.write_bytes(text.encode("cp932"))
    assert not any("UTF-8" in w for w in sniff_csv(p).warnings)


# ---- F4: CSV の読み込みエラーに英語の例外文を出さない -------------------------------------------

def test_csv_error_message_has_no_english_exception_text(tmp_path):
    import csv
    import re

    from app.tables import UploadError

    p = tmp_path / "unclosed.csv"
    p.write_bytes(('管理No,発生日,現象\r\nTR-001,2026/08/01,"閉じていない\r\n' + "続き,x\r\n" * 50).encode("utf-8"))
    old = csv.field_size_limit(100)
    try:
        from app.tables import CsvSource

        src = CsvSource(p, "unclosed.csv", {"encoding": "utf-8", "delimiter": ","})
        with pytest.raises(UploadError) as e:
            list(src.rows())
    finally:
        csv.field_size_limit(old)
    msg = str(e.value)
    assert "値が大きすぎます" in msg and not re.search(r"[A-Za-z]{4,}", msg.replace("CSV", ""))


# ---- T3-6: 「対応」「作業」の文章の列を人名の列にしない ------------------------------------------

def test_prose_columns_similar_to_person_headers_stay_in_the_body(tmp_path):
    lines = ["管理No,発生日,対応,作業,担当"]
    for i in range(20):
        lines.append(f"TR-{i:03d},2026/08/{i + 1:02d},"
                     f"モーターの異音を確認したためベアリングを交換し、試運転で異常がないことを確かめた{i},"
                     f"カバーを外して内部を清掃し、潤滑油を補充したうえで締め付けを点検した{i},石川 陸")
    src = _csv(tmp_path, "prose.csv", lines)
    layout = guess_layout(src, "prose.csv")
    by_header = {s.header: s for s in suggest_columns(layout.headers, sample_data_rows(src, "prose.csv", layout))}
    for h in ("対応", "作業"):
        assert by_header[h].role != "person" and by_header[h].md != "omit"
    assert by_header["担当"].role == "person"


# ---- C3-2: 同じ設定の別の取り込みの AI の結果で、下書きを作り直さない ------------------------------

def test_ai_results_of_another_import_leave_the_preview_signature_alone(app, client):
    from app import aiproc as ai_items
    from app import tables

    import_id = _confirmed_import(app, client)
    with app.app_context():
        imp = tables.get_import(import_id)
        spec = tables.spec_for_import(imp)
        before = tables._preview_signature(import_id, imp, spec)
        ai_items.upsert_item(imp["template_id"], "log", "row-x", status="ok", import_id=import_id + 1000)
        assert tables._preview_signature(import_id, imp, spec) == before
        ai_items.upsert_item(imp["template_id"], "log", "row-x", status="ok", import_id=import_id)
        assert tables._preview_signature(import_id, imp, spec) != before


# ---- F2: Excel で計算されていない数式 -------------------------------------------------------------

def test_uncached_formulas_are_warned(tmp_path):
    """プログラムが書いて Excel で保存していない数式は値がない。空欄として読むことを警告で知らせる。"""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["管理No", "発生日", "設備", "停止時間(分)", "対応内容"])
    for i in range(12):
        ws.append([f"TR-{i:03d}", "2026-08-01", f"設備{i}", f"={i}+1", f"内容{i}"])
    p = tmp_path / "formula.xlsx"
    wb.save(p)
    src = open_source(p, "formula.xlsx")
    _layout, spec, records, stats = _auto_tables_fixes3(src, ws.title)
    assert sum(stats.uncached_formulas.values()) == 12
    issues = run_checks(records, spec, stats)
    assert any(i.code == "uncached_formula" and "12個" in i.message and "停止時間" in i.message for i in issues)


def test_cached_formula_values_are_not_warned(tmp_path):
    """Excel が保存した数式（値あり、空文字列の結果は t="str" と空の v）は警告しない。"""
    import zipfile

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["管理No", "発生日", "停止時間(分)", "備考"])
    for i in range(12):
        ws.append([f"TR-{i:03d}", "2026-08-01", 5, "x"])
    p = tmp_path / "cached.xlsx"
    wb.save(p)
    # 「Excel で保存した」形に書き換える: C列は値つきの数式、D列は空文字列を返す数式
    with zipfile.ZipFile(p) as z:
        items = {n: z.read(n) for n in z.namelist()}
    xml = items["xl/worksheets/sheet1.xml"].decode("utf-8")
    import re

    xml = re.sub(r'<c r="(C\d+)" t="n"><v>5</v></c>', r'<c r="\1"><f>2+3</f><v>5</v></c>', xml)
    xml = re.sub(r'<c r="(D(?!1")\d+)" t="inlineStr"><is><t>x</t></is></c>', r'<c r="\1" t="str"><f>""</f><v></v></c>', xml)
    items["xl/worksheets/sheet1.xml"] = xml.encode("utf-8")
    with zipfile.ZipFile(p, "w") as z:
        for n, data in items.items():
            z.writestr(n, data)
    src = open_source(p, "cached.xlsx")
    assert src.uncached_formulas(ws.title) == {}
    assert xml.count("<f>") == 24


# ---- F3: 打ち間違えた遠い年の日付 ----------------------------------------------------------------

def test_far_off_date_is_warned():
    from app.tables import render_all
    from app.tables import spec_from_dict

    spec = spec_from_dict(list_spec_dict())
    dates = [f"2025-08-{d:02d}" for d in range(1, 11)] + ["2052-08-11"]
    records = [_rec(f"TR-{i:03d}", record_no=f"TR-{i:03d}", occurred_at=d, equipment_id="CMP-101", symptom="停止")
               for i, d in enumerate(dates)]
    for i, r in enumerate(records):
        r["source"]["row"] = i + 2
    issues = run_checks(records, spec, {})
    outlier = [i for i in issues if i.code == "date_outlier"]
    assert len(outlier) == 1 and "12行目（2052-08-11）" in outlier[0].message
    # 記録ファイルは記録のある月だけ（間の月は作らない）
    assert sorted(f.name for f in render_all(spec, records, {})) == ["トラブル対応一覧_2025-08.md",
                                                                     "トラブル対応一覧_2052-08.md"]


def test_short_gaps_between_months_are_not_warned():
    from app.tables import spec_from_dict

    spec = spec_from_dict(list_spec_dict())
    records = [_rec("a", occurred_at="2025-01-10"), _rec("b", occurred_at="2025-06-10")]
    assert not [i for i in run_checks(records, spec, {}) if i.code == "date_outlier"]


def test_zero_serial_in_a_date_formatted_cell_is_a_type_error():
    from datetime import datetime

    from app.tables import ConvertContext, _convert_date

    out, error, _flag = _convert_date(datetime(1899, 12, 30), "1899-12-30", "date", ConvertContext())
    assert error and "変換できません" in error
    assert _convert_date(datetime(2026, 8, 1), "2026-08-01", "date", ConvertContext())[:2] == ("2026-08-01", None)


# ---- R3B-5: 「出さない」にした理由 -------------------------------------------------------------

def test_only_blank_columns_are_left_out_by_default(tmp_path):
    lines = ["管理No,発生日,起票者,予備,区分コード,現象"]
    for i in range(30):
        lines.append(f"TR-{i:03d},2026/08/{i % 28 + 1:02d},石川 陸,,A{i % 3},ポンプ停止{i}")
    src = _csv(tmp_path, "omit.csv", lines)
    layout = guess_layout(src, "omit.csv")
    by_header = {s.header: s for s in suggest_columns(layout.headers, sample_data_rows(src, "omit.csv", layout))}
    # 人名の列・コードの列も既定で出す（空欄だけの列だけ「出さない」にする）
    assert by_header["起票者"].md != "omit" and by_header["区分コード"].md != "omit"
    assert by_header["予備"].md == "omit" and by_header["予備"].omit_reason == "blank"
    assert by_header["現象"].omit_reason == ""


def test_preview_is_not_queued_behind_a_paused_ai_job(app, client):
    """AI整形が一時停止中のまま「内容の確認」の段を開いても、下書きのジョブを後ろに並べて黙って待たせない。"""
    import time

    from app import core
    from tests.conftest import imported, panel

    import_id = imported(app, client, "トラブル一覧.csv", "停止中テスト", ai_role="log")

    def ai(ctx):
        while ctx.wait_if_paused():
            time.sleep(0.01)

    with app.app_context():
        ai_job = core.start_job("ai_format", "table_import", import_id, ai)
        deadline = time.time() + 5
        while core.get_job(ai_job)["status"] != "running" and time.time() < deadline:
            time.sleep(0.01)
        core.request_pause(ai_job)
        while core.get_job(ai_job)["status"] != "paused" and time.time() < deadline:
            time.sleep(0.01)
    try:
        preview = panel(client, import_id, "preview")
        assert "AI整形が動いています（一時停止中を含む）" in preview["locked"]
        with app.app_context():
            assert core.latest_job("table_import", import_id, kind="table_preview") is None
        assert "確認に進めません" in panel(client, import_id, "ai")["html"]
    finally:
        with app.app_context():
            core.request_cancel(ai_job)
            core.wait_job(ai_job, timeout=5)


# ====================================================================================================
# 元 tests/test_tables_fixes4.py
# 一覧表の不具合修正（4巡目）の確認。
# ====================================================================================================

# ---- T4-1: 計器名（pH計・温度計・電力量累計）の行は小計・合計ではない ----------------------------

_METERS = ["pH計", "温度計", "O2計", "流量計", "湿度計", "圧力計", "CO2計", "導電率計", "照度計", "粘度計"]


def test_meter_names_with_own_date_are_data_rows(tmp_path):
    lines = ["計器,点検日,指示値,備考"] + [f"{m},2026/08/{i + 1:02d},{i + 1}.5," for i, m in enumerate(_METERS)]
    src = _csv(tmp_path, "m1.csv", lines)
    layout = guess_layout(src, src.sheets()[0].name)
    assert layout.counts.get("subtotal", 0) == 0
    assert layout.counts.get("data") == 10


def test_cumulative_meter_rows_stay_data_with_auto_and_manual_range(tmp_path):
    lines = ["計器,点検日,指示値,備考"]
    names = ["pH計", "温度計", "電力量累計", "流量計", "稼働時間累計", "圧力計", "CO2計", "照度計", "粘度計", "湿度計"]
    for i, m in enumerate(names):
        lines.append(f"{m},2026/08/{i + 1:02d},{(i + 1) * 100},良好")
    lines.append("合計,,,1234")
    src = _csv(tmp_path, "m2.csv", lines)
    sheet = src.sheets()[0].name
    layout = guess_layout(src, sheet)
    assert layout.data_end >= 11
    kinds = {row.index: rc.kind for row, rc in classify_rows(src, sheet, layout)}
    assert kinds[4] == "data" and kinds[6] == "data"
    assert kinds[12] == "subtotal"  # 日付のない本当の合計行は今までどおり
    manual = guess_layout(src, sheet, header_rows=[1], data_end=12)
    kinds = {row.index: rc.kind for row, rc in classify_rows(src, sheet, manual)}
    assert kinds[4] == "data" and kinds[6] == "data" and kinds[12] == "subtotal"


def test_plain_subtotal_row_is_still_subtotal(tmp_path):
    lines = ["部署,日付,件数"] + [f"製造{i},2026/08/{i + 1:02d},{i}" for i in range(6)] + ["部署計,,123"]
    src = _csv(tmp_path, "s.csv", lines)
    sheet = src.sheets()[0].name
    layout = guess_layout(src, sheet, header_rows=[1], data_end=8)
    kinds = {row.index: rc.kind for row, rc in classify_rows(src, sheet, layout)}
    assert kinds[8] == "subtotal"


# ---- T4-2 / T4-3 / T4-5: 列の対応づけの保存で、値から決めた内容を壊さない ------------------------

def _editor_body(app, client, text=None):
    """取り込みを置いて「列の対応づけ」の段を開き、画面が送る形の表を返す。"""
    from tests.conftest import SPEC_CSV, editor_body, spec_import

    import_id = spec_import(app, client, text=text or SPEC_CSV)
    return import_id, editor_body(client, import_id)


def _saved(app, client, import_id, body):
    from app import tables
    from tests.conftest import save_columns, wait_import_job

    res = save_columns(client, import_id, body)
    assert res.status_code == 200, res.get_json()
    wait_import_job(app, import_id)
    with app.app_context():
        return tables.spec_for_import(tables.get_import(import_id))


def test_the_editor_only_asks_for_use_and_role(app, client):
    """画面に出すのは「使う・見出し・役割」だけ（キー・型・単位・出し方は値から決める）。"""
    from tests.conftest import panel_html, spec_import

    import_id = spec_import(app, client)
    html = panel_html(client, import_id, "columns")
    for gone in ('data-field="key"', 'data-field="type"', 'data-field="unit"', 'data-field="md"',
                 'data-field="display"', 'data-field="description"', 'data-field="fill_down_blank"',
                 'data-field="ai"', 'data-setting="file_prefix"', 'data-setting="description"',
                 'data-setting="group_by"'):
        assert gone not in html, gone
    assert 'data-field="use"' in html and 'data-field="role"' in html and 'data-setting="name"' in html


def test_unticked_column_is_still_removed(app, client):
    import_id, body = _editor_body(app, client)
    for r in body["columns"]:
        if r["header"].startswith("停止時間"):
            r["use"] = False
    after = _saved(app, client, import_id, body)
    assert after.column("downtime") is None
    assert after.column("record_no") is not None


def test_the_record_key_follows_the_role_picked_on_screen(app, client):
    from tests.conftest import set_role

    import_id, body = _editor_body(app, client)
    after = _saved(app, client, import_id, body)
    assert after.record["key"] == ["record_no"]

    set_role(body, "管理No", "attribute")
    set_role(body, "設備名", "key")
    after = _saved(app, client, import_id, body)
    assert after.record["key"] == [after.first_role("key").key]
    assert after.first_role("key").display == "設備名"


def test_the_table_name_decides_the_file_names(app, client):
    """「ファイル名の先頭」は無くなった。md の名前は「表の名前」を使う。"""
    import_id, body = _editor_body(app, client)
    body["name"] = "新しい名前"
    after = _saved(app, client, import_id, body)
    assert after.markdown["file_prefix"] == "" and after.file_prefix == "新しい名前"


# ---- T4-4 / S4-1: JSON で取り込む設定の AI整形の項目を確かめる ----------------------------------------

def _log_spec(**stage):
    from app.tables import spec_from_dict

    return spec_from_dict({"name": "x", "columns": [{"key": "log", "display": "ログ", "type": "text", "role": "log"}],
                           "log_stage": {"column": "log", **stage}})


def test_mask_given_as_a_string_is_read_as_rules():
    from app.tables import parse_log_cell
    from app.tables import validate_spec

    text = "8/1 10:00 田中: 連絡先 090-1234-5678 / taro@example.com に電話"
    for mask in ("email", "phone,email", "phone、email"):
        spec = _log_spec(mask=mask)
        assert validate_spec(spec) == []
        out = repr(parse_log_cell(spec, {"log": text}))
        assert "taro@example.com" not in out and "［メール］" in out
    assert "090-1234-5678" not in repr(parse_log_cell(_log_spec(mask="phone,email"), {"log": text}))


def test_malformed_log_stage_and_checks_are_refused():
    from app.tables import spec_from_dict, validate_spec

    assert any("伏せ字" in e for e in validate_spec(_log_spec(mask=["phone", "住所"])))
    assert any("人名一覧" in e for e in validate_spec(_log_spec(people=["田中"])))
    assert any("用語集" in e for e in validate_spec(_log_spec(glossary=["a"])))
    assert any("区切り" in e for e in validate_spec(_log_spec(splitter="x")))
    assert validate_spec(_log_spec(people=[{"name": "田中 一郎", "aliases": ["田中"]}])) == []
    spec = spec_from_dict({"name": "x", "columns": [{"key": "a", "display": "A"}],
                           "checks": {"reconcile_tolerance": "なし"}})
    assert any("許容差" in e for e in validate_spec(spec))


def test_bad_or_slow_split_patterns_are_refused():
    from app.tables import validate_spec

    assert any("正しくありません" in e for e in validate_spec(_log_spec(splitter={"extra_anchors": ["["]})))
    assert any("遅く" in e for e in validate_spec(_log_spec(splitter={"extra_anchors": ["(.+)+X"]})))
    assert any("遅く" in e for e in validate_spec(_log_spec(splitter={"not_date_patterns": ["(a*)*b"]})))
    assert validate_spec(_log_spec(splitter={"extra_anchors": [r"^【\d+】"]})) == []


# ---- R4-MD-1: 「〃」「同上」は直前の行の値で補う ----------------------------------------------------

def test_ditto_marks_are_filled_from_the_row_above(tmp_path):
    from app.tables import render_all

    lines = ["管理No,発生日,設備ID,設備名,現象,停止時間(分)",
             "TR-1,2025-03-01,CMP-101,CMP研磨機1号機,異音,30",
             "TR-2,〃,〃,〃,振動,20",
             "TR-3,2025-03-02,ETC-301,エッチャ1号機,停止,40",
             "TR-4,2025-03-02,同上,″,停止,50",
             "TR-5,2025-03-03,CVD-201,CVD1号機,警報,10",
             "TR-6,2025-03-04,CVD-202,CVD2号機,警報,10"]
    src = _csv(tmp_path, "d.csv", lines)
    layout, spec, records, stats = _auto_tables_fixes3(src, src.sheets()[0].name)
    by_key = {r.key: r.values for r in records}
    date_key = spec.date_key
    eq = spec.first_role("entity").key
    assert str(by_key["TR-2"][date_key]).startswith("2025-03-01") and by_key["TR-2"][eq] == "CMP-101"
    assert by_key["TR-4"][eq] == "ETC-301"
    assert stats.ditto_filled
    files = render_all(spec, [r.to_dict() for r in records], {})
    assert not any("〃" in f.name or "同上" in f.name for f in files)
    text = "\n".join(f.text for f in files)
    assert "〃" not in text and "日付なし" not in "".join(f.name for f in files)


def test_ditto_without_a_row_above_is_not_an_equipment(tmp_path):
    from app.tables import entity_value

    lines = ["管理No,発生日,設備ID,現象,停止時間(分)",
             "TR-1,2025-03-01,〃,異音,30",
             "TR-2,2025-03-02,CMP-101,振動,20",
             "TR-3,2025-03-02,ETC-301,停止,40",
             "TR-4,2025-03-03,CVD-201,警報,10"]
    src = _csv(tmp_path, "d2.csv", lines)
    _layout, spec, records, stats = _auto_tables_fixes3(src, src.sheets()[0].name)
    first = next(r for r in records if r.key == "TR-1")
    assert spec.first_role("entity") is not None
    assert entity_value(first.values, spec) == ("", "")
    assert not stats.ditto_filled


# ---- ux4-3: 変換できなかった日付は日付の範囲に入れない -------------------------------------------------

def test_date_range_ignores_unconverted_date_text(tmp_path):

    lines = ["管理No,発生日,現象"] + [f"TR-{i},{d},停止" for i, d in enumerate(
        ["2026-08-01", "2026/13/45", "不明", "2026-08-04", "2026-08-05", "2026-08-06"])]
    src = _csv(tmp_path, "u.csv", lines)
    _layout, _spec, _records, stats = _auto_tables_fixes3(src, src.sheets()[0].name)
    assert (stats.date_min, stats.date_max) == ("2026-08-01", "2026-08-06")


# ---- R4-MD-3: 時系列の記入者は、担当者の列から補った表記で出す ------------------------------------------

def _person_spec():
    from app.tables import spec_from_dict

    return spec_from_dict({
        "name": "x", "columns": [
            {"key": "no", "display": "管理No", "type": "code", "role": "key"},
            {"key": "occurred_at", "display": "発生日", "type": "date", "role": "date"},
            {"key": "log", "display": "対応内容", "type": "text", "role": "log"},
            {"key": "worker", "display": "担当者", "type": "string", "role": "person"}],
        "log_stage": {"column": "log"}})


def _timeline(log: str) -> str:
    from app.tables import people_index_for, record_block

    spec = _person_spec()
    rec = {"key": "A-1", "values": {"no": "A-1", "occurred_at": "2026-07-01", "log": log, "worker": "井上 亮"},
           "originals": {}, "source": {"file": "a.csv", "row": 2}, "warnings": []}
    return "\n".join(record_block(rec, spec, None, people_index_for(spec, [rec])))


def test_the_person_column_is_written_and_fills_in_the_timeline_author():
    """人名の列は出す。時系列の記入者も担当者の列から補う（内容を隠さない）。"""
    log = "7/1 2:47 井上:連絡あり\n7/1 3:10 部品を交換した"
    text = _timeline(log)
    assert "- 担当者: 井上 亮" in text
    assert "02:47 井上 亮: 連絡あり" in text and "井上 亮（推定）" in text


# ---- R4-1: AIの試し実行中は取り込みを削除・ダウンロードできない -------------------------------------------



def test_delete_is_refused_while_an_ai_trial_is_running(ai_app_endpoints_ai, ai_client, monkeypatch):
    from app import aiproc
    from app import tables

    import_id = _read_import(ai_app_endpoints_ai, ai_client)
    seen = {}

    def slow_trial(*_a, **_k):
        # 同じブラウザの別のタブから消しに来る（別のブラウザからは、そもそもこの取り込みが見えない）
        seen["delete"] = ai_client.post(f"/tables/imports/{import_id}/delete").get_json()
        raise aiproc.AIJobError("止めました")

    monkeypatch.setattr(aiproc, "trial_row", slow_trial)
    res = ai_client.post(f"/tables/imports/{import_id}/ai/trial", json={"row_key": "TR-001"})
    assert res.status_code == 400
    assert "AIの試し実行中は削除・ダウンロード・保存できません" in seen["delete"]["error"]
    with ai_app_endpoints_ai.app_context():
        assert tables.get_import(import_id) is not None
    # 試し実行が終われば削除できる
    ai_client.post(f"/tables/imports/{import_id}/delete")
    with ai_app_endpoints_ai.app_context():
        assert tables.get_import(import_id) is None


# ---- C4-1: AI整形の実行中・一時停止中は読み込み直さない ------------------------------------------------

def test_reread_is_refused_while_ai_format_is_running_or_paused(ai_app_endpoints_ai, ai_client, monkeypatch):
    from app import core
    from app import tables
    from app import views as views_tables

    import_id = _read_import(ai_app_endpoints_ai, ai_client)
    with ai_app_endpoints_ai.app_context():
        before = tables.get_import(import_id)
    monkeypatch.setattr(views_tables, "_ai_running", lambda _id: True)
    res = ai_client.post(f"/tables/imports/{import_id}/read")
    assert res.status_code == 409 and "処理中は変更できません" in res.get_json()["error"]
    with ai_app_endpoints_ai.app_context():
        after = tables.get_import(import_id)
        assert after["status"] == before["status"] and after["job_id"] == before["job_id"]
        job = core.latest_job("table_import", import_id, kind="table_read")
        assert job is None or job["id"] == before["job_id"]


# ---- R4-FUZZ-1 / R4-FUZZ-2: 列数の上限・読めない値はアップロードの時点で断る ---------------------------------

def _no_import_left(app):
    from pathlib import Path

    from app import tables

    with app.app_context():
        assert tables.list_imports(limit=10) == []
    assert [p for p in Path(app.config["UPLOAD_DIR"]).rglob("*") if p.is_file()] == []


def test_csv_with_a_huge_first_line_is_refused_at_upload(app, client):
    from tests.conftest import upload

    header = ",".join(f"c{i}" for i in range(400_000))
    body = (header + "\r\n" + ",".join("1" for _ in range(3)) + "\r\n").encode("utf-8")
    res = upload(client, body, "wide.csv")
    assert res.status_code == 400 and "列数が上限" in res.get_json()["error"]
    _no_import_left(app)


def test_csv_rows_over_the_column_limit_are_refused_when_read(tmp_path):
    import pytest

    from app.core import UploadError
    from app.tables import MAX_COLUMNS, CsvSource

    p = tmp_path / "w.csv"
    p.write_bytes(("a,b\r\n1,2\r\n" + ",".join("x" for _ in range(MAX_COLUMNS + 5)) + "\r\n").encode("utf-8"))
    src = CsvSource(p, "w.csv", {"encoding": "utf-8", "delimiter": ","})
    with pytest.raises(UploadError, match="3行目の列数が上限"):
        list(src.rows())


def test_xlsx_with_an_unreadable_value_far_down_is_refused_at_upload(app, client, tmp_path):
    import io
    import re
    import zipfile

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["管理No", "発生日", "現象", "処置", "停止時間"])
    for i in range(300):
        ws.append([f"TR-{i}", "2026-08-01", "停止", "交換", i + 1000])
    wb.save(tmp_path / "ok.xlsx")
    src = zipfile.ZipFile(tmp_path / "ok.xlsx")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb"<v>1250</v>", b"<v>NaN</v>", data, count=1)
            z.writestr(item, data)
    from tests.conftest import upload

    res = upload(client, out.getvalue(), "nan.xlsx")
    assert res.status_code == 400 and "行目付近の値を読めません" in res.get_json()["error"]
    _no_import_left(app)


# ---- BR4-3: AI整形の対象を2列にして保存しない ------------------------------------------------------------

def test_two_ai_columns_are_refused_on_save(app, client):
    from tests.conftest import save_columns, set_role

    import_id, body = _editor_body(app, client)
    set_role(body, "現象", "log")
    set_role(body, "対応内容", "log")
    res = save_columns(client, import_id, body)
    assert res.status_code == 400 and res.get_json()["error"] == "経過の記録の列は1つだけにしてください"


# ---- R4-2: 見出し行のないCSVは、見出しがデータのように見えると知らせる ------------------------------------

def test_headerless_csv_is_warned(tmp_path):
    lines = [f"TR-{i:03d},2026-08-{i:02d},EQ-0{i % 3},山田太郎: 主軸ベアリング損傷のため交換した。取引先に部品を手配し、翌日に復旧を確認した,{i}"
             for i in range(1, 9)]
    src = _csv(tmp_path, "nohead.csv", lines)
    layout = guess_layout(src, src.sheets()[0].name)
    assert any("見出し行がデータのように見えます" in w for w in layout.warnings)


def test_normal_header_is_not_warned(tmp_path):
    lines = ["管理No,発生日,設備,現象,停止時間(分)"] + [f"TR-{i},2026-08-{i:02d},EQ-1,停止,{i}" for i in range(1, 9)]
    src = _csv(tmp_path, "head.csv", lines)
    layout = guess_layout(src, src.sheets()[0].name)
    assert not any("見出し行がデータ" in w for w in layout.warnings)


# ====================================================================================================
# 元 tests/test_tables_fixes5.py
# 一覧表の不具合修正（5巡目）の確認。
# ====================================================================================================

# ---- R5T-1: 大文字・小文字だけ違う設備のファイル名が Windows で上書きし合わない ------------------------------

def test_names_differing_only_in_case_get_a_suffix():
    names = _Names()
    a = names.make(["a", "ETC-302号機"])
    b = names.make(["a", "Etc-302号機"])
    assert a.casefold() != b.casefold()


def test_entities_differing_only_in_case_keep_every_record(tmp_path):
    d = list_spec_dict(group_by="entity_month")
    for c in d["columns"]:
        if c["key"] == "equipment_id":
            c["normalize"] = ["nfkc"]   # 大文字にそろえない設定
    d["columns"] = [c for c in d["columns"] if c["key"] != "equipment_name"]
    spec = spec_from_dict(d)
    records = [_rec("A-1", record_no="A-1", occurred_at="2026-08-03", equipment_id="ETC-302号機"),
               _rec("A-2", record_no="A-2", occurred_at="2026-08-04", equipment_id="Etc-302号機")]
    files = render_all(spec, records, {})
    assert len({f.name.casefold() for f in files}) == len(files)
    target = tmp_path / "imp"
    target.mkdir()
    tables._write_md_dir(target / "md", files)
    written = list((target / "md").glob("*.md"))
    assert len(written) == len(files)
    text = "".join(p.read_text(encoding="utf-8") for p in written)
    assert "【A-1】" in text and "【A-2】" in text


# ---- R5T-2: CSV の「1:30」（[h]:mm を書き出した値）を分として読む ------------------------------------------

# ---- R5-MD-4: 1行目が札と日付だけの現象は、次の行を見出しに使う ----------------------------------------------

@pytest.mark.parametrize("symptom", [
    "発生:2024-04-28 14:50(休日)\n設備:CMP-103\n現象:MESとの通信が断続的に切断",
    "【発生】R06.04.28 14:50\n【現象】MESとの通信が断続的に切断",
])
def test_title_uses_the_next_line_when_the_first_is_only_a_date(symptom):
    spec = spec_from_dict(list_spec_dict())
    title = record_title({"record_no": "CA-1", "occurred_at": "2024-05-01", "equipment_id": "CMP-103",
                          "symptom": symptom}, spec)
    assert "MESとの通信が断続的に切断" in title and "設備:" not in title


def test_title_of_a_plain_symptom_is_unchanged():
    spec = spec_from_dict(list_spec_dict())
    title = record_title({"record_no": "CA-1", "occurred_at": "2024-05-01", "equipment_id": "CMP-103",
                          "symptom": "MES通信断\n詳細は別紙"}, spec)
    assert "MES通信断" in title and "詳細" not in title


# ---- SEC5-3: 選択（|）を含むグループに量指定子が付く区切りの正規表現を断る -------------------------------------

def test_overlapping_alternation_pattern_is_refused():

    assert any("遅く" in e for e in validate_spec(_log_spec(splitter={"extra_anchors": [r"(?:\d|\d)*年"]})))
    assert any("遅く" in e for e in validate_spec(_log_spec(splitter={"not_date_patterns": [r"(a|ab)+c"]})))
    assert validate_spec(_log_spec(splitter={"not_date_patterns": [r"\d+\.\d+\s*(?:mm|MPa)"]})) == []




# ---- R5-FUZZ-1: 遠くのセル1つ（XFD1 / XFD1048576）で読み込みが止まらない ---------------------------------------

def _far_book(path: Path, far: str, rows: int) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["管理No", "発生日", "設備", "現象", "処置"])
    for i in range(rows):
        ws.append([f"A-{i}", f"2026/08/{i % 28 + 1:02d}", "CMP-1", f"現象{i}", "処置"])
    ws[far] = "メモ"
    wb.save(path)
    return path


def test_far_stray_header_cell_does_not_widen_the_table(tmp_path):
    from app.tables import guess_layout
    from app.tables import open_source

    path = _far_book(tmp_path / "far.xlsx", "XFD1", 3000)
    started = time.monotonic()
    source = open_source(path, path.name)
    assert sum(1 for _ in source.rows("S")) == 3001
    layout = guess_layout(source, "S")
    assert time.monotonic() - started < 10
    assert len(layout.headers) == 5 and any("XFD" in w for w in layout.warnings)


def test_far_last_cell_in_a_near_column_does_not_pad_every_row(tmp_path):
    from app.tables import open_source

    # 256列目（PAD_MAX_COLUMNS 以下）の遠いセル。行×列で埋めると100秒近くかかっていた
    path = _far_book(tmp_path / "far3.xlsx", "IV1048576", 20)
    started = time.monotonic()
    source = open_source(path, path.name)
    widths = {len(row.cells) for row in source.rows("S")}
    assert time.monotonic() - started < 15
    assert widths == {0, 5, 256}


def test_far_last_cell_upload_finishes_and_reads_the_table(app, client, tmp_path):
    from tests.conftest import upload

    path = _far_book(tmp_path / "far2.xlsx", "XFD1048576", 50)
    started = time.monotonic()
    res = upload(client, path.read_bytes(), "far2.xlsx")
    assert time.monotonic() - started < 30
    assert res.status_code == 200 and res.get_json()["import_id"]


# ---- R5C-1: 確定の処理の途中で渡し終えて消えた取り込みのフォルダを作り直さない ---------------------------------

def test_render_does_not_recreate_a_purged_import(app, monkeypatch):
    from app import core

    with app.app_context():
        _spec, import_id = _new_import(app)
        tables.run_read(FakeCtx(), import_id)
        original = tables.load_rows

        def load_then_purge(*a, **k):
            rows = original(*a, **k)
            core.purge_table_import(import_id)   # 別のタブの保存・ダウンロードが渡し終えた
            return rows

        monkeypatch.setattr(tables, "load_rows", load_then_purge)
        with pytest.raises(tables.PipelineError):
            tables.run_render(FakeCtx(), import_id)
        assert not tables.import_dir(import_id).exists()


def test_write_md_dir_does_not_create_a_missing_import_folder(tmp_path):
    with pytest.raises(OSError):
        tables._write_md_dir(tmp_path / "gone" / "md", [])
    assert not (tmp_path / "gone").exists()


# ====================================================================================================
# 元 tests/test_tables_fixes6.py
# 一覧表の不具合修正（6巡目）の確認。
# ====================================================================================================

# ---- R6T-1: 期間の日付の列は、画面で選んだ「日付」の役割から決める -------------------------------------------

def test_the_period_date_column_follows_the_date_role(app, client):
    from tests.conftest import SPEC_CSV, set_role

    # 完了日の列もある CSV（日付の列が2つある表）
    text = "\r\n".join(line + ("完了日" if i == 0 else "2026-08-28") for i, line in
                       enumerate(x + "," for x in SPEC_CSV.strip("\r\n").split("\r\n"))) + "\r\n"
    import_id, body = _editor_body(app, client, text=text)
    after = _saved(app, client, import_id, body)
    assert after.period["date_column"] == "occurred_at"

    # 最初の日付の列を「その他」にすると、次の日付の列が期間の列になる
    set_role(body, "発生日", "attribute")
    set_role(body, "完了日", "date")
    after = _saved(app, client, import_id, body)
    assert after.column(after.period["date_column"]).display == "完了日"


# ---- R6T-2: 80行目より下の見出し行を手で指定しても見出しが読める ------------------------------------------------

def _deep_book(path: Path, header_row: int, sheet: str = "一覧") -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.cell(1, 1, "トラブル対応一覧（前置き）")
    ws.cell(header_row, 1, "管理No")
    ws.cell(header_row, 2, "発生日")
    ws.cell(header_row, 3, "現象")
    for i in range(1, 40):
        r = header_row + i
        ws.cell(r, 1, f"TR-{i:03d}")
        ws.cell(r, 2, f"2026/08/{(i % 28) + 1:02d}")
        ws.cell(r, 3, f"搬送ロボット{i}号機が停止")
    wb.save(path)
    return path


def test_header_row_below_the_head_rows_is_read(tmp_path):
    from app.tables import guess_layout
    from app.tables import ExcelSource

    src = ExcelSource(_deep_book(tmp_path / "deep90.xlsx", 91))
    layout = guess_layout(src, "一覧", header_rows=[91])
    assert layout.header_rows == [91] and layout.data_start == 92 and layout.data_end == 130
    assert layout.headers[:3] == ["管理No", "発生日", "現象"]
    single = guess_layout(src, "一覧", header_row=91)
    assert single.headers[:3] == ["管理No", "発生日", "現象"]


# ---- R6T-3: 見出しが自動で見つからなくても、指定した見出し行で列の対応づけができる --------------------------------

def test_the_columns_panel_uses_the_saved_header_row(app, client, tmp_path):
    from tests.conftest import panel_html, save_layout, upload_bytes

    path = _deep_book(tmp_path / "deep40.xlsx", 41)
    import_id = upload_bytes(client, path.read_bytes(), "deep40.xlsx")
    assert save_layout(client, import_id, header_rows="41").status_code == 200
    page = panel_html(client, import_id, "columns")
    assert "管理No" in page and "現象" in page   # 保存した見出し行（41行目）で見出しを読んでいる
    # 読み取り方の段には取り込み設定の選択を出さない（設定は保存しない）
    assert 'name="template"' not in panel_html(client, import_id, "source")


# ---- R6T-4: 見出し行の入力は画面の下見とサーバーで同じように読む ---------------------------------------------------

_ROW_INPUTS = ["６,７", "3-4", "1，2", "1、2 3", "3a", "5-3", "1-20", "", "0,2", "１２０", "+5", "2-2"]


def test_row_inputs_are_read_the_same_way_on_the_server():
    from app.views import _int_list, _row_no

    assert _int_list("６,７") == [6, 7]
    assert _int_list("3-4") == [3, 4]
    assert _int_list("1，2") == [1, 2]
    assert _int_list("3a") == [] and _int_list("5-3") == [] and _int_list("1-20") == []
    assert _int_list([3, 1]) == [1, 3]
    assert _row_no("１２０") == 120 and _row_no("12x") is None and _row_no("") is None and _row_no(None) is None


@pytest.mark.skipif(shutil.which("node") is None, reason="node がない")
def test_row_inputs_are_read_the_same_way_in_the_browser():
    from app.views import _int_list, _row_no

    js = _tables_js()
    start = js.index("  const parseEnd = ")
    end = js.index("  function applyLayout(")
    script = js[start:end] + f"""
const inputs = {json.dumps(_ROW_INPUTS, ensure_ascii=False)};
process.stdout.write(JSON.stringify({{rows: inputs.map(parseRows), ends: inputs.map(parseEnd)}}));
"""
    out = subprocess.run(["node", "-e", script], capture_output=True, check=True, timeout=30, encoding="utf-8")
    got = json.loads(out.stdout)
    assert got["rows"] == [_int_list(x) for x in _ROW_INPUTS]
    assert got["ends"] == [_row_no(x) for x in _ROW_INPUTS]


# ---- R6T-5: 既定の候補で「経過の記録」（role=log）は1列だけ -----------------------------------------------------

def test_default_suggestions_tick_only_one_log_column():
    from app.tables import suggest_columns

    log = "4/1 10:00 停止を確認。\n4/2 11:00 センサーを交換。\n4/3 復旧を確認した。"
    headers = ["管理No", "対応内容", "対応内容_2", "対応内容_3"]
    rows = [[f"TR-{i}", log, log + "追記", log + "再発"] for i in range(10)]
    sugg = suggest_columns(headers, rows)
    logs = [s for s in sugg if s.role == "log"]
    assert [s.header for s in logs] == ["対応内容"]
    others = [s for s in sugg if s.header in ("対応内容_2", "対応内容_3")]
    assert all(s.role == "text" and not s.log and s.md in ("body", "omit") for s in others)


def test_untouched_columns_panel_with_two_log_like_columns_can_be_saved(app, client):
    from tests.conftest import editor_body, csv_source, save_columns, save_layout, upload_csv

    log = "4/1 10:00 停止を確認。\n4/2 11:00 センサーを交換。\n4/3 復旧を確認した。"
    text = "管理No,発生日,対応内容,対応内容_2\r\n" + "".join(
        f"TR-{i},2026/08/{i + 1:02d},\"{log}\",\"{log}追記\"\r\n" for i in range(10))
    import_id = upload_csv(client, "log2.csv", text, encoding="utf-8")
    csv_source(client, import_id, encoding="utf-8")
    assert save_layout(client, import_id).status_code == 200
    body = editor_body(client, import_id)
    assert sum(1 for r in body["columns"] if r["role"] == "log") == 1
    res = save_columns(client, import_id, body)
    assert res.status_code == 200, res.get_json()


# ---- R6T-6: JSON の AI整形の上限・実行条件が数値でなければ保存の前に断る -------------------------------------------

def _log_spec_tables_fixes6(**stage):
    return spec_from_dict({"name": "x", "columns": [{"key": "log", "display": "ログ", "type": "text", "role": "log"}],
                           "log_stage": {"column": "log", **stage}})


@pytest.mark.parametrize("stage", [
    {"limits": {"max_segments": "多い"}},
    {"limits": {"max_input_tokens": 0}},
    {"limits": {"max_segments": True}},
    {"run_if": {"min_chars": "x"}},
    {"run_if": {"any": [{"min_segments": 2}, {"min_chars": "x"}]}},
    {"run_if": {"all": "min_chars"}},
])
def test_non_numeric_limits_and_run_if_are_refused(stage):
    assert validate_spec(_log_spec_tables_fixes6(**stage))


@pytest.mark.parametrize("stage", [
    {"limits": {"max_segments": 30, "max_input_tokens": "4000"}},
    {"run_if": {"any": [{"min_segments": 2}, {"min_chars": 80}], "contains": "停止"}},
    {"run_if": {}},
])
def test_numeric_limits_and_run_if_are_accepted(stage):
    assert validate_spec(_log_spec_tables_fixes6(**stage)) == []


# ---- R6-MD-1: 番号だけの行が先にあっても、集計の表示名は同じ番号の行の名前から付ける ---------------------------

def _code_entity_spec():
    return spec_from_dict({"name": "T", "columns": [
        {"key": "record_no", "display": "管理No", "type": "code", "role": "key"},
        {"key": "occurred_at", "display": "発生日", "type": "date", "role": "date"},
        {"key": "equipment", "display": "設備", "type": "code", "role": "entity"},
    ], "period": {"date_column": "occurred_at"}})


# ---- R6-MD-4: 「0:28以降、…」の時刻は見出しに残す ------------------------------------------------------------------

@pytest.mark.parametrize("text, expected", [
    ("0:28以降、装置がオフライン表示。", "0:28以降、装置がオフライン表示。"),
    ("12:30から停止", "12:30から停止"),
    ("12:30以降、装置がオフライン表示", "12:30以降、装置がオフライン表示"),
    ("R05.04.01 11:45(休日)、搬送停止", "搬送停止"),
    ("0:28、搬送停止", "搬送停止"),
])
def test_title_keeps_a_time_that_is_part_of_the_sentence(text, expected):
    from app.tables import _title_text_line

    assert _title_text_line(text) == expected


# ---- R6-1: 渡し終えた取り込みに、読み込み直しが行データ・問題一覧を書き戻さない ------------------------------------

def test_reread_does_not_write_back_into_a_purged_import(app, monkeypatch):
    from app import core

    with app.app_context():
        __spec, import_id = _new_import(app)
        real_checks = tables.run_checks

        def purge_midway(records, spec, stats):
            core.purge_table_import(import_id)   # 読み込みの途中で渡し終えて消えた
            return real_checks(records, spec, stats)

        monkeypatch.setattr(tables, "run_checks", purge_midway)
        with pytest.raises(tables.PipelineError):
            tables.run_read(FakeCtx(), import_id)
        assert tables.get_import(import_id) is None
        assert not tables.import_dir(import_id).exists()


def test_write_helpers_do_not_create_a_missing_folder(tmp_path):
    with pytest.raises(OSError):
        tables.write_atomic(tmp_path / "gone" / "issues.csv", b"x")
    with pytest.raises(OSError):
        tables.write_rows(tmp_path / "gone" / "rows.jsonl.gz", [])
    assert not (tmp_path / "gone").exists()


# ---- R6-SEC-1: 全角の ＝ ＋ － ＠ で始まるシート名・ファイル名も式にしない --------------------------------------------

def test_full_width_formula_prefixes_are_guarded():
    from app.tables import Issue
    from app.tables import guard_formula, issues_csv

    for s in ("＝SUM(1,1)", "＋1", "－1", "＠SUM(1,1)"):
        assert guard_formula(s) == "'" + s
    assert guard_formula("設備") == "設備"
    rows = list(csv.reader(io.StringIO(issues_csv([Issue("warning", "x", "＝SUM(1,1)", row=2, column="＋1")])
                                       .decode("utf-8-sig"))))
    assert rows[1][2:] == ["2", "'＋1", "'＝SUM(1,1)"]




# ---- FUZZ: 手で作った Excel（大きな非表示列・上限を超える行・列が多すぎる表） ---------------------------------------

def _rewrite_sheet(path: Path, fn) -> None:
    buf = io.BytesIO()
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                data = fn(data)
            dst.writestr(info, data)
    path.write_bytes(buf.getvalue())


def _small_book_tables_fixes6(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "一覧"
    ws.append(["管理No", "発生日", "現象"])
    for i in range(5):
        ws.append([f"TR-{i}", "2026/08/01", "停止"])
    wb.save(path)
    return path


def test_huge_hidden_column_span_is_clamped(tmp_path):
    from app.tables import ExcelSource

    path = _small_book_tables_fixes6(tmp_path / "cols.xlsx")
    _rewrite_sheet(path, lambda d: d.replace(
        b"<sheetData>", b'<cols><col min="1" max="20000000" hidden="1" width="5"/></cols><sheetData>', 1))
    t = time.monotonic()
    info = ExcelSource(path).sheets()[0]
    assert time.monotonic() - t < 5
    assert len(info.hidden_columns) <= 16384


def test_cell_beyond_the_excel_row_limit_is_refused_quickly(tmp_path):
    from app.tables import ExcelSource
    from app.tables import UploadError

    path = _small_book_tables_fixes6(tmp_path / "far.xlsx")
    _rewrite_sheet(path, lambda d: d.replace(
        b"</sheetData>",
        b'<row r="4294967296"><c r="D4294967296" t="inlineStr"><is><t>z</t></is></c></row></sheetData>', 1))
    t = time.monotonic()
    with pytest.raises(UploadError, match="Excel の上限"):
        ExcelSource(path)
    assert time.monotonic() - t < 5


def test_xlsx_with_too_many_columns_is_refused(tmp_path):
    from app.tables import MAX_COLUMNS
    from app.tables import ExcelSource
    from app.tables import UploadError

    wb = Workbook()
    ws = wb.active
    ws.title = "一覧"
    ws.append([f"項目{i}" for i in range(MAX_COLUMNS + 1)])
    ws.append([f"v{i}" for i in range(MAX_COLUMNS + 1)])
    path = tmp_path / "wide.xlsx"
    wb.save(path)
    src = ExcelSource(path)
    with pytest.raises(UploadError, match="列数が上限"):
        list(src.rows("一覧", 1, 2))


def test_far_memo_cell_is_not_counted_as_many_columns(tmp_path):
    from app.tables import ExcelSource

    path = _small_book_tables_fixes6(tmp_path / "memo.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "一覧"
    ws.append(["管理No", "発生日", "現象"])
    ws.append(["TR-1", "2026/08/01", "停止"])
    ws["XFD1"] = "メモ"
    wb.save(path)
    rows = list(ExcelSource(path).rows("一覧", 1, 2))
    assert len(rows) == 2


# ---- R6C-1: 確定した取り込みでは試し実行を断る（zip・保存と下書きが食い違わないように） ---------------------------------

def test_trial_is_refused_after_confirm(app, client):
    with app.app_context():
        __spec, import_id = _new_import(app)
        tables.run_read(FakeCtx(), import_id)
        tables.run_render(FakeCtx(), import_id)
        assert tables.get_import(import_id)["status"] == "confirmed"
    res = client.post(f"/tables/imports/{import_id}/ai/trial", json={"row_key": "TR-001"})
    assert res.status_code == 400
    assert "確定後は試し実行できません" in res.get_json()["error"]


# ---- R6C-2: 渡している間に別のタブで読み込み直しが始まったら、「ダウンロード済み」とは言わない ----------------------------

@pytest.mark.parametrize("method, url, builder", [
    ("get", "download.zip", "build_download"),
])
def test_reread_during_hand_out_goes_back_to_the_preview(app, client, monkeypatch, tmp_path, method, url, builder):
    with app.app_context():
        __spec, import_id = _new_import(app)
        tables.run_read(FakeCtx(), import_id)
        tables.run_render(FakeCtx(), import_id)

    def reread_started(*_a, **_k):
        tables.update_import(import_id, status="preview")   # 別のタブの読み込み直しが md を消した
        raise FileNotFoundError("md")

    monkeypatch.setattr(views, builder, reread_started)
    res = getattr(client, method)(f"/tables/imports/{import_id}/{url}")
    # 404「ダウンロード済み」ではなく、画面に戻して読み込み直しが始まったことを知らせる
    assert res.status_code == 302 and res.headers["Location"].endswith("/tables/new")
    with client.session_transaction() as session:
        assert any("読み込み直しが始まった" in text for _level, text in session["_flashes"])
    with app.app_context():
        assert tables.get_import(import_id) is not None


# ---- R6B-2: md のパスが Windows の上限を超えるときは分かる言葉で止める --------------------------------------------------

def test_too_long_md_path_gives_a_clear_message(tmp_path, monkeypatch):
    from app import core as core
    from app.tables import MdFile

    monkeypatch.setattr(core, "_long_paths_enabled", lambda: False)
    target = tmp_path / "md"
    tmp_path.mkdir(exist_ok=True)
    name = "あ" * (core.MAX_PATH_CHARS - len(str(target)) + 10) + ".md"
    files = [MdFile(name=name, text="x", kind="records")]
    with pytest.raises(tables.PipelineError, match="ファイル名が長すぎ"):
        tables._write_md_dir(target, files)
    assert not target.exists()


# ---- R6-T-01: 記録番号・担当者の列の「〃」「同上」も直前の行の値で補う ------------------------------------------

def test_ditto_in_the_record_number_and_person_columns_is_filled(tmp_path):
    from app.tables import render_all

    lines = ["管理No,発生日,設備ID,現象,対応者,停止時間(分)",
             "TR-001,2025-03-01,CMP-101,異音,田中,30",
             "〃,〃,〃,追加調査,〃,20",
             "〃,〃,〃,部品交換,同上,10",
             "TR-002,2025-03-02,ETC-301,停止,佐藤,40",
             "〃,2025-03-02,ETC-301,復旧,佐藤,5"]
    src = _csv(tmp_path, "ditto_key.csv", lines)
    _layout, spec, records, stats = _auto_tables_fixes3(src, src.sheets()[0].name)
    assert {c.key: c.role for c in spec.columns}["record_no"] == "key"
    # 2行目以降は直前の伝票番号で補われ、同じ伝票の続きとして #2・#3 になる
    assert [r.key for r in records] == ["TR-001", "TR-001#2", "TR-001#3", "TR-002", "TR-002#2"]
    assert stats.ditto_filled.get("record_no") == 3 and stats.ditto_filled.get("worker") == 2
    assert "〃" not in stats.duplicate_keys
    text = "\n".join(f.text for f in render_all(spec, [r.to_dict() for r in records], {}))
    assert "〃" not in text and "同上" not in text


# ---- R6-T-02: 記録キーの「列:文字数」を認める／記録キーの代わりの列も確かめる ------------------------------------

def _key_spec(**record):
    return spec_from_dict({"name": "x", "columns": [
        {"key": "occurred_at", "display": "発生日", "type": "date", "role": "date"},
        {"key": "equipment_id", "display": "設備番号", "type": "code", "role": "entity"},
        {"key": "symptom", "display": "現象", "type": "text", "role": "text"},
    ], "record": record})


def test_record_key_accepts_a_truncated_column():
    assert validate_spec(_key_spec(key=["occurred_at", "equipment_id", "symptom:20"])) == []
    assert validate_spec(_key_spec(key=["symptom:20"], fallback_key=["occurred_at"])) == []


def test_record_key_with_a_missing_column_or_a_bad_length_is_refused():
    assert validate_spec(_key_spec(key=["missing"])) == ["記録キーの列「missing」がありません"]
    assert validate_spec(_key_spec(key=["missing:20"])) == ["記録キーの列「missing」がありません"]
    errors = validate_spec(_key_spec(key=["symptom:あ"]))
    assert errors and "記録キーの「symptom:あ」の文字数" in errors[0]


def test_fallback_key_columns_are_checked_when_they_are_written_by_hand():
    errors = validate_spec(_key_spec(key=["equipment_id"], fallback_key=["equipment_no_typo"]))
    assert errors == ["記録キーの代わりの列「equipment_no_typo」がありません"]
    # 既定のままの代わりのキーは、その表に無い列名でも問題にしない（既定は一般的な列名）
    assert validate_spec(_key_spec(key=["equipment_id"])) == []


# ---- R6-T-03: 列の範囲は指定できないので、Excel 側で直す案内にする --------------------------------------------

def test_right_hand_table_warning_tells_what_can_actually_be_done(tmp_path):
    from app.tables import guess_layout

    lines = ["管理No,発生日,現象,,設備No,点検日,結果"]
    for i in range(10):
        lines.append(f"TR-{i:03d},2025-03-{i + 1:02d},停止,,EQ-{i:03d},2025-03-{i + 1:02d},良")
    src = _csv(tmp_path, "two_tables.csv", lines)
    layout = guess_layout(src, src.sheets()[0].name)
    warning = next(w for w in layout.warnings if "別の表があるようです" in w)
    assert "別のシートか別のファイルに分けて" in warning and "範囲を指定して取り込んで" not in warning
    assert layout.headers == ["管理No", "発生日", "現象"]


def test_far_memo_column_warning_tells_what_can_actually_be_done(tmp_path):
    from app.tables import guess_layout

    blanks = "," * 22
    lines = [f"管理No,発生日,現象{blanks},メモ"] + [f"TR-{i:03d},2025-03-{i + 1:02d},停止{blanks}," for i in range(10)]
    src = _csv(tmp_path, "far_memo.csv", lines)
    layout = guess_layout(src, src.sheets()[0].name)
    warning = next(w for w in layout.warnings if "大きく離れた" in w)
    assert "表の右隣に移してから" in warning and "範囲を指定して取り込んで" not in warning


# ---- R6-T-04: 判定に使っていない項目は取り込み設定に持たない ----------------------------------------------------

def test_settings_that_never_changed_the_reading_are_dropped():
    from app.tables import spec_to_dict

    col = {"key": "a", "display": "a", "type": "string", "role": "attribute"}
    spec = spec_from_dict({"name": "x", "columns": [col],
                           "header": {"rows": 1, "anchors": ["管理No"], "search_rows": 50},
                           "data_end": {"blank_rows": 9, "stop_first_col": ["おわり"]},
                           "exclude": {"aggregate_keywords": ["計"], "hidden_rows": "include"}})
    assert spec.header == {"rows": 1, "anchors": ["管理No"]}
    assert not hasattr(spec, "data_end")
    assert spec.exclude == {"hidden_rows": "include", "strike_rows": "exclude_with_warning"}
    d = spec_to_dict(spec)
    assert "data_end" not in d and "search_rows" not in d["header"] and "aggregate_keywords" not in d["exclude"]
    assert validate_spec(spec) == []


# ---- ux6-2: データの行が無いときは、見出し行ではなくデータの範囲のことを言う ----------------------------------------

def _csv_import(client, name: str, text: str) -> int:
    from tests.conftest import csv_source, upload_csv

    import_id = upload_csv(client, name, text, encoding="utf-8")
    csv_source(client, import_id, encoding="utf-8")
    return import_id


def test_a_table_with_no_data_rows_does_not_blame_the_header_row(client):
    from tests.conftest import save_layout

    import_id = _csv_import(client, "empty.csv", "管理No,発生日,設備番号,現象,対応内容\r\n")
    res = save_layout(client, import_id)
    assert res.status_code == 400
    error = res.get_json()["error"]
    assert "データの行が1行もありません" in error
    assert "見出し行の番号を指定してください" not in error


def test_a_data_end_above_the_header_is_echoed_back_with_its_own_message(client):
    from tests.conftest import panel_html, save_layout

    text = "管理No,発生日,現象\r\n" + "".join(f"TR-{i},2026/08/{i + 1:02d},停止\r\n" for i in range(10))
    import_id = _csv_import(client, "rows.csv", text)
    res = save_layout(client, import_id, data_end_row="1")
    assert res.status_code == 400
    error = res.get_json()["error"]
    assert "データの終わりに1行目を指定すると、データの行が1行もありません" in error
    assert "見出し行の番号を指定してください" not in error
    # 断られた入力は保存しないので、段を開き直しても自動判定のまま（画面は入力した値をそのまま残す）
    assert 'name="data_end_row" class="input" value=""' in panel_html(client, import_id, "layout")


# ---- r6-c1: 成功して終わったばかりの読み込みを［中止］が巻き戻さない --------------------------------------------

def test_cancel_does_not_undo_a_read_that_just_finished(app, client, monkeypatch):
    from app import core
    from app import database

    with app.app_context():
        __spec, import_id = _new_import(app)
        tables.run_read(FakeCtx(), import_id)
        db = database.get_db()
        ts = database.now()
        cur = db.execute(
            "INSERT INTO jobs (kind, ref_type, ref_id, status, params_json, progress_json, created_at, updated_at) "
            "VALUES ('table_read', 'table_import', ?, 'running', '{}', '{}', ?, ?)", (import_id, ts, ts))
        job_id = cur.lastrowid
        db.commit()
        tables.update_import(import_id, status="reading", job_id=job_id)   # 画面が読んだときは「読み込み中」

    def cancel_after_the_job_finished(jid):
        # 中止を要求した瞬間に、ジョブ本体は成功して status=preview を書き終えていた
        with app.app_context():
            tables.update_import(import_id, status="preview")
            db = database.get_db()
            db.execute("UPDATE jobs SET status = 'cancelled' WHERE id = ?", (jid,))
            db.commit()
        return True

    monkeypatch.setattr(views, "request_cancel", cancel_after_the_job_finished)
    res = client.post(f"/tables/imports/{import_id}/cancel")
    assert res.status_code == 200 and res.get_json()["ok"] is True
    with app.app_context():
        assert tables.get_import(import_id)["status"] == "preview"


# ---- ux6-6: ［取り込みを削除］の確認文は取り込みのどの画面でも同じ ------------------------------------------------

def test_the_delete_confirm_text_is_the_same_on_every_table_screen():
    root = Path(__file__).resolve().parents[1] / "app" / "templates"
    shared = "この取り込みを削除します。読み込んだ内容と作成した Markdown も消えます。元に戻せません。"
    found = [line.strip() for path in root.rglob("*.html")
             for line in path.read_text(encoding="utf-8").splitlines() if "この取り込みを削除します" in line]
    assert found  # 文言は共通のマクロ（と done.html の「ダウンロードせずに」）だけ
    for line in found:
        assert shared in line and "アップロードしたファイル" not in line


# ---- R6T-6: 見出しも値も無い列は「列の対応づけ」に出さない ------------------------------------------------------

_EMPTY_COL_CSV = "\r\n".join(
    [",,管理No,発生日,,対応内容,,設備名,,"]
    + [f",,TR-{i:03d},2026-08-{i:02d},補足{i},8/{i} 確認した。,,搬送ロボット{i % 2 + 1}号機,"
       + (f"後から入る{i}" if i >= 6 else "") + "," for i in range(1, 13)]) + "\r\n"


def _columns_of(tmp_path, text: str, name: str = "空列.csv"):
    """CSV を読み、列の対応づけに出る候補を返す。"""
    from app.tables import guess_layout, sample_data_rows
    from app.tables import suggest_columns
    from app.tables import open_source

    path = tmp_path / name
    path.write_bytes(text.encode("cp932"))
    source = open_source(path, name, {"encoding": "cp932", "delimiter": ","})
    layout = guess_layout(source, name)
    return layout, suggest_columns(layout.headers, sample_data_rows(source, name, layout))


def test_columns_empty_in_header_and_values_are_not_listed(tmp_path):
    """左端・途中・右端の空の列は一覧に出さない。見出しだけ空で値がある列は残す。"""
    layout, sugg = _columns_of(tmp_path, _EMPTY_COL_CSV)
    # 表の幅は今までどおり（見出しの検出は変えていない）。仮の名前「列N」も今までどおり付く
    assert layout.header_rows == [1] and layout.data_start == 2
    assert layout.headers == ["列1", "列2", "管理No", "発生日", "列5", "対応内容", "列7", "設備名", "列9"]
    # 列1・列2（左端）と列7（途中）は消える。列5（見出しだけ空）と列9（あとから値が入る）は残る
    assert [s.header for s in sugg] == ["管理No", "発生日", "列5", "対応内容", "設備名", "列9"]
    assert [s.index for s in sugg] == [2, 3, 4, 5, 7, 8]  # 列の位置は元のまま（キー col3 などが変わらない）


def test_a_column_with_values_only_in_later_rows_stays_in_the_list(tmp_path):
    """先頭の行だけ空で、あとから値が入る列は残す（空欄だけの列として消さない）。"""
    _layout, sugg = _columns_of(tmp_path, _EMPTY_COL_CSV, "後から.csv")
    later = next(s for s in sugg if s.header == "列9")
    assert later.examples and later.blank_rate < 1.0


def test_an_excel_table_that_starts_at_c3_lists_only_its_own_columns(tmp_path):
    """表が C3 から始まる Excel でも、A・B列は一覧に出さない（Excel も CSV と同じ扱い）。"""
    from app.tables import guess_layout, sample_data_rows
    from app.tables import suggest_columns
    from app.tables import open_source

    wb = Workbook()
    ws = wb.active
    for col, head in enumerate(["カルテNo", "受付日", "設備名", "対応内容"], start=3):
        ws.cell(row=3, column=col, value=head)
    for i in range(1, 13):
        ws.cell(row=3 + i, column=3, value=f"K-{i:03d}")
        ws.cell(row=3 + i, column=4, value=f"2026-08-{i:02d}")
        ws.cell(row=3 + i, column=5, value=f"搬送ロボット{i % 2 + 1}号機")
        ws.cell(row=3 + i, column=6, value=f"8/{i} 確認した。")
    path = tmp_path / "C3から.xlsx"
    wb.save(path)
    source = open_source(path, path.name)
    layout = guess_layout(source, ws.title)
    assert layout.headers[:2] == ["列1", "列2"] and layout.data_start == 4
    sugg = suggest_columns(layout.headers, sample_data_rows(source, ws.title, layout))
    assert [s.header for s in sugg] == ["カルテNo", "受付日", "設備名", "対応内容"]


@pytest.mark.samples
def test_sample_t6_lists_20_columns_instead_of_22():
    """T6（表が C3 から始まる）の列の一覧は 22 列ではなく 20 列。"""
    from app.tables import guess_layout, sample_data_rows
    from app.tables import suggest_columns
    from app.tables import open_source

    path = Path(__file__).resolve().parents[1] / "samples" / "tables" / "T6_装置トラブルカルテ.xlsx"
    if not path.exists():
        pytest.skip(f"{path.name} がありません")
    source = open_source(path, path.name)
    sheet = source.sheets()[0].name
    layout = guess_layout(source, sheet)
    assert len(layout.headers) == 22 and layout.headers[:2] == ["列1", "列2"]
    sugg = suggest_columns(layout.headers, sample_data_rows(source, sheet, layout))
    assert len(sugg) == 20 and sugg[0].header == "カルテNo"


# ====================================================================================================
# 元 tests/test_tables_fixes7.py
# 一覧表で渡すファイルを「RAG に入れる Markdown だけ」にした変更（7巡目）の確認。
#
# 利用者の判断（2026-09-20）: 件数・順位・推移のような集計は RAG の仕組みに向いていないので作らない。
# データセット説明も作らない。管理用の CSV も渡さない。
# ====================================================================================================

# ---- R7-1: 作る md は記録ファイルだけ ---------------------------------------------------------

def test_only_record_files_are_made():
    """集計（月次・設備別年度）とデータセット説明は作らない。"""
    spec = spec_from_dict(list_spec_dict())
    records = [_rec("A-1", record_no="A-1", occurred_at="2026-08-03", equipment_id="CMP-101", downtime=30),
               _rec("A-2", record_no="A-2", occurred_at="2026-09-04", equipment_id="CVD-201", downtime=60)]
    files = render_all(spec, records, {})
    assert [f.name for f in files] == ["トラブル対応一覧_2026-08.md", "トラブル対応一覧_2026-09.md"]
    assert {f.kind for f in files} == {"records"}
    text = "\n".join(f.text for f in files)
    for word in ("集計", "データセット説明", "合計は", "件です。", "上位"):
        assert word not in text, word


def test_saved_settings_for_the_summaries_are_read_and_dropped():
    """前の版の取り込み設定（dataset_card・summaries）が残っていても、読み飛ばして記録ファイルだけを作る。"""
    d = list_spec_dict(dataset_card=True, summaries=[{"id": "month", "metrics": ["count"]}])
    spec = spec_from_dict(d)
    assert "dataset_card" not in spec.markdown and "summaries" not in spec.markdown
    assert validate_spec(spec) == []
    records = [_rec("A-1", record_no="A-1", occurred_at="2026-08-03", equipment_id="CMP-101")]
    assert [f.name for f in render_all(spec, records, {})] == ["トラブル対応一覧_2026-08.md"]


# ---- R7-2: zip は md だけ（フォルダ分けも管理用CSVも無い） ---------------------------------------

def test_the_zip_is_flat_markdown_only(app, client):
    import_id = confirmed(app, client, "トラブル一覧.csv", "トラブル対応一覧")
    res = client.get(f"/tables/imports/{import_id}/download.zip")
    assert res.status_code == 200
    names = zipfile.ZipFile(io.BytesIO(res.data)).namelist()
    assert names and all(n.endswith(".md") and "/" not in n for n in names)


def test_the_screens_do_not_offer_a_csv_download(app, client):
    """確認・確定の段に正規化CSVのボタンを出さない（渡すのは zip だけ）。"""
    import_id = confirmed(app, client, "トラブル一覧.csv", "トラブル対応一覧")
    for name in ("preview", "done"):
        html = panel_html(client, import_id, name)
        assert "正規化CSV" not in html and "normalized.csv" not in html
    assert "RAG に入れる Markdown（.md）だけ" in panel_html(client, import_id, "done")
    assert client.get(f"/tables/imports/{import_id}/normalized.csv").status_code == 404


# ====================================================================================================
# 元 tests/test_tables_fixes8.py
# 表の取り込みの直し（記録0件のまま確定できる・日付の列が選べない・見出しと並び順）。
#
# このまわりで見つかった不具合:
#   - 記録0件でも「確定」でき、④のダウンロードを押すと空の画面に戻されて取り込みが消える
#   - 「日付の列を1つ選んでください」と言うのに、年月だけの列を選ぶと保存が400で拒否され直せない
#   - 識別番号・日付・設備・長文のどれも無い表は、全記録の見出しが「（見出しなし）」になる
# ====================================================================================================

def _struck_book(path):
    """見出し1行＋データ5行。データ行は全部取り消し線（＝取り込める行が1件も無い表）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "一覧"
    for col, header in enumerate(["管理No", "発生日", "設備番号", "現象"], start=1):
        ws.cell(row=1, column=col, value=header).font = Font(bold=True)
    for i in range(1, 6):
        for col, value in enumerate([f"TR-{i:03d}", f"2026-08-{i:02d}", "EQ-01", f"アラーム{i}"], start=1):
            ws.cell(row=1 + i, column=col, value=value).font = Font(strike=True)
    wb.save(path)
    return path


# 「実施月」が年月だけ（2025-03）なので、日付として読めない表
MONTH_CSV = "管理No,実施月,設備番号,現象,対応内容\r\n" + "".join(
    f"K-{i:03d},2025-0{i},EQ-0{i},現象{i},対応{i}\r\n" for i in range(1, 6))

def _month_columns(name: str, date_role: str = "attribute") -> dict:
    roles = ["key", date_role, "attribute", "attribute", "attribute"]
    return {"name": name, "columns": [{"index": i, "use": True, "role": r} for i, r in enumerate(roles)]}


def _prepare(client, csv_name: str, text: str) -> int:
    import_id = upload_csv(client, csv_name, text)
    csv_source(client, import_id)
    assert save_layout(client, import_id).status_code == 200
    return import_id


def _struck_import(app, client, tmp_path, name: str) -> int:
    """全部の行が取り消し線の表を、④の保存まで通した取り込み。"""
    path = _struck_book(tmp_path / name)
    import_id = upload_bytes(client, path.read_bytes(), name)
    assert client.post(f"/tables/imports/{import_id}/source", json={"sheet": "一覧"}).status_code == 200
    assert save_layout(client, import_id).status_code == 200
    payload = {"name": "取り消し線の表",
               "columns": [{"index": i, "use": True, "role": r}
                           for i, r in enumerate(["key", "date", "entity", "attribute"])]}
    assert save_columns(client, import_id, payload).status_code == 200, save_columns
    wait_import_job(app, import_id)
    return import_id


# ---- 記録0件のまま確定させない -----------------------------------------------------------------

def test_a_table_with_no_usable_rows_cannot_be_confirmed():
    """取り込める行が1件も無いときは、確定を止めて理由を出す（警告のまま通さない）。

    警告のままだと⑥で［確定してMarkdownを作成］が押せてしまい、⑦のダウンロードで
    「先に［確定してMarkdownを作成］を押してください」と言われて空の画面に戻されていた。
    """
    spec = spec_from_dict(list_spec_dict())
    issues = run_checks([], spec, {"excluded": {"取り消し線の行": 5}})
    no_records = [i for i in issues if i.code == "no_records"]
    assert len(no_records) == 1
    assert no_records[0].level == "error"
    assert "取り消し線・非表示などで5行を除外しました" in no_records[0].message
    assert "表の範囲か元のファイルを見直してください" in no_records[0].message
    assert has_blocking(issues) is True


def test_the_preview_step_blocks_the_confirm_button_when_nothing_can_be_read(app, client, tmp_path):
    """⑥の［確定してMarkdownを作成］が押せなくなり、理由が画面に出る。"""
    import_id = _struck_import(app, client, tmp_path, "取り消し線.xlsx")

    html = preview_panel(app, client, import_id)["html"]
    assert "取り込める行がありません" in html
    assert "取り消し線・非表示などで5行を除外しました" in html
    assert "data-confirm-run disabled" in html
    assert "エラーが残っているため確定できません" in html
    # 確定も断る
    res = client.post(f"/tables/imports/{import_id}/confirm", json={})
    assert res.status_code == 400 and "エラーが残っているため確定できません" in res.get_json()["error"]


def test_a_confirmed_import_with_no_markdown_does_not_offer_a_download(app, client, tmp_path):
    """記録0件のまま確定された取り込みは、⑦でダウンロードのボタンを出さない。"""
    from app import tables

    import_id = _struck_import(app, client, tmp_path, "取り消し線2.xlsx")
    with app.app_context():
        tables.update_import(import_id, status="confirmed")
    data = panel(client, import_id, "done")
    assert "この取り込みから作られた Markdown はありません" in data["locked"]
    assert "download.zip" not in data["html"]

    res = client.get(f"/tables/imports/{import_id}/download.zip")
    assert res.status_code == 302
    page = client.get("/tables/new").get_data(as_text=True)
    assert "作成された Markdown がありません" in page
    assert "先に [確定してMarkdownを作成] を押してください" not in page


# ---- 日付にできない列は「日付」の選択肢を出さない -------------------------------------------------

def test_a_column_that_is_not_a_date_has_no_date_role_to_choose(app, client):
    """年月だけの列（2025-03）に「日付」の選択肢を出さない（選ぶと保存が必ず400になるため）。"""
    import_id = _prepare(client, "月報.csv", MONTH_CSV)
    html = panel_html(client, import_id, "columns")

    rows = html.split("<tr data-col")[1:]
    assert len(rows) == 5
    month = next(r for r in rows if 'data-header="実施月"' in r)
    assert '<option value="date"' not in month          # 日付として読めないので出さない
    assert '<option value="key"' in month and '<option value="attribute"' in month


def test_a_table_without_any_date_column_is_not_asked_to_pick_one(app, client):
    """日付として読める列が1つも無い表では、「日付の列を1つ選んでください」と求めない。"""
    import_id = _prepare(client, "月報2.csv", MONTH_CSV)
    html = panel_html(client, import_id, "columns")
    assert "日付の列が決まっていません。1つ選んでください" not in html


def test_choosing_a_date_role_for_a_non_date_column_says_what_is_wrong(app, client):
    """古い画面から送られても、「型を日付にしてください」（直せない指示）とは返さない。"""
    import_id = _prepare(client, "月報3.csv", MONTH_CSV)
    res = save_columns(client, import_id, _month_columns("月報", date_role="date"))
    assert res.status_code == 400
    error = res.get_json()["error"]
    assert "列「実施月」の値は日付として読めないので、日付の列にはできません" in error
    assert "型を日付にしてください" not in error


def test_a_real_date_column_can_still_be_chosen(app, client):
    """型の推定が当たっている日付の列には、これまでどおり「日付」を出す。"""
    from tests.conftest import CSV_TEXT

    import_id = _prepare(client, "一覧.csv", CSV_TEXT)
    html = panel_html(client, import_id, "columns")
    rows = html.split("<tr data-col")[1:]
    occurred = next(r for r in rows if 'data-header="発生日"' in r)
    assert '<option value="date"' in occurred
    assert save_columns(client, import_id, columns_payload("トラブル対応一覧")).status_code == 200


# ---- 見出しと並び順（見出しの材料が何も無い表） --------------------------------------------------

def _code_spec():
    d = list_spec_dict()
    d["columns"] = [
        {"key": "kind", "display": "コード種別", "type": "code", "role": "attribute", "md": "attribute"},
        {"key": "code", "display": "コード", "type": "code", "role": "attribute", "md": "attribute"},
        {"key": "name", "display": "名称", "type": "string", "role": "attribute", "md": "attribute"},
    ]
    d["record"] = {"key": ["code"], "fallback_key": ["code"]}
    d["period"] = {"date_column": ""}
    d["markdown"] = {k: v for k, v in (d.get("markdown") or {}).items() if k != "title_columns"}
    return spec_from_dict(d)


def test_records_without_a_title_column_get_a_heading_from_their_values():
    """見出しの材料が無い表でも、記録ごとに違う見出しにする（全部「（見出しなし）」にしない）。"""
    spec = _code_spec()
    first = record_title({"kind": "区分1", "code": "C001", "name": "停止"}, spec, {"row": 2})
    second = record_title({"kind": "区分2", "code": "C002", "name": "故障"}, spec, {"row": 3})
    assert first != second and "（見出しなし）" not in (first, second)
    assert "C001" in first and "C002" in second


def test_a_completely_empty_record_still_says_which_row_it_came_from():
    """値が1つも無い行だけ、行番号を見出しにする（同じ見出しが並ばない）。"""
    spec = _code_spec()
    assert record_title({}, spec, {"row": 42}) == "42行目の記録"
    assert record_title({}, spec) == "（見出しなし）"


def test_records_without_a_date_keep_the_order_of_the_source_table():
    """日付の列が無い表は、元の表の行の順に並べる（記録キーの文字くらべにしない）。"""
    from app.tables import render_all

    spec = _code_spec()
    records = []
    for row in (2, 10, 11, 100, 101):
        rec = _rec(f"行{row}", code=f"C{row:03d}", kind="区分", name=f"名称{row}")
        rec["source"] = {"file": "コード表.csv", "row": row}
        records.append(rec)
    files = render_all(spec, list(reversed(records)), {})
    assert len(files) == 1
    rows = [int(line.split("（")[1].split("行目")[0])
            for line in files[0].text.splitlines() if line.startswith("- 出典:")]
    assert rows == [2, 10, 11, 100, 101]


def test_records_with_a_date_are_still_ordered_by_date(app, client):
    """日付のある表の並びは今までどおり（同じ日付のときだけ元の行の順になる）。"""
    from app.tables import render_all

    spec = spec_from_dict(list_spec_dict())
    early = _rec("A-2", record_no="A-2", occurred_at="2026-08-01", equipment_id="EQ-1")
    late = _rec("A-1", record_no="A-1", occurred_at="2026-08-09", equipment_id="EQ-1")
    early["source"], late["source"] = {"row": 9}, {"row": 2}
    text = "\n".join(f.text for f in render_all(spec, [late, early], {}))
    assert text.index("A-2") < text.index("A-1")
