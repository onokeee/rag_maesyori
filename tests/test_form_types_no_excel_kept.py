"""帳票登録は Excel を預からない（利用者の指示 2026-09-21「見本のExcelは置かずに、設定だけ保持する」）。

置いた Excel は受け取った要求の中で読み取るだけで、uploads/ にも DB にも残さない。
次の操作（セルのクリック・項目の削除・見出しの手直し・読み取りテスト）では、ブラウザが同じ
ファイルを送り直す。サーバーは読み取った結果だけを core.workbook_cache に短い間だけ覚えておく
（メモリの中だけ。ディスクには何も書かない）。

ここで確かめること:
  - 置いても・クリックしても・使用開始しても、Excel のバイトはどこにも残らない
  - 残った設定だけで、開き直した画面の項目一覧と見出しの手直しができる
  - 同じ帳票の Excel をもう一度置けば、シートを見てクリックする作業に戻れる（種類は増えない）
  - 覚えている結果は、その人のもの・短い間だけ・数に上限がある
"""
import io

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from core import files, purge, workbook_cache
from core.files import UploadError
from models import database as db
from tests.test_forms_flow import (activate, add_field, book_part, create_type, panel_html, read_form,
                                   upload_forms)


@pytest.fixture(autouse=True)
def _clean_cache():
    """テストごとに、覚えているブックを空にする（プロセスで1つの置き場なので）。"""
    workbook_cache.clear()
    yield
    workbook_cache.clear()


def _report(path, report_id="R-001"):
    """見出しと値が並ぶ、ふつうの帳票。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "報告書"
    fill = PatternFill("solid", fgColor="FFD9E1F2")
    for coord, value in {"A1": "報告番号", "B1": report_id, "A2": "設備番号", "B2": "EQ-001",
                         "A3": "所見", "B3": "異音あり"}.items():
        ws[coord] = value
    for coord in ("A1", "A2", "A3"):
        ws[coord].fill = fill
    wb.save(path)
    return path


def _stored_files(app) -> list[str]:
    from pathlib import Path

    return [str(p) for p in Path(app.config["UPLOAD_DIR"]).rglob("*") if p.is_file()]


def _tables(app) -> set[str]:
    with app.app_context():
        return {row[0] for row in db.get_db().execute("SELECT name FROM sqlite_master WHERE type = 'table'")}


# ---- 置いても残らない ---------------------------------------------------------------------

def test_dropping_an_excel_keeps_no_file_and_no_row(app, client, tmp_path):
    """Excel を置いて種類を作っても、ファイルは1つも保存されず、控えの行もできない。"""
    path = _report(tmp_path / "設備修理報告書.xlsx")
    pattern_id = create_type(client, path, "設備修理報告書")
    add_field(client, pattern_id, "報告書", "A1", "B1")

    assert _stored_files(app) == []
    assert "pattern_samples" not in _tables(app)
    with app.app_context():
        # Excel を指す行（stored_path）は、どの表にもできていない
        conn = db.get_db()
        for (table,) in conn.execute("SELECT name FROM sqlite_master WHERE type = 'table'"):
            columns = {row[1] for row in conn.execute(f'PRAGMA table_info("{table}")')}
            if "stored_path" in columns:
                assert conn.execute(f'SELECT COUNT(*) FROM "{table}"').fetchone()[0] == 0, table
        # 残るのは設定だけ（シート名・見出しのセル・値のセル・向き・項目名）
        field = db.load_pattern(pattern_id).fields[0]
        assert (field.sheet_name, field.label_cell, field.cell, field.direction) == ("報告書", "A1", "B1", "right")
        assert field.display_name == "報告番号"


def test_the_whole_click_flow_needs_no_stored_excel(app, client, tmp_path):
    """置く → クリック → 見出しを直す → 読み取りテスト → 使用開始 → 帳票取り込み、が通る。"""
    path = _report(tmp_path / "設備修理報告書.xlsx", "R-777")
    pattern_id = create_type(client, path, "設備修理報告書")
    assert "「報告番号」を項目にしました" in add_field(client, pattern_id, "報告書", "A1", "B1")["message"]
    add_field(client, pattern_id, "報告書", "A2", "B2")

    # 見出しを手で直す（画面と同じく、いま置いている Excel も一緒に送る）
    res = client.post(f"/form-types/{pattern_id}/fields/report_id/label",
                      data={"name": "受付番号", "book": book_part(path)},
                      content_type="multipart/form-data")
    assert res.status_code == 200
    panel = res.get_json()["html"]
    assert "- 受付番号: R-777" in panel and "2項目中 <strong>2</strong>項目が見つかりました" in panel

    activate(client, pattern_id)
    doc_id, = upload_forms(client, path)
    assert read_form(client, doc_id, pattern_id, ["報告書"]).status_code == 200
    with app.app_context():
        import json

        read = {f["display_name"]: f["value"] for f in json.loads(db.get_document(doc_id)["data_json"])["fields"]}
    assert read == {"受付番号": "R-777", "設備番号": "EQ-001"}
    # 取り込んだ帳票のファイルは今までどおり残る（ダウンロードで消える）。帳票登録の分は1つも無い
    assert len(_stored_files(app)) == 1
    assert "samples" not in _stored_files(app)[0].replace("\\", "/")


def test_activation_does_not_claim_to_delete_anything(app, client, tmp_path):
    """使用開始の知らせで「見本のExcelを消しました」とは言わない（もともと置いていない）。"""
    path = _report(tmp_path / "点検表.xlsx")
    pattern_id = create_type(client, path, "点検表")
    add_field(client, pattern_id, "報告書", "A1", "B1")

    res = client.post(f"/form-types/{pattern_id}/status", json={"status": "active"})
    message = res.get_json()["message"]
    assert "使用を開始しました" in message and "消しました" not in message
    assert _stored_files(app) == []


# ---- 設定だけで開き直せる -----------------------------------------------------------------

def test_reopening_a_type_without_the_excel_shows_the_settings(app, client, tmp_path):
    """Excel を置いていない画面でも、項目の一覧と見出しの手直しはできる。"""
    path = _report(tmp_path / "点検表.xlsx")
    pattern_id = create_type(client, path, "点検表")
    add_field(client, pattern_id, "報告書", "A1", "B1")
    workbook_cache.clear()                       # 時間がたって忘れたあと

    panel = panel_html(client, pattern_id)
    assert "報告番号" in panel                   # 項目は残っている
    assert "サーバーに残していません" in panel   # 残していないことと、置き直せることを言う
    assert "この帳票のExcelを置く" in panel
    assert "いま Excel を置いていないので、読み取りテストはできません" in panel
    assert "—（Excelを置くと、この設定で読んだ値が出ます）" in panel

    # Excel が無くても見出しは直せる（Markdown に書く名前だけが変わる）
    res = client.post(f"/form-types/{pattern_id}/fields/report_id/label", json={"name": "受付番号"})
    assert res.status_code == 200 and "見出しを「受付番号」にしました" in res.get_json()["message"]
    with app.app_context():
        field = db.load_pattern(pattern_id).fields[0]
    assert field.display_name == "受付番号" and field.candidates[0] == "報告番号" and field.cell == "B1"


def test_dropping_the_same_excel_again_shows_the_sheet(app, client, tmp_path):
    """同じ帳票の Excel をもう一度置くと、シートが出て続きの作業ができる（種類は増えない）。"""
    path = _report(tmp_path / "点検表.xlsx")
    pattern_id = create_type(client, path, "点検表")
    add_field(client, pattern_id, "報告書", "A1", "B1")
    workbook_cache.clear()

    res = client.post(f"/form-types/{pattern_id}/panel", data={"book": book_part(path)},
                      content_type="multipart/form-data")
    assert res.status_code == 200
    body = res.get_json()
    assert "点検表.xlsx" in body["html"] and 'data-cell="A1"' in body["html"]
    assert "R-001" in body["html"]                 # 置いた帳票での値も出る
    assert "サーバーに残しません" in body["message"]
    with app.app_context():
        assert len(db.list_patterns()) == 1        # 置き直しで新しい種類は作らない
    assert _stored_files(app) == []

    # そのままセルをクリックして項目を足せる
    assert "「設備番号」を項目にしました" in add_field(client, pattern_id, "報告書", "A2", "B2")["message"]


def test_a_click_without_the_excel_asks_for_it_again(app, client, tmp_path):
    """Excel を送らずにセルのクリックだけ届いたら、置き直してくださいと断る。"""
    path = _report(tmp_path / "点検表.xlsx")
    pattern_id = create_type(client, path, "点検表")
    workbook_cache.clear()

    res = client.post(f"/form-types/{pattern_id}/fields",
                      data={"sheet": "報告書", "label_cell": "A1", "value_cell": "B1"})
    assert res.status_code == 400
    assert "もう一度置いてください" in res.get_json()["error"]
    with app.app_context():
        assert db.load_pattern(pattern_id).fields == []


def test_a_big_excel_is_refused(app, client, tmp_path, monkeypatch):
    """置く Excel には大きさの上限がある（保存せずにメモリで読むので、無制限にはできない）。"""
    from views import form_types

    monkeypatch.setattr(form_types, "BOOK_MAX_BYTES", 1000)
    path = _report(tmp_path / "大きい.xlsx")
    res = client.post("/form-types/new", data={"name": "大きい", "book": book_part(path)},
                      content_type="multipart/form-data")
    assert res.status_code == 400 and "大きすぎます" in res.get_json()["error"]
    with app.app_context():
        assert db.list_patterns() == []      # 読めなかったときは帳票の種類も作らない
    assert _stored_files(app) == []
    # 画面にも上限を書いてある
    assert "1ファイル" in client.get("/form-types/").get_data(as_text=True)


def test_deleting_a_type_leaves_nothing(app, client, tmp_path):
    """種類を削除しても、消すファイルはもともと無く、設定の行だけが消える。"""
    path = _report(tmp_path / "点検表.xlsx")
    pattern_id = create_type(client, path, "点検表")
    add_field(client, pattern_id, "報告書", "A1", "B1")

    assert client.post(f"/form-types/{pattern_id}/delete").status_code == 200
    with app.app_context():
        assert db.load_pattern(pattern_id) is None
        assert db.get_db().execute("SELECT COUNT(*) FROM pattern_fields").fetchone()[0] == 0
    assert _stored_files(app) == []


# ---- 送り直さずに済ませる合図（sha256）-------------------------------------------------------

def test_the_sheet_can_be_asked_for_by_the_hash_of_the_book(app, client, tmp_path):
    """一度読んだブックは、ファイルを送らずに合図（sha256）だけで開き直せる。"""
    path = _report(tmp_path / "点検表.xlsx")
    pattern_id = create_type(client, path, "点検表")
    panel = panel_html(client, pattern_id, book=path)
    book_hash = panel.split('data-book="')[1].split('"')[0]
    assert len(book_hash) == 64

    res = client.post(f"/form-types/{pattern_id}/fields",
                      data={"sheet": "報告書", "label_cell": "A1", "value_cell": "B1",
                            "book_hash": book_hash})
    assert res.status_code == 200 and "「報告番号」を項目にしました" in res.get_json()["message"]

    # 忘れたあとは合図だけでは足りない（ブラウザが持っているファイルを送り直す）
    workbook_cache.clear()
    res = client.post(f"/form-types/{pattern_id}/fields",
                      data={"sheet": "報告書", "label_cell": "A2", "value_cell": "B2",
                            "book_hash": book_hash})
    assert res.status_code == 400 and "もう一度置いてください" in res.get_json()["error"]


def test_another_browser_cannot_use_the_book_of_the_first_one(app, tmp_path):
    """覚えているブックはその作業場所（ブラウザ）のもの。ほかの人の合図では出てこない。"""
    from views import SESSION_ID_KEY

    path = _report(tmp_path / "点検表.xlsx")
    first, second = app.test_client(), app.test_client()
    for client, sid in ((first, "a" * 32), (second, "b" * 32)):
        with client.session_transaction() as cookie:
            cookie[SESSION_ID_KEY] = sid

    pattern_id = create_type(first, path, "点検表")
    panel = panel_html(first, pattern_id, book=path)
    book_hash = panel.split('data-book="')[1].split('"')[0]

    res = second.post(f"/form-types/{pattern_id}/fields",
                      data={"sheet": "報告書", "label_cell": "A1", "value_cell": "B1",
                            "book_hash": book_hash})
    assert res.status_code == 400 and "もう一度置いてください" in res.get_json()["error"]


# ---- 覚えておく置き場（メモリだけ）------------------------------------------------------------

def _fake_book(name: str, size: int = 10) -> workbook_cache.Book:
    return workbook_cache.Book(file_name=name, file_hash=name, size=size, info=object())


def test_the_cache_forgets_old_books():
    """しばらく使われなかったブックは忘れる（覚えているのは短い間だけ）。"""
    book = workbook_cache.put("sid", _fake_book("a"))
    assert workbook_cache.get("sid", "a") is not None
    book.used_at -= workbook_cache.TTL_SECONDS + 1      # 時間がたったことにする
    assert workbook_cache.get("sid", "a") is None
    assert workbook_cache.count() == 0
    # 新しく置き直したブックは、古いものを片付けるときに巻き込まれない
    old = workbook_cache.put("sid", _fake_book("b"))
    old.used_at -= workbook_cache.TTL_SECONDS + 1
    workbook_cache.put("sid", _fake_book("c"))
    assert workbook_cache.get("sid", "b") is None and workbook_cache.get("sid", "c") is not None


def test_the_cache_keeps_only_the_newest_books(monkeypatch):
    """数と大きさに上限があり、あふれたら古いものから落とす。"""
    monkeypatch.setattr(workbook_cache, "MAX_ENTRIES", 2)
    for name in ("a", "b", "c"):
        workbook_cache.put("sid", _fake_book(name))
    assert workbook_cache.count() == 2
    assert workbook_cache.get("sid", "a") is None      # いちばん古いものが落ちる
    assert workbook_cache.get("sid", "c") is not None

    monkeypatch.setattr(workbook_cache, "MAX_BYTES", 100)
    workbook_cache.clear()
    workbook_cache.put("sid", _fake_book("big", size=80))
    workbook_cache.put("sid", _fake_book("big2", size=80))
    assert workbook_cache.count() == 1 and workbook_cache.get("sid", "big") is None


def test_the_same_book_read_twice_is_only_parsed_once(app, client, tmp_path, monkeypatch):
    """同じ Excel を送り直しても、ブックを開き直さない（覚えている結果を使う）。"""
    path = _report(tmp_path / "点検表.xlsx")
    pattern_id = create_type(client, path, "点検表")

    from views import form_types

    opened = []
    real = form_types.load_workbook_info

    def counting(source):
        opened.append(1)
        return real(source)

    monkeypatch.setattr(form_types, "load_workbook_info", counting)
    add_field(client, pattern_id, "報告書", "A1", "B1")
    add_field(client, pattern_id, "報告書", "A2", "B2")
    assert opened == []          # 1回目（種類を作ったとき）に読んだ結果を使い回す


# ---- メモリに読むだけの受け取り（core.files.read_upload）----------------------------------------

class _Storage:
    def __init__(self, data: bytes, filename: str):
        self.stream = io.BytesIO(data)
        self.filename = filename


def test_read_upload_keeps_nothing_on_disk(app, tmp_path):
    """アップロードは保存先を作らずにメモリへ読む（中身と sha256 だけ返す）。"""
    import hashlib

    data = _report(tmp_path / "点検表.xlsx").read_bytes()
    with app.app_context():
        got = files.read_upload(_Storage(data, "点検表.xlsx"), {".xlsx", ".xlsm"}, 10_000_000)
    assert got.data == data and got.size == len(data)
    assert got.file_hash == hashlib.sha256(data).hexdigest()
    assert got.file_name == "点検表.xlsx"
    assert _stored_files(app) == []

    with app.app_context():
        with pytest.raises(UploadError, match="ファイルを選んでください"):
            files.read_upload(_Storage(data, "点検表.csv"), {".xlsx"}, 10_000_000)
        with pytest.raises(UploadError, match="空です"):
            files.read_upload(_Storage(b"", "空.xlsx"), {".xlsx"}, 10_000_000)
        with pytest.raises(UploadError, match="大きすぎます"):
            files.read_upload(_Storage(data, "点検表.xlsx"), {".xlsx"}, 10)
    assert _stored_files(app) == []


def test_precheck_reads_the_bytes_without_a_file(tmp_path):
    """事前チェックは、ファイルにしていない中身（bytes）でも同じように断る。"""
    files.precheck_excel(_report(tmp_path / "ok.xlsx").read_bytes())      # ふつうのブックは通る
    with pytest.raises(UploadError, match="Excelファイル"):
        files.precheck_excel(b"not an excel file")
    with pytest.raises(UploadError, match="セル数が上限"):
        files.precheck_excel(_report(tmp_path / "ok2.xlsx").read_bytes(), max_cells=1)


# ---- 前の版が置いた見本の片付け ---------------------------------------------------------------

def test_the_old_sample_folder_is_cleaned_up(app):
    """前の版が uploads/samples に置いたままの Excel は、起動時にフォルダごと捨てる。"""
    from pathlib import Path

    folder = Path(app.config["UPLOAD_DIR"]) / "samples"
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "0123456789abcdef0123456789abcdef.xlsx").write_bytes(b"old sample")

    with app.app_context():
        assert purge.purge_old_sample_files() == 1
        assert purge.purge_old_sample_files() == 0     # 2回目は何もしない
    assert not folder.exists()
    assert _stored_files(app) == []


def test_the_migration_drops_the_sample_table(tmp_path):
    """前の版の DB（pattern_samples がある）を開くと、その表を落とす。"""
    import sqlite3

    from tests.test_core_db import _make_old_db, make_app

    path = tmp_path / "app.db"
    _make_old_db(path)
    conn = sqlite3.connect(path)
    conn.execute("INSERT INTO pattern_samples (pattern_id, file_name, file_hash, stored_path, created_at) "
                 "VALUES (1, '見本.xlsx', 'h', 'samples/x.xlsx', '2026-01-01T00:00:00')")
    conn.commit()
    conn.close()

    app = make_app(tmp_path)
    with app.app_context():
        conn = db.get_db()
        tables = {row[0] for row in conn.execute("SELECT name FROM sqlite_master WHERE type = 'table'")}
        assert "pattern_samples" not in tables
        assert db.load_pattern(1) is not None          # 設定（帳票の種類）は残る
