"""帳票まわりの修正（3巡目の確認で見つかった不具合）の回帰テスト。"""
import copy
import json
from datetime import time, timedelta

from excel.extractor import apply_manual_values, refresh_summary
from excel.text import cell_text, format_unit, to_date, to_number
from tests.test_extraction import LABEL_FILL, _book, _fields, _values


def _number_field(**kw) -> dict:
    f = {"field_name": "downtime", "display_name": "停止時間", "data_type": "number", "required": False,
         "value": None, "unit": "", "spec_unit": "", "warning": None, "edited": False, "ai_filled": False}
    f.update(kw)
    return f


# ---- F3-1: 時:分の時間・範囲・「.5」 ------------------------------------------------------------

def test_time_cell_is_converted_to_the_field_unit_not_read_as_its_hour():
    assert to_number(time(1, 30), "01:30", "時間")[0] == 1.5
    assert to_number(time(1, 30), "01:30", "分")[0] == 90
    assert to_number(time(1, 30), "01:30", "")[0] == 90  # 単位の無い項目は分


def test_hmm_duration_cell_over_a_day_is_converted():
    td = timedelta(hours=25, minutes=30)
    assert cell_text(td) == "25:30"  # 「1 day, 1:30:00」にしない
    assert to_number(td, cell_text(td), "分")[0] == 1530
    assert to_number(td, cell_text(td), "時間")[0] == 25.5


def test_hmm_text_is_a_duration():
    assert to_number(None, "2:45", "時間")[0] == 2.75
    assert to_number(None, "2:45", "")[0] == 165
    value, warning = to_number(None, "2:45", "円")
    assert value == "2:45" and warning  # 分・時間でない項目では数値にしない


def test_range_with_a_different_unit_is_not_read_as_its_first_number():
    value, warning = to_number(None, "10～20分", "時間")
    assert value == "10～20分" and "範囲" in warning


def test_leading_dot_decimal_and_h_m_notation():
    assert to_number(None, ".5", "時間") == (0.5, None)
    assert to_number(None, "約.5時間", "時間")[0] == 0.5
    assert to_number(None, "1h30m", "")[0] == 90


def test_manual_partial_number_stays_to_be_checked():
    """手で「約90分」と入れても「数値の部分だけ」の警告は要確認のまま（黙って確定できない）。"""
    from views.forms import _field_status

    ex = {"fields": [_number_field(value=30, unit="分", spec_unit="分")]}
    apply_manual_values(ex, {"value-downtime": "約90分くらい"})
    f = ex["fields"][0]
    assert f["edited"] and "数値の部分だけ" in f["warning"]
    assert _field_status(f)["issue"] is True

    apply_manual_values(ex, {"value-downtime": "2:45"})
    assert f["value"] == 165 and f["unit"] == "分"


# ---- F3-2 / md3-3: 日付のあとの時刻・範囲・平成 ---------------------------------------------------

def test_time_after_japanese_date_era_and_iso_is_kept():
    assert to_date(None, "2024年7月29日 13:41") == ("2024-07-29 13:41", None)
    assert to_date(None, "２０２５年２月２４日 ２０:４０") == ("2025-02-24 20:40", None)
    assert to_date(None, "2023年7月10日 23時08分") == ("2023-07-10 23:08", None)
    assert to_date(None, "2023-07-10T23:08") == ("2023-07-10 23:08", None)
    assert to_date(None, "R5.11.16 12:11") == ("2023-11-16 12:11", None)
    assert to_date(None, "2024/1/5 (金) 9:05") == ("2024-01-05 09:05", None)
    assert to_date(None, "2024年1月5日") == ("2024-01-05", None)


def test_text_after_the_date_is_flagged():
    value, warning = to_date(None, "2023/7/10～7/12")
    assert value == "2023-07-10" and warning and "7/12" in warning


def test_heisei_dates_are_read():
    assert to_date(None, "H30.4.1") == ("2018-04-01", None)
    assert to_date(None, "平成30年4月1日") == ("2018-04-01", None)


# ---- md3-1: 表示形式の単位 ------------------------------------------------------------------

def test_format_unit():
    assert format_unit('#,##0"分"') == "分"
    assert format_unit('#,##0" 分";[Red]-#,##0" 分"') == "分"
    assert format_unit('"¥"#,##0') == "円"
    assert format_unit("#,##0") == "" and format_unit("General") == "" and format_unit('0" - "') == ""


def test_number_unit_from_the_cell_format(tmp_path):
    from openpyxl import load_workbook

    path_info = _book(tmp_path, "fmt.xlsx", {"報告書": ({"A1": "ダウンタイム", "B1": 264},
                                                      {"A1": LABEL_FILL}, [])})
    wb = load_workbook(path_info.path)
    wb["報告書"]["B1"].number_format = '#,##0"分"'
    wb.save(path_info.path)
    from excel.workbook import load_workbook_info

    info = load_workbook_info(path_info.path)
    values, fields = _values(info, _fields(("downtime", "number", "ダウンタイム")))
    assert values["downtime"] == 264 and fields["downtime"]["unit"] == "分" and not fields["downtime"]["warning"]


# ---- md3-2: 水平展開先の一覧の列見出しを1つの値のラベルにしない --------------------------------------

def test_list_under_heading_with_a_row_between_is_not_a_single_value(tmp_path):
    cells = {"A1": "７．水平展開", "A2": "区分", "B2": "☑同型機　□類似設備",
             "A3": "確認", "B3": "対象設備", "C3": "設備名", "D3": "実施日/予定", "E3": "結果・備考",
             "A4": "☑", "B4": "IMP-603", "C4": "高電流注入 3号機", "D4": "2023/11/18", "E4": "展開済み",
             "A5": "☑", "B5": "IMP-604", "C5": "高エネルギー注入 4号機", "D5": "2023/11/17", "E5": "展開済み",
             "A7": "影響", "B7": "停止時間", "C7": "影響ロット", "A8": "搬送停止", "B8": "120", "C8": "3"}
    fills = {c: LABEL_FILL for c in ("A1", "A2", "A3", "B3", "C3", "D3", "E3", "A7", "B7", "C7")}
    info = _book(tmp_path, "yoko.xlsx", {"報告書": (cells, fills, [])})
    values, _ = _values(info, _fields(("target", "string", "対象設備"), ("stop", "string", "停止時間")))
    assert values["target"] is None  # 水平展開先の1行目（IMP-603）を報告書の対象設備にしない
    assert values["stop"] == "120"  # 見出しの行＋値の行1つは、今までどおり項目として読む


# ---- F1: Excel のエラー値 ---------------------------------------------------------------

def test_excel_error_values_are_not_taken_as_values(tmp_path):
    from views.forms import _field_status

    info = _book(tmp_path, "err.xlsx", {"報告書": (
        {"A1": "報告番号", "B1": "#REF!", "A2": "作業時間", "B2": "#DIV/0!", "A3": "設備名", "B3": "CMP 1号機"},
        {"A1": LABEL_FILL, "A2": LABEL_FILL, "A3": LABEL_FILL}, [])})
    values, fields = _values(info, _fields(("report_id", "string", "報告番号"), ("work", "number", "作業時間"),
                                           ("name", "string", "設備名")))
    assert values["report_id"] is None and "#REF!" in fields["report_id"]["warning"]
    assert values["work"] is None and fields["work"]["unit"] == "" and "#DIV/0!" in fields["work"]["warning"]
    assert _field_status(fields["report_id"])["issue"] and _field_status(fields["work"])["issue"]
    assert values["name"] == "CMP 1号機"


# ---- F3-3: 空の明細表で行を足して消しただけ -------------------------------------------------------

def test_empty_table_add_and_remove_row_is_not_an_edit():
    from views.forms import _apply_values, _field_status

    f = {"field_name": "parts", "display_name": "使用部品", "data_type": "table", "required": False,
         "value": None, "table_columns": ["部品名", "数量"], "warning": "表に行がありません",
         "edited": False, "ai_filled": False, "unit": ""}
    ex = {"fields": [f], "pattern": {"id": 1, "version_no": 1}, "sheets": ["S"]}
    refresh_summary(ex)  # 保存済みの読み取り結果と同じ形にする
    confirmed = copy.deepcopy(ex)
    changed = _apply_values(ex, {"parts": json.dumps({"columns": ["部品名", "数量"], "rows": []})}, confirmed)
    assert changed is False
    assert ex["fields"][0]["value"] is None and not ex["fields"][0]["edited"]
    assert _field_status(ex["fields"][0])["source"] == "blank"


# ---- R3B-3: 途中保存の応答が届く前に画面を離れた（同じ画面の beacon） --------------------------------

def _working_doc(app):
    from models import database as db
    from tests.test_forms_fixes import _confirmed_doc

    doc_id = _confirmed_doc(app)
    with app.app_context():
        db.update_document(doc_id, confirmed_json=None)  # 確認中の帳票
    return doc_id


def test_a_beacon_from_the_same_page_is_not_refused_as_another_page(app, client):
    doc_id = _working_doc(app)
    old = client.get(f"/forms/{doc_id}/review").get_json()["version"]
    first = client.post(f"/forms/{doc_id}/draft",
                        json={"values": {"equipment_name": "A"}, "version": old, "page_token": "page-1"})
    assert first.status_code == 204
    # 応答（新しい版）を受け取る前に送った beacon は古い版のまま。同じ画面なら受け付ける
    second = client.post(f"/forms/{doc_id}/draft",
                         json={"values": {"equipment_name": "A2"}, "version": old, "page_token": "page-1"})
    assert second.status_code == 204
    # 別の画面（目印が違う）の古い版は止める
    other = client.post(f"/forms/{doc_id}/draft",
                        json={"values": {"equipment_name": "B"}, "version": old, "page_token": "page-2"})
    assert other.status_code == 409
    # 目印を送らない古い版も止める
    assert client.post(f"/forms/{doc_id}/draft",
                       json={"values": {"equipment_name": "C"}, "version": old}).status_code == 409


def test_a_beacon_is_refused_after_another_page_saved(app, client):
    doc_id = _working_doc(app)
    old = client.get(f"/forms/{doc_id}/review").get_json()["version"]
    res = client.post(f"/forms/{doc_id}/draft",
                      json={"values": {"equipment_name": "A"}, "version": old, "page_token": "page-a"})
    new = res.headers["X-Doc-Version"]
    assert client.post(f"/forms/{doc_id}/draft",
                       json={"values": {"equipment_name": "B"}, "version": new, "page_token": "page-b"}).status_code == 204
    # 画面Aの版は、画面Bの保存でできた版ではないので止める
    assert client.post(f"/forms/{doc_id}/draft",
                       json={"values": {"equipment_name": "A2"}, "version": old, "page_token": "page-a"}).status_code == 409


def test_draft_page_marks_are_forgotten_when_the_form_is_purged(app, client):
    """消した帳票の途中保存の目印（id）をメモリに残さない（design.md 3.3 データを残さない）。"""
    from views import forms

    doc_id = _working_doc(app)
    old = client.get(f"/forms/{doc_id}/review").get_json()["version"]
    assert client.post(f"/forms/{doc_id}/draft",
                       json={"values": {"equipment_name": "A"}, "version": old, "page_token": "page-1"}).status_code == 204
    assert doc_id in forms._DRAFT_TOKENS
    assert client.post(f"/forms/{doc_id}/delete").get_json()["ok"] is True
    assert doc_id not in forms._DRAFT_TOKENS
