"""一覧表の不具合修正（2巡目）の確認。"""
from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from core import jobs
from tables import pipeline, store
from tables.checks import has_blocking, run_checks
from tables.csv_source import MAX_COLUMNS, CsvSource, sniff_csv
from tables.detect import guess_layout, sample_data_rows
from tables.mapping import suggest_columns
from tables.normalize import parse_number_text, read_records
from tables.source import UploadError, open_source
from tables.spec import spec_from_dict, spec_from_suggestions
from tests.tables_helpers import (
    columns_payload, confirmed, csv_source, panel, preview_panel, save_columns, save_layout, upload, upload_csv,
    wait_import_job,
)
from tests.test_tables_pipe import _rec, list_spec_dict


def _auto(src, sheet):
    layout = guess_layout(src, sheet)
    spec = spec_from_suggestions("テスト", layout, suggest_columns(layout.headers, sample_data_rows(src, sheet, layout)), {})
    records, _issues, stats = read_records(src, {"sheet": sheet}, layout, spec)
    return layout, spec, records, stats


# ---- 読み込み（CSV） ----------------------------------------------------------------------------

def test_cp932_file_with_one_bad_byte_stays_cp932(tmp_path):
    """CP932 の拡張文字（髙・﨑）があるファイルに読めないバイトが1つあっても、Shift_JIS-2004 にして字を化けさせない。"""
    rows = "".join(f"TR-00{i},2026-08-01,髙砂ポンプ{i},㈱山﨑製作所で停止\r\n" for i in range(5))
    p = tmp_path / "a.csv"
    p.write_bytes(("管理No,発生日,設備名,内容\r\n" + rows).encode("cp932") + b"TR-999,2026-08-02,\x86\x9f,x\r\n")
    sniff = sniff_csv(p)
    assert sniff.encoding == "cp932" and sniff.decode_error_line == 7 and sniff.warnings
    src = CsvSource(p, "a.csv", {"encoding": "cp932", "delimiter": ",", "errors": "replace"})
    texts = [c.text for r in src.rows() for c in r.cells]
    assert "髙砂ポンプ0" in texts and "㈱山﨑製作所で停止" in texts


def test_long_but_closed_quoted_value_is_a_warning_not_an_error(tmp_path):
    """" が正しく閉じた長い値（600行の貼り付けログ）は確定を止めない。閉じ忘れの疑いとして警告だけ出す。"""
    log = "\n".join(f"{i}: 確認" for i in range(600))
    text = ("管理No,発生日,対応内容\r\nTR-001,2026-08-01,a\r\nTR-002,2026-08-02,b\r\nTR-003,2026-08-03,c\r\n"
            f'TR-004,2026-08-04,"{log}"\r\nTR-005,2026-08-05,e\r\n')
    p = tmp_path / "long.csv"
    p.write_bytes(text.encode("utf-8"))
    src = open_source(p, "long.csv")
    assert len(list(src.rows())) == 6 and src.unclosed_quote_row is None and src.long_record_row == 5
    layout, spec, records, stats = _auto(src, "long.csv")
    assert len(records) == 5 and stats.long_record_row == 5
    issues = run_checks(records, spec, stats)
    assert not has_blocking(issues)
    assert any(i.code == "long_record" and i.level == "warning" and "5行目" in i.message for i in issues)


def test_csv_wider_than_the_column_limit_is_refused(tmp_path):
    p = tmp_path / "wide.csv"
    header = ",".join(f"列{i}" for i in range(MAX_COLUMNS + 1))
    p.write_text(header + "\r\n" + ",".join("1" for _ in range(MAX_COLUMNS + 1)) + "\r\n", encoding="utf-8")
    with pytest.raises(UploadError, match="列数が上限"):
        open_source(p, "wide.csv")


# ---- 表の範囲・見出し ----------------------------------------------------------------------------

def test_row_with_its_own_no_and_date_is_not_a_total_row(tmp_path):
    """点検項目が「稼働時間累計」でも、自分の点検No・点検日がある行はデータ（合計行で表を切らない）。"""
    lines = ["点検No,点検日,点検項目,測定値,判定"]
    for i in range(40):
        item = "稼働時間累計" if i % 4 == 3 else f"振動{i}"
        lines.append(f"P-{i:03d},2026/04/{i % 28 + 1:02d},{item},{i * 10},{'' if i % 4 == 3 else 'OK'}")
    p = tmp_path / "check.csv"
    p.write_bytes(("\r\n".join(lines) + "\r\n").encode("utf-8"))
    src = open_source(p, "check.csv")
    layout = guess_layout(src, "check.csv")
    assert layout.counts.get("data") == 40 and not layout.counts.get("subtotal") and layout.warnings == []
    # 番号・日付のない本当の合計行は、これまでどおり合計
    p2 = tmp_path / "total.csv"
    p2.write_bytes(("\r\n".join(lines[:11] + [",,合計,450,"]) + "\r\n").encode("utf-8"))
    src2 = open_source(p2, "total.csv")
    assert guess_layout(src2, "total.csv").counts.get("data") == 10


def test_unlabeled_last_column_is_kept(tmp_path):
    """右端の見出しが空欄でも値のある列は捨てない（「列6」として取り込み、Markdown にも出る）。"""
    text = "管理No,発生日,設備名,現象,停止時間(分),\r\n" + "".join(
        f"T{i},2026/04/{i + 1:02d},ポンプ,停止,{i},ベアリング交換済み{i}\r\n" for i in range(20))
    p = tmp_path / "blank.csv"
    p.write_bytes(text.encode("utf-8"))
    src = open_source(p, "blank.csv")
    layout, spec, records, _stats = _auto(src, "blank.csv")
    assert layout.headers[-1] == "列6"
    col = next(c for c in spec.columns if "列6" in (c.headers or [c.display]))
    assert records[3].values.get(col.key) == "ベアリング交換済み3"


def test_repeated_headers_keep_their_name_as_display():
    """同じ見出し「備考」が3つあっても、表示名が「2」「3」にならない。"""
    from tables.detect import _dedupe
    from tables.mapping import _display_name

    headers = _dedupe(["管理No", "発生日", "設備名", "備考", "備考", "備考"])
    assert headers[3:] == ["備考", "備考(2)", "備考(3)"]
    assert [_display_name(h) for h in headers[3:]] == ["備考", "備考(2)", "備考(3)"]
    sugg = suggest_columns(headers, [["T1", "2026/04/01", "ポンプ", "a", "b", "c"]])
    displays = [s["display"] if isinstance(s, dict) else s.display for s in (sugg.values() if isinstance(sugg, dict) else sugg)]
    assert "2" not in displays and "3" not in displays


# ---- 数値の変換 --------------------------------------------------------------------------------

def test_excel_number_under_an_unconvertible_unit_is_reported():
    from tables.normalize import _convert_number
    from tables.spec import ColumnSpec

    col = ColumnSpec(key="downtime", display="停止時間", type="number", role="measure", unit="分")
    assert _convert_number(7200, "7200", col, None, "s") == (120, None, None)   # 秒→分
    assert _convert_number(3, "3", col, None, "時") == (180, None, None)        # 時→分
    value, error, _ = _convert_number(5, "5", col, None, "kg")
    assert value == 5 and error and "換算できません" in error


def test_huge_exponent_is_a_cell_issue_not_a_crash(tmp_path):
    assert parse_number_text("1e309") == (None, "")
    assert parse_number_text("9" * 400) == (None, "")
    assert parse_number_text("1e308") == (1e308, "")
    text = "管理No,発生日,停止時間(分)\r\n" + "".join(f"T{i},2026/04/{i + 1:02d},{i}\r\n" for i in range(10)) + "T99,2026/04/20,1e999\r\n"
    p = tmp_path / "big.csv"
    p.write_bytes(text.encode("utf-8"))
    src = open_source(p, "big.csv")
    layout = guess_layout(src, "big.csv")
    spec = spec_from_suggestions("テスト", layout, suggest_columns(layout.headers, sample_data_rows(src, "big.csv", layout)), {})
    num = next(c for c in spec.columns if c.display.startswith("停止時間"))
    num.type, num.role = "number", "measure"
    records, issues, _stats = read_records(src, {"sheet": "big.csv"}, layout, spec)
    assert len(records) == 11
    assert any("1e999" in i.message and "数値に変換できません" in i.message for i in issues)


def test_overflowing_sum_renders_without_error():
    from tables.markdown import render_all
    from tables.records import fmt_number

    assert fmt_number(float("inf")) == "inf"
    spec = spec_from_dict(list_spec_dict(group_by="entity_month"))
    records = [_rec("A-1", record_no="A-1", occurred_at="2026-08-03", equipment_id="CMP-101", cost=1e308),
               _rec("A-2", record_no="A-2", occurred_at="2026-08-10", equipment_id="CMP-101", cost=1e308)]
    assert render_all(spec, records, {})


def test_placeholder_equipment_is_not_an_equipment():
    """対象設備が「調査中」だけの記録は、設備別のファイル分けに入れない（本文には原文のまま出す）。"""
    from tables.markdown import render_all
    from tables.records import entity_value

    d = list_spec_dict(group_by="entity_month")
    d["columns"] = [c for c in d["columns"] if c["key"] != "equipment_name"]
    spec = spec_from_dict(d)
    assert entity_value({"equipment_id": "調査中"}, spec) == ("", "")
    assert entity_value({"equipment_id": "CMP-101"}, spec) == ("CMP-101", "")
    records = [_rec("A-1", record_no="A-1", occurred_at="2026-08-03", equipment_id="調査中", symptom="停止"),
               _rec("A-2", record_no="A-2", occurred_at="2026-08-10", equipment_id="CMP-101", symptom="異音")]
    files = render_all(spec, records, {})
    assert not any("調査中" in f.name for f in files)
    assert "- 設備番号: 調査中" in "".join(f.text for f in files if f.kind == "records")


# ---- 画面（取り込み・削除・ダウンロード） -----------------------------------------------------------------

def _confirmed_import(app, client) -> int:
    return confirmed(app, client, "トラブル一覧.csv", "トラブル対応一覧")


def test_the_import_carries_its_own_spec_and_takes_it_along_when_downloaded(app, client):
    """取り込み設定は保存しない。設定は取り込みの行が持ち、ダウンロードすると一緒に消える。"""
    import_id = _confirmed_import(app, client)
    with app.app_context():
        imp = store.get_import(import_id)
        assert imp["spec"] is not None and imp["spec"].name == "トラブル対応一覧"
        assert imp["spec_hash"] and imp["template_version_id"] == import_id
    res = client.get(f"/tables/imports/{import_id}/download.zip")
    assert res.status_code == 200 and res.mimetype == "application/zip"
    with app.app_context():
        assert store.get_import(import_id) is None


def test_second_download_after_purge_is_404_not_500(app, client, monkeypatch):
    import_id = _confirmed_import(app, client)

    def gone(*args, **kwargs):
        raise FileNotFoundError("purged")

    monkeypatch.setattr(pipeline, "build_download", gone)
    assert client.get(f"/tables/imports/{import_id}/download.zip").status_code == 404


def test_delete_is_refused_while_the_preview_draft_is_being_made(app, client):
    from tests.test_tables_flow import _queued_job

    import_id = upload_csv(client, "間違い.csv")
    _queued_job(app, import_id, "table_preview")
    res = client.post(f"/tables/imports/{import_id}/delete")
    assert res.status_code == 400 and "Markdownの下書きを作っている間は削除できません" in res.get_json()["error"]
    with app.app_context():
        assert store.get_import(import_id) is not None


def test_failed_preview_draft_can_be_made_again(app, client, monkeypatch):
    import_id = upload_csv(client, "t.csv")
    csv_source(client, import_id)
    save_layout(client, import_id)
    save_columns(client, import_id, columns_payload("再作成テスト"))
    wait_import_job(app, import_id)

    real = pipeline.render_files
    calls = {"n": 0}

    def flaky(*args, **kwargs):
        calls["n"] += 1
        if calls["n"] == 1:
            raise OSError("一時的に使えない")
        return real(*args, **kwargs)

    monkeypatch.setattr(pipeline, "render_files", flaky)

    def wait_preview():
        with app.app_context():
            jobs.wait_job(jobs.latest_job("table_import", import_id, kind="table_preview")["id"], timeout=60)

    panel(client, import_id, "preview")
    wait_preview()
    data = panel(client, import_id, "preview")
    assert data["failed"] is True and "もう一度作る" in data["html"]
    # 下書きの失敗を「表を読み込めませんでした」と取り違えない（読み込みは終わっている）
    assert "表を読み込めませんでした" not in data["html"]
    res = client.post(f"/tables/imports/{import_id}/preview")
    assert res.status_code == 200 and res.get_json()["building"] is True
    wait_preview()
    assert "確定してMarkdownを作成" in panel(client, import_id, "preview")["html"]


def test_import_status_is_written_before_the_job_can_finish(app, monkeypatch):
    """ジョブがすぐ終わって書いた状態を、あとから reading / confirming で上書きしない。"""
    with app.app_context():
        import_id = store.create_import("x.csv", "h" * 64, "tables/x.csv", {"encoding": "utf-8"})

        def instant(kind, ref_type, ref_id, fn, params):
            # ワーカーが先に終わった場合と同じ: 呼び出し元が状態を書く前にジョブが状態を書く
            store.update_import(ref_id, status="preview" if kind == "table_read" else "confirmed")
            return 12345

        monkeypatch.setattr(pipeline.jobs, "start_job", instant)
        assert pipeline.start_read_job(import_id) == 12345
        imp = store.get_import(import_id)
        assert imp["status"] == "preview" and imp["job_id"] == 12345
        pipeline.start_render_job(import_id)
        assert store.get_import(import_id)["status"] == "confirmed"


# ---- 読み込み（Excel） --------------------------------------------------------------------------

def _small_book(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append(["管理No", "発生日", "設備名", "現象", "停止時間"])
    for i in range(10):
        ws.append([f"T{i}", f"2026/04/{i + 1:02d}", "ポンプ", "停止", i + 1])
    return wb, ws


def test_format_only_far_cell_does_not_block_upload(app, client, tmp_path):
    from tables.excel_source import ExcelSource

    wb, ws = _small_book(tmp_path)
    ws["XFD1048576"].fill = PatternFill("solid", fgColor="FFFF00")
    path = tmp_path / "far.xlsx"
    wb.save(path)
    ExcelSource(path).close()
    res = upload(client, path.read_bytes(), "far.xlsx")
    assert res.status_code == 200 and res.get_json()["import_id"]


def test_workbook_with_unreadable_values_is_refused_at_upload(app, client, tmp_path):
    wb, _ws = _small_book(tmp_path)
    wb.save(tmp_path / "ok.xlsx")
    src = zipfile.ZipFile(tmp_path / "ok.xlsx")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb"<v>3</v>", b"<v>NaN</v>", data, count=1)
            z.writestr(item, data)
    res = upload(client, out.getvalue(), "nan.xlsx")
    assert res.status_code == 400 and res.get_json()["error"]
    with app.app_context():
        assert store.list_imports(limit=10) == []
    assert [p for p in Path(app.config["UPLOAD_DIR"]).rglob("*") if p.is_file()] == []


def test_retention_note_only_says_the_data_is_deleted():
    """渡すのは zip だけになったので、案内も「消える」ことだけにする（CSVの順番の話は残さない）。"""
    from views.tables import DELETE_ON_DOWNLOAD_NOTE

    assert "より先" not in DELETE_ON_DOWNLOAD_NOTE and "CSV" not in DELETE_ON_DOWNLOAD_NOTE
    assert "管理用" not in DELETE_ON_DOWNLOAD_NOTE and "サーバーから消えます" in DELETE_ON_DOWNLOAD_NOTE


def test_confirmed_import_can_be_deleted_without_downloading(app, client):
    """ダウンロードせずに消す道を、確定したあとの段にも残す。"""
    import_id = _confirmed_import(app, client)
    assert "data-import-delete" in panel(client, import_id, "done")["html"]
    res = client.post(f"/tables/imports/{import_id}/delete")
    assert res.status_code == 200 and res.get_json()["ok"] is True
    with app.app_context():
        assert store.get_import(import_id) is None


def test_preview_headers_show_the_unit(app, client):
    import_id = _confirmed_import(app, client)
    assert "<th>停止時間（分）</th>" in preview_panel(app, client, import_id)["html"]


def test_a_panel_of_a_missing_import_is_404(client):
    assert client.get("/tables/imports/999/panel/done").status_code == 404
    assert client.get("/tables/imports/999/panel/unknown").status_code == 404
