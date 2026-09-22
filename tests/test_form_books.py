"""帳票登録の「覚えたシート」（forms.WorkbookInfo.to_json / database.pattern_books）。

利用者の指示（2026-09-22）:
  「再度シートを置かなくても、登録したときにシートのセル番地と文字情報を記憶しておけばだせるはず」
  「セル色の情報は不要」
Excel のファイルは今までどおり保存しない（2026-09-21「見本のExcelは置かずに、設定だけ保持する」）。
覚えるのはシートの中身だけ: シート名・非表示・結合・セルの番地と文字（型のある値）・太字・塗りの有る無し・
塗りの組の番号（色そのものではない）・画像の場所・数式の結果が無いセル・1904年基準。

ここで確かめること:
  - JSON にして戻しても、セル・結合・クリックで作る項目・読み取りの結果が同じ
  - 色の値は JSON に無い（塗りの組の番号だけ。同じ色は同じ番号）
  - 種類を作ると覚え、別の Excel を置くと入れ替わり、種類を消せば消える。取り込みの片付けでは消えない
  - Excel を送らずに、開き直し・クリック・読み取りテスト・再起動が通る
  - 大きすぎるブックは覚えずに断る（種類も作らない）
"""
from __future__ import annotations

import hashlib
import json
from dataclasses import asdict
from datetime import date, datetime, time, timedelta
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.styles.colors import Color

from app import core, create_app
from app import database as db
from app.forms import (
    WorkbookInfo,
    click_field,
    extract_document,
    load_workbook_info,
    match_pattern,
    table_cells,
)
from tests.conftest import add_confirmed_document, make_config, confirmed_import
from tests.test_forms import activate, add_field, book_part, create_type, panel_html

SAMPLE_F1 = Path("samples/forms/F1_設備修理報告書/Rev1_2019制定/修理報告書_INS-701_20230704.xlsx")


@pytest.fixture(autouse=True)
def _clean_cache():
    """テストごとに、メモリに覚えているブックを空にする（プロセスで1つの置き場なので）。"""
    core.clear()
    yield
    core.clear()


def _f1_path() -> Path:
    if not SAMPLE_F1.exists():
        pytest.skip(f"サンプルがありません: {SAMPLE_F1}")
    return SAMPLE_F1


def _assert_same_info(info: WorkbookInfo, back: WorkbookInfo) -> None:
    """ブックを開いた結果と、JSON から戻した結果が、読み取りに関わる全部で同じ。"""
    assert list(back.grids) == list(info.grids)
    assert back.date1904 == info.date1904
    assert back.uncached_formulas == info.uncached_formulas
    assert back.images == info.images_in(info.sheet_names)
    for name, grid in info.grids.items():
        other = back.grids[name]
        assert (other.hidden, other.max_row, other.max_col) == (grid.hidden, grid.max_row, grid.max_col), name
        assert other.merged_bounds() == grid.merged_bounds(), name
        assert other._bounds == grid._bounds, name
        assert list(other.cells) == list(grid.cells), name
        for key, cell in grid.cells.items():
            assert asdict(other.cells[key]) == asdict(cell), (name, key)
            assert type(other.cells[key].value) is type(cell.value), (name, key)
        assert table_cells(other) == table_cells(grid), name


# ---- JSON にして戻しても同じ ----------------------------------------------------------------

def test_a_snapshot_round_trips_every_cell_of_the_small_samples(sample_dir):
    """見本4種（修理報告書3種・点検記録表）を JSON にして戻すと、セル・結合・表の見出しが全部同じ。"""
    for name in ("standard.xlsx", "shifted.xlsx", "table.xlsx", "inspection.xlsx"):
        info = load_workbook_info(sample_dir / name)
        text = info.to_json()
        back = WorkbookInfo.from_json(text)
        _assert_same_info(info, back)
        assert back.path is None
        # もう一度 JSON にしても同じ文字になる（戻す→書く、で変わらない）
        assert back.to_json() == text, name


@pytest.mark.samples
def test_a_snapshot_round_trips_the_real_repair_report_with_its_images():
    """実物の修理報告書（画像あり・結合多数）でも同じ。画像は場所（シート・セル番地・名前）だけ覚える。"""
    info = load_workbook_info(_f1_path())
    assert info.images, "この見本には画像があるはず"
    back = WorkbookInfo.from_json(info.to_json())
    _assert_same_info(info, back)
    assert back.images_in(back.sheet_names) == info.images_in(info.sheet_names)
    # 覚えるのは中身の文字。1枚の帳票なら数KB（画面・README に書く目安）
    size = len(info.to_json().encode("utf-8"))
    assert 1_000 < size < 64 * 1024, size


def test_a_snapshot_keeps_the_type_of_each_value(tmp_path):
    """日時・日付・時刻・時間・真偽・整数・小数・文字は、戻すと同じ型になる。表示の文字（text）も同じ。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "値の型"
    values = {
        "B1": datetime(2026, 9, 22, 10, 30), "B2": date(2026, 9, 22), "B3": time(9, 5),
        "B4": timedelta(hours=25, minutes=30), "B5": True, "B6": 42, "B7": 3.5, "B8": "文字",
        "B9": 1.0, "B10": 0,
    }
    for coord, value in values.items():
        ws["A" + coord[1:]] = "ラベル" + coord[1:]
        ws[coord] = value
    ws["B6"].number_format = '#,##0"分"'
    wb.save(tmp_path / "型.xlsx")

    info = load_workbook_info(tmp_path / "型.xlsx")
    back = WorkbookInfo.from_json(info.to_json())
    _assert_same_info(info, back)
    grid = back.grids["値の型"]
    assert isinstance(grid.cells[(1, 2)].value, datetime) and grid.cells[(1, 2)].value == values["B1"]
    assert isinstance(grid.cells[(4, 2)].value, timedelta) and grid.cells[(4, 2)].text == "25:30"
    assert grid.cells[(5, 2)].value is True
    assert isinstance(grid.cells[(6, 2)].value, int) and grid.cells[(6, 2)].fmt_unit == "分"
    assert isinstance(grid.cells[(7, 2)].value, float)
    assert grid.cells[(8, 2)].value == "文字"
    # 「1.0」は openpyxl が開いた時点で整数 1 になる（xlsx には "1" と書かれる）。開いたときと同じ型のまま戻り、
    # 文字は「1」。0 は値のあるセル
    assert type(grid.cells[(9, 2)].value) is type(info.grids["値の型"].cells[(9, 2)].value)
    assert grid.cells[(9, 2)].text == "1"
    assert grid.cells[(10, 2)].value == 0 and grid.cells[(10, 2)].text == "0"


def test_a_snapshot_keeps_hidden_sheets_merges_bold_and_uncached_formulas(tmp_path):
    """非表示のシート・結合・太字・「計算結果の無い数式」も戻る。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "表"
    ws["A1"], ws["A2"], ws["B2"] = "■ 見出し", "設備番号", "EQ-001"
    ws["A1"].font = Font(bold=True)
    ws.merge_cells("A1:C1")
    ws["A3"], ws["B3"] = "合計", "=SUM(B2:B2)"      # openpyxl で書いた数式は計算結果が無い
    hidden = wb.create_sheet("控え")
    hidden["A1"] = "控えの値"
    hidden.sheet_state = "hidden"
    wb.save(tmp_path / "形.xlsx")

    info = load_workbook_info(tmp_path / "形.xlsx")
    assert info.uncached_formulas == {"表": {(3, 2)}}
    back = WorkbookInfo.from_json(info.to_json())
    _assert_same_info(info, back)
    assert back.grids["控え"].hidden is True
    assert back.grids["表"].cells[(1, 1)].bold is True
    assert back.grids["表"].cells[(1, 1)].coord == "A1:C1"
    assert back.grids["表"].bounds(1, 3) == (1, 1, 1, 3)
    assert back.uncached_formulas == {"表": {(3, 2)}}


def test_a_snapshot_can_be_read_back_even_when_cells_are_short_or_odd():
    """壊れかけの JSON（項目の足りないセル・知らない型）でも、読めるものだけで戻す。"""
    data = {"version": 1, "date1904": True, "images": [{"type": "image", "sheet": "S", "location": "A9", "name": "x"},
                                                        {"type": "image", "sheet": "無いシート", "location": "A1"}],
            "sheets": [{"name": "S", "hidden": 0, "merged": [[1, 1, 1, 2]],
                        "cells": [[1, 1, 1, 2, "見出し"], [2, 1, 2, 1, {"t": "謎", "v": "そのまま"}],
                                  [2, 2, 2, 2, {"t": "date", "v": "not-a-date"}], [3, 1, 3, 1, "  "],
                                  [1, 2, 1, 2, "結合の中"]],
                        "uncached": [[5, 5]]}]}
    back = WorkbookInfo.from_snapshot(data)
    grid = back.grids["S"]
    assert back.date1904 is True
    assert back.images == [{"type": "image", "sheet": "S", "location": "A9", "name": "x"}]   # 無いシートの画像は落とす
    assert grid.cells[(1, 1)].coord == "A1:B1" and grid.cells[(1, 1)].bold is False
    assert grid.cells[(2, 1)].value == "そのまま"          # 知らない型はその文字
    assert grid.cells[(2, 2)].value == "not-a-date"        # 戻せない日付もその文字
    assert (3, 1) not in grid.cells                        # 空白だけのセルは無いのと同じ
    assert (1, 2) not in grid.cells                        # 結合範囲の左上以外は落とす
    assert back.uncached_formulas == {"S": {(5, 5)}}


# ---- 色は覚えない ----------------------------------------------------------------------------

def test_a_snapshot_has_no_colour_only_fill_groups(tmp_path):
    """JSON に色の値（RGB・テーマ色）は無い。同じ色のセルは同じ番号、違う色は違う番号（見つけた順に 1, 2, …）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "色"
    blue, green = PatternFill("solid", fgColor="FFD9E1F2"), PatternFill("solid", fgColor="FFE2EFDA")
    theme = PatternFill("solid", fgColor=Color(theme=4, tint=0.6))
    ws["A1"], ws["B1"] = "報告番号", "R-1"
    ws["A2"], ws["B2"] = "設備番号", "EQ-1"
    ws["A3"], ws["B3"] = "■ 明細", "x"
    ws["A4"], ws["B4"] = "テーマ色", "y"
    ws["A1"].fill = blue
    ws["A2"].fill = blue
    ws["A3"].fill = green
    ws["A4"].fill = theme
    ws["B1"].fill = PatternFill("solid", fgColor="FFD9E1F2")   # A1 と同じ色
    second = wb.create_sheet("2枚目")
    second["A1"] = "2枚目の見出し"
    second["A1"].fill = green                                  # 別のシートでも同じ色なら同じ番号
    wb.save(tmp_path / "色.xlsx")

    info = load_workbook_info(tmp_path / "色.xlsx")
    text = info.to_json()
    for forbidden in ("D9E1F2", "E2EFDA", "rgb", "theme", "tint", "FFFF"):
        assert forbidden not in text, forbidden
    cells = info.grids["色"].cells
    assert cells[(1, 1)].fill == cells[(2, 1)].fill == cells[(1, 2)].fill == 1
    assert cells[(3, 1)].fill == 2
    assert cells[(4, 1)].fill == 3
    assert cells[(2, 2)].fill == 0 and cells[(2, 2)].filled is False
    assert all(c.filled for c in (cells[(1, 1)], cells[(3, 1)], cells[(4, 1)]))
    assert info.grids["2枚目"].cells[(1, 1)].fill == 2

    back = WorkbookInfo.from_json(text)
    _assert_same_info(info, back)
    # 覚えた形でも「同じ塗りか」の比較はそのまま（表の終わり・ラベル欄の判定が使う）
    kept = back.grids["色"].cells
    assert kept[(1, 1)].fill == kept[(1, 2)].fill and kept[(1, 1)].fill != kept[(3, 1)].fill
    # 番号は JSON の中では小さな整数だけ
    for sheet in json.loads(text)["sheets"]:
        for item in sheet["cells"]:
            assert isinstance(item[7], int) and 0 <= item[7] <= 3


# ---- クリックと読み取りが同じ ------------------------------------------------------------------

def _rows_from_clicks(grid, clicks):
    used: set[str] = set()
    rows = []
    for label, value in clicks:
        row, error = click_field(grid, label, value, used)
        assert row is not None, (label, value, error)
        used.add(row["field_name"])
        rows.append(row)
    return rows


def test_clicks_and_extraction_agree_between_the_book_and_its_snapshot(app, client, sample_dir):
    """開いたブックと覚えたシートで、クリックで作る項目も、その設定での読み取り結果も同じ。"""
    path = sample_dir / "inspection.xlsx"       # 見出し＋値の欄と、1回のクリックで項目になる明細表（A7:C7）がある
    info = load_workbook_info(path)
    back = WorkbookInfo.from_json(info.to_json())
    sheet = info.sheet_names[0]
    grid, kept = info.grids[sheet], back.grids[sheet]

    # 見出し＋値のクリックと、明細表の列見出しの1回クリック
    heads = sorted(table_cells(grid))
    assert heads == ["A7", "B7", "C7"] and heads == sorted(table_cells(kept))
    clicks = [(c.coord.split(":")[0], "") for c in grid.text_cells() if c.filled and c.coord.split(":")[0] not in heads][:4]
    clicks = [(label, _value_right_of(grid, label)) for label, _ in clicks] + [(heads[0], "")]
    assert _rows_from_clicks(grid, clicks) == _rows_from_clicks(kept, clicks)

    # 画面と同じ道すじで種類を作り、その設定で両方を読む
    pattern_id = create_type(client, path, "同じ読み取り")
    for label, value in clicks:
        add_field(client, pattern_id, sheet, label, value)
    with app.app_context():
        pattern = db.load_pattern(pattern_id)
        assert pattern.fields, "クリックで項目ができているはず"
        first = extract_document(info, pattern, [sheet])
        second = extract_document(back, pattern, [sheet])
    assert first["fields"] == second["fields"]
    assert first["values"] == second["values"]
    assert first["attachments"] == second["attachments"]
    assert match_pattern(info, pattern).sheet_names == match_pattern(back, pattern).sheet_names


def _value_right_of(grid, label_coord: str) -> str:
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string

    letters, row = coordinate_from_string(label_coord)
    cell = grid.cells.get((row, column_index_from_string(letters)))
    return f"{get_column_letter(cell.max_col + 1)}{row}" if cell is not None else ""


# ---- 種類と一緒に覚える -----------------------------------------------------------------------

def _book_row(app, pattern_id: int) -> dict | None:
    with app.app_context():
        return db.load_pattern_book(pattern_id)


def _report(path, report_id="R-001"):
    wb = Workbook()
    ws = wb.active
    ws.title = "報告書"
    for coord, value in {"A1": "報告番号", "B1": report_id, "A2": "設備番号", "B2": "EQ-001",
                         "A3": "所見", "B3": "異音あり"}.items():
        ws[coord] = value
    for coord in ("A1", "A2", "A3"):
        ws[coord].fill = PatternFill("solid", fgColor="FFD9E1F2")
    wb.save(path)
    return path


def test_the_migration_adds_the_table(app):
    """user_version が 14 になり、pattern_books の表がある（種類を消せば一緒に消える外部キー）。"""
    with app.app_context():
        conn = db.get_db()
        assert conn.execute("PRAGMA user_version").fetchone()[0] == 14
        columns = [row[1] for row in conn.execute('PRAGMA table_info("pattern_books")')]
    assert columns == ["pattern_id", "file_name", "file_hash", "book_json", "saved_at"]
    assert "pattern_books" in core.SETTINGS_TABLES


def test_creating_a_type_remembers_the_sheet_but_not_the_file(app, client, tmp_path):
    """種類を作ると、覚えたシート（JSON）が種類と一緒に入る。Excel のバイトはどこにも無い。"""
    path = _report(tmp_path / "点検表.xlsx", "R-123")
    pattern_id = create_type(client, path, "点検表")

    row = _book_row(app, pattern_id)
    assert row is not None
    assert row["file_name"] == "点検表.xlsx"
    assert row["file_hash"] == hashlib.sha256(path.read_bytes()).hexdigest()
    assert row["saved_at"][:4].isdigit()
    data = json.loads(row["book_json"])           # 文字の JSON であって、zip（xlsx）ではない
    assert not row["book_json"].startswith("PK")
    assert data["sheets"][0]["name"] == "報告書"
    assert ["報告番号", "R-123"] == [c[4] for c in data["sheets"][0]["cells"][:2]]
    assert not list(Path(app.config["UPLOAD_DIR"]).rglob("*.xlsx"))
    with app.app_context():
        meta = db.pattern_book_meta(pattern_id)
        assert meta["file_name"] == "点検表.xlsx" and meta["size"] == len(row["book_json"])
        # DB のどこにも xlsx のバイト（zip の印）は無い
        conn = db.get_db()
        for (table,) in conn.execute("SELECT name FROM sqlite_master WHERE type = 'table'"):
            for col in [r[1] for r in conn.execute(f'PRAGMA table_info("{table}")')]:
                hits = conn.execute(f'SELECT COUNT(*) FROM "{table}" WHERE CAST("{col}" AS BLOB) LIKE ?',
                                    (b"PK\x03\x04%",)).fetchone()[0]
                assert hits == 0, (table, col)


def test_reopening_clicking_and_testing_need_no_excel(app, client, tmp_path):
    """開き直し → クリック → 削除 → 見出しの手直し → 読み取りテスト → 使用開始、を Excel を送らずに通す。"""
    path = _report(tmp_path / "点検表.xlsx", "R-777")
    pattern_id = create_type(client, path, "点検表")
    core.clear()

    panel = panel_html(client, pattern_id)
    assert 'data-cell="A1"' in panel and "R-777" in panel
    assert "登録したときに覚えたシートです（点検表.xlsx、" in panel

    # クリック（Excel も合図も送らない）
    res = client.post(f"/form-types/{pattern_id}/fields",
                      data={"sheet": "報告書", "label_cell": "A1", "value_cell": "B1"})
    assert res.status_code == 200 and "「報告番号」を項目にしました" in res.get_json()["message"]
    res = client.post(f"/form-types/{pattern_id}/fields",
                      data={"sheet": "報告書", "label_cell": "A2", "value_cell": "B2"})
    assert res.status_code == 200
    html = res.get_json()["html"]
    assert "2項目中 <strong>2</strong>項目が見つかりました" in html     # 読み取りテストも覚えたシートで動く
    assert "- 報告番号: R-777" in html
    assert "登録したときに覚えたシートを実際に読み取った結果です" in html

    # 削除・見出しの手直し・使用開始のあともシートは出たまま
    res = client.post(f"/form-types/{pattern_id}/fields/equipment_id/delete")
    assert res.status_code == 200 and 'data-cell="A1"' in res.get_json()["html"]
    res = client.post(f"/form-types/{pattern_id}/fields/report_id/label", json={"name": "受付番号"})
    assert res.status_code == 200 and "- 受付番号: R-777" in res.get_json()["html"]
    res = client.post(f"/form-types/{pattern_id}/status", json={"status": "active"})
    assert res.status_code == 200 and 'data-cell="A1"' in res.get_json()["html"]
    assert not list(Path(app.config["UPLOAD_DIR"]).rglob("*"))


def test_the_sheet_is_still_there_after_a_restart(tmp_path):
    """アプリを立ち上げ直しても（メモリの覚えが全部消えても）、覚えたシートで同じ画面が出てクリックできる。"""
    config = make_config(tmp_path)
    first = create_app(config)
    client = first.test_client()
    path = _report(tmp_path / "点検表.xlsx", "R-9")
    pattern_id = create_type(client, path, "点検表")
    add_field(client, pattern_id, "報告書", "A1", "B1")
    core.clear()

    second = create_app(config)                       # 同じ DB で立ち上げ直す
    again = second.test_client()
    panel = again.get(f"/form-types/{pattern_id}/panel").get_json()["html"]
    assert 'data-cell="A1"' in panel and "R-9" in panel and "1項目中 <strong>1</strong>項目" in panel
    res = again.post(f"/form-types/{pattern_id}/fields",
                     data={"sheet": "報告書", "label_cell": "A3", "value_cell": "B3"})
    assert res.status_code == 200 and "「所見」を項目にしました" in res.get_json()["message"]
    with second.app_context():
        fields = db.load_pattern(pattern_id).fields
    assert [(f.label_cell, f.cell) for f in fields] == [("A1", "B1"), ("A3", "B3")]


def test_placing_another_excel_replaces_the_remembered_sheet(app, client, tmp_path):
    """「別のExcelに替える」（POST panel にファイル）で、覚えたシートはその Excel のものに入れ替わる。

    同じ Excel を置き直しただけなら、置いた日時も変えない。セルのクリックに添えたファイルでは入れ替えない。
    """
    first = _report(tmp_path / "一番.xlsx", "R-1")
    second = _report(tmp_path / "二番.xlsx", "R-2")
    pattern_id = create_type(client, first, "置き替え")
    before = _book_row(app, pattern_id)

    # クリックに添えた別のファイルでは、覚えたシートは変わらない
    add_field(client, pattern_id, "報告書", "A1", "B1", book=second)
    assert _book_row(app, pattern_id)["file_hash"] == before["file_hash"]

    panel = panel_html(client, pattern_id, book=second)
    assert "二番.xlsx" in panel and "R-2" in panel
    row = _book_row(app, pattern_id)
    assert row["file_name"] == "二番.xlsx" and row["file_hash"] != before["file_hash"]
    assert "R-2" in row["book_json"] and "R-1" not in row["book_json"]
    core.clear()
    reopened = panel_html(client, pattern_id)
    assert "二番.xlsx" in reopened and "R-2" in reopened and "R-1" not in reopened

    # 同じ Excel をもう一度置いても、置いた日時はそのまま
    panel_html(client, pattern_id, book=second)
    assert _book_row(app, pattern_id)["saved_at"] == row["saved_at"]
    with app.app_context():
        assert len(db.list_patterns()) == 1


def test_the_remembered_sheet_goes_with_the_type(app, client, tmp_path):
    """種類を消すと覚えたシートも消える（ほかの種類の分は残る）。"""
    one = create_type(client, _report(tmp_path / "一.xlsx"), "一")
    two = create_type(client, _report(tmp_path / "二.xlsx"), "二")
    assert client.post(f"/form-types/{one}/delete").status_code == 200
    with app.app_context():
        assert db.load_pattern_book(one) is None
        assert db.load_pattern_book(two) is not None
        assert db.get_db().execute("SELECT COUNT(*) FROM pattern_books").fetchone()[0] == 1


def test_the_remembered_sheet_survives_every_purge(app, client, tmp_path):
    """取り込みを捨てる片付け（その人の分・時間切れ・起動時）のどれでも、覚えたシートは消えない。"""
    from app.views import SESSION_ID_KEY

    with client.session_transaction() as cookie:
        cookie[SESSION_ID_KEY] = "a" * 32
    pattern_id = create_type(client, _report(tmp_path / "点検表.xlsx"), "点検表")
    add_field(client, pattern_id, "報告書", "A1", "B1")
    activate(client, pattern_id)
    add_confirmed_document(app, "済み.xlsx")
    confirmed_import(app, client)

    with app.app_context():
        assert core.purge_session("a" * 32, include_busy=True)[1] >= 1
        core.sweep_stale(0)
        core.purge_all_pending()
        assert db.get_db().execute("SELECT COUNT(*) FROM documents").fetchone()[0] == 0
        assert db.get_db().execute("SELECT COUNT(*) FROM table_imports").fetchone()[0] == 0
        row = db.load_pattern_book(pattern_id)
        assert row is not None and row["file_name"] == "点検表.xlsx"
        assert db.load_pattern(pattern_id).fields[0].field_name == "report_id"
    core.clear()
    assert 'data-cell="A1"' in panel_html(client, pattern_id)


def test_a_workbook_too_big_to_remember_is_refused(app, client, tmp_path, monkeypatch):
    """覚えられない大きさ（JSON の上限）のブックは、種類を作らずに断る。置き替えでも前の分を守る。"""
    from app import views

    path = _report(tmp_path / "点検表.xlsx")
    monkeypatch.setattr(views, "SNAPSHOT_MAX_BYTES", 100)
    res = client.post("/form-types/new", data={"name": "大きい", "book": book_part(path)},
                      content_type="multipart/form-data")
    assert res.status_code == 400
    error = res.get_json()["error"]
    assert "大きすぎて覚えられません" in error and "置き直してください" in error and "上限 0MB" in error
    with app.app_context():
        assert db.list_patterns() == []
        assert db.get_db().execute("SELECT COUNT(*) FROM pattern_books").fetchone()[0] == 0

    monkeypatch.setattr(views, "SNAPSHOT_MAX_BYTES", 4 * 1024 * 1024)
    pattern_id = create_type(client, path, "点検表")
    before = _book_row(app, pattern_id)
    monkeypatch.setattr(views, "SNAPSHOT_MAX_BYTES", 100)
    core.clear()
    res = client.post(f"/form-types/{pattern_id}/panel", data={"book": book_part(_report(tmp_path / "大.xlsx", "R-9"))},
                      content_type="multipart/form-data")
    assert res.status_code == 400 and "大きすぎて覚えられません" in res.get_json()["error"]
    assert _book_row(app, pattern_id)["file_hash"] == before["file_hash"]     # 前の覚えはそのまま


def test_a_broken_snapshot_falls_back_to_the_drop_zone(app, client, tmp_path):
    """覚えたシートが読めなくなっていても（壊れた JSON）、画面は落ちずに置き場を出す。"""
    pattern_id = create_type(client, _report(tmp_path / "点検表.xlsx"), "点検表")
    with app.app_context():
        db.get_db().execute("UPDATE pattern_books SET book_json = '{broken' WHERE pattern_id = ?", (pattern_id,))
        db.get_db().commit()
    core.clear()
    panel = panel_html(client, pattern_id)
    assert 'id="drop-book-edit"' in panel and 'data-cell="A1"' not in panel


def test_the_page_and_the_panel_say_what_is_remembered(client, tmp_path):
    """画面のどこでも「ファイルは保存しない・シートの中身は覚える・色は覚えない」を同じ言い方で言う。"""
    html = client.get("/form-types/").get_data(as_text=True)
    # 画面の上の知らせ（先頭の一文は <strong>）と、「新しく登録する」の置き場の下（views.SNAPSHOT_NOTE）
    assert html.count("Excel のファイルは保存しません。") >= 2
    assert html.count("登録したときのシートの中身（セルの番地と文字）だけを覚えておき、あとから開いたときにその画面を出します。") >= 2
    assert html.count("セルの色は覚えません（塗りの有る無しだけを、見出しの判定のために覚えます）") >= 2
    pattern_id = create_type(client, _report(tmp_path / "点検表.xlsx"), "点検表")
    res = client.post(f"/form-types/{pattern_id}/panel", data={"book": book_part(tmp_path / "点検表.xlsx")},
                      content_type="multipart/form-data")
    assert "シートの中身だけを覚えました。次からは置かずに開けます" in res.get_json()["message"]
    panel = res.get_json()["html"]
    assert "置き替えると、覚えておくシートもその Excel のものに入れ替わります" in panel
    assert "セルの番地と文字だけを覚えています（セルの色は覚えません）" in panel
