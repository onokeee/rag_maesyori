"""一覧表パイプライン（spec / normalize / checks / summaries / markdown / outputs / pipeline / store）のテスト。"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import random
import zipfile
from collections import Counter
from datetime import datetime, time as dtime, timedelta
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font

from tables import outputs, pipeline, store
from tables.checks import Issue, has_blocking, run_checks
from tables.detect import guess_layout, sample_data_rows
from tables.mapping import suggest_columns
from core.mdtext import estimate_tokens
from core.naming import md_filename
from tables.markdown import record_block, render_all
from tables.summaries import split_entity_code
from tables.normalize import ConvertContext, convert_cell, find_year_context, read_records
from tables.source import open_source
from tables.spec import (
    ColumnSpec, resolve_columns, spec_from_dict, spec_from_suggestions, spec_hash, spec_to_dict, validate_spec,
)

SAMPLES = Path(__file__).resolve().parent.parent / "samples" / "tables"
HINT = ".[legacy-R(chunk_ts=1500,chunk_ol=0)].md"


# ---- 合成データ ---------------------------------------------------------------------------

LIST_HEADERS = ["管理No", "発生日", "設備番号", "設備名", "故障区分", "現象", "処置", "停止時間(h)", "費用(円)", "担当者", "対応内容"]


def make_list_book(path: Path, include_tr004: bool = True, tr001_hours=1.5) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "2026年8月"
    ws["A1"] = "トラブル対応一覧 2026年度"
    ws.merge_cells("A1:K1")
    for i, h in enumerate(LIST_HEADERS, start=1):
        ws.cell(row=3, column=i, value=h).font = Font(bold=True)
    rows = [
        (4, ["TR-001", datetime(2026, 8, 3), "ｃｍｐ－101", "CMP研磨装置1号機", "機械", "スラリー流量低下", "フィルター交換",
             tr001_hours, "1,234", "田中", "8/3 10:00 田中：ライン停止の連絡あり。\n8/3 10:20 原点復帰で復旧。"]),
        (5, ["TR-002", "R8.8.10", "CMP-101", None, "電気", "#1 アラーム", "- 再起動", "90分", "▲500", "佐藤", None]),
        (6, [None, None, None, None, None, None, "続き：ケーブル交換", None, None, None, None]),
        (7, ["TR-003", "8/15", "CVD-201", "CVD 1号機", "電気", "真空度低下", "Oリング交換", 2, "300-", "高橋", None]),
        (8, ["TR-004", 20260820, "CVD-201", None, "機械", "+停止(セル先頭が記号)", "@cmd", 0.5, 100, "鈴木", None]),
        (9, ["TR-005", datetime(2026, 8, 21), "CVD-201", "CVD 1号機", "機械", "非表示の行", "x", 9, 0, "", None]),
        (10, ["TR-006", datetime(2026, 8, 22), "CVD-201", "CVD 1号機", "機械", "取り消し線の行", "x", 9, 0, "", None]),
        (11, ["TR-007", datetime(2026, 7, 30), "ETC-301", "エッチャー1号機", "その他", "7月の記録", "確認", 1, 0, "", None]),
    ]
    if not include_tr004:
        rows = [r for r in rows if r[1][0] != "TR-004"]
    hours = sum(1.5 if r[1][0] == "TR-002" else r[1][7] for r in rows
                if r[1][0] not in (None, "TR-005", "TR-006"))
    for row_no, values in rows:
        for col, v in enumerate(values, start=1):
            if v is not None:
                ws.cell(row=row_no, column=col, value=v)
    ws.merge_cells("D4:D5")  # 設備名の縦結合
    ws.row_dimensions[9].hidden = True
    for col in range(1, 12):
        ws.cell(row=10, column=col).font = Font(strike=True)
    ws.cell(row=12, column=1, value="小計")
    ws.cell(row=12, column=8, value=hours)
    ws.cell(row=13, column=1, value="合計")
    ws.cell(row=13, column=8, value=hours)
    ws.cell(row=15, column=1, value="※停止時間は時間で記入")
    wb.save(path)
    return path


def list_spec_dict(**markdown) -> dict:
    md = {"file_prefix": "トラブル対応一覧", "group_by": "month", "max_records_per_file": 300}
    md.update(markdown)
    return {
        "name": "トラブル対応一覧",
        "description": "保全課のトラブル対応記録",
        "columns": [
            {"key": "record_no", "display": "管理No", "headers": ["管理No"], "type": "code", "role": "key", "required": True},
            {"key": "occurred_at", "display": "発生日", "headers": ["発生日"], "type": "date", "role": "date", "required": True},
            {"key": "equipment_id", "display": "設備番号", "headers": ["設備番号"], "type": "code", "role": "entity",
             "normalize": ["nfkc", "upper"]},
            {"key": "equipment_name", "display": "設備名", "headers": ["設備名"], "type": "string", "role": "entity_label",
             "fill_down_blank": True},
            {"key": "failure_category", "display": "故障区分", "headers": ["故障区分"], "type": "enum", "role": "category",
             "allowed": ["機械", "電気"]},
            {"key": "symptom", "display": "現象", "headers": ["現象"], "type": "text", "role": "text", "md": "body"},
            {"key": "action", "display": "処置", "headers": ["処置"], "type": "text", "role": "text", "md": "body"},
            {"key": "downtime", "display": "停止時間", "headers": ["停止時間"], "type": "number", "role": "measure", "unit": "分"},
            {"key": "cost", "display": "費用", "headers": ["費用"], "type": "number", "role": "measure", "unit": "円"},
            {"key": "worker", "display": "担当者", "headers": ["担当者"], "type": "string", "role": "person"},
            {"key": "response_log", "display": "対応内容", "headers": ["対応内容"], "type": "text", "role": "log", "md": "body"},
        ],
        "record": {"key": ["record_no"], "fallback_key": ["occurred_at", "equipment_id", "symptom:20"]},
        "period": {"date_column": "occurred_at"},
        "log_stage": {"column": "response_log", "mask": []},
        "markdown": md,
    }


def read_list(path: Path, spec):
    source = open_source(path, path.name)
    layout = guess_layout(source, "2026年8月")
    return read_records(source, {"sheet": "2026年8月"}, layout, spec)


def august(records) -> list[dict]:
    return [r.to_dict() for r in records if str(r.values.get("occurred_at", "")).startswith("2026-08")]


def _no_blank_inside_records(text: str) -> bool:
    lines = text.split("\n")
    for i, line in enumerate(lines[:-1]):
        if line == "":
            nxt = lines[i + 1]
            prev = lines[i - 1] if i else ""
            if not (nxt.startswith("## ") or prev.startswith("# ")):
                return False
    return True


# ---- 値の変換 -------------------------------------------------------------------------------

@pytest.mark.parametrize("type_, value, text, expected, flag", [
    ("date", None, "令和8年8月3日", "2026-08-03", None),
    ("date", None, "R8.8.3", "2026-08-03", None),
    ("date", None, "H31.4.30", "2019-04-30", None),
    ("date", None, "昭和64年1月7日", "1989-01-07", None),
    ("date", None, "令和元年5月1日", "2019-05-01", None),
    ("date", None, "20260803", "2026-08-03", None),
    ("date", 20260803, "20260803", "2026-08-03", None),
    ("date", None, "2026/8/3(月)", "2026-08-03", None),
    ("date", None, "８/５", "2026-08-05", "year_inferred"),
    ("date", None, "2/5", "2027-02-05", "year_inferred"),
    ("date", 46237.0, "46237", "2026-08-03", None),
    ("datetime", None, "2026/08/03 14:20:00", "2026-08-03 14:20", None),
    ("datetime", None, "20260803 142000", "2026-08-03 14:20", None),
    ("datetime", None, "2026年8月3日 14時20分", "2026-08-03 14:20", None),
    ("datetime", datetime(2026, 8, 3, 0, 0), "", "2026-08-03", None),
    ("time", None, "142000", "14:20", None),
    ("time", None, "1420", "14:20", None),
    ("time", 0.5, "0.5", "12:00", None),
    ("time", timedelta(hours=1, minutes=5), "1:05", "01:05", None),
    ("time", dtime(9, 5), "09:05", "09:05", None),
])
def test_convert_dates_and_times(type_, value, text, expected, flag):
    col = ColumnSpec("c", "列", type=type_)
    cctx = ConvertContext(fiscal_year=2026, fiscal_start=4)
    out, error, got_flag = convert_cell(value, text, col, cctx)
    assert error is None
    assert out == expected
    assert got_flag == flag


@pytest.mark.parametrize("value, text, unit, header_unit, fmt, expected", [
    (None, "1,234", "", "", None, 1234),
    (None, "１，２３４", "", "", None, 1234),
    (None, "△500", "", "", None, -500),
    (None, "▲1,000", "円", "", None, -1000),
    (None, "300-", "", "", None, -300),
    (None, "98.5%", "%", "", None, 98.5),
    (0.985, "0.985", "%", "", "0.0%", 98.5),
    (None, "1.5h", "分", "", None, 90),
    (None, "2時間", "分", "", None, 120),
    (1.5, "1.5", "分", "h", None, 90),
    (None, "90分", "分", "h", None, 90),
    (timedelta(hours=2, minutes=30), "2:30", "分", "", None, 150),
    (timedelta(minutes=90), "1:30", "時間", "", None, 1.5),
    (None, "¥12,000", "円", "", None, 12000),
    (None, "1.23E+3", "", "", None, 1230),
    (3.0000000001, "3", "", "", None, 3),
])
def test_convert_numbers(value, text, unit, header_unit, fmt, expected):
    col = ColumnSpec("c", "列", type="number", unit=unit)
    out, error, _flag = convert_cell(value, text, col, ConvertContext(), fmt, header_unit)
    assert error is None
    assert out == expected


def test_convert_blank_na_errors_codes_and_maps():
    cctx = ConvertContext(na_tokens=frozenset({"-", "N/A"}))
    num = ColumnSpec("n", "数", type="number", unit="分")
    assert convert_cell(None, "－", num, cctx) == (None, None, None)
    assert convert_cell("#N/A", "#N/A", num, cctx) == (None, None, "error_value")
    out, error, _ = convert_cell(None, "約3", num, cctx)
    assert error and out == "約3"
    out, error, _ = convert_cell(None, "3件", num, cctx)
    assert error  # 分に換算できない単位
    date_col = ColumnSpec("d", "日付", type="date")
    out, error, flag = convert_cell(None, "8/5", date_col, ConvertContext())
    assert flag == "year_missing" and error
    code = ColumnSpec("k", "コード", type="code", normalize=["nfkc", "upper"])
    assert convert_cell(123, "123", code, cctx, "00000")[0] == "00123"
    assert convert_cell(None, "ｃｍｐ－１０１", code, cctx)[0] == "CMP-101"
    enum = ColumnSpec("e", "区分", type="enum", value_map={"電気系": "電気"})
    assert convert_cell(None, "電気系", enum, cctx)[0] == "電気"
    text = ColumnSpec("t", "文章", type="text")
    assert convert_cell(None, "  ＡＢＣ　 1\r\n  2行目  ", text, cctx)[0] == "ABC 1\n2行目"


def test_year_context():
    assert find_year_context([("タイトル行", "2026年度 故障一覧")])["fiscal_year"] == 2026
    assert find_year_context([("シート名", "FY25")])["fiscal_year"] == 2025
    assert find_year_context([("x", "令和8年度")])["fiscal_year"] == 2026
    ym = find_year_context([("ファイル名", "故障履歴_2026-08")])
    assert (ym["year"], ym["month"], ym["source"]) == (2026, 8, "ファイル名")
    assert find_year_context([("ファイル名", "T1_トラブル対応一覧_2023-2026")])["year"] is None
    assert find_year_context([("a", ""), ("b", "202608")])["month"] == 8


# ---- spec ------------------------------------------------------------------------------------

def test_spec_roundtrip_validate_and_resolve():
    spec = spec_from_dict(list_spec_dict())
    assert validate_spec(spec) == []
    again = spec_from_dict(json.loads(json.dumps(spec_to_dict(spec), ensure_ascii=False)))
    assert spec_hash(again) == spec_hash(spec)
    assert again.summaries()[0].id == "month"
    assert spec.date_key == "occurred_at" and spec.file_prefix == "トラブル対応一覧"

    bad = spec_from_dict({"name": "", "columns": [
        {"key": "1bad", "display": "x", "type": "weird", "role": "person"},
        {"key": "a", "display": "", "type": "number", "md": "nope"},
        {"key": "a", "display": "dup"},
    ], "record": {"key": ["missing"]}, "period": {"date_column": "a"},
        "markdown": {"group_by": "entity_month", "summaries": [{"id": "month", "metrics": ["sum:nothing", "median"]}]},
        "log_stage": {"column": "zzz"}})
    errors = validate_spec(bad)
    joined = "\n".join(errors)
    for fragment in ("設定名", "1bad", "重複", "表示名", "型「weird」", "mdでの扱い", "記録キーの列「missing」",
                     "日付", "設備×月", "sum:nothing", "median", "zzz"):
        assert fragment in joined, fragment

    res = resolve_columns(spec, ["管理No", "発生日", "設備番号", "設備名", "故障区分", "現象", "処置内容", "停止時間(h)",
                                 "費用(円)", "担当者", "新しい列"])
    assert res.positions["downtime"] == 7 and res.positions["cost"] == 8
    assert "action" not in res.positions and "処置" in res.missing
    assert res.unused_headers == ["処置内容", "新しい列"]
    assert res.missing_required == []


def test_spec_from_suggestions(tmp_path):
    path = make_list_book(tmp_path / "list.xlsx")
    source = open_source(path, path.name)
    layout = guess_layout(source, "2026年8月")
    suggestions = suggest_columns(layout.headers, sample_data_rows(source, "2026年8月", layout))
    spec = spec_from_suggestions("トラブル", layout, suggestions, {"description": "説明", "group_by": "entity_month"})
    assert validate_spec(spec) == []
    keys = [c.key for c in spec.columns]
    assert keys[:4] == ["record_no", "occurred_at", "equipment_id", "equipment_name"]
    assert spec.record["key"] == ["record_no"]
    assert spec.record["fallback_key"][0] == "occurred_at"
    assert spec.period == {"date_column": "occurred_at"}
    assert spec.log_stage is not None and spec.log_stage.column == "response_log"
    assert spec.markdown["group_by"] == "entity_month"


# ---- 一覧: 読み込み → チェック → md → zip ---------------------------------------------------------

def test_list_read_normalize_and_checks(tmp_path):
    path = make_list_book(tmp_path / "list.xlsx")
    spec = spec_from_dict(list_spec_dict())
    records, row_issues, stats = read_list(path, spec)
    by_key = {r.key: r for r in records}
    assert list(by_key) == ["TR-001", "TR-002", "TR-003", "TR-004", "TR-007"]
    r1, r2, r3, r4 = (by_key[k].values for k in ("TR-001", "TR-002", "TR-003", "TR-004"))
    assert r1["equipment_id"] == "CMP-101" and by_key["TR-001"].originals["equipment_id"] == "ｃｍｐ－101"
    assert r1["downtime"] == 90 and r1["cost"] == 1234 and r1["occurred_at"] == "2026-08-03"
    assert r2["occurred_at"] == "2026-08-10" and r2["equipment_name"] == "CMP研磨装置1号機"  # 縦結合を埋める
    assert r2["action"] == "- 再起動\n続き:ケーブル交換"  # 継続行を連結（NFKC）
    assert r2["downtime"] == 90 and r2["cost"] == -500
    assert r3["occurred_at"] == "2026-08-15" and "年を" in by_key["TR-003"].warnings[0]
    assert r3["cost"] == -300 and r3["equipment_name"] == "CVD 1号機"
    assert r4["occurred_at"] == "2026-08-20" and r4["equipment_name"] == "CVD 1号機"  # 空欄＝上と同じ
    assert stats.filled_down == {"equipment_name": 1}
    assert stats.continuation_merged == 1
    assert stats.excluded == {"非表示の行": 1, "取り消し線の行": 1, "小計": 1, "合計": 1}
    assert stats.year_inferred == 1 and stats.year_context["fiscal_year"] == 2026
    assert stats.records == 5 and stats.months == {"2026-07": 1, "2026-08": 4}
    assert len(stats.reconcile) == 2 and all(rc["expected"] == rc["actual"] == 390 for rc in stats.reconcile)
    assert row_issues == []

    issues = run_checks(records, spec, stats)
    codes = [i.code for i in issues]
    assert "year_inferred" in codes and "not_allowed" in codes
    assert codes.count("excluded_rows") == 2
    assert not has_blocking(issues)

    # 見出しが足りない・型エラーが多い → 確定を止める
    broken = spec_from_dict(list_spec_dict())
    broken.column("symptom").required = True
    broken.column("symptom").headers = ["不具合内容"]
    broken.column("symptom").display = "不具合"
    broken.column("failure_category").type = "number"
    _recs, _row_issues, st2 = read_list(path, broken)
    issues2 = run_checks(_recs, broken, st2)
    assert {"required_missing", "type_error_rate"} <= {i.code for i in issues2 if i.level == "error"}
    assert len(_row_issues) == 5 and _row_issues[0].column == "故障区分"


def test_list_markdown_and_determinism(tmp_path):
    path = make_list_book(tmp_path / "list.xlsx")
    spec = spec_from_dict(list_spec_dict())
    records, _issues, _stats = read_list(path, spec)
    aug = august(records)
    ai = {
        "TR-001": {"status": "ok", "result": {
            "entries": [{"id": "e1", "segs": ["s1"], "t": ["連絡", "初動"]}],
            "incident": {"root_cause": {"v": "フィルター目詰まり", "certainty": "確定"},
                         "permanent_actions": [{"v": "フィルター交換"}],
                         "parts": [{"name": "フィルター", "model": "FL-10", "qty_q": "×1"}]}}},
        "TR-002": {"status": "flagged", "result": {"incident": {"root_cause": {"v": "出してはいけない"}}}},
    }
    files = render_all(spec, aug, ai, {"coverage": {"start": "2026-08", "end": "2026-08"}})
    names = [f.name for f in files]
    assert names == [
        "トラブル対応一覧_00_データセット説明.md",
        "トラブル対応一覧_2026-08" + HINT,   # 分割ヒントは新しい取り込み設定の既定（記録ファイルだけに付く）
        "トラブル対応一覧_集計_月次_2026-08.md",
        "トラブル対応一覧_集計_設備別_CMP-101_2026年度.md",
        "トラブル対応一覧_集計_設備別_CVD-201_2026年度.md",
    ]
    rec = files[1].text
    assert rec.startswith("# トラブル対応一覧 2026年8月の記録\n\n- データ種別: トラブル対応一覧（1行＝1件）の記録\n"
                          "- 対象期間: 2026-08-01〜2026-08-31\n- このファイルの記録: 4件（2026年8月の全件）\n\n## ")
    block1 = rec.split("\n\n")[2]
    assert block1.split("\n")[0] == "## 【TR-001】CMP研磨装置1号機（CMP-101）スラリー流量低下｜2026-08-03"
    for line in ("- 管理No: TR-001", "- 発生日: 2026-08-03（2026年8月）", "- 設備: CMP研磨装置1号機（CMP-101）",
                 "- 停止時間: 90分", "- 費用: 1,234円", "- 対応の要点（AI抽出）:", "  原因: フィルター目詰まり（確定）",
                 "  恒久処置: フィルター交換", "  使用部品: フィルター FL-10 ×1", "- 対応の時系列:",
                 "- 出典: list.xlsx（管理No TR-001）"):
        assert line in block1.split("\n"), line
    timeline = [ln for ln in block1.split("\n") if ln.startswith("  1. ")]
    assert timeline and timeline[0].startswith("  1. 2026-08-03 10:00［連絡・初動］") and "CMP研磨装置1号機" in timeline[0]
    assert "  2. 2026-08-03 10:20" in block1
    assert "担当者" not in rec and "田中｜" in rec  # 担当者列は出さない（時系列の記入者は原文由来）
    assert "出してはいけない" not in rec
    assert "- 処置:\n  \\- 再起動\n  続き:ケーブル交換" in rec
    assert "- 現象: +停止(セル先頭が記号)" in rec
    assert "|" not in rec.replace("｜", "")  # パイプ表を使わない
    assert _no_blank_inside_records(rec)
    assert "\r" not in rec and rec.endswith("\n") and not rec.endswith("\n\n")

    month = files[2].text
    assert "- 2026年8月のトラブル対応一覧の記録は4件です。" in month
    assert "- 2026年8月の停止時間の合計は330分（5.5時間）です。" in month
    assert "- 1位: CMP研磨装置1号機（CMP-101） 停止時間 180分、2件、主な故障区分: 機械" in month
    assert "- 2位: CVD 1号機（CVD-201） 停止時間 150分、2件、主な故障区分: 機械" in month
    fy = files[4].text
    assert "- CVD-201の2026年度（2026年8月〜2026年8月）の記録は2件です。" in fy
    assert "- 停止時間の合計は150分（2.5時間）、1件あたり平均75分です。" in fy
    card = files[0].text
    assert "- 取り込み範囲: 2026-08-03〜2026-08-20（記録 4件）" in card and "- 担当者" not in card

    # 決定的: 入力の順番を変えても同じバイト列
    shuffled = list(aug)
    random.Random(3).shuffle(shuffled)
    again = render_all(spec, shuffled, ai, {"coverage": {"start": "2026-08", "end": "2026-08"}})
    assert [hashlib.sha256(f.data).hexdigest() for f in again] == [hashlib.sha256(f.data).hexdigest() for f in files]

    # 分割・人名を出す設定、設備×月（ヒントは既定のオンのまま）
    spec2 = spec_from_dict(list_spec_dict(max_records_per_file=3, omit_person=False))
    files2 = render_all(spec2, aug, {}, {})
    part1 = next(f for f in files2 if f.name == "トラブル対応一覧_2026-08" + HINT)
    assert "トラブル対応一覧_2026-08_part2" + HINT in [f.name for f in files2]
    assert "- このファイルの記録: 3件（2026年8月の全 4件のうち 1〜3件目）" in part1.text
    assert "- 担当者: 田中" in part1.text
    files3 = render_all(spec_from_dict(list_spec_dict(group_by="entity_month")), aug, {}, {})
    ent = next(f for f in files3 if f.name == "トラブル対応一覧_CMP-101_2026-08" + HINT)
    assert "- このファイルの記録: 2件（CMP-101の2026年8月の全件）" in ent.text

    # 7月と8月の両方がある取り込みは月ごとのファイルになる
    all_names = [f.name for f in render_all(spec, [r.to_dict() for r in records], {}, {})]
    assert "トラブル対応一覧_2026-07" + HINT in all_names and "トラブル対応一覧_2026-08" + HINT in all_names


def test_zip_csv_and_formula_guard():
    assert outputs.guard_formula("=HYPERLINK(\"x\")") == "'=HYPERLINK(\"x\")"
    for s in ("+1", "-2", "@SUM", "\tx", "\rx"):
        assert outputs.guard_formula(s).startswith("'")
    assert outputs.guard_formula("通常の文字") == "通常の文字"
    assert outputs.guard_formula(-5) == "-5" and outputs.guard_formula(None) == ""
    data = outputs.issues_csv([Issue("warning", "x", "=1+1", row=3, column="-列")])
    assert data.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    assert rows[1] == ["警告", "x", "3", "'-列", "'=1+1"]
    zipped = outputs.build_zip([("b.md", b"B\n"), ("a.md", b"A\n")], {"問題一覧.csv": data})
    assert zipped == outputs.build_zip([("a.md", b"A\n"), ("b.md", b"B\n")], {"問題一覧.csv": data})
    with zipfile.ZipFile(io.BytesIO(zipped)) as zf:
        assert zf.namelist() == ["RAG投入用/a.md", "RAG投入用/b.md", "管理用_RAGには入れない/問題一覧.csv"]


# ---- CSV ---------------------------------------------------------------------------------------

def test_system_csv_quirks(tmp_path):
    text = ("出力日時,2026/09/01 10:00:00\n抽出条件,期間=2026/08\n\n"
            "管理番号,発生日,発生時刻,設備コード,停止時間(分),費用(円),現象\n"
            '="00123",20260803,142000,cmp-101,95,"1,234",スラリー流量低下\n'
            '="00124",20260804,090500,CMP-101,30-,500,"複数行の\n現象"\n'
            "合計件数,2\n")
    path = tmp_path / "システム出力_2026-08.csv"
    path.write_bytes(text.encode("cp932"))
    spec = spec_from_dict({
        "name": "故障", "columns": [
            {"key": "record_no", "display": "管理番号", "headers": ["管理番号"], "type": "code", "role": "key"},
            {"key": "occurred_at", "display": "発生日", "headers": ["発生日"], "type": "date", "role": "date"},
            {"key": "occurred_time", "display": "発生時刻", "headers": ["発生時刻"], "type": "time"},
            {"key": "equipment_id", "display": "設備コード", "headers": ["設備コード"], "type": "code", "role": "entity",
             "normalize": ["nfkc", "upper"]},
            {"key": "downtime", "display": "停止時間", "headers": ["停止時間"], "type": "number", "role": "measure", "unit": "分"},
            {"key": "cost", "display": "費用", "headers": ["費用"], "type": "number", "role": "measure", "unit": "円"},
            {"key": "symptom", "display": "現象", "headers": ["現象"], "type": "text", "role": "text"},
        ],
        "record": {"key": ["record_no"]},
    })
    source = open_source(path, path.name)
    layout = guess_layout(source, path.name)
    records, row_issues, stats = read_records(source, {}, layout, spec)
    assert [r.values for r in records] == [
        {"record_no": "00123", "occurred_at": "2026-08-03", "occurred_time": "14:20", "equipment_id": "CMP-101",
         "downtime": 95, "cost": 1234, "symptom": "スラリー流量低下"},
        {"record_no": "00124", "occurred_at": "2026-08-04", "occurred_time": "09:05", "equipment_id": "CMP-101",
         "downtime": -30, "cost": 500, "symptom": "複数行の\n現象"},
    ]
    assert records[1].source == {"file": path.name, "sheet": "", "row": 6}
    assert stats.excluded == {"合計": 1} and row_issues == []
    files = render_all(spec, [r.to_dict() for r in records], {}, {})
    rec = next(f for f in files if f.name == "故障_2026-08" + HINT).text
    assert "- 発生日: 2026-08-03 14:20（2026年8月）" in rec and "発生時刻" not in rec
    assert "## 【00124】CMP-101 複数行の｜2026-08-04" in rec
    assert "- 現象:\n  複数行の\n  現象\n" in rec


# ---- DB・ジョブを通した流れ ------------------------------------------------------------------------

class FakeCtx:
    def __init__(self):
        self.progress_calls = []

    def progress(self, **kw):
        self.progress_calls.append(kw)

    def check_cancel(self):
        pass


def _new_import(app, spec_dict=None, template_id=None, name="list.xlsx"):
    upload_dir = Path(app.config["UPLOAD_DIR"]) / "tables"
    upload_dir.mkdir(parents=True, exist_ok=True)
    path = make_list_book(upload_dir / name)
    if template_id is None:
        spec = spec_from_dict(spec_dict or list_spec_dict())
        template_id, version_id = store.create_template(spec.name, spec)
    else:
        version_id = store.get_template(template_id)["current_version_id"]
    file_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    import_id = store.create_import(name, file_hash, f"tables/{name}", {"sheet": "2026年8月"}, template_id, version_id)
    return template_id, import_id


def test_pipeline_read_render_download(app):
    with app.app_context():
        template_id, import_id = _new_import(app)
        ctx = FakeCtx()
        result = pipeline.run_read(ctx, import_id)
        assert result["records"] == 5 and result["error"] == 0
        assert any(c.get("phase") == "読み込み" for c in ctx.progress_calls)
        imp = store.get_import(import_id)
        assert imp["status"] == "preview" and imp["stats"]["records"] == 5
        assert Path(imp["rows_path"]).exists() and Path(imp["issues_path"]).exists()
        assert len(pipeline.load_rows(import_id)) == 5 and len(pipeline.load_rows(import_id, 1, 2)) == 2

        listing = pipeline.preview_files(import_id, imp, pipeline.spec_for_import(imp))
        assert "トラブル対応一覧_2026-07" + HINT in [f["name"] for f in listing]
        preview_dir = pipeline.import_files(import_id)["preview"]
        assert pipeline.md_text(preview_dir, "トラブル対応一覧_2026-08" + HINT).startswith("# トラブル対応一覧 2026年8月")
        assert pipeline.md_text(preview_dir, "../x.md") is None

        done = pipeline.run_render(FakeCtx(), import_id)
        imp = store.get_import(import_id)
        assert imp["status"] == "confirmed" and imp["confirmed_at"] and done["records"] == 5
        assert len(pipeline.md_paths(import_id)) == done["files"]
        assert store.get_version(imp["template_version_id"])["used"] == 1
        data = pipeline.build_download(import_id, imp, pipeline.spec_for_import(imp))
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            names = zf.namelist()
            assert len([n for n in names if n.startswith("RAG投入用/")]) == done["files"]
            for extra in ("正規化データ.csv", "問題一覧.csv", "取込レポート.csv"):
                assert f"管理用_RAGには入れない/{extra}" in names
            normalized = list(csv.reader(io.StringIO(zf.read("管理用_RAGには入れない/正規化データ.csv").decode("utf-8-sig"))))
            assert normalized[0][:3] == ["記録キー", "管理No", "発生日"]
            assert "'+停止(セル先頭が記号)" in next(r for r in normalized if r[0] == "TR-004")

        # 確定に使った版は上書きせず、新しい版にする
        spec = store.get_template(template_id)["spec"]
        spec.description = "v2"
        v2 = store.save_template_version(template_id, spec)
        assert v2 != imp["template_version_id"] and store.save_template_version(template_id, spec) == v2


def test_pipeline_blocking_and_broken_file(app):
    with app.app_context():
        spec_dict = list_spec_dict()
        spec_dict["columns"][5].update(required=True, headers=["不具合内容"], display="不具合内容")
        template_id, import_id = _new_import(app, spec_dict)
        pipeline.run_read(FakeCtx(), import_id)
        assert store.get_import(import_id)["stats"]["issue_counts"]["error"] == 1
        assert has_blocking(pipeline.load_issues(import_id))

        bad_path = Path(app.config["UPLOAD_DIR"]) / "tables" / "broken.xlsx"
        bad_path.write_bytes(b"not a zip")
        bad = store.create_import("broken.xlsx", "x", "tables/broken.xlsx", {}, template_id,
                                  store.get_template(template_id)["current_version_id"])
        with pytest.raises(Exception):
            pipeline.run_read(FakeCtx(), bad)
        failed = store.get_import(bad)
        assert failed["status"] == "failed" and "Excel" in failed["stats"]["error"]


def test_read_job_runs_in_worker(app):
    from core import jobs

    with app.app_context():
        _tid, import_id = _new_import(app)
        job_id = pipeline.start_read_job(import_id)
        job = jobs.wait_job(job_id, timeout=60)
        assert job["status"] == "done", job
        assert job["result"]["records"] == 5
        imp = store.get_import(import_id)
        assert imp["status"] == "preview" and imp["job_id"] == job_id


# ---- 取り込みの控え（source_cache）と、作り直さない工夫 ----------------------------------------------

def _must_not_open(_stats):
    raise AssertionError("元のファイルを開き直した")


def test_scan_memo_gives_same_layout(tmp_path):
    path = make_list_book(tmp_path / "list.xlsx")
    source = open_source(path, path.name)
    memo: dict = {}
    auto = guess_layout(source, "2026年8月", scan_memo=memo)
    assert auto == guess_layout(source, "2026年8月") and len(memo) == 1
    manual = guess_layout(source, "2026年8月", header_rows=auto.header_rows, scan_memo=memo)
    assert manual == guess_layout(source, "2026年8月", header_rows=auto.header_rows) and len(memo) == 1
    ended = guess_layout(source, "2026年8月", header_rows=auto.header_rows, data_end=8, scan_memo=memo)
    assert ended == guess_layout(source, "2026年8月", header_rows=auto.header_rows, data_end=8) and len(memo) == 2


def test_import_source_cache_reuses_results_without_reopening(tmp_path):
    from tables.detect import looks_like_list
    from tables.source_cache import CACHE_FILE, ImportSource

    path = make_list_book(tmp_path / "list.xlsx")
    opened = []

    def opener(stats):
        opened.append(stats)
        return open_source(path, path.name, {"sheet_stats": stats} if stats else None)

    direct = open_source(path, path.name)
    sheet = "2026年8月"
    first = ImportSource(tmp_path / "imp", "key1", "excel", path.name, opener)
    auto = first.layout(sheet)
    assert auto == guess_layout(direct, sheet)
    assert first.layout(sheet, max_scan_rows=200) == guess_layout(direct, sheet, max_scan_rows=200)
    assert first.looks_like_list(sheet) == looks_like_list(direct, sheet)
    samples = first.sample_rows(sheet, auto, 200)
    assert samples == sample_data_rows(direct, sheet, auto, 200)
    assert list(first.rows(sheet, 1, 60)) == list(direct.rows(sheet, 1, 60)) and first.sheets() == direct.sheets()
    assert len(opened) == 1 and (tmp_path / "imp" / CACHE_FILE).exists()

    # 同じ鍵なら元のファイルを開かずに同じ結果（自動判定と同じ見出し行を指定し直しても全行をなめ直さない）
    again = ImportSource(tmp_path / "imp", "key1", "excel", path.name, _must_not_open)
    assert again.layout(sheet) == auto and again.sample_rows(sheet, auto, 200) == samples
    assert list(again.rows(sheet, 5, 3)) == list(direct.rows(sheet, 5, 3)) and again.sheets() == direct.sheets()
    assert again.looks_like_list(sheet) == looks_like_list(direct, sheet)
    assert again.layout(sheet, header_rows=auto.header_rows) == guess_layout(direct, sheet, header_rows=auto.header_rows)

    # 鍵（ファイル・読み込み設定）が変われば作り直す。Excel を開き直すときはシートの大きさの控えを渡す
    changed = ImportSource(tmp_path / "imp", "key2", "excel", path.name, opener)
    assert changed.layout(sheet) == auto and len(opened) == 2 and opened[1] is None
    reopened = ImportSource(tmp_path / "imp", "key2", "excel", path.name, opener)
    assert list(reopened.rows(sheet)) == list(direct.rows(sheet)) and opened[2] == direct.sheet_stats()


def test_import_source_does_not_recreate_a_deleted_folder(tmp_path):
    """取り込みを消したあとに、開いたままの画面の処理が控えを書いてもフォルダごと復活させない。

    控えには元の表の先頭行・列の見本の行がそのまま入るので、復活すると「ダウンロードしたら消える」
    （design.md 3.3）が破れる。
    """
    import shutil

    from tables.source_cache import CACHE_FILE, ImportSource

    path = make_list_book(tmp_path / "list.xlsx")
    directory = tmp_path / "imports" / "1"
    source = ImportSource(directory, "key1", "excel", path.name, lambda stats: open_source(path, path.name))
    source.sheets()
    assert (directory / CACHE_FILE).exists()

    shutil.rmtree(directory)                 # ダウンロード（purge_table_import）でフォルダごと消えた
    source.layout("2026年8月")                # 別のタブの画面処理が続きを控えに書こうとする
    assert not directory.exists()            # 復活しない


def test_rows_page_and_render_reuses_preview(app, monkeypatch):
    import tables.source_cache as source_cache

    with app.app_context():
        _tid, import_id = _new_import(app)
        pipeline.run_read(FakeCtx(), import_id)
        rows = pipeline.load_rows(import_id)
        assert pipeline.load_rows_page(import_id, 1, 2) == (rows[1:3], 5)
        assert pipeline.load_rows_page(import_id, 100, 10) == ([], 5)

        # 2回目の読み込みは、控えた表の形を使う
        monkeypatch.setattr(source_cache, "guess_layout", lambda *a, **k: pytest.fail("表の形を推定し直した"))
        pipeline.run_read(FakeCtx(), import_id)
        monkeypatch.undo()
        assert pipeline.load_rows(import_id) == rows

        imp = store.get_import(import_id)
        spec = pipeline.spec_for_import(imp)
        listing = pipeline.preview_files(import_id, imp, spec)
        preview = pipeline.import_files(import_id)["preview"]
        monkeypatch.setattr(pipeline, "render_files", lambda *a, **k: pytest.fail("Markdown を作り直した"))
        done = pipeline.run_render(FakeCtx(), import_id)
        monkeypatch.undo()
        md = {p.name: p.read_bytes() for p in pipeline.md_paths(import_id)}
        assert done["files"] == len(listing) == len(md)
        assert md == {f["name"]: (preview / f["name"]).read_bytes() for f in listing}
        assert md == {f.name: f.data for f in pipeline.render_files(import_id, store.get_import(import_id), spec)}

        # 読み込み直したあとはプレビューが消えるので、確定のときに作る
        pipeline.run_read(FakeCtx(), import_id)
        calls = []
        original = pipeline.render_files
        monkeypatch.setattr(pipeline, "render_files", lambda *a, **k: calls.append(1) or original(*a, **k))
        pipeline.run_render(FakeCtx(), import_id)
        assert calls == [1] and {p.name: p.read_bytes() for p in pipeline.md_paths(import_id)} == md


# ---- samples（実ファイル） ------------------------------------------------------------------------

@pytest.mark.samples
def test_sample_t1_end_to_end():
    path = SAMPLES / "T1_トラブル対応一覧_2023-2026.xlsx"
    if not path.exists():
        pytest.skip(f"{path.name} がありません")
    source = open_source(path, path.name)
    layout = guess_layout(source, "トラブル一覧")
    spec = spec_from_suggestions("トラブル対応一覧", layout,
                                 suggest_columns(layout.headers, sample_data_rows(source, "トラブル一覧", layout)))
    assert validate_spec(spec) == []
    records, row_issues, stats = read_records(source, {"sheet": "トラブル一覧"}, layout, spec)
    assert stats.records == 8095 and len({r.key for r in records}) == 8095
    assert stats.type_errors == {} and row_issues == []
    assert Counter(str(r.values.get("occurred_at"))[:4] for r in records) == {"2023": 1534, "2024": 2195, "2025": 2553, "2026": 1813}
    assert not has_blocking(run_checks(records, spec, stats))
    files = render_all(spec, [r.to_dict() for r in records], {}, {})
    assert sum(f.text.count("\n## ") for f in files if f.kind == "records") == 8095
    first = next(f for f in files if f.name == "トラブル対応一覧_2023-04" + HINT)
    assert "- 発生日: 2023-04-01 02:21（2023年4月）" in first.text and "担当者" not in first.text
    assert _no_blank_inside_records(first.text)


# ---- LightRAG オフライン評価の反映（時系列の重複・記録の上限・コード列・設備の値の揺れ） ----

def _rec(key: str, **values) -> dict:
    return {"key": key, "values": values, "source": {"file": "x.xlsx", "row": 4}}


def test_timeline_drops_sentences_that_repeat_other_columns():
    """時系列の本文から、同じ記録の原因・処置内容と同じ文を省く（全部同じなら「（処置と同じ）」に縮める）。"""
    spec = spec_from_dict(list_spec_dict())
    log = ("8/3 10:00 田中：スラリー流量低下アラームで研磨停止。保全へ連絡\n"
           "8/3 11:00 田中：フィルター交換、流量の再校正を実施\n"
           "8/3 12:00 田中：復旧を確認。流量の再校正を実施\n")
    rec = _rec("TR-001", record_no="TR-001", occurred_at="2026-08-03", equipment_id="CMP-101",
               symptom="スラリー流量低下アラームで研磨停止",
               action="フィルター交換、流量の再校正を実施", response_log=log)
    lines = record_block(rec, spec)
    timeline = [ln.strip() for ln in lines[lines.index("- 対応の時系列:") + 1:]]
    assert timeline[0].endswith(": 保全へ連絡")          # 現象と同じ文は省く
    assert timeline[1].endswith(": （処置と同じ）")      # 全部が処置と同じ
    assert timeline[2].endswith(": 復旧を確認。")        # 重複しない文は残す
    # 元の列はそのまま残る
    assert "- 現象: スラリー流量低下アラームで研磨停止" in lines and "- 処置: フィルター交換、流量の再校正を実施" in lines

    # 短い文（6文字未満）は偶然の一致を避けるため省かない
    rec2 = _rec("TR-002", record_no="TR-002", occurred_at="2026-08-04", equipment_id="CMP-101",
                action="完了", response_log="8/4 10:00 田中：完了\n")
    assert any("完了" in ln for ln in record_block(rec2, spec)[-4:])


def test_record_is_cut_to_the_token_budget_including_the_ai_points():
    """推定トークンが上限（ヒントの chunk_ts − 100）を超える記録は、要点も含めた合計で判定して時系列を切る。"""
    from tables.markdown import RECORD_TOKEN_BUDGET

    assert RECORD_TOKEN_BUDGET == 1400
    log = "".join(f"8/3 {9 + i // 6:02d}:{(i % 6) * 10:02d} 田中：{'対応の記録です。' * 12}\n" for i in range(40))
    rec = _rec("TR-003", record_no="TR-003", occurred_at="2026-08-03", equipment_id="CMP-101", response_log=log)
    lines = record_block(rec, spec_from_dict(list_spec_dict()))
    assert estimate_tokens("\n".join(lines)) <= RECORD_TOKEN_BUDGET
    note = [ln for ln in lines if "管理用の正規化CSVに収録" in ln]
    assert len(note) == 1
    kept = len([ln for ln in lines if ln.startswith("  ") and ln.strip()[0].isdigit()])
    assert kept < 20  # 設定の20件よりさらに減らして収める


def test_opaque_code_columns_are_suggested_as_not_output():
    """「状態コード: 9」のような、置き換え表なしでは意味の分からない列は既定で md=omit にする。"""
    headers = ["管理番号", "状態コード", "再発フラグ", "発見区分コード", "ラインコード", "現象", "設備名"]
    rows = [[f"MS-{i:04d}", "9", "0", "H01", "L2", f"アラーム{i}が出た", "CMP研磨装置1号機"] for i in range(30)]
    by_header = {s.header: s for s in suggest_columns(headers, rows)}
    assert [h for h, s in by_header.items() if s.md == "omit"] == ["状態コード", "再発フラグ", "発見区分コード", "ラインコード"]
    assert by_header["管理番号"].md != "omit" and by_header["現象"].md != "omit"

    # 値の種類が多い（20種類超）・見出しがコードらしくない列は出す
    rows2 = [[f"MS-{i:04d}", f"S{i:03d}", "処置済み", "H01", "L2", "x", "y"] for i in range(30)]
    by_header2 = {s.header: s for s in suggest_columns(["管理番号", "状態コード", "状況", "発見区分コード", "工程", "a", "b"], rows2)}
    assert by_header2["状態コード"].md != "omit"   # S000〜S029 は30種類あるのでコード表がなくても見分けが付く
    assert by_header2["状況"].md != "omit" and by_header2["工程"].md != "omit"


def test_entity_code_and_name_in_one_column_are_split_before_grouping():
    """設備名の列がない台帳で「ETC-302(OXIDEエッチャ 2号機)」と「ETC-302」が別設備に割れないようにする。"""
    assert split_entity_code("ETC-302(OXIDEエッチャ 2号機)") == ("ETC-302", "OXIDEエッチャ 2号機")
    assert split_entity_code("CVD-203 W-CVD 3号機") == ("CVD-203", "W-CVD 3号機")
    assert split_entity_code("CMP-101") == ("CMP-101", "")
    assert split_entity_code("CMP-101 / CMP-102") == ("CMP-101 / CMP-102", "")   # 名前の側も番号だけ
    assert split_entity_code("超純水製造装置 2系") == ("超純水製造装置 2系", "")  # 番号らしくない

    d = list_spec_dict(group_by="entity_month")
    d["columns"] = [c for c in d["columns"] if c["key"] != "equipment_name"]
    spec = spec_from_dict(d)
    records = [_rec("A-1", record_no="A-1", occurred_at="2026-08-03", equipment_id="ETC-302(OXIDEエッチャ 2号機)"),
               _rec("A-2", record_no="A-2", occurred_at="2026-08-10", equipment_id="ETC-302")]
    files = render_all(spec, records, {}, {})
    names = [f.name for f in files]
    assert "トラブル対応一覧_ETC-302_2026-08" + HINT in names      # 1つのファイルにまとまる
    assert not any("OXIDE" in n for n in names)
    rec_file = next(f for f in files if f.name == "トラブル対応一覧_ETC-302_2026-08" + HINT)
    assert "- このファイルの記録: 2件" in rec_file.text
    assert "- 設備: OXIDEエッチャ 2号機（ETC-302）" in rec_file.text
    assert len([f for f in files if f.kind == "summary" and "設備別" in f.name]) == 1


def test_entity_name_before_code_is_also_split():
    """「Oxideエッチャ 2号機（ETC-302）」のように番号が後ろにある並びも、番号にそろえる（集計が二重にならない）。"""
    assert split_entity_code("Oxideエッチャ 2号機（ETC-302）") == ("ETC-302", "Oxideエッチャ 2号機")
    assert split_entity_code("ARFスキャナ 2号機(LIT-402)") == ("LIT-402", "ARFスキャナ 2号機")
    assert split_entity_code("超純水製造装置（2系）") == ("超純水製造装置（2系）", "")   # 番号らしくない
    assert split_entity_code("CMP-101（CMP-102）") == ("CMP-101（CMP-102）", "")      # 名前の側も番号だけ

    d = list_spec_dict(group_by="entity_month")
    d["columns"] = [c for c in d["columns"] if c["key"] != "equipment_name"]
    spec = spec_from_dict(d)
    records = [_rec("A-1", record_no="A-1", occurred_at="2026-08-03", equipment_id="Oxideエッチャ 2号機（ETC-302）"),
               _rec("A-2", record_no="A-2", occurred_at="2026-08-10", equipment_id="ETC-302")]
    names = [f.name for f in render_all(spec, records, {}, {})]
    assert "トラブル対応一覧_ETC-302_2026-08" + HINT in names
    assert not any("Oxide" in n for n in names)


def test_code_column_upper_keeps_the_equipment_name_as_written():
    """code 列の upper は番号だけに効かせる（同じセルの設備名まで大文字にすると帳票側の表記と割れる）。"""
    col = ColumnSpec(key="equipment_id", display="設備番号", type="code", role="entity", normalize=["nfkc", "upper"])
    cctx = ConvertContext(na_tokens=set())
    for raw, expected in (("etc-302", "ETC-302"),
                          ("Oxideエッチャ 2号機（ETC-302）", "Oxideエッチャ 2号機(ETC-302)"),  # 括弧は NFKC で半角
                          ("etc-302 Oxideエッチャ 2号機", "ETC-302 Oxideエッチャ 2号機"),
                          ("Oxideエッチャ 2号機", "Oxideエッチャ 2号機")):
        assert convert_cell(raw, raw, col, cctx)[0] == expected


def test_category_column_that_is_not_in_the_records_is_not_used_for_breakdowns():
    """記録に出していない列（意味の分からないコード値）は、集計の内訳の軸にしない。"""
    from tables.summaries import category_column

    spec = spec_from_dict(list_spec_dict())
    assert category_column(spec).key == "failure_category"
    spec.column("failure_category").md = "omit"
    assert category_column(spec) is None

    records = [_rec("A-1", record_no="A-1", occurred_at="2026-08-03", equipment_id="CMP-101",
                    failure_category="F01", downtime=30),
               _rec("A-2", record_no="A-2", occurred_at="2026-08-04", equipment_id="CMP-101",
                    failure_category="F01", downtime=60)]
    files = render_all(spec, records, {}, {})
    text = "\n".join(f.text for f in files)
    assert "F01" not in text and "故障区分の内訳" not in text and "主な故障区分" not in text
    card = next(f for f in files if f.kind == "dataset").text
    assert "- 記録に出していない列: 故障区分" in card


def test_omitted_person_column_is_explained_separately_from_code_columns():
    """人名の列に「意味の分からないコード値のため」という理由を付けない。"""
    spec = spec_from_dict(list_spec_dict())
    spec.column("failure_category").md = "omit"
    records = [_rec("A-1", record_no="A-1", occurred_at="2026-08-03", equipment_id="CMP-101", worker="田中")]
    card = next(f for f in render_all(spec, records, {}, {}) if f.kind == "dataset").text
    assert "- 記録に出していない列: 故障区分（意味の分からないコード値" in card
    assert "- 記録に出していない列（人名）: 担当者（人名のため出していません" in card


def test_record_title_is_cut_at_a_readable_place():
    """見出しの切り詰めは、閉じない括弧を残さず、切ったことが分かるようにする。"""
    from tables.markdown import TITLE_TEXT_CHARS, record_title

    spec = spec_from_dict(list_spec_dict())
    long = "ドライポンプが過負荷停止(A-3501)、チャンバー圧力上昇(同型機CVD-205は正常なので単体の不具合)"
    values = {"record_no": "TR-1", "occurred_at": "2026-08-03", "equipment_id": "CMP-101", "symptom": long}
    title = record_title(values, spec)
    assert title.endswith("…｜2026-08-03")
    body = title.split("】")[1].split("｜")[0]
    assert body.count("(") == body.count(")") and len(body) <= TITLE_TEXT_CHARS + 1
    # 40字に収まる現象はそのまま（「…」を付けない）
    short = dict(values, symptom="ドライポンプが過負荷停止")
    assert record_title(short, spec).endswith("ドライポンプが過負荷停止｜2026-08-03")


def test_record_title_drops_a_date_only_preamble():
    """本文の1行目が「【発生】R05.04.01 11:45(休日)」なら、見出しの日付と同じものを繰り返さない。"""
    from tables.markdown import record_title

    spec = spec_from_dict(list_spec_dict())
    text = "【発生】R05.04.01 11:45(休日)\n【設備】ETC-305 Polyエッチャ 5号機\n【内容】PM4のVppが範囲外"
    values = {"record_no": "CA-1", "occurred_at": "2023-04-03", "equipment_id": "ETC-305", "symptom": text}
    assert record_title(values, spec) == "【CA-1】ETC-305｜2023-04-03"
    # 日付のあとに中身が続くときは、日付だけ落として中身を見出しに使う
    values2 = dict(values, symptom="R5.4.2 16:58(休日)、ビーム電流が低下し停止")
    assert record_title(values2, spec) == "【CA-1】ETC-305 ビーム電流が低下し停止｜2023-04-03"
    # 札も日付も無い現象はそのまま
    values3 = dict(values, symptom="スラリー流量低下")
    assert record_title(values3, spec) == "【CA-1】ETC-305 スラリー流量低下｜2023-04-03"


def test_timeline_keeps_the_sentence_end_of_a_dropped_sentence():
    """重複で省いた文が「。」で終わっていたら、前後がつながって元にない1文にならないようにする。"""
    spec = spec_from_dict(list_spec_dict())
    log = ("8/3 10:00 田中：過負荷で停止、チャンバー圧力上昇（同型機は正常）。保全へ連絡\n"
           "8/3 11:00 田中：メーカー手配\n")
    rec = _rec("TR-9", record_no="TR-9", occurred_at="2026-08-03", equipment_id="CMP-101",
               symptom="チャンバー圧力上昇（同型機は正常）", response_log=log)
    lines = record_block(rec, spec)
    timeline = [ln.strip() for ln in lines[lines.index("- 対応の時系列:") + 1:]]
    assert timeline[0].endswith(": 過負荷で停止。保全へ連絡")


def test_record_body_shows_the_split_entity_like_the_heading():
    """設備名の列がない台帳でも、本文の設備の行を見出し・集計と同じ書き方にそろえる。"""
    d = list_spec_dict()
    d["columns"] = [c for c in d["columns"] if c["key"] != "equipment_name"]
    spec = spec_from_dict(d)
    rec = _rec("A-1", record_no="A-1", occurred_at="2026-08-03", equipment_id="ETC-302(Oxideエッチャ 2号機)")
    lines = record_block(rec, spec)
    assert "- 設備番号: Oxideエッチャ 2号機（ETC-302）" in lines
    assert lines[0].startswith("## 【A-1】Oxideエッチャ 2号機（ETC-302）")
    # 番号だけの記録は今までどおり（同じ列の名前のまま）
    rec2 = _rec("A-2", record_no="A-2", occurred_at="2026-08-04", equipment_id="ETC-302")
    assert "- 設備番号: ETC-302" in record_block(rec2, spec)


# ---- ユーザーが決めた出力の変更（ヒント既定オン・時系列の重複削除の選択・丸数字） ----

def test_new_spec_turns_the_filename_hint_on_and_saved_specs_keep_their_value():
    """分割ヒントは新しい取り込み設定の既定オン。保存済みの設定（JSON に値がある）は書き換えない。"""
    from tables.spec import TableSpec

    assert TableSpec("新しい設定").markdown["lightrag_hint"] is True
    assert spec_from_suggestions("新しい設定", {"header_rows": [1]}, []).markdown["lightrag_hint"] is True

    saved = list_spec_dict(lightrag_hint=False)          # 前に「付けない」で保存した設定
    assert spec_from_dict(saved).markdown["lightrag_hint"] is False
    assert spec_from_dict(spec_to_dict(spec_from_dict(saved))).markdown["lightrag_hint"] is False

    # 既定のまま出すと、記録ファイルだけにヒントが付く（集計・データセット説明には付けない）
    files = render_all(spec_from_dict(list_spec_dict()),
                       [_rec("TR-1", record_no="TR-1", occurred_at="2026-08-03", equipment_id="CMP-101")], {}, {})
    for f in files:
        assert f.name.endswith(HINT) == (f.kind == "records")


def test_timeline_dedupe_can_be_turned_off_in_the_import_settings():
    """「対応の時系列から、他の列と同じ内容の文を省く」を外すと、重複する文も原文のまま出す。"""
    log = "8/3 10:00 田中：スラリー流量低下アラームで研磨停止。保全へ連絡\n8/3 11:00 田中：フィルター交換、流量の再校正を実施\n"
    values = dict(record_no="TR-001", occurred_at="2026-08-03", equipment_id="CMP-101",
                  symptom="スラリー流量低下アラームで研磨停止",
                  action="フィルター交換、流量の再校正を実施", response_log=log)

    def timeline(spec):
        lines = record_block(_rec("TR-001", **values), spec)
        return [ln.strip() for ln in lines[lines.index("- 対応の時系列:") + 1:]]

    assert spec_from_dict(list_spec_dict()).markdown["dedupe_timeline"] is True   # 既定は今までの動き
    on = timeline(spec_from_dict(list_spec_dict()))
    assert on[0].endswith(": 保全へ連絡") and on[1].endswith(": （処置と同じ）")

    off = timeline(spec_from_dict(list_spec_dict(dedupe_timeline=False)))
    assert off[0].endswith(": スラリー流量低下アラームで研磨停止。保全へ連絡")
    assert off[1].endswith(": フィルター交換、流量の再校正を実施")


@pytest.mark.parametrize("type_, text, expected", [
    ("text", "①分解 ②清掃 ③組立", "①分解 ②清掃 ③組立"),          # 本文は丸数字のまま
    ("text", "Ⓐ系統の㋐弁を閉", "Ⓐ系統の㋐弁を閉"),
    ("string", "①機械", "①機械"),
    ("text", "㈱山田製作所へ連絡", "(株)山田製作所へ連絡"),          # 区切りが残る表記は NFKC のまま
    ("text", "ＡＢＣ　１２３", "ABC 123"),                          # 全角の英数字・空白は今までどおり
    ("code", "①-２", "1-2"),                                        # コードは突き合わせに使うので NFKC のみ
])
def test_enclosed_characters_are_kept_in_the_text_that_goes_into_markdown(type_, text, expected):
    col = ColumnSpec("c", "列", type=type_)
    out, error, flag = convert_cell(None, text, col, ConvertContext(na_tokens=frozenset()))
    assert (out, error, flag) == (expected, None, None)


def test_enclosed_characters_do_not_change_dates_numbers_or_na_tokens():
    """比較・解析に使う正規化は今までどおり NFKC（丸数字を残すのは md に出す文章だけ）。"""
    cctx = ConvertContext(fiscal_year=2026, fiscal_start=4, na_tokens=frozenset(["-", "該当なし"]))
    assert convert_cell(None, "２０２６/８/３", ColumnSpec("d", "発生日", type="date"), cctx)[0] == "2026-08-03"
    assert convert_cell(None, "１，２３４", ColumnSpec("n", "費用", type="number"), cctx)[0] == 1234
    assert convert_cell(None, "－", ColumnSpec("t", "現象", type="text"), cctx)[0] is None       # NA 判定
    assert convert_cell(None, "①", ColumnSpec("t", "現象", type="text"), cctx)[0] == "①"


def test_record_markdown_and_filename_for_a_circled_number(tmp_path):
    """md の本文は丸数字のまま、ファイル名は今までどおり NFKC（OS・LightRAG 側で揺れないように）。"""
    spec = spec_from_dict(list_spec_dict())
    rec = _rec("TR-①", record_no="TR-001", occurred_at="2026-08-03", equipment_id="CMP-101",
               equipment_name="CMP研磨装置1号機", action="①フィルター交換 ②流量再校正")
    assert "- 処置: ①フィルター交換 ②流量再校正" in record_block(rec, spec)
    assert md_filename(["トラブル対応一覧", "①"]) == "トラブル対応一覧_1.md"


# ---- R1: 番号の分け方・平均の分母・9999/12/31 ----

def test_split_entity_code_prefers_trailing_code_and_drops_qualifiers():
    """「Fab1 OHTシステム（OHT-801）」の番号は OHT-801。「IMP-602（推定）」の「推定」は設備名にしない。"""
    assert split_entity_code("Fab1 OHTシステム（OHT-801）") == ("OHT-801", "Fab1 OHTシステム")
    assert split_entity_code("Fab2 自動倉庫（ストッカ）（STK-831）") == ("STK-831", "Fab2 自動倉庫（ストッカ）")
    assert split_entity_code("IMP-602（推定）") == ("IMP-602", "")
    assert split_entity_code("CVD-203 W-CVD 3号機") == ("CVD-203", "W-CVD 3号機")
    assert split_entity_code("ETC-302(OXIDEエッチャ 2号機)") == ("ETC-302", "OXIDEエッチャ 2号機")

    d = list_spec_dict(group_by="entity_month")
    d["columns"] = [c for c in d["columns"] if c["key"] != "equipment_name"]
    spec = spec_from_dict(d)
    records = [_rec("A-1", record_no="A-1", occurred_at="2026-08-03", equipment_id="IMP-602（推定）"),
               _rec("A-2", record_no="A-2", occurred_at="2026-08-10", equipment_id="IMP-602")]
    names = [f.name for f in render_all(spec, records, {}, {})]
    assert "トラブル対応一覧_IMP-602_2026-08" + HINT in names
    assert not any("推定" in n for n in names)


def test_fiscal_year_average_uses_only_records_with_a_value():
    """設備別年度集計の平均は値のある記録だけで割り、整数の列でも 0 に丸めない。"""
    from tables.summaries import fmt_average

    assert fmt_average(112.0, 375, True) == 0.3 and fmt_average(0, 3, True) == 0
    spec = spec_from_dict(list_spec_dict())
    records = [_rec("A-1", record_no="A-1", occurred_at="2026-08-03", equipment_id="CMP-101", downtime=30),
               _rec("A-2", record_no="A-2", occurred_at="2026-08-04", equipment_id="CMP-101", downtime=60),
               _rec("A-3", record_no="A-3", occurred_at="2026-08-05", equipment_id="CMP-101"),
               _rec("A-4", record_no="A-4", occurred_at="2026-08-06", equipment_id="CMP-101")]
    fy = next(f for f in render_all(spec, records, {}, {}) if "設備別_CMP-101" in f.name).text
    assert "- 停止時間の合計は90分（1.5時間）、1件あたり平均（値のある2件）45分です。" in fy


def test_far_future_date_does_not_break_summaries():
    """「期限なし」の 9999/12/31 があっても、月末日の計算で落ちずに md を作れる。"""
    from tables.summaries import month_last_day

    assert month_last_day("9999-12") == "9999-12-31" and month_last_day("2024-02") == "2024-02-29"
    spec = spec_from_dict(list_spec_dict())
    records = [_rec("A-1", record_no="A-1", occurred_at="2026-08-03", equipment_id="CMP-101", downtime=30),
               _rec("A-2", record_no="A-2", occurred_at="9999-12-31", equipment_id="CMP-101", downtime=10)]
    files = render_all(spec, records, {}, {})
    assert any("設備別_CMP-101_9999年度" in f.name for f in files)
