"""帳票まわりの修正（6巡目の確認で見つかった不具合）の回帰テスト。"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from excel.extractor import extract_document
from excel.tables import section_of, sections_of
from excel.workbook import load_workbook_info
from pattern.builder import suggest_rows
from pattern.forms import _section_value
from pattern.model import FieldDef, PatternDef, SheetDef
from tests.test_forms_fixes import _confirmed_doc

FILL = PatternFill("solid", fgColor="DDDDDD")
BOLD = Font(bold=True)


def _head(ws, coord, text):
    ws[coord] = text
    ws[coord].font = BOLD
    ws[coord].fill = FILL


def _label(ws, coord, text, value_coord, value):
    ws[coord] = text
    ws[coord].fill = FILL
    ws[value_coord] = value


def _pattern(*fields):
    return PatternDef(id=1, name="T", version="v1", description="", status="active", image_processing="none",
                      sheets=[SheetDef("報告書")], fields=list(fields))


def _field(ex, name):
    return next(f for f in ex["fields"] if f["field_name"] == name)


# ---- R6F-1: 右上の区画・区画の中の小見出し ------------------------------------------------------

def _side_by_side(path, *, issuer=True):
    """右上に「【回答欄】」、その下の行の左に「■ 発行部署記入欄」がある版。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "報告書"
    ws["A1"] = "工程異常連絡書"
    ws["A1"].font = BOLD
    ws["N1"] = "【回答欄】"
    ws["N1"].font = BOLD
    _head(ws, "A3", "■ 発行部署記入欄")
    if issuer:
        _label(ws, "A6", "処置", "B6", "発行側の処置（仮）")
    _label(ws, "N6", "処置", "O6", "回答側の処置（確定）")
    wb.save(path)
    return load_workbook_info(path)


def test_a_right_hand_section_started_higher_keeps_its_columns(tmp_path):
    info = _side_by_side(tmp_path / "cols.xlsx")
    grid = info.grids["報告書"]
    assert section_of(grid, grid.cells[(6, 14)]) == "回答"
    assert section_of(grid, grid.cells[(6, 1)]) == "発行部署記入"
    ex = extract_document(info, _pattern(FieldDef("action", "処置", ["処置"], section="回答")), ["報告書"])
    assert _field(ex, "action")["value"] == "回答側の処置（確定）"


def test_learning_finds_the_right_hand_answer_section(tmp_path):
    from pattern.builder import _learn_sections

    infos = [_side_by_side(tmp_path / "both.xlsx"), _side_by_side(tmp_path / "answer.xlsx", issuer=False)]
    row = {"use": True, "data_type": "string", "field_name": "action", "display_name": "処置",
           "candidates": "処置", "section": ""}
    _learn_sections(infos, {"報告書"}, [row])
    assert row["section"] == "回答"


def _nested(path):
    """「▼ 回答欄」の下に小見出し「1. 暫定対策」があり、その中に回答側の欄がある版。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "報告書"
    _head(ws, "A1", "■ 発行側")
    _label(ws, "A2", "処置内容", "B2", "発行側の処置")
    _head(ws, "A4", "▼ 回答欄")
    _head(ws, "A5", "1. 暫定対策")
    _label(ws, "A6", "処置内容", "B6", "回答側の処置")
    wb.save(path)
    return load_workbook_info(path)


def test_a_sub_heading_inside_the_answer_section_is_still_the_answer_section(tmp_path):
    info = _nested(tmp_path / "nested.xlsx")
    grid = info.grids["報告書"]
    assert sections_of(grid, grid.cells[(6, 1)]) == ["暫定対策", "回答"]
    assert sections_of(grid, grid.cells[(2, 1)]) == ["発行側"]   # 同じ階層の見出しで上の区画は終わる
    field = FieldDef("action", "処置内容", ["処置内容"], section=_section_value("回答欄"))
    assert _field(extract_document(info, _pattern(field), ["報告書"]), "action")["value"] == "回答側の処置"
    # 外側の区画の名前で決めた項目は、内側の小見出しの中の欄だけを見ない（発行側の欄は区画の外）
    issuer = FieldDef("action", "処置内容", ["処置内容"], section="発行側")
    assert _field(extract_document(info, _pattern(issuer), ["報告書"]), "action")["value"] == "発行側の処置"


# ---- R6F-2: 明細表の項目の区画 -----------------------------------------------------------------

def _two_tables(path, *, answer_label="使用部品"):
    wb = Workbook()
    ws = wb.active
    ws.title = "報告書"
    _head(ws, "A1", "■ 発行側")
    _head(ws, "A3", "使用部品")
    for c, t in zip("ABC", ["品番", "品名", "数量"]):
        _head(ws, f"{c}4", t)
    ws["A5"], ws["B5"], ws["C5"] = "P-1", "発行側の部品", 2
    _head(ws, "A8", "▼ 回答欄")
    _head(ws, "A10", answer_label)
    for c, t in zip("ABC", ["品番", "品名", "数量"]):
        _head(ws, f"{c}11", t)
    ws["A12"], ws["B12"], ws["C12"] = "P-9", "回答側の部品", 5
    wb.save(path)
    return load_workbook_info(path)


def _parts(section):
    return FieldDef("parts", "使用部品", ["使用部品"], data_type="table", table_columns=["品番", "品名", "数量"],
                    section=section)


def test_a_table_field_reads_the_table_in_its_section(tmp_path):
    info = _two_tables(tmp_path / "sec.xlsx")
    parts = _field(extract_document(info, _pattern(_parts(_section_value("▼ 回答欄"))), ["報告書"]), "parts")
    assert parts["value"]["rows"] == [["P-9", "回答側の部品", "5"]]
    # 区画を決めていなければ今までどおり上の表
    parts = _field(extract_document(info, _pattern(_parts("")), ["報告書"]), "parts")
    assert parts["value"]["rows"] == [["P-1", "発行側の部品", "2"]]


def test_a_table_found_by_its_columns_prefers_the_one_in_the_section(tmp_path):
    # 回答欄の表の見出しの書き方が違う版: 列見出しで探すときも区画の中の表を選ぶ
    info = _two_tables(tmp_path / "cols.xlsx", answer_label="交換部品")
    field = _parts(_section_value("▼ 回答欄"))
    field.candidates = ["部品明細"]   # どちらの見出しにも当たらない
    parts = _field(extract_document(info, _pattern(field), ["報告書"]), "parts")
    assert parts["value"]["rows"] == [["P-9", "回答側の部品", "5"]]


# ---- R6-F2: 1件だけになったまとめ取り込み -------------------------------------------------------

def test_a_batch_left_with_one_form_is_no_longer_a_batch(app, client):
    """取り込み失敗・削除で帳票が1件だけになったまとまりは、1件の帳票として扱う。

    「確定済みの分だけ zip にします」「残りの帳票は…」はどれも嘘になるため。
    """
    import html

    first = _confirmed_doc(app, "1.xlsx", batch_id="B9", order=0)
    second = _confirmed_doc(app, "2.xlsx", batch_id="B9", order=1)

    def finish(ids):
        url = "/forms/finish?ids=" + ",".join(str(i) for i in ids) + f"&current={first}"
        return html.unescape(client.get(url).get_json()["html"])

    page = finish([first, second])
    assert "zip" in page and "/forms/batches/B9/download.zip" in page

    client.post(f"/forms/{second}/delete")
    page = finish([first, second])
    assert "zip" not in page and "/forms/batches/" not in page
    assert "Markdown をダウンロード（.md）" in page and f"/forms/{first}/download.md" in page


# ---- R6-B1: 見出しからは日付と分からない欄（「発生」）の型 ----------------------------------------

def test_a_date_value_under_a_non_date_label_becomes_a_date_field(tmp_path):
    """値が日付だけで書かれていれば日付型にする（文字列のままだと md の日付の書き方が混ざる）。"""
    from pattern.builder import _guess_type

    wb = Workbook()
    ws = wb.active
    ws.title = "報告書"
    _label(ws, "A1", "発生", "B1", "2024年7月28日 22:46")
    _label(ws, "A2", "復旧", "B2", "R5.11.16")
    path = tmp_path / "occurred.xlsx"
    wb.save(path)
    types = {r["display_name"]: r["data_type"] for r in suggest_rows([load_workbook_info(path)])[1]}
    assert types["発生"] == "date" and types["復旧"] == "date"

    # 日付を含むだけの値・年の無い値・識別番号は日付にしない
    assert _guess_type("R2026-00123", "報告No") == "string"
    assert _guess_type("2026-09-14-3", "ロットNo") == "string"
    assert _guess_type("2026-09-14 に復旧", "処置") == "string"
    assert _guess_type("2/12", "発生") == "string"
    assert _guess_type("12:30", "発生") == "string"
