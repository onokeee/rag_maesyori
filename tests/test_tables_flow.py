"""一覧表の1画面（/tables）の通し確認。

画面は1枚で、段（読み取り方→表の範囲→列の対応づけ→AI整形→内容の確認→確定してダウンロード）を
fetch で出し入れする。テストも画面と同じ JSON のやりとりで進める。
"""
import io
import zipfile

import pytest

from core import jobs
from tables import pipeline, store
from tests.tables_helpers import (  # noqa: F401  (COLUMNS/CSV_TEXT は他のテストからも読まれる)
    COLUMNS, CSV_TEXT, columns_payload, name_source, panel, panel_html, preview_panel, save_columns, save_layout,
    save_source, upload, upload_csv, wait_import_job,
)


def test_csv_import_flow(app, client, monkeypatch):
    import_id = upload_csv(client, "トラブル一覧.csv")

    source = panel(client, import_id, "source")
    # 取り込み設定の名前は残り続けるので、アップロードしたファイル名を初期値にしない（design.md 3.3）
    assert 'name="new_template_name" class="input" value=""' in source["html"]
    assert name_source(client, import_id, "トラブル対応一覧")["next"] == "layout"

    layout = panel(client, import_id, "layout")
    assert "row-header" in layout["html"]
    detect = client.post(f"/tables/imports/{import_id}/layout/detect", json={"header_rows": [1]}).get_json()
    assert detect["data_end"] == 13 and detect["table_kind"] == "list"
    res = save_layout(client, import_id)
    assert res.status_code == 200 and res.get_json()["next"] == "columns"

    assert "対応内容" in panel_html(client, import_id, "columns")
    # 2回目からは取り込みの控えで画面を作る（CSV を開き直さない）
    monkeypatch.setattr(pipeline, "open_import_source", lambda *a, **k: pytest.fail("CSV を開き直した"))
    assert "row-header" in panel_html(client, import_id, "layout")
    assert "対応内容" in panel_html(client, import_id, "columns")
    assert panel(client, import_id, "source")["html"]
    monkeypatch.undo()

    res = save_columns(client, import_id, columns_payload("トラブル対応一覧", ai_role="log"))
    assert res.status_code == 200, res.get_json()
    assert res.get_json()["next"] == "ai"
    imp = wait_import_job(app, import_id)
    assert imp["status"] == "preview" and imp["stats"]["records"] == 12

    assert client.get(f"/api/jobs/{imp['job_id']}").get_json()["status"] == "done"
    assert "分割プレビュー" in panel_html(client, import_id, "ai")
    split = client.post(f"/tables/imports/{import_id}/ai/split-preview", json={"row_key": "TR-001"}).get_json()
    assert len(split["segments"]) == 2

    preview = preview_panel(app, client, import_id)
    assert "確定してMarkdownを作成" in preview["html"] and "トラブル対応一覧_2026-08.md" in preview["html"]
    text = client.get(f"/tables/imports/{import_id}/preview/file",
                      query_string={"name": "トラブル対応一覧_2026-08.md"}).get_json()["text"]
    assert "TR-001" in text and "対応の時系列" in text

    res = client.post(f"/tables/imports/{import_id}/confirm")
    assert res.status_code == 200 and res.get_json()["next"] == "done"
    imp = wait_import_job(app, import_id)
    assert imp["status"] == "confirmed"
    assert "Markdownをまとめてダウンロード" in panel_html(client, import_id, "done")

    # 正規化CSV は単独でも取れる（zip の 管理用_RAGには入れない フォルダにも入る）。
    # zip をダウンロードするとこの取り込みのデータは消えるので、単独で取るのは zip の前だけ
    assert client.get(f"/tables/imports/{import_id}/normalized.csv").status_code == 200
    zdata = client.get(f"/tables/imports/{import_id}/download.zip").data
    names = zipfile.ZipFile(io.BytesIO(zdata)).namelist()
    assert "RAG投入用/トラブル対応一覧_2026-08.md" in names
    assert {"管理用_RAGには入れない/正規化データ.csv", "管理用_RAGには入れない/問題一覧.csv",
            "管理用_RAGには入れない/取込レポート.csv"} <= set(names)
    with app.app_context():
        assert store.get_import(import_id) is None      # ダウンロードしたら残さない
        assert store.get_template(imp["template_id"]) is not None   # 取り込み設定は残る

    # 2回目: 同じ設定を選べば列の対応づけを飛ばす
    second = upload_csv(client, "トラブル一覧2.csv")
    tid = imp["template_id"]
    assert "必須列" in panel_html(client, second, "source")
    save_source(client, second, encoding="cp932", delimiter=",", template=str(tid))
    res = save_layout(client, second)
    body = res.get_json()
    assert res.status_code == 200 and body["next"] == "ai"
    # 飛ばした段は灰色のままにせず、使った設定の名前を出す
    assert "トラブル対応一覧" in body["done"]["columns"]
    assert wait_import_job(app, second)["status"] == "preview"


# ---- 読み込み中の中止（段に出る［中止］） ------------------------------------------------

def _pending_import(app, status="reading"):
    """読み込み中に見える取り込みを1件つくる（ジョブは待機中のまま）。"""
    from models import database
    with app.app_context():
        import_id = store.create_import("大きい表.csv", "h" * 64, "uploads/大きい表.csv", {"encoding": "utf-8"})
        conn = database.get_db()
        ts = database.now()
        cur = conn.execute("""INSERT INTO jobs (kind, ref_type, ref_id, status, params_json, progress_json,
                              created_at, updated_at) VALUES ('table_read', 'table_import', ?, 'queued', '{}', '{}', ?, ?)""",
                           (import_id, ts, ts))
        conn.commit()
        store.update_import(import_id, status=status, job_id=cur.lastrowid)
        return import_id, cur.lastrowid


def test_panels_show_the_running_read_job_with_a_cancel_url(app, client):
    import_id, _job_id = _pending_import(app)
    data = panel(client, import_id, "preview")
    assert data["reading"] is True
    assert data["job"]["kind"] == "table_read"
    assert data["job"]["cancel_url"] == f"/tables/imports/{import_id}/cancel"


def test_cancel_read_job_puts_the_import_back(app, client):
    import_id, job_id = _pending_import(app)
    res = client.post(f"/tables/imports/{import_id}/cancel")
    assert res.status_code == 200 and res.get_json()["ok"] is True
    with app.app_context():
        assert jobs.get_job(job_id)["status"] == "cancelled"
        # 待機中のまま中止したときは取り込みの状態も戻す（読み込み中のままにしない）
        assert store.get_import(import_id)["status"] == "uploaded"
    # 待ち画面ではなく、読み込みからやり直せる段が出る
    assert panel(client, import_id, "preview").get("reading") is None


def test_cancel_render_job_returns_to_preview(app, client):
    from models import database

    import_id, job_id = _pending_import(app, status="confirming")
    with app.app_context():
        conn = database.get_db()
        conn.execute("UPDATE jobs SET kind = 'table_render' WHERE id = ?", (job_id,))
        conn.commit()
    assert client.post(f"/tables/imports/{import_id}/cancel").status_code == 200
    with app.app_context():
        assert jobs.get_job(job_id)["status"] == "cancelled"
        assert store.get_import(import_id)["status"] == "preview"


def test_cancel_without_running_job_is_reported(app, client):
    with app.app_context():
        import_id = store.create_import("表.csv", "h" * 64, "uploads/表.csv", {})
    res = client.post(f"/tables/imports/{import_id}/cancel")
    assert res.status_code == 400 and "中止できる処理がありません" in res.get_json()["error"]


# ---- 表の範囲の段の表示 ----------------------------------------------------------------

def test_layout_panel_without_headers(app, client):
    import_id = upload_csv(client, "1列.csv", "あ\r\nい\r\n", encoding="utf-8")
    name_source(client, import_id, "1列", encoding="utf-8")
    data = panel(client, import_id, "layout")
    # 見出し行が見つからないときに「1〜0行目」のような存在しない行番号を出さない
    assert "データの行が見つかりません" in data["html"] and "〜0行目" not in data["html"]
    assert data["note"] == ""
    # 見出し行を指定すれば進める
    res = save_layout(client, import_id, header_rows="1")
    assert res.status_code == 200 and res.get_json()["next"] == "columns"


# ---- 取り込みの削除・段の表示 ----------------------------------------------------------

def test_delete_import_removes_the_file_and_the_read_rows(app, client):
    """間違えて取り込んだファイルを消せる（元のファイルも読み込んだ内容も残さない）。"""
    import_id = upload_csv(client, "間違い.csv")
    assert "data-import-delete" in panel_html(client, import_id, "layout")
    with app.app_context():
        base = pipeline.import_dir(import_id)
        base.mkdir(parents=True, exist_ok=True)
        (base / "rows.jsonl.gz").write_bytes(b"x")
        stored = store.get_import(import_id)["stored_path"]
    res = client.post(f"/tables/imports/{import_id}/delete")
    assert res.status_code == 200 and res.get_json()["ok"] is True
    with app.app_context():
        from core.files import upload_path
        assert store.get_import(import_id) is None
        assert not pipeline.import_dir(import_id).exists()
        assert not upload_path(stored).exists()
        assert store.list_imports(limit=10) == []


def test_delete_is_refused_while_a_job_is_running(app, client):
    import_id, _job_id = _pending_import(app)
    res = client.post(f"/tables/imports/{import_id}/delete")
    assert res.status_code == 400 and "処理中の取り込みは削除できません" in res.get_json()["error"]
    with app.app_context():
        assert store.get_import(import_id) is not None


def test_source_rejects_an_encoding_that_is_not_offered(app, client):
    """画面にない文字コードを保存させない（保存できると読み込みで落ちて段が開けなくなる）。"""
    import_id = upload_csv(client, "文字コード.csv")
    res = save_source(client, import_id, encoding="rot13", delimiter=",", template="new", new_template_name="x")
    assert res.status_code == 400 and "この画面にない文字コード" in res.get_json()["error"]
    with app.app_context():
        assert store.get_import(import_id)["source"].get("encoding") != "rot13"
    # 画面にある文字コードは保存できる
    assert name_source(client, import_id, "x", encoding="utf-16")["next"] == "layout"
    # 中身と合わない文字コードでも、500 ではなく段を出して選び直せる（読めなければ理由を1行出す）
    assert panel(client, import_id, "source")["html"]
    layout = panel(client, import_id, "layout")
    assert layout["html"] or layout.get("locked")


def test_column_editor_shows_japanese_type_names(app, client):
    """型が合っていないかもしれないという注意書きを、英語キー（date / number / enum）のまま出さない。"""
    import re

    text = "管理No,発生日,設備番号,設備名,状態コード,対応内容,停止時間\r\n" + "".join(
        f"MS-{i:04d},2026-08-{i:02d},EQ-01,搬送ロボット1号機,{i % 3},点検した,{i * 5}\r\n" for i in range(1, 29))
    import_id = upload_csv(client, "型.csv", text)
    name_source(client, import_id, "型")
    save_layout(client, import_id)
    html = panel_html(client, import_id, "columns")
    shown = re.findall(r"値は「([^」]+)」らしい", html)
    assert shown, "型が合っていない列の注意書きが出ていない"
    assert not ({"date", "number", "enum", "string", "text", "code"} & set(shown)), shown


def test_preview_panel_says_it_is_already_confirmed(app, client):
    """確定済みの取り込みを開き直したときに、黙って作り直せるボタンだけを出さない。"""
    import_id = upload_csv(client, "確定済み.csv")
    name_source(client, import_id, "確定済み")
    save_layout(client, import_id)
    assert save_columns(client, import_id, columns_payload("確定済み")).status_code == 200
    wait_import_job(app, import_id)
    preview_panel(app, client, import_id)
    client.post(f"/tables/imports/{import_id}/confirm")
    assert wait_import_job(app, import_id)["status"] == "confirmed"
    data = preview_panel(app, client, import_id)
    assert data["confirmed"] is True
    assert "この取り込みは確定済みです" in data["html"]
    assert "確定し直してMarkdownを作り直す" in data["html"] and "data-confirm" in data["html"]


def test_source_panel_marks_a_crosstab_sheet(app, client):
    """クロス集計のシートに「一覧表らしい」と出さない（範囲の段で初めて止められると分かりにくい）。"""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "FY2023"
    ws.append(["設備"] + [f"2023-{m:02d}" for m in range(4, 13)] + [f"2024-{m:02d}" for m in range(1, 4)])
    for i in range(1, 16):
        ws.append([f"EQ-{i:02d}"] + [i * m for m in range(1, 13)])
    buf = io.BytesIO()
    wb.save(buf)
    res = upload(client, buf.getvalue(), "月別集計.xlsx")
    import_id = res.get_json()["import_id"]
    html = panel_html(client, import_id, "source")
    assert "クロス集計らしい（対応していません）" in html and "・一覧表らしい" not in html
    # 範囲の段の判定は今までどおり（この範囲では進めない）
    res = save_layout(client, import_id)
    assert res.status_code == 400 and "クロス集計" in res.get_json()["error"]


# ---- 処理中の変更・削除を止める -----------------------------------------------------------

def _queued_job(app, import_id, kind):
    from models import database
    with app.app_context():
        conn = database.get_db()
        ts = database.now()
        cur = conn.execute("""INSERT INTO jobs (kind, ref_type, ref_id, status, params_json, progress_json,
                              created_at, updated_at) VALUES (?, 'table_import', ?, 'queued', '{}', '{}', ?, ?)""",
                           (kind, import_id, ts, ts))
        conn.commit()
        return cur.lastrowid


def test_source_is_not_changed_while_reading(app, client):
    """読み込み中に読み取り方を保存しても、文字コード・区切り文字を書き換えない
    （古い設定の読み込み結果で確定させない）。"""
    import_id, _job_id = _pending_import(app)
    res = save_source(client, import_id, encoding="utf-8", delimiter=";", template="new", new_template_name="x")
    assert res.status_code == 409 and "処理中は変更できません" in res.get_json()["error"]
    assert save_layout(client, import_id, header_rows="2").status_code == 409
    assert save_columns(client, import_id, {"name": "x", "columns": []}).status_code == 409
    with app.app_context():
        imp = store.get_import(import_id)
        assert imp["status"] == "reading" and imp["source"] == {"encoding": "utf-8"}


def test_utf16_without_bom_can_pass_the_source_panel(app, client):
    """BOMなしの UTF-16 と判定された CSV も、読み取り方をそのまま保存して先へ進める。"""
    import_id = upload_csv(client, "u16.csv", encoding="utf-16-le")
    with app.app_context():
        assert store.get_import(import_id)["source"]["encoding"] == "utf-16-le"
    assert 'value="utf-16-le"' in panel_html(client, import_id, "source")
    assert name_source(client, import_id, "u16", encoding="utf-16-le")["next"] == "layout"
    assert panel(client, import_id, "layout").get("locked") is None


def test_delete_and_confirm_are_refused_while_ai_formatting_runs(app, client, monkeypatch):
    """AI整形の実行中に削除すると、動いている呼び出しが消したあとの DB に結果を書き戻すので、止める。"""
    from types import SimpleNamespace

    from views import tables as tables_view

    import_id = upload_csv(client, "AI中.csv")
    with app.app_context():
        store.update_import(import_id, status="preview")
    _queued_job(app, import_id, "ai_format")
    res = client.post(f"/tables/imports/{import_id}/delete")
    assert res.status_code == 400 and "AI整形の実行中は削除できません" in res.get_json()["error"]
    with app.app_context():
        assert store.get_import(import_id) is not None

    monkeypatch.setattr(tables_view, "_spec_for", lambda imp: SimpleNamespace(log_stage=object()))
    monkeypatch.setattr(pipeline, "start_render_job", lambda *a, **k: pytest.fail("AI整形の実行中に確定した"))
    res = client.post(f"/tables/imports/{import_id}/confirm")
    assert res.status_code == 400 and "AI整形の実行中は確定できません" in res.get_json()["error"]


def test_trial_run_needs_the_external_confirmation(app, client, monkeypatch):
    """試し実行も、外部の AI に送るときは画面の確認のチェックがなければ送らない（サーバー側で確かめる）。"""
    from types import SimpleNamespace

    from aiproc import runner
    from services import llm
    from views import tables as tables_view

    import_id = upload_csv(client, "試し.csv")
    monkeypatch.setattr(tables_view, "_spec_for", lambda imp: SimpleNamespace(log_stage=object()))
    monkeypatch.setattr(llm, "job_client_settings", lambda *a, **k: {"chat_url": "https://api.example.com/v1/chat"})
    monkeypatch.setattr(runner, "trial_row", lambda *a, **k: pytest.fail("確認なしで外部に送った"))
    monkeypatch.setattr(runner, "load_rows_for_ai", lambda *a, **k: pytest.fail("確認なしで外部に送った"))
    res = client.post(f"/tables/imports/{import_id}/ai/trial", json={"row_key": "TR-001"})
    assert res.status_code == 400 and "外部のAIサービス" in res.get_json()["error"]


def test_workbook_without_sheets_is_rejected(app, client, tmp_path):
    """シートのないブックは取り込みにせず、選び直しの案内を出す（読み取り方の段が 500 にならない）。"""
    import re

    from openpyxl import Workbook

    wb = Workbook()
    wb.save(tmp_path / "s.xlsx")
    src = zipfile.ZipFile(tmp_path / "s.xlsx")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/workbook.xml":
                data = re.sub(rb"<sheets>.*?</sheets>", b"<sheets/>", data, flags=re.S)
            z.writestr(item, data)
    res = upload(client, out.getvalue(), "シートなし.xlsx")
    assert res.status_code == 400 and "シート" in res.get_json()["error"]
    with app.app_context():
        assert store.list_imports(limit=10) == []


@pytest.mark.parametrize("url, target", [("/tables/upload", "views.tables.open_source"),
                                         ("/forms/upload", "views.forms.precheck_excel")])
def test_unexpected_error_during_upload_leaves_no_file(app, client, monkeypatch, tmp_path, url, target):
    """読み込みの途中で思わぬエラーが出ても、アップロードしたファイルは残さない（design.md 3.3）。"""
    from pathlib import Path

    from openpyxl import Workbook

    def boom(*args, **kwargs):
        raise RuntimeError("想定外")

    monkeypatch.setattr(target, boom)
    wb = Workbook()
    wb.active["A1"] = "管理No"
    wb.save(tmp_path / "a.xlsx")
    with pytest.raises(RuntimeError):
        client.post(url, data={"file": (io.BytesIO((tmp_path / "a.xlsx").read_bytes()), "a.xlsx")},
                    content_type="multipart/form-data")
    assert [p for p in Path(app.config["UPLOAD_DIR"]).rglob("*") if p.is_file()] == []
