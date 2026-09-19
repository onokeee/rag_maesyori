from excel.extractor import extract_document
from excel.text import normalize_label, to_date, to_number
from excel.workbook import load_workbook_info
from pattern.builder import suggest_rows
from pattern.forms import rows_to_pattern
from pattern.matcher import rank_patterns

META = {"name": "設備修理報告書", "version": "v1", "description": "", "image_processing": "none"}


def build_repair_pattern(infos):
    sheets, fields = suggest_rows(infos)
    return rows_to_pattern(1, META, sheets, fields)


def test_normalize_label_absorbs_notation_differences():
    assert normalize_label("設備№：") == normalize_label(" 設備 No. ") == normalize_label("【設備NO】") == "設備no"


def test_value_conversion():
    assert to_date("2026/9/1", "2026/9/1") == ("2026-09-01", None)
    assert to_date("令和8年9月14日", "令和8年9月14日") == ("2026-09-14", None)
    assert to_number("1.5時間", "1.5時間")[0] == 1.5
    assert to_number(3.0, "3") == (3, None)


def test_builder_suggests_standard_fields_from_samples(repair_infos):
    sheets, fields = suggest_rows(repair_infos)
    used_sheets = [s["sheet_name"] for s in sheets if s["use"]]
    assert used_sheets == ["修理報告書"]  # マスタ・参考資料・(2) の重複は選ばない

    used = {f["field_name"] for f in fields if f["use"]}
    assert {"report_id", "equipment_id", "equipment_name", "occurred_date", "reporter",
            "symptom", "cause", "repair", "result"} <= used
    equipment = next(f for f in fields if f["field_name"] == "equipment_id")
    assert equipment["seen"] == "3/3"
    assert "装置番号" in equipment["candidates"].splitlines()


def test_same_json_from_different_layouts(repair_infos):
    pattern = build_repair_pattern(repair_infos)
    expected = [
        {"report_id": "R2026-00123", "equipment_id": "EQ-001", "occurred_date": "2026-09-14", "work_hours": 2.5},
        {"report_id": "R2026-00124", "equipment_id": "EQ-002", "occurred_date": "2026-09-10", "work_hours": 1.5,
         "equipment_name": "CVD装置"},
        {"report_id": "R2026-00125", "equipment_id": "EQ-003", "occurred_date": "2026-08-30", "cause": "リニアガイドの潤滑不足。"},
    ]
    for info, values in zip(repair_infos, expected):
        match = rank_patterns(info, [pattern])[0]
        extraction = extract_document(info, pattern, match.sheet_names)
        for key, value in values.items():
            assert extraction["values"][key] == value, (info.path.name, key)
        assert "Robot Position Error" in extraction["values"]["symptom"] or info is not repair_infos[0]


def test_images_are_detected_with_location(repair_infos):
    standard, shifted, table = repair_infos
    assert [(i["sheet"], i["location"]) for i in standard.images] == [("修理報告書", "J3")]
    assert len(shifted.images) == 2
    assert table.images == []


def test_matcher_prefers_correct_pattern(repair_infos, sample_dir):
    inspection_info = load_workbook_info(sample_dir / "inspection.xlsx")
    repair = build_repair_pattern(repair_infos)
    sheets, fields = suggest_rows([inspection_info])
    inspection = rows_to_pattern(2, {**META, "name": "点検記録表"}, sheets, fields)

    assert rank_patterns(repair_infos[0], [inspection, repair])[0].pattern.name == "設備修理報告書"
    ranked = rank_patterns(inspection_info, [repair, inspection])
    assert ranked[0].pattern.name == "点検記録表"
    assert ranked[1].confidence < 50


# ---- 種類定義の拡張（単位・Markdownでの扱い・タイトル項目） ----

def test_to_date_supports_1904_workbooks():
    assert to_date(45000, "45000") == ("2023-03-15", None)
    assert to_date(45000 - 1462, "43538", date1904=True) == ("2023-03-15", None)


def test_split_label_unit_and_value_unit():
    from excel.text import nfkc_value, split_label_unit, value_unit

    assert split_label_unit("作業時間(h)") == ("作業時間", "時間")
    assert split_label_unit("停止時間（分）") == ("停止時間", "分")
    assert split_label_unit("設備番号") == ("設備番号", "")
    assert value_unit("1.5時間") == "時間" and value_unit("EQ-001") == ""
    assert nfkc_value(" ＣＭＰ　　装置 \r\nﾎﾟﾝﾌﾟ") == "CMP 装置\nポンプ"
    # ラベル比較用の正規化は従来どおり空白を消す
    assert normalize_label("ＣＭＰ　装置") == "cmp装置"


def test_builder_suggests_units_and_title_fields(repair_infos, tmp_path):
    from openpyxl import Workbook
    from pattern.builder import suggest_title_fields

    sheets, fields = suggest_rows(repair_infos)
    work_hours = next(f for f in fields if f["field_name"] == "work_hours")
    assert work_hours["unit"] == "時間"  # 値「1.5時間」から
    assert all(f["rag_output"] == "show" for f in fields)
    assert suggest_title_fields(fields) == ["report_id", "equipment_id", "equipment_name", "occurred_date"]

    wb = Workbook()
    ws = wb.active
    ws.title = "報告書"
    ws["A1"], ws["B1"] = "報告番号", "R-1"
    ws["A2"], ws["B2"] = "作業時間(h)", 3
    ws["A3"], ws["B3"] = "停止時間（分）", 45
    wb.save(tmp_path / "unit.xlsx")
    _, rows = suggest_rows([load_workbook_info(tmp_path / "unit.xlsx")])
    by_name = {r["field_name"]: r for r in rows}
    assert by_name["work_hours"]["unit"] == "時間" and by_name["work_hours"]["display_name"] == "作業時間"
    assert "作業時間(h)" in by_name["work_hours"]["candidates"].splitlines()
    assert by_name["downtime"]["unit"] == "分"


def test_form_rows_roundtrip_unit_rag_output_and_title_fields():
    from pattern.forms import parse_pattern_form, pattern_to_meta, pattern_to_rows

    form = {
        "name": "設備修理報告書", "version": "v1", "title_fields": "report_id, equipment_id, unknown",
        "md_options_form": "1",
        "fields-0-use": "on", "fields-0-field_name": "report_id", "fields-0-display_name": "報告番号",
        "fields-0-candidates": "報告番号", "fields-0-data_type": "string", "fields-0-direction": "auto",
        "fields-1-use": "on", "fields-1-field_name": "equipment_id", "fields-1-display_name": "設備番号",
        "fields-1-candidates": "設備番号", "fields-1-data_type": "string", "fields-1-direction": "auto",
        "fields-2-use": "on", "fields-2-field_name": "work_hours", "fields-2-display_name": "作業時間",
        "fields-2-candidates": "作業時間", "fields-2-data_type": "number", "fields-2-direction": "auto",
        "fields-2-unit": "時間", "fields-2-rag_output": "omit",
    }
    meta, sheet_rows, field_rows, errors = parse_pattern_form(form)
    assert errors == []
    assert meta["title_fields"] == ["report_id", "equipment_id"]
    assert meta["md_options"]["omit_person_fields"] is False  # チェックなしで送信された
    pattern = rows_to_pattern(3, meta, sheet_rows, field_rows)
    work = pattern.fields[2]
    assert (work.unit, work.rag_output) == ("時間", "omit")
    assert pattern.title_fields == ["report_id", "equipment_id"] and pattern.version_no == 1
    _, rows = pattern_to_rows(pattern)
    assert rows[2]["unit"] == "時間" and rows[2]["rag_output"] == "omit"
    assert pattern_to_meta(pattern)["title_fields"] == ["report_id", "equipment_id"]

    meta2, _, _, _ = parse_pattern_form({k: v for k, v in form.items() if k != "md_options_form"})
    assert meta2["md_options"]["omit_person_fields"] is True  # 設定欄のない画面からは既定のまま


def test_extraction_carries_unit_rag_output_and_title_fields(repair_infos):
    pattern = build_repair_pattern(repair_infos)
    pattern.title_fields = ["report_id", "occurred_date"]
    pattern.version_no = 4
    pattern.fields[0].rag_output = "omit"
    info = repair_infos[0]
    extraction = extract_document(info, pattern, rank_patterns(info, [pattern])[0].sheet_names)
    assert extraction["pattern"]["title_fields"] == ["report_id", "occurred_date"]
    assert extraction["pattern"]["version_no"] == 4
    work = next(f for f in extraction["fields"] if f["field_name"] == "work_hours")
    assert work["unit"] == "時間" and work["rag_output"] == "show"
    assert extraction["fields"][0]["rag_output"] == "omit"


# ---- 一覧表らしさ ----

def test_looks_like_table_sheet(repair_infos, sample_dir, tmp_path):
    from openpyxl import Workbook
    from pattern.matcher import looks_like_table_sheet, table_like_sheets

    for info in repair_infos + [load_workbook_info(sample_dir / "inspection.xlsx")]:
        assert not looks_like_table_sheet(info), info.path.name

    wb = Workbook()
    ws = wb.active
    ws.title = "一覧"
    ws["A1"] = "トラブル対応一覧"
    ws.append([])
    ws.append(["管理No", "発生日", "設備番号", "現象", "停止時間(分)"])
    for i in range(12):
        ws.append([f"TR-{i:03d}", f"2026-08-{i + 1:02d}", "EQ-001", "停止" if i % 3 else None, i * 5])
    notes = wb.create_sheet("記入要領")
    notes["A1"] = "この様式の書き方"
    wb.save(tmp_path / "list.xlsx")
    info = load_workbook_info(tmp_path / "list.xlsx")
    assert looks_like_table_sheet(info)
    assert table_like_sheets(info) == ["一覧"]
    assert looks_like_table_sheet(info.grids["一覧"]) and not looks_like_table_sheet(info, ["記入要領"])

    # 9行しかない表は一覧表とみなさない
    ws.delete_rows(13, 3)
    wb.save(tmp_path / "short.xlsx")
    assert not looks_like_table_sheet(load_workbook_info(tmp_path / "short.xlsx"))


# ---- 見出し欄の色・チェックボックス・項番・続きのシート（帳票サンプルの不一致の分析から） ----

LABEL_FILL = "FFD9E1F2"
HEADING_FILL = "FFFFF2CC"


def _book(tmp_path, name: str, sheets: dict):
    """sheets: {シート名: (cells, fills, merges)}"""
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill

    wb = Workbook()
    wb.remove(wb.active)
    for title, (cells, fills, merges) in sheets.items():
        ws = wb.create_sheet(title)
        for coord, value in cells.items():
            ws[coord] = value
        for coord, color in fills.items():
            ws[coord].fill = PatternFill("solid", fgColor=color)
        for rng in merges:
            ws.merge_cells(rng)
    wb.save(tmp_path / name)
    return load_workbook_info(tmp_path / name)


def _fields(*defs):
    from pattern.model import FieldDef, PatternDef

    fields = [FieldDef(name, labels[0], list(labels), data_type=data_type) for name, data_type, *labels in defs]
    return PatternDef(name="テスト", fields=fields)


def _values(info, pattern, sheets=None):
    extraction = extract_document(info, pattern, sheets or info.sheet_names)
    return extraction["values"], {f["field_name"]: f for f in extraction["fields"]}


def test_same_fill_neighbour_is_a_label_so_value_is_read_below(tmp_path):
    info = _book(tmp_path, "stamp.xlsx", {"報告書": (
        {"A1": "承認", "B1": "確認", "C1": "作成", "A2": "高橋", "B2": "中村", "C2": "松本",
         "A4": "報告番号", "B4": "発生日時", "A5": "MR-001", "B5": "2026-09-01 10:00",
         "A7": "故障内容", "A8": "搬送アームが停止した", "A9": "アラーム E-100", "A10": "写真"},
        {"A1": LABEL_FILL, "B1": LABEL_FILL, "C1": LABEL_FILL, "A4": LABEL_FILL, "B4": LABEL_FILL,
         "A7": HEADING_FILL, "A10": HEADING_FILL},
        [],
    )})
    pattern = _fields(("approver", "string", "承認"), ("checker", "string", "確認"), ("report_id", "string", "報告番号"),
                      ("occurred", "string", "発生日時"), ("symptom", "text", "故障内容"))
    values, _ = _values(info, pattern)
    assert values["approver"] == "高橋" and values["checker"] == "中村"  # 右隣の「確認」を値にしない
    assert values["report_id"] == "MR-001" and values["occurred"] == "2026-09-01 10:00"
    assert values["symptom"] == "搬送アームが停止した\nアラーム E-100"  # 同じ色の見出し「写真」で止まる


def test_checkbox_notation_is_read_as_the_checked_option(tmp_path):
    from excel.text import pick_checked

    assert pick_checked("□重大　■大　□中") == ("大", None)
    assert pick_checked("☑同型機　□類似設備　☑他Fab") == ("同型機、他Fab", None)
    assert pick_checked("□有　□無") == (None, "チェックの入った選択肢がありません")
    assert pick_checked("■ 交換部品") is None and pick_checked("□良\n■要観察") is None

    info = _book(tmp_path, "check.xlsx", {"報告書": ({"A1": "重要度", "B1": "□重大　■大　□中",
                                                     "A2": "流出", "B2": "□有　□無"}, {}, [])})
    values, fields = _values(info, _fields(("severity", "string", "重要度"), ("outflow", "string", "流出")))
    assert values["severity"] == "大"
    assert values["outflow"] is None and fields["outflow"]["warning"] == "チェックの入った選択肢がありません"


def test_label_notation_section_numbers_and_parentheses(tmp_path):
    from excel.text import section_stripped

    assert normalize_label("停止時間(分)") == "停止時間(分)"  # 閉じ括弧だけを消さない
    assert normalize_label("【設備NO】") == "設備no"
    assert section_stripped("３．暫定対策") == "暫定対策" and section_stripped("①何が") == "何が"
    assert section_stripped("D2 問題の記述") == "問題の記述"
    assert section_stripped("2号機") == "" and section_stripped("95.0%以上") == "" and section_stripped("1.5") == ""

    info = _book(tmp_path, "labels.xlsx", {"報告書": (
        {"A1": "３．暫定対策", "B1": "ロットをHOLD", "A2": "停止時間（分）", "B2": 45,
         "A3": "原因（推定）", "B3": "摩耗の可能性", "A4": "原因（確定）", "B4": "ベアリング摩耗"}, {}, [])})
    pattern = _fields(("repair", "text", "暫定対策"), ("downtime", "number", "停止時間"), ("cause", "text", "原因（確定）"))
    values, _ = _values(info, pattern)
    assert values["repair"] == "ロットをHOLD"
    assert values["downtime"] == 45  # 候補どおりのラベルが無いときだけ括弧書きの違いを許す
    assert values["cause"] == "ベアリング摩耗"  # 「原因（推定）」と取り違えない


def test_label_in_a_table_header_is_used_only_when_there_is_no_other(tmp_path):
    cells = {"A1": "■ 時系列", "A2": "日時", "B2": "対応内容", "C2": "担当",
             "A3": "9/1 10:00", "B3": "アラーム発報", "C3": "山田", "A4": "9/1 10:30", "B4": "部品交換", "C4": "鈴木",
             "A6": "対応内容", "B6": "ベアリングを交換し、動作確認した"}
    fills = {"A2": LABEL_FILL, "B2": LABEL_FILL, "C2": LABEL_FILL, "A6": LABEL_FILL}
    info = _book(tmp_path, "timeline.xlsx", {"報告書": (cells, fills, [])})
    values, _ = _values(info, _fields(("repair", "text", "対応内容")))
    assert values["repair"] == "ベアリングを交換し、動作確認した"


def test_continuation_sheet_is_selected_with_the_main_sheet(tmp_path):
    from pattern.matcher import match_pattern
    from pattern.model import SheetDef

    first = ({"A1": "管理番号", "B1": "8D-001", "A2": "現象", "B2": "搬送停止", "A3": "原因", "B3": "摩耗"}, {}, [])
    second = ({"A1": "管理番号", "B1": "8D-001", "A2": "効果判定", "B2": "再発なし", "A3": "再発防止策", "B3": "月次点検"}, {}, [])
    notes = ({"A1": "記入要領", "A2": "効果判定", "A3": "再発防止策"}, {}, [])
    info = _book(tmp_path, "8d.xlsx", {"記入要領": notes, "8D報告(1)": first, "8D報告(2)": second})
    pattern = _fields(("report_id", "string", "管理番号"), ("symptom", "text", "現象"), ("cause", "text", "原因"),
                      ("result", "text", "効果判定"), ("prevention", "text", "再発防止策"))
    pattern.sheets = [SheetDef("8D報告(1)")]
    match = match_pattern(info, pattern)
    assert match.sheet_names == ["8D報告(1)", "8D報告(2)"]  # 値の無い記入要領は足さない
    values, _ = _values(info, pattern, match.sheet_names)
    assert values["result"] == "再発なし" and values["prevention"] == "月次点検"


# ---- 明細表 ----

def _parts_book(tmp_path, name="parts.xlsx", extra: dict | None = None):
    cells = {"A1": "報告番号", "B1": "MR-001",
             "A3": "■ 交換部品", "A4": "No.", "B4": "品番", "C4": "品名", "D4": "数量",
             "A5": 1, "B5": "PW48-1591", "C5": "ベアリング", "D5": 2,
             "A6": 2, "B6": "PW35-1577", "C6": "スピンモータ", "D6": 1,
             "A7": 3, "A8": 4,  # 様式の空き行（No だけ）
             "A9": "合計", "D9": 3,
             "A10": "所見", "B10": "異常なし"}
    cells.update(extra or {})
    fills = {c: LABEL_FILL for c in ("A4", "B4", "C4", "D4", "A9")}
    fills["A10"] = HEADING_FILL
    return _book(tmp_path, name, {"報告書": (cells, fills, ["A9:C9"])})


def test_table_field_reads_rows_until_total_row(tmp_path):
    info = _parts_book(tmp_path)
    values, fields = _values(info, _fields(("report_id", "string", "報告番号"), ("parts", "table", "交換部品")))
    assert values["report_id"] == "MR-001"
    assert values["parts"] == {
        "columns": ["品番", "品名", "数量"],  # 連番だけの No 列は落とす
        "rows": [["PW48-1591", "ベアリング", "2"], ["PW35-1577", "スピンモータ", "1"], ["合計", "", "3"]],
    }
    assert fields["parts"]["label_cell"] == "A3" and fields["parts"]["value_cell"] == "A4:D9"


def test_table_with_vertical_anchor_merged_cells_and_heading_stop(tmp_path):
    info = _book(tmp_path, "check.xlsx", {"点検": (
        {"A1": "使用部品", "B1": "品番", "C1": "品名", "B2": "PX84-2134", "C2": "スイッチ", "A4": "結果", "B4": "良",
         "A6": "Ａ．機構部", "A7": "点検部位", "B7": "点検項目", "C7": "判定",
         "A8": "ステージ", "B8": "位置偏差", "C8": "○", "B9": "冷却水流量", "C9": "×", "B10": "温度", "C10": "○",
         "A11": "総合判定", "B11": "要観察"},
        {"A1": LABEL_FILL, "B1": LABEL_FILL, "C1": LABEL_FILL, "A4": LABEL_FILL, "A7": LABEL_FILL, "B7": LABEL_FILL,
         "C7": LABEL_FILL, "A11": HEADING_FILL},
        ["A1:A3", "A8:A10"],
    )})
    values, _ = _values(info, _fields(("parts", "table", "使用部品"), ("check", "table", "機構部")))
    # 縦結合の見出しは右の列見出しから、見出しの範囲の行だけを読む
    assert values["parts"] == {"columns": ["品番", "品名"], "rows": [["PX84-2134", "スイッチ"]]}
    # 縦結合のデータセル（点検部位）は各行に同じ値。表の左端の色付きの「総合判定」で止まる
    assert values["check"]["rows"] == [["ステージ", "位置偏差", "○"], ["ステージ", "冷却水流量", "×"],
                                       ["ステージ", "温度", "○"]]


def test_table_none_rows_and_column_fallback(tmp_path):
    from excel.tables import find_table_by_columns

    info = _parts_book(tmp_path, "none.xlsx", {"B5": "なし", "C5": None, "D5": None, "A6": None, "B6": None,
                                                 "C6": None, "D6": None, "A9": None, "D9": None})
    values, fields = _values(info, _fields(("parts", "table", "交換部品")))
    assert values["parts"] is None and fields["parts"]["warning"] == "表に行がありません"

    other = _parts_book(tmp_path, "renamed.xlsx", {"A3": "■ 部品交換実績"})
    pattern = _fields(("parts", "table", "交換部品"))
    assert _values(other, pattern)[0]["parts"] is None
    pattern.fields[0].table_columns = ["No.", "品番", "品名", "数量", "単価"]
    values, fields = _values(other, pattern)  # 見出しが違っても列見出しが似た表を読む
    assert values["parts"]["rows"][0] == ["PW48-1591", "ベアリング", "2"] and fields["parts"]["label_cell"] == "A3"
    assert find_table_by_columns(other.grids["報告書"], ["品番"]) is None  # 列見出し1つでは探さない


def test_builder_suggests_table_field_and_skips_column_header_fields(tmp_path):
    infos = [_parts_book(tmp_path, "a.xlsx"), _parts_book(tmp_path, "b.xlsx", {"A3": "■ 使用部品"})]
    _, rows = suggest_rows(infos)
    tables = [r for r in rows if r["data_type"] == "table"]
    assert len(tables) == 1  # 見出しの書き方が違っても、列見出しが同じ表は1つの候補
    parts = tables[0]
    assert parts["use"] and parts["field_name"] == "parts" and parts["display_name"] == "交換部品"
    assert {"■ 交換部品", "■ 使用部品"} <= set(parts["candidates"].splitlines())
    assert parts["table_columns"].splitlines() == ["No.", "品番", "品名", "数量"]
    assert parts["examples"][0].startswith("品番: PW48-1591／品名: ベアリング")
    used_labels = {r["display_name"] for r in rows if r["use"]}
    assert not used_labels & {"品番", "品名", "数量", "No."}  # 列見出しは1つの値の項目にしない

    # 見出しの行＋値の行1つ（「影響｜停止時間｜影響ロット」）は明細表にしない
    single = _book(tmp_path, "impact.xlsx", {"報告書": (
        {"A1": "影響", "B1": "停止時間", "C1": "影響ロット", "B2": 355, "C2": "PX2321022.2"},
        {"A1": LABEL_FILL, "B1": LABEL_FILL, "C1": LABEL_FILL}, ["A1:A2"])})
    _, rows = suggest_rows([single])
    assert not any(r["use"] for r in rows if r["data_type"] == "table")
    assert next(r for r in rows if r["field_name"] == "downtime")["use"]


def test_table_columns_roundtrip_through_form_rows():
    from pattern.forms import parse_pattern_form, pattern_to_rows

    form = {"name": "点検報告書", "fields-0-use": "on", "fields-0-field_name": "parts", "fields-0-display_name": "交換部品",
            "fields-0-candidates": "■ 交換部品", "fields-0-data_type": "table", "fields-0-direction": "auto",
            "fields-0-table_columns": "品番\n品名\n数量"}
    meta, sheets, rows, errors = parse_pattern_form(form)
    assert errors == []
    pattern = rows_to_pattern(1, meta, sheets, rows)
    assert pattern.fields[0].data_type == "table" and pattern.fields[0].table_columns == ["品番", "品名", "数量"]
    assert pattern_to_rows(pattern)[1][0]["table_columns"] == "品番\n品名\n数量"


def test_manual_table_edit_is_parsed_from_json(tmp_path):
    import json

    from excel.extractor import apply_manual_values

    info = _parts_book(tmp_path)
    extraction = extract_document(info, _fields(("parts", "table", "交換部品")), ["報告書"])
    edited = {"columns": ["品番", "品名", "数量"], "rows": [["PW48-1591", "ベアリング", "4"], ["", "", ""]]}
    apply_manual_values(extraction, {"value-parts": json.dumps(edited, ensure_ascii=False)})
    field = extraction["fields"][0]
    assert field["value"] == {"columns": ["品番", "品名", "数量"], "rows": [["PW48-1591", "ベアリング", "4"]]}
    assert field["edited"] and extraction["values"]["parts"] == field["value"]
    apply_manual_values(extraction, {"value-parts": "{壊れた"})
    assert field["value"]["rows"][0][2] == "4" and "形式" in field["warning"]
    apply_manual_values(extraction, {"value-parts": ""})
    assert field["value"] is None


def test_evaluator_compares_table_rows_and_treats_empty_marks_as_empty():
    from scripts.samples.evaluate_forms import column_value, compare, compare_table

    assert compare("string", "なし", "") == "empty_ok" and compare("string", None, "－") == "empty_ok"
    assert compare("string", "佐々木\n5/17", "佐々木") == "normalized"
    expected = [{"no": 1, "part_no": "PW-1", "qty": "１", "amount": "21,000"}, {"no": 2, "part_no": "PW-2", "qty": "2"}]
    got = {"columns": ["品番", "数量", "金額"], "rows": [["PW-1", "1", "21000"], ["PW-3", "2", ""]]}
    assert compare_table(got, expected) == {"files": 1, "expected_rows": 2, "extracted_rows": 2, "matched_rows": 1,
                                            "exact_tables": 0}
    lots = {"columns": ["ロットNo.", "投入数"], "rows": [["L1", "25"], ["L2", "13"], ["合計", "38"]]}
    assert column_value(lots, "ロットNo.") == "L1\nL2" and column_value(lots, "投入数") == "38"
    assert column_value(lots, "品番") is False


def test_combined_label_takes_the_matching_part_of_the_value(tmp_path):
    from excel.text import label_parts

    assert label_parts("ライン／工程") == ("ライン", "工程") and label_parts("L/min") == ()
    info = _book(tmp_path, "combined.xlsx", {"報告書": (
        {"A1": "ライン／工程", "B1": "L6／STI-CMP", "A2": "設備名/号機", "B2": "CMP 8号機"}, {}, [])})
    values, _ = _values(info, _fields(("line", "string", "ライン"), ("process", "string", "工程"),
                                      ("name", "string", "設備名")))
    assert values["line"] == "L6" and values["process"] == "STI-CMP"
    assert values["name"] is None  # 値の区切りの数が合わなければ読まない（取り違えるより空欄）


def test_table_header_below_a_heading_is_not_read_as_a_value(tmp_path):
    info = _parts_book(tmp_path, "heading.xlsx", {"A3": "暫定対策"})
    values, _ = _values(info, _fields(("repair", "text", "暫定対策")))
    assert values["repair"] is None  # 「No｜品番…」の列を文章として読まない


def test_duration_text_is_converted_to_the_field_unit():
    assert to_number("3時間40分", "3時間40分", "分") == (220, "「3時間40分」を分に換算しました")
    assert to_number("１時間３０分", "１時間３０分", "h")[0] == 1.5
    assert to_number("3時間40分", "3時間40分")[0] == 220  # 単位の無い項目は分にする（3 と読まない）


# ---- 年の無い日付・単位の無い数値（2026-09-19 の利用者の判断） ----

def test_date_without_a_year_keeps_the_original_text_and_warns():
    """「2/12 3時17分」に年は補わない（外れたときに誤った日付を Markdown に書くことになるため）。"""
    for text in ("2/12 3時17分", "2月12日 14:05", "10/16", "12-25 08:00"):
        value, warning = to_date(text, text)
        assert value == text, text
        assert warning and warning.startswith("年が書かれていません"), text
    # 年のある日付・年が2桁の書き方・そもそも日付でない記入は、今までどおり
    # 年のある日付は時刻（「3時17分」）も残す（F3-2: 以前は時刻を黙って落としていた）
    assert to_date("2026/2/12 3時17分", "2026/2/12 3時17分") == ("2026-02-12 03:17", None)
    assert to_date("24/8/25", "24/8/25")[1] == "日付として解釈できません"
    assert to_date("未復旧（対応中）", "未復旧（対応中）")[1] == "日付として解釈できません"
    assert to_date("13/40", "13/40")[1] == "日付として解釈できません"  # 月日として成り立たない数


def test_numeric_unit_comes_from_the_field_or_the_written_value():
    from excel.text import numeric_unit

    assert numeric_unit("390", "分") == "分" and numeric_unit("390", "h") == "時間"
    assert numeric_unit("390分") == "分" and numeric_unit("1,032分") == "分" and numeric_unit("2.5h") == "時間"
    assert numeric_unit("3時間40分") == "分"  # to_number が分に換算する
    assert numeric_unit("390") == "" and numeric_unit("CMP-101") == ""


def test_number_without_a_unit_is_flagged_only_when_the_unit_changes_the_meaning(tmp_path):
    """単位の無い数値は「単位不明」で要確認。ただし件数・回数のように単位が要らない項目は警告しない。"""
    from pattern.dictionary import ambiguous_unit_hint
    from pattern.model import FieldDef, PatternDef

    assert ambiguous_unit_hint("downtime", "停止時間") and ambiguous_unit_hint("work_hours", "作業時間")
    assert ambiguous_unit_hint("cost", "修理費用") and ambiguous_unit_hint("f1", "部品の外径")
    assert not ambiguous_unit_hint("f2", "発生件数") and not ambiguous_unit_hint("f3", "対応回数")
    assert not ambiguous_unit_hint("f4", "立会人数") and not ambiguous_unit_hint("f5", "不良率")
    assert not ambiguous_unit_hint("report_id", "報告番号")

    info = _book(tmp_path, "unit.xlsx", {"報告書": (
        {"A1": "停止時間", "B1": "390", "A2": "発生件数", "B2": "3", "A3": "作業時間", "B3": "2.5h"},
        {"A1": LABEL_FILL, "A2": LABEL_FILL, "A3": LABEL_FILL}, [])})
    pattern = PatternDef(name="テスト", fields=[
        FieldDef("downtime", "停止時間", ["停止時間"], data_type="number"),
        FieldDef("count", "発生件数", ["発生件数"], data_type="number"),
        FieldDef("work_hours", "作業時間", ["作業時間"], data_type="number")])
    values, fields = _values(info, pattern)
    assert values["downtime"] == 390 and fields["downtime"]["unit"] == ""
    assert "単位が書かれていません" in fields["downtime"]["warning"]
    assert "「帳票の種類」の画面" in fields["downtime"]["warning"]  # 直し方を書く
    # 読み取り済みの帳票はその場で直せる（値に単位を付けて入力）ことを先に書く。種類の単位は読み取り済みに反映されない
    assert "値に単位を付けて入力してください（例: 390分）" in fields["downtime"]["warning"]
    assert "読み取り済みの帳票には反映されません" in fields["downtime"]["warning"]
    assert values["count"] == 3 and not fields["count"]["warning"]      # 件数は単位が無くて当たり前
    assert values["work_hours"] == 2.5 and fields["work_hours"]["unit"] == "時間"  # セルの「h」から補う
    assert not fields["work_hours"]["warning"]


def test_written_unit_wins_over_the_field_unit_and_is_flagged(tmp_path):
    """「停止時間（分）」の欄に「14.9h」と書かれた帳票。勝手に分として出さず、書かれたとおりにして要確認にする。"""
    from pattern.model import FieldDef, PatternDef

    info = _book(tmp_path, "unit2.xlsx", {"報告書": (
        {"A1": "停止時間", "B1": "14.9h"}, {"A1": LABEL_FILL}, [])})
    pattern = PatternDef(name="テスト", fields=[
        FieldDef("downtime", "停止時間", ["停止時間"], data_type="number", unit="分")])
    _, fields = _values(info, pattern)
    assert fields["downtime"]["value"] == 14.9 and fields["downtime"]["unit"] == "時間"
    assert "「分」ですが" in fields["downtime"]["warning"] and "「時間」" in fields["downtime"]["warning"]


# ---- 表記違いのラベル・隣の値の優先・一覧の列見出し・設備番号と設備名のまとめ欄（帳票サンプルの不一致の分析 2回目） ----

def test_label_variants_are_used_when_the_exact_label_gives_no_value(tmp_path):
    from excel.text import label_base, split_label_unit

    assert split_label_unit("発生原因（推定）") == ("発生原因（推定）", "")  # 「推定」は単位でない
    assert split_label_unit("金額（千円）") == ("金額", "千円")
    assert label_base("応急処置内容") == "応急処置" and label_base("復旧完了日時") == "復旧完了"
    assert label_base("発生原因(なぜ起きたか)") == "発生原因" and label_base("処置") == ""

    cells = {"A1": "発生原因（なぜ起きたか）", "B1": "ベアリング摩耗",
             "A2": "応急処置内容", "B2": "ベアリング交換",
             "A3": "復旧完了", "B3": "2026/9/2",
             "A4": "原因（推定）", "B4": "摩耗の可能性", "A5": "原因（確定）",
             "A6": "状況", "B6": "完了", "B8": "次回確認は9/30"}
    fills = {c: LABEL_FILL for c in ("A1", "A2", "A3", "A4", "A5", "A6")}
    info = _book(tmp_path, "variants.xlsx", {"報告書": (cells, fills, [])})
    values, _ = _values(info, _fields(("cause", "text", "発生原因"), ("repair", "text", "応急処置"),
                                      ("completed", "date", "復旧完了日時", "完了日時"),
                                      ("confirmed", "text", "原因（確定）")))
    assert values["cause"] == "ベアリング摩耗" and values["repair"] == "ベアリング交換"
    assert values["completed"] == "2026-09-02"  # 値の「完了」（塗りなし）は「完了日時」の表記違いにしない
    assert values["confirmed"] is None  # 「原因（推定）」は括弧書きが違う別の項目


def test_adjacent_value_first_and_headings_are_not_values(tmp_path):
    cells = {"C1": "発信\n部署", "D1": "承認", "D2": "佐々木", "C5": "2024/2/1",
             "A6": "発信部署", "B6": "製造課 A班",
             "A8": "備考", "A10": "▼ 回答欄（受信部署で記入）", "A11": "回答者", "B11": "木村",
             "A13": "品証確認", "A14": "様式QA-05 Rev.2"}
    fills = {c: LABEL_FILL for c in ("C1", "D1", "A6", "A8", "A11", "A13")}
    fills["A10"] = HEADING_FILL
    info = _book(tmp_path, "stamp.xlsx", {"報告書": (cells, fills, ["C1:C3"])})
    values, _ = _values(info, _fields(("department", "string", "発信部署"), ("remarks", "text", "備考"),
                                      ("approver", "string", "承認"), ("qa", "string", "品証確認")))
    assert values["department"] == "製造課 A班"  # 押印欄の縦書き「発信部署」の2行下の日付ではない
    assert values["remarks"] is None  # 「▼ 回答欄」の見出しは値にしない
    assert values["approver"] == "佐々木"
    assert values["qa"] is None  # 欄外の様式番号は値にしない


def test_list_column_headers_and_seq_no_are_not_single_value_labels(tmp_path):
    cells = {"A1": "No.", "B1": "工異24-021", "C1": "承認", "D1": "確認", "E1": "作成",
             "C2": "高橋", "D2": "斎藤", "E2": "森", "C3": "6/5", "D3": "6/2", "E3": "5/31",
             "A5": "対象設備", "B5": "CMP-108　STI-CMP 8号機",
             "A7": "確認", "B7": "設備No", "C7": "設備名", "D7": "結果",
             "A8": "□", "B8": "CMP-101", "C8": "酸化膜CMP 1号機",
             "A9": "☑", "B9": "CMP-103", "C9": "W-CMP 3号機", "D9": "異常なし",
             "A10": "□", "B10": "CMP-104", "C10": "Cu-CMP 4号機",
             "A12": "■ 対象ロット", "A13": "No.", "B13": "ロットNo.", "C13": "数量",
             "A14": 1, "B14": "SR2503056.2", "C14": 12}
    fills = {c: LABEL_FILL for c in ("A1", "C1", "D1", "E1", "A5", "A7", "B7", "C7", "D7", "A13", "B13", "C13")}
    info = _book(tmp_path, "list.xlsx", {"報告書": (cells, fills, [])})
    values, _ = _values(info, _fields(("report_id", "string", "No", "管理番号"), ("approver", "string", "承認"),
                                      ("equipment_id", "string", "設備No"), ("equipment_name", "string", "設備名")))
    assert values["report_id"] == "工異24-021"  # 表の連番の「No.」の下の 1 ではない
    assert values["approver"] == "高橋"  # 押印と日付の2行の押印欄は、一覧の列見出しとみなさない
    # 一覧（水平展開先）の列見出しでなく、番号と名前をまとめた「対象設備」を分けて読む
    assert values["equipment_id"] == "CMP-108" and values["equipment_name"] == "STI-CMP 8号機"


def test_split_code_name_and_builder_suggests_equipment_from_combined_cell(tmp_path):
    from excel.text import split_code_name

    assert split_code_name("ROB-821（ウェーハソーター 1号機）") == ("ROB-821", "ウェーハソーター 1号機")
    assert split_code_name("W-CMP 3号機　CMP-103") == ("CMP-103", "W-CMP 3号機")
    assert split_code_name("EQ-001：プレス機") == ("EQ-001", "プレス機")
    assert split_code_name("CMP-101 / CMP-102") is None and split_code_name("L6／STI-CMP") is None
    assert split_code_name("CMP-108") is None and split_code_name("プレス機") is None

    info = _book(tmp_path, "combined_equipment.xlsx", {"報告書": (
        {"A1": "報告番号", "B1": "R-001", "A2": "使用設備", "B2": "CVD-202（TEOS CVD 2号機）", "A3": "設備", "B3": "空調"},
        {"A1": LABEL_FILL, "A2": LABEL_FILL, "A3": LABEL_FILL}, [])})
    _, rows = suggest_rows([info])
    by_name = {r["field_name"]: r for r in rows}
    assert by_name["equipment_id"]["use"] and by_name["equipment_id"]["examples"] == ["CVD-202"]
    assert by_name["equipment_name"]["examples"] == ["TEOS CVD 2号機"]
    assert "使用設備" in by_name["equipment_name"]["candidates"].splitlines()
    pattern = rows_to_pattern(1, META, *suggest_rows([info]))
    values, _ = _values(info, pattern)
    assert values["equipment_id"] == "CVD-202" and values["equipment_name"] == "TEOS CVD 2号機"


def test_evaluator_offset_picks_other_sample_files():
    from scripts.samples.evaluate_forms import pick_samples

    entries = [{"file": f"{v}{i}", "layout_version": v} for v in ("a", "b") for i in range(3)]
    assert [e["file"] for e in pick_samples(entries, 3)] == ["a0", "b0", "a1"]
    assert [e["file"] for e in pick_samples(entries, 3, offset=1)] == ["b1", "a1", "b2"]  # 版の順も1つずらす
    assert len({e["file"] for e in pick_samples(entries, 6, offset=4)}) == 6


def test_form_number_note_is_not_a_value_even_in_the_middle_of_a_cell(tmp_path):
    """欄外の様式番号は、先頭になくても項目の値にしない（design.md 6.1「値にしないもの」）。"""
    from excel.extractor import is_form_number_note

    assert is_form_number_note("様式MT-031 Rev.1")
    assert is_form_number_note("製造部 設備保全課 様式MT-031 Rev.2(2025.04改訂)")
    assert not is_form_number_note("様式を見直す")  # 番号が続かない文は注記ではない
    assert not is_form_number_note("旧様式AからBへ移行したが、記入の手順は変えていないので、" * 3)

    info = _book(tmp_path, "form_no.xlsx", {"報告書": (
        {"A1": "設備点検・保全作業報告書", "F1": "製造部 設備保全課 様式MT-031 Rev.2(2025.04改訂)",
         "A3": "作業No.", "B3": "WK-001"},
        {"A1": LABEL_FILL, "A3": LABEL_FILL}, [])})
    values, _ = _values(info, _fields(("title", "string", "設備点検・保全作業報告書"), ("work_no", "string", "作業No.")))
    assert values["title"] is None and values["work_no"] == "WK-001"


def test_stacked_detail_table_blocks_are_all_read(tmp_path):
    """同じ形の列見出しが縦に積み重なる表（特性要因図）は、組ごとに読んで1つの明細表にまとめる。"""
    from pattern.model import FieldDef, PatternDef

    cells = {"A1": "特性要因図", "A2": "人", "B2": "機械", "A3": "教育不足", "B3": "電極の摩耗",
             "A4": "材料", "B4": "方法", "A5": "ガスの純度", "B5": "点検手順の不備",
             "A6": "測定", "B6": "環境", "A7": "校正ずれ", "B7": "室温の変動"}
    fills = {c: LABEL_FILL for c in ("A1", "A2", "B2", "A4", "B4", "A6", "B6")}
    info = _book(tmp_path, "fishbone.xlsx", {"報告書": (cells, fills, [])})
    pattern = PatternDef(name="テスト", fields=[FieldDef("cause_table", "特性要因図", ["特性要因図"], data_type="table")])
    values, fields = _values(info, pattern)
    value = values["cause_table"]
    assert value["columns"] == ["人", "機械", "材料", "方法", "測定", "環境"]
    # 行ごとに自分の組の列だけ埋まる（どの要因がどの見出しのものか、1行だけで分かる）
    assert value["rows"] == [["教育不足", "電極の摩耗", "", "", "", ""],
                             ["", "", "ガスの純度", "点検手順の不備", "", ""],
                             ["", "", "", "", "校正ずれ", "室温の変動"]]
    assert fields["cause_table"]["table_blocks"] == 3
    assert not fields["cause_table"]["warning"]  # 取りこぼしが無くなったので要確認にしない


def test_stacked_blocks_with_the_same_columns_just_add_rows(tmp_path):
    """列見出しが同じまま繰り返される様式（ページごとに見出しを書く表）は、行が増えるだけにする。"""
    from pattern.model import FieldDef, PatternDef

    cells = {"A1": "■ 交換部品", "A2": "品番", "B2": "品名", "A3": "PW48-1591", "B3": "ベアリング",
             "A4": "品番", "B4": "品名", "A5": "PW35-1577", "B5": "スピンドル"}
    fills = {c: LABEL_FILL for c in ("A1", "A2", "B2", "A4", "B4")}
    info = _book(tmp_path, "parts2.xlsx", {"報告書": (cells, fills, [])})
    pattern = PatternDef(name="テスト", fields=[FieldDef("parts", "交換部品", ["交換部品"], data_type="table")])
    values, fields = _values(info, pattern)
    assert values["parts"] == {"columns": ["品番", "品名"],
                               "rows": [["PW48-1591", "ベアリング"], ["PW35-1577", "スピンドル"]]}
    assert fields["parts"]["table_blocks"] == 2


def test_builder_leaves_notes_legends_and_dashes_unchecked(tmp_path):
    """凡例・注記・「－」だけの欄は候補には残すが、既定ではチェックを付けない（design.md 6.1）。"""
    info = _book(tmp_path, "legend.xlsx", {"報告書": (
        {"A1": "判定　○", "B1": "良好　△：要観察　×：不良　－：対象外",
         "A2": "関連TR No.", "B2": "※故障発生日時・停止時間は事後保全時に記入",
         "A3": "品名", "B3": "－",
         "A4": "作業区分", "B4": "定期点検"},
        {c: LABEL_FILL for c in ("A1", "A2", "A3", "A4")}, [])})
    _, rows = suggest_rows([info])
    by_label = {r["display_name"]: r for r in rows}
    assert by_label["作業区分"]["use"]
    for label in ("判定　○", "関連TR No.", "品名"):
        assert not by_label[label]["use"], label
        assert label in by_label  # 候補としては残し、人が選べるようにする


def test_builder_reads_date_like_labels_written_as_text(tmp_path):
    """見本が文字列で日付を持つ欄（「2025/12/19(金)」「令和5年12月5日」）も日付型にする。"""
    info = _book(tmp_path, "dates.xlsx", {"報告書": (
        {"A1": "作業日", "B1": "2025/12/19(金)", "A2": "回答日", "B2": "令和5年12月5日",
         "A3": "曜日", "B3": "金", "A4": "作業日数", "B4": "3"},
        {c: LABEL_FILL for c in ("A1", "A2", "A3", "A4")}, [])})
    _, rows = suggest_rows([info])
    by_label = {r["display_name"]: r["data_type"] for r in rows}
    assert by_label["作業日"] == "date" and by_label["回答日"] == "date"
    assert by_label["曜日"] == "string"  # 値が日付として読めないものは日付にしない

    pattern = rows_to_pattern(1, META, *suggest_rows([info]))
    values, _ = _values(info, pattern)
    assert values[next(f.field_name for f in pattern.fields if f.display_name == "作業日")] == "2025-12-19"


def test_builder_reads_prefixed_number_label_as_report_id(tmp_path):
    """「8D No.」のように短い接頭辞が付いた識別番号の見出しも報告番号として読む。"""
    info = _book(tmp_path, "8d.xlsx", {"報告書": (
        {"A1": "8D No.", "B1": "8D-2023-011", "A2": "設備名", "B2": "P-SiN CVD 5号機"},
        {"A1": LABEL_FILL, "A2": LABEL_FILL}, [])})
    _, rows = suggest_rows([info])
    report = next(r for r in rows if r["field_name"] == "report_id")
    assert report["use"] and "8D No." in report["candidates"].splitlines()
    values, _ = _values(info, rows_to_pattern(1, META, *suggest_rows([info])))
    assert values["report_id"] == "8D-2023-011"


def test_empty_detail_table_keeps_columns_for_manual_entry(tmp_path):
    """明細表が読めなくても、列見出しは確認画面に残して人が行を入れられるようにする。"""
    from excel.tables import clean_table_value
    from pattern.model import FieldDef, PatternDef

    info = _book(tmp_path, "no_table.xlsx", {"報告書": (
        {"A1": "■ 交換部品", "A2": "なし"}, {"A1": HEADING_FILL}, [])})
    pattern = PatternDef(name="テスト", fields=[
        FieldDef("parts", "交換部品", ["■ 交換部品"], data_type="table", table_columns=["品番", "品名", "数量"])])
    _, fields = _values(info, pattern)
    assert fields["parts"]["value"] is None and fields["parts"]["table_columns"] == ["品番", "品名", "数量"]

    # 全部の行を消しても列見出しは残る（表の入力欄が消えない）。Markdown 側は行が無ければ出さない
    assert clean_table_value({"columns": ["品番", "品名"], "rows": []}) == {"columns": ["品番", "品名"], "rows": []}
    assert clean_table_value({"columns": [], "rows": []}) is None


def test_label_that_is_also_a_column_header_keeps_the_real_header_row(tmp_path):
    """見出し語が表の列見出しと同じでも、1行下のデータ行を列見出しにしない（F3 の D6）。

    見本と違う書き方の見出し（「D6．恒久対策の実施と効果の確認」）だと、候補ラベルのうち表の列見出しと
    同じ「実施内容」だけが当たる。そこから下に表を探すと先頭の明細が列見出しになり、1件目と
    完了日・状況の列が落ちていた。
    """
    from pattern.model import FieldDef, PatternDef

    cells = {"A1": "D6．恒久対策の実施と効果の確認",
             "A2": "No", "B2": "実施内容", "C2": "担当", "D2": "完了日", "E2": "状況",
             "A3": 1, "B3": "樹脂交換時期を実負荷ベースで再設定", "C3": "前田", "D3": "2026-09-01", "E3": "完了",
             "A4": 2, "B4": "比抵抗の早期ワーニング閾値追加", "C4": "前田", "D4": "2026-09-05", "E4": "完了",
             "A5": 3, "B5": "供給水比抵抗をFDCで監視", "C5": "中島", "D5": "2026-09-10", "E5": "実施中"}
    fills = {c: LABEL_FILL for c in ("A1", "A2", "B2", "C2", "D2", "E2")}
    info = _book(tmp_path, "8d_d6.xlsx", {"報告書": (cells, fills, [])})
    fd = FieldDef("d6", "恒久対策の実施・効果確認", ["D6：恒久対策の実施・効果確認", "D6", "実施内容"], data_type="table")
    fd.table_columns = ["No", "実施内容", "担当", "完了日", "状況"]  # 見本から作った列見出し
    values, fields = _values(info, PatternDef(name="テスト", fields=[fd]))
    assert values["d6"]["columns"] == ["実施内容", "担当", "完了日", "状況"]  # 連番だけの No 列は落とす
    assert values["d6"]["rows"] == [
        ["樹脂交換時期を実負荷ベースで再設定", "前田", "2026-09-01", "完了"],
        ["比抵抗の早期ワーニング閾値追加", "前田", "2026-09-05", "完了"],
        ["供給水比抵抗をFDCで監視", "中島", "2026-09-10", "実施中"],
    ]
    assert not fields["d6"]["warning"]
