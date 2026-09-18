"""一覧表の読み取り（表ソース・見出し帯・行分類・列の対応づけ候補）のテスト。"""
from __future__ import annotations

from datetime import datetime, time
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils.datetime import CALENDAR_MAC_1904

from tables.csv_source import CsvSource, describe_sniff, sniff_csv
from tables.detect import classify_rows, guess_layout, is_month_label, looks_like_list, sample_data_rows, split_header_unit
from tables.dictionary import BY_KEY, lookup_header
from tables.excel_source import ExcelSource
from tables.mapping import match_templates, suggest_columns
from tables.source import UploadError, open_source, value_kind

SAMPLES = Path(__file__).resolve().parents[1] / "samples" / "tables"


# ---- CSV ----

def _write(path: Path, text: str, encoding: str) -> Path:
    path.write_bytes(text.encode(encoding))
    return path


def _system_csv(n: int = 3) -> str:
    lines = [
        "出力日時,2026/09/14 18:05:33",
        "出力者,M10544,中村 浩二",
        '抽出条件,期間=2023/04/01〜2026/08/31,"記録区分=1:設備故障,2:チョコ停",部門=21000',
        "",
        "管理番号,設備コード,発生日時,現象,処置内容,停止時間(分)",
    ]
    for i in range(1, n + 1):
        lines.append(f'TR-2026-{i:05d},CMP-10{i},2026/08/0{i} 10:00:00,髙橋さん連絡①,"1. 点検\n2. 交換",{i * 10}')
    lines.append(f"合計件数,{n}")
    return "\r\n".join(lines) + "\r\n"


def test_sniff_cp932_preamble_multiline_and_trailer(tmp_path):
    path = _write(tmp_path / "故障履歴.csv", _system_csv(), "cp932")
    sniff = sniff_csv(path)
    assert sniff.encoding == "cp932" and not sniff.bom
    assert sniff.delimiter == ","
    assert sniff.header_row == 5 and sniff.preamble_rows == 4
    assert sniff.trailer_rows == 1
    assert sniff.warnings == []
    assert describe_sniff(sniff) == "文字コード: CP932（Shift_JIS） / 区切り: カンマ / 前置き行: 4行"

    src = open_source(path, "故障履歴.csv")
    rows = list(src.rows("故障履歴.csv"))
    assert [r.index for r in rows] == list(range(1, 10))  # 引用符内の改行は1レコード
    assert rows[5].cells[4].text == "1. 点検\n2. 交換"
    assert rows[5].cells[3].text == "髙橋さん連絡①"
    assert rows[3].is_blank and rows[3].hidden is None

    layout = guess_layout(src, "故障履歴.csv")
    assert layout.table_kind == "list"
    assert layout.header_rows == [5] and layout.data_start == 6
    assert layout.headers == ["管理番号", "設備コード", "発生日時", "現象", "処置内容", "停止時間(分)"]
    assert layout.data_end == 9  # 末尾の件数行は合計行として範囲に含め、分類で除く
    assert layout.counts["data"] == 3 and layout.counts["subtotal"] == 1
    kinds = {rc.index: rc.kind for rc in layout.row_classes}
    assert kinds[1] == "title" and kinds[4] == "blank" and kinds[5] == "header" and kinds[9] == "subtotal"


def test_sniff_utf8_bom_embedded_newlines(tmp_path):
    text = '台帳No,起票日,不具合内容\n' + "".join(
        f'CA-2026-{i:04d},R8.4.{i},"■発生：R8.4.{i}\n■設備：CMP-101, ""A""系"\n' for i in range(1, 8)
    )
    path = tmp_path / "台帳.csv"
    path.write_bytes(b"\xef\xbb\xbf" + text.encode("utf-8"))
    sniff = sniff_csv(path)
    assert (sniff.encoding, sniff.bom, sniff.delimiter, sniff.header_row) == ("utf-8-sig", True, ",", 1)
    src = CsvSource(path)
    rows = list(src.rows())
    assert len(rows) == 8
    assert rows[0].cells[0].text == "台帳No"  # BOM は見出しに残らない
    assert rows[1].cells[2].text == '■発生：R8.4.1\n■設備：CMP-101, "A"系'
    layout = guess_layout(src, src.file_name)
    assert layout.header_rows == [1] and layout.counts["data"] == 7


def test_sniff_prefers_consistent_tab_over_comma(tmp_path):
    lines = ["管理No\t発生日\t処置"]
    for i in range(20):
        commas = "," * (i % 4)
        lines.append(f"A-{i:03d}\t2026/08/{i % 28 + 1:02d}\t交換{commas}、確認,完了")
    path = _write(tmp_path / "data.txt", "\n".join(lines) + "\n", "utf-8")
    sniff = sniff_csv(path)
    assert sniff.delimiter == "\t" and sniff.encoding == "utf-8"
    assert sniff.confidence == 1.0


def test_sniff_shift_jis_2004_and_utf16(tmp_path):
    body = "設備名,担当\n" + "".join(f"森鷗外{i}号機,剝離{i}\n" for i in range(5))
    sj = _write(tmp_path / "sj.csv", body, "shift_jis_2004")
    assert sniff_csv(sj).encoding == "shift_jis_2004"
    rows = list(CsvSource(sj).rows())
    assert rows[1].cells[0].text == "森鷗外0号機"

    u16 = tmp_path / "u16.csv"
    u16.write_bytes(body.encode("utf-16"))
    sniff = sniff_csv(u16)
    assert sniff.encoding == "utf-16" and sniff.bom
    assert list(CsvSource(u16).rows())[2].cells[1].text == "剝離1"


def test_wrong_encoding_choice_is_a_message_not_a_500(tmp_path):
    """画面の選択肢から文字コードを選び直せるように、読めない指定は UploadError にする。

    UTF-8 のファイルに UTF-16 を選ぶと、デコーダは UnicodeDecodeError ではなく UnicodeError（BOM が無い）を出す。
    """
    body = "設備名,担当\n" + "".join(f"CMP-10{i},田中\n" for i in range(5))
    path = _write(tmp_path / "utf8.csv", body, "utf-8")
    with pytest.raises(UploadError, match="文字コード"):
        list(CsvSource(path, options={"encoding": "utf-16", "delimiter": ","}).rows())
    with pytest.raises(UploadError, match="文字コード utf-8-x は使えません"):
        list(CsvSource(path, options={"encoding": "utf-8-x", "delimiter": ","}).rows())
    with pytest.raises(UploadError, match="区切り文字は1文字"):
        CsvSource(path, options={"encoding": "utf-8", "delimiter": "ABC"})
    # 正しい指定はそのまま読める
    assert list(CsvSource(path, options={"encoding": "utf-8", "delimiter": ","}).rows())[1].cells[0].text == "CMP-100"


def test_ascii_head_then_cp932_is_detected_by_full_decode(tmp_path):
    head = "code,name,memo\n" + "".join(f"C{i:06d},item{i},ok\n" for i in range(60000))  # 1MB超のASCII
    path = tmp_path / "big.csv"
    path.write_bytes(head.encode("ascii") + "C999999,設備,異常\n".encode("cp932"))
    assert path.stat().st_size > 1024 * 1024
    assert sniff_csv(path).encoding == "cp932"


def test_mixed_encoding_warns_strict_raises_and_replace_marks_rows(tmp_path):
    good = "管理No,現象\n" + "".join(f"T-{i},停止{i}\n" for i in range(5))
    path = tmp_path / "mixed.csv"
    path.write_bytes(good.encode("cp932") + "T-9,".encode("cp932") + b"\x82\xff\n" + "T-10,復旧\n".encode("cp932"))
    sniff = sniff_csv(path)
    assert sniff.encoding == "cp932"
    assert sniff.decode_error_line == 7
    assert "読めない文字" in sniff.warnings[0] and "〓" in sniff.warnings[0]

    with pytest.raises(UploadError, match="読めない文字"):
        list(CsvSource(path, options={"encoding": "cp932", "delimiter": ","}).rows())

    src = CsvSource(path, options={"encoding": "cp932", "delimiter": ",", "errors": "replace"})
    rows = list(src.rows())
    assert "〓" in rows[6].cells[1].text
    assert src.replaced_rows == [7]
    assert rows[7].cells[1].text == "復旧"


def test_csv_nul_removal_and_excel_wrapped_values(tmp_path):
    text = '設備コード,部品番号,数量\n="00123",="0045",1\nAB\x00C,"=""0099""",2\n'
    path = _write(tmp_path / "codes.csv", text, "utf-8")
    rows = list(CsvSource(path).rows())
    assert rows[1].cells[0].text == "00123" and rows[1].cells[1].text == "0045"
    assert rows[2].cells[0].text == "ABC"
    assert rows[2].cells[1].text == "0099"
    assert value_kind(rows[1].cells[0].value, rows[1].cells[0].text) == "code"  # 先頭ゼロは数値にしない
    info = CsvSource(path).sheets()[0]
    assert (info.name, info.max_row, info.max_col) == ("codes.csv", 3, 3)


def test_csv_start_and_limit(tmp_path):
    path = _write(tmp_path / "s.csv", "a,b\n" + "".join(f"{i},{i}\n" for i in range(10)), "utf-8")
    rows = list(CsvSource(path).rows(None, start=3, limit=2))
    assert [r.index for r in rows] == [3, 4]


def test_csv_two_row_header_without_merge(tmp_path):
    lines = ["不具合,,,対策,", "台帳No,起票日,不具合内容,暫定対策,期限"]
    lines += [f"CA-{i},2026/04/0{i},停止{i},交換,2026/05/0{i}" for i in range(1, 8)]
    path = _write(tmp_path / "two.csv", "\n".join(lines) + "\n", "utf-8")
    src = CsvSource(path)
    layout = guess_layout(src, src.file_name)
    assert layout.header_rows == [1, 2]
    assert layout.headers == ["不具合_台帳No", "不具合_起票日", "不具合_不具合内容", "対策_暫定対策", "対策_期限"]


def test_open_source_rejects_unknown_extension(tmp_path):
    p = tmp_path / "a.xls"
    p.write_bytes(b"x")
    with pytest.raises(UploadError, match="対応していない形式"):
        open_source(p, "a.xls")


# ---- Excel ----

def _list_workbook(path: Path) -> tuple[Path, int]:
    wb = Workbook()
    ws = wb.active
    ws.title = "故障履歴"
    ws["A1"] = "2026年8月 故障履歴一覧"
    ws["A1"].font = Font(bold=True)
    ws.merge_cells("A1:F1")
    ws["A2"] = "作成：保全1係"
    headers = ["管理No", "発生日", "設備番号", "設備名", "現象", "停止時間(分)"]
    for c, h in enumerate(headers, start=1):
        ws.cell(4, c, h).font = Font(bold=True)
    r = 5
    for i in range(1, 13):
        ws.cell(r, 1, f"TR-2026-{i:05d}")
        ws.cell(r, 2, datetime(2026, 8, i))
        ws.cell(r, 3, f"CMP-10{i % 3}")
        ws.cell(r, 4, "CMP研磨装置")
        ws.cell(r, 5, f"スラリー流量低下 {i}")
        ws.cell(r, 6, i * 5)
        r += 1
    # 継続行（キー列が空で文章だけ）
    ws.cell(r, 5, "（追記）フィルター交換後に再発なし")
    r += 1
    # 非表示の行と取り消し線の行
    for label, hide, strike in (("hidden", True, False), ("strike", False, True)):
        ws.cell(r, 1, f"TR-2026-9{r:04d}")
        ws.cell(r, 2, datetime(2026, 8, 20))
        ws.cell(r, 3, "CVD-201")
        ws.cell(r, 5, label)
        ws.cell(r, 6, 1)
        if hide:
            ws.row_dimensions[r].hidden = True
        if strike:
            for c in range(1, 7):
                ws.cell(r, c).font = Font(strike=True)
        r += 1
    ws.cell(r, 1, "合計").font = Font(bold=True)
    ws.cell(r, 6, 999)
    total_row = r
    ws.cell(r + 2, 1, "※停止時間は生産停止から復旧確認までの時間")
    wb.save(path)
    return path, total_row


def test_excel_list_rows_classes_and_end(tmp_path):
    path, total_row = _list_workbook(tmp_path / "list.xlsx")
    src = open_source(path, "list.xlsx")
    assert src.kind == "excel"
    info = src.sheets()[0]
    assert info.name == "故障履歴" and info.max_col == 6 and not info.date1904

    first = next(src.rows("故障履歴", 1, 1))
    assert first.cells[0].bold and first.cells[0].merged_anchor == (1, 1)
    assert first.cells[3].merged_anchor == (1, 1) and first.cells[3].text == ""

    layout = guess_layout(src, "故障履歴")
    assert layout.table_kind == "list"
    assert layout.header_rows == [4] and layout.data_start == 5
    assert layout.headers == ["管理No", "発生日", "設備番号", "設備名", "現象", "停止時間(分)"]
    assert layout.data_end == total_row
    kinds = {rc.index: (rc.kind, rc.reason) for rc in layout.row_classes}
    assert kinds[1][0] == "title" and kinds[4][0] == "header"
    assert kinds[17] == ("continuation", "キー列が空で文章だけの行")
    assert kinds[18] == ("excluded", "非表示の行")
    assert kinds[19] == ("excluded", "取り消し線の行")
    assert kinds[total_row] == ("subtotal", "合計")
    assert kinds[total_row + 2][0] == "note"
    assert layout.counts == {"data": 12, "continuation": 1, "excluded": 2, "subtotal": 1}
    assert any("除外した行が2行" in w for w in layout.warnings)

    classified = list(classify_rows(src, "故障履歴", layout))
    assert classified[0][0].index == 5 and classified[-1][1].kind == "subtotal"
    assert classified[0][0].cells[1].value == datetime(2026, 8, 1)
    assert classified[0][0].cells[1].text == "2026-08-01"
    assert len(sample_data_rows(src, "故障履歴", layout, n=5)) == 5
    assert looks_like_list(src, "故障履歴")


def test_excel_anchor_and_manual_header_row(tmp_path):
    path, _ = _list_workbook(tmp_path / "list.xlsx")
    src = ExcelSource(path)
    by_anchor = guess_layout(src, "故障履歴", anchors=["発生日", "設備番号", "現象"])
    assert by_anchor.header_rows == [4]
    manual = guess_layout(src, "故障履歴", header_row=4, data_end=10)
    assert manual.data_end == 10 and manual.counts["data"] == 6


def _crosstab_workbook(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "FY2025"
    ws["A1"] = "FY2025 月別・設備別 停止時間集計表"
    ws.merge_cells("A1:N1")
    ws["A2"] = "作成日"
    ws["B2"] = "R8.4.8"
    ws["A4"] = "集計期間：2025/4/1～2026/3/31（単位：分）"
    for col, text in ((1, "ライン"), (2, "設備番号"), (3, "設備名"), (16, "合計")):
        ws.cell(6, col, text).font = Font(bold=True)
        ws.merge_cells(start_row=6, start_column=col, end_row=7, end_column=col)
    ws.cell(6, 4, "2025年")
    ws.merge_cells("D6:L6")
    ws.cell(6, 13, "2026年")
    ws.merge_cells("M6:O6")
    months = [4, 5, 6, 7, 8, 9, 10, 11, 12, 1, 2, 3]
    for i, m in enumerate(months):
        ws.cell(7, 4 + i, f"{m}月").font = Font(bold=True)
    r = 8
    for line, equipments in (("L1", ["CMP-101", "CMP-102", "膜厚計-1"]), ("L2", ["CVD-201", "CVD-202"])):
        start = r
        for eq in equipments:
            ws.cell(r, 2, eq)
            ws.cell(r, 3, f"{eq} 号機")
            for i in range(12):
                ws.cell(r, 4 + i, 100 + i if eq != "CVD-202" or i < 6 else "－")
            ws.cell(r, 16, 1234)
            r += 1
        ws.cell(start, 1, line)
        ws.merge_cells(start_row=start, start_column=1, end_row=r - 1, end_column=1)
        ws.cell(r, 1, f"{line}小計").font = Font(bold=True)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        for i in range(13):
            ws.cell(r, 4 + i, 999)
        r += 1
    ws.cell(r, 1, "総合計")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    ws.cell(r + 2, 1, "※注記")
    ws.cell(r + 3, 1, "※1 前年比は年度合計どうしの比")
    wb.save(path)
    return path


def test_excel_crosstab_two_row_merged_header(tmp_path):
    src = ExcelSource(_crosstab_workbook(tmp_path / "ct.xlsx"))
    layout = guess_layout(src, "FY2025")
    assert layout.table_kind == "crosstab"
    assert layout.header_rows == [6, 7] and layout.data_start == 8
    assert layout.headers[:5] == ["ライン", "設備番号", "設備名", "2025年_4月", "2025年_5月"]
    assert layout.headers[12:16] == ["2026年_1月", "2026年_2月", "2026年_3月", "合計"]
    assert layout.header_levels[0][4] == "2025年" and layout.header_levels[1][4] == "5月"
    assert layout.data_end == 15  # 総合計の行まで
    kinds = {rc.index: (rc.kind, rc.reason) for rc in layout.row_classes}
    assert kinds[11] == ("subtotal", "小計") and kinds[14] == ("subtotal", "小計")
    assert kinds[15] == ("subtotal", "合計")
    assert kinds[10][0] == "data"  # 「膜厚計-1」は小計ではない
    assert kinds[9][0] == "data"  # 縦結合の中（ライン列が空）でも継続行にしない
    assert layout.counts == {"data": 5, "subtotal": 3}
    assert not any("下にも表" in w for w in layout.warnings)


def test_excel_multi_block_reads_first_block_and_warns(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "2025年3月"
    ws["A1"] = "2025年3月 部品交換・保全作業記録"
    r = 3
    for block in ("■ L1ライン", "■ L2ライン"):
        ws.cell(r, 1, block)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        r += 1
        for c, h in enumerate(["日付", "設備ID", "部品名", "数量", "金額(円)"], start=1):
            ws.cell(r, c, h)
        r += 1
        for i in range(6):
            ws.cell(r, 1, datetime(2025, 3, i + 1))
            ws.cell(r, 2, "CMP-101" if i % 2 else "〃")
            ws.cell(r, 3, "POUフィルタ")
            ws.cell(r, 4, "2個")
            ws.cell(r, 5, 37000)
            r += 1
        ws.cell(r, 3, "小計（6件）")
        ws.cell(r, 5, 222000)
        r += 1
        ws.cell(r, 1, "※金額は税抜")
        r += 2
    wb.save(tmp_path / "blocks.xlsx")
    src = ExcelSource(tmp_path / "blocks.xlsx")
    layout = guess_layout(src, "2025年3月")
    assert layout.header_rows == [4] and layout.data_end == 11
    assert layout.counts == {"data": 6, "subtotal": 1}
    assert any("下にも表らしい行" in w for w in layout.warnings)
    kinds = {rc.index: rc.kind for rc in layout.row_classes}
    assert kinds[3] == "title" and kinds[12] == "note" and kinds[14] == "title"


def test_excel_left_right_tables_and_repeated_header(tmp_path):
    wb = Workbook()
    ws = wb.active
    for c, h in enumerate(["日付", "設備ID", "部品名", None, "日付", "設備ID", "部品名"], start=1):
        if h:
            ws.cell(1, c, h)
    for r in range(2, 10):
        ws.cell(r, 1, datetime(2025, 3, r))
        ws.cell(r, 2, f"CMP-10{r}")
        ws.cell(r, 3, "パッド")
        ws.cell(r, 5, datetime(2025, 3, r))
        ws.cell(r, 6, f"CVD-20{r}")
        ws.cell(r, 7, "Oリング")
    ws.cell(10, 1, "日付")
    ws.cell(10, 2, "設備ID")
    ws.cell(10, 3, "部品名")
    ws.cell(11, 1, datetime(2025, 4, 1))
    ws.cell(11, 2, "ETC-301")
    ws.cell(11, 3, "リング")
    wb.save(tmp_path / "lr.xlsx")
    layout = guess_layout(ExcelSource(tmp_path / "lr.xlsx"), "Sheet")
    assert layout.headers == ["日付", "設備ID", "部品名"]
    assert layout.data_end == 9
    assert any("右側" in w for w in layout.warnings)
    assert any("見出しが再び出た" in w for w in layout.warnings)


def test_excel_cell_limit_hidden_columns_and_date1904(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "data"
    for r in range(1, 31):
        for c in range(1, 5):
            ws.cell(r, c, "見出し" if r == 1 else r * c)
    ws.column_dimensions["B"].hidden = True
    wb.epoch = CALENDAR_MAC_1904
    wb.save(tmp_path / "w.xlsx")

    with pytest.raises(UploadError, match="セル数が上限"):
        ExcelSource(tmp_path / "w.xlsx", options={"max_cells": 50})
    src = ExcelSource(tmp_path / "w.xlsx")
    info = src.sheets()[0]
    assert info.date1904 and info.hidden_columns == [2]
    assert (info.max_row, info.max_col) == (30, 4)


def test_excel_invalid_file_is_upload_error(tmp_path):
    p = tmp_path / "broken.xlsx"
    p.write_bytes(b"not a zip")
    with pytest.raises(UploadError, match="Excelファイルとして開けません"):
        ExcelSource(p)


def test_form_like_sheet_is_not_list(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "設備修理報告書"
    rows = [("報告番号", "R2026-00123"), ("発生日", datetime(2026, 9, 14)), ("設備", "CMP-101"),
            ("故障内容", "搬送アーム停止"), ("原因", "エンコーダ不良"), ("処置", "交換"), ("作業時間", 2.5)]
    for i, (label, value) in enumerate(rows, start=3):
        ws.cell(i, 1, label)
        ws.cell(i, 3, value)
    wb.save(tmp_path / "form.xlsx")
    src = ExcelSource(tmp_path / "form.xlsx")
    layout = guess_layout(src, "Sheet")
    assert layout.table_kind in ("form_like", "unknown")
    assert not looks_like_list(src, "Sheet")


# ---- 見出し・辞書・対応づけ ----

@pytest.mark.parametrize("header, expected", [
    ("停止時間(分)", ("停止時間", "分")),
    ("単価（円）", ("単価", "円")),
    ("作業工数(H)", ("作業工数", "H")),
    ("数量 [個]", ("数量", "個")),
    ("真因(なぜなぜ要約)", ("真因(なぜなぜ要約)", "")),
    ("作業内容（詳細）", ("作業内容（詳細）", "")),
    ("管理No", ("管理No", "")),
    ("(分)", ("(分)", "")),
])
def test_split_header_unit(header, expected):
    assert split_header_unit(header) == expected


def test_month_labels():
    for text in ("4月", "2025年_4月", "2026/04", "R8.4", "2026-04-01", "１２月"):
        assert is_month_label(text), text
    for text in ("合計", "前年比", "設備名", "2026年度"):
        assert not is_month_label(text), text


def test_dictionary_lookup():
    std, how = lookup_header("対応内容")
    assert (std.key, std.role, std.type, how) == ("response_log", "log", "text", "dictionary")
    assert lookup_header("停止時間(分)")[0].key == "downtime"
    assert lookup_header("交換部品_品番")[0].key == "part_no"
    assert lookup_header("管理　NO")[0].key == "record_no"
    std, how = lookup_header("作業内容（詳細）")
    assert (std.key, how) == ("action", "similar")
    assert lookup_header("2025年_4月") is None
    for key in ("record_no", "occurred_at", "equipment_id", "equipment_name", "line", "process", "failure_category",
                "severity", "symptom", "cause", "action", "response_log", "downtime", "work_hours", "cost", "status",
                "worker", "part_name", "quantity", "unit_price"):
        assert key in BY_KEY, key
    assert BY_KEY["worker"].md == "omit" and BY_KEY["action"].log_candidate


def test_suggest_columns_types_rates_and_log():
    headers = ["管理No", "発生日", "停止時間(分)", "対応内容", "区分", "メモ欄", "担当者", "担当"]
    rows = []
    for i in range(20):
        rows.append([
            f"TR-{i:03d}",
            datetime(2026, 8, i + 1) if i < 19 else "8月末",
            str(i * 10) if i % 5 else "－",
            f"8/{i + 1} 10:00 連絡あり\n8/{i + 1} 11:00 交換完了（田中）",
            ["機械", "電気"][i % 2],
            None,
            "田中",
            "佐藤",
        ])
    result = {s.header: s for s in suggest_columns(headers, rows)}
    assert result["管理No"].key == "record_no" and result["管理No"].matched_by == "dictionary"
    assert result["発生日"].type == "datetime" and result["発生日"].type_error_rate == 0.05
    assert result["停止時間(分)"].unit == "分" and result["停止時間(分)"].display == "停止時間"
    assert result["停止時間(分)"].blank_rate == 0.2 and result["停止時間(分)"].type_error_rate == 0.0
    assert result["対応内容"].role == "log" and result["対応内容"].log and result["対応内容"].md == "body"
    assert result["区分"].key == "failure_category" and result["区分"].inferred_type == "enum"
    assert result["メモ欄"].blank_rate == 1.0 and result["メモ欄"].md == "omit"
    assert result["担当者"].key == "worker" and result["担当者"].md == "omit"
    assert result["担当"].key == "worker_2"  # 同じ標準キーは後の列に番号を付ける
    assert result["管理No"].examples == ["TR-000", "TR-001", "TR-002"]


def test_suggest_columns_with_template_spec():
    spec = {"name": "故障履歴一覧", "columns": [
        {"key": "equipment_id", "display": "設備", "headers": ["号機", "設備番号"], "type": "code", "role": "entity",
         "md": "attribute", "required": True},
    ]}
    s = suggest_columns(["号機", "謎の列"], [["CMP-1", "x"], ["CMP-2", "y"]], template_spec=spec)
    assert (s[0].key, s[0].display, s[0].matched_by) == ("equipment_id", "設備", "template")
    assert (s[1].key, s[1].matched_by, s[1].type) == (None, "none", "string")


def test_match_templates_uses_header_band_and_name():
    trouble = {"name": "トラブル一覧", "name_patterns": ["トラブル"], "columns": [
        {"key": "record_no", "headers": ["管理No"], "required": True},
        {"key": "occurred_at", "headers": ["発生日"], "required": True},
        {"key": "symptom", "headers": ["現象"], "required": True},
    ]}
    parts = {"name": "部品交換", "name_patterns": ["部品"], "columns": [
        {"key": "part_name", "headers": ["部品名"], "required": True},
        {"key": "quantity", "headers": ["数量"], "required": True},
    ]}
    downtime = {"name": "月別停止時間", "name_patterns": ["停止時間"], "crosstab": {"value_key": "downtime"}, "columns": [
        {"key": "equipment_id", "headers": ["設備番号"], "required": True},
    ]}
    headers = ["管理No", "発生日", "設備番号", "現象", "部品名"]
    ranked = match_templates(headers, "T1_トラブル対応一覧.xlsx", [parts, downtime, trouble])
    assert ranked[0] == (trouble, 3, 3)
    assert (ranked[1][0]["name"], ranked[1][1], ranked[1][2]) in {("部品交換", 1, 2), ("月別停止時間", 1, 1)}
    assert {r[0]["name"]: (r[1], r[2]) for r in ranked}["部品交換"] == (1, 2)


def test_value_kind():
    assert value_kind(None, "") == "blank"
    assert value_kind(3.5, "3.5") == "number"
    assert value_kind(datetime(2026, 1, 1), "2026-01-01") == "date"
    assert value_kind(datetime(2026, 1, 1, 9), "2026-01-01 09:00") == "datetime"
    assert value_kind(time(9, 30), "09:30") == "time"
    assert value_kind("1,234", "1,234") == "number"
    assert value_kind("△120", "△120") == "number"
    assert value_kind("20260803", "20260803") == "date"
    assert value_kind("R5.4.12", "R5.4.12") == "date"
    assert value_kind("令和5年4月12日", "令和5年4月12日") == "date"
    assert value_kind("00123", "00123") == "code"
    assert value_kind("CMP-101", "CMP-101") == "code"
    assert value_kind("#N/A", "#N/A") == "error"
    assert value_kind("設備名", "設備名") == "string"
    assert value_kind("a\nb", "a\nb") == "text"


# ---- samples/tables の実ファイル ----

def _sample(name: str) -> Path:
    path = SAMPLES / name
    if not path.exists():
        pytest.skip(f"サンプルがありません: {name}")
    return path


@pytest.mark.samples
def test_sample_t1_clean_list():
    path = _sample("T1_トラブル対応一覧_2023-2026.xlsx")
    src = open_source(path, path.name)
    names = [s.name for s in src.sheets()]
    assert names[:3] == ["トラブル一覧", "設備マスタ", "集計メモ"]
    layout = guess_layout(src, "トラブル一覧")
    assert layout.table_kind == "list"
    assert layout.header_rows == [1] and (layout.data_start, layout.data_end) == (2, 8096)
    assert layout.counts == {"data": 8095}
    assert layout.headers[:4] == ["管理No", "発生日", "発生時刻", "設備番号"]
    assert "停止時間(分)" in layout.headers
    cols = {s.header: s for s in suggest_columns(layout.headers, sample_data_rows(src, "トラブル一覧", layout))}
    assert cols["停止時間(分)"].key == "downtime" and cols["停止時間(分)"].unit == "分"
    assert cols["処置内容"].key == "action" and cols["担当者"].md == "omit"
    assert looks_like_list(src, "トラブル一覧")


@pytest.mark.samples
def test_sample_t2_system_csv():
    path = _sample("T2_設備故障履歴_システム出力.csv")
    sniff = sniff_csv(path)
    assert (sniff.encoding, sniff.delimiter, sniff.preamble_rows, sniff.header_row, sniff.trailer_rows) == (
        "cp932", ",", 4, 5, 1)
    src = open_source(path, path.name)
    layout = guess_layout(src, path.name)
    assert layout.table_kind == "list" and layout.header_rows == [5]
    assert len(layout.headers) == 45 and layout.headers[0] == "管理番号" and layout.headers[-1] == "更新者ID"
    assert layout.counts == {"data": 33095, "subtotal": 1}
    assert layout.row_classes[-1].kind == "subtotal" and layout.row_classes[-1].index == layout.data_end


@pytest.mark.samples
def test_sample_t2_code_table():
    path = _sample("T2_コード表.csv")
    src = open_source(path, path.name)
    layout = guess_layout(src, path.name)
    assert layout.header_rows == [1]
    assert layout.headers == ["コード種別", "コード", "名称", "関連コード", "補足", "使用件数"]


@pytest.mark.samples
def test_sample_t3_crosstab():
    path = _sample("T3_月別停止時間集計.xlsx")
    src = open_source(path, path.name)
    expected = {"FY2023": ([6, 7], 73, 11), "FY2024": ([6, 7], 80, 12), "FY2025": ([6, 7], 80, 12),
                "FY2026": ([6, 7], 80, 12), "全期間": ([5, 6], 79, 12)}
    for sheet, (header_rows, end, subtotals) in expected.items():
        layout = guess_layout(src, sheet)
        assert layout.table_kind == "crosstab", sheet
        assert layout.header_rows == header_rows, sheet
        assert layout.data_end == end, sheet
        assert layout.counts.get("subtotal") == subtotals, sheet
        assert layout.headers[:3] == ["ライン", "設備番号", "設備名"], sheet
        assert layout.counts["data"] in (55, 61), sheet
    fy2025 = guess_layout(src, "FY2025")
    assert fy2025.headers[4] == "2025年_4月" and fy2025.headers[13] == "2026年_1月"
    counts_sheet = guess_layout(src, "件数")
    assert counts_sheet.data_end == 80 and any("下にも表らしい行" in w for w in counts_sheet.warnings)


@pytest.mark.samples
def test_sample_t4_multi_block():
    path = _sample("T4_保全作業記録_部品交換.xlsx")
    src = open_source(path, path.name)
    assert guess_layout(src, "記入要領").table_kind in ("unknown", "form_like")
    march = guess_layout(src, "2025年3月")
    assert march.table_kind == "list" and march.header_rows == [7]
    assert march.headers[:3] == ["実施日", "装置No", "装置名"]
    assert any("右側" in w for w in march.warnings)
    assert any("下にも表らしい行" in w for w in march.warnings)
    assert any("非表示の列" in w for w in march.warnings)
    aug = guess_layout(src, "2026年8月")
    assert aug.header_rows == [5, 6]
    assert "交換部品_品番" in aug.headers
    for sheet in [s.name for s in src.sheets()][1:]:
        layout = guess_layout(src, sheet)
        assert layout.table_kind == "list", sheet
        assert layout.counts.get("data", 0) >= 30, sheet
        assert layout.counts.get("subtotal") == 1, sheet


@pytest.mark.samples
def test_sample_t5_csv_and_two_row_header_xlsx():
    csv_path = _sample("T5_是正処置管理台帳.csv")
    sniff = sniff_csv(csv_path)
    assert (sniff.encoding, sniff.bom, sniff.header_row) == ("utf-8-sig", True, 1)
    src = open_source(csv_path, csv_path.name)
    layout = guess_layout(src, csv_path.name)
    assert layout.counts == {"data": 5043} and layout.headers[0] == "台帳No"

    xlsx = _sample("T5_是正処置管理台帳_2026上期.xlsx")
    xsrc = open_source(xlsx, xlsx.name)
    xl = guess_layout(xsrc, "2026上期")
    assert xl.header_rows == [1, 2] and xl.data_start == 3
    assert xl.headers[0] == "不具合_台帳No" and xl.headers[7] == "対策_暫定対策"
    assert xl.counts == {"data": 728}
    cols = suggest_columns(xl.headers, sample_data_rows(xsrc, "2026上期", xl))
    assert cols[0].key == "record_no" and cols[0].display == "台帳No" and cols[0].matched_by == "dictionary"

# ---- Excel: XML を順に読む実装と、通常モードで開いたときの一致 ----

_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml"


def _raw_workbook(path: Path) -> Path:
    """Excel が書くような形（共有文字列・書式・結合・リンク・コメント・テーブル）と、まれな形を手で組んだブック。"""
    import zipfile

    sheet1 = f"""<worksheet xmlns="{_MAIN}" xmlns:r="{_REL}"><dimension ref="A1:K12"/>
<sheetViews><sheetView workbookViewId="0"><pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/></sheetView></sheetViews>
<cols><col min="2" max="3" width="10" hidden="1" customWidth="1"/></cols>
<sheetData>
<row r="1"><c r="A1" s="1" t="s"><v>0</v></c><c r="B1" s="1" t="s"><v>1</v></c><c r="C1" s="1" t="s"><v>3</v></c>
<c r="D1" t="inlineStr"><is><r><t>リッチ</t></r><r><rPr><b/></rPr><t xml:space="preserve"> 文字</t></r><rPh sb="0" eb="1"><t>ヨミ</t></rPh></is></c><c r="E1" t="s"><v>6</v></c></row>
<row r="2" spans="1:5"><c r="A2" t="s"><v>2</v></c><c r="B2" s="2"><v>46237</v></c><c r="C2" t="str"><f>A1</f><v></v></c><c r="D2" t="b"><v>0</v></c><c r="E2" t="e"><v>#N/A</v></c></row>
<row r="4" hidden="1"><c r="A4"><v>1.5</v></c><c s="3"><v>2</v></c><c t="inlineStr"><is><t/></is></c><c t="inlineStr"/></row>
<row r="5" hidden="0" ht="20" customHeight="1"><c r="A5" s="4"><v>46237.5</v></c><c r="F5" s="3"/></row>
<row r="7"><c r="A7" t="s"><v>3</v></c><c r="B7"><v>99</v></c></row>
<row r="9"><c r="A9" t="d"><v>2026-08-03T10:00:00</v></c><c r="B9" t="s"><v>7</v></c></row>
<row r="10" hidden="true"><c r="G10" t="s"><v>5</v></c></row>
</sheetData>
<autoFilter ref="A1:E5"/>
<mergeCells count="1"><mergeCell ref="A7:C8"/></mergeCells>
<hyperlinks><hyperlink ref="K9" r:id="rId1"/><hyperlink ref="B12" location="Sheet2!A1"/><hyperlink ref="A2" location="x!A1"/><hyperlink ref="B8" location="y!A1"/></hyperlinks>
<tableParts count="1"><tablePart r:id="rId3"/></tableParts>
</worksheet>"""
    sheet2 = f"""<x:worksheet xmlns:x="{_MAIN}"><x:sheetData>
<x:row r="3"><x:c r="A3" t="inlineStr"><x:is><x:t>三</x:t></x:is></x:c></x:row>
<x:row r="2"><x:c r="A2"><x:v>2</x:v></x:c><x:c r="B2"><x:v>5</x:v></x:c></x:row>
<x:row r="3"><x:c r="B3"><x:v>7</x:v></x:c><x:c r="A3"><x:v>8</x:v></x:c></x:row>
<x:row r="5"><x:c r="A6" t="inlineStr"><x:is><x:t>行がずれたセル</x:t></x:is></x:c></x:row>
</x:sheetData></x:worksheet>"""
    sheet3 = f"""<worksheet xmlns="{_MAIN}"><sheetData/><mergeCells><mergeCell ref="B2:C3"/></mergeCells></worksheet>"""
    strings = ["管理No", "発生日", "", "設備", "リンク", "  ", "備考\r\n2行目", "=\"001\""]
    shared = f'<sst xmlns="{_MAIN}" count="{len(strings)}" uniqueCount="{len(strings)}">' + "".join(
        f'<si><t xml:space="preserve">{s}</t></si>' for s in strings) + "</sst>"
    styles = f"""<styleSheet xmlns="{_MAIN}"><numFmts count="1"><numFmt numFmtId="164" formatCode="yyyy/mm/dd hh:mm"/></numFmts>
<fonts count="3"><font><sz val="11"/><name val="Calibri"/></font><font><b/><sz val="11"/><name val="Calibri"/></font><font><strike/><sz val="11"/><name val="Calibri"/></font></fonts>
<fills count="3"><fill><patternFill patternType="none"/></fill><fill><patternFill patternType="gray125"/></fill><fill><patternFill patternType="solid"><fgColor rgb="FFFFFF00"/></patternFill></fill></fills>
<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>
<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>
<cellXfs count="5"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/><xf numFmtId="0" fontId="1" fillId="0" borderId="0" xfId="0" applyFont="1"/><xf numFmtId="14" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/><xf numFmtId="0" fontId="2" fillId="2" borderId="0" xfId="0" applyFont="1" applyFill="1"/><xf numFmtId="164" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/></cellXfs>
</styleSheet>"""
    files = {
        "[Content_Types].xml": f"""<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/xl/workbook.xml" ContentType="{_CT}.sheet.main+xml"/>
<Override PartName="/xl/worksheets/sheet1.xml" ContentType="{_CT}.worksheet+xml"/>
<Override PartName="/xl/worksheets/sheet2.xml" ContentType="{_CT}.worksheet+xml"/>
<Override PartName="/xl/worksheets/sheet3.xml" ContentType="{_CT}.worksheet+xml"/>
<Override PartName="/xl/styles.xml" ContentType="{_CT}.styles+xml"/>
<Override PartName="/xl/sharedStrings.xml" ContentType="{_CT}.sharedStrings+xml"/>
<Override PartName="/xl/comments1.xml" ContentType="{_CT}.comments+xml"/>
<Override PartName="/xl/tables/table1.xml" ContentType="{_CT}.table+xml"/></Types>""",
        "_rels/.rels": f'<Relationships xmlns="{_PKG_REL}"><Relationship Id="rId1" Type="{_REL}/officeDocument" Target="xl/workbook.xml"/></Relationships>',
        "xl/workbook.xml": f"""<workbook xmlns="{_MAIN}" xmlns:r="{_REL}"><sheets>
<sheet name="一覧" sheetId="1" r:id="rId1"/><sheet name="並び" sheetId="2" r:id="rId2"/><sheet name="空" sheetId="3" state="hidden" r:id="rId3"/></sheets></workbook>""",
        "xl/_rels/workbook.xml.rels": f"""<Relationships xmlns="{_PKG_REL}">
<Relationship Id="rId1" Type="{_REL}/worksheet" Target="worksheets/sheet1.xml"/><Relationship Id="rId2" Type="{_REL}/worksheet" Target="worksheets/sheet2.xml"/>
<Relationship Id="rId3" Type="{_REL}/worksheet" Target="worksheets/sheet3.xml"/><Relationship Id="rId4" Type="{_REL}/styles" Target="styles.xml"/>
<Relationship Id="rId5" Type="{_REL}/sharedStrings" Target="sharedStrings.xml"/></Relationships>""",
        "xl/worksheets/sheet1.xml": sheet1,
        "xl/worksheets/sheet2.xml": sheet2,
        "xl/worksheets/sheet3.xml": sheet3,
        "xl/worksheets/_rels/sheet1.xml.rels": f"""<Relationships xmlns="{_PKG_REL}">
<Relationship Id="rId1" Type="{_REL}/hyperlink" Target="https://example.com/a" TargetMode="External"/>
<Relationship Id="rId2" Type="{_REL}/comments" Target="../comments1.xml"/>
<Relationship Id="rId3" Type="{_REL}/table" Target="../tables/table1.xml"/></Relationships>""",
        "xl/comments1.xml": f"""<comments xmlns="{_MAIN}"><authors><author>a</author></authors><commentList>
<comment ref="C3" authorId="0"><text><t>メモ</t></text></comment></commentList></comments>""",
        "xl/tables/table1.xml": f"""<table xmlns="{_MAIN}" id="1" name="表1" displayName="表1" ref="A1:E5"><autoFilter ref="A1:E5"/>
<tableColumns count="5"><tableColumn id="1" name="管理No"/><tableColumn id="2" name="発生日"/><tableColumn id="3" name="設備"/>
<tableColumn id="4" name="列4"/><tableColumn id="5" name="列5"/></tableColumns></table>""",
        "xl/styles.xml": styles,
        "xl/sharedStrings.xml": shared,
    }
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, text in files.items():
            zf.writestr(name, '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n' + text)
    return path


def _normal_mode(path: Path):
    """通常モード（load_workbook）で開いたときの SheetInfo と行（以前の ExcelSource と同じ手順）。"""
    import warnings

    from openpyxl import load_workbook
    from openpyxl.utils import column_index_from_string

    from tables.source import CellInfo, SheetInfo, SourceRow, cell_text

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(path, data_only=True)
    infos, rows = [], {}
    for ws in wb.worksheets:
        max_row = max_col = 0
        for (r, c), cell in ws._cells.items():
            if cell.value is not None and cell.value != "":
                max_row, max_col = max(max_row, r), max(max_col, c)
        anchors = {}
        for rng in ws.merged_cells.ranges:
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    anchors[(r, c)] = (rng.min_row, rng.min_col)
        hidden = set()
        for key, dim in ws.column_dimensions.items():
            if dim.hidden:
                lo = dim.min or column_index_from_string(key)
                hidden.update(range(lo, (dim.max or lo) + 1))
        infos.append(SheetInfo(ws.title, ws.sheet_state != "visible", max_row, max_col, [t.ref for t in ws.tables.values()],
                               wb.epoch == CALENDAR_MAC_1904, sorted(hidden), ws.freeze_panes, ws.auto_filter.ref or None))
        out = []
        for r in range(1, max_row + 1):
            cells = []
            for c in range(1, max_col + 1):
                xl = ws._cells.get((r, c))
                if xl is None:
                    cells.append(CellInfo(None, "", merged_anchor=anchors.get((r, c))))
                    continue
                cells.append(CellInfo(xl.value, cell_text(xl.value), xl.number_format, bool(xl.font.b), bool(xl.font.strike),
                                      xl.fill.fill_type not in (None, "none"), anchors.get((r, c))))
            dim = ws.row_dimensions.get(r)
            out.append(SourceRow(r, cells, bool(dim is not None and dim.hidden)))
        rows[ws.title] = out
    return infos, rows


def test_excel_streaming_reader_matches_normal_mode(tmp_path):
    path = _raw_workbook(tmp_path / "raw.xlsx")
    infos, expected = _normal_mode(path)
    src = ExcelSource(path)
    assert src.sheets() == infos
    for info in infos:
        assert list(src.rows(info.name)) == expected[info.name], info.name
        assert list(src.rows(info.name, 3, 4)) == expected[info.name][2:6], info.name
    # 手で組んだ形がねらいどおり読めていること
    first = {r.index: r for r in expected["一覧"]}
    assert infos[0].max_col == 11 and infos[0].max_row == 12  # リンク先の値が入った K9・B12 まで
    assert first[9].cells[10].value == "https://example.com/a" and first[12].cells[1].value == "Sheet2!A1"
    assert first[2].cells[0].value == "" and first[1].cells[3].value == "リッチ 文字"
    assert first[7].cells[1].value is None and first[7].cells[1].merged_anchor == (7, 1)  # 結合範囲の値は消える
    assert first[3].cells[2].number_format == "General" and first[3].cells[1].number_format is None  # コメントのセル
    assert first[4].hidden and not first[5].hidden and first[10].hidden
    assert infos[0].table_ranges == ["A1:E5"] and infos[0].hidden_columns == [2, 3] and infos[0].freeze_panes == "A2"
    assert [r.cells[0].value for r in expected["並び"]][:3] == [None, 2, 8] and infos[2].hidden
    # 一度数えたシートの大きさを渡すと数え直さずに同じ結果
    again = ExcelSource(path, options={"sheet_stats": src.sheet_stats()})
    assert again.sheets() == infos and list(again.rows("一覧")) == expected["一覧"]


@pytest.mark.samples
def test_sample_excel_streaming_reader_matches_normal_mode():
    for name in ("T5_是正処置管理台帳_2026上期.xlsx", "T3_月別停止時間集計.xlsx"):
        path = _sample(name)
        infos, expected = _normal_mode(path)
        src = ExcelSource(path)
        assert src.sheets() == infos
        for info in infos:
            assert list(src.rows(info.name)) == expected[info.name], (name, info.name)
