from __future__ import annotations

import io
import json
import re
import shutil
import subprocess
import tempfile
import time
import zipfile
from datetime import date
from pathlib import Path

import openpyxl
import pytest
from openpyxl.drawing.image import Image as XLImage

import core
import database as db
import llm
import tables
from aiproc import _ensure_trial_import
from core import JobError
from forms import detect_images, _is_total, Cell, suggest_rows, rows_to_pattern, load_workbook_info
from logproc import (
    PeopleIndex,
    PeopleIndex as PeopleIndex_logproc_people,
    SplitOptions,
    SplitOptions as SplitOptions_logproc_models,
    parse_log,
    parse_log as parse_log_logproc_segment,
)
from tables import CsvSource, clean_text, spec_from_dict, validate_spec
from tests.conftest import add_confirmed_document, confirmed_import
from tests.test_core import _rewrite_merges, _xlsx_with_merge



# ====================================================================================================
# 元 tests/test_round4_cross.py
# 4巡目の各担当の持ち越し（担当範囲の外の変更）の確認。
# ====================================================================================================

# ---- R4-FUZZ-3（一覧表側）: 見えない文字を消す -------------------------------------------------

def test_clean_text_removes_invisible_characters():
    assert clean_text("EQ-01⁠") == "EQ-01"
    assert clean_text("﻿管理No") == "管理No"
    assert clean_text("A​B­C‮D⁦") == "ABCD"
    assert clean_text("改行は\n残す") == "改行は\n残す"


def test_csv_cells_do_not_keep_invisible_characters(tmp_path):
    path = tmp_path / "t.csv"
    path.write_text("管理No,設備\n1,EQ-01\n2,EQ-01⁠\n﻿管理No,設備\n", encoding="utf-8")
    rows = list(CsvSource(path).rows())
    assert rows[1].text(1) == rows[2].text(1) == "EQ-01"
    assert rows[3].text(0) == rows[0].text(0) == "管理No"


# ---- S4-1: 検証より前に保存された、コンパイルできない正規表現で処理を止めない -----------------------

def test_parse_log_skips_patterns_that_do_not_compile():
    text = "4/1 10:00 停止を確認\n4/2 部品交換"
    good = parse_log(text, date(2024, 4, 1), PeopleIndex(), SplitOptions(extra_anchors=[], not_date_patterns=[]))
    bad = parse_log(text, date(2024, 4, 1), PeopleIndex(),
                    SplitOptions(extra_anchors=["[", "("], not_date_patterns=["(?P<"]))
    assert len(good.segments) == 2
    assert [s.raw for s in bad.segments] == [s.raw for s in good.segments]


# ---- F4-1 の続き: 塗りつぶしの「温度計」行で明細表の読み取りを止めない ------------------------------

def _cell(text: str, filled: bool = True) -> Cell:
    return Cell(row=1, col=1, max_row=1, max_col=1, value=text, text=text, norm=text, inline=None, filled=filled)


def test_filled_meter_names_are_not_total_rows():
    for name in ("温度計", "圧力計", "設計", "膜厚計", "pH計"):
        assert not _is_total(_cell(name)), name
    for name in ("合計", "小計", "部品費計", "工数計", "部品費合計", "計"):
        assert _is_total(_cell(name)), name
    assert _is_total(_cell("合計", filled=False))


# ---- S4-2 の続き: 複数シートが同じ drawing を指しても、シートごとに画像を数える ----------------------

def _png() -> bytes:
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (2, 2), "red").save(buf, format="PNG")
    return buf.getvalue()


def test_detect_images_with_a_drawing_shared_by_two_sheets(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "A"
    wb.create_sheet("B")
    wb["A"].add_image(XLImage(io.BytesIO(_png())), "C3")
    path = tmp_path / "shared.xlsx"
    wb.save(path)
    # シートBのリレーションをシートAと同じ drawing に向ける
    src = zipfile.ZipFile(path)
    out = tmp_path / "shared2.xlsx"
    with zipfile.ZipFile(out, "w") as dst:
        for info in src.infolist():
            dst.writestr(info, src.read(info.filename))
        rels = src.read("xl/worksheets/_rels/sheet1.xml.rels")
        dst.writestr("xl/worksheets/_rels/sheet2.xml.rels", rels)
    src.close()
    found = detect_images(out)
    assert sorted(i["sheet"] for i in found) == ["A", "B"]
    assert all(i["location"].startswith("C3") for i in found)


# ====================================================================================================
# 元 tests/test_round5_cross.py
# 5巡目の修正で、担当の範囲をまたいで直したところの確認。
#
# - 帳票・帳票の見本のアップロードは、結合セルの面積の上限を帳票用（FORM_MAX_MERGED_CELLS）で確かめる（R5-FUZZ-2）
# - ログの区切りの目印・日付ではない書き方の正規表現には、行の先頭の一部だけを渡す（SEC5-3）
# - AI接続の保存で設定ファイルが書けなくても 500 の HTML にしない（R5C-2）
# ====================================================================================================

# ---- R5-FUZZ-2 ------------------------------------------------------------------------

def _row_merged_book(tmp_path, rows: int):
    path = _xlsx_with_merge(tmp_path / "rows.xlsx", "A1:XFD1")
    path.write_bytes(_rewrite_merges(path, [f"A{r}:XFD{r}" for r in range(1, rows + 1)]))
    return path


def test_form_upload_refuses_many_whole_row_merges_quickly(app, client, tmp_path):
    path = _row_merged_book(tmp_path, 120)
    started = time.monotonic()
    res = client.post("/forms/upload", data={"file": (io.BytesIO(path.read_bytes()), "rows.xlsx")},
                      content_type="multipart/form-data")
    assert time.monotonic() - started < 10
    assert res.status_code == 400 and "結合セルの範囲が大きすぎます" in res.get_json()["error"]


def test_form_type_book_refuses_many_whole_row_merges(app, client, sample_dir, tmp_path):
    """帳票登録で置いた Excel も、帳票と同じ上限で断る（読み書きできるように開くため）。"""
    res = client.post("/form-types/new",
                      data={"name": "設備修理報告書",
                            "book": ((sample_dir / "standard.xlsx").open("rb"), "standard.xlsx")},
                      content_type="multipart/form-data")
    assert res.status_code == 200, res.get_json()
    pattern_id = res.get_json()["pattern_id"]

    path = _row_merged_book(tmp_path, 120)
    res = client.post(f"/form-types/{pattern_id}/panel",
                      data={"book": (io.BytesIO(path.read_bytes()), "rows.xlsx")},
                      content_type="multipart/form-data")
    assert res.status_code == 400
    assert "結合セルの範囲が大きすぎます" in res.get_json()["error"]
    assert core.FORM_MAX_MERGED_CELLS < core.MAX_MERGED_CELLS   # 一覧表の上限（読み取り専用で開く）は別


# ---- SEC5-3 ----------------------------------------------------------------------------

@pytest.mark.parametrize("field", ["extra_anchors", "not_date_patterns"])
def test_slow_user_pattern_on_a_long_line_finishes_quickly(field):
    text = "2024/04/01 開始\n" + "1" * 5000 + "\n2024/04/02 終了"
    started = time.monotonic()
    parsed = parse_log_logproc_segment(text, date(2024, 4, 1), PeopleIndex_logproc_people(), SplitOptions_logproc_models(**{field: [r"\d*\d*\d*x"]}))
    assert time.monotonic() - started < 5
    assert len(parsed.segments) == 2


def test_user_anchor_still_matches_at_the_line_start():
    parsed = parse_log_logproc_segment("◎ 1件目\n◎ 2件目", date(2024, 4, 1), PeopleIndex_logproc_people(), SplitOptions_logproc_models(extra_anchors=[r"◎"]))
    assert len(parsed.segments) == 2


# ---- R5C-2（AI接続の保存。設定画面は無くなり、表の取り込み画面の AI整形の段から送る） ------------------

def test_ai_settings_that_cannot_be_saved_return_a_japanese_json_error(client, monkeypatch):
    def locked(*a, **k):
        raise OSError("locked")

    monkeypatch.setattr(llm, "write_yaml", locked)
    res = client.post("/tables/ai-connection", json={"models": ["m1"], "default": "m1", "api_key": ""})
    assert res.status_code == 400
    error = res.get_json()["error"]
    assert "設定ファイルに書き込めませんでした" in error and "もう一度保存してください" in error
    assert "Traceback" not in res.get_data(as_text=True)


# ====================================================================================================
# 元 tests/test_round6_cross.py
# 6巡目の修正の、担当をまたぐ残り。
#
# - R6-2 の続き: 消し切れなかったときは「消しました」と言い切らない（一覧表）／
#   消し切れなかったことを画面に伝える（帳票。画面は fetch の答えで文面を決める）
# ====================================================================================================

def test_table_delete_says_files_were_emptied_when_not_fully_removed(app, client, monkeypatch):
    import_id = confirmed_import(app, client)
    monkeypatch.setattr(core.shutil, "rmtree", lambda *a, **k: None)
    res = client.post(f"/tables/imports/{import_id}/delete")
    assert res.status_code == 200
    message = res.get_json()["message"]
    assert "消し切れず" in message and "作った Markdown を消しました" not in message


def test_table_delete_normal_wording_is_kept(app, client):
    import_id = confirmed_import(app, client)
    message = client.post(f"/tables/imports/{import_id}/delete").get_json()["message"]
    assert "作った Markdown を消しました" in message and "消し切れず" not in message


def test_form_delete_reports_that_files_were_not_fully_removed(app, client, monkeypatch):
    """帳票の削除でも、消し切れなかったことを画面に伝える（core.purge.purge_incomplete の目印）。"""
    doc_id, _path = add_confirmed_document(app, "報告書.xlsx")
    monkeypatch.setattr(core, "remove_upload", lambda _p: False)
    res = client.post(f"/forms/{doc_id}/delete")
    assert res.status_code == 200 and res.get_json() == {"ok": True, "incomplete": True}


def test_form_delete_reports_a_clean_removal(app, client):
    doc_id, _path = add_confirmed_document(app, "報告書.xlsx")
    res = client.post(f"/forms/{doc_id}/delete")
    assert res.status_code == 200 and res.get_json() == {"ok": True, "incomplete": False}


# ---- R6C-3 の続き（core/jobs.py）: 進捗の書き込みがロック中でもジョブを失敗にしない -----------------------

def _ctx(app, monkeypatch):
    import sqlite3

    import core
    monkeypatch.setattr(core, "SOFT_UPDATE_BACKOFF", 0)
    with app.app_context():
        import database
        conn = database.connect()
        conn.execute("INSERT INTO jobs (kind, status, params_json, progress_json, created_at, updated_at) "
                     "VALUES ('test', 'running', '{}', '{}', '2026-01-01', '2026-01-01')")
        conn.commit()
        job_id = conn.execute("SELECT max(id) FROM jobs").fetchone()[0]
        conn.close()
        return core.JobContext(job_id), sqlite3


def test_progress_retries_then_writes_when_the_lock_clears(app, monkeypatch):
    ctx, sqlite3 = _ctx(app, monkeypatch)
    real, calls = ctx._update, []

    def flaky(*a):
        calls.append(1)
        if len(calls) <= 2:
            raise sqlite3.OperationalError("database is locked")
        return real(*a)
    monkeypatch.setattr(ctx, "_update", flaky)
    ctx.progress(done=5)
    assert len(calls) == 3
    assert '"done": 5' in ctx._row()["progress_json"]
    ctx.close()


def test_progress_is_skipped_not_raised_while_the_db_stays_locked(app, monkeypatch):
    ctx, sqlite3 = _ctx(app, monkeypatch)

    def locked(*a):
        raise sqlite3.OperationalError("database is locked")
    monkeypatch.setattr(ctx, "_update", locked)
    ctx.progress(done=1)
    ctx.heartbeat()
    ctx.message("途中")
    ctx.close()


def test_other_db_errors_still_raise(app, monkeypatch):
    import pytest
    ctx, sqlite3 = _ctx(app, monkeypatch)

    def broken(*a):
        raise sqlite3.OperationalError("no such column: x")
    monkeypatch.setattr(ctx, "_update", broken)
    with pytest.raises(sqlite3.OperationalError):
        ctx.progress(done=1)
    ctx.close()


# ====================================================================================================
# 元 tests/test_round6_crossfix.py
# 6巡目の「別の担当の領域だったので残した」修正の確認。
#
# - 試し実行の片付けが、まだある取り込みが払った応答まで消さないこと（core.purge と同じ持ち主の確認）
# - 列の対応づけの保存エラーを1件だけにしないこと（static/app.js の postJson → static/tables.js の表示）
# - AI整形の［見積もる］を連打できないこと（static/tables.js）
# ====================================================================================================

STATIC = Path(__file__).resolve().parents[1] / "static"


# ---- 試し実行の片付け: 持ち主のいる応答は消さない ---------------------------------------------------

def _llm_call(conn, cache_key: str, import_id: int | None = None) -> None:
    conn.execute("INSERT INTO llm_calls (cache_key, raw_text, created_at, import_id) "
                 "VALUES (?, 'あ', '2026-09-19', ?)", (cache_key, import_id))


def _llm_keys(conn) -> list[str]:
    return sorted(r[0] for r in conn.execute("SELECT cache_key FROM llm_calls"))


def test_a_trial_on_a_deleted_import_keeps_a_response_another_import_paid_for(app):
    """試し実行の間に取り込みが消えても、まだある別の取り込みが払った応答は消さない。

    試し実行は同じ文面ならほかの取り込みの応答をキャッシュとして引く。巻き添えで消すと、払った
    取り込みの再開・再実行で同じ応答をもう一度買うことになる（design.md 3.3）。
    持ち主のいない応答（この試し実行が払った分）は、これまでどおりすぐ消す。
    """
    with app.app_context():
        alive = tables.create_import("作業中.csv", "1" * 64, "tables/other.csv", {"kind": "csv"})
        conn = db.get_db()
        _llm_call(conn, "K-paid-by-alive", alive)   # まだある取り込みが払った（誰も参照していない）
        _llm_call(conn, "K-no-owner", None)         # 持ち主のいない分
        conn.commit()

        gone = alive + 1000   # もう無い取り込み（試し実行の最中に消された）
        with pytest.raises(JobError):
            _ensure_trial_import(gone, keys=["K-paid-by-alive", "K-no-owner"])

        assert _llm_keys(db.get_db()) == ["K-paid-by-alive"]




# ---- 画面側（node で確かめる） ------------------------------------------------------------------
# 画面は3つの1枚ページになり、段の中身は fetch で入れ替わる（2026-09-20 の作り直し）。
# static/tables.js は window.ragFetch（static/app.js）を通してサーバとやりとりするので、
# ragFetch は本物を切り出して使い、DOM だけを最小限の代わりで置き換える。

def _run_node(script: str):
    """node -e はコマンドラインの長さに上限があるので、ファイルに書いてから動かす。"""
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "check.mjs"
        path.write_text(script, encoding="utf-8")
        out = subprocess.run(["node", str(path)], capture_output=True, check=True, timeout=30, encoding="utf-8")
    return json.loads(out.stdout)


def _rag_fetch() -> str:
    """static/app.js の ragFetch の部分だけ取り出す（断りの errors をそのまま渡すところ）。"""
    js = (STATIC / "app.js").read_text(encoding="utf-8")
    return js[js.index("async function ragFetch("):js.index("// ---- ragSections")]


def _tables_js() -> str:
    return (STATIC / "tables.js").read_text(encoding="utf-8")


# 画面を持たない node で static/tables.js を動かすための、最小限の DOM の代わり
_DOM_STUB = """
class N {
  constructor(tag) {
    this.tagName = tag; this.children = []; this.dataset = {}; this.handlers = {};
    this._text = ""; this.disabled = false; this.className = ""; this.value = ""; this.hidden = false;
    this.sel = new Set(); this.parent = null; this.one = {}; this.many = {};
    this.classList = { add() {}, remove() {}, toggle() {}, contains: () => false };
  }
  get textContent() { return this.children.length ? this.children.map((c) => c.textContent).join("") : this._text; }
  set textContent(v) { this._text = v; this.children = []; }
  setAttribute(k, v) { this.dataset[k] = v; }
  _adopt(kids) { kids.forEach((k) => { if (k instanceof N) k.parent = this; }); }
  append(...kids) { this._adopt(kids); this.children.push(...kids); }
  replaceChildren(...kids) { this._adopt(kids); this.children = kids; this._text = ""; }
  addEventListener(type, fn) { (this.handlers[type] = this.handlers[type] || []).push(fn); }
  fire(type, event) { return (this.handlers[type] || []).map((fn) => fn(event)); }
  querySelector(s) { return this.one[s] || null; }
  querySelectorAll(s) { return this.many[s] || []; }
  closest(s) { let n = this; while (n) { if (n.sel.has(s)) return n; n = n.parent; } return null; }
}
globalThis.Node = N;
globalThis.toasts = [];
globalThis.window = {
  location: {},
  App: { toast: (m) => globalThis.toasts.push(m), bindProgressBox() {} },
  ragSections: { el: () => null, body: () => new N("div"), open() {}, close() {}, done() {}, reset() {},
                 working() {} },
};
globalThis.document = {
  createElement: (tag) => new N(tag),
  createTextNode: (t) => { const n = new N("#text"); n.textContent = t; return n; },
  querySelector: (s) => globalThis.roots[s] || null,
  querySelectorAll: (s) => [],
  addEventListener() {},
};
globalThis.roots = {};
function node(tag, selectors = [], dataset = {}) {
  const n = new N(tag);
  selectors.forEach((s) => n.sel.add(s));
  Object.assign(n.dataset, dataset);
  return n;
}
"""


def _columns_editor_script(fetch_body: str) -> str:
    """列の対応づけの段だけを置いた画面で、［この対応づけで読み込む］を1回押す。"""
    return _DOM_STUB + f"""
{fetch_body}
{_rag_fetch()}
globalThis.window.ragFetch = ragFetch;
const errors = node("ul");
const editor = node("div", ["[data-columns-editor]"],
                     {{ saveUrl: "/tables/imports/1/columns", headerRowsCount: "1" }});
editor.one = {{ "[data-editor-errors]": errors }};
editor.many = {{ "[data-setting]": [], "tr[data-col]": [] }};
const saveButton = node("button", ["[data-editor-save]"]);
saveButton.parent = editor;
const page = node("div", ["[data-tables-page]"], {{ newUrl: "/tables/new" }});
page.one = {{ "[data-columns-editor]": editor }};
globalThis.roots["[data-tables-page]"] = page;
{_tables_js()}
(async () => {{
  await Promise.all(page.fire("click", {{ target: saveButton }}));
  process.stdout.write(JSON.stringify({{
    shown: errors.children.map((c) => c.textContent),
    toast: globalThis.toasts[globalThis.toasts.length - 1],
    reenabled: saveButton.disabled === false,
  }}));
}})();
"""


def _json_response(status: int, body: str) -> str:
    return f"""
globalThis.fetch = async () => ({{
  ok: {str(status == 200).lower()}, status: {status},
  headers: {{ get: () => "application/json" }},
  json: async () => ({body}),
}});
"""


@pytest.mark.skipif(shutil.which("node") is None, reason="node がない")
def test_every_save_error_of_the_column_editor_is_shown():
    """列の対応づけの保存が複数の理由で断られたら、全部並べる（1件ずつ直させない）。"""
    body = ('{ error: "キーの列を1つ選んでください", '
            'errors: ["キーの列を1つ選んでください", "日付の列がありません", "見出しが重なっています"] }')
    got = _run_node(_columns_editor_script(_json_response(400, body)))
    assert got["shown"] == ["キーの列を1つ選んでください", "日付の列がありません", "見出しが重なっています"]
    assert got["toast"] == "3件の問題があります"
    assert got["reenabled"]   # 直して保存し直せる


@pytest.mark.skipif(shutil.which("node") is None, reason="node がない")
def test_the_column_editor_shows_a_single_error_as_before():
    """理由が1つだけ（errors を返さない断り方）なら、今までどおりその1件を出す。"""
    got = _run_node(_columns_editor_script(_json_response(409, '{ error: "ほかの処理の最中です" }')))
    assert got["shown"] == ["ほかの処理の最中です"]
    assert got["toast"] == "ほかの処理の最中です"


@pytest.mark.skipif(shutil.which("node") is None, reason="node がない")
@pytest.mark.parametrize("ok", [True, False])
def test_the_estimate_button_cannot_be_pressed_twice(ok):
    """［見積もる］は答えが返るまで押せない（連打すると同じ計算が何度も走る）。終わったら押せる。"""
    script = _DOM_STUB + f"""
let release;
const pending = new Promise((r) => {{ release = r; }});
globalThis.fetch = async () => {{
  await pending;
  return {{
    ok: {"true" if ok else "false"}, status: {"200" if ok else "400"},
    headers: {{ get: () => "application/json" }},
    json: async () => ({{ ai_rows: 3, calls: 3, cached: 0, rule_only_rows: 1, tokens_in: 10, tokens_out: 5,
                        minutes: 1, error: "見積もれません" }}),
  }};
}};
{_rag_fetch()}
globalThis.window.ragFetch = ragFetch;
const button = node("button", ["[data-estimate-run]"]);
const out = node("p");
const aiPage = node("div", ["[data-ai-page]"], {{ estimateUrl: "/tables/imports/1/ai/estimate" }});
aiPage.one = {{ "[data-estimate-run]": button, "[data-estimate]": out }};
button.parent = aiPage;
const page = node("div", ["[data-tables-page]"], {{ newUrl: "/tables/new" }});
page.one = {{ "[data-ai-page]": aiPage }};
globalThis.roots["[data-tables-page]"] = page;
{_tables_js()}
(async () => {{
  const running = Promise.all(page.fire("click", {{ target: button }}));
  await new Promise((r) => setTimeout(r, 0));
  const whileRunning = button.disabled;      // 答えを待っている間
  release();
  await running;
  process.stdout.write(JSON.stringify({{ whileRunning, after: button.disabled }}));
}})();
"""
    got = _run_node(script)
    assert got["whileRunning"] is True    # 待っている間は押せない
    assert got["after"] is False          # 成功でも失敗でも、終われば押せる


# ====================================================================================================
# 元 tests/test_three_changes_together.py
# 2026-09-21 の3つの直しが、3画面そろって噛み合っているかの通しテスト。
#
# 直した3つ:
#   1. 帳票サンプルの版フォルダ名を「版の名前＋いつから」にそろえた（scripts/samples/*）
#   2. 帳票登録は Excel を預からず、設定だけ残す（views/form_types.py・core/workbook_cache.py）
#   3. 「列の対応づけ」の言葉づかい（設備 → 対象、追記ログ → 経過の記録、「知らせ」の列をやめた）
#
# ここで見るのは、3つの担当が別々に直したあとに残りやすい「つなぎ目」:
#   - 帳票取り込みの画面が、まだ「見本の Excel から…」と案内していないか（帳票登録はもう預からない）
#   - 一覧表の検証メッセージの中で「設備」と「対象」が混ざっていないか
#   - 版フォルダを丸ごと置くと、1つの帳票の種類・1組のシートでまとまって読めるか（1ファイル＝1つの .md）
# ====================================================================================================

SAMPLES = Path(__file__).resolve().parents[1] / "samples" / "forms"

# 置いた Excel を預からなくなったので、画面のどこにも出さない言葉（帳票側）
NO_SAMPLE_WORDS = ("見本の Excel", "見本のExcel", "見本ファイル")
# 役割の名前を変えたので、画面のどこにも出さない言葉（一覧表側）
NO_ROLE_WORDS = ("追記ログ", "AI整形の対象", ">設備</option>", ">知らせ<")


# ---- つなぎ目1: 帳票取り込みの画面が「見本」と言わない -------------------------------------

def test_forms_page_does_not_promise_a_kept_sample_excel(client):
    """帳票取り込みの入口（種類がまだ無いとき）が「見本の Excel から」と案内しない。

    帳票登録は Excel を預からなくなったので、「見本」と書くと預かるように読めてしまう。
    """
    html = client.get("/forms/", follow_redirects=True).get_data(as_text=True)
    assert "帳票の種類がまだありません" in html
    for word in NO_SAMPLE_WORDS:
        assert word not in html, word


def test_type_fragment_without_active_types_does_not_say_sample(client, sample_dir):
    """帳票を置いたのに使用中の種類が無いときの案内も「見本」と言わない。"""
    from tests.test_forms import book_part

    res = client.post("/forms/upload", data={"file": book_part(sample_dir / "standard.xlsx")},
                      content_type="multipart/form-data")
    doc_id = res.get_json()["docs"][0]["id"]
    html = client.get(f"/forms/{doc_id}/type").get_json()["html"]
    assert "使用中の帳票の種類がありません" in html
    for word in NO_SAMPLE_WORDS:
        assert word not in html, word


def test_no_screen_says_sample_or_log_words(client, sample_dir):
    """3画面とも、やめた言葉を1つも出さない（帳票登録で種類を1つ作ったあとでも）。"""
    from tests.test_forms import add_field, create_type, panel_html

    pattern_id = create_type(client, sample_dir / "standard.xlsx", name="通し確認")
    add_field(client, pattern_id, "修理報告書", "A4", "B4")
    pages = [client.get(p, follow_redirects=True).get_data(as_text=True)
             for p in ("/forms/", "/tables", "/form-types/")]
    pages.append(panel_html(client, pattern_id, book=sample_dir / "standard.xlsx"))
    pages.append(panel_html(client, pattern_id))          # Excel を置いていない画面
    for html in pages:
        for word in NO_SAMPLE_WORDS + NO_ROLE_WORDS:
            assert word not in html, word


# ---- つなぎ目2: 一覧表の検証メッセージで役割の呼び方をそろえる -----------------------------

def test_spec_errors_use_the_new_role_names_only():
    """④の保存から出る検証メッセージが「設備」と「対象」を混ぜない。

    役割 entity の画面名は「対象（設備・製品・顧客など）」。まとめ方の説明だけ「設備×月」のままだと、
    同じものを2つの名前で呼ぶことになる。
    """
    bad = spec_from_dict({"name": "", "columns": [{"key": "a", "display": "あ", "type": "string"}],
                          "markdown": {"group_by": "entity_month"},
                          "log_stage": {"column": "無い列"}})
    joined = "\n".join(validate_spec(bad))
    assert "対象×月でまとめるには" in joined
    assert "役割「対象（設備・製品・顧客など）」" in joined
    assert "経過の記録の列「無い列」がありません" in joined
    assert "設備×月" not in joined
    assert "役割「entity" not in joined and "AI整形の対象列" not in joined


def test_two_entity_columns_are_refused_with_the_new_name():
    """役割「対象」を2列に付けたときの断り文も新しい名前で言う。"""
    spec = spec_from_dict({"name": "一覧", "columns": [
        {"key": "equipment_id", "display": "設備番号", "type": "code", "role": "entity"},
        {"key": "product_id", "display": "製品番号", "type": "code", "role": "entity"},
    ]})
    joined = "\n".join(validate_spec(spec))
    assert "役割「対象（設備・製品・顧客など）」の列は1つだけにしてください" in joined


# ---- つなぎ目3: 版フォルダを丸ごと置く ---------------------------------------------------

def _version_dirs() -> list[Path]:
    """samples/forms の版フォルダ全部（samples を作っていなければ空）。"""
    if not SAMPLES.is_dir():
        return []
    return sorted(v for family in sorted(SAMPLES.iterdir()) if family.is_dir()
                  for v in family.iterdir() if v.is_dir())


def _register(app, family: Path) -> int:
    """その様式の帳票の種類を作る（版が混ざるように、版ごとに1件ずつ見て候補を作る）。"""
    picks = [sorted(v.glob("*.xlsx"))[0] for v in sorted(p for p in family.iterdir() if p.is_dir())][:3]
    sheet_rows, field_rows = suggest_rows([load_workbook_info(p) for p in picks])
    with app.app_context():
        pattern_id = db.create_pattern(family.name)
        db.save_pattern(rows_to_pattern(pattern_id, {"name": family.name, "version": ""}, sheet_rows, field_rows),
                        "active")
    return pattern_id


@pytest.mark.samples
@pytest.mark.parametrize("family_name", ["F1_設備修理報告書", "F5_工程異常連絡票"])
def test_a_whole_version_folder_reads_as_one_batch(app, client, family_name):
    """版フォルダの .xlsx を丸ごと置くと、1つの帳票の種類・1組のシートで全件読めて、
    zip には1ファイルにつき1つの .md が入る（版ごとにフォルダを分けた狙い）。"""
    family = SAMPLES / family_name
    if not family.is_dir():
        pytest.skip("samples/forms がありません")
    pattern_id = _register(app, family)
    version = sorted(p for p in family.iterdir() if p.is_dir())[0]
    files = sorted(version.glob("*.xlsx"))

    parts = [(io.BytesIO(p.read_bytes()), p.name) for p in files]   # 画面でフォルダを丸ごと置いたのと同じ
    res = client.post("/forms/upload", data={"file": parts}, content_type="multipart/form-data")
    assert res.status_code == 200, res.get_data(as_text=True)
    body = res.get_json()
    ids = [d["id"] for d in body["docs"]]
    assert len(ids) == len(files), f"{version.name}: 置いた {len(files)} 件のうち {len(ids)} 件しか取り込めていない"
    batch_id = body["batch_id"]

    # 帳票の種類とシートは1回だけ選ぶ（置いた全部に同じものを使う）
    html = client.get("/forms/type", query_string={"ids": ",".join(map(str, ids))}).get_json()["html"]
    chosen = re.search(r'name="pattern_id" value="(\d+)" checked', html)
    assert chosen is not None and int(chosen.group(1)) == pattern_id, "その様式の種類が選ばれていない"
    sheets = re.findall(r'name="sheets" value="([^"]+)" checked', html)
    assert sheets, "読み取るシートにチェックが付いていない"

    form = {"ids": ",".join(map(str, ids)), "pattern_id": str(pattern_id), "sheets": sheets}
    assert client.post("/forms/read", data=form).status_code == 200
    for doc_id in ids:
        assert client.post(f"/forms/{doc_id}/confirm", json={}).status_code == 200, doc_id

    res = client.get(f"/forms/batches/{batch_id}/download.zip")
    assert res.status_code == 200
    with zipfile.ZipFile(io.BytesIO(res.get_data())) as zf:
        names = zf.namelist()
    assert len(names) == len(files) == len(set(names)), f"{version.name}: {len(files)}件 → {names}"
    assert all(n.endswith(".md") for n in names)


@pytest.mark.samples
def test_every_version_folder_holds_only_xlsx_so_it_can_be_dropped_whole():
    """どの版フォルダも中身は .xlsx だけ（_README.md などが混ざると丸ごと置けない）。"""
    dirs = _version_dirs()
    if not dirs:
        pytest.skip("samples/forms がありません")
    for version in dirs:
        kids = sorted(version.iterdir())
        assert kids and all(k.is_file() and k.suffix == ".xlsx" for k in kids), \
            f"{version.parent.name}/{version.name}: {[k.name for k in kids]}"


# ====================================================================================================
# 元 tests/test_ui_fixes6.py
# 画面まわりの不具合修正（6巡目）の確認。
#
# 画面が3つになり、取り込みの一覧は無くなった（2026-09-20 の作り直し）。
# ［削除］はその取り込みの段の中にしか無いので、「処理中は削除させない」は送り先で確かめる。
# ====================================================================================================

# ---- UX6-4: 処理中（読み込み中・確定処理中）の取り込みは削除できない --------------------------------

@pytest.mark.parametrize("status", ["reading", "confirming"])
def test_a_busy_import_cannot_be_deleted(app, client, status):
    with app.app_context():
        import_id = tables.create_import("一覧.xlsx", "hash", "tables/一覧.xlsx", {"sheet": "一覧"})
        tables.update_import(import_id, status=status)
    res = client.post(f"/tables/imports/{import_id}/delete")
    assert res.status_code == 400
    assert "処理中の取り込みは削除できません" in res.get_json()["error"]
    with app.app_context():
        assert tables.get_import(import_id) is not None   # 消えていない


@pytest.mark.parametrize("status", ["uploaded", "preview", "failed"])
def test_an_idle_import_can_be_deleted(app, client, status):
    with app.app_context():
        import_id = tables.create_import("一覧.xlsx", "hash", "tables/一覧.xlsx", {"sheet": "一覧"})
        tables.update_import(import_id, status=status)
    res = client.post(f"/tables/imports/{import_id}/delete")
    assert res.status_code == 200 and res.get_json()["ok"] is True
    with app.app_context():
        assert tables.get_import(import_id) is None


# ---- R6-C3: 段のポーリングは、ジョブが消えた（404）ら止まって画面を読み直す ----------------------------



@pytest.mark.skipif(shutil.which("node") is None, reason="node がない")
def test_polling_stops_and_reloads_when_the_job_is_gone():
    js = (Path(__file__).resolve().parents[1] / "static" / "app.js").read_text(encoding="utf-8")
    start = js.index("async function getJson(")
    end = js.index("const JOB_LABELS")
    script = """
let calls = 0, reloads = 0, status = 404;
globalThis.fetch = async () => {
  calls++;
  return { ok: status === 200, status, json: async () => ({ error: "このジョブはありません" }) };
};
globalThis.window = { location: { reload() { reloads++; } } };
""" + js[start:end] + """
(async () => {
  const run = async (code) => {
    calls = 0; reloads = 0; status = code;
    const poll = pollJob("/api/jobs/4", () => {}, { interval: 5, maxInterval: 20 });
    await new Promise((r) => setTimeout(r, 150));
    poll.stop();
    return { calls, reloads };
  };
  const gone = await run(404);
  const broken = await run(500);
  process.stdout.write(JSON.stringify({ gone, broken }));
})();
"""
    got = _run_node(script)
    # 404 は一度で止めて読み直す
    assert got["gone"] == {"calls": 1, "reloads": 1}
    # ほかの通信エラーは今までどおり間隔を広げて再試行する（読み直さない）
    assert got["broken"]["calls"] > 1 and got["broken"]["reloads"] == 0


# ====================================================================================================
# 元 tests/test_ui_fixes7.py
# 画面まわりの直し（7巡目）: ③の移動ボタン・押しても何も起きないボタン・隠れるシートタブ。
#
# 読み取り結果（③）は帳票を縦に並べるだけで、1件ずつ確定するボタンが無い。確定した件数で
# 動かしていた［次の未確定の帳票へ］は、どれも未確定のままなので必ず1件目に戻っていた。
# 静的ファイルはサーバを通さないので、node で本物の static/review.js の一部を動かして確かめる。
# ====================================================================================================

_mark_ui_fixes7 = pytest.mark.skipif(shutil.which("node") is None, reason="node が無い")


def _run_node_ui_fixes7(script: str):
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "check.mjs"
        path.write_text(script, encoding="utf-8")
        out = subprocess.run(["node", str(path)], capture_output=True, check=True, timeout=30, encoding="utf-8")
    return json.loads(out.stdout)


def _batch_bar_source() -> str:
    """static/review.js の「進み具合の1行」の部分だけ取り出す（本物のコードを動かす）。"""
    js = (STATIC / "review.js").read_text(encoding="utf-8")
    start = js.index("  // ---- まとめて置いたときの進み具合")
    return js[start:js.index("  function gotoBlock(")]


# 画面を持たない node で動かすための、最小限の DOM の代わり。
# 帳票は縦に GAP px ずつ並び、［次の帳票へ］で飛ぶとその帳票が画面のいちばん上に来る（＝実際の動き）
_STUB = """
const GAP = 900;
function textNode() { return { textContent: "" }; }
function button() { return { disabled: false, textContent: "" }; }
function reviewBody(text, next) {
  const handlers = [];
  return {
    addEventListener: (type, fn) => handlers.push(fn),
    querySelector: (s) => (s === "[data-batch-text]" ? text : s === "[data-next-doc]" ? next : null),
    click: () => handlers.forEach((fn) => fn({ target: { closest: () => ({}) } })),
  };
}
"""


def _script(count: int, at: int, presses: int, body: str) -> str:
    """count 件が並ぶ③の at 件目を見ているとき、［次の帳票へ］を presses 回押した動き。"""
    docs = [{"id": 10 + i, "state": "reviewing"} for i in range(count)]
    return _STUB + """
const docs = %s;
let scroll = %d * GAP;
const blocks = new Map(docs.map((d, i) => [d.id, {
  id: d.id, index: i, root: { getBoundingClientRect: () => ({ top: i * GAP - scroll }) },
}]));
const text = textNode();
const next = button();
const el = { reviewBody: reviewBody(text, next) };
const issueCount = () => 0;
const jumped = [];
function gotoBlock(id) { jumped.push(id); scroll = blocks.get(id).index * GAP; }
const api = new Function("docs", "blocks", "el", "issueCount", "gotoBlock",
  %s + "\\nreturn { updateBar };")(docs, blocks, el, issueCount, gotoBlock);
api.updateBar();
const first = { text: text.textContent, label: next.textContent, disabled: next.disabled };
for (let i = 0; i < %d; i += 1) el.reviewBody.click();
console.log(JSON.stringify({ first, jumped, after: text.textContent, label: next.textContent }));
""" % (json.dumps(docs), at, json.dumps(body), presses)


# ---- ③ の［次の帳票へ］は、押すたびに次へ進む ---------------------------------------------------

@_mark_ui_fixes7
def test_the_next_button_walks_through_the_forms_instead_of_returning_to_the_first():
    """3件のうち1件目を見ているとき、押すたびに 2件目 → 3件目 → 1件目 と進む。

    まとめ取り込みの③・④には1件ずつ確定するボタンが無いので、確定済みかどうかで次を探すと
    どの帳票も「未確定」のままになり、いつも1件目へ戻っていた。
    """
    result = _run_node_ui_fixes7(_script(3, at=0, presses=3, body=_batch_bar_source()))
    assert result["jumped"] == [11, 12, 10]          # 最後まで行ったら先頭に戻る
    assert result["first"]["text"] == "3件中1件目を表示中"
    assert result["first"]["disabled"] is False and result["first"]["label"] == "次の帳票へ"
    assert result["after"] == "3件中1件目を表示中"   # 帯の数字も付いていく


@_mark_ui_fixes7
def test_the_next_button_starts_from_the_form_at_the_top_of_the_screen():
    """2件目まで手でスクロールしていたら、次は3件目（1件目には戻らない）。"""
    result = _run_node_ui_fixes7(_script(3, at=1, presses=1, body=_batch_bar_source()))
    assert result["jumped"] == [12] and result["after"] == "3件中3件目を表示中"


@_mark_ui_fixes7
def test_the_bar_says_where_you_are_and_the_button_wraps_at_the_end():
    """最後の帳票にいるときは［先頭の帳票へ］になる（「すべて確定しました」で止めない）。"""
    result = _run_node_ui_fixes7(_script(3, at=2, presses=1, body=_batch_bar_source()))
    assert result["first"]["text"] == "3件中3件目を表示中"
    assert result["first"]["label"] == "先頭の帳票へ" and result["first"]["disabled"] is False
    assert result["jumped"] == [10]


@_mark_ui_fixes7
def test_the_next_button_is_off_when_only_one_form_is_shown():
    """③に1件しか並んでいなければ押せない（行き先が無い）。"""
    result = _run_node_ui_fixes7(_script(1, at=0, presses=2, body=_batch_bar_source()))
    assert result["first"]["disabled"] is True and result["jumped"] == []


# ---- ③に並んでいない帳票のボタンは、理由を出す ---------------------------------------------------

@_mark_ui_fixes7
def test_going_to_a_form_that_was_not_read_says_why():
    """読み取れなかった帳票の［読み取り結果を見る］は、黙って何も起きないのではなく理由を出す。"""
    js = (STATIC / "review.js").read_text(encoding="utf-8")
    start = js.index("  function gotoBlock(")
    body = js[start:js.index("\n  }", start) + 4]
    script = """
const blocks = new Map();
const docs = [{ id: 7, file_name: "点検記録表.xlsx" }];
const docOf = (id) => docs.find((d) => d.id === id) || null;
const toasts = [];
const toast = (m, kind) => toasts.push([m, kind]);
const sections = { open() {} };
const api = new Function("blocks", "docOf", "toast", "sections",
  %s + "\\nreturn { gotoBlock };")(blocks, docOf, toast, sections);
api.gotoBlock(7);
console.log(JSON.stringify({ toasts }));
""" % json.dumps(body)
    result = _run_node_ui_fixes7(script)
    assert len(result["toasts"]) == 1
    message, kind = result["toasts"][0]
    assert "点検記録表.xlsx" in message and "読み取れていません" in message and kind == "err"


# ---- ダウンロードのあと、渡しに行った分も片付けの対象にする ---------------------------------------

@_mark_ui_fixes7
def test_the_download_keeps_the_handed_over_ids_for_the_close_time_cleanup():
    """通信が切れても「消えました」と言い切らず、画面を閉じるときに捨てられるようにする。

    purge_after_send は本文を渡しきれなかったときは消さない（tests/test_retention.py）ので、
    渡した番号を docs から抜いたままにすると、どの片付けにも載らなくなる。
    """
    js = (STATIC / "review.js").read_text(encoding="utf-8")
    assert "let handedOver = [];" in js
    # 画面を閉じるときの送り先に、渡しに行った分も載せる
    assert "doc_ids: docs.map((d) => d.id).concat(handedOver)" in js
    # 渡しに行った分を控える（ダウンロードのボタンを押したところ）
    assert "handedOver = handedOver.concat(docs.filter(isDone).map((d) => d.id));" in js
    # 言い切らない文言にする
    assert "渡した分のデータはサーバーから消えました" not in js
    assert "ダウンロードしました。渡し終えた分はサーバーから消えます" in js


# ---- 帳票1件のとき、ファイル名の帯がシートタブを隠さない ------------------------------------------

@_mark_ui_fixes7
def test_the_sticky_sheet_pane_clears_the_file_name_bar_for_a_single_form():
    """1件だけ置いたときも「元のシート」の枠をファイル名の帯の分だけ下げる。

    まとめ置き（.review-stack.is-batch）のときだけずらしてあり、いちばん普通の使い方である
    1件のときはシートのタブの上3分の1が帯に隠れて押せなかった。
    """
    css = (STATIC / "style.css").read_text(encoding="utf-8")
    single = ".review-stack .split-pane.sticky { top: calc(var(--topbar-h) + 3rem);"
    assert single in css
    # まとめ置きのずらし（より細かい指定）はそのまま残す
    batch = ".review-stack.is-batch .split-pane.sticky { top: calc(var(--topbar-h) + 5.6rem);"
    assert batch in css
    assert css.index(single) < css.index(batch)


# ---- 登録画面: 項目を削除するとき、いま見ている Excel を送る -----------------------------------------

@_mark_ui_fixes7
def test_deleting_a_field_posts_the_book_that_is_open():
    """置いた Excel はサーバーに残らないので、削除のときも一緒に送ってシートを出したままにする。"""
    js = (STATIC / "form_types.js").read_text(encoding="utf-8")
    handler = js[js.index('const field = event.target.closest("[data-delete-field]");'):]
    handler = handler[:handler.index("async function openType")]
    assert "postForm(field.dataset.deleteField, withBook())" in handler
    # withBook は、持っているファイルを送る（無ければ、さっき読んでもらったブックの合図を送る）
    book = js[js.index("function withBook(data)"):]
    book = book[:book.index("// ---- 返ってきた断片")]
    assert 'form.append("book", bookFile, bookFile.name)' in book
    assert 'form.append("book_hash", hash)' in book


# ---- 登録画面: 失敗したあとに置き直したファイルの名前を使う ---------------------------------------

@_mark_ui_fixes7
def test_dropping_another_file_after_a_failure_replaces_the_auto_filled_name():
    """1回目に断られたあと別のファイルを置いたら、前のファイル名のまま登録しない。"""
    js = (STATIC / "form_types.js").read_text(encoding="utf-8")
    block = js[js.index('file.addEventListener("change"'):]
    block = block[:block.index("newForm.addEventListener")]
    # 自動で入れた名前かどうかを覚えておき、そのときだけ置き直したファイルの名前に入れ替える
    assert 'name.dataset.autofill' in block
    assert "if (!name.value.trim() || name.value === name.dataset.autofill) {" in block
    assert 'name.dataset.autofill = ""' in block      # ファイルを外したら覚えも消す
