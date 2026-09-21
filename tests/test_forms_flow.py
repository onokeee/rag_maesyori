"""帳票の通しテスト（帳票登録 → 使用開始 → 取り込み → 読み取り → 確定 → ダウンロード）。

画面は /form-types と /forms の1枚ずつしかなく、どの操作も fetch で JSON か HTML の断片を受け取る。
ほかの帳票テストからも使えるよう、画面を1つ進める手続きをこのファイルにまとめてある。
"""
import io
import json
from pathlib import Path

from models import database as db


# ---- 画面を1つ進める（ほかの帳票テストからも使う） -------------------------------------------------
# 置いた Excel はサーバーに残らない（2026-09-21 の利用者の指示）。ブラウザが選んだファイルを
# 持ち続けて毎回送るのと同じように、テストでも「その種類で置いたファイル」を覚えて送り直す。

def book_part(path):
    """multipart で送る Excel（画面が毎回送っているものと同じ）。"""
    path = Path(path)
    return (io.BytesIO(path.read_bytes()), path.name)


def _books(client) -> dict:
    books = getattr(client, "form_books", None)
    if books is None:
        books = {}
        client.form_books = books
    return books


def create_type(client, path, name=None) -> int:
    """帳票の Excel を1つ置いて帳票の種類を作る（名前を省くとファイル名になる）。"""
    data = {"book": book_part(path)}
    if name is not None:
        data["name"] = name
    res = client.post("/form-types/new", data=data, content_type="multipart/form-data")
    assert res.status_code == 200, res.get_data(as_text=True)
    pattern_id = res.get_json()["pattern_id"]
    _books(client)[pattern_id] = Path(path)
    return pattern_id


def add_field(client, pattern_id: int, sheet: str, label_cell: str, value_cell: str = "", book=None) -> dict:
    """見出しのセル（と値のセル）をクリックして項目を1つ作る（画面と同じく Excel も一緒に送る）。"""
    data = {"sheet": sheet, "label_cell": label_cell, "value_cell": value_cell}
    path = book if book is not None else _books(client).get(pattern_id)
    if path is not None:
        data["book"] = book_part(path)
    res = client.post(f"/form-types/{pattern_id}/fields", data=data, content_type="multipart/form-data")
    assert res.status_code == 200, res.get_data(as_text=True)
    return res.get_json()


def panel_html(client, pattern_id: int, book=None) -> str:
    """登録中の欄（HTML の断片）。Excel を渡すとシートも出る（渡さなければ設定だけの画面）。"""
    if book is None:
        return client.get(f"/form-types/{pattern_id}/panel").get_json()["html"]
    res = client.post(f"/form-types/{pattern_id}/panel", data={"book": book_part(book)},
                      content_type="multipart/form-data")
    assert res.status_code == 200, res.get_data(as_text=True)
    return res.get_json()["html"]


def activate(client, pattern_id: int) -> None:
    res = client.post(f"/form-types/{pattern_id}/status", json={"status": "active"})
    assert res.status_code == 200 and res.get_json()["status"] == "active"


def upload_forms(client, *paths) -> list[int]:
    """帳票のファイルを置く（複数ならまとまり）。戻り値: 帳票ID の並び。"""
    files = [(io.BytesIO(p.read_bytes()), p.name) for p in paths]
    res = client.post("/forms/upload", data={"file": files}, content_type="multipart/form-data")
    assert res.status_code == 200, res.get_data(as_text=True)
    return [d["id"] for d in res.get_json()["docs"]]


def read_form(client, doc_id: int, pattern_id: int, sheets: list[str], **extra):
    return client.post(f"/forms/{doc_id}/read", data={"pattern_id": pattern_id, "sheets": sheets, **extra})


def finish(client, ids, current=None):
    query = ",".join(str(i) for i in ids)
    url = f"/forms/finish?ids={query}" + (f"&current={current}" if current else "")
    return client.get(url).get_json()


# ---- 通し -------------------------------------------------------------------------------

def test_form_flow_happy_path(app, client, sample_dir):
    # 1. 帳票登録: 帳票の Excel を置き、見出しと値のセルをクリックして項目を作る（Excel は残らない）
    path = sample_dir / "standard.xlsx"
    pattern_id = create_type(client, path, "設備修理報告書")
    page = client.get("/form-types/").get_data(as_text=True)
    assert "設備修理報告書" in page and "作成中" in page
    panel = panel_html(client, pattern_id, book=path)
    # クリックで作る画面: 置いた帳票のシートが出て、型やキー名の入力欄は無い
    assert 'data-cell="B3"' in panel and "まだ項目がありません" in panel
    assert "RAGに出す" not in panel and "キー名" not in panel

    assert "「報告番号」を項目にしました" in add_field(client, pattern_id, "修理報告書", "A3", "B3")["message"]
    add_field(client, pattern_id, "修理報告書", "E4", "F4")
    add_field(client, pattern_id, "修理報告書", "A7", "A8")
    with app.app_context():
        pattern = db.load_pattern(pattern_id)
        assert [f.field_name for f in pattern.fields] == ["report_id", "equipment_name", "symptom"]
        assert pattern.status == "draft"          # 保存しても使用中にはならない

    # 読み取りテスト: いまの設定で、置いた帳票を読んだ結果が同じ欄に出る
    panel = panel_html(client, pattern_id, book=path)
    assert "R2026-00123" in panel and "CMP装置" in panel and "3項目中 <strong>3</strong>項目" in panel

    # 使用開始を押すまでは帳票取り込みの候補に出ない
    assert client.get(f"/forms/1/type").status_code in (404, 409)
    activate(client, pattern_id)
    with app.app_context():
        assert db.load_pattern(pattern_id).status == "active"

    # 2. 帳票取り込み: ファイルを置く → 帳票の種類とシート
    assert client.get("/forms/").status_code == 200
    doc_id, = upload_forms(client, path)
    body = client.get(f"/forms/{doc_id}/type").get_json()
    assert body["has_types"] and body["file_name"] == "standard.xlsx"
    assert "3項目中3項目が見つかりました" in body["html"] and "設備修理報告書" in body["html"]

    # 3. 読み取り結果
    body = read_form(client, doc_id, pattern_id, ["修理報告書"]).get_json()
    assert 'data-cell="B3"' in body["html"] and 'name="value-report_id"' in body["html"]
    assert body["summary"]["markdown"].startswith("# 設備修理報告書 R2026-00123")
    version = body["version"]

    # 途中保存（204・新しい版を返す）とプレビュー
    res = client.post(f"/forms/{doc_id}/draft",
                      json={"values": {"equipment_name": "CMP研磨装置"}, "version": version})
    assert res.status_code == 204
    version = res.headers["X-Doc-Version"]
    summary = client.post(f"/forms/{doc_id}/preview", json={"values": {}}).get_json()
    assert "CMP研磨装置" in summary["markdown"] and summary["counts"]["manual"] == 1

    # 4. 確定してダウンロード
    state = finish(client, [doc_id])
    assert state["ready"] is False and state["read_yet"] is True and "確定して Markdown を作る" in state["html"]
    assert client.post(f"/forms/{doc_id}/confirm", json={"version": version}).get_json()["ok"] is True
    with app.app_context():
        assert db.get_document(doc_id)["state"] == "confirmed"
    state = finish(client, [doc_id])
    assert state["ready"] is True and state["confirmed"] == 1
    assert "ダウンロードすると、この帳票の元のファイルと読み取り結果はサーバーから消えます" in state["html"]
    assert f"/forms/{doc_id}/download.md" in state["html"]

    # 確定後に直すと「修正中」。確定し直せば確定済みに戻る
    client.post(f"/forms/{doc_id}/draft", json={"values": {"equipment_name": "別の名前"}})
    with app.app_context():
        assert db.get_document(doc_id)["state"] == "modified"
    assert client.post(f"/forms/{doc_id}/preview", json={"values": {}}).get_json()["state"] == "modified"

    # 確定済みの読み取り直しは、手の修正が消えることの確認が要る
    res = read_form(client, doc_id, pattern_id, ["修理報告書"])
    assert res.status_code == 400 and "確認のチェック" in res.get_json()["error"]
    assert read_form(client, doc_id, pattern_id, ["修理報告書"], acknowledge="on").status_code == 200

    # 5. ダウンロードが最後の手順（ダウンロードするとサーバーからデータが消える）
    client.post(f"/forms/{doc_id}/confirm", json={})
    md = client.get(f"/forms/{doc_id}/download.md")
    assert md.status_code == 200
    assert md.headers["Content-Type"] == "text/markdown; charset=utf-8"   # charset は1つだけ
    text = md.get_data(as_text=True)
    assert text.startswith("# 設備修理報告書 R2026-00123") and "CMP装置" in text
    assert "別の名前" not in text          # 確定済みの版で作る
    with app.app_context():
        assert db.get_document(doc_id) is None
    assert client.get(f"/forms/{doc_id}/type").status_code == 404
    assert client.get(f"/forms/{doc_id}/download.md").status_code == 404

    # 帳票の種類は設定なので残る
    assert "設備修理報告書" in client.get("/form-types/").get_data(as_text=True)


def test_an_empty_value_does_not_stop_the_confirm(app, client, sample_dir):
    """空欄でも確定は止まらない（「必須」の設定は帳票登録の画面に無く、どこにも無い）。

    空欄は今までどおり「空欄」の印が付くだけで、確定の前に消さなければならない案内は出ない。
    """
    path = sample_dir / "standard.xlsx"
    pattern_id = create_type(client, path, "設備修理報告書")
    add_field(client, pattern_id, "修理報告書", "A3", "B3")
    activate(client, pattern_id)

    doc_id, = upload_forms(client, path)
    version = read_form(client, doc_id, pattern_id, ["修理報告書"]).get_json()["version"]
    res = client.post(f"/forms/{doc_id}/draft", json={"values": {"report_id": ""}, "version": version})
    version = res.headers["X-Doc-Version"]
    state = client.post(f"/forms/{doc_id}/preview", json={"version": version}).get_json()
    assert state["fields"]["report_id"]["blank"] is True   # 空欄の印は今までどおり

    html = finish(client, [doc_id])["html"]
    assert "必須" not in html and "空欄のまま確定する" not in html
    res = client.post(f"/forms/{doc_id}/confirm", json={"version": version})
    assert res.status_code == 200 and res.get_json()["ok"] is True
    with app.app_context():
        assert db.get_document(doc_id)["state"] == "confirmed"


def _parts_report(path):
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "点検報告書"
    fill = PatternFill("solid", fgColor="FFD9E1F2")
    for coord, value in {"A1": "報告番号", "B1": "IN-2026-001", "A2": "設備番号", "B2": "CMP-101",
                         "A4": "■ 交換部品", "A5": "No.", "B5": "品番", "C5": "品名", "D5": "数量",
                         "A6": 1, "B6": "PW48-1591", "C6": "ベアリング", "D6": 2,
                         "A7": 2, "B7": "PW35-1577", "C7": "スピンモータ", "D7": 1,
                         "A9": "所見", "B9": "異常なし"}.items():
        ws[coord] = value
    for coord in ("A1", "A2", "A5", "B5", "C5", "D5", "A9"):
        ws[coord].fill = fill
    wb.save(path)
    return path


def test_form_flow_with_table_field(app, client, tmp_path):
    """明細表の項目: 列見出しを1回クリック → 取り込み → 読み取り結果で行を直す → Markdown。"""
    path = _parts_report(tmp_path / "点検報告書_CMP-101.xlsx")
    pattern_id = create_type(client, path, "点検報告書")
    add_field(client, pattern_id, "点検報告書", "A1", "B1")
    add_field(client, pattern_id, "点検報告書", "A5")        # 列見出しを1回クリックするだけ
    with app.app_context():
        parts = next(f for f in db.load_pattern(pattern_id).fields if f.data_type == "table")
    assert parts.table_columns == ["No.", "品番", "品名", "数量"]
    panel = panel_html(client, pattern_id, book=path)
    assert "品番: PW35-1577／品名: スピンモータ" in panel
    activate(client, pattern_id)

    doc_id, = upload_forms(client, path)
    html = read_form(client, doc_id, pattern_id, ["点検報告書"]).get_json()["html"]
    assert "data-table-editor" in html and "スピンモータ</textarea>" in html
    assert f'name="value-{parts.field_name}"' in html

    # 画面の値（hidden の JSON）が表示中と同じなら「手で修正」にしない
    summary = client.post(f"/forms/{doc_id}/preview", json={"values": {}}).get_json()
    assert summary["counts"]["manual"] == 0
    assert "## 交換部品\n- 品番: PW48-1591／品名: ベアリング／数量: 2\n- 品番: PW35-1577／品名: スピンモータ／数量: 1\n" \
        in summary["markdown"]
    with app.app_context():
        current = json.loads(db.get_document(doc_id)["data_json"])
    value = next(f for f in current["fields"] if f["field_name"] == parts.field_name)["value"]
    assert client.post(f"/forms/{doc_id}/draft",
                       json={"values": {parts.field_name: json.dumps(value, ensure_ascii=False)}}).status_code == 204
    assert client.post(f"/forms/{doc_id}/preview", json={"values": {}}).get_json()["counts"]["manual"] == 0

    # 読み取った表の列は「No.」（行番号の列）を落とした3列。その形のまま直す
    assert value["columns"] == ["品番", "品名", "数量"]
    edited = {"columns": value["columns"], "rows": [["PW48-1591", "ベアリング", "4"]]}
    client.post(f"/forms/{doc_id}/draft",
                json={"values": {parts.field_name: json.dumps(edited, ensure_ascii=False)}})
    assert client.post(f"/forms/{doc_id}/confirm", json={}).get_json()["ok"] is True
    assert "品番: PW48-1591／品名: ベアリング／数量: 4" in finish(client, [doc_id])["html"]

    # 行を全部消しても列見出しと表の入力欄は残る（人が入れ直せる）。Markdown には行の無い明細表を出さない
    empty = {"columns": value["columns"], "rows": []}
    client.post(f"/forms/{doc_id}/draft",
                json={"values": {parts.field_name: json.dumps(empty, ensure_ascii=False)}})
    with app.app_context():
        saved = json.loads(db.get_document(doc_id)["data_json"])
    kept = next(f for f in saved["fields"] if f["field_name"] == parts.field_name)
    assert kept["value"] == {"columns": value["columns"], "rows": []}   # 列見出しは残る
    assert "## 交換部品" not in client.post(f"/forms/{doc_id}/preview", json={"values": {}}).get_json()["markdown"]

    # ダウンロードは、直したまま確定していない（修正中）ときは直した値から作る（design.md 3.3。
    # 最後の手順。ここでデータは消える）。行を全部消したので明細表は出ない
    md = client.get(f"/forms/{doc_id}/download.md").get_data(as_text=True)
    assert "## 交換部品" not in md and "スピンモータ" not in md
    assert "- 報告番号: IN-2026-001" in md
