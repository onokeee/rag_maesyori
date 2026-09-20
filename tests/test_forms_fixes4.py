"""帳票まわりの修正（4巡目の確認で見つかった不具合）の回帰テスト。"""
import io
import json
import zipfile
from types import SimpleNamespace

from excel.extractor import apply_manual_values
from models import database as db
from tests.test_forms_fixes import _confirmed_doc, _extraction, _state


# ---- F4-1: 「〜計」で終わる品名（温度計・圧力計）を合計行にしない ------------------------------------

def test_instrument_names_ending_in_kei_are_not_total_rows():
    from export.formats import table_markdown_lines

    value = {"columns": ["品名", "型式", "数量"],
             "rows": [["ベアリング", "6205ZZ", "2"], ["圧力計", "GV-50", "1"], ["温度計", "", ""],
                      ["設計", "", ""], ["膜厚計", "", ""], ["部品費計", "", "12000"], ["合計", "", "3"]]}
    lines = table_markdown_lines(value)
    assert "- 品名: 圧力計／型式: GV-50／数量: 1" in lines
    assert "- 品名: 温度計" in lines and "- 品名: 設計" in lines and "- 品名: 膜厚計" in lines
    assert "- 部品費計: 数量: 12000" in lines and "- 合計: 数量: 3" in lines
    # 数字の無い合計行は今までどおり出さない
    assert table_markdown_lines({"columns": ["品名"], "rows": [["ノギス"], ["小計"]]}) == ["- 品名: ノギス"]


# ---- F4-3: AIが入れた値を手で消すと、確定済みの状態に戻る -----------------------------------------------

def test_clearing_an_ai_filled_value_returns_the_field_to_the_confirmed_one():
    from views.forms import _apply_values

    confirmed = json.loads(_extraction())
    confirmed["fields"][0].update(value=None, sheet=None, value_cell=None, warning="ラベルはありますが値が空です")
    working = json.loads(json.dumps(confirmed))
    working["fields"][0].update(value="AIの値", sheet="報告書", value_cell="Z99", ai_filled=True,
                                warning="AI（m）が入力しました。元のファイルと照合してください。")
    assert _apply_values(working, {"equipment_name": ""}, confirmed)
    assert working["fields"][0] == confirmed["fields"][0]


# ---- F4-4 / BR4-2: まとめ取り込みの「次の帳票へ」と、途中保存後の zip の確認文 ---------------------------

def _reviewing_doc(app, name, order):
    doc_id = _confirmed_doc(app, name, batch_id="B", order=order)
    with app.app_context():
        db.update_document(doc_id, confirmed_json=None)  # 確認中（手の修正があるかもしれない）
    return doc_id


def test_the_next_form_of_a_batch_is_the_one_that_is_not_confirmed_yet(app, client):
    first = _confirmed_doc(app, "1.xlsx", batch_id="B", order=0)
    second = _reviewing_doc(app, "2.xlsx", 1)
    assert _state(app, second) == "reviewing"
    body = client.get(f"/forms/finish?ids={first},{second}&current={first}").get_json()
    assert body["next_id"] == second and body["confirmed"] == 1
    assert f'data-open-doc="{second}">次の帳票へ（残り1件）' in body["html"]

    # まだ確定していない帳票は zip に入らない（確定済みの分だけを渡す）
    res = client.get("/forms/batches/B/download.zip")
    assert res.status_code == 200 and res.mimetype == "application/zip"
    with zipfile.ZipFile(io.BytesIO(res.data)) as zf:
        assert len(zf.namelist()) == 1
    with app.app_context():
        assert db.get_document(first) is None and db.get_document(second) is not None
    assert client.get("/forms/batches/B/download.zip").status_code == 404   # 確定済みが無ければ渡さない


# ---- R4-MD-2: 明細表の「〃」は上の行の値にする ---------------------------------------------------------

def _cells(*texts):
    return [SimpleNamespace(text=t) for t in texts]


def test_ditto_marks_in_a_detail_table_take_the_value_above():
    from excel.tables import Table
    from export.formats import table_markdown_lines

    table = Table(anchor=None, header=_cells("点検箇所", "点検内容", "判定"),
                  rows=[_cells("〃", "外観", "○"), _cells("ステージ", "位置偏差", "○"), _cells("〃", "振動", "×"),
                        _cells("同上", "異音", "○"), _cells("", "清掃", "○"), _cells("〃", "給油", "○")])
    value = table.to_value()
    assert [r[0] for r in value["rows"]] == ["〃", "ステージ", "ステージ", "ステージ", "", "〃"]
    assert "- 点検箇所: ステージ／点検内容: 振動／判定: ×" in table_markdown_lines(value)
    # 「〃」だけの行（明細なしの記入）は今までどおり落とす
    only = Table(anchor=None, header=_cells("部位", "内容"), rows=[_cells("軸", "給油"), _cells("〃", "")])
    assert only.to_value()["rows"] == [["軸", "給油"]]


# ---- ux4-1: 同じ値に読める手の入力でも、警告の出る入力は手で修正にする ----------------------------------

def test_same_value_input_with_a_warning_is_kept_as_an_edit():
    ex = {"fields": [
        {"field_name": "w", "display_name": "作業時間", "data_type": "number", "required": False, "value": 1.5,
         "unit": "時間", "spec_unit": "時間", "warning": None, "edited": False, "ai_filled": False},
        {"field_name": "d", "display_name": "発生日", "data_type": "date", "required": False,
         "value": "2026-09-14", "warning": None, "edited": False, "ai_filled": False}]}
    apply_manual_values(ex, {"value-w": "1.5～3時間", "value-d": "2026-09-14 24:30"})
    w, d = ex["fields"]
    assert w["edited"] and w["warning"]
    assert d["edited"] and d["warning"]
    # 警告の消える入力（同じ値を書き直した）も手で修正にして、警告を消す
    apply_manual_values(ex, {"value-d": "2026/9/14"})
    assert d["value"] == "2026-09-14" and not d["warning"] and d["edited"]


def test_same_value_without_warning_is_still_not_an_edit():
    ex = {"fields": [{"field_name": "w", "display_name": "作業時間", "data_type": "number", "required": False,
                      "value": 1.5, "unit": "時間", "spec_unit": "時間", "warning": None, "edited": False,
                      "ai_filled": False}]}
    apply_manual_values(ex, {"value-w": "1.50"})
    assert not ex["fields"][0]["edited"]


# ---- ux4-4: 読み取りテストの単位不明の警告は、種類での直し方にする ------------------------------------

def test_form_type_panel_explains_how_to_set_the_unit(app, client, tmp_path):
    from openpyxl import Workbook

    from tests.test_forms_flow import add_field, create_type

    wb = Workbook()
    ws = wb.active
    ws.title = "報告書"
    ws["A1"], ws["B1"], ws["A2"], ws["B2"] = "報告番号", "R-001", "作業時間", 2.5
    path = tmp_path / "unit.xlsx"
    wb.save(path)
    pattern_id = create_type(client, path, "単位なし")
    add_field(client, pattern_id, "報告書", "A1", "B1")
    add_field(client, pattern_id, "報告書", "A2", "B2")
    with app.app_context():
        assert next(f for f in db.load_pattern(pattern_id).fields if f.display_name == "作業時間").data_type == "number"
    panel = client.get(f"/form-types/{pattern_id}/panel").get_json()["html"]
    assert "帳票に単位が書かれていません" in panel and "値に単位（分・時間など）を付けて入力できます" in panel


# ---- R4-FUZZ-3: 見えない文字（ゼロ幅スペース・BOM）を消す -------------------------------------------------

def test_invisible_characters_are_removed_from_labels_and_values(tmp_path):
    from excel.text import cell_text, nfkc_value, normalize_label
    from tests.test_extraction import _book, _fields, _values

    assert normalize_label("設備名​") == normalize_label("設備名")
    assert cell_text("﻿EQ-01⁠") == "EQ-01"
    assert nfkc_value("CMP​-108") == "CMP-108"
    info = _book(tmp_path, "zw.xlsx", {"報告書": ({"A1": "設備名​", "B1": "CMP⁠研磨装置"}, {}, [])})
    values, _ = _values(info, _fields(("equipment_name", "string", "設備名")))
    assert values["equipment_name"] == "CMP研磨装置"


# ---- R4-FUZZ-4: 計算結果の無い数式を「値が空」と言わない --------------------------------------------------

def test_formula_without_a_cached_value_gets_its_own_warning(tmp_path):
    from openpyxl import Workbook

    from excel.extractor import UNCACHED_FORMULA_WARNING
    from excel.workbook import load_workbook_info
    from tests.test_extraction import _fields, _values

    wb = Workbook()
    ws = wb.active
    ws.title = "報告書"
    ws["A1"], ws["B1"] = "設備名", '="CMP"&"装置"'  # openpyxl は計算結果を保存しない
    ws["A2"], ws["B2"] = "報告番号", None
    path = tmp_path / "formula.xlsx"
    wb.save(path)
    info = load_workbook_info(path)
    assert info.uncached_formulas == {"報告書": {(1, 2)}}
    values, fields = _values(info, _fields(("equipment_name", "string", "設備名"), ("report_id", "string", "報告番号")))
    assert values["equipment_name"] is None
    assert fields["equipment_name"]["warning"] == UNCACHED_FORMULA_WARNING
    assert fields["report_id"]["warning"] == "ラベルはありますが値が空です"
