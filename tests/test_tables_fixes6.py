"""一覧表の不具合修正（6巡目）の確認。"""
from __future__ import annotations

import csv
import io
import json
import re
import shutil
import subprocess
import time
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook

from tables import pipeline, store
from tables.spec import spec_from_dict, validate_spec
from tests.test_tables_pipe import FakeCtx, _new_import, _rec


# ---- R6T-1: 設定の編集を何も変えずに保存しても、JSON で選んだ期間の列（2列目の日付）を置き換えない -------------

def test_noop_save_keeps_the_period_date_column(app, client):
    from tables.spec import ColumnSpec
    from tests.tables_helpers import SPEC_CSV, json_only_spec
    from tests.test_tables_fixes4 import _editor_body, _saved

    spec = json_only_spec()
    spec.columns.append(ColumnSpec("completed_at", "完了日", headers=["完了日"], type="date", role="date"))
    spec.period["date_column"] = "completed_at"
    assert validate_spec(spec) == []
    # 完了日の列もある CSV（画面にその行が出るように）
    text = "\r\n".join(line + ("完了日" if i == 0 else "2026-08-28") for i, line in
                       enumerate(x + "," for x in SPEC_CSV.strip("\r\n").split("\r\n"))) + "\r\n"
    import_id, template_id, body = _editor_body(app, client, spec, text=text)
    after = _saved(app, client, import_id, template_id, body)
    assert after.period["date_column"] == "completed_at"

    # 日付の役割の列を変えたときは画面の役割から決め直す
    for r in body["columns"]:
        if r.get("key") == "completed_at":
            r["role"] = "attribute"
    after = _saved(app, client, import_id, template_id, body)
    assert after.period["date_column"] == "occurred_at"


# ---- R6T-2: 80行目より下の見出し行を手で指定しても見出しが読める ------------------------------------------------

def _deep_book(path: Path, header_row: int, sheet: str = "一覧") -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.cell(1, 1, "トラブル対応一覧（前置き）")
    ws.cell(header_row, 1, "管理No")
    ws.cell(header_row, 2, "発生日")
    ws.cell(header_row, 3, "現象")
    for i in range(1, 40):
        r = header_row + i
        ws.cell(r, 1, f"TR-{i:03d}")
        ws.cell(r, 2, f"2026/08/{(i % 28) + 1:02d}")
        ws.cell(r, 3, f"搬送ロボット{i}号機が停止")
    wb.save(path)
    return path


def test_header_row_below_the_head_rows_is_read(tmp_path):
    from tables.detect import guess_layout
    from tables.excel_source import ExcelSource

    src = ExcelSource(_deep_book(tmp_path / "deep90.xlsx", 91))
    layout = guess_layout(src, "一覧", header_rows=[91])
    assert layout.header_rows == [91] and layout.data_start == 92 and layout.data_end == 130
    assert layout.headers[:3] == ["管理No", "発生日", "現象"]
    single = guess_layout(src, "一覧", header_row=91)
    assert single.headers[:3] == ["管理No", "発生日", "現象"]


# ---- R6T-3: 見出しが自動で見つからなくても、保存した取り込み設定を選べる ------------------------------------------

def test_the_source_panel_lists_saved_settings_and_checks_the_imports_own(app, client, tmp_path):
    from tests.tables_helpers import panel_html, upload_bytes

    path = _deep_book(tmp_path / "deep40.xlsx", 41)
    import_id = upload_bytes(client, path.read_bytes(), "deep40.xlsx")
    spec_a = spec_from_dict({"name": "トラブル一覧", "columns": [
        {"key": "record_no", "display": "管理No", "headers": ["管理No"], "type": "code", "role": "key"}]})
    spec_b = spec_from_dict({"name": "別名", "columns": [
        {"key": "no", "display": "番号", "headers": ["番号"], "type": "code", "role": "key"}]})
    with app.app_context():
        store.create_template("トラブル一覧", spec_a, "")
        t2, v2 = store.create_template("別名", spec_b, "")
        imp = store.get_import(import_id)
        src = dict(imp["source"] or {})
        src["header_rows"] = [41]
        store.update_import(import_id, source=src, template_id=t2, template_version_id=v2)
    page = panel_html(client, import_id, "source")
    radios = re.findall(r'<input type="radio" name="template" value="([^"]+)"([^>]*)>', page)
    values = [v for v, _rest in radios]
    assert str(t2) in values and "new" in values and len(values) == 3
    assert [v for v, rest in radios if "checked" in rest] == [str(t2)]
    assert "管理No" in page   # 保存した見出し行（41行目）で見出しを読んでいる


# ---- R6T-4: 見出し行の入力は画面の下見とサーバーで同じように読む ---------------------------------------------------

_ROW_INPUTS = ["６,７", "3-4", "1，2", "1、2 3", "3a", "5-3", "1-20", "", "0,2", "１２０", "+5", "2-2"]


def test_row_inputs_are_read_the_same_way_on_the_server():
    from views.tables import _int_list, _row_no

    assert _int_list("６,７") == [6, 7]
    assert _int_list("3-4") == [3, 4]
    assert _int_list("1，2") == [1, 2]
    assert _int_list("3a") == [] and _int_list("5-3") == [] and _int_list("1-20") == []
    assert _int_list([3, 1]) == [1, 3]
    assert _row_no("１２０") == 120 and _row_no("12x") is None and _row_no("") is None and _row_no(None) is None


@pytest.mark.skipif(shutil.which("node") is None, reason="node がない")
def test_row_inputs_are_read_the_same_way_in_the_browser():
    from views.tables import _int_list, _row_no

    js = (Path(__file__).resolve().parents[1] / "static" / "tables.js").read_text(encoding="utf-8")
    start = js.index("  const parseEnd = ")
    end = js.index("  function applyLayout(")
    script = js[start:end] + f"""
const inputs = {json.dumps(_ROW_INPUTS, ensure_ascii=False)};
process.stdout.write(JSON.stringify({{rows: inputs.map(parseRows), ends: inputs.map(parseEnd)}}));
"""
    out = subprocess.run(["node", "-e", script], capture_output=True, check=True, timeout=30, encoding="utf-8")
    got = json.loads(out.stdout)
    assert got["rows"] == [_int_list(x) for x in _ROW_INPUTS]
    assert got["ends"] == [_row_no(x) for x in _ROW_INPUTS]


# ---- R6T-5: 既定の候補で AI整形の対象（追記ログ）は1列だけ ---------------------------------------------------------

def test_default_suggestions_tick_only_one_log_column():
    from tables.mapping import suggest_columns

    log = "4/1 10:00 停止を確認。\n4/2 11:00 センサーを交換。\n4/3 復旧を確認した。"
    headers = ["管理No", "対応内容", "対応内容_2", "対応内容_3"]
    rows = [[f"TR-{i}", log, log + "追記", log + "再発"] for i in range(10)]
    sugg = suggest_columns(headers, rows)
    logs = [s for s in sugg if s.role == "log"]
    assert [s.header for s in logs] == ["対応内容"]
    others = [s for s in sugg if s.header in ("対応内容_2", "対応内容_3")]
    assert all(s.role == "text" and not s.log and s.md in ("body", "omit") for s in others)


def test_untouched_columns_panel_with_two_log_like_columns_can_be_saved(app, client):
    from tests.tables_helpers import editor_body, name_source, save_columns, save_layout, upload_csv

    log = "4/1 10:00 停止を確認。\n4/2 11:00 センサーを交換。\n4/3 復旧を確認した。"
    text = "管理No,発生日,対応内容,対応内容_2\r\n" + "".join(
        f"TR-{i},2026/08/{i + 1:02d},\"{log}\",\"{log}追記\"\r\n" for i in range(10))
    import_id = upload_csv(client, "log2.csv", text, encoding="utf-8")
    name_source(client, import_id, "T", encoding="utf-8")
    assert save_layout(client, import_id).status_code == 200
    body = editor_body(client, import_id)
    assert sum(1 for r in body["columns"] if r.get("ai")) == 1
    res = save_columns(client, import_id, body)
    assert res.status_code == 200, res.get_json()


# ---- R6T-6: JSON の AI整形の上限・実行条件が数値でなければ保存の前に断る -------------------------------------------

def _log_spec(**stage):
    return spec_from_dict({"name": "x", "columns": [{"key": "log", "display": "ログ", "type": "text", "role": "log"}],
                           "log_stage": {"column": "log", **stage}})


@pytest.mark.parametrize("stage", [
    {"limits": {"max_segments": "多い"}},
    {"limits": {"max_input_tokens": 0}},
    {"limits": {"max_segments": True}},
    {"run_if": {"min_chars": "x"}},
    {"run_if": {"any": [{"min_segments": 2}, {"min_chars": "x"}]}},
    {"run_if": {"all": "min_chars"}},
])
def test_non_numeric_limits_and_run_if_are_refused(stage):
    assert validate_spec(_log_spec(**stage))


@pytest.mark.parametrize("stage", [
    {"limits": {"max_segments": 30, "max_input_tokens": "4000"}},
    {"run_if": {"any": [{"min_segments": 2}, {"min_chars": 80}], "contains": "停止"}},
    {"run_if": {}},
])
def test_numeric_limits_and_run_if_are_accepted(stage):
    assert validate_spec(_log_spec(**stage)) == []


# ---- R6-MD-1: 番号だけの行が先にあっても、集計の表示名は同じ番号の行の名前から付ける ---------------------------

def _code_entity_spec():
    return spec_from_dict({"name": "T", "columns": [
        {"key": "record_no", "display": "管理No", "type": "code", "role": "key"},
        {"key": "occurred_at", "display": "発生日", "type": "date", "role": "date"},
        {"key": "equipment", "display": "設備", "type": "code", "role": "entity"},
    ], "period": {"date_column": "occurred_at"}})


def test_summary_display_name_comes_from_the_group_not_the_first_row():
    from tables.spec import SummarySpec
    from tables.summaries import entity_fiscal_year_summaries, month_summaries

    spec = _code_entity_spec()
    records = [_rec("A-1", record_no="A-1", occurred_at="2026-08-03", equipment="CLN-502"),
               _rec("A-2", record_no="A-2", occurred_at="2026-08-04", equipment="枚葉洗浄 2号機（CLN-502）"),
               _rec("A-3", record_no="A-3", occurred_at="2026-09-04", equipment="CLN-502")]
    fy = entity_fiscal_year_summaries(records, spec, None, SummarySpec("entity_fiscal_year"))
    assert [(s["entity"], s["display"]) for s in fy] == [("CLN-502", "枚葉洗浄 2号機（CLN-502）")]
    months = month_summaries(records, spec, None, SummarySpec("month"))
    assert [m["top"][0]["display"] for m in months] == ["枚葉洗浄 2号機（CLN-502）"] * 2


# ---- R6-MD-2: 整数だけの列の平均を整数に丸めない ----------------------------------------------------------------

def test_average_of_integer_values_keeps_one_decimal():
    from tables.summaries import fmt_average

    assert fmt_average(9, 18, True) == 0.5
    assert fmt_average(57, 39, True) == 1.5
    assert fmt_average(150, 2, True) == 75 and isinstance(fmt_average(150, 2, True), int)
    assert fmt_average(0, 3, True) == 0


# ---- R6-MD-4: 「0:28以降、…」の時刻は見出しに残す ------------------------------------------------------------------

@pytest.mark.parametrize("text, expected", [
    ("0:28以降、装置がオフライン表示。", "0:28以降、装置がオフライン表示。"),
    ("12:30から停止", "12:30から停止"),
    ("12:30以降、装置がオフライン表示", "12:30以降、装置がオフライン表示"),
    ("R05.04.01 11:45(休日)、搬送停止", "搬送停止"),
    ("0:28、搬送停止", "搬送停止"),
])
def test_title_keeps_a_time_that_is_part_of_the_sentence(text, expected):
    from tables.markdown import _title_text_line

    assert _title_text_line(text) == expected


# ---- R6-1: 渡し終えた取り込みに、読み込み直しが行データ・問題一覧を書き戻さない ------------------------------------

def test_reread_does_not_write_back_into_a_purged_import(app, monkeypatch):
    from core import purge

    with app.app_context():
        _template_id, import_id = _new_import(app)
        real_checks = pipeline.run_checks

        def purge_midway(records, spec, stats):
            purge.purge_table_import(import_id)   # 読み込みの途中で渡し終えて消えた
            return real_checks(records, spec, stats)

        monkeypatch.setattr(pipeline, "run_checks", purge_midway)
        with pytest.raises(pipeline.PipelineError):
            pipeline.run_read(FakeCtx(), import_id)
        assert store.get_import(import_id) is None
        assert not pipeline.import_dir(import_id).exists()


def test_write_helpers_do_not_create_a_missing_folder(tmp_path):
    with pytest.raises(OSError):
        pipeline.write_atomic(tmp_path / "gone" / "issues.csv", b"x")
    with pytest.raises(OSError):
        pipeline.write_rows(tmp_path / "gone" / "rows.jsonl.gz", [])
    assert not (tmp_path / "gone").exists()


# ---- R6-SEC-1: 全角の ＝ ＋ － ＠ で始まるシート名・ファイル名も式にしない --------------------------------------------

def test_full_width_formula_prefixes_are_guarded():
    from tables.outputs import guard_formula, import_report_items, normalized_csv, report_csv

    for s in ("＝SUM(1,1)", "＋1", "－1", "＠SUM(1,1)"):
        assert guard_formula(s) == "'" + s
    assert guard_formula("設備") == "設備"
    spec = _code_entity_spec()
    rec = {"key": "A-1", "values": {"record_no": "A-1"}, "source": {"file": "＋1.xlsx", "sheet": "＝SUM(1,1)", "row": 2}}
    rows = list(csv.reader(io.StringIO(normalized_csv(spec, [rec]).decode("utf-8-sig"))))
    assert rows[1][-3:] == ["'＋1.xlsx", "'＝SUM(1,1)", "2"]
    imp = {"id": 1, "file_name": "＋1.xlsx", "stats": {"sheet": "＝SUM(1,1)"}}
    report = dict(csv.reader(io.StringIO(report_csv(import_report_items(imp, spec, 1)).decode("utf-8-sig"))))
    assert report["ファイル名"] == "'＋1.xlsx" and report["シート"] == "'＝SUM(1,1)"




# ---- FUZZ: 手で作った Excel（大きな非表示列・上限を超える行・列が多すぎる表） ---------------------------------------

def _rewrite_sheet(path: Path, fn) -> None:
    buf = io.BytesIO()
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                data = fn(data)
            dst.writestr(info, data)
    path.write_bytes(buf.getvalue())


def _small_book(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "一覧"
    ws.append(["管理No", "発生日", "現象"])
    for i in range(5):
        ws.append([f"TR-{i}", "2026/08/01", "停止"])
    wb.save(path)
    return path


def test_huge_hidden_column_span_is_clamped(tmp_path):
    from tables.excel_source import ExcelSource

    path = _small_book(tmp_path / "cols.xlsx")
    _rewrite_sheet(path, lambda d: d.replace(
        b"<sheetData>", b'<cols><col min="1" max="20000000" hidden="1" width="5"/></cols><sheetData>', 1))
    t = time.monotonic()
    info = ExcelSource(path).sheets()[0]
    assert time.monotonic() - t < 5
    assert len(info.hidden_columns) <= 16384


def test_cell_beyond_the_excel_row_limit_is_refused_quickly(tmp_path):
    from tables.excel_source import ExcelSource
    from tables.source import UploadError

    path = _small_book(tmp_path / "far.xlsx")
    _rewrite_sheet(path, lambda d: d.replace(
        b"</sheetData>",
        b'<row r="4294967296"><c r="D4294967296" t="inlineStr"><is><t>z</t></is></c></row></sheetData>', 1))
    t = time.monotonic()
    with pytest.raises(UploadError, match="Excel の上限"):
        ExcelSource(path)
    assert time.monotonic() - t < 5


def test_xlsx_with_too_many_columns_is_refused(tmp_path):
    from tables.csv_source import MAX_COLUMNS
    from tables.excel_source import ExcelSource
    from tables.source import UploadError

    wb = Workbook()
    ws = wb.active
    ws.title = "一覧"
    ws.append([f"項目{i}" for i in range(MAX_COLUMNS + 1)])
    ws.append([f"v{i}" for i in range(MAX_COLUMNS + 1)])
    path = tmp_path / "wide.xlsx"
    wb.save(path)
    src = ExcelSource(path)
    with pytest.raises(UploadError, match="列数が上限"):
        list(src.rows("一覧", 1, 2))


def test_far_memo_cell_is_not_counted_as_many_columns(tmp_path):
    from tables.excel_source import ExcelSource

    path = _small_book(tmp_path / "memo.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "一覧"
    ws.append(["管理No", "発生日", "現象"])
    ws.append(["TR-1", "2026/08/01", "停止"])
    ws["XFD1"] = "メモ"
    wb.save(path)
    rows = list(ExcelSource(path).rows("一覧", 1, 2))
    assert len(rows) == 2


# ---- R6C-1: 確定した取り込みでは試し実行を断る（zip・保存と下書きが食い違わないように） ---------------------------------

def test_trial_is_refused_after_confirm(app, client):
    with app.app_context():
        _template_id, import_id = _new_import(app)
        pipeline.run_read(FakeCtx(), import_id)
        pipeline.run_render(FakeCtx(), import_id)
        assert store.get_import(import_id)["status"] == "confirmed"
    res = client.post(f"/tables/imports/{import_id}/ai/trial", json={"row_key": "TR-001"})
    assert res.status_code == 400
    assert "確定後は試し実行できません" in res.get_json()["error"]


# ---- R6C-2: 渡している間に別のタブで読み込み直しが始まったら、「ダウンロード済み」とは言わない ----------------------------

@pytest.mark.parametrize("method, url, builder", [
    ("get", "download.zip", "build_download"),
])
def test_reread_during_hand_out_goes_back_to_the_preview(app, client, monkeypatch, tmp_path, method, url, builder):
    with app.app_context():
        _template_id, import_id = _new_import(app)
        pipeline.run_read(FakeCtx(), import_id)
        pipeline.run_render(FakeCtx(), import_id)

    def reread_started(*_a, **_k):
        store.update_import(import_id, status="preview")   # 別のタブの読み込み直しが md を消した
        raise FileNotFoundError("md")

    monkeypatch.setattr(pipeline, builder, reread_started)
    res = getattr(client, method)(f"/tables/imports/{import_id}/{url}")
    # 404「ダウンロード済み」ではなく、画面に戻して読み込み直しが始まったことを知らせる
    assert res.status_code == 302 and res.headers["Location"].endswith("/tables/new")
    with client.session_transaction() as session:
        assert any("読み込み直しが始まった" in text for _level, text in session["_flashes"])
    with app.app_context():
        assert store.get_import(import_id) is not None


# ---- R6B-1: 入力したキーが重なったら黙って「_2」にせず断る -------------------------------------------------------

def test_duplicate_typed_key_is_refused_and_does_not_carry_another_columns_settings():
    from tables.spec import spec_from_suggestions
    from views.tables import _build_spec

    base = spec_from_suggestions("T", {"header_rows": [1]}, [
        {"header": "台帳No", "key": "record_no", "role": "key", "type": "code"},
        {"header": "起票日", "key": "date", "role": "date", "type": "datetime"}])
    base.columns[0].value_map = {"A": "B"}
    base.columns[0].headers = ["台帳No", "管理番号"]
    rows = [{"use": 1, "header": "台帳No", "key": "record_no", "role": "key", "type": "code"},
            {"use": 1, "header": "起票日", "key": "record_no", "role": "date", "type": "datetime"}]
    _spec, errors = _build_spec({"name": "T", "columns": rows}, base)
    assert "キー「record_no」が重複しています" in errors

    # キーを入れ替えただけなら保存できるが、別の列の見出しの別名・値の置き換えは引き継がない
    swapped = [{"use": 1, "header": "台帳No", "key": "date", "role": "key", "type": "code"},
               {"use": 1, "header": "起票日", "key": "record_no", "role": "date", "type": "datetime"}]
    spec, errors = _build_spec({"name": "T", "columns": swapped}, base)
    assert errors == []
    dated = spec.column("record_no")
    assert dated.value_map == {} and "台帳No" not in dated.headers and "管理番号" not in dated.headers


def test_duplicate_typed_key_is_refused_by_the_editor(app, client):
    from tests.tables_helpers import json_only_spec, save_columns
    from tests.test_tables_fixes4 import _editor_body

    import_id, _template_id, body = _editor_body(app, client, json_only_spec())
    keys = [r["key"] for r in body["columns"] if r.get("use")]
    body["columns"][1]["key"] = keys[0]
    res = save_columns(client, import_id, body)
    assert res.status_code == 400
    assert f"キー「{keys[0]}」が重複しています" in json.dumps(res.get_json(), ensure_ascii=False)


# ---- R6B-2: md のパスが Windows の上限を超えるときは分かる言葉で止める --------------------------------------------------

def test_too_long_md_path_gives_a_clear_message(tmp_path, monkeypatch):
    from core import files as core_files
    from tables.markdown import MdFile

    monkeypatch.setattr(core_files, "_long_paths_enabled", lambda: False)
    target = tmp_path / "md"
    tmp_path.mkdir(exist_ok=True)
    name = "あ" * (core_files.MAX_PATH_CHARS - len(str(target)) + 10) + ".md"
    files = [MdFile(name=name, text="x", kind="records")]
    with pytest.raises(pipeline.PipelineError, match="ファイル名が長すぎ"):
        pipeline._write_md_dir(target, files)
    assert not target.exists()


# ---- R6-T-01: 記録番号・担当者の列の「〃」「同上」も直前の行の値で補う ------------------------------------------

def test_ditto_in_the_record_number_and_person_columns_is_filled(tmp_path):
    from tables.markdown import render_all
    from tests.test_tables_fixes3 import _auto
    from tests.test_tables_fixes4 import _csv

    lines = ["管理No,発生日,設備ID,現象,対応者,停止時間(分)",
             "TR-001,2025-03-01,CMP-101,異音,田中,30",
             "〃,〃,〃,追加調査,〃,20",
             "〃,〃,〃,部品交換,同上,10",
             "TR-002,2025-03-02,ETC-301,停止,佐藤,40",
             "〃,2025-03-02,ETC-301,復旧,佐藤,5"]
    src = _csv(tmp_path, "ditto_key.csv", lines)
    _layout, spec, records, stats = _auto(src, src.sheets()[0].name)
    assert {c.key: c.role for c in spec.columns}["record_no"] == "key"
    # 2行目以降は直前の伝票番号で補われ、同じ伝票の続きとして #2・#3 になる
    assert [r.key for r in records] == ["TR-001", "TR-001#2", "TR-001#3", "TR-002", "TR-002#2"]
    assert stats.ditto_filled.get("record_no") == 3 and stats.ditto_filled.get("worker") == 2
    assert "〃" not in stats.duplicate_keys
    text = "\n".join(f.text for f in render_all(spec, [r.to_dict() for r in records], {}, {}))
    assert "〃" not in text and "同上" not in text


# ---- R6-T-02: 記録キーの「列:文字数」を認める／記録キーの代わりの列も確かめる ------------------------------------

def _key_spec(**record):
    return spec_from_dict({"name": "x", "columns": [
        {"key": "occurred_at", "display": "発生日", "type": "date", "role": "date"},
        {"key": "equipment_id", "display": "設備番号", "type": "code", "role": "entity"},
        {"key": "symptom", "display": "現象", "type": "text", "role": "text"},
    ], "record": record})


def test_record_key_accepts_a_truncated_column():
    assert validate_spec(_key_spec(key=["occurred_at", "equipment_id", "symptom:20"])) == []
    assert validate_spec(_key_spec(key=["symptom:20"], fallback_key=["occurred_at"])) == []


def test_record_key_with_a_missing_column_or_a_bad_length_is_refused():
    assert validate_spec(_key_spec(key=["missing"])) == ["記録キーの列「missing」がありません"]
    assert validate_spec(_key_spec(key=["missing:20"])) == ["記録キーの列「missing」がありません"]
    errors = validate_spec(_key_spec(key=["symptom:あ"]))
    assert errors and "記録キーの「symptom:あ」の文字数" in errors[0]


def test_fallback_key_columns_are_checked_when_they_are_written_by_hand():
    errors = validate_spec(_key_spec(key=["equipment_id"], fallback_key=["equipment_no_typo"]))
    assert errors == ["記録キーの代わりの列「equipment_no_typo」がありません"]
    # 既定のままの代わりのキーは、その表に無い列名でも問題にしない（既定は一般的な列名）
    assert validate_spec(_key_spec(key=["equipment_id"])) == []


# ---- R6-T-03: 列の範囲は指定できないので、Excel 側で直す案内にする --------------------------------------------

def test_right_hand_table_warning_tells_what_can_actually_be_done(tmp_path):
    from tables.detect import guess_layout
    from tests.test_tables_fixes4 import _csv

    lines = ["管理No,発生日,現象,,設備No,点検日,結果"]
    for i in range(10):
        lines.append(f"TR-{i:03d},2025-03-{i + 1:02d},停止,,EQ-{i:03d},2025-03-{i + 1:02d},良")
    src = _csv(tmp_path, "two_tables.csv", lines)
    layout = guess_layout(src, src.sheets()[0].name)
    warning = next(w for w in layout.warnings if "別の表があるようです" in w)
    assert "別のシートか別のファイルに分けて" in warning and "範囲を指定して取り込んで" not in warning
    assert layout.headers == ["管理No", "発生日", "現象"]


def test_far_memo_column_warning_tells_what_can_actually_be_done(tmp_path):
    from tables.detect import guess_layout
    from tests.test_tables_fixes4 import _csv

    blanks = "," * 22
    lines = [f"管理No,発生日,現象{blanks},メモ"] + [f"TR-{i:03d},2025-03-{i + 1:02d},停止{blanks}," for i in range(10)]
    src = _csv(tmp_path, "far_memo.csv", lines)
    layout = guess_layout(src, src.sheets()[0].name)
    warning = next(w for w in layout.warnings if "大きく離れた" in w)
    assert "表の右隣に移してから" in warning and "範囲を指定して取り込んで" not in warning


# ---- R6-T-04: 判定に使っていない項目は取り込み設定に持たない ----------------------------------------------------

def test_settings_that_never_changed_the_reading_are_dropped():
    from tables.spec import spec_to_dict

    col = {"key": "a", "display": "a", "type": "string", "role": "attribute"}
    spec = spec_from_dict({"name": "x", "columns": [col],
                           "header": {"rows": 1, "anchors": ["管理No"], "search_rows": 50},
                           "data_end": {"blank_rows": 9, "stop_first_col": ["おわり"]},
                           "exclude": {"aggregate_keywords": ["計"], "hidden_rows": "include"}})
    assert spec.header == {"rows": 1, "anchors": ["管理No"]}
    assert not hasattr(spec, "data_end")
    assert spec.exclude == {"hidden_rows": "include", "strike_rows": "exclude_with_warning"}
    d = spec_to_dict(spec)
    assert "data_end" not in d and "search_rows" not in d["header"] and "aggregate_keywords" not in d["exclude"]
    assert validate_spec(spec) == []


# ---- ux6-2: データの行が無いときは、見出し行ではなくデータの範囲のことを言う ----------------------------------------

def _csv_import(client, name: str, text: str) -> int:
    from tests.tables_helpers import name_source, upload_csv

    import_id = upload_csv(client, name, text, encoding="utf-8")
    name_source(client, import_id, name, encoding="utf-8")
    return import_id


def test_a_table_with_no_data_rows_does_not_blame_the_header_row(client):
    from tests.tables_helpers import save_layout

    import_id = _csv_import(client, "empty.csv", "管理No,発生日,設備番号,現象,対応内容\r\n")
    res = save_layout(client, import_id)
    assert res.status_code == 400
    error = res.get_json()["error"]
    assert "データの行が1行もありません" in error
    assert "見出し行の番号を指定してください" not in error


def test_a_data_end_above_the_header_is_echoed_back_with_its_own_message(client):
    from tests.tables_helpers import panel_html, save_layout

    text = "管理No,発生日,現象\r\n" + "".join(f"TR-{i},2026/08/{i + 1:02d},停止\r\n" for i in range(10))
    import_id = _csv_import(client, "rows.csv", text)
    res = save_layout(client, import_id, data_end_row="1")
    assert res.status_code == 400
    error = res.get_json()["error"]
    assert "データの終わりに1行目を指定すると、データの行が1行もありません" in error
    assert "見出し行の番号を指定してください" not in error
    # 断られた入力は保存しないので、段を開き直しても自動判定のまま（画面は入力した値をそのまま残す）
    assert 'name="data_end_row" class="input" value=""' in panel_html(client, import_id, "layout")


# ---- r6-c1: 成功して終わったばかりの読み込みを［中止］が巻き戻さない --------------------------------------------

def test_cancel_does_not_undo_a_read_that_just_finished(app, client, monkeypatch):
    from core import jobs
    from models import database

    with app.app_context():
        _template_id, import_id = _new_import(app)
        pipeline.run_read(FakeCtx(), import_id)
        db = database.get_db()
        ts = database.now()
        cur = db.execute(
            "INSERT INTO jobs (kind, ref_type, ref_id, status, params_json, progress_json, created_at, updated_at) "
            "VALUES ('table_read', 'table_import', ?, 'running', '{}', '{}', ?, ?)", (import_id, ts, ts))
        job_id = cur.lastrowid
        db.commit()
        store.update_import(import_id, status="reading", job_id=job_id)   # 画面が読んだときは「読み込み中」

    def cancel_after_the_job_finished(jid):
        # 中止を要求した瞬間に、ジョブ本体は成功して status=preview を書き終えていた
        with app.app_context():
            store.update_import(import_id, status="preview")
            db = database.get_db()
            db.execute("UPDATE jobs SET status = 'cancelled' WHERE id = ?", (jid,))
            db.commit()
        return True

    monkeypatch.setattr(jobs, "request_cancel", cancel_after_the_job_finished)
    res = client.post(f"/tables/imports/{import_id}/cancel")
    assert res.status_code == 200 and res.get_json()["ok"] is True
    with app.app_context():
        assert store.get_import(import_id)["status"] == "preview"


# ---- r6-c2: 版番号は INSERT の中で数える（同時に保存しても衝突しない） -------------------------------------------

def test_new_versions_are_numbered_without_a_race(app):
    from tables.spec import ColumnSpec

    with app.app_context():
        col = {"key": "a", "display": "a", "type": "string", "role": "attribute"}
        spec = spec_from_dict({"name": "版の確認", "columns": [col]})
        template_id, version_id = store.create_template(spec.name, spec)
        for i in range(2, 5):
            store.mark_version_used(version_id)   # 確定に使われた版は上書きしない＝新しい版になる
            spec.columns.append(ColumnSpec(f"c{i}", f"列{i}"))
            version_id = store.save_template_version(template_id, spec)
            assert store.get_template(template_id)["version"] == i


def test_a_version_conflict_is_not_explained_as_a_duplicate_name():
    import sqlite3

    from views.tables import _save_conflict_message

    name_taken = sqlite3.IntegrityError("UNIQUE constraint failed: table_templates.name")
    version_clash = sqlite3.IntegrityError("UNIQUE constraint failed: table_template_versions.template_id, "
                                           "table_template_versions.version")
    assert "別の名前にしてください" in _save_conflict_message(name_taken, "トラブル対応一覧")
    assert _save_conflict_message(version_clash, "トラブル対応一覧") == "保存が他の操作と重なりました。もう一度保存してください"


# ---- ux6-6: ［取り込みを削除］の確認文は取り込みのどの画面でも同じ ------------------------------------------------

def test_the_delete_confirm_text_is_the_same_on_every_table_screen():
    root = Path(__file__).resolve().parents[1] / "templates"
    shared = "この取り込みを削除します。読み込んだ内容と作成した Markdown も消えます（取り込み設定は残ります）。元に戻せません。"
    found = [line.strip() for path in root.rglob("*.html")
             for line in path.read_text(encoding="utf-8").splitlines() if "この取り込みを削除します" in line]
    assert found  # 文言は共通のマクロ（と done.html の「ダウンロードせずに」）だけ
    for line in found:
        assert shared in line and "アップロードしたファイル" not in line
