from __future__ import annotations

import copy
import hashlib
import io
import json
import re
import zipfile
from datetime import date, time, timedelta, datetime
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

import core
import database as db
from core import UploadError
from forms import (
    extract_document,
    normalize_label,
    to_date,
    to_number,
    load_workbook_info,
    suggest_rows,
    rows_to_pattern,
    rank_patterns,
    refresh_summary,
    build_markdown,
    markdown_filename,
    suggest_title_fields,
    detect_tables,
    find_table,
    section_of,
    sections,
    pattern_to_rows,
    FieldDef,
    PatternDef,
    apply_manual_values,
    cell_text,
    format_unit,
    sections_of,
    _section_value,
    SheetDef,
    click_field,
    table_cells,
)
from scripts.samples import (
    f1_repair_report as f1,
    f2_trouble_report as f2,
    f3_8d_report as f3,
    f4_inspection_report as f4,
    f5_process_abnormality as f5,
)
from tests.conftest import add_confirmed_document
from views import SESSION_ID_KEY



# ====================================================================================================
# 元 tests/test_extraction.py
# ====================================================================================================

META = {"name": "設備修理報告書", "version": "v1", "description": "", "image_processing": "none"}


def build_repair_pattern(infos):
    sheets, fields = suggest_rows(infos)
    return rows_to_pattern(1, META, sheets, fields)


def test_normalize_label_absorbs_notation_differences():
    assert normalize_label("設備№：") == normalize_label(" 設備 No. ") == normalize_label("【設備NO】") == "設備no"


def test_value_conversion():
    assert to_date("2026/9/1", "2026/9/1") == ("2026-09-01", None)
    assert to_date("令和8年9月14日", "令和8年9月14日") == ("2026-09-14", None)
    assert to_number("1.5時間", "1.5時間")[0] == 1.5
    assert to_number(3.0, "3") == (3, None)


def test_builder_suggests_standard_fields_from_samples(repair_infos):
    sheets, fields = suggest_rows(repair_infos)
    used_sheets = [s["sheet_name"] for s in sheets if s["use"]]
    assert used_sheets == ["修理報告書"]  # マスタ・参考資料・(2) の重複は選ばない

    used = {f["field_name"] for f in fields if f["use"]}
    assert {"report_id", "equipment_id", "equipment_name", "occurred_date", "reporter",
            "symptom", "cause", "repair", "result"} <= used
    equipment = next(f for f in fields if f["field_name"] == "equipment_id")
    assert equipment["seen"] == "3/3"
    assert "装置番号" in equipment["candidates"].splitlines()


def test_same_json_from_different_layouts(repair_infos):
    pattern = build_repair_pattern(repair_infos)
    expected = [
        {"report_id": "R2026-00123", "equipment_id": "EQ-001", "occurred_date": "2026-09-14", "work_hours": 2.5},
        {"report_id": "R2026-00124", "equipment_id": "EQ-002", "occurred_date": "2026-09-10", "work_hours": 1.5,
         "equipment_name": "CVD装置"},
        {"report_id": "R2026-00125", "equipment_id": "EQ-003", "occurred_date": "2026-08-30", "cause": "リニアガイドの潤滑不足。"},
    ]
    for info, values in zip(repair_infos, expected):
        match = rank_patterns(info, [pattern])[0]
        extraction = extract_document(info, pattern, match.sheet_names)
        for key, value in values.items():
            assert extraction["values"][key] == value, (info.path.name, key)
        assert "Robot Position Error" in extraction["values"]["symptom"] or info is not repair_infos[0]


def test_images_are_detected_with_location(repair_infos):
    standard, shifted, table = repair_infos
    assert [(i["sheet"], i["location"]) for i in standard.images] == [("修理報告書", "J3")]
    assert len(shifted.images) == 2
    assert table.images == []


def test_matcher_prefers_correct_pattern(repair_infos, sample_dir):
    inspection_info = load_workbook_info(sample_dir / "inspection.xlsx")
    repair = build_repair_pattern(repair_infos)
    sheets, fields = suggest_rows([inspection_info])
    inspection = rows_to_pattern(2, {**META, "name": "点検記録表"}, sheets, fields)

    assert rank_patterns(repair_infos[0], [inspection, repair])[0].pattern.name == "設備修理報告書"
    ranked = rank_patterns(inspection_info, [repair, inspection])
    assert ranked[0].pattern.name == "点検記録表"
    assert ranked[1].confidence < 50


# ---- 種類定義の拡張（単位・Markdownでの扱い・タイトル項目） ----

def test_to_date_supports_1904_workbooks():
    assert to_date(45000, "45000") == ("2023-03-15", None)
    assert to_date(45000 - 1462, "43538", date1904=True) == ("2023-03-15", None)


def test_split_label_unit_and_value_unit():
    from forms import nfkc_value, split_label_unit, value_unit

    assert split_label_unit("作業時間(h)") == ("作業時間", "時間")
    assert split_label_unit("停止時間（分）") == ("停止時間", "分")
    assert split_label_unit("設備番号") == ("設備番号", "")
    assert value_unit("1.5時間") == "時間" and value_unit("EQ-001") == ""
    assert nfkc_value(" ＣＭＰ　　装置 \r\nﾎﾟﾝﾌﾟ") == "CMP 装置\nポンプ"
    # ラベル比較用の正規化は従来どおり空白を消す
    assert normalize_label("ＣＭＰ　装置") == "cmp装置"


def test_builder_suggests_units_and_title_fields(repair_infos, tmp_path):
    from openpyxl import Workbook
    from forms import suggest_title_fields

    sheets, fields = suggest_rows(repair_infos)
    work_hours = next(f for f in fields if f["field_name"] == "work_hours")
    assert work_hours["unit"] == "時間"  # 値「1.5時間」から
    assert all(f["rag_output"] == "show" for f in fields)
    assert suggest_title_fields(fields) == ["report_id", "equipment_id", "equipment_name", "occurred_date"]

    wb = Workbook()
    ws = wb.active
    ws.title = "報告書"
    ws["A1"], ws["B1"] = "報告番号", "R-1"
    ws["A2"], ws["B2"] = "作業時間(h)", 3
    ws["A3"], ws["B3"] = "停止時間（分）", 45
    wb.save(tmp_path / "unit.xlsx")
    _, rows = suggest_rows([load_workbook_info(tmp_path / "unit.xlsx")])
    by_name = {r["field_name"]: r for r in rows}
    assert by_name["work_hours"]["unit"] == "時間" and by_name["work_hours"]["display_name"] == "作業時間"
    assert "作業時間(h)" in by_name["work_hours"]["candidates"].splitlines()
    assert by_name["downtime"]["unit"] == "分"


def _field_row(field_name, display_name, data_type="string", **extra):
    row = {"use": True, "field_name": field_name, "display_name": display_name, "candidates": display_name,
           "data_type": data_type, "direction": "auto", "unit": "", "rag_output": "show",
           "table_columns": "", "section": ""}
    row.update(extra)
    return row


def test_rows_roundtrip_unit_rag_output_and_title_fields():
    from forms import pattern_to_meta, pattern_to_rows

    meta = {"name": "設備修理報告書", "version": "v1",
            "title_fields": ["report_id", "equipment_id", "unknown"]}
    sheet_rows = []
    field_rows = [_field_row("report_id", "報告番号"), _field_row("equipment_id", "設備番号"),
                  _field_row("work_hours", "作業時間", "number", unit="時間", rag_output="omit")]
    pattern = rows_to_pattern(3, meta, sheet_rows, field_rows)
    work = pattern.fields[2]
    assert (work.unit, work.rag_output) == ("時間", "omit")
    assert pattern.title_fields == ["report_id", "equipment_id"] and pattern.version_no == 1
    _, rows = pattern_to_rows(pattern)
    assert rows[2]["unit"] == "時間" and rows[2]["rag_output"] == "omit"
    assert pattern_to_meta(pattern)["title_fields"] == ["report_id", "equipment_id"]


def test_extraction_carries_unit_rag_output_and_title_fields(repair_infos):
    pattern = build_repair_pattern(repair_infos)
    pattern.title_fields = ["report_id", "occurred_date"]
    pattern.version_no = 4
    pattern.fields[0].rag_output = "omit"
    info = repair_infos[0]
    extraction = extract_document(info, pattern, rank_patterns(info, [pattern])[0].sheet_names)
    assert extraction["pattern"]["title_fields"] == ["report_id", "occurred_date"]
    assert extraction["pattern"]["version_no"] == 4
    work = next(f for f in extraction["fields"] if f["field_name"] == "work_hours")
    assert work["unit"] == "時間" and work["rag_output"] == "show"
    assert extraction["fields"][0]["rag_output"] == "omit"


# ---- 一覧表らしさ ----

def test_table_like_sheets(repair_infos, sample_dir, tmp_path):
    from openpyxl import Workbook
    from forms import table_like_sheets

    for info in repair_infos + [load_workbook_info(sample_dir / "inspection.xlsx")]:
        assert not table_like_sheets(info), info.path.name

    wb = Workbook()
    ws = wb.active
    ws.title = "一覧"
    ws["A1"] = "トラブル対応一覧"
    ws.append([])
    ws.append(["管理No", "発生日", "設備番号", "現象", "停止時間(分)"])
    for i in range(12):
        ws.append([f"TR-{i:03d}", f"2026-08-{i + 1:02d}", "EQ-001", "停止" if i % 3 else None, i * 5])
    notes = wb.create_sheet("記入要領")
    notes["A1"] = "この様式の書き方"
    wb.save(tmp_path / "list.xlsx")
    info = load_workbook_info(tmp_path / "list.xlsx")
    assert table_like_sheets(info) == ["一覧"]
    assert not table_like_sheets(info, ["記入要領"])

    # 9行しかない表は一覧表とみなさない
    ws.delete_rows(13, 3)
    wb.save(tmp_path / "short.xlsx")
    assert not table_like_sheets(load_workbook_info(tmp_path / "short.xlsx"))


# ---- 見出し欄の色・チェックボックス・項番・続きのシート（帳票サンプルの不一致の分析から） ----

LABEL_FILL = "FFD9E1F2"
HEADING_FILL = "FFFFF2CC"


def _book(tmp_path, name: str, sheets: dict):
    """sheets: {シート名: (cells, fills, merges)}"""
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill

    wb = Workbook()
    wb.remove(wb.active)
    for title, (cells, fills, merges) in sheets.items():
        ws = wb.create_sheet(title)
        for coord, value in cells.items():
            ws[coord] = value
        for coord, color in fills.items():
            ws[coord].fill = PatternFill("solid", fgColor=color)
        for rng in merges:
            ws.merge_cells(rng)
    wb.save(tmp_path / name)
    return load_workbook_info(tmp_path / name)


def _fields(*defs):
    from forms import FieldDef, PatternDef

    fields = [FieldDef(name, labels[0], list(labels), data_type=data_type) for name, data_type, *labels in defs]
    return PatternDef(name="テスト", fields=fields)


def _values(info, pattern, sheets=None):
    extraction = extract_document(info, pattern, sheets or info.sheet_names)
    return extraction["values"], {f["field_name"]: f for f in extraction["fields"]}


def test_same_fill_neighbour_is_a_label_so_value_is_read_below(tmp_path):
    info = _book(tmp_path, "stamp.xlsx", {"報告書": (
        {"A1": "承認", "B1": "確認", "C1": "作成", "A2": "高橋", "B2": "中村", "C2": "松本",
         "A4": "報告番号", "B4": "発生日時", "A5": "MR-001", "B5": "2026-09-01 10:00",
         "A7": "故障内容", "A8": "搬送アームが停止した", "A9": "アラーム E-100", "A10": "写真"},
        {"A1": LABEL_FILL, "B1": LABEL_FILL, "C1": LABEL_FILL, "A4": LABEL_FILL, "B4": LABEL_FILL,
         "A7": HEADING_FILL, "A10": HEADING_FILL},
        [],
    )})
    pattern = _fields(("approver", "string", "承認"), ("checker", "string", "確認"), ("report_id", "string", "報告番号"),
                      ("occurred", "string", "発生日時"), ("symptom", "text", "故障内容"))
    values, _ = _values(info, pattern)
    assert values["approver"] == "高橋" and values["checker"] == "中村"  # 右隣の「確認」を値にしない
    assert values["report_id"] == "MR-001" and values["occurred"] == "2026-09-01 10:00"
    assert values["symptom"] == "搬送アームが停止した\nアラーム E-100"  # 同じ色の見出し「写真」で止まる


def test_checkbox_notation_is_read_as_the_checked_option(tmp_path):
    from forms import pick_checked

    assert pick_checked("□重大　■大　□中") == ("大", None)
    assert pick_checked("☑同型機　□類似設備　☑他Fab") == ("同型機、他Fab", None)
    assert pick_checked("□有　□無") == (None, "チェックの入った選択肢がありません")
    assert pick_checked("■ 交換部品") is None and pick_checked("□良\n■要観察") is None

    info = _book(tmp_path, "check.xlsx", {"報告書": ({"A1": "重要度", "B1": "□重大　■大　□中",
                                                     "A2": "流出", "B2": "□有　□無"}, {}, [])})
    values, fields = _values(info, _fields(("severity", "string", "重要度"), ("outflow", "string", "流出")))
    assert values["severity"] == "大"
    assert values["outflow"] is None and fields["outflow"]["warning"] == "チェックの入った選択肢がありません"


def test_label_notation_section_numbers_and_parentheses(tmp_path):
    from forms import section_stripped

    assert normalize_label("停止時間(分)") == "停止時間(分)"  # 閉じ括弧だけを消さない
    assert normalize_label("【設備NO】") == "設備no"
    assert section_stripped("３．暫定対策") == "暫定対策" and section_stripped("①何が") == "何が"
    assert section_stripped("D2 問題の記述") == "問題の記述"
    assert section_stripped("2号機") == "" and section_stripped("95.0%以上") == "" and section_stripped("1.5") == ""

    info = _book(tmp_path, "labels.xlsx", {"報告書": (
        {"A1": "３．暫定対策", "B1": "ロットをHOLD", "A2": "停止時間（分）", "B2": 45,
         "A3": "原因（推定）", "B3": "摩耗の可能性", "A4": "原因（確定）", "B4": "ベアリング摩耗"}, {}, [])})
    pattern = _fields(("repair", "text", "暫定対策"), ("downtime", "number", "停止時間"), ("cause", "text", "原因（確定）"))
    values, _ = _values(info, pattern)
    assert values["repair"] == "ロットをHOLD"
    assert values["downtime"] == 45  # 候補どおりのラベルが無いときだけ括弧書きの違いを許す
    assert values["cause"] == "ベアリング摩耗"  # 「原因（推定）」と取り違えない


def test_label_in_a_table_header_is_used_only_when_there_is_no_other(tmp_path):
    cells = {"A1": "■ 時系列", "A2": "日時", "B2": "対応内容", "C2": "担当",
             "A3": "9/1 10:00", "B3": "アラーム発報", "C3": "山田", "A4": "9/1 10:30", "B4": "部品交換", "C4": "鈴木",
             "A6": "対応内容", "B6": "ベアリングを交換し、動作確認した"}
    fills = {"A2": LABEL_FILL, "B2": LABEL_FILL, "C2": LABEL_FILL, "A6": LABEL_FILL}
    info = _book(tmp_path, "timeline.xlsx", {"報告書": (cells, fills, [])})
    values, _ = _values(info, _fields(("repair", "text", "対応内容")))
    assert values["repair"] == "ベアリングを交換し、動作確認した"


def test_continuation_sheet_is_selected_with_the_main_sheet(tmp_path):
    from forms import match_pattern
    from forms import SheetDef

    first = ({"A1": "管理番号", "B1": "8D-001", "A2": "現象", "B2": "搬送停止", "A3": "原因", "B3": "摩耗"}, {}, [])
    second = ({"A1": "管理番号", "B1": "8D-001", "A2": "効果判定", "B2": "再発なし", "A3": "再発防止策", "B3": "月次点検"}, {}, [])
    notes = ({"A1": "記入要領", "A2": "効果判定", "A3": "再発防止策"}, {}, [])
    info = _book(tmp_path, "8d.xlsx", {"記入要領": notes, "8D報告(1)": first, "8D報告(2)": second})
    pattern = _fields(("report_id", "string", "管理番号"), ("symptom", "text", "現象"), ("cause", "text", "原因"),
                      ("result", "text", "効果判定"), ("prevention", "text", "再発防止策"))
    pattern.sheets = [SheetDef("8D報告(1)")]
    match = match_pattern(info, pattern)
    assert match.sheet_names == ["8D報告(1)", "8D報告(2)"]  # 値の無い記入要領は足さない
    values, _ = _values(info, pattern, match.sheet_names)
    assert values["result"] == "再発なし" and values["prevention"] == "月次点検"


# ---- 明細表 ----

def _parts_book(tmp_path, name="parts.xlsx", extra: dict | None = None):
    cells = {"A1": "報告番号", "B1": "MR-001",
             "A3": "■ 交換部品", "A4": "No.", "B4": "品番", "C4": "品名", "D4": "数量",
             "A5": 1, "B5": "PW48-1591", "C5": "ベアリング", "D5": 2,
             "A6": 2, "B6": "PW35-1577", "C6": "スピンモータ", "D6": 1,
             "A7": 3, "A8": 4,  # 様式の空き行（No だけ）
             "A9": "合計", "D9": 3,
             "A10": "所見", "B10": "異常なし"}
    cells.update(extra or {})
    fills = {c: LABEL_FILL for c in ("A4", "B4", "C4", "D4", "A9")}
    fills["A10"] = HEADING_FILL
    return _book(tmp_path, name, {"報告書": (cells, fills, ["A9:C9"])})


def test_table_field_reads_rows_until_total_row(tmp_path):
    info = _parts_book(tmp_path)
    values, fields = _values(info, _fields(("report_id", "string", "報告番号"), ("parts", "table", "交換部品")))
    assert values["report_id"] == "MR-001"
    assert values["parts"] == {
        "columns": ["品番", "品名", "数量"],  # 連番だけの No 列は落とす
        "rows": [["PW48-1591", "ベアリング", "2"], ["PW35-1577", "スピンモータ", "1"], ["合計", "", "3"]],
    }
    assert fields["parts"]["label_cell"] == "A3" and fields["parts"]["value_cell"] == "A4:D9"


def test_table_with_vertical_anchor_merged_cells_and_heading_stop(tmp_path):
    info = _book(tmp_path, "check.xlsx", {"点検": (
        {"A1": "使用部品", "B1": "品番", "C1": "品名", "B2": "PX84-2134", "C2": "スイッチ", "A4": "結果", "B4": "良",
         "A6": "Ａ．機構部", "A7": "点検部位", "B7": "点検項目", "C7": "判定",
         "A8": "ステージ", "B8": "位置偏差", "C8": "○", "B9": "冷却水流量", "C9": "×", "B10": "温度", "C10": "○",
         "A11": "総合判定", "B11": "要観察"},
        {"A1": LABEL_FILL, "B1": LABEL_FILL, "C1": LABEL_FILL, "A4": LABEL_FILL, "A7": LABEL_FILL, "B7": LABEL_FILL,
         "C7": LABEL_FILL, "A11": HEADING_FILL},
        ["A1:A3", "A8:A10"],
    )})
    values, _ = _values(info, _fields(("parts", "table", "使用部品"), ("check", "table", "機構部")))
    # 縦結合の見出しは右の列見出しから、見出しの範囲の行だけを読む
    assert values["parts"] == {"columns": ["品番", "品名"], "rows": [["PX84-2134", "スイッチ"]]}
    # 縦結合のデータセル（点検部位）は各行に同じ値。表の左端の色付きの「総合判定」で止まる
    assert values["check"]["rows"] == [["ステージ", "位置偏差", "○"], ["ステージ", "冷却水流量", "×"],
                                       ["ステージ", "温度", "○"]]


def test_table_none_rows_and_column_fallback(tmp_path):
    from forms import find_table_by_columns

    info = _parts_book(tmp_path, "none.xlsx", {"B5": "なし", "C5": None, "D5": None, "A6": None, "B6": None,
                                                 "C6": None, "D6": None, "A9": None, "D9": None})
    values, fields = _values(info, _fields(("parts", "table", "交換部品")))
    assert values["parts"] is None and fields["parts"]["warning"] == "表に行がありません"

    other = _parts_book(tmp_path, "renamed.xlsx", {"A3": "■ 部品交換実績"})
    pattern = _fields(("parts", "table", "交換部品"))
    assert _values(other, pattern)[0]["parts"] is None
    pattern.fields[0].table_columns = ["No.", "品番", "品名", "数量", "単価"]
    values, fields = _values(other, pattern)  # 見出しが違っても列見出しが似た表を読む
    assert values["parts"]["rows"][0] == ["PW48-1591", "ベアリング", "2"] and fields["parts"]["label_cell"] == "A3"
    assert find_table_by_columns(other.grids["報告書"], ["品番"]) is None  # 列見出し1つでは探さない


def test_builder_suggests_table_field_and_skips_column_header_fields(tmp_path):
    infos = [_parts_book(tmp_path, "a.xlsx"), _parts_book(tmp_path, "b.xlsx", {"A3": "■ 使用部品"})]
    _, rows = suggest_rows(infos)
    tables = [r for r in rows if r["data_type"] == "table"]
    assert len(tables) == 1  # 見出しの書き方が違っても、列見出しが同じ表は1つの候補
    parts = tables[0]
    assert parts["use"] and parts["field_name"] == "parts" and parts["display_name"] == "交換部品"
    assert {"■ 交換部品", "■ 使用部品"} <= set(parts["candidates"].splitlines())
    assert parts["table_columns"].splitlines() == ["No.", "品番", "品名", "数量"]
    assert parts["examples"][0].startswith("品番: PW48-1591／品名: ベアリング")
    used_labels = {r["display_name"] for r in rows if r["use"]}
    assert not used_labels & {"品番", "品名", "数量", "No."}  # 列見出しは1つの値の項目にしない

    # 見出しの行＋値の行1つ（「影響｜停止時間｜影響ロット」）は明細表にしない
    single = _book(tmp_path, "impact.xlsx", {"報告書": (
        {"A1": "影響", "B1": "停止時間", "C1": "影響ロット", "B2": 355, "C2": "PX2321022.2"},
        {"A1": LABEL_FILL, "B1": LABEL_FILL, "C1": LABEL_FILL}, ["A1:A2"])})
    _, rows = suggest_rows([single])
    assert not any(r["use"] for r in rows if r["data_type"] == "table")
    assert next(r for r in rows if r["field_name"] == "downtime")["use"]


def test_table_columns_roundtrip_through_rows():
    from forms import pattern_to_rows

    rows = [_field_row("parts", "交換部品", "table", candidates="■ 交換部品", table_columns="品番\n品名\n数量")]
    pattern = rows_to_pattern(1, {"name": "点検報告書"}, [], rows)
    assert pattern.fields[0].data_type == "table" and pattern.fields[0].table_columns == ["品番", "品名", "数量"]
    assert pattern_to_rows(pattern)[1][0]["table_columns"] == "品番\n品名\n数量"


def test_manual_table_edit_is_parsed_from_json(tmp_path):
    import json

    from forms import apply_manual_values

    info = _parts_book(tmp_path)
    extraction = extract_document(info, _fields(("parts", "table", "交換部品")), ["報告書"])
    edited = {"columns": ["品番", "品名", "数量"], "rows": [["PW48-1591", "ベアリング", "4"], ["", "", ""]]}
    apply_manual_values(extraction, {"value-parts": json.dumps(edited, ensure_ascii=False)})
    field = extraction["fields"][0]
    assert field["value"] == {"columns": ["品番", "品名", "数量"], "rows": [["PW48-1591", "ベアリング", "4"]]}
    assert field["edited"] and extraction["values"]["parts"] == field["value"]
    apply_manual_values(extraction, {"value-parts": "{壊れた"})
    assert field["value"]["rows"][0][2] == "4" and "形式" in field["warning"]
    apply_manual_values(extraction, {"value-parts": ""})
    assert field["value"] is None


def test_evaluator_compares_table_rows_and_treats_empty_marks_as_empty():
    from scripts.samples.evaluate_forms import column_value, compare, compare_table

    assert compare("string", "なし", "") == "empty_ok" and compare("string", None, "－") == "empty_ok"
    assert compare("string", "佐々木\n5/17", "佐々木") == "normalized"
    expected = [{"no": 1, "part_no": "PW-1", "qty": "１", "amount": "21,000"}, {"no": 2, "part_no": "PW-2", "qty": "2"}]
    got = {"columns": ["品番", "数量", "金額"], "rows": [["PW-1", "1", "21000"], ["PW-3", "2", ""]]}
    assert compare_table(got, expected) == {"files": 1, "expected_rows": 2, "extracted_rows": 2, "matched_rows": 1,
                                            "exact_tables": 0}
    lots = {"columns": ["ロットNo.", "投入数"], "rows": [["L1", "25"], ["L2", "13"], ["合計", "38"]]}
    assert column_value(lots, "ロットNo.") == "L1\nL2" and column_value(lots, "投入数") == "38"
    assert column_value(lots, "品番") is False


def test_evaluator_accepts_the_normalized_value_and_checkbox_booleans():
    """正解の *_norm（ISO の日時）も照合に使い、真偽値のチェック欄は ☑ / □ とみなす。"""
    from scripts.samples.evaluate_forms import compare_table

    timeline = [{"datetime": "5/23 12:07", "datetime_norm": "2025-05-23 12:07", "event": "停止"}]
    got = {"columns": ["日時", "内容"], "rows": [["2025-05-23 12:07", "停止"]]}
    assert compare_table(got, timeline)["matched_rows"] == 1
    targets = [{"target": "2号機", "checked": True}, {"target": "3号機", "checked": False}]
    got = {"columns": ["", "展開先"], "rows": [["☑", "2号機"], ["□", "3号機"]]}
    assert compare_table(got, targets)["matched_rows"] == 2
    assert compare_table({"columns": ["", "展開先"], "rows": [["□", "2号機"], ["☑", "3号機"]]}, targets)["matched_rows"] == 0


def test_combined_label_takes_the_matching_part_of_the_value(tmp_path):
    from forms import label_parts

    assert label_parts("ライン／工程") == ("ライン", "工程") and label_parts("L/min") == ()
    info = _book(tmp_path, "combined.xlsx", {"報告書": (
        {"A1": "ライン／工程", "B1": "L6／STI-CMP", "A2": "設備名/号機", "B2": "CMP 8号機"}, {}, [])})
    values, _ = _values(info, _fields(("line", "string", "ライン"), ("process", "string", "工程"),
                                      ("name", "string", "設備名")))
    assert values["line"] == "L6" and values["process"] == "STI-CMP"
    assert values["name"] is None  # 値の区切りの数が合わなければ読まない（取り違えるより空欄）


def test_table_header_below_a_heading_is_not_read_as_a_value(tmp_path):
    info = _parts_book(tmp_path, "heading.xlsx", {"A3": "暫定対策"})
    values, _ = _values(info, _fields(("repair", "text", "暫定対策")))
    assert values["repair"] is None  # 「No｜品番…」の列を文章として読まない


def test_duration_text_is_converted_to_the_field_unit():
    assert to_number("3時間40分", "3時間40分", "分") == (220, "「3時間40分」を分に換算しました")
    assert to_number("１時間３０分", "１時間３０分", "h")[0] == 1.5
    assert to_number("3時間40分", "3時間40分")[0] == 220  # 単位の無い項目は分にする（3 と読まない）


# ---- 年の無い日付・単位の無い数値（2026-09-19 の利用者の判断） ----

def test_date_without_a_year_keeps_the_original_text_and_warns():
    """「2/12 3時17分」に年は補わない（外れたときに誤った日付を Markdown に書くことになるため）。"""
    for text in ("2/12 3時17分", "2月12日 14:05", "10/16", "12-25 08:00"):
        value, warning = to_date(text, text)
        assert value == text, text
        assert warning and warning.startswith("年が書かれていません"), text
    # 年のある日付・年が2桁の書き方・そもそも日付でない記入は、今までどおり
    # 年のある日付は時刻（「3時17分」）も残す（F3-2: 以前は時刻を黙って落としていた）
    assert to_date("2026/2/12 3時17分", "2026/2/12 3時17分") == ("2026-02-12 03:17", None)
    assert to_date("24/8/25", "24/8/25")[1] == "日付として解釈できません"
    assert to_date("未復旧（対応中）", "未復旧（対応中）")[1] == "日付として解釈できません"
    assert to_date("13/40", "13/40")[1] == "日付として解釈できません"  # 月日として成り立たない数


def test_numeric_unit_comes_from_the_field_or_the_written_value():
    from forms import numeric_unit

    assert numeric_unit("390", "分") == "分" and numeric_unit("390", "h") == "時間"
    assert numeric_unit("390分") == "分" and numeric_unit("1,032分") == "分" and numeric_unit("2.5h") == "時間"
    assert numeric_unit("3時間40分") == "分"  # to_number が分に換算する
    assert numeric_unit("390") == "" and numeric_unit("CMP-101") == ""


def test_number_without_a_unit_is_flagged_only_when_the_unit_changes_the_meaning(tmp_path):
    """単位の無い数値は「単位不明」で要確認。ただし件数・回数のように単位が要らない項目は警告しない。"""
    from forms import ambiguous_unit_hint
    from forms import FieldDef, PatternDef

    assert ambiguous_unit_hint("downtime", "停止時間") and ambiguous_unit_hint("work_hours", "作業時間")
    assert ambiguous_unit_hint("cost", "修理費用") and ambiguous_unit_hint("f1", "部品の外径")
    assert not ambiguous_unit_hint("f2", "発生件数") and not ambiguous_unit_hint("f3", "対応回数")
    assert not ambiguous_unit_hint("f4", "立会人数") and not ambiguous_unit_hint("f5", "不良率")
    assert not ambiguous_unit_hint("report_id", "報告番号")

    info = _book(tmp_path, "unit.xlsx", {"報告書": (
        {"A1": "停止時間", "B1": "390", "A2": "発生件数", "B2": "3", "A3": "作業時間", "B3": "2.5h"},
        {"A1": LABEL_FILL, "A2": LABEL_FILL, "A3": LABEL_FILL}, [])})
    pattern = PatternDef(name="テスト", fields=[
        FieldDef("downtime", "停止時間", ["停止時間"], data_type="number"),
        FieldDef("count", "発生件数", ["発生件数"], data_type="number"),
        FieldDef("work_hours", "作業時間", ["作業時間"], data_type="number")])
    values, fields = _values(info, pattern)
    assert values["downtime"] == 390 and fields["downtime"]["unit"] == ""
    assert "単位が書かれていません" in fields["downtime"]["warning"]
    assert "「帳票の種類」の画面" in fields["downtime"]["warning"]  # 直し方を書く
    # 読み取り済みの帳票はその場で直せる（値に単位を付けて入力）ことを先に書く。種類の単位は読み取り済みに反映されない
    assert "値に単位を付けて入力してください（例: 390分）" in fields["downtime"]["warning"]
    assert "読み取り済みの帳票には反映されません" in fields["downtime"]["warning"]
    assert values["count"] == 3 and not fields["count"]["warning"]      # 件数は単位が無くて当たり前
    assert values["work_hours"] == 2.5 and fields["work_hours"]["unit"] == "時間"  # セルの「h」から補う
    assert not fields["work_hours"]["warning"]


def test_written_unit_wins_over_the_field_unit_and_is_flagged(tmp_path):
    """「停止時間（分）」の欄に「14.9h」と書かれた帳票。勝手に分として出さず、書かれたとおりにして要確認にする。"""
    from forms import FieldDef, PatternDef

    info = _book(tmp_path, "unit2.xlsx", {"報告書": (
        {"A1": "停止時間", "B1": "14.9h"}, {"A1": LABEL_FILL}, [])})
    pattern = PatternDef(name="テスト", fields=[
        FieldDef("downtime", "停止時間", ["停止時間"], data_type="number", unit="分")])
    _, fields = _values(info, pattern)
    assert fields["downtime"]["value"] == 14.9 and fields["downtime"]["unit"] == "時間"
    assert "「分」ですが" in fields["downtime"]["warning"] and "「時間」" in fields["downtime"]["warning"]


# ---- 表記違いのラベル・隣の値の優先・一覧の列見出し・設備番号と設備名のまとめ欄（帳票サンプルの不一致の分析 2回目） ----

def test_label_variants_are_used_when_the_exact_label_gives_no_value(tmp_path):
    from forms import label_base, split_label_unit

    assert split_label_unit("発生原因（推定）") == ("発生原因（推定）", "")  # 「推定」は単位でない
    assert split_label_unit("金額（千円）") == ("金額", "千円")
    assert label_base("応急処置内容") == "応急処置" and label_base("復旧完了日時") == "復旧完了"
    assert label_base("発生原因(なぜ起きたか)") == "発生原因" and label_base("処置") == ""

    cells = {"A1": "発生原因（なぜ起きたか）", "B1": "ベアリング摩耗",
             "A2": "応急処置内容", "B2": "ベアリング交換",
             "A3": "復旧完了", "B3": "2026/9/2",
             "A4": "原因（推定）", "B4": "摩耗の可能性", "A5": "原因（確定）",
             "A6": "状況", "B6": "完了", "B8": "次回確認は9/30"}
    fills = {c: LABEL_FILL for c in ("A1", "A2", "A3", "A4", "A5", "A6")}
    info = _book(tmp_path, "variants.xlsx", {"報告書": (cells, fills, [])})
    values, _ = _values(info, _fields(("cause", "text", "発生原因"), ("repair", "text", "応急処置"),
                                      ("completed", "date", "復旧完了日時", "完了日時"),
                                      ("confirmed", "text", "原因（確定）")))
    assert values["cause"] == "ベアリング摩耗" and values["repair"] == "ベアリング交換"
    assert values["completed"] == "2026-09-02"  # 値の「完了」（塗りなし）は「完了日時」の表記違いにしない
    assert values["confirmed"] is None  # 「原因（推定）」は括弧書きが違う別の項目


def test_adjacent_value_first_and_headings_are_not_values(tmp_path):
    cells = {"C1": "発信\n部署", "D1": "承認", "D2": "佐々木", "C5": "2024/2/1",
             "A6": "発信部署", "B6": "製造課 A班",
             "A8": "備考", "A10": "▼ 回答欄（受信部署で記入）", "A11": "回答者", "B11": "木村",
             "A13": "品証確認", "A14": "様式QA-05 Rev.2"}
    fills = {c: LABEL_FILL for c in ("C1", "D1", "A6", "A8", "A11", "A13")}
    fills["A10"] = HEADING_FILL
    info = _book(tmp_path, "stamp.xlsx", {"報告書": (cells, fills, ["C1:C3"])})
    values, _ = _values(info, _fields(("department", "string", "発信部署"), ("remarks", "text", "備考"),
                                      ("approver", "string", "承認"), ("qa", "string", "品証確認")))
    assert values["department"] == "製造課 A班"  # 押印欄の縦書き「発信部署」の2行下の日付ではない
    assert values["remarks"] is None  # 「▼ 回答欄」の見出しは値にしない
    assert values["approver"] == "佐々木"
    assert values["qa"] is None  # 欄外の様式番号は値にしない


def test_list_column_headers_and_seq_no_are_not_single_value_labels(tmp_path):
    cells = {"A1": "No.", "B1": "工異24-021", "C1": "承認", "D1": "確認", "E1": "作成",
             "C2": "高橋", "D2": "斎藤", "E2": "森", "C3": "6/5", "D3": "6/2", "E3": "5/31",
             "A5": "対象設備", "B5": "CMP-108　STI-CMP 8号機",
             "A7": "確認", "B7": "設備No", "C7": "設備名", "D7": "結果",
             "A8": "□", "B8": "CMP-101", "C8": "酸化膜CMP 1号機",
             "A9": "☑", "B9": "CMP-103", "C9": "W-CMP 3号機", "D9": "異常なし",
             "A10": "□", "B10": "CMP-104", "C10": "Cu-CMP 4号機",
             "A12": "■ 対象ロット", "A13": "No.", "B13": "ロットNo.", "C13": "数量",
             "A14": 1, "B14": "SR2503056.2", "C14": 12}
    fills = {c: LABEL_FILL for c in ("A1", "C1", "D1", "E1", "A5", "A7", "B7", "C7", "D7", "A13", "B13", "C13")}
    info = _book(tmp_path, "list.xlsx", {"報告書": (cells, fills, [])})
    values, _ = _values(info, _fields(("report_id", "string", "No", "管理番号"), ("approver", "string", "承認"),
                                      ("equipment_id", "string", "設備No"), ("equipment_name", "string", "設備名")))
    assert values["report_id"] == "工異24-021"  # 表の連番の「No.」の下の 1 ではない
    assert values["approver"] == "高橋"  # 押印と日付の2行の押印欄は、一覧の列見出しとみなさない
    # 一覧（水平展開先）の列見出しでなく、番号と名前をまとめた「対象設備」を分けて読む
    assert values["equipment_id"] == "CMP-108" and values["equipment_name"] == "STI-CMP 8号機"


def test_split_code_name_and_builder_suggests_equipment_from_combined_cell(tmp_path):
    from forms import split_code_name

    assert split_code_name("ROB-821（ウェーハソーター 1号機）") == ("ROB-821", "ウェーハソーター 1号機")
    assert split_code_name("W-CMP 3号機　CMP-103") == ("CMP-103", "W-CMP 3号機")
    assert split_code_name("EQ-001：プレス機") == ("EQ-001", "プレス機")
    assert split_code_name("CMP-101 / CMP-102") is None and split_code_name("L6／STI-CMP") is None
    assert split_code_name("CMP-108") is None and split_code_name("プレス機") is None

    info = _book(tmp_path, "combined_equipment.xlsx", {"報告書": (
        {"A1": "報告番号", "B1": "R-001", "A2": "使用設備", "B2": "CVD-202（TEOS CVD 2号機）", "A3": "設備", "B3": "空調"},
        {"A1": LABEL_FILL, "A2": LABEL_FILL, "A3": LABEL_FILL}, [])})
    _, rows = suggest_rows([info])
    by_name = {r["field_name"]: r for r in rows}
    assert by_name["equipment_id"]["use"] and by_name["equipment_id"]["examples"] == ["CVD-202"]
    assert by_name["equipment_name"]["examples"] == ["TEOS CVD 2号機"]
    assert "使用設備" in by_name["equipment_name"]["candidates"].splitlines()
    pattern = rows_to_pattern(1, META, *suggest_rows([info]))
    values, _ = _values(info, pattern)
    assert values["equipment_id"] == "CVD-202" and values["equipment_name"] == "TEOS CVD 2号機"


def test_evaluator_offset_picks_other_sample_files():
    from scripts.samples.evaluate_forms import pick_samples

    entries = [{"file": f"{v}{i}", "layout_version": v} for v in ("a", "b") for i in range(3)]
    assert [e["file"] for e in pick_samples(entries, 3)] == ["a0", "b0", "a1"]
    assert [e["file"] for e in pick_samples(entries, 3, offset=1)] == ["b1", "a1", "b2"]  # 版の順も1つずらす
    assert len({e["file"] for e in pick_samples(entries, 6, offset=4)}) == 6


def test_form_number_note_is_not_a_value_even_in_the_middle_of_a_cell(tmp_path):
    """欄外の様式番号は、先頭になくても項目の値にしない（design.md 6.1「値にしないもの」）。"""
    from forms import is_form_number_note

    assert is_form_number_note("様式MT-031 Rev.1")
    assert is_form_number_note("製造部 設備保全課 様式MT-031 Rev.2(2025.04改訂)")
    assert not is_form_number_note("様式を見直す")  # 番号が続かない文は注記ではない
    assert not is_form_number_note("旧様式AからBへ移行したが、記入の手順は変えていないので、" * 3)

    info = _book(tmp_path, "form_no.xlsx", {"報告書": (
        {"A1": "設備点検・保全作業報告書", "F1": "製造部 設備保全課 様式MT-031 Rev.2(2025.04改訂)",
         "A3": "作業No.", "B3": "WK-001"},
        {"A1": LABEL_FILL, "A3": LABEL_FILL}, [])})
    values, _ = _values(info, _fields(("title", "string", "設備点検・保全作業報告書"), ("work_no", "string", "作業No.")))
    assert values["title"] is None and values["work_no"] == "WK-001"


def test_stacked_detail_table_blocks_are_all_read(tmp_path):
    """同じ形の列見出しが縦に積み重なる表（特性要因図）は、組ごとに読んで1つの明細表にまとめる。"""
    from forms import FieldDef, PatternDef

    cells = {"A1": "特性要因図", "A2": "人", "B2": "機械", "A3": "教育不足", "B3": "電極の摩耗",
             "A4": "材料", "B4": "方法", "A5": "ガスの純度", "B5": "点検手順の不備",
             "A6": "測定", "B6": "環境", "A7": "校正ずれ", "B7": "室温の変動"}
    fills = {c: LABEL_FILL for c in ("A1", "A2", "B2", "A4", "B4", "A6", "B6")}
    info = _book(tmp_path, "fishbone.xlsx", {"報告書": (cells, fills, [])})
    pattern = PatternDef(name="テスト", fields=[FieldDef("cause_table", "特性要因図", ["特性要因図"], data_type="table")])
    values, fields = _values(info, pattern)
    value = values["cause_table"]
    assert value["columns"] == ["人", "機械", "材料", "方法", "測定", "環境"]
    # 行ごとに自分の組の列だけ埋まる（どの要因がどの見出しのものか、1行だけで分かる）
    assert value["rows"] == [["教育不足", "電極の摩耗", "", "", "", ""],
                             ["", "", "ガスの純度", "点検手順の不備", "", ""],
                             ["", "", "", "", "校正ずれ", "室温の変動"]]
    assert fields["cause_table"]["table_blocks"] == 3
    assert not fields["cause_table"]["warning"]  # 取りこぼしが無くなったので要確認にしない


def test_stacked_blocks_with_the_same_columns_just_add_rows(tmp_path):
    """列見出しが同じまま繰り返される様式（ページごとに見出しを書く表）は、行が増えるだけにする。"""
    from forms import FieldDef, PatternDef

    cells = {"A1": "■ 交換部品", "A2": "品番", "B2": "品名", "A3": "PW48-1591", "B3": "ベアリング",
             "A4": "品番", "B4": "品名", "A5": "PW35-1577", "B5": "スピンドル"}
    fills = {c: LABEL_FILL for c in ("A1", "A2", "B2", "A4", "B4")}
    info = _book(tmp_path, "parts2.xlsx", {"報告書": (cells, fills, [])})
    pattern = PatternDef(name="テスト", fields=[FieldDef("parts", "交換部品", ["交換部品"], data_type="table")])
    values, fields = _values(info, pattern)
    assert values["parts"] == {"columns": ["品番", "品名"],
                               "rows": [["PW48-1591", "ベアリング"], ["PW35-1577", "スピンドル"]]}
    assert fields["parts"]["table_blocks"] == 2


def test_builder_leaves_notes_legends_and_dashes_unchecked(tmp_path):
    """凡例・注記・「－」だけの欄は候補には残すが、既定ではチェックを付けない（design.md 6.1）。"""
    info = _book(tmp_path, "legend.xlsx", {"報告書": (
        {"A1": "判定　○", "B1": "良好　△：要観察　×：不良　－：対象外",
         "A2": "関連TR No.", "B2": "※故障発生日時・停止時間は事後保全時に記入",
         "A3": "品名", "B3": "－",
         "A4": "作業区分", "B4": "定期点検"},
        {c: LABEL_FILL for c in ("A1", "A2", "A3", "A4")}, [])})
    _, rows = suggest_rows([info])
    by_label = {r["display_name"]: r for r in rows}
    assert by_label["作業区分"]["use"]
    for label in ("判定　○", "関連TR No.", "品名"):
        assert not by_label[label]["use"], label
        assert label in by_label  # 候補としては残し、人が選べるようにする


def test_builder_reads_date_like_labels_written_as_text(tmp_path):
    """見本が文字列で日付を持つ欄（「2025/12/19(金)」「令和5年12月5日」）も日付型にする。"""
    info = _book(tmp_path, "dates.xlsx", {"報告書": (
        {"A1": "作業日", "B1": "2025/12/19(金)", "A2": "回答日", "B2": "令和5年12月5日",
         "A3": "曜日", "B3": "金", "A4": "作業日数", "B4": "3"},
        {c: LABEL_FILL for c in ("A1", "A2", "A3", "A4")}, [])})
    _, rows = suggest_rows([info])
    by_label = {r["display_name"]: r["data_type"] for r in rows}
    assert by_label["作業日"] == "date" and by_label["回答日"] == "date"
    assert by_label["曜日"] == "string"  # 値が日付として読めないものは日付にしない

    pattern = rows_to_pattern(1, META, *suggest_rows([info]))
    values, _ = _values(info, pattern)
    assert values[next(f.field_name for f in pattern.fields if f.display_name == "作業日")] == "2025-12-19"


def test_builder_reads_prefixed_number_label_as_report_id(tmp_path):
    """「8D No.」のように短い接頭辞が付いた識別番号の見出しも報告番号として読む。"""
    info = _book(tmp_path, "8d.xlsx", {"報告書": (
        {"A1": "8D No.", "B1": "8D-2023-011", "A2": "設備名", "B2": "P-SiN CVD 5号機"},
        {"A1": LABEL_FILL, "A2": LABEL_FILL}, [])})
    _, rows = suggest_rows([info])
    report = next(r for r in rows if r["field_name"] == "report_id")
    assert report["use"] and "8D No." in report["candidates"].splitlines()
    values, _ = _values(info, rows_to_pattern(1, META, *suggest_rows([info])))
    assert values["report_id"] == "8D-2023-011"


def test_empty_detail_table_keeps_columns_for_manual_entry(tmp_path):
    """明細表が読めなくても、列見出しは確認画面に残して人が行を入れられるようにする。"""
    from forms import clean_table_value
    from forms import FieldDef, PatternDef

    info = _book(tmp_path, "no_table.xlsx", {"報告書": (
        {"A1": "■ 交換部品", "A2": "なし"}, {"A1": HEADING_FILL}, [])})
    pattern = PatternDef(name="テスト", fields=[
        FieldDef("parts", "交換部品", ["■ 交換部品"], data_type="table", table_columns=["品番", "品名", "数量"])])
    _, fields = _values(info, pattern)
    assert fields["parts"]["value"] is None and fields["parts"]["table_columns"] == ["品番", "品名", "数量"]

    # 全部の行を消しても列見出しは残る（表の入力欄が消えない）。Markdown 側は行が無ければ出さない
    assert clean_table_value({"columns": ["品番", "品名"], "rows": []}) == {"columns": ["品番", "品名"], "rows": []}
    assert clean_table_value({"columns": [], "rows": []}) is None


def test_label_that_is_also_a_column_header_keeps_the_real_header_row(tmp_path):
    """見出し語が表の列見出しと同じでも、1行下のデータ行を列見出しにしない（F3 の D6）。

    見本と違う書き方の見出し（「D6．恒久対策の実施と効果の確認」）だと、候補ラベルのうち表の列見出しと
    同じ「実施内容」だけが当たる。そこから下に表を探すと先頭の明細が列見出しになり、1件目と
    完了日・状況の列が落ちていた。
    """
    from forms import FieldDef, PatternDef

    cells = {"A1": "D6．恒久対策の実施と効果の確認",
             "A2": "No", "B2": "実施内容", "C2": "担当", "D2": "完了日", "E2": "状況",
             "A3": 1, "B3": "樹脂交換時期を実負荷ベースで再設定", "C3": "前田", "D3": "2026-09-01", "E3": "完了",
             "A4": 2, "B4": "比抵抗の早期ワーニング閾値追加", "C4": "前田", "D4": "2026-09-05", "E4": "完了",
             "A5": 3, "B5": "供給水比抵抗をFDCで監視", "C5": "中島", "D5": "2026-09-10", "E5": "実施中"}
    fills = {c: LABEL_FILL for c in ("A1", "A2", "B2", "C2", "D2", "E2")}
    info = _book(tmp_path, "8d_d6.xlsx", {"報告書": (cells, fills, [])})
    fd = FieldDef("d6", "恒久対策の実施・効果確認", ["D6：恒久対策の実施・効果確認", "D6", "実施内容"], data_type="table")
    fd.table_columns = ["No", "実施内容", "担当", "完了日", "状況"]  # 見本から作った列見出し
    values, fields = _values(info, PatternDef(name="テスト", fields=[fd]))
    assert values["d6"]["columns"] == ["実施内容", "担当", "完了日", "状況"]  # 連番だけの No 列は落とす
    assert values["d6"]["rows"] == [
        ["樹脂交換時期を実負荷ベースで再設定", "前田", "2026-09-01", "完了"],
        ["比抵抗の早期ワーニング閾値追加", "前田", "2026-09-05", "完了"],
        ["供給水比抵抗をFDCで監視", "中島", "2026-09-10", "実施中"],
    ]
    assert not fields["d6"]["warning"]


# ====================================================================================================
# 元 tests/test_forms_flow.py
# 帳票の通しテスト（帳票登録 → 使用開始 → 取り込み → 読み取り → 確定 → ダウンロード）。
#
# 画面は /form-types と /forms の1枚ずつしかなく、どの操作も fetch で JSON か HTML の断片を受け取る。
# ほかの帳票テストからも使えるよう、画面を1つ進める手続きをこのファイルにまとめてある。
# ====================================================================================================

# ---- 画面を1つ進める（ほかの帳票テストからも使う） -------------------------------------------------
# 置いた Excel はサーバーに残らない（2026-09-21 の利用者の指示）。ブラウザが選んだファイルを
# 持ち続けて毎回送るのと同じように、テストでも「その種類で置いたファイル」を覚えて送り直す。

def book_part(path):
    """multipart で送る Excel（画面が毎回送っているものと同じ）。"""
    path = Path(path)
    return (io.BytesIO(path.read_bytes()), path.name)


def _books(client) -> dict:
    books = getattr(client, "form_books", None)
    if books is None:
        books = {}
        client.form_books = books
    return books


def create_type(client, path, name=None) -> int:
    """帳票の Excel を1つ置いて帳票の種類を作る（名前を省くとファイル名になる）。"""
    data = {"book": book_part(path)}
    if name is not None:
        data["name"] = name
    res = client.post("/form-types/new", data=data, content_type="multipart/form-data")
    assert res.status_code == 200, res.get_data(as_text=True)
    pattern_id = res.get_json()["pattern_id"]
    _books(client)[pattern_id] = Path(path)
    return pattern_id


def add_field(client, pattern_id: int, sheet: str, label_cell: str, value_cell: str = "", book=None) -> dict:
    """見出しのセル（と値のセル）をクリックして項目を1つ作る（画面と同じく Excel も一緒に送る）。"""
    data = {"sheet": sheet, "label_cell": label_cell, "value_cell": value_cell}
    path = book if book is not None else _books(client).get(pattern_id)
    if path is not None:
        data["book"] = book_part(path)
    res = client.post(f"/form-types/{pattern_id}/fields", data=data, content_type="multipart/form-data")
    assert res.status_code == 200, res.get_data(as_text=True)
    return res.get_json()


def panel_html(client, pattern_id: int, book=None) -> str:
    """登録中の欄（HTML の断片）。Excel を渡すとシートも出る（渡さなければ設定だけの画面）。"""
    if book is None:
        return client.get(f"/form-types/{pattern_id}/panel").get_json()["html"]
    res = client.post(f"/form-types/{pattern_id}/panel", data={"book": book_part(book)},
                      content_type="multipart/form-data")
    assert res.status_code == 200, res.get_data(as_text=True)
    return res.get_json()["html"]


def activate(client, pattern_id: int) -> None:
    res = client.post(f"/form-types/{pattern_id}/status", json={"status": "active"})
    assert res.status_code == 200 and res.get_json()["status"] == "active"


def upload_forms(client, *paths) -> list[int]:
    """帳票のファイルを置く（複数ならまとまり）。戻り値: 帳票ID の並び。"""
    files = [(io.BytesIO(p.read_bytes()), p.name) for p in paths]
    res = client.post("/forms/upload", data={"file": files}, content_type="multipart/form-data")
    assert res.status_code == 200, res.get_data(as_text=True)
    return [d["id"] for d in res.get_json()["docs"]]


def read_form(client, doc_id: int, pattern_id: int, sheets: list[str], **extra):
    return client.post(f"/forms/{doc_id}/read", data={"pattern_id": pattern_id, "sheets": sheets, **extra})


def finish(client, ids, current=None):
    query = ",".join(str(i) for i in ids)
    url = f"/forms/finish?ids={query}" + (f"&current={current}" if current else "")
    return client.get(url).get_json()


# ---- 通し -------------------------------------------------------------------------------

def test_form_flow_happy_path(app, client, sample_dir):
    # 1. 帳票登録: 帳票の Excel を置き、見出しと値のセルをクリックして項目を作る（Excel は残らない）
    path = sample_dir / "standard.xlsx"
    pattern_id = create_type(client, path, "設備修理報告書")
    page = client.get("/form-types/").get_data(as_text=True)
    assert "設備修理報告書" in page and "作成中" in page
    panel = panel_html(client, pattern_id, book=path)
    # クリックで作る画面: 置いた帳票のシートが出て、型やキー名の入力欄は無い
    assert 'data-cell="B3"' in panel and "まだ項目がありません" in panel
    assert "RAGに出す" not in panel and "キー名" not in panel

    assert "「報告番号」を項目にしました" in add_field(client, pattern_id, "修理報告書", "A3", "B3")["message"]
    add_field(client, pattern_id, "修理報告書", "E4", "F4")
    add_field(client, pattern_id, "修理報告書", "A7", "A8")
    with app.app_context():
        pattern = db.load_pattern(pattern_id)
        assert [f.field_name for f in pattern.fields] == ["report_id", "equipment_name", "symptom"]
        assert pattern.status == "draft"          # 保存しても使用中にはならない

    # 読み取りテスト: いまの設定で、置いた帳票を読んだ結果が同じ欄に出る
    panel = panel_html(client, pattern_id, book=path)
    assert "R2026-00123" in panel and "CMP装置" in panel and "3項目中 <strong>3</strong>項目" in panel

    # 使用開始を押すまでは帳票取り込みの候補に出ない
    assert client.get(f"/forms/1/type").status_code in (404, 409)
    activate(client, pattern_id)
    with app.app_context():
        assert db.load_pattern(pattern_id).status == "active"

    # 2. 帳票取り込み: ファイルを置く → 帳票の種類とシート
    assert client.get("/forms/").status_code == 200
    doc_id, = upload_forms(client, path)
    body = client.get(f"/forms/{doc_id}/type").get_json()
    assert body["has_types"] and body["file_name"] == "standard.xlsx"
    assert "3項目中3項目が見つかりました" in body["html"] and "設備修理報告書" in body["html"]

    # 3. 読み取り結果
    body = read_form(client, doc_id, pattern_id, ["修理報告書"]).get_json()
    assert 'data-cell="B3"' in body["html"] and 'name="value-report_id"' in body["html"]
    assert body["summary"]["markdown"].startswith("# 設備修理報告書 R2026-00123")
    version = body["version"]

    # 途中保存（204・新しい版を返す）とプレビュー
    res = client.post(f"/forms/{doc_id}/draft",
                      json={"values": {"equipment_name": "CMP研磨装置"}, "version": version})
    assert res.status_code == 204
    version = res.headers["X-Doc-Version"]
    summary = client.post(f"/forms/{doc_id}/preview", json={"values": {}}).get_json()
    assert "CMP研磨装置" in summary["markdown"] and summary["counts"]["manual"] == 1

    # 4. 確定してダウンロード
    state = finish(client, [doc_id])
    assert state["ready"] is False and state["read_yet"] is True and "確定して Markdown を作る" in state["html"]
    assert client.post(f"/forms/{doc_id}/confirm", json={"version": version}).get_json()["ok"] is True
    with app.app_context():
        assert db.get_document(doc_id)["state"] == "confirmed"
    state = finish(client, [doc_id])
    assert state["ready"] is True and state["confirmed"] == 1
    assert "ダウンロードすると、この帳票の元のファイルと読み取り結果はサーバーから消えます" in state["html"]
    assert f"/forms/{doc_id}/download.md" in state["html"]

    # 確定後に直すと「修正中」。確定し直せば確定済みに戻る
    client.post(f"/forms/{doc_id}/draft", json={"values": {"equipment_name": "別の名前"}})
    with app.app_context():
        assert db.get_document(doc_id)["state"] == "modified"
    assert client.post(f"/forms/{doc_id}/preview", json={"values": {}}).get_json()["state"] == "modified"

    # 確定済みの読み取り直しは、手の修正が消えることの確認が要る
    res = read_form(client, doc_id, pattern_id, ["修理報告書"])
    assert res.status_code == 400 and "確認のチェック" in res.get_json()["error"]
    assert read_form(client, doc_id, pattern_id, ["修理報告書"], acknowledge="on").status_code == 200

    # 5. ダウンロードが最後の手順（ダウンロードするとサーバーからデータが消える）
    client.post(f"/forms/{doc_id}/confirm", json={})
    md = client.get(f"/forms/{doc_id}/download.md")
    assert md.status_code == 200
    assert md.headers["Content-Type"] == "text/markdown; charset=utf-8"   # charset は1つだけ
    text = md.get_data(as_text=True)
    assert text.startswith("# 設備修理報告書 R2026-00123") and "CMP装置" in text
    assert "別の名前" not in text          # 確定済みの版で作る
    with app.app_context():
        assert db.get_document(doc_id) is None
    assert client.get(f"/forms/{doc_id}/type").status_code == 404
    assert client.get(f"/forms/{doc_id}/download.md").status_code == 404

    # 帳票の種類は設定なので残る
    assert "設備修理報告書" in client.get("/form-types/").get_data(as_text=True)


def test_an_empty_value_does_not_stop_the_confirm(app, client, sample_dir):
    """空欄でも確定は止まらない（「必須」の設定は帳票登録の画面に無く、どこにも無い）。

    空欄は今までどおり「空欄」の印が付くだけで、確定の前に消さなければならない案内は出ない。
    """
    path = sample_dir / "standard.xlsx"
    pattern_id = create_type(client, path, "設備修理報告書")
    add_field(client, pattern_id, "修理報告書", "A3", "B3")
    activate(client, pattern_id)

    doc_id, = upload_forms(client, path)
    version = read_form(client, doc_id, pattern_id, ["修理報告書"]).get_json()["version"]
    res = client.post(f"/forms/{doc_id}/draft", json={"values": {"report_id": ""}, "version": version})
    version = res.headers["X-Doc-Version"]
    state = client.post(f"/forms/{doc_id}/preview", json={"version": version}).get_json()
    assert state["fields"]["report_id"]["blank"] is True   # 空欄の印は今までどおり

    html = finish(client, [doc_id])["html"]
    assert "必須" not in html and "空欄のまま確定する" not in html
    res = client.post(f"/forms/{doc_id}/confirm", json={"version": version})
    assert res.status_code == 200 and res.get_json()["ok"] is True
    with app.app_context():
        assert db.get_document(doc_id)["state"] == "confirmed"


def _parts_report(path):
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "点検報告書"
    fill = PatternFill("solid", fgColor="FFD9E1F2")
    for coord, value in {"A1": "報告番号", "B1": "IN-2026-001", "A2": "設備番号", "B2": "CMP-101",
                         "A4": "■ 交換部品", "A5": "No.", "B5": "品番", "C5": "品名", "D5": "数量",
                         "A6": 1, "B6": "PW48-1591", "C6": "ベアリング", "D6": 2,
                         "A7": 2, "B7": "PW35-1577", "C7": "スピンモータ", "D7": 1,
                         "A9": "所見", "B9": "異常なし"}.items():
        ws[coord] = value
    for coord in ("A1", "A2", "A5", "B5", "C5", "D5", "A9"):
        ws[coord].fill = fill
    wb.save(path)
    return path


def test_form_flow_with_table_field(app, client, tmp_path):
    """明細表の項目: 列見出しを1回クリック → 取り込み → 読み取り結果で行を直す → Markdown。"""
    path = _parts_report(tmp_path / "点検報告書_CMP-101.xlsx")
    pattern_id = create_type(client, path, "点検報告書")
    add_field(client, pattern_id, "点検報告書", "A1", "B1")
    add_field(client, pattern_id, "点検報告書", "A5")        # 列見出しを1回クリックするだけ
    with app.app_context():
        parts = next(f for f in db.load_pattern(pattern_id).fields if f.data_type == "table")
    assert parts.table_columns == ["No.", "品番", "品名", "数量"]
    panel = panel_html(client, pattern_id, book=path)
    assert "品番: PW35-1577／品名: スピンモータ" in panel
    activate(client, pattern_id)

    doc_id, = upload_forms(client, path)
    html = read_form(client, doc_id, pattern_id, ["点検報告書"]).get_json()["html"]
    assert "data-table-editor" in html and "スピンモータ</textarea>" in html
    assert f'name="value-{parts.field_name}"' in html

    # 画面の値（hidden の JSON）が表示中と同じなら「手で修正」にしない
    summary = client.post(f"/forms/{doc_id}/preview", json={"values": {}}).get_json()
    assert summary["counts"]["manual"] == 0
    assert "## 交換部品\n- 品番: PW48-1591／品名: ベアリング／数量: 2\n- 品番: PW35-1577／品名: スピンモータ／数量: 1\n" \
        in summary["markdown"]
    with app.app_context():
        current = json.loads(db.get_document(doc_id)["data_json"])
    value = next(f for f in current["fields"] if f["field_name"] == parts.field_name)["value"]
    assert client.post(f"/forms/{doc_id}/draft",
                       json={"values": {parts.field_name: json.dumps(value, ensure_ascii=False)}}).status_code == 204
    assert client.post(f"/forms/{doc_id}/preview", json={"values": {}}).get_json()["counts"]["manual"] == 0

    # 読み取った表の列は「No.」（行番号の列）を落とした3列。その形のまま直す
    assert value["columns"] == ["品番", "品名", "数量"]
    edited = {"columns": value["columns"], "rows": [["PW48-1591", "ベアリング", "4"]]}
    client.post(f"/forms/{doc_id}/draft",
                json={"values": {parts.field_name: json.dumps(edited, ensure_ascii=False)}})
    assert client.post(f"/forms/{doc_id}/confirm", json={}).get_json()["ok"] is True
    assert "品番: PW48-1591／品名: ベアリング／数量: 4" in finish(client, [doc_id])["html"]

    # 行を全部消しても列見出しと表の入力欄は残る（人が入れ直せる）。Markdown には行の無い明細表を出さない
    empty = {"columns": value["columns"], "rows": []}
    client.post(f"/forms/{doc_id}/draft",
                json={"values": {parts.field_name: json.dumps(empty, ensure_ascii=False)}})
    with app.app_context():
        saved = json.loads(db.get_document(doc_id)["data_json"])
    kept = next(f for f in saved["fields"] if f["field_name"] == parts.field_name)
    assert kept["value"] == {"columns": value["columns"], "rows": []}   # 列見出しは残る
    assert "## 交換部品" not in client.post(f"/forms/{doc_id}/preview", json={"values": {}}).get_json()["markdown"]

    # ダウンロードは、直したまま確定していない（修正中）ときは直した値から作る（design.md 3.3。
    # 最後の手順。ここでデータは消える）。行を全部消したので明細表は出ない
    md = client.get(f"/forms/{doc_id}/download.md").get_data(as_text=True)
    assert "## 交換部品" not in md and "スピンモータ" not in md
    assert "- 報告番号: IN-2026-001" in md


# ====================================================================================================
# 元 tests/test_forms_md.py
# 帳票の Markdown（docs/design.md 6.1）: タイトル・ファイル名・定型文なし・NFKC・単位・出さない項目・決定性。
# ====================================================================================================

def _pattern_forms_md(infos, **meta):
    sheets, fields = suggest_rows(infos)
    return rows_to_pattern(1, {**META, "title_fields": suggest_title_fields(fields), **meta}, sheets, fields)


def _doc(info, doc_id=1, file_hash="0123456789abcdef"):
    return {"id": doc_id, "file_name": "修理報告書_標準.xlsx" if doc_id else info.path.name, "file_hash": file_hash}


@pytest.fixture
def standard(repair_infos):
    info = repair_infos[0]
    pattern = _pattern_forms_md(repair_infos)
    extraction = extract_document(info, pattern, rank_patterns(info, [pattern])[0].sheet_names)
    return info, pattern, extraction


def _field(extraction, name):
    return next(f for f in extraction["fields"] if f["field_name"] == name)


def test_title_basic_block_and_source(standard):
    info, _, extraction = standard
    md = build_markdown(_doc(info), extraction)
    lines = md.split("\n")
    assert lines[0] == "# 設備修理報告書 R2026-00123｜CMP装置（EQ-001）｜2026-09-14"
    assert "- 帳票の種類: 設備修理報告書" in lines
    assert "- 報告番号: R2026-00123" in lines
    assert "- 設備: CMP装置（EQ-001）" in lines
    assert "- 発生日: 2026-09-14（2026年9月）" in lines
    assert "## 故障内容" in lines  # 短い帳票は見出しに識別子を入れない
    assert "- 添付画像: 1枚" in lines
    assert lines[-2] == "- 出典: 修理報告書_標準.xlsx（報告番号 R2026-00123）"


def test_title_falls_back_to_identifier_keys_when_not_configured(repair_infos):
    info = repair_infos[2]
    pattern = _pattern_forms_md(repair_infos, title_fields=[])
    extraction = extract_document(info, pattern, rank_patterns(info, [pattern])[0].sheet_names)
    md = build_markdown(_doc(info), extraction)
    assert md.startswith("# 設備修理報告書 R2026-00125｜露光装置（EQ-003）｜2026-08-30\n")


def test_configured_title_fields_are_used_in_order(standard):
    info, _, extraction = standard
    extraction["pattern"]["title_fields"] = ["equipment_id", "occurred_date"]
    md = build_markdown(_doc(info), extraction)
    assert md.startswith("# 設備修理報告書 EQ-001｜2026-09-14\n")
    # 識別番号（report_id）がタイトルに入らないので、同名を避けるため file_hash の先頭8桁を足す
    assert markdown_filename(_doc(info), extraction) == "設備修理報告書_EQ-001_2026-09-14_01234567.md"


def test_no_boilerplate_internal_ids_or_cell_coordinates(standard):
    info, _, extraction = standard
    md = build_markdown(_doc(info, doc_id=42), extraction)
    for banned in ("文書ID", "42", "解析していません", "## 原本", "## 基本情報", "テンプレート", "v1", "J3",
                   "修理報告書シート", "A8", "## 添付画像", "原本を参照"):
        assert banned not in md, banned


def test_filename_is_stable_and_independent_of_document_id(standard):
    info, _, extraction = standard
    a = markdown_filename(_doc(info, doc_id=1), extraction)
    b = markdown_filename(_doc(info, doc_id=999, file_hash="ffffffffffffffff"), copy.deepcopy(extraction))
    assert a == b == "設備修理報告書_R2026-00123_EQ-001_CMP装置_2026-09-14.md"


def test_filename_uses_file_hash_when_title_values_are_empty(standard):
    info, _, extraction = standard
    for name in ("report_id", "equipment_id", "equipment_name", "occurred_date"):
        _field(extraction, name)["value"] = None
    refresh_summary(extraction)
    assert markdown_filename(_doc(info), extraction) == "設備修理報告書_01234567.md"


def test_filename_sanitizes_hint_brackets_and_unsafe_chars(standard):
    info, _, extraction = standard
    extraction["pattern"]["name"] = "修理報告書.[legacy-R(chunk_ts=800)]"
    extraction["pattern"]["title_fields"] = ["report_id"]
    _field(extraction, "report_id")["value"] = "R/2026:00123 [改]\t*?"
    name = markdown_filename(_doc(info), extraction)
    assert name.endswith(".md") and name.count(".") == 1
    for ch in ('.[', '[', ']', '/', '\\', ':', '*', '?', '"', '<', '>', '|', ' ', '\t'):
        assert ch not in name, (ch, name)
    assert name.startswith("修理報告書")
    # 全角英数は NFKC で半角に
    _field(extraction, "report_id")["value"] = "Ｒ２０２６－００１２３"
    assert "R2026-00123" in markdown_filename(_doc(info), extraction)


def test_values_are_nfkc_normalized(standard):
    info, _, extraction = standard
    _field(extraction, "equipment_name")["value"] = "ＣＭＰ　　装置"
    _field(extraction, "equipment_id")["value"] = "ＥＱ－００１"
    _field(extraction, "cause")["value"] = "ﾎﾟﾝﾌﾟの  ｺﾈｸﾀ緩み。\n\n\n再締結した。"
    md = build_markdown(_doc(info), extraction)
    assert md.startswith("# 設備修理報告書 R2026-00123｜CMP 装置（EQ-001）｜2026-09-14\n")
    assert "- 設備: CMP 装置（EQ-001）" in md
    assert "## 原因\nポンプの コネクタ緩み。\n再締結した。\n" in md  # レコード内に空行を入れない


def test_units_are_written_with_the_number(standard):
    info, _, extraction = standard
    assert "- 作業時間: 2.5時間" in build_markdown(_doc(info), extraction)
    work = _field(extraction, "work_hours")
    work["unit"], work["value"] = "h", 3.0
    assert "- 作業時間: 3h" in build_markdown(_doc(info), extraction)
    work["value"] = "3時間くらい"  # 数値にできなかった値には単位を足さない
    assert "- 作業時間: 3時間くらい" in build_markdown(_doc(info), extraction)


def test_person_fields_are_written(standard):
    """人名の項目も出す（読み取った内容は削らない。2026-09-20 の利用者の指示）。"""
    info, _, extraction = standard
    assert "- 報告者: 山田 太郎" in build_markdown(_doc(info), extraction)


def test_rag_output_omit_fields_are_not_written(standard):
    info, _, extraction = standard
    _field(extraction, "cause")["rag_output"] = "omit"
    _field(extraction, "work_hours")["rag_output"] = "omit"
    md = build_markdown(_doc(info), extraction)
    assert "## 原因" not in md and "コネクタ接触不良" not in md and "作業時間" not in md
    # 読み取り結果（画面で見る値）には残す
    assert _field(extraction, "cause")["value"] and _field(extraction, "work_hours")["unit"] == "時間"


def test_long_documents_get_identifier_headings(standard):
    info, _, extraction = standard
    _field(extraction, "symptom")["value"] = "搬送アームが停止した。\n" * 120
    md = build_markdown(_doc(info), extraction)
    assert "## 故障内容（R2026-00123／EQ-001 CMP装置／2026-09-14）" in md
    assert "## 原因（R2026-00123／EQ-001 CMP装置／2026-09-14）" in md


def test_markdown_syntax_in_values_is_escaped(standard):
    info, _, extraction = standard
    _field(extraction, "cause")["value"] = "# 見出しではない\n---\n> 引用ではない"
    md = build_markdown(_doc(info), extraction)
    assert "\n# 見出し" not in md and "\n---\n" not in md and "\n> 引用" not in md
    assert md.count("\n# ") == 0 and md.startswith("# ")


def test_output_is_deterministic_bytes(repair_infos):
    digests = set()
    for _ in range(2):
        info = repair_infos[1]
        pattern = _pattern_forms_md(repair_infos)
        extraction = extract_document(info, pattern, rank_patterns(info, [pattern])[0].sheet_names)
        md = build_markdown(_doc(info, doc_id=None), extraction)
        digests.add(hashlib.sha256(md.encode("utf-8")).hexdigest())
        assert "\r" not in md and md.endswith("\n") and not md.endswith("\n\n") and "\n\n\n" not in md
        assert md.startswith("# 設備修理報告書 R2026-00124｜CVD装置（EQ-002）｜2026-09-10\n")
        assert "- 添付画像: 2枚" in md
    assert len(digests) == 1


# ---- 明細表 ----

def _table_field(value, name="parts", display="交換部品", **extra):
    return {"field_name": name, "display_name": display, "data_type": "table", "value": value,
            "unit": "", "rag_output": "show", "edited": False, "warning": None, **extra}


def test_table_field_is_written_one_line_per_row(standard):
    info, _, extraction = standard
    value = {"columns": ["品番", "品名", "数量", "備考"],
             "rows": [["ＰＷ４８－１５９１", "ベアリング\n（軸受）", "2", ""], ["", "", "", ""],
                      ["# 見出しではない", "スピンモータ", "1", "予備品から"], ["部品費計", "", "3", ""]]}
    extraction["fields"].insert(5, _table_field(value))
    refresh_summary(extraction)
    md = build_markdown(_doc(info), extraction)
    assert ("\n## 交換部品\n"
            "- 品番: PW48-1591／品名: ベアリング (軸受)／数量: 2\n"
            "- 品番: # 見出しではない／品名: スピンモータ／数量: 1／備考: 予備品から\n"
            "- 部品費計: 数量: 3\n\n") in md
    assert "|" not in md  # パイプ表は使わない
    assert "- 交換部品" not in md  # 基本の箇条書きには入れない

    _field(extraction, "parts")["value"] = None
    refresh_summary(extraction)
    md = build_markdown(_doc(info), extraction)
    assert "## 交換部品" not in md


def test_table_field_is_not_used_as_title(standard):
    info, _, extraction = standard
    extraction["fields"].append(_table_field({"columns": ["品番"], "rows": [["PW-1"]]}))
    extraction["pattern"]["title_fields"] = ["parts", "report_id"]
    assert build_markdown(_doc(info), extraction).startswith("# 設備修理報告書 R2026-00123\n")
    assert markdown_filename(_doc(info), extraction) == "設備修理報告書_R2026-00123.md"


# ---- LightRAG オフライン評価の反映（見出し語が値になった行・ファイル名の一意性・タイトルの手がかり） ----

def test_label_as_value_lines_are_not_written(standard):
    """読み取り誤りで値が別の欄の見出し語になった項目は md・タイトル・ファイル名に出さない（画面には残す）。"""
    info, _, extraction = standard
    labels = set(extraction["pattern"]["labels"])
    assert "報告番号" in labels and "発生日" in labels  # 候補ラベル・表示名から作られている

    quantity = {"field_name": "quantity", "display_name": "数量", "data_type": "string",
                "value": "発生日", "unit": "", "rag_output": "show", "edited": False, "warning": None}
    extraction["fields"].append(quantity)
    extraction["fields"].append({**quantity, "field_name": "part_name", "display_name": "品名", "value": "報告番号"})
    refresh_summary(extraction)
    md = build_markdown(_doc(info), extraction)
    assert "- 数量: 発生日" not in md and "- 品名: 報告番号" not in md
    assert extraction["values"]["quantity"] == "発生日"  # 読み取り結果には残す

    # 人が直した値は消さない。ラベルでない値はそのまま出す
    quantity["edited"] = True
    assert "- 数量: 発生日" in build_markdown(_doc(info), extraction)
    quantity["edited"], quantity["value"] = False, "3個"
    assert "- 数量: 3個" in build_markdown(_doc(info), extraction)


def test_filename_gets_file_hash_when_the_title_has_no_report_id(standard):
    """報告番号がタイトルに入らない帳票は同名になりやすい（LightRAG 1.5.x は同名だと HTTP 409）。"""
    info, _, extraction = standard
    assert markdown_filename(_doc(info), extraction) == "設備修理報告書_R2026-00123_EQ-001_CMP装置_2026-09-14.md"
    extraction["pattern"]["title_fields"] = ["equipment_id", "occurred_date"]
    assert markdown_filename(_doc(info), extraction) == "設備修理報告書_EQ-001_2026-09-14_01234567.md"
    extraction["pattern"]["title_fields"] = ["occurred_date"]
    assert markdown_filename(_doc(info), extraction) == "設備修理報告書_2026-09-14_01234567.md"


def test_title_adds_the_source_file_name_when_it_has_no_identifier(standard):
    """識別番号も設備も入らないタイトル（工程異常連絡票のような様式）は、元ファイル名で帳票を特定できるようにする。"""
    info, _, extraction = standard
    extraction["pattern"]["title_fields"] = ["occurred_date"]
    assert build_markdown(_doc(info), extraction).startswith("# 設備修理報告書 2026-09-14｜修理報告書_標準\n")
    extraction["pattern"]["title_fields"] = ["equipment_id", "occurred_date"]
    assert build_markdown(_doc(info), extraction).startswith("# 設備修理報告書 EQ-001｜2026-09-14\n")


def test_person_columns_are_written_and_empty_total_rows_are_not():
    """明細表の人名の列（担当・氏名）も出す。数字のない合計行は記録にならないので書かない（design.md 6.1）。"""
    from forms import table_markdown_lines
    from forms import is_person_field, is_person_label

    value = {"columns": ["日時", "対応内容", "担当"],
             "rows": [["9:10", "電極を交換", "中村"], ["合計", "", ""]]}
    assert table_markdown_lines(value) == ["- 日時: 9:10／対応内容: 電極を交換／担当: 中村"]
    # 数字のある合計行はこれまでどおり「- 合計: …」で書く
    assert table_markdown_lines({"columns": ["ロットNo.", "投入数"], "rows": [["合計", "50"]]}) == ["- 合計: 投入数: 50"]

    # 押印欄の「確認」「作成」も人名の項目。「効果確認」「作成日」は違う
    assert is_person_field("field_3", "確認") and is_person_field("field_4", "作成")
    assert not is_person_field("field_5", "効果確認") and not is_person_field("field_6", "作成日")
    assert is_person_label("担当") and is_person_label("氏名") and not is_person_label("確認")  # 表の「確認」は判定の列


# ---- 利用者の判断（2026-09-19）: 丸数字を残す・積み重なった列見出しをすべて出す ----

def test_enclosed_numbers_are_kept_as_written(standard):
    """丸数字（①②）は NFKC で囲みを外さない。「①破損…」が「1破損…」になると番号と本文の区切りが消える。"""
    info, _, extraction = standard
    _field(extraction, "repair")["value"] = "①破損ウェーハ片を回収\n②Ｈｅａｄ３ メンブレン交換"
    _field(extraction, "cause")["value"] = "㈱テスト製 ⑴ ﾎﾟﾝﾌﾟの劣化"
    md = build_markdown(_doc(info), extraction)
    assert "## 修理内容\n①破損ウェーハ片を回収\n②Head3 メンブレン交換\n" in md
    # 囲みが外れても区切りが残る表記（㈱・⑴）は、これまでどおり NFKC でそろえる
    assert "## 原因\n(株)テスト製 (1) ポンプの劣化\n" in md
    assert "\\" not in md.split("## 修理内容\n")[1].split("\n")[0]  # 行頭の①はエスケープしない


def test_stacked_table_rows_keep_their_own_column_headings(standard):
    """積み重なった列見出しをまとめた明細表は、行ごとに自分の組の列見出しだけを書く（空の欄は書かない）。"""
    info, _, extraction = standard
    value = {"columns": ["人", "機械", "材料", "方法", "測定", "環境"],
             "rows": [["①日常点検での見落とし", "軸受の摩耗", "", "", "", ""],
                      ["", "", "", "点検手順に記載なし", "", "室温の変動"]]}
    extraction["fields"].insert(5, _table_field(value, name="fishbone", display="特性要因（4M+2）"))
    refresh_summary(extraction)
    md = build_markdown(_doc(info), extraction)
    assert ("\n## 特性要因(4M+2)\n"  # 見出しの全角かっこは今までどおり NFKC で半角に
            "- 人: ①日常点検での見落とし／機械: 軸受の摩耗\n"
            "- 方法: 点検手順に記載なし／環境: 室温の変動\n") in md
    assert "|" not in md


# ---- R6-F1: 手で入れた明細表の連番「No」列 ----

def test_a_hand_entered_row_drops_the_sequence_no_column(standard):
    """読み取れなかった明細表に手で入れた行も、読み取った表と同じ形で書く（「- No.: 1」を出さない）。

    読み取った表は Table.to_value で連番の No 列を落とすが、読めなかった表の入力欄には
    帳票の種類の列見出し（No を含む）がそのまま出るので、出すときに形をそろえる。
    """
    info, _, extraction = standard
    value = {"columns": ["No.", "品名", "品番", "数量"],
             "rows": [["1", "走行部", "目視", "-"], ["2", "駆動部", "PW-1", "1"], ["合計", "", "", "2"]]}
    extraction["fields"].insert(5, _table_field(value))
    refresh_summary(extraction)
    md = build_markdown(_doc(info), extraction)
    assert ("\n## 交換部品\n"
            "- 品名: 走行部／品番: 目視／数量: -\n"
            "- 品名: 駆動部／品番: PW-1／数量: 1\n"
            "- 合計: 数量: 2\n") in md
    assert "No.:" not in md


def test_a_no_column_that_is_not_a_sequence_is_kept(standard):
    """連番でない「No」列（本当の項番・品番）は、読み取り側と同じく残す。"""
    info, _, extraction = standard
    value = {"columns": ["No.", "品名"], "rows": [["12", "走行部"], ["7", "駆動部"]]}
    extraction["fields"].insert(5, _table_field(value))
    refresh_summary(extraction)
    md = build_markdown(_doc(info), extraction)
    assert "- No.: 12／品名: 走行部\n- No.: 7／品名: 駆動部\n" in md


# ---- R6-MD-02: 長い明細表の断片にも識別子を入れる ----

def test_a_long_detail_table_is_split_so_every_part_carries_the_identifier(standard):
    """明細表1節が固定窓（1,200トークン）を超えると、識別番号も設備名も無い断片ができる。

    識別子を入れる帳票（長い帳票）では、明細表を「（続き）」の見出しで分ける。
    """
    from forms import _estimate_tokens

    info, _, extraction = standard
    rows = [[f"部位{i}", "目視", "○", "異常なし。次回も同じ手順で確認すること"] for i in range(1, 41)]
    extraction["fields"].insert(5, _table_field({"columns": ["点検部位", "方法", "判定", "所見"], "rows": rows},
                                                name="checks", display="点検結果"))
    refresh_summary(extraction)
    md = build_markdown(_doc(info), extraction)

    identifier = "R2026-00123／EQ-001 CMP装置／2026-09-14"
    heads = [ln for ln in md.split("\n") if ln.startswith("## ")]
    assert len(heads) > 1 and all(identifier in h for h in heads)
    assert f"## 点検結果（続き）（{identifier}）" in md
    # 見出しから次の見出しまでは固定窓に収まる（＝どの断片にも識別子付きの見出しが入る）
    assert all(_estimate_tokens(part) < 1200 for part in md.split("\n## ")[1:])
    # 行は1つも落ちず、順番も変わらない
    assert all(f"- 点検部位: 部位{i}／" in md for i in range(1, 41))
    assert md.index("部位1／") < md.index("部位40／") and "|" not in md


# ====================================================================================================
# 元 tests/test_forms_sections.py
# 帳票の読み取り: 区画（発行側と回答側）・日付の右の時刻・見出しと列見出しの間に1行ある明細表 のテスト。
# ====================================================================================================

HEAD_FILL = "FFFCE4D6"


def _pattern_forms_sections(*fields: FieldDef) -> PatternDef:
    return PatternDef(name="テスト", fields=list(fields))


def _renraku(tmp_path, name: str, answered: bool = True, answer_heading: str = "▼ 回答欄（宛先部署にて記入し返却）"):
    """「■ 異常連絡」の下（発行側）に「発生原因（推定）」「処置内容」、下の回答欄に「原因（確定）」「暫定対策（処置）」がある連絡票。"""
    cells = {"A1": "連絡No.", "B1": "PA-001",
             "A2": "■ 異常連絡",
             "A3": "発生原因（推定）", "B3": "研磨パッドの異常と思われる",
             "A4": "処置内容", "B4": "該当ロットをHOLD",
             "A5": answer_heading,
             "A6": "原因（確定）", "A7": "暫定対策（処置）"}
    if answered:
        cells.update({"B6": "パッド溝の摩耗", "B7": "研磨パッド交換"})
    fills = {c: LABEL_FILL for c in ("A1", "A3", "A4", "A6", "A7")}
    fills["A2"] = fills["A5"] = HEAD_FILL
    return _book(tmp_path, name, {"連絡票": (cells, fills, ["A2:B2", "A5:B5"])})


def _side_by_side(tmp_path, name: str):
    """左が発行側、右上の「【回答欄】」の下が回答側（A3横の版）。"""
    cells = {"A1": "工程異常連絡票", "D1": "【回答欄】",
             "A3": "連絡No.", "B3": "工異-001", "D3": "真因", "E3": "減速機の損傷",
             "D4": "応急処置", "E4": "減速機交換",
             "A6": "推定原因", "B6": "衝突の可能性", "A7": "処置", "B7": "全数選別"}
    fills = {c: LABEL_FILL for c in ("A3", "D3", "D4", "A6", "A7")}
    info = _book(tmp_path, name, {"連絡票": (cells, fills, [])})
    return info


# ---- 区画 ---------------------------------------------------------------------------------

def test_sections_split_the_sheet_by_marked_headings(tmp_path):
    info = _renraku(tmp_path, "a.xlsx")
    grid = info.grids["連絡票"]
    assert [s.key for s in sections(grid)] == ["異常連絡", "回答"]  # 印・括弧書き・末尾の「欄」を除いた名前
    assert section_of(grid, grid.cells[(1, 1)]) == ""          # 見出しより上
    assert section_of(grid, grid.cells[(4, 1)]) == "異常連絡"  # 処置内容（発行側）
    assert section_of(grid, grid.cells[(7, 1)]) == "回答"      # 暫定対策（処置）


def test_a_heading_in_the_right_half_makes_a_column_section(tmp_path, monkeypatch):
    from openpyxl import load_workbook
    from openpyxl.styles import Font

    info = _side_by_side(tmp_path, "b.xlsx")
    wb = load_workbook(info.path)
    wb["連絡票"]["D1"].font = Font(bold=True)  # 「【回答欄】」は太字だけ（塗りつぶしなし）
    wb.save(info.path)
    from forms import load_workbook_info

    grid = load_workbook_info(info.path).grids["連絡票"]
    assert section_of(grid, grid.cells[(4, 4)]) == "回答"  # 右側の応急処置
    assert section_of(grid, grid.cells[(7, 1)]) == ""      # 左側の処置


def test_field_with_a_section_reads_the_answer_side(tmp_path):
    info = _renraku(tmp_path, "a.xlsx")
    repair = FieldDef("repair", "処置", ["処置内容", "暫定対策", "応急処置"], data_type="text")
    cause = FieldDef("cause", "原因", ["原因（確定）", "発生原因", "推定原因"], data_type="text")
    values, _ = _values(info, _pattern_forms_sections(repair, cause))
    assert values == {"repair": "該当ロットをHOLD", "cause": "パッド溝の摩耗"}  # 区画なし: 上にある欄（今までどおり）

    repair.section = cause.section = "回答"
    values, fields = _values(info, _pattern_forms_sections(repair, cause))
    assert values == {"repair": "研磨パッド交換", "cause": "パッド溝の摩耗"}
    assert fields["repair"]["label_cell"] == "A7"


def test_unanswered_form_does_not_take_the_issuing_side_value(tmp_path):
    info = _renraku(tmp_path, "a.xlsx", answered=False)
    repair = FieldDef("repair", "処置", ["処置内容", "暫定対策"], data_type="text", section="回答")
    cause = FieldDef("cause", "原因", ["原因（確定）", "発生原因"], data_type="text", section="回答")
    values, fields = _values(info, _pattern_forms_sections(repair, cause))
    assert values == {"repair": None, "cause": None}  # 発行側の「処置内容」「発生原因（推定）」を読まない
    assert fields["cause"]["label_found"] and fields["cause"]["label_cell"] == "A6"


def test_layout_without_the_section_is_read_as_before(tmp_path):
    cells = {"A1": "処置内容", "B1": "該当ロットをHOLD"}
    info = _book(tmp_path, "c.xlsx", {"連絡票": (cells, {"A1": LABEL_FILL}, [])})
    fd = FieldDef("repair", "処置", ["処置内容"], data_type="text", section="回答")
    values, _ = _values(info, _pattern_forms_sections(fd))
    assert values["repair"] == "該当ロットをHOLD"


def test_builder_learns_the_answer_section_from_samples(tmp_path):
    both = _renraku(tmp_path, "a.xlsx")
    side = _side_by_side(tmp_path, "b.xlsx")
    from openpyxl import load_workbook
    from openpyxl.styles import Font

    wb = load_workbook(side.path)
    wb["連絡票"]["D1"].font = Font(bold=True)
    wb.save(side.path)
    from forms import load_workbook_info

    _, rows = suggest_rows([both, load_workbook_info(side.path)])
    by_name = {r["field_name"]: r for r in rows}
    assert by_name["repair"]["use"] and by_name["repair"]["section"] == "回答"
    assert by_name["cause"]["use"] and by_name["cause"]["section"] == "回答"
    assert by_name["report_id"]["section"] == ""  # 1か所にしかない項目には区画を付けない

    pattern = rows_to_pattern(1, {"name": "連絡票"}, [{"use": True, "sheet_name": "連絡票", "required": True}], rows)
    unanswered = _renraku(tmp_path, "d.xlsx", answered=False)
    values = extract_document(unanswered, pattern, ["連絡票"])["values"]
    assert values["repair"] is None and values["cause"] is None


def test_builder_leaves_the_section_empty_when_both_sides_are_common(tmp_path):
    """押印欄の「確認」のように、どの見本でも区画の外と回答欄の両方にある項目は、区画を決めない
    （どちら側の欄かを見本から決められない。今までどおり上にある欄を読み、人が画面で決める）。"""
    def book(name):
        cells = {"A1": "確認", "A2": "長谷川", "C1": "【回答欄】", "C2": "確認", "C3": "小川"}
        return _book(tmp_path, name, {"連絡票": (cells, {"A1": LABEL_FILL, "C1": HEAD_FILL, "C2": LABEL_FILL}, [])})

    _, rows = suggest_rows([book("a.xlsx"), book("b.xlsx")])
    assert all(r["section"] == "" for r in rows)


def test_section_round_trips_through_the_rows():
    fd = FieldDef("repair", "処置", ["暫定対策"], data_type="text", section="回答")
    _, rows = pattern_to_rows(_pattern_forms_sections(fd))
    assert rows[0]["section"] == "回答"
    # 見出しのとおり入力しても比較用の名前にそろえる
    rows[0]["section"] = "▼ 回答欄（宛先記入）"
    assert rows_to_pattern(1, {"name": "連絡票"}, [], rows).fields[0].section == "回答"


def test_section_with_two_suffixes_is_stable_through_learning_and_saving(tmp_path):
    """「■ 処置内容欄」のように末尾の語が重なる見出し: 何度そろえても同じ名前になり、区画の中の値を読む。"""
    from forms import section_name

    for text in ("■ 処置内容欄", "■ 回答内容欄（記入）", "■ 作業日時欄"):
        assert section_name(section_name(text)) == section_name(text)

    def book(name, rows):
        cells, fills = {"A1": "報告番号", "B1": name}, {"A1": LABEL_FILL}
        for r, (label, value) in enumerate(rows, start=3):
            cells[f"A{r}"] = label
            if value is not None:
                cells[f"B{r}"] = value
            fills[f"A{r}"] = HEAD_FILL if value is None else LABEL_FILL
        return _book(tmp_path, name + ".xlsx", {"報告書": (cells, fills, [])})

    a = book("a", [("■ 依頼内容欄", None), ("担当者", "田中"), ("依頼事項", "ポンプ異音"),
                   ("■ 処置内容欄", None), ("担当者", "鈴木"), ("処置", "ベアリング交換")])
    b = book("b", [("■ 処置内容欄", None), ("担当者", "佐藤"), ("処置", "清掃")])
    sheets, rows = suggest_rows([a, b])
    pattern = rows_to_pattern(1, {"name": "報告書"}, sheets, rows)
    reporter = next(f for f in pattern.fields if f.display_name == "担当者")
    assert reporter.section == section_name("■ 処置内容欄")
    assert extract_document(a, pattern, ["報告書"])["values"][reporter.field_name] == "鈴木"

    # 画面から変更せずに保存し直しても区画は変わらない
    meta, sheet_rows = {"name": "報告書"}, [{"use": True, "sheet_name": "報告書", "required": True}]
    _, again = pattern_to_rows(pattern)
    assert rows_to_pattern(1, meta, sheet_rows, again).fields == pattern.fields


# ---- ラベルの下の値: 左の見出しの値の欄を取らない ----------------------------------------------

def test_value_below_does_not_take_the_wide_value_of_the_label_on_the_left(tmp_path):
    # 品証コメント（A2）の値の欄（B2:D3）が、空のクローズ判定（C1）の真下まで広がっている
    cells = {"C1": "クローズ判定", "A2": "品証コメント", "B2": "効果確認未了。7/8頃に再確認"}
    info = _book(tmp_path, "e.xlsx", {"連絡票": (cells, {"C1": LABEL_FILL, "A2": LABEL_FILL}, ["B2:D3"])})
    values, _ = _values(info, _pattern_forms_sections(FieldDef("close", "クローズ判定", ["クローズ判定"]),
                                       FieldDef("qa", "品証コメント", ["品証コメント"], data_type="text")))
    assert values["close"] is None
    assert values["qa"] == "効果確認未了。7/8頃に再確認"


# ---- 日付の右のセルの時刻（md3-3）------------------------------------------------------------

def _date_time_book(tmp_path, name: str, cells: dict):
    base = {"A1": "発生日時", "E1": "作成者"}
    base.update(cells)
    return _book(tmp_path, name, {"報告書": (base, {"A1": LABEL_FILL, "E1": LABEL_FILL}, [])})


def test_time_in_the_next_cell_is_joined_to_the_date(tmp_path):
    info = _date_time_book(tmp_path, "t1.xlsx", {"B1": date(2023, 5, 23), "C1": "12:07"})
    values, fields = _values(info, _pattern_forms_sections(FieldDef("occurred", "発生日時", ["発生日時"], data_type="date")))
    assert values["occurred"] == "2023-05-23 12:07"
    assert fields["occurred"]["value_cell"] == "B1:C1" and not fields["occurred"]["warning"]

    info = _date_time_book(tmp_path, "t2.xlsx", {"B1": "2023/5/23", "C1": time(9, 5)})
    values, _ = _values(info, _pattern_forms_sections(FieldDef("occurred", "発生日時", ["発生日時"], data_type="date")))
    assert values["occurred"] == "2023-05-23 09:05"

    info = _date_time_book(tmp_path, "t3.xlsx", {"B1": "2023/5/23", "C1": "9時5分"})
    values, _ = _values(info, _pattern_forms_sections(FieldDef("occurred", "発生日時", ["発生日時"], data_type="date")))
    assert values["occurred"] == "2023-05-23 09:05"


def test_time_is_not_joined_when_it_is_ambiguous(tmp_path):
    fd = FieldDef("occurred", "発生日時", ["発生日時"], data_type="date")
    cases = {
        "range.xlsx": {"B1": "2023/5/23", "C1": "9:00", "D1": "17:00"},    # 時刻が2つ（範囲）
        "gap.xlsx": {"B1": "2023/5/23", "D1": "12:07"},                    # すぐ右ではない
        "clock.xlsx": {"B1": "2023/5/23 10:00", "C1": "12:07"},            # 日付に時刻が書いてある
        "about.xlsx": {"B1": "2023/5/23", "C1": "12:07頃"},                # 時刻だけではない
    }
    expected = {"range.xlsx": "2023-05-23", "gap.xlsx": "2023-05-23", "clock.xlsx": "2023-05-23 10:00",
                "about.xlsx": "2023-05-23"}
    for name, cells in cases.items():
        values, _ = _values(_date_time_book(tmp_path, name, cells), _pattern_forms_sections(fd))
        assert values["occurred"] == expected[name], name
    # 日付の項目でなければつながない
    values, _ = _values(_date_time_book(tmp_path, "s.xlsx", {"B1": "2023/5/23", "C1": "12:07"}),
                        _pattern_forms_sections(FieldDef("occurred", "発生日時", ["発生日時"])))
    assert values["occurred"] == "2023/5/23"


# ---- 見出しと列見出しの間に「項目｜値」の1行がある明細表 ----------------------------------------

def _yokoten(tmp_path, name: str, heading: str = "■ 水平展開"):
    cells = {"A1": heading, "A2": "展開区分", "B2": "☑同型機　□類似設備",
             "A3": "確認", "B3": "設備No", "C3": "設備名", "D3": "結果",
             "A4": "☑", "B4": "CMP-103", "C4": "W-CMP 3号機", "D4": "異常なし",
             "A5": "□", "B5": "CMP-104", "C5": "Cu-CMP 4号機",
             "A6": "展開先・内容", "B6": "同型機を点検"}
    fills = {c: LABEL_FILL for c in ("A2", "A3", "B3", "C3", "D3", "A6")}
    fills["A1"] = HEAD_FILL
    return _book(tmp_path, name, {"報告書": (cells, fills, ["A1:D1", "B2:D2", "B6:D6"])})


def test_table_under_a_heading_with_a_label_row_between_gets_the_heading(tmp_path):
    info = _yokoten(tmp_path, "y.xlsx")
    grid = info.grids["報告書"]
    [table] = [t for t in detect_tables(grid) if t.header[0].text == "確認"]
    assert table.anchor is not None and table.anchor.text == "■ 水平展開"
    assert table.to_value()["rows"] == [["☑", "CMP-103", "W-CMP 3号機", "異常なし"], ["□", "CMP-104", "Cu-CMP 4号機", ""]]
    # 見出しから探しても同じ表（「展開区分」の行を飛ばす）
    found = find_table(grid, grid.cells[(1, 1)])
    assert found is not None and [h.text for h in found.header] == ["確認", "設備No", "設備名", "結果"]
    values, _ = _values(info, _pattern_forms_sections(FieldDef("targets", "水平展開", ["■ 水平展開"], data_type="table"),
                                       FieldDef("kubun", "展開区分", ["展開区分"])))
    assert len(values["targets"]["rows"]) == 2 and values["kubun"] == "同型機"


def test_builder_suggests_the_table_under_a_heading_with_a_label_row_between(tmp_path):
    _, rows = suggest_rows([_yokoten(tmp_path, "y1.xlsx"), _yokoten(tmp_path, "y2.xlsx", "７．水平展開")])
    [row] = [r for r in rows if r["data_type"] == "table"]
    assert row["use"] and row["display_name"] == "水平展開"
    assert set(row["candidates"].splitlines()) >= {"■ 水平展開", "７．水平展開"}


def test_label_row_is_skipped_only_under_a_marked_heading(tmp_path):
    """見出しに「■」「1.」の印が無ければ、間の1行を飛ばして見出しにしない（今までどおり見出しの無い表）。"""
    info = _yokoten(tmp_path, "y.xlsx", heading="水平展開")
    grid = info.grids["報告書"]
    [table] = [t for t in detect_tables(grid) if t.header[0].text == "確認"]
    assert table.anchor is None


# ====================================================================================================
# 元 tests/test_forms_batch_review.py
# まとめ取り込み: 同じフォームの帳票を一度に読み取り、読み取り結果を縦に全部並べる。
#
# 利用者の指示（2026-09-20）:
#   「帳票取り込みですが、同じフォームのExcelをまとめて一気に処理できるようにしたい。
#     ③の読み取り結果の表示を②タブ選択で切り替えるのではなく、全部のシートをまとめて表示する
#     （スクロールして全部確認していく）イメージで」
# ====================================================================================================

def _pattern_forms_batch_review(client, sample_dir) -> int:
    pattern_id = create_type(client, sample_dir / "standard.xlsx", "設備修理報告書")
    add_field(client, pattern_id, "修理報告書", "A3", "B3")     # 報告番号
    add_field(client, pattern_id, "修理報告書", "E4", "F4")     # 設備名
    activate(client, pattern_id)
    return pattern_id


def _copies(sample_dir, tmp_path, count: int) -> list:
    """同じフォームの帳票を名前だけ変えて count 件（まとめて置くファイル）。"""
    data = (sample_dir / "standard.xlsx").read_bytes()
    paths = []
    for i in range(1, count + 1):
        path = tmp_path / f"修理報告書_{i}.xlsx"
        path.write_bytes(data)
        paths.append(path)
    return paths


def read_all(client, ids, pattern_id, sheets, **extra):
    return client.post("/forms/read", data={"ids": ",".join(str(i) for i in ids),
                                            "pattern_id": pattern_id, "sheets": sheets, **extra})


# ---- ② 帳票の種類とシートは、置かれた分すべてで1回だけ決める ------------------------------------

def test_the_type_and_sheets_are_decided_once_for_the_whole_batch(app, client, sample_dir, tmp_path):
    pattern_id = _pattern_forms_batch_review(client, sample_dir)
    ids = upload_forms(client, *_copies(sample_dir, tmp_path, 3))

    body = client.get("/forms/type?ids=" + ",".join(str(i) for i in ids)).get_json()
    html = body["html"]
    assert body["has_types"] and [d["id"] for d in body["docs"]] == ids
    # 1組の選択（種類のラジオとシートのチェック）が全部のファイルに効く。タブで1件ずつ選ばない
    assert html.count("data-type-form") == 1 and "data-doc-tabs" not in html
    assert f'value="{pattern_id}" checked' in html and html.count('name="sheets"') == 3
    assert f'name="ids" value="{",".join(str(i) for i in ids)}"' in html
    assert "置かれた<strong>3件</strong>すべてに" in html and "3件をまとめて読み取る" in html
    # 合わないファイルを外せるよう、ファイルごとに見つかった項目数と ✕（外す）を出す
    for doc_id in ids:
        assert f'data-file-count="{doc_id}"' in html and f'data-drop-doc="{doc_id}"' in html


# ---- ③ 読み取り結果は全部の帳票を縦に並べる ---------------------------------------------------

def test_three_forms_are_read_and_confirmed_from_the_stacked_view(app, client, sample_dir, tmp_path):
    pattern_id = _pattern_forms_batch_review(client, sample_dir)
    ids = upload_forms(client, *_copies(sample_dir, tmp_path, 3))

    body = read_all(client, ids, pattern_id, ["修理報告書"]).get_json()
    html = body["html"]
    assert [d["id"] for d in body["docs"]] == ids
    # 3件とも塊として並ぶ（見出しにファイル名と状態）
    assert html.count('class="review-doc"') == 3
    for no, doc_id in enumerate(ids, start=1):
        assert f'id="doc-{doc_id}"' in html and f'id="reviewForm-{doc_id}"' in html
        assert f"修理報告書_{no}.xlsx" in html
    assert html.count("要確認 <strong") == 3 and "未確定" in html
    # 進み具合の1行（まとめ取り込みでは1件ずつ確定しないので、いま見ている場所を出す）
    assert "3件中1件目を表示中" in html and "data-next-doc" in html

    # 3件とも確定でき、zip には3件ぶんの .md が入る
    for doc_id in ids:
        assert client.post(f"/forms/{doc_id}/confirm", json={}).get_json()["ok"] is True
    state = finish(client, ids)
    assert state["confirmed"] == 3 and state["total"] == 3
    assert "確定済み <strong>3</strong> / 3 件" in state["html"]
    assert "まとめて Markdown をダウンロード（zip）" in state["html"]

    with app.app_context():
        batch_id = db.get_document(ids[0])["batch_id"]
    res = client.get(f"/forms/batches/{batch_id}/download.zip")
    assert res.status_code == 200
    with zipfile.ZipFile(io.BytesIO(res.data)) as zf:
        names = zf.namelist()
    assert len(names) == 3 and all(n.endswith(".md") for n in names)
    with app.app_context():
        assert all(db.get_document(i) is None for i in ids)    # 渡したら消える


def test_the_progress_line_counts_the_confirmed_forms(app, client, sample_dir, tmp_path):
    pattern_id = _pattern_forms_batch_review(client, sample_dir)
    ids = upload_forms(client, *_copies(sample_dir, tmp_path, 3))
    read_all(client, ids, pattern_id, ["修理報告書"])

    client.post(f"/forms/{ids[0]}/confirm", json={})
    html = read_all(client, ids, pattern_id, ["修理報告書"], acknowledge="on").get_json()["html"]
    assert "3件中1件目を表示中" in html
    assert html.count('class="badge badge-green" data-doc-state>確定済み') == 1
    assert finish(client, ids)["confirmed"] == 1


def test_an_edit_in_the_second_block_is_saved_to_that_document(app, client, sample_dir, tmp_path):
    """2つ目の塊で直した値は、その帳票だけに入る（ほかの帳票は触らない）。"""
    pattern_id = _pattern_forms_batch_review(client, sample_dir)
    ids = upload_forms(client, *_copies(sample_dir, tmp_path, 3))
    body = read_all(client, ids, pattern_id, ["修理報告書"]).get_json()
    second = body["docs"][1]

    res = client.post(f"/forms/{second['id']}/draft",
                      json={"values": {"equipment_name": "CMP研磨装置（2号機）"}, "version": second["version"]})
    assert res.status_code == 204
    with app.app_context():
        saved = [json.loads(db.get_document(i)["data_json"]) for i in ids]
    values = [s["values"]["equipment_name"] for s in saved]
    assert values[1] == "CMP研磨装置（2号機）"
    assert values[0] != "CMP研磨装置（2号機）" and values[2] != "CMP研磨装置（2号機）"
    assert saved[1]["fields"][1]["edited"] and not saved[0]["fields"][1]["edited"]


def test_only_the_first_sheet_grid_comes_with_the_read(app, client, sample_dir, tmp_path):
    """12件置いても重くならないよう、元のシートの表は先頭だけ作り、残りは後から読み込む。"""
    pattern_id = _pattern_forms_batch_review(client, sample_dir)
    ids = upload_forms(client, *_copies(sample_dir, tmp_path, 3))
    html = read_all(client, ids, pattern_id, ["修理報告書"]).get_json()["html"]

    assert html.count('data-cell="B3"') == 1          # 表は先頭の帳票だけ
    assert html.count("data-grid-wait") == 2          # 残りは画面に入ってから
    for doc_id in ids:
        assert f'data-grid-url="/forms/{doc_id}/grid"' in html

    grid = client.get(f"/forms/{ids[1]}/grid").get_json()
    assert grid["doc_id"] == ids[1] and 'data-cell="B3"' in grid["html"]
    assert client.get(f"/forms/{ids[1]}/grid", headers={}).status_code == 200


def test_a_file_that_does_not_match_can_be_left_out(app, client, sample_dir, tmp_path):
    """合わないファイルは外して、残りをそのまままとめて読み取れる。"""
    pattern_id = _pattern_forms_batch_review(client, sample_dir)
    paths = _copies(sample_dir, tmp_path, 3)
    ids = upload_forms(client, *paths)

    assert client.post(f"/forms/{ids[1]}/delete").get_json()["ok"] is True
    rest = [ids[0], ids[2]]
    body = client.get("/forms/type?ids=" + ",".join(str(i) for i in ids)).get_json()
    assert [d["id"] for d in body["docs"]] == rest    # 外したファイルはもう出ない

    html = read_all(client, rest, pattern_id, ["修理報告書"]).get_json()["html"]
    assert html.count('class="review-doc"') == 2 and "2件中1件目を表示中" in html


# ---- 1件だけのときは今までどおり（まとまりの1行も出さない） ---------------------------------------

def test_a_single_file_review_has_one_block_and_no_batch_line(app, client, sample_dir):
    pattern_id = _pattern_forms_batch_review(client, sample_dir)
    doc_id, = upload_forms(client, sample_dir / "standard.xlsx")

    body = client.get(f"/forms/{doc_id}/type").get_json()
    assert "2項目中2項目が見つかりました" in body["html"] and "置かれた<strong>" not in body["html"]
    assert "この帳票はやめる" in body["html"]

    body = read_all(client, [doc_id], pattern_id, ["修理報告書"]).get_json()
    html = body["html"]
    assert html.count('class="review-doc"') == 1
    assert "data-batch-bar" not in html and "件目を表示中" not in html
    assert 'data-cell="B3"' in html and "data-grid-wait" not in html   # 1件のときは元のシートもすぐ出す
    assert 'name="value-report_id"' in html and body["version"]

    # 確定して .md をダウンロードするところまで今までどおり
    assert client.post(f"/forms/{doc_id}/confirm", json={"version": body["version"]}).get_json()["ok"] is True
    state = finish(client, [doc_id])
    assert "zip" not in state["html"] and f"/forms/{doc_id}/download.md" in state["html"]
    assert client.get(f"/forms/{doc_id}/download.md").status_code == 200


# ---- 確定したあとに直した帳票は、ダウンロードの前に確定し直す ---------------------------------------

def test_a_modified_form_is_confirmed_again_before_the_md_download(app, client, sample_dir):
    """1件だけのときも zip と同じで、直した値が .md に入らずに消えることが無いようにする。

    確定したあとに読み取り結果を直すと状態は「修正中」になる。このとき .md は確定済みの古い版から
    作られるので、ダウンロードのボタンは zip と同じ data-confirm-all を持ち（app.js の帳票取り込みの部分が先に確定し直す）、
    確認文にもそのことを書く。
    """
    pattern_id = _pattern_forms_batch_review(client, sample_dir)
    doc_id, = upload_forms(client, sample_dir / "standard.xlsx")
    body = read_all(client, [doc_id], pattern_id, ["修理報告書"]).get_json()

    client.post(f"/forms/{doc_id}/confirm", json={"version": body["version"]})
    state = finish(client, [doc_id])
    assert "data-confirm-all" in state["html"]                       # 確定済みでも同じ通り道を使う
    assert "確定し直してから" not in state["html"]                    # 直していなければ断らない

    # 確定したあとに直す → 修正中。このままでは .md は確定済みの古い版から作られる
    res = client.post(f"/forms/{doc_id}/draft",
                      json={"values": {"equipment_name": "直した設備名"}, "version": body["version"]})
    assert res.status_code == 204
    with app.app_context():
        assert db.get_document(doc_id)["state"] == "modified"
        assert "直した設備名" not in db.get_document(doc_id)["confirmed_json"]

    state = finish(client, [doc_id])
    assert "data-confirm-all" in state["html"] and "確定し直してから、直した値で Markdown を作ります。" in state["html"]

    # app.js（帳票取り込み）が先に確定し直すので、渡す .md には直した値が入る
    client.post(f"/forms/{doc_id}/confirm", json={"version": res.headers["X-Doc-Version"]})
    assert "直した設備名" in client.get(f"/forms/{doc_id}/download.md").get_data(as_text=True)


# ====================================================================================================
# 元 tests/test_forms_fixes.py
# 帳票まわりの修正（確認で見つかった不具合）の回帰テスト。
# ====================================================================================================

def _number_field(**kw) -> dict:
    f = {"field_name": "downtime", "display_name": "停止時間", "data_type": "number",
         "value": None, "unit": "", "warning": None, "edited": False}
    f.update(kw)
    return f


def test_manual_number_is_judged_by_the_form_type_unit_not_the_read_unit():
    """「停止時間（分）」の欄が「14.9h」と読まれたあと、手で「894分」と直したら単位は分に戻る。

    数字だけの「894」は、画面のラベルに出ている読み取った単位（時間）のまま要確認にする（黙って分に変えない）。
    """
    ex = {"fields": [_number_field(value=14.9, unit="時間", spec_unit="分", warning="この項目の単位は「分」ですが…")]}
    apply_manual_values(ex, {"value-downtime": "894"})
    f = ex["fields"][0]
    assert f["value"] == 894 and f["unit"] == "時間" and f["edited"] and f["warning"].startswith("この項目の単位は")
    apply_manual_values(ex, {"value-downtime": "894分"})
    assert f["value"] == 894 and f["unit"] == "分" and f["edited"] and not f["warning"]


def test_manual_number_with_another_written_unit_keeps_it_and_stays_to_be_checked():
    from views import _field_status

    ex = {"fields": [_number_field(value=30, unit="分", spec_unit="分")]}
    apply_manual_values(ex, {"value-downtime": "2.5h"})
    f = ex["fields"][0]
    assert f["value"] == 2.5 and f["unit"] == "時間" and f["edited"]
    assert f["warning"].startswith("この項目の単位は")
    assert _field_status(f) == {"source": "manual", "issue": True, "blank": False}


def test_manual_number_takes_the_written_unit_when_the_form_type_has_none():
    """単位の無い「作業時間」に「150分」「2.5時間」と入れたら、その単位で出す（数値だけにしない）。"""
    from views import _field_status

    ex = {"fields": [_number_field(field_name="work_hours", display_name="作業時間", value=2.5, spec_unit="",
                                   warning="単位が書かれていません（時間か分か）")]}
    apply_manual_values(ex, {"value-work_hours": "2.5時間"})
    f = ex["fields"][0]
    assert f["value"] == 2.5 and f["unit"] == "時間" and f["edited"] and not f["warning"]

    apply_manual_values(ex, {"value-work_hours": "150分"})
    assert f["value"] == 150 and f["unit"] == "分" and not f["warning"]

    # 入力欄には数字だけが出る（単位はラベル）。数字だけ直したら、ラベルに出ていた単位のまま
    apply_manual_values(ex, {"value-work_hours": "160"})
    assert f["value"] == 160 and f["unit"] == "分" and not f["warning"] and not _field_status(f)["issue"]


def test_old_extractions_without_spec_unit_fall_back_to_the_unit():
    ex = {"fields": [_number_field(value=30, unit="分")]}
    apply_manual_values(ex, {"value-downtime": "45"})
    assert ex["fields"][0]["value"] == 45 and ex["fields"][0]["unit"] == "分"


def test_extraction_keeps_the_form_type_unit(tmp_path):
    from openpyxl import Workbook

    from forms import extract_document
    from forms import load_workbook_info
    from forms import FieldDef, PatternDef

    wb = Workbook()
    wb.active.title = "報告書"
    wb.active["A1"], wb.active["B1"] = "停止時間", "14.9h"
    path = tmp_path / "u.xlsx"
    wb.save(path)
    pattern = PatternDef(name="テスト", fields=[
        FieldDef("downtime", "停止時間", ["停止時間"], data_type="number", unit="分")])
    ex = extract_document(load_workbook_info(path), pattern, ["報告書"])
    f = ex["fields"][0]
    assert f["unit"] == "時間" and f["spec_unit"] == "分"
    apply_manual_values(ex, {"value-downtime": "894分"})
    assert f["value"] == 894 and f["unit"] == "分"


# ---- 確認画面・まとめ取り込み（views/forms.py） ------------------------------------------------




def _extraction(name="CMP研磨装置") -> str:
    field = {"field_name": "equipment_name", "display_name": "設備名", "data_type": "string",
             "value": name, "sheet": "報告書", "label_cell": "A1", "value_cell": "B1", "label_found": True,
             "warning": None, "edited": False, "unit": "", "spec_unit": "",
             "rag_output": "show", "table_columns": [], "table_blocks": 1}
    data = {"pattern": {"id": 1, "name": "設備修理報告書", "version": "v1"}, "sheets": ["報告書"],
            "fields": [field], "attachments": [], "values": {"equipment_name": name}}
    return json.dumps(data, ensure_ascii=False)


def _confirmed_doc(app, name="1.xlsx", *, batch_id="", order=0) -> int:
    with app.app_context():
        stored = f"documents/{name}"
        path = Path(app.config["UPLOAD_DIR"]) / stored
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(b"dummy-excel")
        doc_id = db.create_document(name, "0" * 64, stored, batch_id=batch_id, batch_order=order)
        db.update_document(doc_id, data_json=_extraction(), confirmed_json=_extraction(), title=f"{name} 確定")
    return doc_id


def _state(app, doc_id):
    with app.app_context():
        return db.get_document(doc_id)["state"]


def _doc_version(app, doc_id) -> str:
    """いまの作業データの版（画面が読み取り結果と一緒に受け取る値）。"""
    from views import _version

    with app.app_context():
        return _version(db.get_document(doc_id))


def _review_html(app, doc_id) -> str:
    """読み取り結果の欄のHTML（読み取りの応答が返すのと同じもの）。"""
    from views import _review_response

    with app.test_request_context():
        return _review_response(doc_id).get_json()["html"]


def test_typing_the_confirmed_value_back_returns_the_form_to_confirmed(app, client):
    doc_id = _confirmed_doc(app)
    client.post(f"/forms/{doc_id}/draft", json={"values": {"equipment_name": "別の名前"}})
    assert _state(app, doc_id) == "modified"
    client.post(f"/forms/{doc_id}/draft", json={"values": {"equipment_name": "CMP研磨装置"}})
    assert _state(app, doc_id) == "confirmed"
    summary = client.post(f"/forms/{doc_id}/preview", json={"values": {}}).get_json()
    assert summary["state"] == "confirmed" and summary["counts"]["manual"] == 0


def test_a_stale_review_page_cannot_save_or_confirm(app, client):
    """別のタブ（開いたままの古い画面）から、見ていない内容を上書き・確定しない。"""
    doc_id = _confirmed_doc(app)
    with app.app_context():
        db.update_document(doc_id, confirmed_json=None)  # 確認中の帳票
    old = _doc_version(app, doc_id)
    assert f'name="version" value="{old}"' in _review_html(app, doc_id)

    # タブAの保存: 新しい版が返り、それを送れば続けて保存できる
    res = client.post(f"/forms/{doc_id}/draft", json={"values": {"equipment_name": "A"}, "version": old})
    assert res.status_code == 204
    new = res.headers["X-Doc-Version"]
    assert new and new != old
    assert client.post(f"/forms/{doc_id}/draft",
                       json={"values": {"equipment_name": "A2"}, "version": new}).status_code == 204

    # タブB（古い版のまま）の保存・プレビュー・確定は止める
    res = client.post(f"/forms/{doc_id}/draft", json={"values": {"equipment_name": "B"}, "version": old})
    assert res.status_code == 409 and "別の画面で内容が変わりました" in res.get_json()["error"]
    assert client.post(f"/forms/{doc_id}/preview", json={"values": {}, "version": old}).status_code == 409
    res = client.post(f"/forms/{doc_id}/confirm", json={"values": {"equipment_name": "B"}, "version": old})
    assert res.status_code == 409 and "別の画面で内容が変わりました" in res.get_json()["error"]
    with app.app_context():
        doc = db.get_document(doc_id)
        assert doc["state"] == "reviewing" and "A2" in doc["data_json"] and '"B"' not in doc["data_json"]


def _finish(client, ids, current=None) -> dict:
    url = "/forms/finish?ids=" + ",".join(str(i) for i in ids) + (f"&current={current}" if current else "")
    return client.get(url).get_json()


def test_a_partly_confirmed_batch_does_not_say_everything_is_deleted(app, client):
    first = _confirmed_doc(app, "1.xlsx", batch_id="B", order=0)
    with app.app_context():
        pending = db.create_document("2.xlsx", "0" * 64, "documents/2.xlsx", batch_id="B", batch_order=1)
        db.update_document(pending, data_json=_extraction())   # 読み取り済み・未確定
    body = _finish(client, [first, pending], current=first)
    assert body["confirmed"] == 1 and body["total"] == 2 and body["next_id"] == pending
    page = body["html"]
    # 未確定が残っているときは、その分もこのボタンで確定してから渡すと書く
    assert "まだ確定していない1件も確定してから、2件をまとめて zip でダウンロードします" in page
    assert "残り1件も確定して、まとめてダウンロード（zip）" in page

    # すべて確定すれば、まとめてのダウンロードだけになる
    with app.app_context():
        db.update_document(pending, data_json=_extraction(), confirmed_json=_extraction())
    page = _finish(client, [first, pending], current=first)["html"]
    assert "まとめて Markdown をダウンロード（zip）" in page and "残り1件も確定して" not in page
    assert "このまとまりの帳票のデータはサーバーからすべて消えます" in page


def test_batch_upload_error_names_the_file_that_was_skipped(app, client, sample_dir):
    files = [(io.BytesIO((sample_dir / "standard.xlsx").read_bytes()), "standard.xlsx"),
             (io.BytesIO(b"this is not an excel file"), "broken.xlsx"),
             (io.BytesIO((sample_dir / "shifted.xlsx").read_bytes()), "shifted.xlsx")]
    res = client.post("/forms/upload", data={"file": files}, content_type="multipart/form-data")
    body = res.get_json()
    # どのファイルかは選んだ順の位置で示す（ファイル名そのものは画面の外に出さない。design.md 3.3）
    assert res.status_code == 200 and len(body["docs"]) == 2 and body["batch_id"]
    assert body["errors"] and body["errors"][0].startswith("2件目のファイル: ")
    assert "broken.xlsx" not in " ".join(body["errors"])


# ---- 時刻の範囲の「作業時間」（excel/text.py） ------------------------------------------------------

def test_time_range_is_not_read_as_its_start_hour():
    from forms import numeric_unit, to_number

    # 添えた時間数（工数）を項目の単位で使う
    value, warning = to_number(None, "09:30-12:45（3.2h）", "時間")
    assert value == 3.2 and "時間数" in warning
    assert to_number(None, "09:30-12:45（3.2h）", "分")[0] == 192
    assert to_number(None, "９:３０～１３:５５（工数 ８.７５h）", "時間")[0] == 8.75
    # 時間数が無ければ範囲の長さ（日をまたぐ作業も）
    value, warning = to_number(None, "9:30-12:45", "分")
    assert value == 195 and "時刻の範囲" in warning
    assert to_number(None, "22:00～02:30", "時間")[0] == 4.5
    # 単位の決まっていない項目は分にする（「3時間40分」と同じ）
    assert to_number(None, "09:30-12:45（3.2h）")[0] == 192 and numeric_unit("09:30-12:45（3.2h）") == "分"
    # 時間でない単位の項目では読まない（開始時刻の 9 を返さない）
    assert to_number(None, "9:30-12:45", "円") == ("9:30-12:45", "時刻の範囲です。数値として読み取れません")


# ---- 設備だけのタイトル（export/formats.py） ----------------------------------------------------

def _md_field(name, display, value, data_type="string"):
    return {"field_name": name, "display_name": display, "data_type": data_type, "value": value,
            "unit": "", "rag_output": "show", "edited": False}


def _inspection(work_no, work_date, finding="端子の緩みを増し締めした。" * 200):
    fields = [_md_field("equipment_id", "設備番号", "IMP-603"), _md_field("equipment_name", "設備名", "高電流注入 3号機"),
              _md_field("field_13", "作業日", work_date, "date"), _md_field("field_6", "作業No.", work_no),
              _md_field("finding", "所見", finding, "text")]
    ex = {"pattern": {"id": 4, "name": "点検保全作業報告書", "title_fields": [], "md_options": {}, "labels": []},
          "sheets": ["報告書"], "fields": fields, "attachments": []}
    refresh_summary(ex)
    return ex


def test_reports_for_the_same_equipment_get_different_titles_and_headings():
    from forms import build_markdown

    doc = {"id": 1, "file_name": "点検.xlsx", "file_hash": "0" * 64}
    a = build_markdown(doc, _inspection("W-0101", "2026-04-01"))
    b = build_markdown(doc, _inspection("W-0102", "2026-04-08"))
    title_a, title_b = a.split("\n", 1)[0], b.split("\n", 1)[0]
    assert title_a == "# 点検保全作業報告書 高電流注入 3号機（IMP-603）｜W-0101" and title_a != title_b
    heading_a = next(line for line in a.split("\n") if line.startswith("## "))
    heading_b = next(line for line in b.split("\n") if line.startswith("## "))
    assert heading_a == "## 所見（IMP-603 高電流注入 3号機／W-0101）" and heading_a != heading_b

    # 番号らしい項目が無ければ日付、それも無ければ元ファイル名
    ex = _inspection("", "2026-04-08")
    assert build_markdown(doc, ex).startswith("# 点検保全作業報告書 高電流注入 3号機（IMP-603）｜2026-04-08\n")
    ex = _inspection("", None)
    assert build_markdown(doc, ex).startswith("# 点検保全作業報告書 高電流注入 3号機（IMP-603）｜点検\n")


# ---- 押印欄の人名を見出しにした項目（pattern/builder.py・export/formats.py） ----------------------------

def test_a_person_name_used_as_a_label_is_not_written(tmp_path):
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill

    from forms import load_workbook_info
    from forms import build_markdown
    from forms import suggest_rows

    # 押印欄: 「作成｜確認」の下に人名。builder は「斎藤」を見出しの候補にしない
    wb = Workbook()
    ws = wb.active
    ws.title = "報告書"
    ws["A1"], ws["B1"] = "件名", "搬送エラー"
    ws["D1"], ws["E1"] = "作成", "確認"
    ws["D2"], ws["E2"] = "斎藤", "森"
    for coord in ("A1", "D1", "E1"):  # 見出しの欄は色付き（帳票の普通の形）
        ws[coord].fill = PatternFill("solid", fgColor="DDEBF7")
    path = tmp_path / "stamp.xlsx"
    wb.save(path)
    _, rows = suggest_rows([load_workbook_info(path)])
    assert "斎藤" not in {r["display_name"] for r in rows}

    # すでにそう作られた帳票の種類（人名が見出しになったもの）は、読み取ったとおりに出す
    fields = [_md_field("subject", "件名", "搬送エラー"), _md_field("field_60", "作成", "斎藤"),
              _md_field("field_29", "斎藤", "森"), _md_field("field_61", "確認", "森")]
    ex = {"pattern": {"id": 2, "name": "トラブル報告書", "title_fields": [], "md_options": {}, "labels": []},
          "sheets": ["報告書"], "fields": fields, "attachments": []}
    refresh_summary(ex)
    md = build_markdown({"id": 1, "file_name": "t.xlsx", "file_hash": "0" * 64}, ex)
    assert "- 件名: 搬送エラー" in md and "- 作成: 斎藤" in md and "- 斎藤: 森" in md


def test_upload_page_shows_the_batch_limits(client):
    page = client.get("/forms/new").get_data(as_text=True)
    assert "50ファイル・合計200MBまで" in page


# ---- 2回目の修正（帳票） -----------------------------------------------------------------------

def _read_standard(app, client, sample_dir) -> tuple[int, int]:
    """standard.xlsx を取り込んで読み取った帳票（確認中）を作る。戻り値: (帳票ID, 種類ID)"""

    path = sample_dir / "standard.xlsx"
    pattern_id = create_type(client, path, "設備修理報告書")
    for label_cell, value_cell in (("A3", "B3"), ("A4", "B4"), ("E4", "F4"), ("A7", "A8")):
        add_field(client, pattern_id, "修理報告書", label_cell, value_cell)
    activate(client, pattern_id)
    doc_id, = upload_forms(client, path)
    read_form(client, doc_id, pattern_id, ["修理報告書"])
    return doc_id, pattern_id


def _doc_extraction(app, doc_id) -> dict:
    with app.app_context():
        return json.loads(db.get_document(doc_id)["data_json"])


def test_editing_a_reread_confirmed_form_keeps_the_new_form_type_settings(app, client, sample_dir):
    """確定後に種類で「Markdownに出さない」にした項目は、読み直したあと別の項目を直しても出ないまま。"""
    doc_id, pattern_id = _read_standard(app, client, sample_dir)
    ex = _doc_extraction(app, doc_id)
    shown = [f for f in ex["fields"] if f["data_type"] == "string" and f["value"]]
    omitted, edited = shown[0], shown[1]
    with app.app_context():
        doc = db.get_document(doc_id)
        db.update_document(doc_id, confirmed_json=doc["data_json"])  # 確定済み
        pattern = db.load_pattern(pattern_id)
        for fd in pattern.fields:
            if fd.field_name == omitted["field_name"]:
                fd.rag_output = "omit"
        db.save_pattern(pattern, "active")
    client.post(f"/forms/{doc_id}/read", data={"pattern_id": pattern_id, "sheets": ex["sheets"],
                                               "acknowledge": "on"})
    line = f"- {omitted['display_name']}: "
    assert line not in client.post(f"/forms/{doc_id}/preview", json={"values": {}}).get_json()["markdown"]

    client.post(f"/forms/{doc_id}/draft", json={"values": {edited["field_name"]: "別の値"}})
    ex = _doc_extraction(app, doc_id)
    field = next(f for f in ex["fields"] if f["field_name"] == omitted["field_name"])
    assert field["rag_output"] == "omit"
    assert line not in client.post(f"/forms/{doc_id}/preview", json={"values": {}}).get_json()["markdown"]


def test_confirmed_field_is_not_restored_when_its_settings_differ():
    from views import _restore_confirmed_fields

    old = {"field_name": "line", "display_name": "ライン", "data_type": "string",
           "value": "L4", "unit": "", "rag_output": "show"}
    new = dict(old, display_name="製造ライン", rag_output="omit", edited=True)
    confirmed = {"pattern": {"id": 1, "version_no": 2}, "sheets": ["s"], "fields": [old]}
    ex = {"pattern": {"id": 1, "version_no": 2}, "sheets": ["s"], "fields": [new]}
    _restore_confirmed_fields(ex, confirmed)
    assert ex["fields"][0]["rag_output"] == "omit" and ex["fields"][0]["display_name"] == "製造ライン"
    # 設定が同じなら、確定済みの中身に戻す（従来どおり）
    ex = {"pattern": {"id": 1, "version_no": 2}, "sheets": ["s"], "fields": [dict(old, edited=True)]}
    _restore_confirmed_fields(ex, confirmed)
    assert ex["fields"][0] == old


def test_triangle_minus_sign_is_read_as_negative():
    from forms import number_unit
    from forms import to_number

    assert to_number("▲50万円", "▲50万円", "万円")[0] == -50
    assert to_number("△0.8", "△0.8", "%") == (-0.8, None)
    assert to_number("▲5", "▲5", "") == (-5, None)
    value, warning = to_number("▲50万円", "▲50万円", "万円")
    assert number_unit(value, "▲50万円", "万円", "cost", "コスト削減効果", warning) == ("万円", None)
    # 手で「▲50万円」と入れ直しても負の数
    ex = {"fields": [_number_field(field_name="cost", display_name="コスト削減効果", value=50, unit="万円",
                                   spec_unit="万円")]}
    apply_manual_values(ex, {"value-cost": "▲50万円"})
    assert ex["fields"][0]["value"] == -50 and ex["fields"][0]["unit"] == "万円"


def test_dated_time_range_is_not_read_as_its_month():
    from forms import numeric_unit, to_number

    s = "12/24 21:53-12/25 11:09（13.3h）"
    assert to_number(s, s, "")[0] == 798 and numeric_unit(s) == "分"
    assert to_number(s, s, "時間")[0] == 13.3
    s = "3/19 11:19～3/20 0:48（工数 27.5h）"
    assert to_number(s, s, "時間")[0] == 27.5
    # 時間数の無い月日付きの範囲は、日数が分からないので数値にしない
    s = "12/24 21:53-12/25 11:09"
    value, warning = to_number(s, s, "")
    assert value == s and "時刻の範囲" in warning


def test_unit_written_next_to_an_annotated_number_is_kept():
    from forms import number_unit
    from forms import to_number

    def judge(text, spec):
        value, warning = to_number(text, text, spec)
        return (value, *number_unit(value, text, spec, "stop_minutes", "停止時間", warning))

    value, unit, warning = judge("約90分", "時間")
    assert value == 90 and unit == "分" and warning.startswith("この項目の単位は「時間」")
    value, unit, warning = judge("約2時間", "分")
    assert value == 2 and unit == "時間" and warning.startswith("この項目の単位は「分」")
    value, unit, warning = judge("595分（9.9h）", "")
    assert value == 595 and unit == "分" and "数値の部分だけ" in warning  # 括弧書きは残っているので要確認のまま
    # 「3時間40分」は換算するので、最初の「時間」を書かれた単位として扱わない
    assert judge("3時間40分", "分")[:2] == (220, "分")
    from forms import written_unit

    assert written_unit("約2号機") == "" and written_unit("約90分ぐらいかかった") == ""
    assert written_unit("約2号機 30分") == ""  # to_number が読む最初の数値（2）の単位だけを見る  # 単位らしくない言葉は単位にしない


def test_date_keeps_the_time_of_day():
    from datetime import datetime

    from forms import to_date
    from forms import _format_value

    assert to_date(datetime(2023, 7, 10, 23, 8), "") == ("2023-07-10 23:08", None)
    assert to_date(datetime(2023, 7, 10), "") == ("2023-07-10", None)
    assert to_date("2023/7/10 23:08", "2023/7/10 23:08") == ("2023-07-10 23:08", None)
    f = {"data_type": "date", "value": "2023-07-10 23:08"}
    assert _format_value(f) == "2023-07-10 23:08（2023年7月）"


def test_hand_typed_date_without_a_year_stays_to_be_checked():
    from views import _field_status

    ex = {"fields": [{"field_name": "d", "display_name": "発生日", "data_type": "date",
                      "value": "2/12 3時17分", "warning": "年が書かれていません。", "edited": False}]}
    apply_manual_values(ex, {"value-d": "2/12 3:17"})
    assert ex["fields"][0]["edited"] and _field_status(ex["fields"][0])["issue"]
    apply_manual_values(ex, {"value-d": "あした"})
    assert _field_status(ex["fields"][0])["issue"]
    apply_manual_values(ex, {"value-d": "2026-02-12 3:17"})
    assert not ex["fields"][0]["warning"] and not _field_status(ex["fields"][0])["issue"]


def test_batch_progress_counts_the_confirmed_forms(app, client):
    ids = [_confirmed_doc(app, f"{i}.xlsx", batch_id="b1", order=i) for i in range(2)]
    body = _finish(client, ids, current=ids[0])
    assert body["confirmed"] == 2 and body["next_id"] is None
    assert "確定済み <strong>2</strong> / 2 件" in body["html"]
    with app.app_context():
        db.update_document(ids[1], confirmed_json=None)
    body = _finish(client, ids, current=ids[0])
    assert body["confirmed"] == 1 and body["next_id"] == ids[1]
    assert "確定済み <strong>1</strong> / 2 件" in body["html"]
    assert "残り1件も確定して、まとめてダウンロード（zip）" in body["html"]


def test_upload_errors_do_not_put_the_file_name_in_the_session(app, client, sample_dir):
    files = [(io.BytesIO((sample_dir / "standard.xlsx").read_bytes()), "standard.xlsx"),
             (io.BytesIO(b"PK\x03\x04 broken"), "X社_社外秘.xlsx")]
    res = client.post("/forms/upload", data={"file": files}, content_type="multipart/form-data")
    errors = " ".join(res.get_json()["errors"])
    assert "2件目のファイル" in errors and "社外秘" not in errors
    # 画面を離れても残るところ（セッションのクッキー）にはファイル名を置かない
    assert "社外秘" not in str(res.headers.get("Set-Cookie", ""))
    with client.session_transaction() as session:
        assert not session.get("_flashes")


def _sheetless_workbook(sample_dir) -> bytes:
    import zipfile

    src = zipfile.ZipFile(sample_dir / "standard.xlsx")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/workbook.xml":
                data = re.sub(rb"<sheets>.*?</sheets>", b"", data, flags=re.S)
            dst.writestr(item, data)
    return out.getvalue()


def test_workbook_without_a_sheet_list_is_refused(app, client, sample_dir):
    res = client.post("/forms/upload", data={"file": (io.BytesIO(_sheetless_workbook(sample_dir)), "s.xlsx")},
                      content_type="multipart/form-data")
    assert res.status_code == 400 and "シートがないブックです" in res.get_json()["error"]
    with app.app_context():
        assert db.get_db().execute("SELECT COUNT(*) FROM documents").fetchone()[0] == 0
    stored = Path(app.config["UPLOAD_DIR"]) / "documents"
    assert not stored.exists() or not any(stored.iterdir())


def test_workbook_with_too_many_cells_is_refused_before_reading(app, client, tmp_path):
    """展開すると大量のセルがある小さなブックは、帳票でも帳票登録でも読み込む前に断り、何も残さない。"""
    from tests.test_core import _xlsx_with_cells

    app.config["EXCEL_MAX_CELLS"] = 1000
    data = _xlsx_with_cells(tmp_path / "many.xlsx", 1001).read_bytes()
    res = client.post("/forms/upload", data={"file": (io.BytesIO(data), "many.xlsx")},
                      content_type="multipart/form-data")
    assert res.status_code == 400 and "セル数が上限（1,000 セル）を超えています" in res.get_json()["error"]
    res = client.post("/form-types/new", data={"name": "多すぎ", "book": (io.BytesIO(data), "many.xlsx")},
                      content_type="multipart/form-data")
    assert res.status_code == 400 and "セル数が上限（1,000 セル）を超えています" in res.get_json()["error"]
    with app.app_context():
        assert db.get_db().execute("SELECT COUNT(*) FROM documents").fetchone()[0] == 0
        assert db.list_patterns() == []      # 読めなかったブックでは帳票の種類も作らない
    assert [p for p in Path(app.config["UPLOAD_DIR"]).rglob("*") if p.is_file()] == []


# ====================================================================================================
# 元 tests/test_forms_fixes3.py
# 帳票まわりの修正（3巡目の確認で見つかった不具合）の回帰テスト。
# ====================================================================================================

def _number_field_forms_fixes3(**kw) -> dict:
    f = {"field_name": "downtime", "display_name": "停止時間", "data_type": "number",
         "value": None, "unit": "", "spec_unit": "", "warning": None, "edited": False}
    f.update(kw)
    return f


# ---- F3-1: 時:分の時間・範囲・「.5」 ------------------------------------------------------------

def test_time_cell_is_converted_to_the_field_unit_not_read_as_its_hour():
    assert to_number(time(1, 30), "01:30", "時間")[0] == 1.5
    assert to_number(time(1, 30), "01:30", "分")[0] == 90
    assert to_number(time(1, 30), "01:30", "")[0] == 90  # 単位の無い項目は分


def test_hmm_duration_cell_over_a_day_is_converted():
    td = timedelta(hours=25, minutes=30)
    assert cell_text(td) == "25:30"  # 「1 day, 1:30:00」にしない
    assert to_number(td, cell_text(td), "分")[0] == 1530
    assert to_number(td, cell_text(td), "時間")[0] == 25.5


def test_hmm_text_is_a_duration():
    assert to_number(None, "2:45", "時間")[0] == 2.75
    assert to_number(None, "2:45", "")[0] == 165
    value, warning = to_number(None, "2:45", "円")
    assert value == "2:45" and warning  # 分・時間でない項目では数値にしない


def test_range_with_a_different_unit_is_not_read_as_its_first_number():
    value, warning = to_number(None, "10～20分", "時間")
    assert value == "10～20分" and "範囲" in warning


def test_leading_dot_decimal_and_h_m_notation():
    assert to_number(None, ".5", "時間") == (0.5, None)
    assert to_number(None, "約.5時間", "時間")[0] == 0.5
    assert to_number(None, "1h30m", "")[0] == 90


def test_manual_partial_number_stays_to_be_checked():
    """手で「約90分」と入れても「数値の部分だけ」の警告は要確認のまま（黙って確定できない）。"""
    from views import _field_status

    ex = {"fields": [_number_field_forms_fixes3(value=30, unit="分", spec_unit="分")]}
    apply_manual_values(ex, {"value-downtime": "約90分くらい"})
    f = ex["fields"][0]
    assert f["edited"] and "数値の部分だけ" in f["warning"]
    assert _field_status(f)["issue"] is True

    apply_manual_values(ex, {"value-downtime": "2:45"})
    assert f["value"] == 165 and f["unit"] == "分"


# ---- F3-2 / md3-3: 日付のあとの時刻・範囲・平成 ---------------------------------------------------

def test_time_after_japanese_date_era_and_iso_is_kept():
    assert to_date(None, "2024年7月29日 13:41") == ("2024-07-29 13:41", None)
    assert to_date(None, "２０２５年２月２４日 ２０:４０") == ("2025-02-24 20:40", None)
    assert to_date(None, "2023年7月10日 23時08分") == ("2023-07-10 23:08", None)
    assert to_date(None, "2023-07-10T23:08") == ("2023-07-10 23:08", None)
    assert to_date(None, "R5.11.16 12:11") == ("2023-11-16 12:11", None)
    assert to_date(None, "2024/1/5 (金) 9:05") == ("2024-01-05 09:05", None)
    assert to_date(None, "2024年1月5日") == ("2024-01-05", None)


def test_text_after_the_date_is_flagged():
    value, warning = to_date(None, "2023/7/10～7/12")
    assert value == "2023-07-10" and warning and "7/12" in warning


def test_heisei_dates_are_read():
    assert to_date(None, "H30.4.1") == ("2018-04-01", None)
    assert to_date(None, "平成30年4月1日") == ("2018-04-01", None)


# ---- md3-1: 表示形式の単位 ------------------------------------------------------------------

def test_format_unit():
    assert format_unit('#,##0"分"') == "分"
    assert format_unit('#,##0" 分";[Red]-#,##0" 分"') == "分"
    assert format_unit('"¥"#,##0') == "円"
    assert format_unit("#,##0") == "" and format_unit("General") == "" and format_unit('0" - "') == ""


def test_number_unit_from_the_cell_format(tmp_path):
    from openpyxl import load_workbook

    path_info = _book(tmp_path, "fmt.xlsx", {"報告書": ({"A1": "ダウンタイム", "B1": 264},
                                                      {"A1": LABEL_FILL}, [])})
    wb = load_workbook(path_info.path)
    wb["報告書"]["B1"].number_format = '#,##0"分"'
    wb.save(path_info.path)
    from forms import load_workbook_info

    info = load_workbook_info(path_info.path)
    values, fields = _values(info, _fields(("downtime", "number", "ダウンタイム")))
    assert values["downtime"] == 264 and fields["downtime"]["unit"] == "分" and not fields["downtime"]["warning"]


# ---- md3-2: 水平展開先の一覧の列見出しを1つの値のラベルにしない --------------------------------------

def test_list_under_heading_with_a_row_between_is_not_a_single_value(tmp_path):
    cells = {"A1": "７．水平展開", "A2": "区分", "B2": "☑同型機　□類似設備",
             "A3": "確認", "B3": "対象設備", "C3": "設備名", "D3": "実施日/予定", "E3": "結果・備考",
             "A4": "☑", "B4": "IMP-603", "C4": "高電流注入 3号機", "D4": "2023/11/18", "E4": "展開済み",
             "A5": "☑", "B5": "IMP-604", "C5": "高エネルギー注入 4号機", "D5": "2023/11/17", "E5": "展開済み",
             "A7": "影響", "B7": "停止時間", "C7": "影響ロット", "A8": "搬送停止", "B8": "120", "C8": "3"}
    fills = {c: LABEL_FILL for c in ("A1", "A2", "A3", "B3", "C3", "D3", "E3", "A7", "B7", "C7")}
    info = _book(tmp_path, "yoko.xlsx", {"報告書": (cells, fills, [])})
    values, _ = _values(info, _fields(("target", "string", "対象設備"), ("stop", "string", "停止時間")))
    assert values["target"] is None  # 水平展開先の1行目（IMP-603）を報告書の対象設備にしない
    assert values["stop"] == "120"  # 見出しの行＋値の行1つは、今までどおり項目として読む


# ---- F1: Excel のエラー値 ---------------------------------------------------------------

def test_excel_error_values_are_not_taken_as_values(tmp_path):
    from views import _field_status

    info = _book(tmp_path, "err.xlsx", {"報告書": (
        {"A1": "報告番号", "B1": "#REF!", "A2": "作業時間", "B2": "#DIV/0!", "A3": "設備名", "B3": "CMP 1号機"},
        {"A1": LABEL_FILL, "A2": LABEL_FILL, "A3": LABEL_FILL}, [])})
    values, fields = _values(info, _fields(("report_id", "string", "報告番号"), ("work", "number", "作業時間"),
                                           ("name", "string", "設備名")))
    assert values["report_id"] is None and "#REF!" in fields["report_id"]["warning"]
    assert values["work"] is None and fields["work"]["unit"] == "" and "#DIV/0!" in fields["work"]["warning"]
    assert _field_status(fields["report_id"])["issue"] and _field_status(fields["work"])["issue"]
    assert values["name"] == "CMP 1号機"


# ---- F3-3: 空の明細表で行を足して消しただけ -------------------------------------------------------

def test_empty_table_add_and_remove_row_is_not_an_edit():
    from views import _apply_values, _field_status

    f = {"field_name": "parts", "display_name": "使用部品", "data_type": "table",
         "value": None, "table_columns": ["部品名", "数量"], "warning": "表に行がありません",
         "edited": False, "unit": ""}
    ex = {"fields": [f], "pattern": {"id": 1, "version_no": 1}, "sheets": ["S"]}
    refresh_summary(ex)  # 保存済みの読み取り結果と同じ形にする
    confirmed = copy.deepcopy(ex)
    changed = _apply_values(ex, {"parts": json.dumps({"columns": ["部品名", "数量"], "rows": []})}, confirmed)
    assert changed is False
    assert ex["fields"][0]["value"] is None and not ex["fields"][0]["edited"]
    assert _field_status(ex["fields"][0])["source"] == "blank"


# ---- R3B-3: 途中保存の応答が届く前に画面を離れた（同じ画面の beacon） --------------------------------

def _working_doc(app):
    import database as db

    doc_id = _confirmed_doc(app)
    with app.app_context():
        db.update_document(doc_id, confirmed_json=None)  # 確認中の帳票
    return doc_id


def test_a_beacon_from_the_same_page_is_not_refused_as_another_page(app, client):
    doc_id = _working_doc(app)
    old = _doc_version(app, doc_id)
    first = client.post(f"/forms/{doc_id}/draft",
                        json={"values": {"equipment_name": "A"}, "version": old, "page_token": "page-1"})
    assert first.status_code == 204
    # 応答（新しい版）を受け取る前に送った beacon は古い版のまま。同じ画面なら受け付ける
    second = client.post(f"/forms/{doc_id}/draft",
                         json={"values": {"equipment_name": "A2"}, "version": old, "page_token": "page-1"})
    assert second.status_code == 204
    # 別の画面（目印が違う）の古い版は止める
    other = client.post(f"/forms/{doc_id}/draft",
                        json={"values": {"equipment_name": "B"}, "version": old, "page_token": "page-2"})
    assert other.status_code == 409
    # 目印を送らない古い版も止める
    assert client.post(f"/forms/{doc_id}/draft",
                       json={"values": {"equipment_name": "C"}, "version": old}).status_code == 409


def test_a_beacon_is_refused_after_another_page_saved(app, client):
    doc_id = _working_doc(app)
    old = _doc_version(app, doc_id)
    res = client.post(f"/forms/{doc_id}/draft",
                      json={"values": {"equipment_name": "A"}, "version": old, "page_token": "page-a"})
    new = res.headers["X-Doc-Version"]
    assert client.post(f"/forms/{doc_id}/draft",
                       json={"values": {"equipment_name": "B"}, "version": new, "page_token": "page-b"}).status_code == 204
    # 画面Aの版は、画面Bの保存でできた版ではないので止める
    assert client.post(f"/forms/{doc_id}/draft",
                       json={"values": {"equipment_name": "A2"}, "version": old, "page_token": "page-a"}).status_code == 409


def test_draft_page_marks_are_forgotten_when_the_form_is_purged(app, client):
    """消した帳票の途中保存の目印（id）をメモリに残さない（design.md 3.3 データを残さない）。"""
    import views

    doc_id = _working_doc(app)
    old = _doc_version(app, doc_id)
    assert client.post(f"/forms/{doc_id}/draft",
                       json={"values": {"equipment_name": "A"}, "version": old, "page_token": "page-1"}).status_code == 204
    assert doc_id in views._DRAFT_TOKENS
    assert client.post(f"/forms/{doc_id}/delete").get_json()["ok"] is True
    assert doc_id not in views._DRAFT_TOKENS


# ====================================================================================================
# 元 tests/test_forms_fixes4.py
# 帳票まわりの修正（4巡目の確認で見つかった不具合）の回帰テスト。
# ====================================================================================================

# ---- F4-1: 「〜計」で終わる品名（温度計・圧力計）を合計行にしない ------------------------------------

def test_instrument_names_ending_in_kei_are_not_total_rows():
    from forms import table_markdown_lines

    value = {"columns": ["品名", "型式", "数量"],
             "rows": [["ベアリング", "6205ZZ", "2"], ["圧力計", "GV-50", "1"], ["温度計", "", ""],
                      ["設計", "", ""], ["膜厚計", "", ""], ["部品費計", "", "12000"], ["合計", "", "3"]]}
    lines = table_markdown_lines(value)
    assert "- 品名: 圧力計／型式: GV-50／数量: 1" in lines
    assert "- 品名: 温度計" in lines and "- 品名: 設計" in lines and "- 品名: 膜厚計" in lines
    assert "- 部品費計: 数量: 12000" in lines and "- 合計: 数量: 3" in lines
    # 数字の無い合計行は今までどおり出さない
    assert table_markdown_lines({"columns": ["品名"], "rows": [["ノギス"], ["小計"]]}) == ["- 品名: ノギス"]


# ---- F4-3: 手で入れた値を消すと、確定済みの状態に戻る -------------------------------------------------

def test_clearing_a_hand_typed_value_returns_the_field_to_the_confirmed_one():
    from views import _apply_values

    confirmed = json.loads(_extraction())
    confirmed["fields"][0].update(value=None, sheet=None, value_cell=None, warning="ラベルはありますが値が空です")
    working = json.loads(json.dumps(confirmed))
    working["fields"][0].update(value="手で入れた値", sheet="報告書", value_cell="Z99", edited=True, warning=None)
    assert _apply_values(working, {"equipment_name": ""}, confirmed)
    assert working["fields"][0] == confirmed["fields"][0]


# ---- F4-4 / BR4-2: まとめ取り込みの「次の帳票へ」と、途中保存後の zip の確認文 ---------------------------

def _reviewing_doc(app, name, order):
    doc_id = _confirmed_doc(app, name, batch_id="B", order=order)
    with app.app_context():
        db.update_document(doc_id, confirmed_json=None)  # 確認中（手の修正があるかもしれない）
    return doc_id


def test_the_next_form_of_a_batch_is_the_one_that_is_not_confirmed_yet(app, client):
    first = _confirmed_doc(app, "1.xlsx", batch_id="B", order=0)
    second = _reviewing_doc(app, "2.xlsx", 1)
    assert _state(app, second) == "reviewing"
    body = client.get(f"/forms/finish?ids={first},{second}&current={first}").get_json()
    assert body["next_id"] == second and body["confirmed"] == 1
    # 読み取り結果は全部並んでいるので、④からはその帳票の塊へ移動できる
    assert f'data-goto-doc="{second}"' in body["html"]
    assert "残り1件も確定して、まとめてダウンロード（zip）" in body["html"]

    # まだ確定していない帳票は zip に入らない（確定済みの分だけを渡す）
    res = client.get("/forms/batches/B/download.zip")
    assert res.status_code == 200 and res.mimetype == "application/zip"
    with zipfile.ZipFile(io.BytesIO(res.data)) as zf:
        assert len(zf.namelist()) == 1
    with app.app_context():
        assert db.get_document(first) is None and db.get_document(second) is not None
    assert client.get("/forms/batches/B/download.zip").status_code == 404   # 確定済みが無ければ渡さない


# ---- R4-MD-2: 明細表の「〃」は上の行の値にする ---------------------------------------------------------

def _cells(*texts):
    return [SimpleNamespace(text=t) for t in texts]


def test_ditto_marks_in_a_detail_table_take_the_value_above():
    from forms import Table
    from forms import table_markdown_lines

    table = Table(anchor=None, header=_cells("点検箇所", "点検内容", "判定"),
                  rows=[_cells("〃", "外観", "○"), _cells("ステージ", "位置偏差", "○"), _cells("〃", "振動", "×"),
                        _cells("同上", "異音", "○"), _cells("", "清掃", "○"), _cells("〃", "給油", "○")])
    value = table.to_value()
    assert [r[0] for r in value["rows"]] == ["〃", "ステージ", "ステージ", "ステージ", "", "〃"]
    assert "- 点検箇所: ステージ／点検内容: 振動／判定: ×" in table_markdown_lines(value)
    # 「〃」だけの行（明細なしの記入）は今までどおり落とす
    only = Table(anchor=None, header=_cells("部位", "内容"), rows=[_cells("軸", "給油"), _cells("〃", "")])
    assert only.to_value()["rows"] == [["軸", "給油"]]


# ---- ux4-1: 同じ値に読める手の入力でも、警告の出る入力は手で修正にする ----------------------------------

def test_same_value_input_with_a_warning_is_kept_as_an_edit():
    ex = {"fields": [
        {"field_name": "w", "display_name": "作業時間", "data_type": "number", "value": 1.5,
         "unit": "時間", "spec_unit": "時間", "warning": None, "edited": False},
        {"field_name": "d", "display_name": "発生日", "data_type": "date",
         "value": "2026-09-14", "warning": None, "edited": False}]}
    apply_manual_values(ex, {"value-w": "1.5～3時間", "value-d": "2026-09-14 24:30"})
    w, d = ex["fields"]
    assert w["edited"] and w["warning"]
    assert d["edited"] and d["warning"]
    # 警告の消える入力（同じ値を書き直した）も手で修正にして、警告を消す
    apply_manual_values(ex, {"value-d": "2026/9/14"})
    assert d["value"] == "2026-09-14" and not d["warning"] and d["edited"]


def test_same_value_without_warning_is_still_not_an_edit():
    ex = {"fields": [{"field_name": "w", "display_name": "作業時間", "data_type": "number",
                      "value": 1.5, "unit": "時間", "spec_unit": "時間", "warning": None, "edited": False}]}
    apply_manual_values(ex, {"value-w": "1.50"})
    assert not ex["fields"][0]["edited"]


# ---- ux4-4: 読み取りテストの単位不明の警告は、種類での直し方にする ------------------------------------

def test_form_type_panel_explains_how_to_set_the_unit(app, client, tmp_path):
    from openpyxl import Workbook


    wb = Workbook()
    ws = wb.active
    ws.title = "報告書"
    ws["A1"], ws["B1"], ws["A2"], ws["B2"] = "報告番号", "R-001", "作業時間", 2.5
    path = tmp_path / "unit.xlsx"
    wb.save(path)
    pattern_id = create_type(client, path, "単位なし")
    add_field(client, pattern_id, "報告書", "A1", "B1")
    add_field(client, pattern_id, "報告書", "A2", "B2")
    with app.app_context():
        assert next(f for f in db.load_pattern(pattern_id).fields if f.display_name == "作業時間").data_type == "number"

    panel = panel_html(client, pattern_id, book=path)
    assert "帳票に単位が書かれていません" in panel and "値に単位（分・時間など）を付けて入力できます" in panel


# ---- R4-FUZZ-3: 見えない文字（ゼロ幅スペース・BOM）を消す -------------------------------------------------

def test_invisible_characters_are_removed_from_labels_and_values(tmp_path):
    from forms import cell_text, nfkc_value, normalize_label

    assert normalize_label("設備名​") == normalize_label("設備名")
    assert cell_text("﻿EQ-01⁠") == "EQ-01"
    assert nfkc_value("CMP​-108") == "CMP-108"
    info = _book(tmp_path, "zw.xlsx", {"報告書": ({"A1": "設備名​", "B1": "CMP⁠研磨装置"}, {}, [])})
    values, _ = _values(info, _fields(("equipment_name", "string", "設備名")))
    assert values["equipment_name"] == "CMP研磨装置"


# ---- R4-FUZZ-4: 計算結果の無い数式を「値が空」と言わない --------------------------------------------------

def test_formula_without_a_cached_value_gets_its_own_warning(tmp_path):
    from openpyxl import Workbook

    from forms import UNCACHED_FORMULA_WARNING
    from forms import load_workbook_info

    wb = Workbook()
    ws = wb.active
    ws.title = "報告書"
    ws["A1"], ws["B1"] = "設備名", '="CMP"&"装置"'  # openpyxl は計算結果を保存しない
    ws["A2"], ws["B2"] = "報告番号", None
    path = tmp_path / "formula.xlsx"
    wb.save(path)
    info = load_workbook_info(path)
    assert info.uncached_formulas == {"報告書": {(1, 2)}}
    values, fields = _values(info, _fields(("equipment_name", "string", "設備名"), ("report_id", "string", "報告番号")))
    assert values["equipment_name"] is None
    assert fields["equipment_name"]["warning"] == UNCACHED_FORMULA_WARNING
    assert fields["report_id"]["warning"] == "ラベルはありますが値が空です"


# ====================================================================================================
# 元 tests/test_forms_fixes5.py
# 帳票まわりの修正（5巡目の確認で見つかった不具合）の回帰テスト。
# ====================================================================================================

# ---- R5F-2: 読み取り結果の入力欄で Enter を押しても、何も送信されない ----------------------------------

def test_enter_in_the_review_form_submits_nothing(app):
    doc_id = _confirmed_doc(app, "1.xlsx")
    html = _review_html(app, doc_id)
    start = html.index(f'id="reviewForm-{doc_id}"')
    form = html[start:html.index("</form>", start)]
    # 暗黙の送信（Enter）を止める。送信ボタンも action も持たせない（値は途中保存で送る）
    assert f'<form id="reviewForm-{doc_id}" data-review-form onsubmit="return false">' in html
    assert 'type="submit"' not in form and "formaction" not in form and "action=" not in form


# ---- R5-MD-5: タイトル項目が設備だけのとき、出典にもタイトルに足した番号を書く ------------------------

def test_source_line_names_the_work_number_when_the_title_is_equipment_only():
    from forms import build_markdown

    doc = {"id": 1, "file_name": "点検.xlsx", "file_hash": "0" * 64}
    md = build_markdown(doc, _inspection("W-0101", "2026-04-01"))
    assert md.split("\n", 1)[0].endswith("｜W-0101")
    assert "- 出典: 点検.xlsx（作業No. W-0101）" in md.split("\n")
    # 番号が無ければ日付、どちらも無ければ今までどおり元ファイル名だけ
    assert "- 出典: 点検.xlsx（作業日 2026-04-08）" in build_markdown(doc, _inspection("", "2026-04-08")).split("\n")
    assert "- 出典: 点検.xlsx" in build_markdown(doc, _inspection("", None)).split("\n")


# ====================================================================================================
# 元 tests/test_forms_fixes6.py
# 帳票まわりの修正（6巡目の確認で見つかった不具合）の回帰テスト。
# ====================================================================================================

FILL = PatternFill("solid", fgColor="DDDDDD")
BOLD = Font(bold=True)


def _head(ws, coord, text):
    ws[coord] = text
    ws[coord].font = BOLD
    ws[coord].fill = FILL


def _label(ws, coord, text, value_coord, value):
    ws[coord] = text
    ws[coord].fill = FILL
    ws[value_coord] = value


def _pattern_forms_fixes6(*fields):
    return PatternDef(id=1, name="T", version="v1", description="", status="active", image_processing="none",
                      sheets=[SheetDef("報告書")], fields=list(fields))


def _field_forms_fixes6(ex, name):
    return next(f for f in ex["fields"] if f["field_name"] == name)


# ---- R6F-1: 右上の区画・区画の中の小見出し ------------------------------------------------------

def _side_by_side_forms_fixes6(path, *, issuer=True):
    """右上に「【回答欄】」、その下の行の左に「■ 発行部署記入欄」がある版。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "報告書"
    ws["A1"] = "工程異常連絡書"
    ws["A1"].font = BOLD
    ws["N1"] = "【回答欄】"
    ws["N1"].font = BOLD
    _head(ws, "A3", "■ 発行部署記入欄")
    if issuer:
        _label(ws, "A6", "処置", "B6", "発行側の処置（仮）")
    _label(ws, "N6", "処置", "O6", "回答側の処置（確定）")
    wb.save(path)
    return load_workbook_info(path)


def test_a_right_hand_section_started_higher_keeps_its_columns(tmp_path):
    info = _side_by_side_forms_fixes6(tmp_path / "cols.xlsx")
    grid = info.grids["報告書"]
    assert section_of(grid, grid.cells[(6, 14)]) == "回答"
    assert section_of(grid, grid.cells[(6, 1)]) == "発行部署記入"
    ex = extract_document(info, _pattern_forms_fixes6(FieldDef("action", "処置", ["処置"], section="回答")), ["報告書"])
    assert _field_forms_fixes6(ex, "action")["value"] == "回答側の処置（確定）"


def test_learning_finds_the_right_hand_answer_section(tmp_path):
    from forms import _learn_sections

    infos = [_side_by_side_forms_fixes6(tmp_path / "both.xlsx"), _side_by_side_forms_fixes6(tmp_path / "answer.xlsx", issuer=False)]
    row = {"use": True, "data_type": "string", "field_name": "action", "display_name": "処置",
           "candidates": "処置", "section": ""}
    _learn_sections(infos, {"報告書"}, [row])
    assert row["section"] == "回答"


def _nested(path):
    """「▼ 回答欄」の下に小見出し「1. 暫定対策」があり、その中に回答側の欄がある版。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "報告書"
    _head(ws, "A1", "■ 発行側")
    _label(ws, "A2", "処置内容", "B2", "発行側の処置")
    _head(ws, "A4", "▼ 回答欄")
    _head(ws, "A5", "1. 暫定対策")
    _label(ws, "A6", "処置内容", "B6", "回答側の処置")
    wb.save(path)
    return load_workbook_info(path)


def test_a_sub_heading_inside_the_answer_section_is_still_the_answer_section(tmp_path):
    info = _nested(tmp_path / "nested.xlsx")
    grid = info.grids["報告書"]
    assert sections_of(grid, grid.cells[(6, 1)]) == ["暫定対策", "回答"]
    assert sections_of(grid, grid.cells[(2, 1)]) == ["発行側"]   # 同じ階層の見出しで上の区画は終わる
    field = FieldDef("action", "処置内容", ["処置内容"], section=_section_value("回答欄"))
    assert _field_forms_fixes6(extract_document(info, _pattern_forms_fixes6(field), ["報告書"]), "action")["value"] == "回答側の処置"
    # 外側の区画の名前で決めた項目は、内側の小見出しの中の欄だけを見ない（発行側の欄は区画の外）
    issuer = FieldDef("action", "処置内容", ["処置内容"], section="発行側")
    assert _field_forms_fixes6(extract_document(info, _pattern_forms_fixes6(issuer), ["報告書"]), "action")["value"] == "発行側の処置"


# ---- R6F-2: 明細表の項目の区画 -----------------------------------------------------------------

def _two_tables(path, *, answer_label="使用部品"):
    wb = Workbook()
    ws = wb.active
    ws.title = "報告書"
    _head(ws, "A1", "■ 発行側")
    _head(ws, "A3", "使用部品")
    for c, t in zip("ABC", ["品番", "品名", "数量"]):
        _head(ws, f"{c}4", t)
    ws["A5"], ws["B5"], ws["C5"] = "P-1", "発行側の部品", 2
    _head(ws, "A8", "▼ 回答欄")
    _head(ws, "A10", answer_label)
    for c, t in zip("ABC", ["品番", "品名", "数量"]):
        _head(ws, f"{c}11", t)
    ws["A12"], ws["B12"], ws["C12"] = "P-9", "回答側の部品", 5
    wb.save(path)
    return load_workbook_info(path)


def _parts(section):
    return FieldDef("parts", "使用部品", ["使用部品"], data_type="table", table_columns=["品番", "品名", "数量"],
                    section=section)


def test_a_table_field_reads_the_table_in_its_section(tmp_path):
    info = _two_tables(tmp_path / "sec.xlsx")
    parts = _field_forms_fixes6(extract_document(info, _pattern_forms_fixes6(_parts(_section_value("▼ 回答欄"))), ["報告書"]), "parts")
    assert parts["value"]["rows"] == [["P-9", "回答側の部品", "5"]]
    # 区画を決めていなければ今までどおり上の表
    parts = _field_forms_fixes6(extract_document(info, _pattern_forms_fixes6(_parts("")), ["報告書"]), "parts")
    assert parts["value"]["rows"] == [["P-1", "発行側の部品", "2"]]


def test_a_table_found_by_its_columns_prefers_the_one_in_the_section(tmp_path):
    # 回答欄の表の見出しの書き方が違う版: 列見出しで探すときも区画の中の表を選ぶ
    info = _two_tables(tmp_path / "cols.xlsx", answer_label="交換部品")
    field = _parts(_section_value("▼ 回答欄"))
    field.candidates = ["部品明細"]   # どちらの見出しにも当たらない
    parts = _field_forms_fixes6(extract_document(info, _pattern_forms_fixes6(field), ["報告書"]), "parts")
    assert parts["value"]["rows"] == [["P-9", "回答側の部品", "5"]]


# ---- R6-F2: 1件だけになったまとめ取り込み -------------------------------------------------------

def test_a_batch_left_with_one_form_is_no_longer_a_batch(app, client):
    """取り込み失敗・削除で帳票が1件だけになったまとまりは、1件の帳票として扱う。

    「確定済みの分だけ zip にします」「残りの帳票は…」はどれも嘘になるため。
    """
    import html

    first = _confirmed_doc(app, "1.xlsx", batch_id="B9", order=0)
    second = _confirmed_doc(app, "2.xlsx", batch_id="B9", order=1)

    def finish(ids):
        url = "/forms/finish?ids=" + ",".join(str(i) for i in ids) + f"&current={first}"
        return html.unescape(client.get(url).get_json()["html"])

    page = finish([first, second])
    assert "zip" in page and "/forms/batches/B9/download.zip" in page

    client.post(f"/forms/{second}/delete")
    page = finish([first, second])
    assert "zip" not in page and "/forms/batches/" not in page
    assert "Markdown をダウンロード（.md）" in page and f"/forms/{first}/download.md" in page


# ---- R6-B1: 見出しからは日付と分からない欄（「発生」）の型 ----------------------------------------

def test_a_date_value_under_a_non_date_label_becomes_a_date_field(tmp_path):
    """値が日付だけで書かれていれば日付型にする（文字列のままだと md の日付の書き方が混ざる）。"""
    from forms import _guess_type

    wb = Workbook()
    ws = wb.active
    ws.title = "報告書"
    _label(ws, "A1", "発生", "B1", "2024年7月28日 22:46")
    _label(ws, "A2", "復旧", "B2", "R5.11.16")
    path = tmp_path / "occurred.xlsx"
    wb.save(path)
    types = {r["display_name"]: r["data_type"] for r in suggest_rows([load_workbook_info(path)])[1]}
    assert types["発生"] == "date" and types["復旧"] == "date"

    # 日付を含むだけの値・年の無い値・識別番号は日付にしない
    assert _guess_type("R2026-00123", "報告No") == "string"
    assert _guess_type("2026-09-14-3", "ロットNo") == "string"
    assert _guess_type("2026-09-14 に復旧", "処置") == "string"
    assert _guess_type("2/12", "発生") == "string"
    assert _guess_type("12:30", "発生") == "string"


# ====================================================================================================
# 元 tests/test_forms_fixes7.py
# 帳票取り込みの直し（まとめ取り込みの件数・シートのチェック・見回り・ダウンロード）。
#
# このまわりで見つかった不具合:
#   - ②へ戻るとシートのチェックが先頭ファイル分だけになり、読み取り直すと帳票が黙って消える
#   - ④が③に並んでいない帳票まで数え、案内の件数と zip の中身が合わない
#   - 見回りが「取り込んだ時刻」で切るので、2時間かけて確認すると編集中の帳票が消える
#   - 修正中の帳票をリンクから直接ダウンロードすると、直した値が入らないまま消える
# ====================================================================================================

def _type_fragment(client, ids) -> str:
    return client.get("/forms/type?ids=" + ",".join(str(i) for i in ids)).get_json()["html"]


def _checked_sheets(html: str) -> list[str]:
    """②「読み取るシート」でチェックが入っているシートの名前。"""
    out = []
    for part in html.split('name="sheets"')[1:]:
        tag = part.split(">")[0]
        if " checked" in tag:
            out.append(tag.split('value="')[1].split('"')[0])
    return out


# ---- ② シートのチェックは、読み取ったファイル全部を合わせる -------------------------------------

def test_reopening_the_type_step_keeps_every_sheet_that_was_read(app, client, sample_dir, tmp_path):
    """②へ戻っても、読み取ったシートのチェックが先頭ファイル分に減らない。

    1件分の読み取り結果には、そのファイルに在ったシートしか残らない。先頭のファイルだけを見ると
    ほかのファイルのシートのチェックが外れ、そのまま読み取り直すと大半の帳票が③から消えていた。
    """
    pattern_id = _pattern_forms_batch_review(client, sample_dir)
    # シートの名前が違う2件（standard は「修理報告書」、shifted は「修理報告書(2)」）
    ids = upload_forms(client, sample_dir / "standard.xlsx", sample_dir / "shifted.xlsx")
    chosen = ["修理報告書", "修理報告書(2)"]
    assert read_all(client, ids, pattern_id, chosen).status_code == 200

    checked = _checked_sheets(_type_fragment(client, ids))
    assert sorted(checked) == sorted(chosen)

    # そのまま読み取り直しても2件とも読める（前は先頭ファイルのシートだけになり1件に減っていた）
    body = read_all(client, ids, pattern_id, checked, acknowledge="on").get_json()
    assert [d["id"] for d in body["docs"]] == ids and not body.get("errors")


def test_a_sheet_the_user_unchecked_stays_unchecked(app, client, sample_dir, tmp_path):
    """手で外したシートは、②へ戻ったときにチェックが戻らない。"""
    pattern_id = _pattern_forms_batch_review(client, sample_dir)
    ids = upload_forms(client, *_copies(sample_dir, tmp_path, 2))

    first = _type_fragment(client, ids)
    assert "参考資料" in first                      # 見本のブックには参考資料シートもある
    assert read_all(client, ids, pattern_id, ["修理報告書"]).status_code == 200

    assert _checked_sheets(_type_fragment(client, ids)) == ["修理報告書"]


# ---- ④ 数えるのは zip に入る帳票だけ -----------------------------------------------------------

def test_the_finish_step_counts_only_the_forms_that_can_reach_the_zip(app, client, sample_dir, tmp_path):
    """読み取れなかったファイルを件数に入れない（案内の件数と zip の中身を合わせる）。"""
    pattern_id = _pattern_forms_batch_review(client, sample_dir)
    paths = _copies(sample_dir, tmp_path, 2)
    inspection = tmp_path / "点検記録表.xlsx"
    inspection.write_bytes((sample_dir / "inspection.xlsx").read_bytes())
    ids = upload_forms(client, inspection, *paths)

    body = read_all(client, ids, pattern_id, ["修理報告書"]).get_json()
    assert body["errors"] == ["点検記録表.xlsx: 選んだシートがファイルにありません"]

    state = finish(client, ids)
    assert state["total"] == 2                       # 置いたのは3件だが、zip に入るのは2件
    page = state["html"]
    assert "確定済み <strong>0</strong> / 2 件" in page
    assert "残り2件も確定して、まとめてダウンロード（zip）" in page
    # 読み取れなかったファイルは、入らないことを名前を出して知らせる
    assert "まだ読み取れていない1件（zip には入りません）" in page and "点検記録表.xlsx" in page
    assert "サーバーからすべて消えます" not in page and "サーバーに残ります" in page
    # ③に並んでいない帳票には［読み取り結果を見る］を出さない（押しても何も起きないボタンを出さない）
    assert page.count("data-goto-doc") == 2

    for doc_id in [d["id"] for d in body["docs"]]:
        assert client.post(f"/forms/{doc_id}/confirm", json={}).get_json()["ok"] is True
    with app.app_context():
        batch_id = db.get_document(ids[0])["batch_id"]
    res = client.get(f"/forms/batches/{batch_id}/download.zip")
    assert res.status_code == 200
    assert len(zipfile.ZipFile(io.BytesIO(res.data)).namelist()) == 2   # 案内どおりの件数


def test_a_form_that_fails_a_second_read_drops_out_of_the_count(app, client, sample_dir, tmp_path):
    """一度読めた帳票が読み取り直しで読めなくなったら、③から消えるのと一緒に件数からも外す。"""
    pattern_id = _pattern_forms_batch_review(client, sample_dir)
    ids = upload_forms(client, sample_dir / "standard.xlsx", sample_dir / "shifted.xlsx")
    assert read_all(client, ids, pattern_id, ["修理報告書", "修理報告書(2)"]).status_code == 200
    assert finish(client, ids)["total"] == 2

    # 「修理報告書」だけで読み直す → shifted はそのシートが無いので読めない
    body = read_all(client, ids, pattern_id, ["修理報告書"], acknowledge="on").get_json()
    assert [d["id"] for d in body["docs"]] == [ids[0]]
    with app.app_context():
        assert db.get_document(ids[1])["data_json"] is None     # 前の読み取り結果は残さない
    assert finish(client, ids)["total"] == 1


# ---- 見回りは「最後にさわった時刻」で切る -------------------------------------------------------

def _age(app, doc_ids, hours: float, column: str = "created_at") -> None:
    stamp = (datetime.now() - timedelta(hours=hours)).isoformat(timespec="seconds")
    with app.app_context():
        for doc_id in doc_ids:
            db.update_document(doc_id, **{column: stamp})


def test_the_sweeper_keeps_a_form_that_is_still_being_corrected(app, client, sample_dir, tmp_path):
    """途中保存をしている帳票は、取り込みから2時間たっても捨てない。

    帳票には updated_at が無く、取り込んだ時刻で切られていたので、50件を2時間かけて確認すると
    手で直した値ごと消えていた（一覧表は updated_at を持っていて、約束どおり動いていた）。
    """
    pattern_id = _pattern_forms_batch_review(client, sample_dir)
    ids = upload_forms(client, *_copies(sample_dir, tmp_path, 3))
    read_all(client, ids, pattern_id, ["修理報告書"])
    _age(app, ids, 2.5)                                   # 2時間半前に取り込んだ

    for doc_id in ids:                                    # ずっと手で直している
        assert client.post(f"/forms/{doc_id}/draft",
                           json={"values": {"equipment_name": "直した設備名"}}).status_code == 204
    with app.app_context():
        assert core.sweep_stale(2) == (0, 0)
        assert all(db.get_document(i) is not None for i in ids)


def test_the_sweeper_keeps_a_batch_together_while_one_form_is_fresh(app, client, sample_dir, tmp_path):
    """まとまりのどれか1件でもさわられていれば、まとまりごと残す（zip が欠けないように）。"""
    pattern_id = _pattern_forms_batch_review(client, sample_dir)
    ids = upload_forms(client, *_copies(sample_dir, tmp_path, 3))
    read_all(client, ids, pattern_id, ["修理報告書"])
    _age(app, ids, 2.5)
    _age(app, ids, 2.5, column="updated_at")

    # いちばん下の1件だけをいま直した（上から順に見ていくと、まだ手が届いていない分がこうなる）
    assert client.post(f"/forms/{ids[2]}/draft", json={"values": {"equipment_name": "直した"}}).status_code == 204
    with app.app_context():
        assert core.sweep_stale(2) == (0, 0)
        assert all(db.get_document(i) is not None for i in ids)

    # まとまりのどれもさわられなくなれば、これまでどおり捨てる
    _age(app, ids, 2.5, column="updated_at")
    with app.app_context():
        assert core.sweep_stale(2) == (3, 0)
        assert all(db.get_document(i) is None for i in ids)


# ---- 修正中の帳票は、直した値で渡す -------------------------------------------------------------

def test_a_modified_form_downloads_with_the_corrected_value(app, client, sample_dir):
    """画面の JS を通さずにリンクを開いても、直した値が .md に入る（design.md 3.3）。"""
    pattern_id = create_type(client, sample_dir / "standard.xlsx", "設備修理報告書")
    add_field(client, pattern_id, "修理報告書", "E4", "F4")        # 設備名
    activate(client, pattern_id)
    doc_id, = upload_forms(client, sample_dir / "standard.xlsx")
    read_all(client, [doc_id], pattern_id, ["修理報告書"])

    assert client.post(f"/forms/{doc_id}/confirm", json={}).get_json()["ok"] is True
    assert client.post(f"/forms/{doc_id}/draft",
                       json={"values": {"equipment_name": "直した設備名"}}).status_code == 204
    with app.app_context():
        assert db.get_document(doc_id)["state"] == "modified"

    md = client.get(f"/forms/{doc_id}/download.md").get_data(as_text=True)
    assert "直した設備名" in md


def test_a_modified_form_in_a_batch_zip_carries_the_corrected_value(app, client, sample_dir, tmp_path):
    """まとまりの zip も同じ（中クリックなどで JS を通さずに開いたとき）。"""
    pattern_id = _pattern_forms_batch_review(client, sample_dir)
    ids = upload_forms(client, *_copies(sample_dir, tmp_path, 2))
    read_all(client, ids, pattern_id, ["修理報告書"])
    for doc_id in ids:
        assert client.post(f"/forms/{doc_id}/confirm", json={}).get_json()["ok"] is True
    assert client.post(f"/forms/{ids[0]}/draft",
                       json={"values": {"equipment_name": "直した設備名"}}).status_code == 204

    with app.app_context():
        batch_id = db.get_document(ids[0])["batch_id"]
    res = client.get(f"/forms/batches/{batch_id}/download.zip")
    assert res.status_code == 200
    with zipfile.ZipFile(io.BytesIO(res.data)) as zf:
        bodies = [zf.read(n).decode("utf-8") for n in zf.namelist()]
    assert any("直した設備名" in b for b in bodies)


# ---- ② は置かれたブックを1件ずつ開く（50件でメモリを食いつぶさない） -------------------------------

def test_the_type_step_opens_one_workbook_at_a_time(app, client, sample_dir, tmp_path, monkeypatch):
    """置かれたファイル全部のブックを同時にメモリへ広げない（50件でサーバーが落ちないように）。

    1件ぶんの WorkbookInfo はセルごとの控えを持つので大きい。全部を持ち続けると件数に比例して
    増え、50件で数百MB〜数GBになっていた。
    """
    import gc
    import weakref

    import views as forms_view

    seen: list[weakref.ref] = []
    alive_during: list[int] = []
    original = forms_view._load_info

    def watching(doc):
        gc.collect()
        alive_during.append(sum(1 for ref in seen if ref() is not None))
        info = original(doc)
        seen.append(weakref.ref(info))
        return info

    _pattern_forms_batch_review(client, sample_dir)
    ids = upload_forms(client, *_copies(sample_dir, tmp_path, 4))
    monkeypatch.setattr(forms_view, "_load_info", watching)
    body = client.get("/forms/type?ids=" + ",".join(str(i) for i in ids)).get_json()

    assert len(seen) == 4 and len(body["docs"]) == 4
    # 次の1件を開くとき、前に開いたブックはもう残っていない
    assert alive_during == [0, 0, 0, 0]
    # 画面に出す中身はこれまでどおり（シートは名前ごとに1行にまとまる。4件ぶん並ばない）
    assert body["html"].count('name="sheets"') == 3 and "4件をまとめて読み取る" in body["html"]


# ====================================================================================================
# 元 tests/test_form_types_clicks.py
# 帳票の種類をセルのクリックで作る: クリック→項目の組み立てと、その画面の動き。
# ====================================================================================================

def _book_form_types_clicks(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "報告書"
    fill = PatternFill("solid", fgColor="FFD9E1F2")
    for coord, value in {"A1": "報告番号", "B1": "R-001", "A2": "設備番号", "B2": "EQ-001",
                         "A3": "作業時間(h)", "B3": 2.5, "A4": "所見", "A5": "異音あり",
                         "C1": "ライン：L6", "A7": "■ 交換部品",
                         "A8": "品番", "B8": "品名", "C8": "数量",
                         "A9": "PW-1", "B9": "ベアリング", "C9": 2,
                         "A10": "PW-2", "B10": "モータ", "C10": 1}.items():
        ws[coord] = value
    for coord in ("A1", "A2", "A3", "A4", "A7", "A8", "B8", "C8"):
        ws[coord].fill = fill
    wb.save(path)
    return path


@pytest.fixture()
def grid(tmp_path):
    return load_workbook_info(_book_form_types_clicks(tmp_path / "click.xlsx")).grids["報告書"]


def test_click_label_then_value_makes_a_field(grid):
    row, error = click_field(grid, "A1", "B1")
    assert error == "" and row["field_name"] == "report_id" and row["display_name"] == "報告番号"
    # 探す見出しはクリックした語＋辞書の同じ意味の語。値の位置・セル番地も控える
    assert row["candidates"].splitlines()[0] == "報告番号" and len(row["candidates"].splitlines()) > 1
    assert row["direction"] == "right" and row["label_cell"] == "A1" and row["cell"] == "B1"
    assert row["sheet_name"] == "報告書" and row["examples"] == ["R-001"]


def test_click_derives_type_and_unit_from_the_value(grid):
    row, _ = click_field(grid, "A3", "B3")
    assert row["data_type"] == "number" and row["unit"] == "時間" and row["display_name"] == "作業時間"
    below, _ = click_field(grid, "A4", "A5")
    assert below["direction"] == "below" and below["data_type"] == "string"


def test_clicking_the_same_cell_twice(grid):
    inline, _ = click_field(grid, "C1", "C1")   # 「ライン：L6」はセル内の見出しと値
    assert inline["direction"] == "same_cell" and inline["display_name"] == "ライン"
    only, _ = click_field(grid, "A5", "A5")     # 見出しのない値だけの項目
    assert only["candidates"] == "" and only["cell"] == "A5" and "A5" in only["display_name"]


def test_clicking_a_table_header_makes_one_table_field(grid):
    assert {"A7", "A8", "B8", "C8"} <= table_cells(grid)
    row, error = click_field(grid, "A8")
    assert error == "" and row["data_type"] == "table"
    assert row["table_columns"].splitlines() == ["品番", "品名", "数量"]


def test_a_table_filled_in_once_in_the_sample_is_still_a_table(tmp_path):
    """見本での記入が1行だけの表も、列見出しのすぐ上の見出しを指せば明細表になる。"""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "報告書"
    for coord, value in {"A1": "設備修理報告書", "F1": "承認", "G1": "確認",   # 帳票のタイトルと押印欄
                         "F2": "高橋", "G2": "中村", "F3": "7/5", "G3": "7/5",
                         "A5": "使用部品", "A6": "品番", "B6": "品名", "C6": "数量",
                         "A7": "PM26-1778", "B7": "センサ", "C7": 1}.items():
        ws[coord] = value
    fill = PatternFill("solid", fgColor="FFD9E1F2")
    for coord in ("A5", "A6", "B6", "C6", "F1", "G1"):
        ws[coord].fill = fill
    ws.merge_cells("A5:C5")   # 表の見出しは表の幅いっぱいに結合されている
    wb.save(tmp_path / "one_row.xlsx")
    grid = load_workbook_info(tmp_path / "one_row.xlsx").grids["報告書"]

    row, error = click_field(grid, "A5")
    assert error == "" and row["data_type"] == "table" and row["display_name"] == "使用部品"
    assert row["table_columns"].splitlines() == ["品番", "品名", "数量"]
    assert "A5" in table_cells(grid)
    # 帳票のタイトルは、下の押印欄から離れているので表の見出しにしない
    assert "A1" not in table_cells(grid)
    assert click_field(grid, "A1", "B1")[0]["data_type"] != "table"


def test_clicking_an_empty_cell_as_a_label_is_refused(grid):
    row, error = click_field(grid, "F20")
    assert row is None and "見出し" in error
    assert click_field(grid, "zz")[0] is None


def test_label_without_a_value_can_be_registered(grid):
    row, error = click_field(grid, "A2", "")
    assert error == "" and row["cell"] == "" and row["candidates"].startswith("設備番号")


# ---- 画面 ---------------------------------------------------------------------------

# 置いた Excel はサーバーに残らないので、画面（ブラウザ）と同じように毎回一緒に送る

def _part(path):
    return (io.BytesIO(path.read_bytes()), path.name)


def _create(client, path, name="点検報告書"):
    res = client.post("/form-types/new", data={"name": name, "book": _part(path)},
                      content_type="multipart/form-data")
    assert res.status_code == 200, res.get_data(as_text=True)
    return res.get_json()["pattern_id"]


def _panel(client, pattern_id, path=None) -> str:
    """項目の一覧（HTML の断片）。Excel を渡すとシートと読み取りテストも出る。"""
    if path is None:
        return client.get(f"/form-types/{pattern_id}/panel").get_json()["html"]
    res = client.post(f"/form-types/{pattern_id}/panel", data={"book": _part(path)},
                      content_type="multipart/form-data")
    assert res.status_code == 200, res.get_data(as_text=True)
    return res.get_json()["html"]


def test_build_panel_adds_and_deletes_fields_by_clicking(app, client, tmp_path):
    path = _book_form_types_clicks(tmp_path / "click.xlsx")
    pattern_id = _create(client, path)
    page = client.get("/form-types/").get_data(as_text=True)
    assert "点検報告書" in page and "app.js" in page
    panel = _panel(client, pattern_id, path)
    assert 'data-cell="A1"' in panel and "まだ項目がありません" in panel
    # クリックを送るための仕掛け（app.js の帳票登録の部分が使う）と、明細表の印
    assert 'id="cellBuilder"' in panel and "data-click-hint" in panel
    assert 'data-table-head="1"' in panel

    data = {"sheet": "報告書", "label_cell": "A1", "value_cell": "B1", "book": _part(path)}
    body = client.post(f"/form-types/{pattern_id}/fields", data=data,
                       content_type="multipart/form-data").get_json()
    assert "「報告番号」を項目にしました" in body["message"] and "R-001" in body["html"]
    assert "点検報告書" in body["list_html"]
    with app.app_context():
        pattern = db.load_pattern(pattern_id)
    assert [f.field_name for f in pattern.fields] == ["report_id"]
    assert pattern.title_fields == ["report_id"]          # タイトル項目は自動で決める
    assert [s.sheet_name for s in pattern.sheets] == ["報告書"]
    assert pattern.fields[0].cell == "B1" and pattern.fields[0].sheet_name == "報告書"

    # 同じセルをもう一度クリックしても増えない
    data["book"] = _part(path)
    body = client.post(f"/form-types/{pattern_id}/fields", data=data,
                       content_type="multipart/form-data").get_json()
    assert body["message"] == "そのセルはもう項目になっています"

    # 明細表は列見出しを1回クリックするだけ
    client.post(f"/form-types/{pattern_id}/fields", content_type="multipart/form-data",
                data={"sheet": "報告書", "label_cell": "A8", "value_cell": "", "book": _part(path)})
    with app.app_context():
        table = [f for f in db.load_pattern(pattern_id).fields if f.data_type == "table"]
    assert len(table) == 1 and table[0].table_columns == ["品番", "品名", "数量"]

    body = client.post(f"/form-types/{pattern_id}/fields/report_id/delete").get_json()
    assert body["message"] == "項目を削除しました"
    with app.app_context():
        assert [f.field_name for f in db.load_pattern(pattern_id).fields] == ["parts"]
    assert client.post(f"/form-types/{pattern_id}/fields/zzz/delete").status_code == 404

    # 使用開始を押すまでは作成中のまま
    with app.app_context():
        assert db.load_pattern(pattern_id).status == "draft"
    assert client.post(f"/form-types/{pattern_id}/status", json={"status": "active"}).get_json()["status"] == "active"


def test_a_type_without_fields_cannot_be_used(app, client, tmp_path):
    pattern_id = _create(client, _book_form_types_clicks(tmp_path / "click.xlsx"))
    res = client.post(f"/form-types/{pattern_id}/status", json={"status": "active"})
    assert res.status_code == 400 and "読み取る項目がありません" in res.get_json()["error"]
    with app.app_context():
        assert db.load_pattern(pattern_id).status == "draft"


def test_the_name_comes_from_the_file_name_and_can_be_changed(app, client, tmp_path):
    path = _book_form_types_clicks(tmp_path / "設備点検表.xlsx")
    res = client.post("/form-types/new", data={"name": "", "book": _part(path)},
                      content_type="multipart/form-data")
    pattern_id = res.get_json()["pattern_id"]
    with app.app_context():
        assert db.load_pattern(pattern_id).name == "設備点検表"
    assert client.post(f"/form-types/{pattern_id}/name", json={"name": "点検表"}).get_json()["ok"] is True
    with app.app_context():
        assert db.load_pattern(pattern_id).name == "点検表"
    res = client.post(f"/form-types/{pattern_id}/name", json={"name": " "})
    assert res.status_code == 400 and "名前を入れてください" in res.get_json()["error"]


def test_a_field_label_can_be_typed_by_hand(app, client, tmp_path):
    """読み取る項目の見出しは手で直せる。直るのは書き出す名前だけで、探す見出しとセルは変えない。"""
    path = _book_form_types_clicks(tmp_path / "click.xlsx")
    pattern_id = _create(client, path)
    _click(client, pattern_id, path, "A1", "B1")
    _click(client, pattern_id, path, "A2", "B2")

    body = client.post(f"/form-types/{pattern_id}/fields/report_id/label",
                       data={"name": "受付番号", "book": _part(path)},
                       content_type="multipart/form-data").get_json()
    assert body["message"] == "見出しを「受付番号」にしました"
    assert "- 受付番号: R-001" in body["html"]      # 読み取りテストの Markdown に新しい名前で書き出す
    with app.app_context():
        field = db.load_pattern(pattern_id).fields[0]
    assert field.display_name == "受付番号"
    assert field.candidates[0] == "報告番号" and field.cell == "B1"   # 探す先は変わらない

    # Excel を置いていなくても、直した見出しのまま。置いた帳票での書き方は「探す見出し」として残る
    panel = _panel(client, pattern_id)
    assert 'value="受付番号"' in panel and "探す見出し: 報告番号" in panel

    # ほかの項目が書き出す名前とは重ねられない（空の見出しも断る）。どちらも元の見出しのまま残る
    res = client.post(f"/form-types/{pattern_id}/fields/report_id/label", json={"name": "設備番号"})
    assert res.status_code == 400 and "ほかの項目が使っています" in res.get_json()["error"]
    assert client.post(f"/form-types/{pattern_id}/fields/report_id/label", json={"name": " "}).status_code == 400
    with app.app_context():
        assert [f.display_name for f in db.load_pattern(pattern_id).fields] == ["受付番号", "設備番号"]
    assert client.post(f"/form-types/{pattern_id}/fields/zzz/label", json={"name": "x"}).status_code == 404


def test_a_field_found_by_its_cell_when_the_label_is_missing(app, client, tmp_path):
    """見出しの無い「値だけ」の項目は、クリックしたセルの番地から読む。"""
    path = _book_form_types_clicks(tmp_path / "click.xlsx")
    pattern_id = _create(client, path)
    client.post(f"/form-types/{pattern_id}/fields", content_type="multipart/form-data",
                data={"sheet": "報告書", "label_cell": "A5", "value_cell": "A5", "book": _part(path)})
    assert "異音あり" in _panel(client, pattern_id, path)     # 読み取りテストは同じ欄に出る


# ---- 版によって書き方が違う帳票 -------------------------------------------------------

def _versions(tmp_path):
    """見本（v1: 見出しの右に値）と、書き方の違う同じ帳票（v2: セル内に「設備番号：…」）。"""
    from openpyxl import Workbook
    v1 = Workbook()
    ws = v1.active
    ws.title = "報告書"
    for coord, value in {"A1": "報告番号", "B1": "R-001",
                         "A2": "使用設備", "B2": "CMP-108　STI-CMP 8号機",
                         "A3": "所見", "B3": "異音あり"}.items():
        ws[coord] = value
    v1.save(tmp_path / "v1.xlsx")

    v2 = Workbook()
    ws2 = v2.active
    ws2.title = "報告書"
    # 同じ番地には別の欄の見出しが来る。設備は1つのセルに「設備番号：…」と書かれている
    for coord, value in {"A1": "報告番号", "B1": "R-002",
                         "A2": "所見", "B2": "発生日時",
                         "A4": "設備番号　：IMP-603", "A5": "所見", "B5": "振動あり"}.items():
        ws2[coord] = value
    v2.save(tmp_path / "v2.xlsx")
    return tmp_path / "v1.xlsx", tmp_path / "v2.xlsx"


def test_one_cell_holding_an_equipment_number_and_name_becomes_two_fields(tmp_path):
    """「使用設備：CMP-108　STI-CMP 8号機」を1回クリックすると、設備番号と設備名の2項目になる。"""
    from forms import split_rows

    v1, _ = _versions(tmp_path)
    grid = load_workbook_info(v1).grids["報告書"]
    row, error = click_field(grid, "A2", "B2")
    assert error == ""
    parts = split_rows(row)
    assert [p["field_name"] for p in parts] == ["equipment_id", "equipment_name"]
    # 文字は消さず、同じセルの読む場所を分けるだけ。書き方の違う見出しは全部入れておく
    assert [p["examples"] for p in parts] == [["CMP-108"], ["STI-CMP 8号機"]]
    assert all(p["cell"] == "B2" for p in parts)
    assert "使用設備" in parts[0]["candidates"] and "設備番号" in parts[0]["candidates"]
    # 番号と名前に分けられない欄はそのまま1項目
    plain, _ = click_field(grid, "A1", "B1")
    assert split_rows(plain) == [plain]


def test_a_differently_written_version_is_read_by_the_label_not_by_the_cell_address(tmp_path):
    """見本と書き方が違う帳票でも、セル内の見出しから読む。別の欄の見出しは値にしない。"""
    from forms import extract_document
    from forms import split_rows
    from forms import rows_to_pattern

    v1, v2 = _versions(tmp_path)
    grid = load_workbook_info(v1).grids["報告書"]
    rows = []
    for label, value in (("A2", "B2"), ("A3", "B3")):
        row, _ = click_field(grid, label, value, {r["field_name"] for r in rows})
        rows += split_rows(row, {r["field_name"] for r in rows})
    pattern = rows_to_pattern(1, {"name": "報告書", "version": "v1", "description": "",
                                  "image_processing": "none", "title_fields": [], "md_options": {}},
                              [{"use": True, "sheet_name": "報告書", "required": False}], rows)
    fields = {f["field_name"]: f for f in extract_document(load_workbook_info(v2), pattern, ["報告書"])["fields"]}
    # 見本では B2 だった番地に、この版では別の欄の見出し「発生日時」が来ている。それは値にしない
    assert fields["equipment_id"]["value"] == "IMP-603"
    assert fields["equipment_name"]["value"] in (None, "")
    assert fields["field_1"]["value"] == "振動あり"


def test_the_clicked_label_wins_over_a_dictionary_synonym_on_the_same_sheet(tmp_path):
    """「報告者」と「担当者」が並ぶ帳票で、クリックした方の値を読む。

    辞書は「担当者」の言い換えとして「報告者」「記入者」…も探す。その言い換えがシートの先に
    出てくると、クリックした欄ではなく言い換えの欄の値が読まれてしまっていた。
    """
    from forms import extract_document
    from forms import rows_to_pattern

    wb = Workbook()
    ws = wb.active
    ws.title = "報告書"
    fill = PatternFill("solid", fgColor="FFD9E1F2")
    for coord, value in {"A1": "報告者", "B1": "長谷川 聡", "C1": "担当者", "D1": "清水 彩花"}.items():
        ws[coord] = value
    for coord in ("A1", "C1"):
        ws[coord].fill = fill
    path = tmp_path / "二人.xlsx"
    wb.save(path)

    info = load_workbook_info(path)
    grid = info.grids["報告書"]
    meta = {"name": "報告書", "version": "v1", "description": "", "image_processing": "none",
            "title_fields": [], "md_options": {}}
    sheets = [{"use": True, "sheet_name": "報告書", "required": False}]
    for label_cell, value_cell, expected in (("C1", "D1", "清水 彩花"), ("A1", "B1", "長谷川 聡")):
        row, error = click_field(grid, label_cell, value_cell)
        assert error == ""
        # 辞書がもう一方の見出しも探す候補に入れている（この状況でクリックした方を読めること）
        assert "報告者" in row["candidates"] and "担当者" in row["candidates"]
        pattern = rows_to_pattern(1, meta, sheets, [row])
        field = extract_document(info, pattern, ["報告書"])["fields"][0]
        assert field["value"] == expected, (label_cell, field)


# ---- 同じ意味になる2つの欄（辞書の名前が同じだけの別の欄） -------------------------------------

def _two_people_book(path, equipment_label="設備No", equipment_row=4):
    """「担当者」と「報告者」が別のセルに並ぶ帳票（設備の見出しは版によって書き方も場所も違う）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "報告書"
    fill = PatternFill("solid", fgColor="FFD9E1F2")
    cells = {"A1": "報告番号", "B1": "R-001", "A2": "担当者", "B2": "清水 彩花",
             "A3": "報告者", "B3": "長谷川 聡",
             f"A{equipment_row}": equipment_label, f"B{equipment_row}": "EQ-001"}
    for coord, value in cells.items():
        ws[coord] = value
    for coord in ("A1", "A2", "A3", f"A{equipment_row}"):
        ws[coord].fill = fill
    wb.save(path)
    return path


def _click(client, pattern_id, path, label_cell, value_cell="", sheet="報告書"):
    """画面と同じクリック（いま見ている Excel を一緒に送る）。"""
    return client.post(f"/form-types/{pattern_id}/fields", content_type="multipart/form-data",
                       data={"sheet": sheet, "label_cell": label_cell, "value_cell": value_cell,
                             "book": _part(path)}).get_json()


def test_two_labels_of_the_same_meaning_on_one_sheet_become_two_fields(app, client, tmp_path):
    """同じ帳票の別のセル（「担当者」と「報告者」）は、辞書の名前が同じでもそれぞれ項目になる。"""
    path = _two_people_book(tmp_path / "二人.xlsx")
    pattern_id = _create(client, path)
    assert "「担当者」を項目にしました" in _click(client, pattern_id, path, "A2", "B2")["message"]
    body = _click(client, pattern_id, path, "A3", "B3")
    assert "「報告者」を「担当者」とは別の項目にしました" in body["message"]

    with app.app_context():
        fields = db.load_pattern(pattern_id).fields
    assert [f.field_name for f in fields] == ["reporter", "reporter_2"]
    # それぞれ自分の見出しを持ち、自分のセルの値を読む
    assert [f.candidates[0] for f in fields] == ["担当者", "報告者"]
    assert [f.display_name for f in fields] == ["担当者", "報告者"]
    assert [f.label_cell for f in fields] == ["A2", "A3"]
    panel = _panel(client, pattern_id, path)
    assert "清水 彩花" in panel and "長谷川 聡" in panel


def test_the_same_field_written_differently_in_another_book_is_still_merged(app, client, tmp_path):
    """書き方の違う同じ帳票（「設備No」と「設備番号」）に置き替えてクリックすると、見出しに足す。

    見本を何枚も預かる代わりに、画面で Excel を置き替えて同じ欄をクリックする。
    """
    first = _two_people_book(tmp_path / "v1.xlsx", "設備No")
    pattern_id = _create(client, first)
    other = _two_people_book(tmp_path / "v2.xlsx", "設備番号", equipment_row=6)
    # 別の書き方の帳票に置き替える（新しい種類は作らない）
    res = client.post(f"/form-types/{pattern_id}/panel", data={"book": _part(other)},
                      content_type="multipart/form-data")
    assert res.status_code == 200, res.get_data(as_text=True)
    assert "v2.xlsx" in res.get_json()["html"]

    _click(client, pattern_id, first, "A4", "B4")
    body = _click(client, pattern_id, other, "A6", "B6")
    assert "の見出しに「設備番号」を足しました" in body["message"]

    with app.app_context():
        fields = db.load_pattern(pattern_id).fields
    assert [f.field_name for f in fields] == ["equipment_id"]
    assert fields[0].candidates[:2] == ["設備No", "設備番号"]


# ====================================================================================================
# 元 tests/test_form_types_fixes7.py
# 帳票登録の直し（2枚目のシートの項目・手で直した見出し・項目を全部消したとき・Excel を置いていないとき）。
#
# このまわりで見つかった不具合:
#   - 同じ見出し語が2枚のシートにあると、2つ目の項目が1枚目の値をそのまま読む
#   - 2枚目のシートに登録した項目が、記入位置がずれた帳票では1枚目の同じ番地の別の欄になる
#   - 手で直した見出しが、次にセルをクリックしたときに元の見出しへ勝手に戻る
#   - 使用中の種類から項目を全部削除でき、中身の無い Markdown を作り続ける
#   - 使用開始の直後、登録した項目が全部「—（見つかりません）」になる
#
# 置いた Excel はサーバーに残らない（2026-09-21 の利用者の指示）ので、どの操作でも画面と同じように
# その Excel を一緒に送る。
# ====================================================================================================

def _create_form_types_fixes7(client, path, name="二次報告書") -> int:
    res = client.post("/form-types/new", data={"name": name, "book": book_part(path)},
                      content_type="multipart/form-data")
    assert res.status_code == 200, res.get_data(as_text=True)
    return res.get_json()["pattern_id"]


def _click_form_types_fixes7(client, pattern_id, path, sheet, label_cell, value_cell=""):
    res = client.post(f"/form-types/{pattern_id}/fields", content_type="multipart/form-data",
                      data={"sheet": sheet, "label_cell": label_cell, "value_cell": value_cell,
                            "book": book_part(path)})
    assert res.status_code == 200, res.get_data(as_text=True)
    return res.get_json()


def _panel_form_types_fixes7(client, pattern_id, path=None) -> str:
    return panel_html(client, pattern_id, book=path)


def _fill(ws, *coords) -> None:
    for coord in coords:
        ws[coord].fill = PatternFill("solid", fgColor="FFD9E1F2")


# ---- 同じ見出し語が2枚のシートにある帳票 ----------------------------------------------------

def _two_sheet_book(path, second_row: int = 2):
    """「独自見出し」が1次報告と2次報告の両方にある帳票（辞書に無い語なので別々の項目になる）。"""
    wb = Workbook()
    first = wb.active
    first.title = "1次報告"
    first["A3"], first["B3"] = "独自見出し", "1次の内容"
    _fill(first, "A3")
    second = wb.create_sheet("2次報告")
    second[f"A{second_row}"], second[f"B{second_row}"] = "独自見出し", "2次の内容"
    _fill(second, f"A{second_row}")
    wb.save(path)
    return path


def test_a_field_registered_on_the_second_sheet_reads_its_own_sheet(app, client, tmp_path):
    """2枚目のシートに登録した項目は、1枚目の同じ見出しの値で埋まらない。

    渡された順にシートを回り、最初に見出しが当たったシートで返していたので、Markdown に
    同じ「- 見出し: 値」が2行並び、2枚目の値はどこにも出なかった。
    """
    path = _two_sheet_book(tmp_path / "二枚.xlsx")
    pattern_id = _create_form_types_fixes7(client, path)
    _click_form_types_fixes7(client, pattern_id, path, "1次報告", "A3", "B3")
    _click_form_types_fixes7(client, pattern_id, path, "2次報告", "A2", "B2")

    with app.app_context():
        fields = db.load_pattern(pattern_id).fields
    assert len(fields) == 2 and [f.sheet_name for f in fields] == ["1次報告", "2次報告"]

    # 読み取りテストの Markdown に、1次と2次の値がそれぞれ出る
    panel = _panel_form_types_fixes7(client, pattern_id, path)
    assert "1次の内容" in panel and "2次の内容" in panel

    # 使用開始して取り込んでも同じ（2つ目の項目が自分のセルを読む）
    activate(client, pattern_id)
    doc_id, = upload_forms(client, path)
    res = client.post(f"/forms/{doc_id}/read",
                      data={"pattern_id": pattern_id, "sheets": ["1次報告", "2次報告"]})
    assert res.status_code == 200
    with app.app_context():
        extraction = db.get_document(doc_id)["data_json"]
    import json
    read = {f["field_name"]: f for f in json.loads(extraction)["fields"]}
    second = read[[f.field_name for f in fields][1]]
    assert second["value"] == "2次の内容" and second["sheet"] == "2次報告" and second["value_cell"] == "B2"


def test_a_lone_field_on_the_second_sheet_is_not_filled_from_the_first(app, client, tmp_path):
    """項目が1つでも同じ。2枚目に登録した項目が1枚目の同名の欄を読まない。"""
    path = _two_sheet_book(tmp_path / "二枚単独.xlsx")
    pattern_id = _create_form_types_fixes7(client, path)
    _click_form_types_fixes7(client, pattern_id, path, "2次報告", "A2", "B2")
    activate(client, pattern_id)

    doc_id, = upload_forms(client, path)
    assert client.post(f"/forms/{doc_id}/read",
                       data={"pattern_id": pattern_id, "sheets": ["1次報告", "2次報告"]}).status_code == 200
    import json
    with app.app_context():
        field = json.loads(db.get_document(doc_id)["data_json"])["fields"][0]
    assert field["value"] == "2次の内容" and field["sheet"] == "2次報告"


# ---- 2枚目のシートが選ばれない・別の欄の値になる ------------------------------------------------

def _attachment_book(path, note_row: int = 1):
    """「修理報告書」に項目が集まり、「別紙」に1項目だけある帳票（別紙の記入位置は版で変わる）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "修理報告書"
    cells = {"A1": "報告番号", "B1": "R-003", "A2": "設備番号", "B2": "EQ-9",
             "A3": "発生日", "B3": "2026-02-03", "A4": "報告者", "B4": "山田",
             "A5": "作業時間", "B5": 2, "A6": "処置内容", "B6": "交換した"}
    for coord, value in cells.items():
        ws[coord] = value
    _fill(ws, "A1", "A2", "A3", "A4", "A5", "A6")
    other = wb.create_sheet("別紙")
    other[f"A{note_row}"], other[f"B{note_row}"] = "再発防止策", "月次点検の項目に追加する"
    _fill(other, f"A{note_row}")
    wb.save(path)
    return path


def test_the_second_sheet_is_chosen_by_its_name_even_with_few_fields(app, client, tmp_path):
    """名前がそのまま同じシートは、項目の数で負けても「読み取るシート」に選ばれる。

    項目の数で競わせていたので、項目の少ない「別紙」が「修理報告書」に負けてチェックが外れ、
    そのまま読むと再発防止策に報告番号の値が入っていた（警告も出なかった）。
    """
    sample = _attachment_book(tmp_path / "見本.xlsx")
    pattern_id = _create_form_types_fixes7(client, sample, name="設備修理報告書")
    for coord in ("A1", "A2", "A3", "A4", "A5", "A6"):
        _click_form_types_fixes7(client, pattern_id, sample, "修理報告書", coord, f"B{coord[1:]}")
    _click_form_types_fixes7(client, pattern_id, sample, "別紙", "A1", "B1")
    activate(client, pattern_id)

    # 別紙の記入位置が2行下にずれた帳票
    target = _attachment_book(tmp_path / "対象.xlsx", note_row=3)
    doc_id, = upload_forms(client, target)
    html = client.get(f"/forms/{doc_id}/type").get_json()["html"]
    assert '<input type="checkbox" name="sheets" value="別紙" checked' in html
    assert "7項目中7項目が見つかりました" in html


def test_a_field_is_never_filled_from_the_same_cell_of_another_sheet(app, client, tmp_path):
    """その項目のシートでセルが空でも、別のシートの同じ番地は読まない（空のままにする）。"""
    sample = _attachment_book(tmp_path / "見本2.xlsx")
    pattern_id = _create_form_types_fixes7(client, sample, name="設備修理報告書2")
    for coord in ("A1", "A2", "A3", "A4", "A5", "A6"):
        _click_form_types_fixes7(client, pattern_id, sample, "修理報告書", coord, f"B{coord[1:]}")
    _click_form_types_fixes7(client, pattern_id, sample, "別紙", "A1", "B1")
    activate(client, pattern_id)

    target = _attachment_book(tmp_path / "対象2.xlsx", note_row=3)
    doc_id, = upload_forms(client, target)
    # 「修理報告書」だけを選んで読む（別紙のチェックを外したとき）
    assert client.post(f"/forms/{doc_id}/read",
                       data={"pattern_id": pattern_id, "sheets": ["修理報告書"]}).status_code == 200
    import json
    with app.app_context():
        fields = json.loads(db.get_document(doc_id)["data_json"])["fields"]
    note = next(f for f in fields if f["display_name"] == "再発防止策")
    assert note["value"] in (None, "") and "R-003" != note["value"]
    assert note["warning"] == "ラベルが見つかりません"


# ---- 手で直した見出しは、次のクリックで戻らない --------------------------------------------------

def _two_people_book_form_types_fixes7(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "報告書"
    ws["A1"], ws["B1"] = "報告者", "長谷川 聡"
    ws["A2"], ws["B2"] = "担当者", "清水 彩花"
    _fill(ws, "A1", "A2")
    wb.save(path)
    return path


def test_a_hand_typed_label_survives_the_next_click(app, client, tmp_path):
    """見出しを手で「担当者」に直したあと、担当者の欄をクリックしても手の直しが消えない。"""
    path = _two_people_book_form_types_fixes7(tmp_path / "二人.xlsx")
    pattern_id = _create_form_types_fixes7(client, path, name="人の帳票")
    _click_form_types_fixes7(client, pattern_id, path, "報告書", "A1", "B1")

    res = client.post(f"/form-types/{pattern_id}/fields/reporter/label", json={"name": "担当者"})
    assert res.status_code == 200 and "見出しを「担当者」にしました" in res.get_json()["message"]

    body = _click_form_types_fixes7(client, pattern_id, path, "報告書", "A2", "B2")
    with app.app_context():
        fields = db.load_pattern(pattern_id).fields
    names = {f.field_name: f.display_name for f in fields}
    assert names["reporter"] == "担当者"                      # 手で付けた見出しはそのまま
    assert names["reporter_2"] == "担当者（A2）"              # 新しいほうを見分けられる名前にする
    assert "「担当者（A2）」を「担当者」とは別の項目にしました" in body["message"]


def test_two_auto_named_fields_are_still_told_apart_by_their_labels(app, client, tmp_path):
    """手で直していない項目どうしは、これまでどおりクリックした見出しで見分ける。"""
    path = _two_people_book_form_types_fixes7(tmp_path / "二人2.xlsx")
    pattern_id = _create_form_types_fixes7(client, path, name="人の帳票2")
    _click_form_types_fixes7(client, pattern_id, path, "報告書", "A1", "B1")
    _click_form_types_fixes7(client, pattern_id, path, "報告書", "A2", "B2")

    with app.app_context():
        fields = db.load_pattern(pattern_id).fields
    assert [f.display_name for f in fields] == ["報告者", "担当者"]


# ---- 項目を全部消したら、使用を停止する ---------------------------------------------------------

def test_deleting_the_last_field_stops_an_active_type(app, client, sample_dir):
    """使用中の種類から項目を全部消したら、帳票取り込みの候補に出なくなる。

    そのままだと「0項目中0項目が見つかりました」で読み取れてしまい、中身の無い .md ができていた。
    """
    pattern_id = create_type(client, sample_dir / "standard.xlsx", "設備修理報告書")
    add_field(client, pattern_id, "修理報告書", "A3", "B3")
    activate(client, pattern_id)

    res = client.post(f"/form-types/{pattern_id}/fields/report_id/delete")
    assert res.status_code == 200
    assert "使用を停止しました" in res.get_json()["message"]
    with app.app_context():
        assert db.load_pattern(pattern_id).status == "inactive"
        assert db.load_active_patterns() == []

    doc_id, = upload_forms(client, sample_dir / "standard.xlsx")
    body = client.get(f"/forms/{doc_id}/type").get_json()
    assert body["has_types"] is False


def test_deleting_one_of_several_fields_keeps_the_type_in_use(app, client, sample_dir):
    """項目が残っているなら、使用中のまま（削除のたびに止めない）。"""
    pattern_id = create_type(client, sample_dir / "standard.xlsx", "設備修理報告書")
    for label, value in (("A3", "B3"), ("E4", "F4")):
        add_field(client, pattern_id, "修理報告書", label, value)
    activate(client, pattern_id)

    res = client.post(f"/form-types/{pattern_id}/fields/report_id/delete")
    assert res.get_json()["message"] == "項目を削除しました"
    with app.app_context():
        assert db.load_pattern(pattern_id).status == "active"


# ---- Excel を置いていないときの文言（サーバーには残していない） ------------------------------------

def test_the_panel_without_a_book_says_the_excel_is_not_kept(app, client, sample_dir):
    """Excel を置いていない画面では、「見つかりません」「項目を1つ以上作ると」とは言わない。

    置いた Excel は残さないので、種類を開き直したときはシートも読み取りテストも出せない。
    そのことと「設定は残っている」ことを画面で言う。
    """
    path = sample_dir / "standard.xlsx"
    pattern_id = create_type(client, path, "設備修理報告書")
    add_field(client, pattern_id, "修理報告書", "A3", "B3")
    before = _panel_form_types_fixes7(client, pattern_id, path)
    assert "1項目中 <strong>1</strong>項目が見つかりました" in before

    # 使用開始しても、置いている Excel はそのまま見られる（消すものがもう無い）
    res = client.post(f"/form-types/{pattern_id}/status",
                      data={"status": "active", "book": book_part(path)},
                      content_type="multipart/form-data")
    assert "1項目中 <strong>1</strong>項目が見つかりました" in res.get_json()["html"]
    assert "見本のExcelはサーバーから消しました" not in res.get_json()["message"]

    # 開き直したとき（Excel を置いていない）は、設定だけの画面になる
    panel = _panel_form_types_fixes7(client, pattern_id)
    assert "—（見つかりません）" not in panel
    assert "—（Excelを置くと、この設定で読んだ値が出ます）" in panel
    assert "項目を1つ以上作ると、ここに読み取り結果が出ます。" not in panel
    assert "いま Excel を置いていないので、読み取りテストはできません" in panel
    assert "サーバーに残していません" in panel


# ---- 置いた Excel を替えたあと、項目を削除しても表示が戻らない ---------------------------------------

def test_deleting_a_field_keeps_the_book_that_is_being_looked_at(app, client, tmp_path):
    """2つ目の Excel を見ているときに項目を削除しても、その Excel の表示のままにする。"""
    first = _two_people_book_form_types_fixes7(tmp_path / "一番.xlsx")
    pattern_id = _create_form_types_fixes7(client, first, name="Excelを置き替える")
    second = _two_people_book_form_types_fixes7(tmp_path / "二番.xlsx")

    _click_form_types_fixes7(client, pattern_id, first, "報告書", "A1", "B1")
    assert "二番.xlsx" in _panel_form_types_fixes7(client, pattern_id, second)

    # 画面の JS と同じように、いま見ている Excel を一緒に送る
    res = client.post(f"/form-types/{pattern_id}/fields/reporter/delete",
                      data={"book": book_part(second)}, content_type="multipart/form-data")
    assert res.status_code == 200
    html = res.get_json()["html"]
    assert "二番.xlsx" in html and 'data-cell="A1"' in html


# ====================================================================================================
# 元 tests/test_form_types_no_excel_kept.py
# 帳票登録は Excel を預からない（利用者の指示 2026-09-21「見本のExcelは置かずに、設定だけ保持する」）。
#
# 置いた Excel は受け取った要求の中で読み取るだけで、uploads/ にも DB にも残さない。
# 次の操作（セルのクリック・項目の削除・見出しの手直し・読み取りテスト）では、ブラウザが同じ
# ファイルを送り直す。サーバーは読み取った結果だけを core.workbook_cache に短い間だけ覚えておく
# （メモリの中だけ。ディスクには何も書かない）。
#
# ここで確かめること:
#   - 置いても・クリックしても・使用開始しても、Excel のバイトはどこにも残らない
#   - 残った設定だけで、開き直した画面の項目一覧と見出しの手直しができる
#   - 同じ帳票の Excel をもう一度置けば、シートを見てクリックする作業に戻れる（種類は増えない）
#   - 覚えている結果は、その人のもの・短い間だけ・数に上限がある
# ====================================================================================================

@pytest.fixture(autouse=True)
def _clean_cache():
    """テストごとに、覚えているブックを空にする（プロセスで1つの置き場なので）。"""
    core.clear()
    yield
    core.clear()


def _report(path, report_id="R-001"):
    """見出しと値が並ぶ、ふつうの帳票。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "報告書"
    fill = PatternFill("solid", fgColor="FFD9E1F2")
    for coord, value in {"A1": "報告番号", "B1": report_id, "A2": "設備番号", "B2": "EQ-001",
                         "A3": "所見", "B3": "異音あり"}.items():
        ws[coord] = value
    for coord in ("A1", "A2", "A3"):
        ws[coord].fill = fill
    wb.save(path)
    return path


def _stored_files(app) -> list[str]:
    from pathlib import Path

    return [str(p) for p in Path(app.config["UPLOAD_DIR"]).rglob("*") if p.is_file()]


def _tables(app) -> set[str]:
    with app.app_context():
        return {row[0] for row in db.get_db().execute("SELECT name FROM sqlite_master WHERE type = 'table'")}


# ---- 置いても残らない ---------------------------------------------------------------------

def test_dropping_an_excel_keeps_no_file_and_no_row(app, client, tmp_path):
    """Excel を置いて種類を作っても、ファイルは1つも保存されず、控えの行もできない。"""
    path = _report(tmp_path / "設備修理報告書.xlsx")
    pattern_id = create_type(client, path, "設備修理報告書")
    add_field(client, pattern_id, "報告書", "A1", "B1")

    assert _stored_files(app) == []
    assert "pattern_samples" not in _tables(app)
    with app.app_context():
        # Excel を指す行（stored_path）は、どの表にもできていない
        conn = db.get_db()
        for (table,) in conn.execute("SELECT name FROM sqlite_master WHERE type = 'table'"):
            columns = {row[1] for row in conn.execute(f'PRAGMA table_info("{table}")')}
            if "stored_path" in columns:
                assert conn.execute(f'SELECT COUNT(*) FROM "{table}"').fetchone()[0] == 0, table
        # 残るのは設定だけ（シート名・見出しのセル・値のセル・向き・項目名）
        field = db.load_pattern(pattern_id).fields[0]
        assert (field.sheet_name, field.label_cell, field.cell, field.direction) == ("報告書", "A1", "B1", "right")
        assert field.display_name == "報告番号"


def test_the_whole_click_flow_needs_no_stored_excel(app, client, tmp_path):
    """置く → クリック → 見出しを直す → 読み取りテスト → 使用開始 → 帳票取り込み、が通る。"""
    path = _report(tmp_path / "設備修理報告書.xlsx", "R-777")
    pattern_id = create_type(client, path, "設備修理報告書")
    assert "「報告番号」を項目にしました" in add_field(client, pattern_id, "報告書", "A1", "B1")["message"]
    add_field(client, pattern_id, "報告書", "A2", "B2")

    # 見出しを手で直す（画面と同じく、いま置いている Excel も一緒に送る）
    res = client.post(f"/form-types/{pattern_id}/fields/report_id/label",
                      data={"name": "受付番号", "book": book_part(path)},
                      content_type="multipart/form-data")
    assert res.status_code == 200
    panel = res.get_json()["html"]
    assert "- 受付番号: R-777" in panel and "2項目中 <strong>2</strong>項目が見つかりました" in panel

    activate(client, pattern_id)
    doc_id, = upload_forms(client, path)
    assert read_form(client, doc_id, pattern_id, ["報告書"]).status_code == 200
    with app.app_context():
        import json

        read = {f["display_name"]: f["value"] for f in json.loads(db.get_document(doc_id)["data_json"])["fields"]}
    assert read == {"受付番号": "R-777", "設備番号": "EQ-001"}
    # 取り込んだ帳票のファイルは今までどおり残る（ダウンロードで消える）。帳票登録の分は1つも無い
    assert len(_stored_files(app)) == 1
    assert "samples" not in _stored_files(app)[0].replace("\\", "/")


def test_activation_does_not_claim_to_delete_anything(app, client, tmp_path):
    """使用開始の知らせで「見本のExcelを消しました」とは言わない（もともと置いていない）。"""
    path = _report(tmp_path / "点検表.xlsx")
    pattern_id = create_type(client, path, "点検表")
    add_field(client, pattern_id, "報告書", "A1", "B1")

    res = client.post(f"/form-types/{pattern_id}/status", json={"status": "active"})
    message = res.get_json()["message"]
    assert "使用を開始しました" in message and "消しました" not in message
    assert _stored_files(app) == []


# ---- 設定だけで開き直せる -----------------------------------------------------------------

def test_reopening_a_type_without_the_excel_shows_the_settings(app, client, tmp_path):
    """Excel を置いていない画面でも、項目の一覧と見出しの手直しはできる。"""
    path = _report(tmp_path / "点検表.xlsx")
    pattern_id = create_type(client, path, "点検表")
    add_field(client, pattern_id, "報告書", "A1", "B1")
    core.clear()                       # 時間がたって忘れたあと

    panel = panel_html(client, pattern_id)
    assert "報告番号" in panel                   # 項目は残っている
    assert "サーバーに残していません" in panel   # 残していないことと、置き直せることを言う
    assert "この帳票のExcelを置く" in panel
    assert "いま Excel を置いていないので、読み取りテストはできません" in panel
    assert "—（Excelを置くと、この設定で読んだ値が出ます）" in panel

    # Excel が無くても見出しは直せる（Markdown に書く名前だけが変わる）
    res = client.post(f"/form-types/{pattern_id}/fields/report_id/label", json={"name": "受付番号"})
    assert res.status_code == 200 and "見出しを「受付番号」にしました" in res.get_json()["message"]
    with app.app_context():
        field = db.load_pattern(pattern_id).fields[0]
    assert field.display_name == "受付番号" and field.candidates[0] == "報告番号" and field.cell == "B1"


def test_dropping_the_same_excel_again_shows_the_sheet(app, client, tmp_path):
    """同じ帳票の Excel をもう一度置くと、シートが出て続きの作業ができる（種類は増えない）。"""
    path = _report(tmp_path / "点検表.xlsx")
    pattern_id = create_type(client, path, "点検表")
    add_field(client, pattern_id, "報告書", "A1", "B1")
    core.clear()

    res = client.post(f"/form-types/{pattern_id}/panel", data={"book": book_part(path)},
                      content_type="multipart/form-data")
    assert res.status_code == 200
    body = res.get_json()
    assert "点検表.xlsx" in body["html"] and 'data-cell="A1"' in body["html"]
    assert "R-001" in body["html"]                 # 置いた帳票での値も出る
    assert "サーバーに残しません" in body["message"]
    with app.app_context():
        assert len(db.list_patterns()) == 1        # 置き直しで新しい種類は作らない
    assert _stored_files(app) == []

    # そのままセルをクリックして項目を足せる
    assert "「設備番号」を項目にしました" in add_field(client, pattern_id, "報告書", "A2", "B2")["message"]


def test_a_click_without_the_excel_asks_for_it_again(app, client, tmp_path):
    """Excel を送らずにセルのクリックだけ届いたら、置き直してくださいと断る。"""
    path = _report(tmp_path / "点検表.xlsx")
    pattern_id = create_type(client, path, "点検表")
    core.clear()

    res = client.post(f"/form-types/{pattern_id}/fields",
                      data={"sheet": "報告書", "label_cell": "A1", "value_cell": "B1"})
    assert res.status_code == 400
    assert "もう一度置いてください" in res.get_json()["error"]
    with app.app_context():
        assert db.load_pattern(pattern_id).fields == []


def test_a_big_excel_is_refused(app, client, tmp_path, monkeypatch):
    """置く Excel には大きさの上限がある（保存せずにメモリで読むので、無制限にはできない）。"""
    import views

    monkeypatch.setattr(views, "BOOK_MAX_BYTES", 1000)
    path = _report(tmp_path / "大きい.xlsx")
    res = client.post("/form-types/new", data={"name": "大きい", "book": book_part(path)},
                      content_type="multipart/form-data")
    assert res.status_code == 400 and "大きすぎます" in res.get_json()["error"]
    with app.app_context():
        assert db.list_patterns() == []      # 読めなかったときは帳票の種類も作らない
    assert _stored_files(app) == []
    # 画面にも上限を書いてある
    assert "1ファイル" in client.get("/form-types/").get_data(as_text=True)


def test_deleting_a_type_leaves_nothing(app, client, tmp_path):
    """種類を削除しても、消すファイルはもともと無く、設定の行だけが消える。"""
    path = _report(tmp_path / "点検表.xlsx")
    pattern_id = create_type(client, path, "点検表")
    add_field(client, pattern_id, "報告書", "A1", "B1")

    assert client.post(f"/form-types/{pattern_id}/delete").status_code == 200
    with app.app_context():
        assert db.load_pattern(pattern_id) is None
        assert db.get_db().execute("SELECT COUNT(*) FROM pattern_fields").fetchone()[0] == 0
    assert _stored_files(app) == []


# ---- 送り直さずに済ませる合図（sha256）-------------------------------------------------------

def test_the_sheet_can_be_asked_for_by_the_hash_of_the_book(app, client, tmp_path):
    """一度読んだブックは、ファイルを送らずに合図（sha256）だけで開き直せる。"""
    path = _report(tmp_path / "点検表.xlsx")
    pattern_id = create_type(client, path, "点検表")
    panel = panel_html(client, pattern_id, book=path)
    book_hash = panel.split('data-book="')[1].split('"')[0]
    assert len(book_hash) == 64

    res = client.post(f"/form-types/{pattern_id}/fields",
                      data={"sheet": "報告書", "label_cell": "A1", "value_cell": "B1",
                            "book_hash": book_hash})
    assert res.status_code == 200 and "「報告番号」を項目にしました" in res.get_json()["message"]

    # 忘れたあとは合図だけでは足りない（ブラウザが持っているファイルを送り直す）
    core.clear()
    res = client.post(f"/form-types/{pattern_id}/fields",
                      data={"sheet": "報告書", "label_cell": "A2", "value_cell": "B2",
                            "book_hash": book_hash})
    assert res.status_code == 400 and "もう一度置いてください" in res.get_json()["error"]


def test_another_browser_cannot_use_the_book_of_the_first_one(app, tmp_path):
    """覚えているブックはその作業場所（ブラウザ）のもの。ほかの人の合図では出てこない。"""
    from views import SESSION_ID_KEY

    path = _report(tmp_path / "点検表.xlsx")
    first, second = app.test_client(), app.test_client()
    for client, sid in ((first, "a" * 32), (second, "b" * 32)):
        with client.session_transaction() as cookie:
            cookie[SESSION_ID_KEY] = sid

    pattern_id = create_type(first, path, "点検表")
    panel = panel_html(first, pattern_id, book=path)
    book_hash = panel.split('data-book="')[1].split('"')[0]

    res = second.post(f"/form-types/{pattern_id}/fields",
                      data={"sheet": "報告書", "label_cell": "A1", "value_cell": "B1",
                            "book_hash": book_hash})
    assert res.status_code == 400 and "もう一度置いてください" in res.get_json()["error"]


# ---- 覚えておく置き場（メモリだけ）------------------------------------------------------------

def _fake_book(name: str, size: int = 10) -> core.Book:
    return core.Book(file_name=name, file_hash=name, size=size, info=object())


def test_the_cache_forgets_old_books():
    """しばらく使われなかったブックは忘れる（覚えているのは短い間だけ）。"""
    book = core.put("sid", _fake_book("a"))
    assert core.get("sid", "a") is not None
    book.used_at -= core.TTL_SECONDS + 1      # 時間がたったことにする
    assert core.get("sid", "a") is None
    assert core.count() == 0
    # 新しく置き直したブックは、古いものを片付けるときに巻き込まれない
    old = core.put("sid", _fake_book("b"))
    old.used_at -= core.TTL_SECONDS + 1
    core.put("sid", _fake_book("c"))
    assert core.get("sid", "b") is None and core.get("sid", "c") is not None


def test_the_cache_keeps_only_the_newest_books(monkeypatch):
    """数と大きさに上限があり、あふれたら古いものから落とす。"""
    monkeypatch.setattr(core, "MAX_ENTRIES", 2)
    for name in ("a", "b", "c"):
        core.put("sid", _fake_book(name))
    assert core.count() == 2
    assert core.get("sid", "a") is None      # いちばん古いものが落ちる
    assert core.get("sid", "c") is not None

    monkeypatch.setattr(core, "MAX_BYTES", 100)
    core.clear()
    core.put("sid", _fake_book("big", size=80))
    core.put("sid", _fake_book("big2", size=80))
    assert core.count() == 1 and core.get("sid", "big") is None


def test_the_same_book_read_twice_is_only_parsed_once(app, client, tmp_path, monkeypatch):
    """同じ Excel を送り直しても、ブックを開き直さない（覚えている結果を使う）。"""
    path = _report(tmp_path / "点検表.xlsx")
    pattern_id = create_type(client, path, "点検表")

    import views

    opened = []
    real = views.load_workbook_info

    def counting(source):
        opened.append(1)
        return real(source)

    monkeypatch.setattr(views, "load_workbook_info", counting)
    add_field(client, pattern_id, "報告書", "A1", "B1")
    add_field(client, pattern_id, "報告書", "A2", "B2")
    assert opened == []          # 1回目（種類を作ったとき）に読んだ結果を使い回す


# ---- メモリに読むだけの受け取り（core.files.read_upload）----------------------------------------

class _Storage:
    def __init__(self, data: bytes, filename: str):
        self.stream = io.BytesIO(data)
        self.filename = filename


def test_read_upload_keeps_nothing_on_disk(app, tmp_path):
    """アップロードは保存先を作らずにメモリへ読む（中身と sha256 だけ返す）。"""
    import hashlib

    data = _report(tmp_path / "点検表.xlsx").read_bytes()
    with app.app_context():
        got = core.read_upload(_Storage(data, "点検表.xlsx"), {".xlsx", ".xlsm"}, 10_000_000)
    assert got.data == data and got.size == len(data)
    assert got.file_hash == hashlib.sha256(data).hexdigest()
    assert got.file_name == "点検表.xlsx"
    assert _stored_files(app) == []

    with app.app_context():
        with pytest.raises(UploadError, match="ファイルを選んでください"):
            core.read_upload(_Storage(data, "点検表.csv"), {".xlsx"}, 10_000_000)
        with pytest.raises(UploadError, match="空です"):
            core.read_upload(_Storage(b"", "空.xlsx"), {".xlsx"}, 10_000_000)
        with pytest.raises(UploadError, match="大きすぎます"):
            core.read_upload(_Storage(data, "点検表.xlsx"), {".xlsx"}, 10)
    assert _stored_files(app) == []


def test_precheck_reads_the_bytes_without_a_file(tmp_path):
    """事前チェックは、ファイルにしていない中身（bytes）でも同じように断る。"""
    core.precheck_excel(_report(tmp_path / "ok.xlsx").read_bytes())      # ふつうのブックは通る
    with pytest.raises(UploadError, match="Excelファイル"):
        core.precheck_excel(b"not an excel file")
    with pytest.raises(UploadError, match="セル数が上限"):
        core.precheck_excel(_report(tmp_path / "ok2.xlsx").read_bytes(), max_cells=1)


# ---- 前の版が置いた見本の片付け ---------------------------------------------------------------

def test_the_old_sample_folder_is_cleaned_up(app):
    """前の版が uploads/samples に置いたままの Excel は、起動時にフォルダごと捨てる。"""
    from pathlib import Path

    folder = Path(app.config["UPLOAD_DIR"]) / "samples"
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "0123456789abcdef0123456789abcdef.xlsx").write_bytes(b"old sample")

    with app.app_context():
        assert core.purge_old_sample_files() == 1
        assert core.purge_old_sample_files() == 0     # 2回目は何もしない
    assert not folder.exists()
    assert _stored_files(app) == []


def test_the_migration_drops_the_sample_table(tmp_path):
    """前の版の DB（pattern_samples がある）を開くと、その表を落とす。"""
    import sqlite3

    from tests.test_core import _make_old_db, make_app

    path = tmp_path / "app.db"
    _make_old_db(path)
    conn = sqlite3.connect(path)
    conn.execute("INSERT INTO pattern_samples (pattern_id, file_name, file_hash, stored_path, created_at) "
                 "VALUES (1, '見本.xlsx', 'h', 'samples/x.xlsx', '2026-01-01T00:00:00')")
    conn.commit()
    conn.close()

    app = make_app(tmp_path)
    with app.app_context():
        conn = db.get_db()
        tables = {row[0] for row in conn.execute("SELECT name FROM sqlite_master WHERE type = 'table'")}
        assert "pattern_samples" not in tables
        assert db.load_pattern(1) is not None          # 設定（帳票の種類）は残る


# ====================================================================================================
# 元 tests/test_form_sample_version_dirs.py
# 帳票サンプル（F1〜F5）の版フォルダ名と、その中身の置き方のテスト。
#
# フォルダ名の付け方は5つの生成スクリプトで共通にしてある。
# 「版の名前＋いつから」を日本語で書き、区別に要るもの（様式番号・用紙サイズ・シート枚数）だけを足す。
# ローマ字は使わず、Windows のフォルダ名に使えない文字も使わない。
# ====================================================================================================

SAMPLES = Path(__file__).resolve().parents[1] / "samples" / "forms"

# 帳票フォルダ名 → そのスクリプトが作る版フォルダ名
FAMILIES: dict[str, list[str]] = {
    f1.FORM_DIR: [f1._version_dir(v) for v in f1.REV_INFO],
    f2.FOLDER: list(f2.VERSION_DIR.values()),
    f3.FOLDER: list(f3.VERSION_DIRS.values()),
    f4.FORM_DIR: [f4._rev_dir(v) for v in f4.REV_INFO],
    f5.FORM_FOLDER: list(f5.VERSION_FOLDERS.values()),
}

NG_CHARS = '.:/\\*?"<>|'
# 名前に出てよい英数字は、様式番号（QA-F-021）・版の名前（Rev3）・用紙サイズ（A3/A4）・年だけ
ALLOWED_TOKENS = re.compile(r"QA-F-\d{3}|Rev\d+|A[34]|\d+|_")


def _families():
    return [(family, name) for family, names in FAMILIES.items() for name in names]


def test_version_dir_names_are_windows_safe():
    """版フォルダ名に Windows で使えない文字・前後の空白・末尾のピリオドが無い。"""
    for family, name in _families():
        assert name, family
        assert not set(name) & set(NG_CHARS), f"{family}/{name}"
        assert name == name.strip() and not name.endswith(("。", "、")), f"{family}/{name}"
        assert len(name) <= 60, f"{family}/{name}"


def test_version_dir_names_have_no_romaji():
    """版フォルダ名にローマ字（hougan・cols・sheet など）が残っていない。"""
    for family, name in _families():
        rest = ALLOWED_TOKENS.sub("", name)
        assert not re.search(r"[A-Za-z]", rest), f"ローマ字が残っている: {family}/{name}"


def test_version_dir_names_say_the_revision_and_when():
    """版フォルダ名は「版の名前＋いつから」。版の名前と年があり、制定／改訂／まで で終わる。"""
    for family, name in _families():
        assert re.search(r"Rev\d+", name), f"版の名前が無い: {family}/{name}"
        assert re.search(r"(19|20)\d{2}", name), f"年が無い: {family}/{name}"
        assert name.endswith(("制定", "改訂", "まで")), f"いつからが無い: {family}/{name}"


def test_version_dir_names_are_unique_in_a_family():
    """同じ帳票の中で版フォルダ名がぶつからない（別レイアウトを同じ名前にしていない）。"""
    for family, names in FAMILIES.items():
        assert len(set(names)) == len(names), family
    assert [len(v) for v in FAMILIES.values()] == [3, 2, 5, 2, 3]


def test_every_layout_version_has_a_folder():
    """生成スクリプトが作りうる版（layout_version）には、必ず入れ先のフォルダがある。"""
    assert set(f1.REV_INFO) == {"v1", "v2", "v3"} and set(f4.REV_INFO) == {"v1", "v2"}
    assert set(f2.VERSION_DIR) == {f2.OLD, f2.NEW}
    assert set(f5.VERSION_FOLDERS) == {"A3横_Rev.1", "A3横_Rev.2", "A4縦_Rev.3"}
    made = {f3._make_style(rank, claim, rank).layout_version for claim in (False, True) for rank in range(20)}
    assert made == set(f3.VERSION_DIRS), made ^ set(f3.VERSION_DIRS)


# ---- samples/forms の実ファイル ----

def _family_dir(name: str) -> Path:
    path = SAMPLES / name
    if not path.exists():
        pytest.skip(f"サンプルがありません: {name}")
    return path


@pytest.mark.samples
def test_sample_tree_has_one_folder_per_version():
    """帳票フォルダの直下は版フォルダと _README.md・_expected.jsonl だけで、版フォルダの中は .xlsx だけ。"""
    for family, names in FAMILIES.items():
        root = _family_dir(family)
        assert sorted(p.name for p in root.iterdir() if p.is_dir()) == sorted(names), family
        assert sorted(p.name for p in root.iterdir() if p.is_file()) == ["_README.md", "_expected.jsonl"], family
        for name in names:
            kids = list((root / name).iterdir())
            assert kids, f"{family}/{name} が空"
            assert all(k.is_file() and k.suffix == ".xlsx" for k in kids), f"{family}/{name}"


@pytest.mark.samples
def test_sample_expected_file_is_the_path_under_the_version_folder():
    """_expected.jsonl の file は「版フォルダ/ファイル名」で、実在する .xlsx を指している。"""
    for family, names in FAMILIES.items():
        root = _family_dir(family)
        rows = [json.loads(line) for line in (root / "_expected.jsonl").read_text(encoding="utf-8").splitlines() if line.strip()]
        assert len(rows) == 30, family
        for row in rows:
            folder, _, filename = row["file"].partition("/")
            assert folder in names, f"{family}: {row['file']}"
            assert filename.endswith(".xlsx") and "/" not in filename, f"{family}: {row['file']}"
            assert (root / folder / filename).exists(), f"{family}: {row['file']}"
        # 版フォルダの .xlsx は全部 _expected.jsonl に載っている
        listed = {row["file"] for row in rows}
        actual = {f"{d.name}/{p.name}" for d in root.iterdir() if d.is_dir() for p in d.iterdir()}
        assert listed == actual, family


@pytest.mark.samples
def test_sample_readme_names_every_version_folder():
    """_README.md がフォルダ名を載せていて、フォルダごとまとめて取り込めると書いてある。"""
    for family, names in FAMILIES.items():
        text = _family_dir(family).joinpath("_README.md").read_text(encoding="utf-8")
        assert "フォルダ" in text.split("\n##")[0] or "フォルダ" in text, family
        for name in names:
            assert name in text, f"{family}: {name} が _README.md に無い"
        assert "帳票取り込み" in text and "まとめて" in text, family
        assert "版の名前＋いつから" in text, family


def test_rerun_removes_the_folders_of_old_names(tmp_path):
    """作り直すと、前の名前の版フォルダと中の .xlsx が残らない（F4 で確かめる）。"""
    out = tmp_path / "forms" / f4.FORM_DIR
    stale = out / "Rev1_2018seitei"          # 前の名前のフォルダ
    stale.mkdir(parents=True)
    (stale / "古い報告書.xlsx").write_bytes(b"old")
    (stale / "メモ.txt").write_text("消えてよい", encoding="utf-8")
    (out / "直下の残り.xlsx").write_bytes(b"old")

    paths = f4.generate(tmp_path)

    assert not stale.exists() and not (out / "直下の残り.xlsx").exists()
    assert sorted(p.name for p in out.iterdir() if p.is_dir()) == sorted(FAMILIES[f4.FORM_DIR])
    assert {p.parent.name for p in paths if p.suffix == ".xlsx"} == set(FAMILIES[f4.FORM_DIR])
    rows = [json.loads(line) for line in (out / "_expected.jsonl").read_text(encoding="utf-8").splitlines()]
    assert all((out / row["file"]).exists() for row in rows)


# ====================================================================================================
# 元 tests/test_keep_nothing.py
# 「その場でダウンロードしない限り、その場ですぐ捨てる」（利用者の指示 2026-09-20）の確認。
#
# 社内LANに置いて数人が同時に使うので、捨てるのは「その人の分」だけで、ほかの人の作業には触らない。
# 捨てる機会は4つ（core/purge.py）:
#   - 画面を離れたとき（ブラウザが /forms/discard・/tables/discard に伝える）… purge_session / discard_*
#   - 新しいファイルを置いたとき（同じ人の前の分）                          … discard_documents / discard_table_imports
#   - しばらくさわられていないとき（IDLE_HOURS）                            … sweep_stale
#   - 起動時                                                                … purge_all_pending
# あとの2つは、動いているジョブが付いているものには手を出さない。
# ====================================================================================================

# views.current_session_id() が配る id と同じ形（32桁）にする（そうでないとブラウザ側で作り直される）
SESSION_A = "a" * 32
SESSION_B = "b" * 32


# ---- 下ごしらえ ---------------------------------------------------------------------------

def _ensure_session_column(app) -> None:
    """持ち主の列（views.current_session_id() が入れる）がまだ無い DB でも試せるようにする。"""
    with app.app_context():
        conn = db.get_db()
        for table in ("documents", "table_imports"):
            names = {row[1] for row in conn.execute(f'PRAGMA table_info("{table}")')}
            if "session_id" not in names:
                conn.execute(f"ALTER TABLE {table} ADD COLUMN session_id TEXT")
        conn.commit()


def _own(app, table: str, row_id: int, session_id: str) -> None:
    with app.app_context():
        conn = db.get_db()
        conn.execute(f"UPDATE {table} SET session_id = ? WHERE id = ?", (session_id, row_id))
        conn.commit()


def _document(app, name: str, session_id: str) -> tuple[int, Path]:
    doc_id, path = add_confirmed_document(app, name)
    _own(app, "documents", doc_id, session_id)
    return doc_id, path


def _import(app, session_id: str, *, file_name: str = "一覧.csv") -> int:
    """一覧表の取り込みを1件作る（読み込みまでは進めない。捨てられるかどうかだけを見る）。"""
    with app.app_context():
        stored = f"tables/{session_id}-{file_name}"
        path = Path(app.config["UPLOAD_DIR"]) / stored
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(b"a,b\r\n1,2\r\n")
        conn = db.get_db()
        now = datetime.now().isoformat(timespec="seconds")
        import_id = conn.execute(
            "INSERT INTO table_imports (file_name, file_hash, stored_path, created_at, updated_at, session_id) "
            "VALUES (?, ?, ?, ?, ?, ?)", (file_name, "0" * 64, stored, now, now, session_id)).lastrowid
        conn.commit()
    folder = _import_folder(app, import_id)
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "rows.jsonl").write_text('{"a": 1}\n', encoding="utf-8")
    return import_id


def _import_folder(app, import_id: int) -> Path:
    with app.app_context():
        return core.import_dir(import_id)


def _running_job(app, ref_type: str, ref_id: int, kind: str = "table_read") -> int:
    with app.app_context():
        conn = db.get_db()
        now = datetime.now().isoformat(timespec="seconds")
        job_id = conn.execute(
            "INSERT INTO jobs (kind, ref_type, ref_id, status, created_at, updated_at) "
            "VALUES (?, ?, ?, 'running', ?, ?)", (kind, ref_type, ref_id, now, now)).lastrowid
        conn.commit()
        return job_id


def _age_keep_nothing(app, table: str, row_id: int, hours: float) -> None:
    """その行を hours 時間前にさわられたことにする。"""
    old = (datetime.now() - timedelta(hours=hours)).isoformat(timespec="seconds")
    with app.app_context():
        conn = db.get_db()
        names = {row[1] for row in conn.execute(f'PRAGMA table_info("{table}")')}
        sets = ", ".join(f"{c} = ?" for c in ("created_at", "updated_at", "confirmed_at") if c in names)
        args = [old] * sets.count("?") + [row_id]
        conn.execute(f"UPDATE {table} SET {sets} WHERE id = ?", args)
        conn.commit()


def _alive(app, table: str, row_id: int) -> bool:
    with app.app_context():
        return db.get_db().execute(f"SELECT 1 FROM {table} WHERE id = ?", (row_id,)).fetchone() is not None


@pytest.fixture
def sessions(app):
    _ensure_session_column(app)
    return app


# ---- 1 画面を離れたら捨てる ------------------------------------------------------------------

def test_leaving_the_page_discards_that_persons_work_only(sessions):
    """画面を離れた人の分だけを捨て、同時に使っているほかの人の分は残す。"""
    app = sessions
    mine, my_file = _document(app, "mine.xlsx", SESSION_A)
    yours, your_file = _document(app, "yours.xlsx", SESSION_B)
    my_import = _import(app, SESSION_A)
    your_import = _import(app, SESSION_B)
    my_folder = _import_folder(app, my_import)

    with app.app_context():
        assert core.purge_session(SESSION_A) == (1, 1)

    assert not _alive(app, "documents", mine)
    assert not _alive(app, "table_imports", my_import)
    assert not my_file.exists()
    assert not my_folder.exists()          # imports/<id>/（読み込んだ行・控え・作った md）ごと消える
    assert _alive(app, "documents", yours)
    assert _alive(app, "table_imports", your_import)
    assert your_file.exists()


def test_discarding_twice_does_nothing_the_second_time(sessions):
    """画面を閉じる合図は2回届くことがある（pagehide と visibilitychange）。2回目は何もしない。"""
    app = sessions
    doc_id, _ = _document(app, "mine.xlsx", SESSION_A)
    import_id = _import(app, SESSION_A)
    with app.app_context():
        assert core.discard_documents([doc_id], SESSION_A) == 1
        assert core.discard_documents([doc_id], SESSION_A) == 0
        assert core.discard_table_imports([import_id], SESSION_A) == 1
        assert core.discard_table_imports([import_id], SESSION_A) == 0
        assert core.purge_session(SESSION_A) == (0, 0)


def test_one_person_cannot_discard_another_persons_work(sessions):
    """番号を知っていても、ほかの人の取り込みは捨てられない。"""
    app = sessions
    yours, _ = _document(app, "yours.xlsx", SESSION_B)
    your_import = _import(app, SESSION_B)
    with app.app_context():
        assert core.discard_documents([yours], SESSION_A) == 0
        assert core.discard_table_imports([your_import], SESSION_A) == 0
    assert _alive(app, "documents", yours)
    assert _alive(app, "table_imports", your_import)


def test_rows_without_an_owner_are_never_swept_by_session(app):
    """持ち主の分からない行（この仕組みを入れる前のDB・直接作った行）は、まとめては捨てない。

    誰のものか分からないものを「自分の分」として巻き込むと、同時に使っている人の作業が消える。
    こういう行は起動時の片付け（purge_all_pending）と時間切れ（sweep_stale）で消える。
    """
    doc_id, _ = add_confirmed_document(app, "mine.xlsx")
    with app.app_context():
        assert core.purge_session(SESSION_A) == (0, 0)
    assert _alive(app, "documents", doc_id)


# ---- 2 新しいファイルを置いたら前の分を捨てる ----------------------------------------------------

def test_a_new_upload_discards_the_previous_unfinished_one(sessions):
    """同じ人が新しいファイルを置いたら、前の（ダウンロードしていない）分はその場で消える。"""
    app = sessions
    first, first_file = _document(app, "first.xlsx", SESSION_A)
    first_import = _import(app, SESSION_A, file_name="first.csv")

    # 画面（static/app.js（帳票取り込み）・static/app.js（表の取り込み））は新しいファイルを送る前にこれを呼ぶ
    with app.app_context():
        assert core.discard_documents([first], SESSION_A) == 1
        assert core.discard_table_imports([first_import], SESSION_A) == 1

    second, second_file = _document(app, "second.xlsx", SESSION_A)
    assert not _alive(app, "documents", first)
    assert not first_file.exists()
    assert not _alive(app, "table_imports", first_import)
    assert _alive(app, "documents", second)
    assert second_file.exists()


# ---- 3 処理中のものには手を出さない ----------------------------------------------------------

def test_a_running_job_is_not_discarded(sessions):
    """AI整形や読み込みが動いている取り込みは、画面を離れても・時間切れでも捨てない。"""
    app = sessions
    doc_id, _ = _document(app, "busy.xlsx", SESSION_A)
    import_id = _import(app, SESSION_A)
    _running_job(app, "document", doc_id, kind="form_read")
    _running_job(app, "table_import", import_id, kind="ai_format")

    with app.app_context():
        assert core.purge_session(SESSION_A) == (0, 0)
        assert core.discard_documents([doc_id], SESSION_A) == 0
        assert core.discard_table_imports([import_id], SESSION_A) == 0
    assert _alive(app, "documents", doc_id)
    assert _alive(app, "table_imports", import_id)

    # 時間切れの掃除（IDLE_HOURS）も起動時の片付けも、動いているジョブには触らない
    _age_keep_nothing(app, "documents", doc_id, core.IDLE_HOURS + 1)
    _age_keep_nothing(app, "table_imports", import_id, core.IDLE_HOURS + 1)
    with app.app_context():
        assert core.sweep_stale(core.IDLE_HOURS) == (0, 0)
        assert core.purge_all_pending() == (0, 0)
    assert _alive(app, "documents", doc_id)
    assert _alive(app, "table_imports", import_id)


def test_the_idle_sweep_discards_what_is_left(sessions):
    """合図が届かなかった分（ブラウザの強制終了・LANの切断）は、IDLE_HOURS で片付く。"""
    app = sessions
    old_doc, old_file = _document(app, "old.xlsx", SESSION_A)
    old_import = _import(app, SESSION_A, file_name="old.csv")
    fresh_doc, _ = _document(app, "fresh.xlsx", SESSION_B)
    _age_keep_nothing(app, "documents", old_doc, core.IDLE_HOURS + 1)
    _age_keep_nothing(app, "table_imports", old_import, core.IDLE_HOURS + 1)

    with app.app_context():
        assert core.sweep_stale(core.IDLE_HOURS) == (1, 1)
    assert not _alive(app, "documents", old_doc)
    assert not old_file.exists()
    assert not _alive(app, "table_imports", old_import)
    assert _alive(app, "documents", fresh_doc)   # いま使っている人の分は残る


def test_the_idle_timeout_is_short_and_in_one_place():
    """放っておかれた分を捨てるまでの時間は定数1つ（変えやすくする）。"""
    assert core.IDLE_HOURS <= 4
    assert core.STALE_HOURS == core.IDLE_HOURS   # 旧名（app.py が参照している）


# ---- 4 ダウンロードした分はもう無い -----------------------------------------------------------

def test_a_downloaded_document_is_already_gone(sessions, client):
    """ダウンロードが終わった時点で消えている（purge_after_send）ので、捨てるものは残らない。"""
    app = sessions
    doc_id, path = _document(app, "download.xlsx", SESSION_A)
    with client.session_transaction() as cookie:
        cookie[SESSION_ID_KEY] = SESSION_A          # この帳票を取り込んだブラウザとして開く
    res = client.get(f"/forms/{doc_id}/download.md")
    assert res.status_code == 200

    assert not _alive(app, "documents", doc_id)
    assert not path.exists()
    with app.app_context():
        assert core.purge_session(SESSION_A) == (0, 0)
        assert core.discard_documents([doc_id], SESSION_A) == 0


# ---- 5 画面に1行で書いてある ----------------------------------------------------------------

@pytest.mark.parametrize("path", ["/forms", "/tables"])
def test_every_page_says_what_happens_when_you_leave(client, path):
    res = client.get(path, follow_redirects=True)
    assert res.status_code == 200
    html = res.get_data(as_text=True)
    assert "画面を閉じると" in html
    assert "捨てられます" in html


def test_the_form_type_page_says_the_excel_is_never_kept(client):
    """帳票登録は預からない: 置いた Excel はその場で読むだけ（2026-09-21 の利用者の指示）。"""
    html = client.get("/form-types", follow_redirects=True).get_data(as_text=True)
    assert "サーバーに残しません" in html
    assert "残るのは登録した帳票の種類（読み取りの設定）だけです" in html


@pytest.mark.parametrize("path,url", [("/forms", "/forms/discard"), ("/tables", "/tables/discard")])
def test_the_pages_know_where_to_send_the_discard(client, path, url):
    """画面を離れたときの送り先（static/app.js の ragDiscard が使う）。"""
    assert url in client.get(path, follow_redirects=True).get_data(as_text=True)


# ---- 6 画面からの合図（/forms/discard・/tables/discard） ------------------------------------
# ブラウザは navigator.sendBeacon で送る。中身の型は text/plain で、応答は読めず、やり直しもできない。
# だからサーバーはいつでも 204 を返し、捨てられないもの（もう無い・ほかの人の・処理中）は黙って外す。

def _as_beacon(client, url: str, body: str):
    """sendBeacon と同じ送り方（text/plain・同一サイトからの POST）。"""
    return client.post(url, data=body, content_type="text/plain;charset=UTF-8",
                       headers={"Sec-Fetch-Site": "same-origin"})


def test_the_leave_signal_discards_only_the_senders_work(sessions, client):
    """画面を離れた合図で、その番号の帳票だけを捨てる。ほかの人の分は番号を指しても残る。"""
    app = sessions
    mine, my_file = _document(app, "mine.xlsx", SESSION_A)
    yours, your_file = _document(app, "yours.xlsx", SESSION_B)
    with client.session_transaction() as cookie:
        cookie[SESSION_ID_KEY] = SESSION_A

    res = _as_beacon(client, "/forms/discard", f'{{"doc_ids": [{mine}, {yours}]}}')
    assert res.status_code == 204
    assert res.get_data() == b""

    assert not _alive(app, "documents", mine)
    assert not my_file.exists()
    assert _alive(app, "documents", yours)      # ほかの人の分は、番号を書かれても捨てない
    assert your_file.exists()


def test_the_leave_signal_discards_the_senders_table_import(sessions, client):
    app = sessions
    mine = _import(app, SESSION_A)
    yours = _import(app, SESSION_B, file_name="ほかの人.csv")
    folder = _import_folder(app, mine)
    with client.session_transaction() as cookie:
        cookie[SESSION_ID_KEY] = SESSION_A

    assert _as_beacon(client, "/tables/discard", f'{{"import_ids": [{mine}, {yours}]}}').status_code == 204
    assert not _alive(app, "table_imports", mine)
    assert not folder.exists()
    assert _alive(app, "table_imports", yours)


def test_the_leave_signal_never_fails(sessions, client):
    """もう無い番号・中身の無い body・壊れた body でも 204（ブラウザは答えを読めず、やり直せない）。"""
    app = sessions
    with client.session_transaction() as cookie:
        cookie[SESSION_ID_KEY] = SESSION_A
    for url, body in [("/forms/discard", '{"doc_ids": [999999]}'), ("/tables/discard", '{"import_ids": [999999]}'),
                      ("/forms/discard", ""), ("/tables/discard", "これはJSONではない")]:
        assert _as_beacon(client, url, body).status_code == 204, (url, body)


def test_the_forms_signal_does_not_touch_the_tables_tab(sessions, client):
    """同じブラウザで表の取り込みも開いていることがある。帳票の合図で表の分まで捨てない（逆も同じ）。"""
    app = sessions
    doc_id, _ = _document(app, "mine.xlsx", SESSION_A)
    import_id = _import(app, SESSION_A)
    with client.session_transaction() as cookie:
        cookie[SESSION_ID_KEY] = SESSION_A

    assert _as_beacon(client, "/forms/discard", "{}").status_code == 204   # 番号なし＝「自分の分ぜんぶ」
    assert not _alive(app, "documents", doc_id)
    assert _alive(app, "table_imports", import_id)

    assert _as_beacon(client, "/tables/discard", "{}").status_code == 204
    assert not _alive(app, "table_imports", import_id)


def test_the_leave_signal_does_not_pull_a_running_job_out(sessions, client):
    """読み込み中に画面を閉じても、動いているジョブの取り込みは捨てない（時間切れの片付けに任せる）。"""
    app = sessions
    import_id = _import(app, SESSION_A)
    _running_job(app, "table_import", import_id)
    with client.session_transaction() as cookie:
        cookie[SESSION_ID_KEY] = SESSION_A

    assert _as_beacon(client, "/tables/discard", f'{{"import_ids": [{import_id}]}}').status_code == 204
    assert _alive(app, "table_imports", import_id)
