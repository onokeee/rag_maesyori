"""帳票の読み取り: 区画（発行側と回答側）・日付の右の時刻・見出しと列見出しの間に1行ある明細表 のテスト。"""
from datetime import date, time

from excel.extractor import extract_document
from excel.tables import detect_tables, find_table, section_of, sections
from pattern.builder import merge_with_existing, suggest_rows
from pattern.forms import parse_pattern_form, pattern_to_rows, rows_to_pattern
from pattern.model import FieldDef, PatternDef
from tests.test_extraction import LABEL_FILL, _book, _values

HEAD_FILL = "FFFCE4D6"


def _pattern(*fields: FieldDef) -> PatternDef:
    return PatternDef(name="テスト", fields=list(fields))


def _renraku(tmp_path, name: str, answered: bool = True, answer_heading: str = "▼ 回答欄（宛先部署にて記入し返却）"):
    """「■ 異常連絡」の下（発行側）に「発生原因（推定）」「処置内容」、下の回答欄に「原因（確定）」「暫定対策（処置）」がある連絡票。"""
    cells = {"A1": "連絡No.", "B1": "PA-001",
             "A2": "■ 異常連絡",
             "A3": "発生原因（推定）", "B3": "研磨パッドの異常と思われる",
             "A4": "処置内容", "B4": "該当ロットをHOLD",
             "A5": answer_heading,
             "A6": "原因（確定）", "A7": "暫定対策（処置）"}
    if answered:
        cells.update({"B6": "パッド溝の摩耗", "B7": "研磨パッド交換"})
    fills = {c: LABEL_FILL for c in ("A1", "A3", "A4", "A6", "A7")}
    fills["A2"] = fills["A5"] = HEAD_FILL
    return _book(tmp_path, name, {"連絡票": (cells, fills, ["A2:B2", "A5:B5"])})


def _side_by_side(tmp_path, name: str):
    """左が発行側、右上の「【回答欄】」の下が回答側（A3横の版）。"""
    cells = {"A1": "工程異常連絡票", "D1": "【回答欄】",
             "A3": "連絡No.", "B3": "工異-001", "D3": "真因", "E3": "減速機の損傷",
             "D4": "応急処置", "E4": "減速機交換",
             "A6": "推定原因", "B6": "衝突の可能性", "A7": "処置", "B7": "全数選別"}
    fills = {c: LABEL_FILL for c in ("A3", "D3", "D4", "A6", "A7")}
    info = _book(tmp_path, name, {"連絡票": (cells, fills, [])})
    return info


# ---- 区画 ---------------------------------------------------------------------------------

def test_sections_split_the_sheet_by_marked_headings(tmp_path):
    info = _renraku(tmp_path, "a.xlsx")
    grid = info.grids["連絡票"]
    assert [s.key for s in sections(grid)] == ["異常連絡", "回答"]  # 印・括弧書き・末尾の「欄」を除いた名前
    assert section_of(grid, grid.cells[(1, 1)]) == ""          # 見出しより上
    assert section_of(grid, grid.cells[(4, 1)]) == "異常連絡"  # 処置内容（発行側）
    assert section_of(grid, grid.cells[(7, 1)]) == "回答"      # 暫定対策（処置）


def test_a_heading_in_the_right_half_makes_a_column_section(tmp_path, monkeypatch):
    from openpyxl import load_workbook
    from openpyxl.styles import Font

    info = _side_by_side(tmp_path, "b.xlsx")
    wb = load_workbook(info.path)
    wb["連絡票"]["D1"].font = Font(bold=True)  # 「【回答欄】」は太字だけ（塗りつぶしなし）
    wb.save(info.path)
    from excel.workbook import load_workbook_info

    grid = load_workbook_info(info.path).grids["連絡票"]
    assert section_of(grid, grid.cells[(4, 4)]) == "回答"  # 右側の応急処置
    assert section_of(grid, grid.cells[(7, 1)]) == ""      # 左側の処置


def test_field_with_a_section_reads_the_answer_side(tmp_path):
    info = _renraku(tmp_path, "a.xlsx")
    repair = FieldDef("repair", "処置", ["処置内容", "暫定対策", "応急処置"], data_type="text")
    cause = FieldDef("cause", "原因", ["原因（確定）", "発生原因", "推定原因"], data_type="text")
    values, _ = _values(info, _pattern(repair, cause))
    assert values == {"repair": "該当ロットをHOLD", "cause": "パッド溝の摩耗"}  # 区画なし: 上にある欄（今までどおり）

    repair.section = cause.section = "回答"
    values, fields = _values(info, _pattern(repair, cause))
    assert values == {"repair": "研磨パッド交換", "cause": "パッド溝の摩耗"}
    assert fields["repair"]["label_cell"] == "A7"


def test_unanswered_form_does_not_take_the_issuing_side_value(tmp_path):
    info = _renraku(tmp_path, "a.xlsx", answered=False)
    repair = FieldDef("repair", "処置", ["処置内容", "暫定対策"], data_type="text", section="回答")
    cause = FieldDef("cause", "原因", ["原因（確定）", "発生原因"], data_type="text", section="回答")
    values, fields = _values(info, _pattern(repair, cause))
    assert values == {"repair": None, "cause": None}  # 発行側の「処置内容」「発生原因（推定）」を読まない
    assert fields["cause"]["label_found"] and fields["cause"]["label_cell"] == "A6"


def test_layout_without_the_section_is_read_as_before(tmp_path):
    cells = {"A1": "処置内容", "B1": "該当ロットをHOLD"}
    info = _book(tmp_path, "c.xlsx", {"連絡票": (cells, {"A1": LABEL_FILL}, [])})
    fd = FieldDef("repair", "処置", ["処置内容"], data_type="text", section="回答")
    values, _ = _values(info, _pattern(fd))
    assert values["repair"] == "該当ロットをHOLD"


def test_builder_learns_the_answer_section_from_samples(tmp_path):
    both = _renraku(tmp_path, "a.xlsx")
    side = _side_by_side(tmp_path, "b.xlsx")
    from openpyxl import load_workbook
    from openpyxl.styles import Font

    wb = load_workbook(side.path)
    wb["連絡票"]["D1"].font = Font(bold=True)
    wb.save(side.path)
    from excel.workbook import load_workbook_info

    _, rows = suggest_rows([both, load_workbook_info(side.path)])
    by_name = {r["field_name"]: r for r in rows}
    assert by_name["repair"]["use"] and by_name["repair"]["section"] == "回答"
    assert by_name["cause"]["use"] and by_name["cause"]["section"] == "回答"
    assert by_name["report_id"]["section"] == ""  # 1か所にしかない項目には区画を付けない

    pattern = rows_to_pattern(1, {"name": "連絡票"}, [{"use": True, "sheet_name": "連絡票", "required": True}], rows)
    unanswered = _renraku(tmp_path, "d.xlsx", answered=False)
    values = extract_document(unanswered, pattern, ["連絡票"])["values"]
    assert values["repair"] is None and values["cause"] is None


def test_builder_leaves_the_section_empty_when_both_sides_are_common(tmp_path):
    """押印欄の「確認」のように、どの見本でも区画の外と回答欄の両方にある項目は、区画を決めない
    （どちら側の欄かを見本から決められない。今までどおり上にある欄を読み、人が画面で決める）。"""
    def book(name):
        cells = {"A1": "確認", "A2": "長谷川", "C1": "【回答欄】", "C2": "確認", "C3": "小川"}
        return _book(tmp_path, name, {"連絡票": (cells, {"A1": LABEL_FILL, "C1": HEAD_FILL, "C2": LABEL_FILL}, [])})

    _, rows = suggest_rows([book("a.xlsx"), book("b.xlsx")])
    assert all(r["section"] == "" for r in rows)


def test_section_round_trips_through_the_form():
    fd = FieldDef("repair", "処置", ["暫定対策"], data_type="text", section="回答")
    _, rows = pattern_to_rows(_pattern(fd))
    assert rows[0]["section"] == "回答"
    form = {"name": "連絡票", "fields-0-use": "on", "fields-0-field_name": "repair", "fields-0-display_name": "処置",
            "fields-0-candidates": "暫定対策", "fields-0-data_type": "text", "fields-0-section": "▼ 回答欄（宛先記入）"}
    meta, sheets, field_rows, errors = parse_pattern_form(form)
    assert not errors and field_rows[0]["section"] == "回答"  # 見出しのとおり入力しても比較用の名前にそろえる
    assert rows_to_pattern(1, meta, sheets, field_rows).fields[0].section == "回答"


def test_adding_samples_fills_an_empty_section_but_keeps_a_set_one():
    """見本の追加で見つかった区画は、登録済みの項目の区画が空のときだけ入れる（手で決めた区画は変えない）。"""
    existing = _pattern(FieldDef("repair", "処置", ["暫定対策"], data_type="text"),
                        FieldDef("cause", "原因", ["原因"], data_type="text", section="手入力"))
    found = [{"field_name": f, "display_name": d, "candidates": c, "data_type": "text", "examples": "", "seen": 2,
              "unit": "", "section": "回答", "use": True} for f, d, c in (("repair", "処置", "暫定対策"), ("cause", "原因", "原因"))]
    _, rows = merge_with_existing(existing, [], found)
    assert {r["field_name"]: r["section"] for r in rows} == {"repair": "回答", "cause": "手入力"}


# ---- ラベルの下の値: 左の見出しの値の欄を取らない ----------------------------------------------

def test_value_below_does_not_take_the_wide_value_of_the_label_on_the_left(tmp_path):
    # 品証コメント（A2）の値の欄（B2:D3）が、空のクローズ判定（C1）の真下まで広がっている
    cells = {"C1": "クローズ判定", "A2": "品証コメント", "B2": "効果確認未了。7/8頃に再確認"}
    info = _book(tmp_path, "e.xlsx", {"連絡票": (cells, {"C1": LABEL_FILL, "A2": LABEL_FILL}, ["B2:D3"])})
    values, _ = _values(info, _pattern(FieldDef("close", "クローズ判定", ["クローズ判定"]),
                                       FieldDef("qa", "品証コメント", ["品証コメント"], data_type="text")))
    assert values["close"] is None
    assert values["qa"] == "効果確認未了。7/8頃に再確認"


# ---- 日付の右のセルの時刻（md3-3）------------------------------------------------------------

def _date_time_book(tmp_path, name: str, cells: dict):
    base = {"A1": "発生日時", "E1": "作成者"}
    base.update(cells)
    return _book(tmp_path, name, {"報告書": (base, {"A1": LABEL_FILL, "E1": LABEL_FILL}, [])})


def test_time_in_the_next_cell_is_joined_to_the_date(tmp_path):
    info = _date_time_book(tmp_path, "t1.xlsx", {"B1": date(2023, 5, 23), "C1": "12:07"})
    values, fields = _values(info, _pattern(FieldDef("occurred", "発生日時", ["発生日時"], data_type="date")))
    assert values["occurred"] == "2023-05-23 12:07"
    assert fields["occurred"]["value_cell"] == "B1:C1" and not fields["occurred"]["warning"]

    info = _date_time_book(tmp_path, "t2.xlsx", {"B1": "2023/5/23", "C1": time(9, 5)})
    values, _ = _values(info, _pattern(FieldDef("occurred", "発生日時", ["発生日時"], data_type="date")))
    assert values["occurred"] == "2023-05-23 09:05"

    info = _date_time_book(tmp_path, "t3.xlsx", {"B1": "2023/5/23", "C1": "9時5分"})
    values, _ = _values(info, _pattern(FieldDef("occurred", "発生日時", ["発生日時"], data_type="date")))
    assert values["occurred"] == "2023-05-23 09:05"


def test_time_is_not_joined_when_it_is_ambiguous(tmp_path):
    fd = FieldDef("occurred", "発生日時", ["発生日時"], data_type="date")
    cases = {
        "range.xlsx": {"B1": "2023/5/23", "C1": "9:00", "D1": "17:00"},    # 時刻が2つ（範囲）
        "gap.xlsx": {"B1": "2023/5/23", "D1": "12:07"},                    # すぐ右ではない
        "clock.xlsx": {"B1": "2023/5/23 10:00", "C1": "12:07"},            # 日付に時刻が書いてある
        "about.xlsx": {"B1": "2023/5/23", "C1": "12:07頃"},                # 時刻だけではない
    }
    expected = {"range.xlsx": "2023-05-23", "gap.xlsx": "2023-05-23", "clock.xlsx": "2023-05-23 10:00",
                "about.xlsx": "2023-05-23"}
    for name, cells in cases.items():
        values, _ = _values(_date_time_book(tmp_path, name, cells), _pattern(fd))
        assert values["occurred"] == expected[name], name
    # 日付の項目でなければつながない
    values, _ = _values(_date_time_book(tmp_path, "s.xlsx", {"B1": "2023/5/23", "C1": "12:07"}),
                        _pattern(FieldDef("occurred", "発生日時", ["発生日時"])))
    assert values["occurred"] == "2023/5/23"


# ---- 見出しと列見出しの間に「項目｜値」の1行がある明細表 ----------------------------------------

def _yokoten(tmp_path, name: str, heading: str = "■ 水平展開"):
    cells = {"A1": heading, "A2": "展開区分", "B2": "☑同型機　□類似設備",
             "A3": "確認", "B3": "設備No", "C3": "設備名", "D3": "結果",
             "A4": "☑", "B4": "CMP-103", "C4": "W-CMP 3号機", "D4": "異常なし",
             "A5": "□", "B5": "CMP-104", "C5": "Cu-CMP 4号機",
             "A6": "展開先・内容", "B6": "同型機を点検"}
    fills = {c: LABEL_FILL for c in ("A2", "A3", "B3", "C3", "D3", "A6")}
    fills["A1"] = HEAD_FILL
    return _book(tmp_path, name, {"報告書": (cells, fills, ["A1:D1", "B2:D2", "B6:D6"])})


def test_table_under_a_heading_with_a_label_row_between_gets_the_heading(tmp_path):
    info = _yokoten(tmp_path, "y.xlsx")
    grid = info.grids["報告書"]
    [table] = [t for t in detect_tables(grid) if t.header[0].text == "確認"]
    assert table.anchor is not None and table.anchor.text == "■ 水平展開"
    assert table.to_value()["rows"] == [["☑", "CMP-103", "W-CMP 3号機", "異常なし"], ["□", "CMP-104", "Cu-CMP 4号機", ""]]
    # 見出しから探しても同じ表（「展開区分」の行を飛ばす）
    found = find_table(grid, grid.cells[(1, 1)])
    assert found is not None and [h.text for h in found.header] == ["確認", "設備No", "設備名", "結果"]
    values, _ = _values(info, _pattern(FieldDef("targets", "水平展開", ["■ 水平展開"], data_type="table"),
                                       FieldDef("kubun", "展開区分", ["展開区分"])))
    assert len(values["targets"]["rows"]) == 2 and values["kubun"] == "同型機"


def test_builder_suggests_the_table_under_a_heading_with_a_label_row_between(tmp_path):
    _, rows = suggest_rows([_yokoten(tmp_path, "y1.xlsx"), _yokoten(tmp_path, "y2.xlsx", "７．水平展開")])
    [row] = [r for r in rows if r["data_type"] == "table"]
    assert row["use"] and row["display_name"] == "水平展開"
    assert set(row["candidates"].splitlines()) >= {"■ 水平展開", "７．水平展開"}


def test_label_row_is_skipped_only_under_a_marked_heading(tmp_path):
    """見出しに「■」「1.」の印が無ければ、間の1行を飛ばして見出しにしない（今までどおり見出しの無い表）。"""
    info = _yokoten(tmp_path, "y.xlsx", heading="水平展開")
    grid = info.grids["報告書"]
    [table] = [t for t in detect_tables(grid) if t.header[0].text == "確認"]
    assert table.anchor is None
