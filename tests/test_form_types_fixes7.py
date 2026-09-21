"""帳票登録の直し（2枚目のシートの項目・手で直した見出し・項目を全部消したとき・Excel を置いていないとき）。

このまわりで見つかった不具合:
  - 同じ見出し語が2枚のシートにあると、2つ目の項目が1枚目の値をそのまま読む
  - 2枚目のシートに登録した項目が、記入位置がずれた帳票では1枚目の同じ番地の別の欄になる
  - 手で直した見出しが、次にセルをクリックしたときに元の見出しへ勝手に戻る
  - 使用中の種類から項目を全部削除でき、中身の無い Markdown を作り続ける
  - 使用開始の直後、登録した項目が全部「—（見つかりません）」になる

置いた Excel はサーバーに残らない（2026-09-21 の利用者の指示）ので、どの操作でも画面と同じように
その Excel を一緒に送る。
"""
import io

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from models import database as db
from tests.test_forms_flow import activate, add_field, book_part, create_type, panel_html, upload_forms


def _create(client, path, name="二次報告書") -> int:
    res = client.post("/form-types/new", data={"name": name, "book": book_part(path)},
                      content_type="multipart/form-data")
    assert res.status_code == 200, res.get_data(as_text=True)
    return res.get_json()["pattern_id"]


def _click(client, pattern_id, path, sheet, label_cell, value_cell=""):
    res = client.post(f"/form-types/{pattern_id}/fields", content_type="multipart/form-data",
                      data={"sheet": sheet, "label_cell": label_cell, "value_cell": value_cell,
                            "book": book_part(path)})
    assert res.status_code == 200, res.get_data(as_text=True)
    return res.get_json()


def _panel(client, pattern_id, path=None) -> str:
    return panel_html(client, pattern_id, book=path)


def _fill(ws, *coords) -> None:
    for coord in coords:
        ws[coord].fill = PatternFill("solid", fgColor="FFD9E1F2")


# ---- 同じ見出し語が2枚のシートにある帳票 ----------------------------------------------------

def _two_sheet_book(path, second_row: int = 2):
    """「独自見出し」が1次報告と2次報告の両方にある帳票（辞書に無い語なので別々の項目になる）。"""
    wb = Workbook()
    first = wb.active
    first.title = "1次報告"
    first["A3"], first["B3"] = "独自見出し", "1次の内容"
    _fill(first, "A3")
    second = wb.create_sheet("2次報告")
    second[f"A{second_row}"], second[f"B{second_row}"] = "独自見出し", "2次の内容"
    _fill(second, f"A{second_row}")
    wb.save(path)
    return path


def test_a_field_registered_on_the_second_sheet_reads_its_own_sheet(app, client, tmp_path):
    """2枚目のシートに登録した項目は、1枚目の同じ見出しの値で埋まらない。

    渡された順にシートを回り、最初に見出しが当たったシートで返していたので、Markdown に
    同じ「- 見出し: 値」が2行並び、2枚目の値はどこにも出なかった。
    """
    path = _two_sheet_book(tmp_path / "二枚.xlsx")
    pattern_id = _create(client, path)
    _click(client, pattern_id, path, "1次報告", "A3", "B3")
    _click(client, pattern_id, path, "2次報告", "A2", "B2")

    with app.app_context():
        fields = db.load_pattern(pattern_id).fields
    assert len(fields) == 2 and [f.sheet_name for f in fields] == ["1次報告", "2次報告"]

    # 読み取りテストの Markdown に、1次と2次の値がそれぞれ出る
    panel = _panel(client, pattern_id, path)
    assert "1次の内容" in panel and "2次の内容" in panel

    # 使用開始して取り込んでも同じ（2つ目の項目が自分のセルを読む）
    activate(client, pattern_id)
    doc_id, = upload_forms(client, path)
    res = client.post(f"/forms/{doc_id}/read",
                      data={"pattern_id": pattern_id, "sheets": ["1次報告", "2次報告"]})
    assert res.status_code == 200
    with app.app_context():
        extraction = db.get_document(doc_id)["data_json"]
    import json
    read = {f["field_name"]: f for f in json.loads(extraction)["fields"]}
    second = read[[f.field_name for f in fields][1]]
    assert second["value"] == "2次の内容" and second["sheet"] == "2次報告" and second["value_cell"] == "B2"


def test_a_lone_field_on_the_second_sheet_is_not_filled_from_the_first(app, client, tmp_path):
    """項目が1つでも同じ。2枚目に登録した項目が1枚目の同名の欄を読まない。"""
    path = _two_sheet_book(tmp_path / "二枚単独.xlsx")
    pattern_id = _create(client, path)
    _click(client, pattern_id, path, "2次報告", "A2", "B2")
    activate(client, pattern_id)

    doc_id, = upload_forms(client, path)
    assert client.post(f"/forms/{doc_id}/read",
                       data={"pattern_id": pattern_id, "sheets": ["1次報告", "2次報告"]}).status_code == 200
    import json
    with app.app_context():
        field = json.loads(db.get_document(doc_id)["data_json"])["fields"][0]
    assert field["value"] == "2次の内容" and field["sheet"] == "2次報告"


# ---- 2枚目のシートが選ばれない・別の欄の値になる ------------------------------------------------

def _attachment_book(path, note_row: int = 1):
    """「修理報告書」に項目が集まり、「別紙」に1項目だけある帳票（別紙の記入位置は版で変わる）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "修理報告書"
    cells = {"A1": "報告番号", "B1": "R-003", "A2": "設備番号", "B2": "EQ-9",
             "A3": "発生日", "B3": "2026-02-03", "A4": "報告者", "B4": "山田",
             "A5": "作業時間", "B5": 2, "A6": "処置内容", "B6": "交換した"}
    for coord, value in cells.items():
        ws[coord] = value
    _fill(ws, "A1", "A2", "A3", "A4", "A5", "A6")
    other = wb.create_sheet("別紙")
    other[f"A{note_row}"], other[f"B{note_row}"] = "再発防止策", "月次点検の項目に追加する"
    _fill(other, f"A{note_row}")
    wb.save(path)
    return path


def test_the_second_sheet_is_chosen_by_its_name_even_with_few_fields(app, client, tmp_path):
    """名前がそのまま同じシートは、項目の数で負けても「読み取るシート」に選ばれる。

    項目の数で競わせていたので、項目の少ない「別紙」が「修理報告書」に負けてチェックが外れ、
    そのまま読むと再発防止策に報告番号の値が入っていた（警告も出なかった）。
    """
    sample = _attachment_book(tmp_path / "見本.xlsx")
    pattern_id = _create(client, sample, name="設備修理報告書")
    for coord in ("A1", "A2", "A3", "A4", "A5", "A6"):
        _click(client, pattern_id, sample, "修理報告書", coord, f"B{coord[1:]}")
    _click(client, pattern_id, sample, "別紙", "A1", "B1")
    activate(client, pattern_id)

    # 別紙の記入位置が2行下にずれた帳票
    target = _attachment_book(tmp_path / "対象.xlsx", note_row=3)
    doc_id, = upload_forms(client, target)
    html = client.get(f"/forms/{doc_id}/type").get_json()["html"]
    assert '<input type="checkbox" name="sheets" value="別紙" checked' in html
    assert "7項目中7項目が見つかりました" in html


def test_a_field_is_never_filled_from_the_same_cell_of_another_sheet(app, client, tmp_path):
    """その項目のシートでセルが空でも、別のシートの同じ番地は読まない（空のままにする）。"""
    sample = _attachment_book(tmp_path / "見本2.xlsx")
    pattern_id = _create(client, sample, name="設備修理報告書2")
    for coord in ("A1", "A2", "A3", "A4", "A5", "A6"):
        _click(client, pattern_id, sample, "修理報告書", coord, f"B{coord[1:]}")
    _click(client, pattern_id, sample, "別紙", "A1", "B1")
    activate(client, pattern_id)

    target = _attachment_book(tmp_path / "対象2.xlsx", note_row=3)
    doc_id, = upload_forms(client, target)
    # 「修理報告書」だけを選んで読む（別紙のチェックを外したとき）
    assert client.post(f"/forms/{doc_id}/read",
                       data={"pattern_id": pattern_id, "sheets": ["修理報告書"]}).status_code == 200
    import json
    with app.app_context():
        fields = json.loads(db.get_document(doc_id)["data_json"])["fields"]
    note = next(f for f in fields if f["display_name"] == "再発防止策")
    assert note["value"] in (None, "") and "R-003" != note["value"]
    assert note["warning"] == "ラベルが見つかりません"


# ---- 手で直した見出しは、次のクリックで戻らない --------------------------------------------------

def _two_people_book(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "報告書"
    ws["A1"], ws["B1"] = "報告者", "長谷川 聡"
    ws["A2"], ws["B2"] = "担当者", "清水 彩花"
    _fill(ws, "A1", "A2")
    wb.save(path)
    return path


def test_a_hand_typed_label_survives_the_next_click(app, client, tmp_path):
    """見出しを手で「担当者」に直したあと、担当者の欄をクリックしても手の直しが消えない。"""
    path = _two_people_book(tmp_path / "二人.xlsx")
    pattern_id = _create(client, path, name="人の帳票")
    _click(client, pattern_id, path, "報告書", "A1", "B1")

    res = client.post(f"/form-types/{pattern_id}/fields/reporter/label", json={"name": "担当者"})
    assert res.status_code == 200 and "見出しを「担当者」にしました" in res.get_json()["message"]

    body = _click(client, pattern_id, path, "報告書", "A2", "B2")
    with app.app_context():
        fields = db.load_pattern(pattern_id).fields
    names = {f.field_name: f.display_name for f in fields}
    assert names["reporter"] == "担当者"                      # 手で付けた見出しはそのまま
    assert names["reporter_2"] == "担当者（A2）"              # 新しいほうを見分けられる名前にする
    assert "「担当者（A2）」を「担当者」とは別の項目にしました" in body["message"]


def test_two_auto_named_fields_are_still_told_apart_by_their_labels(app, client, tmp_path):
    """手で直していない項目どうしは、これまでどおりクリックした見出しで見分ける。"""
    path = _two_people_book(tmp_path / "二人2.xlsx")
    pattern_id = _create(client, path, name="人の帳票2")
    _click(client, pattern_id, path, "報告書", "A1", "B1")
    _click(client, pattern_id, path, "報告書", "A2", "B2")

    with app.app_context():
        fields = db.load_pattern(pattern_id).fields
    assert [f.display_name for f in fields] == ["報告者", "担当者"]


# ---- 項目を全部消したら、使用を停止する ---------------------------------------------------------

def test_deleting_the_last_field_stops_an_active_type(app, client, sample_dir):
    """使用中の種類から項目を全部消したら、帳票取り込みの候補に出なくなる。

    そのままだと「0項目中0項目が見つかりました」で読み取れてしまい、中身の無い .md ができていた。
    """
    pattern_id = create_type(client, sample_dir / "standard.xlsx", "設備修理報告書")
    add_field(client, pattern_id, "修理報告書", "A3", "B3")
    activate(client, pattern_id)

    res = client.post(f"/form-types/{pattern_id}/fields/report_id/delete")
    assert res.status_code == 200
    assert "使用を停止しました" in res.get_json()["message"]
    with app.app_context():
        assert db.load_pattern(pattern_id).status == "inactive"
        assert db.load_active_patterns() == []

    doc_id, = upload_forms(client, sample_dir / "standard.xlsx")
    body = client.get(f"/forms/{doc_id}/type").get_json()
    assert body["has_types"] is False


def test_deleting_one_of_several_fields_keeps_the_type_in_use(app, client, sample_dir):
    """項目が残っているなら、使用中のまま（削除のたびに止めない）。"""
    pattern_id = create_type(client, sample_dir / "standard.xlsx", "設備修理報告書")
    for label, value in (("A3", "B3"), ("E4", "F4")):
        add_field(client, pattern_id, "修理報告書", label, value)
    activate(client, pattern_id)

    res = client.post(f"/form-types/{pattern_id}/fields/report_id/delete")
    assert res.get_json()["message"] == "項目を削除しました"
    with app.app_context():
        assert db.load_pattern(pattern_id).status == "active"


# ---- Excel を置いていないときの文言（サーバーには残していない） ------------------------------------

def test_the_panel_without_a_book_says_the_excel_is_not_kept(app, client, sample_dir):
    """Excel を置いていない画面では、「見つかりません」「項目を1つ以上作ると」とは言わない。

    置いた Excel は残さないので、種類を開き直したときはシートも読み取りテストも出せない。
    そのことと「設定は残っている」ことを画面で言う。
    """
    path = sample_dir / "standard.xlsx"
    pattern_id = create_type(client, path, "設備修理報告書")
    add_field(client, pattern_id, "修理報告書", "A3", "B3")
    before = _panel(client, pattern_id, path)
    assert "1項目中 <strong>1</strong>項目が見つかりました" in before

    # 使用開始しても、置いている Excel はそのまま見られる（消すものがもう無い）
    res = client.post(f"/form-types/{pattern_id}/status",
                      data={"status": "active", "book": book_part(path)},
                      content_type="multipart/form-data")
    assert "1項目中 <strong>1</strong>項目が見つかりました" in res.get_json()["html"]
    assert "見本のExcelはサーバーから消しました" not in res.get_json()["message"]

    # 開き直したとき（Excel を置いていない）は、設定だけの画面になる
    panel = _panel(client, pattern_id)
    assert "—（見つかりません）" not in panel
    assert "—（Excelを置くと、この設定で読んだ値が出ます）" in panel
    assert "項目を1つ以上作ると、ここに読み取り結果が出ます。" not in panel
    assert "いま Excel を置いていないので、読み取りテストはできません" in panel
    assert "サーバーに残していません" in panel


# ---- 置いた Excel を替えたあと、項目を削除しても表示が戻らない ---------------------------------------

def test_deleting_a_field_keeps_the_book_that_is_being_looked_at(app, client, tmp_path):
    """2つ目の Excel を見ているときに項目を削除しても、その Excel の表示のままにする。"""
    first = _two_people_book(tmp_path / "一番.xlsx")
    pattern_id = _create(client, first, name="Excelを置き替える")
    second = _two_people_book(tmp_path / "二番.xlsx")

    _click(client, pattern_id, first, "報告書", "A1", "B1")
    assert "二番.xlsx" in _panel(client, pattern_id, second)

    # 画面の JS と同じように、いま見ている Excel を一緒に送る
    res = client.post(f"/form-types/{pattern_id}/fields/reporter/delete",
                      data={"book": book_part(second)}, content_type="multipart/form-data")
    assert res.status_code == 200
    html = res.get_json()["html"]
    assert "二番.xlsx" in html and 'data-cell="A1"' in html
