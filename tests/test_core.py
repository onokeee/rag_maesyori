from __future__ import annotations

import hashlib
import io
import json
import os
import re
import socket
import sqlite3
import struct
import threading
import time
import zipfile
from datetime import datetime, timedelta, date
from pathlib import Path

import openpyxl
import pytest
import yaml
from flask import Flask, current_app, send_file, render_template_string
from openpyxl import Workbook
from openpyxl.comments import Comment
from waitress.server import create_server
from werkzeug.datastructures import FileStorage

import app as app_module
import core
import database as db
import tables
from app import create_app
from core import (
    UploadError,
    precheck_excel,
    remove_upload,
    save_upload,
    upload_path,
    escape_md_line,
    estimate_tokens,
    join_blocks,
    md_bullet,
    nfkc_value,
    md_filename,
    safe_filename_part,
    FORM_MAX_MERGED_CELLS,
)
from forms import FieldDef, PatternDef, SheetDef
from logproc import (
    PeopleIndex,
    SplitOptions,
    apply_glossary,
    glossary_hits,
    mask_text,
    parse_log,
    render_timeline,
    review_notes,
    parse_when_at,
    resolve_whens,
    extract_identifiers,
    extract_plans,
    extract_quantities,
    shadow,
)
from tables import get_import
from tests.conftest import (
    add_confirmed_document,
    confirmed_import,
    CSV_TEXT,
    confirmed as _confirm_table,
    upload_csv,
    BufferedClient,
    imported,
    panel_html,
    make_config,
    OPENAI_KEY,
    FakeServer,
    columns_payload,
    editor_body,
    panel,
    save_columns,
    save_layout,
    save_source,
    upload_bytes,
    wait_import_job,
)
from tests.test_ai import ROWS, _make_import, ai_app, fake
from tests.test_forms import upload_forms



# ====================================================================================================
# 元 tests/test_core_db.py
# models/database.py: マイグレーション・PRAGMA・帳票のデータアクセス。
# ====================================================================================================

def make_app(tmp_path) -> Flask:
    """DB だけを初期化した最小のアプリ（app.py の画面構成に依存しない）。"""
    app = Flask(__name__)
    app.config.update(TESTING=True, DATABASE=tmp_path / "app.db", UPLOAD_DIR=tmp_path / "uploads")
    db.init_app(app)
    return app


@pytest.fixture
def core_app(tmp_path):
    return make_app(tmp_path)


# 作り直し前（user_version=0）のスキーマ
OLD_SCHEMA = """
CREATE TABLE patterns (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL, version TEXT NOT NULL DEFAULT 'v1',
    description TEXT NOT NULL DEFAULT '', status TEXT NOT NULL DEFAULT 'draft', image_processing TEXT NOT NULL DEFAULT 'none',
    created_at TEXT NOT NULL, updated_at TEXT NOT NULL);
CREATE TABLE pattern_sheets (id INTEGER PRIMARY KEY AUTOINCREMENT,
    pattern_id INTEGER NOT NULL REFERENCES patterns(id) ON DELETE CASCADE, sheet_name TEXT NOT NULL,
    required INTEGER NOT NULL DEFAULT 1);
CREATE TABLE pattern_fields (id INTEGER PRIMARY KEY AUTOINCREMENT,
    pattern_id INTEGER NOT NULL REFERENCES patterns(id) ON DELETE CASCADE, sort_order INTEGER NOT NULL DEFAULT 0,
    field_name TEXT NOT NULL, display_name TEXT NOT NULL, candidates TEXT NOT NULL DEFAULT '[]',
    required INTEGER NOT NULL DEFAULT 0, data_type TEXT NOT NULL DEFAULT 'string',
    extraction_rule TEXT NOT NULL DEFAULT '{}');
CREATE TABLE pattern_samples (id INTEGER PRIMARY KEY AUTOINCREMENT,
    pattern_id INTEGER NOT NULL REFERENCES patterns(id) ON DELETE CASCADE, file_name TEXT NOT NULL,
    file_hash TEXT NOT NULL, stored_path TEXT NOT NULL, created_at TEXT NOT NULL);
CREATE TABLE documents (id INTEGER PRIMARY KEY AUTOINCREMENT, file_name TEXT NOT NULL, file_hash TEXT NOT NULL,
    stored_path TEXT NOT NULL, pattern_id INTEGER REFERENCES patterns(id) ON DELETE SET NULL,
    status TEXT NOT NULL DEFAULT 'uploaded', data_json TEXT, markdown TEXT, created_at TEXT NOT NULL,
    registered_at TEXT);
"""

DATA = json.dumps({"values": {"report_id": "R-1"}}, ensure_ascii=False)


def _make_old_db(path):
    conn = sqlite3.connect(path)
    conn.executescript(OLD_SCHEMA)
    ts = "2026-01-01T00:00:00"
    conn.execute("INSERT INTO patterns (name, status, created_at, updated_at) VALUES (?, 'active', ?, ?)",
                 ("設備修理報告書", ts, ts))
    conn.execute("INSERT INTO pattern_fields (pattern_id, field_name, display_name, candidates, extraction_rule) "
                 "VALUES (1, 'report_id', '報告番号', ?, ?)", ('["報告番号"]', '{"direction": "right"}'))
    conn.execute("INSERT INTO documents (file_name, file_hash, stored_path, pattern_id, status, data_json, markdown, "
                 "created_at, registered_at) VALUES (?, 'h1', 'documents/a.xlsx', 1, 'registered', ?, '# md', ?, ?)",
                 ("登録済み.xlsx", DATA, "2026-02-01T10:00:00", "2026-02-02T11:00:00"))
    conn.execute("INSERT INTO documents (file_name, file_hash, stored_path, pattern_id, status, data_json, created_at) "
                 "VALUES (?, 'h2', 'documents/b.xlsx', 1, 'extracted', ?, ?)", ("確認中.xlsx", DATA, "2026-03-01T10:00:00"))
    conn.execute("INSERT INTO documents (file_name, file_hash, stored_path, status, created_at) "
                 "VALUES (?, 'h3', 'documents/c.xlsx', 'uploaded', ?)", ("読み取り前.xlsx", "2026-04-01T10:00:00"))
    conn.commit()
    conn.close()


def test_migrates_old_db_with_registered_document(tmp_path):
    _make_old_db(tmp_path / "app.db")
    app = make_app(tmp_path)
    with app.app_context():
        conn = db.get_db()
        assert conn.execute("PRAGMA user_version").fetchone()[0] == len(db.MIGRATIONS)
        assert conn.execute("PRAGMA journal_mode").fetchone()[0].lower() == "wal"
        assert conn.execute("PRAGMA foreign_keys").fetchone()[0] == 1
        assert conn.execute("PRAGMA busy_timeout").fetchone()[0] == db.BUSY_TIMEOUT_MS
        tables = {r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        assert {"table_imports", "jobs", "llm_calls", "ai_items"} <= tables
        # 使っていない表は残さない（取り込みを指す列が無く、行が入ると purge の探索から漏れる。design.md 3.3）
        assert tables.isdisjoint({"table_outputs", "table_downloads", "table_template_samples", "alias_entries"})
        # 取り込み設定は保存しない（表もビューも落とし、設定は取り込みの行が持つ）
        objects = {r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type IN ('table', 'view')")}
        assert objects.isdisjoint({"table_templates", "table_template_versions"})
        assert {"spec_json", "spec_hash"} <= {r[1] for r in conn.execute("PRAGMA table_info(table_imports)")}
        # AI整形の控えは取り込み単位で消せる（同じ設定の別の取り込みを巻き添えにしない）
        assert "import_id" in {r[1] for r in conn.execute("PRAGMA table_info(ai_items)")}

        doc = db.get_document(1)
        assert doc["state"] == "confirmed"
        assert doc["confirmed_json"] == DATA
        assert doc["confirmed_at"] == "2026-02-02T11:00:00"
        assert doc["title"] == ""
        assert db.get_document(2)["state"] == "reviewing"
        assert db.get_document(3)["state"] == "unread"

        pattern = db.load_pattern(1)
        assert pattern.title_fields == [] and pattern.md_options == {} and pattern.version_no == 1
        assert pattern.fields[0].unit == "" and pattern.fields[0].rag_output == "show"
        assert pattern.fields[0].direction == "right"
        assert db.list_patterns()[0]["document_count"] == 1
        assert db.find_confirmed_by_hash("h1")["id"] == 1

        # 古いDBに残っている pattern_fields.required（もう使わない「必須」）は書かない。
        # 列があっても既定値のまま保存できる
        assert "required" in {r[1] for r in conn.execute("PRAGMA table_info(pattern_fields)")}
        db.save_pattern(pattern, "active")
        assert [f.field_name for f in db.load_pattern(1).fields] == ["report_id"]
        assert conn.execute("SELECT required FROM pattern_fields WHERE pattern_id = 1").fetchone()[0] == 0

    # 2回目の起動でも壊れない（冪等）
    app2 = make_app(tmp_path)
    with app2.app_context():
        assert db.get_document(1)["state"] == "confirmed"
        assert db.get_db().execute("SELECT COUNT(*) FROM documents").fetchone()[0] == 3


def test_new_db_document_states_and_filters(core_app):
    with core_app.app_context():
        pid = db.create_pattern("設備修理報告書")
        d1 = db.create_document("修理報告書_標準.xlsx", "hash1", "forms/x.xlsx", pattern_id=pid)
        d2 = db.create_document("点検記録表.xlsx", "hash2", "forms/y.xlsx")
        assert db.get_document(d1)["state"] == "unread"

        db.save_draft(d1, {"values": {"report_id": "R2026-00123"}}, title="R2026-00123 CMP-101")
        assert db.get_document(d1)["state"] == "reviewing"
        assert db.find_confirmed_by_hash("hash1") is None
        assert db.confirm_document(d2) is False  # 読み取り前は確定できない

        assert db.confirm_document(d1, title="R2026-00123 CMP研磨装置1号機")
        doc = db.get_document(d1)
        assert doc["state"] == "confirmed" and doc["confirmed_at"]
        assert db.find_confirmed_by_hash("hash1")["id"] == d1
        assert db.find_confirmed_by_hash("hash1", exclude_id=d1) is None

        db.save_draft(d1, {"values": {"report_id": "R2026-00124"}})
        assert db.get_document(d1)["state"] == "modified"
        assert json.loads(db.get_document(d1)["confirmed_json"])["values"]["report_id"] == "R2026-00123"

        assert [d["id"] for d in db.list_confirmed_documents()] == [d1]
        assert db.list_confirmed_documents([]) == []

        # まとめ取り込み（同じ batch_id を選んだ順に返す）
        b1 = db.create_document("1.xlsx", "h1", "documents/1.xlsx", batch_id="B", batch_order=0)
        b2 = db.create_document("2.xlsx", "h2", "documents/2.xlsx", batch_id="B", batch_order=1)
        assert [d["id"] for d in db.list_batch_documents("B")] == [b1, b2]
        assert db.list_batch_documents("") == [] and db.list_batch_documents("ない") == []
        assert db.get_document(b1)["batch_id"] == "B"


def test_save_pattern_new_fields_and_version_no(core_app):
    with core_app.app_context():
        pid = db.create_pattern("設備修理報告書")
        fd = FieldDef("work_hours", "作業時間", ["作業時間"], data_type="number")
        fd.unit = "時間"
        fd.rag_output = "omit"
        pattern = PatternDef(name="設備修理報告書", id=pid, sheets=[SheetDef("修理報告書")], fields=[fd])
        pattern.title_fields = ["report_id", "equipment"]
        pattern.md_options = {"domain_context_fields": ["equipment"]}
        db.save_pattern(pattern, "active")
        assert pattern.version_no == 2

        loaded = db.load_pattern(pid)
        assert loaded.title_fields == ["report_id", "equipment"]
        assert loaded.md_options == {"domain_context_fields": ["equipment"]}
        assert loaded.version_no == 2 and loaded.status == "active"
        assert loaded.fields[0].unit == "時間" and loaded.fields[0].rag_output == "omit"

        # 新しい属性を持たない PatternDef / FieldDef でも保存できる
        db.save_pattern(PatternDef(name="別名", id=pid, fields=[FieldDef("a", "A")]), "draft")
        loaded = db.load_pattern(pid)
        assert loaded.version_no == 3 and loaded.title_fields == [] and loaded.fields[0].rag_output == "show"


def test_field_section_survives_save_and_load(core_app):
    """探す区画は extraction_rule に入れて保存する。保存し直すと消える、では画面で決めた区画が効かない。"""
    with core_app.app_context():
        pid = db.create_pattern("不具合連絡票")
        fields = [FieldDef("repair", "処置", ["暫定対策"], data_type="text", section="回答"), FieldDef("no", "No.", ["No."])]
        db.save_pattern(PatternDef(name="不具合連絡票", id=pid, fields=fields), "active")
        loaded = db.load_pattern(pid)
        assert [f.section for f in loaded.fields] == ["回答", ""]
        rules = [json.loads(r["extraction_rule"]) for r in db.get_db().execute(
            "SELECT extraction_rule FROM pattern_fields WHERE pattern_id = ? ORDER BY sort_order", (pid,))]
        assert rules[0]["section"] == "回答" and "section" not in rules[1]


def test_foreign_keys_cascade_with_background_connection(core_app):
    with core_app.app_context():
        conn = db.connect()
        try:
            ts = db.now()
            pid = conn.execute("INSERT INTO patterns (name, created_at, updated_at) VALUES (?, ?, ?)",
                               ("不具合連絡票", ts, ts)).lastrowid
            conn.execute("INSERT INTO pattern_sheets (pattern_id, sheet_name) VALUES (?, '報告書')", (pid,))
            conn.execute("INSERT INTO pattern_fields (pattern_id, field_name, display_name) VALUES (?, 'a', 'A')", (pid,))
            did = conn.execute("INSERT INTO documents (file_name, file_hash, stored_path, pattern_id, created_at) "
                               "VALUES ('a.xlsx', 'h', 'forms/a.xlsx', ?, ?)", (pid, ts)).lastrowid
            conn.commit()
            conn.execute("DELETE FROM patterns WHERE id = ?", (pid,))
            conn.commit()
            assert conn.execute("SELECT COUNT(*) FROM pattern_fields").fetchone()[0] == 0
            assert conn.execute("SELECT COUNT(*) FROM pattern_sheets").fetchone()[0] == 0
            assert conn.execute("SELECT pattern_id FROM documents WHERE id = ?", (did,)).fetchone()[0] is None
        finally:
            conn.close()


def test_the_import_spec_can_be_written_after_the_templates_are_gone(core_app):
    """table_templates を落としたあとも、取り込みの行は書ける（外部キーの参照先が消えていない）。"""
    with core_app.app_context():
        conn = db.connect()
        try:
            ts = db.now()
            iid = conn.execute("INSERT INTO table_imports (file_name, file_hash, stored_path, spec_json, spec_hash, "
                               "created_at, updated_at) VALUES ('a.csv', 'h', 'tables/a.csv', '{}', 'x', ?, ?)",
                               (ts, ts)).lastrowid
            conn.execute("UPDATE table_imports SET template_id = id, template_version_id = id WHERE id = ?", (iid,))
            conn.commit()
            # aiproc.runner.load_spec は取り込みの行から設定を読む
            row = conn.execute("SELECT spec_json FROM table_imports WHERE id = ?", (iid,)).fetchone()
            assert row["spec_json"] == "{}"
        finally:
            conn.close()


# ====================================================================================================
# 元 tests/test_core_files.py
# core/files.py: 保存（分割読み sha256）と Excel の事前チェック。
# ====================================================================================================

@pytest.fixture
def core_app_core_files(tmp_path):
    app = Flask(__name__)
    app.config.update(TESTING=True, UPLOAD_DIR=tmp_path / "uploads")
    return app


def _storage(data: bytes, name: str) -> FileStorage:
    return FileStorage(stream=io.BytesIO(data), filename=name)


def _xlsx(path):
    wb = Workbook()
    wb.active["A1"] = "管理No"
    wb.save(path)
    return path


def test_save_upload_hash_and_remove(core_app_core_files, monkeypatch):
    monkeypatch.setattr(core, "CHUNK_SIZE", 7)  # 分割読みを確かめる
    data = "管理No,設備\nTR-1,CMP-101\n".encode("cp932") * 50
    with core_app_core_files.app_context():
        stored = save_upload(_storage(data, "C:\\fakepath\\故障履歴.csv"), "tables", {".csv", "xlsx"}, 10_000)
        assert stored.file_name == "故障履歴.csv"
        assert stored.stored_path.startswith("tables/") and stored.stored_path.endswith(".csv")
        assert stored.file_hash == hashlib.sha256(data).hexdigest()
        assert stored.size == len(data)
        path = upload_path(stored.stored_path)
        assert path.read_bytes() == data
        remove_upload(stored.stored_path)
        assert not path.exists()
        remove_upload(stored.stored_path)  # 2回目も例外にしない


def test_save_upload_rejects(core_app_core_files):
    with core_app_core_files.app_context():
        with pytest.raises(UploadError, match="保存し直して"):
            save_upload(_storage(b"x", "old.xls"), "forms", {".xlsx"}, 100)
        with pytest.raises(UploadError, match="大きすぎます"):
            save_upload(_storage(b"x" * 101, "big.csv"), "tables", {".csv"}, 100)
        with pytest.raises(UploadError, match="空です"):
            save_upload(_storage(b"", "empty.csv"), "tables", {".csv"}, 100)
        with pytest.raises(UploadError, match="選択"):
            save_upload(_storage(b"x", ""), "tables", {".csv"}, 100)
        # 失敗したファイルは残さない
        folder = core_app_core_files.config["UPLOAD_DIR"] / "tables"
        assert not folder.exists() or not any(folder.iterdir())
        with pytest.raises(UploadError):
            upload_path("../outside.txt")


def test_precheck_accepts_normal_xlsx(tmp_path):
    precheck_excel(_xlsx(tmp_path / "ok.xlsx"))


def test_precheck_ole_header(tmp_path):
    path = tmp_path / "locked.xlsx"
    path.write_bytes(bytes.fromhex("D0CF11E0A1B11AE1") + b"\x00" * 504)
    with pytest.raises(UploadError, match="パスワード付き"):
        precheck_excel(path)


def test_precheck_not_zip(tmp_path):
    path = tmp_path / "text.xlsx"
    path.write_text("管理No,設備", encoding="utf-8")
    with pytest.raises(UploadError, match="読み込めません"):
        precheck_excel(path)


def test_precheck_xlsb(tmp_path):
    path = tmp_path / "book.xlsx"
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("xl/workbook.bin", b"\x83\x01\x00")
    with pytest.raises(UploadError, match="xlsb"):
        precheck_excel(path)


def test_precheck_strict(tmp_path):
    path = tmp_path / "strict.xlsx"
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("xl/workbook.xml", '<?xml version="1.0"?><workbook '
                    'xmlns="http://purl.oclc.org/ooxml/spreadsheetml/main"><sheets/></workbook>')
    with pytest.raises(UploadError, match="Strict"):
        precheck_excel(path)


def test_precheck_zip_bomb_ratio(tmp_path):
    path = tmp_path / "bomb.xlsx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("xl/workbook.xml", "<workbook/>")
        zf.writestr("xl/worksheets/sheet1.xml", b"\x00" * (core.ZIP_RATIO_MIN_BYTES + 1024))
    assert path.stat().st_size < 1024 * 1024
    with pytest.raises(UploadError, match="圧縮率"):
        precheck_excel(path)


def test_precheck_zip_limits(tmp_path, monkeypatch):
    path = tmp_path / "large.xlsx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("xl/workbook.xml", "<workbook/>")
        zf.writestr("xl/worksheets/sheet1.xml", b"a" * 5000)
        zf.writestr("xl/worksheets/sheet2.xml", b"b" * 5000)
    monkeypatch.setattr(core, "ZIP_MAX_PART", 4000)
    with pytest.raises(UploadError, match="1つの部品"):
        precheck_excel(path)
    monkeypatch.setattr(core, "ZIP_MAX_PART", 6000)
    monkeypatch.setattr(core, "ZIP_MAX_TOTAL", 8000)
    with pytest.raises(UploadError, match="合計"):
        precheck_excel(path)


def test_remove_upload_empties_a_locked_file_and_reports_it(core_app_core_files, monkeypatch):
    """他のプロセスに掴まれていて消せないときは、例外を出さずに中身を空にして False を返す。

    「ダウンロードしたらサーバーから消えます」（design.md 3.3）と言い切っているので、ファイル名が残っても
    元の Excel/CSV の中身は残さない（空になったファイルは次の起動時に片付く）。
    """
    with core_app_core_files.app_context():
        stored = save_upload(_storage("社外秘の中身".encode("utf-8") * 10, "掴まれたブック.xlsx"),
                             "documents", {".xlsx"}, 10_000)
        path = upload_path(stored.stored_path)

        def locked(self, missing_ok=False):
            raise PermissionError(32, "別のプロセスが使用中です")

        monkeypatch.setattr(Path, "unlink", locked)
        monkeypatch.setattr(core, "REMOVE_RETRY_WAIT", 0)
        assert remove_upload(stored.stored_path) is False   # 例外を出さず、消せなかったことを返す
        assert path.exists() and path.read_bytes() == b""   # 中身は残さない

        monkeypatch.undo()
        assert remove_upload(stored.stored_path) is True
        assert not path.exists()


def test_remove_orphan_uploads(core_app_core_files):
    """DB から参照されていない取り込み済みファイルだけを片付ける。"""
    with core_app_core_files.app_context():
        used = save_upload(_storage(b"a" * 100, "使用中.xlsx"), "documents", {".xlsx"}, 10_000)
        orphan = save_upload(_storage(b"b" * 100, "残骸.xlsx"), "tables", {".xlsx"}, 10_000)
        fresh = save_upload(_storage(b"c" * 100, "保存したばかり.xlsx"), "documents", {".xlsx"}, 10_000)
        other = upload_path("documents/メモ.txt")
        other.write_text("save_upload が作った名前ではないファイル", encoding="utf-8")
        old = time.time() - 3600
        for stored in (used, orphan):
            os.utime(upload_path(stored.stored_path), (old, old))

        removed = core.remove_orphan_uploads(core_app_core_files.config["UPLOAD_DIR"], [used.stored_path])

        assert removed == 1
        assert upload_path(used.stored_path).exists()
        assert not upload_path(orphan.stored_path).exists()
        assert other.exists()
        # できたばかりのファイルは、別に起動しているアプリが DB の行を作る前かもしれないので消さない
        assert upload_path(fresh.stored_path).exists()


# ---- 結合セルの面積（openpyxl で開く前に止める） ------------------------------------------------

def _xlsx_with_merge(path, ref: str, *, pad: int = 0):
    """ふつうの xlsx を作り、シートの XML に <mergeCell ref="..."> を書き足す（pad: その前に入れる空白の量）。"""
    src = _xlsx(path.with_name("src_" + path.name))
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                merge = f'{" " * pad}<mergeCells count="1"><mergeCell ref="{ref}"/></mergeCells>'.encode()
                data = data.replace(b"</sheetData>", b"</sheetData>" + merge, 1)
                if b"<sheetData/>" in data:
                    data = data.replace(b"<sheetData/>", b"<sheetData/>" + merge, 1)
            zout.writestr(info, data)
    return path


def test_precheck_refuses_a_whole_sheet_merge_before_openpyxl(tmp_path):
    import time
    path = _xlsx_with_merge(tmp_path / "merge_full.xlsx", "A1:XFD1048576")
    started = time.monotonic()
    with pytest.raises(UploadError, match="結合セルの範囲が大きすぎます"):
        precheck_excel(path)
    assert time.monotonic() - started < 2


def test_precheck_accepts_whole_row_and_whole_column_merges(tmp_path):
    precheck_excel(_xlsx_with_merge(tmp_path / "row.xlsx", "A1:XFD1"))
    precheck_excel(_xlsx_with_merge(tmp_path / "col.xlsx", "A1:A1048576"))


def test_precheck_for_forms_refuses_many_whole_row_merges_quickly(tmp_path):
    """帳票は画面を開くたびに通常モードで開き直す。行全体の結合 120 個（約200万セル）は1回に約16秒かかるので断る。"""
    many = _xlsx_with_merge(tmp_path / "rows120.xlsx", "A1:XFD1")
    many.write_bytes(_rewrite_merges(many, [f"A{r}:XFD{r}" for r in range(1, 121)]))
    precheck_excel(many)   # 一覧表（読み取り専用で開き、結合を展開しない）はこれまでどおり通す
    started = time.monotonic()
    with pytest.raises(UploadError, match="結合セルの範囲が大きすぎます"):
        precheck_excel(many, max_merged=core.FORM_MAX_MERGED_CELLS)
    assert time.monotonic() - started < 2
    with pytest.raises(UploadError, match="結合セル"):   # 列全体の結合（罫線付きなら1分以上かかる）も
        precheck_excel(_xlsx_with_merge(tmp_path / "col.xlsx", "A1:A1048576"), max_merged=core.FORM_MAX_MERGED_CELLS)
    # 行全体の結合が数個の帳票は通す
    few = _xlsx_with_merge(tmp_path / "rows12.xlsx", "A1:XFD1")
    few.write_bytes(_rewrite_merges(few, [f"A{r}:XFD{r}" for r in range(1, 13)]))
    precheck_excel(few, max_merged=core.FORM_MAX_MERGED_CELLS)


def _rewrite_merges(path, refs):
    """_xlsx_with_merge で作ったブックの結合範囲を refs に置き換えた中身を返す。"""
    out = io.BytesIO()
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                merges = "".join(f'<mergeCell ref="{ref}"/>' for ref in refs)
                data = re.sub(rb"<mergeCells\b.*?</mergeCells>",
                              f'<mergeCells count="{len(refs)}">{merges}</mergeCells>'.encode(), data, flags=re.S)
            zout.writestr(info, data)
    return out.getvalue()


def test_precheck_finds_a_merge_split_across_read_chunks(tmp_path, monkeypatch):
    monkeypatch.setattr(core, "CHUNK_SIZE", 64)   # タグがチャンクの境目で切れるようにする
    for pad in range(0, 64, 7):
        path = _xlsx_with_merge(tmp_path / f"split{pad}.xlsx", "C1:Z200000", pad=pad)
        with pytest.raises(UploadError, match="結合セル"):
            precheck_excel(path)


def test_precheck_counts_each_merge_once(tmp_path, monkeypatch):
    monkeypatch.setattr(core, "CHUNK_SIZE", 64)
    monkeypatch.setattr(core, "MAX_MERGED_CELLS", 100)
    precheck_excel(_xlsx_with_merge(tmp_path / "exact.xlsx", "A1:J10"))   # ちょうど100セルは通す


def upload_error(client, url: str, data: bytes, name: str) -> str:
    """ファイルを置いたときの断りの文（画面は fetch で受け取って、そのまま出す）。"""
    res = client.post(url, data={"file": (io.BytesIO(data), name)}, content_type="multipart/form-data")
    assert res.status_code == 400, (url, res.status_code)
    return res.get_json()["error"]


def test_uploads_with_a_whole_sheet_merge_are_refused_in_japanese(app, client, tmp_path):
    path = _xlsx_with_merge(tmp_path / "merge_full.xlsx", "A1:XFD1048576")
    for url in ("/forms/upload", "/tables/upload"):
        assert "結合セルの範囲が大きすぎます" in upload_error(client, url, path.read_bytes(), "結合.xlsx"), url
    assert [p for p in Path(app.config["UPLOAD_DIR"]).rglob("*") if p.is_file()] == []


def test_precheck_reports_a_damaged_compressed_part(tmp_path):
    """圧縮データが壊れている（zlib.error）ときも UploadError（500 にしない）。"""
    path = _xlsx(tmp_path / "ok.xlsx")
    data = bytearray(path.read_bytes())
    with zipfile.ZipFile(path) as zf:
        info = zf.getinfo("xl/workbook.xml")
    start = info.header_offset + 30 + len(info.filename.encode()) + len(info.extra)
    for i in range(start, start + min(info.compress_size, 64)):
        data[i] ^= 0xFF
    broken = tmp_path / "broken.xlsx"
    broken.write_bytes(bytes(data))
    with pytest.raises(UploadError, match="壊れている"):
        precheck_excel(broken)


def test_sheetless_workbook_is_refused_for_both_flows(app, client, tmp_path):
    """シートの無いブックは帳票でも一覧表でも受け付けず、アップロードしたファイルも残さない。"""
    import re

    path = _xlsx(tmp_path / "s.xlsx")
    out = io.BytesIO()
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(out, "w") as z:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/workbook.xml":
                data = re.sub(rb"<sheets>.*?</sheets>", b"<sheets/>", data, flags=re.S)
            z.writestr(item, data)
    sheetless = tmp_path / "sheetless.xlsx"
    sheetless.write_bytes(out.getvalue())
    with pytest.raises(UploadError, match="シートがないブック"):
        precheck_excel(sheetless)
    precheck_excel(path)   # ふつうのブックは通る
    for url in ("/forms/upload", "/tables/upload"):
        assert "シートがないブックです" in upload_error(client, url, out.getvalue(), "シートなし.xlsx"), url
    assert [p for p in Path(app.config["UPLOAD_DIR"]).rglob("*") if p.is_file()] == []


def _xlsx_with_cells(path, count: int):
    """ふつうの xlsx のシートに、書式だけのセル <c/> を count 個書き足す（圧縮すると小さい）。"""
    src = _xlsx(path.with_name("src_" + path.name))
    rows = b'<row r="9999">' + b'<c s="0"/>' * count + b"</row>"
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                data, n = re.subn(rb"<sheetData\s*/>", b"<sheetData>" + rows + b"</sheetData>", data, 1)
                if not n:
                    data = data.replace(b"</sheetData>", rows + b"</sheetData>", 1)
            zout.writestr(info, data)
    return path


def test_precheck_refuses_too_many_cells_in_a_small_file(tmp_path):
    """圧縮すると小さいが展開すると大量のセルがあるブックは、openpyxl で開く前に断る。"""
    path = _xlsx_with_cells(tmp_path / "many.xlsx", core.MAX_CELLS + 1)
    assert path.stat().st_size < 200 * 1024
    with pytest.raises(UploadError, match="セル数が上限"):
        precheck_excel(path)
    precheck_excel(path, max_cells=core.MAX_CELLS + 100)   # 上限を上げれば通る


def test_precheck_counts_cells_split_across_read_chunks_once(tmp_path, monkeypatch):
    monkeypatch.setattr(core, "CHUNK_SIZE", 64)
    path = _xlsx_with_cells(tmp_path / "exact.xlsx", 500)
    with zipfile.ZipFile(path) as zf:
        existing = len(re.findall(rb"<(?:\w+:)?c(?=[\s/>])", zf.read("xl/worksheets/sheet1.xml")))
    precheck_excel(path, max_cells=existing)          # ちょうど上限は通す（二重に数えない）
    with pytest.raises(UploadError, match="セル数が上限"):
        precheck_excel(path, max_cells=existing - 1)  # 1つでも多ければ断る（数え落とさない）


def _xlsx_with_sheet_xml(path, edit, *, rename: str | None = None):
    """ふつうの xlsx のシートの XML を edit(bytes) で書き換える（rename: シートの部品の名前を変える）。"""
    src = _xlsx(path.with_name("src_" + path.name))
    sheet = "xl/worksheets/sheet1.xml"
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            name = info.filename
            if name == sheet:
                data = edit(data)
                name = rename or name
            elif rename and name in ("xl/_rels/workbook.xml.rels", "[Content_Types].xml"):
                data = data.replace(b"worksheets/sheet1.xml", rename.removeprefix("xl/").encode())
            zout.writestr(name, data)
    return path


def _add_merge(ref_attr: bytes):
    def edit(data: bytes) -> bytes:
        merge = b"<mergeCells count=\"1\"><mergeCell " + ref_attr + b"/></mergeCells>"
        if b"<sheetData/>" in data:
            return data.replace(b"<sheetData/>", b"<sheetData/>" + merge, 1)
        return data.replace(b"</sheetData>", b"</sheetData>" + merge, 1)
    return edit


def _add_cells(row: bytes):
    """<row> を1つシートに書き足す（sheetData が空要素でも中身があっても）。"""
    def edit(data: bytes) -> bytes:
        data, n = re.subn(rb"<sheetData\s*/>", b"<sheetData>" + row + b"</sheetData>", data, count=1)
        return data if n else data.replace(b"</sheetData>", row + b"</sheetData>", 1)
    return edit


def test_precheck_reads_a_sheet_part_whatever_its_name(tmp_path):
    """シートの部品の名前は workbook.xml.rels で決まる（.xml で終わらなくても openpyxl は開く）。"""
    import openpyxl

    small = _xlsx_with_sheet_xml(tmp_path / "small.xlsx", _add_merge(b'ref="C3:E5"'),
                                 rename="xl/worksheets/sheet1.dat")
    precheck_excel(small)
    assert [str(r) for r in openpyxl.load_workbook(small).active.merged_cells.ranges] == ["C3:E5"]
    big = _xlsx_with_sheet_xml(tmp_path / "big.xlsx", _add_merge(b'ref="C3:XFD1048576"'),
                               rename="xl/worksheets/sheet1.dat")
    with pytest.raises(UploadError, match="結合セルの範囲が大きすぎます"):
        precheck_excel(big)
    cells = _xlsx_with_sheet_xml(tmp_path / "cells.xlsx", _add_cells(b'<row r="9">' + b'<c s="0"/>' * 50 + b"</row>"),
                                 rename="xl/worksheets/sheet1.dat")
    with pytest.raises(UploadError, match="セル数が上限"):
        precheck_excel(cells, max_cells=40)


def test_precheck_reads_merge_refs_the_way_an_xml_parser_does(tmp_path):
    """= の前後の空白・文字参照（&#88; = X）も、XML として正しい書き方なので openpyxl は読む。数え落とさない。"""
    import openpyxl

    for i, attr in enumerate((b'ref = "C3:XFD1048576"', b'ref="C3:&#88;FD1048576"', b"ref\n=\n'C3:XFD1048576'")):
        path = _xlsx_with_sheet_xml(tmp_path / f"m{i}.xlsx", _add_merge(attr))
        with pytest.raises(UploadError, match="結合セルの範囲が大きすぎます"):
            precheck_excel(path)
    small = _xlsx_with_sheet_xml(tmp_path / "small.xlsx", _add_merge(b'ref = "C3:&#69;5"'))
    precheck_excel(small)
    assert [str(r) for r in openpyxl.load_workbook(small).active.merged_cells.ranges] == ["C3:E5"]


def test_precheck_counts_cells_under_any_namespace_prefix(tmp_path):
    """<x.y:c> のような接頭辞でも、SpreadsheetML の名前空間の <c> ならセルとして数える。"""
    main = b"http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    row = b'<row r="9" xmlns:x.y-z="' + main + b'">' + b'<x.y-z:c s="0"/>' * 50 + b"</row>"
    path = _xlsx_with_sheet_xml(tmp_path / "prefixed.xlsx", _add_cells(row))
    with pytest.raises(UploadError, match="セル数が上限"):
        precheck_excel(path, max_cells=40)
    precheck_excel(path, max_cells=100)


def test_save_upload_and_precheck_messages_carry_no_file_name_or_class_name(app, tmp_path):
    """flash に入るメッセージにはファイル名も例外の種類名も入れない（design.md 3.3）。"""
    from werkzeug.datastructures import FileStorage

    with app.app_context():
        for data, allowed in ((b"x", {".csv"}), (b"", {".xlsx"})):
            with pytest.raises(UploadError) as err:
                save_upload(FileStorage(io.BytesIO(data), "秘密の名前.xlsx"), "documents", allowed, 1024)
            assert "秘密の名前" not in str(err.value)
    broken = tmp_path / "broken.xlsx"
    broken.write_bytes(b"PK\x03\x04" + b"\x00" * 64)
    with pytest.raises(UploadError) as err:
        precheck_excel(broken)
    assert "BadZipFile" not in str(err.value) and "壊れている" in str(err.value)


def test_precheck_refuses_entity_definitions(tmp_path):
    """DTD で実体を定義して展開させる細工（正しいブックには無い）は、数える前に断る。"""
    def edit(data: bytes) -> bytes:
        return b'<!DOCTYPE worksheet [<!ENTITY a "aaaaaaaaaa">]>' + data.split(b"?>", 1)[-1]

    with pytest.raises(UploadError, match="不正なファイル"):
        precheck_excel(_xlsx_with_sheet_xml(tmp_path / "entity.xlsx", edit))


_XDR = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
_DRAWING_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing"


def _xlsx_sharing_one_drawing(path, sheets: int, anchors: int, absolute_target: bool = False):
    """多数のシートが同じ1つの描画部品（大量の図形）を参照するブック（圧縮後は小さい）。"""
    wb = Workbook()
    wb.active.title = "S1"
    for i in range(2, sheets + 1):
        wb.create_sheet(f"S{i}")
    wb.save(path)
    anchor = ('<xdr:absoluteAnchor><xdr:pos x="0" y="0"/><xdr:ext cx="1" cy="1"/>'
              '<xdr:clientData/></xdr:absoluteAnchor>')
    drawing = f'<?xml version="1.0"?><xdr:wsDr xmlns:xdr="{_XDR}">{anchor * anchors}</xdr:wsDr>'
    target = "/xl/drawings/drawing1.xml" if absolute_target else "../drawings/drawing1.xml"
    rels = ('<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            f'<Relationship Id="rIdD" Type="{_DRAWING_REL}" Target="{target}"/></Relationships>')
    src = path.with_suffix(".src.xlsx")
    path.rename(src)
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", item.filename):
                text = data.decode("utf-8").replace(
                    "<worksheet ", '<worksheet xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" ', 1)
                data = text.replace("</worksheet>", '<drawing r:id="rIdD"/></worksheet>').encode("utf-8")
            zout.writestr(item, data)
        for i in range(1, sheets + 1):
            zout.writestr(f"xl/worksheets/_rels/sheet{i}.xml.rels", rels)
        zout.writestr("xl/drawings/drawing1.xml", drawing)
    src.unlink()
    return path


def test_precheck_refuses_many_sheets_sharing_one_large_drawing_quickly(tmp_path):
    """図形の多い描画を多数のシートが共有するブックは、openpyxl が開く前に断る（開くたびに数十秒かかるため）。"""
    path = _xlsx_sharing_one_drawing(tmp_path / "shared.xlsx", sheets=20, anchors=40_000)
    assert path.stat().st_size < 200_000
    started = time.monotonic()
    with pytest.raises(UploadError, match="図形や画像の数が多すぎる"):
        precheck_excel(path)
    assert time.monotonic() - started < 2


def test_precheck_multiplies_drawing_anchors_by_referencing_sheets(tmp_path, monkeypatch):
    monkeypatch.setattr(core, "MAX_DRAWING_ANCHORS", 100)
    precheck_excel(_xlsx_sharing_one_drawing(tmp_path / "ok.xlsx", sheets=4, anchors=25))       # 100 は通す
    with pytest.raises(UploadError, match="図形や画像"):
        precheck_excel(_xlsx_sharing_one_drawing(tmp_path / "ng.xlsx", sheets=5, anchors=25))   # 125
    with pytest.raises(UploadError, match="図形や画像"):                                          # 絶対パスの Target も数える
        precheck_excel(_xlsx_sharing_one_drawing(tmp_path / "abs.xlsx", sheets=5, anchors=25, absolute_target=True))


def test_precheck_limits_drawing_bytes_times_referencing_sheets(tmp_path, monkeypatch):
    monkeypatch.setattr(core, "MAX_DRAWING_BYTES", 10_000)
    with pytest.raises(UploadError, match="図形や画像"):
        precheck_excel(_xlsx_sharing_one_drawing(tmp_path / "big.xlsx", sheets=10, anchors=20))


def test_precheck_accepts_a_workbook_with_a_few_images(tmp_path):
    from openpyxl.drawing.image import Image as XLImage
    from PIL import Image as PILImage

    png = tmp_path / "p.png"
    PILImage.new("RGB", (4, 4), "red").save(png)
    wb = Workbook()
    for i in range(3):
        ws = wb.active if i == 0 else wb.create_sheet(f"S{i}")
        ws["A1"] = "管理No"
        ws.add_image(XLImage(str(png)), "C3")
        ws.add_image(XLImage(str(png)), "E5")
    path = tmp_path / "images.xlsx"
    wb.save(path)
    precheck_excel(path)


_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_WORKSHEET_REL = f"{_R_NS}/worksheet"
_COMMENTS_REL = f"{_R_NS}/comments"
_PKG_RELS = "http://schemas.openxmlformats.org/package/2006/relationships"


def _rels(entries: str) -> str:
    return f'<?xml version="1.0"?><Relationships xmlns="{_PKG_RELS}">{entries}</Relationships>'


def _xlsx_sharing_one_sheet_part(path, sheets: int, merge_ref: str):
    """複数の <sheet> が同じ1つのシートの部品を指すブック（Excel は作らないが、手で書けば openpyxl は開く）。

    openpyxl は <sheet> の数だけその部品を読み直すので、結合セルもセルもその回数だけ作られる。
    """
    sheet_xml = (f'<?xml version="1.0"?><worksheet xmlns="{_MAIN_NS}"><sheetData/>'
                 f'<mergeCells count="1"><mergeCell ref="{merge_ref}"/></mergeCells></worksheet>')
    tabs = "".join(f'<sheet name="S{i}" sheetId="{i}" r:id="rId{i}"/>' for i in range(1, sheets + 1))
    links = "".join(f'<Relationship Id="rId{i}" Type="{_WORKSHEET_REL}" Target="worksheets/sheet1.xml"/>'
                    for i in range(1, sheets + 1))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("_rels/.rels", _rels(""))
        zf.writestr("xl/workbook.xml", f'<?xml version="1.0"?><workbook xmlns="{_MAIN_NS}" xmlns:r="{_R_NS}">'
                                       f"<sheets>{tabs}</sheets></workbook>")
        zf.writestr("xl/_rels/workbook.xml.rels", _rels(links))
        zf.writestr("xl/worksheets/sheet1.xml", sheet_xml)
    return path


def test_precheck_multiplies_one_sheet_part_by_the_sheets_that_share_it(tmp_path):
    """同じシートの部品を多数の <sheet> が指すブックは、部品1つ分だけ数えると上限をすり抜ける（1シートあたり約4秒）。"""
    one = _xlsx_sharing_one_sheet_part(tmp_path / "one.xlsx", sheets=1, merge_ref="A1:XFD1")   # 16,384 セル
    precheck_excel(one, max_merged=core.FORM_MAX_MERGED_CELLS)          # 1シート分は通す
    many = _xlsx_sharing_one_sheet_part(tmp_path / "many.xlsx", sheets=20, merge_ref="A1:XFD1")
    assert many.stat().st_size < 10_000
    started = time.monotonic()
    with pytest.raises(UploadError, match="結合セルの範囲が大きすぎます"):
        precheck_excel(many, max_merged=core.FORM_MAX_MERGED_CELLS)     # 20シート分＝約33万セル
    assert time.monotonic() - started < 2


def test_precheck_multiplies_one_comments_part_by_the_sheets_that_share_it(tmp_path):
    """1つのコメントの部品を多数のシートが参照すると、コメントの範囲のセルもシートの数だけ作られる。"""
    comments = (f'<?xml version="1.0"?><comments xmlns="{_MAIN_NS}"><commentList>'
                f'<comment ref="A1:D2500" authorId="0"/></commentList></comments>')   # 10,000 セル
    sheets = 6
    tabs = "".join(f'<sheet name="S{i}" sheetId="{i}" r:id="rId{i}"/>' for i in range(1, sheets + 1))
    links = "".join(f'<Relationship Id="rId{i}" Type="{_WORKSHEET_REL}" Target="worksheets/sheet{i}.xml"/>'
                    for i in range(1, sheets + 1))
    path = tmp_path / "comments.xlsx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("_rels/.rels", _rels(""))
        zf.writestr("xl/workbook.xml", f'<?xml version="1.0"?><workbook xmlns="{_MAIN_NS}" xmlns:r="{_R_NS}">'
                                       f"<sheets>{tabs}</sheets></workbook>")
        zf.writestr("xl/_rels/workbook.xml.rels", _rels(links))
        for i in range(1, sheets + 1):
            zf.writestr(f"xl/worksheets/sheet{i}.xml", f'<?xml version="1.0"?><worksheet xmlns="{_MAIN_NS}">'
                                                       "<sheetData/></worksheet>")
            zf.writestr(f"xl/worksheets/_rels/sheet{i}.xml.rels",
                        _rels(f'<Relationship Id="rIdC" Type="{_COMMENTS_REL}" Target="../comments1.xml"/>'))
        zf.writestr("xl/comments1.xml", comments)
    with pytest.raises(UploadError, match="コメントの範囲が大きすぎます"):
        precheck_excel(path)   # 6シート分＝60,000 セル（MAX_LINKED_CELLS 50,000 超え）


def test_precheck_refuses_a_small_file_that_expands_to_hundreds_of_megabytes(tmp_path, monkeypatch):
    """展開後が大きくなりすぎるブックは断る（部品ごとの圧縮率の確認は、小さい部品を見ないためすり抜ける）。"""
    monkeypatch.setattr(core, "ZIP_RATIO_MIN_BYTES", 200_000)   # 本物は 16MB。同じ形を小さく作る
    path = tmp_path / "many_small_parts.xlsx"
    part = b"<si/>" * 20_000                                     # 100,000 バイト（1部品ずつは上の値未満）
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("xl/workbook.xml", f'<?xml version="1.0"?><workbook xmlns="{_MAIN_NS}"><sheets/></workbook>')
        for i in range(30):
            zf.writestr(f"xl/pad{i}.xml", b"<sst>" + part + b"</sst>")
    assert path.stat().st_size < 20_000                          # 20KB が 3MB に展開される（150倍）
    with pytest.raises(UploadError, match="圧縮率"):
        precheck_excel(path)


# ====================================================================================================
# 元 tests/test_core_jobs.py
# core/jobs.py: 進捗・中止・一時停止・中断からの回復。
# ====================================================================================================

@pytest.fixture
def core_app_core_jobs(tmp_path):
    app = Flask(__name__)
    app.config.update(TESTING=True, DATABASE=tmp_path / "app.db", MARK="この取り込み")
    db.init_app(app)
    return app


def _wait_until(predicate, timeout=10.0):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        value = predicate()
        if value:
            return value
        time.sleep(0.02)
    raise AssertionError("時間内に条件を満たしませんでした")


def test_progress_and_result(core_app_core_jobs):
    def work(ctx):
        conn = db.connect()  # バックグラウンドスレッドでも DB と app 設定が使える
        try:
            count = conn.execute("SELECT COUNT(*) FROM jobs").fetchone()[0]
        finally:
            conn.close()
        for i in range(1, 4):
            ctx.progress(done=i, total=3, phase="読み込み")
        return {"rows": 3, "mark": current_app.config["MARK"], "jobs": count, "params": ctx.params}

    with core_app_core_jobs.app_context():
        job_id = core.start_job("table_read", "table_import", 7, work, params={"scope": "all"})
        job = core.wait_job(job_id, timeout=10)
        assert job["status"] == "done" and job["finished"] and job["status_label"] == "完了"
        assert job["progress"]["done"] == 3 and job["progress"]["phase"] == "読み込み"
        assert job["result"] == {"rows": 3, "mark": "この取り込み", "jobs": 1, "params": {"scope": "all"}}
        assert job["kind"] == "table_read" and job["ref_type"] == "table_import" and job["ref_id"] == 7
        assert job["heartbeat_at"]
        assert core.latest_job("table_import", 7)["id"] == job_id
        assert core.latest_job("table_import", 7, kind="other") is None


def test_failure_is_recorded(core_app_core_jobs):
    """利用者に見せるエラー（JobError）のメッセージは、そのまま画面に出る。"""
    def work(ctx):
        ctx.progress(done=1)
        raise core.JobError("列が見つかりません")

    with core_app_core_jobs.app_context():
        job = core.wait_job(core.start_job("table_read", "table_import", 1, work))
        assert job["status"] == "failed"
        assert job["message"] == "列が見つかりません" and job["progress"]["done"] == 1


def test_unexpected_failure_does_not_show_the_python_exception_name(core_app_core_jobs):
    """想定外の例外は Python の例外名（KeyError など）を画面に出さず、日本語の案内にする。"""
    def work(ctx):
        raise KeyError("equipment_no")

    with core_app_core_jobs.app_context():
        job = core.wait_job(core.start_job("table_read", "table_import", 1, work))
        assert job["status"] == "failed"
        assert "KeyError" not in job["message"] and "equipment_no" not in job["message"]
        assert "処理中にエラーが発生しました" in job["message"] and "ログ" in job["message"]


def test_ai_and_table_errors_are_shown_to_the_user(core_app_core_jobs):
    """AI整形・一覧表の処理エラーは日本語のメッセージを持つので、そのまま出す（JobError を継承している）。"""
    from aiproc import AIJobError
    from tables import PipelineError

    assert issubclass(AIJobError, core.JobError) and issubclass(PipelineError, core.JobError)
    assert core.error_message(AIJobError("AI接続が設定されていません")) == "AI接続が設定されていません"
    assert core.error_message(PipelineError("見出しの行が見つかりません")) == "見出しの行が見つかりません"


def test_cancel_running_job(core_app_core_jobs):
    started = threading.Event()

    def work(ctx):
        i = 0
        while True:
            i += 1
            ctx.progress(done=i)
            started.set()
            ctx.check_cancel()
            time.sleep(0.01)

    with core_app_core_jobs.app_context():
        job_id = core.start_job("ai", "table_import", 1, work)
        assert started.wait(10)
        assert core.request_cancel(job_id)
        job = core.wait_job(job_id)
        assert job["status"] == "cancelled" and job["message"] == "中止しました"
        assert job["cancel_requested"] == 1
        assert core.request_cancel(job_id) is False  # 終わったジョブには効かない


def test_pause_and_resume(core_app_core_jobs):
    finish = threading.Event()
    observed = {}

    def work(ctx):
        i = 0
        while not finish.is_set():
            i += 1
            ctx.progress(done=i)
            if ctx.should_stop():
                observed["should_stop"] = True
            if not ctx.wait_if_paused():
                return {"stopped": True}
            time.sleep(0.01)
        return {"done": i}

    with core_app_core_jobs.app_context():
        job_id = core.start_job("ai", "table_import", 1, work)
        _wait_until(lambda: (core.get_job(job_id) or {}).get("status") == "running")
        assert core.request_pause(job_id)
        paused = _wait_until(lambda: (j := core.get_job(job_id))["status"] == "paused" and j)
        assert paused["status_label"] == "一時停止中"
        time.sleep(0.5)
        assert core.get_job(job_id)["progress"]["done"] == paused["progress"]["done"]  # 止まっている
        assert observed.get("should_stop")

        assert core.request_resume(job_id)
        _wait_until(lambda: core.get_job(job_id)["progress"]["done"] > paused["progress"]["done"])
        assert core.get_job(job_id)["status"] == "running"
        finish.set()
        job = core.wait_job(job_id)
        assert job["status"] == "done" and job["result"]["done"] > paused["progress"]["done"]
        assert job["pause_requested"] == 0


def test_cancel_while_paused_and_queued(core_app_core_jobs):
    release = threading.Event()
    ran = []

    def blocker(ctx):
        while not release.is_set():
            if not ctx.wait_if_paused():
                return None
            time.sleep(0.01)
        return None

    def never(ctx):
        ran.append(ctx.job_id)

    with core_app_core_jobs.app_context():
        first = core.start_job("ai", "table_import", 1, blocker)
        # 同じ取り込みのジョブは前のジョブが終わるまで動かない（ワーカーが何本あっても）
        second = core.start_job("ai", "table_import", 1, never)
        _wait_until(lambda: core.get_job(first)["status"] == "running")
        assert core.get_job(second)["status"] == "queued"
        assert core.request_cancel(second)
        assert core.get_job(second)["status"] == "cancelled"  # 待機中はその場で中止

        core.request_pause(first)
        _wait_until(lambda: core.get_job(first)["status"] == "paused")
        core.request_cancel(first)
        assert core.wait_job(first)["status"] == "cancelled"  # fn が正常終了しても中止扱い
        release.set()
        # 後ろのジョブが実行されないことを、次のジョブの完了で確かめる
        third = core.start_job("ai", "table_import", 3, lambda ctx: {"ok": True})
        assert core.wait_job(third)["status"] == "done"
        assert ran == []


def test_recover_interrupted(core_app_core_jobs):
    old = (datetime.now() - timedelta(minutes=5)).isoformat(timespec="seconds")
    fresh = datetime.now().isoformat(timespec="seconds")
    rows = [
        ("running", old, old),     # 止まったまま → 中断
        ("paused", old, old),      # 一時停止中のままアプリが終わった → 中断
        ("queued", None, old),     # 待機中のまま → 中断
        ("running", fresh, old),   # heartbeat が新しい → そのまま
        ("done", old, old),        # 終わったジョブ → そのまま
    ]
    with core_app_core_jobs.app_context():
        conn = db.connect()
        ids = []
        for status, beat, updated in rows:
            ids.append(conn.execute(
                "INSERT INTO jobs (kind, ref_type, ref_id, status, heartbeat_at, created_at, updated_at) "
                "VALUES ('table_confirm', 'table_import', 1, ?, ?, ?, ?)", (status, beat, updated, updated)).lastrowid)
        conn.commit()
        conn.close()

        assert core.recover_interrupted() == 3
        statuses = [core.get_job(i)["status"] for i in ids]
        assert statuses == ["interrupted", "interrupted", "interrupted", "running", "done"]
        job = core.get_job(ids[0])
        assert job["status_label"] == "中断" and "中断" in job["message"] and job["finished"]
        assert core.request_resume(ids[0]) is False
        assert core.recover_interrupted() == 0


def _insert_job(status, beat, kind="table_render", ref_id=1):
    conn = db.connect()
    job_id = conn.execute(
        "INSERT INTO jobs (kind, ref_type, ref_id, status, heartbeat_at, created_at, updated_at) "
        "VALUES (?, 'table_import', ?, ?, ?, ?, ?)", (kind, ref_id, status, beat, beat, beat)).lastrowid
    conn.commit()
    conn.close()
    return job_id


def test_a_job_left_by_a_previous_process_is_interrupted_once_its_heartbeat_is_stale(core_app_core_jobs):
    """閉じてすぐ起動し直すと起動時の回復では拾えない（2分未満）。参照したときに持ち主がいなければ中断にする。"""
    recent = (datetime.now() - timedelta(seconds=40)).isoformat(timespec="seconds")
    old = (datetime.now() - timedelta(minutes=3)).isoformat(timespec="seconds")
    with core_app_core_jobs.app_context():
        left = _insert_job("running", recent)
        assert core.recover_interrupted() == 0          # 起動時: まだ新しいので残す
        assert core.get_job(left)["status"] == "running"
        conn = db.connect()
        conn.execute("UPDATE jobs SET heartbeat_at = ? WHERE id = ?", (old, left))   # 2分以上たった
        conn.commit()
        conn.close()
        job = core.latest_job("table_import", 1)
        assert job["status"] == "interrupted" and job["finished"] and "中断" in job["message"]


def test_a_queued_job_of_this_process_is_not_interrupted_while_it_waits(core_app_core_jobs, monkeypatch):
    release = threading.Event()
    with core_app_core_jobs.app_context():
        first = core.start_job("table_read", "table_import", 1, lambda ctx: release.wait(10) and None)
        # 同じ取り込みなので、前のジョブが終わるまで待機中のまま
        second = core.start_job("table_render", "table_import", 1, lambda ctx: {"ok": True})
        _wait_until(lambda: core.get_job(first)["status"] == "running")
        monkeypatch.setattr(core, "STALE_AFTER", timedelta(seconds=-60))   # どれも「古い」とみなす
        assert core.get_job(second)["status"] == "queued"   # このプロセスのジョブなので中断にしない
        release.set()
        assert core.wait_job(second)["status"] == "done"


def test_a_long_queued_job_is_not_interrupted_by_a_second_launch(core_app_core_jobs, monkeypatch):
    """誤って2つ目を起動しても、その起動時の回復（別プロセス）が、1つ目で2分以上待っているジョブを中断にしない。"""
    monkeypatch.setattr(core, "HEARTBEAT_INTERVAL", 0.05)
    release = threading.Event()
    old = (datetime.now() - timedelta(minutes=5)).isoformat(timespec="seconds")
    with core_app_core_jobs.app_context():
        first = core.start_job("table_read", "table_import", 1, lambda ctx: release.wait(10) and None)
        second = core.start_job("table_render", "table_import", 1, lambda ctx: {"ok": True})
        _wait_until(lambda: core.get_job(first)["status"] == "running")
        conn = db.connect()
        conn.execute("UPDATE jobs SET heartbeat_at = NULL, created_at = ?, updated_at = ? WHERE id = ?",
                     (old, old, second))   # 5分前から待っている
        conn.commit()
        conn.close()
        # 動いているジョブの Ticker が、待機中のジョブの heartbeat も更新する
        _wait_until(lambda: (core.get_job(second)["heartbeat_at"] or "") > old)
        assert core.recover_interrupted() == 0   # 2つ目の起動時の回復（このプロセスのことは知らない）
        assert core.get_job(second)["status"] == "queued"
        release.set()
        assert core.wait_job(second)["status"] == "done"


def test_a_paused_ai_job_does_not_hold_up_other_imports(core_app_core_jobs):
    """AI整形を一時停止しても、ほかの取り込みの読み込みは進む。同じ取り込みの分は AI整形が終わるまで待つ。"""
    def ai(ctx):
        while True:
            if not ctx.wait_if_paused():
                return None
            time.sleep(0.01)

    with core_app_core_jobs.app_context():
        ai_job = core.start_job("ai_format", "table_import", 1, ai)
        _wait_until(lambda: core.get_job(ai_job)["status"] == "running")
        core.request_pause(ai_job)
        _wait_until(lambda: core.get_job(ai_job)["status"] == "paused")

        other = core.start_job("table_read", "table_import", 2, lambda ctx: {"rows": 3})
        assert core.wait_job(other, timeout=5)["status"] == "done"

        same = core.start_job("table_render", "table_import", 1, lambda ctx: {"files": 1})
        time.sleep(0.5)
        assert core.get_job(same)["status"] == "queued"    # 同じ取り込みの AI整形が終わるまで動かさない
        later = core.start_job("table_read", "table_import", 3, lambda ctx: {"rows": 1})
        assert core.wait_job(later, timeout=5)["status"] == "done"   # 後回しの分がほかを止めない

        core.request_cancel(ai_job)
        assert core.wait_job(ai_job)["status"] == "cancelled"
        assert core.wait_job(same, timeout=5)["status"] == "done"


def test_a_job_waiting_behind_a_paused_ai_job_says_what_it_waits_for(core_app_core_jobs):
    """一時停止中の AI整形の後ろで待つジョブは、理由（何を待っているか・どうすれば動くか）を出す。"""
    def ai(ctx):
        while True:
            if not ctx.wait_if_paused():
                return None
            time.sleep(0.01)

    with core_app_core_jobs.app_context():
        ai_job = core.start_job("ai_format", "table_import", 11, ai)
        _wait_until(lambda: core.get_job(ai_job)["status"] == "running")
        core.request_pause(ai_job)
        _wait_until(lambda: core.get_job(ai_job)["status"] == "paused")

        same = core.start_job("table_preview", "table_import", 11, lambda ctx: {"files": 1})
        # 別の取り込みの AI整形は待たされない（列にワーカーが複数あるため。design.md 同時利用）
        other_ai = core.start_job("ai_format", "table_import", 12, lambda ctx: {"rows": 1})
        assert core.wait_job(other_ai, timeout=10)["status"] == "done"
        waiting = core.get_job(same)
        assert waiting["status"] == "queued"
        assert "この取り込みのAI整形が一時停止中" in waiting["message"] and "再開するか中止" in waiting["message"]
        assert waiting["waiting_for"]["job_id"] == ai_job and waiting["waiting_for"]["same_ref"]
        assert "message" not in core.get_job(ai_job) or "待っています" not in core.get_job(ai_job)["message"]

        core.request_cancel(ai_job)
        assert core.wait_job(same, timeout=5)["status"] == "done"
        assert core.get_job(same).get("waiting_for") is None


def test_a_job_waits_when_every_worker_of_its_lane_is_busy(core_app_core_jobs, monkeypatch):
    """ワーカーが全部ふさがっている列で待つジョブにも、何を待っているかを出す（JOB_WORKERS の効き目）。"""
    monkeypatch.setitem(core.LANE_OF_KIND, "slow_test", "test-lane-1")
    core_app_core_jobs.config["JOB_WORKERS"] = 1     # この列は初めて使うので、この本数で立つ
    release = threading.Event()

    with core_app_core_jobs.app_context():
        first = core.start_job("slow_test", "table_import", 21, lambda ctx: release.wait(10) and None)
        _wait_until(lambda: core.get_job(first)["status"] == "running")
        second = core.start_job("slow_test", "table_import", 22, lambda ctx: {"ok": True})
        time.sleep(0.3)
        waiting = core.get_job(second)
        assert waiting["status"] == "queued" and waiting["waiting_for"]["job_id"] == first
        assert "待っています" in waiting["message"]
        release.set()
        assert core.wait_job(second, timeout=10)["status"] == "done"


def test_a_job_whose_last_status_write_fails_ends_as_failed(core_app_core_jobs, monkeypatch):
    import sqlite3

    real = core.JobContext.progress

    def locked(self, **kw):
        if "result" in kw:
            raise sqlite3.OperationalError("database is locked")
        return real(self, **kw)

    monkeypatch.setattr(core.JobContext, "progress", locked)
    with core_app_core_jobs.app_context():
        job_id = core.start_job("table_render", "table_import", 1, lambda ctx: {"files": 1})
        job = core.wait_job(job_id)
        assert job["status"] == "failed" and "エラー" in job["message"] and job["finished"]


# ====================================================================================================
# 元 tests/test_core_mdtext.py
# core/mdtext.py: Markdown テキスト処理。
# ====================================================================================================

def test_nfkc_value():
    assert nfkc_value("ＣＭＰ－１０１　 研磨") == "CMP-101 研磨"
    assert nfkc_value("フィルター  交換\r\n\r\n  流量　再校正  ") == "フィルター 交換\n\n流量 再校正"
    assert nfkc_value("原点 復帰") == "原点 復帰"   # 日本語文字間の半角空白は残す
    assert nfkc_value(None) == ""
    assert nfkc_value(95) == "95"
    # 丸数字は NFKC で囲みが外れると「①破損」が「1破損」になり番号と本文の区切りが消えるので、そのまま残す
    assert nfkc_value("①破損ウェーハ片を回収\n②Head3 メンブレン交換") == "①破損ウェーハ片を回収\n②Head3 メンブレン交換"
    assert nfkc_value("⑳Ⓐ㋐㊤") == "⑳Ⓐ㋐㊤"
    # 区切りが残る表記は今までどおり NFKC で正規化する
    assert nfkc_value("㈱テスト ⑴ ⒈ ﾎﾟﾝﾌﾟ") == "(株)テスト (1) 1. ポンプ"


def test_escape_md_line():
    assert escape_md_line("# 見出し") == "\\# 見出し"
    assert escape_md_line("#123 は番号") == "#123 は番号"
    assert escape_md_line("- 項目") == "\\- 項目"
    assert escape_md_line("-5℃") == "-5℃"
    assert escape_md_line("* 注") == "\\* 注"
    assert escape_md_line("+ 追加") == "\\+ 追加"
    assert escape_md_line("> 引用") == "\\> 引用"
    assert escape_md_line("1. 手順") == "1\\. 手順"
    assert escape_md_line("2) 手順") == "2\\) 手順"
    assert escape_md_line("2026.08.03 対応") == "2026.08.03 対応"
    assert escape_md_line("---") == "\\---"
    assert escape_md_line("= = =") == "\\= = ="
    assert escape_md_line("```python") == "\\`\\`\\`python"
    assert escape_md_line("  # 字下げ") == "  \\# 字下げ"
    assert escape_md_line("普通の文") == "普通の文"
    assert escape_md_line("") == ""


def test_md_bullet():
    assert md_bullet("停止時間", "95分") == ["- 停止時間: 95分"]
    assert md_bullet("停止時間", 95) == ["- 停止時間: 95"]
    assert md_bullet("処置", "フィルター交換\n\n流量再校正") == ["- 処置:", "  フィルター交換", "  流量再校正"]
    assert md_bullet("処置", "# 手順\n1. 交換") == ["- 処置:", "  \\# 手順", "  1\\. 交換"]
    assert md_bullet("時系列", ["1. 2026-08-03 14:20［連絡・初動］田中", "2. 復旧"]) == \
        ["- 時系列:", "  1\\. 2026-08-03 14:20［連絡・初動］田中", "  2\\. 復旧"]
    assert md_bullet("原因", "") == []
    assert md_bullet("原因", None) == []
    assert md_bullet("原因", "\n  \n") == []


def test_estimate_tokens():
    # 実トークン（tiktoken o200k_base）以上になる見積もり: 非ASCII 1文字=1.1、ASCII 2文字=1
    assert estimate_tokens("") == 0
    assert estimate_tokens("故障") == 3
    assert estimate_tokens("abc") == 2
    assert estimate_tokens("abcd") == 2
    assert estimate_tokens("CMP研磨") == 4
    assert estimate_tokens("あ" * 100) == 110
    # ASCII をまとめて数えても、1文字ずつ数えたときと同じ（U+007F/U+0080 の境目・絵文字・孤立サロゲート）
    assert estimate_tokens("\x7f\x80") == 2
    assert estimate_tokens("a\U0001F600b") == 3
    assert estimate_tokens("\ud800x") == 2


def test_estimate_tokens_is_not_below_the_real_count_for_dates_and_part_numbers():
    """日時・品番・計測値は o200k_base で細かく区切られる。右の数は tiktoken（gpt-4o-mini）で数えた実トークン数。"""
    real = {
        "2023-09-01 09:44": 10,
        "- 発生日時: 2023-09-01 09:44": 16,
        "8/3 10:07 田中：FDC ch3 V=1039V I=2.34A P=1.11kW 0x0B73 R=11.1ohm OK": 47,
        "PN-A12345-B67 SN:0x1F3A9C": 16,
        "1,234,567.89": 7,
        "2026/08/03 10:07〜2026/08/03 11:45": 21,
        "ロット番号 L2308-0412-07 の不良率 0.35%": 20,
    }
    for text, count in real.items():
        assert estimate_tokens(text) >= count, text
    # 数字のまとまりは、前が ASCII でない文字・記号・先頭でも数える（「1年2月」は2つ）
    assert estimate_tokens("1年2月") == estimate_tokens("x年y月") + 1
    # 40行の計測ログ。実 1,661 トークン（チャンク 1,500 を超える）を、旧式は 1,252 と見積もり記録の上限 1,400 を通していた
    log = "\n".join(f"8/3 {10 + i * 7 // 60}:{i * 7 % 60:02d} 田中：FDC ch{i % 8} V={1000 + i * 13}V "
                    f"I={2.31 + i / 100:.2f}A 0x{i * 977:04X} dP={i * 0.013:.3f}kPa OK" for i in range(40))
    assert estimate_tokens(log) >= 1661


def test_join_blocks():
    text = join_blocks([["# タイトル"], [], ["- a: 1", "", "- b: 2\r\n  続き"], ["- 出典: x.xlsx"]])
    assert text == "# タイトル\n\n- a: 1\n- b: 2\n  続き\n\n- 出典: x.xlsx\n"
    assert join_blocks([]) == ""
    assert join_blocks([[""], []]) == ""


# ====================================================================================================
# 元 tests/test_core_naming.py
# core/naming.py: 出力ファイル名。
# ====================================================================================================

def test_safe_filename_part_replaces_unsafe_chars():
    assert safe_filename_part('a\\b/c:d*e?f"g<h>i|j') == "a_b_c_d_e_f_g_h_i_j"
    assert safe_filename_part("CMP研磨装置 1号機") == "CMP研磨装置_1号機"
    assert safe_filename_part("ＣＭＰ－１０１") == "CMP-101"          # NFKC
    assert safe_filename_part("行1\n行2\t\x01") == "行1_行2"
    assert safe_filename_part("[重要] 報告") == "重要_報告"
    assert safe_filename_part("報告.[legacy]") == "報告legacy"       # '.[' は除去
    assert safe_filename_part("..報告書__") == "報告書"
    assert safe_filename_part(None) == ""
    assert safe_filename_part("   ") == ""


def test_safe_filename_part_length_and_reserved():
    assert safe_filename_part("あ" * 100) == "あ" * 60
    assert safe_filename_part("あ" * 10 + "_" + "い" * 10, max_len=11) == "あ" * 10
    assert safe_filename_part("con") == "con_"
    assert safe_filename_part("COM1") == "COM1_"


def test_md_filename():
    assert md_filename(["設備修理報告書", "R2026-00123", "CMP-101"]) == "設備修理報告書_R2026-00123_CMP-101.md"
    assert md_filename(["トラブル対応一覧", "", None, "2026-08"]) == "トラブル対応一覧_2026-08.md"
    # LightRAG のファイル名ヒント（.[...]）は付けない。ヒントらしい文字列は名前から外す
    assert md_filename(["故障履歴.[legacy-R(chunk_ts=800)]", "2026-08"]) == "故障履歴legacy-R(chunk_ts=800)_2026-08.md"
    assert md_filename([]) == "無題.md"


# ====================================================================================================
# 元 tests/test_core_purge.py
# ダウンロードで消す処理（core/purge.py）の、本番のサーバー（waitress）での動きと番号の続き（design.md 3.3）。
# ====================================================================================================

BIG = 1024 * 1024   # 溜められる上限（OUTBUF_HIGH_WATERMARK）より十分大きい本文


def _serve(app):
    """本番と同じ waitress の設定で、このテストの間だけ空いているポートで動かす。"""
    server = create_server(app, host="127.0.0.1", port=0, **app_module.WAITRESS_OPTIONS)
    threading.Thread(target=server.run, daemon=True).start()
    return server


def _add_big_download(app, purged: list, finished: threading.Event):
    body = os.urandom(BIG)

    @app.get("/_test/big.zip")
    def _big():
        res = send_file(io.BytesIO(body), mimetype="application/zip", as_attachment=True, download_name="big.zip",
                        conditional=False)
        res = core.purge_after_send(res, purged.append, "消した")
        res.call_on_close(finished.set)   # 応答を閉じた（送り終えた・切れた）ことをテストに知らせる
        return res


def _request(port: int) -> socket.socket:
    sock = socket.create_connection(("127.0.0.1", port))
    sock.sendall(f"GET /_test/big.zip HTTP/1.1\r\nHost: 127.0.0.1:{port}\r\nConnection: close\r\n\r\n".encode())
    return sock


def test_the_server_does_not_buffer_a_whole_download_ahead():
    """waitress の既定（16MB）だと、16MB 未満の zip は受け取られる前にアプリからは送り終えたように見える。"""
    assert app_module.WAITRESS_OPTIONS["outbuf_high_watermark"] <= 64 * 1024


def test_a_large_download_cut_off_under_waitress_keeps_the_data(app):
    purged, finished = [], threading.Event()
    _add_big_download(app, purged, finished)
    server = _serve(app)
    try:
        sock = _request(server.effective_port)
        assert sock.recv(1)   # 本文の最初だけ受け取って、ブラウザを閉じたように切る（RST）
        sock.setsockopt(socket.SOL_SOCKET, socket.SO_LINGER, struct.pack("ii", 1, 0))
        sock.close()
        assert finished.wait(10)
        time.sleep(0.2)
        assert purged == []
    finally:
        server.close()


def test_a_large_download_received_to_the_end_under_waitress_removes_the_item(app):
    purged, finished = [], threading.Event()
    _add_big_download(app, purged, finished)
    server = _serve(app)
    try:
        sock = _request(server.effective_port)
        received = b""
        while chunk := sock.recv(65536):
            received += chunk
        sock.close()
        assert len(received.split(b"\r\n\r\n", 1)[1]) == BIG
        assert finished.wait(10)
        time.sleep(0.2)
        assert purged == ["消した"]
    finally:
        server.close()


# ---- 番号の続き（sqlite_sequence）を残さない ---------------------------------------------

def _sequence(app, table: str):
    with app.app_context():
        row = db.get_db().execute("SELECT seq FROM sqlite_sequence WHERE name = ?", (table,)).fetchone()
        return None if row is None else row[0]


def test_the_id_counter_does_not_record_how_many_forms_were_taken_in(app, client):
    ids = [add_confirmed_document(app, f"報告書{i}.xlsx")[0] for i in range(3)]
    assert _sequence(app, "documents") == max(ids)
    for doc_id in ids[:-1]:
        assert client.get(f"/forms/{doc_id}/download.md").status_code == 200
    assert _sequence(app, "documents") == max(ids)   # まだ作業中の帳票があるうちはそのまま

    assert client.get(f"/forms/{ids[-1]}/download.md").status_code == 200
    seq = _sequence(app, "documents")
    assert core._ID_BASE_MIN <= seq < core._ID_BASE_MAX   # 空になったら乱数に置き換わる

    # 次の帳票は消した番号を使い回さないので、開いたままの古い画面は新しい帳票を指さない
    new_id, _ = add_confirmed_document(app, "次の報告書.xlsx")
    assert new_id == seq + 1 and new_id not in ids
    for doc_id in ids:
        assert client.get(f"/forms/{doc_id}/type").status_code == 404
        assert client.get(f"/forms/{doc_id}/download.md").status_code == 404
    assert new_id < 2 ** 53   # 画面の JavaScript でも正確に扱える


def test_the_id_counter_of_table_imports_and_jobs_is_forgotten_too(app, client):
    import_id = confirmed_import(app, client)
    assert client.get(f"/tables/imports/{import_id}/download.zip").status_code == 200
    for table in ("table_imports", "jobs"):
        seq = _sequence(app, table)
        assert core._ID_BASE_MIN <= seq < core._ID_BASE_MAX, table
    # 開いたままの画面が段を取りに来ても、もう無い（ダウンロード済み）
    assert client.get(f"/tables/imports/{import_id}/panel/preview").status_code == 404


def test_the_counter_is_randomised_each_time_the_table_becomes_empty(app, client):
    seen = set()
    for i in range(3):
        doc_id, _ = add_confirmed_document(app, f"報告書{i}.xlsx")
        assert doc_id not in seen
        seen.add(doc_id)
        assert client.get(f"/forms/{doc_id}/download.md").status_code == 200
        seen.add(_sequence(app, "documents"))
    assert len(seen) == 6


def test_a_new_database_still_starts_at_one(app):
    with app.app_context():
        core.forget_id_counters(db.get_db())   # 一度も帳票を入れていない表には番号の続きが無く、そのまま
    doc_id, _ = add_confirmed_document(app, "最初.xlsx")
    assert doc_id == 1


def test_startup_forgets_a_counter_left_by_an_older_version(tmp_path):
    from tests.conftest import make_config

    first = app_module.create_app(make_config(tmp_path))
    with first.app_context():
        conn = db.get_db()
        conn.execute("INSERT INTO documents (file_name, file_hash, stored_path, created_at) "
                     "VALUES ('a.xlsx', '0', 'documents/a.xlsx', '2026-09-19')")
        conn.execute("DELETE FROM documents")   # 前の版は消しても番号の続き（1）を残していた
        conn.commit()
        assert conn.execute("SELECT seq FROM sqlite_sequence WHERE name = 'documents'").fetchone()[0] == 1

    again = app_module.create_app(make_config(tmp_path))
    assert core._ID_BASE_MIN <= _sequence(again, "documents") < core._ID_BASE_MAX


def test_the_counter_is_kept_while_the_table_still_has_rows(app):
    first, _ = add_confirmed_document(app, "作業中.xlsx")
    with app.app_context():
        core.forget_id_counters(db.get_db())
    assert _sequence(app, "documents") == first


# ====================================================================================================
# 元 tests/test_core_fixes6.py
# 6巡目の修正（core）。
#
# - R6-FUZZ-1: シート全体を指すハイパーリンク・コメントの範囲は、openpyxl で開く前に断る
# - R6-SEC-2: 書式だけの空の行（<row ht customHeight/>）が大量にあるブックは、帳票では開く前に断る
# - R6-2: imports/<id>/ を消し切れなかったときは中身を 0 バイトにする
# ====================================================================================================

def _rewrite(src: Path, dest: Path, edit) -> Path:
    """xlsx の部品を edit(名前, 中身) で書き換えて dest に保存する。"""
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            zout.writestr(info, edit(info.filename, zin.read(info.filename)))
    return dest


def _book(path: Path, *, comment=False) -> Path:
    wb = Workbook()
    wb.active["A1"] = "管理No"
    wb.active["B1"] = "設備"
    if comment:
        wb.active["A1"].comment = Comment("メモ", "作成者")
    wb.save(path)
    return path


def _with_hyperlink(tmp_path, ref: str) -> Path:
    def edit(name, data):
        if name == "xl/worksheets/sheet1.xml":
            link = f'<hyperlinks><hyperlink ref="{ref}" location="A1"/></hyperlinks>'.encode()
            data = data.replace(b"<pageMargins", link + b"<pageMargins", 1)
            assert link in data
        return data
    return _rewrite(_book(tmp_path / "base.xlsx"), tmp_path / f"link_{ref.replace(':', '_')}.xlsx", edit)


def _with_comment(tmp_path, ref: str) -> Path:
    def edit(name, data):
        if re.fullmatch(r"xl/comments/?\w*\.xml", name):
            data, n = re.subn(rb'ref="A1"', f'ref="{ref}"'.encode(), data)
            assert n == 1
        return data
    return _rewrite(_book(tmp_path / "base_c.xlsx", comment=True), tmp_path / f"comment_{ref.replace(':', '_')}.xlsx",
                    edit)


def _with_empty_rows(tmp_path, count: int) -> Path:
    def edit(name, data):
        if name == "xl/worksheets/sheet1.xml":
            extra = "".join(f'<row r="{n}" ht="20" customHeight="1"/>' for n in range(2, count + 2)).encode()
            data = data.replace(b"</row></sheetData>", b"</row>" + extra + b"</sheetData>", 1)
            assert extra[:20] in data
        return data
    return _rewrite(_book(tmp_path / "base_r.xlsx"), tmp_path / f"rows_{count}.xlsx", edit)


# ---- R6-FUZZ-1 -------------------------------------------------------------------------

@pytest.mark.parametrize("make", [_with_hyperlink, _with_comment])
def test_a_link_or_comment_over_the_whole_sheet_is_refused_before_opening(tmp_path, make):
    with pytest.raises(UploadError, match="ハイパーリンクまたはコメントの範囲が大きすぎます"):
        precheck_excel(make(tmp_path, "A1:XFD1048576"))
    with pytest.raises(UploadError, match="ハイパーリンクまたはコメントの範囲が大きすぎます"):
        precheck_excel(make(tmp_path, "A1:XFD50"), max_merged=FORM_MAX_MERGED_CELLS)
    precheck_excel(make(tmp_path, "A1"))            # ふつうのリンク・コメント（1セル）は通す
    precheck_excel(make(tmp_path, "A1:C20"))        # 数セルの範囲も通す


@pytest.mark.parametrize("url", ["/forms/upload", "/tables/upload"])
@pytest.mark.parametrize("make", [_with_hyperlink, _with_comment])
def test_both_upload_routes_refuse_a_whole_sheet_link_or_comment(app, client, tmp_path, url, make):
    data = make(tmp_path, "A1:XFD1048576").read_bytes()
    assert "ハイパーリンクまたはコメントの範囲が大きすぎます" in upload_error(client, url, data, "リンク.xlsx")
    assert [p for p in Path(app.config["UPLOAD_DIR"]).rglob("*") if p.is_file()] == []


# ---- R6-SEC-2 --------------------------------------------------------------------------

def test_many_empty_formatted_rows_are_refused_for_forms(tmp_path):
    path = _with_empty_rows(tmp_path, FORM_MAX_MERGED_CELLS + 1)
    assert path.stat().st_size < 2 * 1024 * 1024
    with pytest.raises(UploadError, match="行数が上限"):
        precheck_excel(path, max_merged=FORM_MAX_MERGED_CELLS)
    precheck_excel(path)   # 一覧表（読み取り専用で開く。行の書式を作らない）の上限では通る


def test_row_limit_counts_every_row_exactly(tmp_path):
    path = _with_empty_rows(tmp_path, 99)   # 1行目（値の入った行）＋99行 = 100行
    precheck_excel(path, max_rows=100)
    with pytest.raises(UploadError, match="行数が上限（99 行）"):
        precheck_excel(path, max_rows=99)


def test_forms_upload_refuses_many_empty_formatted_rows(app, client, tmp_path):
    data = _with_empty_rows(tmp_path, FORM_MAX_MERGED_CELLS + 1).read_bytes()
    assert "行数が上限" in upload_error(client, "/forms/upload", data, "行だらけ.xlsx")
    assert [p for p in Path(app.config["UPLOAD_DIR"]).rglob("*") if p.is_file()] == []


# ---- R6-2 ------------------------------------------------------------------------------

def _leave_files(monkeypatch):
    """rmtree が掴まれたファイルを消せなかったときと同じく、フォルダとファイルを残す。"""
    monkeypatch.setattr(core.shutil, "rmtree", lambda *a, **k: None)


def test_leftover_import_files_are_emptied_and_flagged(app, client, monkeypatch):
    import_id = confirmed_import(app, client)
    _leave_files(monkeypatch)
    with app.test_request_context("/"):
        folder = core.import_dir(import_id)
        assert any(p.stat().st_size for p in folder.rglob("*") if p.is_file())
        assert not core.purge_incomplete()
        core.purge_table_import(import_id)
        leftovers = [p for p in folder.rglob("*") if p.is_file()]
        assert leftovers and all(p.stat().st_size == 0 for p in leftovers)
        assert core.purge_incomplete()


def test_a_clean_purge_is_not_flagged(app, client):
    import_id = confirmed_import(app, client)
    with app.test_request_context("/"):
        core.purge_table_import(import_id)
        assert not core.import_dir(import_id).exists()
        assert not core.purge_incomplete()


# ====================================================================================================
# 元 tests/test_retention.py
# データを残さないこと（design.md 3.3）の確認。
#
# ダウンロードが終わったら、その取り込みに属するものは何も残らない:
# アップロードしたファイル・imports/<id>/ のファイル・DB の行（どの表でも）。
# 残るのは帳票の種類と一覧表の取り込み設定（＝設定であってデータではない）。
# ====================================================================================================

EXTRACTION = {
    "pattern": {"id": 1, "name": "設備修理報告書", "version": "v1"},
    "values": {"equipment_id": "EQ-001"},
    "fields": [{"field_name": "equipment_id", "display_name": "設備番号", "data_type": "string", "value": "EQ-001",
                "sheet": "修理報告書", "label_cell": "A4", "value_cell": "B4", "edited": False}],
    "attachments": [], "sheets": ["修理報告書"],
}


# ---- 共通 -------------------------------------------------------------------------

def _extraction(value="EQ-001") -> str:
    data = json.loads(json.dumps(EXTRACTION))
    data["values"]["equipment_id"] = value
    data["fields"][0]["value"] = value
    return json.dumps(data, ensure_ascii=False)


def _add_confirmed_document(app, name, *, value="EQ-001", batch_id="", order=0) -> tuple[int, Path]:
    """確定済みの帳票を1件作る（アップロードしたファイルの実体も置く）。"""
    with app.app_context():
        stored = f"documents/{name}"
        path = Path(app.config["UPLOAD_DIR"]) / stored
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(b"dummy-excel")
        doc_id = db.create_document(name, "0" * 64, stored, batch_id=batch_id, batch_order=order)
        db.update_document(doc_id, data_json=_extraction(value), confirmed_json=_extraction(value),
                           title=f"{value} 確定")
    return doc_id, path


def _rows_for(app, table_of_id: str, ref_columns: tuple[str, ...], value) -> dict[str, int]:
    """その取り込みを指す行が残っている表を {表.列: 件数} で返す。

    表の名前を決め打ちせず sqlite_master を見るので、あとから増えた表も見つかる。
    """
    hits: dict[str, int] = {}
    with app.app_context():
        conn = db.get_db()
        for (table,) in conn.execute("SELECT name FROM sqlite_master WHERE type = 'table'").fetchall():
            if table.startswith("sqlite_"):
                continue
            columns = {row[1] for row in conn.execute(f'PRAGMA table_info("{table}")')}
            checks = [(c, value) for c in ref_columns if c in columns]
            if table == table_of_id and "id" in columns:
                checks.append(("id", value))
            for column, val in checks:
                count = conn.execute(f'SELECT COUNT(*) FROM "{table}" WHERE "{column}" = ?', (val,)).fetchone()[0]
                if count:
                    hits[f"{table}.{column}"] = count
    return hits


def _uploaded_files(app) -> list[Path]:
    base = Path(app.config["UPLOAD_DIR"])
    return [p for p in base.rglob("*") if p.is_file()]


# ---- 帳票（1件ずつ） ---------------------------------------------------------------------

def test_downloading_a_form_removes_its_file_and_every_row(app, client):
    doc_id, path = _add_confirmed_document(app, "報告書.xlsx")
    assert _rows_for(app, "documents", ("document_id",), doc_id)

    res = client.get(f"/forms/{doc_id}/download.md")
    assert res.status_code == 200 and "EQ-001" in res.get_data(as_text=True)

    assert not path.exists()
    assert _uploaded_files(app) == []
    assert _rows_for(app, "documents", ("document_id",), doc_id) == {}
    assert client.get(f"/forms/{doc_id}/type").status_code == 404
    assert client.get(f"/forms/{doc_id}/download.md").status_code == 404


def test_a_form_that_cannot_be_built_is_not_deleted(app, client):
    """確定していない帳票はダウンロードできず、そのときデータも消さない。"""
    with app.app_context():
        doc_id = db.create_document("未確定.xlsx", "0" * 64, "documents/未確定.xlsx")
        db.save_draft(doc_id, _extraction())
    assert client.get(f"/forms/{doc_id}/download.md").status_code == 404
    with app.app_context():
        assert db.get_document(doc_id) is not None


def test_deleting_a_form_by_hand_removes_the_file_too(app, client):
    doc_id, path = _add_confirmed_document(app, "消す.xlsx")
    res = client.post(f"/forms/{doc_id}/delete")
    assert res.status_code == 200 and res.get_json()["ok"] is True
    assert not path.exists() and _rows_for(app, "documents", ("document_id",), doc_id) == {}


def test_the_form_type_is_kept_when_the_form_is_deleted(app, client):
    """帳票の種類は設定なので、データを消しても残る。"""
    with app.app_context():
        pattern_id = db.create_pattern("設備修理報告書")
    doc_id, _path = _add_confirmed_document(app, "種類つき.xlsx")
    with app.app_context():
        db.update_document(doc_id, pattern_id=pattern_id)
    assert client.get(f"/forms/{doc_id}/download.md").status_code == 200
    with app.app_context():
        assert db.load_pattern(pattern_id) is not None


# ---- 帳票（まとめ取り込み） ----------------------------------------------------------------

def _upload_many(client, sample_dir, names):
    files = [(io.BytesIO((sample_dir / n).read_bytes()), n) for n in names]
    return client.post("/forms/upload", data={"file": files}, content_type="multipart/form-data")


def _finish(client, ids, current=None) -> dict:
    url = "/forms/finish?ids=" + ",".join(str(i) for i in ids) + (f"&current={current}" if current else "")
    return client.get(url).get_json()


def test_uploading_several_forms_makes_one_batch(app, client, sample_dir):
    res = _upload_many(client, sample_dir, ["standard.xlsx", "shifted.xlsx", "table.xlsx"])
    body = res.get_json()
    assert res.status_code == 200 and len(body["docs"]) == 3 and body["batch_id"]
    with app.app_context():
        rows = [db.get_document(d["id"]) for d in body["docs"]]
        batch_ids = {r["batch_id"] for r in rows}
        assert len(batch_ids) == 1 and "" not in batch_ids
    # 読み取る前は件数も zip のボタンも出さない（読み取れた分だけが zip に入るため）
    ids = [d["id"] for d in body["docs"]]
    state = _finish(client, ids)
    assert state["total"] == 0 and state["read_yet"] is False
    assert "確定済み" not in state["html"] and "zip" not in state["html"]
    assert "帳票の種類とシートを選び" in state["html"]


def test_a_single_file_upload_still_has_no_batch(app, client, sample_dir):
    res = client.post("/forms/upload", data={"file": (io.BytesIO((sample_dir / "standard.xlsx").read_bytes()),
                                                      "standard.xlsx")}, content_type="multipart/form-data")
    body = res.get_json()
    assert res.status_code == 200 and len(body["docs"]) == 1 and body["batch_id"] == ""
    with app.app_context():
        assert db.get_document(body["docs"][0]["id"])["batch_id"] == ""
    assert "zip" not in _finish(client, [body["docs"][0]["id"]])["html"]


def test_downloading_the_batch_zip_removes_the_whole_batch(app, client):
    first, path1 = _add_confirmed_document(app, "1.xlsx", value="EQ-001", batch_id="B", order=0)
    second, path2 = _add_confirmed_document(app, "2.xlsx", value="EQ-002", batch_id="B", order=1)

    # すべて確定したら zip でまとめてダウンロードできる（押す前に消えることを知らせる）
    page = _finish(client, [first, second], current=second)["html"]
    assert "確定済み <strong>2</strong> / 2 件" in page and "まとめて Markdown をダウンロード（zip）" in page
    assert "/forms/batches/B/download.zip" in page
    assert "もう一度ダウンロードすることはできません" in page

    res = client.get("/forms/batches/B/download.zip")
    assert res.status_code == 200 and res.mimetype == "application/zip"
    with zipfile.ZipFile(io.BytesIO(res.data)) as zf:
        names = zf.namelist()
        texts = [zf.read(n).decode("utf-8") for n in names]
    assert len(names) == 2 and all(n.endswith(".md") for n in names)
    assert any("EQ-001" in t for t in texts) and any("EQ-002" in t for t in texts)

    assert not path1.exists() and not path2.exists()
    assert _uploaded_files(app) == []
    for doc_id in (first, second):
        assert _rows_for(app, "documents", ("document_id",), doc_id) == {}
    assert client.get("/forms/batches/B/download.zip").status_code == 404


# ---- 一覧表 ------------------------------------------------------------------------------

def _confirmed_import(app, client) -> int:
    """CSV を取り込んで確定まで進める（AI整形は使わない）。"""
    return _confirm_table(app, client, "トラブル一覧.csv", "トラブル対応一覧")


def test_downloading_the_table_zip_removes_everything(app, client):
    import_id = _confirmed_import(app, client)
    with app.app_context():
        imp = tables.get_import(import_id)
        template_id = imp["template_id"]
        stored = Path(app.config["UPLOAD_DIR"]) / imp["stored_path"]
        folder = tables.import_dir(import_id)
        # AI整形の控え（行ごとの結果と生の応答）も取り込みと一緒に消える
        conn = db.get_db()
        conn.execute("INSERT INTO llm_calls (cache_key, raw_text, created_at) VALUES ('K', 'あ', '2026-09-19')")
        conn.execute("""INSERT INTO ai_items (template_id, stage_id, row_key, cache_key, status, updated_at)
                        VALUES (?, 'log', 'TR-001', 'K', 'ok', '2026-09-19')""", (template_id,))
        conn.commit()
    assert stored.exists() and folder.exists()

    res = client.get(f"/tables/imports/{import_id}/download.zip")
    assert res.status_code == 200 and res.mimetype == "application/zip"
    assert "トラブル対応一覧_2026-08.md" in zipfile.ZipFile(io.BytesIO(res.data)).namelist()

    assert not stored.exists() and not folder.exists()
    assert _uploaded_files(app) == []
    assert _rows_for(app, "table_imports", ("import_id", "table_import_id"), import_id) == {}
    with app.app_context():
        conn = db.get_db()
        assert conn.execute("SELECT COUNT(*) FROM ai_items").fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM llm_calls").fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM jobs WHERE ref_id = ?", (import_id,)).fetchone()[0] == 0
        # 取り込み設定も取り込みの行ごと消える（保存しない）
        assert conn.execute("SELECT COUNT(*) FROM table_imports").fetchone()[0] == 0
    assert client.get(f"/tables/imports/{import_id}/download.zip").status_code == 404


def test_a_table_zip_that_cannot_be_built_deletes_nothing(app, client):
    """確定していない取り込みは zip を作れない。そのときデータも消さない。"""
    import_id = upload_csv(client, "途中.csv", CSV_TEXT)
    res = client.get(f"/tables/imports/{import_id}/download.zip")
    assert res.status_code == 302
    with app.app_context():
        imp = tables.get_import(import_id)
        assert imp is not None and (Path(app.config["UPLOAD_DIR"]) / imp["stored_path"]).exists()


# ---- 起動時の片付け -------------------------------------------------------------------------

def test_startup_removes_orphan_import_folders_but_keeps_work_in_progress(app, client, tmp_path):
    from app import create_app
    from tests.conftest import make_config

    import_id = upload_csv(client, "作業中.csv", CSV_TEXT)
    with app.app_context():
        working = tables.import_dir(import_id)
        working.mkdir(parents=True, exist_ok=True)
        (working / "rows.jsonl.gz").write_bytes(b"x")
        orphan = working.parent / "99999"
        orphan.mkdir(parents=True, exist_ok=True)
        (orphan / "rows.jsonl.gz").write_bytes(b"x")
        config = {"DATABASE": app.config["DATABASE"], "UPLOAD_DIR": app.config["UPLOAD_DIR"],
                  "DATA_DIR": app.config["DATA_DIR"], "TABLES_DIR": app.config["TABLES_DIR"]}

    create_app(make_config(tmp_path, **config))  # もう一度起動する
    assert not orphan.exists()          # DB に無い取り込みのフォルダは消える
    assert working.exists()             # 作業中のものは残る


def test_purge_is_safe_when_there_is_nothing_left(app):
    """二重にダウンロードされても（すでに消えていても）落ちない。"""
    with app.app_context():
        assert core.purge_documents([]) == 0
        assert core.purge_documents([999]) == 0
        assert core.purge_batch("") == 0
        assert core.purge_table_import(999) == 0
        assert core.recover_interrupted() == 0


# ---- 消したデータの残りかす（DBファイル・アップロードファイル） -----------------------------------

SECRET = "ヒミツ商事_山田一郎"


def _db_bytes(app) -> bytes:
    """instance/app.db 本体と、まだ書き戻されていない -wal の生バイト。"""
    base = Path(app.config["DATABASE"])
    data = b""
    for path in (base, base.with_name(base.name + "-wal")):
        if path.exists():
            data += path.read_bytes()
    return data


def test_downloaded_values_do_not_stay_in_the_database_file(app, client):
    """行を消すだけでは、解放されたページに帳票の値が平文で残る（design.md 3.3）。"""
    doc_id, _path = _add_confirmed_document(app, "残りかす.xlsx", value=SECRET)
    assert SECRET.encode("utf-8") in _db_bytes(app)      # 確定した時点では当然ある

    assert client.get(f"/forms/{doc_id}/download.md").status_code == 200
    assert SECRET.encode("utf-8") not in _db_bytes(app)  # ダウンロードしたら DB ファイルからも読めない


def test_downloaded_table_values_do_not_stay_in_the_database_file(app, client):
    import_id = _confirmed_import(app, client)
    with app.app_context():
        conn = db.get_db()
        conn.execute("UPDATE table_imports SET file_name = ? WHERE id = ?", (f"{SECRET}.csv", import_id))
        conn.commit()
    assert SECRET.encode("utf-8") in _db_bytes(app)

    assert client.get(f"/tables/imports/{import_id}/download.zip").status_code == 200
    assert SECRET.encode("utf-8") not in _db_bytes(app)


# ---- 送信できなかったダウンロード ----------------------------------------------------------

def test_a_download_that_is_cut_off_keeps_the_data(app, client):
    """本文を受け取れなかったとき（ブラウザを閉じた・通信が切れた）は消さない。消したら取り戻せないため。"""
    doc_id, path = _add_confirmed_document(app, "切れた.xlsx")
    res = client.get(f"/forms/{doc_id}/download.md", buffered=False)
    res.close()  # 本文を1バイトも受け取らずに切れた
    assert path.exists()
    with app.app_context():
        assert db.get_document(doc_id) is not None

    # もう一度ダウンロードすれば受け取れる（そのときは消える）
    res = client.get(f"/forms/{doc_id}/download.md")
    assert res.status_code == 200 and "EQ-001" in res.get_data(as_text=True)
    assert not path.exists()


def test_a_table_download_that_is_cut_off_keeps_the_data(app, client):
    import_id = _confirmed_import(app, client)
    client.get(f"/tables/imports/{import_id}/download.zip", buffered=False).close()
    with app.app_context():
        assert tables.get_import(import_id) is not None
    assert client.get(f"/tables/imports/{import_id}/download.zip").status_code == 200
    with app.app_context():
        assert tables.get_import(import_id) is None


# ---- まとめ取り込み: 確定済みだけをダウンロードする ------------------------------------------------

def test_only_the_confirmed_forms_of_a_batch_can_be_downloaded(app, client):
    first, path1 = _add_confirmed_document(app, "1.xlsx", value="EQ-001", batch_id="B", order=0)
    with app.app_context():
        pending = db.create_document("2.xlsx", "0" * 64, "documents/2.xlsx", batch_id="B", batch_order=1)

    # まだ読み取っていない帳票は確定できない＝zip に入らないので、件数にも数えない。
    # 「すべて消えます」とも言わず、残ることをそのまま書く
    page = _finish(client, [first, pending], current=first)["html"]
    assert "確定済み <strong>1</strong> / 1 件" in page
    assert "まとめて Markdown をダウンロード（zip）" in page
    assert "まだ読み取れていない1件（zip には入りません）" in page and "2.xlsx" in page

    res = client.get("/forms/batches/B/download.zip")
    assert res.status_code == 200 and res.mimetype == "application/zip"
    assert len(zipfile.ZipFile(io.BytesIO(res.data)).namelist()) == 1

    assert not path1.exists()                       # 渡した分だけ消える
    with app.app_context():
        assert db.get_document(first) is None
        assert db.get_document(pending) is not None  # 未確定の帳票は残る
    assert client.get("/forms/batches/B/download.zip").status_code == 404


# ---- 消えたあとの画面・ファイル名・メッセージ ------------------------------------------------------

def test_json_requests_after_the_data_is_gone_get_a_json_error(app, client):
    doc_id, _path = _add_confirmed_document(app, "消えた.xlsx")
    assert client.get(f"/forms/{doc_id}/download.md").status_code == 200
    res = client.post(f"/forms/{doc_id}/draft", json={"values": {}}, headers={"Accept": "application/json"})
    assert res.status_code == 404 and res.mimetype == "application/json"
    assert "残っていません" in res.get_json()["error"]


def test_downloads_have_a_readable_ascii_file_name(app, client):
    import_id = _confirmed_import(app, client)
    disposition = client.get(f"/tables/imports/{import_id}/download.zip").headers["Content-Disposition"]
    # filename* を読めない相手にも、何の zip か分かる名前を渡す（「_2026-09-19.zip」にしない）
    assert "filename=\"records_" in disposition and "filename*=UTF-8''" in disposition


def test_delete_messages_do_not_carry_the_file_name(app, client):
    """消したあとの応答にファイル名を載せない（クッキーにも画面にも残さない）。"""
    doc_id, _path = _add_confirmed_document(app, "社外秘_取引先A.xlsx")
    res = client.post(f"/forms/{doc_id}/delete")
    assert res.get_json()["ok"] is True
    assert "社外秘" not in str(res.headers.get("Set-Cookie", ""))
    assert "社外秘" not in res.get_data(as_text=True)


# ---- AI整形の控え（取り込み単位で消す） ----------------------------------------------------

def _ai_row(conn, template_id: int, import_id: int, row_key: str, cache_key: str) -> None:
    conn.execute("INSERT INTO llm_calls (cache_key, raw_text, created_at) VALUES (?, 'あ', '2026-09-19')",
                 (cache_key,))
    conn.execute("""INSERT INTO ai_items (template_id, import_id, stage_id, row_key, cache_key, status, updated_at)
                    VALUES (?, ?, 'log', ?, ?, 'ok', '2026-09-19')""",
                 (template_id, import_id, row_key, cache_key))


def test_purging_one_import_keeps_the_ai_results_of_another_import(app, client):
    """同じ取り込み設定で作業中の別の取り込みの AI整形結果（課金済み）まで消さない。"""
    import_id = _confirmed_import(app, client)
    with app.app_context():
        template_id = tables.get_import(import_id)["template_id"]
        other = tables.create_import("作業中.csv", "1" * 64, "tables/other.csv", {"kind": "csv"})
        conn = db.get_db()
        _ai_row(conn, template_id, import_id, "TR-001", "K1")
        _ai_row(conn, other, other, "TR-900", "K2")
        conn.commit()

    assert client.get(f"/tables/imports/{import_id}/download.zip").status_code == 200

    with app.app_context():
        conn = db.get_db()
        rows = conn.execute("SELECT row_key, import_id FROM ai_items").fetchall()
        assert [(r["row_key"], r["import_id"]) for r in rows] == [("TR-900", other)]
        assert [r[0] for r in conn.execute("SELECT cache_key FROM llm_calls")] == ["K2"]
        assert tables.get_import(other) is not None


def test_ai_results_of_an_import_that_is_gone_are_swept(app, client, tmp_path):
    """消した取り込みを指す AI整形の結果（消したあとに動いていた AI が書いた分など）も残さない。"""
    from app import create_app
    from tests.conftest import make_config

    import_id = _confirmed_import(app, client)
    with app.app_context():
        template_id = tables.get_import(import_id)["template_id"]
        conn = db.get_db()
        _ai_row(conn, template_id, 777, "TR-777", "K7")   # もう無い取り込みの分
        conn.commit()
    assert client.get(f"/tables/imports/{import_id}/download.zip").status_code == 200
    with app.app_context():
        conn = db.get_db()
        assert conn.execute("SELECT COUNT(*) FROM ai_items").fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM llm_calls").fetchone()[0] == 0
        # 起動時の片付けでも消える
        _ai_row(conn, template_id, 778, "TR-778", "K8")
        conn.commit()
        config = {"DATABASE": app.config["DATABASE"], "UPLOAD_DIR": app.config["UPLOAD_DIR"],
                  "DATA_DIR": app.config["DATA_DIR"], "TABLES_DIR": app.config["TABLES_DIR"]}
    restarted = create_app(make_config(tmp_path, **config))
    with restarted.app_context():
        conn = db.get_db()
        assert conn.execute("SELECT COUNT(*) FROM ai_items").fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM llm_calls").fetchone()[0] == 0


def _llm_call(conn, cache_key: str, import_id: int | None = None) -> None:
    conn.execute("INSERT INTO llm_calls (cache_key, raw_text, created_at, import_id) "
                 "VALUES (?, 'あ', '2026-09-19', ?)", (cache_key, import_id))


def _llm_keys(conn) -> list[str]:
    return sorted(r[0] for r in conn.execute("SELECT cache_key FROM llm_calls"))


def test_purging_one_import_keeps_the_unreferenced_responses_another_import_still_uses(app, client):
    """再依頼で直した行の1回目の応答は ai_items から参照されないが、再実行でキャッシュとして引く。
    別の取り込みを消したときに巻き添えで消すと、再実行で再課金になる。その取り込みが無くなれば消える。"""
    import_id = _confirmed_import(app, client)
    with app.app_context():
        template_id = tables.get_import(import_id)["template_id"]
        other = tables.create_import("作業中.csv", "1" * 64, "tables/other.csv", {"kind": "csv"})
        conn = db.get_db()
        _ai_row(conn, template_id, import_id, "TR-001", "K1")
        _ai_row(conn, other, other, "TR-900", "K2-repair")
        _llm_call(conn, "K2-first")   # 別の取り込みの、再依頼で直した行の1回目の応答（どこからも参照されない）
        conn.commit()

    assert client.get(f"/tables/imports/{import_id}/download.zip").status_code == 200
    with app.app_context():
        assert _llm_keys(db.get_db()) == ["K2-first", "K2-repair"]   # 消した取り込みの K1 だけ消える
        core.purge_table_import(other)
        assert _llm_keys(db.get_db()) == []


def test_purging_one_import_keeps_a_response_another_import_paid_for(app, client):
    """消す取り込みがキャッシュとして引いていただけの応答（払ったのは、まだある別の取り込み）は消さない。

    同じ文面の行は同じキーになるので、あとの取り込みは AI を呼ばずに前の応答を使う（llm_calls.import_id は
    払った取り込みのまま）。払った側は、一時停止・中断で ai_items の行を書く前に終わると、その応答を
    どこからも参照していない。消す取り込みの使ったキーだからと巻き添えで消すと、払った取り込みの
    再開・再実行で同じ応答をもう一度買うことになる（design.md 3.3）。
    """
    import_id = _confirmed_import(app, client)
    with app.app_context():
        template_id = tables.get_import(import_id)["template_id"]
        other = tables.create_import("作業中.csv", "1" * 64, "tables/other.csv", {"kind": "csv"})
        conn = db.get_db()
        _ai_row(conn, template_id, import_id, "TR-001", "K")      # 消す取り込みはキャッシュとして引いただけ
        conn.execute("UPDATE llm_calls SET import_id = ? WHERE cache_key = 'K'", (other,))   # 払ったのは別の取り込み
        _ai_row(conn, other, other, "TR-900", "K9")         # 別の取り込みは AI整形の作業中
        conn.commit()

    assert client.get(f"/tables/imports/{import_id}/download.zip").status_code == 200
    with app.app_context():
        assert _llm_keys(db.get_db()) == ["K", "K9"]
        core.purge_table_import(other)      # 払った取り込みを消せば、その応答も残らない
        assert _llm_keys(db.get_db()) == []


def test_purging_an_import_deletes_its_own_unreferenced_responses(app, client, tmp_path):
    """消した取り込みが払った応答は、どの ai_items からも参照されていなくても一緒に消す（design.md 3.3）。

    再依頼で直した行は ai_items が再依頼の応答のキーしか覚えていないので、1回目の応答はどこからも
    参照されない。ほかの取り込みで AI整形が動いていると、これまでは残ったままだった。
    ほかの取り込みが払った分（再実行で引くキャッシュ）は、これまでどおり残す。
    """
    from app import create_app
    from tests.conftest import make_config

    import_id = _confirmed_import(app, client)
    with app.app_context():
        template_id = tables.get_import(import_id)["template_id"]
        other = tables.create_import("作業中.csv", "1" * 64, "tables/other.csv", {"kind": "csv"})
        conn = db.get_db()
        _ai_row(conn, template_id, import_id, "TR-001", "K1-repair")
        _llm_call(conn, "K1-first", import_id)          # 消す取り込みの、再依頼で直した行の1回目の応答
        _ai_row(conn, other, other, "TR-900", "K2-repair")
        _llm_call(conn, "K2-first", other)              # 別の取り込みの分（再実行で引く。残す）
        conn.execute("UPDATE llm_calls SET import_id = ? WHERE cache_key = 'K1-repair'", (import_id,))
        conn.execute("UPDATE llm_calls SET import_id = ? WHERE cache_key = 'K2-repair'", (other,))
        conn.commit()
        config = {"DATABASE": app.config["DATABASE"], "UPLOAD_DIR": app.config["UPLOAD_DIR"],
                  "DATA_DIR": app.config["DATA_DIR"], "TABLES_DIR": app.config["TABLES_DIR"]}

    assert client.get(f"/tables/imports/{import_id}/download.zip").status_code == 200
    with app.app_context():
        assert _llm_keys(db.get_db()) == ["K2-first", "K2-repair"]

    restarted = create_app(make_config(tmp_path, **config))   # 起動時の片付けでも戻らない・巻き添えにしない
    with restarted.app_context():
        assert _llm_keys(db.get_db()) == ["K2-first", "K2-repair"]
        core.purge_table_import(other)
        assert _llm_keys(db.get_db()) == []


def test_the_ai_run_records_which_import_paid_for_each_response(ai_app, fake):
    """実際に AI を呼んだとき、生の応答に持ち主の取り込みを記録する（design.md 3.3）。

    型番違いの行は照合で引っかかって再依頼になり、ai_items は再依頼の応答のキーだけを覚える。
    1回目の応答も持ち主が分かるので、その取り込みを消せば残らない。
    """
    import aiproc

    iid = _make_import(ai_app, rows={"R2": ROWS["R2"]})
    with ai_app.app_context():
        aiproc.trial_row(iid, "R2", stage_ids=["log"])
        conn = db.get_db()
        rows = conn.execute("SELECT cache_key, import_id FROM llm_calls").fetchall()
        assert len(rows) == 2 and {r["import_id"] for r in rows} == {iid}   # 1回目と再依頼
        assert len({r["cache_key"] for r in rows} -
                   {r[0] for r in conn.execute("SELECT cache_key FROM ai_items")}) == 1   # 1回目は参照されない
        core.purge_table_import(iid)
        assert conn.execute("SELECT COUNT(*) FROM llm_calls").fetchone()[0] == 0


def test_a_response_saved_just_before_a_pause_survives_the_startup_sweep(app, client, tmp_path):
    """一時停止の直前に受け取った応答は ai_items の行がまだ無い。起動時の片付けで消すと、再開で再課金になる。"""
    from app import create_app
    from tests.conftest import make_config

    import_id = _confirmed_import(app, client)
    with app.app_context():
        conn = db.get_db()
        conn.execute("INSERT INTO jobs (kind, ref_type, ref_id, status, created_at, updated_at) "
                     "VALUES ('ai_format', 'table_import', ?, 'interrupted', '2026-09-19', '2026-09-19')", (import_id,))
        _llm_call(conn, "PAUSED")
        conn.commit()
        config = {"DATABASE": app.config["DATABASE"], "UPLOAD_DIR": app.config["UPLOAD_DIR"],
                  "DATA_DIR": app.config["DATA_DIR"], "TABLES_DIR": app.config["TABLES_DIR"]}
    restarted = create_app(make_config(tmp_path, **config))
    with restarted.app_context():
        assert _llm_keys(db.get_db()) == ["PAUSED"]
        core.purge_table_import(import_id)
        assert _llm_keys(db.get_db()) == []


def test_old_ai_rows_without_an_import_are_swept_once_their_template_has_no_import(app, client):
    """古いDBの import_id の無い ai_items（と、その応答）は、その取り込み設定に取り込みが残っていなければ消す。"""
    import_id = _confirmed_import(app, client)
    with app.app_context():
        template_id = tables.get_import(import_id)["template_id"]
        conn = db.get_db()
        conn.execute("""INSERT INTO ai_items (template_id, stage_id, row_key, cache_key, status, updated_at)
                        VALUES (?, 'log', 'TR-OLD', 'OLD', 'ok', '2026-09-19')""", (template_id,))
        _llm_call(conn, "OLD")
        conn.commit()
        core.sweep_orphan_ai(conn)
        assert _llm_keys(conn) == ["OLD"]   # 取り込みが残っている間は、持ち主かもしれないので残す
        conn.execute("UPDATE table_imports SET template_id = NULL WHERE id = ?", (import_id,))
        conn.commit()
        core.sweep_orphan_ai(conn)
        assert conn.execute("SELECT COUNT(*) FROM ai_items").fetchone()[0] == 0
        assert _llm_keys(conn) == []


def test_purge_does_not_vacuum_while_a_job_is_running(app):
    """VACUUM は DB 全体の書き込みを止める。動いているジョブの書き込みを「database is locked」で落とさない。"""
    with app.app_context():
        conn = db.get_db()
        seen: list[str] = []
        conn.set_trace_callback(seen.append)
        conn.execute("INSERT INTO jobs (kind, ref_type, ref_id, status, created_at, updated_at) "
                     "VALUES ('ai_format', 'table_import', 1, 'running', '2026-09-19', '2026-09-19')")
        conn.commit()
        core._shrink(conn)
        assert any("wal_checkpoint" in s for s in seen) and not any(s.strip() == "VACUUM" for s in seen)
        conn.execute("UPDATE jobs SET status = 'done'")
        conn.commit()
        seen.clear()
        core._shrink(conn)
        assert any(s.strip() == "VACUUM" for s in seen)
        conn.set_trace_callback(None)


# ---- 他サイトのページからのダウンロード（＝削除）を断る -----------------------------------------------
# ダウンロードはデータを消すので、GET でも他サイト発なら断る（<img> や別サイトのリンクで消されないように）。

CROSS_SITE_IMG = {"Sec-Fetch-Site": "cross-site", "Sec-Fetch-Mode": "no-cors", "Sec-Fetch-Dest": "image",
                  "Referer": "http://evil.example/"}


def test_a_cross_site_form_download_is_refused_and_keeps_the_data(app, client):
    doc_id, path = _add_confirmed_document(app, "他サイト.xlsx")
    res = client.get(f"/forms/{doc_id}/download.md", headers=CROSS_SITE_IMG)
    assert res.status_code == 403 and "EQ-001" not in res.get_data(as_text=True)
    assert path.exists() and _rows_for(app, "documents", ("document_id",), doc_id)


def test_a_form_download_with_a_foreign_referer_and_no_fetch_headers_is_refused(app, client):
    """Sec-Fetch-Site を付けないブラウザでも、Referer が別サイトなら断る。"""
    doc_id, path = _add_confirmed_document(app, "古いブラウザ.xlsx")
    res = client.get(f"/forms/{doc_id}/download.md", headers={"Referer": "http://evil.example/page"})
    assert res.status_code == 403
    assert path.exists() and _rows_for(app, "documents", ("document_id",), doc_id)


def test_a_cross_site_batch_zip_download_is_refused_and_keeps_the_data(app, client):
    first, path1 = _add_confirmed_document(app, "b1.xlsx", value="EQ-001", batch_id="X", order=0)
    second, path2 = _add_confirmed_document(app, "b2.xlsx", value="EQ-002", batch_id="X", order=1)
    res = client.get("/forms/batches/X/download.zip", headers=CROSS_SITE_IMG)
    assert res.status_code == 403 and res.mimetype != "application/zip"
    assert path1.exists() and path2.exists()
    for doc_id in (first, second):
        assert _rows_for(app, "documents", ("document_id",), doc_id)


def test_a_cross_site_table_zip_download_is_refused_and_keeps_the_data(app, client):
    import_id = _confirmed_import(app, client)
    with app.app_context():
        stored = Path(app.config["UPLOAD_DIR"]) / tables.get_import(import_id)["stored_path"]
    res = client.get(f"/tables/imports/{import_id}/download.zip", headers=CROSS_SITE_IMG)
    assert res.status_code == 403 and res.mimetype != "application/zip"
    with app.app_context():
        assert stored.exists() and tables.import_dir(import_id).exists()
    assert _rows_for(app, "table_imports", ("import_id", "table_import_id"), import_id)


def test_same_origin_and_typed_downloads_still_work(app, client):
    """アプリ内のクリック（same-origin）とアドレス欄への入力（none）は、これまでどおりダウンロードして消す。"""
    doc_id, path = _add_confirmed_document(app, "同じサイト.xlsx")
    res = client.get(f"/forms/{doc_id}/download.md",
                     headers={"Sec-Fetch-Site": "same-origin", "Referer": "http://localhost/forms/"})
    assert res.status_code == 200 and "EQ-001" in res.get_data(as_text=True)
    assert not path.exists()

    doc_id, path = _add_confirmed_document(app, "手入力.xlsx")
    res = client.get(f"/forms/{doc_id}/download.md", headers={"Sec-Fetch-Site": "none"})
    assert res.status_code == 200 and not path.exists()


def test_other_cross_site_gets_are_still_allowed(app, client):
    """消さない GET（画面の表示）は他サイトからのリンクでも開ける。"""
    assert client.get("/forms/", headers=CROSS_SITE_IMG).status_code == 200


# ---- 一部だけのダウンロード（Range）・削除の順番・取り込み設定の削除 -----------------------------------

def test_a_ranged_form_download_keeps_the_data(app, client):
    """Range 付き（ダウンロードマネージャー・再開）で一部だけ渡したときは消さない。全部渡したときだけ消す。"""
    doc_id, path = _add_confirmed_document(app, "一部だけ.xlsx")
    res = client.get(f"/forms/{doc_id}/download.md", headers={"Range": "bytes=0-9"})
    assert res.status_code in (200, 206)
    if res.status_code == 206:
        assert path.exists()
        with app.app_context():
            assert db.get_document(doc_id) is not None
        res = client.get(f"/forms/{doc_id}/download.md")
        assert res.status_code == 200 and "EQ-001" in res.get_data(as_text=True)
    assert not path.exists()
    with app.app_context():
        assert db.get_document(doc_id) is None


def test_a_ranged_table_download_keeps_the_data(app, client):
    import_id = _confirmed_import(app, client)
    res = client.get(f"/tables/imports/{import_id}/download.zip", headers={"Range": "bytes=0-99"})
    assert res.status_code in (200, 206)
    if res.status_code == 206:
        with app.app_context():
            assert tables.get_import(import_id) is not None
        assert client.get(f"/tables/imports/{import_id}/download.zip").status_code == 200
    with app.app_context():
        assert tables.get_import(import_id) is None


def _failing_delete(*_args, **_kwargs):
    raise RuntimeError("database is locked")


def test_files_stay_when_deleting_the_form_rows_fails(app, client, monkeypatch):
    """DB の削除に失敗したら、ファイルも残す（ファイルの無い帳票が作業中として残らないように）。"""
    doc_id, path = _add_confirmed_document(app, "消せない.xlsx")
    monkeypatch.setattr(core, "_delete_by_columns", _failing_delete)
    assert client.get(f"/forms/{doc_id}/download.md").status_code == 200
    assert path.exists()
    with app.app_context():
        assert db.get_document(doc_id) is not None


def test_files_stay_when_deleting_the_table_rows_fails(app, client, monkeypatch):
    import_id = _confirmed_import(app, client)
    with app.app_context():
        stored = Path(app.config["UPLOAD_DIR"]) / tables.get_import(import_id)["stored_path"]
        folder = tables.import_dir(import_id)
    monkeypatch.setattr(core, "_delete_by_columns", _failing_delete)
    assert client.get(f"/tables/imports/{import_id}/download.zip").status_code == 200
    assert stored.exists() and folder.exists()
    with app.app_context():
        assert tables.get_import(import_id) is not None
    # 失敗が解消すれば、もう一度ダウンロードして消せる
    monkeypatch.undo()
    assert client.get(f"/tables/imports/{import_id}/download.zip").status_code == 200
    assert not stored.exists() and not folder.exists()


def test_ai_responses_left_by_a_template_delete_are_swept_at_startup(app, client):
    """AI整形の控え（ai_items）が消えたのに生の応答（llm_calls）が残る場面は、起動時の片付けで消す。"""
    from app import _cleanup_leftovers

    import_id = _confirmed_import(app, client)
    with app.app_context():
        conn = db.get_db()
        _ai_row(conn, import_id, import_id, "TR-001", "K")
        conn.execute("DELETE FROM ai_items WHERE import_id = ?", (import_id,))
        conn.commit()
        assert conn.execute("SELECT COUNT(*) FROM llm_calls").fetchone()[0] == 1
    _cleanup_leftovers(app)
    with app.app_context():
        assert db.get_db().execute("SELECT COUNT(*) FROM llm_calls").fetchone()[0] == 0


def test_sweep_orphan_ai_removes_responses_nobody_uses(app):
    with app.app_context():
        conn = db.get_db()
        conn.execute("INSERT INTO llm_calls (cache_key, raw_text, created_at) VALUES ('Z', ?, '2026-09-19')",
                     (SECRET,))
        conn.commit()
        assert core.sweep_orphan_ai(conn) == 1
        assert conn.execute("SELECT COUNT(*) FROM llm_calls").fetchone()[0] == 0
    assert SECRET.encode("utf-8") not in _db_bytes(app)


def test_deleting_an_import_removes_the_raw_ai_responses_at_once(app, client):
    """取り込みを消したら、生の応答（llm_calls）も起動を待たずにすぐ消える。"""
    import_id = _confirmed_import(app, client)
    with app.app_context():
        conn = db.get_db()
        _ai_row(conn, import_id, import_id, "TR-001", "K")
        conn.commit()
        assert conn.execute("SELECT COUNT(*) FROM llm_calls").fetchone()[0] == 1
        core.purge_table_import(import_id)
        assert conn.execute("SELECT COUNT(*) FROM llm_calls").fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM ai_items").fetchone()[0] == 0


def test_a_ranged_download_gets_the_whole_file_and_removes_the_item(app, client):
    """Range 付きでも全体を 200 で返す（一部だけ渡して消すことが無い）。全部渡したので消える。"""
    doc_id, path = _add_confirmed_document(app, "範囲指定.xlsx")
    res = client.get(f"/forms/{doc_id}/download.md", headers={"Range": "bytes=0-9"})
    assert res.status_code == 200 and "EQ-001" in res.get_data(as_text=True)
    assert not path.exists()
    import_id = _confirmed_import(app, client)
    res = client.get(f"/tables/imports/{import_id}/download.zip", headers={"Range": "bytes=0-99"})
    assert res.status_code == 200
    assert "トラブル対応一覧_2026-08.md" in zipfile.ZipFile(io.BytesIO(res.data)).namelist()
    with app.app_context():
        assert tables.get_import(import_id) is None


def test_purge_after_send_keeps_the_item_for_a_partial_response(app):
    """206（一部だけ）・304 の応答では、送り終えても消さない。"""
    from flask import Response

    called = []
    with app.app_context():
        for status in (206, 304):
            res = core.purge_after_send(Response(b"abc", status=status), called.append, status)
            list(res.response)
            res.close()
        assert called == []
        res = core.purge_after_send(Response(b"abc", status=200), called.append, 200)
        list(res.response)
        res.close()
    assert called == [200]


def test_a_folder_that_cannot_be_removed_is_logged(app, client, monkeypatch, caplog):
    """imports/<id>/ を消し切れなかったら、黙って成功扱いにせず記録する。"""
    import logging

    import_id = _confirmed_import(app, client)
    monkeypatch.setattr(core.shutil, "rmtree", lambda *a, **k: None)
    with app.app_context(), caplog.at_level(logging.WARNING, logger="core"):
        assert core.import_dir(import_id).exists()
        core.purge_table_import(import_id)
        assert tables.get_import(import_id) is None
    assert f"imports/{import_id}" in caplog.text


# ====================================================================================================
# 元 tests/test_sessions.py
# 社内LANで数人が同時に使うときの分かれ方（2026-09-20 の利用者の指示）。
#
# - 取り込んだ帳票・一覧表は、置いたブラウザ（セッションのクッキー）のものだけが見える・触れる。
#   ほかのブラウザからは「無い」として扱う（403 ではなく 404。あることも知らせない）。
# - ジョブの進み具合も、その取り込みを持っているブラウザにしか見せない。
# - 「設定」（帳票の種類・一覧表の取り込み設定）はみんなで使うので分けない。
# - 読み込みは同時に動く（誰かの大きい表で、ほかの人が待たされない）。
# ====================================================================================================

@pytest.fixture
def other_client(app):
    """もう1台のPC（別のブラウザ）。クッキーが別なので作業場所も別になる。"""
    app.test_client_class = BufferedClient
    return app.test_client()


# ---- 一覧表の取り込み --------------------------------------------------------------

def test_another_browser_cannot_see_or_delete_an_import(app, client, other_client):
    import_id = upload_csv(client, "他人の一覧.csv")

    for path in (f"/tables/imports/{import_id}/panel/source", f"/tables/imports/{import_id}/panel/preview",
                 f"/tables/imports/{import_id}/download.zip", f"/tables/imports/{import_id}/issues.csv"):
        assert other_client.get(path).status_code == 404, path
    for path in (f"/tables/imports/{import_id}/source", f"/tables/imports/{import_id}/layout",
                 f"/tables/imports/{import_id}/columns", f"/tables/imports/{import_id}/read",
                 f"/tables/imports/{import_id}/confirm", f"/tables/imports/{import_id}/preview",
                 f"/tables/imports/{import_id}/cancel", f"/tables/imports/{import_id}/ai/run",
                 f"/tables/imports/{import_id}/ai/trial", f"/tables/imports/{import_id}/delete"):
        assert other_client.post(path, json={}).status_code == 404, path

    with app.app_context():   # 消されていない（持ち主はそのまま使える）
        assert tables.get_import(import_id) is not None
    assert client.get(f"/tables/imports/{import_id}/panel/source").status_code == 200


def test_another_browsers_job_is_invisible(app, client, other_client):
    """ほかのブラウザの取り込みのジョブは、進み具合も見せない。"""
    import_id = imported(app, client, "進み具合.csv", "同時実行テスト")
    with app.app_context():
        job_id = tables.get_import(import_id)["job_id"]

    assert client.get(f"/api/jobs/{job_id}").status_code == 200
    assert other_client.get(f"/api/jobs/{job_id}").status_code == 404


# ---- 帳票の取り込み ----------------------------------------------------------------

def test_another_browser_cannot_see_or_delete_a_form(app, client, other_client, sample_dir):
    doc_id = upload_forms(client, sample_dir / "standard.xlsx")[0]

    assert other_client.get(f"/forms/{doc_id}/type").status_code == 404
    assert other_client.get(f"/forms/{doc_id}/type").status_code == 404
    assert other_client.get(f"/forms/{doc_id}/download.md").status_code == 404
    assert other_client.get(f"/forms/{doc_id}/original").status_code == 404
    assert other_client.post(f"/forms/{doc_id}/draft", json={}).status_code == 404
    assert other_client.post(f"/forms/{doc_id}/delete").status_code == 404
    # 番号を並べて送っても、ほかのブラウザの帳票は「確定してダウンロード」の欄に出てこない
    assert other_client.get(f"/forms/finish?ids={doc_id}").get_json()["total"] == 0

    import database as db
    with app.app_context():
        assert db.get_document(doc_id) is not None
    assert client.get(f"/forms/{doc_id}/type").status_code == 200


# ---- 設定は分けない ----------------------------------------------------------------

def test_settings_are_shared_between_browsers(app, client, other_client, sample_dir):
    """帳票の種類はみんなで使う設定なので、別のブラウザからも見える。表の取り込み設定は保存しない。"""
    from tests.test_forms import add_field, create_type

    import_id = imported(app, client, "みんなの一覧.csv", "この取り込みだけの名前")
    pattern_id = create_type(client, sample_dir / "standard.xlsx", name="共有の帳票の種類")
    add_field(client, pattern_id, "修理報告書", "A4")

    other_import = upload_csv(other_client, "別の人の一覧.csv")
    # 表の設定は取り込みの中にしかないので、ほかのブラウザには名前も出ない
    assert "この取り込みだけの名前" not in panel_html(other_client, other_import, "source")
    assert other_client.get(f"/tables/imports/{import_id}/panel/columns").status_code == 404
    panel = other_client.get(f"/form-types/{pattern_id}/panel")
    assert panel.status_code == 200 and "共有の帳票の種類" in panel.get_json()["html"]


# ---- 同時に動く --------------------------------------------------------------------

def test_two_reads_run_at_the_same_time(app):
    """別々の取り込みの読み込みは同時に動く（1本ずつなら待ち合わせが成立せず時間切れになる）。"""
    both_started = threading.Barrier(2, timeout=10)

    def work(ctx):
        both_started.wait()      # もう1つが動き出すまで待つ
        return {"ok": True}

    with app.app_context():
        first = core.start_job("table_read", "table_import", 101, work)
        second = core.start_job("table_read", "table_import", 102, work)
        assert core.wait_job(first, timeout=20)["status"] == "done"
        assert core.wait_job(second, timeout=20)["status"] == "done"


def test_two_jobs_never_touch_the_same_import(app):
    """同じ取り込みのジョブは、ワーカーが増えても同時には動かさない。"""
    inside, overlapped = threading.Semaphore(0), []
    release = threading.Event()

    def work(ctx):
        inside.release()
        if not release.wait(5):
            return {"ok": False}
        return {"ok": True}

    def peek(ctx):
        overlapped.append(core.get_job(first)["status"])
        return {"ok": True}

    with app.app_context():
        first = core.start_job("table_read", "table_import", 201, work)
        assert inside.acquire(timeout=10)
        second = core.start_job("table_render", "table_import", 201, peek)
        assert core.get_job(second)["status"] == "queued"
        release.set()
        assert core.wait_job(second, timeout=20)["status"] == "done"
        assert overlapped and overlapped[0] in ("done", "cancelled")   # 前のジョブが終わってから動いた


# ====================================================================================================
# 元 tests/test_shell.py
# 画面の骨組み: 3つの画面へのナビ、エラー画面、Host の確認、キャッシュ、他サイトからの書き込み。
#
# 画面は「帳票取り込み・表の取り込み・帳票登録」の3つだけになった（2026-09-20 の作り直し）。
# ホーム画面・設定の画面・取り込み履歴・作業中の一覧は無い。AI接続は「表の取り込み」画面の
# AI整形の段の中（/tables/ai-connection）へ移した。
# ====================================================================================================

# ---- 共通 -------------------------------------------------------------------------

@pytest.fixture
def fake_shell():
    with FakeServer() as server:
        yield server


@pytest.fixture
def ai_app_shell(tmp_path, fake_shell):
    return create_app(make_config(tmp_path, OPENAI_BASE_URL=f"{fake_shell.url}/v1", OPENAI_API_KEY="",
                                  OPENAI_MODELS=[], OPENAI_MODEL="gpt-test"))


@pytest.fixture
def ai_client(ai_app_shell):
    return ai_app_shell.test_client()


def _save_ai_settings(client, **overrides):
    """AI接続の保存（「表の取り込み」画面の AI整形の段が送る fetch と同じ）。"""
    body = {"models": ["gpt-test", "picky-model"], "default": "gpt-test", "api_key": OPENAI_KEY,
            "chat_url": "", "models_url": "", "add_models": ""}
    body.update(overrides)
    return client.post("/tables/ai-connection", json=body)


# ---- ナビ（3つの画面しかない） ---------------------------------------------------------------

def test_root_goes_to_the_form_import_screen(client):
    """ホーム画面は無い。「/」（古いブックマーク）は帳票取り込みへ送る。"""
    res = client.get("/")
    assert res.status_code == 302 and res.headers["Location"] == "/forms/new"


def test_nav_has_only_the_three_screens(client):
    page = client.get("/forms/new").get_data(as_text=True)
    for href, label in (('href="/forms/new"', "帳票取り込み"), ('href="/tables/new"', "表の取り込み"),
                        ('href="/form-types/"', "帳票登録")):
        assert href in page and label in page, href
    for url in ("/forms/new", "/tables/new", "/form-types/"):
        assert client.get(url).status_code == 200, url
    # 無くした画面の言葉はヘッダーに出さない
    for word in ("設定", "取り込み履歴", "LightRAGへの入れ方", "保存先フォルダ", "名寄せ辞書"):
        assert word not in page, word


def test_the_removed_screens_are_gone(client):
    """設定・履歴・ホームの一覧・LightRAG案内・保存先フォルダの URL は残っていない。"""
    for url in ("/home", "/history/", "/documents", "/patterns",
                "/settings/ai", "/settings/output", "/settings/lightrag", "/settings/lightrag/entity-types.yml",
                "/settings/table-templates", "/settings/form-types/", "/settings/models", "/settings/aliases",
                "/api/models", "/forms/1/ai-classify", "/forms/1/ai-fill"):
        assert client.get(url).status_code == 404, url


def test_old_deep_links_come_back_to_the_one_page_screens(client):
    """作り直す前の URL（お気に入り・開いたままの画面）は、その画面の1枚ページへ送る。"""
    for url, target in (("/forms/1", "/forms/"), ("/forms/1/done", "/forms/"),
                        ("/form-types/new", "/form-types/"), ("/form-types/1/edit", "/form-types/")):
        res = client.get(url)
        assert res.status_code == 302 and res.headers["Location"] == target, url


def test_ui_macros_render(app):
    grid = [{"name": "修理報告書", "letters": ["A", "B"], "truncated": False, "images": 1, "click_cells": ["A1"],
             "rows": [{"index": 1, "cells": [{"coord": "A1", "text": "報告番号", "rowspan": 1, "colspan": 2,
                                              "label": True},
                                             {"coord": "B1", "text": "R-1", "rowspan": 1, "colspan": 1,
                                              "label": False}]}]}]
    source = """{% import "ui.html" as ui %}
    {% call ui.step("sheet", 2, "帳票の種類とシート", open=True) %}本文{% endcall %}
    {% call ui.card("見出し") %}本文{% endcall %}
    {{ ui.badge("modified") }}{{ ui.badge("active") }}
    {{ ui.work_line("read") }}
    {{ ui.progress({"status": "running", "progress": {"done": 3, "total": 10}, "message": "処理中"}, url="/api/jobs/1") }}
    {{ ui.data_grid([{"index": 1, "cells": ["管理No", "設備"], "kind": "header"},
                     {"index": 2, "cells": ["TR-1", "CMP-101"], "kind": "data", "strike": True}] + [{"index": 3, "cells": ["x"] * 28}]) }}
    {{ ui.grid_legend() }}
    {{ ui.sheet_grid(grids, click_cells=True) }}
    {{ ui.file_drop("file", ".xlsx,.xlsm") }}"""
    with app.test_request_context("/forms/new"):
        html = render_template_string(source, grids=grid)
    assert 'data-step="sheet"' in html and "is-open" in html
    assert "修正中" in html and "使用中" in html
    assert 'data-work="read"' in html
    assert 'data-job-url="/api/jobs/1"' in html and 'aria-valuenow="30"' in html
    assert "row-header" in html and "row-strike" in html and ">AB<" in html
    assert 'data-cell="A1"' in html and 'colspan="2"' in html and 'class="is-label"' in html
    assert 'data-table-head="1"' in html and "画像が1件" in html
    assert 'accept=".xlsx,.xlsm"' in html


# ---- AI接続（設定画面は無い。「表の取り込み」画面の AI整形の段の中） -----------------------------------

def test_ai_connection_is_saved_from_the_table_screen(ai_client, ai_app_shell):
    res = _save_ai_settings(ai_client)
    assert res.status_code == 200
    body = res.get_json()
    assert body["ok"] is True and "AI接続の設定を保存しました" in body["message"]
    assert OPENAI_KEY not in res.get_data(as_text=True)   # キーの値は返さない

    saved = yaml.safe_load((ai_app_shell.config["DATA_DIR"] / "model_settings.yaml").read_text(encoding="utf-8"))
    assert saved["models"] == ["gpt-test", "picky-model"] and saved["api_key"] == OPENAI_KEY
    assert "chat_url" not in saved   # env と同じURLは上書きとして持たない

    catalog = ai_client.post("/tables/ai-connection/models").get_json()
    assert catalog["ok"] is True and "catalog-only" in catalog["models"]   # APIの models.list() から取得


def test_ai_connection_validation(ai_client):
    assert "APIキーの長さが不自然です" in _save_ai_settings(ai_client, api_key="short").get_json()["error"]
    assert "/chat/completions で終わるフルパス" in \
        _save_ai_settings(ai_client, chat_url="https://example.com/v1").get_json()["error"]
    assert "候補に入っていません" in _save_ai_settings(ai_client, default="not-in-list").get_json()["error"]


def test_ai_connection_test(ai_client, fake_shell):
    result = ai_client.post("/tables/ai-connection/test").get_json()
    assert result["ok"] is False and "未設定" in result["steps"][0]["detail"]

    _save_ai_settings(ai_client)
    fake_shell.chat_replies = ["OK"]
    result = ai_client.post("/tables/ai-connection/test").get_json()
    assert result["ok"] is True
    assert [s["ok"] for s in result["steps"]] == [True, True]
    assert "3件のモデル" in result["steps"][0]["detail"] and "OK" in result["steps"][1]["detail"]
    assert fake_shell.requests[-1]["path"] == "/v1/chat/completions"


# ---- エラー画面・Host の確認 ----------------------------------------------------------

def test_not_found_page_is_japanese(client):
    """消した帳票の URL を開いても、英語の既定ページではなく日本語の案内と戻り道を出す。"""
    for url in ("/forms/999/original", "/tables/imports/999/panel/preview", "/form-types/999/panel"):
        res = client.get(url)
        assert res.status_code == 404, url
        page = res.get_data(as_text=True)
        assert "見つかりません" in page, url
        assert 'href="/forms/new"' in page, url


def test_not_found_answers_json_when_the_screen_asked_for_json(client):
    """開いたままの画面が、消えた帳票へ途中保存・プレビューを送ったとき（app.js の postJson）。"""
    res = client.post("/forms/999/draft", json={"values": {}})
    assert res.status_code == 404 and "サーバーに残っていません" in res.get_json()["error"]


def test_method_not_allowed_page_is_japanese(client):
    """送信専用の URL をアドレス欄から開いても、Werkzeug の英語の画面を出さない（design.md 2）。"""
    res = client.get("/forms/1/delete")
    assert res.status_code == 405
    page = res.get_data(as_text=True)
    assert "Method Not Allowed" not in page and "取り込みの画面から開き直してください" in page
    assert 'href="/forms/new"' in page
    res = client.get("/forms/1/delete", headers={"Accept": "application/json"})
    assert res.status_code == 405 and "取り込みの画面から" in res.get_json()["error"]


def test_bad_request_page_is_japanese(client):
    res = client.get("/forms/new", headers={"Host": "evil.example"})
    assert res.status_code == 400 and "Bad Request" not in res.get_data(as_text=True)


def test_other_host_is_refused(client):
    """待ち受けていない名前で届いたリクエストは断る（DNSリバインディング対策）。"""
    assert client.get("/forms/new", headers={"Host": "127.0.0.1:5000"}).status_code == 200
    assert client.get("/forms/new", headers={"Host": "localhost:5000"}).status_code == 200
    assert client.get("/forms/new", headers={"Host": "evil.example"}).status_code == 400
    assert client.get("/tables/new", headers={"Host": "evil.example:5000"}).status_code == 400


def test_a_refused_host_is_told_which_address_works(client):
    """別の名前（PC名・hosts の別名）で開かれたときは、同じアドレスへ戻さず、開けるアドレスへ案内する。"""
    res = client.get("/forms/new", headers={"Host": "mypc:5123"})
    page = res.get_data(as_text=True)
    assert res.status_code == 400
    assert "http://127.0.0.1:5123/" in page and 'href="http://127.0.0.1:5123/"' in page
    assert "このページは再読み込みや古いアドレスからは開けません" not in page   # 断られた理由に合う文だけ出す
    res = client.post("/forms/1/delete", headers={"Host": "mypc", "Accept": "application/json"})
    assert res.status_code == 400 and "http://127.0.0.1:5000/" in res.get_json()["error"]


# ---- 社内LANで動かす（2026-09-20 の運用変更） ------------------------------------------
# サーバ（JupyterLab のターミナルなど）で起動し、他のPCから開いて数人で使う。
# 待ち受け先は環境変数（HOST・PORT）で決め、受け付ける宛先の名前は app.allowed_hosts が決める。

def test_the_operator_chooses_where_to_listen(monkeypatch):
    """PORT は環境変数から読み、数字でなければ理由を出して起動しない。"""
    import app as app_module

    monkeypatch.setenv("PORT", "8080")
    assert app_module._env_port() == 8080
    monkeypatch.setenv("PORT", "")
    assert app_module._env_port() == 5000          # 未設定なら既定
    monkeypatch.setenv("PORT", "ポート")
    with pytest.raises(SystemExit, match="PORT"):
        app_module._env_port()


def test_lan_addresses_are_accepted_only_when_we_listen_on_the_lan(monkeypatch):
    """LAN に出したときだけ、このサーバ自身の名前・アドレスでも開ける（出していなければループバックだけ）。"""
    import socket

    import app as app_module

    monkeypatch.delenv("ALLOWED_HOSTS", raising=False)
    own = socket.gethostname().split(".")[0].lower()
    assert app_module.allowed_hosts("127.0.0.1") == {"127.0.0.1", "localhost", "::1"}
    lan = app_module.allowed_hosts("0.0.0.0")
    assert {"127.0.0.1", "localhost", "::1"} <= lan and own in lan
    assert "evil.example" not in lan
    # 特定のアドレスを指定したときは、そのアドレスでも開ける
    assert "192.168.10.20" in app_module.allowed_hosts("192.168.10.20")


def test_the_operator_can_add_names_with_allowed_hosts(monkeypatch):
    """別名（社内DNSの名前）で開くときは ALLOWED_HOSTS で足す。`*` はすべて受け付ける。"""
    import app as app_module

    monkeypatch.setenv("ALLOWED_HOSTS", "rag.example.local:5000, RAG-SERVER")
    names = app_module.allowed_hosts("0.0.0.0")
    assert "rag.example.local" in names and "rag-server" in names   # ポートは外し、小文字に揃える
    monkeypatch.setenv("ALLOWED_HOSTS", "*")
    assert app_module.allowed_hosts("0.0.0.0") == {"*"}


def test_a_request_to_an_allowed_name_is_served(tmp_path):
    """運用者が意図した名前で届いた要求は通し、それ以外は 400 のまま。"""
    other = create_app(make_config(tmp_path, ALLOWED_HOSTS={"rag-server", "192.168.10.20"}))
    with other.test_client() as c:
        assert c.get("/forms/new", headers={"Host": "RAG-Server:5000"}).status_code == 200
        assert c.get("/forms/new", headers={"Host": "192.168.10.20:5000"}).status_code == 200
        assert c.get("/forms/new", headers={"Host": "127.0.0.1:5000"}).status_code == 200
        assert c.get("/forms/new", headers={"Host": "evil.example"}).status_code == 400


def test_the_startup_notice_says_there_is_no_login(monkeypatch):
    """LAN に出して起動したときは、ログインが無いことと消えることを必ず知らせる。"""
    import app as app_module

    monkeypatch.setattr(app_module, "HOST", "0.0.0.0")
    notice = app_module.startup_notice(5000)
    assert "ログインはありません" in notice and "誰でも使えます" in notice
    assert "ダウンロードするとサーバーからデータが消えます" in notice
    assert "http://" in notice and ":5000/" in notice
    monkeypatch.setattr(app_module, "HOST", "127.0.0.1")
    assert app_module.startup_notice(5000) == "[app] http://127.0.0.1:5000/ で起動しました（このサーバの中からだけ開けます）"


def test_pages_and_json_are_not_kept_in_the_browser_cache(app, client):
    """取り込んだ値の載る画面・JSON はブラウザに保存させない（ダウンロードで消したあと、戻るで出さない）。"""
    doc_id, _path = add_confirmed_document(app, "キャッシュ.xlsx")
    for url, code in (("/forms/new", 200), ("/tables/new", 200), (f"/forms/finish?ids={doc_id}", 200),
                      ("/forms/999/original", 404), ("/api/jobs/999", 404)):
        res = client.get(url)
        assert res.status_code == code, url
        assert res.headers.get("Cache-Control") == "no-store", url
    assert client.get("/forms/new", headers={"Host": "evil.example"}).headers.get("Cache-Control") == "no-store"
    static = client.get("/static/app.js")
    assert static.status_code == 200 and "no-store" not in (static.headers.get("Cache-Control") or "")


def test_no_response_can_be_shown_inside_another_sites_frame(client):
    """ほかのサイトの iframe に入れて削除ボタンなどを押させない（クリックジャッキング）。静的ファイル・エラーも同じ。"""
    responses = {url: client.get(url) for url in ("/forms/new", "/tables/new", "/static/app.js", "/forms/999/original")}
    responses["他のホスト名"] = client.get("/forms/new", headers={"Host": "evil.example"})
    responses["他のサイトからの書き込み"] = client.post("/forms/1/delete", headers={"Origin": "http://evil.example"})
    for name, res in responses.items():
        assert res.headers.get("X-Frame-Options") == "DENY", name
        assert res.headers.get("Content-Security-Policy") == "frame-ancestors 'none'", name
    assert responses["/forms/new"].status_code == 200 and responses["/tables/new"].status_code == 200
    assert responses["他のサイトからの書き込み"].status_code == 403


def test_server_error_page_is_japanese(tmp_path):
    """想定外の例外でも、英語の既定ページではなく日本語の案内と戻り道を出す。"""
    app = create_app(make_config(tmp_path, PROPAGATE_EXCEPTIONS=False))

    @app.route("/_raise_for_test")
    def _raise_for_test():
        raise RuntimeError("テスト用の例外")

    res = app.test_client().get("/_raise_for_test")
    assert res.status_code == 500
    page = res.get_data(as_text=True)
    assert "エラーが発生しました" in page and 'href="/forms/new"' in page


def test_too_large_upload_page_is_japanese(tmp_path):
    """MAX_CONTENT_LENGTH を超えた送信は Werkzeug の英語の画面ではなく、日本語の案内と戻り道。"""
    small = create_app(make_config(tmp_path, MAX_CONTENT_LENGTH=1024))
    res = small.test_client().post("/forms/upload", data={"file": (io.BytesIO(b"x" * 5000), "大きい.xlsx")},
                                   content_type="multipart/form-data")
    body = res.get_data(as_text=True)
    assert res.status_code == 413
    assert "ファイルが大きすぎます" in body and "合計 1KB まで" in body and "分けて" in body
    assert "Request Entity Too Large" not in body and 'href="/forms/new"' in body
    res = small.test_client().post("/forms/upload", data=b"x" * 5000, headers={"Accept": "application/json"},
                                   content_type="application/json")
    assert res.status_code == 413 and "大きすぎます" in res.get_json()["error"]


# ---- 起動時の片付け ---------------------------------------------------------------------

def test_startup_removes_orphan_uploads(tmp_path):
    """起動時に、DB から参照されていない取り込み途中の残骸だけを片付ける。"""
    config = make_config(tmp_path)
    app = create_app(config)
    documents = app.config["UPLOAD_DIR"] / "documents"
    documents.mkdir(parents=True, exist_ok=True)
    orphan = documents / f"{'a' * 32}.xlsx"
    orphan.write_bytes(b"broken")
    old = time.time() - 3600
    os.utime(orphan, (old, old))   # 取り込みの途中ではない（できてから時間がたった）残骸
    fresh = documents / f"{'c' * 32}.xlsx"
    fresh.write_bytes(b"uploading")   # 別に起動しているアプリが保存したばかり（DB の行を作る前）
    used = documents / f"{'b' * 32}.xlsx"
    used.write_bytes(b"used")
    with app.app_context():
        db.create_document("点検表.xlsx", "0" * 64, f"documents/{used.name}")

    create_app(config)   # 起動し直す

    assert not orphan.exists()
    assert used.exists()
    assert fresh.exists()   # できたばかりのファイルは消さない


# ---- 他サイトからの書き込みを断る ------------------------------------------------------------

CROSS_SITE = {"Origin": "https://evil.example", "Sec-Fetch-Site": "cross-site"}


def test_cross_site_post_is_refused(ai_client):
    res = ai_client.post("/tables/ai-connection", json={"chat_url": "http://evil.example/v1/chat/completions"},
                         headers=CROSS_SITE)
    assert res.status_code == 403
    assert ai_client.post("/tables/ai-connection/test", headers=CROSS_SITE).status_code == 403
    # Origin だけ・Sec-Fetch-Site だけでも断る
    assert ai_client.post("/tables/ai-connection/test", headers={"Origin": "https://evil.example"}).status_code == 403
    assert ai_client.post("/tables/ai-connection/test", headers={"Sec-Fetch-Site": "cross-site"}).status_code == 403


def test_same_origin_post_still_works(ai_client):
    """この画面から送られた（同じサイトの）書き込みは通す。"""
    res = ai_client.post("/tables/ai-connection",
                         json={"models": ["gpt-test"], "default": "gpt-test", "api_key": ""},
                         headers={"Origin": "http://localhost", "Sec-Fetch-Site": "same-origin"})
    assert res.status_code == 200 and res.get_json()["ok"] is True


def test_cross_site_get_is_allowed(client):
    assert client.get("/forms/new", headers=CROSS_SITE).status_code == 200


def test_refused_write_shows_a_japanese_page(ai_client):
    """断ったことが日本語で分かり、取り込みの画面に戻れる（Flask の英語の403ページを出さない）。"""
    res = ai_client.post("/tables/ai-connection/test", headers=CROSS_SITE)
    assert res.status_code == 403
    page = res.get_data(as_text=True)
    assert "ほかのサイトのページから送られてきた操作" in page
    assert "データは変わっていません" in page and 'href="/forms/new"' in page
    assert "Forbidden" not in page


def test_refused_write_answers_json_when_the_screen_asked_for_json(ai_client):
    """画面の JSON 送信（app.js の postJson）には JSON で返す。HTML だと理由が出ない。"""
    res = ai_client.post("/tables/ai-connection/test", headers={**CROSS_SITE, "Accept": "application/json"})
    assert res.status_code == 403
    assert "ほかのサイト" in res.get_json()["error"]


# ====================================================================================================
# 元 tests/test_endpoints_settings.py
# 画面の経路: 列の対応づけの保存、帳票の種類の見本と状態、
# 帳票の元ファイルのダウンロード、Excel のシートの選択。
# ====================================================================================================

# ---- 列の対応づけ（取り込みごとに決める。設定だけを開く画面は無い） ------------------------------------

def test_saving_the_columns_decides_everything_but_use_and_role_from_the_values(app, client):
    """画面が送るのは「使う・役割」だけ。キー・型・単位・出し方は見出しと値から決める。"""
    import_id = upload_csv(client, "a.csv")
    assert save_source(client, import_id, encoding="cp932", delimiter=",").status_code == 200
    assert save_layout(client, import_id).status_code == 200

    body = editor_body(client, import_id)
    assert set(body["columns"][0]) == {"index", "header", "use", "role"}
    assert body["name"] == "a"   # 表の名前の初期値はファイル名
    body["name"] = "トラブル対応一覧"
    assert save_columns(client, import_id, body).status_code == 200
    wait_import_job(app, import_id)

    with app.app_context():
        spec = get_import(import_id)["spec"]
    assert spec.name == "トラブル対応一覧"
    assert spec.column("occurred_at").type in ("date", "datetime")   # 型は値から
    assert spec.column("downtime").unit == "分"                # 単位は見出しから
    assert spec.column("symptom").md in ("body", "attribute")  # md での扱いも自動
    assert spec.markdown["group_by"] == "month"               # まとめ方は常に月
    assert spec.markdown["file_prefix"] == "" and spec.file_prefix == "トラブル対応一覧"


def test_saving_the_columns_again_replaces_the_imports_spec(app, client):
    """保存し直しても版は作らず、その取り込みの設定を上書きする。"""
    import_id = upload_csv(client, "a.csv")
    save_source(client, import_id, encoding="cp932", delimiter=",")
    save_layout(client, import_id)
    assert save_columns(client, import_id, columns_payload("1回目")).status_code == 200
    wait_import_job(app, import_id)
    with app.app_context():
        first = get_import(import_id)
    assert save_columns(client, import_id, columns_payload("2回目")).status_code == 200
    wait_import_job(app, import_id)
    with app.app_context():
        second = get_import(import_id)
    assert first["spec"].name == "1回目" and second["spec"].name == "2回目"
    assert second["spec_hash"] != first["spec_hash"]
    assert second["template_version_id"] == first["template_version_id"] == import_id


def test_the_table_name_is_required(app, client):
    import_id = upload_csv(client, "a.csv")
    save_source(client, import_id, encoding="cp932", delimiter=",")
    save_layout(client, import_id)
    res = save_columns(client, import_id, columns_payload(""))
    assert res.status_code == 400 and res.get_json()["error"] == "表の名前を入力してください"


def test_there_is_no_saved_settings_endpoint(client):
    """取り込み設定の経路（一覧・削除）は残していない。"""
    assert client.post("/tables/templates/1/delete").status_code == 404


# ---- 帳票の種類: 状態・削除、帳票の JSON と元ファイル ----------------------------------------------

def _make_form_type(app, client, sample_dir) -> int:
    from tests.test_forms import add_field, create_type

    pattern_id = create_type(client, sample_dir / "standard.xlsx", "設備修理報告書")
    add_field(client, pattern_id, "修理報告書", "A3", "B3")
    add_field(client, pattern_id, "修理報告書", "A4", "B4")
    return pattern_id


def _upload_files(tmp_path) -> set[str]:
    return {p.name for p in (tmp_path / "uploads").rglob("*") if p.is_file()}


def test_form_type_status_and_delete(app, client, sample_dir, tmp_path):
    """使用開始・停止・削除。置いた Excel はどこにも残らない（2026-09-21 の利用者の指示）。"""
    from tests.test_forms import book_part

    pattern_id = _make_form_type(app, client, sample_dir)
    other_id = _make_form_type(app, client, sample_dir)

    # 種類を作っても、セルをクリックしても、Excel はサーバーに残らない
    assert _upload_files(tmp_path) == set()
    with app.app_context():
        tables = {row[0] for row in db.get_db().execute(
            "SELECT name FROM sqlite_master WHERE type = 'table'")}
    assert "pattern_samples" not in tables

    res = client.post(f"/form-types/{pattern_id}/status", json={"status": "active"})
    assert res.status_code == 200 and "使用を開始しました" in res.get_json()["message"]
    assert "見本" not in res.get_json()["message"]
    with app.app_context():
        assert db.load_pattern(pattern_id).status == "active"
    assert client.post(f"/form-types/{pattern_id}/status", json={"status": "zzz"}).status_code == 400
    assert client.post("/form-types/9999/status", json={"status": "active"}).status_code == 404

    # Excel を置き直すとシートが出る（種類は増えない）。Excel でないファイルは断る
    res = client.post(f"/form-types/{pattern_id}/panel",
                      data={"book": book_part(sample_dir / "shifted.xlsx")},
                      content_type="multipart/form-data")
    assert res.status_code == 200 and 'data-cell="A1"' in res.get_json()["html"]
    assert _upload_files(tmp_path) == set()
    res = client.post(f"/form-types/{pattern_id}/panel",
                      data={"book": (io.BytesIO(b"not excel"), "x.xlsx")},
                      content_type="multipart/form-data")
    assert res.status_code == 400 and "Excelファイル" in res.get_json()["error"]
    assert _upload_files(tmp_path) == set()
    with app.app_context():
        assert len(db.list_patterns()) == 2      # 置き直しで種類は増えない

    res = client.post(f"/form-types/{pattern_id}/status", json={"status": "inactive"})
    assert "使用を停止しました" in res.get_json()["message"]
    with app.app_context():
        assert db.load_pattern(pattern_id).status == "inactive"

    res = client.post(f"/form-types/{pattern_id}/delete")
    assert "を削除しました" in res.get_json()["message"]
    with app.app_context():
        assert db.load_pattern(pattern_id) is None
        assert db.load_pattern(other_id) is not None
    assert client.post(f"/form-types/{pattern_id}/delete").status_code == 404


def test_the_original_file_can_be_downloaded_without_losing_the_form(app, client, sample_dir):
    """元のファイルのダウンロードでは何も消さない（消えるのは Markdown を渡したときだけ）。"""
    from tests.test_forms import activate, read_form, upload_forms

    pattern_id = _make_form_type(app, client, sample_dir)
    activate(client, pattern_id)
    doc_id, = upload_forms(client, sample_dir / "standard.xlsx")
    read_form(client, doc_id, pattern_id, ["修理報告書"])

    res = client.get(f"/forms/{doc_id}/original")
    assert res.status_code == 200 and res.data == (sample_dir / "standard.xlsx").read_bytes()
    assert "standard.xlsx" in res.headers["Content-Disposition"]
    with app.app_context():
        assert db.get_document(doc_id) is not None
    assert client.get("/forms/9999/original").status_code == 404
    # 確定していない帳票の Markdown は渡さない
    assert client.get(f"/forms/{doc_id}/download.md").status_code == 404


# ---- 一覧表: Excel のシートの選択 ---------------------------------------------------------------

def test_save_source_uses_the_chosen_excel_sheet(app, client, tmp_path):
    book = openpyxl.Workbook()
    memo = book.active
    memo.title = "メモ"
    memo["A1"] = "このブックの説明"
    sheet = book.create_sheet("一覧")
    sheet.append(["管理No", "発生日", "現象"])
    for i in range(1, 6):
        sheet.append([f"TR-{i:03d}", f"2026-08-{i:02d}", f"アラーム停止{i}"])
    path = tmp_path / "二つのシート.xlsx"
    book.save(path)

    import_id = upload_bytes(client, path.read_bytes(), path.name)
    page = panel_html(client, import_id, "source")
    assert "メモ" in page and "一覧" in page

    res = save_source(client, import_id, sheet="一覧")
    assert res.status_code == 200 and res.get_json()["next"] == "layout"
    with app.app_context():
        assert get_import(import_id)["source"]["sheet"] == "一覧"
    page = panel_html(client, import_id, "layout")
    assert "TR-001" in page and "このブックの説明" not in page
    detect = client.post(f"/tables/imports/{import_id}/layout/detect", json={"header_rows": [1]}).get_json()
    assert detect["data_end"] == 6
    res = save_layout(client, import_id)
    assert res.status_code == 200 and res.get_json()["next"] == "columns"
    assert "発生日" in panel_html(client, import_id, "columns")

    # シートを選び直すと、前のシートで決めた見出し行・範囲は使わない
    save_source(client, import_id, sheet="メモ")
    with app.app_context():
        src = get_import(import_id)["source"]
    assert src["sheet"] == "メモ" and "header_rows" not in src
    assert "このブックの説明" in panel_html(client, import_id, "layout")


def test_the_section_of_a_clicked_cell_is_saved_with_the_field(app, client, tmp_path):
    """回答欄の中のセルをクリックした項目は「探す区画」を持つ（消えると回答側の欄を読めない）。"""
    from openpyxl.styles import Font, PatternFill

    from tests.test_forms import add_field, create_type

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "報告書"
    for coord, text in (("A1", "■ 発行側"), ("A4", "▼ 回答欄")):
        ws[coord] = text
        ws[coord].font = Font(bold=True)
        ws[coord].fill = PatternFill("solid", fgColor="DDDDDD")
    for coord, text in (("A2", "処置内容"), ("A5", "処置内容")):
        ws[coord] = text
        ws[coord].fill = PatternFill("solid", fgColor="DDDDDD")
    ws["B2"], ws["B5"] = "発行側の処置", "回答側の処置"
    path = tmp_path / "回答欄.xlsx"
    wb.save(path)

    pattern_id = create_type(client, path, "工程異常連絡書")
    add_field(client, pattern_id, "報告書", "A5", "B5")
    with app.app_context():
        field = db.load_pattern(pattern_id).fields[0]
    assert field.section == "回答"
    # 同じ帳票の Excel を置き直しても区画は残り、読み取りテストは回答側の値を出す
    # （シートには両方の値が写っているので、Markdown の中身で確かめる）
    from tests.test_forms import panel_html

    panel = panel_html(client, pattern_id, book=path)
    markdown = panel[panel.index("md-preview"):]
    assert "## 処置内容\n回答側の処置" in markdown and "発行側の処置" not in markdown
    with app.app_context():
        assert db.load_pattern(pattern_id).fields[0].section == "回答"


# ====================================================================================================
# 元 tests/test_logproc.py
# logproc（「経過の記録」の列のルール処理）の単体テスト。
#
# 実例セルは docs/対応内容セル分析.json の realistic_examples を使う。
# ====================================================================================================

CELLS_JSON = Path(__file__).resolve().parents[1] / "docs" / "対応内容セル分析.json"


@pytest.fixture(scope="module")
def cells() -> list[str]:
    if not CELLS_JSON.exists():
        pytest.skip("対応内容セル分析.json がありません")
    return json.loads(CELLS_JSON.read_text(encoding="utf-8"))["realistic_examples"]


@pytest.fixture(scope="module")
def people() -> PeopleIndex:
    return PeopleIndex(
        [
            {"name": "高橋 圭太", "aliases": ["K.T", "TK"], "org": "保全"},
            {"name": "佐藤 誠", "aliases": ["佐藤(保)"], "org": "保全"},
            {"name": "佐藤 美咲", "aliases": ["佐藤(製)"], "org": "製造"},
        ],
        column_names=["佐藤", "田中/山本"],
    )


def _dates(parse):
    return [s.when.date if s.when else None for s in parse.segments]


def _authors(parse):
    return [(s.author.name if s.author else None) for s in parse.segments]


# ---- 実例セル全体 ----

def test_all_realistic_cells_parse_deterministically(cells, people):
    for text in cells:
        a = parse_log(text, date(2024, 4, 1), people)
        b = parse_log(text, date(2024, 4, 1), people)
        assert a == b
        assert a.kind in ("log", "single", "header_cell")
        ids = [s.id for s in a.segments]
        assert ids == [f"s{i}" for i in range(1, len(ids) + 1)]
        for s in a.segments:
            # 位置は分割に使ったテキスト上で原文と一致し、重ならない
            assert a.text[s.start:s.end] == s.raw
            assert s.raw.strip() == s.raw
        for x, y in zip(a.segments, a.segments[1:]):
            assert x.end <= y.start
        render_timeline(a, "設備")
        review_notes(a)


def test_representative_cell_412(cells, people):
    p = parse_log(cells[0], date(2024, 4, 1), people)
    assert p.kind == "log" and p.order == "asc"
    assert len(p.segments) == 6
    assert _dates(p) == ["2024-04-01", "2024-04-01", "2024-04-02", "2024-04-03", "2024-04-08", "2024-04-19"]
    s1, s2, s3, s4, s5, s6 = p.segments
    assert (s1.when.time, s1.author.name) == ("10:00", "田中")
    assert s1.body.startswith("ライン停止の連絡あり")
    assert s1.identifiers == ["ALM-2031"]
    # 時刻だけの行：日付は直前から、記入者は（推定）
    assert s2.when.time == "10:20" and s2.when.how == "inherited" and not s2.when.estimated
    assert s2.author.name == "田中" and s2.author.estimated
    # 【4/2 夜勤】と行末の（K.T）＝人物一覧の別名
    assert s3.when.shift == "夜勤"
    assert s3.author.name == "高橋 圭太" and s3.author.raw == "K.T"
    assert s3.body == "同エラー再発2回、都度リセット"
    assert s3.quantities == ["2回"]
    # 同姓2人は特定しない
    assert s4.author.name == "佐藤" and "2人" in s4.author.note
    assert s4.identifiers == ["RB-ENC-05M"]
    assert "×1" in s4.quantities and s4.plans == ["納期1週間"]
    # 翌週は範囲＋推定
    assert (s5.when.date, s5.when.date_to, s5.when.estimated) == ("2024-04-08", "2024-04-14", True)
    assert s6.body == "再発なし→クローズ"

    lines = render_timeline(p, "搬送ロボット2号機", types={"s1": ["連絡", "初動"]})
    assert lines[0] == ("1. 2024-04-01 10:00［連絡・初動］田中｜搬送ロボット2号機: "
                        "ライン停止の連絡あり（製造 大野さん）。現場確認、搬送ロボのアーム原点復帰エラー（ALM-2031）。")
    assert lines[1] == "2. 2024-04-01 10:20 田中（推定）｜搬送ロボット2号機: 原点復帰→再起動で復旧。様子見。"
    assert lines[2].startswith("3. 2024-04-02 夜勤 高橋 圭太（K.T）｜搬送ロボット2号機: ")
    assert lines[4].startswith("5. 2024-04-08〜2024-04-14（原文「翌週」、推定） 佐藤｜")
    notes = review_notes(p)
    assert "5の日付は原文「翌週」から推定した（基準は直前の記録の2024-04-03）。" in notes
    assert "2の記入者は原文になく、直前の田中を引き継いだ。" in notes
    assert any("佐藤 誠／佐藤 美咲" in n for n in notes)


def test_glossary_in_render(cells, people):
    p = parse_log(cells[0], date(2024, 4, 1), people)
    lines = render_timeline(p, "搬送ロボット2号機", glossary={"様子見": "経過観察"})
    assert lines[1].endswith("原点復帰→再起動で復旧。経過観察。")


def test_slash_joined_single_line(cells, people):
    p = parse_log(cells[1], None, people)
    assert len(p.segments) == 4
    assert _dates(p) == ["2024-04-01", "2024-04-01", "2024-04-02", "2024-04-08"]
    assert _authors(p) == ["山本", "山本", "鈴木", "山本"]
    s3 = p.segments[2]
    assert s3.body == "部品手配 6205ZZ ×2 納期4/8"
    assert s3.identifiers == ["6205ZZ"] and s3.plans == ["納期4/8"]


def test_wareki_and_group_author(cells, people):
    p = parse_log(cells[2], None, people)
    assert _dates(p) == ["2024-04-01", "2024-04-03", "2024-04-10"]
    assert p.segments[0].author.name == "高橋 圭太"
    assert p.segments[0].quantities == ["0.35MPa", "0.22MPa", "0.33MPa"]
    assert p.segments[2].author.estimated


def test_bracket_shift_cells(cells, people):
    p = parse_log(cells[3], date(2024, 4, 2), people)
    assert [s.when.shift for s in p.segments] == ["夜勤", "日勤", "夜勤", "日勤", None]
    assert _dates(p)[-1] == "2024-04-11"
    assert p.segments[0].author is None
    assert "5〜6回" in p.segments[0].quantities


def test_relative_days(cells, people):
    p = parse_log(cells[4], date(2024, 4, 1), people)
    assert p.order == "asc"
    w = [s.when for s in p.segments]
    assert (w[0].date, w[0].shift) == ("2024-04-01", "夜")
    assert (w[1].date, w[1].estimated) == ("2024-04-02", True)
    assert (w[2].date, w[2].time, w[2].estimated) == ("2024-04-02", "15:00", False)   # 同日は推定にしない
    assert (w[3].date, w[3].estimated) == ("2024-04-05", True)
    assert p.segments[0].body.startswith("夜勤班長より連絡")   # 「夜勤班長」は勤務帯にしない


def test_handoff_arrow(cells, people):
    p = parse_log(cells[5], date(2024, 4, 1), people)
    assert p.kind == "single"
    a = p.segments[0].author
    assert a.name == "田中" and a.raw == "田中→佐藤"
    assert p.segments[0].when.how == "base" and p.segments[0].when.estimated


def test_email_block_is_one_segment(cells, people):
    p = parse_log(cells[6], date(2024, 4, 8), people)
    assert len(p.segments) == 3
    email = p.segments[1]
    assert "email" in email.marks
    assert email.raw.startswith("-----Original Message-----") and email.raw.endswith("-----")
    assert "4/10に弊社FEが伺います" in email.raw
    assert _dates(p) == ["2024-04-08", "2024-04-08", "2024-04-09"]
    assert any("メール" in w for w in p.warnings)


def test_no_anchor_paragraph_sentence_split(cells, people):
    text = cells[7]
    p = parse_log(text, date(2024, 4, 1), people)
    assert p.kind == "single"   # 120字以下は分けない
    p2 = parse_log(text, date(2024, 4, 1), people, SplitOptions(sentence_split_min_chars=40))
    assert len(p2.segments) >= 4
    assert all("sentence_split" in s.marks for s in p2.segments)
    assert "".join(s.raw for s in p2.segments) == text.replace("\n", "")
    assert p2.segments[0].raw.startswith("朝一で立ち上げ時にエラー。リセットで復帰。")   # 短い文は前につなぐ
    long_text = text + "工事は来月の連休に実施予定で、それまでは電源電圧を毎日測定して記録する運用とした。"
    p3 = parse_log(long_text, date(2024, 4, 1), people)
    assert len(p3.segments) > 1 and "工事は来月の連休に実施予定" in p3.segments[-1].raw


def test_desc_order_same_year(cells, people):
    p = parse_log(cells[8], date(2024, 5, 9), people)
    assert p.order == "desc"
    assert _dates(p) == ["2024-05-20", "2024-05-13", "2024-05-12", "2024-05-10", "2024-05-09"]
    lines = render_timeline(p, "クランプ")
    assert lines[0].startswith("1. 2024-05-09 渡辺｜クランプ: クランプ動作遅い")
    assert lines[-1].startswith("5. 2024-05-20 木村｜クランプ: クローズ")
    assert p.segments[2].identifiers == ["CDQ2B32-50D"]
    assert p.segments[3].plans == ["納期5/12"]


def test_desc_order_across_year_end(people):
    text = "1/7 木村 漏れなし確認 完了\n1/6 木村 オイルシール交換\n12/29 木村 応急でシール剤塗布\n12/28 木村 油漏れ発見"
    p = parse_log(text, date(2024, 12, 28), people)
    assert p.order == "desc"
    assert _dates(p) == ["2025-01-07", "2025-01-06", "2024-12-29", "2024-12-28"]
    assert render_timeline(p, "")[0].startswith("1. 2024-12-28 木村: 油漏れ発見")


def test_asc_across_year_end(cells, people):
    p = parse_log(cells[20], date(2024, 12, 28), people)
    assert p.order == "asc"
    assert _dates(p) == ["2024-12-28", "2024-12-29", "2025-01-06", "2025-01-07"]


def test_desc_is_not_pushed_to_next_year(people):
    # 新しい順の 5/20→5/9 を翌年にしない（発生日なし・セル内に年あり）
    text = "5/20 木村 クローズ\n5/13 木村 異常なし\n2024/5/9 渡辺 クランプ動作遅い"
    p = parse_log(text, None, people)
    assert p.order == "desc"
    assert _dates(p) == ["2024-05-20", "2024-05-13", "2024-05-09"]


def test_initials_unregistered_and_registered(cells, people):
    p = parse_log(cells[9], date(2024, 4, 15), people)
    assert _authors(p) == ["高橋 圭太", "高橋 圭太", None, None]
    assert p.segments[2].author.raw == "M.S"
    assert p.segments[0].quantities == ["5.0E-3Pa", "1.0E-4Pa"]
    assert any("M.S" in n for n in review_notes(p))


def test_bullets_and_trailing_paren_author(cells, people):
    p = parse_log(cells[10], date(2024, 4, 3), people)
    assert len(p.segments) == 4
    assert all("bullet" in s.marks for s in p.segments)
    assert _authors(p) == ["小林", "小林", "小林", "加藤"]
    assert p.segments[0].body == "設備停止、非常停止ボタン押下されたまま"


def test_no_base_date_year_unknown(cells, people):
    p = parse_log(cells[11], None, people)
    assert _dates(p) == [None, None, None, None]
    assert any("年を決められない" in w for w in p.warnings)
    assert render_timeline(p, "")[0].startswith("1. 6/3（年不明） 松本: ")
    assert p.segments[2].quantities == ["2.2kW", "4P"]


def test_note_line_joins_previous(cells, people):
    p = parse_log(cells[14], date(2024, 10, 2), people)
    assert len(p.segments) == 3
    last = p.segments[-1]
    assert "note" in last.marks and last.body.endswith("※暫定で手動運転にて生産継続中")
    assert "¥385,000" in p.segments[1].quantities and "納期6週間" in p.segments[1].plans


def test_packed_times_split_on_fullwidth_space(cells, people):
    p = parse_log(cells[15], date(2024, 4, 1), people)
    assert [s.when.time for s in p.segments] == ["10:05", "10:15", "10:40", "11:00"]
    assert [s.body for s in p.segments] == ["停止", "現場着", "ﾁｪｰﾝ外れ　復旧", "生産再開"]


def test_short_single(cells, people):
    p = parse_log(cells[16], date(2024, 4, 1), people)
    assert p.kind == "single" and p.segments[0].body == "再起動で復旧"


def test_group_authors_and_identifiers(cells, people):
    p = parse_log(cells[17], date(2024, 4, 22), people)
    assert _authors(p) == ["保全G", "保全G", "メーカーFE", "保全G"]
    assert p.segments[0].identifiers == ["OC2"]
    assert p.segments[2].identifiers == ["FR-A840-7.5K"]
    assert p.segments[3].quantities == ["3時間"]


def test_arrow_chain_split(cells, people):
    p = parse_log(cells[18], date(2024, 4, 1), people)
    assert [s.body for s in p.segments] == ["停止→リセット復帰", "再発→センサ位置ずれ→調整", "以降OK"]
    assert _dates(p) == ["2024-04-01", "2024-04-02", "2024-04-03"]


def test_time_shift_and_lot(cells, people):
    p = parse_log(cells[19], date(2024, 11, 12), people)
    s1 = p.segments[0]
    assert (s1.when.date, s1.when.time, s1.when.shift, s1.author.name) == ("2024-11-12", "21:30", "夜勤", "渡部")
    assert s1.identifiers == ["L2411-0123"]
    assert "50枚" in s1.quantities
    assert p.segments[1].author.raw == "品証 小川" and p.segments[1].author.name == "小川"


def test_mmdd_and_m_dot_d_at_line_start(cells, people):
    p = parse_log(cells[21], None, people)   # 基準はセル内の 2024-04-04
    assert _dates(p) == ["2024-04-01", "2024-04-02", "2024-04-03", "2024-04-04"]
    assert p.segments[2].identifiers == ["D4N-4120"]
    assert all(s.author is None for s in p.segments)   # 「調査」は記入者にしない


def test_correction_mark(cells, people):
    p = parse_log(cells[22], date(2024, 5, 7), people)
    assert "correction" in p.segments[1].marks


def test_meeting_bullets_and_plan(cells, people):
    p = parse_log(cells[23], date(2024, 6, 10), people)
    assert len(p.segments) == 6
    assert _dates(p)[:3] == ["2024-06-10", "2024-06-10", "2024-06-10"]
    assert p.segments[1].author is None   # 「暫定：」は人名ではない
    assert p.segments[2].plans == ["6月末目処"]
    assert _authors(p)[3:] == ["佐々木", "山田", "村上"]


def test_same_day_and_part_of_day_lines(cells, people):
    p = parse_log(cells[24], date(2024, 8, 20), people)
    assert len(p.segments) == 4
    w = [s.when for s in p.segments]
    assert (w[0].date, w[0].shift) == ("2024-08-20", "午前中")
    assert (w[1].date, w[1].time) == ("2024-08-20", "15:00")
    assert (w[2].date, w[2].shift, w[2].how) == ("2024-08-20", "夕方", "inherited")
    assert p.segments[1].identifiers == ["F3"] and p.segments[1].quantities == ["2A"]


def test_header_cell(cells, people):
    p = parse_log(cells[25], date(2024, 4, 1), people)
    assert p.kind == "header_cell"
    assert [s.label for s in p.segments] == ["現象", "原因", "対応", "再発防止"]
    assert all("header_cell" in s.marks and s.when is None for s in p.segments)
    assert render_timeline(p, "設備") == ["現象: 起動しない", "原因: ブレーカーOFF（清掃時に誤ってOFF）", "対応: ON復帰",
                                          "再発防止: ブレーカーにカバー取付 4/12 済"]
    off = parse_log(cells[25], date(2024, 4, 1), people, SplitOptions(header_cells="off"))
    assert off.kind == "log"


def test_checklist_with_date_in_paren(cells, people):
    p = parse_log(cells[26], date(2024, 4, 1), people)
    assert len(p.segments) == 3
    assert all("checklist" in s.marks for s in p.segments)
    assert _dates(p)[:2] == ["2024-04-03", "2024-04-03"]
    assert p.segments[0].author.name == "西村" and p.segments[0].body == "冷却ファン交換　済"
    assert p.segments[2].raw.endswith("→③は予算化待ち")


def test_reference_and_plans(cells, people):
    p = parse_log(cells[27], date(2024, 4, 16), people)
    assert "明日来場予定" in p.segments[0].plans
    assert "reference" in p.segments[1].marks
    p32 = parse_log(cells[32], date(2024, 4, 1), people)
    assert "reference" in p32.segments[0].marks and p32.segments[0].identifiers == ["TR-24-00398"]


def test_x000d_and_romaji(cells, people):
    p = parse_log(cells[28], date(2024, 4, 1), people)
    assert "_x000D_" not in p.text
    assert p.segments[0].author.name == "Tanaka" and p.segments[0].body == "E-stop tripped. Reset OK."
    assert p.segments[1].author.name == "田中"


def test_time_range_and_alias_with_paren(cells, people):
    p = parse_log(cells[29], date(2024, 9, 3), people)
    assert p.segments[0].when.time == "08:15"
    assert "85分" in p.segments[0].quantities
    p30 = parse_log(cells[30], date(2024, 3, 4), people)
    assert _authors(p30) == ["佐藤 誠", "佐藤 美咲", "佐藤 誠"]


def test_fullwidth_digits(cells, people):
    p = parse_log(cells[31], None, people)
    assert [(s.when.date, s.when.time) for s in p.segments] == [("2024-04-01", "10:00"), ("2024-04-01", "11:30")]
    assert _authors(p) == ["田中", "田中"]


def test_weekday_and_note_after_date(cells, people):
    p = parse_log(cells[33], date(2024, 4, 5), people)
    assert _dates(p) == ["2024-04-05", "2024-04-08", "2024-04-08"]
    assert "週明け確認予定" in p.segments[0].plans


def test_era_change(cells, people):
    p = parse_log(cells[34], date(2019, 4, 26), people)
    assert _dates(p) == ["2019-04-26", "2019-05-07", "2019-05-08"]
    assert p.segments[0].author.name == "小西"
    assert p.segments[0].body == "連休前に油圧作動油交換（46番 200L）"
    assert p.segments[1].body.startswith("連休明け立上げ時")   # 日付の後の「連休明け」は相対日にしない


# ---- 日付の個別ケース ----

@pytest.mark.parametrize("text,expected", [
    ("令和6年4月1日", (2024, 4, 1)), ("R6/4/1", (2024, 4, 1)), ("H31.4.26", (2019, 4, 26)), ("R元.5.1", (2019, 5, 1)),
    ("平成30年12月1日", (2018, 12, 1)), ("2024-04-01", (2024, 4, 1)), ("2024年4月1日", (2024, 4, 1)), ("24/4/1", (2024, 4, 1)),
])
def test_full_dates(text, expected):
    w = parse_when_at(shadow(text), 0)
    assert (w.year, w.month, w.day, w.full) == (*expected, True)


@pytest.mark.parametrize("text", ["1.5mm 隙間あり", "0.35MPa", "納期4/8 手配", "NG 2/50", "4/32 不明", "1時間停止"])
def test_not_event_dates(text):
    w = parse_when_at(shadow(text), 0, not_date_res=[__import__("re").compile(p) for p in SplitOptions().not_date_patterns])
    assert w is None or not w.has_date


def test_plan_date_in_body_is_not_split(people):
    p = parse_log("4/2 鈴木 部品手配 納期4/8\n4/9 鈴木 交換済", date(2024, 4, 2), people)
    assert len(p.segments) == 2 and _dates(p) == ["2024-04-02", "2024-04-09"]


def test_next_week_range_from_friday():
    heads = [parse_when_at(shadow("4/5(金)"), 0), parse_when_at(shadow("翌週"), 0), parse_when_at(shadow("週明け"), 0)]
    whens, order, _ = resolve_whens(heads, ["4/5(金)", "翌週", "週明け"], date(2024, 4, 5))
    assert (whens[1].date, whens[1].date_to, whens[1].estimated) == ("2024-04-08", "2024-04-14", True)
    assert whens[2].date == "2024-04-15"   # 週明けは直前（翌週の開始日）からの次の月曜
    assert order == "asc"


def test_unresolvable_relative(people):
    p = parse_log("昨日 夜勤者より連絡あり\n4/3 田中 確認", date(2024, 4, 3), people)
    assert p.segments[0].when.date is None and p.segments[0].when.how == "unresolved"
    assert any("昨日" in w for w in p.warnings)


def test_order_anomaly_within_tolerance(people):
    p = parse_log("4/1 田中 停止\n4/5 田中 交換\n4/3 田中 追記：部品到着", date(2024, 4, 1), people)
    assert _dates(p) == ["2024-04-01", "2024-04-05", "2024-04-03"]
    assert any("前後" in w for w in p.warnings)


def test_time_only_lines_join_option(people):
    text = "4/1 10:00 田中：停止\n10:20 再起動で復旧"
    assert len(parse_log(text, date(2024, 4, 1), people).segments) == 2
    joined = parse_log(text, date(2024, 4, 1), people, SplitOptions(time_only_lines="join"))
    assert len(joined.segments) == 1


def test_numbered_lines(people):
    p = parse_log("1. マッチャー取外し\n2. （継続対応中）", date(2024, 4, 1), people)
    assert [s.body for s in p.segments] == ["マッチャー取外し", "（継続対応中）"]
    assert all("checklist" in s.marks for s in p.segments)


def test_extra_anchor(people):
    text = "4/1 田中 停止\n◎追記 部品到着"
    assert len(parse_log(text, date(2024, 4, 1), people).segments) == 1
    p = parse_log(text, date(2024, 4, 1), people, SplitOptions(extra_anchors=[r"◎"]))
    assert len(p.segments) == 2


def test_empty_cells(people):
    for t in ("", None, "-", "－", "  ", "なし"):
        p = parse_log(t, date(2024, 4, 1), people)
        assert p.kind == "empty" and p.segments == []
        assert render_timeline(p, "設備") == []


def test_split_options_from_dict():
    o = SplitOptions.from_dict({"order": "desc", "sentence_split": {"min_chars": 80}, "header_cells": "off",
                                "extra_anchors": ["^◎"], "not_date_patterns": []})
    assert (o.order, o.sentence_split_min_chars, o.header_cells, o.extra_anchors, o.not_date_patterns) == \
        ("desc", 80, "off", ["^◎"], [])


def test_forced_order_option(people):
    p = parse_log("4/2 田中 交換\n4/1 田中 停止", date(2024, 4, 1), people, SplitOptions(order="asc"))
    assert p.order == "asc"


# ---- 記入者 ----

def test_people_index_resolution(people):
    assert people.resolve("K.T").name == "高橋 圭太"
    assert people.resolve("高橋").name == "高橋 圭太"
    amb = people.resolve("佐藤")
    assert amb.name == "佐藤" and "佐藤 誠／佐藤 美咲" in amb.note
    assert people.resolve("佐藤(製)").name == "佐藤 美咲"
    assert people.resolve("保全G").name == "保全G"
    assert people.is_known("山本")   # 担当列の「田中/山本」を分けて登録


def test_known_name_without_separator(people):
    p = parse_log("4/3佐藤(保)エンコーダ交換", date(2024, 4, 3), people)
    assert p.segments[0].author.name == "佐藤 誠" and p.segments[0].body == "エンコーダ交換"


def test_unknown_word_is_not_author():
    p = parse_log("4/3 旧品 保管\n4/4 交換 完了（在庫あり）", date(2024, 4, 3), PeopleIndex())
    assert all(s.author is None for s in p.segments)


# ---- 識別子・数量・予定句 ----

def test_identifiers_exclude_dates_units_and_masks():
    text = "R6.4.1 ALM-2031 発生。0.35MPa、198V、2A。No.3 AM10:00 ［電話番号］ FE-100 6205ZZ Ver4.1 ＰＬＣ２"
    assert extract_identifiers(text) == ["ALM-2031", "FE-100", "6205ZZ", "Ver4.1", "PLC2"]


def test_quantities_and_plans():
    text = "ケーブル手配（RB-ENC-05M ×1、納期1週間）。再発2回。見積¥385,000。1時間に5〜6回。6月末目処で対策予定"
    q = extract_quantities(text)
    assert q[:2] == ["×1", "1週間"] and "2回" in q and "¥385,000" in q and "5〜6回" in q
    assert "05M" not in "".join(q)
    assert extract_plans(text) == ["納期1週間", "6月末目処で対策予定"]


# ---- マスク ----

@pytest.mark.parametrize("phone", [
    "090-1234-5678", "09012345678", "03-1234-5678", "(03)1234-5678", "（03）1234-5678", "0312345678",
    "0120-123-456", "0800-123-4567", "０９０－１２３４－５６７８", "+81-90-1234-5678", "内線1234", "内線：567",
    "03 1234 5678", "080ー1234ー5678",
])
def test_mask_phone_variants(phone):
    masked, spans = mask_text(f"担当 上田様 携帯{phone} まで", ["電話番号"])
    assert masked == "担当 上田様 携帯［電話番号］ まで"
    assert spans[0].kind == "phone" and spans[0].text == phone


@pytest.mark.parametrize("text", ["2024-04-01 10:00", "0401 停止", "ロット L2411-0123", "D4N-4120", "0.35MPa", "1234-5678"])
def test_mask_does_not_touch_non_phone_numbers(text):
    assert mask_text(text, ["phone"])[0] == text


def test_mask_email_amount_person(people):
    text = "yamaguchi@example.co.jp から回答。見積 ¥385,000、追加 12,000円。佐藤 誠さん・高橋さんに連絡（K.T）"
    masked, spans = mask_text(text)
    assert masked.startswith("［メール］ から回答")
    assert "385,000" in masked   # 既定は電話・メールだけ
    masked2, spans2 = mask_text(text, ["email", "金額", "人名"], names=people.names())
    assert "［金額］" in masked2 and "385,000" not in masked2 and "12,000円" not in masked2
    assert "佐藤" not in masked2 and "高橋" not in masked2 and "K.T" not in masked2
    for s in spans2:
        assert text[s.start:s.end] == s.text


# ---- 用語集 ----

GLOSSARY = {
    "様子見": "経過観察",
    "チョコ停": "短時間停止（チョコ停）",
    "FE": {"to": "メーカーのフィールドエンジニア（FE）", "ascii_boundary": True},
    "TEL済": "電話連絡済み",
    "取替": "交換",
    "交換済": "交換済み",
}


def test_glossary_protects_identifiers():
    text = "FE来場、FE-100 の基板取替。チョコ停は様子見。TEL済"
    out = apply_glossary(text, GLOSSARY)
    assert out == "メーカーのフィールドエンジニア（FE）来場、FE-100 の基板交換。短時間停止（チョコ停）は経過観察。電話連絡済み"


def test_glossary_protected_spans_explicit_and_boundary():
    text = "XFE-100 と FE-100"
    # 保護範囲を渡さなくても ascii_boundary で前後が英数字・「-」なら置き換えない
    assert apply_glossary(text, {"FE": {"to": "エンジニア", "ascii_boundary": True}}, protected_spans=[]) == text
    # ascii_boundary なしでも識別子の範囲は保護される
    assert apply_glossary("FE-100 FE", {"FE": {"to": "E", "ascii_boundary": False}}) == "FE-100 E"
    assert apply_glossary("FE-100", {"FE": {"to": "E", "ascii_boundary": False}}, protected_spans=[]) == "E-100"


def test_glossary_longest_first_and_idempotent():
    g = {"交換": "取り替え", "交換済": "交換済み"}
    assert apply_glossary("交換済、交換", g) == "交換済み、取り替え"
    once = apply_glossary("チョコ停", GLOSSARY)
    assert apply_glossary(once, GLOSSARY) == once
    hits = glossary_hits("ﾁｮｺ停と様子見", GLOSSARY)   # 半角カナも照合
    assert [(h.term, h.to) for h in hits] == [("ﾁｮｺ停", "短時間停止（チョコ停）"), ("様子見", "経過観察")]


def test_glossary_skips_mask_tokens():
    assert apply_glossary("［電話番号］へTEL済", {"電話": "でんわ", "TEL済": "電話連絡済み"}) == "［電話番号］へ電話連絡済み"


# ---- 速さのための控え（結果は変えない） ----

def test_shadow_matches_per_character_nfkc():
    """shadow は表と控えを使うが、1文字ずつ NFKC をかけた写し（長さを保つ・丸数字は残す）と同じ。"""
    import unicodedata

    def reference(text):
        out = []
        for ch in text:
            code = ord(ch)
            if code < 0x80 or 0x2460 <= code <= 0x24FF or 0x2776 <= code <= 0x2793:
                out.append(ch)
                continue
            n = unicodedata.normalize("NFKC", ch)
            out.append(n if len(n) == 1 else ch)
        return "".join(out)

    sample = "ＡＢＣ１２３　ｶﾀｶﾅ①②❶ ㍻ ﬁ ㌔ Ⅻ ¥１,０００ ＠ｘ．ｊｐ 全角／半角→ “引用” \t\n漢字"
    everything = "".join(chr(c) for c in range(0x80, 0x10000) if not 0xD800 <= c <= 0xDFFF)
    for text in (sample, everything, "", "ascii only 123"):
        assert shadow(text) == reference(text) and len(shadow(text)) == len(text)
        assert shadow(text) == reference(text)          # 2回目（控えから）も同じ


def test_parse_head_memo_does_not_leak_between_cells(people):
    """同じ位置・同じ長さの別のセルを続けて分けても、前のセルの読み取り結果を使わない。"""
    a = parse_log("4/1 田中：ライン停止\n4/2 佐藤：復旧確認", date(2024, 4, 1), people)
    b = parse_log("【現象】ラインの停止\n【原因】センサー汚れ", date(2024, 4, 1), people)
    c = parse_log("4/1 田中：ライン停止\n4/2 佐藤：復旧確認", date(2024, 4, 1), people)
    assert a.kind == "log" and b.kind == "header_cell"
    assert [s.when.date for s in a.segments] == ["2024-04-01", "2024-04-02"] and repr(a) == repr(c)


def test_extractors_are_unchanged_by_the_cached_spans():
    text = "RB-ENC-05M ×1 手配、納期1週間。ALM-2031 表示、5mm ずれ、N2パージ"
    ids = extract_identifiers(text)
    ids.append("書き換え")                               # 戻り値を変えても控えは変わらない
    assert extract_identifiers(text) == ["RB-ENC-05M", "ALM-2031"]
    assert "×1" in extract_quantities(text) and "5mm" in extract_quantities(text)
    assert extract_plans(text) == ["納期1週間"] and extract_plans("") == [] and extract_plans(None) == []


# ---- R6-AI-1: 単位の後ろに英数字・「-」・カタカナが続くときは寸法とみなさず、日付として読む ----------------------
@pytest.mark.parametrize("text", ["4.3 AGV 停止", "4.3 ALM-2031 発生", "4.3 Aライン停止", "4.3 Vベルト交換"])
def test_unit_letter_followed_by_word_is_still_a_date(text):
    w = parse_when_at(shadow(text), 0, not_date_res=[__import__("re").compile(p) for p in SplitOptions().not_date_patterns])
    assert w is not None and w.has_date and (w.month, w.day) == (4, 3)


@pytest.mark.parametrize("text", ["2.5A 流れた", "3.3V", "1.2 A", "2.5Aの電流"])
def test_plain_units_are_still_not_dates(text):
    w = parse_when_at(shadow(text), 0, not_date_res=[__import__("re").compile(p) for p in SplitOptions().not_date_patterns])
    assert w is None or not w.has_date
