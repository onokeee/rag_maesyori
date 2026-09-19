"""4巡目の各担当の持ち越し（担当範囲の外の変更）の確認。"""
from __future__ import annotations

import io
import zipfile
from datetime import date

import openpyxl
from openpyxl.drawing.image import Image as XLImage

from excel.image_detector import detect_images
from excel.tables import _is_total
from excel.workbook import Cell
from logproc import PeopleIndex, SplitOptions, parse_log
from tables import store
from tables.csv_source import CsvSource
from tables.source import clean_text


# ---- R4-FUZZ-3（一覧表側）: 見えない文字を消す -------------------------------------------------

def test_clean_text_removes_invisible_characters():
    assert clean_text("EQ-01⁠") == "EQ-01"
    assert clean_text("﻿管理No") == "管理No"
    assert clean_text("A​B­C‮D⁦") == "ABCD"
    assert clean_text("改行は\n残す") == "改行は\n残す"


def test_csv_cells_do_not_keep_invisible_characters(tmp_path):
    path = tmp_path / "t.csv"
    path.write_text("管理No,設備\n1,EQ-01\n2,EQ-01⁠\n﻿管理No,設備\n", encoding="utf-8")
    rows = list(CsvSource(path).rows())
    assert rows[1].text(1) == rows[2].text(1) == "EQ-01"
    assert rows[3].text(0) == rows[0].text(0) == "管理No"


# ---- S4-1: 深く入れ子のJSONは500ではなく「取り込み設定のJSONではありません」 ----------------------

def test_import_of_deeply_nested_json_is_refused_without_error(client):
    body = b"[" * 200000 + b"]" * 200000
    res = client.post("/settings/table-templates/import", data={"file": (io.BytesIO(body), "deep.json")},
                      content_type="multipart/form-data")
    assert res.status_code == 302
    with client.application.app_context():
        assert store.list_templates() == []


# ---- S4-1: 検証より前に保存された、コンパイルできない正規表現で処理を止めない -----------------------

def test_parse_log_skips_patterns_that_do_not_compile():
    text = "4/1 10:00 停止を確認\n4/2 部品交換"
    good = parse_log(text, date(2024, 4, 1), PeopleIndex(), SplitOptions(extra_anchors=[], not_date_patterns=[]))
    bad = parse_log(text, date(2024, 4, 1), PeopleIndex(),
                    SplitOptions(extra_anchors=["[", "("], not_date_patterns=["(?P<"]))
    assert len(good.segments) == 2
    assert [s.raw for s in bad.segments] == [s.raw for s in good.segments]


# ---- F4-1 の続き: 塗りつぶしの「温度計」行で明細表の読み取りを止めない ------------------------------

def _cell(text: str, filled: bool = True) -> Cell:
    return Cell(row=1, col=1, max_row=1, max_col=1, value=text, text=text, norm=text, inline=None, filled=filled)


def test_filled_meter_names_are_not_total_rows():
    for name in ("温度計", "圧力計", "設計", "膜厚計", "pH計"):
        assert not _is_total(_cell(name)), name
    for name in ("合計", "小計", "部品費計", "工数計", "部品費合計", "計"):
        assert _is_total(_cell(name)), name
    assert _is_total(_cell("合計", filled=False))


# ---- S4-2 の続き: 複数シートが同じ drawing を指しても、シートごとに画像を数える ----------------------

def _png() -> bytes:
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (2, 2), "red").save(buf, format="PNG")
    return buf.getvalue()


def test_detect_images_with_a_drawing_shared_by_two_sheets(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "A"
    wb.create_sheet("B")
    wb["A"].add_image(XLImage(io.BytesIO(_png())), "C3")
    path = tmp_path / "shared.xlsx"
    wb.save(path)
    # シートBのリレーションをシートAと同じ drawing に向ける
    src = zipfile.ZipFile(path)
    out = tmp_path / "shared2.xlsx"
    with zipfile.ZipFile(out, "w") as dst:
        for info in src.infolist():
            dst.writestr(info, src.read(info.filename))
        rels = src.read("xl/worksheets/_rels/sheet1.xml.rels")
        dst.writestr("xl/worksheets/_rels/sheet2.xml.rels", rels)
    src.close()
    found = detect_images(out)
    assert sorted(i["sheet"] for i in found) == ["A", "B"]
    assert all(i["location"].startswith("C3") for i in found)
