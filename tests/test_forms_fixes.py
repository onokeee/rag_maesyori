"""帳票まわりの修正（確認で見つかった不具合）の回帰テスト。"""
from excel.extractor import apply_manual_values, refresh_summary


def _number_field(**kw) -> dict:
    f = {"field_name": "downtime", "display_name": "停止時間", "data_type": "number", "required": False,
         "value": None, "unit": "", "warning": None, "edited": False}
    f.update(kw)
    return f


def test_manual_number_is_judged_by_the_form_type_unit_not_the_read_unit():
    """「停止時間（分）」の欄が「14.9h」と読まれたあと、手で「894分」と直したら単位は分に戻る。

    数字だけの「894」は、画面のラベルに出ている読み取った単位（時間）のまま要確認にする（黙って分に変えない）。
    """
    ex = {"fields": [_number_field(value=14.9, unit="時間", spec_unit="分", warning="この項目の単位は「分」ですが…")]}
    apply_manual_values(ex, {"value-downtime": "894"})
    f = ex["fields"][0]
    assert f["value"] == 894 and f["unit"] == "時間" and f["edited"] and f["warning"].startswith("この項目の単位は")
    apply_manual_values(ex, {"value-downtime": "894分"})
    assert f["value"] == 894 and f["unit"] == "分" and f["edited"] and not f["warning"]


def test_manual_number_with_another_written_unit_keeps_it_and_stays_to_be_checked():
    from views.forms import _field_status

    ex = {"fields": [_number_field(value=30, unit="分", spec_unit="分")]}
    apply_manual_values(ex, {"value-downtime": "2.5h"})
    f = ex["fields"][0]
    assert f["value"] == 2.5 and f["unit"] == "時間" and f["edited"]
    assert f["warning"].startswith("この項目の単位は")
    assert _field_status(f) == {"source": "manual", "issue": True, "blank": False, "missing_required": False}


def test_manual_number_takes_the_written_unit_when_the_form_type_has_none():
    """単位の無い「作業時間」に「150分」「2.5時間」と入れたら、その単位で出す（数値だけにしない）。"""
    from views.forms import _field_status

    ex = {"fields": [_number_field(field_name="work_hours", display_name="作業時間", value=2.5, spec_unit="",
                                   warning="単位が書かれていません（時間か分か）")]}
    apply_manual_values(ex, {"value-work_hours": "2.5時間"})
    f = ex["fields"][0]
    assert f["value"] == 2.5 and f["unit"] == "時間" and f["edited"] and not f["warning"]

    apply_manual_values(ex, {"value-work_hours": "150分"})
    assert f["value"] == 150 and f["unit"] == "分" and not f["warning"]

    # 入力欄には数字だけが出る（単位はラベル）。数字だけ直したら、ラベルに出ていた単位のまま
    apply_manual_values(ex, {"value-work_hours": "160"})
    assert f["value"] == 160 and f["unit"] == "分" and not f["warning"] and not _field_status(f)["issue"]


def test_old_extractions_without_spec_unit_fall_back_to_the_unit():
    ex = {"fields": [_number_field(value=30, unit="分")]}
    apply_manual_values(ex, {"value-downtime": "45"})
    assert ex["fields"][0]["value"] == 45 and ex["fields"][0]["unit"] == "分"


def test_extraction_keeps_the_form_type_unit(tmp_path):
    from openpyxl import Workbook

    from excel.extractor import extract_document
    from excel.workbook import load_workbook_info
    from pattern.model import FieldDef, PatternDef

    wb = Workbook()
    wb.active.title = "報告書"
    wb.active["A1"], wb.active["B1"] = "停止時間", "14.9h"
    path = tmp_path / "u.xlsx"
    wb.save(path)
    pattern = PatternDef(name="テスト", fields=[
        FieldDef("downtime", "停止時間", ["停止時間"], data_type="number", unit="分")])
    ex = extract_document(load_workbook_info(path), pattern, ["報告書"])
    f = ex["fields"][0]
    assert f["unit"] == "時間" and f["spec_unit"] == "分"
    apply_manual_values(ex, {"value-downtime": "894分"})
    assert f["value"] == 894 and f["unit"] == "分"


# ---- 確認画面・まとめ取り込み（views/forms.py） ------------------------------------------------

import io  # noqa: E402
import json  # noqa: E402
import re  # noqa: E402
from pathlib import Path  # noqa: E402

from models import database as db  # noqa: E402


def _extraction(name="CMP研磨装置") -> str:
    field = {"field_name": "equipment_name", "display_name": "設備名", "data_type": "string", "required": False,
             "value": name, "sheet": "報告書", "label_cell": "A1", "value_cell": "B1", "label_found": True,
             "warning": None, "edited": False, "unit": "", "spec_unit": "",
             "rag_output": "show", "table_columns": [], "table_blocks": 1}
    data = {"pattern": {"id": 1, "name": "設備修理報告書", "version": "v1"}, "sheets": ["報告書"],
            "fields": [field], "attachments": [], "values": {"equipment_name": name}, "missing_required": []}
    return json.dumps(data, ensure_ascii=False)


def _confirmed_doc(app, name="1.xlsx", *, batch_id="", order=0) -> int:
    with app.app_context():
        stored = f"documents/{name}"
        path = Path(app.config["UPLOAD_DIR"]) / stored
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(b"dummy-excel")
        doc_id = db.create_document(name, "0" * 64, stored, batch_id=batch_id, batch_order=order)
        db.update_document(doc_id, data_json=_extraction(), confirmed_json=_extraction(), title=f"{name} 確定")
    return doc_id


def _state(app, doc_id):
    with app.app_context():
        return db.get_document(doc_id)["state"]


def _doc_version(app, doc_id) -> str:
    """いまの作業データの版（画面が読み取り結果と一緒に受け取る値）。"""
    from views.forms import _version

    with app.app_context():
        return _version(db.get_document(doc_id))


def _review_html(app, doc_id) -> str:
    """読み取り結果の欄のHTML（読み取りの応答が返すのと同じもの）。"""
    from views.forms import _review_response

    with app.test_request_context():
        return _review_response(doc_id).get_json()["html"]


def test_typing_the_confirmed_value_back_returns_the_form_to_confirmed(app, client):
    doc_id = _confirmed_doc(app)
    client.post(f"/forms/{doc_id}/draft", json={"values": {"equipment_name": "別の名前"}})
    assert _state(app, doc_id) == "modified"
    client.post(f"/forms/{doc_id}/draft", json={"values": {"equipment_name": "CMP研磨装置"}})
    assert _state(app, doc_id) == "confirmed"
    summary = client.post(f"/forms/{doc_id}/preview", json={"values": {}}).get_json()
    assert summary["state"] == "confirmed" and summary["counts"]["manual"] == 0


def test_a_stale_review_page_cannot_save_or_confirm(app, client):
    """別のタブ（開いたままの古い画面）から、見ていない内容を上書き・確定しない。"""
    doc_id = _confirmed_doc(app)
    with app.app_context():
        db.update_document(doc_id, confirmed_json=None)  # 確認中の帳票
    old = _doc_version(app, doc_id)
    assert f'name="version" value="{old}"' in _review_html(app, doc_id)

    # タブAの保存: 新しい版が返り、それを送れば続けて保存できる
    res = client.post(f"/forms/{doc_id}/draft", json={"values": {"equipment_name": "A"}, "version": old})
    assert res.status_code == 204
    new = res.headers["X-Doc-Version"]
    assert new and new != old
    assert client.post(f"/forms/{doc_id}/draft",
                       json={"values": {"equipment_name": "A2"}, "version": new}).status_code == 204

    # タブB（古い版のまま）の保存・プレビュー・確定は止める
    res = client.post(f"/forms/{doc_id}/draft", json={"values": {"equipment_name": "B"}, "version": old})
    assert res.status_code == 409 and "別の画面で内容が変わりました" in res.get_json()["error"]
    assert client.post(f"/forms/{doc_id}/preview", json={"values": {}, "version": old}).status_code == 409
    res = client.post(f"/forms/{doc_id}/confirm", json={"values": {"equipment_name": "B"}, "version": old})
    assert res.status_code == 409 and "別の画面で内容が変わりました" in res.get_json()["error"]
    with app.app_context():
        doc = db.get_document(doc_id)
        assert doc["state"] == "reviewing" and "A2" in doc["data_json"] and '"B"' not in doc["data_json"]


def _finish(client, ids, current=None) -> dict:
    url = "/forms/finish?ids=" + ",".join(str(i) for i in ids) + (f"&current={current}" if current else "")
    return client.get(url).get_json()


def test_a_partly_confirmed_batch_does_not_say_everything_is_deleted(app, client):
    first = _confirmed_doc(app, "1.xlsx", batch_id="B", order=0)
    with app.app_context():
        pending = db.create_document("2.xlsx", "0" * 64, "documents/2.xlsx", batch_id="B", batch_order=1)
    body = _finish(client, [first, pending], current=first)
    assert body["confirmed"] == 1 and body["total"] == 2 and body["next_id"] == pending
    page = body["html"]
    assert "このまとまりの帳票のデータはサーバーからすべて消えます" not in page
    assert "確定済みの1件だけを zip でダウンロードします" in page and "未確定の1件は残ります" in page
    assert "確定済み1件だけをダウンロード（zip）" in page

    # すべて確定すれば、まとめてのダウンロードだけになる
    with app.app_context():
        db.update_document(pending, data_json=_extraction(), confirmed_json=_extraction())
    page = _finish(client, [first, pending], current=first)["html"]
    assert "まとめて Markdown をダウンロード（zip）" in page and "確定済み1件だけ" not in page
    assert "このまとまりの帳票のデータはサーバーからすべて消えます" in page


def test_batch_upload_error_names_the_file_that_was_skipped(app, client, sample_dir):
    files = [(io.BytesIO((sample_dir / "standard.xlsx").read_bytes()), "standard.xlsx"),
             (io.BytesIO(b"this is not an excel file"), "broken.xlsx"),
             (io.BytesIO((sample_dir / "shifted.xlsx").read_bytes()), "shifted.xlsx")]
    res = client.post("/forms/upload", data={"file": files}, content_type="multipart/form-data")
    body = res.get_json()
    # どのファイルかは選んだ順の位置で示す（ファイル名そのものは画面の外に出さない。design.md 3.3）
    assert res.status_code == 200 and len(body["docs"]) == 2 and body["batch_id"]
    assert body["errors"] and body["errors"][0].startswith("2件目のファイル: ")
    assert "broken.xlsx" not in " ".join(body["errors"])


# ---- 時刻の範囲の「作業時間」（excel/text.py） ------------------------------------------------------

def test_time_range_is_not_read_as_its_start_hour():
    from excel.text import numeric_unit, to_number

    # 添えた時間数（工数）を項目の単位で使う
    value, warning = to_number(None, "09:30-12:45（3.2h）", "時間")
    assert value == 3.2 and "時間数" in warning
    assert to_number(None, "09:30-12:45（3.2h）", "分")[0] == 192
    assert to_number(None, "９:３０～１３:５５（工数 ８.７５h）", "時間")[0] == 8.75
    # 時間数が無ければ範囲の長さ（日をまたぐ作業も）
    value, warning = to_number(None, "9:30-12:45", "分")
    assert value == 195 and "時刻の範囲" in warning
    assert to_number(None, "22:00～02:30", "時間")[0] == 4.5
    # 単位の決まっていない項目は分にする（「3時間40分」と同じ）
    assert to_number(None, "09:30-12:45（3.2h）")[0] == 192 and numeric_unit("09:30-12:45（3.2h）") == "分"
    # 時間でない単位の項目では読まない（開始時刻の 9 を返さない）
    assert to_number(None, "9:30-12:45", "円") == ("9:30-12:45", "時刻の範囲です。数値として読み取れません")


# ---- 設備だけのタイトル（export/formats.py） ----------------------------------------------------

def _md_field(name, display, value, data_type="string"):
    return {"field_name": name, "display_name": display, "data_type": data_type, "required": False, "value": value,
            "unit": "", "rag_output": "show", "edited": False}


def _inspection(work_no, work_date, finding="端子の緩みを増し締めした。" * 200):
    fields = [_md_field("equipment_id", "設備番号", "IMP-603"), _md_field("equipment_name", "設備名", "高電流注入 3号機"),
              _md_field("field_13", "作業日", work_date, "date"), _md_field("field_6", "作業No.", work_no),
              _md_field("finding", "所見", finding, "text")]
    ex = {"pattern": {"id": 4, "name": "点検保全作業報告書", "title_fields": [], "md_options": {}, "labels": []},
          "sheets": ["報告書"], "fields": fields, "attachments": []}
    refresh_summary(ex)
    return ex


def test_reports_for_the_same_equipment_get_different_titles_and_headings():
    from export.formats import build_markdown

    doc = {"id": 1, "file_name": "点検.xlsx", "file_hash": "0" * 64}
    a = build_markdown(doc, _inspection("W-0101", "2026-04-01"))
    b = build_markdown(doc, _inspection("W-0102", "2026-04-08"))
    title_a, title_b = a.split("\n", 1)[0], b.split("\n", 1)[0]
    assert title_a == "# 点検保全作業報告書 高電流注入 3号機（IMP-603）｜W-0101" and title_a != title_b
    heading_a = next(line for line in a.split("\n") if line.startswith("## "))
    heading_b = next(line for line in b.split("\n") if line.startswith("## "))
    assert heading_a == "## 所見（IMP-603 高電流注入 3号機／W-0101）" and heading_a != heading_b

    # 番号らしい項目が無ければ日付、それも無ければ元ファイル名
    ex = _inspection("", "2026-04-08")
    assert build_markdown(doc, ex).startswith("# 点検保全作業報告書 高電流注入 3号機（IMP-603）｜2026-04-08\n")
    ex = _inspection("", None)
    assert build_markdown(doc, ex).startswith("# 点検保全作業報告書 高電流注入 3号機（IMP-603）｜点検\n")


# ---- 押印欄の人名を見出しにした項目（pattern/builder.py・export/formats.py） ----------------------------

def test_a_person_name_used_as_a_label_is_not_written(tmp_path):
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill

    from excel.workbook import load_workbook_info
    from export.formats import build_markdown
    from pattern.builder import suggest_rows

    # 押印欄: 「作成｜確認」の下に人名。builder は「斎藤」を見出しの候補にしない
    wb = Workbook()
    ws = wb.active
    ws.title = "報告書"
    ws["A1"], ws["B1"] = "件名", "搬送エラー"
    ws["D1"], ws["E1"] = "作成", "確認"
    ws["D2"], ws["E2"] = "斎藤", "森"
    for coord in ("A1", "D1", "E1"):  # 見出しの欄は色付き（帳票の普通の形）
        ws[coord].fill = PatternFill("solid", fgColor="DDEBF7")
    path = tmp_path / "stamp.xlsx"
    wb.save(path)
    _, rows = suggest_rows([load_workbook_info(path)])
    assert "斎藤" not in {r["display_name"] for r in rows}

    # すでにそう作られた帳票の種類（人名が見出しになったもの）は、読み取ったとおりに出す
    fields = [_md_field("subject", "件名", "搬送エラー"), _md_field("field_60", "作成", "斎藤"),
              _md_field("field_29", "斎藤", "森"), _md_field("field_61", "確認", "森")]
    ex = {"pattern": {"id": 2, "name": "トラブル報告書", "title_fields": [], "md_options": {}, "labels": []},
          "sheets": ["報告書"], "fields": fields, "attachments": []}
    refresh_summary(ex)
    md = build_markdown({"id": 1, "file_name": "t.xlsx", "file_hash": "0" * 64}, ex)
    assert "- 件名: 搬送エラー" in md and "- 作成: 斎藤" in md and "- 斎藤: 森" in md


def test_upload_page_shows_the_batch_limits(client):
    page = client.get("/forms/new").get_data(as_text=True)
    assert "50ファイル・合計200MBまで" in page


# ---- 2回目の修正（帳票） -----------------------------------------------------------------------

def _read_standard(app, client, sample_dir) -> tuple[int, int]:
    """standard.xlsx を取り込んで読み取った帳票（確認中）を作る。戻り値: (帳票ID, 種類ID)"""
    from tests.test_forms_flow import activate, add_field, create_type, read_form, upload_forms

    path = sample_dir / "standard.xlsx"
    pattern_id = create_type(client, path, "設備修理報告書")
    for label_cell, value_cell in (("A3", "B3"), ("A4", "B4"), ("E4", "F4"), ("A7", "A8")):
        add_field(client, pattern_id, "修理報告書", label_cell, value_cell)
    activate(client, pattern_id)
    doc_id, = upload_forms(client, path)
    read_form(client, doc_id, pattern_id, ["修理報告書"])
    return doc_id, pattern_id


def _doc_extraction(app, doc_id) -> dict:
    with app.app_context():
        return json.loads(db.get_document(doc_id)["data_json"])


def test_editing_a_reread_confirmed_form_keeps_the_new_form_type_settings(app, client, sample_dir):
    """確定後に種類で「Markdownに出さない」にした項目は、読み直したあと別の項目を直しても出ないまま。"""
    doc_id, pattern_id = _read_standard(app, client, sample_dir)
    ex = _doc_extraction(app, doc_id)
    shown = [f for f in ex["fields"] if f["data_type"] == "string" and f["value"]]
    omitted, edited = shown[0], shown[1]
    with app.app_context():
        doc = db.get_document(doc_id)
        db.update_document(doc_id, confirmed_json=doc["data_json"])  # 確定済み
        pattern = db.load_pattern(pattern_id)
        for fd in pattern.fields:
            if fd.field_name == omitted["field_name"]:
                fd.rag_output = "omit"
        db.save_pattern(pattern, "active")
    client.post(f"/forms/{doc_id}/read", data={"pattern_id": pattern_id, "sheets": ex["sheets"],
                                               "acknowledge": "on"})
    line = f"- {omitted['display_name']}: "
    assert line not in client.post(f"/forms/{doc_id}/preview", json={"values": {}}).get_json()["markdown"]

    client.post(f"/forms/{doc_id}/draft", json={"values": {edited["field_name"]: "別の値"}})
    ex = _doc_extraction(app, doc_id)
    field = next(f for f in ex["fields"] if f["field_name"] == omitted["field_name"])
    assert field["rag_output"] == "omit"
    assert line not in client.post(f"/forms/{doc_id}/preview", json={"values": {}}).get_json()["markdown"]


def test_confirmed_field_is_not_restored_when_its_settings_differ():
    from views.forms import _restore_confirmed_fields

    old = {"field_name": "line", "display_name": "ライン", "data_type": "string", "required": False,
           "value": "L4", "unit": "", "rag_output": "show"}
    new = dict(old, display_name="製造ライン", rag_output="omit", edited=True)
    confirmed = {"pattern": {"id": 1, "version_no": 2}, "sheets": ["s"], "fields": [old]}
    ex = {"pattern": {"id": 1, "version_no": 2}, "sheets": ["s"], "fields": [new]}
    _restore_confirmed_fields(ex, confirmed)
    assert ex["fields"][0]["rag_output"] == "omit" and ex["fields"][0]["display_name"] == "製造ライン"
    # 設定が同じなら、確定済みの中身に戻す（従来どおり）
    ex = {"pattern": {"id": 1, "version_no": 2}, "sheets": ["s"], "fields": [dict(old, edited=True)]}
    _restore_confirmed_fields(ex, confirmed)
    assert ex["fields"][0] == old


def test_triangle_minus_sign_is_read_as_negative():
    from excel.extractor import number_unit
    from excel.text import to_number

    assert to_number("▲50万円", "▲50万円", "万円")[0] == -50
    assert to_number("△0.8", "△0.8", "%") == (-0.8, None)
    assert to_number("▲5", "▲5", "") == (-5, None)
    value, warning = to_number("▲50万円", "▲50万円", "万円")
    assert number_unit(value, "▲50万円", "万円", "cost", "コスト削減効果", warning) == ("万円", None)
    # 手で「▲50万円」と入れ直しても負の数
    ex = {"fields": [_number_field(field_name="cost", display_name="コスト削減効果", value=50, unit="万円",
                                   spec_unit="万円")]}
    apply_manual_values(ex, {"value-cost": "▲50万円"})
    assert ex["fields"][0]["value"] == -50 and ex["fields"][0]["unit"] == "万円"


def test_dated_time_range_is_not_read_as_its_month():
    from excel.text import numeric_unit, to_number

    s = "12/24 21:53-12/25 11:09（13.3h）"
    assert to_number(s, s, "")[0] == 798 and numeric_unit(s) == "分"
    assert to_number(s, s, "時間")[0] == 13.3
    s = "3/19 11:19～3/20 0:48（工数 27.5h）"
    assert to_number(s, s, "時間")[0] == 27.5
    # 時間数の無い月日付きの範囲は、日数が分からないので数値にしない
    s = "12/24 21:53-12/25 11:09"
    value, warning = to_number(s, s, "")
    assert value == s and "時刻の範囲" in warning


def test_unit_written_next_to_an_annotated_number_is_kept():
    from excel.extractor import number_unit
    from excel.text import to_number

    def judge(text, spec):
        value, warning = to_number(text, text, spec)
        return (value, *number_unit(value, text, spec, "stop_minutes", "停止時間", warning))

    value, unit, warning = judge("約90分", "時間")
    assert value == 90 and unit == "分" and warning.startswith("この項目の単位は「時間」")
    value, unit, warning = judge("約2時間", "分")
    assert value == 2 and unit == "時間" and warning.startswith("この項目の単位は「分」")
    value, unit, warning = judge("595分（9.9h）", "")
    assert value == 595 and unit == "分" and "数値の部分だけ" in warning  # 括弧書きは残っているので要確認のまま
    # 「3時間40分」は換算するので、最初の「時間」を書かれた単位として扱わない
    assert judge("3時間40分", "分")[:2] == (220, "分")
    from excel.text import written_unit

    assert written_unit("約2号機") == "" and written_unit("約90分ぐらいかかった") == ""
    assert written_unit("約2号機 30分") == ""  # to_number が読む最初の数値（2）の単位だけを見る  # 単位らしくない言葉は単位にしない


def test_date_keeps_the_time_of_day():
    from datetime import datetime

    from excel.text import to_date
    from export.formats import _format_value

    assert to_date(datetime(2023, 7, 10, 23, 8), "") == ("2023-07-10 23:08", None)
    assert to_date(datetime(2023, 7, 10), "") == ("2023-07-10", None)
    assert to_date("2023/7/10 23:08", "2023/7/10 23:08") == ("2023-07-10 23:08", None)
    f = {"data_type": "date", "value": "2023-07-10 23:08"}
    assert _format_value(f) == "2023-07-10 23:08（2023年7月）"


def test_hand_typed_date_without_a_year_stays_to_be_checked():
    from views.forms import _field_status

    ex = {"fields": [{"field_name": "d", "display_name": "発生日", "data_type": "date", "required": False,
                      "value": "2/12 3時17分", "warning": "年が書かれていません。", "edited": False}]}
    apply_manual_values(ex, {"value-d": "2/12 3:17"})
    assert ex["fields"][0]["edited"] and _field_status(ex["fields"][0])["issue"]
    apply_manual_values(ex, {"value-d": "あした"})
    assert _field_status(ex["fields"][0])["issue"]
    apply_manual_values(ex, {"value-d": "2026-02-12 3:17"})
    assert not ex["fields"][0]["warning"] and not _field_status(ex["fields"][0])["issue"]


def test_batch_progress_counts_the_confirmed_forms(app, client):
    ids = [_confirmed_doc(app, f"{i}.xlsx", batch_id="b1", order=i) for i in range(2)]
    body = _finish(client, ids, current=ids[0])
    assert body["confirmed"] == 2 and body["next_id"] is None
    assert "確定済み <strong>2</strong> / 2 件" in body["html"]
    with app.app_context():
        db.update_document(ids[1], confirmed_json=None)
    body = _finish(client, ids, current=ids[0])
    assert body["confirmed"] == 1 and body["next_id"] == ids[1]
    assert "確定済み <strong>1</strong> / 2 件" in body["html"] and "次の帳票へ（残り1件）" in body["html"]


def test_upload_errors_do_not_put_the_file_name_in_the_session(app, client, sample_dir):
    files = [(io.BytesIO((sample_dir / "standard.xlsx").read_bytes()), "standard.xlsx"),
             (io.BytesIO(b"PK\x03\x04 broken"), "X社_社外秘.xlsx")]
    res = client.post("/forms/upload", data={"file": files}, content_type="multipart/form-data")
    errors = " ".join(res.get_json()["errors"])
    assert "2件目のファイル" in errors and "社外秘" not in errors
    # 画面を離れても残るところ（セッションのクッキー）にはファイル名を置かない
    assert "社外秘" not in str(res.headers.get("Set-Cookie", ""))
    with client.session_transaction() as session:
        assert not session.get("_flashes")


def _sheetless_workbook(sample_dir) -> bytes:
    import zipfile

    src = zipfile.ZipFile(sample_dir / "standard.xlsx")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/workbook.xml":
                data = re.sub(rb"<sheets>.*?</sheets>", b"", data, flags=re.S)
            dst.writestr(item, data)
    return out.getvalue()


def test_workbook_without_a_sheet_list_is_refused(app, client, sample_dir):
    res = client.post("/forms/upload", data={"file": (io.BytesIO(_sheetless_workbook(sample_dir)), "s.xlsx")},
                      content_type="multipart/form-data")
    assert res.status_code == 400 and "シートがないブックです" in res.get_json()["error"]
    with app.app_context():
        assert db.get_db().execute("SELECT COUNT(*) FROM documents").fetchone()[0] == 0
    stored = Path(app.config["UPLOAD_DIR"]) / "documents"
    assert not stored.exists() or not any(stored.iterdir())


def test_workbook_with_too_many_cells_is_refused_before_reading(app, client, tmp_path):
    """展開すると大量のセルがある小さなブックは、帳票でも見本でも読み込む前に断り、何も残さない。"""
    from tests.test_core_files import _xlsx_with_cells

    app.config["EXCEL_MAX_CELLS"] = 1000
    data = _xlsx_with_cells(tmp_path / "many.xlsx", 1001).read_bytes()
    res = client.post("/forms/upload", data={"file": (io.BytesIO(data), "many.xlsx")},
                      content_type="multipart/form-data")
    assert res.status_code == 400 and "セル数が上限（1,000 セル）を超えています" in res.get_json()["error"]
    res = client.post("/form-types/new", data={"name": "多すぎ", "samples": (io.BytesIO(data), "many.xlsx")},
                      content_type="multipart/form-data")
    assert res.status_code == 400 and "セル数が上限（1,000 セル）を超えています" in res.get_json()["error"]
    with app.app_context():
        assert db.get_db().execute("SELECT COUNT(*) FROM documents").fetchone()[0] == 0
    assert [p for p in Path(app.config["UPLOAD_DIR"]).rglob("*") if p.is_file()] == []
