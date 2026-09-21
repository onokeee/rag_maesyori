"""画面（blueprint）。帳票取り込み（forms_bp）・帳票登録（form_types_bp）・表の取り込み（tables_bp）の3つと、
使っている人ごとの作業場所・ダウンロード名・他サイトからの書き込みの拒否などの共通の小物。
"""
from __future__ import annotations

import hashlib
import io
import json
import re
import threading
import time
import unicodedata
import zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace
from urllib.parse import quote, urlsplit
from uuid import uuid4

from flask import (
    request,
    session,
    Blueprint,
    abort,
    current_app,
    jsonify,
    redirect,
    render_template,
    send_file,
    url_for,
    flash,
)

import database as db
from core import (
    FORM_MAX_MERGED_CELLS,
    UploadError,
    original_name,
    precheck_excel,
    remove_upload,
    save_upload,
    upload_path,
    read_upload,
    Book,
    safe_filename_part,
    discard_documents,
    discard_table_imports,
    get,
    get_job,
    latest_job,
    on_documents_purged,
    purge_after_send,
    purge_batch,
    purge_documents,
    purge_incomplete,
    purge_session,
    purge_table_import,
    put,
    request_cancel,
    request_pause,
    request_resume,
)
from forms import (
    apply_manual_values,
    extract_document,
    is_blank_value,
    refresh_summary,
    clean_table_value,
    is_table_value,
    parse_table_text,
    table_text_lines,
    EXCEL_ERROR_WARNING,
    WorkbookInfo,
    load_workbook_info,
    build_markdown,
    markdown_filename,
    rank_patterns,
    table_like_sheets,
    suggest_title_fields,
    click_field,
    merge_labels,
    merge_target,
    same_sheet_field,
    separate_names,
    split_rows,
    table_cells,
    pattern_to_meta,
    pattern_to_rows,
    rows_to_pattern,
    match_pattern,
    PatternDef,
)
from tables import (
    count_levels,
    has_blocking,
    suggest_columns,
    ai_point_lines,
    people_index_for,
    record_block,
    open_source,
    spec_from_suggestions,
    validate_spec,
    build_download,
    create_import,
    get_import,
    import_files,
    import_source,
    issues_csv as build_issues_csv,
    load_issues,
    load_rows,
    load_rows_page,
    md_paths,
    md_text,
    preview_signature,
    ready_preview_files,
    save_spec,
    spec_for_import,
    start_preview_job,
    start_read_job,
    start_render_job,
    update_import,
)



# ====================================================================================================
# 元 views/__init__.py
# ====================================================================================================

# ---- 使っている人ごとの作業場所 ---------------------------------------------------------
# 社内LANのサーバーで動かし、数人が同時に別々のPCから使う（2026-09-20 の利用者の指示）。
# ログインは無いので「誰か」は分からないが、「どのブラウザか」はセッションのクッキーで分かる。
# 取り込んだ帳票・一覧表はそのブラウザのものとして持ち主を記録し、ほかのブラウザからは
# 見えない・触れないようにする（views/forms.py・views/tables.py の 404）。
# 「設定」（帳票の種類・取り込み設定・AI接続）はみんなで使うものなので分けない。

SESSION_ID_KEY = "sid"


def current_session_id() -> str:
    """このブラウザの作業場所の id（クッキーに無ければ作る）。"""
    sid = session.get(SESSION_ID_KEY)
    if not isinstance(sid, str) or len(sid) != 32:
        sid = uuid4().hex
        session[SESSION_ID_KEY] = sid
    return sid


def owns(row) -> bool:
    """帳票・取り込みの行がこのブラウザのものか。

    session_id が空の行は持ち主が分からない（この仕組みを入れる前のDBの行・テストで直接作った行）ので、
    これまでどおり誰からでも扱えるものとして扱う。起動時の片付けで消えるので、普段は残らない。
    """
    owner = (row or {}).get("session_id") if hasattr(row, "get") else None
    return not owner or owner == current_session_id()


# ---- ダウンロードのファイル名 ---------------------------------------------------------
# Content-Disposition には UTF-8 の名前（filename*）と、それを読めない相手向けの ASCII の名前
# （filename）が並ぶ。日本語だけの名前だと ASCII 側が「_2026-09-19.zip」のように意味を失い、
# ログや古い環境では何のファイルか分からなくなるので、英字の既定名にそろえる。

_ASCII_UNSAFE = re.compile(r"[^A-Za-z0-9._-]+")


def ascii_download_name(name: str, default_stem: str) -> str:
    """name から作る ASCII のファイル名。英字が残らなければ default_stem を使う（拡張子は保つ）。"""
    stem = _ASCII_UNSAFE.sub("_", Path(name).stem).strip("._-")
    if not re.search(r"[A-Za-z]", stem):
        stem = f"{default_stem}_{stem}".strip("._-") if stem else default_stem
    return f"{stem}{Path(name).suffix}"


def set_download_name(response, name: str, default_stem: str):
    """添付ファイル名を「意味のある ASCII 名 ＋ 日本語の名前」にそろえる。"""
    fallback = ascii_download_name(name, default_stem)
    encoded = quote(name, safe="")
    response.headers["Content-Disposition"] = (
        f"attachment; filename=\"{fallback}\"; filename*=UTF-8''{encoded}")
    return response


# ---- 書き込みの発火元を確かめる（他サイトからの操作を断る） ----------------------------
# 127.0.0.1 だけで待ち受けても、「利用者が開いた別のサイトのページが、そのブラウザから
# このアプリへ POST する」ことは防げない（AI接続先の書き換え・帳票の削除などができてしまう）。
# ブラウザは POST に必ず Origin を付け、最近のブラウザは Sec-Fetch-Site も付けるので、
# 他サイト発と分かる書き込みだけを断る（curl などヘッダの無い要求はアプリを直接たたく操作として許す）。

SAFE_METHODS = ("GET", "HEAD", "OPTIONS")
SAME_SITE_FETCH = ("same-origin", "none")

# GET でもデータを消すルート（Markdown をダウンロードすると、その帳票・取り込みを消す。design.md 3.3）。
# 他サイトのページの <img>・リンク・window.open からでも GET は出せるので、書き込みと同じく発火元を確かめる。
# アドレス欄に打った URL（Sec-Fetch-Site: none）とアプリ内のクリック（same-origin）は通す。
PURGING_ENDPOINTS = frozenset({"forms.download_md", "forms.download_batch", "tables.download_zip"})


def _origin_of(url: str) -> str:
    parts = urlsplit(url)
    return f"{parts.scheme}://{parts.netloc}" if parts.scheme and parts.netloc else ""


def is_cross_site_write(req=None) -> bool:
    """他サイトのページから出された書き込み要求（データを消すダウンロードを含む）なら True。"""
    req = req if req is not None else request
    purging = req.endpoint in PURGING_ENDPOINTS
    if req.method in SAFE_METHODS and not purging:
        return False
    site = (req.headers.get("Sec-Fetch-Site") or "").strip().lower()
    if site and site not in SAME_SITE_FETCH:
        return True
    host = req.host_url.rstrip("/")
    origin = (req.headers.get("Origin") or "").strip()
    if origin and origin.rstrip("/") != host:
        return True
    if purging and not site:
        # Sec-Fetch-Site を付けない古いブラウザ向け。Referer が別サイトなら断る（無ければ手入力として通す）
        referer = (req.headers.get("Referer") or "").strip()
        return bool(referer) and _origin_of(referer) != host
    return False


def render_part(template: str, part: str, **ctx) -> str:
    """画面のテンプレート（forms.html など）の中のマクロ part_* を1つだけ描いて、HTML の断片を返す。

    段の中身は fetch でそのつど取りに来るので、ページ全体ではなく断片だけを描く。マクロは render_template と同じ
    文脈（url_for・request・ctx の変数）を見る。マクロに引数があれば ctx の同じ名前の値を渡す。
    """
    current_app.update_template_context(ctx)
    module = current_app.jinja_env.get_template(template).make_module(ctx)
    fn = getattr(module, part)
    return str(fn(**{k: ctx[k] for k in fn.arguments if k in ctx}))


# ====================================================================================================
# 元 views/forms.py
# 帳票取り込み（1ファイル＝1件）。画面は /forms の1枚だけ。
#
# 上から順に「ファイルを置く」「帳票の種類とシート」「読み取り結果」「確定してダウンロード」の
# 欄が現れる。画面の移動はなく、どの操作も fetch でこのファイルのルートを呼び、返ってきた
# HTML の断片をその場に入れ替える（views/forms.py のルートは JSON か HTML の断片を返す）。
#
# Markdown は常にデータから作る（読み取り結果のプレビュー＝作業中の値、ダウンロード＝確定済みの値）。
# 複数ファイルをまとめて置くと「取り込みのまとまり（batch）」になり、同じフォームの帳票として
# まとめて1回だけ種類とシートを決め、読み取り結果は全部の帳票を縦に並べて見ていく（タブで1件ずつ
# 切り替えない。利用者の指示 2026-09-20）。最後に zip でまとめて渡す。
# ダウンロードしたデータはその場で消す（design.md 3.3）。消すのは Markdown を作り終え、本文を送り終えたあとだけ
# （途中で切れたときは消さない: core.purge.purge_after_send）。
# ====================================================================================================

forms_bp = Blueprint("forms", __name__, url_prefix="/forms")

# 左のシートプレビューに出す範囲の上限（大きいシートで画面が重くならないように）
GRID_MAX_ROWS = 300
GRID_MAX_COLS = 60
CONFIRMED_STATES = ("confirmed", "modified")
MAX_BATCH_FILES = 50
# ダウンロードでデータが消えることの案内（画面の文言・確認ダイアログで使う）
FORMS_DELETE_ON_DOWNLOAD_NOTE = ("ダウンロードすると、この帳票の元のファイルと読み取り結果はサーバーから消えます。"
                           "同じものをもう一度ダウンロードすることはできません。")
FORMS_DELETE_ON_DOWNLOAD_CONFIRM = "ダウンロードすると、この帳票のデータはサーバーから消えます。もう一度ダウンロードすることはできません。"
# 確定したあとに直した（修正中の）帳票は、直した値が .md に入るよう確定し直してから渡す
MODIFIED_DOWNLOAD_CONFIRM = "確定し直してから、直した値で Markdown を作ります。" + FORMS_DELETE_ON_DOWNLOAD_CONFIRM
BATCH_DELETE_CONFIRM = ("ダウンロードすると、このまとまりの帳票のデータはサーバーからすべて消えます。"
                        "もう一度ダウンロードすることはできません。")
LOST_WORK_MESSAGE = "読み取り直すと、手で修正した値は失われます"
# 別のタブ・古い画面から保存・確定されたときの案内
STALE_MESSAGE = "別の画面で内容が変わりました。画面を読み込み直してください"


# ---- 共通 -------------------------------------------------------------------------

def _get_document(doc_id: int) -> dict:
    """この画面（このブラウザ）の帳票を返す。ほかの人の帳票は「無い」として扱う（404）。

    403 にすると「その番号の帳票はある」ことが分かってしまうので、404 にそろえる。
    """
    doc = db.get_document(doc_id)
    if doc is None or not owns(doc):
        abort(404)
    return doc


def _load_info(doc: dict) -> WorkbookInfo | None:
    """元のファイルを読む。消えている・壊れている場合は None。"""
    try:
        return load_workbook_info(upload_path(doc["stored_path"]))
    except Exception:
        return None


def _data(doc: dict, column: str = "data_json") -> dict | None:
    raw = doc.get(column)
    if not raw:
        return None
    try:
        return json.loads(raw)
    except ValueError:
        return None


def _dumps(extraction: dict) -> str:
    # 状態（修正中かどうか）は文字列の比較で決まるので、保存は常にこの形にそろえる
    return json.dumps(extraction, ensure_ascii=False)


def _search_title(doc: dict, extraction: dict) -> str:
    """一覧に出す見出し（Markdown の1行目）。"""
    first = build_markdown(doc, extraction).split("\n", 1)[0]
    return first[2:].strip() if first.startswith("# ") else first.strip()


def _display_text(value) -> str:
    if is_table_value(value):
        return table_json(value)
    return "" if value is None else str(value)


@forms_bp.app_template_filter("table_json")
def table_json(value) -> str:
    """明細表の値を画面の入力欄（hidden）に入れる JSON。"""
    return json.dumps(value, ensure_ascii=False) if is_table_value(value) else ""


@forms_bp.app_template_filter("field_text")
def field_text(value) -> str:
    """一覧・読み取りテスト用の表示。明細表は1行1明細の「品番: X／品名: Y」。"""
    if is_table_value(value):
        return "\n".join(table_text_lines(value))
    return "" if value is None else str(value)


def _same_text(a: str, b: str) -> bool:
    norm = lambda s: s.replace("\r\n", "\n").replace("\r", "\n").strip()  # noqa: E731
    return norm(a) == norm(b)


def _apply_values(extraction: dict, values: dict, confirmed: dict | None = None) -> bool:
    """画面の入力値を反映する。表示中の値と同じ文字列の項目は触らない（勝手に「手で修正」にしない）。

    confirmed（確定済みの版）を渡すと、確定済みと同じ値に戻した項目は確定済みの項目そのものに戻す。
    戻り値: 変わった項目があれば True。
    """
    before = _dumps(extraction)
    fields = {f["field_name"]: f for f in extraction["fields"]}
    changed = {}
    for name, text in values.items():
        f = fields.get(name)
        if f is None or text is None:
            continue
        text = text if isinstance(text, str) else json.dumps(text, ensure_ascii=False)
        if f["data_type"] == "table":
            parsed, warning = parse_table_text(text)
            # 行の無い表（列見出しだけ）と空の値は同じもの
            if warning is None and (parsed == clean_table_value(f["value"])
                                    or (is_blank_value(parsed) and is_blank_value(f["value"]))):
                continue
        elif _same_text(text, _display_text(f["value"])):
            continue
        changed[f"value-{name}"] = text
    if changed:
        apply_manual_values(extraction, changed)
        _restore_confirmed_fields(extraction, confirmed)
    else:
        refresh_summary(extraction)
    return _dumps(extraction) != before


# 手で直すと変わる項目のキー。これ以外（表示名・Markdownへの出し方など）が違う項目は戻さない
_EDIT_KEYS = {"value", "unit", "warning", "edited"}
_LOCATION_KEYS = {"sheet", "value_cell"}


def _restore_confirmed_fields(extraction: dict, confirmed: dict | None) -> None:
    """確定済みと同じ値（数値は単位も）になった項目を、確定済みの項目の中身に戻す。"""
    new_pattern, old_pattern = extraction.get("pattern", {}), (confirmed or {}).get("pattern", {})
    if (not confirmed or old_pattern.get("id") != new_pattern.get("id")
            or old_pattern.get("version_no") != new_pattern.get("version_no")
            or confirmed.get("sheets") != extraction.get("sheets")):
        return  # 別の種類・版・シートで読み直した版は比べない
    before = {f["field_name"]: f for f in confirmed.get("fields", [])}
    for i, f in enumerate(extraction["fields"]):
        old = before.get(f["field_name"])
        same_value = old is not None and (old.get("value") == f.get("value") or (
            f.get("data_type") == "table" and is_blank_value(old.get("value")) and is_blank_value(f.get("value"))))
        if (old is not None and old != f and same_value
                and (old.get("unit") or "") == (f.get("unit") or "")
                and {k: v for k, v in old.items() if k not in _EDIT_KEYS | _LOCATION_KEYS}
                == {k: v for k, v in f.items() if k not in _EDIT_KEYS | _LOCATION_KEYS}):
            extraction["fields"][i] = json.loads(json.dumps(old))
    refresh_summary(extraction)


def _is_blank(value) -> bool:
    return is_blank_value(value)


_UNIT_WARNING_PREFIXES = ("この項目の単位は", "単位が書かれていません")


def _field_status(f: dict) -> dict:
    """項目の状態タグ: 値の出どころ（auto/manual/blank）と要確認かどうか。"""
    blank = _is_blank(f["value"])
    if f.get("edited"):
        source = "manual"
    elif blank:
        source = "blank"
    else:
        source = "auto"
    # 単位の食い違い・単位なしの警告は、手で直した値でも Markdown に誤った単位で出るので要確認のまま
    unit_issue = f["data_type"] == "number" and str(f.get("warning") or "").startswith(_UNIT_WARNING_PREFIXES)
    date_issue = f["data_type"] == "date" and bool(f.get("warning"))
    value = f["value"]
    number_issue = f["data_type"] == "number" and (
        "数値の部分だけ" in str(f.get("warning") or "")
        or not isinstance(value, (int, float)) or isinstance(value, bool))
    error_value = source == "blank" and str(f.get("warning") or "").startswith(EXCEL_ERROR_WARNING)
    issue = error_value or (
        not blank and bool(f.get("warning")) and (source == "auto" or unit_issue or date_issue or number_issue))
    return {"source": source, "issue": bool(issue), "blank": blank}


def _summary(doc: dict, extraction: dict) -> dict:
    """読み取り結果の欄の数値（チップ・Markdownプレビュー・項目ごとの状態）。"""
    statuses = {f["field_name"]: _field_status(f) for f in extraction["fields"]}
    return {
        "markdown": build_markdown(doc, extraction),
        "file_name": markdown_filename(doc, extraction),
        "state": doc["state"],
        "counts": {
            "issue": sum(1 for s in statuses.values() if s["issue"]),
            "manual": sum(1 for s in statuses.values() if s["source"] == "manual"),
            "blank": sum(1 for s in statuses.values() if s["blank"]),
        },
        "fields": {name: {**s, "warning": f.get("warning") or ""}
                   for name, s, f in ((f["field_name"], statuses[f["field_name"]], f) for f in extraction["fields"])},
    }


def _sheet_grids(info: WorkbookInfo | None, sheet_names: list[str]) -> list[dict]:
    """元のシートを HTML の表にするための行データ（結合セルは rowspan/colspan）。"""
    if info is None:
        return []
    from openpyxl.utils import get_column_letter

    sheets = []
    for name in sheet_names:
        grid = info.grids.get(name)
        if grid is None:
            continue
        max_row, max_col = min(grid.max_row, GRID_MAX_ROWS), min(grid.max_col, GRID_MAX_COLS)
        rows = []
        for r in range(1, max_row + 1):
            cells = []
            for c in range(1, max_col + 1):
                top, left, bottom, right = grid.bounds(r, c)
                if (top, left) != (r, c):
                    continue  # 結合範囲の左上以外は描かない
                cell = grid.cells.get((r, c))
                cells.append({
                    "coord": f"{get_column_letter(c)}{r}",
                    "text": cell.text if cell else "",
                    "rowspan": min(bottom, max_row) - r + 1,
                    "colspan": min(right, max_col) - c + 1,
                    "label": bool(cell and (cell.bold or cell.filled)),
                })
            rows.append({"index": r, "cells": cells})
        sheets.append({
            "name": name,
            "rows": rows,
            "letters": [get_column_letter(c) for c in range(1, max_col + 1)],
            "truncated": grid.max_row > GRID_MAX_ROWS or grid.max_col > GRID_MAX_COLS,
            "images": len(info.images_in([name])),
        })
    return sheets


def _version(doc: dict) -> str:
    """読み取り結果を出したときの作業データの版（古い画面からの保存・確定を見分ける）。"""
    return hashlib.sha1(str(doc.get("data_json") or "").encode("utf-8")).hexdigest()[:16]


def _forms_payload() -> dict:
    payload = request.get_json(force=True, silent=True)
    return payload if isinstance(payload, dict) else {}


def _is_stale(doc: dict) -> bool:
    """画面が送ってきた版が今の作業データと違うか。版を送らない呼び出しは比べない。"""
    sent = _forms_payload().get("version") or request.form.get("version")
    return bool(sent) and str(sent) != _version(doc)


# ---- 1枚の画面 ----------------------------------------------------------------------

@forms_bp.get("/", endpoint="index")
@forms_bp.get("/new", endpoint="new")
def forms_page():
    """帳票取り込みの1枚の画面。ここから先はすべて fetch で欄が増えていく。"""
    current_session_id()   # 画面を開いた時点で作業場所（クッキー）を決めておく（同時に置かれても取り違えない）
    return render_template("forms.html",
                           active_pattern_count=db.count_active_patterns(),
                           pattern_count=len(db.list_patterns()),
                           max_files=MAX_BATCH_FILES,
                           max_mb=(current_app.config.get("MAX_CONTENT_LENGTH") or 0) // (1024 * 1024))


# ---- ファイルを置く ------------------------------------------------------------------

def upload_error_text(storage, exc: Exception, position: int | None = None) -> str:
    """取り込めなかったファイルのメッセージ。まとめて置いたときは「3件目のファイル」と位置で示す。"""
    message = str(exc)
    name = original_name(storage)
    if name and message.startswith(f"{name}: "):
        message = message[len(name) + 2:]
    return f"{position}件目のファイル: {message}" if position else message


def _store_document(storage, batch_id: str = "", order: int = 0) -> int:
    """1ファイルを保存して帳票を作る。読めないファイルは保存先から消して UploadError。"""
    cfg = current_app.config
    stored = save_upload(storage, "documents", cfg["ALLOWED_EXTENSIONS"], cfg["MAX_CONTENT_LENGTH"])
    try:
        path = upload_path(stored.stored_path)
        precheck_excel(path, cfg.get("EXCEL_MAX_CELLS"), max_merged=FORM_MAX_MERGED_CELLS)
        try:
            info = load_workbook_info(path)
        except Exception as exc:
            current_app.logger.warning("帳票を読み込めませんでした: %s", exc.__class__.__name__)
            raise UploadError("Excelファイルとして読み込めませんでした") from exc
        if not info.grids:
            raise UploadError("シートがないブックです。シートのあるブックを選んでください")
    except Exception:
        remove_upload(stored.stored_path)   # 思わぬエラーでもアップロードしたファイルを残さない（design.md 3.3）
        raise
    return db.create_document(stored.file_name, stored.file_hash, stored.stored_path,
                              batch_id=batch_id, batch_order=order, session_id=current_session_id())


@forms_bp.post("/upload", endpoint="upload")
def forms_upload():
    """置かれたファイルを取り込む（fetch）。複数なら1つのまとまり（batch）にして、同じ画面で1件ずつ読む。"""
    storages = [s for s in request.files.getlist("file") if s is not None and s.filename]
    if not storages:
        return jsonify(error="ファイルを置いてください"), 400
    if len(storages) > MAX_BATCH_FILES:
        return jsonify(error=f"一度に置けるのは{MAX_BATCH_FILES}ファイルまでです（置かれたのは{len(storages)}ファイル）"), 400
    batch_id = uuid4().hex if len(storages) > 1 else ""
    docs, errors = [], []
    for order, storage in enumerate(storages):
        try:
            doc_id = _store_document(storage, batch_id=batch_id, order=order)
        except UploadError as exc:
            errors.append(upload_error_text(storage, exc, order + 1 if len(storages) > 1 else None))
            continue
        doc = db.get_document(doc_id)
        docs.append({"id": doc_id, "file_name": doc["file_name"] if doc else ""})
    if not docs:
        return jsonify(error=errors[0] if errors else "取り込めるファイルがありませんでした", errors=errors), 400
    return jsonify(docs=docs, batch_id=batch_id, errors=errors)


# ---- 画面を離れたので捨てる ------------------------------------------------------------
# 「その場でダウンロードしない限り、その場ですぐ捨てる」（利用者の指示 2026-09-20）。
# 画面を閉じた・隠したときに static/app.js の ragDiscard がここへ「捨てて」と送ってくる。
# 送り主は navigator.sendBeacon なので、
#   - 中身の型は text/plain（get_json(force=True) で読む）
#   - 応答は読めず、やり直しもできない（いつでも 204 を返し、4xx にしない）
# もう無い番号・ほかの人の番号・処理中のものは core.purge 側で黙って外れる。

@forms_bp.post("/discard", endpoint="discard")
def forms_discard():
    """この画面（このブラウザ）の、まだダウンロードしていない帳票を捨てる。"""
    payload = request.get_json(force=True, silent=True) or {}
    ids = payload.get("doc_ids") or []
    sid = current_session_id()
    try:
        if ids:
            discard_documents(ids, sid)
        else:
            purge_session(sid, tables=False)   # 表の取り込み（別のタブ）は巻き込まない
    except Exception as exc:   # 捨て損ねてもブラウザには伝えられない。時間切れの片付けに任せる
        current_app.logger.warning("帳票の片付けに失敗しました: %s", exc.__class__.__name__)
    return "", 204


# ---- 2 帳票の種類とシート ---------------------------------------------------------------

def _id_list(raw: str) -> list[int]:
    return [int(part) for part in raw.split(",") if part.strip().isdigit()]


@forms_bp.get("/type", endpoint="type_all")
def type_all_fragment():
    """帳票の種類とシートの欄（まとめて置いた分すべてに同じ設定を使う）。?ids=1,2,3"""
    docs = _docs_of(_id_list(request.args.get("ids", "")))
    if not docs:
        return jsonify(error="取り込んだ帳票がありません。ファイルを置き直してください"), 404
    return _type_response(docs)


@forms_bp.get("/<int:doc_id>/type")
def type_fragment(doc_id: int):
    """帳票1件の「帳票の種類とシート」（まとまりでも同じ欄を使う）。"""
    return _type_response([_get_document(doc_id)])


def _batch_matches(ranked: list[dict]) -> list[SimpleNamespace]:
    """置かれたファイル全部をまとめた帳票の種類の候補（見つかった項目が多い順）。

    同じフォームの帳票をまとめて置く前提なので、種類は1つだけ選ぶ。件数で違う分は
    「3項目中2〜3項目」のように幅で出し、ファイルごとの数は下のファイルの行に出す。
    """
    out = []
    for pattern_id in (ranked[0] if ranked else {}):
        ms = [r[pattern_id] for r in ranked if pattern_id in r]
        if not ms:
            continue
        sheets: list[str] = []
        for m in ms:
            sheets.extend(n for n in m.sheet_names if n not in sheets)
        lo, hi = min(m.found_fields for m in ms), max(m.found_fields for m in ms)
        total = ms[0].total_fields
        out.append(SimpleNamespace(
            pattern=ms[0].pattern, total_fields=total, found_fields=lo, sheet_names=sheets,
            confidence=sum(m.confidence for m in ms) / len(ms),
            found_label=(f"{total}項目中{lo}項目が見つかりました" if lo == hi
                         else f"{total}項目中{lo}〜{hi}項目が見つかりました")))
    out.sort(key=lambda m: m.confidence, reverse=True)
    return out


def _type_response(docs: list[dict]):
    patterns = db.load_active_patterns()
    ranked: list[dict] = []
    sheets: list[dict] = []
    by_name: dict[str, dict] = {}
    table_sheets: list[str] = []
    # ブックは1件ずつ開いて、必要な数だけ取り出したらすぐ手放す。50件を同時にメモリへ広げると
    # サーバーが落ちるので、置かれたファイル全部の WorkbookInfo を持ち続けない
    for doc in docs:
        info = _load_info(doc)
        if info is None:
            return jsonify(error="元のファイルを読み込めませんでした。ファイルを置き直してください"), 409
        ranked.append({m.pattern.id: m for m in rank_patterns(info, patterns)})
        for name, grid in info.grids.items():
            sheet = by_name.get(name)
            if sheet is None:
                sheet = {"name": name, "cells": len(grid.cells), "images": len(info.images_in([name])),
                         "hidden": grid.hidden, "files": 0}
                by_name[name] = sheet
                sheets.append(sheet)
            sheet["files"] += 1
        table_sheets.extend(n for n in table_like_sheets(info) if n not in table_sheets)
        info = grid = None   # 次の1件を開く前にブックを手放す

    matches = _batch_matches(ranked)
    suggested = {str(m.pattern.id): m.sheet_names for m in matches}
    first = docs[0]
    selected_id = first["pattern_id"] if any(m.pattern.id == first["pattern_id"] for m in matches) else None
    selected_id = selected_id or (matches[0].pattern.id if matches else None)
    # 読み取り済みの帳票が読んだシートを全部合わせる。1件分の "sheets" には、そのファイルに
    # 在ったシートしか残らないので、先頭ファイルだけを見ると選んだチェックが消えてしまう
    selected_sheets: list[str] = []
    for d in docs:
        data = _data(d) if d["pattern_id"] == selected_id else None
        for name in (data or {}).get("sheets") or []:
            if name not in selected_sheets:
                selected_sheets.append(name)
    if not selected_sheets:
        selected_sheets = suggested.get(str(selected_id)) or []

    files = [{"id": d["id"], "file_name": d["file_name"],
              "found": ranked[i].get(selected_id).found_fields if selected_id in ranked[i] else 0}
             for i, d in enumerate(docs)]
    # 種類を選び直したときにファイルごとの項目数を出し直すための表（画面の JS が使う）
    file_counts = {str(pid): {str(d["id"]): ranked[i][pid].found_fields for i, d in enumerate(docs)}
                   for pid in (ranked[0] if ranked else {})}
    html = render_part("forms.html", "part_type", doc=first,
        docs=docs,
        files=files,
        file_counts=file_counts,
        matches=matches,
        sheets=sheets,
        suggested=suggested,
        selected_id=selected_id,
        selected_sheets=selected_sheets,
        table_sheets=table_sheets,
        duplicate=any(db.find_confirmed_by_hash(d["file_hash"], exclude_id=d["id"],
                                                session_id=current_session_id()) for d in docs),
        lost_work=any(d["state"] in CONFIRMED_STATES for d in docs),
        form_types_url=url_for("form_types.index"),
    )
    return jsonify(html=html, file_name=first["file_name"], has_types=bool(matches),
                   docs=[{"id": d["id"], "file_name": d["file_name"]} for d in docs])


@forms_bp.post("/read", endpoint="read_all")
def read_all():
    """置かれた帳票を1つの種類・シートでまとめて読み取り、全部の読み取り結果を返す。"""
    return _read_documents(_docs_of(_id_list(request.form.get("ids", ""))))


@forms_bp.post("/<int:doc_id>/read")
def read(doc_id: int):
    """帳票1件を読み取る（まとまりでも同じ道すじを通る）。"""
    return _read_documents([_get_document(doc_id)])


def _read_documents(docs: list[dict]):
    if not docs:
        return jsonify(error="取り込んだ帳票がありません。ファイルを置き直してください"), 404
    pattern = db.load_pattern(request.form.get("pattern_id", type=int))
    sheets = request.form.getlist("sheets")
    if pattern is None or not sheets:
        return jsonify(error="帳票の種類と読み取るシートを選んでください"), 400
    if any(d["state"] in CONFIRMED_STATES for d in docs) and request.form.get("acknowledge") != "on":
        return jsonify(error=f"{LOST_WORK_MESSAGE}。確認のチェックを入れてから読み取り直してください"), 400

    read_ids, errors = [], []

    def failed(doc: dict, reason: str) -> None:
        """読み取れなかった帳票。③に並ばないので、前の読み取り結果も残さない
        （残すと、画面に出ていない帳票を④が数えて zip の件数と合わなくなる）。"""
        errors.append(f"{doc['file_name']}: {reason}")
        if doc["state"] not in CONFIRMED_STATES and doc.get("data_json"):
            db.reset_document(doc["id"], doc["pattern_id"], None)

    for doc in docs:
        info = _load_info(doc)
        if info is None:
            failed(doc, "元のファイルを読み込めませんでした")
            continue
        use = [s for s in sheets if s in info.grids]
        if not use:
            failed(doc, "選んだシートがファイルにありません")
            continue
        extraction = extract_document(info, pattern, use)
        db.reset_document(doc["id"], pattern.id, _dumps(extraction))
        if doc["state"] not in CONFIRMED_STATES:
            db.update_document(doc["id"], title=_search_title(doc, extraction))
        read_ids.append(doc["id"])
    if not read_ids:
        return jsonify(error=errors[0] if errors else "読み取れる帳票がありませんでした", errors=errors), 400
    return _review_response(read_ids, errors=errors)


def _state_label(doc: dict, summary: dict) -> str:
    """読み取り結果の見出しに出す、その帳票の今の状態。"""
    if doc["state"] == "confirmed":
        return "確定済み"
    if doc["state"] == "modified":
        return "修正中"
    issue = summary["counts"]["issue"]
    return f"要確認 {issue}件" if issue else "未確定"


STATE_COLORS = {"確定済み": "green", "修正中": "violet", "未確定": "gray"}


def _review_response(doc_ids, errors=None):
    """読み取り結果の欄（帳票を縦に並べた HTML の断片）。

    重くならないよう、元のシートの表を作るのは先頭の帳票だけにする。残りは画面に入った時点で
    GET /forms/<id>/grid を呼んで読み込む（static/review.js）。
    """
    ids = [doc_ids] if isinstance(doc_ids, int) else list(doc_ids)
    items, docs = [], []
    for no, doc_id in enumerate(ids, start=1):
        doc = _get_document(doc_id)
        extraction = _data(doc)
        if extraction is None:
            if len(ids) == 1:
                return jsonify(error="まだ読み取りをしていません"), 409
            continue
        summary = _summary(doc, extraction)
        info = _load_info(doc) if no == 1 else None
        label = _state_label(doc, summary)
        items.append({
            "no": no, "doc": doc, "extraction": extraction, "summary": summary,
            "grids": _sheet_grids(info, extraction["sheets"]) if no == 1 else None,
            "info_missing": no == 1 and info is None,
            "version": _version(doc), "state_label": label,
            "state_color": STATE_COLORS.get(label, "amber"),
        })
        docs.append({"id": doc["id"], "file_name": doc["file_name"], "state": doc["state"],
                     "version": _version(doc), "summary": summary, "state_label": label})
    if not items:
        return jsonify(error="まだ読み取りをしていません"), 409
    html = render_part("forms.html", "part_review", items=items,
        confirmed=sum(1 for d in docs if d["state"] in CONFIRMED_STATES),
    )
    first = items[0]
    return jsonify(html=html, version=first["version"], summary=first["summary"], doc_id=first["doc"]["id"],
                   docs=docs, errors=errors or [])


@forms_bp.get("/<int:doc_id>/grid")
def grid_fragment(doc_id: int):
    """1件の帳票の「元のシート」（読み取り結果の欄が画面に入ったときに読み込む）。"""
    doc = _get_document(doc_id)
    extraction = _data(doc)
    if extraction is None:
        return jsonify(error="まだ読み取りをしていません"), 409
    info = _load_info(doc)
    html = render_part("forms.html", "part_grid", grids=_sheet_grids(info, extraction["sheets"]),
                           info_missing=info is None)
    return jsonify(html=html, doc_id=doc_id)


# ---- 3 読み取り結果（その場で直す・途中保存） -----------------------------------------------

# 途中保存をした画面の目印（doc_id → (保存後の版, 画面の目印)）。
_DRAFT_TOKENS: dict[int, tuple[str, str]] = {}


@on_documents_purged
def _forget_draft_tokens(doc_ids) -> None:
    for doc_id in doc_ids:
        _DRAFT_TOKENS.pop(int(doc_id), None)


def _page_token() -> str:
    token = _forms_payload().get("page_token")
    return str(token)[:64] if token else ""


def _saved_by_same_page(doc_id: int, doc: dict) -> bool:
    token = _page_token()
    saved = _DRAFT_TOKENS.get(doc_id)
    return bool(token) and saved is not None and saved == (_version(doc), token)


def _json_values() -> dict:
    payload = _forms_payload()
    if payload:
        values = payload.get("values", payload)
        return {str(k): v for k, v in values.items()} if isinstance(values, dict) else {}
    return {key[len("value-"):]: request.form[key] for key in request.form.keys() if key.startswith("value-")}


@forms_bp.post("/<int:doc_id>/draft")
def draft(doc_id: int):
    """入力内容の途中保存（fetch）。変わった項目だけ保存し、204 を返す。"""
    doc = _get_document(doc_id)
    extraction = _data(doc)
    if extraction is None:
        return jsonify(error="まだ読み取りをしていません"), 409
    if _is_stale(doc) and not _saved_by_same_page(doc_id, doc):
        return jsonify(error=STALE_MESSAGE), 409
    version = _version(doc)
    if _apply_values(extraction, _json_values(), _data(doc, "confirmed_json")):
        title = None if doc["state"] in CONFIRMED_STATES else _search_title(doc, extraction)
        new_json = _dumps(extraction)
        db.save_draft(doc_id, new_json, title=title)
        version = _version({"data_json": new_json})
        token = _page_token()
        if token:
            _DRAFT_TOKENS[doc_id] = (version, token)
    return "", 204, {"X-Doc-Version": version}


@forms_bp.post("/<int:doc_id>/preview")
def preview(doc_id: int):
    """入力中の値で作った Markdown と状態（保存はしない）。"""
    doc = _get_document(doc_id)
    extraction = _data(doc)
    if extraction is None:
        return jsonify(error="まだ読み取りをしていません"), 409
    if _is_stale(doc):
        return jsonify(error=STALE_MESSAGE), 409
    _apply_values(extraction, _json_values(), _data(doc, "confirmed_json"))
    return jsonify(_summary(doc, extraction))


# ---- 4 確定してダウンロード ---------------------------------------------------------------

@forms_bp.post("/<int:doc_id>/confirm", endpoint="confirm")
def forms_confirm(doc_id: int):
    """読み取り結果を確定する（fetch）。値は途中保存で入っているので、ここでは版だけ確かめる。"""
    doc = _get_document(doc_id)
    extraction = _data(doc)
    if extraction is None:
        return jsonify(error="まだ読み取りをしていません"), 409
    if _is_stale(doc):
        return jsonify(error=STALE_MESSAGE), 409
    values = _json_values()
    if values and _apply_values(extraction, values, _data(doc, "confirmed_json")):
        db.save_draft(doc_id, _dumps(extraction))
    db.confirm_document(doc_id, title=_search_title(doc, extraction))
    return jsonify(ok=True, doc_id=doc_id)


def _docs_of(ids: list[int]) -> list[dict]:
    """この画面の帳票だけ（ほかのブラウザの帳票は、番号を送られても無いものとして外す）。"""
    return [d for d in (db.get_document(i) for i in ids) if d is not None and owns(d)]


@forms_bp.get("/finish", endpoint="finish")
def finish_fragment():
    """確定してダウンロードの欄（HTML の断片）。?ids=1,2,3 は同じ画面で扱っている帳票。"""
    docs = _docs_of(_id_list(request.args.get("ids", "")))
    if not docs:
        return jsonify(html="", ready=False, confirmed=0, total=0)
    # 読み取れていない帳票は確定できない＝zip に入らないので、件数には数えない
    # （数えると「残り12件も確定して…」と書いてあるのに .md が4件しか入らない zip になる）
    ready = [d for d in docs if d.get("data_json")]
    unread = [d for d in docs if not d.get("data_json")]
    current_id = request.args.get("current", type=int)
    current = next((d for d in ready if d["id"] == current_id), None) or (ready[0] if ready else docs[0])
    confirmed = [d for d in ready if d["state"] in CONFIRMED_STATES]
    pending = [d for d in ready if d["state"] not in CONFIRMED_STATES]
    batch_id = current.get("batch_id") or ""
    working = _data(current)
    read_yet = working is not None or any(_data(d) is not None for d in docs)
    extraction = _data(current, "confirmed_json") or working
    html = render_part("forms.html", "part_finish", docs=docs,
        ready=ready,
        unread=unread,
        current=current,
        confirmed=confirmed,
        pending=pending,
        batch_id=batch_id if len(docs) > 1 else "",
        read_yet=read_yet,
        to_confirm=_to_confirm(ready),
        confirmed_states=CONFIRMED_STATES,
        file_name=markdown_filename(current, extraction) if extraction else "",
        markdown=build_markdown(current, _data(current, "confirmed_json")) if current["state"] in CONFIRMED_STATES else "",
        delete_note=FORMS_DELETE_ON_DOWNLOAD_NOTE,
        delete_confirm=(MODIFIED_DOWNLOAD_CONFIRM if current["state"] == "modified"
                        else FORMS_DELETE_ON_DOWNLOAD_CONFIRM),
        batch_confirm=_batch_zip_confirm(ready, unread),
    )
    return jsonify(html=html, ready=bool(confirmed), confirmed=len(confirmed), total=len(ready),
                   read_yet=read_yet,
                   next_id=(pending[0]["id"] if pending else None))


def _to_confirm(docs: list[dict]) -> list[dict]:
    """まとめてのダウンロードの前に確定し直す帳票（未確定と、直したまま確定していない修正中）。"""
    return [d for d in docs if d["state"] != "confirmed"]


def _batch_zip_confirm(docs: list[dict], unread: list[dict] | None = None) -> str:
    """まとまりの zip ダウンロードの確認文（確定していない分はこのボタンで確定してから渡す）。

    docs は zip に入る帳票（読み取り済み）だけ。読み取れていない帳票はサーバーに残るので、
    「すべて消えます」とは言わない。
    """
    rest = _to_confirm(docs)
    head = (f"まだ確定していない{len(rest)}件も確定してから、{len(docs)}件をまとめて zip でダウンロードします。"
            if rest else f"{len(docs)}件をまとめて zip でダウンロードします。")
    if unread:
        return (head + "ダウンロードすると、渡したこの"
                f"{len(docs)}件のデータはサーバーから消えます。"
                f"まだ読み取れていない{len(unread)}件はサーバーに残ります（②に戻ってもう一度読み取ってください）。")
    if rest:
        return head + "ダウンロードすると、このまとまりの帳票のデータはサーバーからすべて消えます。"
    return BATCH_DELETE_CONFIRM


# ---- ダウンロード -------------------------------------------------------------------

@forms_bp.get("/<int:doc_id>/download.md")
def download_md(doc_id: int):
    """Markdown を渡し、渡し終えた帳票のデータを消す（design.md 3.3）。"""
    name, body = _single_markdown(doc_id)
    response = send_file(io.BytesIO(body), mimetype="text/markdown", as_attachment=True, download_name=name,
                         conditional=False)   # Range でも全体を返す
    set_download_name(response, name, "form")
    return purge_after_send(response, purge_documents, [doc_id])


def _single_markdown(doc_id: int) -> tuple[str, bytes]:
    doc = _get_document(doc_id)
    if doc["confirmed_json"] is None:
        abort(404)   # まだ確定していない帳票は渡さない
    # 直したまま確定していない（修正中）帳票は、直した値で作る（design.md 3.3）。画面の JS を
    # 通さずにリンクを開いたとき（中クリック・「名前を付けてリンク先を保存」）でも同じにする
    extraction = _data(doc) if doc["state"] == "modified" else _data(doc, "confirmed_json")
    if extraction is None:
        abort(404)
    return markdown_filename(doc, extraction), build_markdown(doc, extraction).encode("utf-8")


def _unique_name(name: str, used: set[str]) -> str:
    """まとまりの中で重ならない名前（大文字・小文字だけの違いも重なりとみなす）。"""
    stem, suffix = Path(name).stem, Path(name).suffix or ".md"
    candidate, n = name, 2
    while candidate.casefold() in used:
        candidate = f"{stem}_{n}{suffix}"
        n += 1
    used.add(candidate.casefold())
    return candidate


def _batch_markdown_files(confirmed_ids: list[int]) -> list[tuple[str, bytes]]:
    files, used = [], set()
    for doc in db.list_confirmed_documents(confirmed_ids, session_id=current_session_id()):
        # 修正中（確定したあとに直した）帳票は、直した値で作る（1件の .md と同じ）
        column = "data_json" if doc["state"] == "modified" else "confirmed_json"
        try:
            extraction = json.loads(doc[column])
        except (TypeError, ValueError):
            continue
        files.append((_unique_name(markdown_filename(doc, extraction), used),
                      build_markdown(doc, extraction).encode("utf-8")))
    return files


@forms_bp.get("/batches/<batch_id>/download.zip")
def download_batch(batch_id: str):
    """まとめ取り込み1回分の Markdown を zip で渡し、渡し終えた分のデータを消す。

    確定済みの帳票だけを zip にして、その分だけ消す（未確定の帳票はサーバーに残る）。
    """
    docs = db.list_batch_documents(batch_id, session_id=current_session_id())
    if not docs:
        abort(404)   # ほかのブラウザのまとまりも「無い」として扱う
    pending = [d for d in docs if d["state"] not in CONFIRMED_STATES]
    confirmed_ids = [d["id"] for d in docs if d["state"] in CONFIRMED_STATES]
    if not confirmed_ids:
        abort(404)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, body in _batch_markdown_files(confirmed_ids):
            zf.writestr(name, body)
    buffer.seek(0)
    zip_name = f"帳票Markdown_{datetime.now():%Y%m%d_%H%M%S}.zip"
    response = send_file(buffer, mimetype="application/zip", as_attachment=True, download_name=zip_name,
                         conditional=False)   # Range でも全体を返す
    set_download_name(response, zip_name, "forms_markdown")
    if pending:
        return purge_after_send(response, purge_documents, confirmed_ids)
    return purge_after_send(response, purge_batch, batch_id)


@forms_bp.get("/<int:doc_id>/original")
def original(doc_id: int):
    doc = _get_document(doc_id)
    try:
        path = upload_path(doc["stored_path"])
    except UploadError:
        abort(404)
    if not path.exists():
        abort(404)
    return send_file(path, as_attachment=True, download_name=doc["file_name"])


@forms_bp.post("/<int:doc_id>/delete", endpoint="delete")
def forms_delete(doc_id: int):
    """この帳票の取り込みをやめる（元のファイルと読み取り結果を消す）。"""
    _get_document(doc_id)
    purge_documents([doc_id])
    incomplete = purge_incomplete()
    return jsonify(ok=True, incomplete=incomplete)


# 画面を作り直す前の URL（お気に入り・古いリンク）は1枚の画面へ送る
@forms_bp.get("/<int:doc_id>", endpoint="legacy")
@forms_bp.get("/<int:doc_id>/done", endpoint="legacy")
def forms_legacy(doc_id: int):
    return redirect(url_for(".index"))


# ====================================================================================================
# 元 views/form_types.py
# 帳票登録。画面は /form-types の1枚だけ。
#
# 上から「登録済みの帳票の種類」、その下に「新しく登録する」。Excel を1つ置くと、同じ画面に
# 名前（ファイル名から入れる）・置いた Excel のシート・読み取る項目・読み取りテストの結果・
# ［使用開始］が現れる。画面の移動はなく、どの操作も fetch でこのファイルのルートを呼び、
# HTML の断片を入れ替える。
#
# 項目は「見出しのセル → 値のセル」をクリックするだけで作る。キー名・型・単位・探す見出しは
# pattern.clicks が見本の値から決めるので、画面には出さない。
# 保存しても使用中にはしない。使用中になるのは［使用開始］を押したときだけ。
#
# 置いた Excel はサーバーに残さない（利用者の指示 2026-09-21「見本のExcelは置かずに、設定だけ
# 保持するようにしてほしい」）。受け取った要求の中で読み取り、中身はそのまま捨てる。ブラウザは
# 選んだファイルを持ったままなので、セルのクリック・項目の作り直し・読み取りテストのたびに同じ
# Excel を送り直してくる。開き直すのが遅いので、読み取った結果だけを core.workbook_cache が
# 短い間メモリに覚えておく（ディスクには書かない）。
# 残すのは設定だけ: シート名・見出しのセル・値のセル・読み取る向き・項目名。
# ====================================================================================================

form_types_bp = Blueprint("form_types", __name__, url_prefix="/form-types")

# excel/extractor.number_unit の単位不明の警告の書き出し
_NO_UNIT_WARNING = "単位が書かれていません"
TEST_NO_UNIT_WARNING = ("帳票に単位が書かれていません。取り込んだあとの「読み取り結果」で、"
                        "値に単位（分・時間など）を付けて入力できます")


def _get_pattern(pattern_id: int) -> PatternDef:
    pattern = db.load_pattern(pattern_id)
    if pattern is None:
        abort(404)
    return pattern


def _confirmed_count(pattern_id: int) -> int:
    row = db.get_db().execute(
        "SELECT COUNT(*) FROM documents WHERE pattern_id = ? AND confirmed_json IS NOT NULL", (pattern_id,)
    ).fetchone()
    return row[0] if row else 0


# ---- ブラウザが置いた Excel（保存しない） -----------------------------------------------
# 置かれた Excel は、この要求の中で読み取って中身を捨てる。次の操作のときはブラウザが同じ
# ファイルを送り直してくるので、2回目からは読み取った結果（core.workbook_cache）を使い回す。

BOOK_FIELD = "book"            # ブラウザが送ってくる Excel（<input type=file name=book>）
BOOK_HASH_FIELD = "book_hash"  # 送り直さずに、さっき読んだブックを指すとき（sha256）
NO_BOOK_ERROR = "この帳票のExcelをもう一度置いてください（サーバーには残していません）"
# ここで受け取る Excel の大きさの上限。保存せずにメモリで読む（数人が同時に置く）ので、
# 取り込みの上限（MAX_CONTENT_LENGTH＝まとめ置きの合計）より小さくしておく。
# 帳票は1枚の紙なので、写真付きでもこの大きさに収まる（見本のいちばん大きいもので約0.1MB）。
BOOK_MAX_BYTES = 50 * 1024 * 1024


def _read_book(storage) -> Book:
    """置かれた Excel をメモリで読み、読み取った結果だけを覚える。読めなければ UploadError。

    受け取ったファイルは werkzeug が 500KB まではメモリに、それを超える分だけ OS の一時ファイルに
    置く（要求が終わると消える、名前の無いファイル）。こちらからディスクに書くことはしない。
    """
    cfg = current_app.config
    limit = min(cfg["MAX_CONTENT_LENGTH"] or BOOK_MAX_BYTES, BOOK_MAX_BYTES)
    memory = read_upload(storage, cfg["ALLOWED_EXTENSIONS"], limit)
    known = get(current_session_id(), memory.file_hash)
    if known is not None:
        known.file_name = memory.file_name   # 同じ中身を別の名前で置き直したとき
        return known
    precheck_excel(memory.data, cfg.get("EXCEL_MAX_CELLS"), max_merged=FORM_MAX_MERGED_CELLS)
    try:
        # BytesIO で渡すので、どこにもファイルを作らずに読める
        info = load_workbook_info(io.BytesIO(memory.data))
    except Exception as exc:
        current_app.logger.warning("置かれたExcelを読み込めませんでした: %s", exc.__class__.__name__)
        raise UploadError("Excelファイルとして読み込めませんでした") from exc
    if not info.grids:
        raise UploadError("シートがないブックです。シートのあるブックを選んでください")
    return put(current_session_id(),
                              Book(file_name=memory.file_name, file_hash=memory.file_hash,
                                   size=memory.size, info=info))


def _request_book() -> tuple[Book | None, str]:
    """この操作で見ている Excel。戻り値: (ブック, エラー文)。

    ブラウザが送ってきたファイルを読む。ファイルが無いときは、さっき読んだブック（sha256）を探す。
    どちらも無ければ (None, "")＝Excel を置いていない画面（項目の一覧と見出しの手直しはできる）。
    """
    storage = request.files.get(BOOK_FIELD)
    if storage is not None and storage.filename:
        try:
            return _read_book(storage), ""
        except UploadError as exc:
            return None, upload_error_text(storage, exc)
    file_hash = (request.form.get(BOOK_HASH_FIELD) or request.args.get(BOOK_HASH_FIELD) or "").strip()
    if file_hash:
        return get(current_session_id(), file_hash), ""
    return None, ""


# ---- 1枚の画面 --------------------------------------------------------------------

@form_types_bp.get("/", endpoint="index")
def form_types_page():
    # 帳票の種類は「設定」なのでみんなで使う（ブラウザごとに分けない）。
    # 作業場所のクッキーだけは、ここで開いたときにも決めておく（ほかの画面での取り違えを防ぐ）
    current_session_id()
    return render_template("form_types.html", list_html=_list_html(),
                           max_mb=BOOK_MAX_BYTES // (1024 * 1024))


def _list_html() -> str:
    return render_part("form_types.html", "part_list", patterns=db.list_patterns())


# 画面を作り直す前の URL（お気に入り・古いリンク）は1枚の画面へ送る
@form_types_bp.get("/new", endpoint="legacy")
@form_types_bp.get("/<int:pattern_id>/build", endpoint="legacy")
@form_types_bp.get("/<int:pattern_id>/edit", endpoint="legacy")
@form_types_bp.get("/<int:pattern_id>/review", endpoint="legacy")
@form_types_bp.get("/<int:pattern_id>/test", endpoint="legacy")
def form_types_legacy(pattern_id: int | None = None):
    return redirect(url_for(".index"))


# ---- 新しく登録する ---------------------------------------------------------------

@form_types_bp.post("/new")
def create():
    """Excel を1つ置いて帳票の種類を作る（fetch）。名前は画面でファイル名から入れてある。

    作るのは設定だけ。置かれた Excel は読み取るだけで、サーバーには残さない。
    """
    storage = request.files.get(BOOK_FIELD)
    if storage is None or not storage.filename:
        return jsonify(error="帳票のExcelファイルを置いてください"), 400
    name = request.form.get("name", "").strip() or Path(storage.filename).stem.strip()
    if not name:
        return jsonify(error="帳票の種類の名前を入れてください"), 400
    try:
        book = _read_book(storage)
    except UploadError as exc:
        return jsonify(error=upload_error_text(storage, exc)), 400
    pattern_id = db.create_pattern(name)
    return jsonify(pattern_id=pattern_id, html=_build_html(pattern_id, book), list_html=_list_html(),
                   message=f"「{name}」を作りました。読み取りたい欄の見出しと値をクリックしてください")


# ---- 読み取る欄をクリックして決める ＋ 読み取りテスト ------------------------------------

def _build_html(pattern_id: int, book: Book | None = None, notes: list[str] | None = None) -> str:
    """登録中の帳票の種類の欄（HTML の断片）。book が無ければシートの無い（設定だけの）画面。"""
    pattern = _get_pattern(pattern_id)
    info = book.info if book is not None else None
    grids = _sheet_grids(info, list(info.grids)) if info is not None else []
    for g in grids:
        g["click_cells"] = table_cells(info.grids[g["name"]])
    return render_part("form_types.html", "part_build", pattern=pattern,
        book=book,
        grids=grids,
        rows=_field_view_rows(pattern, info),
        test=_test_result(pattern, book),
        confirmed_count=_confirmed_count(pattern_id),
        notes=notes or [],
    )


@form_types_bp.get("/<int:pattern_id>/panel")
@form_types_bp.post("/<int:pattern_id>/panel")
def build_fragment(pattern_id: int):
    """項目の一覧を出す（GET）。Excel を一緒に置くと（POST）、そのシートを見ながら直せる。

    保存済みの種類を開き直したときは Excel が無いので、項目の一覧と見出しの手直しだけができる。
    同じ帳票の Excel をもう一度置くと、シートが出てセルをクリックできるようになる（種類は増えない）。
    """
    _get_pattern(pattern_id)
    book, error = _request_book()
    if error:
        return jsonify(error=error), 400
    message = ""
    if request.method == "POST" and book is not None:
        message = (f"「{book.file_name}」を読み込みました。読み取りたい欄の見出しと値をクリックしてください"
                   "（このExcelはサーバーに残しません）")
    return jsonify(html=_build_html(pattern_id, book), list_html=_list_html(), message=message)


def _soften_unit_warning(f: dict) -> None:
    """この欄では値を直せないので、単位なしの警告はどこで直せるかを書く。"""
    if f.get("data_type") == "number" and str(f.get("warning") or "").startswith(_NO_UNIT_WARNING):
        f["warning"] = TEST_NO_UNIT_WARNING


def _field_view_rows(pattern: PatternDef, info) -> list[dict]:
    """項目の一覧（見出し・置いた Excel で見つかった値・セル）。値はいま登録されている設定で読み直す。

    Excel を置いていないとき（info が None）は値の欄を空にする。見出しの手直しはそれでもできる。
    """
    found: dict[str, dict] = {}
    if info is not None and pattern.fields:
        sheets = [s.sheet_name for s in pattern.sheets if s.sheet_name in info.grids] or list(info.grids)[:1]
        found = {f["field_name"]: f for f in extract_document(info, pattern, sheets)["fields"]}
    for f in found.values():
        _soften_unit_warning(f)
    rows = []
    for fd in pattern.fields:
        f = found.get(fd.field_name) or {}
        rows.append({
            "field": fd,
            "label": (fd.candidates[0] if fd.candidates else "") or "（見出しなし）",
            "value": f.get("value"),
            "warning": f.get("warning") or "",
            "sheet": f.get("sheet") or fd.sheet_name,
            "label_cell": f.get("label_cell") or fd.label_cell,
            "value_cell": f.get("value_cell") or fd.cell,
        })
    return rows


def _test_result(pattern: PatternDef, book: Book | None) -> dict | None:
    """いま置いている Excel を、この設定で読み取った結果（Markdown・見つかった件数）。"""
    if book is None or not pattern.fields:
        return None
    info = book.info
    match = match_pattern(info, pattern)
    sheets = match.sheet_names or info.sheet_names[:1]
    extraction = extract_document(info, pattern, sheets)
    doc = {"id": 0, "file_name": book.file_name, "file_hash": book.file_hash}
    return {
        "sheets": sheets,
        "found": sum(1 for f in extraction["fields"] if f["value"] not in (None, "")),
        "total": len(extraction["fields"]),
        "markdown": build_markdown(doc, extraction),
        "file_name": markdown_filename(doc, extraction),
    }


@form_types_bp.post("/<int:pattern_id>/fields")
def add_field(pattern_id: int):
    """クリックした見出しセル（と値セル）から項目を1つ作る（fetch）。

    どのセルを指しているかは、ブラウザが一緒に送ってくる Excel を読み直して確かめる
    （サーバーには置いていないため）。
    """
    pattern = _get_pattern(pattern_id)
    book, error = _request_book()
    if error:
        return jsonify(error=error), 400
    if book is None:
        return jsonify(error=NO_BOOK_ERROR), 400
    info = book.info
    sheet = request.form.get("sheet", "")
    grid = info.grids.get(sheet)
    if grid is None:
        return jsonify(error="シートが見つかりません。同じ帳票のExcelを置き直してください"), 400

    label_cell = request.form.get("label_cell", "")
    value_cell = request.form.get("value_cell", "")
    row, error = click_field(grid, label_cell, value_cell, {f.field_name for f in pattern.fields})
    if row is None:
        return jsonify(error=error), 400
    if any(f.sheet_name == sheet and f.label_cell == row["label_cell"] and f.cell == row["cell"]
           for f in pattern.fields):
        return jsonify(html=_build_html(pattern_id, book), list_html=_list_html(),
                       message="そのセルはもう項目になっています")

    sheet_rows, field_rows = pattern_to_rows(pattern)
    # 番号と名前を1つのセルにまとめた「使用設備」欄は、設備番号・設備名の2項目になる
    added, separated, merged = [], [], []
    for part in split_rows(row, {f.field_name for f in pattern.fields}):
        # 別の見本で書き方の違う同じ欄（「設備No」と「設備番号」）をクリックしたときは、新しい項目にせず
        # その項目の探す見出しに足す
        same = merge_target(field_rows, part, grid)
        if same is not None:
            merge_labels(same, part)
            merged.append(same["display_name"])
            continue
        # 同じ見本の別のセル（「担当者」と「報告者」）なら、辞書の名前が同じでも別の項目にする
        twin = same_sheet_field(field_rows, part, grid)
        if twin is not None:
            separate_names(twin, part)
            separated.append((part["display_name"], twin["display_name"]))
        else:
            added.append(part["display_name"])
        field_rows.append(part)
    message = "。".join(_add_messages(row, added, separated, merged))
    if not any(r["sheet_name"] == sheet for r in sheet_rows):
        sheet_rows.append({"use": True, "sheet_name": sheet})
    _save_rows(pattern, sheet_rows, field_rows)
    return jsonify(html=_build_html(pattern_id, book), list_html=_list_html(), message=message)


def _add_messages(row: dict, added: list[str], separated: list[tuple[str, str]], merged: list[str]) -> list[str]:
    """クリックの結果の知らせ（項目にした／別の項目にした／見出しに足した、のどれをしたか）。"""
    out = []
    if added:
        out.append(f"「{'」「'.join(added)}」を項目にしました")
    for name, twin in separated:
        out.append(f"「{name}」を「{twin}」とは別の項目にしました（同じ帳票の別のセルなので、両方を読み取ります）")
    if merged:
        label = (row["candidates"].splitlines() or [""])[0]
        out.append(f"「{'」「'.join(merged)}」の見出しに「{label}」を足しました"
                   "（書き方の違う同じ欄なので、1つの項目として読み取ります）")
    return out


@form_types_bp.post("/<int:pattern_id>/fields/<field_name>/delete")
def delete_field(pattern_id: int, field_name: str):
    pattern = _get_pattern(pattern_id)
    sheet_rows, field_rows = pattern_to_rows(pattern)
    rest = [r for r in field_rows if r["field_name"] != field_name]
    if len(rest) == len(field_rows):
        abort(404)
    kept = {r["sheet_name"] for r in rest if r["sheet_name"]}
    sheet_rows = [s for s in sheet_rows if not kept or s["sheet_name"] in kept]
    # 読み取る項目が無くなった使用中の種類は、使用を停止する。そのままだと帳票取り込みの候補に出て、
    # 中身の無い Markdown ができてしまう（［使用開始］も同じ決まりで断っている）
    stopped = not rest and pattern.status == "active"
    _save_rows(pattern, sheet_rows, rest, status="inactive" if stopped else None)
    message = "項目を削除しました"
    if stopped:
        message += "。読み取る項目が無くなったので、この種類の使用を停止しました（帳票取り込みの候補に出なくなります）"
    # ブラウザが Excel を一緒に送ってきていれば、シートを出したままにする
    return jsonify(html=_build_html(pattern_id, _request_book()[0]), list_html=_list_html(), message=message)


@form_types_bp.post("/<int:pattern_id>/fields/<field_name>/label")
def rename_field(pattern_id: int, field_name: str):
    """読み取る項目の見出しを手で直す（fetch）。

    直すのは Markdown に書き出す名前だけ。探す見出し（クリックしたときの見出しの言葉）と
    読み取るセルはそのままにする。書き出す名前は項目どうしで重ならないようにする。
    """
    pattern = _get_pattern(pattern_id)
    payload = request.get_json(force=True, silent=True) or {}
    name = " ".join(str(payload.get("name") or request.form.get("name", "")).split())
    if not name:
        return jsonify(error="見出しを入れてください"), 400
    sheet_rows, field_rows = pattern_to_rows(pattern)
    target = next((r for r in field_rows if r["field_name"] == field_name), None)
    if target is None:
        abort(404)
    if any(r["display_name"] == name for r in field_rows if r is not target):
        return jsonify(error=f"「{name}」はほかの項目が使っています。別の見出しにしてください"), 400
    # 探す見出しを持たない項目（見出しのない表など）は、書き出す名前をそのまま探していた。
    # 書き替えで探す先が変わらないよう、いまの見出しを探す見出しとして控えてから名前を変える
    if not target["candidates"] and not target["cell"]:
        target["candidates"] = target["display_name"]
    target["display_name"] = name
    target["renamed"] = True   # このあと別の欄をクリックしても、手で付けた見出しに戻さない
    _save_rows(pattern, sheet_rows, field_rows)
    return jsonify(html=_build_html(pattern_id, _request_book()[0]), list_html=_list_html(),
                   message=f"見出しを「{name}」にしました")


@form_types_bp.post("/<int:pattern_id>/name")
def rename(pattern_id: int):
    pattern = _get_pattern(pattern_id)
    payload = request.get_json(force=True, silent=True) or {}
    name = str(payload.get("name") or request.form.get("name", "")).strip()
    if not name:
        return jsonify(error="帳票の種類の名前を入れてください"), 400
    sheet_rows, field_rows = pattern_to_rows(pattern)
    _save_rows(pattern, sheet_rows, field_rows, name=name)
    return jsonify(ok=True, list_html=_list_html(), message="名前を変えました")


def _save_rows(pattern: PatternDef, sheet_rows: list[dict], field_rows: list[dict], name: str | None = None,
               status: str | None = None) -> None:
    """状態は変えずに保存する（使用開始は［使用開始］を押したときだけ）。タイトル項目は自動で決める。

    status を渡したときだけ状態も変える（読み取る項目が無くなったら使用を停止する）。
    """
    meta = pattern_to_meta(pattern)
    if name:
        meta["name"] = name
    meta["title_fields"] = suggest_title_fields(field_rows)
    db.save_pattern(rows_to_pattern(pattern.id, meta, sheet_rows, field_rows), status or pattern.status)


# ---- 使用開始・停止・削除 --------------------------------------------------------------

@form_types_bp.post("/<int:pattern_id>/status")
def change_status(pattern_id: int):
    pattern = _get_pattern(pattern_id)
    payload = request.get_json(force=True, silent=True) or {}
    status = payload.get("status") or request.form.get("status")
    if status not in ("active", "inactive"):
        abort(400)
    if status == "active" and not pattern.fields:
        return jsonify(error="読み取る項目がありません。シートで見出しのセルと値のセルをクリックしてください"), 400
    db.set_pattern_status(pattern_id, status)
    if status == "active":
        # 残すのは設定だけ（Excel はもともと置いていない）。画面はそのまま続けて使える
        message = f"「{pattern.name}」の使用を開始しました。帳票取り込みの候補に出ます"
    else:
        message = f"「{pattern.name}」の使用を停止しました。帳票取り込みの候補に出なくなります"
    return jsonify(ok=True, status=status, html=_build_html(pattern_id, _request_book()[0]),
                   list_html=_list_html(), message=message)


@form_types_bp.post("/<int:pattern_id>/delete", endpoint="delete")
def form_types_delete(pattern_id: int):
    pattern = _get_pattern(pattern_id)
    db.delete_pattern(pattern_id)
    return jsonify(ok=True, list_html=_list_html(), message=f"帳票の種類「{pattern.name}」を削除しました")


# ====================================================================================================
# 元 views/tables.py
# 表の取り込み（design.md 2.4）。1画面で全部できる（利用者の指示 2026-09-20）。
#
# 画面は /tables の1枚だけ。ファイルを置く → 読み取り方 → 表の範囲 → 列の対応づけ → AI整形（任意）
# → 内容の確認 → 確定してダウンロード（zip）を、同じ画面の「段」として順に開く。
# 段の中身はこの blueprint が HTML の断片（panel）として返し、保存・実行は JSON でやりとりする。
# 画面の移動（①→②→③）は無く、URL は変わらない（static/tables.js）。
# 範囲の決まり（利用者の判断）: 期間の置き換え・投入済みとの差分・取り消しはしない。取り込みごとに
# その取り込みの記録だけから全 Markdown を作り、全ファイルを zip で渡す。クロス集計と名寄せ辞書は扱わない。
# ダウンロードした取り込みのデータは、zip を送り終えたあとに消す（design.md 3.3・core.purge.purge_after_send）。
# 読み込み・Markdown 作成・AI整形は core.jobs のジョブ（tables.pipeline / aiproc.runner）で動かし、
# 同じ画面の中に進み具合を出す。表の範囲・列の段は取り込みの控え（tables.source_cache）を通して読む。
# ====================================================================================================

tables_bp = Blueprint("tables", __name__, url_prefix="/tables")

EXCEL_EXT = {".xlsx", ".xlsm"}
GRID_ROWS = 60
GRID_COLS = 30
GRID_CELL_CHARS = 40
DATA_PAGE = 100
ISSUES_SHOWN = 200

# 画面で選べる役割（列の対応づけ）。ここに無い役割（人名・数値・区分など）は候補のまま使い、画面では「その他」に見せる。
# 名前はどんな表にも当てはまる言い方にする（利用者の指示 2026-09-21。entity は「設備」、log は
# 「AI整形の対象（追記ログ）」と呼んでいたが、設備の記録以外の表には当てはまらなかった）。
# 中で使う名前（key/date/entity/log/attribute）は変えない（前に保存した取り込み設定がそのまま読める）。
SCREEN_ROLES = [("key", "識別番号"), ("date", "日付"), ("entity", "対象（設備・製品・顧客など）"),
                ("log", "経過の記録（1つのセルに日付ごとに書き足した列）"), ("attribute", "その他")]
SCREEN_ROLE_KEYS = {role for role, _label in SCREEN_ROLES}
KIND_LABELS = {"header": "見出し", "data": "データ", "subtotal": "小計・合計", "note": "注記", "continuation": "継続行",
               "excluded": "除外", "title": "表題", "blank": "空行"}
TABLE_KIND_LABELS = {"list": "一覧表", "crosstab": "クロス集計", "form_like": "帳票らしい", "unknown": "不明"}
SCOPE_LABELS = {"pending": "まだ整形していない行と、内容が変わった行", "errors": "エラーになった行だけ",
                "flagged": "要確認の行だけ", "all": "すべての行をやり直す"}
# CSV の文字コード・区切り文字は画面の選択肢だけを受け付ける（他の値は読み込みで落ちて画面が開けなくなるため）
# ダウンロードでデータが消えることの案内（design.md 3.3）
TABLES_DELETE_ON_DOWNLOAD_NOTE = ("zip をダウンロードすると、この取り込みのデータはサーバーから消えます"
                           "（もう一度ダウンロードすることはできません）。")
TABLES_DELETE_ON_DOWNLOAD_CONFIRM = ("ダウンロードすると、この取り込みのデータはサーバーから消えます。"
                              "もう一度ダウンロードすることはできません。")
ENCODING_CHOICES = [("utf-8-sig", "UTF-8（BOM付き）"), ("utf-8", "UTF-8"), ("cp932", "CP932（Shift_JIS）"),
                    ("shift_jis_2004", "Shift_JIS 2004"), ("utf-16", "UTF-16"),
                    ("utf-16-le", "UTF-16LE（BOMなし）"), ("utf-16-be", "UTF-16BE（BOMなし）")]
BUSY_MESSAGE = "処理中は変更できません。終わるか中止してから変更してください"
DELIMITER_CHOICES = [(",", "カンマ"), ("\t", "タブ"), (";", "セミコロン"), ("|", "縦棒")]


# ---- 共通 ---------------------------------------------------------------------------------

def _load_import(import_id: int) -> dict:
    """この画面（このブラウザ）の取り込みを返す。ほかの人の取り込みは「無い」として扱う（404）。

    社内LANで数人が同時に使うので、番号を打ち替えただけでほかの人の表を読めてはいけない。
    403 にすると「その番号の取り込みはある」ことが分かってしまうので、404 にそろえる。
    """
    imp = get_import(import_id)
    if imp is None or not owns(imp):
        abort(404)
    return imp


def _spec_for(imp: dict):
    return spec_for_import(imp)


def _is_excel(imp: dict) -> bool:
    return Path(imp["file_name"]).suffix.lower() in EXCEL_EXT


def _open(imp: dict):
    """控え付きの表ソース。元のファイルは控えにない範囲を読むときだけ開く。"""
    return import_source(imp)


def _sheet(imp: dict, source) -> str:
    if not _is_excel(imp):
        return source.file_name  # CSV はファイル名が1つのシート（全行を数えずに済ませる）
    src = imp.get("source") or {}
    names = [s.name for s in source.sheets()]
    if src.get("sheet") in names:
        return src["sheet"]
    visible = [s.name for s in source.sheets() if not s.hidden]
    if not (visible or names):
        raise UploadError("シートがないブックです。シートのあるブックを選んでください")
    return (visible or names)[0]


def _layout(imp: dict, source, sheet: str, header_rows=None, data_end=None, use_saved=True):
    src = imp.get("source") or {}
    if use_saved:
        header_rows = header_rows or [int(r) for r in src.get("header_rows") or [] if r] or None
        data_end = data_end or src.get("data_end_row") or None
    return source.layout(sheet, header_rows=header_rows or None, data_end=data_end or None)


def _after_read_panel(spec) -> str:
    """読み込みのあとに開く段。経過の記録の列があれば AI整形、なければ内容の確認。"""
    return "ai" if spec is not None and spec.log_stage is not None else "preview"


def _json_error(message: str, status: int = 400, **extra):
    return jsonify({"error": message, **extra}), status


def _tables_payload() -> dict:
    data = request.get_json(silent=True)
    return data if isinstance(data, dict) else {}


def _int(value, default=None):
    try:
        return int(str(value).strip())
    except (TypeError, ValueError):
        return default


# 行番号の範囲「3-4」を広げる上限（見出し行の帯なので広くはならない）。static/tables.js の parseRows と同じ規則
_MAX_ROW_SPAN = 10


def _row_no(value) -> int | None:
    """行番号の入力（全角数字も可）。数字だけのときだけ読む（static/tables.js の parseEnd と同じ規則）。"""
    text = unicodedata.normalize("NFKC", str(value if value is not None else "")).strip()
    return int(text) if text.isascii() and text.isdigit() and int(text) > 0 else None


def _int_list(value) -> list[int]:
    """見出し行の入力（「1,2」「1、2」「1 2」「１，２」「1-2」）。static/tables.js の parseRows と同じ規則で読む。"""
    if isinstance(value, (list, tuple)):
        items = [str(x) for x in value]
    else:
        items = re.split(r"[,、\s]+", unicodedata.normalize("NFKC", str(value or "")))
    out: set[int] = set()
    for item in items:
        m = re.fullmatch(r"(\d+)-(\d+)", item.strip())
        if m and item.isascii():
            a, b = int(m.group(1)), int(m.group(2))
            if 0 < a <= b and b - a < _MAX_ROW_SPAN:
                out.update(range(a, b + 1))
            continue
        n = _row_no(item)
        if n:
            out.add(n)
    return sorted(out)


def _ai_running(import_id: int) -> bool:
    job = latest_job("table_import", import_id, kind="ai_format")
    return bool(job and not job.get("finished"))


# AIの試し実行（同期の要求でジョブが無い）をしている取り込み → 実行中の数。実行中に取り込みを消すと、
# 終わった試し実行が AI の結果・応答を書き戻して残してしまう（design.md 3.3）
_TRIALS: Counter[int] = Counter()
_TRIALS_LOCK = threading.Lock()
TRIAL_BUSY_MESSAGE = "AIの試し実行中は削除・ダウンロード・保存できません。試し実行が終わってからもう一度押してください"


def _trial_running(import_id: int) -> bool:
    with _TRIALS_LOCK:
        return _TRIALS[import_id] > 0


def _processing(imp: dict) -> bool:
    """読み込み中・Markdown作成中・AI整形の実行中か。途中で設定やデータを変えると、古い設定の結果が残ったり、
    消したはずの AI の結果が書き戻されたりする（design.md 3.3）。"""
    return imp.get("status") in ("reading", "confirming") or _ai_running(imp["id"])


# ---- 1画面のやりとり（段の URL と、段の中身） ------------------------------------------------------

def _urls(import_id: int) -> dict:
    """画面（static/tables.js）が使う URL。panel は末尾の NAME を段の名前に置き換えて使う。

    列の対応づけ・AI整形・プレビュー・ダウンロードの URL は、それぞれの段の HTML が
    data-* 属性や <a href> で持っているので、ここには入れない。
    """
    def u(endpoint: str, **kw) -> str:
        return url_for(endpoint, import_id=import_id, **kw)

    return {
        "panel": url_for("tables.panel", import_id=import_id, name="NAME"),
        "source": u("tables.save_source"), "detect": u("tables.layout_detect"), "layout": u("tables.save_layout"),
        "read": u("tables.reread"), "cancel": u("tables.cancel_job"), "delete": u("tables.delete_import"),
        "preview_start": u("tables.start_preview"), "confirm": u("tables.confirm"),
    }


def _job_info(job: dict | None, import_id: int | None = None) -> dict | None:
    if job is None:
        return None
    return {"id": job["id"], "kind": job.get("kind"), "status": job.get("status"),
            "finished": bool(job.get("finished")), "message": job.get("message") or "",
            "progress": job.get("progress") or {},
            "url": url_for("tables.api_job", job_id=job["id"]),
            "cancel_url": url_for("tables.cancel_job", import_id=import_id) if import_id else None}


def _running_job(imp: dict) -> dict | None:
    """読み込み中・確定処理中なら、その進み具合。終わっているのに状態が残っていれば直す。"""
    if imp.get("status") not in ("reading", "confirming"):
        return None
    job = get_job(imp["job_id"]) if imp.get("job_id") else None
    if job is None or job.get("finished"):
        fresh = get_import(imp["id"])
        if fresh and fresh.get("status") in ("reading", "confirming"):
            if fresh["status"] == "reading":
                stats = {**(fresh.get("stats") or {}), "error": "読み込みが途中で止まりました。もう一度読み込んでください"}
                update_import(imp["id"], status="failed", stats=stats)
            else:
                update_import(imp["id"], status="preview")
        return None
    return job


def _panel(html: str, **extra):
    return jsonify({"html": html, **extra})


def _locked(reason: str, **extra):
    """まだ使えない段（灰色で1行だけ理由を出す）。"""
    return jsonify({"html": "", "locked": reason, **extra})


@tables_bp.get("/")
@tables_bp.get("/new")
def new():
    """表の取り込みの画面（1枚）。段の中身はここでは出さず、ファイルを置いたあとに取りに来る。"""
    current_session_id()   # 画面を開いた時点で作業場所（クッキー）を決めておく（同時に置かれても取り違えない）
    return render_template("tables.html")


@tables_bp.get("/imports/<int:import_id>/panel/<name>")
def panel(import_id: int, name: str):
    imp = _load_import(import_id)
    handler = _PANELS.get(name)
    if handler is None:
        abort(404)
    return handler(imp)


# ---- ファイルを置く ------------------------------------------------------------------------

@tables_bp.post("/upload", endpoint="upload")
def tables_upload():
    cfg = current_app.config
    try:
        stored = save_upload(request.files.get("file"), "tables", set(cfg["TABLE_ALLOWED_EXTENSIONS"]),
                             int(cfg["TABLE_MAX_UPLOAD_BYTES"]))
    except UploadError as exc:
        return _json_error(str(exc))
    path = upload_path(stored.stored_path)
    source_info: dict = {}
    try:
        if Path(stored.file_name).suffix.lower() in EXCEL_EXT:
            precheck_excel(path)
            source = open_source(path, stored.file_name, {"max_cells": cfg.get("EXCEL_MAX_CELLS", 500000)})
            sheets = [s for s in source.sheets() if not s.hidden] or source.sheets()
            if not sheets:
                raise UploadError("シートがないブックです。シートのあるブックを選んでください")
            # 開けても値を読めないブック（<v>NaN</v> など）は、取り込みを作る前にここで断る。
            # 先頭だけでなく全行を読む（下の方の値が読めないと、取り込みが読み込みの段から先へ進めない）
            for _row in source.rows(sheets[0].name):
                pass
            source_info = {"kind": "excel", "sheet": sheets[0].name}
        else:
            source = open_source(path, stored.file_name)
            sniff = source.sniff
            source_info = {"kind": "csv", "encoding": source.encoding, "delimiter": source.delimiter,
                           "errors": "strict"}
            if sniff is not None:
                source_info.update(bom=sniff.bom, preamble_rows=sniff.preamble_rows, sniff_warnings=sniff.warnings,
                                   confidence=sniff.confidence)
    except UploadError as exc:
        remove_upload(stored.stored_path)
        return _json_error(str(exc))
    except Exception:
        remove_upload(stored.stored_path)   # 思わぬエラーでもアップロードしたファイルを残さない（design.md 3.3）
        raise
    import_id = create_import(stored.file_name, stored.file_hash, stored.stored_path, source=source_info,
                                    session_id=current_session_id())
    if source_info.get("kind") == "excel":
        import_source(get_import(import_id), real=source).sheets()  # シート一覧を控えに入れる
    return jsonify({"import_id": import_id, "file_name": stored.file_name, "urls": _urls(import_id)})


# ---- 画面を離れたので捨てる ------------------------------------------------------------------
# 「その場でダウンロードしない限り、その場ですぐ捨てる」（利用者の指示 2026-09-20）。
# 画面を閉じた・隠したときに static/app.js の ragDiscard がここへ「捨てて」と送ってくる。
# navigator.sendBeacon で届くので中身の型は text/plain（get_json(force=True) で読む）。
# 応答は読めず、やり直しもできないので、いつでも 204 を返す（もう無い番号・ほかの人の番号・
# 処理中のものは core.purge 側で黙って外れる）。

@tables_bp.post("/discard", endpoint="discard")
def tables_discard():
    """この画面（このブラウザ）の、まだダウンロードしていない取り込みを捨てる。"""
    payload = request.get_json(force=True, silent=True) or {}
    ids = payload.get("import_ids") or []
    sid = current_session_id()
    try:
        if ids:
            discard_table_imports(ids, sid)
        else:
            purge_session(sid, documents=False)   # 帳票取り込み（別のタブ）は巻き込まない
    except Exception as exc:   # 捨て損ねてもブラウザには伝えられない。時間切れの片付けに任せる
        current_app.logger.warning("取り込みの片付けに失敗しました: %s", exc.__class__.__name__)
    return "", 204


# ---- 読み取り方（文字コード・区切り・シート） ---------------------------------------------------------

def _source_note(imp: dict, src: dict) -> str:
    """段の見出しに出す1行（どう読むか）。"""
    if _is_excel(imp):
        return f"シート: {src.get('sheet') or ''}"
    delimiters = dict(DELIMITER_CHOICES)
    return f"{src.get('encoding') or ''}／{delimiters.get(src.get('delimiter'), src.get('delimiter') or '')}"


def _panel_source(imp: dict):
    import_id = imp["id"]
    src = imp.get("source") or {}
    error = None
    sheets, sheet, list_like = [], None, {}
    try:
        source_obj = _open(imp)
        sheets = source_obj.sheets() if _is_excel(imp) else []
        sheet = _sheet(imp, source_obj)
        if _is_excel(imp) and len([s for s in sheets if not s.hidden]) <= 8:
            # 「一覧表らしい」「クロス集計らしい（対応していません）」を出し分ける（表の範囲の判定と合わせる）
            list_like = {s.name: source_obj.list_kind(s.name) for s in sheets if not s.hidden}
    except UploadError as exc:
        error = str(exc)
    html = render_part("tables.html", "part_source", imp=imp, src=src, error=error, sheets=sheets, sheet=sheet, list_like=list_like,
        encodings=ENCODING_CHOICES, delimiters=DELIMITER_CHOICES)
    note = _source_note(imp, {**src, "sheet": sheet or src.get("sheet")})
    return _panel(html, note=note, error=error, import_id=import_id)


@tables_bp.post("/imports/<int:import_id>/source")
def save_source(import_id: int):
    """読み取り方（シート・文字コード・区切り）。変えるたびに保存する。"""
    imp = _load_import(import_id)
    if _processing(imp):
        return _json_error(BUSY_MESSAGE, 409)
    form = _tables_payload() or request.form
    src = dict(imp.get("source") or {})
    before = (src.get("sheet"), src.get("encoding"), src.get("delimiter"), src.get("errors"))
    if _is_excel(imp):
        if form.get("sheet"):
            src["sheet"] = form["sheet"]
    else:
        allowed = {"encoding": {v for v, _label in ENCODING_CHOICES},
                   "delimiter": {v for v, _label in DELIMITER_CHOICES}}
        for name in ("encoding", "delimiter"):
            value = form.get(name)
            if value and value not in allowed[name]:
                return _json_error("この画面にない文字コード・区切り文字は選べません")
            if value:
                src[name] = value
        src["errors"] = "replace" if form.get("replace_errors") in ("on", True, "true", "1") else "strict"
    columns: dict = {"source": src, "status": "uploaded"}   # 読み取り方を保存したら読み込みからやり直す
    if (src.get("sheet"), src.get("encoding"), src.get("delimiter"), src.get("errors")) != before:
        src.pop("header_rows", None)
        src.pop("data_end_row", None)
        # 見出しが変わるので、この取り込みの列の対応づけは捨てて決め直す
        columns.update(spec_json="", spec_hash="")
    update_import(import_id, **columns)
    return jsonify({"ok": True, "note": _source_note(imp, src), "next": "layout",
                    "reset": ["layout", "columns", "ai", "preview", "done"]})


# ---- 表の範囲（見出し行・データの終わり） ----------------------------------------------------------

def _grid(source, sheet: str, layout) -> tuple[list[dict], int]:
    classes = {rc.index: rc for rc in layout.row_classes}
    rows, width = [], 0
    for row in source.rows(sheet, 1, GRID_ROWS):
        cells = [(c.text or "")[:GRID_CELL_CHARS] for c in row.cells[:GRID_COLS]]
        while cells and not cells[-1]:
            cells.pop()
        width = max(width, len(cells))
        rc = classes.get(row.index)
        rows.append({"index": row.index, "cells": cells, "kind": rc.kind if rc else ("blank" if row.is_blank else "title"),
                     "reason": rc.reason if rc else "", "hidden": bool(row.hidden),
                     "strike": any(getattr(c, "strike", False) for c in row.cells if c.text)})
    return rows, max(width, 1)


def _layout_json(layout) -> dict:
    return {
        "sheet": layout.sheet, "table_kind": layout.table_kind,
        "table_kind_label": TABLE_KIND_LABELS.get(layout.table_kind, layout.table_kind),
        "header_rows": layout.header_rows, "data_start": layout.data_start, "data_end": layout.data_end,
        "headers": layout.headers, "warnings": layout.warnings,
        "counts": {KIND_LABELS.get(k, k): v for k, v in (layout.counts or {}).items()},
        "rows": {str(rc.index): {"kind": rc.kind, "reason": rc.reason} for rc in layout.row_classes},
    }


def _panel_layout(imp: dict):
    job = _running_job(imp)
    if job is not None and job.get("kind") == "table_read":
        return _panel("", job=_job_info(job, imp["id"]), reading=True)
    try:
        source_obj = _open(imp)
        sheet = _sheet(imp, source_obj)
        guess = _layout(imp, source_obj, sheet)
        rows, width = _grid(source_obj, sheet, guess)
    except UploadError as exc:
        return _locked(str(exc))
    src = imp.get("source") or {}
    html = render_part("tables.html", "part_layout", imp=imp, layout=guess, info=_layout_json(guess), rows=rows,
                           width=width, sheet=sheet, data_end_saved=src.get("data_end_row") or "")
    return _panel(html, note=f"{guess.data_start}〜{guess.data_end}行目" if guess.data_end >= guess.data_start else "")


@tables_bp.post("/imports/<int:import_id>/layout/detect")
def layout_detect(import_id: int):
    imp = _load_import(import_id)
    data = _tables_payload()
    try:
        source_obj = _open(imp)
        sheet = _sheet(imp, source_obj)
        guess = _layout(imp, source_obj, sheet, header_rows=_int_list(data.get("header_rows")) or None,
                        data_end=_row_no(data.get("data_end")), use_saved=False)
    except UploadError as exc:
        return _json_error(str(exc))
    return jsonify(_layout_json(guess))


@tables_bp.post("/imports/<int:import_id>/layout")
def save_layout(import_id: int):
    """この範囲で読み込む。設定が無い・見出しが変わったときは列の対応づけへ、そうでなければ読み込みを始める。"""
    imp = _load_import(import_id)
    if _processing(imp):
        return _json_error(BUSY_MESSAGE, 409)
    data = _tables_payload() or request.form
    header_rows = _int_list(data.get("header_rows"))
    data_end = _row_no(data.get("data_end_row"))
    try:
        source_obj = _open(imp)
        sheet = _sheet(imp, source_obj)
        guess = _layout(imp, source_obj, sheet, header_rows=header_rows or None, data_end=data_end, use_saved=False)
    except UploadError as exc:
        return _json_error(str(exc))
    if guess.table_kind == "crosstab":
        return _json_error("月別集計のようなクロス集計の表には対応していません。1行＝1件の一覧表を選んでください")
    if not guess.headers:
        return _json_error("見出し行とデータの範囲が見つかりません。見出し行の番号を指定してください")
    if guess.data_end < guess.data_start:
        # 見出しは見つかっている。直すのは見出し行ではなくデータの範囲なので、そう伝える
        last_header = guess.header_rows[-1] if guess.header_rows else 0
        if data_end:
            return _json_error(f"データの終わりに{data_end}行目を指定すると、データの行が1行もありません"
                               f"（見出しは{last_header}行目）。空欄にするか、見出し行より下の行を指定してください")
        return _json_error(f"見出し行（{last_header}行目）より下にデータの行が1行もありません。"
                           "見出し行の番号を確かめるか、データのある表を選んでください")
    src = dict(imp.get("source") or {})
    src.update(sheet=sheet, header_rows=guess.header_rows, header_row=None, data_end_row=data_end or None)
    update_import(import_id, source=src)
    # 取り込み設定は保存しないので、範囲を決めたら必ず「列の対応づけ」へ進む（利用者の指示 2026-09-20）
    return jsonify({"ok": True, "next": "columns", "reset": ["columns", "ai", "preview", "done"]})


# ---- 列の対応づけ ------------------------------------------------------------------------------
# 画面で決めるのは「使う・役割」だけ。キー・型・単位・md での扱い・空欄＝上と同じは、見出しと値から
# 候補づくり（tables.mapping.suggest_columns・tables.dictionary）が決める（利用者の指示 2026-09-20）。


def _screen_role(role: str) -> str:
    """候補の役割を画面の選択肢に寄せる（画面に出さない役割＝人名・数値・区分などは「その他」）。"""
    return role if role in SCREEN_ROLE_KEYS else "attribute"


def _suggest(imp: dict):
    """保存した範囲で見出しと先頭のデータを読み、列ごとの候補を作る。戻り値: (layout, suggestions)"""
    source_obj = _open(imp)
    sheet = _sheet(imp, source_obj)
    guess = _layout(imp, source_obj, sheet)
    samples = source_obj.sample_rows(sheet, guess, 200)
    return guess, suggest_columns(guess.headers, samples)


def _unused_note(sugg) -> str:
    """はじめから「使わない」にしてある列の、その理由（見出しの下に小さく出す。ふつうの列は空）。

    理由が読めないと、チェックの外れている列を入れ直してよいのか分からない。
    「知らせ」の列は 2026-09-21 にやめた（ほとんどの行で空で、表の上の行と同じことを書いていた）。
    型エラー・空欄の割合は表の上（_columns_todo）に出すので、ここには書かない。
    「値は「日付」らしい」も出さない（日付の役割は日付として読める列にしか出ないので、読んでも直せない）。
    """
    if sugg.md != "omit":
        return ""
    return ("空欄だけなので、はじめから使わない設定にしています" if sugg.omit_reason == "blank"
            else "記録に不要な管理用の列らしいので、はじめから使わない設定にしています")


def _column_row(sugg) -> dict:
    """画面の1行（使う・見出し・役割・値の例と、使わない設定にしている理由）。"""
    return {"index": sugg.index, "header": sugg.header, "use": sugg.md != "omit",
            "role": _screen_role(sugg.role), "type": sugg.type, "examples": list(sugg.examples or [])[:3],
            "note": _unused_note(sugg)}


DATE_TYPES = ("date", "datetime")


def _roles_for(row: dict) -> list[tuple[str, str]]:
    """その列で選べる役割。

    「日付」は値が日付として読める列にだけ出す。年月だけの列（2025-03 など）を日付にしても
    保存が通らないので、画面で選べてしまうと直しようのない行き止まりになる。
    """
    return [(role, label) for role, label in SCREEN_ROLES
            if role != "date" or row["type"] in DATE_TYPES or row["role"] == "date"]


def _has_date_column(pairs) -> bool:
    """日付として読める列があるか（1つも無ければ「日付の列を選べ」とは言わない）。"""
    return any(sugg.type in DATE_TYPES for _row, sugg in pairs)


# 「列の対応づけは決まっている」の決まり（利用者の問い 2026-09-20「列の対応付けを行う意味は？」）。
# この段で決められるのは「出す／出さない」と四つの役割だけで、ふつうの一覧表ならそのどちらも
# 候補づくり（tables.mapping）が見出しと値から決めている。決めることが無いのに22行の表を出すのは
# 意味がないので、次の条件をすべて満たすときは表をたたんで要約1行だけ出す。
#   1. 識別番号の列がちょうど1つで、見出しが辞書と完全一致している（matched_by == "dictionary"）
#   2. 日付の列がちょうど1つで、同じく完全一致している
#   3. 対象の列は0か1つ。1つなら完全一致している（0なら要約に「対象の列はありません」と書く）
#   4. 経過の記録の列は0か1つ。1つなら完全一致している（0なら要約にそう書く）
#   5. 出す列のどれにも、出すかどうかを決め直す理由が無い
#      ＝ 読めない値がある（type_error_rate > 0）／ほとんど空欄（blank_rate >= UNSURE_BLANK_RATE）
# 似た語で当たっただけ（matched_by == "similar"）や、値の並びから当てた（"none"）列が四つの役割に
# 付いていると 1〜4 で外れる。役割が本当に合っているかは人にしか決められないので、表を開く。
# 半分くらい空欄なのは決め直す理由にしない（出しても困らないため。ここに出さなければ画面のどこにも出ない）。
UNSURE_BLANK_RATE = 0.9

# 四つの役割（画面で決められるもの）と、要約・上の行に出す短い名前（プルダウンの但し書きまでは書かない）。
# 識別番号と日付は無いと決まらない
_DECIDED_ROLES = [("key", "識別番号", True), ("date", "日付", True),
                  ("entity", "対象", False), ("log", "経過の記録", False)]


def _role_columns(pairs, role: str) -> list[tuple[dict, object]]:
    return [pair for pair in pairs if pair[0]["use"] and pair[0]["role"] == role]


def _columns_todo(pairs) -> list[str]:
    """表を開いて決めてもらうことを並べる（空なら要約だけでよい）。pairs: [(画面の1行, 候補)]"""
    todo: list[str] = []
    for role, label, required in _DECIDED_ROLES:
        found = _role_columns(pairs, role)
        if len(found) > 1:
            todo.append(f"{label}の列が{len(found)}つあります。1つにしてください")
        elif not found:
            # 日付として読める列が1つも無い表では、選びようがないので求めない
            if required and (role != "date" or _has_date_column(pairs)):
                todo.append(f"{label}の列が決まっていません。1つ選んでください")
        elif found[0][1].matched_by != "dictionary":
            todo.append(f"「{found[0][0]['header']}」を{label}として読み取ります。これでよいか確かめてください")
    for row, sugg in pairs:
        if not row["use"]:
            continue
        if sugg.type_error_rate:
            todo.append(f"列「{row['header']}」に読み取れない値があります"
                        f"（{round(sugg.type_error_rate * 100, 1)}%）。出すかどうか決めてください")
        elif sugg.blank_rate >= UNSURE_BLANK_RATE:
            todo.append(f"列「{row['header']}」はほとんど空欄です"
                        f"（{int(round(sugg.blank_rate * 100))}%）。出すかどうか決めてください")
    return todo


def _columns_summary(pairs) -> str:
    """決まっているときに出す1行。見つけた役割と、出す列・出さない列の数を正直に書く。"""
    named, missing = [], []
    for role, label, _required in _DECIDED_ROLES:
        found = _role_columns(pairs, role)
        if found:
            named.append(f"{found[0][0]['header']}＝{label}")
        elif role == "date":
            missing.append("日付の列はありません。記録は「日付なし」の1ファイルにまとめます。")
        else:
            missing.append(f"{label}の列はありません。")
    used = sum(1 for row, _s in pairs if row["use"])
    left = len(pairs) - used
    return ("、".join(named) + "として読み取ります。" + "".join(missing)
            + f"{len(pairs)}列のうち{used}列を Markdown に出します"
            + (f"（残る{left}列は出しません）。" if left else "（出さない列はありません）。"))


def _default_table_name(imp: dict) -> str:
    """表の名前の初期値。ファイル名から作る（設定は残らないので、名前もこの取り込みだけのもの）。"""
    return Path(imp["file_name"]).stem


def _build_spec(payload: dict, guess, suggestions):
    """画面の選択（使う・役割）と列の候補から取り込み設定を作る。戻り値: (spec, errors)"""
    choices: dict[int, dict] = {}
    for row in payload.get("columns") or []:
        if isinstance(row, dict):
            choices[_int(row.get("index"), -1)] = row
    used: list[dict] = []
    date_errors: list[str] = []
    for sugg in suggestions:
        choice = choices.get(sugg.index)
        if not (bool(choice.get("use")) if choice else sugg.md != "omit"):
            continue
        role = str((choice or {}).get("role") or "")
        if role not in SCREEN_ROLE_KEYS:
            role = _screen_role(sugg.role)
        # 画面でそのままなら候補の役割（人名・数値・区分など）を活かす。選び直したときだけその役割にする
        role = sugg.role if role == _screen_role(sugg.role) else role
        if role == "date" and sugg.type not in DATE_TYPES:
            # 画面には出さない選び方だが、古い画面から送られたときに「型を日付にしてください」
            # （利用者には直せない指示）ではなく、何が起きているかを返す
            date_errors.append(f"列「{sugg.header}」の値は日付として読めないので、日付の列にはできません")
        type_ = "text" if role == "log" else sugg.type
        # 「使う」列は必ず md に出す（候補が「出さない」でも、チェックを入れたのだから出す）
        md = sugg.md if sugg.md != "omit" else ("body" if type_ == "text" else "attribute")
        used.append({"index": sugg.index, "header": sugg.header, "key": sugg.key, "display": sugg.display,
                     "type": type_, "role": role, "unit": sugg.unit, "md": md,
                     "fill_down_blank": bool(sugg.fill_down_blank)})
    name = str(payload.get("name") or "").strip()
    header_rows = list(getattr(guess, "header_rows", None) or [1])
    spec = spec_from_suggestions(name, {"table_kind": "list", "header_rows": header_rows}, used)
    errors = validate_spec(spec)
    if date_errors:
        # 「型を日付にしてください」は画面に直す場所が無いので、こちらの言い方に置き換える
        errors = date_errors + [e for e in errors if not e.startswith("日付の列「")]
    if sum(1 for u in used if u["role"] == "log") > 1:
        # 黙って最初の列だけを経過の記録にしない（2列目は日付ごとに分けられず1行につながれて出てしまう）
        errors.insert(0, "経過の記録の列は1つだけにしてください")
    return spec, errors


def _panel_columns(imp: dict):
    import_id = imp["id"]
    job = _running_job(imp)
    if job is not None and job.get("kind") == "table_read":
        return _panel("", job=_job_info(job, import_id), reading=True)
    try:
        guess, suggestions = _suggest(imp)
    except UploadError as exc:
        return _locked(str(exc))
    if not guess.headers:
        return _locked("見出しが見つかりません。上の「表の範囲」で見出し行を指定してください")
    rows = [_column_row(s) for s in suggestions]
    spec = _spec_for(imp)
    if spec is not None:
        # この取り込みで一度保存していれば、そのときの「使う・役割」を残す（開き直しても選び直さずに済む）
        by_header = {}
        for col in spec.columns:
            for header in (col.headers or [col.display]):
                by_header[header] = col
        for row in rows:
            col = by_header.get(row["header"])
            row["use"] = col is not None
            if col is not None:
                row["role"] = _screen_role(col.role)
    for row in rows:
        row["roles"] = _roles_for(row)
    # 決めることが無ければ表をたたんで要約1行にする（表は隠すだけで残すので、保存で送る中身は同じ）
    pairs = list(zip(rows, suggestions))
    todo = _columns_todo(pairs)
    html = render_part("tables.html", "part_columns", rows=rows, todo=todo,
        summary="" if todo else _columns_summary(pairs),
        name=spec.name if spec is not None else _default_table_name(imp),
        save_url=url_for("tables.save_columns", import_id=import_id))
    return _panel(html, note=f"{sum(1 for r in rows if r['use'])}／{len(rows)}列")


@tables_bp.post("/imports/<int:import_id>/columns")
def save_columns(import_id: int):
    imp = _load_import(import_id)
    if _processing(imp):
        return _json_error(BUSY_MESSAGE, 409)
    try:
        guess, suggestions = _suggest(imp)
    except UploadError as exc:
        return _json_error(str(exc))
    spec, errors = _build_spec(_tables_payload(), guess, suggestions)
    if errors:
        return _json_error(errors[0], errors=errors)
    save_spec(import_id, spec)
    job_id = start_read_job(import_id)
    return jsonify({"ok": True, "next": _after_read_panel(spec), "reset": ["ai", "preview", "done"],
                    "job": _job_info(get_job(job_id), import_id), "reading": True})


@tables_bp.post("/imports/<int:import_id>/read")
def reread(import_id: int):
    """読み込み直し（失敗したとき・もう一度読むとき）。"""
    imp = _load_import(import_id)
    spec = _spec_for(imp)
    if spec is None:
        return _json_error("先に列の対応づけを保存してください")
    if _processing(imp):
        # AI整形の実行中・一時停止中に読み込み直すと、読み込みが AI整形の後ろで待ち続ける
        return _json_error(BUSY_MESSAGE, 409)
    job_id = start_read_job(import_id)
    return jsonify({"ok": True, "next": _after_read_panel(spec), "reset": ["ai", "preview", "done"],
                    "job": _job_info(get_job(job_id), import_id), "reading": True})


# 中止できるジョブ（取り込みの job_id に入るもの）と、中止したときに戻す状態
CANCELLABLE_JOBS = {"table_read": "uploaded", "table_render": "preview"}


@tables_bp.post("/imports/<int:import_id>/cancel")
def cancel_job(import_id: int):
    """［中止］: 読み込み・Markdown作成のジョブを止める。"""
    imp = _load_import(import_id)
    job = get_job(imp["job_id"]) if imp.get("job_id") else None
    if job is None or job["kind"] not in CANCELLABLE_JOBS or job.get("finished"):
        return _json_error("中止できる処理がありません（すでに終わっている可能性があります）")
    request_cancel(job["id"])
    after = get_job(job["id"])
    # 取り込みの状態は読み直す（この間にジョブが成功して preview/confirmed を書いていたら、巻き戻してはいけない）
    fresh = get_import(import_id)
    if (after and after["status"] == "cancelled" and fresh is not None
            and fresh["status"] in ("reading", "confirming")):
        # 待機中のまま中止されたときは本体が動かないので、ここで取り込みの状態を戻す
        update_import(import_id, status=CANCELLABLE_JOBS[job["kind"]])
    return jsonify({"ok": True, "message": "処理を中止しました"})


@tables_bp.post("/imports/<int:import_id>/delete")
def delete_import(import_id: int):
    """間違えて取り込んだファイルを消す（読み込んだ内容・作った Markdown も消える）。"""
    imp = _load_import(import_id)
    if imp["status"] in ("reading", "confirming"):
        return _json_error("処理中の取り込みは削除できません。終わるか中止してから削除してください")
    if _ai_running(import_id):
        return _json_error("AI整形の実行中は削除できません。AI整形を中止してから削除してください")
    if _trial_running(import_id):
        return _json_error(TRIAL_BUSY_MESSAGE)
    preview_job = latest_job("table_import", import_id, kind="table_preview")
    if preview_job is not None and not preview_job.get("finished"):
        # 下書きの作成中に消すと、開いているファイルが消え残ったり、消したあとに記録の md が書き戻されたりする
        return _json_error("Markdownの下書きを作っている間は削除できません。終わってから削除してください")
    purge_table_import(import_id)
    # ファイル名は出さない: 取引先名や「社外秘」を含みうる名前を画面に残さない（design.md 3.3）
    message = "取り込みを削除しました（元のファイル・読み込んだ内容・作った Markdown を消しました）"
    if purge_incomplete():
        # 使用中などで消し切れなかったファイルは中身を 0 バイトにしてある（次の起動時に片付く。design.md 3.3）
        message = "取り込みを削除しました（一部のファイルは使用中で消し切れず、中身を空にしました。次の起動時に片付きます）"
    return jsonify({"ok": True, "message": message, "reload": True})


def _not_ready_reason(imp: dict, spec) -> str:
    """読み込みが終わっていない・失敗したときに、段に出す1行。"""
    if spec is None:
        return "先に「列の対応づけ」を保存してください"
    if imp["status"] == "failed":
        return (imp.get("stats") or {}).get("error") or "表を読み込めませんでした"
    if imp["status"] == "uploaded":
        return "先に「表の範囲」で読み込んでください"
    return ""


# ---- AI整形（任意。経過の記録の列があるときだけ） --------------------------------------------------

def _ai_job(import_id: int) -> dict | None:
    return latest_job("table_import", import_id, kind="ai_format")


def _ai_connection_ctx() -> dict:
    """AI接続の状態（段の中の「AI接続」パネル）。"""
    import llm

    status = llm.admin_status()
    ready = llm.is_configured()
    external, model = False, ""
    if ready:
        try:
            settings = llm.job_client_settings()
            external = not llm.is_local_endpoint(settings.get("chat_url") or "")
            model = settings.get("model") or ""
        except Exception:
            ready = False
    return {"ai_status": status, "ai_ready": ready, "external": external, "model": model,
            "save_url": url_for("tables.save_ai_connection"), "test_url": url_for("tables.test_ai_connection"),
            "models_url": url_for("tables.refresh_ai_models")}


def _panel_ai(imp: dict):
    import_id = imp["id"]
    job = _running_job(imp)
    if job is not None:
        return _panel("", job=_job_info(job, import_id), reading=True)
    spec = _spec_for(imp)
    if spec is not None and spec.log_stage is None:
        return _locked("経過の記録の列がないので、この取り込みでは使いません")
    reason = _not_ready_reason(imp, spec)
    if reason:
        return _locked(reason)
    import aiproc as ai_items

    log_key = spec.log_stage.column
    col = spec.column(log_key)
    row_choices = []
    for rec in load_rows(import_id):
        text = str((rec.get("values") or {}).get(log_key) or "").strip()
        if text:
            label = rec["key"] + "｜" + " ".join(text.split())[:40]
            row_choices.append((rec["key"], label))
            if len(row_choices) >= 200:
                break
    ai_job = _ai_job(import_id)
    html = render_part("tables.html", "part_ai", imp=imp, log_display=col.display if col else log_key, row_choices=row_choices,
        trial_keys=[k for k, _ in row_choices[:10]], job=ai_job, job_active=bool(ai_job and not ai_job.get("finished")),
        job_url=url_for("tables.api_job", job_id=ai_job["id"]) if ai_job else None,
        counts=ai_items.counts(imp["template_id"], "log", import_id=import_id),
        scopes=SCOPE_LABELS, item_labels=ai_items.STATUS_LABELS, **_ai_connection_ctx())
    return _panel(html, note=f"経過の記録の列: {col.display if col else log_key}",
                  job=_job_info(ai_job, import_id) if ai_job and not ai_job.get("finished") else None)


@tables_bp.post("/imports/<int:import_id>/ai/split-preview")
def ai_split_preview(import_id: int):
    import aiproc
    from logproc import format_author, format_when, review_notes
    from logproc import render_timeline
    import llm

    _load_import(import_id)
    row_key = str(_tables_payload().get("row_key") or "")
    try:
        data = aiproc.load_rows_for_ai(import_id)
        works = aiproc.prepare_works(data, ["log"], [row_key])
    except aiproc.AIJobError as exc:
        return _json_error(str(exc))
    if not works:
        return _json_error("行が見つかりません")
    w = works[0]
    parse = w.parse
    segments = [{
        "id": s.id, "body": s.body, "when": format_when(s.when), "when_estimated": bool(s.when and s.when.estimated),
        "author": format_author(s.author), "author_estimated": bool(s.author and s.author.estimated),
        "marks": list(s.marks or []), "identifiers": list(s.identifiers or []), "quantities": list(s.quantities or []),
        "plans": list(s.plans or []),
    } for s in parse.segments]
    return jsonify({
        "row_key": row_key, "kind": parse.kind, "order": parse.order, "text": parse.text, "segments": segments,
        "route": w.route, "reason": w.reason, "ai_ready": llm.is_configured(), "notes": review_notes(parse),
        "timeline": render_timeline(parse, w.entity_label, glossary=(w.stage.glossary or None) if w.stage else None),
        "sent_text": aiproc.messages_text(w.messages) if w.messages else "",
    })


@tables_bp.post("/imports/<int:import_id>/ai/trial")
def ai_trial(import_id: int):
    import aiproc
    import llm

    imp = _load_import(import_id)
    spec = _spec_for(imp)
    if spec is None or spec.log_stage is None:
        return _json_error("経過の記録の列がありません")
    if imp["status"] == "confirmed":
        # 試し実行の結果は下書きに入るが、zip は確定したときの md を渡す。食い違わないよう断る
        return _json_error("確定後は試し実行できません（結果が確定した Markdown に入らないため）。"
                           "試すときは列の対応づけからもう一度読み込んでください")
    payload = _tables_payload()
    row_key = str(payload.get("row_key") or "")
    try:
        settings = llm.job_client_settings()
        if not llm.is_local_endpoint(settings.get("chat_url") or "") and not payload.get("confirm_external"):
            return _json_error("対応内容が外部のAIサービスに送信されます。確認のチェックを入れてください")
        with _TRIALS_LOCK:
            _TRIALS[import_id] += 1
        try:
            data = aiproc.load_rows_for_ai(import_id)
            result = aiproc.trial_row(import_id, row_key, stage_ids=["log"], data=data)
        finally:
            with _TRIALS_LOCK:
                _TRIALS[import_id] -= 1
                if _TRIALS[import_id] <= 0:
                    del _TRIALS[import_id]
    except llm.LLMNotConfigured:
        return _json_error("AI接続が設定されていません。この段の「AI接続」で設定してください")
    except llm.LLMCallError as exc:
        return _json_error(str(exc))
    except aiproc.AIJobError as exc:
        return _json_error(str(exc))
    stage = next((s for s in result["stages"] if s.get("stage_id") == "log"), None) or {}
    accepted = stage.get("result") if stage.get("status") in ("ok", "flagged") else None
    record = next((r for r in data.rows if r["key"] == row_key), None)
    md = ""
    if record is not None:
        ai_results = {row_key: {"status": "ok", "result": accepted}} if accepted else {}
        md = "\n".join(record_block(record, spec, ai_results, people_index_for(spec, data.rows)))
    checks = stage.get("checks") or {}
    issues = [i.get("message", "") for i in (checks.get("issues") or []) if isinstance(i, dict)]
    return jsonify({
        "row_key": row_key, "status": stage.get("status"), "route": stage.get("route"), "reason": stage.get("reason"),
        "error": stage.get("error") or "", "points": ai_point_lines(accepted, spec.log_stage.glossary) if accepted else [],
        "timeline": stage.get("timeline") or [], "issues": issues, "markdown": md, "cached": stage.get("cached"),
        "stats": aiproc.trial_stats([result]),
    })


@tables_bp.post("/imports/<int:import_id>/ai/estimate")
def ai_estimate(import_id: int):
    import aiproc as ai_estimate_mod
    import aiproc
    import llm

    _load_import(import_id)
    data = _tables_payload()
    scope = data.get("scope") if data.get("scope") in SCOPE_LABELS else "pending"
    try:
        settings = llm.job_client_settings()
        result = ai_estimate_mod.estimate(import_id, [s for s in data.get("trials") or [] if isinstance(s, dict)],
                                          scope=scope, concurrency=_int(data.get("concurrency")) or None,
                                          settings=settings, stage_ids=["log"])
    except llm.LLMNotConfigured:
        return _json_error("AI接続が設定されていません")
    except aiproc.AIJobError as exc:
        return _json_error(str(exc))
    return jsonify(result)


@tables_bp.post("/imports/<int:import_id>/ai/run")
def ai_run(import_id: int):
    import aiproc
    import llm

    imp = _load_import(import_id)
    data = _tables_payload()
    if imp["status"] != "preview":
        return _json_error("表の読み込みが終わってから実行してください")
    job = _ai_job(import_id)
    if job and not job.get("finished"):
        return _json_error("AI整形はすでに実行中です")
    scope = data.get("scope") if data.get("scope") in SCOPE_LABELS else "pending"
    concurrency = min(16, max(1, _int(data.get("concurrency"), 1) or 1))
    try:
        settings = llm.job_client_settings()
        if not llm.is_local_endpoint(settings.get("chat_url") or "") and not data.get("confirm_external"):
            return _json_error("対応内容が外部のAIサービスに送信されます。確認のチェックを入れてください")
        job_id = aiproc.start_ai_job(import_id, scope=scope, concurrency=concurrency, stage_ids=["log"])
    except llm.LLMNotConfigured:
        return _json_error("AI接続が設定されていません")
    return jsonify({"ok": True, "job_id": job_id, "job_url": url_for("tables.api_job", job_id=job_id)})


@tables_bp.post("/imports/<int:import_id>/ai/<action>")
def ai_control(import_id: int, action: str):
    _load_import(import_id)
    handlers = {"pause": request_pause, "resume": request_resume, "cancel": request_cancel}
    if action not in handlers:
        abort(404)
    job = _ai_job(import_id)
    ok = bool(job) and handlers[action](job["id"])
    if not ok:
        return _json_error("AI整形のジョブを操作できませんでした（すでに終わっている可能性があります）")
    return jsonify({"ok": True})


# ---- AI接続（AI整形の段の中。設定画面は無い） ---------------------------------------------------

@tables_bp.post("/ai-connection")
def save_ai_connection():
    import llm

    data = _tables_payload()
    models = [str(m).strip() for m in (data.get("models") or []) if str(m).strip()]
    models += [m.strip() for m in str(data.get("add_models") or "").splitlines() if m.strip()]
    try:
        llm.save_admin({
            "models": list(dict.fromkeys(models)),
            "default": str(data.get("default") or ""),
            "chat_url": str(data.get("chat_url") or ""),
            "models_url": str(data.get("models_url") or ""),
            "api_key": str(data.get("api_key") or ""),
            "api_key_clear": bool(data.get("api_key_clear")),
        })
    except ValueError as exc:
        return _json_error(str(exc))
    except OSError as exc:
        return _json_error(f"設定ファイルに書き込めませんでした（{exc.strerror or exc.__class__.__name__}）。"
                           "少し待ってから、もう一度保存してください")
    return jsonify({"ok": True, "message": "AI接続の設定を保存しました"})


@tables_bp.post("/ai-connection/models")
def refresh_ai_models():
    """APIからモデル一覧を取得する。"""
    import llm

    try:
        catalog = llm.model_catalog(refresh=True)
    except Exception as exc:
        return _json_error(f"モデル一覧を取得できませんでした: {llm.friendly_error(exc)}")
    return jsonify({"ok": True, "models": catalog, "message": f"APIからモデル一覧を取得しました（{len(catalog)}件）"})


@tables_bp.post("/ai-connection/test")
def test_ai_connection():
    """接続テスト: モデル一覧の取得と、1回の短いチャット。"""
    import llm

    if not llm.is_configured():
        return jsonify({"ok": False, "steps": [{"name": "設定", "ok": False,
                                                "detail": "APIキーまたは接続先が未設定です。"}]})
    steps = []
    started = time.monotonic()
    try:
        names = llm.fetch_api_models(refresh=True)
        steps.append({"name": "モデル一覧の取得", "ok": True,
                      "detail": f"{len(names)}件のモデルが見つかりました（{_ms(started)}ミリ秒）"})
    except Exception as exc:
        steps.append({"name": "モデル一覧の取得", "ok": False, "detail": llm.friendly_error(exc)})
    model = llm.current_model()
    started = time.monotonic()
    try:
        # AI整形のジョブと同じ呼び出し口（再試行なし・明示のタイムアウト）で確かめる
        result = llm.chat_raw(llm.job_client_settings(),
                              [{"role": "user", "content": "接続テストです。「OK」とだけ返してください。"}], max_tokens=20)
        reply = (result.text or "").strip()
        steps.append({"name": f"チャット（{model}）", "ok": True,
                      "detail": f"応答あり: {reply[:40] or '（本文なし）'}（{_ms(started)}ミリ秒）"})
    except Exception as exc:
        steps.append({"name": f"チャット（{model}）", "ok": False, "detail": llm.friendly_error(exc)})
    # チャットが通れば AI 整形は使える（モデル一覧APIが無い互換サーバもある）
    return jsonify({"ok": steps[-1]["ok"], "steps": steps})


def _ms(started: float) -> int:
    return int((time.monotonic() - started) * 1000)


# ---- 内容の確認 ---------------------------------------------------------------------------

def _preview_job(import_id: int, imp: dict, spec, retry: bool = False) -> dict:
    """確認の段の md を作るジョブ。同じ入力のジョブがあればそれを見せ、なければ始める。

    retry: 同じ入力で失敗したジョブを作り直す（［もう一度作る］。ファイルの一時的なロックなどで失敗したとき）。
    """
    signature = preview_signature(import_id, imp, spec)
    job = latest_job("table_import", import_id, kind="table_preview")
    if job is not None and (job.get("params") or {}).get("signature") == signature:
        if not job.get("finished") or (job["status"] == "failed" and not retry):
            return job  # 動いている途中、または同じ入力で失敗したまま（開き直すたびに作り直さない）
    return get_job(start_preview_job(import_id, signature))


def _display_with_unit(col) -> str:
    """データ表の見出し。Markdown は「停止時間: 5分」と単位を付けて書くので、見出しにも単位を添える。"""
    unit = (col.unit or "").strip()
    if unit and unit not in col.display:
        return f"{col.display}（{unit}）"
    return col.display


def _panel_preview(imp: dict, page: int = 1):
    import_id = imp["id"]
    job = _running_job(imp)
    if job is not None:
        return _panel("", job=_job_info(job, import_id), reading=True)
    spec = _spec_for(imp)
    reason = _not_ready_reason(imp, spec)
    if reason:
        return _locked(reason, failed=imp["status"] == "failed",
                       retry_url=url_for("tables.reread", import_id=import_id) if spec is not None else None)
    if _ai_running(import_id):
        return _locked("AI整形が動いています（一時停止中を含む）。再開して終わらせるか、中止してから確認してください")
    files = ready_preview_files(import_id, imp, spec)
    if files is None:
        # 件数分の md を作るのに時間がかかるので、ジョブにして同じ画面に進み具合を出す
        draft = _preview_job(import_id, imp, spec)
        if draft.get("finished") and draft["status"] != "done":
            return _panel(render_part("tables.html", "part_preview_failed", error=draft.get("message") or "Markdownの下書きを作れませんでした"),
                          failed=True)
        return _panel("", job=_job_info(draft, import_id), building=True)
    stats = imp.get("stats") or {}
    issues = load_issues(import_id)
    page = max(1, page)
    data_rows, total = load_rows_page(import_id, (page - 1) * DATA_PAGE, DATA_PAGE)
    total_pages = max(1, (total + DATA_PAGE - 1) // DATA_PAGE)
    if page > total_pages:
        page = total_pages
        data_rows, _total = load_rows_page(import_id, (page - 1) * DATA_PAGE, DATA_PAGE)
    blocking = has_blocking(issues)
    html = render_part("tables.html", "part_preview", imp=imp, spec=spec, stats=stats, issues=issues[:ISSUES_SHOWN],
        issue_total=len(issues), counts=count_levels(issues), blocking=blocking, files=files, data_rows=data_rows,
        page=page, total_pages=total_pages, columns=[(c.key, _display_with_unit(c)) for c in spec.columns],
        confirmed=imp["status"] == "confirmed", delete_note=TABLES_DELETE_ON_DOWNLOAD_NOTE)
    return _panel(html, note=f"{stats.get('records') or 0}件・{len(files)}ファイル", blocking=blocking,
                  confirmed=imp["status"] == "confirmed")


@tables_bp.post("/imports/<int:import_id>/preview")
def start_preview(import_id: int):
    """確認の段の Markdown の下書きを作り直す（失敗したときの［もう一度作る］）。"""
    imp = _load_import(import_id)
    spec = _spec_for(imp)
    if spec is None or imp["status"] not in ("preview", "confirmed"):
        return _json_error("表の読み込みが終わってから確認してください")
    job = _preview_job(import_id, imp, spec, retry=True)
    return jsonify({"ok": True, "job": _job_info(job, import_id), "building": True})


@tables_bp.get("/imports/<int:import_id>/preview/file")
def preview_file(import_id: int):
    _load_import(import_id)
    text = md_text(import_files(import_id)["preview"], request.args.get("name", ""))
    if text is None:
        return _json_error("ファイルが見つかりません", 404)
    return jsonify({"name": request.args.get("name"), "text": text})


@tables_bp.get("/imports/<int:import_id>/issues.csv")
def issues_csv(import_id: int):
    imp = _load_import(import_id)
    data = build_issues_csv(load_issues(import_id))
    name = f"{Path(imp['file_name']).stem}_問題一覧.csv"
    return set_download_name(
        send_file(io.BytesIO(data), mimetype="text/csv", as_attachment=True, download_name=name), name, "issues")


@tables_bp.post("/imports/<int:import_id>/confirm", endpoint="confirm")
def tables_confirm(import_id: int):
    """確定して Markdown を作る（できたら同じ画面にダウンロードのボタンを出す）。"""
    imp = _load_import(import_id)
    spec = _spec_for(imp)
    if spec is None or imp["status"] not in ("preview", "confirmed"):
        return _json_error("この取り込みはまだ確定できる状態ではありません")
    if _ai_running(import_id):
        return _json_error("AI整形の実行中は確定できません。終わるか中止してから確定してください")
    if has_blocking(load_issues(import_id)):
        return _json_error("エラーが残っているため確定できません。問題一覧を確認して、範囲や列の対応づけを直してください")
    job_id = start_render_job(import_id)
    return jsonify({"ok": True, "next": "done", "job": _job_info(get_job(job_id), import_id), "building": True})


# ---- 確定してダウンロード（zip） ------------------------------------------------------------------

def _download_base(imp: dict, spec) -> str:
    stem = safe_filename_part(spec.file_prefix if spec is not None else Path(imp["file_name"]).stem)
    return f"{stem}_{datetime.now().strftime('%Y-%m-%d')}"


def _panel_done(imp: dict):
    import_id = imp["id"]
    job = _running_job(imp)
    if job is not None:
        return _panel("", job=_job_info(job, import_id), building=True)
    if imp["status"] != "confirmed":
        return _locked("上の「内容の確認」で［確定してMarkdownを作成］を押してください")
    spec = _spec_for(imp)
    files = [{"name": p.name, "size": p.stat().st_size} for p in md_paths(import_id)]
    if not files:
        # 記録0件のまま確定された取り込み（渡せるものが無いので、ダウンロードのボタンは出さない）
        return _locked("この取り込みから作られた Markdown はありません（取り込める行がありませんでした）。"
                       "表の範囲か元のファイルを見直してください")
    html = render_part("tables.html", "part_done", imp=imp, files=files, stats=imp.get("stats") or {},
                           delete_note=TABLES_DELETE_ON_DOWNLOAD_NOTE, delete_confirm=TABLES_DELETE_ON_DOWNLOAD_CONFIRM)
    return _panel(html, note=f"{len(files)}ファイル")


@tables_bp.get("/imports/<int:import_id>/download.zip")
def download_zip(import_id: int):
    """zip を渡し、渡し終えた取り込みのデータを消す（design.md 3.3）。

    zip は全体をメモリに作ってから消すので、消す処理で中身が欠けることはない。作れなかったときは何も消さない。
    """
    imp = _load_import(import_id)
    spec = _spec_for(imp)
    if imp["status"] != "confirmed" or spec is None:
        flash("先に [確定してMarkdownを作成] を押してください", "error")
        return redirect(url_for("tables.new"))
    if not md_paths(import_id):
        # 確定はしたが記録が0件だった（押したばかりのボタンをもう一度押せ、とは言わない）
        flash("作成された Markdown がありません（取り込める行がありませんでした）。表の範囲か元のファイルを見直してください",
              "error")
        return redirect(url_for("tables.new"))
    if _ai_running(import_id):
        flash("AI整形の実行中はダウンロードできません。終わるか中止してからダウンロードしてください", "error")
        return redirect(url_for("tables.new"))
    if _trial_running(import_id):
        flash(TRIAL_BUSY_MESSAGE, "error")
        return redirect(url_for("tables.new"))
    try:
        data = build_download(import_id)
    except FileNotFoundError:
        # 同時に押した別のダウンロードが渡し終えてデータを消した（ダブルクリックなど）。500 ではなく「ダウンロード済み」。
        return _handout_lost(import_id, "ダウンロード")
    name = f"{_download_base(imp, spec)}.zip"
    response = send_file(io.BytesIO(data), mimetype="application/zip", as_attachment=True, download_name=name,
                         conditional=False)   # Range でも全体を返す（一部だけ渡して消すことが無いように）
    set_download_name(response, name, "records")
    return purge_after_send(response, purge_table_import, import_id)


def _handout_lost(import_id: int, what: str):
    """渡す途中で md が消えたとき。取り込みが残っていて確定済みでなければ、別のタブで読み込み直しが始まった
    （「ダウンロード済み」の 404 は事実と違うので、画面に戻して知らせる）。消えていれば渡し終えた・削除した。"""
    now = get_import(import_id)
    if now is None or now["status"] == "confirmed":
        abort(404)
    flash(f"読み込み直しが始まったため、{what}を取りやめました。確定し直してからもう一度押してください", "error")
    return redirect(url_for("tables.new"))


# ---- 段の割り当て ---------------------------------------------------------------------------

def _panel_rows(imp: dict):
    """データ（100行ずつ）のページ送りだけ作り直す。"""
    return _panel_preview(imp, page=max(1, _int(request.args.get("page"), 1) or 1))


_PANELS = {"source": _panel_source, "layout": _panel_layout, "columns": _panel_columns, "ai": _panel_ai,
           "preview": _panel_rows, "done": _panel_done}


# ---- ジョブの進捗（JSON） ------------------------------------------------------------------------------

def _job_visible(job: dict) -> bool:
    """このブラウザの取り込みのジョブか（ほかの人のジョブは進み具合も見せない）。

    取り込みが消えている（ダウンロード済み・削除済み）ジョブは、持ち主が分からないので見せない。
    """
    if (job.get("ref_type") or "") != "table_import" or job.get("ref_id") is None:
        return False
    imp = get_import(job["ref_id"])
    return imp is not None and owns(imp)


def api_job(job_id: int):
    job = get_job(job_id)
    if job is None or not _job_visible(job):
        return _json_error("ジョブが見つかりません", 404)
    return jsonify({k: job.get(k) for k in ("id", "kind", "status", "status_label", "progress", "message", "result",
                                            "finished")})


@tables_bp.record_once
def _register_api(setup_state) -> None:
    # /api/jobs は /tables の外に置く（design.md 2.4）
    setup_state.app.add_url_rule("/api/jobs/<int:job_id>", "tables.api_job", api_job)
