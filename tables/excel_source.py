"""Excel(.xlsx/.xlsm) の読み込み。シートの XML を先頭から順に読み、ブック全体をメモリに載せない。

通常モード（load_workbook(data_only=True)）で開いたときと同じ CellInfo を返す:
- セルの値の変換は openpyxl の WorkSheetParser をそのまま使う（日付・時刻・共有文字列・インライン文字列）
- 結合セル・列の非表示・ウィンドウ枠・オートフィルタ・テーブル・ハイパーリンク・コメントは、
  sheetData を除いたシートの XML を openpyxl の WorksheetReader で読む（結合範囲の値の消去やリンク先の値もそのまま）
- 値のある最終行・最終列とセル数は、開くときに expat で1回数える（セル数の上限もここで確かめる）
行が昇順に並んでいないシート（まれ）は、そのシートだけ全行を読んでから並べ直す。
"""
from __future__ import annotations

import io
import re
import warnings
import zipfile
from pathlib import Path
from typing import Iterator
from xml.parsers import expat

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.comments.comment_sheet import CommentSheet
from openpyxl.packaging.relationship import RelationshipList, get_dependents, get_rels_path
from openpyxl.styles.numbers import BUILTIN_FORMATS, BUILTIN_FORMATS_MAX_SIZE
from openpyxl.utils import column_index_from_string, range_boundaries
from openpyxl.utils.cell import coordinate_to_tuple
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from openpyxl.worksheet._reader import INLINE_STRING, WorksheetReader, WorkSheetParser
from openpyxl.worksheet.table import Table
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.xml.constants import COMMENTS_NS
from openpyxl.xml.functions import fromstring

from tables.source import CellInfo, SheetInfo, SourceRow, UploadError, cell_text

DEFAULT_MAX_CELLS = 500_000
_DIMENSION_RE = re.compile(rb'<dimension ref="([A-Z]+\d+(?::[A-Z]+\d+)?)"')
_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_SHEETDATA_START = re.compile(rb"<((?:[A-Za-z_][\w.\-]*:)?)sheetData\b[^>]*?(/?)>")
_SHEETDATA_END = re.compile(rb"</(?:[A-Za-z_][\w.\-]*:)?sheetData\s*>")
_CHUNK = 1024 * 1024
_STAT_KEYS = ("cell_count", "max_row", "max_col", "ordered")
_OPEN_ERROR = "Excelファイルとして開けませんでした。.xlsx 形式で保存し直してください"


def _too_many_cells(count: int, limit: int) -> UploadError:
    return UploadError(
        f"セル数が上限（{limit:,}セル）を超えています（約{count:,}セル）。"
        "CSVで保存してから取り込んでください（書式による判定は使えなくなります）"
    )


def estimate_cell_count(path) -> int:
    """シートXMLの <dimension> から、読み込み前にセル数の目安を出す。"""
    total = 0
    try:
        with zipfile.ZipFile(path) as zf:
            for name in zf.namelist():
                if not (name.startswith("xl/worksheets/sheet") and name.endswith(".xml")):
                    continue
                with zf.open(name) as f:
                    head = f.read(4096)
                m = _DIMENSION_RE.search(head)
                if not m:
                    continue
                ref = m.group(1).decode()
                if ":" not in ref:
                    total += 1
                    continue
                min_col, min_row, max_col, max_row = range_boundaries(ref)
                total += (max_row - min_row + 1) * (max_col - min_col + 1)
    except (zipfile.BadZipFile, OSError, ValueError):
        return 0
    return total


# ---- シートの XML: sheetData の外（結合・列・リンクなど） ----

def _split_sheet_xml(f) -> bytes:
    """sheetData の中身を除いたシートの XML（行が大きくても前後だけを持つ）。"""
    buf = b""
    while True:
        chunk = f.read(_CHUNK)
        buf += chunk
        m = _SHEETDATA_START.search(buf)
        if m:
            break
        if not chunk:
            return buf  # sheetData がない
    before, prefix = buf[:m.start()], m.group(1)
    empty = b"<" + prefix + b"sheetData/>"
    if m.group(2):  # <sheetData/>
        return before + empty + buf[m.end():] + f.read()
    rest = buf[m.end():]
    while True:
        e = _SHEETDATA_END.search(rest)
        if e:
            return before + empty + rest[e.end():] + f.read()
        chunk = f.read(_CHUNK)
        if not chunk:
            return before + empty  # 閉じタグがない（openpyxl でも読めない形）
        rest = rest[-64:] + chunk


def _sheet_extras(wb, zf: zipfile.ZipFile, ws_ro, shared_strings, table_names: set[str]) -> dict:
    """sheetData 以外を通常モードと同じ手順で読む。_cells には結合範囲・リンク・コメントで作られたセルだけが入る。"""
    member = ws_ro._worksheet_path
    with zf.open(member) as f:
        stripped = _split_sheet_xml(f)
    ws = Worksheet(wb, ws_ro.title)
    rels_path = get_rels_path(member)
    ws._rels = get_dependents(zf, rels_path) if rels_path in zf.namelist() else RelationshipList()
    reader = WorksheetReader(ws, io.BytesIO(stripped), shared_strings, True, False)
    reader.bind_all()
    for r in ws._rels.find(COMMENTS_NS):
        comment_sheet = CommentSheet.from_tree(fromstring(zf.read(r.target)))
        for ref, comment in comment_sheet.comments:
            try:
                ws[ref].comment = comment
            except AttributeError:
                continue
    tables = []
    for target in reader.tables:
        table = Table.from_tree(fromstring(zf.read(target)))
        name = table.name.lower()
        if name in table_names or name in wb.defined_names:
            raise ValueError(f"Table with name {table.name} already exists")
        table_names.add(name)
        tables.append(table)
    anchors: dict[tuple[int, int], tuple[int, int]] = {}
    for rng in ws.merged_cells.ranges:
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                anchors[(r, c)] = (rng.min_row, rng.min_col)
    hidden_columns: list[int] = []
    for key, dim in ws.column_dimensions.items():
        if dim.hidden:
            lo = dim.min or column_index_from_string(key)
            hi = dim.max or lo
            hidden_columns.extend(range(lo, hi + 1))
    by_name = {}
    for t in tables:
        by_name[t.name] = t  # TableList と同じく名前で上書き
    return {
        "extra_cells": dict(ws._cells), "anchors": anchors, "hidden_columns": sorted(set(hidden_columns)),
        "freeze_panes": ws.freeze_panes, "auto_filter": ws.auto_filter.ref or None,
        "table_ranges": [t.ref for t in by_name.values()],
    }


# ---- シートの XML: 値のある範囲とセル数（expat で1回なめる） ----

def _scan_cells(f, extra: dict, shared_strings) -> dict:
    """WorkSheetParser と同じ規則で、値のある最終行・最終列・セル数・行の並びを調べる。"""
    ROW, V, IS = f"{_MAIN_NS} row", f"{_MAIN_NS} v", f"{_MAIN_NS} is"
    depth = 0
    row_depth = cell_depth = is_depth = r_depth = 0
    row_no = col_no = 0
    last_row = 0
    ordered = True
    count = max_row = max_col = 0
    seen_extra: set = set()
    # セル1つぶんの状態
    cell_row = cell_col = 0
    cell_type = "n"
    v_state = 0  # 0: 未読 1: 読み中 2: 読んだ（findtext と同じく最初の v だけ）
    v_text = None
    is_found = False
    plain = None
    rich: list = []
    text: list[str] = []
    capture = 0  # 1: v 2: is/t 3: is/r/t

    def local(name: str) -> str:
        return name[name.rfind(" ") + 1:]

    def start(name, attrs):
        nonlocal depth, row_depth, cell_depth, is_depth, r_depth, row_no, col_no, last_row, ordered
        nonlocal cell_row, cell_col, cell_type, v_state, v_text, is_found, plain, rich, capture
        depth += 1
        if cell_depth:
            if depth == cell_depth + 1:
                if name == V and v_state == 0:
                    v_state, capture = 1, 1
                    text.clear()
                elif name == IS and not is_found:
                    is_found, is_depth = True, depth
            elif is_depth:
                if depth == is_depth + 1:
                    tag = local(name)
                    if tag == "t":
                        capture = 2
                        text.clear()
                    elif tag == "r":
                        r_depth = depth
                        rich.append(None)
                elif r_depth and depth == r_depth + 1 and local(name) == "t":
                    capture = 3
                    text.clear()
            return
        if row_depth:
            if depth == row_depth + 1:
                cell_depth = depth
                ref = attrs.get("r")
                if ref:
                    cell_row, cell_col = coordinate_to_tuple(ref)
                    col_no = cell_col
                    if cell_row != row_no:
                        ordered = False
                else:
                    col_no += 1
                    cell_row, cell_col = row_no, col_no
                cell_type = attrs.get("t", "n")
                v_state, v_text, is_found, plain, capture = 0, None, False, None, 0
                rich = []
            return
        if name == ROW:
            row_depth = depth
            ref = attrs.get("r")
            if ref:
                try:
                    row_no = int(ref)
                except ValueError:
                    row_no = int(float(ref))
            else:
                row_no += 1
            col_no = 0
            if row_no <= last_row:
                ordered = False
            last_row = row_no

    def end(name):
        nonlocal depth, row_depth, cell_depth, is_depth, r_depth, count, max_row, max_col
        nonlocal v_state, v_text, plain, capture
        if capture and depth == (cell_depth + 1 if capture == 1 else is_depth + 1 if capture == 2 else r_depth + 1):
            value = "".join(text) if text else None  # ElementTree の .text と同じく、文字がなければ None
            if capture == 1:
                v_state, v_text = 2, value
            elif capture == 2:
                plain = value
            else:
                rich[-1] = value
            capture = 0
        elif r_depth and depth == r_depth:
            r_depth = 0
        elif is_depth and depth == is_depth:
            is_depth = 0
        elif cell_depth and depth == cell_depth:
            count += 1
            if cell_type == "inlineStr":
                value = (plain or "") + "".join(t for t in rich if t) if is_found else None
            else:
                value = (v_text or None) if v_state == 2 else None
                if value is not None and cell_type == "s":
                    value = shared_strings[int(value)]
            key = (cell_row, cell_col)
            m = extra.get(key)
            if m is not None:
                seen_extra.add(key)
                if isinstance(m, MergedCell):
                    value = None
                elif value is None:
                    value = m._value
            if value is not None and value != "":
                if cell_row > max_row:
                    max_row = cell_row
                if cell_col > max_col:
                    max_col = cell_col
            cell_depth = 0
        elif row_depth and depth == row_depth:
            row_depth = 0
        depth -= 1

    def chars(data):
        if capture:
            text.append(data)

    parser = expat.ParserCreate(namespace_separator=" ")
    parser.buffer_text = True
    parser.StartElementHandler = start
    parser.EndElementHandler = end
    parser.CharacterDataHandler = chars
    parser.ParseFile(f)
    for key, m in extra.items():
        if key in seen_extra or isinstance(m, MergedCell) or m._value is None or m._value == "":
            continue
        max_row, max_col = max(max_row, key[0]), max(max_col, key[1])
    return {"cell_count": count + len(extra) - len(seen_extra), "max_row": max_row, "max_col": max_col, "ordered": ordered}


# ---- シートの XML: セルの値（WorkSheetParser） ----

def _inline_text(element) -> str:
    """Text.from_tree(element).content と同じ文字列（直下の t は最後のもの＋各 r の t。ふりがな rPh は除く）。"""
    plain = None
    runs = []
    for el in element:
        tag = el.tag.rpartition("}")[2]
        if tag == "t":
            plain = el.text
        elif tag == "r":
            text = None
            for sub in el:
                if sub.tag.rpartition("}")[2] == "t":
                    text = sub.text
            if text is not None:
                runs.append(text)
    return (plain or "") + "".join(runs)


class _SheetParser(WorkSheetParser):
    """インライン文字列だけを速く読む WorkSheetParser（値は同じ。Serialisable を作らない）。"""

    def parse_cell(self, element):
        if element.get("t") != "inlineStr" or not self.data_only:
            return super().parse_cell(element)
        coordinate = element.get("r")
        style_id = element.get("s", 0)
        if style_id:
            style_id = int(style_id)
        if coordinate:
            row, column = coordinate_to_tuple(coordinate)
            self.col_counter = column
        else:
            self.col_counter += 1
            row, column = self.row_counter, self.col_counter
        value, data_type = None, "inlineStr"
        child = element.find(INLINE_STRING)
        if child is not None:
            value, data_type = _inline_text(child), "s"
        return {"row": row, "column": column, "value": value, "data_type": data_type, "style_id": style_id}


def _row_hidden(attrs: dict | None) -> bool:
    """RowDimension(hidden=...) と同じ変換（'false' 'f' '0' と空は False）。"""
    value = (attrs or {}).get("hidden")
    return bool(value) and value not in ("false", "f", "0")


class ExcelSource:
    kind = "excel"

    def __init__(self, path, file_name: str | None = None, options: dict | None = None):
        self.path = Path(path)
        self.file_name = file_name or self.path.name
        options = dict(options or {})
        self.max_cells = int(options.get("max_cells") or DEFAULT_MAX_CELLS)
        # 同じファイルを前に開いたときの sheet_stats()（取り込みの控え）。あれば値のある範囲を数え直さない
        known_stats = options.get("sheet_stats") or {}
        self._shared_strings = []

        # 範囲情報（dimension）は書式だけの行で大きく出ることがあるので、目安の4倍を超えるときだけ先に止める
        estimated = estimate_cell_count(self.path)
        if estimated > self.max_cells * 4:
            raise _too_many_cells(estimated, self.max_cells)
        wb = None
        try:
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                wb = load_workbook(self.path, read_only=True, data_only=True)
                self._sheets = self._read_sheets(wb, known_stats)
        except UploadError:
            raise
        except Exception as e:  # openpyxl・XML は形式ごとに様々な例外を出す
            raise UploadError(_OPEN_ERROR) from e
        finally:
            if wb is not None:
                wb.close()

        actual = sum(s["cell_count"] for s in self._sheets.values())
        if actual > self.max_cells:
            raise _too_many_cells(actual, self.max_cells)
        self.date1904 = wb.epoch == CALENDAR_MAC_1904
        self._epoch = wb.epoch
        self._date_formats = wb._date_formats
        self._timedelta_formats = wb._timedelta_formats
        self._styles: dict[int, tuple] = {}
        self._wb_styles = (wb._cell_styles, wb._fonts, wb._fills, wb._number_formats)
        self._default_style = self._style_tuple(None)

    def _read_sheets(self, wb, known_stats: dict) -> dict[str, dict]:
        sheets: dict[str, dict] = {}
        table_names: set[str] = set()
        zf = wb._archive
        use_known = bool(known_stats) and set(known_stats) == set(wb.sheetnames)
        for ws in wb.worksheets:
            self._shared_strings = ws._shared_strings
            extras = _sheet_extras(wb, zf, ws, ws._shared_strings, table_names)
            if use_known:
                scan = {k: known_stats[ws.title][k] for k in _STAT_KEYS}
            else:
                with zf.open(ws._worksheet_path) as f:
                    scan = _scan_cells(f, extras["extra_cells"], ws._shared_strings)
            sheets[ws.title] = {**extras, **scan, "member": ws._worksheet_path, "state": ws.sheet_state}
        return sheets

    def sheet_stats(self) -> dict[str, dict]:
        """シートごとの値のある範囲・セル数・行の並び（JSON にできる形。次に開くときの options["sheet_stats"]）。"""
        return {name: {k: meta[k] for k in _STAT_KEYS} for name, meta in self._sheets.items()}

    def close(self) -> None:
        pass

    def _style_tuple(self, style_array) -> tuple:
        """(表示形式, 太字, 取り消し線, 塗り)。style_array が None なら既定（全部0）のスタイル。"""
        _cell_styles, fonts, fills, number_formats = self._wb_styles
        num_id = style_array.numFmtId if style_array is not None else 0
        font = fonts[style_array.fontId if style_array is not None else 0]
        fill = fills[style_array.fillId if style_array is not None else 0]
        if num_id < BUILTIN_FORMATS_MAX_SIZE:
            number_format = BUILTIN_FORMATS.get(num_id, "General")
        else:
            number_format = number_formats[num_id - BUILTIN_FORMATS_MAX_SIZE]
        return (number_format, bool(font is not None and font.b), bool(font is not None and font.strike),
                bool(fill is not None and fill.fill_type not in (None, "none")))

    def _style(self, style_id) -> tuple:
        found = self._styles.get(style_id)
        if found is None:
            found = self._styles[style_id] = self._style_tuple(self._wb_styles[0][style_id])
        return found

    def sheets(self) -> list[SheetInfo]:
        return [SheetInfo(
            name=name,
            hidden=meta["state"] != "visible",
            max_row=meta["max_row"],
            max_col=meta["max_col"],
            table_ranges=list(meta["table_ranges"]),
            date1904=self.date1904,
            hidden_columns=list(meta["hidden_columns"]),
            freeze_panes=meta["freeze_panes"],
            auto_filter=meta["auto_filter"],
        ) for name, meta in self._sheets.items()]

    def _meta(self, sheet: str) -> dict:
        if sheet not in self._sheets:
            raise UploadError(f"シート「{sheet}」が見つかりません")
        return self._sheets[sheet]

    def _parsed_rows(self, zf: zipfile.ZipFile, meta: dict, end: int) -> Iterator[tuple[int, list[dict], dict | None]]:
        """(行番号, セルの dict, 行の属性) を行番号の昇順に返す。end より後は読まない。"""
        with zf.open(meta["member"]) as f:
            parser = _SheetParser(f, self._shared_strings, data_only=True, epoch=self._epoch,
                                     date_formats=self._date_formats, timedelta_formats=self._timedelta_formats)
            if meta["ordered"]:
                for idx, cells in parser.parse():
                    if idx > end:
                        return
                    yield idx, cells, parser.row_dimensions.pop(str(idx), None)
                return
            # 行が昇順でないシート: 全行を読んでから並べる（通常モードと同じく後のセルで上書き）
            rows: dict[int, dict[int, dict]] = {}
            for _idx, cells in parser.parse():
                for d in cells:
                    rows.setdefault(d["row"], {})[d["column"]] = d
            for idx in sorted(rows):
                if idx > end:
                    return
                yield idx, list(rows[idx].values()), parser.row_dimensions.get(str(idx))

    def rows(self, sheet: str, start: int = 1, limit: int | None = None) -> Iterator[SourceRow]:
        meta = self._meta(sheet)
        anchors = meta["anchors"]
        extra = meta["extra_cells"]
        extra_rows = {r for r, _c in extra}
        anchor_rows = {r for r, _c in anchors}
        max_col = meta["max_col"]
        begin = max(1, start)
        end = meta["max_row"] if limit is None else min(meta["max_row"], start + limit - 1)
        if end < begin:
            return
        default = self._default_style
        style_of = self._style
        try:
            with zipfile.ZipFile(self.path) as zf:
                parsed = self._parsed_rows(zf, meta, end)
                pending = next(parsed, None)
                for r in range(1, end + 1):
                    while pending is not None and pending[0] < r:
                        pending = next(parsed, None)
                    if pending is not None and pending[0] == r:
                        _idx, xml_cells, attrs = pending
                        pending = next(parsed, None)
                    else:
                        xml_cells, attrs = (), None
                    if r < begin:
                        continue
                    by_col = {}
                    for d in xml_cells:
                        if d["column"] <= max_col:
                            by_col[d["column"]] = d
                    has_extra = r in extra_rows
                    has_anchor = r in anchor_rows
                    row_cells = []
                    for c in range(1, max_col + 1):
                        anchor = anchors.get((r, c)) if has_anchor else None
                        m = extra.get((r, c)) if has_extra else None
                        d = by_col.get(c)
                        if m is not None and (isinstance(m, MergedCell) or d is None):
                            # 結合範囲の左上以外（値は消える）や、リンク・コメントで作られたセルは既定のスタイル
                            value = None if isinstance(m, MergedCell) else m._value
                            number_format, bold, strike, fill = default
                        elif d is not None:
                            value = d["value"]
                            if value is None and m is not None:
                                value = m._value
                            number_format, bold, strike, fill = style_of(d["style_id"])
                        else:
                            row_cells.append(CellInfo(value=None, text="", merged_anchor=anchor))
                            continue
                        row_cells.append(CellInfo(value=value, text=cell_text(value), number_format=number_format,
                                                  bold=bold, strike=strike, fill=fill, merged_anchor=anchor))
                    yield SourceRow(index=r, cells=row_cells, hidden=_row_hidden(attrs))
        except (UploadError, GeneratorExit):
            raise
        except Exception as e:
            raise UploadError(_OPEN_ERROR) from e
