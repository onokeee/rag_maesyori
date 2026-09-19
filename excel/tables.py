"""帳票の中の明細表（交換部品・時系列・チェックシートなど）を読む。

見出しのセル（アンカー: 「■ 交換部品」「使用部品」など）の下、または縦に結合した見出しの右にある
「列見出しの行」を見つけ、その下の行を空行・見出し欄（列見出しと同じ塗りつぶし色のセル）まで読む。
1行 = 1明細。値は {"columns": [列見出し...], "rows": [[セルの文字列...], ...]} で持つ。
"""
from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass

from excel.text import pick_checked
from excel.workbook import Cell, SheetGrid

MAX_TABLE_ROWS = 200
MAX_HEADER_CHARS = 30
MIN_COLUMNS = 2
# 自動検出（見本からの候補づくり・見出しの優先順位）で、アンカーの無い表とみなす最小の列数・行数
MIN_COLUMNS_WITHOUT_ANCHOR = 3
MIN_ROWS_WITHOUT_ANCHOR = 2
# 列見出しが「同じ表」とみなす一致率（共通の列見出し ÷ どちらかにある列見出し）
SAME_COLUMNS_RATIO = 0.5

_NUMBER_LIKE = re.compile(r"^[\d\s,.\-/:+%]+$")
# 判定記号と、その凡例の1組（「◎：主要因」「○=要因の可能性あり」「－：対象外」）
_JUDGE_MARKS = "◎○◯〇△▲▽×✕✖"
_LEGEND_PAIR = re.compile(rf"[{_JUDGE_MARKS}ー―－\-]\s*[:：=＝]")
MAX_LEGEND_CHARS = 60
_SEQ_HEADERS = {"no", "№", "#", "項番", "番号", "順", "順番"}
_TOTAL_WORDS = {"合計", "小計", "総計", "計", "合計数", "総合計"}
# 合計行の先頭（「合計」「小計」「部品費計」「工数計」「部品費合計」）。
# 「〜計」を全部合計とみなすと「温度計」「圧力計」「設計」などの行で読み取りが止まり、Markdownでも消えるので、
# 合計の言い方だけにする（export.formats も同じ決まりを使う）
TOTAL_LABEL_RE = re.compile(r"^(?:合計|小計|総計|総合計|合計数|計|.{0,6}(?:合計|小計)|"
                            r".{0,6}(?:費|工数|個数|件数|台数|本数|枚数|金額|額)計)$")
_NONE_MARKS = {"なし", "無し", "該当なし", "特になし", "-", "ー", "―", "‐", "/", "〃"}
# 「上の行と同じ」の記号（″ は NFKC で ′′ になる）。明細表の値では上の行の同じ列の値に置き換える
_DITTO_MARKS = {"〃", "′′", "同上", "仝"}
# 見出しらしい書き出し: 「■ 交換部品」「【時系列】」「▼ 回答欄」「1. 時系列」「A. 機構部」
_HEADING_START = re.compile(r"^(?:[■□◆◇●▼▽▶►【\[]|\d{1,2}[.)、](?!\d)|[a-z][.)、](?![a-z])|d[1-8](?=\D))")


@dataclass
class Table:
    anchor: Cell | None
    header: list[Cell]
    rows: list[list[Cell | None]]
    blank_rows: int = 0  # 飛ばした「No だけの空き行」の数

    @property
    def has_room(self) -> bool:
        """行を足せる明細表の形か（2行以上ある／空き行がある／「■」「1.」の見出しの下／縦結合の見出しに余りの行がある）。

        「影響｜停止時間｜影響ロット」のような見出しの行＋値の行1つは、明細表でなく項目の並びとみなす。
        """
        if len(self.rows) >= 2 or self.blank_rows:
            return True
        anchor = self.anchor
        if anchor is None:
            return False
        if _HEADING_START.match(_compact(anchor.text)):
            return True
        return anchor.max_row > anchor.row and anchor.max_row > self.last_row

    @property
    def last_row(self) -> int:
        """読み取った範囲の最後の行。"""
        return max([c.max_row for row in self.rows for c in row if c is not None]
                   or [max(h.max_row for h in self.header)])

    @property
    def coord(self) -> str:
        first = self.header[0]
        right = max(c.max_col for c in self.header)
        from openpyxl.utils import get_column_letter

        return f"{get_column_letter(first.col)}{first.row}:{get_column_letter(right)}{self.last_row}"

    def to_value(self) -> dict | None:
        """保存用の値。連番だけの No 列は落とす。行が無ければ None。"""
        columns = [" ".join(h.text.split()) for h in self.header]
        rows = [[c.text if c is not None else "" for c in row] for row in self.rows]
        if len(columns) > 1 and _is_seq_column(columns[0], [r[0] for r in rows]):
            columns = columns[1:]
            rows = [_drop_seq_cell(r) for r in rows]
        # 空の行と「なし」「－」だけの行（明細なしの記入）は落とす
        rows = [r for r in rows if any(v.strip() and _compact(v) not in _NONE_MARKS for v in r)]
        if not rows:
            return None
        _resolve_ditto(rows)
        return {"columns": columns, "rows": rows}


def _resolve_ditto(rows: list[list[str]]) -> None:
    """「〃」「同上」のセルを、上の行の同じ列の値にする（縦結合のセルと同じ扱い。1行だけ読んでも意味が通るように）。

    上の行の値が空か、最初の行なら書かれたままにする。
    """
    for above, row in zip(rows, rows[1:]):
        for i, v in enumerate(row):
            if _compact(v) in _DITTO_MARKS and i < len(above) and above[i].strip()                     and _compact(above[i]) not in _DITTO_MARKS:
                row[i] = above[i]


def find_table(grid: SheetGrid, anchor: Cell, direction: str = "auto", stop_labels: set[str] | None = None) -> Table | None:
    """アンカーのセルから明細表を探す。

    - 縦に結合したアンカー（「使用部品」が4行分など）は、右隣の列見出しの行を先に見る（読む行はアンカーの範囲内）
    - それ以外は、アンカーの下2行以内で、アンカーの列から始まる列見出しの行を探す
    """
    stop_labels = stop_labels or set()
    vertical = anchor.max_row > anchor.row
    if direction in ("auto", "right") and vertical:
        header = header_cells(grid, anchor.row, anchor.max_col + 1)
        if len(header) >= MIN_COLUMNS:
            return read_table(grid, anchor, header, stop_labels, last_row=anchor.max_row)
    if direction in ("auto", "below"):
        marked = bool(_HEADING_START.match(_compact(anchor.text)))
        row, last = anchor.max_row + 1, anchor.max_row + 2
        while row <= last:
            starts = [c for c in _cells_in_row(grid, row) if anchor.col <= c.col <= anchor.max_col]
            if starts:
                header = header_cells(grid, row, starts[0].col)
                if len(header) >= MIN_COLUMNS:
                    return read_table(grid, anchor, header, stop_labels)
                if marked and last == anchor.max_row + 2 and _label_value_row(grid, row):
                    # 「■ 水平展開」の下に「展開区分｜☑同型機…」の1行があり、その下に列見出しが来る様式
                    last, row = last + 1, row + 1
                    continue
                break
            if any(c.col <= anchor.max_col and c.max_col >= anchor.col for c in _cells_in_row(grid, row)):
                break
            row += 1
    return None


def header_cells(grid: SheetGrid, row: int, start_col: int) -> list[Cell]:
    """row 行目の start_col から右へ、すき間なく並ぶ見出しらしいセル（高さが先頭のセルと同じもの）。"""
    cells: list[Cell] = []
    col = start_col
    while col <= grid.max_col:
        cell = grid.cells.get((row, col))
        if cell is None or not _header_like(cell) or (cells and cell.max_row != cells[0].max_row):
            break
        cells.append(cell)
        col = cell.max_col + 1
    return cells


def read_table(grid: SheetGrid, anchor: Cell | None, header: list[Cell], stop_labels: set[str],
               last_row: int | None = None) -> Table:
    """列見出しの下の行を読む。空行、見出し欄（列見出しと同じ色の文字列・ラベル）、合計行の後で止める。

    縦に結合したデータセル（点検部位など）は、結合範囲の各行に同じ値を入れる。
    連番（No）だけが書かれた空き行（様式の固定行数の余り）は飛ばして続ける。
    """
    fills = {h.fill for h in header if h.fill}
    row = max(h.max_row for h in header) + 1
    end = min(grid.max_row, last_row or grid.max_row, row + MAX_TABLE_ROWS - 1)
    rows: list[list[Cell | None]] = []
    blank_rows = 0
    seq_header = _norm_header(header[0]) in _SEQ_HEADERS
    while row <= end:
        cells: list[Cell | None] = []
        seen: set[tuple[int, int]] = set()
        for h in header:
            found = None
            for col in range(h.col, h.max_col + 1):
                found = grid.cell_at(row, col)
                if found is not None:
                    break
            if found is not None and (found.row, found.col) in seen:
                found = None
            if found is not None:
                seen.add((found.row, found.col))
            cells.append(found)
        fresh = [c for c in cells if c is not None and c.row == row]
        if not fresh:
            break
        total = any(_is_total(c) for c in fresh)
        if not total and any(_is_stop(c, fills, stop_labels, header[0].col) for c in fresh):
            break
        if seq_header and all(c is cells[0] for c in fresh) and _NUMBER_LIKE.match(fresh[0].norm):
            row += 1  # No だけの空き行
            blank_rows += 1
            continue
        rows.append(cells)
        if total and any(c.filled for c in fresh):
            break
        row = min(c.max_row for c in fresh) + 1
    return Table(anchor, header, rows, blank_rows)


def detect_tables(grid: SheetGrid) -> list[Table]:
    """シート内の明細表を自動で見つける（見本からの候補づくり・見出しの優先順位に使う）。

    列見出し = 塗りつぶしのある短い文字列セルが2つ以上すき間なく並ぶ行。
    アンカー（真上の見出し、または左の縦結合の見出し）があれば1行以上、無ければ3列以上・2行以上の表だけを採る。
    """
    cached = getattr(grid, "_tables", None)
    if cached is not None:
        return cached
    tables: list[Table] = []
    used: set[tuple[int, int]] = set()
    for cell in grid.text_cells():
        if (cell.row, cell.col) in used or not cell.filled or not _header_like(cell):
            continue
        # 行の先頭の縦結合セル（「使用部品」）は列見出しにならず（高さが違う）、右隣からの並びのアンカーになる
        header = _contiguous([c for c in header_cells(grid, cell.row, cell.col) if c.filled])
        if len(header) < MIN_COLUMNS:
            continue
        used.update((c.row, c.col) for c in header)
        anchor, last_row = _find_anchor(grid, header)
        table = read_table(grid, anchor, header, set(), last_row=last_row)
        if anchor is not None and table.rows:
            tables.append(table)
        elif len(header) >= MIN_COLUMNS_WITHOUT_ANCHOR and len(table.rows) >= MIN_ROWS_WITHOUT_ANCHOR:
            tables.append(table)
    grid._tables = tables
    return tables


def find_table_by_columns(grid: SheetGrid, columns: list[str], keep=None) -> Table | None:
    """列見出しの半分以上が columns と同じ、行のある明細表（最も似ているもの）。keep(表) が偽の表は使わない。"""
    from excel.text import normalize_label

    wanted = {normalize_label(c) for c in columns} - {""}
    if len(wanted) < MIN_COLUMNS:
        return None
    best, best_score = None, 0.0
    for table in detect_tables(grid):
        have = {h.norm for h in table.header}
        score = len(have & wanted) / len(have | wanted)
        if (score >= SAME_COLUMNS_RATIO and score > best_score and table.to_value() is not None
                and (keep is None or keep(table))):
            best, best_score = table, score
    return best


def stacked_tables(grid: SheetGrid, table: Table, stop_labels: set[str] | None = None,
                   limit: int = 10) -> list[Table]:
    """読み取った範囲のすぐ下に続く、同じ形（同じ列位置・同じ塗りつぶし色）の列見出しの行とその行。

    特性要因図のように「人｜機械」「材料｜方法」「測定｜環境」と同じ形の小さな表が縦に並ぶ帳票では、
    組ごとに列見出しが変わるため、1つの表として読むと2組目から先が落ちる。組ごとに読んで
    merge_table_values でまとめる（design.md 6.1「積み重なった列見出し」）。
    """
    stop_labels = stop_labels or set()
    if not table.header:
        return []
    fills = {h.fill for h in table.header if h.fill}
    shape = [(h.col, h.max_col) for h in table.header]
    last = table.last_row
    blocks: list[Table] = []
    while len(blocks) < limit:
        header = None
        for row in (last + 1, last + 2):
            if row > grid.max_row:
                break
            found = _contiguous([c for c in header_cells(grid, row, shape[0][0]) if c.filled])
            if [(h.col, h.max_col) for h in found] == shape \
                    and (not fills or {h.fill for h in found if h.fill} & fills):
                header = found
                break
        if header is None:
            break
        block = read_table(grid, None, header, stop_labels)
        blocks.append(block)
        last = max(block.last_row, header[0].max_row)
    return blocks


def stacked_header_rows(grid: SheetGrid, table: Table, limit: int = 10) -> int:
    """読み取った範囲のすぐ下に、同じ形の列見出しの行が何組続くか（stacked_tables の組数）。"""
    return len(stacked_tables(grid, table, limit=limit))


def merge_table_values(values: list[dict | None]) -> dict | None:
    """積み重なった組（同じ形の列見出しが縦に並ぶ表）の値を、1つの明細表の値にまとめる。

    列見出しは出てきた順に並べ、同じ見出しは同じ列にそろえる（組ごとに見出しが変わる特性要因図は
    「人｜機械｜材料｜方法｜測定｜環境」の6列になり、各行は自分の組の列だけ埋まる）。
    同じ見出しが続くだけの組（ページをまたいで見出しを繰り返す様式）は、そのまま行が増える。
    """
    parts = [v for v in values if is_table_value(v) and v.get("rows")]
    if not parts:
        return next((v for v in values if is_table_value(v)), None)
    if len(parts) == 1:
        return parts[0]
    from excel.text import normalize_label

    columns: list[str] = []
    keys: list[str] = []
    rows: list[list[str]] = []
    for part in parts:
        used: set[int] = set()
        slots: list[int] = []
        for i, name in enumerate(part["columns"]):
            key = normalize_label(name) or f"\x00{len(keys)}"  # 空の見出しは他とまとめない
            slot = next((n for n, k in enumerate(keys) if k == key and n not in used), None)
            if slot is None:
                keys.append(key)
                columns.append(name)
                slot = len(keys) - 1
            used.add(slot)
            slots.append(slot)
        for row in part["rows"]:
            cells = [""] * len(keys)
            for i, cell in enumerate(row):
                if i < len(slots):
                    cells[slots[i]] = cell
            rows.append(cells)
    return {"columns": columns, "rows": [row + [""] * (len(columns) - len(row)) for row in rows]}


# ---- 区切りの見出し（区画）------------------------------------------------------------------
# 発行側と回答側で同じ意味の欄が並ぶ帳票（「処置内容」と「▼ 回答欄」の下の「暫定対策（処置）」）で、
# 項目がどちら側の欄かを見分けるための区画。区画 = 「■」「▼」「【】」「1.」などで始まる見出しのセルから、
# 右は同じ高さにある次の見出しの手前まで（無ければシートの右端まで）、下は次の見出しまで。
MAX_SECTION_CHARS = 60
_SECTION_TAIL = re.compile(r"[(（].*$")


@dataclass
class Section:
    key: str        # 区画の名前（比較用。「▼ 回答欄（宛先部署にて記入…）」「【回答欄】」→「回答」）
    cell: Cell
    right: int      # 区画の右端の列


def section_key(cell: Cell) -> str:
    """区画の見出しのセルの比較用の名前（section_name）。"""
    return section_name(cell.text)


def section_name(text) -> str:
    """区画の見出しの比較用の名前。印・項番・括弧書き・末尾の「欄」「内容」を除く（版ごとの書き方の違いを吸収する）。

    「▼ 回答欄（宛先部署にて記入し…）」「【回答欄】」「回答欄」→「回答」。帳票の種類の画面で入力された区画もこれでそろえる。
    何度通しても同じ結果になる（「■ 処置内容欄」も「処置内容」も「処置」）。
    """
    from excel.text import label_base, normalize_label, section_stripped

    title = _title_text(text)
    norm = normalize_label(_SECTION_TAIL.sub("", title)) or normalize_label(title)
    norm = section_stripped(norm) or norm
    # 「処置内容欄」→「処置内容」→「処置」のように末尾の語が重なることがあるので、変わらなくなるまで除く。
    # 保存した区画をもう一度この関数に通しても同じ名前になる（見出しと保存値が必ずそろう）。
    while base := label_base(norm):
        norm = base
    return norm


def _section_start(cell: Cell) -> bool:
    """区画の見出しか（塗りつぶし・太字で、「■」「▼」「【】」「1.」などの印で始まる1行の短い文字列）。"""
    if not isinstance(cell.value, str) or "\n" in cell.text or cell.inline is not None:
        return False
    if len(cell.norm) > MAX_SECTION_CHARS or pick_checked(cell.text) is not None:
        return False
    return (cell.filled or cell.bold) and bool(_HEADING_START.match(_compact(cell.text)))


def sections(grid: SheetGrid) -> list[Section]:
    """シートの区画（見出しの上→下、左→右の順）。"""
    cached = getattr(grid, "_sections", None)
    if cached is not None:
        return cached
    heads = [c for c in grid.text_cells() if _section_start(c)]
    out: list[Section] = []
    for head in heads:
        # 同じ高さ（行が重なる）で右にある次の見出しの手前までを、この区画の横幅にする
        right_heads = [h.col for h in heads if h.col > head.max_col and h.row <= head.max_row and h.max_row >= head.row]
        # 上の行で右側に始まった区画（右上の「【回答欄】」の下に左の「■ 発行部署記入欄」がある版）の列は、
        # その区画のまま（左の区画の横幅に入れない）。右上の区画の列に、あとから別の見出しが無いときだけ
        for h in heads:
            if h.col > head.max_col and h.max_row < head.row and not any(
                    h.max_row < k.row < head.row and k.max_col >= h.col for k in heads):
                right_heads.append(h.col)
        out.append(Section(section_key(head), head, min(right_heads, default=grid.max_col + 1) - 1))
    grid._sections = out
    return out


def _enclosing(grid: SheetGrid, cell: Cell) -> list[Section]:
    """セルを含む区画（見出しがセルより上か同じ行で、横幅にセルの列が入るもの）。近い見出し（下・右）から順に。"""
    found = [sec for sec in sections(grid) if sec.cell.row <= cell.row and sec.cell.col <= cell.col <= sec.right]
    return sorted(found, key=lambda sec: (sec.cell.row, sec.cell.col), reverse=True)


def section_of(grid: SheetGrid, cell: Cell) -> str:
    """セルが入っている区画の名前。どの区画にも入らなければ ""。区画の見出しのセル自身はその区画に入る。"""
    found = _enclosing(grid, cell)
    return found[0].key if found else ""


def _heading_level(cell: Cell) -> int:
    """区画の見出しの階層。「■」「▼」「【】」などの印=0、「1.」「D1」=1、「a.」=2（数字の小見出しは印の見出しの中）。"""
    head = _compact(cell.text)
    if re.match(r"^[a-z][.)、]", head):
        return 2
    return 1 if re.match(r"^(?:\d|d[1-8])", head) else 0


def sections_of(grid: SheetGrid, cell: Cell) -> list[str]:
    """セルが入っている区画の名前を、内側（一番近い見出し）から外側へ。

    「▼ 回答欄」の下の「1. 暫定対策」の中のセルは ["暫定対策", "回答"]。外側の見出しは、それより内側の見出しより
    上の階層（印の見出しは数字の小見出しの外側）のものだけ。同じ階層の見出しが下にあれば、上の区画はそこで終わっている。
    """
    out: list[str] = []
    level = None
    for sec in _enclosing(grid, cell):
        lv = _heading_level(sec.cell)
        if level is None or lv < level:
            out.append(sec.key)
            level = lv
    return out


def table_header_keys(grid: SheetGrid) -> set[tuple[int, int]]:
    """明細表の列見出しのセル位置。"""
    return {(c.row, c.col) for t in detect_tables(grid) for c in t.header}


def list_header_keys(grid: SheetGrid) -> set[tuple[int, int]]:
    """行を足せる明細表の列見出しのうち、その列だけに収まる値が2行以上並ぶもののセル位置。1つの値の項目のラベルにはしない。

    押印欄（承認｜確認｜作成 の下に印と日付、その下に横長の注記）は含めない。見出し（アンカー）の無い表は、
    押印と日付の2行と区別するため3行以上並ぶときだけ。
    """
    keys = set()
    for table in detect_tables(grid):
        if not table.has_room:
            continue
        # 見出しの無い表でも、日付だけの行の無い表（水平展開先の「確認｜対象設備｜設備名…」が2行）は一覧とみなす
        need = 2 if table.anchor is not None or not _has_date_row(table) else 3
        for i, h in enumerate(table.header):
            cells = [row[i] for row in table.rows if i < len(row) and row[i] is not None]
            if sum(1 for c in cells if c.col >= h.col and c.max_col <= h.max_col) >= need:
                keys.add((h.row, h.col))
    return keys


def _has_date_row(table: Table) -> bool:
    """押印欄（承認｜確認｜作成 の下に印、その下に「6/5」のような日付）らしい、日付・数字だけの行があるか。"""
    for row in table.rows:
        cells = [c for c in row if c is not None]
        if cells and all(not isinstance(c.value, str) or _NUMBER_LIKE.match(c.norm) for c in cells):
            return True
    return False


def seq_header(cell: Cell) -> bool:
    """連番の列見出し（No / № / 項番）。表の中にあれば、報告書の「No.」欄とは別物。"""
    return _norm_header(cell) in _SEQ_HEADERS


def section_heading(cell: Cell) -> bool:
    """塗りつぶしのある「■ 現象・処置」「▼ 回答欄（…）」のような区切りの見出し。値にはしない。"""
    return (cell.filled and isinstance(cell.value, str) and "\n" not in cell.text and cell.inline is None
            and bool(_HEADING_START.match(_compact(cell.text))) and pick_checked(cell.text) is None)


def heading_like(cell: Cell) -> bool:
    """見出しらしいセル（塗りつぶし・太字・「■」「1.」などで始まる短い文字列）。"""
    if not isinstance(cell.value, str) or len(cell.norm) > MAX_HEADER_CHARS + 10:
        return False
    marked = bool(_HEADING_START.match(_compact(cell.text)))
    if cell.inline and not marked:  # 「区分：同型機」はラベルと値。「D1：チームの結成」は見出し
        return False
    return cell.filled or cell.bold or marked


def legend_text(text) -> bool:
    """判定記号の凡例か（「◎：主要因　○：影響あり　×：検証の結果 要因でない」「（◎主要因／○寄与要因／×否定）」）。

    値でも見出しでもない飾りなので、項目の値の候補にも、表の見出しを探すときの「右にある値」にもしない。
    """
    s = " ".join(str(text or "").split())
    if not s or len(s) > MAX_LEGEND_CHARS:
        return False
    return sum(s.count(m) for m in _JUDGE_MARKS) >= 3 or len(_LEGEND_PAIR.findall(s)) >= 2


def table_title(cell: Cell) -> str:
    """アンカーの表示名。「■ 経緯（時系列）」→「経緯（時系列）」。"""
    return _title_text(cell.text)


def _title_text(raw) -> str:
    text = _one_line(raw)
    text = re.sub(r"^[■□◆◇●▼▽▶►・*※\s]+", "", text)
    if text.startswith(("【", "[")) and text.endswith(("】", "]")):
        text = text[1:-1]
    return text.strip(" :：") or _one_line(raw)


# ---- 値（{"columns": [...], "rows": [[...]]}）の扱い -----------------------------------------

def is_table_value(value) -> bool:
    return isinstance(value, dict) and isinstance(value.get("rows"), list)


def clean_table_value(value) -> dict | None:
    """画面から戻ってきた表の値をそろえる。列数に合わせて行を詰め、空の行を落とす。

    行が無くても列見出しがあれば {"columns": [...], "rows": []} を返す（確認・修正画面の表を残し、行を入れ直せるように）。
    列見出しも行も無ければ None。Markdown 側は行の無い明細表を出さない。
    """
    if not is_table_value(value):
        return None
    columns = [str(c if c is not None else "").strip() for c in (value.get("columns") or [])]
    rows = []
    for row in value["rows"]:
        if not isinstance(row, list):
            continue
        cells = ["" if c is None else str(c).replace("\r\n", "\n").strip() for c in row]
        while len(cells) > len(columns):
            columns.append(f"列{len(columns) + 1}")
        cells += [""] * (len(columns) - len(cells))
        if any(cells):
            rows.append(cells)
    if not rows and not any(c for c in columns):
        return None
    return {"columns": columns, "rows": rows}


def parse_table_text(text: str) -> tuple[dict | None, str | None]:
    """確認画面の入力（JSON）を表の値にする。戻り値: (値, 警告)。空なら (None, None)。"""
    import json

    if not (text or "").strip():
        return None, None
    try:
        data = json.loads(text)
    except ValueError:
        return None, "表の形式が正しくないため、変更を反映できませんでした"
    if not is_table_value(data):
        return None, "表の形式が正しくないため、変更を反映できませんでした"
    return clean_table_value(data), None


def table_row_items(value, row: list) -> list[tuple[str, str]]:
    """1行を (列見出し, 値) の組にする。空のセルは除く。"""
    columns = list(value.get("columns") or [])
    items = []
    for i, cell in enumerate(row):
        text = "" if cell is None else str(cell).strip()
        if text:
            items.append((columns[i] if i < len(columns) and columns[i] else f"列{i + 1}", text))
    return items


def table_text_lines(value) -> list[str]:
    """画面表示用: 1行を「品番: X／品名: Y」の1行にする。"""
    if not is_table_value(value):
        return []
    return ["／".join(f"{k}: {' '.join(v.split())}" for k, v in table_row_items(value, row))
            for row in value["rows"]]


# ---- 内部 ---------------------------------------------------------------------------

def _find_anchor(grid: SheetGrid, header: list[Cell]) -> tuple[Cell | None, int | None]:
    first, right = header[0], header[-1].max_col
    row = first.row
    if first.col > 1:
        left = grid.cell_at(row, first.col - 1)
        if left is not None and left.max_col == first.col - 1 and left.max_row > row and left.row <= row \
                and heading_like(left):
            return left, left.max_row
    above, last = row - 1, row - 2
    skipped = False
    while above >= max(1, last):
        # 判定記号の凡例（「◎：主要因　○：影響あり　×：否定」）は飾りなので、見出しの右にあっても数に入れない
        in_span = [c for c in _cells_in_row(grid, above)
                   if c.col <= right and c.max_col >= first.col and not legend_text(c.text)]
        if not in_span:
            above -= 1
            continue
        cell = in_span[0]
        marked = bool(_HEADING_START.match(_compact(cell.text)))
        # 右に値が並ぶセル（「区分｜☑同型機」）はアンカーにしない。「■」「1.」で始まる見出しは右に凡例があってもよい
        alone = len(in_span) == 1 or marked
        if alone and cell.col <= first.col + 1 and heading_like(cell) and (marked or not skipped) \
                and len(header_cells(grid, above, cell.col)) < MIN_COLUMNS:
            return cell, None
        if not skipped and _label_value_row(grid, above):
            # 「■ 水平展開」と列見出しの間に「展開区分｜☑同型機…」の1行がある様式: その上の「■」「1.」の見出しを見る
            skipped, last, above = True, last - 1, above - 1
            continue
        break
    return None, None


def _label_value_row(grid: SheetGrid, row: int) -> bool:
    """「展開区分｜☑同型機　□類似設備…」のような、見出し欄1つとその値だけの1行か（表の見出しと列見出しの間に挟まる行）。"""
    cells = _cells_in_row(grid, row)
    if len(cells) != 2:
        return False
    label, value = cells
    return (heading_like(label) and label.filled and not _HEADING_START.match(_compact(label.text))
            and not value.filled and value.col == label.max_col + 1 and value.max_row == label.max_row)


def _cells_in_row(grid: SheetGrid, row: int) -> list[Cell]:
    """row 行目から始まるセル（左上がその行にあるもの）を左から順に。"""
    by_row = getattr(grid, "_cells_by_row", None)
    if by_row is None:
        by_row = {}
        for cell in grid.text_cells():
            by_row.setdefault(cell.row, []).append(cell)
        grid._cells_by_row = by_row
    return by_row.get(row, [])


def _header_like(cell: Cell) -> bool:
    if not isinstance(cell.value, str) or cell.inline is not None:
        return False
    if len(cell.norm) > MAX_HEADER_CHARS or _NUMBER_LIKE.match(cell.norm):
        return False
    if legend_text(cell.text):  # 見出しの右に置かれた判定記号の凡例は列見出しではない
        return False
    return pick_checked(cell.text) is None


def _contiguous(cells: list[Cell]) -> list[Cell]:
    out: list[Cell] = []
    for c in cells:
        if out and c.col != out[-1].max_col + 1:
            break
        out.append(c)
    return out


def _is_stop(cell: Cell, fills: set[str], stop_labels: set[str], first_col: int) -> bool:
    """表の終わりを示すセル: 列見出しと同じ色の文字列、表の左端の色付きの短い文字列（「確認期間」などの項目欄）、
    色付き・太字のラベル。表の途中の色付きの値（NG の強調など）では止めない。"""
    if not isinstance(cell.value, str):
        return False
    if cell.fill and cell.fill in fills:
        return True
    if cell.filled and cell.col <= first_col and len(cell.norm) <= MAX_HEADER_CHARS:
        return True
    return (cell.filled or cell.bold) and cell.is_label(stop_labels)


def _is_total(cell: Cell) -> bool:
    if not isinstance(cell.value, str):
        return False
    return cell.norm in _TOTAL_WORDS or (cell.filled and len(cell.norm) <= 8 and bool(TOTAL_LABEL_RE.match(cell.norm)))


def _drop_seq_cell(row: list[str]) -> list[str]:
    """No 列を落とした行。合計行の「合計」が No 列にあれば、次の列が空のときだけそこへ移す。"""
    first, rest = row[0].strip(), row[1:]
    if first and not first.isdigit() and not rest[0].strip():
        return [first, *rest[1:]]
    return rest


def _is_seq_column(header: str, values: list[str]) -> bool:
    if _compact(header).strip(".:：") not in _SEQ_HEADERS:
        return False
    numbers = [v.strip() for v in values if v.strip() and _compact(v) not in _TOTAL_WORDS]
    if not numbers or not all(n.isdigit() for n in numbers):
        return False
    return [int(n) for n in numbers] == list(range(1, len(numbers) + 1))


def _norm_header(cell: Cell) -> str:
    return cell.norm.strip(".:：")


def _compact(text: str) -> str:
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", str(text))).lower()


def _one_line(text) -> str:
    return " ".join(unicodedata.normalize("NFKC", str(text or "")).split())
