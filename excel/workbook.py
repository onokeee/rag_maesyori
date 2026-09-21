"""Excelの構造解析。結合セルを考慮した「値のあるセル」のグリッドを作る。"""
from __future__ import annotations

import io
import warnings
import xml.etree.ElementTree as ET
import zipfile
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import coordinate_from_string
from openpyxl.utils.datetime import MAC_EPOCH

from excel.image_detector import NS, R_ID, _rels, detect_images
from excel.text import (MAX_LABEL_LENGTH, cell_text, format_unit, label_base, label_parts, normalize_label, paren_stripped,
                        section_stripped, split_inline)


@dataclass
class Cell:
    """値を持つセル。結合セルの場合は左上セルが範囲全体を代表する。"""

    row: int
    col: int
    max_row: int
    max_col: int
    value: object
    text: str
    norm: str
    inline: tuple[str, str] | None
    bold: bool = False
    filled: bool = False
    fill: str = ""  # 塗りつぶし色の識別キー（"rgb:FFDDEEFF" / "theme:4:0.6" など。塗りなしは ""）
    alt_norm: str = ""  # 先頭の項番を除いたラベル（「３．暫定対策」→「暫定対策」）。無ければ ""
    part_norms: tuple[str, ...] = ()  # 「ライン／工程」のように2つのラベルをまとめた見出しの各部分
    fmt_unit: str = ""  # 数値のセルの表示形式に書かれた単位（「#,##0"分"」→「分」）。無ければ ""

    @property
    def coord(self) -> str:
        start = f"{get_column_letter(self.col)}{self.row}"
        if (self.max_row, self.max_col) == (self.row, self.col):
            return start
        return f"{start}:{get_column_letter(self.max_col)}{self.max_row}"

    @property
    def label_keys(self) -> tuple[str, ...]:
        """セル全体をラベルとみなすときの比較キー。"""
        return (self.norm, self.alt_norm) if self.alt_norm else (self.norm,)

    def matches(self, label_norms: set[str]) -> bool:
        """セル全体がラベル候補に一致するか（セル内ラベル「設備番号：EQ-001」は含まない）。"""
        return self.norm in label_norms or (bool(self.alt_norm) and self.alt_norm in label_norms)

    def is_label(self, label_norms: set[str]) -> bool:
        return self.matches(label_norms) or (self.inline is not None and self.inline[0] in label_norms)


class SheetGrid:
    def __init__(self, ws):
        self.name: str = ws.title
        self.hidden: bool = getattr(ws, "sheet_state", "visible") != "visible"
        self._bounds: dict[tuple[int, int], tuple[int, int, int, int]] = {}
        for rng in ws.merged_cells.ranges:
            bounds = (rng.min_row, rng.min_col, rng.max_row, rng.max_col)
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    self._bounds[(r, c)] = bounds

        self.cells: dict[tuple[int, int], Cell] = {}
        self.max_row = self.max_col = 0
        for (r, c), xl in ws._cells.items():
            text = cell_text(xl.value)
            if not text:
                continue
            top, left, bottom, right = self.bounds(r, c)
            if (top, left) != (r, c):
                continue
            self.cells[(r, c)] = Cell(
                row=r, col=c, max_row=bottom, max_col=right,
                value=xl.value, text=text, norm=normalize_label(text), inline=split_inline(text),
                bold=bool(xl.font and xl.font.b),
                filled=bool(xl.fill and xl.fill.fill_type == "solid"),
                fill=fill_key(xl),
                alt_norm=section_stripped(text) if _label_like(xl.value, text) else "",
                part_norms=label_parts(text) if _label_like(xl.value, text) else (),
                fmt_unit=(format_unit(xl.number_format)
                          if isinstance(xl.value, (int, float)) and not isinstance(xl.value, bool) else ""),
            )
            self.max_row = max(self.max_row, bottom)
            self.max_col = max(self.max_col, right)

        self._by_norm: dict[str, list[Cell]] = defaultdict(list)
        self._by_inline: dict[str, list[Cell]] = defaultdict(list)
        self._by_base: dict[str, list[str]] = defaultdict(list)  # 括弧書きを除いたラベル → 元のキー
        self._by_variant: dict[str, list[str]] = defaultdict(list)  # 括弧書き・「内容」を除いたラベル → 元のキー
        self._by_part: dict[str, list[Cell]] = defaultdict(list)  # 「ライン／工程」の「ライン」「工程」
        for cell in self.text_cells():
            for key in cell.part_norms:
                self._by_part[key].append(cell)
            for key in cell.label_keys:
                self._by_norm[key].append(cell)
            if cell.inline:
                self._by_inline[cell.inline[0]].append(cell)
            for key in (*cell.label_keys, *(cell.inline[:1] if cell.inline else ())):
                base = paren_stripped(key)
                if base:
                    self._by_base[base].append(key)
                # 表記違いの照合は見出しらしいセル（塗りつぶし・太字・「ラベル：値」）だけ。値の「完了」を「完了日時」と見ない
                if cell.filled or cell.bold or (cell.inline and key == cell.inline[0]):
                    self._by_variant[label_base(key) or key].append(key)

    def bounds(self, row: int, col: int) -> tuple[int, int, int, int]:
        """(top, left, bottom, right) を返す。結合されていなければ自セルのみ。"""
        return self._bounds.get((row, col), (row, col, row, col))

    def cell_at(self, row: int, col: int) -> Cell | None:
        top, left, _, _ = self.bounds(row, col)
        return self.cells.get((top, left))

    def text_cells(self) -> list[Cell]:
        return [self.cells[key] for key in sorted(self.cells)]

    def resolve_labels(self, label_norms: set[str]) -> set[str]:
        """シートで探すラベルの集合。候補どおりのラベルが無いときだけ、末尾の括弧書きの違いを許す。

        例: 候補「停止時間」でシートに「停止時間(分)」だけがある（または逆）。「原因(推定)」と「原因(確定)」の
        ように括弧書きで区別しているラベルは、候補どおりのものがあればそちらだけを使う。
        """
        if any(n in self._by_norm or n in self._by_inline or n in self._by_part for n in label_norms):
            return label_norms
        bases = {paren_stripped(n) or n for n in label_norms}
        extra = {key for base in bases for key in self._by_base.get(base, [])}
        extra |= {base for base in bases if base in self._by_norm or base in self._by_inline}
        return label_norms | extra if extra else label_norms

    def label_variants(self, label_norms: set[str]) -> set[str]:
        """シートにある、候補と表記だけが違うラベル（候補どおりのラベルから値が取れなかったときに使う）。

        「発生原因」と「発生原因（なぜ起きたか）」、「応急処置」と「応急処置内容」。
        括弧書きどうしが違うもの（「原因（推定）」と「原因（確定）」）は別の項目なので含めない。
        """
        found = set()
        for norm in label_norms:
            qualified = bool(paren_stripped(norm))
            for key in self._by_variant.get(label_base(norm) or norm, []):
                if key in label_norms or (qualified and paren_stripped(key)):
                    continue
                found.add(key)
        return found

    def find_labels(self, label_norms: set[str]) -> list[Cell]:
        """ラベル候補に一致するセルを上→下、左→右の順で返す（セル内ラベルも含む）。

        明細表の列見出し（時系列表の「対応内容」、部品表の「備考」など）は、同じラベルが表の外にもあれば後回しにする。
        """
        from excel.tables import table_header_keys  # excel.tables が SheetGrid を使うため

        label_norms = self.resolve_labels(label_norms)
        hits = {}
        for norm in label_norms:
            for cell in self._by_norm.get(norm, []) + self._by_inline.get(norm, []) + self._by_part.get(norm, []):
                hits[(cell.row, cell.col)] = cell
        if len(hits) > 1:
            headers = table_header_keys(self)
            return [hits[key] for key in sorted(hits, key=lambda k: (k in headers, k))]
        return [hits[key] for key in sorted(hits)]


def _label_like(value, text: str) -> bool:
    return isinstance(value, str) and "\n" not in text and len(text) <= MAX_LABEL_LENGTH + 4


def fill_key(xl) -> str:
    """セルの塗りつぶし色を比較用の文字列にする（ラベル欄と同じ色かの判定に使う）。"""
    fill = getattr(xl, "fill", None)
    if fill is None or fill.fill_type != "solid":
        return ""
    color = fill.fgColor
    try:
        if color.type == "theme":
            return f"theme:{color.theme}:{round(color.tint or 0, 3)}"
        return f"{color.type}:{color.value}"
    except (AttributeError, TypeError, ValueError):
        return "solid"


@dataclass
class WorkbookInfo:
    path: Path | None   # 読んだファイル（メモリから読んだときは None）
    grids: dict[str, SheetGrid]
    images: list[dict]
    date1904: bool = False  # 1904年基準のブック（日付シリアル値の起点が違う）
    # 計算結果が保存されていない数式のセル {シート名: {(行, 列)}}（openpyxl などで書いたブック）。
    # data_only で読むと空になるので、「値が空」でなく「数式の結果が無い」と知らせるのに使う
    uncached_formulas: dict[str, set[tuple[int, int]]] = field(default_factory=dict)

    @property
    def sheet_names(self) -> list[str]:
        return list(self.grids)

    def images_in(self, sheet_names: list[str]) -> list[dict]:
        return [img for img in self.images if img["sheet"] in sheet_names]


def load_workbook_info(path: str | Path | io.BytesIO) -> WorkbookInfo:
    """ブックを読む。パスのほか、メモリの中のブック（BytesIO）も渡せる。

    読み方は同じで、どこから読むかだけが違う。帳票登録は見本の Excel をサーバーに置かずに読むので
    BytesIO を渡す（2026-09-21 の利用者の指示「見本のExcelは置かずに、設定だけ保持する」）。
    """
    def source():
        if isinstance(path, (str, Path)):
            return path
        path.seek(0)   # 同じブックを3回読む（openpyxl・画像・数式）ので、そのつど先頭へ戻す
        return path

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(source(), data_only=True)
    try:
        grids = {ws.title: SheetGrid(ws) for ws in wb.worksheets}
        date1904 = wb.epoch == MAC_EPOCH
    finally:
        wb.close()
    return WorkbookInfo(path=Path(path) if isinstance(path, (str, Path)) else None, grids=grids,
                        images=detect_images(source()), date1904=date1904,
                        uncached_formulas=uncached_formula_cells(source()))


def uncached_formula_cells(path: str | Path) -> dict[str, set[tuple[int, int]]]:
    """数式（<f>）があって計算結果（<v>）が無いセルを、シートごとに返す。読めないブックは空。"""
    out: dict[str, set[tuple[int, int]]] = {}
    try:
        with zipfile.ZipFile(path) as zf:
            names = set(zf.namelist())
            workbook_part = next((target for rtype, target in _rels(zf, names, "").values()
                                  if rtype.endswith("/officeDocument")), "xl/workbook.xml")
            if workbook_part not in names:
                return out
            workbook_rels = _rels(zf, names, workbook_part)
            for sheet in ET.fromstring(zf.read(workbook_part)).findall("main:sheets/main:sheet", NS):
                part = workbook_rels.get(sheet.get(R_ID), ("", ""))[1]
                if part not in names:
                    continue
                data = zf.read(part)
                if b"<f" not in data:  # 数式の無いシート（ほとんど）は読み直さない
                    continue
                cells = _uncached_in_sheet(data)
                if cells:
                    out[sheet.get("name")] = cells
    except (OSError, zipfile.BadZipFile, ET.ParseError, KeyError, ValueError):
        return {}
    return out


def _uncached_in_sheet(xml: bytes) -> set[tuple[int, int]]:
    cells: set[tuple[int, int]] = set()
    main = "{" + NS["main"] + "}"
    for _, elem in ET.iterparse(io.BytesIO(xml)):
        if elem.tag != main + "c":
            continue
        formula, value = elem.find(main + "f"), elem.find(main + "v")
        # 空文字の結果（t="str" の空の <v>）は計算済み
        if formula is not None and (value is None or (not value.text and elem.get("t") != "str")) and elem.get("r"):
            letters, row = coordinate_from_string(elem.get("r"))
            cells.add((row, column_index_from_string(letters)))
        elem.clear()
    return cells
