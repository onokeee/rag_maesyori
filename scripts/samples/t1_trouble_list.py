"""T1 トラブル対応一覧（きれいな一覧表・ベースライン）を作る。

    python -m scripts.samples.t1_trouble_list      # samples/tables/ に出力

出力:
- samples/tables/T1_トラブル対応一覧_2023-2026.xlsx
    シート「トラブル一覧」: standard_incidents() 全件（約8,000行）。1行目が見出しの素直な表。
        「対応内容」列は複数人が1セルに追記した対応ログ（_t1_log.py で Incident の日時から作る）。
        実日付セル・時刻セル・数値書式・ウィンドウ枠固定・Excelテーブル（ListObject、オートフィルタ付き）・
        列幅・長文列の折り返し・行高さの概算設定。
    シート「設備マスタ」: equipment_master() 全61台＋期間内の件数・停止時間（シートのオートフィルタ）。
    シート「集計メモ」: 保全課の担当者が書いた体裁の自由記述メモ（結合セル・小さな集計表が混在）。
- samples/tables/T1_トラブル対応一覧_2023-2026_README.md
    構成・文字コード・行数・意図的なイレギュラーの説明（件数は生成時に実データから計算）。

値はすべて domain.py の正準データから作るので、帳票系サンプルと管理No・設備番号・人名が一致する。
openpyxl は保存時刻を zip とプロパティに埋め込むため、保存後に固定時刻へ書き換えて
何度実行しても同一バイトのファイルになるようにしている。
"""
from __future__ import annotations

import collections
import math
import re
import time
import unicodedata
import zipfile
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from . import _t1_log
from . import domain as dm

STEM = "T1_トラブル対応一覧_2023-2026"
TABLE_NAME = "tblトラブル一覧"

# 保存時に埋め込む固定時刻（抽出日の翌朝に保存した体裁）
_FIXED_ZIP_TIME = (2026, 9, 2, 8, 30, 0)
_FIXED_PROP_TIME = datetime(2026, 9, 2, 8, 30, 0)

# ---------------------------------------------------------------------------
# 列定義: (見出し, 列幅, 種別)
#   種別 text=短い文字列 / wrap=長文（折り返し） / date / time / int / hours
# ---------------------------------------------------------------------------
COLUMNS = [
    ("管理No", 15.5, "text"),
    ("発生日", 11.5, "date"),
    ("発生時刻", 8.5, "time"),
    ("設備番号", 10.0, "text"),
    ("設備名", 22.0, "text"),
    ("ライン", 8.5, "text"),
    ("工程", 18.0, "text"),
    ("故障区分", 10.0, "text"),
    ("重要度", 7.5, "text"),
    ("現象", 44.0, "wrap"),
    ("アラーム", 26.0, "wrap"),
    ("原因", 36.0, "wrap"),
    ("処置内容", 50.0, "wrap"),
    ("対応内容", 60.0, "wrap"),
    ("使用部品", 34.0, "wrap"),
    ("停止時間(分)", 11.0, "int"),
    ("作業工数(h)", 10.5, "hours"),
    ("担当者", 20.0, "wrap"),
    ("状態", 9.0, "text"),
    ("完了日", 11.5, "date"),
    ("再発", 6.0, "text"),
    ("備考", 30.0, "wrap"),
]

NUMBER_FORMATS = {"date": "yyyy/mm/dd", "time": "hh:mm", "int": "#,##0", "hours": "0.00"}

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")
TOP_CENTER = Alignment(horizontal="center", vertical="top")
THIN = Side(style="thin", color="A6A6A6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MEMO_HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")


# ---------------------------------------------------------------------------
# 共通ヘルパー
# ---------------------------------------------------------------------------
def _disp_width(s: str) -> int:
    """表示幅（全角=2, 半角=1）の概算。"""
    return sum(2 if unicodedata.east_asian_width(c) in ("W", "F", "A") else 1 for c in s)


def _line_count(s: str, col_width: float) -> int:
    """折り返し表示したときの行数の概算（列幅は半角文字数相当）。"""
    usable = max(col_width - 1.5, 4)
    return sum(max(1, math.ceil(_disp_width(seg) / usable)) for seg in s.split("\n"))


def _parts_text(parts: list) -> str:
    """使用部品セル: 「部品名（品番）×数量」をセル内改行で並べる。"""
    return "\n".join(f"{p.name}（{p.part_no}）×{n}" for p, n in parts)


def _person_by_id(eid: str) -> dm.Person:
    return next(p for p in dm.people() if p.employee_id == eid)


def fiscal_year(d: date | datetime) -> int:
    """4月始まりの年度。"""
    return d.year if d.month >= 4 else d.year - 1


def save_workbook_deterministic(wb: Workbook, path: Path) -> Path:
    """zip エントリ時刻と docProps/core.xml の更新日時を固定して保存する（再実行で同一バイト）。"""
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    path.parent.mkdir(parents=True, exist_ok=True)
    stamp = _FIXED_PROP_TIME.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    with zipfile.ZipFile(buf) as src, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)", rb"\g<1>" + stamp + rb"\g<2>", data)
                data = re.sub(rb"(<dcterms:created[^>]*>)[^<]*(</dcterms:created>)", rb"\g<1>" + stamp + rb"\g<2>", data)
            zi = zipfile.ZipInfo(info.filename, date_time=_FIXED_ZIP_TIME)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            dst.writestr(zi, data)
    return path


# ---------------------------------------------------------------------------
# シート1: トラブル一覧
# ---------------------------------------------------------------------------
def incident_row(inc: dm.Incident, log: str | None = None) -> list:
    """Incident 1件を一覧の1行（COLUMNS と同順）に変換する。log は「対応内容」セル（_t1_log.response_logs の値）。"""
    eq = inc.equipment
    alarm = f"{inc.alarm.code} {inc.alarm.message}" if inc.alarm else None
    return [
        inc.incident_id,
        inc.occurred_at.date(),
        inc.occurred_at.time().replace(second=0, microsecond=0),
        eq.equipment_id,
        eq.name,
        eq.line,
        eq.process,
        inc.category,
        inc.severity,
        inc.symptom,
        alarm,
        inc.cause or None,
        inc.action,
        log,
        _parts_text(inc.parts_used) or None,
        inc.downtime_min,
        inc.work_hours,
        "、".join(p.name for p in inc.assignees),
        inc.status,
        inc.completed_at.date() if inc.completed_at else None,
        "有" if inc.recurrence else "無",
        inc.remarks or None,
    ]


def head_is_center(col: int) -> bool:
    """中央寄せにする短いコード系の列。"""
    return COLUMNS[col - 1][0] in ("設備番号", "ライン", "故障区分", "重要度", "状態", "再発")


def _write_trouble_sheet(ws, incidents: list[dm.Incident], logs: dict) -> int:
    ws.title = "トラブル一覧"
    kinds = [k for _, _, k in COLUMNS]
    widths = [w for _, w, _ in COLUMNS]

    # 見出し
    for c, (head, width, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=c, value=head)
        cell.font, cell.fill, cell.alignment = HEADER_FONT, HEADER_FILL, HEADER_ALIGN
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[1].height = 30

    # 明細
    for r, inc in enumerate(incidents, start=2):
        lines = 1
        for c, (value, kind, width) in enumerate(zip(incident_row(inc, logs[inc.incident_id]), kinds, widths), start=1):
            cell = ws.cell(row=r, column=c, value=value)
            if kind in NUMBER_FORMATS:
                cell.number_format = NUMBER_FORMATS[kind]
            if kind == "wrap":
                cell.alignment = WRAP_TOP
                if value:
                    lines = max(lines, _line_count(value, width))
            elif kind in ("date", "time") or head_is_center(c):
                cell.alignment = TOP_CENTER
            else:
                cell.alignment = TOP
        # 折り返し列の行数に合わせて行高さを概算（Excel は開いたときに自動調整しないため）
        ws.row_dimensions[r].height = min(15.0 * lines, 409.0)   # 409pt は Excel の行高さの上限

    last_row = len(incidents) + 1
    ref = f"A1:{get_column_letter(len(COLUMNS))}{last_row}"
    # Excelテーブル（ListObject）。オートフィルタはテーブル側の autoFilter 要素で持つ
    # （シートの autoFilter をテーブル範囲に重ねると Excel が修復ダイアログを出すため二重に付けない）
    table = Table(displayName=TABLE_NAME, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    ws.add_table(table)

    ws.freeze_panes = "B2"
    ws.sheet_view.zoomScale = 85
    ws.print_title_rows = "1:1"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    return last_row


# ---------------------------------------------------------------------------
# シート2: 設備マスタ
# ---------------------------------------------------------------------------
MASTER_COLUMNS = [
    ("設備番号", 10.0), ("設備名", 24.0), ("設備区分", 14.0), ("メーカー", 22.0), ("型式", 22.0), ("ライン", 10.0),
    ("工程", 20.0), ("設置場所", 24.0), ("設置日", 11.5), ("重要度ランク", 8.0), ("経過年数", 9.0),
    ("期間内トラブル件数", 10.0), ("期間内停止時間(h)", 11.0),
]


def _write_master_sheet(ws, incidents: list[dm.Incident]) -> int:
    counts = collections.Counter(i.equipment.equipment_id for i in incidents)
    down = collections.defaultdict(int)
    for i in incidents:
        down[i.equipment.equipment_id] += i.downtime_min

    for c, (head, width) in enumerate(MASTER_COLUMNS, start=1):
        cell = ws.cell(row=1, column=c, value=head)
        cell.font, cell.fill, cell.alignment, cell.border = HEADER_FONT, HEADER_FILL, HEADER_ALIGN, BOX
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[1].height = 30

    for r, eq in enumerate(dm.equipment_master(), start=2):
        age = round((dm.PERIOD_END - eq.installed_date).days / 365.25, 1)
        values = [eq.equipment_id, eq.name, eq.category, eq.maker, eq.model, eq.line, eq.process, eq.area,
                  eq.installed_date, eq.criticality, age, counts.get(eq.equipment_id, 0),
                  round(down.get(eq.equipment_id, 0) / 60, 1)]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BOX
        ws.cell(row=r, column=9).number_format = "yyyy/mm/dd"
        ws.cell(row=r, column=11).number_format = "0.0"
        ws.cell(row=r, column=12).number_format = "#,##0"
        ws.cell(row=r, column=13).number_format = "#,##0.0"
        ws.cell(row=r, column=10).alignment = Alignment(horizontal="center")

    last_row = len(dm.equipment_master()) + 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(MASTER_COLUMNS))}{last_row}"
    ws.freeze_panes = "B2"
    # 注記は表から1行空けて下に置く（経過年数・件数の基準）
    note = ws.cell(row=last_row + 2, column=1,
                   value=f"※経過年数は{dm.PERIOD_END:%Y/%m/%d}時点。件数・停止時間は「トラブル一覧」シートの集計値（瞬低の同時多発分を含む）。")
    note.font = Font(size=9, color="595959")
    return last_row


# ---------------------------------------------------------------------------
# シート3: 集計メモ（自由記述）
# ---------------------------------------------------------------------------
def _memo_table(ws, top: int, headers: list[str], rows: list[list], formats: dict[int, str] | None = None) -> int:
    """メモ内の小さな集計表を B列から書く。次に書ける行番号を返す。"""
    formats = formats or {}
    for c, h in enumerate(headers, start=2):
        cell = ws.cell(row=top, column=c, value=h)
        cell.font, cell.fill, cell.border = Font(bold=True), MEMO_HEAD_FILL, BOX
        cell.alignment = Alignment(horizontal="center")
    for r, row in enumerate(rows, start=top + 1):
        for c, v in enumerate(row, start=2):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BOX
            if (c - 2) in formats:
                cell.number_format = formats[c - 2]
    return top + len(rows) + 2


def _write_memo_sheet(ws, incidents: list[dm.Incident], extract_date: date) -> None:
    author = _person_by_id("M10544")      # 中村 浩二（保全1係 係長）
    updater = _person_by_id("M10688")     # 斎藤 健太郎（保全2係 係長）
    for col, width in zip("ABCDEFGHI", (3, 16, 24, 10, 10, 10, 10, 10, 14)):
        ws.column_dimensions[col].width = width

    ws["A1"] = "トラブル対応一覧　集計メモ（保全課内用・社外秘）"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:I1")
    ws["A2"] = (f"作成：{dm.surname(author)}（{author.section.split()[-1]}）{extract_date.year}/{extract_date.month}/{extract_date.day}"
                f"　／　追記：{dm.surname(updater)} {(extract_date + timedelta(days=2)).month}/{(extract_date + timedelta(days=2)).day}")
    ws["A2"].font = Font(size=9, color="595959")
    ws.merge_cells("A2:I2")

    n_all = len(incidents)
    status = collections.Counter(i.status for i in incidents)
    open_n = sum(1 for i in incidents if i.completed_at is None)
    first, last = incidents[0].occurred_at, incidents[-1].occurred_at

    row = 4
    ws.cell(row=row, column=1, value="■抽出条件").font = Font(bold=True)
    notes = [
        f"・保全DBから {extract_date:%Y/%m/%d} に抽出。発生日 {dm.PERIOD_START:%Y/%m/%d}〜{dm.PERIOD_END:%Y/%m/%d} の全{n_all:,}件"
        f"（実データの最初 {first:%Y/%m/%d}、最後 {last:%Y/%m/%d}）",
        f"・状態内訳：完了 {status['完了']:,}件／経過観察 {status['経過観察']:,}件／対応中 {status['対応中']:,}件／保留 {status['保留']:,}件"
        f"（完了日なし {open_n:,}件）",
        "・停止時間は「設備停止〜生産復帰」。対応中・保留は抽出時点までの暫定値なので、集計に使うときは注意",
        "・再発フラグはシステムが同一設備・同一故障モード30日以内で機械的に付けている。人が判断したものではない",
    ]
    for text in notes:
        row += 1
        ws.cell(row=row, column=1, value=text)
    row += 2

    # 年度×重要度の件数
    ws.cell(row=row, column=1, value="■年度別件数（重要度別）").font = Font(bold=True)
    row += 1
    by_fy = collections.defaultdict(collections.Counter)
    down_fy = collections.Counter()
    for i in incidents:
        fy = fiscal_year(i.occurred_at)
        by_fy[fy][i.severity] += 1
        down_fy[fy] += i.downtime_min
    rows = []
    for fy in sorted(by_fy):
        c = by_fy[fy]
        label = f"{fy}年度" + ("（〜8月）" if fy == fiscal_year(dm.PERIOD_END) else "")
        rows.append([label, c["重大"], c["大"], c["中"], c["小"], sum(c.values()), round(down_fy[fy] / 60, 1)])
    row = _memo_table(ws, row, ["年度", "重大", "大", "中", "小", "計", "停止時間(h)"], rows,
                      {1: "#,##0", 2: "#,##0", 3: "#,##0", 4: "#,##0", 5: "#,##0", 6: "#,##0.0"})

    # 停止時間ワースト10
    ws.cell(row=row, column=1, value="■停止時間ワースト10（期間累計）").font = Font(bold=True)
    row += 1
    per_eq = collections.defaultdict(list)
    for i in incidents:
        per_eq[i.equipment.equipment_id].append(i)
    worst = sorted(per_eq.items(), key=lambda kv: (-sum(x.downtime_min for x in kv[1]), kv[0]))[:10]
    rows = []
    for rank, (eid, lst) in enumerate(worst, start=1):
        top_sub, top_n = collections.Counter(x.subsystem for x in lst).most_common(1)[0]
        rows.append([rank, f"{eid} {lst[0].equipment.name}", len(lst), round(sum(x.downtime_min for x in lst) / 60, 1),
                     f"{top_sub}（{top_n}件）"])
    row = _memo_table(ws, row, ["順位", "設備", "件数", "停止時間(h)", "最多サブシステム"], rows,
                      {2: "#,##0", 3: "#,##0.0"})
    # 最多サブシステム列は幅が足りないので結合して読みやすくする
    for r in range(row - len(rows) - 2, row - 1):
        ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=8)

    # 気づき（数値は実データから計算し、事実と食い違わないようにする）
    ws.cell(row=row, column=1, value="■気づき・メモ").font = Font(bold=True)
    memo = []
    w_eid, w_lst = worst[0]
    w_sub, w_n = collections.Counter(x.subsystem for x in w_lst).most_common(1)[0]
    memo.append(f"・{w_eid}が停止時間ダントツ。{w_sub}起因が{w_n}件。{dm.MAKER_SHORT.get(w_lst[0].equipment.maker, '')}FEと定例で対策協議中")

    cvd = [i for i in incidents if i.equipment.equipment_id == "CVD-205" and i.subsystem == "ドライポンプ"]
    fix = date(2025, 4, 15)
    before = [i for i in cvd if i.occurred_at.date() <= fix]
    after = [i for i in cvd if i.occurred_at.date() > fix]
    m_before = (fix - dm.PERIOD_START).days / 30.44
    m_after = (dm.PERIOD_END - fix).days / 30.44
    if before and len(after) / m_after < len(before) / m_before:
        memo.append(f"・CVD-205 ドライポンプ：2025/4 DP更新以降は {len(before) / m_before:.1f}件/月 → {len(after) / m_after:.1f}件/月 に減少。効果あり")

    l6 = collections.Counter(fiscal_year(i.occurred_at) for i in incidents if i.equipment.line == "L6")
    if l6:
        memo.append("・L6立上げ機の初期故障：" + "、".join(f"{fy}年度 {n}件" for fy, n in sorted(l6.items()))
                    + "。立上げ後1年は保全2係で重点フォロー")

    external = [i for i in incidents if i.category == "外部要因"]
    power = [i for i in external if i.subsystem == "電源/受電"]
    summer = sum(1 for i in power if i.occurred_at.month in (7, 8, 9))
    if power:
        memo.append(f"・外部要因{len(external)}件のうち電源/受電（瞬低・停電）が{len(power)}件、その{summer}件が7〜9月の雷シーズン。"
                    "夏場前に施設係と復電手順を再確認")

    stale = [i for i in incidents if i.completed_at is None and (extract_date - i.occurred_at.date()).days > 60]
    memo.append(f"・完了日なしで発生から60日超の案件が{len(stale)}件。月末までに棚卸し（担当係長確認）")
    memo.append("・重要度の付け方は記入者でばらつきあり（特に中／大の境界）。停止時間と合わせて見ること")
    memo.append(f"　→ 追記({dm.surname(updater)})：備考欄の「品証へ連絡済み」はQA側の記録と突合せ未実施")
    for text in memo:
        row += 1
        ws.cell(row=row, column=1, value=text)


# ---------------------------------------------------------------------------
# README
# ---------------------------------------------------------------------------
_FULLWIDTH_RE = re.compile(r"[０-９Ａ-Ｚａ-ｚ]")
_HANKANA_RE = re.compile(r"[ｦ-ﾟ]")


def _log_section(incidents: list[dm.Incident], logs: dict, log_meta) -> str:
    """README の「対応内容列」の節（件数・割合は生成時の集計から作る）。"""
    n = len(incidents)
    m = log_meta
    filled = [logs[i.incident_id] for i in incidents if logs[i.incident_id]]
    nl = sum(1 for t in filled if "\n" in t)
    lens = sorted(len(t) for t in filled)
    ent = {int(k[2:]): v for k, v in m.items() if k.startswith("n:")}
    zen_cells = sum(1 for t in filled if _FULLWIDTH_RE.search(t))
    open_cells = m["open"]

    def pct(x: int) -> str:
        return f"{x:,}件（{x / n * 100:.1f}%）"

    dates = sorted(((k[5:], v) for k, v in m.items() if k.startswith("date:")), key=lambda kv: -kv[1])
    authors = sorted(((k[7:], v) for k, v in m.items() if k.startswith("author:")), key=lambda kv: -kv[1])
    joiners = {"newline": "セル内改行（1行1エントリ）", "slash": "「／」で1行に詰める", "maru": "改行なしで「。」だけでつなぐ",
               "arrow": "「→」で短くつなぐ（4/1 停止連絡→復旧→完了）"}
    joined = sorted(((k[7:], v) for k, v in m.items() if k.startswith("joiner:")), key=lambda kv: -kv[1])
    lines = [
        "## 対応内容列（追記型の対応ログ）",
        "",
        "保全・製造・生技・メーカーFEなど複数人が、発生から完了までの対応を1セルに追記していった体裁の列。"
        "アプリの AI整形／ログ分割（`logproc/`、`docs/research/対応内容AI整形設計.md` 2章）の主な入力を想定している。",
        "`scripts/samples/_t1_log.py` が各 Incident の 発生（occurred_at）・連絡（reported_at）・初動（response_started_at）・調査・原因・"
        "処置・使用部品・メーカー対応・結果（completed_at）・状態から作るので、他の列（発生日・原因・処置内容・使用部品・担当者・状態・完了日）と"
        "日時・内容が矛盾しない。",
        "",
        f"- セルの種類: 時系列ログ {pct(m['cell:時系列'])}、見出し型（【現象】【原因】【処置】…）{pct(m['cell:見出し型'])}、"
        f"「同上」{pct(m['cell:同上'])}（瞬低など同時刻の外部要因の2件目以降）、空欄・「-」{pct(m['cell:空欄'])}",
        "- 時系列ログのエントリ数（生成時の件数。重要度で 小2〜4／中3〜6／大4〜9／重大6〜12 を目標にし、時間の近いエントリは結合）: "
        + "、".join(f"{k}件 {v:,}セル" for k, v in sorted(ent.items())),
        f"- 文字数: 中央値 {lens[len(lens) // 2]:,}、90%点 {lens[int(len(lens) * 0.9)]:,}、最大 {lens[-1]:,}。セル内改行を含むセル {nl:,}",
        "- つなぎ方: " + "、".join(f"{joiners[k]} {v:,}" for k, v in joined),
        f"- 記入順: 新しい順（上へ追記）のセル {m['order:新しい順']:,}（全エントリに日付あり）。ほかは古い順",
        "- 締め方: 完了・経過観察は完了日時のエントリ（完了／クローズ／様子見）で締め、経過観察はその後の様子見の追記が付くことがある。"
        f"対応中・保留は締めのないまま「引き続き調査中」「部品入荷待ち」などで終わる（{open_cells:,}セル）",
        f"- メール転記（`-----Original Message-----` から `-----` まで。From/Sent/To/Subject・宛名・挨拶・結びを含む）: {m['email']:,}セル",
        "- 日付の書き方（エントリ単位の出現数）: " + "、".join(f"{k} {v:,}" for k, v in dates),
        "- 記入者の書き方（エントリ単位の出現数）: " + "、".join(f"{k} {v:,}" for k, v in authors),
        f"- 略語への置き換え（様子見・TEL済・交換済・FE来場）{m['abbr']:,} 箇所、チョコ停の記述 {m['abbr:チョコ停']:,} エントリ、"
        f"誤変換（様子身・異常ナシ・以上 など）{m['typo']:,} 箇所",
        f"- 全角英数字を含むセル {zen_cells:,}（記入者の癖。数値・日付だけで、管理No・品番・アラームコード・ロットIDの中は常に半角）",
        "",
        "書き方は記入者（`domain.people()` の社員番号）ごとに固定した癖で決まる。",
        "",
        "- 日付: `4/1 10:00`／`2024/04/01 10:00`／`4月1日 10時`／`R6.4.1`（まれ）／`【4/2 夜勤】`（夜勤の明け方は前日の日付）／全角 `４／１ １０：００`。"
        "同じ日の2件目以降は `10:20` のような時刻だけの行や `同日15時`、`翌日`・`翌朝`・`翌々日`・`翌週`・`3日後` も使う。日付のない行もある",
        "- 記入者: `田中：`／`田中:`／`田中 本文`／行末の `（田中）`・`(田中)`／イニシャル `K.T`（人物一覧に別名がないと特定できない）／"
        "`保全G 田中：`・`製造 長谷川：`・`生技：`／`メーカーFE：`・`扶桑FE 金子：`／引継ぎ `田中→佐藤 引継ぎ。`／直前と同じ人なら省略",
        "- 大・重大の完了セルは、最後のエントリの下に日付なしの `※再発防止：…`・`※水平展開：…` 行が付くことがある",
        "- 関連する管理Noへの言及（`TR-2024-00369と同じ現象、再発`、`同上（TR-…と同じ処置）`）",
        "",
    ]
    return "\n".join(lines)


def _readme_text(xlsx: Path, incidents: list[dm.Incident], stats: dict, logs: dict, log_meta) -> str:
    rows = [incident_row(i, logs[i.incident_id]) for i in incidents]
    idx = {h: k for k, (h, _, _) in enumerate(COLUMNS)}
    text_cols = ["現象", "原因", "処置内容", "備考"]

    def blank(col: str) -> int:
        return sum(1 for r in rows if r[idx[col]] in (None, ""))

    def distinct(col: str) -> int:
        return len({r[idx[col]] for r in rows if r[idx[col]]})

    zen = sum(1 for r in rows if any(_FULLWIDTH_RE.search(r[idx[c]] or "") for c in text_cols))
    hankana = sum(1 for r in rows if any(_HANKANA_RE.search(r[idx[c]] or "") for c in text_cols))
    nl_cols = [h for h, _, k in COLUMNS if k == "wrap"]
    nl = {c: sum(1 for r in rows if "\n" in (r[idx[c]] or "")) for c in nl_cols}
    status = collections.Counter(i.status for i in incidents)
    years = collections.Counter(i.occurred_at.year for i in incidents)
    numbering = collections.Counter()
    for i in incidents:
        head = i.action.lstrip()
        for mark in ("1.", "(1)", "1)", "①", "・"):
            if head.startswith(mark):
                numbering[mark] += 1
                break
    dup = sum(1 for i in incidents if "確認確認" in i.symptom + i.cause + i.action + i.remarks)
    ongoing = sum(1 for i in incidents if i.action.endswith("（継続対応中）"))
    simul = sum(1 for n in collections.Counter(i.occurred_at for i in incidents if i.category == "外部要因").values() if n > 1)

    col_lines = []
    for h, _, k in COLUMNS:
        desc = {
            "管理No": "トラブル管理番号 `TR-YYYY-NNNNN`（年ごとの連番。domain.py の incident_id そのままで、他のサンプルと共通）",
            "発生日": "Excel日付セル（書式 `yyyy/mm/dd`）",
            "発生時刻": "Excel時刻セル（書式 `hh:mm`、秒は切り捨て）",
            "設備番号": "`CMP-101` 形式。「設備マスタ」シートの設備番号と対応",
            "設備名": "設備マスタの設備名",
            "ライン": "`L1`〜`L6`。搬送系は `L1-L3` など範囲、ユーティリティは `Fab1共通` / `全Fab共通`",
            "工程": "設備マスタの工程名",
            "故障区分": "機械/電気/制御/ソフト/ユーティリティ/人為/品質/外部要因",
            "重要度": "重大/大/中/小",
            "現象": "記入者の書き方そのまま（長文）",
            "アラーム": "`コード メッセージ`。装置アラームが出ていない案件は空欄",
            "原因": "対応中の案件は空欄",
            "対応内容": "複数人が追記した対応ログ（日時・記入者・本文。書き方は記入者ごとに違う）。セル内改行あり。詳細は「対応内容列」の節",
            "処置内容": "番号付き手順（`1.` / `(1)` / `1)` / `①` / `・` など記入者で書式が違う）。セル内改行あり",
            "使用部品": "`部品名（品番）×数量` をセル内改行で列挙。部品を使っていなければ空欄",
            "停止時間(分)": "整数（書式 `#,##0`）",
            "作業工数(h)": "小数（書式 `0.00`、0.25h 刻み）",
            "担当者": "複数名は `、` 区切り（メーカーFEを含む）",
            "状態": "完了/経過観察/対応中/保留",
            "完了日": "Excel日付セル。対応中・保留は空欄",
            "再発": "`有` / `無`",
            "備考": "空欄が多い",
        }[h]
        col_lines.append(f"| {h} | {k} | {desc} |")

    return f"""# {xlsx.name}

製造部 設備保全課のトラブル対応一覧（2023年4月〜2026年8月発生分）。RAG前処理の **表データ（Excel一覧）** 用サンプルのうち、
いちばん素直な「きれいな一覧表」のベースライン。`scripts/samples/t1_trouble_list.py` が生成する（固定シードで決定的。再実行しても同じバイト列になる）。

## ファイル

| 項目 | 内容 |
|---|---|
| 形式 | Excel ブック（.xlsx、Office Open XML） |
| 文字コード | xlsx 内部の XML は UTF-8。この README も UTF-8（BOMなし、改行LF） |
| シート数 | 3（トラブル一覧 / 設備マスタ / 集計メモ） |
| 元データ | `scripts/samples/domain.py` の `standard_incidents()` と `equipment_master()` |

## シート「トラブル一覧」

- 1行目が見出し、2行目以降が1件1行。データ行 **{len(incidents):,} 行**（見出し含め {len(incidents) + 1:,} 行 × {len(COLUMNS)} 列、範囲 `{stats['table_ref']}`）
- 並び順は発生日時の昇順（同時刻は管理No順）
- Excelテーブル（ListObject）`{TABLE_NAME}`、スタイル TableStyleMedium2。オートフィルタはテーブルの autoFilter として付与
  （シート側の `ws.auto_filter` は設定していない）
- ウィンドウ枠固定 `B2`（見出し行と管理No列）、長文列は折り返し＋上揃え、行高さは折り返し行数から概算して設定
- 数式・結合セル・画像はなし
- 年別件数: {", ".join(f"{y}年 {n:,}件" for y, n in sorted(years.items()))}
- 状態別件数: {", ".join(f"{k} {v:,}件" for k, v in status.most_common())}

| 列 | 種別 | 内容 |
|---|---|---|
{chr(10).join(col_lines)}

{_log_section(incidents, logs, log_meta)}
## シート「設備マスタ」

- 1行目見出し、データ **{stats['master_rows']} 行 × {len(MASTER_COLUMNS)} 列**。シートのオートフィルタ `A1:{get_column_letter(len(MASTER_COLUMNS))}{stats['master_rows'] + 1}`、枠固定 `B2`
- 列: {" / ".join(h for h, _ in MASTER_COLUMNS)}
- 「期間内トラブル件数」「期間内停止時間(h)」は一覧シートの集計値（数式ではなく値）
- 表の2行下（{stats['master_rows'] + 3}行目）に注記テキストが1セルある

## シート「集計メモ」

保全課の係長が書いた体裁の自由記述メモ。**表ではない**。

- タイトル・作成者行は `A:I` の結合セル。見出し（■）と箇条書きは A 列に1セル1行
- 途中に小さな集計表が2つ（年度別件数、停止時間ワースト10）。ワースト10の「最多サブシステム」列は F:H を行ごとに結合
- 数値はすべて一覧シートから計算した値なので一覧と矛盾しない

## 意図的なイレギュラー（ベースラインなので構造は素直、崩れは主に中身）

1. **空欄**: 原因 {blank('原因'):,} 行（対応中）、対応内容 {blank('対応内容'):,} 行、完了日 {blank('完了日'):,} 行（対応中＋保留）、アラーム {blank('アラーム'):,} 行、使用部品 {blank('使用部品'):,} 行、備考 {blank('備考'):,} 行
2. **セル内改行**: {", ".join(f"{c} {n:,}行" for c, n in nl.items() if n)}
3. **全角英数字の混在**: 現象・原因・処置内容・備考のいずれかに全角英数字を含む行が {zen:,} 行（記入者の癖。管理No・設備番号・アラームコードは常に半角）
4. **半角カナ**: 同4列に半角カナ（ｱﾗｰﾑ、ﾎﾟﾝﾌﾟ など）を含む行が {hankana:,} 行
5. **表記ゆれ・誤変換**: 処置内容の手順番号の書式が記入者ごとに違う（`1.` 始まり {numbering['1.']:,} 行、`(1)` {numbering['(1)']:,} 行、`1)` {numbering['1)']:,} 行、`①` {numbering['①']:,} 行、`・` {numbering['・']:,} 行、ほかに全角 `(１)` など）。「異常→以上」「回収→改修」などの誤変換や、「確認確認」のような重複入力（{dup:,} 行）がまれに含まれる
6. **未完了案件の書き方**: 対応中 {status['対応中']:,} 件は原因が空欄で、処置内容の末尾が「（継続対応中）」（{ongoing:,} 行）。停止時間・工数は抽出時点の暫定値
7. **発生日と発生時刻が別セル**（日付シリアルと時刻シリアル）。完了日は日付のみで時刻なし
8. **ライン列の値域が混在**: 生産設備は `L1`〜`L6`、搬送系は `L1-L3`、ユーティリティは `Fab1共通` など
9. **瞬低などの同時多発**: 外部要因の案件は同じ発生時刻で複数設備に並ぶことがある（同時刻に2件以上並ぶ発生時刻が {simul:,} 組）
10. **1ブックに性質の違うシート**: 一覧（テーブル）・マスタ（オートフィルタのみ）・自由記述メモが同居

## 文章のばらつき（ユニーク値の数）

| 列 | ユニーク値 |
|---|---|
| 現象 | {distinct('現象'):,} |
| 原因 | {distinct('原因'):,} |
| 処置内容 | {distinct('処置内容'):,} |
| 使用部品 | {distinct('使用部品'):,} |
| 備考 | {distinct('備考'):,} |
"""


# ---------------------------------------------------------------------------
# エントリポイント
# ---------------------------------------------------------------------------
def generate(output_root: Path | None = None) -> list[Path]:
    """T1 の xlsx と README を書き出し、出力パスのリストを返す。"""
    root = Path(output_root) if output_root is not None else dm.OUTPUT_ROOT
    out_dir = root / "tables"
    out_dir.mkdir(parents=True, exist_ok=True)

    incidents = sorted(dm.standard_incidents(), key=lambda i: (i.occurred_at, i.incident_id))
    completed = [i.completed_at for i in incidents if i.completed_at]
    extract_date = max(max(completed).date(), dm.PERIOD_END)   # 最後の完了より後に抽出した体裁

    wb = Workbook()
    wb.properties.creator = _person_by_id("M10544").name
    wb.properties.lastModifiedBy = _person_by_id("M10688").name
    wb.properties.title = "トラブル対応一覧 2023-2026"
    wb.properties.created = _FIXED_PROP_TIME

    logs, log_meta = _t1_log.response_logs(incidents, extract_date)
    last_row = _write_trouble_sheet(wb.active, incidents, logs)
    master_last = _write_master_sheet(wb.create_sheet("設備マスタ"), incidents)
    _write_memo_sheet(wb.create_sheet("集計メモ"), incidents, extract_date)

    xlsx = save_workbook_deterministic(wb, out_dir / f"{STEM}.xlsx")
    stats = {"table_ref": f"A1:{get_column_letter(len(COLUMNS))}{last_row}", "master_rows": master_last - 1}
    readme = out_dir / f"{STEM}_README.md"
    with open(readme, "w", encoding="utf-8", newline="\n") as f:
        f.write(_readme_text(xlsx, incidents, stats, logs, log_meta))
    return [xlsx, readme]


if __name__ == "__main__":
    t0 = time.time()
    for p in generate():
        print(p, f"{p.stat().st_size / 1024:,.0f} KB")
    print(f"{time.time() - t0:.1f}s")
