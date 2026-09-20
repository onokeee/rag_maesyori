"""F2 トラブル対応報告書（初動〜恒久対策）のサンプルExcel帳票を生成する。

    python -m scripts.samples.f2_trouble_report      # samples/forms/F2_トラブル対応報告書/ に出力

- 正準トラブル履歴（domain.standard_incidents）から、なぜなぜ分析があり重要度「中」以上の案件を30件選ぶ。
- 様式は2版。報告日が 2025/04/01 より前は「旧様式 Rev.3」（方眼紙レイアウト・押印欄が右上）、
  以降は「新様式 Rev.5」（ラベル左・値右の表形式・承認欄が末尾）。移行後も旧様式を使い回した報告を2件混ぜる。
- xlsx は様式の版ごとのサブフォルダ（旧様式Rev3_2025年03月まで／新様式Rev5_2025年04月改訂）に分けて出力する。
  1つのフォルダの中は同じ様式だけなので、フォルダごと帳票取り込みにまとめて投入できる。
- 各ファイルの正解値を _expected.jsonl に、様式説明と意図的なゆらぎを _README.md に書き出す。
- 乱数は domain.rng() からのみ取り、xlsx の zip タイムスタンプ・文書プロパティも固定するため、
  再実行してもバイト単位で同一のファイルになる。
"""
from __future__ import annotations

import io
import json
import math
import re
import shutil
import time as _time
import unicodedata
import zipfile
from dataclasses import dataclass, field, replace
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.properties import PageSetupProperties
from PIL import Image as PILImage, ImageDraw, ImageFont

from . import domain as D

# ---------------------------------------------------------------------------
# 基本設定
# ---------------------------------------------------------------------------
FOLDER = "F2_トラブル対応報告書"
N_FILES = 30
SWITCH_DATE = date(2025, 4, 1)          # 新様式 Rev.5 の運用開始日
TODAY = date(2026, 9, 15)               # これより後の日付の承認は「未承認」のまま
SEV_QUOTA = {"重大": 6, "大": 12, "中": 12}
N_PHOTO_FILES = 12

OLD, NEW = "旧様式Rev.3", "新様式Rev.5"

# 様式の版ごとのフォルダ（1フォルダ＝1様式。フォルダごと帳票取り込みにまとめて投入できるようにする）
VERSION_DIR = {
    OLD: "旧様式Rev3_2025年03月まで",
    NEW: "新様式Rev5_2025年04月改訂",
}

# 様式ごとの項目ラベル（_expected.jsonl の labels_used にもそのまま使う）
LABELS = {
    OLD: dict(
        banner="設備トラブル報告書", form_no="保全様式 F2-01（Rev.3）", title="件名", report_id="管理No", report_date="作成日",
        occurred_at="発生日時", recovered_at="復旧日時", finder="発見者", reporter="作成者", equipment_id="設備No",
        equipment_name="設備名", line="ライン", process="工程", assignee="対応者", severity="重要度", cause_category="原因区分",
        failure_category="故障区分", impact="影響", downtime_min="停止時間", lots="影響ロット", scrap_wafers="廃棄枚数",
        loss_yen="損失金額", related_report="関連No", recurrence="再発", timeline="１．時系列", symptom="２．現象",
        alarm="発生アラーム", investigation="調査結果", action="３．暫定対策", parts="使用部品", result="復旧確認",
        why_why="４．なぜなぜ分析", event="事象", cause="５．真因", prevention="６．恒久対策", horizontal_section="７．水平展開",
        horizontal_options="区分", horizontal_targets="対象設備", horizontal_deployment="展開内容", impression="８．所感",
        approver="承認", checker="確認", creator_stamp="作成",
    ),
    NEW: dict(
        banner="トラブル対応報告書（初動〜恒久対策）", form_no="様式F2 Rev.5（2025.04改訂）", title="表題", report_id="報告書No",
        report_date="報告日", occurred_at="発生日時", recovered_at="復旧日時", finder="第一発見者", detected_by="検知方法",
        reporter="報告者", equipment_id="対象設備", equipment_name="対象設備", line="ライン／工程", process="ライン／工程",
        assignee="対応者", severity="重要度", failure_category="故障区分", status="対応状況", cause_category="原因区分",
        impact="影響", downtime_min="ダウンタイム", lots="影響ロット", scrap_wafers="スクラップ枚数", loss_yen="損失額",
        related_report="関連報告", timeline="■ 経緯（時系列）", symptom="発生現象", alarm="アラーム",
        investigation="確認・調査結果", action="応急処置", parts="使用部品", result="処置結果", why_why="■ なぜなぜ分析",
        cause="根本原因", prevention="再発防止策（恒久対策）", event="発生事象", horizontal_section="■ 水平展開", horizontal_options="展開区分",
        horizontal_targets="設備No", horizontal_deployment="展開先・内容",
        impression="所感・コメント", supervisor_comment="上長コメント", approver="承認", checker="確認", creator_stamp="作成",
        recurrence="再発有無",
    ),
}

SEVERITIES = ("重大", "大", "中", "小")
FAIL_CATS = ("機械", "電気", "制御", "ソフト", "ユーティリティ", "人為", "品質", "外部要因")
CIRCLED = "①②③④⑤"

# ---------------------------------------------------------------------------
# 罫線・塗り・文字幅
# ---------------------------------------------------------------------------
THIN = Side(style="thin", color="000000")
HAIR = Side(style="hair", color="7F7F7F")
MED = Side(style="medium", color="000000")
DOUBLE = Side(style="double", color="000000")


def _fill(rgb: str) -> PatternFill:
    return PatternFill("solid", fgColor=rgb)


def _dw(s: str) -> int:
    """表示幅（全角=2, 半角=1）。"""
    return sum(2 if unicodedata.east_asian_width(ch) in "WFA" else 1 for ch in s)


class Grid:
    """セル結合・罫線・行高さ調整をまとめた簡易レイアウトヘルパー。"""

    def __init__(self, ws, widths: list[float], font_name: str, font_size: float, line_h: float):
        self.ws = ws
        self.widths = [0.0] + list(widths)
        self.font_name, self.font_size, self.line_h = font_name, font_size, line_h
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    @property
    def ncols(self) -> int:
        return len(self.widths) - 1

    def width(self, c1: int, c2: int) -> float:
        return sum(self.widths[c1:c2 + 1])

    def box(self, r1, c1, r2, c2, value=None, *, fill=None, bold=False, size=None, color=None, h="left", v="center",
            wrap=True, border=THIN, fmt=None, shrink=False, italic=False):
        ws = self.ws
        if (r1, c1) != (r2, c2):
            ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
        cell = ws.cell(r1, c1)
        if value is not None:
            cell.value = value
        cell.font = Font(name=self.font_name, size=size or self.font_size, bold=bold, color=color, italic=italic)
        # 縮小表示と折り返しは両立しない（折り返しが優先される）ため、縮小時は折り返さない
        cell.alignment = Alignment(horizontal=h, vertical=v, wrap_text=wrap and not shrink, shrink_to_fit=shrink)
        if fmt:
            cell.number_format = fmt
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                x = ws.cell(r, c)
                if border is not None:
                    x.border = Border(left=border if c == c1 else x.border.left, right=border if c == c2 else x.border.right,
                                      top=border if r == r1 else x.border.top, bottom=border if r == r2 else x.border.bottom)
                if fill is not None:
                    x.fill = fill
        return cell

    def outline(self, r1, c1, r2, c2, side=MED):
        """範囲の外周だけ太線にする（内側の罫線は維持）。"""
        ws = self.ws
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                if r not in (r1, r2) and c not in (c1, c2):
                    continue
                b = ws.cell(r, c).border
                ws.cell(r, c).border = Border(left=side if c == c1 else b.left, right=side if c == c2 else b.right,
                                              top=side if r == r1 else b.top, bottom=side if r == r2 else b.bottom)

    def lines(self, text, c1, c2) -> int:
        cap = max(4.0, self.width(c1, c2) - 1.5)
        n = 0
        for line in str(text or "").split("\n"):
            n += max(1, math.ceil(_dw(line) * 1.08 / cap))
        return n

    def fit(self, r, text, c1, c2, min_h=None, pad=5.0):
        """結合セル内の文章が収まる行高さにする（Excel は結合セルを自動調整しないため）。"""
        need = self.lines(text, c1, c2) * self.line_h + pad
        h = max(min_h or 0, need)
        cur = self.ws.row_dimensions[r].height or 0
        self.ws.row_dimensions[r].height = min(409, max(cur, h))

    def height(self, r, h):
        self.ws.row_dimensions[r].height = h

    def paginate(self, marks: list[int], last_row: int, margins_lr: float = 0.9):
        """区切り行（セクション先頭など）の単位で改ページを入れ、表や箱がページをまたがないようにする。

        横1ページに合わせる縮小率から、1ページに入るシート上の高さ(pt)を見積もる。
        """
        ws = self.ws
        width_px = sum(w * 7 + 5 for w in self.widths[1:])
        scale = min(1.0, (8.27 - margins_lr) * 72 / (width_px * 0.75))
        cap = (11.69 - 1.2) * 72 / scale * 0.93
        row_h = lambda r: ws.row_dimensions[r].height or 15.0
        bounds = sorted(set(m for m in marks if 1 <= m <= last_row)) + [last_row + 1]
        used = 0.0
        for s, e in zip(bounds, bounds[1:]):
            h = sum(row_h(r) for r in range(s, e))
            if used + h <= cap:
                used += h
                continue
            if h <= cap and used > 0:
                ws.row_breaks.append(Break(id=s - 1))
                used = h
                continue
            for r in range(s, e):          # 1ページに収まらない大きな塊は行単位で区切る
                if used + row_h(r) > cap and used > 0:
                    ws.row_breaks.append(Break(id=r - 1))
                    used = 0.0
                used += row_h(r)


# ---------------------------------------------------------------------------
# 文面の加工
# ---------------------------------------------------------------------------
def _prefix_patterns() -> re.Pattern:
    """domain の現象文プレフィックス（「夜勤帯、」「{time}頃、」など）を除去する正規表現。"""
    pats = []
    for key in ("any", "日勤", "夜勤", "休日"):
        for p in D._SYM_PREFIX[key]:
            if not p:
                continue
            for v in {p, D.to_hankaku_kana(p)}:
                s = re.escape(v)
                s = s.replace(r"\{time\}", r"[0-9０-９]{1,2}:[0-9０-９]{2}")
                s = s.replace(r"\{lot\}", r"[A-Z]{2}[0-9]{7}\.[0-9]")
                s = s.replace(r"\{n\}", r"[0-9０-９]+")
                pats.append(s)
    return re.compile("^(?:" + "|".join(sorted(pats, key=len, reverse=True)) + ")")


_PREFIX_RE = _prefix_patterns()
_NUM_PREFIX = re.compile(r"^\s*(?:[0-9０-９]+[.)．]\s*|\([0-9０-９]+\)\s*|[①-⑳]\s*|・\s*)")
_GENERIC_INV_HEADS = ("アラーム履歴を確認", "前回PMは", "目視点検では", "作業前に装置停止", "過去の類似トラブル", "班長への聞き取り",
                      "発生直前の処理", "同型機")


def short_symptom(inc: D.Incident) -> str:
    """件名・事象欄に使う短い現象（先頭文、時間帯などの前置きと末尾の括弧書きを除く）。"""
    s = inc.symptom.split("。")[0]
    s = _PREFIX_RE.sub("", s)
    s = re.sub(r"（[^）]*）$", "", s).strip("、 ")
    if _dw(s) > 56:
        parts = s.split("、")
        out = parts[0]
        for p in parts[1:]:
            if _dw(out + "、" + p) > 56:
                break
            out += "、" + p
        s = out
    return s


def split_steps(action: str) -> list[str]:
    """番号付きの処置文を1手順ずつに分ける（番号の書き癖は除去）。"""
    out = []
    for line in action.split("\n"):
        t = _NUM_PREFIX.sub("", line).strip()
        if t and t != "（継続対応中）":
            out.append(t)
    return out


# 恒久対策の文に付いてくる「状況」を表す断片（それだけでは対策にならない）: 断片 → (備考欄の書き方, 状況)
_PREV_STATUS_FRAGMENTS = {
    "保全カレンダーへ登録済み": ("保全カレンダー登録済", "計画"),
    "点検チェックシートに反映": ("チェックシート反映済", "完了"),
}


def split_prevention(text: str) -> list[tuple[str, str | None]]:
    """恒久対策の文を表の行 (対策内容, 状況の断片 or None) に分ける。

    - 「／」「。」は対策の区切りなので行を分ける
    - 「→」では分けない（「交換周期を60日→30日に短縮」の変更内容や「A→B」の流れを1つの対策として残す）
    - 対策の後ろに「→保全カレンダーへ登録済み」のような状況の断片が付いている場合や、断片だけが区切られて
      書かれている場合は、直前の対策行の状況・備考として扱い、単独の対策行にはしない
    """
    rows: list[tuple[str, str | None]] = []
    for chunk in re.split(r"\s*(?:／|。)\s*", text or ""):
        chunk = chunk.strip()
        if not chunk:
            continue
        if chunk in _PREV_STATUS_FRAGMENTS:
            if rows and rows[-1][1] is None:
                rows[-1] = (rows[-1][0], chunk)
                continue
            rows.append((chunk, None))         # 断片しか書かれていない場合は書かれたとおり1行にする
            continue
        note = None
        m = re.search(r"\s*→\s*(" + "|".join(map(re.escape, _PREV_STATUS_FRAGMENTS)) + r")$", chunk)
        if m and m.start() > 0:
            note, chunk = m.group(1), chunk[:m.start()].strip()
        rows.append((chunk, note))
    return rows


def _sanitize(name: str) -> str:
    return re.sub(r'[\\/:*?"<>|]', "・", name)


_POLITE_FIXED = [("ていなかった。", "ていませんでした。"), ("だった。", "でした。"), ("得ない。", "得ません。"), ("しまった。", "しまいました。"),
                 ("かかった。", "かかりました。"), ("迷った。", "迷いました。"), ("なった。", "なりました。"), ("たい。", "たいです。"),
                 ("必要。", "必要です。"), ("反省点。", "反省点です。"), ("一つ。", "一つです。"), ("起点。", "起点です。"),
                 ("推定原因。", "推定原因です。"), ("及んだ。", "及びました。"), ("手間取った。", "手間取りました。"), ("欲しい。", "欲しいです。"), ("ある。", "あります。"), ("いる。", "います。"), ("できる。", "できます。"),
                 ("おく。", "おきます。"), ("べき。", "べきと考えます。")]
# 一段動詞・サ変の過去形（「感じた」「防げた」「記録した」「出ていた」）は「〜ました」にする
_POLITE_ICHIDAN = re.compile(r"((?<![まで])し|[いきちにひみりえけせてねへめれげぜでべぺじび])た。")   # 変換済みの「ました」「でした」は除く


def _polite_extra(s: str) -> str:
    """所感を敬体に寄せる（敬体で書く癖のある記入者用）。"""
    for a, b in _POLITE_FIXED:
        s = s.replace(a, b)
    s = _POLITE_ICHIDAN.sub(r"\1ました。", s)
    return s.replace("かった。", "かったです。")      # 形容詞の過去形（大きかった）


def _clean_text(s: str) -> str:
    """元データに残っている編集記号（任意項目マーカーの「?」）を除く。例「清掃・?交換」→「清掃・交換」。"""
    return re.sub(r"(?<=[・、\n（])\?|^\?", "", s or "", flags=re.M)


def _check_line(options, selected, mark="☑", blank="□", sep="　") -> str:
    return sep.join(f"{mark if o in selected else blank}{o}" for o in options)


def _iso(dt) -> str:
    """正規化した日付・日時（_expected.jsonl の *_norm 用）。"""
    if isinstance(dt, datetime):
        return dt.strftime("%Y-%m-%d %H:%M")
    return dt.strftime("%Y-%m-%d")


# Excel の表示形式を適用したときの見え方（_expected.jsonl の値は「セルに表示されている文字列」で持つ）
def _ymd(d: date, pad: bool = False) -> str:
    """yyyy/m/d（pad=True なら yyyy/mm/dd）。"""
    return f"{d.year}/{d.month:02d}/{d.day:02d}" if pad else f"{d.year}/{d.month}/{d.day}"


def _hm(t: datetime | time, pad: bool = False) -> str:
    """h:mm（pad=True なら hh:mm）。"""
    return f"{t.hour:02d}:{t.minute:02d}" if pad else f"{t.hour}:{t.minute:02d}"


def _md(d: date) -> str:
    """m/d。"""
    return f"{d.month}/{d.day}"


# ---------------------------------------------------------------------------
# 報告書の中身（様式に依存しない部分）
# ---------------------------------------------------------------------------
@dataclass
class TLRow:
    at: datetime
    text: str
    who: str


@dataclass
class PermAction:
    content: str
    owner: str
    due: date | None
    due_text: str | None       # 「5月末」など文字列で書く場合
    status: str
    remarks: str = ""          # 備考（「保全カレンダー登録済」など状況の補足）


@dataclass
class Target:
    equipment_id: str
    name: str
    checked: bool
    when: str                  # 実施日／予定の欄に書く文字（2025/02/25、10/24まで、2月PM など）
    result: str
    when_date: date | None = None   # when が具体的な日付を指す場合の日付


@dataclass
class Report:
    inc: D.Incident
    version: str
    report_date: date
    title: str
    event: str
    writer: D.Person
    checker: D.Person
    approver: D.Person
    checked_on: date | None
    approved_on: date | None
    recovered_at: datetime
    timeline: list
    perm: list
    hz_options: list
    targets: list
    impression: str
    supervisor_comment: str
    photos: list = field(default_factory=list)
    style: dict = field(default_factory=dict)
    irregularities: list = field(default_factory=list)
    file_name: str = ""


_OCC_DET = ("装置アラーム", "中央監視", "FDC監視")

# 初動の断片のうち、同じ行動を指すもの（リセット、保全呼出し、DOWN連絡、装置停止）は最初の1つだけ残す
_FR_DUP_GROUPS = (("リセット",), ("保全コール", "保全呼出し", "保全呼び出し", "保全へ連絡", "保全へPHS"),
                  ("ダウン連絡", "DOWN連絡"), ("装置を停止", "着工停止"))
# 時系列上の自然な順番（断片の先頭に近いキーワードで判定）: リセット → 状態確認・記録 → 停止・隔離 → ロット処置 → 報告・連絡
_FR_ORDER = ((0, ("リセット",)),
             (1, ("写真", "記録", "モニタ", "状態を確認", "状態確認", "照合", "リストアップ", "アラーム内容を確認")),
             (2, ("停止", "インヒビット", "LOTO", "立入り制限", "切替")),
             (3, ("退避", "振替", "HOLD")),
             (4, ("報告", "連絡", "コール", "呼出し", "呼び出し", "引継ぎ")))
_HANDOVER_WORDS = ("生産再開", "引渡し", "引き渡し", "DOWN解除", "MES解除", "復旧連絡", "着工再開")


def _tidy_first_response(pieces: list[str]) -> list[str]:
    """初動欄の断片から重複する行動を除き、行動の順に並べ替える（半角カナの書き癖は NFKC で吸収して判定）。"""
    kept, seen = [], set()
    for p in pieces:
        n = unicodedata.normalize("NFKC", p)
        groups = {i for i, words in enumerate(_FR_DUP_GROUPS) if any(w in n for w in words)}
        if groups & seen:
            continue
        seen |= groups
        kept.append(p)

    def order(p: str) -> float:
        n = unicodedata.normalize("NFKC", p)
        hits = [(n.find(w), rank) for rank, words in _FR_ORDER for w in words if w in n]
        return min(hits)[1] if hits else 1.5

    return sorted(kept, key=order)


def _is_handover(step: str) -> bool:
    return any(w in unicodedata.normalize("NFKC", step) for w in _HANDOVER_WORDS)


def build_timeline(inc: D.Incident, r, end_at: datetime, report_date: date) -> list[TLRow]:
    """発生〜初動〜調査〜処置〜復帰〜事後対応の時系列（5〜15行）を作る。"""
    eq = inc.equipment
    mnt = [p for p in inc.assignees if p.role != "FE"]
    fe = [p for p in inc.assignees if p.role == "FE"]
    maker = D.MAKER_SHORT.get(eq.maker, eq.maker)
    rep_sn = D.surname(inc.reporter)
    m_names = "・".join(D.surname(p) for p in mnt[:2]) or "保全"
    short = short_symptom(inc)

    def who_m(i=0):
        return D.surname(mnt[i % len(mnt)]) if mnt else "保全"

    # --- 発生 ---
    if inc.alarm is not None and inc.alarm.code not in short:
        occ = r.choice([f"{short}（{inc.alarm.code}）", f"装置アラーム {inc.alarm.code} {inc.alarm.message} 発報",
                        f"{inc.alarm.code} 発生、{short}", f"{short}。{inc.alarm.code}表示"])
    elif inc.detected_by not in _OCC_DET and inc.detected_by != "オペレーター":
        occ = r.choice([f"{inc.detected_by}にて異常検知：{short}", f"{short}（{inc.detected_by}）"])
    else:
        occ = r.choice([short, f"{short} 発生", f"異常発生：{short}"])
    who_occ = "装置" if inc.detected_by in _OCC_DET and D.kind_of(eq) not in D.UTILITY_KINDS else rep_sn
    if inc.detected_by == "中央監視":
        who_occ = "中央監視"

    # --- 初動（発生〜連絡の間） ---
    fr = inc.first_response
    if "→" in fr:
        fr_pieces = [x.strip() for x in fr.split("→") if x.strip()]
    elif "。" in fr:
        fr_pieces = [x.strip() for x in fr.split("。") if x.strip()]
    else:
        fr_pieces = [fr.strip()] if fr.strip() else []
    # 「保全〇〇が△分後に到着」は到着行と重複するので落とす
    fr_pieces = [x for x in fr_pieces if not re.search(r"分後に到着$", x)] or fr_pieces[:1]
    fr_pieces = _tidy_first_response(fr_pieces)
    fr = "、".join(fr_pieces)

    # --- 連絡 ---
    sect = inc.reporter.section
    rep_is_mnt = inc.reporter in mnt            # 保全巡回・施設係など、保全自身が見つけた案件
    rep_who = rep_sn
    if sect.startswith("設備保全課"):
        if D.kind_of(eq) in D.UTILITY_KINDS:
            # 初動欄に書いた行動（係長への報告、現地確認）と重ならない書き方を選ぶ
            opts = ["係長へ第一報、関係先へ連絡", "施設係内で情報共有、応援要請", "関係部署（製造課・生産管理）へ第一報"]
            if not fr_pieces:
                opts.append("中央監視から施設係へ連絡、現地確認へ")
            rep_text = r.choice([o for o in opts if "係長" not in fr or "係長" not in o])
        else:
            rep_text = r.choice(["保全課内へ第一報", "係長へ報告、応援要請"])
    elif sect in ("生産技術課", "品質保証課"):
        rep_text = r.choice([f"{sect.replace('課', '')}より保全へ連絡", f"{sect}から保全へ調査依頼", f"{sect}より連絡受け（{rep_sn}）"])
        if "受け" in rep_text:
            rep_who = who_m(0)
    elif any(w in unicodedata.normalize("NFKC", fr) for w in ("保全へ", "保全呼出し", "保全コール", "保全呼び出し")):
        # 初動欄で既に保全を呼んでいる場合は、保全側の受付として書く（記入者も保全側）
        rep_text = r.choice(["保全受付、担当者手配", "保全PHSで受付", f"保全係で受付（{sect.replace('製造課 ', '')}より）"])
        rep_who = who_m(0)
    else:
        opts = ["保全へPHS連絡", "保全呼出し", "班長経由で保全へ連絡", f"保全へ連絡（{sect.replace('製造課 ', '')} {rep_sn}）"]
        if "生産管理" not in fr and "DOWN" not in unicodedata.normalize("NFKC", fr).upper():
            opts.append("生産管理へDOWN連絡、保全呼出し")
        rep_text = r.choice(opts)
    others = [p for p in mnt if p != inc.reporter]
    arrive_who = who_m(0)
    if rep_is_mnt:
        # 発見者が保全員なら「到着」ではなく、応援の合流や点検開始として書く
        if others:
            o_names = "・".join(D.surname(p) for p in others[:2])
            arrive = r.choice([f"応援（{o_names}）到着、LOTO実施し点検開始", f"{o_names} 合流、詳細点検開始"])
            arrive_who = D.surname(others[0])
        else:
            arrive = r.choice(["LOTO実施し点検開始", "工具・測定器を準備し点検開始", "詳細点検開始"])
            arrive_who = rep_sn
    else:
        arrive = r.choice([f"保全 {m_names} 到着、現場確認", f"保全対応開始（{m_names}）", f"{m_names} 現場到着。アラーム履歴・状態確認",
                           f"保全到着（{m_names}）、LOTO実施し点検開始"])

    # --- 調査 ---
    sents = [x.strip() for x in inc.investigation.split("。") if x.strip()]
    specific = [x for x in sents if not x.startswith(_GENERIC_INV_HEADS)]
    inv_rows = (specific or sents)[:3]
    has_fe_text = any("FE" in x for x in inv_rows)

    body: list[tuple[str, str, float, str]] = []   # (内容, 担当, 直前の間隔の重み, 種別)
    for k, s in enumerate(inv_rows):
        body.append((s, who_m(k), 1.0, "inv"))
        if k == 0 and fe and not has_fe_text:
            body.append((f"{maker}FEへ連絡、来場依頼", who_m(0), 0.6, "fe_call"))
            body.append((f"{maker}FE（{D.surname(fe[0])}）来場、共同で調査", f"{maker}FE", 4.0, "fe_come"))
    expensive = [p for p, _ in inc.parts_used if p.unit_price_yen >= 30_000]
    parts_wait = bool(expensive) and inc.downtime_min >= 480 and r.random() < 0.75
    if parts_wait:
        pn = expensive[0].name
        body.append((r.choice([f"部品手配（{pn}、在庫なし→{maker}へ緊急手配）", f"{pn}を予備品倉庫から出庫依頼",
                               f"{pn}の手配。納期確認"]), who_m(1), 0.8, "parts"))
    steps = split_steps(inc.action)
    # 引渡し・生産再開の手順は、QC確認などの後ろ（処置の最後）に回す
    steps = [s for s in steps if not _is_handover(s)] + [s for s in steps if _is_handover(s)]
    arrived_part = False
    for k, s in enumerate(steps):
        if parts_wait and not arrived_part and ("交換" in s or k == len(steps) // 2):
            body.append((r.choice(["部品入荷、交換作業開始", "部品到着", "手配部品受領、作業再開"]), who_m(1), 7.0, "parts_in"))
            arrived_part = True
        body.append((s, f"{maker}FE" if fe and "FE" in s else who_m(k % 2), 1.0 + (0.8 if "QC" in s or "確認" in s else 0), "step"))

    # --- 復帰（処置手順と同じ内容を繰り返さない） ---
    resumed = any(w in s for s in steps for w in ("再開", "引渡し", "復旧連絡", "DOWN解除"))
    if inc.status == "保留":
        end_text = r.choice(["部品発注済み、入荷後に本処置予定。監視条件を製造課と申し合わせ", "恒久処置は保留、暫定運用の注意点を引継ぎ簿に記入"])             if resumed else r.choice(["暫定処置にて運転再開（部品入荷待ち）", "暫定条件で生産再開、恒久処置は保留", "代替運用で生産継続（監視強化）"])
    elif inc.status == "経過観察":
        end_text = r.choice(["経過観察としてトレンド監視を開始", "製造課へ経過観察中である旨を連絡、初品確認OK"])             if resumed else r.choice(["生産復帰。トレンド監視を継続", "DOWN解除、経過観察として生産再開", "製造課へ引渡し（経過観察）"])
    elif resumed:
        # 処置手順に引渡しが書かれている場合、復帰行は稼働確認として書く
        end_text = r.choice(["稼働状態を確認、異常なし", "初品処理完了を確認", "生産復帰を確認（製造課）"])
    else:
        end_text = r.choice(["生産復帰（DOWN解除）", "製造課へ引渡し、生産再開", "復旧確認完了、MES解除", "着工再開"])
    end_who = who_m(0)

    # --- 事後対応 ---
    post = []
    if inc.lots_affected and r.random() < 0.7:
        post.append((r.choice(["品証へ影響ロットの判定依頼", "影響ロットのHOLD解除可否を品証・生技と協議"]), who_m(0), 2))
    if inc.severity in ("重大", "大") and r.random() < 0.65:
        post.append((r.choice(["対策会議（保全・生技・品証）", "臨時対策会議にて原因・対策を審議", f"{maker}へ不具合報告書の提出を依頼"]), who_m(0), 20))

    # --- 行数を 5〜15 に収める（書く人によって細かさが違うので上限も案件ごとに変える） ---
    max_rows = r.choices(range(6, 16), weights=[2, 3, 4, 5, 5, 5, 4, 4, 3, 3])[0]

    def total():
        return 4 + len(post) + len(fr_pieces) + len(body)   # 4 = 発生・連絡・到着・復帰

    def merge_steps(min_steps: int) -> bool:
        """隣り合う処置手順を2つずつ1行にまとめる（手順行が min_steps 行になるまで）。"""
        idx = [i for i, b in enumerate(body) if b[3] == "step"]
        pairs = [i for i in idx if i + 1 in idx]
        if len(idx) <= min_steps or not pairs:
            return False
        i = pairs[len(pairs) // 2]
        a, b = body[i], body[i + 1]
        body[i:i + 2] = [(f"{a[0]}、{b[0]}", a[1], a[2], "step")]
        return True

    def drop(kind: str, keep: int) -> bool:
        idx = [i for i, b in enumerate(body) if b[3] == kind]
        if len(idx) > keep:
            body.pop(idx[-1])
            return True
        return False

    def join_first(limit: int) -> bool:
        nonlocal fr_pieces
        if len(fr_pieces) > limit:
            fr_pieces = ["、".join(fr_pieces)]
            return True
        return False

    # 削る順番: 調査の3行目 → 初動の集約 → 手順の集約 → 事後対応 → … 処置の行は最低1行残す
    reducers = [lambda: drop("inv", 2), lambda: join_first(2), lambda: merge_steps(2), lambda: bool(post) and bool(post.pop()),
                lambda: join_first(1), lambda: drop("inv", 1), lambda: drop("parts", 0), lambda: drop("fe_call", 0),
                lambda: merge_steps(1), lambda: drop("parts_in", 0), lambda: drop("fe_come", 0)]
    while total() > max_rows:
        if not any(f() for f in reducers):
            break
    while total() < 5:
        fr_pieces.append(r.choice(["班長へ状況報告", "生産管理へ停止連絡", "仕掛りロットの状態確認"]))

    # --- 時刻の割り付け ---
    rows: list[TLRow] = [TLRow(inc.occurred_at, occ, who_occ)]
    gap = (inc.reported_at - inc.occurred_at).total_seconds() / 60
    for i, p in enumerate(fr_pieces):
        t = inc.occurred_at + timedelta(minutes=int(gap * (i + 1) / (len(fr_pieces) + 1)))
        rows.append(TLRow(t, p, rep_sn))
    rows.append(TLRow(inc.reported_at, rep_text, rep_who))
    rows.append(TLRow(inc.response_started_at, arrive, arrive_who))
    span = (end_at - inc.response_started_at).total_seconds() / 60
    ws = [b[2] * r.uniform(0.6, 1.4) for b in body] + [1.0]
    tot = sum(ws)
    step = 5 if span / (len(body) + 1) >= 20 else 1
    cum = 0.0
    prev = inc.response_started_at
    for (text, who, _, _), w in zip(body, ws):
        cum += w
        t = inc.response_started_at + timedelta(minutes=span * cum / tot)
        t = t.replace(second=0, microsecond=0)
        t -= timedelta(minutes=t.minute % step)
        t = min(max(t, prev), end_at)
        rows.append(TLRow(t, text, who))
        prev = t
    rows.append(TLRow(end_at, end_text, end_who))
    limit = datetime.combine(report_date, time(12, 0))
    t = end_at
    for text, who, hours in post:
        t = t + timedelta(hours=hours * r.uniform(0.5, 1.5))
        if t.hour < 8 or t.hour >= 19:
            t = datetime.combine(t.date() + timedelta(days=1 if t.hour >= 19 else 0), time(r.choice([9, 10, 13]), r.choice([0, 30])))
        if t >= limit:
            break
        rows.append(TLRow(t, text, who))
    return rows


def _owner_for(text: str, inc: D.Incident, r) -> str:
    ppl = D.people()
    if any(k in text for k in ("FDC", "SPC", "レシピ", "生産技術", "承認フロー")):
        return D.surname(r.choice([p for p in ppl if p.section == "生産技術課"]))
    if any(k in text for k in ("メーカー", "申し入れ", "改造", "設計")):
        return f"{D.MAKER_SHORT.get(inc.equipment.maker, inc.equipment.maker)}／{D.surname(inc.assignees[0])}"
    if any(k in text for k in ("標準書", "教育", "手順", "朝会", "共有")):
        chief = [p for p in ppl if p.section == inc.assignees[0].section and p.role in ("係長", "主任")]
        return D.surname(chief[0]) if chief else D.surname(inc.assignees[0])
    mnt = [p for p in inc.assignees if p.role != "FE"]
    return D.surname(r.choice(mnt[:2])) if mnt else "保全"


def _done_day(r, recovered_on: date, latest: date) -> date:
    """実施済みの作業の日付（復旧日〜latest の間。トラブル発生・復旧より前の日付にはしない）。"""
    return recovered_on + timedelta(days=r.randint(0, max(0, (latest - recovered_on).days)))


def _hz_latest(report_date: date, checked_on: date | None, approved_on: date | None) -> date:
    """水平展開の実施日として書ける最も遅い日。

    報告書は作成後も確認・承認までは追記されるため、作成日から2週間後までの実施結果は書き込まれうる。
    ただし押印（確認・承認）後は追記しないので、最後の押印日を上限にする。
    """
    latest = min(report_date + timedelta(days=14), TODAY)
    stamps = [d for d in (checked_on, approved_on) if d is not None]
    if stamps:
        latest = min(latest, max(stamps))
    return max(latest, report_date)


def build_perm(inc: D.Incident, r, report_date: date, recovered_on: date, due_text_style: bool) -> list[PermAction]:
    items = split_prevention(inc.prevention)
    out = []
    for text, note in items:
        if text.startswith(("特になし", "現状の点検で対応可")):
            out.append(PermAction(text, "―", None, "―", "―"))
            continue
        if text.startswith("恒久対策は"):
            out.append(PermAction(text, D.surname(inc.assignees[0]), None, "未定", "保留"))
            continue
        due_text = None
        remarks, fixed_status = _PREV_STATUS_FRAGMENTS[note] if note else ("", None)
        if fixed_status == "完了" or (fixed_status is None and "済" in text):
            # 報告書作成時点で実施済みの対策は、期限＝実施日（復旧日〜報告日）として書く
            due = _done_day(r, recovered_on, report_date)
            status = "完了"
        else:
            due = report_date + timedelta(days=r.randint(7, 80))
            if due_text_style and r.random() < 0.6:
                due = (date(due.year + (due.month == 12), due.month % 12 + 1, 1) - timedelta(days=1))
                due_text = f"{due.month}月末"
            if fixed_status:
                # 保全カレンダーに登録済み＝日程は決まったが未実施
                status = r.choice(["計画", "計画", "実施中"])
            else:
                status = r.choices(["完了", "実施中", "計画"], weights=[25, 35, 40])[0]
        out.append(PermAction(text, _owner_for(text, inc, r), due, due_text, status, remarks))
    return out


HZ_OPTIONS = ("同型機", "類似設備", "他Fab", "課内共有", "展開不要")


def build_targets(inc: D.Incident, r, report_date: date, recovered_on: date, latest: date) -> tuple[list[str], list[Target]]:
    """水平展開欄: 展開区分のチェックと対象設備リストを、記入された展開内容の文から組み立てる。

    実施済みの展開の実施日は、トラブルの復旧日〜 latest（作成日＋14日、押印済みなら最後の押印日まで）の間から選ぶ
    （トラブル発生・復旧より前の日付にならないように）。
    """
    eq = inc.equipment
    text = inc.horizontal_deployment
    kind = D.kind_of(eq)
    ids_in = re.findall(r"[A-Z]{3}-\d{3}", text)
    sibs = [e for e in D.equipment_master() if D.kind_of(e) == kind and e is not eq
            and (e.installed_date <= report_date or e.equipment_id in ids_in)]
    by_id = {e.equipment_id: e for e in D.equipment_master()}

    def done(eid: str, name: str, result: str) -> Target:
        d = _done_day(r, recovered_on, latest)
        return Target(eid, name, True, d.strftime("%Y/%m/%d"), result, d)

    if not text:
        return [], []
    if "異常なし" in text and ids_in:
        tg = [done(e.equipment_id, e.name, "異常なし") if e.equipment_id in ids_in else Target(e.equipment_id, e.name, False, "", "")
              for e in sibs]
        return ["同型機"], tg
    if "展開予定" in text:
        m = re.search(r"（(\d+)月）", text)
        when = f"{m.group(1)}月PM" if m else "次回PM"
        return ["同型機"], [Target(i, by_id[i].name, False, when, "未実施") for i in ids_in]
    if "点検予定" in text:
        m = re.search(r"(\d+)日までに", text)
        mm = re.search(r"担当 (\S+?)）", text)
        day = int(m.group(1)) if m else 28
        y, mo = report_date.year, report_date.month
        if day <= report_date.day:
            y, mo = (y + 1, 1) if mo == 12 else (y, mo + 1)
        when = f"{mo}/{day}まで"
        return ["同型機"], [Target(i, by_id[i].name, False, when, f"担当 {mm.group(1)}" if mm else "", date(y, mo, day)) for i in ids_in]
    if "他Fab" in text:
        other = [e for e in sibs if D.fab_of(e) != D.fab_of(eq)] or sibs
        return ["同型機", "他Fab"], [done(e.equipment_id, e.name, "展開済み") for e in other]
    if "類似構造" in text:
        sim = [e for e in D.equipment_master() if e.category == eq.category and e is not eq and D.fab_of(e) == D.fab_of(eq)][:2]
        return ["類似設備"], [done(e.equipment_id, e.name, "点検実施、異常なし") for e in sim]
    if "本機固有" in text:
        return ["展開不要"], []
    if "定例会" in text:
        return ["課内共有"], []
    if "検討中" in text:
        return [], [Target(e.equipment_id, e.name, False, "", "要否検討中") for e in sibs[:4]]
    return [], []


# ---------------------------------------------------------------------------
# 所感・上長コメント
# ---------------------------------------------------------------------------
_IMP = {
    "long": ["復旧まで{hours}時間を要し、生産への影響が大きかった。", "停止が{hours}時間に及んだ。切り分けに時間がかかった点は反省点。",
             "実作業よりも部品待ち・確認待ちの時間が長く、段取りの悪さを感じた。", "復旧までの時間が長く、途中経過の連絡が製造側に十分届いていなかった。"],
    "short": ["初動の切り分けが早く、比較的短時間で復旧できた。", "アラーム内容から早い段階で原因箇所を絞り込めた。",
              "過去の類似事例を保全DBで引けたので、迷わず対応できた。"],
    "夜勤": ["夜間で人手が少なく、応援要請の判断が遅れた。", "夜勤帯の対応で部品倉庫の開錠に時間を取られた。夜間の部品持ち出しルールを見直したい。",
             "深夜の呼出しで到着まで時間がかかった。夜間の一次対応をオペレーターでどこまでできるか整理が必要。"],
    "休日": ["休日対応のため呼出しから到着まで時間を要した。", "休日で担当者が不在、連絡網が古いままだった。更新が必要。"],
    "recur": ["前回（{rel}）と同一箇所の再発で、前回の対策が不十分だったと言わざるを得ない。",
              "再発案件。前回は暫定処置で止めてしまい、恒久対策が遅れていたことが要因の一つ。",
              "同様の不具合が続いている。部品交換で済ませず、根本的な対策をメーカーと詰める必要がある。"],
    "scrap": ["ウェーハ{scrap}枚の廃棄が発生。検知がもう少し早ければ被害を抑えられた。",
              "品質影響（廃棄{scrap}枚）が出てしまった。インラインでの早期検知の仕組みを生技と検討したい。"],
    "lots": ["影響ロットの判定に品証との調整が必要で、ロット処置の決定まで時間がかかった。", "影響ロットの洗い出しに手間取った。処理履歴をすぐ引ける仕組みが欲しい。"],
    "fe": ["{fe}FEとの連携はスムーズで、ログ解析も迅速だった。", "FE到着まで待機が発生した。一次切り分けを社内でできるよう教育を依頼したい。",
           "FE頼みの対応になった。社内に知見が残るよう、作業内容を写真付きで記録した。"],
    "wear": ["経年劣化による故障。設置{age}年目の設備であり、予防保全の比重を上げるべき。",
             "消耗部品の寿命管理が使用時間ベースになっておらず、交換時期を逃した。", "劣化の兆候はトレンドに出ていた。見る仕組みがあれば防げた。"],
    "dirt": ["汚れ・付着が起点。清掃周期と清掃方法の標準化が必要。", "清掃の品質が作業者によってばらついている印象がある。"],
    "human": ["手順の思い込みによるミス。ダブルチェックの仕組みが形骸化していた。", "作業手順書が現状と合っていない箇所があり、改訂が必要。"],
    "unknown": ["原因を特定しきれていない。再発時に確実にログ・写真を取れるよう準備しておく。", "現時点では推定原因。経過観察で効果を確認したい。"],
    "hold": ["部品納期が長く、暫定運用が続く。予備品の持ち方を見直したい。", "恒久処置が部品待ちで保留。暫定運用中の監視を製造側と申し合わせた。"],
    "good": ["オペレーターの早期発見・連絡に助けられた。", "チョコ停の段階で兆候が出ていたので、情報を拾えていれば未然防止できた。",
             "若手メンバーにとって良い教育の機会になった。", "関係部署の協力で早期復旧できた。感謝。", "写真とログをその場で残したので、報告書作成と原因解析が楽だった。"],
    "extern": ["外部要因で防ぎようがない面もあるが、復旧手順の整備で停止時間は短縮できる。", "複数設備の同時停止で、復旧の優先順位付けに迷った。"],
}
_SUP = ["水平展開の完了まで確実にフォローすること。", "恒久対策の効果確認を{m}月の定例で報告のこと。", "同型機の点検結果を一覧で提出願います。",
        "類似事例がないか保全DBを再確認すること。", "FE依存にならないよう、社内での一次切り分け力の強化を。", "了解。対策期限の遵守をお願いします。",
        "費用対効果も含め、部品の予防交換を検討すること。", "品質影響があったため、品証とも対策内容をすり合わせること。", "対応ご苦労様でした。"]


def build_impression(inc: D.Incident, r, writer: D.Person) -> str:
    keys = []
    hours = round(inc.downtime_min / 60, 1)
    if inc.recurrence and inc.related_incident_id:
        keys.append("recur")
    if inc.scrap_wafers:
        keys.append("scrap")
    elif inc.lots_affected:
        keys.append("lots")
    if inc.status == "保留":
        keys.append("hold")
    if inc.status == "経過観察" or inc.cause_category == "不明":
        keys.append("unknown")
    if inc.shift in ("夜勤", "休日"):
        keys.append(inc.shift)
    if any(p.role == "FE" for p in inc.assignees):
        keys.append("fe")
    if inc.category == "外部要因":
        keys.append("extern")
    if inc.cause_category in ("摩耗", "劣化"):
        keys.append("wear")
    elif inc.cause_category in ("汚れ", "異物"):
        keys.append("dirt")
    elif inc.cause_category in ("設定ミス", "作業ミス") or inc.category == "人為":
        keys.append("human")
    keys.append("long" if hours >= 8 else "short")
    keys.append("good")
    r.shuffle(keys)
    chosen = sorted(keys[: r.choice([2, 2, 3])], key=lambda k: list(_IMP).index(k))
    age = int((inc.occurred_at.date() - inc.equipment.installed_date).days / 365.25) + 1
    params = dict(hours=f"{hours:g}", rel=inc.related_incident_id or "", scrap=inc.scrap_wafers,
                  fe=D.MAKER_SHORT.get(inc.equipment.maker, inc.equipment.maker), age=age)
    text = "".join(r.choice(_IMP[k]).format(**params) for k in chosen)
    if D._writer_habit(writer.employee_id)["polite"]:
        text = _polite_extra(text)
    return text


# ---------------------------------------------------------------------------
# 対象案件の選定
# ---------------------------------------------------------------------------
def _report_date(inc: D.Incident, r) -> date:
    base = (inc.completed_at or inc.occurred_at + timedelta(minutes=inc.downtime_min)).date()
    d = base + timedelta(days=r.randint(1, 9))
    while d.weekday() >= 5:
        d += timedelta(days=1)
    return d


def select_incidents() -> list[D.Incident]:
    """なぜなぜ分析あり・重要度「中」以上から、状態・再発・品質影響・設備種別がばらけるよう30件選ぶ。"""
    r = D.rng("f2-select")
    cands = [i for i in D.standard_incidents()
             if i.why_why and i.severity in SEV_QUOTA and i.status in ("完了", "経過観察", "保留")]
    r.shuffle(cands)
    chosen: list[D.Incident] = []
    quota = dict(SEV_QUOTA)
    used_scen, used_eq, used_cat = set(), {}, {}

    # 報告日が旧様式期間になる案件 13件＋新様式期間 17件（うち2件は旧様式を使い回すので、様式としては15:15）
    period_cap = {True: 13, False: N_FILES - 13}

    def before(i):
        return i.occurred_at.date() < SWITCH_DATE - timedelta(days=14)

    def ok(i):
        if (i.completed_at or i.occurred_at).date() >= SWITCH_DATE - timedelta(days=14) and before(i):
            return False       # 様式切替の直前（報告日が前後どちらになるか曖昧な案件）は使わない
        return (quota[i.severity] > 0 and i.scenario_id not in used_scen and used_eq.get(i.equipment.equipment_id, 0) < 2
                and used_cat.get(i.equipment.category, 0) < 5 and sum(1 for c in chosen if before(c) == before(i)) < period_cap[before(i)])

    def take(i):
        chosen.append(i)
        quota[i.severity] -= 1
        used_scen.add(i.scenario_id)
        used_eq[i.equipment.equipment_id] = used_eq.get(i.equipment.equipment_id, 0) + 1
        used_cat[i.equipment.category] = used_cat.get(i.equipment.category, 0) + 1


    reqs = [(lambda i: i.status == "保留", 2), (lambda i: i.status == "経過観察", 3),
            (lambda i: i.recurrence and i.related_incident_id, 3), (lambda i: i.scrap_wafers > 0, 4),
            (lambda i: len(i.why_why) == 5, 3), (lambda i: len(i.why_why) == 3, 2),
            (lambda i: i.category == "外部要因", 1), (lambda i: i.shift == "休日", 2)]
    for pred, n in reqs:
        got = 0
        for i in cands:
            if got >= n:
                break
            if i not in chosen and pred(i) and ok(i):
                take(i)
                got += 1
    # 残りを補充（期間ごとの上限は ok() で管理）
    for i in cands:
        if len(chosen) >= N_FILES:
            break
        if i not in chosen and ok(i):
            take(i)
    return sorted(chosen, key=lambda i: i.occurred_at)


def build_reports() -> list[Report]:
    incs = select_incidents()
    r = D.rng("f2-reports")
    ppl = D.people()
    boss = next(p for p in ppl if p.role == "課長")
    reports: list[Report] = []
    for inc in incs:
        # 元データの編集記号の残りを除いた写しを使う（正準データ自体は変更しない）
        inc = replace(inc, **{f: _clean_text(getattr(inc, f)) for f in ("symptom", "investigation", "action", "result")})
        rr = D.rng(f"f2:{inc.incident_id}")
        rd = _report_date(inc, rr)
        version = OLD if rd < SWITCH_DATE else NEW
        writer = next((p for p in inc.assignees if p.role != "FE"), inc.assignees[0])
        chief = [p for p in ppl if p.section == writer.section and p.role == "係長" and p != writer]
        if not chief:
            chief = [p for p in ppl if p.section == writer.section and p.role == "主任" and p != writer] or \
                    [p for p in ppl if p.name == "斎藤 健太郎"]
        checker = chief[0]
        checked_on = rd + timedelta(days=rr.randint(0, 2))
        approved_on = checked_on + timedelta(days=rr.randint(0, 5))
        if inc.status == "保留":
            approved_on = None
            if rr.random() < 0.5:
                checked_on = None
        elif inc.status == "経過観察" and rr.random() < 0.5:
            approved_on = None
        if checked_on and checked_on > TODAY:
            checked_on = None
        if approved_on and (approved_on > TODAY or checked_on is None):
            approved_on = None
        end_at = inc.completed_at or inc.occurred_at + timedelta(minutes=inc.downtime_min)
        short = short_symptom(inc)
        eq = inc.equipment
        stop_word = "品質異常" if inc.category == "品質" else "設備停止"
        # 現象が用言（「温度が低い」「上がらない」など）で終わる場合は「〜の件」「〜による」を付けない
        verbal = "ぁ" <= short[-1] <= "ゟ"
        title_opts = [f"{eq.equipment_id} {short}", f"【{eq.name}】{short}{'件' if verbal else 'の件'}",
                      f"{eq.name}（{eq.equipment_id}）{inc.subsystem}トラブル", f"{eq.equipment_id} {inc.subsystem}不具合に伴う{stop_word}"]
        if not verbal and not short.endswith(("停止", "異常", "低下")):
            title_opts.append(f"{short}による{stop_word}（{eq.equipment_id}）")
        rep = Report(
            inc=inc, version=version, report_date=rd, title=rr.choice(title_opts), event=short, writer=writer, checker=checker,
            approver=boss, checked_on=checked_on, approved_on=approved_on, recovered_at=end_at,
            timeline=build_timeline(inc, rr, end_at, rd), perm=[], hz_options=[], targets=[],
            impression=build_impression(inc, rr, writer),
            supervisor_comment=rr.choice(_SUP).format(m=(rd.month + 2) % 12 + 1) if rr.random() < 0.55 else "",
        )
        reports.append(rep)
    # 移行後に旧様式を使い回した報告（2件）
    after = [x for x in reports if x.version == NEW]
    for x in r.sample(after, k=min(2, len(after))):
        x.version = OLD
        x.irregularities.append("新様式運用開始（2025/04）後に旧様式テンプレートを使用")
    # 様式ごとの書き方のゆらぎ（各ゆらぎが最低1件は出るよう、通し番号で割り当てる）
    olds = [x for x in reports if x.version == OLD]
    news = [x for x in reports if x.version == NEW]
    for k, x in enumerate(olds):
        rr = D.rng(f"f2-style:{x.inc.incident_id}")
        x.style = dict(
            date_text=k % 4 == 1, downtime=["num", "text_hm", "num", "text_min", "num"][k % 5], loss_man=k % 7 == 3,
            mark="■" if k % 6 == 4 else "☑", time_text=k % 4 == 2, extra_sheets=k % 3 == 0, label_nl=k % 5 == 2,
            due_text=k % 3 == 1, writer_sect=k % 6 == 5, sheet=["報告書", "トラブル報告", "F2"][k % 3],
            approver_missing=False, name_pat=rr.choice([0, 0, 2, 3]),
        )
    # 完了なのに承認印が漏れている旧様式を1件
    miss = [x for x in olds[3:] if x.inc.status == "完了" and x.approved_on is not None]
    if miss:
        miss[0].style["approver_missing"] = True
    for k, x in enumerate(news):
        rr = D.rng(f"f2-style:{x.inc.incident_id}")
        x.style = dict(
            loss_man=False, mark="☑", due_text=k % 4 == 3, impression_blank=k % 9 == 4, dt_text=k % 5 == 3,
            sheet=["報告書", x.inc.incident_id, "報告書"][k % 3], name_pat=rr.choice([1, 1, 2, 3]),
        )
        if x.style["impression_blank"]:
            x.impression = ""
    for x in reports:
        rr = D.rng(f"f2-detail:{x.inc.incident_id}")
        x.perm = build_perm(x.inc, rr, x.report_date, x.recovered_at.date(), x.style.get("due_text", False))
        if x.style.get("approver_missing"):
            x.approved_on = None
        x.hz_options, x.targets = build_targets(x.inc, rr, x.report_date, x.recovered_at.date(),
                                                _hz_latest(x.report_date, x.checked_on, x.approved_on))
    # 写真を貼る報告（写真キャプションのある案件から N_PHOTO_FILES 件）
    with_photo = [x for x in reports if x.inc.photos]
    for x in D.rng("f2-photo-pick").sample(with_photo, k=min(N_PHOTO_FILES, len(with_photo))):
        x.photos = list(x.inc.photos)
    _collect_irregularities(reports)
    _name_files(reports)
    return reports


def _collect_irregularities(reports: list[Report]):
    for x in reports:
        st, inc, irr = x.style, x.inc, x.irregularities
        if x.version == OLD:
            if st["date_text"]:
                irr.append("発生日・復旧日・作成日が和暦の文字列（例 R6.2.29）")
            irr.append("発生日時が日付セルと時刻セルに分割")
            if st["downtime"] == "text_hm":
                irr.append("停止時間が「○時間○分」の文字列")
            elif st["downtime"] == "text_min":
                irr.append("停止時間が「○分」の文字列")
            if st["loss_man"]:
                irr.append("損失金額が「約○万円」の概算文字列")
            if st["mark"] == "■":
                irr.append("チェック記号が☑ではなく■")
            if st["time_text"]:
                irr.append("時系列の時刻が「3:12頃」の文字列（日付は日付が変わった行だけ記入）")
            else:
                irr.append("時系列の日付は日付が変わった行だけ記入")
            if st["extra_sheets"]:
                irr.append("記入要領シートと非表示のリストシートあり")
            if st["label_nl"]:
                irr.append("ラベルにセル内改行（例「停止\\n時間」）")
            if st["writer_sect"]:
                irr.append("作成者欄が「係名＋姓」表記")
            if st["approver_missing"]:
                irr.append("状態は完了だが承認印が未押印")
            if x.photos:
                irr.append("写真は別シート「写真」に貼付")
        else:
            irr.append("対象設備が「設備No＋設備名」の1セル、ライン／工程も1セル")
            if st["dt_text"]:
                irr.append("時系列の日時が文字列（例 8/11 1:45）")
            if st["impression_blank"]:
                irr.append("所感・コメント欄が未記入")
            if x.photos:
                irr.append("写真は報告書シート末尾の「添付写真」欄に貼付")
        if st.get("due_text"):
            irr.append("恒久対策の期限が「○月末」の文字列")
        if not x.perm:
            irr.append("恒久対策の表が空欄（未記入）")
        if not inc.horizontal_deployment:
            irr.append("水平展開欄が未記入")
        if len(inc.why_why) < 5:
            irr.append(f"なぜなぜ分析は{len(inc.why_why)}段まで記入（残りの枠は空欄）")
        if x.approved_on is None:
            irr.append("承認欄が空欄（未承認）")
        if inc.status == "保留":
            irr.append("対応状況が保留（復旧日時は暫定、恒久対策未定）")
        if not inc.lots_affected:
            irr.append("影響ロットなし（「なし」と記入）")


def _name_files(reports: list[Report]):
    used = set()
    for x in reports:
        inc, eq = x.inc, x.inc.equipment
        pats = [f"{inc.incident_id}_トラブル報告書_{eq.equipment_id}.xlsx",
                f"トラブル対応報告書_{eq.equipment_id}_{x.report_date:%Y%m%d}.xlsx",
                f"【{inc.severity}】{inc.incident_id}({eq.equipment_id} {inc.subsystem}).xlsx",
                f"F2_{inc.incident_id}.xlsx"]
        name = _sanitize(pats[x.style["name_pat"]])
        if name in used:
            name = _sanitize(pats[0])
        used.add(name)
        x.file_name = name


# ---------------------------------------------------------------------------
# 写真（それらしい現場写真をPillowで描く）
# ---------------------------------------------------------------------------
def _font(size: int):
    for f in ("C:/Windows/Fonts/meiryo.ttc", "C:/Windows/Fonts/msgothic.ttc", "C:/Windows/Fonts/YuGothM.ttc"):
        try:
            return ImageFont.truetype(f, size)
        except OSError:
            continue
    return ImageFont.load_default()


def photo_png(caption: str, stamp: datetime, seed_name: str) -> bytes:
    r = D.rng(seed_name)
    W, H = 640, 480
    img = PILImage.new("RGB", (W, H))
    d = ImageDraw.Draw(img)
    base = r.randint(70, 130)
    tint = (r.randint(-12, 8), r.randint(-8, 8), r.randint(-4, 18))
    for y in range(H):
        s = base + int(70 * y / H)
        d.line([(0, y), (W, y)], fill=tuple(max(0, min(255, s + t)) for t in tint))
    # 筐体・部品らしい矩形とボルト
    for _ in range(r.randint(3, 6)):
        x0, y0 = r.randint(-40, W - 120), r.randint(40, H - 120)
        w, h = r.randint(90, 320), r.randint(60, 220)
        col = r.choice([(182, 186, 192), (150, 155, 160), (92, 97, 104), (206, 200, 188), (62, 72, 84), (120, 140, 150)])
        d.rectangle([x0, y0, x0 + w, y0 + h], fill=col, outline=(35, 35, 35), width=2)
        for bx, by in ((x0 + 10, y0 + 10), (x0 + w - 18, y0 + 10), (x0 + 10, y0 + h - 18), (x0 + w - 18, y0 + h - 18)):
            d.ellipse([bx, by, bx + 8, by + 8], fill=(210, 210, 210), outline=(40, 40, 40))
    # 配管・ケーブル
    for _ in range(r.randint(2, 5)):
        pts = [(r.randint(0, W), r.randint(40, H)) for _ in range(r.randint(2, 4))]
        d.line(pts, fill=r.choice([(30, 30, 30), (200, 200, 60), (40, 90, 170), (220, 220, 220)]), width=r.randint(4, 12))
    # 点描で質感
    for _ in range(1800):
        v = r.randint(0, 255)
        d.point((r.randrange(W), r.randrange(H)), fill=(v, v, v))
    # 異常箇所の赤丸と矢印
    cx, cy = r.randint(160, W - 160), r.randint(140, H - 110)
    rx, ry = r.randint(40, 80), r.randint(30, 60)
    d.ellipse([cx - rx, cy - ry, cx + rx, cy + ry], outline=(230, 20, 20), width=6)
    ax, ay = cx + rx + 60, cy - ry - 50
    d.line([(ax, ay), (cx + rx * 0.7, cy - ry * 0.7)], fill=(230, 20, 20), width=5)
    # キャプション帯と撮影日時スタンプ
    d.rectangle([0, 0, W, 42], fill=(255, 255, 255))
    d.text((10, 6), caption, fill=(0, 0, 0), font=_font(24))
    d.text((W - 230, H - 36), stamp.strftime("%Y/%m/%d %H:%M"), fill=(255, 150, 0), font=_font(22))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


# ---------------------------------------------------------------------------
# 共通: 印刷設定・保存
# ---------------------------------------------------------------------------
def _print_setup(ws, last_col: int, last_row: int, footer_left: str, footer_right: str):
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins = PageMargins(left=0.5, right=0.4, top=0.6, bottom=0.6, header=0.3, footer=0.3)
    ws.print_options.horizontalCentered = True
    ws.print_area = f"A1:{get_column_letter(last_col)}{last_row}"
    ws.oddHeader.right.text = "社外秘"
    ws.oddHeader.right.size = 8
    ws.oddFooter.left.text = footer_left
    ws.oddFooter.left.size = 8
    ws.oddFooter.center.text = "&P / &N"
    ws.oddFooter.right.text = footer_right
    ws.oddFooter.right.size = 8
    ws.sheet_view.showGridLines = False


def _save_deterministic(wb: Workbook, path: Path, stamp: datetime):
    """xlsx を保存し、zip内タイムスタンプと文書プロパティの日時を固定して再パックする。"""
    wb.properties.created = stamp
    buf = io.BytesIO()
    wb.save(buf)
    iso = stamp.strftime("%Y-%m-%dT%H:%M:%SZ")
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                txt = data.decode("utf-8")
                txt = re.sub(r"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)", rf"\g<1>{iso}\g<2>", txt)
                txt = re.sub(r"(<dcterms:created[^>]*>)[^<]*(</dcterms:created>)", rf"\g<1>{iso}\g<2>", txt)
                data = txt.encode("utf-8")
            zi = zipfile.ZipInfo(info.filename, date_time=(stamp.year, stamp.month, stamp.day, stamp.hour, stamp.minute, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            z.writestr(zi, data)
    path.write_bytes(out.getvalue())


def _downtime_text(minutes: int, style: str):
    if style == "text_hm":
        h, m = divmod(minutes, 60)
        return f"{h}時間{m}分" if h else f"{m}分"
    if style == "text_min":
        return f"{minutes:,}分"
    return minutes


def _loss_value(yen: int, man: bool):
    """損失金額の表記と、人が読み取る値（概算表記なら丸めた金額）。"""
    if man:
        v = round(yen / 10000)
        return f"約{v:,}万円", v * 10000
    return yen, yen


def _lots_text(lots: list[str], sep: str) -> str:
    return sep.join(lots) if lots else "なし"


def _parts_lines(inc: D.Incident) -> list[str]:
    return [f"{p.part_no} {p.name} ×{q}" for p, q in inc.parts_used]


def _perm_value(p: PermAction, due_shown) -> dict:
    """恒久対策表の1行の正解値。期限は表示どおりの文字列（日付セルは表示形式を適用した文字列）、年月日は due_norm。"""
    if isinstance(due_shown, date):
        due_shown = _ymd(due_shown)          # 表示形式 yyyy/m/d
    return dict(content=p.content, owner=p.owner, due=due_shown, due_norm=_iso(p.due) if p.due else None,
                status=p.status, remarks=p.remarks)


def _target_value(t: Target) -> dict:
    """水平展開の対象設備1行の正解値。when は記入どおり、日付を指す場合は when_norm に年月日。"""
    return dict(equipment_id=t.equipment_id, checked=t.checked, when=t.when,
                when_norm=_iso(t.when_date) if t.when_date else None, result=t.result)


# ---------------------------------------------------------------------------
# 旧様式 Rev.3（方眼紙・右上押印欄）
# ---------------------------------------------------------------------------
OLD_FONT = "ＭＳ Ｐゴシック"
OLD_LABEL = _fill("F2F2F2")
OLD_SEC = _fill("D9D9D9")
NC = 36     # 方眼の列数


def render_old(x: Report) -> tuple[Workbook, dict]:
    inc, eq, st, L = x.inc, x.inc.equipment, x.style, LABELS[OLD]
    wb = Workbook()
    ws = wb.active
    ws.title = st["sheet"]
    g = Grid(ws, [2.75] * NC, OLD_FONT, 10, 13.5)
    vals: dict = {}
    labels: dict = {}
    mark = st["mark"]
    lab = lambda key: L[key].replace("時間", "\n時間") if st["label_nl"] and key in ("downtime_min",) else L[key]

    def dtext(d: date):
        return D.fmt_date(d, style=2) if st["date_text"] else d

    def put_date(r, c1, c2, d) -> str:
        """日付セル（表示形式 yyyy/m/d）か和暦の文字列を書き、画面に表示される文字列を返す。"""
        v = dtext(d)
        g.box(r, c1, r, c2, v, h="center", fmt=None if isinstance(v, str) else "yyyy/m/d")
        return v if isinstance(v, str) else _ymd(d)

    for rr in range(1, 17):
        g.height(rr, 18)
    g.height(1, 12)
    # 様式番号・タイトル・押印欄
    g.box(1, 1, 1, 14, L["form_no"], border=None, size=8)
    g.box(2, 1, 5, 22, L["banner"], border=None, size=18, bold=True, h="center")
    for c in range(3, 21):
        ws.cell(5, c).border = Border(bottom=DOUBLE)
    stamps = [("approver", x.approver, x.approved_on), ("checker", x.checker, x.checked_on), ("creator_stamp", x.writer, x.report_date)]
    for k, (key, person, day) in enumerate(stamps):
        c = 25 + k * 4
        g.box(2, c, 2, c + 3, L[key], fill=OLD_LABEL, h="center", size=9)
        g.box(3, c, 4, c + 3, D.surname(person) if day else None, h="center", size=12, bold=True, color="C00000")
        g.box(5, c, 5, c + 3, _md(day) if day else None, h="center", size=8)
    # 押印欄の日付は年なしの「m/d」の文字列。値は表示どおり、年を補った日付は *_norm に持つ
    if x.approved_on:
        vals["approver"] = D.surname(x.approver)
        vals.update(approved_on=_md(x.approved_on), approved_on_norm=_iso(x.approved_on))
    if x.checked_on:
        vals["checker"] = D.surname(x.checker)
        vals.update(checked_on=_md(x.checked_on), checked_on_norm=_iso(x.checked_on))
    labels.update(approver=L["approver"], checker=L["checker"])

    def lv(r, lc1, lc2, key, vc1, vc2, value, fmt=None, h="left", label_text=None):
        g.box(r, lc1, r, lc2, label_text or lab(key), fill=OLD_LABEL, h="center", size=9)
        g.box(r, vc1, r, vc2, value, fmt=fmt, h=h)
        labels[key] = label_text or lab(key)

    r = 7
    lv(r, 1, 5, "title", 6, NC, x.title)
    vals["title"] = x.title
    r += 1
    lv(r, 1, 5, "report_id", 6, 18, inc.incident_id)
    lv(r, 19, 23, "report_date", 24, NC, None)
    rd_disp = put_date(r, 24, NC, x.report_date)
    vals.update(report_id=inc.incident_id, report_date=rd_disp, report_date_norm=_iso(x.report_date))
    r += 1
    # 発生日時・復旧日時は 日付セル＋時刻セル に分けて書く様式（値は2つのセルを続けて読んだ文字列）
    for key, c0, dt in (("occurred_at", 1, inc.occurred_at), ("recovered_at", 19, x.recovered_at)):
        g.box(r, c0, r, c0 + 4, L[key], fill=OLD_LABEL, h="center", size=9)
        labels[key] = L[key]
        d_disp = put_date(r, c0 + 5, c0 + 12, dt.date())
        if key == "recovered_at" and inc.status == "保留":
            # 保留案件は暫定復旧の時刻を文字列で書く
            t_disp = f"{_hm(dt)}（暫定）"
            g.box(r, c0 + 13, r, c0 + 17, t_disp, h="center", shrink=True)
        else:
            t_disp = _hm(dt)
            g.box(r, c0 + 13, r, c0 + 17, dt.time(), h="center", fmt="h:mm")
        vals[key] = f"{d_disp} {t_disp}"
        vals[f"{key}_norm"] = _iso(dt)
    r += 1
    finder = inc.reporter.name
    writer_txt = f"{inc.assignees[0].section.replace('設備保全課 ', '')} {D.surname(x.writer)}" if st["writer_sect"] else x.writer.name
    lv(r, 1, 5, "finder", 6, 18, f"{finder}（{inc.detected_by}）")
    lv(r, 19, 23, "reporter", 24, NC, writer_txt)
    vals.update(finder=finder, detected_by=inc.detected_by, reporter=writer_txt)
    r += 1
    lv(r, 1, 5, "equipment_id", 6, 18, eq.equipment_id)
    lv(r, 19, 23, "equipment_name", 24, NC, eq.name)
    vals.update(equipment_id=eq.equipment_id, equipment_name=eq.name)
    r += 1
    lv(r, 1, 5, "line", 6, 10, eq.line, h="center")
    lv(r, 11, 13, "process", 14, 18, eq.process, label_text="工程")
    ws.cell(r, 14).font = Font(name=OLD_FONT, size=8)
    ws.cell(r, 14).alignment = Alignment(shrink_to_fit=True, vertical="center")
    names = "、".join(p.name for p in inc.assignees)
    lv(r, 19, 23, "assignee", 24, NC, names)
    g.fit(r, names, 24, NC, 18)
    vals.update(line=eq.line, process=eq.process, assignee=[p.name for p in inc.assignees])
    r += 1
    lv(r, 1, 5, "severity", 6, 18, _check_line(SEVERITIES[:3], {inc.severity}, mark))
    lv(r, 19, 23, "cause_category", 24, NC, inc.cause_category)
    vals.update(severity=inc.severity, cause_category=inc.cause_category)
    r += 1
    lv(r, 1, 5, "failure_category", 6, NC, _check_line(FAIL_CATS, {inc.category}, mark))
    vals["failure_category"] = inc.category
    r += 1
    # 影響
    g.box(r, 1, r + 1, 5, L["impact"], fill=OLD_LABEL, h="center", size=9)
    heads = [("downtime_min", 6, 12), ("lots", 13, 24), ("scrap_wafers", 25, 29), ("loss_yen", 30, NC)]
    for key, c1, c2 in heads:
        g.box(r, c1, r, c2, lab(key), fill=OLD_LABEL, h="center", size=9)
        labels[key] = lab(key)
    if st["label_nl"]:
        g.height(r, 28)
    dt_val = _downtime_text(inc.downtime_min, st["downtime"])
    g.box(r + 1, 6, r + 1, 12, dt_val, h="center", fmt='#,##0"分"' if isinstance(dt_val, int) else None)
    lots_txt = _lots_text(inc.lots_affected, "\n")
    g.box(r + 1, 13, r + 1, 24, lots_txt, h="center")
    g.box(r + 1, 25, r + 1, 29, inc.scrap_wafers, h="center", fmt='0"枚"')
    loss_disp, loss_val = _loss_value(inc.cost_yen, st["loss_man"])
    g.box(r + 1, 30, r + 1, NC, loss_disp, h="right", fmt='"¥"#,##0' if isinstance(loss_disp, int) else None)
    g.fit(r + 1, lots_txt, 13, 24, 20)
    vals.update(downtime_min=inc.downtime_min, lots=list(inc.lots_affected), scrap_wafers=inc.scrap_wafers, loss_yen=loss_val)
    r += 2
    rel = inc.related_incident_id or "―"
    lv(r, 1, 5, "related_report", 6, 18, rel, h="center")
    recur = "有" if inc.recurrence else "無"
    recur_disp = _check_line(("有", "無"), {recur}, mark)
    lv(r, 19, 23, "recurrence", 24, NC, recur_disp)
    g.height(r, 18)
    if inc.related_incident_id:
        vals["related_report"] = inc.related_incident_id
    # 再発有無: 値はチェックの付いた選択肢（有/無）、セルの文字列そのものは recurrence_display
    vals.update(recurrence=recur, recurrence_display=recur_disp)
    g.outline(7, 1, r, NC)
    r += 2

    marks = [1]            # 改ページ判定の区切り（この行から次の区切りまでを1ページ内に収めたい）

    def section(r, key):
        marks.append(r)
        g.box(r, 1, r, NC, L[key], fill=OLD_SEC, bold=True)
        g.height(r, 18)
        labels[key] = L[key]
        return r + 1

    # １．時系列
    r = section(r, "timeline")
    for c1, c2, t in ((1, 4, "日付"), (5, 7, "時刻"), (8, 31, "対応内容"), (32, NC, "担当")):
        g.box(r, c1, r, c2, t, fill=OLD_LABEL, h="center", size=9)
    r += 1
    last_d = None
    tl_vals = []
    n_rows = max(12, len(x.timeline))
    for i in range(n_rows):
        row = x.timeline[i] if i < len(x.timeline) else None
        dv = tv = None
        if row:
            if row.at.date() != last_d:
                dv = _md(row.at) if st["time_text"] or st["date_text"] else row.at.date()
                last_d = row.at.date()
            if st["time_text"] and i == 0:
                tv = f"{_hm(row.at)}頃"
            elif st["time_text"]:
                tv = _hm(row.at)
            else:
                tv = row.at.time()
            # 値は表示どおり: 日付は変わった行だけ「m/d」があり、それ以外の行は時刻（h:mm）だけ。年を補った日時は datetime_norm
            shown = _hm(row.at) if isinstance(tv, time) else tv
            if dv is not None:
                shown = f"{_md(row.at)} {shown}"
            tl_vals.append(dict(datetime=shown, datetime_norm=_iso(row.at), content=row.text, person=row.who))
        g.box(r, 1, r, 4, dv, h="center", fmt="m/d" if isinstance(dv, date) else None, border=THIN)
        g.box(r, 5, r, 7, tv, h="center", fmt="h:mm" if isinstance(tv, time) else None)
        g.box(r, 8, r, 31, row.text if row else None)
        g.box(r, 32, r, NC, row.who if row else None, h="center", shrink=True)
        g.fit(r, row.text if row else "", 8, 31, 18)
        r += 1
    vals["timeline"] = tl_vals
    r += 1

    # ２．現象
    r = section(r, "symptom")
    g.box(r, 1, r, NC, inc.symptom, v="top")
    g.fit(r, inc.symptom, 1, NC, 36)
    vals["symptom"] = inc.symptom
    r += 1
    alarm_txt = f"{inc.alarm.code}　{inc.alarm.message}" if inc.alarm else "なし"
    lv(r, 1, 6, "alarm", 7, NC, alarm_txt)
    g.height(r, 18)
    vals["alarm"] = alarm_txt if inc.alarm else ""
    r += 1
    lv(r, 1, 6, "investigation", 7, NC, inc.investigation)
    ws.cell(r, 7).alignment = Alignment(wrap_text=True, vertical="top")
    g.fit(r, inc.investigation, 7, NC, 36)
    vals["investigation"] = inc.investigation
    r += 2

    # ３．暫定対策
    r = section(r, "action")
    g.box(r, 1, r, NC, inc.action, v="top")
    g.fit(r, inc.action, 1, NC, 40)
    vals["action"] = inc.action
    r += 1
    parts = _parts_lines(inc)
    parts_txt = "\n".join(parts) if parts else "なし"
    lv(r, 1, 6, "parts", 7, NC, parts_txt)
    g.fit(r, parts_txt, 7, NC, 18)
    vals["parts"] = parts
    r += 1
    if inc.result:
        lv(r, 1, 6, "result", 7, NC, inc.result)
        g.fit(r, inc.result, 7, NC, 18)
        vals["result"] = inc.result
        r += 1
    r += 1

    # ４．なぜなぜ分析（階段状の箱を罫線の矢印でつなぐ）
    r = section(r, "why_why")
    g.height(r, 6)
    r += 1
    boxes = [(L["event"], x.event, 1)] + [(f"なぜ{i + 1}", inc.why_why[i] if i < len(inc.why_why) else None, 1 + 2 * (i + 1))
                                      for i in range(5)]
    for k, (name, text, s) in enumerate(boxes):
        if k > 0:
            stem = s + 1            # 次の箱のラベル（2列）の中央 = 右側列の左罫線
            g.height(r, 9)
            ws.cell(r, stem).border = Border(left=MED)
            r += 1
            g.height(r, 11)
            g.box(r, stem - 1, r, stem, "▼", border=None, size=8, h="center", v="bottom")
            r += 1
        g.box(r, s, r, s + 1, name, fill=OLD_LABEL if k else OLD_SEC, bold=True, size=8, h="center")
        g.box(r, s + 2, r, NC, text)
        g.outline(r, s, r, NC, MED)
        g.fit(r, text or "", s + 2, NC, 26)
        r += 1
    vals["why_why"] = list(inc.why_why)
    vals["event"] = x.event
    labels["event"] = L["event"]
    r += 1

    # ５．真因
    r = section(r, "cause")
    g.box(r, 1, r, NC, inc.cause)
    g.outline(r, 1, r, NC, DOUBLE)
    g.fit(r, inc.cause, 1, NC, 30)
    vals["cause"] = inc.cause
    r += 2

    # ６．恒久対策
    r = section(r, "prevention")
    cols = [(1, 2, "No"), (3, 21, "対策内容"), (22, 25, "担当"), (26, 29, "期限"), (30, 32, "状況"), (33, NC, "備考")]
    for c1, c2, t in cols:
        g.box(r, c1, r, c2, t, fill=OLD_LABEL, h="center", size=9)
    r += 1
    pv = []
    for i in range(max(3, len(x.perm))):
        p = x.perm[i] if i < len(x.perm) else None
        due = None
        if p:
            due = p.due_text or (p.due if p.due else "―")
            if isinstance(due, date) and st["date_text"]:
                due = D.fmt_date(due, style=2)
        g.box(r, 1, r, 2, i + 1 if p else None, h="center")
        g.box(r, 3, r, 21, p.content if p else None)
        g.box(r, 22, r, 25, p.owner if p else None, h="center", shrink=True)
        g.box(r, 26, r, 29, due, h="center", fmt="yyyy/m/d" if isinstance(due, date) else None, shrink=True)
        g.box(r, 30, r, 32, p.status if p else None, h="center", shrink=True)
        g.box(r, 33, r, NC, (p.remarks or None) if p else None, size=8)
        g.fit(r, p.content if p else "", 3, 21, 18)
        g.fit(r, p.remarks if p else "", 33, NC, 18)
        if p:
            pv.append(_perm_value(p, due))
        r += 1
    vals["prevention"] = pv
    r += 1

    # ７．水平展開
    r = section(r, "horizontal_section")
    g.box(r, 1, r, 5, L["horizontal_options"], fill=OLD_LABEL, h="center", size=9)
    labels["horizontal_options"] = L["horizontal_options"]
    g.box(r, 6, r, NC, _check_line(HZ_OPTIONS, set(x.hz_options), mark))
    g.height(r, 18)
    r += 1
    for c1, c2, t in ((1, 2, "確認"), (3, 8, L["horizontal_targets"]), (9, 19, "設備名"), (20, 25, "実施日/予定"), (26, NC, "結果・備考")):
        g.box(r, c1, r, c2, t, fill=OLD_LABEL, h="center", size=9)
    labels["horizontal_targets"] = L["horizontal_targets"]
    r += 1
    tv = []
    for i in range(max(3, len(x.targets))):
        t = x.targets[i] if i < len(x.targets) else None
        g.box(r, 1, r, 2, (mark if t.checked else "□") if t else None, h="center")
        g.box(r, 3, r, 8, t.equipment_id if t else None, h="center")
        g.box(r, 9, r, 19, t.name if t else None, shrink=True)
        g.box(r, 20, r, 25, t.when if t else None, h="center", shrink=True)
        g.box(r, 26, r, NC, t.result if t else None, shrink=True)
        g.height(r, 18)
        if t:
            tv.append(_target_value(t))
        r += 1
    g.box(r, 1, r, 5, L["horizontal_deployment"], fill=OLD_LABEL, h="center", size=9)
    labels["horizontal_deployment"] = L["horizontal_deployment"]
    g.box(r, 6, r, NC, inc.horizontal_deployment or None)
    g.fit(r, inc.horizontal_deployment, 6, NC, 18)
    vals["horizontal_deployment"] = inc.horizontal_deployment
    vals["horizontal_options"] = list(x.hz_options)
    vals["horizontal_targets"] = tv
    r += 2

    # ８．所感
    r = section(r, "impression")
    g.box(r, 1, r, NC, x.impression or None, v="top")
    g.fit(r, x.impression, 1, NC, 48)
    vals["impression"] = x.impression
    r += 1
    g.box(r, 1, r, NC, "配布先：製造課長／生産技術課／品質保証課／保全課内回覧", border=None, size=8)
    last = r
    _print_setup(ws, NC, last, f"管理No {inc.incident_id}", L["form_no"])
    g.paginate(marks, last)
    ws.sheet_view.zoomScale = 100

    if st["extra_sheets"]:
        _old_extra_sheets(wb, ws, st)
    if x.photos:
        _photo_sheet(wb, x)
    return wb, dict(values=vals, labels=labels)


def _old_extra_sheets(wb: Workbook, ws, st):
    """記入要領シートと、プルダウン用の非表示リストシート（旧様式の一部に存在）。"""
    guide = wb.create_sheet("記入要領")
    gg = Grid(guide, [4, 18, 70], OLD_FONT, 10, 13.5)
    gg.box(1, 1, 1, 3, "トラブル報告書 記入要領（保全課）", border=None, bold=True, size=14)
    rows = [("1", "対象", "重要度「中」以上、または停止2時間以上のトラブルは本様式で報告する。"),
            ("2", "提出期限", "復旧後5営業日以内に係長へ提出。未完了の場合も暫定報告を出すこと。"),
            ("3", "時系列", "発生から復帰までを時刻順に記入。日付は変わった行だけでよい。"),
            ("4", "なぜなぜ分析", "「なぜ」を3回以上繰り返し、管理・仕組みの問題まで掘り下げる。"),
            ("5", "恒久対策", "担当者と期限を必ず記入。期限は係長と調整。"),
            ("6", "水平展開", "同型機・類似設備の点検要否を判断し、対象設備をチェックする。")]
    for i, row in enumerate(rows, 3):
        for c, v in enumerate(row, 1):
            gg.box(i, c, i, c, v, h="center" if c == 1 else "left")
        gg.fit(i, row[2], 3, 3, 18)
    _print_setup(guide, 3, 2 + len(rows), "保全課", "記入要領")
    lst = wb.create_sheet("リスト")
    for i, v in enumerate(("重要度", *SEVERITIES), 1):
        lst.cell(i, 1, v)
    for i, v in enumerate(("状況", "完了", "実施中", "計画", "保留"), 1):
        lst.cell(i, 2, v)
    lst.sheet_state = "hidden"
    dv = DataValidation(type="list", formula1="=リスト!$B$2:$B$5", allow_blank=True)
    ws.add_data_validation(dv)
    for row in ws.iter_rows(min_col=30, max_col=30):     # 恒久対策表の「状況」列
        c = row[0]
        if isinstance(c.value, str) and c.value in ("完了", "実施中", "計画", "保留"):
            dv.add(c.coordinate)


def _photo_sheet(wb: Workbook, x: Report):
    ws = wb.create_sheet("写真")
    g = Grid(ws, [2] + [11] * 8, OLD_FONT, 10, 13.5)
    g.box(1, 2, 1, 9, f"添付写真　管理No {x.inc.incident_id}　{x.inc.equipment.equipment_id}", border=None, bold=True, size=12)
    n_blocks = (len(x.photos) + 1) // 2
    last = 2 + n_blocks * 15
    for rr in range(2, last + 1):
        g.height(rr, 15)          # 画像（225px ≒ 169pt）が12行分に収まるよう行高さを固定
    for k, cap in enumerate(x.photos):
        col = 2 if k % 2 == 0 else 6
        row = 3 + (k // 2) * 15
        stamp = x.inc.response_started_at + timedelta(minutes=20 + 35 * k)
        img = XLImage(io.BytesIO(photo_png(cap, stamp, f"f2-photo:{x.inc.incident_id}:{k}")))
        img.width, img.height = 300, 225
        ws.add_image(img, f"{get_column_letter(col)}{row}")
        g.box(row + 12, col, row + 12, col + 3, f"写真{k + 1}：{cap}", size=9)
        g.height(row + 12, 18)
    _print_setup(ws, 9, last, f"管理No {x.inc.incident_id}", "写真")


# ---------------------------------------------------------------------------
# 新様式 Rev.5（ラベル左・値右の表形式、承認欄は末尾）
# ---------------------------------------------------------------------------
NEW_FONT = "Meiryo UI"
NEW_LABEL = _fill("DDEBF7")
NEW_SEC = _fill("1F4E78")
NEW_SUB = _fill("F2F2F2")
LC, RC = 2, 11     # 使う列 B〜K


def render_new(x: Report) -> tuple[Workbook, dict]:
    inc, eq, st, L = x.inc, x.inc.equipment, x.style, LABELS[NEW]
    wb = Workbook()
    ws = wb.active
    ws.title = st["sheet"]
    g = Grid(ws, [1.5, 13] + [9.3] * 9 + [1.5], NEW_FONT, 10, 16.5)
    vals: dict = {}
    labels: dict = {}

    def lv(r, lc, key, vc1, vc2, value, fmt=None, h="left", label_text=None):
        text = label_text or L[key]
        narrow = lc != 2 and _dw(text) > 8          # 狭い列のラベルは縮小表示
        g.box(r, lc, r, lc, text, fill=NEW_LABEL, bold=True, size=8 if narrow else 9, h="center", shrink=narrow)
        g.box(r, vc1, r, vc2, value, fmt=fmt, h=h)
        labels[key] = text
        g.fit(r, value if isinstance(value, str) else "", vc1, vc2, 21)

    marks = [1]            # 改ページ判定の区切り

    def section(r, text):
        marks.append(r)
        g.box(r, LC, r, RC, text, fill=NEW_SEC, bold=True, color="FFFFFF", border=None)
        g.height(r, 21)
        return r + 1

    g.height(1, 15)
    g.box(1, 2, 1, 7, "製造部 設備保全課", border=None, size=9)
    g.box(1, 8, 1, RC, L["form_no"], border=None, size=8, h="right")
    g.box(2, LC, 2, RC, L["banner"], border=None, size=16, bold=True, h="center")
    g.height(2, 30)
    for c in range(LC, RC + 1):
        ws.cell(2, c).border = Border(bottom=MED)
    g.height(3, 6)
    r = 4
    lv(r, 2, "title", 3, RC, x.title)
    vals["title"] = x.title
    r += 1
    lv(r, 2, "report_id", 3, 6, inc.incident_id)
    lv(r, 7, "report_date", 8, RC, x.report_date, fmt="yyyy/mm/dd", h="center")
    # 値は表示形式を適用した文字列（Excel の表示どおり）、ISO に正規化した値は *_norm
    vals.update(report_id=inc.incident_id, report_date=_ymd(x.report_date, pad=True), report_date_norm=_iso(x.report_date))
    r += 1
    lv(r, 2, "occurred_at", 3, 6, inc.occurred_at, fmt="yyyy/mm/dd hh:mm", h="center")
    provisional = inc.status == "保留"
    lv(r, 7, "recovered_at", 8, RC, x.recovered_at, fmt="yyyy/mm/dd hh:mm" + (' "(暫定)"' if provisional else ""), h="center")
    vals.update(occurred_at=f"{_ymd(inc.occurred_at, pad=True)} {_hm(inc.occurred_at, pad=True)}", occurred_at_norm=_iso(inc.occurred_at),
                recovered_at=f"{_ymd(x.recovered_at, pad=True)} {_hm(x.recovered_at, pad=True)}" + ("(暫定)" if provisional else ""),
                recovered_at_norm=_iso(x.recovered_at))
    r += 1
    lv(r, 2, "finder", 3, 6, inc.reporter.name)
    lv(r, 7, "detected_by", 8, RC, inc.detected_by)
    vals.update(finder=inc.reporter.name, detected_by=inc.detected_by)
    r += 1
    lv(r, 2, "equipment_id", 3, 6, f"{eq.equipment_id}　{eq.name}")
    lv(r, 7, "line", 8, RC, f"{eq.line}／{eq.process}")
    vals.update(equipment_id=eq.equipment_id, equipment_name=eq.name, line=eq.line, process=eq.process)
    labels.update(equipment_name=L["equipment_name"], process=L["process"])
    r += 1
    lv(r, 2, "severity", 3, 4, inc.severity, h="center")
    lv(r, 5, "failure_category", 6, 7, inc.category, h="center")
    lv(r, 8, "status", 9, RC, inc.status, h="center")
    dv_sev = DataValidation(type="list", formula1='"重大,大,中,小"', allow_blank=True)
    dv_st = DataValidation(type="list", formula1='"完了,経過観察,保留,対応中"', allow_blank=True)
    ws.add_data_validation(dv_sev)
    ws.add_data_validation(dv_st)
    dv_sev.add(f"C{r}")
    dv_st.add(f"I{r}")
    vals.update(severity=inc.severity, failure_category=inc.category, status=inc.status)
    r += 1
    lv(r, 2, "cause_category", 3, 6, inc.cause_category)
    lv(r, 7, "reporter", 8, RC, x.writer.name)
    vals.update(cause_category=inc.cause_category, reporter=x.writer.name)
    r += 1
    names = "、".join(f"{p.name}（{p.section.replace('設備保全課 ', '')}）" if p.role == "FE" or not p.section.startswith("設備保全課")
                     else p.name for p in inc.assignees)
    lv(r, 2, "assignee", 3, RC, names)
    vals["assignee"] = [p.name for p in inc.assignees]
    r += 1
    g.box(r, 2, r + 1, 2, L["impact"], fill=NEW_LABEL, bold=True, size=9, h="center")
    for key, c1, c2 in (("downtime_min", 3, 4), ("lots", 5, 7), ("scrap_wafers", 8, 9), ("loss_yen", 10, RC)):
        g.box(r, c1, r, c2, L[key], fill=NEW_SUB, size=9, h="center")
        labels[key] = L[key]
    g.height(r, 19)
    lots_txt = _lots_text(inc.lots_affected, "、")
    g.box(r + 1, 3, r + 1, 4, inc.downtime_min, fmt='#,##0" 分"', h="center")
    g.box(r + 1, 5, r + 1, 7, lots_txt, h="center")
    g.box(r + 1, 8, r + 1, 9, inc.scrap_wafers, fmt='0" 枚"', h="center")
    g.box(r + 1, 10, r + 1, RC, inc.cost_yen, fmt='"¥"#,##0', h="right")
    g.fit(r + 1, lots_txt, 5, 7, 21)
    vals.update(downtime_min=inc.downtime_min, lots=list(inc.lots_affected), scrap_wafers=inc.scrap_wafers, loss_yen=inc.cost_yen)
    r += 2
    rel = f"{inc.related_incident_id}（{'再発' if inc.recurrence else '関連'}）" if inc.related_incident_id else "なし"
    lv(r, 2, "related_report", 3, 7, rel)
    if inc.related_incident_id:
        vals["related_report"] = inc.related_incident_id
    # 再発有無はリスト入力（有／無）のセル。値は表示どおりの 有/無（旧様式のチェック文字列とそろえて recurrence_display も持つ）
    recur = "有" if inc.recurrence else "無"
    lv(r, 8, "recurrence", 9, RC, recur, h="center")
    dv_rec = DataValidation(type="list", formula1='"有,無"', allow_blank=True)
    ws.add_data_validation(dv_rec)
    dv_rec.add(f"I{r}")
    vals.update(recurrence=recur, recurrence_display=recur)
    g.outline(4, LC, r, RC)
    r += 2

    # 経緯（時系列）
    r = section(r, L["timeline"])
    labels["timeline"] = L["timeline"]
    for c1, c2, t in ((2, 3, "日時"), (4, 10, "対応内容"), (11, 11, "対応者")):
        g.box(r, c1, r, c2, t, fill=NEW_SUB, size=9, h="center")
    r += 1
    tl_vals = []
    for row in x.timeline:
        v = f"{_md(row.at)} {_hm(row.at)}" if st["dt_text"] else row.at
        g.box(r, 2, r, 3, v, fmt="m/d hh:mm" if isinstance(v, datetime) else None, h="center")
        g.box(r, 4, r, 10, row.text)
        g.box(r, 11, r, 11, row.who, h="center", shrink=True)
        g.fit(r, row.text, 4, 10, 20)
        shown = v if isinstance(v, str) else f"{_md(row.at)} {_hm(row.at, pad=True)}"     # 表示形式 m/d hh:mm
        tl_vals.append(dict(datetime=shown, datetime_norm=_iso(row.at), content=row.text, person=row.who))
        r += 1
    vals["timeline"] = tl_vals
    r += 1

    # 現象・処置
    r = section(r, "■ 現象・処置")
    lv(r, 2, "symptom", 3, RC, inc.symptom)
    vals["symptom"] = inc.symptom
    r += 1
    alarm_txt = f"{inc.alarm.code}（{inc.alarm.message}）" if inc.alarm else "―"
    lv(r, 2, "alarm", 3, RC, alarm_txt)
    vals["alarm"] = alarm_txt if inc.alarm else ""
    r += 1
    lv(r, 2, "investigation", 3, RC, inc.investigation)
    vals["investigation"] = inc.investigation
    r += 1
    lv(r, 2, "action", 3, RC, inc.action)
    ws.cell(r, 3).alignment = Alignment(wrap_text=True, vertical="top")
    vals["action"] = inc.action
    r += 1
    marks.append(r)        # 使用部品表と処置結果はまとめて同じページに
    n_parts = max(1, len(inc.parts_used))
    g.box(r, 2, r + n_parts, 2, L["parts"], fill=NEW_LABEL, bold=True, size=9, h="center")
    labels["parts"] = L["parts"]
    for c1, c2, t in ((3, 4, "品番"), (5, 8, "品名"), (9, 9, "数量"), (10, RC, "金額")):
        g.box(r, c1, r, c2, t, fill=NEW_SUB, size=9, h="center")
    r += 1
    if inc.parts_used:
        for p, q in inc.parts_used:
            g.box(r, 3, r, 4, p.part_no, h="center")
            g.box(r, 5, r, 8, p.name, shrink=True)
            g.box(r, 9, r, 9, q, h="center")
            g.box(r, 10, r, RC, p.unit_price_yen * q, fmt='"¥"#,##0', h="right")
            g.height(r, 20)
            r += 1
    else:
        g.box(r, 3, r, RC, "使用部品なし", h="center")
        g.height(r, 20)
        r += 1
    vals["parts"] = _parts_lines(inc)
    if inc.result:
        lv(r, 2, "result", 3, RC, inc.result)
        vals["result"] = inc.result
        r += 1
    r += 1

    # なぜなぜ分析（階段状、矢印はセル文字）
    r = section(r, L["why_why"])
    labels["why_why"] = L["why_why"]
    g.height(r, 6)
    r += 1
    boxes = [(L["event"], x.event, 2)] + [(f"なぜ{CIRCLED[i]}", inc.why_why[i] if i < len(inc.why_why) else None, 3 + i) for i in range(5)]
    for k, (name, text, s) in enumerate(boxes):
        if k > 0:
            g.height(r, 15)
            g.box(r, s, r, s, "↓", border=None, bold=True, size=11, h="center", color="1F4E78")
            r += 1
        g.box(r, s, r, s, name, fill=NEW_LABEL if k else NEW_SUB, bold=True, size=9, h="center")
        g.box(r, s + 1, r, RC, text)
        g.outline(r, s, r, RC, MED if k else THIN)
        g.fit(r, text or "", s + 1, RC, 30)
        r += 1
    vals["why_why"] = list(inc.why_why)
    vals["event"] = x.event
    labels["event"] = L["event"]
    r += 1

    # 根本原因・再発防止
    r = section(r, "■ 根本原因・再発防止")
    lv(r, 2, "cause", 3, RC, inc.cause)
    g.outline(r, 2, r, RC, MED)
    vals["cause"] = inc.cause
    r += 1
    n_perm = max(2, len(x.perm))
    g.box(r, 2, r + n_perm, 2, L["prevention"], fill=NEW_LABEL, bold=True, size=9, h="center")
    labels["prevention"] = L["prevention"]
    for c1, c2, t in ((3, 3, "No"), (4, 7, "対策内容"), (8, 8, "担当"), (9, 9, "期限"), (10, 10, "状況"), (11, 11, "備考")):
        g.box(r, c1, r, c2, t, fill=NEW_SUB, size=9, h="center")
    r += 1
    pv = []
    dv_p = DataValidation(type="list", formula1='"完了,実施中,計画,保留,―"', allow_blank=True)
    ws.add_data_validation(dv_p)
    for i in range(n_perm):
        p = x.perm[i] if i < len(x.perm) else None
        due = (p.due_text or p.due or "―") if p else None
        g.box(r, 3, r, 3, i + 1 if p else None, h="center")
        g.box(r, 4, r, 7, p.content if p else None)
        g.box(r, 8, r, 8, p.owner if p else None, h="center", shrink=True)
        g.box(r, 9, r, 9, due, fmt="yyyy/m/d" if isinstance(due, date) else None, h="center", shrink=True)
        g.box(r, 10, r, 10, p.status if p else None, h="center")
        g.box(r, 11, r, 11, (p.remarks or None) if p else None, size=8)
        dv_p.add(f"J{r}")
        g.fit(r, p.content if p else "", 4, 7, 21)
        g.fit(r, p.remarks if p else "", 11, 11, 21)
        if p:
            pv.append(_perm_value(p, due))
        r += 1
    vals["prevention"] = pv
    r += 1

    # 水平展開
    r = section(r, L["horizontal_section"])
    labels["horizontal_section"] = L["horizontal_section"]
    g.box(r, 2, r, 2, L["horizontal_options"], fill=NEW_LABEL, bold=True, size=9, h="center")
    labels["horizontal_options"] = L["horizontal_options"]
    g.box(r, 3, r, RC, _check_line(HZ_OPTIONS, set(x.hz_options)))
    g.height(r, 21)
    r += 1
    for c1, c2, t in ((2, 2, "確認"), (3, 3, L["horizontal_targets"]), (4, 6, "設備名"), (7, 8, "実施(予定)日"), (9, RC, "結果")):
        g.box(r, c1, r, c2, t, fill=NEW_SUB, size=9, h="center")
    labels["horizontal_targets"] = L["horizontal_targets"]
    r += 1
    tv = []
    for i in range(max(2, len(x.targets))):
        t = x.targets[i] if i < len(x.targets) else None
        g.box(r, 2, r, 2, ("☑" if t.checked else "□") if t else None, h="center", size=11)
        g.box(r, 3, r, 3, t.equipment_id if t else None, h="center")
        g.box(r, 4, r, 6, t.name if t else None, shrink=True)
        g.box(r, 7, r, 8, t.when if t else None, h="center")
        g.box(r, 9, r, RC, t.result if t else None)
        g.height(r, 20)
        if t:
            tv.append(_target_value(t))
        r += 1
    lv(r, 2, "horizontal_deployment", 3, RC, inc.horizontal_deployment or None)
    vals.update(horizontal_deployment=inc.horizontal_deployment, horizontal_options=list(x.hz_options), horizontal_targets=tv)
    r += 2

    # 所感・コメント
    r = section(r, "■ 所感")
    lv(r, 2, "impression", 3, RC, x.impression or None)
    g.fit(r, x.impression, 3, RC, 48)
    vals["impression"] = x.impression
    r += 1
    lv(r, 2, "supervisor_comment", 3, RC, x.supervisor_comment or None)
    if x.supervisor_comment:
        vals["supervisor_comment"] = x.supervisor_comment
    r += 2

    # 承認欄
    r = section(r, "■ 承認欄")
    g.box(r, 2, r, 2, "区分", fill=NEW_SUB, size=9, h="center")
    heads = [("creator_stamp", 3, 5), ("checker", 6, 8), ("approver", 9, RC)]
    for key, c1, c2 in heads:
        g.box(r, c1, r, c2, L[key], fill=NEW_LABEL, bold=True, size=9, h="center")
    labels.update(checker=L["checker"], approver=L["approver"])
    r += 1
    people_row = [(x.writer, x.report_date), (x.checker, x.checked_on), (x.approver, x.approved_on)]
    g.box(r, 2, r, 2, "氏名", fill=NEW_SUB, size=9, h="center")
    g.box(r + 1, 2, r + 1, 2, "日付", fill=NEW_SUB, size=9, h="center")
    g.box(r + 2, 2, r + 2, 2, "状態", fill=NEW_SUB, size=9, h="center")
    for (key, c1, c2), (person, day) in zip(heads, people_row):
        g.box(r, c1, r, c2, person.name if day else None, h="center")
        g.box(r + 1, c1, r + 1, c2, day, fmt="yyyy/mm/dd", h="center")
        g.box(r + 2, c1, r + 2, c2, ("提出済" if key == "creator_stamp" else "承認済" if key == "approver" else "確認済") if day else "未",
              h="center", color=None if day else "C00000")
    for k in range(3):
        g.height(r + k, 21)
    g.outline(r - 1, 2, r + 2, RC)
    if x.approved_on:
        vals.update(approver=x.approver.name, approved_on=_ymd(x.approved_on, pad=True), approved_on_norm=_iso(x.approved_on))
    if x.checked_on:
        vals.update(checker=x.checker.name, checked_on=_ymd(x.checked_on, pad=True), checked_on_norm=_iso(x.checked_on))
    r += 3

    # 添付写真（本シート末尾）
    if x.photos:
        r += 1
        r = section(r, "■ 添付写真")
        # 写真2枚で1段: 画像(280x210px ≒ 157.5pt)を10行×16.5pt に置き、直下の行にキャプション、1行空けて次の段
        for k, cap in enumerate(x.photos):
            col = 2 if k % 2 == 0 else 7
            if k % 2 == 0:
                if k:
                    r += 12
                    marks.append(r)
                for rr in range(r, r + 10):
                    g.height(rr, 16.5)
                g.height(r + 10, 18)
                g.height(r + 11, 8)
            stamp = inc.response_started_at + timedelta(minutes=15 + 40 * k)
            img = XLImage(io.BytesIO(photo_png(cap, stamp, f"f2-photo:{inc.incident_id}:{k}")))
            img.width, img.height = 280, 210
            ws.add_image(img, f"{get_column_letter(col)}{r}")
            g.box(r + 10, col, r + 10, col + 4, f"写真{k + 1}　{cap}", size=9, border=None)
        r += 11
    last = r
    _print_setup(ws, 12, last, f"報告書No {inc.incident_id}", L["form_no"])
    g.paginate(marks, last, margins_lr=0.9)
    ws.sheet_view.zoomScale = 90
    return wb, dict(values=vals, labels=labels)


# ---------------------------------------------------------------------------
# README
# ---------------------------------------------------------------------------
def _readme(reports: list[Report]) -> str:
    n_old = sum(1 for x in reports if x.version == OLD)
    n_new = len(reports) - n_old
    lines = [
        "# F2 トラブル対応報告書（初動〜恒久対策）サンプル",
        "",
        "`python -m scripts.samples.f2_trouble_report` で再生成できます（固定シードのため毎回同一のファイル）。",
        "元データは `scripts/samples/domain.py` の正準トラブル履歴で、なぜなぜ分析があり重要度が「中」以上の案件から30件を選んでいます。",
        "",
        "## フォルダの構成",
        "",
        "xlsx は**様式の版ごとにフォルダを分けて**います。1つのフォルダの中は同じ様式だけなので、"
        "フォルダごと帳票取り込みにドロップすれば、帳票の種類とシートを1回決めるだけで全ファイルをまとめて読み取れます。",
        "",
        f"- `{VERSION_DIR[OLD]}/` … 旧様式 Rev.3 の {n_old} ファイル"
        "（うち2件は2025/04以降の報告で旧様式テンプレートを使い回したもの）",
        f"- `{VERSION_DIR[NEW]}/` … 新様式 Rev.5 の {n_new} ファイル",
        "",
        "`_README.md`（このファイル）と `_expected.jsonl` は版フォルダの外、帳票フォルダの直下に置いています。",
        "",
        "## 帳票の構成",
        "",
        "1. ヘッダ: 件名／管理No／作成日／発生日時／復旧日時／発見者／作成者／設備／対応者／重要度／故障区分／原因区分",
        "2. 影響: 停止時間・影響ロット・廃棄枚数・損失金額（関連No・再発有無）",
        "3. 時系列: 日付・時刻・対応内容・担当（5〜15行。発生→初動→連絡→到着→調査→処置→復帰→事後対応）",
        "4. 現象（発生アラーム・調査結果）",
        "5. 暫定対策（使用部品・復旧確認）",
        "6. なぜなぜ分析: 事象→なぜ1〜なぜ5 を右下がりの階段状の箱で配置し、箱の間を矢印でつなぐ",
        "7. 真因",
        "8. 恒久対策: No／対策内容／担当／期限／状況／備考 の表",
        "9. 水平展開: 展開区分のチェック（☑/□）と対象設備チェックリスト、展開内容",
        "10. 所感（新様式は上長コメントも）",
        "11. 承認（旧様式は右上の押印欄、新様式は末尾の承認欄）",
        "",
        f"## 様式の版（旧様式 {n_old} 件／新様式 {n_new} 件）",
        "",
        "| 項目 | 旧様式 Rev.3（〜2025/03） | 新様式 Rev.5（2025/04〜） |",
        "|---|---|---|",
        f"| フォルダ | `{VERSION_DIR[OLD]}/` | `{VERSION_DIR[NEW]}/` |",
        "| レイアウト | 36列の方眼紙。見出し行の下に本文 | B〜K列の表形式。ラベル左・値右 |",
        "| フォント | ＭＳ Ｐゴシック 10pt | Meiryo UI 10pt |",
    ]
    lo, ln = LABELS[OLD], LABELS[NEW]
    for key, name in [("banner", "タイトル"), ("title", "件名"), ("report_id", "管理番号"), ("report_date", "作成日"), ("finder", "発見者"),
                      ("reporter", "作成者"), ("equipment_id", "設備"), ("downtime_min", "停止時間"), ("scrap_wafers", "廃棄枚数"),
                      ("loss_yen", "損失金額"), ("timeline", "時系列"), ("symptom", "現象"), ("action", "暫定処置"),
                      ("why_why", "なぜなぜ"), ("cause", "真因"), ("prevention", "恒久対策"), ("horizontal_section", "水平展開（見出し）"), ("horizontal_options", "水平展開（区分）"),
                      ("horizontal_targets", "水平展開（対象設備の列）"), ("horizontal_deployment", "水平展開（展開内容）"),
                      ("impression", "所感")]:
        lines.append(f"| {name} | {lo[key]} | {ln[key]} |")
    lines += [
        "| 発生日時 | 日付セル＋時刻セルに分割 | 1セルの日時 |",
        "| 設備 | 設備No・設備名が別セル | 「設備No　設備名」を1セル |",
        "| 重要度・故障区分 | ☑/□ のチェック文字列 | 値セル（重要度・対応状況・恒久対策の状況は入力規則のリスト付き） |",
        f"| 再発 | 「{lo['recurrence']}」欄に ☑有　□無 | 「{ln['recurrence']}」欄に 有／無（入力規則のリスト） |",
        "| なぜなぜの矢印 | 罫線（太線）の縦線＋「▼」を結合セルに配置 | 次の箱のラベル列に「↓」 |",
        "| 承認 | 右上の押印欄（姓を赤字、日付 m/d） | 末尾の承認欄（氏名・日付・状態） |",
        "| 写真 | 別シート「写真」 | 報告書シート末尾の「添付写真」 |",
        "",
        "## 意図的なゆらぎ・不備",
        "",
        "- 2025/04 以降の報告なのに旧様式テンプレートを使い回したファイルが2件ある。",
        "- 旧様式の一部では、日付が和暦文字列（R6.2.29）、停止時間が「4時間36分」「276分」の文字列、損失金額が「約26万円」の概算になっている。",
        "- 旧様式の一部では、チェック記号が ■、時系列の時刻が「3:12頃」の文字列、ラベルにセル内改行（停止\\n時間）がある。",
        "- 旧様式の一部に「記入要領」シートと非表示の「リスト」シート（入力規則の参照先）がある。",
        "- 時系列の日付は、日付が変わった行だけに書いてある（旧様式）。空の予備行が残っている。",
        "- なぜなぜは3〜5段で、使わない枠は空欄のまま。",
        "- 恒久対策の期限が「6月末」のような文字列の場合がある。「特になし」の行は担当・期限が「―」になる。",
        "- 水平展開の展開内容が未記入の報告があり、その場合はチェックリストも空になっている。",
        "- 保留案件では、恒久対策の表が空欄のもの、「恒久対策は原因確定後に検討（期限 未定）」とだけ書いたものがある。",
        "- 対応状況が保留・経過観察の報告は承認欄が空欄の場合がある。完了なのに承認印が漏れている旧様式が1件ある。",
        "- 新様式の所感欄が未記入のファイルや、時系列の日時が文字列になっているファイルがある。",
        "- 本文には記入者ごとの書き癖（全角数字、半角カナ、番号の振り方、敬体、誤変換）がそのまま残っている。",
        "- ファイル名の付け方が4通りある（管理No始まり、設備No＋日付、【重要度】付き、F2_管理No）。",
        "",
        "## _expected.jsonl",
        "",
        "1行が1ファイルで、`file` / `layout_version` / `values` / `labels_used` / `irregularities` を持つ。",
        "",
        "- `file`: この帳票フォルダからの相対パス（`版フォルダ名/ファイル名.xlsx`、区切りは `/`）。並びは下の「ファイル一覧」と同じ。",
        "- `values`: 人が帳票を読んで取り出す値。文章・日付は **Excel に表示されている文字列そのまま**（日付セルは表示形式を適用した文字列、"
        "例 旧様式 `2023/8/24 16:06`・`R5.8.24 16:06`・押印欄 `8/29`、新様式 `2025/04/28 16:06`）。",
        "- 日付・日時には、年を補って ISO 形式（`YYYY-MM-DD` / `YYYY-MM-DD HH:MM`）にした値を `*_norm` キーで併記する"
        "（`report_date_norm` `occurred_at_norm` `recovered_at_norm` `checked_on_norm` `approved_on_norm`、表の中は `datetime_norm` `due_norm` `when_norm`）。"
        "F4 の `judge` / `judge_norm` と同じ考え方。",
        "- 時系列の `datetime` は行に表示されている文字列（旧様式は日付が変わった行だけ `m/d h:mm`、他の行は `h:mm` のみ。`3:12頃` のような文字列もそのまま）。",
        "- 恒久対策の `due` は表示どおり（`2025/5/31` / `R6.5.31` / `5月末` / `未定` / `―`）。「○月末」は `due_norm` に月末日を入れる。",
        "- 停止時間（分）・枚数・金額は整数で、「約26万円」は 260000 とする（数値の読み取り値）。",
        "- `recurrence` は再発有無の選択値（`有` / `無`）。旧様式はチェック文字列（`☑有　□無`）、新様式は「再発有無」のリスト入力セルで、"
        "セルの文字列は `recurrence_display` に持つ。",
        "- 恒久対策の文に付いていた「→保全カレンダーへ登録済み」「→点検チェックシートに反映」のような状況の断片は、独立した対策行にせず、"
        "直前の対策行の `status`（計画／完了）と `remarks`（備考列: 保全カレンダー登録済／チェックシート反映済）に反映している。",
        "- リスト値: `lots`, `assignee`, `parts`（\"品番 品名 ×数量\"）, `why_why`, `timeline`（datetime/datetime_norm/content/person）, "
        "`prevention`（content/owner/due/due_norm/status/remarks）, `horizontal_targets`（equipment_id/checked/when/when_norm/result）。",
        "- 水平展開の実施日（☑の行）は、トラブルの復旧日から作成日＋14日（確認・承認の押印済みなら最後の押印日）までの範囲。",
        "- `labels_used`: その項目に対応するファイル中のラベル文字列（見出しの番号・記号も含めた完全一致）。",
        "- 承認されていない（空欄の）場合は `approver` / `approved_on` を出力しない。",
        "",
        "## ファイル一覧",
        "",
        "「ファイル」は帳票フォルダからの相対パス（`版フォルダ名/ファイル名.xlsx`）で、`_expected.jsonl` の `file` と同じです。",
        "",
        "| ファイル | 様式 | 設備 | 重要度 | 状態 | 写真 |",
        "|---|---|---|---|---|---|",
    ]
    for x in reports:
        lines.append(f"| {VERSION_DIR[x.version]}/{x.file_name} | {x.version} | {x.inc.equipment.equipment_id} | "
                     f"{x.inc.severity} | {x.inc.status} | {len(x.photos) if x.photos else '-'} |")
    return "\n".join(lines) + "\n"


# ---------------------------------------------------------------------------
# エントリポイント
# ---------------------------------------------------------------------------
def generate(output_root: Path | None = None) -> list[Path]:
    root = Path(output_root) if output_root else D.OUTPUT_ROOT
    out_dir = root / "forms" / FOLDER
    out_dir.mkdir(parents=True, exist_ok=True)
    # 名前を変えたファイルや、作らなくなった版のフォルダが残らないように掃除する（_README.md と _expected.jsonl は残す）
    keep_dirs = set(VERSION_DIR.values())
    for old in out_dir.glob("*.xlsx"):          # 版フォルダに分ける前の出力の残り
        old.unlink()
    for sub in out_dir.iterdir():
        if not sub.is_dir():
            continue
        if sub.name in keep_dirs:
            for old in sub.glob("*.xlsx"):
                old.unlink()
        else:
            shutil.rmtree(sub)
    reports = build_reports()
    written: list[Path] = []
    expected = []
    for x in reports:
        wb, meta = (render_old if x.version == OLD else render_new)(x)
        wb.properties.creator = x.writer.name
        wb.properties.lastModifiedBy = x.checker.name if x.checked_on else x.writer.name
        wb.properties.title = x.title
        rel = f"{VERSION_DIR[x.version]}/{x.file_name}"      # _expected.jsonl はフォルダからの相対パスで持つ
        path = out_dir / VERSION_DIR[x.version] / x.file_name
        path.parent.mkdir(exist_ok=True)
        _save_deterministic(wb, path, datetime.combine(x.report_date, time(17, 30)))
        written.append(path)
        expected.append(dict(file=rel, layout_version=x.version, values=meta["values"],
                             labels_used={k: v for k, v in meta["labels"].items() if k in meta["values"]},
                             irregularities=x.irregularities))
    exp_path = out_dir / "_expected.jsonl"
    exp_path.write_text("".join(json.dumps(e, ensure_ascii=False) + "\n" for e in expected), encoding="utf-8", newline="\n")
    readme = out_dir / "_README.md"
    readme.write_text(_readme(reports), encoding="utf-8", newline="\n")
    return written + [exp_path, readme]


if __name__ == "__main__":
    t0 = _time.time()
    paths = generate()
    for p in paths:
        print(p)
    print(f"{len(paths)} files ({_time.time() - t0:.1f}s)")
