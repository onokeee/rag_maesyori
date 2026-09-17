"""T5 不具合連絡・是正処置管理台帳（CSV / Excel）を作る。

    python -m scripts.samples.t5_capa_ledger_csv

出力（samples/tables/）
- T5_是正処置管理台帳.csv            : UTF-8(BOM付き)・LF改行、約5,000行。長文セルに改行・カンマ・ダブルクォートを含む
- T5_是正処置管理台帳_2026上期.xlsx  : 2026年度上期（2026/4/1〜9/30 起票）分。大項目（不具合/対策/確認）を結合した2段ヘッダー
- T5_是正処置管理台帳_README.md       : 構成・文字コード・件数・意図的なイレギュラーの説明

台帳の行は domain.standard_incidents() のうち重大・大・中のトラブルから起票したもの（関連トラブルNo・設備ID・人名が
トラブル報告書と一致）と、品質保証課が起票した後工程・顧客からの品質クレームで構成する。
効果確認・判定は、対策完了後の同一設備・同一故障モードの再発やチョコ停件数を実データから数えて決める。
"""
from __future__ import annotations

import csv
import io
import re
import time as _time
import zipfile
from bisect import bisect_left, bisect_right
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from functools import lru_cache
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .domain import (
    OUTPUT_ROOT,
    TRANSPORT_KINDS,
    UTILITY_KINDS,
    Incident,
    Person,
    equipment_master,
    fab_of,
    kind_of,
    minor_stops,
    people,
    rng,
    standard_incidents,
    surname,
    to_zenkaku,
)

# ---------------------------------------------------------------------------
# 基本設定
# ---------------------------------------------------------------------------
STEM = "T5_是正処置管理台帳"
SNAPSHOT = date(2026, 9, 10)                      # 台帳を抽出した日（これ以降の完了日・判定は存在しない）
H1_START, H1_END = date(2026, 4, 1), date(2026, 9, 30)   # 2026年度上期（4月始まりの年度）
N_LINKED_CLAIMS = 330                             # トラブルと紐付く品質クレーム
N_FREE_CLAIMS = 120                               # トラブルと紐付かない品質クレーム
N_DUPLICATES = 12                                 # 同一案件の重複起票

COLUMNS = ["台帳No", "起票日", "起票部署", "起票者", "対象設備", "関連トラブルNo", "不具合内容", "暫定対策",
           "真因(なぜなぜ要約)", "恒久対策", "水平展開", "期限", "完了日", "効果確認", "判定", "承認者"]
# Excel の大項目: (見出し, 開始列index, 終了列index(含まない))
GROUPS = [("不具合", 0, 7), ("対策", 7, 12), ("確認", 12, 16)]
DATE_COLS = ("起票日", "期限", "完了日")
LONG_COLS = ("不具合内容", "暫定対策", "真因(なぜなぜ要約)", "恒久対策", "水平展開", "効果確認")
VERDICTS = ("有効", "再検討", "未確認")

_ISO_RE = re.compile(r"\d{4}-\d{2}-\d{2}")


@dataclass
class _Row:
    """台帳の1行。values は出力列そのもの、それ以外は並べ替え・採番・集計用の内部情報。"""
    key: str
    filed: date                     # 起票日
    sort_dt: datetime               # 同日内の並び順
    source: str                     # incident / claim / claim_free / duplicate
    incident_id: str = ""
    deadline: date | None = None    # 期限（集計用）
    done: date | None = None        # 完了日（集計用）
    values: dict = field(default_factory=dict)


# ---------------------------------------------------------------------------
# 日付表記（2023年起票は和暦 R5.x.x、2024年以降は ISO。移行期の書き癖を少し残す）
# ---------------------------------------------------------------------------
def _fmt_d(d: date | None, style: str) -> str:
    if d is None:
        return ""
    y = d.year - 2018
    if style == "R":
        return f"R{y}.{d.month}.{d.day}"
    if style == "R0":
        return f"R{y:02d}.{d.month:02d}.{d.day:02d}"
    if style == "RK":
        return f"令和{y}年{d.month}月{d.day}日"
    if style == "SL":
        return f"{d.year}/{d.month}/{d.day}"
    return d.isoformat()


def _row_date_style(filed: date, r) -> str:
    if filed.year <= 2023:
        return r.choices(["R", "R0", "RK"], weights=[88, 9, 3])[0]
    if filed.year == 2024 and filed.month <= 3 and r.random() < 0.18:
        return "R"          # 年が変わっても和暦で書き続けた人
    return "ISO"


def _date_writer(filed: date, r):
    """行ごとの日付表記関数を返す。ISO の行でもまれにスラッシュ表記が混じる。"""
    style = _row_date_style(filed, r)

    def dfmt(d: date | None, allow_noise: bool = True) -> str:
        if style == "ISO" and allow_noise and d is not None and r.random() < 0.015:
            return _fmt_d(d, "SL")
        return _fmt_d(d, style)

    return dfmt


# ---------------------------------------------------------------------------
# 人・書き癖
# ---------------------------------------------------------------------------
@lru_cache(maxsize=None)
def _ppl() -> dict:
    ppl = people()
    return dict(
        all=ppl,
        kacho=next(p for p in ppl if p.role == "課長"),
        chiefs={p.section: p for p in ppl if p.role == "係長"},
        pe=[p for p in ppl if p.section == "生産技術課"],
        pe_lead=next(p for p in ppl if p.section == "生産技術課" and p.role == "主任技師"),
        qa=[p for p in ppl if p.section == "品質保証課"],
    )


def _maint_section(eq) -> str:
    if kind_of(eq) in UTILITY_KINDS:
        return "設備保全課 施設係"
    return "設備保全課 保全1係" if fab_of(eq) == "Fab1" else "設備保全課 保全2係"


def _name_variant(p: Person, r) -> str:
    """氏名の書き方ゆれ（姓名 / 姓のみ / 空白なし / 全角空白 / 役職付き）。"""
    return r.choices([p.name, surname(p), p.name.replace(" ", ""), p.name.replace(" ", "　"), f"{p.name}（{p.role}）"],
                     weights=[66, 14, 8, 6, 6])[0]


def _dept_variant(p: Person, r) -> str:
    sec = p.section
    if sec.startswith("設備保全課 "):
        sub = sec.split(" ")[1]
        return r.choices([sec, "製造部 設備保全課", sub, "設備保全課", sec.replace(" ", "")], weights=[64, 14, 12, 6, 4])[0]
    if sec == "生産技術課":
        return r.choices(["生産技術課", "製造部 生産技術課", "生技"], weights=[75, 20, 5])[0]
    if sec == "品質保証課":
        return r.choices(["品質保証課", "品質保証部 品質保証課", "品証"], weights=[70, 22, 8])[0]
    return sec


@lru_cache(maxsize=None)
def _style(employee_id: str) -> dict:
    """起票者ごとの台帳の書き方（見出し記号・箇条書き記号・詳しさ）。"""
    r = rng(f"t5-style:{employee_id}")
    return dict(
        hdr=r.choice(["bracket", "bracket", "square", "colon", "prose"]),
        bullet=r.choice(["1. ", "1. ", "①", "・", "(1)"]),
        verbose=r.random() < 0.5,
        join=r.choice(["、", " / ", "，"]),
    )


def _numbered(items: list[str], bullet: str) -> str:
    out = []
    for i, s in enumerate(items, 1):
        if bullet == "①":
            out.append(f"{chr(0x2460 + i - 1)}{s}")
        elif bullet == "・":
            out.append(f"・{s}")
        elif bullet == "(1)":
            out.append(f"({i}) {s}")
        else:
            out.append(f"{i}. {s}")
    return "\n".join(out)


def _render_blocks(blocks: list[tuple[str, str]], st: dict) -> str:
    """(見出し, 本文) を起票者の書式で並べる。本文が複数行なら見出しの次の行から書く。"""
    out = []
    for label, body in blocks:
        body = (body or "").strip()
        if not body:
            continue
        multi = "\n" in body
        if st["hdr"] == "bracket":
            out.append(f"【{label}】" + ("\n" if multi else "") + body)
        elif st["hdr"] == "square":
            out.append(f"■{label}" + ("\n" if multi else "：") + body)
        elif st["hdr"] == "colon":
            out.append(f"{label}：" + ("\n" if multi else "") + body)
        else:
            out.append(body)
    return "\n".join(out)


def _sentences(text: str, k: int) -> str:
    parts = [s for s in text.split("。") if s.strip()]
    if len(parts) <= k:
        return text
    return "。".join(parts[:k]) + "。"


def _short(text: str, n: int = 38) -> str:
    """先頭の文を n 文字程度に縮める（括弧の途中で切らない）。"""
    s = re.split(r"[。\n]", text)[0]
    if len(s) <= n:
        return s
    cut = max(s.rfind("（", 0, n), s.rfind("、", 0, n))
    return (s[:cut] if cut >= 12 else s[:n]) + "…"


# ---------------------------------------------------------------------------
# 再発・チョコ停の索引（効果確認に使う）
# ---------------------------------------------------------------------------
class _Index:
    def __init__(self, incidents: list[Incident]):
        self.mode: dict = defaultdict(list)    # (設備ID, シナリオID) -> [Incident]（発生日時順）
        for inc in incidents:
            self.mode[(inc.equipment.equipment_id, inc.scenario_id)].append(inc)
        self.times = {k: [i.occurred_at for i in v] for k, v in self.mode.items()}
        self.ms: dict = defaultdict(list)      # (設備ID, サブシステム) -> [発生日時]
        for m in minor_stops():
            if m.alarm is not None:
                self.ms[(m.equipment.equipment_id, m.alarm.subsystem)].append(m.occurred_at)

    def incidents_in(self, key, a: datetime, b: datetime) -> list[Incident]:
        """a < 発生 <= b の同一設備・同一故障モードのトラブル。"""
        ts = self.times.get(key, [])
        return self.mode.get(key, [])[bisect_right(ts, a):bisect_right(ts, b)]

    def minor_in(self, key, a: datetime, b: datetime) -> int:
        lst = self.ms.get(key, [])
        return bisect_right(lst, b) - bisect_right(lst, a)


def _is_recurrence(inc: Incident) -> bool:
    """効果確認で「再発」と数える案件: 報告書で再発扱い（再発フラグ）かつ 中以上。軽微な同種事象は数えない。"""
    return inc.recurrence and inc.severity != "小"


# ---------------------------------------------------------------------------
# 恒久対策のフレーズ（原因分類ごと）
# ---------------------------------------------------------------------------
_MEASURES = {
    "摩耗": ["{part}の交換周期を{mo1}ヶ月→{mo2}ヶ月に短縮し、保全カレンダーへ登録", "{sub}の摩耗量測定を月次点検に追加（判定基準を点検表に明記）",
           "{part}を予防交換部品に指定、在庫{k}個を常備", "摩耗傾向を記録する管理表を作成し、交換目安の{pct}%で計画交換"],
    "劣化": ["{part}を定期交換部品に追加（{yr}年ごと）", "劣化傾向をFDCで監視（{sub}関連パラメータを監視項目に追加）",
           "使用時間カウンタの管理値を見直し、寿命の{pct}%で計画交換", "同ロットの{part}を{sibs}でも前倒し交換"],
    "汚れ": ["{sub}の清掃周期を{wk1}週→{wk2}週に短縮", "清掃手順書 {sop} を改訂（清掃範囲に\"{sub}周辺\"を追加、Rev.{rev}）",
           "清掃後の拭取り確認・パーティクル測定を追加", "汚れの付着状況を写真で記録し、清掃周期の妥当性を{mo2}ヶ月後に再評価"],
    "異物": ["{sub}周辺の発塵源対策（カバー追加）を実施", "PM後の立上げパーティクル確認を必須化（基準 {k}個以下）",
           "異物混入経路を封止し、フィルタ交換周期を見直し", "部品開梱・取付をクリーンブース内作業に限定"],
    "調整不良": ["調整値の管理幅を作業標準に明記（{sop} Rev.{rev}）", "調整後のダブルチェックを必須化（チェックシートに確認者欄を追加）",
             "調整治具を導入し、作業者間のばらつきを低減", "調整作業を認定者限定とし、スキルマップで管理"],
    "設定ミス": ["パラメータ変更時の承認フローを追加（変更申請書で係長承認）", "レシピ・パラメータの変更履歴を自動記録し、差分チェックを週次で実施",
             "設定値一覧を装置横に掲示し、\"変更禁止\"項目を明示"],
    "作業ミス": ["作業手順書 {sop} に写真付きで注意点を追記", "対象作業を認定者限定とし、教育記録で管理", "作業後の相互確認（指差呼称）を徹底",
             "類似作業のヒヤリハットを集めて課内教育を実施（{mo2}月）"],
    "断線": ["ケーブル屈曲部にプロテクタを追加し、固定方法を変更", "{sub}のケーブルを耐屈曲品（メーカー推奨品）へ変更", "導通・絶縁抵抗測定を半年点検に追加",
           "{part}の使用時間を管理し、寿命前に予防交換", "断線の予兆（抵抗値・電流値）をトレンド監視"],
    "締結緩み": ["締結部のトルク管理（{nm}N・m）と合いマークを追加", "緩み止め（ねじロック剤・ばね座金）を採用", "定期増締めを月次点検表に追加"],
    "ソフト不具合": ["メーカー修正版ソフト（Ver.{ver}）を適用", "既知不具合の回避手順を運用マニュアルに追記", "ソフト更新時の受入れ試験項目を追加（{k}項目）"],
    "設計起因": ["メーカー改造（ECN-{ecn}）を同型機に順次適用", "{sub}の構造変更をメーカーへ依頼（{fy}年度に予算化）", "改造完了までの暫定として点検周期を短縮"],
    "外部要因": ["瞬低補償装置（UPS）の対象回路拡大を施設係で検討", "瞬低発生時の一斉復旧手順書を整備し訓練を実施", "停電・瞬低時の装置自動復帰設定を見直し"],
    "不明": ["再発時に備えてログ自動取得を設定", "{sub}の監視項目を追加し、再発時に原因特定できるようにする", "メーカー解析結果を待って恒久対策を決定"],
    "施工不良": ["施工業者へ是正要求、施工チェックリストを追加", "施工後の立会検査を必須化"],
    "能力不足": ["{sub}の能力増強を設備投資計画に計上（{fy}年度）", "負荷分散の運用を見直し"],
    "前工程起因": ["前工程との異常時連絡ルールを明確化", "受入れ時の確認項目を追加（{k}項目）"],
}
_OUTFLOW_CAUSE = ["インライン検査のサンプリング（1枚/ロット）では検出できなかった", "SPC管理限界が広く、異常の兆候を見逃した",
                  "装置アラーム発生時のロット判定ルールが曖昧だった", "トラブル復旧後の着工前確認（QC）が不十分だった",
                  "影響ロットの特定範囲が狭く、前後ロットの確認が漏れた"]
_OUTFLOW_MEASURE = ["トラブル発生時は前後{k}ロットを全数検査とするルールを制定", "インライン検査のサンプリングを1枚→{k}枚/ロットに増加",
                    "SPC管理限界を見直し（±3σ→±2.5σ）", "復旧後QCの合格判定を品証承認制に変更", "装置トラブル情報をMESのロット判定に自動連携"]


def _measure_params(inc: Incident | None, eq, r) -> dict:
    if inc is not None and inc.parts_used:
        part = re.split(r"\s*[（(]", inc.parts_used[0][0].name)[0]
    else:
        part = f"{inc.subsystem if inc else '該当'}部品"
    k = kind_of(eq) if eq is not None else "GEN"
    sibs = [e.equipment_id for e in equipment_master() if eq is not None and kind_of(e) == k and e != eq]
    mo1 = r.choice([6, 12, 12, 24])
    wk1 = r.choice([4, 8, 12])
    filed_year = inc.occurred_at.year if inc else 2025
    return dict(part=part, sub=inc.subsystem if inc else "当該部位", mo1=mo1, mo2=max(1, mo1 // r.choice([2, 3])), wk1=wk1, wk2=wk1 // 2,
                k=r.randint(2, 10), pct=r.choice([70, 80, 80, 90]), yr=r.choice([1, 2, 3]), sop=f"SOP-{k}-{r.randint(1, 48):03d}",
                rev=r.randint(2, 9), nm=r.choice([4, 6, 8, 12, 25]), ver=f"{r.randint(3, 9)}.{r.randint(0, 12)}.{r.randint(0, 30)}",
                ecn=r.randint(1000, 9999), fy=filed_year + 1, sibs="、".join(sibs[:2]) if sibs else "他設備")


# ---------------------------------------------------------------------------
# トラブル起点の行
# ---------------------------------------------------------------------------
_QUALITY_DET = ("SPC", "QC", "インライン", "後工程", "マクロ", "パーティクル")


def _needs_capa(inc: Incident, r) -> bool:
    """是正処置の起票対象か。重大・大は全件、中は影響の大きいものほど起票されやすい。"""
    if inc.severity in ("重大", "大"):
        return True
    if inc.severity != "中":
        return False
    if inc.recurrence or inc.lots_affected or inc.downtime_min >= 300 or inc.cost_yen >= 300_000:
        return r.random() < 0.97
    return r.random() < 0.72


def _drafter(inc: Incident, r) -> Person:
    P = _ppl()
    sect = _maint_section(inc.equipment)
    internal = [p for p in inc.assignees if p.role != "FE"]
    first = internal[0] if internal else next(p for p in P["all"] if p.section == sect and p.role != "係長")
    chief = P["chiefs"].get(sect, P["kacho"])
    staff = [p for p in P["all"] if p.section == sect and p.role in ("主任", "担当")]
    if (inc.category == "品質" or inc.detected_by.startswith(_QUALITY_DET)) and r.random() < 0.45:
        return inc.reporter if inc.reporter.section == "生産技術課" else r.choice(P["pe"])
    if inc.severity == "重大":
        return r.choices([chief, P["kacho"], first], weights=[62, 8, 30])[0]
    if inc.severity == "大":
        leads = [p for p in staff if p.role == "主任"] or staff
        return r.choices([first, r.choice(leads), chief], weights=[60, 30, 10])[0]
    if inc.reporter.section.startswith("設備保全課") and r.random() < 0.15:
        return inc.reporter
    return first


def _approvers(inc: Incident, drafter: Person, r) -> list[Person]:
    P = _ppl()
    chief = P["chiefs"].get(_maint_section(inc.equipment), P["kacho"])
    if inc.severity == "重大":
        out = [P["kacho"]]
        if inc.lots_affected and r.random() < 0.6:
            out.append(r.choice(P["qa"]))
    elif drafter.section == "生産技術課":
        out = [P["pe_lead"] if drafter != P["pe_lead"] else chief]
    elif inc.severity == "大":
        out = [r.choices([chief, P["kacho"]], weights=[70, 30])[0]]
    else:
        out = [chief]
    return [P["kacho"] if p == drafter else p for p in out]


def _equipment_cell(eq, r) -> str:
    return r.choices([eq.equipment_id, f"{eq.equipment_id} {eq.name}", f"{eq.equipment_id}（{eq.name}）", f"{eq.name}（{eq.equipment_id}）"],
                     weights=[45, 30, 20, 5])[0]


def _defect_text(inc: Incident, st: dict, r, dfmt) -> str:
    eq = inc.equipment
    occ = inc.occurred_at
    when = f"{dfmt(occ.date(), False)} {occ.hour:02d}:{occ.minute:02d}"
    impact = [f"装置停止 {inc.downtime_min:,}分（約{inc.downtime_min / 60:.1f}h）"]
    if inc.lots_affected:
        impact.append(f"影響ロット{len(inc.lots_affected)}件（{', '.join(inc.lots_affected)}）")
    if inc.scrap_wafers:
        impact.append(f"廃棄ウェーハ {inc.scrap_wafers}枚")
    if inc.cost_yen and r.random() < 0.7:
        impact.append(r.choice([f"損失 約{inc.cost_yen:,}円", f"費用 ¥{inc.cost_yen:,}"]))
    inv = inc.investigation if (st["verbose"] or inc.severity != "中") else _sentences(inc.investigation, 2)
    rel = ""
    if inc.related_incident_id:
        if "同時刻に工場内複数設備が停止" in inc.symptom:
            rel = f"瞬低による同時停止（親案件 {inc.related_incident_id}）"
        else:
            rel = r.choice([f"{inc.related_incident_id} の再発", f"前回 {inc.related_incident_id}", f"再発（{inc.related_incident_id}参照）"])
    qa_comment = ""
    if inc.lots_affected and r.random() < 0.3:
        qa = r.choice(_ppl()["qa"])
        qa_comment = f'品証（{surname(qa)}）コメント: "' + r.choice(["該当ロットは出荷保留", "後工程へ注意喚起済み", "ロット判定会議で審議予定",
                                                                    "特性影響なしと判断、条件付き流動可"]) + '"'

    if st["hdr"] == "prose":
        lines = [f"{when}頃（{inc.shift}）、{eq.equipment_id}（{eq.name}）の{inc.subsystem}にて、{inc.symptom.rstrip('。')}。"]
        l2 = ""
        if inc.alarm is not None:
            l2 += f'アラーム {inc.alarm.code} "{inc.alarm.message}" 発報。'
        l2 += f"{inc.detected_by}で検知、{surname(inc.reporter)}より連絡。" + "、".join(impact) + "。"
        lines.append(l2)
        lines.append("調査の結果、" + inv if not inv.startswith(("調査", "原因")) else inv)
        if rel:
            lines.append(f"※{rel}")
        if qa_comment:
            lines.append(qa_comment)
        return "\n".join(lines)

    blocks = [(r.choice(["発生日時", "発生"]), f"{when}（{inc.shift}）"),
              ("設備", f"{eq.equipment_id} {eq.name}／{eq.line} {eq.process}／{inc.subsystem}"),
              (r.choice(["現象", "不具合現象", "内容"]), inc.symptom)]
    if inc.alarm is not None:
        blocks.append(("アラーム", f'{inc.alarm.code} "{inc.alarm.message}"（{inc.alarm.category}）'))
    blocks.append(("検知", f"{inc.detected_by}／連絡者 {inc.reporter.name}（{inc.reporter.section}）"))
    blocks.append(("影響", st["join"].join(impact) if len(impact) < 3 or r.random() < 0.5 else "\n".join(f"・{x}" for x in impact)))
    blocks.append((r.choice(["調査結果", "調査", "状況"]), inv))
    if rel:
        blocks.append(("関連", rel))
    if qa_comment:
        blocks.append(("品証", qa_comment))
    return _render_blocks(blocks, st)


def _temp_text(inc: Incident, st: dict, r) -> str:
    parts = ", ".join(f"{p.name}×{q}" for p, q in inc.parts_used)
    result = inc.result or {"対応中": "継続対応中", "保留": "保留（部品・回答待ち）"}.get(inc.status, "")
    blocks = [(r.choice(["初動", "応急処置"]), inc.first_response), (r.choice(["処置", "処置内容", "実施内容"]), inc.action)]
    if parts and r.random() < 0.65:
        blocks.append(("交換部品", parts))
    if result:
        blocks.append((r.choice(["結果", "復旧"]), result))
    if st["hdr"] == "prose":
        return "\n".join(x for x in [f"初動：{inc.first_response}", inc.action, f"交換部品：{parts}" if parts else "", result] if x)
    return _render_blocks(blocks, st)


def _root_text(inc: Incident, st: dict, r) -> str:
    cause = inc.cause.strip()
    if not cause:
        return r.choices(["", "調査中", "原因調査中（メーカー解析待ち）", "未特定"], weights=[55, 20, 15, 10])[0]
    why = [w.rstrip("。") for w in inc.why_why]
    if not why:
        text = cause
    else:
        mode = r.choices(["chain", "list", "cause_root", "oneline"], weights=[30, 20, 35, 15])[0]
        if mode == "chain":
            text = "\n".join(f"なぜ{i}：{w}" for i, w in enumerate(why, 1)) + f"\n→ 真因：{why[-1]}"
        elif mode == "list":
            text = f"直接原因：{cause}\n" + _numbered(why, st["bullet"])
        elif mode == "cause_root":
            text = f"直接原因：{cause}\n真因：{why[-1]}"
        else:
            text = f"{cause}（{'→'.join(why)}）"
    if r.random() < 0.35:
        text = f"[{inc.cause_category}] " + text
    if (inc.lots_affected or inc.category == "品質") and inc.detected_by.startswith(_QUALITY_DET) and r.random() < 0.45:
        text += f"\n流出原因：{r.choice(_OUTFLOW_CAUSE)}"
    return text


# 対策文に含まれる語 -> その対策が成り立つ条件（トラブル記録にこの語のどれかがある / 設備種別）
_MEASURE_NEEDS_WORD = {"ケーブル": ("ケーブル", "配線", "コネクタ", "ハーネス", "信号線", "ケーブルベア")}
_MEASURE_PROCESS_ONLY = ("FDC", "パーティクル", "PM後の立上げ")


def _applicable_measures(inc: Incident) -> list[str]:
    """原因分類の対策候補から、その案件で意味の通らないもの（ヒーター断線にケーブル対策など）を除く。"""
    text = " ".join([inc.cause, inc.investigation, inc.action, inc.symptom])
    process = kind_of(inc.equipment) not in TRANSPORT_KINDS + UTILITY_KINDS
    out = []
    for m in _MEASURES.get(inc.cause_category, _MEASURES["不明"]):
        if any(w in m and not any(k in text for k in ks) for w, ks in _MEASURE_NEEDS_WORD.items()):
            continue
        if not process and any(w in m for w in _MEASURE_PROCESS_ONLY):
            continue
        out.append(m)
    return out or _MEASURES["不明"][:1]


def _perm_text(inc: Incident, st: dict, r) -> str:
    if not inc.cause:
        return r.choices(["", "原因特定後に立案", "検討中"], weights=[50, 30, 20])[0]
    items = [inc.prevention.strip()] if inc.prevention.strip() else []
    if items and items[0].startswith(("特になし", "現状の点検で対応可")):
        n_extra = r.choice([0, 0, 1])
    else:
        n_extra = r.choice([0, 1, 1, 2]) if items else r.choice([1, 1, 2])
    pool = _applicable_measures(inc)
    prm = _measure_params(inc, inc.equipment, r)
    items += [m.format(**prm) for m in r.sample(pool, k=min(n_extra, len(pool)))]
    if inc.status == "保留" and r.random() < 0.5:
        items.append("部品入荷後に本対策を実施")
    if not items:
        return ""
    return items[0] if len(items) == 1 else _numbered(items, st["bullet"])


_EQ_ID_RE = re.compile(r"(?<![A-Za-z0-9\-])[A-Z]{3}-\d{3}(?!\d)")     # 「SOP-CVD-006」のような文書番号の一部は除く


def _horiz_text(inc: Incident, perm: str, done: date | None, r) -> str:
    """水平展開欄。恒久対策に他の設備ID（同ロット部品の前倒し交換など）が出ていれば、その設備への展開として書く。"""
    known = {e.equipment_id for e in equipment_master()}
    others = []
    for eid in _EQ_ID_RE.findall(perm):
        if eid in known and eid != inc.equipment.equipment_id and eid not in others:
            others.append(eid)
    if others and not all(eid in inc.horizontal_deployment for eid in others):
        # 「展開不要」などと矛盾させない。選択は台帳の乱数列を変えないよう案件ごとの補助乱数で行う
        rh = rng(f"t5:horiz:{inc.incident_id}")
        ids = "、".join(others)
        if done is not None:
            return rh.choice([f"{ids} 前倒し交換済み", f"{ids}も同時期に交換済み（交換記録あり）", f"{ids} は{done.month}月PMで交換完了"])
        return rh.choice([f"{ids} は次回PMで前倒し交換予定", f"{ids} 前倒し交換を計画中（部品手配済み）", f"{ids}へ展開予定（交換時期調整中）"])
    if inc.horizontal_deployment:
        return inc.horizontal_deployment
    return r.choices(["", "－", "要否検討中"], weights=[60, 25, 15])[0]


def _deadline(filed: date, sev: str, r) -> date:
    days = {"重大": r.choice([14, 14, 21]), "大": r.choice([30, 30, 45]), "中": r.choice([30, 45, 60, 60])}.get(sev, 30)
    d = filed + timedelta(days=days)
    if r.random() < 0.3:        # 月末締めにする人
        nxt = date(d.year + (d.month == 12), d.month % 12 + 1, 1)
        d = nxt - timedelta(days=1)
    return d


def _deadline_cell(d: date, dfmt, r) -> str:
    if r.random() < 0.03:
        return r.choice(["次回PM時", "部品入荷後", f"{d.month}月末", "至急"])
    return dfmt(d)


def _completion(inc_done: date | None, filed: date, deadline: date, open_: bool, r) -> date | None:
    if open_:
        return None
    span = max(1, (deadline - filed).days)
    if r.random() < 0.72:
        done = filed + timedelta(days=r.randint(max(1, int(span * 0.3)), span))
    else:
        done = deadline + timedelta(days=r.randint(1, 75))       # 期限超過
    if inc_done is not None:
        done = max(done, inc_done)
    return None if done > SNAPSHOT else done


def _effect(key, sub: str, occurred: datetime, done: date | None, checker: str, r, dfmt, idx: _Index,
            open_texts: list[tuple[str, int]]) -> tuple[str, str]:
    """効果確認欄と判定。対策完了後の同一モード再発・チョコ停件数から決める。"""
    if done is None:
        text = r.choices([t for t, _ in open_texts], weights=[w for _, w in open_texts])[0]
        return text, r.choices(["未確認", ""], weights=[85, 15])[0]
    W = r.choice([60, 90, 90, 90, 180])
    wl = {60: "2ヶ月", 90: "3ヶ月", 180: "半年"}[W]
    d0 = datetime.combine(done, time(23, 59))
    end = d0 + timedelta(days=W)
    snap = datetime.combine(SNAPSHOT, time(0, 0))
    after_all = idx.incidents_in(key, d0, min(end, snap))
    after = [i for i in after_all if _is_recurrence(i)]
    light = len(after_all) - len(after)
    before = len(idx.incidents_in(key, occurred - timedelta(days=W), occurred - timedelta(seconds=1))) + 1   # 今回分を含む
    msb = idx.minor_in((key[0], sub), occurred - timedelta(days=W), occurred)
    msa = idx.minor_in((key[0], sub), d0, min(end, snap))
    by = f"（確認 {checker}）" if checker and r.random() < 0.35 else ""
    if after:
        first_dt, first_id = after[0].occurred_at, after[0].incident_id
        if end <= snap and len(after) == 1 and r.random() < 0.12:
            return f"対策後{wl}で1件発生（{first_id}）したが、発生要因が異なるため効果ありと判断{by}", "有効"
        ids = "、".join(x.incident_id for x in after[:3]) + (" ほか" if len(after) > 3 else "")
        return r.choice([
            f"対策後{(first_dt.date() - done).days}日で同モード再発（{first_id}、{dfmt(first_dt.date())}）。恒久対策の見直し要",
            f"再発{len(after)}件（{ids}）。メーカーと追加対策を協議中",
            f"対策前{wl} {before}件 → 対策後 {len(after)}件。効果不十分のため再検討{by}",
        ]), "再検討"
    if end > snap:
        elapsed = (SNAPSHOT - done).days
        return r.choice([f"{wl}フォロー中（{dfmt(end.date())}判定予定）", f"対策後{elapsed}日経過、現時点で再発なし。{dfmt(end.date())}に最終判定",
                         "効果確認期間中", ""]), "未確認"
    if msb >= 4 and msa > msb and r.random() < 0.5:
        return f"同モードの再発はないが、{sub}関連のチョコ停が{msb}件→{msa}件と減っていない。追加対策を検討", "再検討"
    texts = [f"対策後{wl}（〜{dfmt(end.date())}）同一モードの再発なし。対策前{wl}：{before}件{by}",
             f"再発なし（{dfmt(end.date())} 確認）{by}",
             f"FDC・アラーム履歴を確認し、{wl}間 異常なし。効果あり",
             f'{wl}フォロー完了。"再発なし"、水平展開先も異常なし']
    if msb:
        texts.append(f"{wl}フォロー完了、再発0件。{sub}関連チョコ停 {msb}件→{msa}件")
    text = r.choice(texts)
    if light and r.random() < 0.5:
        text += f"。軽微な同種事象{light}件あり（停止影響小のため再発扱いせず）"
    verdict = "有効" if r.random() > 0.015 else ""        # 判定の記入漏れ
    return text, verdict


_OPEN_TEXTS = [("", 45), ("対策未完了のため未確認", 25), ("恒久対策実施後に確認", 20), ("部品入荷待ち（対策後に確認）", 10)]


def _incident_row(inc: Incident, r, idx: _Index) -> _Row:
    lag = {"重大": r.choice([0, 0, 1]), "大": r.choice([0, 1, 1, 2, 3])}.get(inc.severity, r.choice([1, 2, 3, 5, 7, 10]))
    if r.random() < 0.03:
        lag += r.randint(15, 40)               # 起票忘れで後日まとめて起票
    filed = min(SNAPSHOT, inc.reported_at.date() + timedelta(days=lag))
    dfmt = _date_writer(filed, r)
    drafter = _drafter(inc, r)
    st = _style(drafter.employee_id)
    deadline = _deadline(filed, inc.severity, r)
    open_ = (inc.status == "対応中" or (inc.status == "保留" and r.random() < 0.85)
             or (inc.status == "経過観察" and r.random() < 0.4))
    done = _completion(inc.completed_at.date() if inc.completed_at else None, filed, deadline, open_, r)
    internal = [p for p in inc.assignees if p.role != "FE"]
    effect, verdict = _effect((inc.equipment.equipment_id, inc.scenario_id), inc.subsystem, inc.occurred_at, done,
                              surname(internal[0]) if internal else "", r, dfmt, idx, _OPEN_TEXTS)
    approvers = _approvers(inc, drafter, r)
    appr = r.choice(["／", "、"]).join(_name_variant(p, r) for p in approvers)
    if done is None and r.random() < 0.7:
        appr = ""                               # 未完了は承認前
    related = inc.incident_id
    if inc.related_incident_id and r.random() < 0.6:
        related += r.choice([f" / {inc.related_incident_id}", f"、{inc.related_incident_id}", f"\n{inc.related_incident_id}（前回）"])
    # 乱数を使う順番は従来どおり（列の並び順）。水平展開は恒久対策の内容を見て決める
    filed_cell, dept, name = dfmt(filed, False), _dept_variant(drafter, r), _name_variant(drafter, r)
    eq_cell = _equipment_cell(inc.equipment, r)
    defect = _defect_text(inc, st, r, dfmt)
    temp = _temp_text(inc, st, r)
    root = _root_text(inc, st, r)
    perm = _perm_text(inc, st, r)
    values = {
        "起票日": filed_cell, "起票部署": dept, "起票者": name,
        "対象設備": eq_cell, "関連トラブルNo": related, "不具合内容": defect,
        "暫定対策": temp, "真因(なぜなぜ要約)": root, "恒久対策": perm,
        "水平展開": _horiz_text(inc, perm, done, r), "期限": _deadline_cell(deadline, dfmt, r), "完了日": dfmt(done),
        "効果確認": effect, "判定": verdict, "承認者": appr,
    }
    return _Row(key=f"I:{inc.incident_id}", filed=filed, sort_dt=inc.reported_at, source="incident",
                incident_id=inc.incident_id, deadline=deadline, done=done, values=values)


# ---------------------------------------------------------------------------
# 品質クレーム（後工程・顧客・社内判定からの不具合連絡）
# ---------------------------------------------------------------------------
_CLAIM_SOURCES = [("後工程（組立・テスト工程）", "社内", 34), ("顧客 A社（車載Tier1）", "顧客", 14), ("顧客 B社（民生機器）", "顧客", 10),
                  ("顧客 C社（産業機器）", "顧客", 8), ("WAT判定会議", "社内", 16), ("出荷検査（品証）", "社内", 12), ("信頼性試験（HTOL）", "社内", 6)]
_PRODUCTS = {"MK": "車載PMIC", "SR": "CIS用ロジック", "PX": "汎用MCU", "TQ": "パワーMOSFET", "NB": "NORフラッシュ", "HV": "高耐圧ドライバIC"}
# 設備群 -> [(不良名, 検出工程, 解析所見)]
_CLAIM_DEFECTS = {
    "CMP": [("配線間ショート", "PT（プローブテスト）", "不良品の平面SEM観察で配線間にマイクロスクラッチ（長さ約{um3}μm）を確認"),
            ("ビア抵抗高", "WAT", "断面TEMでW残り・リセス（{nm}nm）を確認"),
            ("残膜厚不足", "出荷前インライン測定", "残膜厚が規格下限を{nm}nm下回り、ウェーハ外周で薄い")],
    "CVD": [("層間絶縁耐圧不良", "信頼性試験（HTOL 168h）", "故障解析で層間膜中にパーティクル（φ{um}μm）を確認"),
            ("W埋込み不良による開放不良", "FT（ファイナルテスト）", "断面SEMでコンタクト内部にボイドを確認"),
            ("膜ストレス起因のクラック", "外観検査", "ウェーハ外周部に放射状クラック、膜ストレス {mpa}MPa（通常 -150MPa前後）")],
    "ETC": [("コンタクトオープン", "PT（プローブテスト）", "断面SEMでコンタクト底部にエッチング残渣を確認"),
            ("リーク電流増", "WAT", "ゲートCDが狙い値比 {nm}nm 細り"),
            ("ゲート酸化膜破壊", "FT（ファイナルテスト）", "アンテナTEGで異常、プラズマダメージ起因と推定")],
    "LIT": [("重ね合わせずれによる開放不良", "PT（プローブテスト）", "重ね合わせ実測 {nm}nm（管理値 ±15nm）"),
            ("繰返し欠陥", "外観検査", "全ショット同一座標に欠陥、レチクル上異物の転写と判明"),
            ("CD異常（デフォーカス）", "WAT", "ウェーハ中央部でCDが{nm}nm太り")],
    "CLN": [("ウォーターマーク起因の外観不良", "外観検査", "乾燥ムラ跡（ウォーターマーク）がウェーハ外周に分布"),
            ("金属汚染によるリーク不良", "WAT", "TXRFでFe {e}E10 atoms/cm2 を検出"),
            ("パーティクル残りによる歩留まり低下", "PT（プローブテスト）", "欠陥マップが洗浄ユニットのノズル軌跡と一致")],
    "IMP": [("Vthシフト", "WAT", "Vthが狙い値比 +{mv}mV、シート抵抗 {pct}%高"),
            ("特性ばらつき（注入角ずれ）", "WAT", "ウェーハ面内でVthが傾斜、注入角 {deg}°ずれ相当"),
            ("ゲート破壊（チャージアップ）", "FT（ファイナルテスト）", "ESD類似の破壊痕、注入時のチャージアップと推定")],
    "INS": [("欠陥の検出漏れによる流出", "顧客受入検査", "流出品の欠陥サイズ {um}μm、設定感度では検出できるはずだった"),
            ("測定値オフセットによる規格外品流出", "出荷前測定", "測定値に{nm}nmのオフセット、標準試料での校正ずれ")],
    "TRN": [("ウェーハ欠け・クラック", "外観検査", "ウェーハエッジ部に欠け（{mm}mm）、搬送時の衝撃痕"),
            ("工程飛ばし（FOUP取違え）", "WAT", "1工程未処理のウェーハが混入、搬送履歴で取違えを確認"),
            ("裏面スクラッチ", "外観検査", "裏面に円弧状スクラッチ、ハンド接触痕と一致")],
    "UTL": [("配線腐食", "信頼性試験（THB）", "Al配線に腐食、洗浄純水の比抵抗低下時期と一致"),
            ("パーティクル増加による歩留まり低下", "PT（プローブテスト）", "複数工程にまたがりランダム欠陥が増加"),
            ("膜質異常", "WAT", "排気圧変動時期の処理ロットで膜質ばらつき")],
}
# 不良名 -> トラブル記録（現象・原因・部位・調査）に現れやすい語。紐付くトラブルと矛盾しない不良を選ぶのに使う
_DEFECT_KEYWORDS = {
    "配線間ショート": ("スクラッチ", "パッド", "異物", "パーティクル", "欠陥", "コンディショナ", "ドレッサ", "スラリー"),
    "ビア抵抗高": ("W-CMP", "リセス", "ディッシング", "エロージョン", "終点", "オーバー研磨", "EPD"),
    "残膜厚不足": ("膜厚", "残膜", "レート", "研磨量", "ヘッド", "圧力", "メンブレン", "リテーナ"),
    "層間絶縁耐圧不良": ("パーティクル", "異物", "フレーク", "剥離", "チャンバー"),
    "W埋込み不良による開放不良": ("WF6", "ボイド", "埋込", "W-CVD"),
    "膜ストレス起因のクラック": ("ストレス", "膜質", "屈折率", "RF", "ヒーター", "温度", "膜厚"),
    "コンタクトオープン": ("残渣", "エッチレート", "E/R", "EPD", "終点", "ガス", "MFC"),
    "リーク電流増": ("CD", "均一性", "ESC", "He", "温度", "チラー"),
    "ゲート酸化膜破壊": ("アーク", "アーキング", "RF", "プラズマ", "Vpp", "マッチング"),
    "重ね合わせずれによる開放不良": ("重ね合わせ", "アライメント", "ステージ", "位置", "干渉計", "エンコーダ"),
    "繰返し欠陥": ("レチクル", "異物", "ペリクル"),
    "CD異常（デフォーカス）": ("フォーカス", "AF", "CD", "露光量", "ドーズ", "光源", "照度"),
    "ウォーターマーク起因の外観不良": ("乾燥", "ウォーターマーク", "IPA", "スピン", "チャック", "N2"),
    "金属汚染によるリーク不良": ("金属", "薬液", "濃度", "汚染", "槽"),
    "パーティクル残りによる歩留まり低下": ("パーティクル", "ノズル", "ブラシ", "滴下", "フィルタ"),
    "Vthシフト": ("ドーズ", "ビーム電流", "ファラデー", "HV", "高圧", "アーク", "ソース"),
    "特性ばらつき（注入角ずれ）": ("角度", "チルト", "スキャン", "プラテン"),
    "ゲート破壊（チャージアップ）": ("チャージ", "フラッド", "PFG", "電子"),
    "欠陥の検出漏れによる流出": ("感度", "検出", "光源", "光量", "ランプ", "画像", "レシピ"),
    "測定値オフセットによる規格外品流出": ("校正", "オフセット", "測定", "膜厚", "重ね合わせ", "ステージ"),
    "ウェーハ欠け・クラック": ("落下", "衝撃", "振動", "ホイスト", "ベルト", "グリッパ", "割れ", "欠け"),
    "工程飛ばし（FOUP取違え）": ("FOUP ID", "BCR", "RFID", "取違", "誤搬送"),
    "裏面スクラッチ": ("ハンド", "吸着", "ロボ", "ティーチング", "接触"),
    "配線腐食": ("比抵抗", "純水", "TOC", "RO", "イオン"),
    "パーティクル増加による歩留まり低下": ("微粒子", "パーティクル", "UF", "フィルタ"),
    "膜質異常": ("排気", "静圧", "ファン", "ダンパ", "冷却水", "温度", "真空"),
}


def _pick_defect(eq, inc: Incident | None, r) -> tuple:
    """設備群の不良候補から、紐付くトラブルの記録と最も語が重なるものを選ぶ（重ならなければランダム）。"""
    cands = _CLAIM_DEFECTS[_group_of(eq)]
    if inc is None:
        return r.choice(cands)
    text = " ".join([inc.symptom, inc.cause, inc.subsystem, inc.investigation, eq.process])
    scored = [(sum(1 for k in _DEFECT_KEYWORDS.get(d[0], ()) if k in text), d) for d in cands]
    best = max(sc for sc, _ in scored)
    return r.choice([d for sc, d in scored if sc == best]) if best > 0 else r.choice(cands)


_CLAIM_TEMP = ["該当ロット（{lots}）を出荷停止", "倉庫在庫 {stock:,}個を全数再テストで選別", "同時期に{eqid}で処理した{m}ロットを追加抜取り検査",
               "顧客在庫 {cstock:,}個の返品を受入れ、代替品を{d1}に出荷", "顧客へ初報（{d0}）、中間報告を{d2}に実施",
               "{eqid}を着工停止し、QCで正常を確認後に再開", "後工程へ同時期ロットの注意喚起を発信", "不良品{ret}個を返却受入れ、解析を開始"]


def _group_of(eq) -> str:
    k = kind_of(eq)
    if k in TRANSPORT_KINDS:
        return "TRN"
    if k in UTILITY_KINDS:
        return "UTL"
    return k


def _lot(r, d: date, prod: str | None = None) -> str:
    prod = prod or r.choice(list(_PRODUCTS))
    return f"{prod}{d.year % 100:02d}{d.isocalendar()[1]:02d}{r.randint(1, 399):03d}.{r.randint(1, 3)}"


def _claim_params(r) -> dict:
    return dict(um3=r.randint(20, 300), nm=r.randint(8, 60), um=round(r.uniform(0.1, 2.5), 1), mpa=r.randint(-420, -260),
                e=round(r.uniform(1.2, 9.5), 1), mv=r.randint(25, 120), pct=r.randint(4, 18), deg=round(r.uniform(0.3, 1.5), 1),
                mm=round(r.uniform(0.2, 3.0), 1))


def _claim_drafter(r) -> Person:
    qa = _ppl()["qa"]
    return r.choices(qa, weights=[35 if p.role == "主任" else 65 for p in qa])[0]


def _claim_approvers(kind: str, drafter: Person, r) -> str:
    P = _ppl()
    names = [P["kacho"]] if kind == "顧客" else []
    lead = next((p for p in P["qa"] if p.role == "主任"), P["qa"][0])
    if drafter != lead:
        names.insert(0, lead)
    if not names:
        names = [P["pe_lead"]]
    return r.choice(["／", "、"]).join(_name_variant(p, r) for p in names)


def _claim_body(src: str, kind: str, recv: date, lots: list[str], defect: tuple, prm: dict, st: dict, r, dfmt,
                investigation: str, due: date) -> str:
    name, test, finding = defect
    prod = "・".join(dict.fromkeys(_PRODUCTS[lot[:2]] for lot in lots))
    rate = round(r.uniform(0.8, 9.5), 1)
    base = round(r.uniform(0.05, 0.6), 2)
    ng = r.randint(40, 4800)
    blocks = [("連絡元", f"{src}　受付 {dfmt(recv, False)}"),
              ("対象製品", f"{prod}（ロット {', '.join(lots)}、計{len(lots)}ロット）"),
              ("不具合内容", f"{test}にて{name}が発生。不良率 {rate}%（通常 {base}%）、不良数 {ng:,}個"),
              ("解析結果", finding.format(**prm)),
              ("当社調査", investigation)]
    if kind == "顧客":
        blocks.append(("顧客要求", r.choice([f"8D報告書を{dfmt(due)}までに提出", f"初回回答（D3まで）を{dfmt(recv + timedelta(days=3))}、最終報告を{dfmt(due)}",
                                          f'"原因と流出防止策" の回答を{dfmt(due)}までに'])))
    if st["hdr"] == "prose":
        return "\n".join(b for _, b in blocks)
    return _render_blocks(blocks, st)


def _linked_claim_row(inc: Incident, capa_ids: set, r, idx: _Index) -> _Row | None:
    src, kind, _ = r.choices(_CLAIM_SOURCES, weights=[w for *_, w in _CLAIM_SOURCES])[0]
    lag_max = (SNAPSHOT - inc.occurred_at.date()).days - 3
    if lag_max < 5:
        return None
    recv = inc.occurred_at.date() + timedelta(days=r.randint(5, min(60, lag_max)))
    filed = min(SNAPSHOT, recv + timedelta(days=r.choice([0, 0, 1, 2])))
    dfmt = _date_writer(filed, r)
    drafter = _claim_drafter(r)
    st = _style(drafter.employee_id)
    eq = inc.equipment
    lots = inc.lots_affected[: r.choice([1, 2, 3, 6])]
    defect = _pick_defect(eq, inc, r)
    prm = _claim_params(r)
    due = filed + timedelta(days=r.choice([14, 30, 30]))
    ca_ref = f"設備側の是正は«CA:{inc.incident_id}»で管理。" if inc.incident_id in capa_ids and r.random() < 0.8 else ""
    inv = (f"ロット履歴を追跡し、{eq.equipment_id}（{eq.name}）の{dfmt(inc.occurred_at.date())}処理分に集中していることを確認。"
           f"同時期に{inc.incident_id}（{inc.subsystem}：{_short(inc.symptom)}）が発生しており関連性大。{ca_ref}")
    body = _claim_body(src, kind, recv, lots, defect, prm, st, r, dfmt, inv, due)

    tctx = dict(lots=", ".join(lots), stock=r.randint(500, 60000), eqid=eq.equipment_id, m=r.randint(3, 30), cstock=r.randint(200, 20000),
                d0=dfmt(recv), d1=dfmt(recv + timedelta(days=r.randint(3, 12))), d2=dfmt(recv + timedelta(days=r.randint(7, 20))),
                ret=r.randint(5, 60))
    temp_pool = [t for t in _CLAIM_TEMP if kind == "顧客" or ("顧客" not in t)]
    temp = _numbered([t.format(**tctx) for t in r.sample(temp_pool, k=r.randint(2, 4))], st["bullet"])

    if inc.cause:
        why = "→".join(w.rstrip("。") for w in inc.why_why) if inc.why_why and r.random() < 0.5 else ""
        root = f"発生原因：{inc.cause}" + (f"（{why}）" if why else "") + f"\n流出原因：{r.choice(_OUTFLOW_CAUSE)}"
    else:
        root = r.choice(["発生原因：調査中\n流出原因：調査中", "解析中（返却品待ち）", ""])
    prm2 = _measure_params(inc, eq, r)
    occ_measure = inc.prevention or r.choice(_applicable_measures(inc)).format(**prm2)
    perm = f"【発生対策】{occ_measure}\n【流出対策】{r.choice(_OUTFLOW_MEASURE).format(**prm2)}" if inc.cause else r.choice(["", "原因確定後に立案"])
    horiz = r.choice([f"同工程を流れる他製品（{'、'.join(r.sample(sorted(_PRODUCTS.values()), 2))}）でも同条件を確認", f"{prm2['sibs']}の処理ロットも遡り確認、異常なし",
                      "全製品の出荷判定基準に反映", "他工場（後工程）へ事例展開", ""])

    deadline = due
    open_ = inc.status in ("対応中", "保留") or r.random() < 0.08
    done = _completion(inc.completed_at.date() if inc.completed_at else None, filed, deadline, open_, r)
    if done is None:
        effect, verdict = r.choice(["", "対策未完了", "顧客回答待ち"]), r.choices(["未確認", ""], weights=[85, 15])[0]
    else:
        key = (eq.equipment_id, inc.scenario_id)
        end = datetime.combine(done, time(23, 59)) + timedelta(days=90)
        after = [i for i in idx.incidents_in(key, datetime.combine(done, time(23, 59)), min(end, datetime.combine(SNAPSHOT, time(0, 0))))
                 if _is_recurrence(i)]
        if after:
            effect, verdict = f"設備側で同モード再発（{after[0].incident_id}）。流出はないが発生対策を再検討", "再検討"
        elif end.date() > SNAPSHOT:
            effect, verdict = r.choice([f"出荷後フォロー中（{dfmt(end.date())}まで）", "効果確認期間中"]), "未確認"
        else:
            effect = r.choice([f"対策後 出荷{r.randint(12, 180)}ロットで同不良なし（不良率 {round(r.uniform(0.01, 0.3), 2)}%）",
                               f"顧客より8D報告書の承認取得（{dfmt(done + timedelta(days=r.randint(5, 30)))}）" if kind == "顧客" else f"WAT判定会議で効果確認済み（{dfmt(end.date())}）",
                               f"後工程の不良率 {round(r.uniform(1.0, 6.0), 1)}%→{round(r.uniform(0.02, 0.4), 2)}% に改善、効果あり"])
            verdict = "有効"
    appr = _claim_approvers(kind, drafter, r) if (done is not None or r.random() < 0.3) else ""
    values = {
        "起票日": dfmt(filed, False), "起票部署": _dept_variant(drafter, r), "起票者": _name_variant(drafter, r),
        "対象設備": _equipment_cell(eq, r), "関連トラブルNo": inc.incident_id, "不具合内容": body, "暫定対策": temp,
        "真因(なぜなぜ要約)": root, "恒久対策": perm, "水平展開": horiz, "期限": _deadline_cell(deadline, dfmt, r), "完了日": dfmt(done),
        "効果確認": effect, "判定": verdict, "承認者": appr,
    }
    return _Row(key=f"C:{inc.incident_id}", filed=filed, sort_dt=datetime.combine(recv, time(9, 0)), source="claim",
                incident_id=inc.incident_id, deadline=deadline, done=done, values=values)


# トラブルと紐付かないクレームの結末: (真因, 対象設備の書き方, 恒久対策, 判定候補)
_FREE_OUTCOMES = [
    ("ウェーハ受入品（基板メーカー）の結晶欠陥起因と判明。当社工程起因ではない", "該当なし", "基板メーカーへ是正要求（回答書受領）、受入検査に抜取りX線トポを追加", "有効"),
    ("顧客実装工程（リフロー温度プロファイル）起因と判明", "該当なし", "顧客へ推奨実装条件を再提示、技術資料を改訂", "有効"),
    ("原因特定できず。同時期ロットの工程履歴・装置ログに異常なし", "", "同製品の不良率を週次でフォロー（監視強化）", "未確認"),
    ("{eid}の{sub}起因と推定（確証なし）。ロット共通工程の絞込みで可能性が最も高い", "{eid}（推定）", "{eid}の{sub}を臨時点検、条件出しで再現確認中", "再検討"),
    ("解析中（顧客からの返却品待ち）", "調査中", "", "未確認"),
    ("測定器（テスタ）の校正ずれによる誤判定。製品自体は良品", "該当なし", "テスタの校正周期を短縮し、日常点検にゴールデンサンプル測定を追加", "有効"),
]


def _free_claim_row(i: int, r) -> _Row:
    span = (SNAPSHOT - date(2023, 4, 10)).days
    recv = date(2023, 4, 10) + timedelta(days=r.randint(0, span - 3))
    filed = min(SNAPSHOT, recv + timedelta(days=r.choice([0, 1, 1, 3])))
    dfmt = _date_writer(filed, r)
    src, kind, _ = r.choices(_CLAIM_SOURCES, weights=[w for *_, w in _CLAIM_SOURCES])[0]
    drafter = _claim_drafter(r)
    st = _style(drafter.employee_id)
    eqs = [e for e in equipment_master() if e.installed_date < recv and kind_of(e) not in UTILITY_KINDS]
    eq = r.choice(eqs)
    prod = r.choice(list(_PRODUCTS))
    lots = [_lot(r, recv - timedelta(days=r.randint(10, 40)), prod) for _ in range(r.choice([1, 1, 2, 3]))]
    defect = _pick_defect(eq, None, r)
    prm = _claim_params(r)
    due = filed + timedelta(days=r.choice([14, 30, 30, 45]))
    cause, eq_cell, perm, verdict = r.choices(_FREE_OUTCOMES, weights=[18, 10, 22, 22, 14, 14])[0]
    fmtp = dict(eid=eq.equipment_id, sub=eq.process)
    cause, eq_cell, perm = cause.format(**fmtp), eq_cell.format(**fmtp), perm.format(**fmtp)
    inv = r.choice([f"該当ロットの工程履歴を全工程で確認中。共通工程は{eq.process}ほか{r.randint(2, 6)}工程",
                    "同時期の装置トラブル記録を照合したが該当なし", f"同一製品の前後{r.randint(3, 12)}ロットを追加測定、異常なし",
                    f"不良品{r.randint(5, 40)}個を解析依頼（社内解析センター）"])
    body = _claim_body(src, kind, recv, lots, defect, prm, st, r, dfmt, inv, due)
    tctx = dict(lots=", ".join(lots), stock=r.randint(500, 60000), eqid=eq.equipment_id, m=r.randint(3, 30), cstock=r.randint(200, 20000),
                d0=dfmt(recv), d1=dfmt(recv + timedelta(days=r.randint(3, 12))), d2=dfmt(recv + timedelta(days=r.randint(7, 20))),
                ret=r.randint(5, 60))
    temp_pool = [t for t in _CLAIM_TEMP if "{eqid}" not in t and (kind == "顧客" or "顧客" not in t)]
    temp = _numbered([t.format(**tctx) for t in r.sample(temp_pool, k=r.randint(1, 3))], st["bullet"])
    open_ = verdict == "未確認" and r.random() < 0.7 or cause.startswith("解析中")
    done = _completion(None, filed, due, open_, r)
    if done is None:
        effect, verdict = r.choice(["", "原因調査中", "返却品の解析結果待ち"]), r.choices(["未確認", ""], weights=[85, 15])[0]
        perm = perm if r.random() < 0.5 else ""
    elif verdict == "有効":
        effect = r.choice([f"対策後の同製品で不良再発なし（{dfmt(min(SNAPSHOT, done + timedelta(days=90)))}時点）", "顧客了承済み",
                           f"{r.randint(20, 150)}ロット流動し同不良ゼロ"])
    elif verdict == "再検討":
        effect = r.choice(["臨時点検で異常見つからず。原因の再調査が必要", f"同製品で再度不良発生（{dfmt(min(SNAPSHOT, done + timedelta(days=r.randint(10, 60))))}）"])
    else:
        effect = r.choice(["監視継続中", "フォロー中（不良率推移を確認）"])
    appr = _claim_approvers(kind, drafter, r) if done is not None else ""
    values = {
        "起票日": dfmt(filed, False), "起票部署": _dept_variant(drafter, r), "起票者": _name_variant(drafter, r),
        "対象設備": eq_cell, "関連トラブルNo": "", "不具合内容": body, "暫定対策": temp,
        "真因(なぜなぜ要約)": cause, "恒久対策": perm,
        "水平展開": r.choice(["", "－", "同製品群の出荷検査で注意喚起", "他製品への影響なし（確認済み）"]),
        "期限": _deadline_cell(due, dfmt, r), "完了日": dfmt(done), "効果確認": effect, "判定": verdict, "承認者": appr,
    }
    return _Row(key=f"F:{i:04d}", filed=filed, sort_dt=datetime.combine(recv, time(10, 0)), source="claim_free",
                deadline=due, done=done, values=values)


# ---------------------------------------------------------------------------
# 重複起票
# ---------------------------------------------------------------------------
def _duplicate_row(orig: _Row, inc: Incident, r) -> _Row:
    """同じトラブルを別の人がもう一度起票してしまった行（後で統合された体裁）。"""
    filed = min(SNAPSHOT, orig.filed + timedelta(days=r.randint(1, 4)))
    dfmt = _date_writer(filed, r)
    P = _ppl()
    cands = [p for p in inc.assignees if p.role != "FE"] + P["pe"]
    drafter = r.choice(cands)
    eq = inc.equipment
    body = (f"{dfmt(inc.occurred_at.date(), False)} {eq.equipment_id} {inc.subsystem} {_short(inc.symptom, 60)}\n"
            f"停止{inc.downtime_min:,}分。詳細はトラブル報告書 {inc.incident_id} 参照")
    values = {
        "起票日": dfmt(filed, False), "起票部署": _dept_variant(drafter, r), "起票者": _name_variant(drafter, r),
        "対象設備": eq.equipment_id, "関連トラブルNo": inc.incident_id, "不具合内容": body,
        "暫定対策": f"※«NO:{orig.key}»と同一案件のため記載省略", "真因(なぜなぜ要約)": "", "恒久対策": "", "水平展開": "",
        "期限": "", "完了日": "", "効果確認": f"重複起票のため«NO:{orig.key}»に統合", "判定": "", "承認者": "",
    }
    return _Row(key=f"D:{inc.incident_id}", filed=filed, sort_dt=inc.reported_at + timedelta(days=1), source="duplicate",
                incident_id=inc.incident_id, values=values)


# ---------------------------------------------------------------------------
# 台帳の組立て
# ---------------------------------------------------------------------------
_BLANKABLE = [("起票者", 0.005), ("対象設備", 0.004), ("期限", 0.008), ("起票部署", 0.003)]
_PLACEHOLDER_RE = re.compile(r"«(CA|NO):([^»]+)»")


def build_rows() -> list[_Row]:
    incidents = standard_incidents()
    idx = _Index(incidents)
    r_sel = rng("t5:select")
    r_inc = rng("t5:incident-rows")
    rows: list[_Row] = []
    by_id = {}
    for inc in incidents:
        if _needs_capa(inc, r_sel):
            row = _incident_row(inc, r_inc, idx)
            rows.append(row)
            by_id[inc.incident_id] = row
    capa_ids = set(by_id)

    # 品質クレーム（トラブル紐付き）: ロット影響のある品質系トラブルから抽出
    r_cl = rng("t5:claims")
    cands = [i for i in incidents if i.lots_affected and i.status != "対応中"
             and (i.category == "品質" or i.scrap_wafers > 0 or i.detected_by.startswith(_QUALITY_DET))]
    for inc in sorted(r_cl.sample(cands, k=min(N_LINKED_CLAIMS, len(cands))), key=lambda x: x.incident_id):
        row = _linked_claim_row(inc, capa_ids, r_cl, idx)
        if row is not None:
            rows.append(row)
    r_free = rng("t5:free-claims")
    rows += [_free_claim_row(i, r_free) for i in range(N_FREE_CLAIMS)]

    # 重複起票
    r_dup = rng("t5:duplicates")
    inc_by_id = {i.incident_id: i for i in incidents}
    heavy = [row for row in rows if row.source == "incident" and inc_by_id[row.incident_id].severity in ("重大", "大")]
    for orig in sorted(r_dup.sample(heavy, k=N_DUPLICATES), key=lambda x: x.key):
        rows.append(_duplicate_row(orig, inc_by_id[orig.incident_id], r_dup))

    # 必須項目の記入漏れ（重複行以外）
    r_blank = rng("t5:blanks")
    for row in rows:
        if row.source == "duplicate":
            continue
        for col, p in _BLANKABLE:
            if row.values[col] and r_blank.random() < p:
                row.values[col] = ""
        # まれに設備IDを全角で書く・トラブルNoのハイフン抜け
        if row.values["対象設備"] and r_blank.random() < 0.004:
            row.values["対象設備"] = to_zenkaku(row.values["対象設備"])
        if row.values["関連トラブルNo"] and r_blank.random() < 0.003:
            row.values["関連トラブルNo"] = row.values["関連トラブルNo"].replace("TR-", "TR", 1)

    # 採番（起票日順・年ごとに連番）と、相互参照の差し込み
    rows.sort(key=lambda x: (x.filed, x.sort_dt, x.key))
    seq: Counter = Counter()
    no_of_key = {}
    for row in rows:
        seq[row.filed.year] += 1
        row.values["台帳No"] = f"CA-{row.filed.year}-{seq[row.filed.year]:04d}"
        no_of_key[row.key] = row.values["台帳No"]
    ca_of_incident = {k[2:]: v for k, v in no_of_key.items() if k.startswith("I:")}

    def _sub(m: re.Match) -> str:
        return ca_of_incident.get(m.group(2), "") if m.group(1) == "CA" else no_of_key.get(m.group(2), "")

    for row in rows:
        for col in LONG_COLS + ("暫定対策",):
            if "«" in row.values[col]:
                row.values[col] = _PLACEHOLDER_RE.sub(_sub, row.values[col])
    return rows


# ---------------------------------------------------------------------------
# 書き出し
# ---------------------------------------------------------------------------
def _write_csv(rows: list[_Row], path: Path) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, lineterminator="\n", quoting=csv.QUOTE_MINIMAL)
        w.writerow(COLUMNS)
        for row in rows:
            w.writerow([row.values[c] for c in COLUMNS])


_FIXED_TS = datetime(2026, 9, 10, 9, 0, 0)


def _save_xlsx_deterministic(wb: Workbook, path: Path) -> None:
    """openpyxl は保存時刻を埋め込むため、ZIP内の時刻と更新日時を固定して毎回同じバイト列にする。"""
    wb.properties.creator = "設備保全課"
    wb.properties.created = _FIXED_TS
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    with zipfile.ZipFile(buf) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)", rb"\g<1>2026-09-10T09:00:00Z\g<2>", data)
            zi = zipfile.ZipInfo(info.filename, date_time=_FIXED_TS.timetuple()[:6])
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = info.external_attr
            zout.writestr(zi, data)


def _write_xlsx(rows: list[_Row], path: Path) -> dict:
    wb = Workbook()
    ws = wb.active
    ws.title = "2026上期"
    thin = Side(style="thin", color="808080")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    group_fill = {"不具合": "F8CBAD", "対策": "C6E0B4", "確認": "BDD7EE"}
    sub_fill = {"不具合": "FCE4D6", "対策": "E2EFDA", "確認": "DDEBF7"}
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for label, c0, c1 in GROUPS:
        cell = ws.cell(row=1, column=c0 + 1, value=label)
        ws.merge_cells(start_row=1, start_column=c0 + 1, end_row=1, end_column=c1)
        for c in range(c0 + 1, c1 + 1):
            ws.cell(row=1, column=c).fill = PatternFill("solid", fgColor=group_fill[label])
            ws.cell(row=1, column=c).border = border
            h = ws.cell(row=2, column=c, value=COLUMNS[c - 1])
            h.fill = PatternFill("solid", fgColor=sub_fill[label])
            h.font = Font(bold=True)
            h.alignment = center
            h.border = border
        cell.font = Font(bold=True, size=12)
        cell.alignment = center
    wrap = Alignment(wrap_text=True, vertical="top")
    top = Alignment(vertical="top")
    n_date_cells = n_text_dates = 0
    for i, row in enumerate(rows, start=3):
        for j, col in enumerate(COLUMNS, start=1):
            v = row.values[col]
            cell = ws.cell(row=i, column=j)
            if col in DATE_COLS and v and _ISO_RE.fullmatch(v):
                cell.value = date.fromisoformat(v)
                cell.number_format = "yyyy/mm/dd"
                n_date_cells += 1
            else:
                cell.value = v if v != "" else None
                if col in DATE_COLS and v:
                    n_text_dates += 1          # 「次回PM時」「2026/5/7」など文字列のまま入った日付欄
            cell.alignment = wrap if col in LONG_COLS or col == "関連トラブルNo" else top
    widths = {"台帳No": 13, "起票日": 11, "起票部署": 16, "起票者": 12, "対象設備": 22, "関連トラブルNo": 16, "不具合内容": 60,
              "暫定対策": 50, "真因(なぜなぜ要約)": 40, "恒久対策": 40, "水平展開": 28, "期限": 11, "完了日": 11, "効果確認": 36,
              "判定": 8, "承認者": 16}
    for j, col in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(j)].width = widths[col]
    last = 2 + len(rows)
    ws.freeze_panes = "C3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(COLUMNS))}{last}"
    dv = DataValidation(type="list", formula1='"有効,再検討,未確認"', allow_blank=True)
    dv.error = "有効・再検討・未確認 から選択"
    ws.add_data_validation(dv)
    vcol = get_column_letter(COLUMNS.index("判定") + 1)
    dv.add(f"{vcol}3:{vcol}{max(3, last)}")
    ws.print_title_rows = "1:2"
    _save_xlsx_deterministic(wb, path)
    return dict(date_cells=n_date_cells, text_dates=n_text_dates, merged=[str(m) for m in ws.merged_cells.ranges])


def _stats(rows: list[_Row]) -> dict:
    st = dict(total=len(rows), by_year=Counter(r.filed.year for r in rows), by_source=Counter(r.source for r in rows),
              verdict=Counter(r.values["判定"] or "（空欄）" for r in rows),
              empty={c: sum(1 for r in rows if not r.values[c]) for c in COLUMNS})
    st["nl"] = sum(1 for r in rows if any("\n" in v for v in r.values.values()))
    st["comma"] = sum(1 for r in rows if any("," in v for v in r.values.values()))
    st["quote"] = sum(1 for r in rows if any('"' in v for v in r.values.values()))
    st["distinct"] = {c: len({r.values[c] for r in rows}) for c in LONG_COLS}
    st["wareki_2023"] = sum(1 for r in rows if r.filed.year == 2023 and r.values["起票日"].startswith(("R", "令和")))
    st["rows_2023"] = sum(1 for r in rows if r.filed.year == 2023)
    st["wareki_2024"] = sum(1 for r in rows if r.filed.year >= 2024 and r.values["起票日"].startswith(("R", "令和")))
    st["slash"] = sum(1 for r in rows for c in DATE_COLS if re.fullmatch(r"\d{4}/\d{1,2}/\d{1,2}", r.values[c]))
    st["free_deadline"] = sum(1 for r in rows if r.values["期限"] and not re.fullmatch(r"\d{4}[-/]\d{1,2}[-/]\d{1,2}|R\d+\.\d+\.\d+|令和\d+年\d+月\d+日", r.values["期限"]))
    closed = [r for r in rows if r.done is not None and r.deadline is not None]
    st["overdue_pct"] = round(100 * sum(1 for r in closed if r.done > r.deadline) / max(1, len(closed)))
    st["zen_eq"] = sum(1 for r in rows if "－" in r.values["対象設備"])
    st["tr_nohyphen"] = sum(1 for r in rows if re.search(r"TR\d{4}-", r.values["関連トラブルNo"]))
    return st


def _write_readme(path: Path, csv_path: Path, xlsx_path: Path, st: dict, st_h1: dict, xinfo: dict) -> None:
    yr = "、".join(f"{y}年 {n:,}行" for y, n in sorted(st["by_year"].items()))
    src_label = {"incident": "トラブル起点（重大・大・中）", "claim": "品質クレーム（トラブル紐付き）", "claim_free": "品質クレーム（紐付きなし）",
                 "duplicate": "重複起票"}
    src = "\n".join(f"| {src_label[k]} | {v:,} |" for k, v in st["by_source"].most_common())
    verdict = "、".join(f"{k} {v:,}" for k, v in st["verdict"].most_common())
    empty = "\n".join(f"| {c} | {n:,} |" for c, n in st["empty"].items() if n)
    distinct = "、".join(f"{c} {n:,}" for c, n in st["distinct"].items())
    text = f"""# T5 不具合連絡・是正処置管理台帳（サンプル）

製造部 設備保全課・品質保証課が共同で管理している「是正処置（CAPA）管理台帳」を模したサンプル表データです。
トラブル報告書（`domain.standard_incidents()`）のうち重大・大・中の案件から起票した行と、後工程・顧客などからの品質クレームで構成しています。
RAG 前処理で「セル内改行・カンマ・ダブルクォートを含む長文CSV」「2段ヘッダーのExcel」を扱えるかの確認に使います。

生成: `python -m scripts.samples.t5_capa_ledger_csv`（固定シードのため再実行しても同一ファイル）

## ファイル

| ファイル | 内容 |
|---|---|
| `{csv_path.name}` | 全期間（起票日 2023-04〜{SNAPSHOT:%Y-%m}）の台帳 {st['total']:,}行（ヘッダー1行を除く） |
| `{xlsx_path.name}` | 2026年度上期（起票日 {H1_START:%Y/%m/%d}〜{H1_END:%Y/%m/%d}）分 {st_h1['total']:,}行。列構成はCSVと同じで、ヘッダーが2段 |
| `{path.name}` | このファイル |

台帳の抽出日は {SNAPSHOT:%Y-%m-%d} の想定です。完了日・判定はそれより後の日付を持ちません。

## CSV の形式

- 文字コード: **UTF-8（BOM付き）**。改行コード: **LF**（`\\n`）。区切り: カンマ。
- クォート: RFC 4180 形式。改行・カンマ・ダブルクォートを含むセルだけを `"` で囲み、セル内の `"` は `""` に二重化しています。
- セル内改行を含む行 {st['nl']:,}、ASCIIカンマを含む行 {st['comma']:,}、ダブルクォートを含む行 {st['quote']:,}。
  物理行数はレコード数よりずっと多いので、行単位の `split("\\n")` では正しく読めません。
- 読み込み例: `csv.reader(open(path, encoding="utf-8-sig", newline=""))`

## 列

| 列 | 内容 |
|---|---|
| 台帳No | `CA-YYYY-NNNN`。起票日の年ごとに起票日順の連番 |
| 起票日 | 2023年起票は和暦（`R5.4.12` が基本。`R05.04.12`・`令和5年4月12日` も混在）、2024年以降は ISO（`2024-04-12`） |
| 起票部署 | 起票者の所属。`設備保全課 保全1係`／`製造部 設備保全課`／`保全1係`／`品証` などの表記ゆれあり |
| 起票者 | 人マスタ（`domain.people()`）の人物。`高橋 誠`／`高橋`／`高橋誠`／全角空白／`（役職）`付きが混在 |
| 対象設備 | 設備ID（`CMP-103` など）単独、または `ID 名称`／`ID（名称）`／`名称（ID）`。クレームでは `該当なし`・`調査中`・`（推定）` もあり |
| 関連トラブルNo | トラブル報告書の番号（`TR-2024-00507`）。再発案件は前回番号を ` / `・`、`・改行で併記。紐付きなしのクレームは空欄 |
| 不具合内容 | 発生日時・設備・現象・アラーム（`コード "メッセージ"`）・検知・影響（停止分数、ロット、損失額 `¥1,234,500`）・調査結果。起票者ごとに `【見出し】`／`■見出し：`／`見出し：`／見出しなしの文章体 |
| 暫定対策 | 初動・処置（番号付き複数行）・交換部品（`部品名×数, …`）・結果。クレームは出荷停止・選別などの番号付き手順 |
| 真因(なぜなぜ要約) | `なぜ1：…`の連鎖、`直接原因：…／真因：…`、1行要約など。品質系は `流出原因：…` 行が付くことがある。調査中の案件は空欄か `調査中` |
| 恒久対策 | トラブル報告書の再発防止策＋原因分類（摩耗・劣化・汚れ…）に応じた管理的対策。複数なら番号付き。クレームは `【発生対策】`／`【流出対策】` |
| 水平展開 | 同型機への展開状況。空欄・`－`・`要否検討中` あり。恒久対策に他の設備IDが出てくる行は、その設備への展開（`CVD-201、CVD-202 前倒し交換済み` など）を書く |
| 期限 | 起票日＋14〜60日（重大ほど短い）。月末締めもあり。まれに `次回PM時`・`部品入荷後`・`8月末`・`至急` の文字列 |
| 完了日 | 対策完了日。未完了は空欄。完了行の約{st['overdue_pct']}%は期限超過で完了 |
| 効果確認 | 完了後2ヶ月／3ヶ月／半年の間に、**同じ設備・同じ故障モードのトラブルが再発したか**と、同じ部位のチョコ停件数（`domain.minor_stops()`）を実データから数えて記載。「再発」はトラブル報告書で再発扱い（`recurrence=True`）かつ重大度が中以上の案件で、小の同種事象は「軽微な同種事象」として別記 |
| 判定 | `有効`／`再検討`／`未確認`（空欄もあり）。再発ありなら原則 `再検討`、確認期間が抽出日を越えるものは `未確認` |
| 承認者 | 重大は課長（品証併記あり）、大・中は係長、生産技術課の起票は主任技師、クレームは品証主任＋課長（顧客案件）。未完了行は多くが空欄 |

## 件数

- 総行数: **{st['total']:,}行**（{yr}）

| 行の由来 | 行数 |
|---|---|
{src}

- 判定: {verdict}
- 長文列の値の種類数: {distinct}

## ほかのサンプルとの整合

- `関連トラブルNo`・設備ID・人名はトラブル報告書（`domain.standard_incidents()`）と同じ値です。台帳の文面（現象・調査・処置・原因・なぜなぜ・再発防止・水平展開）はその案件の記録から転記・要約しています。
- 品質クレームの影響ロットIDは、紐付くトラブルの `lots_affected` から取っています。本文中の「設備側の是正は CA-… で管理」は同じ台帳内の実在する台帳Noです。
- 効果確認の「再発（TR-…）」は、対策完了後に実際に発生した同一設備・同一故障シナリオのトラブル番号です。
- 品質クレームの不良名（配線間ショート・Vthシフトなど）は、紐付くトラブルの現象・原因に現れる語から、設備群ごとの候補の中で矛盾しないものを選んでいます。

## 意図的なイレギュラー

1. **日付表記の混在**: 2023年起票の行は和暦（{st['rows_2023']:,}行中 {st['wareki_2023']:,}行）。2024年1〜3月は和暦で書き続けた行が残っています（2024年以降の和暦行 {st['wareki_2024']:,}行）。ISO の行にもスラッシュ表記（`2025/3/4`）が {st['slash']:,}セル混じります。行内の期限・完了日も、その行の表記に合わせています（2023年起票で完了が2024年なら `R6.1.15`）。
2. **期限の文字列**: `次回PM時`・`部品入荷後`・`8月末`（年なし）・`至急` が {st['free_deadline']:,}行。
3. **必須項目の空欄**（列ごとの空欄数。未完了のため空欄の完了日・承認者、紐付きなしクレームの関連トラブルNo、重複起票行の空欄も含みます。起票者・起票部署・対象設備・期限は約0.3〜0.8%の行でランダムに記入漏れ）:

| 列 | 空欄数 |
|---|---|
{empty}

4. **表記ゆれ**: 起票部署・起票者・承認者の書き方がばらばらです。対象設備を全角で書いた行（`ＣＭＰ－１０３` など）が {st['zen_eq']:,}行、トラブルNoのハイフン抜け（`TR2024-00507`）が {st['tr_nohyphen']:,}行あります。
5. **重複起票**: 同じトラブルを別の人がもう一度起票した行が {st['by_source'].get('duplicate', 0)}行あります。`効果確認` に「重複起票のため CA-… に統合」と書かれ、期限・判定などは空欄です。
6. **記入者の癖**（トラブル報告書から引き継ぎ）: 全角数字、半角カナ（`ｱﾗｰﾑ` など）、誤変換（`異常`→`以上`）、敬体と常体の混在。
7. **判定と文面の食い違い**: 再発1件でも「発生要因が異なる」として `有効` にした行や、再発なしでもチョコ停が減らず `再検討` にした行があります。完了済みで判定が空欄の行もあります。
8. **品質クレームの結末**: 基板メーカー起因・顧客実装起因（対象設備 `該当なし`）や、原因不明のまま監視継続の行があります。

## Excel（{xlsx_path.name}）

- シート `2026上期`。1行目が大項目、2行目が小項目の2段ヘッダーで、データは3行目から。
- 大項目の結合セル: {', '.join(sorted(xinfo['merged']))}（不具合＝台帳No〜不具合内容、対策＝暫定対策〜期限、確認＝完了日〜承認者）。
- ウィンドウ枠固定 `C3`、オートフィルタは2行目、`判定` 列に入力規則（有効,再検討,未確認）、長文列は折り返し表示。
- 日付欄は ISO 表記のものを Excel の日付型（表示形式 `yyyy/mm/dd`）で {xinfo['date_cells']:,}セル格納。文字列のままの日付欄（`2026/5/7`・`次回PM時` など）が {xinfo['text_dates']:,}セルあります。
- 対象期間の判定: 「上期」は4月始まりの年度の上期（4〜9月）と解釈し、起票日が {H1_START:%Y/%m/%d}〜{H1_END:%Y/%m/%d} の行を収録しています。台帳No・値はCSVの該当行と同一です（判定 {'、'.join(f'{k} {v}' for k, v in st_h1['verdict'].most_common())}）。
"""
    path.write_text(text, encoding="utf-8", newline="\n")


def generate(output_root: Path | None = None) -> list[Path]:
    """CSV・Excel・README を書き出し、そのパスを返す。"""
    out_dir = Path(output_root or OUTPUT_ROOT) / "tables"
    out_dir.mkdir(parents=True, exist_ok=True)
    rows = build_rows()
    csv_path = out_dir / f"{STEM}.csv"
    xlsx_path = out_dir / f"{STEM}_2026上期.xlsx"
    readme_path = out_dir / f"{STEM}_README.md"
    _write_csv(rows, csv_path)
    h1 = [row for row in rows if H1_START <= row.filed <= H1_END]
    xinfo = _write_xlsx(h1, xlsx_path)
    _write_readme(readme_path, csv_path, xlsx_path, _stats(rows), _stats(h1), xinfo)
    return [csv_path, xlsx_path, readme_path]


if __name__ == "__main__":
    t0 = _time.time()
    for p in generate():
        print(f"{p}  ({p.stat().st_size / 1024:,.0f} KB)")
    print(f"done in {_time.time() - t0:.1f}s")
