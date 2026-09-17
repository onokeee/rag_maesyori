"""F4 設備点検・保全作業報告書（Excel帳票）のサンプルを作る。

    python -m scripts.samples.f4_inspection_report     # samples/forms/F4_点検保全作業報告書/ に出力

保全課が定期点検・予防保全（PM）・事後保全（故障復旧）のたびに作る「点検・保全作業報告書」を 30 ファイル作る。
帳票の中にチェックシート（点検部位・点検項目・基準値・測定値・判定・処置）を持ち、交換部品表・所見・次回点検予定日・
特記事項・承認欄が続く。様式は2版が混在する。

- Rev.1（v1, A4縦）: チェックシートは1つの縦長の表（点検部位は縦結合、前回値の列あり）
- Rev.2（v2, A4横）: チェックシートを「A.機構部・プロセス部」「B.電装・安全・ユーティリティ」の2表に分けて左右に並べ、
                      列見出しも変更（管理値/単位/実測値/判定、規格/結果/良否/対応）。部位の繰返しは「〃」

点検項目と管理値は _bank_inspection.py の設備種別ごとの定義を使い、測定値は正常分布＋劣化傾向から決定的に作る。
設備・人・部品は domain.py と共通で、次の3種類はトラブル履歴（standard_incidents）と対応付けている。
- 事後保全: トラブルの復旧作業。故障部位の点検項目が×（測定値はトラブル記録の数値を優先）、交換部品はトラブルの使用部品
- 定期点検で異常発見: 保全巡回で検知されたトラブルと同じ日時・部位で×とし、処置欄に TR 番号を記入
- 予防保全の前倒し交換: 持病のある号機のトラブル後、同じ部位の部品を前倒しで交換
一部のファイルには2枚目のシートに「過去6回の測定値」の傾向管理表（折れ線グラフ付きのものもある）を付ける。

あわせて、抽出結果の正解データ _expected.jsonl と説明 _README.md を書き出す。
乱数は domain.rng() からのみ取り、ZIP 内のタイムスタンプも固定するため、何度実行しても同一バイトのファイルになる。
"""
from __future__ import annotations

import json
import math
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from functools import lru_cache
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break
from openpyxl.writer.excel import ExcelWriter
from PIL import Image as PILImage, ImageDraw

from . import _bank_inspection as B
from . import domain as D
from .f1_repair_report import _photo_png as _sketch_photo_png, _pil_font   # 設備スケッチ写真は F1 と同じ描画を使う

# ---------------------------------------------------------------------------
# 基本設定
# ---------------------------------------------------------------------------
FORM_DIR = "F4_点検保全作業報告書"
N_FILES = 30
FORM_NO = "様式MT-031"
FONT_NAME = "ＭＳ ゴシック"
TODAY = date(2026, 9, 14)            # 報告日の上限（サンプル作成時点）
LAST_WORK_DAY = date(2026, 8, 28)
ZIP_TIME = (2026, 9, 1, 0, 0, 0)     # xlsx(ZIP)内エントリの固定タイムスタンプ
V2_START = date(2025, 4, 1)          # Rev.2 の運用開始日

REV_INFO = {"v1": ("Rev.1", "2018.10制定"), "v2": ("Rev.2", "2025.04改訂")}
WORK_TYPES = ("定期点検", "予防保全", "事後保全")
CYCLE_NAME = {("定期点検", 1): "月例点検", ("定期点検", 3): "3ヶ月点検", ("予防保全", 3): "3ヶ月PM",
              ("予防保全", 6): "6ヶ月PM", ("予防保全", 12): "年次PM"}
CYCLE_SHORT = {1: "1ヶ月", 3: "3ヶ月", 6: "6ヶ月", 12: "12ヶ月（年次）"}
WD = "月火水木金土日"

THIN = Side(style="thin", color="000000")
MEDIUM = Side(style="medium", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
STAMP_RED = "C00000"
CHEAP_LIMIT = 300_000                 # これ以下の部品は点検時にその場で交換する（予備品あり）

SECTION_HEAD = {"設備保全課 保全1係": "M10544", "設備保全課 保全2係": "M10688", "設備保全課 施設係": "M10902"}
MANAGER_ID = "M10231"
UNIT_KEYS = ("car", "pn", "fan", "lp", "shelf")


# ---------------------------------------------------------------------------
# データ構造
# ---------------------------------------------------------------------------
@dataclass
class Row:
    """チェックシートの1行。"""
    ci: B.CI
    zone: str
    item: str
    unit_val: tuple | None = None
    state: str = "ok"          # ok / tri / ng / na
    value: float | None = None
    measured: str = ""         # 測定値の表示（単位なし）
    judge: str = "○"
    action: str = ""
    resolved: bool = True
    replaced: str = ""         # "" / "pm"（定期交換） / "early"（前倒し交換）
    link: str = ""             # "incident" / "patrol"
    blank: bool = False        # 測定値の記入漏れ
    ok_text: bool = False      # 数値の代わりに「OK」と記入
    after: str = ""
    past: list = field(default_factory=list)   # 過去6回の値（古い順、None はデータなし）

    @property
    def group(self) -> str:
        return self.ci.grp


@dataclass
class FileSpec:
    key: str
    work_type: str
    eq: D.Equipment
    work_date: date
    cycle: int                    # 1/3/6/12（事後保全は 0）
    inc: D.Incident | None = None
    link: str = ""                # "breakdown" / "patrol" / "early" / ""
    version: str = "v1"
    start: datetime | None = None
    end: datetime | None = None
    report_id: str = ""
    pm_no: str = ""
    report_date: date | None = None
    workers: list = field(default_factory=list)
    witness: D.Person | None = None
    flags: set = field(default_factory=set)
    filename: str = ""

    @property
    def kind(self) -> str:
        return D.kind_of(self.eq)

    @property
    def cycle_name(self) -> str:
        return CYCLE_NAME.get((self.work_type, self.cycle), "事後保全（復旧後点検）")


@dataclass
class Content:
    rows: list
    parts: list                   # [(Part, qty, 備考)]
    findings: str = ""
    remarks: str = ""
    overall: str = "良"
    next_date: date | None = None
    next_text: str = ""
    trend_dates: list = field(default_factory=list)
    params: dict = field(default_factory=dict)
    photos: list = field(default_factory=list)   # [(caption, png bytes)]
    ok_dash: bool = False


@dataclass
class Record:
    values: dict = field(default_factory=dict)
    labels: dict = field(default_factory=dict)
    irregularities: list = field(default_factory=list)
    images: dict = field(default_factory=dict)
    extra_sheets: dict = field(default_factory=dict)
    hidden_sheets: list = field(default_factory=list)
    hidden_columns: dict = field(default_factory=dict)
    main_sheet: str = ""


# ---------------------------------------------------------------------------
# 設備パラメータ・点検項目の組み立て
# ---------------------------------------------------------------------------
@lru_cache(maxsize=None)
def _eq_params(eid: str) -> dict:
    """設備ごとに固定の設定値（スラリー流量SV、チラー温度SV、使用ガスなど）。"""
    eq = D.equipment_by_id(eid)
    r = D.rng(f"f4-eqparam:{eid}")
    k = D.kind_of(eq)
    p: dict = {}
    if k == "CMP":
        p.update(slurry_sv=float(r.choice([150, 180, 200, 220, 250])), dresser_sv=float(r.choice([5, 6, 7])),
                 platen_rpm=float(r.choice([83, 87, 93, 101])))
        p["head_rpm"] = p["platen_rpm"] - r.choice([4, 6])
    elif k == "CVD":
        gas = {"P-SiN": ("SiH4", "NH3"), "パッシベーション": ("SiH4", "NH3"), "TEOS": ("TEOS", "O2"), "W-CVD": ("WF6", "H2"),
               "SiON": ("SiH4", "N2O")}
        g = next(v for kk, v in gas.items() if kk in eq.process)
        p.update(gas1=g[0], gas2=g[1] if g[1] != "H2" else "Ar")
    elif k == "ETC":
        gas = {"Poly": "HBr", "Contact": "C4F8", "Via": "C4F8", "Metal": "Cl2", "SiN": "CHF3", "アッシング": "O2"}
        p.update(gas1=next(v for kk, v in gas.items() if kk in eq.process), chiller_sv=float(r.choice([-10, 0, 20, 40, 60])),
                 tmp_rpm=float(r.choice([33000, 36000, 40000])), vpp_sv=float(r.randint(45, 65) * 20),
                 he_sv=float(r.choice([10, 12, 15, 20])))
    elif k == "CLN":
        chem = "SC1" if "RCA" in eq.process else ("DHF" if "HF" in eq.process else r.choice(["DHF", "APM"]))
        p.update(chem=chem, bath_sv={"SC1": 70.0, "DHF": 23.0, "APM": 60.0}[chem], conc_sv=0.50)
    elif k == "IMP":
        p.update(arc_sv=float(r.choice([3.0, 3.5, 4.0])))
    elif k == "UPW":
        cap = 150 if "150" in eq.model else 200
        p.update(ro_lo=round(cap * 0.80, 1), ro_mean=round(cap * 0.93, 1))
    elif k == "OHT":
        p.update(fleet=48 if "48" in eq.model else 40)
    elif k == "AGV":
        p.update(fleet=8 if eq.equipment_id == "AGV-811" else 12)
    return p


def _applies(ci: B.CI, eq: D.Equipment) -> bool:
    if not ci.only:
        return True
    text = f"{eq.name} {eq.model} {eq.process}"
    return any(w in text for w in ci.only.split("|"))


@lru_cache(maxsize=None)
def _kind_items_cached(eid: str) -> tuple:
    eq = D.equipment_by_id(eid)
    params = _eq_params(eid)
    return tuple(B.resolve_numeric(ci, params) for ci in B.BANK[D.kind_of(eq)] if _applies(ci, eq))


def _kind_items(eq: D.Equipment) -> list[B.CI]:
    return list(_kind_items_cached(eq.equipment_id))


def _common_items(eq: D.Equipment) -> list[B.CI]:
    return list(B.common_items(D.kind_of(eq), D.PROCESS_KINDS, D.TRANSPORT_KINDS))


@lru_cache(maxsize=None)
def _bank_subs(eid: str) -> frozenset:
    eq = D.equipment_by_id(eid)
    return frozenset(ci.sub for ci in _kind_items(eq))


def _unit_key(ci: B.CI) -> str | None:
    for k in UNIT_KEYS:
        tag = "{" + k + "}"
        if tag in ci.zone or tag in ci.item:
            return k
    return None


class _Safe(dict):
    def __missing__(self, key):
        return "{" + key + "}"


def _units(spec: FileSpec, r) -> dict:
    """号車・号機・ポートなど、点検対象として選ぶ単位。トラブル対応の場合はその号車を含める。"""
    eq, k = spec.eq, spec.kind
    p = _eq_params(eq.equipment_id)
    sym = unicodedata.normalize("NFKC", spec.inc.symptom) if spec.inc else ""
    u: dict = {"bay": r.randint(1, 24)}
    if k == "OHT":
        cars = r.sample(range(1, p["fleet"] + 1), r.choice([2, 2, 3]))
        m = re.search(r"(\d+)号車", sym)
        if m and int(m.group(1)) <= p["fleet"] and int(m.group(1)) not in cars:
            cars[0] = int(m.group(1))
        m = re.search(r"Bay(\d+)", sym)
        if m:
            u["bay"] = int(m.group(1))
        u["car"] = [str(c) for c in sorted(cars)]
    elif k == "AGV":
        cars = r.sample(range(1, p["fleet"] + 1), 2)
        m = re.search(r"AGV(\d+)号機", sym)
        if m and int(m.group(1)) <= p["fleet"] and int(m.group(1)) not in cars:
            cars[0] = int(m.group(1))
        u["car"] = [str(c) for c in sorted(cars)]
    elif k == "ROB":
        u["lp"] = [str(x) for x in sorted(r.sample(range(1, 5), 2))]
    elif k == "STK":
        u["lp"] = [str(x) for x in sorted(r.sample(range(1, 7), 2))]
        u["shelf"] = sorted(f"{r.choice('ABCD')}-{r.randint(1, 40):02d}" for _ in range(2))
    elif k in ("UPW", "VAC"):
        u["pn"] = ["1", "2"]
    elif k == "PCW":
        u["pn"] = ["1", "2"]
        u["fan"] = sorted(r.sample(["1", "2", "3"], 2))
    elif k == "EXH":
        u["fan"] = ["A", "B"]
    return u


def _expand(items: list[B.CI], units: dict) -> list[tuple]:
    """号車などの単位つき項目を、連続する項目のまとまりごとに単位の数だけ展開する。"""
    out, i = [], 0
    while i < len(items):
        k = _unit_key(items[i])
        if k is None:
            out.append((items[i], None))
            i += 1
            continue
        j = i
        while j < len(items) and _unit_key(items[j]) == k:
            j += 1
        for uv in units.get(k, ["1"]):
            for ci in items[i:j]:
                out.append((ci, (k, uv)))
        i = j
    return out


def _mk_row(ci: B.CI, uv: tuple | None, spec: FileSpec, units: dict) -> Row:
    ctx = _Safe({k: (str(int(v)) if isinstance(v, float) and v.is_integer() else v) for k, v in _eq_params(spec.eq.equipment_id).items()})
    ctx["bay"] = units.get("bay", 1)
    if uv:
        ctx[uv[0]] = uv[1]
    return Row(ci=ci, zone=ci.zone.format_map(ctx), item=ci.item.format_map(ctx), unit_val=uv)


def _select_rows(spec: FileSpec, r, units: dict) -> list[Row]:
    """周期レベルに応じて点検項目を選び、20〜40行にそろえる。"""
    eq = spec.eq
    kitems, citems = _kind_items(eq), _common_items(eq)
    level = spec.cycle if spec.cycle else 1
    inc_sub = spec.inc.subsystem if (spec.inc and spec.link in ("breakdown", "patrol")) else None

    def pick(items, lvl):
        return [ci for ci in items if ci.cyc <= lvl or (inc_sub and ci.sub == inc_sub)]

    lvl = level
    while True:
        chosen_k, chosen_c = pick(kitems, lvl), pick(citems, lvl)
        rows = [_mk_row(ci, uv, spec, units) for ci, uv in _expand(chosen_k, units)]
        rows += [_mk_row(ci, uv, spec, units) for ci, uv in _expand(chosen_c, units)]
        if len(rows) >= 20 or lvl >= 12:
            break
        lvl = {1: 3, 3: 6, 6: 12}[lvl]
    # 40行を超える場合: 周期の長い共通項目 → 周期の長い設備項目 の順に外す（トラブル部位は残す）
    if len(rows) > 40:
        def removable(rw: Row, pool: str) -> bool:
            is_common = rw.ci in citems
            return (pool == "common") == is_common and rw.ci.sub != inc_sub
        for pool in ("common", "kind"):
            for cyc in (12, 6, 3, 1):
                cands = [i for i, rw in enumerate(rows) if removable(rw, pool) and rw.ci.cyc == cyc]
                r.shuffle(cands)
                drop = set(cands[: max(0, len(rows) - 40)])
                rows = [rw for i, rw in enumerate(rows) if i not in drop]
                if len(rows) <= 40:
                    break
            if len(rows) <= 40:
                break
    # 事後保全は「復旧後点検」なので、故障部位以外の月例項目は一部だけ
    if spec.work_type == "事後保全" and len(rows) > 28:
        others = [i for i, rw in enumerate(rows) if rw.ci.sub != inc_sub and rw.ci.kind != "num"]
        r.shuffle(others)
        drop = set(others[: len(rows) - r.randint(24, 28)])
        rows = [rw for i, rw in enumerate(rows) if i not in drop]
    return rows


# ---------------------------------------------------------------------------
# 測定値・判定
# ---------------------------------------------------------------------------
def _warn(ci: B.CI) -> tuple:
    lo, hi, m = ci.lo, ci.hi, ci.mean
    if lo is not None and hi is not None:
        w = (hi - lo) * 0.15
        return lo + w, hi - w
    if hi is not None:
        return None, m + 0.72 * (hi - m)
    return m - 0.72 * (m - lo), None


def _fmt_num(ci: B.CI, v: float, nd: int | None = None) -> str:
    if ci.sci:
        mant, exp = f"{v:.1E}".split("E")
        return f"{mant}E{int(exp)}"
    nd = ci.nd if nd is None else nd
    s = f"{v:.{nd}f}"
    return s[1:] if s.startswith("-") and float(s) == 0 else s   # -0.00 は 0.00 と表記


def _judge_value(ci: B.CI, v: float) -> str:
    if (ci.lo is not None and v < ci.lo - 1e-12) or (ci.hi is not None and v > ci.hi + 1e-12):
        return "×"
    wlo, whi = _warn(ci)
    if (wlo is not None and v < wlo) or (whi is not None and v > whi):
        return "△"
    return "○"


def _side(ci: B.CI, r) -> int:
    """劣化方向（+1 上限側 / -1 下限側）。"""
    if ci.lo is not None and ci.hi is not None:
        return ci.drift if ci.drift else r.choice([1, -1])
    return 1 if ci.hi is not None else -1


def _sample(ci: B.CI, state: str, r, bias: float = 0.0) -> float:
    target = {"ok": "○", "tri": "△", "ng": "×"}[state]
    wlo, whi = _warn(ci)
    v = ci.mean
    for _ in range(80):
        if state == "ok":
            v = r.gauss(ci.mean + bias, ci.sd)
        else:
            side = _side(ci, r)
            lim = ci.hi if side > 0 else ci.lo
            w = whi if side > 0 else wlo
            if state == "tri":
                v = w + (lim - w) * r.uniform(0.12, 0.92)
            else:
                gap = max(abs(lim - ci.mean), 2 * ci.sd)
                v = lim + side * gap * r.uniform(0.06, 0.55)
        if ci.minv is not None:
            v = max(v, ci.minv)
        v = float(_fmt_num(ci, v))
        if _judge_value(ci, v) == target:
            return v
    return v


def _counter_past(ci: B.CI, cur: float, r, dates: list, installed: date, life_steps: int) -> list:
    """積算値（使用時間など）の過去6回。交換周期ごとに 0 から積み上がる。"""
    L = max(1, life_steps)
    pos = r.randint(0, L - 1) if L > 1 else 0          # 前回交換から何回目の点検か
    inc = cur / (pos + 1)
    out = []
    n = len(dates)
    for k, d in enumerate(dates):
        if d is None or d < installed + timedelta(days=20):
            out.append(None)
            continue
        back = n - k
        p = pos - back
        if p >= 0:
            v = min(inc * (p + 1) * r.uniform(0.95, 1.03), cur * 0.99)
        else:
            v = inc * ((p % L) + 1) * r.uniform(0.85, 1.12)
        if ci.hi is not None and v > ci.hi:
            v = ci.hi * r.uniform(0.85, 0.97)
        out.append(float(_fmt_num(ci, max(v, 0.0))))
    return out


def _past_values(ci: B.CI, cur: float, state: str, r, dates: list, installed: date) -> list:
    """過去6回の測定値（古い順）。劣化傾向の項目は今回値に向かって推移させる。"""
    out = []
    lim_hi = ci.hi if ci.hi is not None else None
    lim_lo = ci.lo if ci.lo is not None else None
    wlo, whi = _warn(ci)
    if ci.drift and state != "ng":
        start = ci.mean - ci.drift * ci.sd * r.uniform(0.8, 1.8)
    elif ci.drift:
        start = ci.mean - ci.drift * ci.sd * r.uniform(0.2, 1.0)
    else:
        start = None
    reset = r.randint(1, 4) if (ci.drift and ci.part and r.random() < 0.3) else None
    for k, d in enumerate(dates):
        if d is None or d < installed + timedelta(days=20):
            out.append(None)
            continue
        if start is None:
            v = r.gauss(ci.mean, ci.sd * 0.8)
            if state == "ng" and k == len(dates) - 1 and r.random() < 0.5:
                v = (v + cur) / 2        # 前回から兆候が出ていたケース
        else:
            frac = (k + 1) / (len(dates) + 1)
            v = start + (cur - start) * frac + r.gauss(0, ci.sd * 0.25)
            if state == "ng" and k < len(dates) - 1:
                v = start + (cur - start) * frac * 0.6 + r.gauss(0, ci.sd * 0.25)
            if reset is not None and k < reset:
                # 前回の部品交換より前: 1サイクル前の劣化カーブ（交換直前は基準寄り）
                v = start + (cur - start) * (k + 1 + (len(dates) - reset)) / (len(dates) + 1) + r.gauss(0, ci.sd * 0.2)
        if ci.minv is not None:
            v = max(v, ci.minv)
        # 過去値は基準外にしない（基準外なら当時処置されているはず）
        if lim_hi is not None and v > lim_hi:
            v = whi if whi is not None else lim_hi
        if lim_lo is not None and v < lim_lo:
            v = wlo if wlo is not None else lim_lo
        out.append(float(_fmt_num(ci, v)))
    return out


# ---------------------------------------------------------------------------
# トラブル記録からの測定値の読み取り
# ---------------------------------------------------------------------------
_NUM = r"(?<![A-Za-z\d.\-])(-?\d[\d,]*(?:\.\d+)?)"


def _extract_value(ci: B.CI, inc: D.Incident) -> tuple | None:
    """トラブル記録（現象・調査）に、この項目の異常値が書かれていれば (値, 表示) を返す。"""
    if ci.kind != "num" or not ci.kw or ci.sci:
        return None
    text = unicodedata.normalize("NFKC", f"{inc.symptom} {inc.investigation}")
    unit = unicodedata.normalize("NFKC", ci.unit)
    for kw in ci.kw:
        kwn = unicodedata.normalize("NFKC", kw)
        if unit:
            pat = re.escape(kwn) + r".{0,20}?" + _NUM + r"\s*" + re.escape(unit)
        else:
            pat = re.escape(kwn) + r"\s?" + _NUM
        for m in re.finditer(pat, text):
            tail = text[m.end():m.end() + 2]
            if tail in ("低下", "減") or tail.startswith("減"):
                continue
            s = m.group(1).replace(",", "")
            try:
                v = float(s)
            except ValueError:
                continue
            if _judge_value(ci, v) == "×":
                nd = len(s.split(".")[1]) if "." in s else 0
                return v, f"{v:.{nd}f}"
    return None


def _bigrams(s: str) -> set:
    s = re.sub(r"[\s（）()、。・:：\d]", "", s)
    return {s[i:i + 2] for i in range(len(s) - 1)}


def _match_row(rows: list[Row], inc: D.Incident) -> Row | None:
    """トラブルの故障部位に最も対応する点検項目を選ぶ。"""
    cands = [rw for rw in rows if rw.ci.sub == inc.subsystem]
    if not cands:
        return None
    part_names = {p.name for p, _ in inc.parts_used}
    text = f"{inc.symptom} {inc.investigation} {inc.cause}"
    tb = _bigrams(text)
    sym = unicodedata.normalize("NFKC", inc.symptom)
    veh = re.search(r"(\d+)号車|AGV(\d+)号機", sym)
    best, best_sc = None, -1.0
    for rw in cands:
        sc = 0.0
        if rw.ci.part and (rw.ci.part in part_names or (rw.ci.part.startswith("MFC (") and any(n.startswith("MFC (") for n in part_names))):
            sc += 6
        if _extract_value(rw.ci, inc):
            sc += 5
        sc += 2 * sum(1 for kw in rw.ci.kw if kw in text)
        sc += min(4, len(_bigrams(rw.item) & tb)) * 0.8
        if rw.ci.kind == "num":
            sc += 0.5
        if veh and rw.unit_val and rw.unit_val[1] == (veh.group(1) or veh.group(2)):
            sc += 3
        if sc > best_sc:
            best, best_sc = rw, sc
    return best


# ---------------------------------------------------------------------------
# 対象の選定・ファイル設計
# ---------------------------------------------------------------------------
def _is_workday(d: date) -> bool:
    return not D._is_holiday(d)


def _next_workday(d: date) -> date:
    while not _is_workday(d):
        d += timedelta(days=1)
    return d


def _add_months(d: date, n: int) -> date:
    y, m = divmod(d.month - 1 + n, 12)
    y += d.year
    m += 1
    last = (date(y + (m == 12), m % 12 + 1, 1) - timedelta(days=1)).day
    return date(y, m, min(d.day, last))


def _maint_section(eq: D.Equipment) -> str:
    if D.kind_of(eq) in D.UTILITY_KINDS:
        return "設備保全課 施設係"
    return "設備保全課 保全1係" if D.fab_of(eq) == "Fab1" else "設備保全課 保全2係"


def _base_ok(i: D.Incident) -> bool:
    return (i.completed_at is not None and i.status in ("完了", "経過観察") and "同時刻に工場内複数設備が停止" not in i.symptom
            and i.subsystem in _bank_subs(i.equipment.equipment_id) and i.completed_at.date() <= LAST_WORK_DAY)


def _select_specs() -> list[FileSpec]:
    r = D.rng("f4-inspection:select")
    incs = D.standard_incidents()
    specs: list[FileSpec] = []
    used_eq: dict[str, int] = {}
    used_kinds: dict[str, int] = {}

    def take(s: FileSpec) -> None:
        specs.append(s)
        used_eq[s.eq.equipment_id] = used_eq.get(s.eq.equipment_id, 0) + 1
        used_kinds[s.kind] = used_kinds.get(s.kind, 0) + 1

    def fresh_kind_first(pool: list[D.Incident]) -> D.Incident:
        if not pool:
            raise RuntimeError("F4: 条件に合うトラブルが見つからない")
        best = min(used_kinds.get(D.kind_of(i.equipment), 0) for i in pool)
        cands = [i for i in pool if used_kinds.get(D.kind_of(i.equipment), 0) == best]
        return r.choice(cands)

    periods = [(date(2023, 5, 1), date(2024, 6, 30)), (date(2024, 7, 1), date(2025, 6, 30)), (date(2025, 7, 1), LAST_WORK_DAY)]

    # 1) 事後保全: 部品交換を伴う中〜重大の故障（人為・外部要因は除く）
    bm_pool = [i for i in incs if _base_ok(i) and i.parts_used and i.severity in ("中", "大", "重大")
               and i.category not in ("人為", "外部要因")
               and all(p.unit_price_yen < 5_000_000 for p, _ in i.parts_used)]
    n = 0
    for pi in [0, 1, 2, 0, 1, 2, 0, 1, 2]:
        lo, hi = periods[pi]
        pool = [i for i in bm_pool if lo <= i.occurred_at.date() <= hi and i.equipment.equipment_id not in used_eq]
        # 部品が点検項目の部品と対応するものを優先
        good = [i for i in pool if any(ci.part in {p.name for p, _ in i.parts_used}
                                       for ci in _kind_items(i.equipment) if ci.sub == i.subsystem)]
        if not pool:
            pool = [i for i in bm_pool if i.equipment.equipment_id not in used_eq]
            good = []
        inc = fresh_kind_first(good or pool)
        n += 1
        take(FileSpec(key=f"BM{n}", work_type="事後保全", eq=inc.equipment, work_date=inc.response_started_at.date(),
                      cycle=0, inc=inc, link="breakdown"))

    # 2) 定期点検中に異常発見（保全巡回で検知したトラブル、日勤の平日）
    pt_pool = [i for i in incs if _base_ok(i) and i.detected_by == "保全巡回" and D.shift_of(i.occurred_at) == "日勤"
               and 9 <= i.occurred_at.hour <= 16 and D.kind_of(i.equipment) not in ("OHT",)]
    for pi in (0, 1, 2):
        lo, hi = periods[pi]
        pool = [i for i in pt_pool if lo <= i.occurred_at.date() <= hi and i.equipment.equipment_id not in used_eq]
        pool = pool or [i for i in pt_pool if i.equipment.equipment_id not in used_eq]
        inc = fresh_kind_first(pool)
        take(FileSpec(key=f"PT{pi}", work_type="定期点検", eq=inc.equipment, work_date=inc.occurred_at.date(),
                      cycle=r.choice([1, 1, 3]), inc=inc, link="patrol"))

    # 3) 予防保全での前倒し交換（持病のある号機のトラブル後）
    bad_scen = {eid: set(v[0]) for eid, v in D._BAD_UNITS.items()}
    ea_pool = []
    for i in incs:
        eid = i.equipment.equipment_id
        if eid not in bad_scen or i.scenario_id not in bad_scen[eid] or not _base_ok(i) or not i.parts_used:
            continue
        names = {p.name for p, _ in i.parts_used}
        if not any(ci.part in names and ci.sub == i.subsystem for ci in _kind_items(i.equipment)):
            continue
        ea_pool.append(i)
    used_bad: set = set()
    for pi in (0, 1, 2):
        lo, hi = periods[pi]
        pool = [i for i in ea_pool if lo <= i.occurred_at.date() <= hi and i.equipment.equipment_id not in used_bad
                and (i.completed_at.date() + timedelta(days=45)) <= LAST_WORK_DAY]
        pool = pool or [i for i in ea_pool if i.equipment.equipment_id not in used_bad
                        and (i.completed_at.date() + timedelta(days=45)) <= LAST_WORK_DAY]
        inc = r.choice(pool)
        used_bad.add(inc.equipment.equipment_id)
        d = _next_workday(inc.completed_at.date() + timedelta(days=r.randint(10, 40)))
        take(FileSpec(key=f"EA{pi}", work_type="予防保全", eq=inc.equipment, work_date=d, cycle=r.choice([3, 6]),
                      inc=inc, link="early"))

    # 4) 通常の定期点検・予防保全（未使用の設備種別を優先）
    eqs = D.equipment_master()
    plan = ["定期点検"] * 9 + ["予防保全"] * 6
    r.shuffle(plan)
    all_kinds = sorted({D.kind_of(e) for e in eqs})
    kind_w = {k: (3.0 if k in D.PROCESS_KINDS else (1.2 if k in D.TRANSPORT_KINDS else 0.8)) for k in all_kinds}
    for n, wt in enumerate(plan):
        missing = [k for k in all_kinds if used_kinds.get(k, 0) == 0]
        if missing:
            kind = r.choice(missing)
        else:
            free = [k for k in all_kinds if any(D.kind_of(e) == k and e.equipment_id not in used_eq for e in eqs)]
            kind = r.choices(free, weights=[kind_w[k] / (1 + used_kinds.get(k, 0)) for k in free])[0]
        cands = [e for e in eqs if D.kind_of(e) == kind]
        cands.sort(key=lambda e: (used_eq.get(e.equipment_id, 0), e.equipment_id))
        eq = r.choice([e for e in cands if used_eq.get(e.equipment_id, 0) == used_eq.get(cands[0].equipment_id, 0)])
        # 日付: 後半（Rev.2 期間）をやや多めに
        first = max(date(2023, 5, 1), eq.installed_date + timedelta(days=240))
        span_lo = first if r.random() < 0.45 else max(first, V2_START)
        d = span_lo + timedelta(days=r.randint(0, max(0, (LAST_WORK_DAY - span_lo).days)))
        d = _next_workday(d)
        if d > LAST_WORK_DAY:
            d = LAST_WORK_DAY
        if wt == "定期点検":
            cyc = r.choice([1, 1, 1, 3, 3])
        else:
            cyc = r.choice([3, 6, 6, 12]) if kind not in D.UTILITY_KINDS else r.choice([6, 12])
        take(FileSpec(key=f"RT{n}", work_type=wt, eq=eq, work_date=d, cycle=cyc))
    specs.sort(key=lambda s: (s.work_date, s.eq.equipment_id))
    return specs


def _pick_workers(spec: FileSpec, r) -> None:
    ppl = D.people()
    eq = spec.eq
    sect = _maint_section(eq)
    staff = [p for p in ppl if p.section == sect and p.role != "係長"]
    if spec.inc is not None and spec.link in ("breakdown", "patrol"):
        internal = [p for p in spec.inc.assignees if p.section.startswith("設備保全課")]
        workers = list(spec.inc.assignees) if spec.link == "breakdown" else internal[:1]
        if spec.link == "patrol" and r.random() < 0.6:
            workers += [p for p in r.sample(staff, 2) if p not in workers][:1]
        if not any(p.section.startswith("設備保全課") for p in workers):
            workers.insert(0, r.choice(staff))
    else:
        n = {1: r.choice([1, 2, 2]), 3: r.choice([2, 2, 3]), 6: r.choice([2, 3]), 12: r.choice([2, 3, 3])}[spec.cycle]
        workers = r.sample(staff, min(n, len(staff)))
    fe = [p for p in ppl if p.role == "FE" and p.section == eq.maker]
    if spec.work_type == "予防保全" and fe and spec.cycle >= 6 and r.random() < 0.6 and fe[0] not in workers:
        workers.append(fe[0])
    spec.workers = workers
    # 立会者: 製造課の班長（シフトの班）、露光・検査は生産技術、ユーティリティは施設係の主任
    k = spec.kind
    if k in D.UTILITY_KINDS:
        cands = [p for p in ppl if p.section == "設備保全課 施設係" and p.role == "主任" and p not in workers]
        cands = cands or [p for p in ppl if p.section == "設備保全課 施設係" and p not in workers]
        spec.witness = cands[0] if cands else None
    elif k in ("LIT", "INS") and r.random() < 0.5:
        spec.witness = r.choice([p for p in ppl if p.section == "生産技術課"])
    else:
        crew = D.crew_of(spec.start)
        spec.witness = next(p for p in ppl if p.section == f"製造課 {crew}班" and p.role == "班長")


def _plan_times(spec: FileSpec, r) -> None:
    if spec.link == "breakdown":
        inc = spec.inc
        spec.start = inc.response_started_at
        spec.end = inc.completed_at
        return
    dur_h = {1: r.uniform(1.5, 3.5), 3: r.uniform(3.0, 5.5), 6: r.uniform(5.5, 9.0), 12: r.uniform(7.5, 10.5)}[spec.cycle]
    if spec.kind in ("OHT", "AGV", "STK", "ROB") and spec.cycle >= 6:
        dur_h *= 0.8
    if spec.link == "patrol":
        occ = spec.inc.occurred_at
        start = occ - timedelta(minutes=r.randint(25, 110))
        start = start.replace(minute=(start.minute // 5) * 5)
        end = max(start + timedelta(hours=dur_h), occ + timedelta(minutes=40))
        if spec.inc.completed_at.date() == occ.date():
            end = max(end, spec.inc.completed_at + timedelta(minutes=r.randint(10, 40)))
    else:
        h = r.choice([8, 8, 9, 9, 13]) if dur_h < 6 else r.choice([8, 8, 9])
        start = datetime.combine(spec.work_date, time(h, r.choice([0, 0, 30])))
        end = start + timedelta(hours=dur_h)
        if end.hour >= 12 and start.hour < 12:
            end += timedelta(minutes=45)     # 昼休憩をはさむ
    end = end.replace(minute=(end.minute // 5) * 5, second=0, microsecond=0)
    spec.start, spec.end = start, end


def _spread(r, specs: list[FileSpec], k: int, allowed=lambda s: True) -> list[FileSpec]:
    """版をまたいで均等になるように k 件を選ぶ。"""
    pools = {v: [s for s in specs if s.version == v and allowed(s)] for v in ("v1", "v2")}
    for v in pools:
        r.shuffle(pools[v])
    out = []
    while len(out) < k and any(pools.values()):
        for v in ("v1", "v2"):
            if pools[v] and len(out) < k:
                out.append(pools[v].pop())
    return out


@lru_cache(maxsize=None)
def _file_specs() -> tuple:
    specs = _select_specs()
    r = D.rng("f4-inspection:flags")
    for s in specs:
        s.version = "v1" if s.work_date < V2_START else "v2"
        _plan_times(s, r)
        _pick_workers(s, r)
    # 改訂後も旧様式を使い続ける人（2件）
    for s in r.sample([s for s in specs if s.version == "v2"], 2):
        s.version = "v1"
        s.flags.add("old_form")

    # 番号: v1 は区分別の月内連番（IN/PM/BM）、v2 は年内通し番号（MW-）。PM指示No は T4 と同じ体系
    used = set()
    for s in specs:
        d = s.work_date
        while True:
            if s.version == "v1":
                prefix = {"定期点検": "IN", "予防保全": "PM", "事後保全": "BM"}[s.work_type]
                rid = f"{prefix}{d.year % 100:02d}{d.month:02d}-{d.day * 2 + r.randint(1, 12):03d}"
            else:
                rid = f"MW-{d.year}-{d.timetuple().tm_yday * 4 + r.randint(0, 9):04d}"
            if rid not in used:
                break
        used.add(rid)
        s.report_id = rid
        if s.work_type == "予防保全":
            s.pm_no = f"PM{d.year % 100:02d}{d.month:02d}-{r.randint(1, 60):03d}"
        s.report_date = min(TODAY, s.end.date() + timedelta(days=r.choice([0, 0, 1, 1, 2, 3, 6])))

    # イレギュラーの割り当て
    def has_history(s: FileSpec) -> bool:
        return _add_months(s.work_date, -(s.cycle or 1) * 4) >= s.eq.installed_date + timedelta(days=20)
    for s in _spread(r, specs, 12, allowed=has_history):
        s.flags.add("trend")
    for s in _spread(r, specs, 6, allowed=lambda s: "trend" in s.flags):
        s.flags.add("chart")
    for s in r.sample([s for s in specs if s.version == "v1"], 3):
        s.flags.add("hide_prev")
    for s in _spread(r, specs, 3, allowed=lambda s: "hide_prev" not in s.flags):
        s.flags.add("zenkaku")
    for s in _spread(r, specs, 3):
        s.flags.add("blank_measure")
    for s in _spread(r, specs, 3, allowed=lambda s: "zenkaku" not in s.flags):
        s.flags.add("ok_text")
    for s in _spread(r, specs, 3):
        s.flags.add("glyph")
    for s in _spread(r, specs, 3, allowed=lambda s: s.work_type != "事後保全"):
        s.flags.add("no_witness")
    for s in r.sample([s for s in specs if s.work_type == "定期点検"], 1):
        s.flags.add("no_next")
    for s in r.sample([s for s in specs if s.version == "v2"], 3):
        s.flags.add("list_sheet")
    for s in _spread(r, specs, 3, allowed=lambda s: s.work_type != "定期点検"):
        s.flags.add("na_row")
    for s in specs:
        if s.report_date >= date(2026, 8, 27):
            s.flags.add("unapproved")
    for s in r.sample([s for s in specs if "unapproved" not in s.flags], 2):
        s.flags.add("unapproved")
    for s in specs:
        if "no_witness" in s.flags:
            s.witness = None

    # ファイル名（現場でありがちな付け方のゆれ）
    names = set()
    for s in specs:
        eid, d = s.eq.equipment_id, s.work_date
        ymd = d.strftime("%Y%m%d")
        if s.version == "v1":
            cands = [f"点検報告書_{eid}_{ymd}", f"{s.report_id}_{eid}_{s.work_type}", f"{eid} {s.cycle_name} {d.month}月",
                     f"保全作業報告書_{s.report_id}"]
            if s.work_type == "事後保全":
                cands = [f"{eid}_事後保全_{s.inc.incident_id}", f"点検保全報告_{eid}_{ymd}", f"{s.report_id}_{eid}_復旧後点検"]
        else:
            cands = [f"{d:%y%m%d}_{eid}_{s.work_type}報告", f"【{s.work_type}】{eid}_{s.report_id}", f"F4_{eid}_{s.cycle_name}_{ymd}"]
            if s.work_type == "事後保全":
                cands = [f"{d:%y%m%d}_{eid}_事後保全（{s.inc.incident_id}）", f"【事後保全】{eid}_{s.report_id}"]
        name = r.choice(cands)
        while name in names:
            name += "_2"
        names.add(name)
        s.filename = name + ".xlsx"
    return tuple(specs)


# ---------------------------------------------------------------------------
# 内容の設計（チェックシート・部品・所見・特記事項）
# ---------------------------------------------------------------------------
def _part_by_name(name: str) -> D.Part:
    return D._part_by_name()[name]


def _mfc_part(eq: D.Equipment, name: str) -> str:
    """MFC はガス種に合わせた型式に置き換える。"""
    if not name.startswith("MFC ("):
        return name
    gas = _eq_params(eq.equipment_id).get("gas1", "")
    for p in D.parts_catalog():
        if p.name.startswith(f"MFC ({gas} ") and eq.category in p.applicable_categories:
            return p.name
    return "MFC (汎用 1slm)"


def _short(name: str) -> str:
    return re.sub(r"\s*[（(][^）)]*[）)]$", "", name)


def _mv(row: Row) -> str:
    """文中で使う測定値（単位つき）。"""
    if row.ci.kind == "num":
        return f"{row.measured}{row.ci.unit}"
    return row.measured


def _std_text(ci: B.CI, style: str) -> str:
    """基準値の表記。v1: 単位込み・以上/以下、v2a: 単位は別列・≧≦、v2b: 単位込み・≧≦。"""
    if ci.kind != "num":
        return ci.std or ""
    nd = ci.snd if ci.snd is not None else ci.nd
    f = lambda x: _fmt_num(ci, x, nd)   # noqa: E731
    u = "" if style == "v2a" else ci.unit
    if ci.center is not None:
        if abs(ci.center) < 1e-12:
            return f"±{f(ci.tol)}{u}以内" if style == "v1" else f"±{f(ci.tol)}{u}"
        return f"{f(ci.center)}±{f(ci.tol)}{u}"
    if ci.lo is not None and ci.hi is not None:
        return f"{f(ci.lo)}～{f(ci.hi)}{u}"
    if ci.hi is not None:
        return f"{f(ci.hi)}{u}以下" if style == "v1" else f"≦{f(ci.hi)}{u}"
    return f"{f(ci.lo)}{u}以上" if style == "v1" else f"≧{f(ci.lo)}{u}"


_GENERIC_TRI_P = ["次回PMで交換予定", "部品手配済み", "経過観察"]
_GENERIC_TRI = ["経過観察", "次回点検で再確認", "監視頻度UP"]


def _assign_states(spec: FileSpec, rows: list[Row], r) -> None:
    """×/△ の項目を決める（トラブル対応の部位、直近トラブルの部位、持病の部位を優先）。"""
    eq, inc = spec.eq, spec.inc
    # トラブルと対応付ける行
    if inc is not None and spec.link in ("breakdown", "patrol"):
        rw = _match_row(rows, inc)
        if rw is not None:
            rw.state, rw.link = "ng", ("incident" if spec.link == "breakdown" else "patrol")
    if inc is not None and spec.link == "early":
        names = {p.name for p, _ in inc.parts_used}
        cands = [rw for rw in rows if rw.ci.sub == inc.subsystem and rw.ci.part in names]
        cands = cands or [rw for rw in rows if rw.ci.sub == inc.subsystem and rw.ci.part]
        if cands:
            rw = cands[0]
            rw.state, rw.replaced = r.choice(["tri", "tri", "ok"]), "early"
    # 予防保全: 定期交換部品
    if spec.work_type == "予防保全":
        done_parts = {(rw.ci.part, rw.unit_val) for rw in rows if rw.replaced}
        for name, iv, qty, only in B.PM_PARTS.get(spec.kind, []):
            if only and not any(w in f"{eq.name} {eq.model} {eq.process}" for w in only.split("|")):
                continue
            if iv > spec.cycle:
                continue
            prob = 0.9 if spec.cycle % iv == 0 else 0.45
            if spec.kind in ("OHT", "AGV"):
                prob = 0.5
            seen_units = set()
            has_unit_rows = any(rw.ci.part == name and rw.unit_val for rw in rows)
            for rw in rows:
                if rw.ci.part != name or rw.unit_val in seen_units or (name, rw.unit_val) in done_parts:
                    continue
                if has_unit_rows and not rw.unit_val:
                    continue
                seen_units.add(rw.unit_val)
                if r.random() < prob:
                    rw.replaced = "pm"
                    if rw.state == "ok" and r.random() < 0.3 and (rw.ci.kind == "num" or rw.ci.tri):
                        rw.state = "tri"
    # 通常の×/△
    recent_subs = set()
    for i in D.standard_incidents():
        if i.equipment.equipment_id == eq.equipment_id and spec.start - timedelta(days=180) <= i.occurred_at < spec.start:
            recent_subs.add(i.subsystem)
    bad = D._BAD_UNITS.get(eq.equipment_id)
    bad_subs = set()
    if bad and spec.work_date <= bad[2]:
        bad_subs = {s["sub"] for s in D._SCEN[spec.kind] if s["id"] in bad[0]}
    age = (spec.work_date - eq.installed_date).days / 365.25

    if spec.link == "breakdown":
        nx, nt = r.choices([0, 1], [88, 12])[0], r.choices([0, 1, 2], [55, 35, 10])[0]
    elif spec.link == "patrol":
        nx, nt = 0, r.choices([0, 1, 2], [50, 35, 15])[0]
    elif spec.work_type == "予防保全":
        nx, nt = r.choices([0, 1], [72, 28])[0], r.choices([0, 1, 2, 3], [42, 33, 18, 7])[0]
    else:
        nx, nt = r.choices([0, 1, 2], [62, 30, 8])[0], r.choices([0, 1, 2, 3], [48, 30, 16, 6])[0]

    def weight(rw: Row) -> float:
        w = 1.0 + 0.03 * age
        if rw.ci.sub in recent_subs:
            w *= 2.5
        if rw.ci.sub in bad_subs:
            w *= 2.0
        if rw.ci.drift:
            w *= 1.5
        if rw.ci.grp == "B":
            w *= 0.6
        return w

    replaced_parts = {(rw.ci.part, rw.unit_val) for rw in rows if rw.replaced}
    linked_parts = {(rw.ci.part, rw.unit_val) for rw in rows if rw.link and rw.ci.part}
    for state, cnt in (("ng", nx), ("tri", nt)):
        for _ in range(cnt):
            if state == "ng":
                pool = [rw for rw in rows if rw.state == "ok" and not rw.replaced and (rw.ci.kind == "num" or rw.ci.ng)
                        and (rw.ci.part, rw.unit_val) not in replaced_parts | linked_parts]
            else:
                pool = [rw for rw in rows if rw.state == "ok" and (rw.ci.kind == "num" or rw.ci.tri)
                        and (rw.replaced or (rw.ci.part, rw.unit_val) not in replaced_parts | linked_parts)]
            if not pool:
                break
            rw = r.choices(pool, weights=[weight(x) for x in pool])[0]
            rw.state = state


def _action_for(spec: FileSpec, rw: Row, r, parts: list) -> None:
    """測定値・判定・処置欄・交換部品を決める。"""
    ci, inc = rw.ci, spec.inc
    fe = D.MAKER_SHORT.get(spec.eq.maker, spec.eq.maker)
    # --- 測定値 ---
    if ci.kind == "num":
        ext = _extract_value(ci, inc) if (rw.link and inc is not None) else None
        if ext:
            rw.value, rw.measured = ext
        else:
            bias = ci.drift * r.uniform(0.0, 1.0) * ci.sd if rw.state == "ok" else 0.0
            rw.value = _sample(ci, rw.state, r, bias)
            rw.measured = _fmt_num(ci, rw.value)
        rw.judge = _judge_value(ci, rw.value)
        if rw.replaced and ci.drift:
            # 新品交換後は劣化方向と逆側（新品相当）の値に戻る
            best = ci.mean - ci.drift * ci.sd * r.uniform(1.6, 2.4)
            if ci.minv is not None:
                best = max(best, ci.minv)
            lo_ok, hi_ok = _warn(ci)
            best = min(best, hi_ok) if hi_ok is not None else best
            best = max(best, lo_ok) if lo_ok is not None else best
            rw.after = _fmt_num(ci, best)
        else:
            rw.after = _fmt_num(ci, _sample(ci, "ok", r, -ci.drift * ci.sd * 0.8))
        if re.search(B.COUNTER_RE, rw.item) and (rw.replaced or rw.state == "ng"):
            rw.after = _fmt_num(ci, 0.0)
    else:
        pool = {"ok": ci.ok, "tri": ci.tri, "ng": ci.ng}[rw.state] or {"ok": ("良",), "tri": ("要観察",), "ng": ("異常あり",)}[rw.state]
        rw.measured = r.choice(pool)
        rw.judge = {"ok": "○", "tri": "△", "ng": "×"}[rw.state]
    unit = ci.unit

    def fill(t: str) -> str:
        return t.replace("{after}", rw.after).replace("{fe}", fe)

    part_name = _mfc_part(spec.eq, ci.part) if ci.part else None
    part = _part_by_name(part_name) if part_name else None
    qty = r.randint(*ci.qty)

    # --- 処置 ---
    if rw.replaced == "pm":
        pm = next((x for x in B.PM_PARTS.get(spec.kind, []) if x[0] == ci.part), None)
        if pm:
            qty = r.randint(*pm[2])
        if ci.kind == "num":
            rw.action = r.choice([f"定期交換（交換後 {rw.after}{unit}）", f"定期交換 → {rw.after}{unit}",
                                  f"{_short(part_name)}交換（定期）→ {rw.after}{unit}"])
        else:
            rw.action = r.choice([f"{_short(part_name)}定期交換", f"{_short(part_name)}交換（定期）"])
        parts.append((part, qty, r.choice(["定期交換", "定期", f"{spec.cycle}M交換" if spec.cycle else "定期交換"])))
        return
    if rw.replaced == "early":
        rw.action = r.choice([f"前倒し交換（{inc.incident_id}再発防止）", f"{_short(part_name)}前倒し交換（{inc.incident_id}）"])
        if ci.kind == "num":
            rw.action += f" → {rw.after}{unit}"
        parts.append((part, qty, r.choice(["前倒し交換", f"予防交換（{inc.incident_id}）"])))
        return
    if rw.state == "ok":
        if ci.ok_act and r.random() < ci.ok_p:
            rw.action = r.choice(ci.ok_act)
        return
    if rw.state == "tri":
        acts = ci.tri_act or (_GENERIC_TRI_P if ci.part else _GENERIC_TRI)
        rw.action = fill(r.choice(acts))
        return
    # --- 不良（×） ---
    if rw.link == "incident":
        names = [p.name for p, _ in inc.parts_used]
        cb = _bigrams(f"{inc.cause} {inc.investigation}")
        ranked = sorted(names, key=lambda nm: -len(_bigrams(_short(nm)) & cb))
        use = part_name if part_name in names else (ranked[0] if ranked else None)
        if use:
            base = f"{_short(use)}交換"
        else:
            base = r.choice(["調整", "清掃・調整"])
        rw.action = base + (f" → {rw.after}{unit}" if ci.kind == "num" else " → 良") + f"（{inc.incident_id}）"
        for p, q in inc.parts_used:
            parts.append((p, q, "故障交換" if p.name == use else r.choice(["同時交換", "故障交換"])))
        return
    if rw.link == "patrol":
        same_day = inc.completed_at.date() == inc.occurred_at.date()
        if same_day and inc.parts_used:
            use = inc.parts_used[0][0].name
            rw.action = f"TR起票（{inc.incident_id}）、{_short(use)}交換" + (f" → {rw.after}{unit}" if ci.kind == "num" else "")
            for p, q in inc.parts_used:
                parts.append((p, q, inc.incident_id))
        else:
            rw.action = r.choice([f"TR起票（{inc.incident_id}）→ 保全対応", f"{inc.incident_id} 起票、別途処置", f"トラブル報告（{inc.incident_id}）"])
            rw.resolved = same_day
        return
    acts = list(ci.ng_act)
    if not acts:
        if part and part.unit_price_yen <= CHEAP_LIMIT:
            acts = [f"{_short(part_name)}交換" + (" → {after}" + unit if ci.kind == "num" else "")]
        elif ci.kind == "num":
            acts = next((list(a) for pat, a in B.NG_FALLBACK if re.search(pat, rw.item)), ["調整 → {after}{unit}"])
        else:
            acts = ["清掃・調整 → 良"]
        if part and part.unit_price_yen > CHEAP_LIMIT and ci.kind == "num":
            acts = [f"{_short(part_name)}交換"]
    act = fill(r.choice(acts)).replace("{unit}", unit)
    unresolved = any(w in act for w in ("手配", "協議", "依頼", "計画", "監視"))
    if part and "交換" in act and not unresolved:
        if part.unit_price_yen > CHEAP_LIMIT:
            act = f"{_short(part_name)}交換手配（{fe}）、暫定処置で運転継続"
            unresolved = True
        else:
            parts.append((part, qty, "不良交換"))
    rw.action = act
    rw.resolved = not unresolved


def _trend_dates(spec: FileSpec, r) -> list:
    step = spec.cycle if spec.cycle else 1
    out = []
    for k in range(6, 0, -1):
        d = _add_months(spec.work_date, -step * k) + timedelta(days=r.randint(-3, 3))
        out.append(_next_workday(d))
    return out


_OPEN = {
    "定期点検": ["{cycle}を実施（点検{n}項目）。", "計画どおり{cycle}を実施した。", "チェックシートに基づき{cycle}（{n}項目）実施。",
             "{name}の{cycle}実施。"],
    "予防保全": ["{cycle}実施（計画停止 {stop}）。", "{cycle}として消耗部品の定期交換と点検を実施。", "計画PM（{cycle}）実施、部品交換{np}品目。",
             "{pm}に基づき{cycle}を実施。"],
    "事後保全": ["{inc}（{occ} 発生）の復旧作業と復旧後点検を実施。", "{occ}に発生した{sub}の不具合（{inc}）の処置および復旧後点検。",
             "{inc} 対応。{sub}の故障復旧後、チェックシートにより点検。"],
}
_SUM_NONE = ["全項目基準内。", "点検結果、異常なし。", "全項目で基準値内、異常は認められなかった。", "判定はすべて○。"]
_NG = ["{zone} {item}が{mv}（基準 {std}）で基準外。{act}。", "{item}：{mv}と基準（{std}）を外れていたため、{act}。",
       "{zone}の{item}で不良（{mv}）。{act}。"]
_NG_VIS = ["{zone}の{item}で「{mv}」を確認。{act}。", "{item}：{mv}。{act}。"]
_TRI = ["{item} {mv}（基準 {std}）。基準内だが余裕が少ないため{act}。", "{zone}の{item}が{mv}で管理値に近い。{act}。"]
_TRI_PREV = ["{item}は前回{prev}→今回{mv}と悪化傾向（基準 {std}）。{act}。", "{item} 前回{prev}、今回{mv}。{act}。"]
_TRI_VIS = ["{zone}：{item}に{mv}。{act}。", "{item}で{mv}を確認、{act}。"]
_CLOSE_PROC = ["作業後、動作確認を行い製造課へ引渡し。", "立上げ後の確認OK、生産へ引渡し（{end}）。", "設備を生産へ引渡し済み。",
               "作業後に{sib}QC確認し問題なし。"]
_CLOSE_UTIL = ["中央監視へ運転再開を連絡。", "系統切替を元に戻し通常運転へ復帰（{end}）。", "作業完了を中央監視・関係部署へ連絡。"]


def _findings(spec: FileSpec, c: Content, r) -> str:
    rows, inc, eq = c.rows, spec.inc, spec.eq
    writer = spec.workers[0]
    hb = D._writer_habit(writer.employee_id)
    lines: list[str] = []
    npart = len({p.name for p, _, _ in c.parts})
    stop_h = (spec.end - spec.start).total_seconds() / 3600
    P = _Safe(cycle=spec.cycle_name, n=len(rows), name=eq.name, stop=f"{stop_h:.1f}h", np=npart, pm=spec.pm_no or "PM計画",
              inc=inc.incident_id if inc else "", occ=f"{inc.occurred_at:%m/%d %H:%M}" if inc else "",
              sub=inc.subsystem if inc else "", end=f"{spec.end:%H:%M}", sib="")
    lines.append(r.choice(_OPEN[spec.work_type]).format_map(P))
    if spec.link == "breakdown":
        cause = re.sub(r"。.*$", "", inc.cause).strip()
        cause = re.sub(r"(と推定|と判断|と思われる|（メーカー見解も同様）)$", "", cause)
        lines.append(r.choice(["原因：{c}。", "原因は{c}。", "推定原因は{c}（詳細は設備修理報告書）。"]).format(c=cause))
    xs = [rw for rw in rows if rw.judge == "×"]
    ts = [rw for rw in rows if rw.judge == "△" and not rw.replaced]
    rep_ts = [rw for rw in rows if rw.judge == "△" and rw.replaced]
    if not xs and not ts and rep_ts:
        lines.append(r.choice(["△{n}件は今回の交換部位、その他は基準内。", "交換対象の△{n}件を除き、全項目基準内。"]).format(n=len(rep_ts)))
    elif not xs and not ts:
        lines.append(r.choice(_SUM_NONE))
    else:
        bits = ([f"×{len(xs)}件"] if xs else []) + ([f"△{len(ts)}件"] if ts else [])
        lines.append(r.choice(["判定 {b}、その他は基準内。", "{b}あり、他は良。", "点検の結果 {b}。"]).format(b="・".join(bits)))
    for rw in xs:
        std = _std_text(rw.ci, "v1")
        act = rw.action.rstrip("。")
        if rw.link == "patrol":
            t = r.choice(["{item}の異常（{mv}）を発見し、{inc}として起票、保全対応とした。",
                          "点検中に{zone}の異常を確認（{mv}）。{inc}を起票し処置。"])
        elif rw.ci.kind == "num":
            t = r.choice(_NG)
        else:
            t = r.choice(_NG_VIS)
        lines.append(t.format_map(_Safe(zone=rw.zone, item=rw.item, mv=_mv(rw), std=std, act=act, inc=inc.incident_id if inc else "")))
        if not rw.resolved:
            lines.append(r.choice(["部品入荷までは監視を強化する。", "暫定運用中、恒久処置は部品入荷後に実施。", "処置完了まで要観察。"]))
    for rw in ts:
        std = _std_text(rw.ci, "v1")
        act = rw.action.rstrip("。") or "経過観察"
        prev = next((v for v in reversed(rw.past) if v is not None), None) if rw.past else None
        if rw.ci.kind != "num":
            t = r.choice(_TRI_VIS)
        elif prev is not None and ((rw.ci.drift > 0 and prev < rw.value) or (rw.ci.drift < 0 and prev > rw.value)):
            t = r.choice(_TRI_PREV)
        else:
            t = r.choice(_TRI)
        lines.append(t.format_map(_Safe(zone=rw.zone, item=rw.item, mv=_mv(rw), std=std, act=act,
                                        prev=(_fmt_num(rw.ci, prev) + rw.ci.unit) if prev is not None else "")))
    pm_rows = [rw for rw in rows if rw.replaced == "pm"]
    pm_parts = [f"{_short(p.name)}×{q}" for p, q, note in c.parts if "定期" in note or note.endswith("M交換")]
    if pm_parts:
        lines.append(r.choice(["定期交換：{p}。", "{p}を定期交換した。", "交換部品は{p}（いずれも定期）。"]).format(p="、".join(pm_parts)))
    elif pm_rows:
        lines.append("定期交換部品は交換部品欄のとおり。")
    early = [rw for rw in rows if rw.replaced == "early"]
    for rw in early:
        lines.append(r.choice(["{inc}（{d}、{sub}）の再発防止として{part}を前倒しで交換。旧品は{cond}。",
                               "{part}を前倒し交換（{inc} の再発防止）。取外し品は{cond}。"]).format(
            inc=inc.incident_id, d=f"{inc.occurred_at.month}/{inc.occurred_at.day}", sub=inc.subsystem,
            part=_short(_mfc_part(eq, rw.ci.part)), cond=r.choice(["摩耗が進行していた", "外観上は軽微な劣化", "亀裂の兆候あり", "使用限度に近い状態"])))
    if spec.link == "breakdown" and inc.result:
        res = re.sub(r"^(処置後、|処置後の確認で)", "", inc.result).rstrip("。")
        lines.append(f"復旧確認：{res}。")
    if "trend" in spec.flags:
        tr = [rw for rw in rows if rw.ci.kind == "num" and rw.ci.drift and rw.past and any(v is not None for v in rw.past)
              and rw.state != "na" and not rw.blank and not rw.ok_text]
        if tr:
            rw = r.choice(tr)
            lines.append(r.choice(["{item}の推移は傾向管理表（別シート）参照。", "{item}は過去6回の推移をみると{d}（別シート）。"]).format(
                item=rw.item, d="劣化傾向" if rw.judge != "○" else "緩やかに変化しているが基準内"))
    close = _CLOSE_UTIL if spec.kind in D.UTILITY_KINDS else _CLOSE_PROC
    if r.random() < 0.75:
        lines.append(r.choice(close).format_map(_Safe(end=f"{spec.end:%H:%M}", sib="ダミーで")))
    lines = [re.sub(r"。。", "。", x) for x in lines if x]
    style = hb["numbering"] if r.random() < 0.7 else "・"
    if hb["polite"]:
        # 予定・計画の文を「〜しました」にしないよう、文ごとに判定して敬体にする
        lines = [x if re.search(r"次回|予定|計画|手配|監視|検討|要観察|待ち", x) else D._polite(x) for x in lines]
    text = D._numbered(lines, style) if len(lines) > 2 else "\n".join(lines)
    return D._noise(text, hb, r)


def _remarks(spec: FileSpec, c: Content, r) -> str:
    eq, inc = spec.eq, spec.inc
    fe = D.MAKER_SHORT.get(eq.maker, eq.maker)
    out = []
    if spec.link == "breakdown":
        out.append(r.choice(["{i}の詳細は設備修理報告書を参照", "本作業は{i}の復旧作業を兼ねる", "停止時間は{i}に計上"]).format(i=inc.incident_id))
    elif spec.link == "patrol":
        out.append(r.choice(["点検中に発見した不具合を{i}として起票", "{i}にて処置、詳細は修理報告書参照"]).format(i=inc.incident_id))
    elif spec.link == "early":
        rw = next((x for x in c.rows if x.replaced == "early"), None)
        if rw:
            out.append(r.choice(["{i}（{d}）の再発防止として前倒し交換", "{p}の交換周期見直しを検討中（{i}）"]).format(
                i=inc.incident_id, d=f"{inc.occurred_at:%Y/%m/%d}", p=_short(_mfc_part(eq, rw.ci.part))))
    # 前回点検以降のトラブル（domain の履歴から）
    if spec.link != "breakdown" and r.random() < 0.55:
        prev = _add_months(spec.work_date, -(spec.cycle or 1))
        trs = [i for i in D.standard_incidents() if i.equipment.equipment_id == eq.equipment_id and prev <= i.occurred_at.date() < spec.work_date
               and (inc is None or i.incident_id != inc.incident_id)]
        if trs:
            ids = "、".join(i.incident_id for i in trs[:3]) + ("ほか" if len(trs) > 3 else "")
            out.append(r.choice(["前回点検以降のトラブル：{ids}", "前回点検以降 {n}件のトラブルあり（{ids}）。関連部位を重点点検"]).format(ids=ids, n=len(trs)))
    unresolved = [rw for rw in c.rows if rw.judge == "×" and not rw.resolved and rw.link != "patrol"]
    for rw in unresolved[:1]:
        due = spec.work_date + timedelta(days=r.randint(5, 21))
        out.append(r.choice(["{item}は部品入荷後に再作業予定（入荷予定 {d}）", "{item}の処置は{fe}回答待ち"]).format(
            item=rw.item, d=f"{due.month}/{due.day}", fe=fe))
    if any(p.role == "FE" for p in spec.workers):
        out.append(r.choice(["{fe}FE立会い（保守契約内）", "{fe}FE作業費は保守契約内"]).format(fe=fe))
    pool = ["作業前にLOTO実施、作業後の解除を確認", "作業中の生産影響なし（計画停止枠内で実施）", "点検結果は保全DBへ登録済み",
            "チェックシート記入後、班長へ口頭報告済み", "作業時間には立上げ・確認待ちを含む"]
    if c.parts:
        p = r.choice(c.parts)[0]
        pool.append(f"予備品在庫：{_short(p.name)} 残{r.randint(0, 5)}個" + r.choice(["（発注済み）", "", "（補充手配中）"]))
    if spec.work_date.month in (7, 8) and spec.kind in D.UTILITY_KINDS:
        pool.append("夏季のため冷却水温度・盤内温度を重点確認")
    if r.random() < 0.6 or not out:
        out.append(r.choice(pool))
    if not out or (r.random() < 0.12 and spec.link == ""):
        return r.choice(["特になし", "なし", ""])
    sep = r.choice(["\n", "。", "\n"])
    return sep.join(out) + ("。" if sep == "。" else "")


def _overall(spec: FileSpec, rows: list[Row]) -> str:
    """チェックシートの判定から総合判定を決める（様式MT-031 の記入ルール）。

    - ×のうち処置が終わっていないもの（部品手配中・メーカー回答待ち・TR起票のみ など）があれば「要処置」
    - ×をその場で処置した（交換・調整で基準内に戻した）場合は、処置後の再発確認のため「要観察」
    - 今回交換しなかった△が残る場合、故障復旧で元トラブルが経過観察中の場合も「要観察」
    - 「良」は全項目○、または△が今回の作業で交換した部位だけの場合に限る
    """
    if any(rw.judge == "×" and not rw.resolved for rw in rows):
        return "要処置"
    if any(rw.judge == "×" for rw in rows):
        return "要観察"
    if any(rw.judge == "△" and not rw.replaced for rw in rows):
        return "要観察"
    if spec.link == "breakdown" and spec.inc.status == "経過観察":
        return "要観察"
    return "良"


def _apply_overall_slip(specs: list[FileSpec], contents: list[Content]) -> None:
    """総合判定の記入誤りを1件だけ作る（ドキュメント化したイレギュラー）。

    Rev.2 の総合判定は選択肢の印字がない自由記入セルのため、事後保全で故障部位を修理し終えた作業者が
    ルール上の「要観察」ではなく「良」と書いてしまう、というありがちな誤りを再現する。
    対象は「×がすべて修理で処置済み・△なし・元トラブルが完了」の v2 事後保全のうち最も古い1件（なければ作らない）。
    """
    cands = [(s, c) for s, c in zip(specs, contents)
             if s.version == "v2" and s.link == "breakdown" and s.inc.status != "経過観察" and c.overall == "要観察"
             and any(rw.judge == "×" for rw in c.rows) and all(rw.resolved for rw in c.rows if rw.judge == "×")
             and not any(rw.judge == "△" for rw in c.rows)]
    if not cands:
        return
    spec, c = min(cands, key=lambda sc: sc[0].work_date)
    c.overall = "良"
    spec.flags.add("overall_slip")


def _meter_png(value: str, unit: str, label: str, eid: str, shot: datetime, seed: str, size=(360, 270)) -> bytes:
    """ハンディ計測器の表示を撮った写真風の画像。"""
    r = D.rng(seed)
    W, H = size
    bg = (r.randint(170, 200), r.randint(175, 205), r.randint(180, 210))
    img = PILImage.new("RGB", (W, H), bg)
    dr = ImageDraw.Draw(img)
    for y in range(H):
        k = 1.0 - 0.18 * (y / H)
        dr.line([(0, y), (W, y)], fill=tuple(int(c * k) for c in bg))
    x0, y0 = 70 + r.randint(-8, 8), 30 + r.randint(-6, 6)
    dr.rounded_rectangle([x0, y0, x0 + 220, y0 + 215], radius=18, fill=(250, 190, 40), outline=(60, 60, 60), width=3)
    dr.rounded_rectangle([x0 + 18, y0 + 20, x0 + 202, y0 + 95], radius=6, fill=(186, 200, 170), outline=(40, 40, 40), width=2)
    f_big, f_s = _pil_font(36), _pil_font(14)
    txt = value
    tw = dr.textlength(txt, font=f_big)
    dr.text((x0 + 190 - tw, y0 + 35), txt, fill=(20, 25, 20), font=f_big)
    dr.text((x0 + 150, y0 + 78), unit[:6], fill=(20, 25, 20), font=f_s)
    for k in range(3):
        dr.ellipse([x0 + 35 + k * 55, y0 + 125, x0 + 75 + k * 55, y0 + 165], fill=(70, 70, 75), outline=(30, 30, 30))
    dr.rectangle([x0 + 40, y0 + 180, x0 + 180, y0 + 202], fill=(245, 245, 245), outline=(30, 30, 30))
    lab = label if dr.textlength(label, font=f_s) < 136 else label[:9] + "…"
    dr.text((x0 + 46, y0 + 184), lab, fill=(10, 10, 10), font=f_s)
    for _ in range(500):
        g = r.randint(90, 240)
        dr.point((r.randrange(W), r.randrange(H)), fill=(g, g, g))
    dr.text((8, 6), shot.strftime("%Y/%m/%d %H:%M"), fill=(255, 150, 40), font=f_s)
    tw = dr.textlength(eid, font=f_s)
    dr.rectangle([W - tw - 16, 4, W - 4, 24], fill=(40, 40, 40))
    dr.text((W - tw - 10, 6), eid, fill=(255, 255, 255), font=f_s)
    buf = BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _design(spec: FileSpec) -> Content:
    r = D.rng(f"f4-inspection:{spec.key}")
    units = _units(spec, r)
    rows = _select_rows(spec, r, units)
    _assign_states(spec, rows, r)
    parts: list = []
    for rw in rows:
        _action_for(spec, rw, r, parts)
    # 同じ部品の重複をまとめる（号車ごとの行が同じ部品を指す場合など）
    def note_kind(note: str) -> str:
        if "定期" in note or note.endswith("M交換"):
            return "pm"
        if "前倒し" in note or "予防交換" in note:
            return "early"
        return note
    merged: dict = {}
    for p, q, note in parts:
        key = (p.name, note_kind(note))
        if key in merged:
            if key[1] in ("pm", "early", "不良交換"):
                merged[key] = (p, merged[key][1] + q, merged[key][2])
        else:
            merged[key] = (p, q, note)
    c = Content(rows=rows, parts=list(merged.values()), params=_eq_params(spec.eq.equipment_id))

    # 過去6回の測定値
    c.trend_dates = _trend_dates(spec, r)
    step = spec.cycle or 1
    for rw in rows:
        if rw.ci.kind != "num":
            continue
        if re.search(B.COUNTER_RE, rw.item):
            pm = next((x for x in B.PM_PARTS.get(spec.kind, []) if x[0] == rw.ci.part), None)
            life = max(1, round(pm[1] / step)) if pm else 12
            rw.past = _counter_past(rw.ci, rw.value, r, c.trend_dates, spec.eq.installed_date, life)
        else:
            rw.past = _past_values(rw.ci, rw.value, rw.state, r, c.trend_dates, spec.eq.installed_date)

    # 記入のイレギュラー
    ok_rows = [rw for rw in rows if rw.state == "ok" and not rw.replaced and not rw.link]
    if "blank_measure" in spec.flags and ok_rows:
        rw = r.choice(ok_rows)
        rw.blank = True
        ok_rows.remove(rw)
    if "ok_text" in spec.flags:
        cands = [rw for rw in ok_rows if rw.ci.kind == "num" and not rw.blank]
        if cands:
            r.choice(cands).ok_text = True
    if "na_row" in spec.flags:
        cands = [rw for rw in ok_rows if rw.ci.kind == "num" and not rw.blank and not rw.ok_text
                 and any(w in rw.item for w in ("QC", "パーティクル", "均一性", "反射波", "Vpp", "流量", "電流", "圧"))]
        if cands:
            rw = r.choice(cands)
            rw.state, rw.judge = "na", "－"
            rw.action = r.choice(["停止中のため測定不可", "立上げ後に測定", "PM中につき未実施"])
    c.ok_dash = r.random() < 0.3

    c.overall = _overall(spec, rows)
    c.findings = _findings(spec, c, r)
    c.remarks = _remarks(spec, c, r)

    # 次回点検予定日
    if spec.work_type == "事後保全":
        if r.random() < 0.5:
            c.next_date = _next_workday(spec.work_date + timedelta(days=r.randint(7, 14)))
            c.next_text = "（再点検）"
        else:
            c.next_text = r.choice(["－", "定期点検に準ずる"])
    elif "no_next" not in spec.flags:
        c.next_date = _next_workday(_add_months(spec.work_date, spec.cycle))

    # 写真（×の行があるファイルの一部）
    xs = [rw for rw in rows if rw.judge == "×"]
    pr = D.rng(f"f4-inspection:photo:{spec.key}")
    if xs and pr.random() < (0.85 if spec.version == "v2" else 0.45):
        for n, rw in enumerate(xs[:2], 1):
            shot = spec.start + timedelta(minutes=30 + 25 * n)
            if rw.ci.kind == "num" and not rw.ci.sci:
                cap = f"{rw.item} {_mv(rw)}"
                png = _meter_png(rw.measured, rw.ci.unit, _short(rw.item)[:10], spec.eq.equipment_id, shot, f"f4-meter:{spec.key}:{n}",
                                 size=(480, 360))
            else:
                cap = f"{rw.zone} {rw.item}：{rw.measured}"
                png = _sketch_photo_png(spec.kind, cap, spec.eq.equipment_id, shot, f"f4-sketch:{spec.key}:{n}")
            c.photos.append((cap, png))
    return c


# ---------------------------------------------------------------------------
# 帳票書き込みヘルパー
# ---------------------------------------------------------------------------
def _disp_width(s: str) -> int:
    return sum(2 if unicodedata.east_asian_width(ch) in "FWA" else 1 for ch in str(s))


def _n_lines(text: str, width_units: float, font_pt: float = 10.0) -> int:
    """折返し後の行数の見積り（全角2・半角1で数え、フォントサイズで補正）。"""
    cap = max(4.0, width_units * 0.95 - 1)
    k = font_pt / 10.5
    return sum(max(1, math.ceil(_disp_width(p) * k / cap)) for p in str(text or "").split("\n"))


class Form:
    """1シート分の帳票を組み立てる。"""

    def __init__(self, ws, widths: dict[str, float], fill: str, rec: Record | None):
        self.ws, self.rec, self.widths = ws, rec, widths
        self.fill = PatternFill("solid", fgColor=fill)
        for col, w in widths.items():
            ws.column_dimensions[col].width = w

    def width_of(self, ref: str) -> float:
        a, b = (ref.split(":") + [ref])[:2]
        c1 = column_index_from_string("".join(ch for ch in a if ch.isalpha()))
        c2 = column_index_from_string("".join(ch for ch in b if ch.isalpha()))
        return sum(self.widths.get(get_column_letter(c), 8.43) for c in range(c1, c2 + 1))

    def put(self, ref: str, value, *, size: float = 10, bold: bool = False, color: str | None = None, fill: PatternFill | None = None,
            h: str = "left", v: str = "center", wrap: bool = True, border: Border | None = BOX, fmt: str | None = None):
        c = self.ws[ref.split(":")[0]]
        c.value = value
        c.font = Font(name=FONT_NAME, size=size, bold=bold, color=color)
        numeric = isinstance(value, (datetime, date, int, float)) and not isinstance(value, bool)
        c.alignment = Alignment(horizontal=h, vertical=v, wrap_text=wrap and not numeric, shrink_to_fit=numeric)
        if fill is not None:
            c.fill = fill
        if border is not None:
            c.border = border
        if fmt:
            c.number_format = fmt
        if ":" in ref:
            self.ws.merge_cells(ref)
        return c

    def label(self, ref: str, text: str, key: str | None = None, *, h: str = "center", size: float = 10, bold: bool = False):
        self.put(ref, text, size=size, bold=bold, fill=self.fill, h=h)
        if key and self.rec is not None:
            self.rec.labels[key] = text

    def value(self, ref: str, display: str, key: str | None = None, *, raw=None, fmt: str | None = None, h: str = "left",
              v: str = "center", size: float = 10, color: str | None = None, bold: bool = False):
        self.put(ref, raw if raw is not None else (display if display != "" else None), size=size, h=h, v=v, fmt=fmt,
                 color=color, bold=bold)
        if key and self.rec is not None:
            self.rec.values[key] = display

    def fit_block(self, ref: str, text: str, *, min_row_pt: float = 15.0, line_pt: float = 13.2):
        a, b = (ref.split(":") + [ref])[:2]
        r1 = int("".join(ch for ch in a if ch.isdigit()))
        r2 = int("".join(ch for ch in b if ch.isdigit()))
        n = r2 - r1 + 1
        total = max(_n_lines(text, self.width_of(ref)) * line_pt + 6, n * min_row_pt)
        for rr in range(r1, r2 + 1):
            self.ws.row_dimensions[rr].height = round(total / n, 1)

    def fit_row(self, row: int, cells: list[tuple[str, str]], min_pt: float = 17.0, line_pt: float = 12.6, font_pt: float = 10.0):
        lines = max([_n_lines(t, self.width_of(ref), font_pt) for ref, t in cells] + [1])
        self.ws.row_dimensions[row].height = max(min_pt, round(lines * line_pt + 4, 1))

    def height(self, row: int, pt: float):
        self.ws.row_dimensions[row].height = pt

    def image(self, png: bytes, col: str, row: int, w: int, h: int, x_off: int = 6, y_off: int = 4):
        img = XLImage(BytesIO(png))
        img.width, img.height = w, h
        marker = AnchorMarker(col=column_index_from_string(col) - 1, colOff=pixels_to_EMU(x_off), row=row - 1, rowOff=pixels_to_EMU(y_off))
        img.anchor = OneCellAnchor(_from=marker, ext=XDRPositiveSize2D(pixels_to_EMU(w), pixels_to_EMU(h)))
        self.ws.add_image(img)

    def page_setup(self, area: str, *, landscape: bool = False, title: str = "", max_pages: int = 3, split_row: int | None = None,
                   force_split: bool = False):
        """A4・横1ページ合わせの印刷設定。

        印刷倍率は環境（プリンタドライバ・DPI）で変わるため、行の高さの合計から大まかにページ数を見積もり、
        - 1.2ページ以内: 縦も1ページに縮小
        - それ以上: split_row（所見・部品などの下段ブロックの先頭）の前で改ページ。上段が1ページに収まらない場合は縦 N ページ合わせ
        """
        ws = self.ws
        ws.print_area = area
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.orientation = "landscape" if landscape else "portrait"
        ws.page_margins.left = ws.page_margins.right = 0.45
        ws.page_margins.top, ws.page_margins.bottom = 0.55, 0.6
        ws.print_options.horizontalCentered = True
        ws.sheet_view.showGridLines = False
        ws.oddFooter.center.text = "&P / &N"
        ws.oddFooter.center.size = 8
        if title:
            ws.oddFooter.left.text = title
            ws.oddFooter.left.size = 8
        c1, c2 = (x.rstrip("0123456789") for x in area.split(":"))
        r_last = int(area.split(":")[1].lstrip("ABCDEFGHIJKLMNOPQRSTUVWXYZ"))
        sheet_w_pt = sum(int(self.widths.get(get_column_letter(c), 8.43) * 7 + 5)
                         for c in range(column_index_from_string(c1), column_index_from_string(c2) + 1)
                         if not ws.column_dimensions[get_column_letter(c)].hidden) * 0.75
        paper_w, paper_h = (11.69, 8.27) if landscape else (8.27, 11.69)
        scale = min(1.0, (paper_w - 0.9) * 72 / sheet_w_pt)
        page_pt = (paper_h - 1.15) * 72 / scale

        def rows_pt(a: int, b: int) -> float:
            return sum((ws.row_dimensions[rr].height or 15.0) for rr in range(a, b + 1))

        total = rows_pt(1, r_last)
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        if total <= page_pt * 1.2 or max_pages == 1:
            ws.page_setup.fitToHeight = 1
        elif split_row and (force_split or rows_pt(1, split_row - 1) <= page_pt * 0.97):
            ws.row_breaks.append(Break(id=split_row - 1))
            ws.page_setup.fitToHeight = 0
        else:
            ws.page_setup.fitToHeight = max(2, min(max_pages, math.ceil(total / page_pt)))


# ---------------------------------------------------------------------------
# 表示用の値
# ---------------------------------------------------------------------------
def _z(spec: FileSpec, s: str) -> str:
    return D.to_zenkaku(s, digits_only=True) if "zenkaku" in spec.flags else s


def _glyph(spec: FileSpec, j: str, n: int) -> str:
    if j == "○" and "glyph" in spec.flags:
        return "◯" if n % 3 else "〇"
    return j


def _stamps(spec: FileSpec) -> tuple[str, str, str]:
    ppl = {p.employee_id: p for p in D.people()}
    internal = [p for p in spec.workers if p.section.startswith("設備保全課")]
    writer = internal[0] if internal else ppl[SECTION_HEAD[_maint_section(spec.eq)]]
    head = ppl[SECTION_HEAD.get(writer.section, SECTION_HEAD[_maint_section(spec.eq)])]
    creator = D.surname(writer)
    checker = "／" if writer.employee_id == head.employee_id else D.surname(head)
    approver = "" if "unapproved" in spec.flags else D.surname(ppl[MANAGER_ID])
    if "unapproved" in spec.flags and spec.report_date >= date(2026, 9, 7):
        checker = ""
    return approver, checker, creator


def _workers_text(spec: FileSpec, style: str) -> str:
    out = []
    for p in spec.workers:
        if p.role == "FE":
            short = D.MAKER_SHORT.get(p.section, p.section)
            out.append(f"{p.name}（{short}FE）" if style == "v1" else f"{D.surname(p)}（{short}）")
        else:
            out.append(p.name if style == "v1" else D.surname(p))
    return "、".join(out) if style == "v1" else "・".join(out)


def _witness_text(spec: FileSpec, style: str) -> str:
    w = spec.witness
    if w is None:
        return ""
    sect = w.section.replace("製造課 ", "製造").replace("設備保全課 ", "")
    if style == "v1":
        return f"{w.name}（{sect} {w.role}）"
    return f"{D.surname(w)}（{sect}）"


def _worktime_text(spec: FileSpec, style: str) -> str:
    s, e = spec.start, spec.end
    hours = (e - s).total_seconds() / 3600
    same = s.date() == e.date()
    if style == "v1":
        body = f"{s.hour}:{s.minute:02d}～{e.hour}:{e.minute:02d}" if same else f"{s.month}/{s.day} {s.hour}:{s.minute:02d}～{e.month}/{e.day} {e.hour}:{e.minute:02d}"
        n = sum(1 for p in spec.workers if p.role != "FE")
        mh = spec.inc.work_hours if spec.link == "breakdown" else round(hours * n * 4) / 4
        return _z(spec, f"{body}（工数 {mh:g}h）")
    body = f"{s:%H:%M}-{e:%H:%M}" if same else f"{s:%m/%d %H:%M}-{e:%m/%d %H:%M}"
    return _z(spec, f"{body}（{hours:.1f}h）")


def _measured_cell(spec: FileSpec, rw: Row, style: str):
    """(表示文字列, セル値, 表示形式)。style: v1 / v2a（単位別列）/ v2b（単位込み文字列）。"""
    if rw.blank:
        return "", None, None
    if rw.state == "na":
        return "－", "－", None
    if rw.ok_text:
        return "OK", "OK", None
    ci = rw.ci
    if ci.kind != "num":
        return rw.measured, rw.measured, None
    if style == "v2b":
        s = _z(spec, f"{rw.measured}{ci.unit}")
        return s, s, None
    if "zenkaku" in spec.flags or ci.sci:
        s = _z(spec, rw.measured)
        return s, s, None
    nd = len(rw.measured.split(".")[1]) if "." in rw.measured else 0
    fmt = "0" if nd == 0 else "0." + "0" * nd
    return rw.measured, rw.value, fmt


def _prev_cell(rw: Row):
    if rw.ci.kind != "num" or not rw.past or rw.past[-1] is None:
        return "", None, None
    v = rw.past[-1]
    if rw.ci.sci:
        s = _fmt_num(rw.ci, v)
        return s, s, None
    fmt = "0" if rw.ci.nd == 0 else "0." + "0" * rw.ci.nd
    return _fmt_num(rw.ci, v), v, fmt


def _date_disp(spec: FileSpec, d: date | None, style: str):
    if d is None:
        return "", None, None
    if "zenkaku" in spec.flags:
        s = _z(spec, f"{d.year}年{d.month}月{d.day}日")
        return s, s, None
    if style == "v1":
        return f"{d.year}/{d.month:02d}/{d.day:02d}", datetime(d.year, d.month, d.day), "yyyy/mm/dd"
    s = f"{d.year}/{d.month:02d}/{d.day:02d}（{WD[d.weekday()]}）"
    return s, s, None


def _check_record(spec: FileSpec, rw: Row, no: int, style: str, table: str | None, n: int) -> dict:
    disp, _, _ = _measured_cell(spec, rw, style)
    d = {"no": no, "zone": rw.zone, "item": rw.item, "standard": _std_text(rw.ci, {"v1": "v1", "v2a": "v2a", "v2b": "v2b"}[style])}
    if style == "v2a":
        d["unit"] = rw.ci.unit if rw.ci.kind == "num" else ""
    if style == "v1" and "hide_prev" not in spec.flags:
        d["previous"] = _prev_cell(rw)[0]
    d["measured"] = disp
    d["judge"] = _glyph(spec, rw.judge, n)
    d["judge_norm"] = rw.judge
    d["action"] = rw.action
    if table:
        d = {"table": table, **d}
    return d


# ---------------------------------------------------------------------------
# v1（Rev.1 / A4縦）: 縦長のチェックシート1表
# ---------------------------------------------------------------------------
def _build_v1(spec: FileSpec, c: Content, wb: Workbook, rec: Record, r):
    ws = wb.active
    ws.title = r.choice(["点検報告書", "報告書", "作業報告"])
    rec.main_sheet = ws.title
    widths = {"A": 4.5, "B": 12, "C": 23, "D": 16, "E": 8.5, "F": 9.5, "G": 5.5, "H": 17}
    f = Form(ws, widths, "DDEBF7", rec)
    head_fill = PatternFill("solid", fgColor="BDD7EE")
    eq = spec.eq
    approver, checker, creator = _stamps(spec)

    f.put("A1:C1", "製造部　設備保全課", size=9, border=None)
    for ref, text, key in (("D1", "承認", "approver"), ("E1:F1", "確認", "checker"), ("G1:H1", "作成", "creator")):
        f.label(ref, text, key, size=9)
    f.put("A2:C3", "設備点検・保全作業報告書", size=16, bold=True, h="center", border=Border(bottom=MEDIUM))
    for ref, who, key in (("D2:D3", approver, "approver"), ("E2:F3", checker, "checker"), ("G2:H3", creator, "creator")):
        f.value(ref, who, key, h="center", size=12, bold=True, color=STAMP_RED)
    f.height(1, 16)
    f.height(2, 22)
    f.height(3, 22)
    f.put("A4:H4", f"{FORM_NO}(1)　{REV_INFO['v1'][0]} {REV_INFO['v1'][1]}", size=8, h="right", border=None)
    f.height(4, 14)

    wt_box = "　".join(("■" if w == spec.work_type else "□") + w for w in WORK_TYPES)
    related = spec.inc.incident_id if (spec.inc and spec.link) else (spec.pm_no if spec.pm_no else "")
    wdisp, wraw, wfmt = _date_disp(spec, spec.work_date, "v1")
    header = [
        (("作業No.", "report_id", spec.report_id, None, None), ("作業区分", "work_type", spec.work_type, wt_box, None)),
        (("設備No.", "equipment_id", eq.equipment_id, None, None), ("設備名", "equipment_name", eq.name, None, None)),
        (("ライン・工程", "line", f"{eq.line}／{eq.process}", None, None), ("メーカー・型式", "maker", f"{eq.maker}　{eq.model}", None, None)),
        (("作業日", "work_date", wdisp, wraw, wfmt), ("作業時間", "work_time", _worktime_text(spec, "v1"), None, None)),
        (("作業者", "assignee", _workers_text(spec, "v1"), None, None), ("立会者", "witness", _witness_text(spec, "v1"), None, None)),
        (("点検周期", "inspection_cycle", CYCLE_SHORT.get(spec.cycle, "－") if spec.cycle else "－（事後保全）", None, None),
         ("関連No.", "related_no", related, None, None)),
    ]
    row = 5
    for left, right in header:
        (ll, lk, ld, lr, lf), (rl, rk, rd, rr_, rf) = left, right
        f.label(f"A{row}:B{row}", ll, lk)
        f.value(f"C{row}", ld, lk, raw=lr, fmt=lf)
        f.label(f"D{row}", rl, rk)
        f.value(f"E{row}:H{row}", rd, rk, raw=rr_, fmt=rf)
        f.fit_row(row, [(f"C{row}", ld), (f"E{row}:H{row}", rd if rr_ is None else str(rr_))], min_pt=21)
        row += 1
    f.height(row, 6)
    row += 1

    # --- チェックシート ---
    f.put(f"A{row}:D{row}", "■ 点検結果（チェックシート）", bold=True, border=None)
    f.put(f"E{row}:H{row}", "判定　○：良　△：要観察　×：不良", size=8, h="right", border=None)
    rec.labels["check_items"] = "点検結果（チェックシート）"
    f.height(row, 18)
    row += 1
    cols = [("A", "No.", "no"), ("B", "点検部位", "zone"), ("C", "点検項目", "item"), ("D", "基準値", "standard"),
            ("E", "前回値", "previous"), ("F", "測定値", "measured"), ("G", "判定", "judge"), ("H", "処置", "action")]
    for col, text, key in cols:
        f.put(f"{col}{row}", text, size=9, bold=True, fill=head_fill, h="center")
        if key != "previous" or "hide_prev" not in spec.flags:
            rec.labels[f"check_items.{key}"] = text
    f.height(row, 18)
    row += 1
    items = []
    zone_start = row
    for n, rw in enumerate(c.rows, 1):
        disp, raw, fmt = _measured_cell(spec, rw, "v1")
        pdisp, praw, pfmt = _prev_cell(rw)
        std = _std_text(rw.ci, "v1")
        act = rw.action or ("－" if c.ok_dash else "")
        f.put(f"A{row}", n, size=9, h="center")
        f.put(f"B{row}", rw.zone, size=9)
        f.put(f"C{row}", rw.item, size=9)
        f.put(f"D{row}", std, size=9, h="center")      # 基準値は様式に印刷済みの文字（全角化しない）
        f.put(f"E{row}", praw if praw is not None else None, size=9, h="right", fmt=pfmt, color="595959")
        f.put(f"F{row}", raw, size=9, h="right" if rw.ci.kind == "num" else "center", fmt=fmt,
              color="C00000" if rw.judge == "×" else None, bold=rw.judge == "×")
        f.put(f"G{row}", _glyph(spec, rw.judge, n), size=10, h="center", bold=rw.judge in ("×", "△"))
        f.put(f"H{row}", act if act else None, size=9)
        f.fit_row(row, [("B1", rw.zone), ("C1", rw.item), ("D1", std), ("F1", disp), ("H1", act)], min_pt=17, line_pt=11.8, font_pt=9)
        rec_d = _check_record(spec, rw, n, "v1", None, n)
        rec_d["action"] = act
        items.append(rec_d)
        # 点検部位: 同じ部位が続く範囲を縦結合
        nxt = c.rows[n].zone if n < len(c.rows) else None
        if nxt != rw.zone:
            if row > zone_start:
                ws.merge_cells(f"B{zone_start}:B{row}")
                ws[f"B{zone_start}"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            zone_start = row + 1
        row += 1
    rec.values["check_items"] = items
    if "hide_prev" in spec.flags:
        ws.column_dimensions["E"].hidden = True
        rec.hidden_columns[ws.title] = ["E（前回値）"]
    else:
        rec.hidden_columns = {}

    # 総合判定・次回点検予定日
    overall_box = "　".join(("■" if o == c.overall else "□") + o for o in ("良", "要観察", "要処置"))
    f.label(f"A{row}:B{row}", "総合判定", "overall_judgement")
    f.value(f"C{row}", c.overall, "overall_judgement", raw=overall_box)
    f.label(f"D{row}", "次回点検予定日", "next_inspection_date", size=9)
    ndisp, nraw, nfmt = _date_disp(spec, c.next_date, "v1")
    if c.next_date is None:
        ndisp = nraw = c.next_text
        nfmt = None
    elif c.next_text:
        ndisp = f"{ndisp}{c.next_text}"
        nraw, nfmt = ndisp, None
    f.value(f"E{row}:H{row}", ndisp, "next_inspection_date", raw=nraw if nraw != "" else None, fmt=nfmt)
    f.height(row, 22)
    row += 1
    f.height(row, 8)
    row += 1

    # --- 交換部品 ---
    lower_start = row
    f.put(f"A{row}:H{row}", "■ 交換部品", bold=True, fill=head_fill)
    rec.labels["parts"] = "交換部品"
    f.height(row, 18)
    row += 1
    for ref, text, key in ((f"A{row}", "No.", None), (f"B{row}:C{row}", "品名", "name"), (f"D{row}", "品番", "part_no"),
                           (f"E{row}", "数量", "qty"), (f"F{row}:H{row}", "備考", "remarks")):
        f.put(ref, text, size=9, fill=f.fill, h="center")
        if key:
            rec.labels[f"parts.{key}"] = text
    f.height(row, 17)
    parts = []
    for n in range(max(3, len(c.parts))):
        row += 1
        f.height(row, 18)
        if n < len(c.parts):
            p, q, note = c.parts[n]
            f.put(f"A{row}", n + 1, size=9, h="center")
            f.put(f"B{row}:C{row}", p.name, size=9)
            f.put(f"D{row}", p.part_no, size=9, h="center")
            f.put(f"E{row}", _z(spec, str(q)) if "zenkaku" in spec.flags else q, size=9, h="center")
            f.put(f"F{row}:H{row}", note, size=9)
            parts.append({"name": p.name, "part_no": p.part_no, "qty": _z(spec, str(q)), "remarks": note})
        else:
            f.put(f"A{row}", None, size=9)
            f.put(f"B{row}:C{row}", "なし" if (n == 0 and not c.parts) else None, size=9)
            f.put(f"D{row}", None, size=9)
            f.put(f"E{row}", None, size=9)
            f.put(f"F{row}:H{row}", None, size=9)
    rec.values["parts"] = parts
    row += 1

    # --- 所見・特記事項 ---
    for title, key, text, nrows in (("■ 所見", "findings", c.findings, 4), ("■ 特記事項", "remarks", c.remarks, 2)):
        f.put(f"A{row}:H{row}", title, bold=True, fill=head_fill)
        rec.labels[key] = title.lstrip("■ ")
        f.height(row, 18)
        ref = f"A{row + 1}:H{row + nrows}"
        f.value(ref, text, key, v="top")
        f.fit_block(ref, text)
        row += nrows + 1

    # --- 写真 ---
    if c.photos:
        f.put(f"A{row}:H{row}", "■ 写真（不良箇所）", bold=True, fill=head_fill)
        rec.labels["photos"] = "写真（不良箇所）"
        f.height(row, 18)
        row += 1
        hidden_e = "hide_prev" in spec.flags
        img_w = 215 if hidden_e else 260
        img_h = int(img_w * 0.75)
        img_rows = math.ceil((img_h * 0.75 + 10) / 15)
        for rr in range(row, row + img_rows):
            f.height(rr, 15)
            for cc in range(1, 9):
                ws.cell(rr, cc).border = Border(left=THIN if cc == 1 else None, right=THIN if cc == 8 else None)
        caps = []
        for n, (cap, png) in enumerate(c.photos):
            col = "A" if n == 0 else ("F" if hidden_e else "E")
            f.image(png, col, row, img_w, img_h, x_off=60 if n == 0 else 10, y_off=6)
            ref = f"A{row + img_rows}:D{row + img_rows}" if n == 0 else f"E{row + img_rows}:H{row + img_rows}"
            f.put(ref, f"写真{n + 1}：{cap}", size=8, h="center")
            caps.append(f"写真{n + 1}：{cap}")
        if len(c.photos) == 1:
            f.put(f"E{row + img_rows}:H{row + img_rows}", None, size=8)
        f.fit_row(row + img_rows, [("A1:D1", caps[0])] + ([("E1:H1", caps[1])] if len(caps) > 1 else []), min_pt=18)
        row += img_rows + 1
        rec.values["photos"] = caps
        rec.images[ws.title] = len(c.photos)

    f.page_setup(f"A1:H{row - 1}", title=f"{spec.report_id}", split_row=lower_start - 1)   # 余白行の前で改ページ
    return ws


# ---------------------------------------------------------------------------
# v2（Rev.2 / A4横）: 左右2表のチェックシート
# ---------------------------------------------------------------------------
def _balance_groups(rows: list[Row]) -> tuple[list[Row], list[Row]]:
    a = [rw for rw in rows if rw.group == "A"]
    b = [rw for rw in rows if rw.group == "B"]
    # 片側が極端に長い場合は、部位のまとまりごとに B 側へ移す（末尾の部位から）
    while len(a) - len(b) > 8:
        last_zone = a[-1].zone
        block = [rw for rw in a if rw.zone == last_zone]
        if len(block) >= len(a) - 2:
            break
        a = [rw for rw in a if rw.zone != last_zone]
        b = block + b
    while len(b) - len(a) > 8:
        first_zone = b[0].zone
        block = [rw for rw in b if rw.zone == first_zone]
        if len(block) >= len(b) - 2:
            break
        b = [rw for rw in b if rw.zone != first_zone]
        a = a + block
    return a, b


def _build_v2(spec: FileSpec, c: Content, wb: Workbook, rec: Record, r):
    ws = wb.active
    ws.title = r.choice(["点検・保全報告", "作業報告書", spec.eq.equipment_id])
    rec.main_sheet = ws.title
    widths = {"A": 1.5, "B": 4, "C": 10.5, "D": 24, "E": 10, "F": 6.5, "G": 8, "H": 5, "I": 14, "J": 1.5,
              "K": 4, "L": 10.5, "M": 24, "N": 13, "O": 9, "P": 5, "Q": 14, "R": 1.5}
    f = Form(ws, widths, "E2EFDA", rec)
    head_fill = PatternFill("solid", fgColor="C6E0B4")
    eq = spec.eq
    approver, checker, creator = _stamps(spec)

    f.put("B1:I1", "設備点検・保全作業報告書", size=16, bold=True, border=None)
    f.height(1, 26)
    for ref, text, key in (("N1", "承認", "approver"), ("O1:P1", "審査", "checker"), ("Q1", "担当", "creator")):
        f.label(ref, text, key, size=9)
    for ref, who, key in (("N2:N3", approver, "approver"), ("O2:P3", checker, "checker"), ("Q2:Q3", creator, "creator")):
        f.value(ref, who, key, h="center", size=12, bold=True, color=STAMP_RED)
    f.put("B2:I2", f"製造部 設備保全課　{FORM_NO} {REV_INFO['v2'][0]}（{REV_INFO['v2'][1]}）", size=9, border=None)
    f.put("B3:I3", "※判定　○：良好　△：要観察（基準内・傾向注意）　×：不良（処置要）　－：対象外", size=8, border=None)
    f.height(2, 18)
    f.height(3, 18)
    f.height(4, 6)

    wd, wraw, wfmt = _date_disp(spec, spec.work_date, "v2")
    rd, rraw, rfmt = _date_disp(spec, spec.report_date, "v2")
    inc = spec.inc if spec.link in ("breakdown", "patrol", "early") else None
    occ = _z(spec, f"{inc.occurred_at:%Y/%m/%d %H:%M}") if (inc and spec.link == "breakdown") else ""
    dt = _z(spec, f"{inc.downtime_min:,}分") if (inc and spec.link == "breakdown") else ""
    wt_cell = spec.work_type
    grid = [
        (5, [("B5:C5", "作業No.", "report_id", "D5", spec.report_id), ("E5", "作業区分", "work_type", "F5:I5", wt_cell),
             ("K5:L5", "作業日", "work_date", "M5", wd), ("N5", "作業時間", "work_time", "O5:Q5", _worktime_text(spec, "v2"))]),
        (6, [("B6:C6", "設備No.", "equipment_id", "D6", eq.equipment_id), ("E6", "設備名", "equipment_name", "F6:I6", eq.name),
             ("K6:L6", "ライン/工程", "line", "M6", f"{eq.line} / {eq.process}"),
             ("N6", "点検周期", "inspection_cycle", "O6:Q6", spec.cycle_name if spec.cycle else "－")]),
        (7, [("B7:C7", "作業者", "assignee", "D7:I7", _workers_text(spec, "v2")), ("K7:L7", "立会者", "witness", "M7", _witness_text(spec, "v2")),
             ("N7", "報告日", "report_date", "O7:Q7", rd)]),
        (8, [("B8:C8", "故障発生日時", "occurred_at", "D8", occ), ("E8", "停止時間", "downtime_min", "F8:I8", dt),
             ("K8:L8", "関連TR No.", "related_no", "M8", inc.incident_id if inc else "")]),
    ]
    for rownum, cells in grid:
        fits = []
        for lref, ltext, key, vref, vtext in cells:
            f.label(lref, ltext, key, size=9)
            f.value(vref, vtext, key, size=9)
            fits.append((vref, vtext))
        f.fit_row(rownum, fits, min_pt=20)
    f.put("N8:Q8", "※故障発生日時・停止時間・関連TRは事後保全時に記入", size=7, border=None)
    f.height(9, 6)

    # --- チェックシート（左右2表） ---
    a_rows, b_rows = _balance_groups(c.rows)
    f.put("B10:I10", "Ａ．機構部・プロセス部", bold=True, border=None)
    f.put("K10:Q10", "Ｂ．電装・安全・ユーティリティ", bold=True, border=None)
    rec.labels["check_items"] = "Ａ．機構部・プロセス部 / Ｂ．電装・安全・ユーティリティ"
    f.height(10, 18)
    heads_a = [("B", "No", "no"), ("C", "点検箇所", "zone"), ("D", "点検内容", "item"), ("E", "管理値", "standard"), ("F", "単位", "unit"),
               ("G", "実測値", "measured"), ("H", "判定", "judge"), ("I", "処置・備考", "action")]
    heads_b = [("K", "No", "no"), ("L", "部位", "zone"), ("M", "チェック項目", "item"), ("N", "規格", "standard"), ("O", "結果", "measured"),
               ("P", "良否", "judge"), ("Q", "対応", "action")]
    for heads, tbl in ((heads_a, "left"), (heads_b, "right")):
        for col, text, key in heads:
            f.put(f"{col}11", text, size=9, bold=True, fill=head_fill, h="center")
            rec.labels[f"check_items.{tbl}.{key}"] = text
    f.height(11, 18)
    n_lines = max(len(a_rows), len(b_rows), 12) + 1      # 様式の固定行数（空行が残る）
    items = []
    use_dv = DataValidation(type="list", formula1="'リスト'!$B$2:$B$5" if "list_sheet" in spec.flags else '"○,△,×,－"', allow_blank=True)
    ws.add_data_validation(use_dv)
    counter = 0
    for side, rows_side, style, cols in (("A", a_rows, "v2a", "BCDEFGHI"), ("B", b_rows, "v2b", "KLMNOPQ")):
        prev_zone = None
        for k in range(n_lines):
            rr = 12 + k
            refs = [f"{col}{rr}" for col in cols]
            if k >= len(rows_side):
                for ref in refs:
                    f.put(ref, None, size=9)
                continue
            rw = rows_side[k]
            counter += 1
            disp, raw, fmt = _measured_cell(spec, rw, style)
            std = _std_text(rw.ci, style)
            zone_cell = "〃" if rw.zone == prev_zone else rw.zone
            prev_zone = rw.zone
            act = rw.action or ("－" if c.ok_dash else "")
            j = _glyph(spec, rw.judge, counter)
            if side == "A":
                vals = [(k + 1, "center"), (zone_cell, "left"), (rw.item, "left"), (std, "center"),
                        (rw.ci.unit if rw.ci.kind == "num" else "", "center"), (raw, "right" if rw.ci.kind == "num" else "center"),
                        (j, "center"), (act or None, "left")]
            else:
                vals = [(k + 1, "center"), (zone_cell, "left"), (rw.item, "left"), (std, "center"), (raw, "center"),
                        (j, "center"), (act or None, "left")]
            for ref, (val, h) in zip(refs, vals):
                is_meas = (side == "A" and ref[0] == "G") or (side == "B" and ref[0] == "O")
                f.put(ref, val, size=9, h=h, fmt=fmt if is_meas else None,
                      color="C00000" if (is_meas and rw.judge == "×") else None, bold=(ref[0] in "HP" and rw.judge in ("×", "△")))
            use_dv.add(f"{'H' if side == 'A' else 'P'}{rr}")
            rd_ = _check_record(spec, rw, k + 1, style, "A" if side == "A" else "B", counter)
            rd_["zone_as_written"] = zone_cell
            rd_["action"] = act
            items.append(rd_)
        # 行の高さ（左右の表のうち高い方）
    for k in range(n_lines):
        rr = 12 + k
        cells = []
        for rows_side, cref, style in ((a_rows, ("C", "D", "E", "I", "G"), "v2a"), (b_rows, ("L", "M", "N", "Q", "O"), "v2b")):
            if k < len(rows_side):
                rw = rows_side[k]
                cells += [(f"{cref[0]}1", rw.zone), (f"{cref[1]}1", rw.item), (f"{cref[2]}1", _std_text(rw.ci, style)),
                          (f"{cref[3]}1", rw.action), (f"{cref[4]}1", _measured_cell(spec, rw, style)[0])]
        f.fit_row(rr, cells, min_pt=17, line_pt=11.8, font_pt=9)
    rec.values["check_items"] = items
    row = 12 + n_lines

    # 判定集計・総合判定
    cnt = {j: sum(1 for rw in c.rows if rw.judge == j) for j in ("○", "△", "×")}
    summary = f"○ {cnt['○']}　△ {cnt['△']}　× {cnt['×']}"
    f.label(f"B{row}:D{row}", "判定集計（Ａ＋Ｂ）", "judge_summary", size=9)
    f.value(f"E{row}:I{row}", summary, "judge_summary", h="center")
    f.label(f"K{row}:M{row}", "総合判定", "overall_judgement", size=9)
    f.value(f"N{row}:Q{row}", c.overall, "overall_judgement", h="center", bold=True)
    wt_dv = DataValidation(type="list", formula1="'リスト'!$A$2:$A$4" if "list_sheet" in spec.flags else '"定期点検,予防保全,事後保全"', allow_blank=False)
    ws.add_data_validation(wt_dv)
    wt_dv.add("F5")
    f.height(row, 20)
    row += 1
    f.height(row, 8)
    row += 1

    # --- 所見（左）と交換部品（右） ---
    lower_start = row
    f.put(f"B{row}:I{row}", "所見", bold=True, fill=head_fill)
    f.put(f"K{row}:Q{row}", "交換部品", bold=True, fill=head_fill)
    rec.labels["findings"] = "所見"
    rec.labels["parts"] = "交換部品"
    f.height(row, 18)
    row += 1
    top = row
    for ref, text, key in ((f"K{row}", "No", None), (f"L{row}:M{row}", "品名", "name"), (f"N{row}", "品番", "part_no"),
                           (f"O{row}", "数量", "qty"), (f"P{row}:Q{row}", "備考", "remarks")):
        f.put(ref, text, size=9, fill=f.fill, h="center")
        if key:
            rec.labels[f"parts.{key}"] = text
    n_part_rows = max(4, len(c.parts))
    parts = []
    for n in range(n_part_rows):
        rr = row + 1 + n
        if n < len(c.parts):
            p, q, note = c.parts[n]
            f.put(f"K{rr}", n + 1, size=9, h="center")
            f.put(f"L{rr}:M{rr}", p.name, size=9)
            f.put(f"N{rr}", p.part_no, size=9, h="center")
            f.put(f"O{rr}", _z(spec, f"{q}個"), size=9, h="center")
            f.put(f"P{rr}:Q{rr}", note, size=9)
            parts.append({"name": p.name, "part_no": p.part_no, "qty": _z(spec, f"{q}個"), "remarks": note})
        else:
            for ref in (f"K{rr}", f"L{rr}:M{rr}", f"N{rr}", f"O{rr}", f"P{rr}:Q{rr}"):
                f.put(ref, "－" if (n == 0 and not c.parts and ref.startswith("L")) else None, size=9, h="center" if ref.startswith("L") else "left")
    rec.values["parts"] = parts
    bottom = row + n_part_rows
    ref = f"B{top}:I{bottom}"
    f.value(ref, c.findings, "findings", v="top", size=9)
    need = _n_lines(c.findings, f.width_of("B1:I1")) * 12.4 + 8
    per = max(17.0, need / (bottom - top + 1))
    for rr in range(top, bottom + 1):
        f.height(rr, round(per, 1))
    row = bottom + 1
    f.height(row, 8)
    row += 1

    # --- 次回点検予定日・特記事項 ---
    f.label(f"B{row}:C{row + 1}", "次回点検予定日", "next_inspection_date", size=9)
    ndisp, _, _ = _date_disp(spec, c.next_date, "v2")
    if c.next_date is None:
        ndisp = c.next_text
    elif c.next_text:
        ndisp = f"{ndisp}{c.next_text}"
    f.value(f"D{row}:E{row + 1}", ndisp, "next_inspection_date", h="center", size=9)
    f.label(f"F{row}:G{row + 1}", "特記事項", "remarks", size=9)
    f.value(f"H{row}:Q{row + 1}", c.remarks, "remarks", v="top", size=9)
    need = max(40.0, _n_lines(c.remarks, f.width_of("H1:Q1")) * 12.4 + 8)
    f.height(row, round(need / 2, 1))
    f.height(row + 1, round(need / 2, 1))
    row += 2
    f.page_setup(f"A1:R{row}", landscape=True, title=spec.report_id, split_row=lower_start - 1, force_split=True)   # 余白行の前で改ページ

    if "list_sheet" in spec.flags:
        ls = wb.create_sheet("リスト")
        ls["A1"], ls["B1"] = "作業区分", "判定"
        for n, w in enumerate(WORK_TYPES, 2):
            ls[f"A{n}"] = w
        for n, j in enumerate(["○", "△", "×", "－"], 2):
            ls[f"B{n}"] = j
        ls.sheet_state = "hidden"
        rec.hidden_sheets.append("リスト")

    if c.photos:
        ps = wb.create_sheet("写真")
        pf = Form(ps, {"A": 2, "B": 26, "C": 26, "D": 2}, "F2F2F2", None)
        pf.put("B1:C1", f"不良箇所写真　{spec.report_id}（{eq.equipment_id} {eq.name}）", size=12, bold=True, border=None)
        pf.height(1, 24)
        caps = []
        for rr in range(3, 11):              # 画像エリア 8行×15pt = 120pt（画像 135px ≒ 101pt）
            pf.height(rr, 15)
        for n, (cap, png) in enumerate(c.photos):
            col = "B" if n == 0 else "C"
            pf.image(png, col, 3, 180, 135, x_off=4, y_off=4)
            pf.put(f"{col}11", f"No.{n + 1}　{cap}", size=9, h="center")
            caps.append(f"No.{n + 1}　{cap}")
        pf.fit_row(11, [("B1", caps[0])] + ([("C1", caps[1])] if len(caps) > 1 else []), min_pt=18)
        pf.page_setup("A1:D12", landscape=False, max_pages=1)
        rec.images["写真"] = len(c.photos)
        rec.extra_sheets["写真"] = {"captions": caps}
        rec.values["photos"] = caps
    return ws


# ---------------------------------------------------------------------------
# 傾向管理（過去6回の測定値）
# ---------------------------------------------------------------------------
def _add_trend_sheet(spec: FileSpec, c: Content, wb: Workbook, rec: Record, r) -> None:
    eq = spec.eq
    name = r.choice(["傾向管理", "過去データ", "トレンド"])
    ws = wb.create_sheet(name)
    widths = {"A": 4, "B": 12, "C": 24, "D": 13, "E": 8.5, "F": 8.5, "G": 8.5, "H": 8.5, "I": 8.5, "J": 8.5, "K": 9, "L": 5, "M": 22}
    f = Form(ws, widths, "FFF2CC", None)
    num_rows = [rw for rw in c.rows if rw.ci.kind == "num" and rw.past and any(v is not None for v in rw.past)
                and not rw.blank and not rw.ok_text and rw.state != "na"]
    prio = sorted(num_rows, key=lambda rw: (0 if rw.judge in ("×", "△") else (1 if rw.ci.drift else 2), c.rows.index(rw)))
    sel = sorted(prio[: r.randint(5, 9)], key=c.rows.index)
    step = CYCLE_SHORT.get(spec.cycle, "1ヶ月") if spec.cycle else "1ヶ月（月例点検値）"
    f.put("A1:M1", f"傾向管理表（過去6回の測定値）　{eq.equipment_id} {eq.name}", size=13, bold=True, border=None)
    f.height(1, 24)
    f.put("A2:M2", f"点検周期：{step}　／　作業No. {spec.report_id}　／　今回：{spec.work_date:%Y/%m/%d}", size=9, border=None)
    f.height(2, 16)
    head = ["No.", "点検部位", "点検項目", "基準値"] + [None] * 6 + ["今回", "判定", "傾向・コメント"]
    hfill = PatternFill("solid", fgColor="FFE699")
    for n, text in enumerate(head):
        col = get_column_letter(n + 1)
        if text is None:
            d = c.trend_dates[n - 4]
            f.put(f"{col}4", datetime(d.year, d.month, d.day), size=9, bold=True, fill=hfill, h="center", fmt="yy/mm/dd")
        else:
            f.put(f"{col}4", text, size=9, bold=True, fill=hfill, h="center")
    f.put("E3:J3", "過去6回", size=9, h="center", fill=hfill)
    f.height(4, 18)
    items = []
    red_font = "C00000"
    for n, rw in enumerate(sel, 1):
        rr = 4 + n
        std = _std_text(rw.ci, "v1")
        f.put(f"A{rr}", n, size=9, h="center")
        f.put(f"B{rr}", rw.zone, size=9)
        f.put(f"C{rr}", rw.item, size=9)
        f.put(f"D{rr}", std, size=9, h="center")
        vals = []
        fmt = "0" if rw.ci.nd == 0 else "0." + "0" * rw.ci.nd
        for k, v in enumerate(rw.past + [rw.value]):
            col = get_column_letter(5 + k)
            if v is None:
                f.put(f"{col}{rr}", "－", size=9, h="center")
                vals.append("－")
                continue
            j = _judge_value(rw.ci, v)
            cell_v = _fmt_num(rw.ci, v) if rw.ci.sci else v
            f.put(f"{col}{rr}", cell_v, size=9, h="right", fmt=None if rw.ci.sci else fmt,
                  color=red_font if j == "×" else None, fill=PatternFill("solid", fgColor="FFFF99") if j == "△" else None,
                  bold=k == 6)
            vals.append(_fmt_num(rw.ci, v))
        f.put(f"L{rr}", rw.judge, size=9, h="center")
        known = [v for v in rw.past if v is not None]
        base = sum(known[-3:]) / len(known[-3:]) if known else rw.value
        jump = (rw.value - base) / rw.ci.sd if rw.ci.sd else 0.0
        if rw.replaced:
            comment = f"今回交換（交換後 {rw.after}{rw.ci.unit}）" if rw.after else "今回交換"
        elif rw.judge == "×":
            comment = r.choice(["今回基準外", "今回急変・基準外"]) + ("（処置済）" if rw.resolved else "（未処置）")
        elif abs(jump) >= 2.0:
            comment = r.choice(["今回上昇、要観察", "前回から上昇"]) if jump > 0 else r.choice(["今回低下、要観察", "前回から低下"])
        elif len(known) >= 2 and rw.ci.drift and ((rw.value - known[0]) * rw.ci.drift) > rw.ci.sd * 0.8:
            comment = r.choice(["悪化傾向", "↑ 悪化傾向、監視継続", "劣化進行中"]) if rw.ci.drift > 0 else r.choice(["低下傾向", "↓ 低下傾向", "劣化進行中"])
        else:
            comment = r.choice(["横ばい", "変化なし", "安定", "→"])
        if any(v is None for v in rw.past):
            comment += "（設置前のデータなし）" if rw.past[0] is None and eq.installed_date > c.trend_dates[0] else ""
        f.put(f"M{rr}", comment, size=9)
        f.fit_row(rr, [("B1", rw.zone), ("C1", rw.item), ("M1", comment)], min_pt=17)
        items.append({"zone": rw.zone, "item": rw.item, "standard": std, "values": vals[:6], "current": vals[6],
                      "judge": rw.judge, "comment": comment})
    last = 4 + len(sel)
    chart_added = False
    if "chart" in spec.flags and sel:
        target = next((rw for rw in sel if rw.judge != "○" and not rw.ci.sci), None) or next((rw for rw in sel if not rw.ci.sci), None)
        if target is not None:
            idx = sel.index(target) + 5
            ch = LineChart()
            ch.title = f"{target.item}（{target.ci.unit}）" if target.ci.unit else target.item
            ch.height, ch.width = 7.0, 16.0
            data = Reference(ws, min_col=5, max_col=11, min_row=idx, max_row=idx)
            ch.add_data(data, from_rows=True, titles_from_data=False)
            cats = Reference(ws, min_col=5, max_col=11, min_row=4, max_row=4)
            ch.set_categories(cats)
            ch.legend = None
            ch.x_axis.delete = False
            ch.y_axis.delete = False
            ch.x_axis.number_format = "yy/mm/dd"
            ch.varyColors = False
            ser = ch.series[0]
            ser.smooth = False
            ser.marker.symbol = "circle"
            ser.marker.size = 6
            ser.graphicalProperties.line.solidFill = "1F4E79"
            ser.graphicalProperties.line.width = 22000
            ws.add_chart(ch, f"B{last + 2}")
            chart_added = True
            rec.extra_sheets.setdefault(name, {})["chart_item"] = target.item
    f.page_setup(f"A1:M{last + (18 if chart_added else 1)}", landscape=True, max_pages=1)
    rec.extra_sheets.setdefault(name, {}).update({
        "dates": [f"{d:%y/%m/%d}" for d in c.trend_dates] + ["今回"], "items": items, "chart": chart_added})


# ---------------------------------------------------------------------------
# ワークブック組み立て・保存
# ---------------------------------------------------------------------------
def _build_workbook(spec: FileSpec, c: Content | None = None) -> tuple[Workbook, Record, Content]:
    c = c if c is not None else _design(spec)
    r = D.rng(f"f4-inspection:render:{spec.key}")
    wb = Workbook()
    rec = Record()
    (_build_v1 if spec.version == "v1" else _build_v2)(spec, c, wb, rec, r)
    if "trend" in spec.flags:
        _add_trend_sheet(spec, c, wb, rec, r)
    if "list_sheet" in spec.flags and "リスト" in wb.sheetnames:
        wb.move_sheet("リスト", offset=len(wb.sheetnames))   # 非表示シートは末尾へ

    # イレギュラーの記録
    irr = rec.irregularities
    if spec.version == "v1":
        irr.append("点検部位が縦結合セル（同じ部位の行をまとめて結合）")
        irr.append("作業区分・総合判定がチェックボックス表記（■定期点検 など）")
    else:
        irr.append("チェックシートが左右2表（A.機構部・プロセス部 / B.電装・安全・ユーティリティ）で列見出しが異なる")
        irr.append("部位欄の繰返しは「〃」（同上）で記入")
        irr.append("判定・作業区分に入力規則（ドロップダウン）")
        irr.append("チェックシートは様式の固定行数のため末尾に空行あり")
    if "old_form" in spec.flags:
        irr.append(f"改訂後（{V2_START:%Y/%m/%d}〜）も旧様式 {REV_INFO['v1'][0]} を使用")
    if "trend" in spec.flags:
        has_chart = any(v.get("chart") for v in rec.extra_sheets.values())
        irr.append("2枚目のシートに傾向管理表（過去6回の測定値）" + ("、折れ線グラフ付き" if has_chart else ""))
    if "hide_prev" in spec.flags:
        irr.append("前回値の列（E列）が非表示")
    if "zenkaku" in spec.flags:
        irr.append("日付・作業時間・測定値（文字列）・数量が全角数字")
    if any(rw.blank for rw in c.rows):
        rw = next(rw for rw in c.rows if rw.blank)
        irr.append(f"測定値の記入漏れ（{rw.zone} {rw.item}、判定は○）")
    if any(rw.ok_text for rw in c.rows):
        rw = next(rw for rw in c.rows if rw.ok_text)
        irr.append(f"数値項目の測定値欄に「OK」と記入（{rw.item}）")
    if any(rw.state == "na" for rw in c.rows):
        rw = next(rw for rw in c.rows if rw.state == "na")
        irr.append(f"測定不可の行（{rw.item}：判定「－」）")
    if "glyph" in spec.flags:
        irr.append("判定の○が「◯」「〇」（別の文字）で記入されている")
    if not _stamps(spec)[0]:
        irr.append("承認欄が未押印")
    if spec.witness is None:
        irr.append("立会者が空欄")
    if c.next_date is None and spec.work_type != "事後保全":
        irr.append("次回点検予定日が空欄")
    if c.ok_dash:
        irr.append("処置欄の「処置なし」を「－」で記入")
    if c.photos:
        irr.append("不良箇所の写真を貼付" + ("（本紙下部）" if spec.version == "v1" else "（別シート「写真」）"))
    if "list_sheet" in spec.flags:
        irr.append("非表示シート「リスト」（入力規則のリスト元）")
    if any(rw.judge == "×" and not rw.resolved for rw in c.rows):
        irr.append("×の項目が未処置（部品手配中など）で総合判定「要処置」")
    if "overall_slip" in spec.flags:
        n_x = sum(1 for rw in c.rows if rw.judge == "×")
        irr.append(f"総合判定の記入誤り: ×{n_x}件（修理で処置済み）があるのに「良」と記入（記入ルールでは「要観察」）")

    writer = spec.workers[0]
    created = datetime.combine(spec.report_date, time(r.randint(9, 18), r.randint(0, 59)))
    wb.properties.creator = writer.name
    wb.properties.lastModifiedBy = writer.name
    wb.properties.title = f"設備点検・保全作業報告書 {spec.report_id}"
    wb.properties.created = created
    wb.properties.modified = created + timedelta(minutes=r.randint(5, 180))
    return wb, rec, c


def _save_deterministic(wb: Workbook, path: Path) -> None:
    """openpyxl の保存結果を、ZIP内タイムスタンプを固定して書き直す（再実行で同一バイトにする）。"""
    buf = BytesIO()
    archive = zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED, allowZip64=True)
    ExcelWriter(wb, archive).save()
    src = zipfile.ZipFile(BytesIO(buf.getvalue()))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for info in src.infolist():
            zi = zipfile.ZipInfo(info.filename, date_time=ZIP_TIME)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            out.writestr(zi, src.read(info.filename))


# ---------------------------------------------------------------------------
# README
# ---------------------------------------------------------------------------
def _readme(specs: list[FileSpec], recs: list[Record]) -> str:
    def files(pred) -> str:
        names = [s.filename for s, rc in zip(specs, recs) if pred(s, rc)]
        return "、".join(f"`{n}`" for n in names) if names else "（なし）"

    cnt_v = {v: sum(1 for s in specs if s.version == v) for v in ("v1", "v2")}
    cnt_w = {w: sum(1 for s in specs if s.work_type == w) for w in WORK_TYPES}
    kinds = sorted({s.kind for s in specs})
    link_lines = []
    for s in specs:
        if s.link:
            label = {"breakdown": "事後保全（故障復旧）", "patrol": "定期点検中に異常発見", "early": "予防保全で前倒し交換"}[s.link]
            link_lines.append(f"| `{s.filename}` | {label} | {s.inc.incident_id} | {s.inc.equipment.equipment_id} {s.inc.subsystem} |")
    lines = [
        "# F4 設備点検・保全作業報告書（サンプル帳票）",
        "",
        f"製造部 設備保全課の「設備点検・保全作業報告書」（{FORM_NO}）を想定した Excel 帳票サンプル {len(specs)} ファイル。",
        "定期点検・予防保全（PM）・事後保全（故障復旧）の作業ごとに1ファイル。帳票の中にチェックシート（20〜40行）を持つ。",
        "設備・人・部品・トラブル番号は `scripts/samples/domain.py` と共通。点検項目・管理値は `scripts/samples/_bank_inspection.py`。",
        "",
        "再生成: `python -m scripts.samples.f4_inspection_report`（固定シードのため何度実行しても同一バイトのファイルになる）",
        "",
        f"- 作業区分: 定期点検 {cnt_w['定期点検']} / 予防保全 {cnt_w['予防保全']} / 事後保全 {cnt_w['事後保全']}",
        f"- 設備種別: {len(kinds)} 種（{', '.join(kinds)}）",
        f"- 作業日: {min(s.work_date for s in specs):%Y/%m/%d} 〜 {max(s.work_date for s in specs):%Y/%m/%d}",
        "",
        "## 様式の版",
        "",
        "| 版 | layout_version | 制定・改訂 | ファイル数 | レイアウト |",
        "|---|---|---|---|---|",
        f"| {REV_INFO['v1'][0]} | v1 | {REV_INFO['v1'][1]} | {cnt_v['v1']} | A4縦・A〜H列。上部右に承認/確認/作成の押印欄、ヘッダは2列組み"
        "（作業No./作業区分、設備No./設備名、ライン・工程/メーカー・型式、作業日/作業時間、作業者/立会者、点検周期/関連No.）。"
        "作業区分は「■定期点検　□予防保全　□事後保全」のチェックボックス表記。"
        "チェックシートは1つの縦長の表で列は No.／点検部位／点検項目／基準値／前回値／測定値／判定／処置。点検部位は同じ部位の行を縦結合。"
        "基準値は単位込み（例: `0.30mm以上`、`150±10ml/min`、`±0.20psi以内`）。測定値は数値セル＋表示形式。"
        "表の下に総合判定（■良 □要観察 □要処置）と次回点検予定日、交換部品表（No./品名/品番/数量/備考）、所見、特記事項、写真（任意）。 |",
        f"| {REV_INFO['v2'][0]} | v2 | {REV_INFO['v2'][1]} | {cnt_v['v2']} | A4横・B〜Q列。チェックシートを左右2表に分割: "
        "左「Ａ．機構部・プロセス部」（No／点検箇所／点検内容／管理値／単位／実測値／判定／処置・備考）、"
        "右「Ｂ．電装・安全・ユーティリティ」（No／部位／チェック項目／規格／結果／良否／対応）。"
        "左表の管理値は単位なし（`≧0.30`、`≦2.0`）で単位は別列、右表の規格・結果は単位込みの文字列（`≦40℃`、`31℃`）。"
        "部位の繰返しは「〃」。様式の固定行数のため短い方の表の下に空行が残る。判定・作業区分に入力規則。"
        "ヘッダに報告日・故障発生日時・停止時間・関連TR No.（事後保全時のみ）あり。承認/審査/担当の押印欄は右上。"
        "表の下に判定集計・総合判定、所見（左）と交換部品（右）を並べ、最下段に次回点検予定日（曜日付き）と特記事項。 |",
        "",
        "## 共通の体裁",
        "",
        f"- フォント: {FONT_NAME}（タイトル16pt、本文9〜10pt）。ラベルセルは塗りつぶし（v1 淡青 / v2 淡緑）、罫線は細線",
        "- 行の高さ: 文章量（全角2・半角1で見積もり）に合わせて調整。所見・特記事項は結合セルで複数行",
        "- 印刷設定: A4（v1 縦 / v2 横）・横1ページに合わせる・印刷範囲・フッタにページ番号と作業No.。"
        "内容が1.2ページ以内なら縦も1ページに縮小、それ以上は下段ブロック（v1 交換部品 / v2 所見・交換部品）の前で改ページ"
        "（v1 でチェックシートが1ページに収まらない場合は縦2〜3ページに合わせる）。傾向管理表・写真シートは1ページ。枠線は非表示",
        "- 押印: 赤字の姓。確認者が作成者本人の場合は「／」",
        "- ブックのプロパティ（作成者・作成日時）は作業者・報告日に合わせてある",
        "",
        "## 内容の作り方（ドメインデータとの整合）",
        "",
        "- 点検項目は設備種別ごとの定義（周期: 月例/3ヶ月/6ヶ月/年次）から、作業の周期以下の項目を選ぶ。"
        "OHT・AGV は号車/号機、ポンプ・ファンは号機、ソーター/ストッカはポート・棚ごとに展開",
        "- 管理値は domain のトラブルシナリオの通常値・異常値と整合（例: ドライポンプ電流 上限9.0A、OHT走行ホイール径 98.5mm以上）",
        "- 測定値は正常分布から生成し、劣化傾向のある項目（摩耗・目詰まり等）は使用に伴い基準側へ寄る。"
        "×/△の件数は作業ごとに抽選し、直近180日にトラブルのあった部位・持病のある号機の部位ほど出やすい",
        "- 予防保全は種別ごとの定期交換部品（T4 の PM 計画と同じ部品・周期）を交換し、処置欄に「定期交換 → 交換後の値」",
        "- 点検時に部品交換で処置した×は交換部品表に載る。30万円を超える部品は「交換手配」とし未処置（総合判定「要処置」）",
        "- 総合判定はチェックシートの判定から決める: 未処置の×（手配中・回答待ち・TR起票のみ）があれば「要処置」、"
        "×をその場で処置した場合・交換しなかった△が残る場合・故障復旧で元トラブルが経過観察中の場合は「要観察」、"
        "「良」は全項目○か、△が今回交換した部位だけの場合",
        "- 前回値・傾向管理表の過去値は今回値に向かう推移として生成（過去値は基準外にしない。設備設置前は「－」）。"
        "使用時間・積算時間・サイクル数は部品の交換周期ごとに0から積み上がる（のこぎり波）。交換後の値は新品相当",
        "",
        "トラブル履歴と対応付けたファイル:",
        "",
        "| ファイル | 種類 | TR No. | 設備・部位 |",
        "|---|---|---|---|",
        *link_lines,
        "",
        "- 事後保全: 作業時間はトラブルの対応開始〜復旧、作業者はトラブルの担当者、交換部品はトラブルの使用部品。"
        "故障部位に対応する点検項目を×とし、トラブル記録（現象・調査）に同じ項目の数値があればその値を測定値にする",
        "- 定期点検中に異常発見: 保全巡回で検知されたトラブルと同じ日の点検。発生時刻をはさむ作業時間とし、×の処置欄に TR 番号",
        "- 前倒し交換: 持病のある号機（domain の _BAD_UNITS）のトラブル10〜40日後の PM で、同じ部位の部品を前倒し交換",
        "- 特記事項の「前回点検以降のトラブル：TR-…」は、前回点検日〜今回の間の同設備のトラブル（domain の履歴）",
        "",
        "## 意図的なイレギュラー",
        "",
        f"- 傾向管理表（2枚目のシート、過去6回＋今回の測定値、基準外は赤字・△は黄色）: {files(lambda s, rc: 'trend' in s.flags)}",
        f"  - うち折れ線グラフ付き: {files(lambda s, rc: any(v.get('chart') for v in rc.extra_sheets.values()))}",
        f"- v1 で前回値の列（E列）が非表示: {files(lambda s, rc: 'hide_prev' in s.flags)}",
        f"- 日付・作業時間・測定値・数量が全角数字の文字列（基準値は様式の印字なので半角のまま）: {files(lambda s, rc: 'zenkaku' in s.flags)}",
        f"- 測定値の記入漏れ（判定は○）: {files(lambda s, rc: any('記入漏れ' in x for x in rc.irregularities))}",
        f"- 数値項目の測定値欄に「OK」: {files(lambda s, rc: any('「OK」' in x for x in rc.irregularities))}",
        f"- 測定不可で判定「－」の行: {files(lambda s, rc: any('測定不可' in x for x in rc.irregularities))}",
        f"- 判定の○が別の文字（◯ U+25EF / 〇 U+3007）: {files(lambda s, rc: 'glyph' in s.flags)}",
        f"- 承認欄が未押印: {files(lambda s, rc: '承認欄が未押印' in rc.irregularities)}",
        f"- 立会者が空欄: {files(lambda s, rc: s.witness is None)}",
        f"- 次回点検予定日が空欄（定期点検）: {files(lambda s, rc: '次回点検予定日が空欄' in rc.irregularities)}",
        f"- 処置なしを「－」で記入: {files(lambda s, rc: any('「－」で記入' in x for x in rc.irregularities))}",
        f"- 不良箇所の写真（計測器表示・設備スケッチの PNG）: {files(lambda s, rc: bool(rc.images))}",
        f"- 非表示シート「リスト」（入力規則のリスト元）: {files(lambda s, rc: 'list_sheet' in s.flags)}",
        f"- ×が未処置で総合判定「要処置」: {files(lambda s, rc: any('未処置' in x for x in rc.irregularities))}",
        f"- 総合判定の記入誤り（Rev.2 の自由記入欄に、修理済みの×があるのに「良」と記入。記入ルールでは「要観察」。この1件だけ）: "
        f"{files(lambda s, rc: 'overall_slip' in s.flags)}",
        f"- 改訂後も旧様式（Rev.1）を使用: {files(lambda s, rc: 'old_form' in s.flags)}",
        "- 事後保全の次回点検予定日は「－」「定期点検に準ずる」または再点検日（`2025/06/20（再点検）` など）",
        "- 所見・特記事項の書き方は記入者ごとの癖（番号の振り方、敬体、全角数字、半角カナ、誤変換）が domain.py 由来で混ざる",
        "- ファイル名の付け方がばらばら（作業No.入り／日付入り／【作業区分】付き／TR番号入り）",
        "",
        "## _expected.jsonl",
        "",
        "1行1ファイルの JSON（UTF-8）。",
        "",
        "```",
        '{"file", "layout_version": "v1"|"v2", "form_revision", "main_sheet", "work_type", "incident_id"（対応付けたTR、なければ null）,',
        ' "values": {キー: 人が読むとおりの値}, "labels_used": {キー: そのファイルで実際に使われているラベル文字列},',
        ' "sheets", "hidden_sheets", "hidden_columns", "images": {シート名: 枚数}, "extra_sheets": {...}, "irregularities": [...],',
        ' "counts": {"check_items", "ng", "watch"}}',
        "```",
        "",
        "values の主なキー: report_id, work_type, equipment_id, equipment_name, line, maker(v1), work_date, work_time, assignee, witness, "
        "inspection_cycle, related_no, report_date(v2), occurred_at(v2), downtime_min(v2), check_items, judge_summary(v2), "
        "overall_judgement, next_inspection_date, parts, findings, remarks, photos, approver, checker, creator",
        "",
        "- `check_items` は行ごとの dict のリスト: no, zone（〃・縦結合を解決した部位名）, item, standard, measured, judge（記入どおりの文字）, "
        "judge_norm（○/△/×/－ に正規化）, action。v1 は previous（前回値列が表示されている場合のみ）、v2 は table（A/B）, unit（A表のみ）, "
        "zone_as_written（〃のまま）",
        "- labels_used の列見出しは v1 が `check_items.<列>`、v2 が `check_items.left.<列>` / `check_items.right.<列>`、部品表は `parts.<列>`",
        "- 空欄は空文字 `\"\"`。チェックボックス表記は選択された値（例: `予防保全`、`要観察`）で記録",
        "- 日付型セルは Excel での表示（例: `2024/05/12`）で記録。測定値の数値セルは表示形式どおりの桁（例: `0.52`、`-78`）",
        "- extra_sheets の傾向管理表: dates（6回分＋今回）、items（部位・項目・基準値・過去6回の値・今回値・判定・コメント）、chart",
        "",
    ]
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# エントリポイント
# ---------------------------------------------------------------------------
def generate(output_root: Path | None = None) -> list[Path]:
    root = Path(output_root) if output_root is not None else D.OUTPUT_ROOT
    out_dir = root / "forms" / FORM_DIR
    out_dir.mkdir(parents=True, exist_ok=True)
    for old in out_dir.glob("*.xlsx"):   # このフォルダは本スクリプト専用。名前の変わった古い生成物を残さない
        old.unlink()

    specs = list(_file_specs())
    # 先に全ファイルの中身を決め、総合判定の記入誤り（1件）を選んでから描画する
    designs = [_design(spec) for spec in specs]
    _apply_overall_slip(specs, designs)
    recs, paths, sheets_list, contents = [], [], [], []
    for spec, design in zip(specs, designs):
        wb, rec, c = _build_workbook(spec, design)
        path = out_dir / spec.filename
        _save_deterministic(wb, path)
        recs.append(rec)
        sheets_list.append(wb.sheetnames)
        contents.append(c)
        paths.append(path)

    exp_path = out_dir / "_expected.jsonl"
    with exp_path.open("w", encoding="utf-8", newline="\n") as fp:
        for spec, rec, sheets, c in zip(specs, recs, sheets_list, contents):
            obj = {
                "file": spec.filename,
                "layout_version": spec.version,
                "form_revision": REV_INFO[spec.version][0],
                "main_sheet": rec.main_sheet,
                "work_type": spec.work_type,
                "incident_id": spec.inc.incident_id if spec.inc else None,
                "values": rec.values,
                "labels_used": rec.labels,
                "sheets": sheets,
                "hidden_sheets": rec.hidden_sheets,
                "hidden_columns": rec.hidden_columns,
                "images": rec.images,
                "extra_sheets": rec.extra_sheets,
                "irregularities": rec.irregularities,
                "counts": {"check_items": len(c.rows), "ng": sum(1 for rw in c.rows if rw.judge == "×"),
                           "watch": sum(1 for rw in c.rows if rw.judge == "△")},
            }
            fp.write(json.dumps(obj, ensure_ascii=False) + "\n")
    readme_path = out_dir / "_README.md"
    readme_path.write_text(_readme(specs, recs), encoding="utf-8", newline="\n")
    return paths + [exp_path, readme_path]


if __name__ == "__main__":
    import time as _time

    t0 = _time.time()
    written = generate()
    for p in written:
        print(p)
    print(f"{len(written)} files ({_time.time() - t0:.1f}s)")
