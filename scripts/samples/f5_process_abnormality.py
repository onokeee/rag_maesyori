"""F5 工程異常連絡票（品質系）のサンプル帳票を生成する。

    python -m scripts.samples.f5_process_abnormality

samples/forms/F5_工程異常連絡票/ に以下を書き出す。
- 版ごとのフォルダ（A3横_Rev1_2023以前 / A3横_Rev2_2024改訂 / A4縦_Rev3_2024改訂）に Excel帳票 30 ファイル。
  1つのフォルダの中は様式・シート構成がそろっているので、そのまま「帳票取り込み」へまとめて投入できる。
- _expected.jsonl : ファイルごとの正解値（人が読んだ値）と、そのファイルで使われているラベル文字列。
  `file` は本フォルダからの相対パス（`A4縦_Rev3_2024改訂/....xlsx`、区切りは `/`）。
- _README.md      : テンプレート・版の違い・意図的に入れた「ゆれ」の説明

元データは domain.standard_incidents() のうち lots_affected（影響ロット）のあるトラブル。
発行部署（製造課・生産技術課・品質保証課）が起票し、宛先部署（設備保全課など）が回答欄を埋める運用を想定する。
乱数は domain.rng() からのみ取り、xlsx の zip タイムスタンプ・文書プロパティも固定するため、
何度実行してもバイト単位で同じファイルになる。
"""
from __future__ import annotations

import io
import json
import math
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.page import PageMargins
from PIL import Image, ImageDraw, ImageFont

from . import domain as D

# ---------------------------------------------------------------------------
# 基本設定
# ---------------------------------------------------------------------------
FORM_FOLDER = "F5_工程異常連絡票"
# 様式の版ごとのフォルダ。同じフォルダの中は様式・シート構成がそろっており、まとめて取り込める。
# 名前は README の版（layout_version）そのままに使用開始時期を足したもの（Windows で使えない文字「. : / \ * ? " < > |」は使わない）
VERSION_FOLDERS = {
    "A3横_Rev.1": "A3横_Rev1_2023以前",
    "A3横_Rev.2": "A3横_Rev2_2024改訂",
    "A4縦_Rev.3": "A4縦_Rev3_2024改訂",
}
N_FILES = 30
SNAPSHOT_DATE = date(2026, 8, 31)          # 帳票フォルダを保存した時点（これより後の回答は存在しない）
REV3_START = date(2024, 10, 1)             # A4縦 Rev.3 への改訂日
REV2_START = date(2024, 1, 1)              # A3横 Rev.1 → Rev.2 のラベル変更日
ZIP_TIME = (2026, 9, 1, 9, 0, 0)           # xlsx 内部 zip エントリの固定タイムスタンプ

BASE_FONT = "ＭＳ Ｐゴシック"
ANSWER_FONT = "ＭＳ 明朝"                   # 回答欄は受信部署が別の書式で記入する
ANSWER_COLOR = "1F3F8F"
STAMP_COLOR = "C00000"
THIN = Side(style="thin", color="000000")
HAIR = Side(style="hair", color="000000")
MED = Side(style="medium", color="000000")
FILL_A = PatternFill("solid", fgColor="E7E6E6")
FILL_B = PatternFill("solid", fgColor="DDEBF7")
FILL_BAR_A = PatternFill("solid", fgColor="595959")
FILL_ANS_A = PatternFill("solid", fgColor="FCE4D6")
FILL_ANS_B = PatternFill("solid", fgColor="FFF2CC")

# 部課長（domain の人員マスタに課長がいない課のみ、本帳票の承認欄用に補う）
MANAGERS = {"設備保全課": "高橋 誠", "製造課": "佐々木 隆之", "生産技術課": "山本 健一", "品質保証課": "吉田 典子"}

# ロットID先頭2文字 -> (品名, 品番の本体)
PRODUCTS = {
    "MK": ("車載マイコン", "MK32A7"), "SR": ("SRAM 16Mb", "SR16M2"), "PX": ("CMOSイメージセンサ", "PX1208C"),
    "TQ": ("電源IC（PMIC）", "TQ5520B"), "NB": ("NORフラッシュ 256Mb", "NB256F3"), "HV": ("高耐圧ドライバIC", "HV800D"),
}
OPTIONS = ("選別", "再加工", "廃棄", "特採", "出荷停止")

# 設備種別ごとの選別検査（異常内容のキーワードで決まらない場合の既定）・再加工・後工程
INSPECT = {
    "CMP": ["欠陥検査", "49点膜厚測定", "欠陥レビューSEM"], "CVD": ["膜厚・屈折率測定", "パーティクル検査"],
    "ETC": ["CD-SEM測定", "欠陥検査"], "LIT": ["重ね合わせ測定", "CD測定", "マクロ検査"],
    "CLN": ["パーティクル検査", "TXRF金属汚染分析"], "IMP": ["TWシート抵抗測定", "サーマウェーブ測定"],
    "INS": ["他号機再測定"], "OHT": ["外観検査", "欠陥検査"], "ROB": ["外観検査", "裏面検査"], "UPW": ["パーティクル検査", "TXRF金属汚染分析"],
    "SCR": ["パーティクル検査"],
}
# 異常内容のキーワード -> (選別に使う検査, マップの名称, マップ下部の表記ラベル)。上から順に判定する
METRIC_RULES = [
    (("重ね合わせ",), ["重ね合わせ測定"], "重ね合わせマップ", "OVL"),
    (("CD",), ["CD-SEM測定", "CD測定"], "CDマップ", "CD"),
    (("フォーカス",), ["CD測定", "マクロ検査"], "CDマップ", "CD"),
    (("スクラッチ",), ["欠陥検査", "欠陥レビューSEM"], "欠陥マップ", "DEF"),
    (("パーティクル", "異物"), ["パーティクル検査"], "パーティクルマップ", "PC"),
    (("ウォーターマーク", "欠陥", "アーキング"), ["欠陥検査"], "欠陥マップ", "DEF"),
    (("残膜", "過研磨", "膜厚", "研磨レート", "ロールオフ", "プロファイル"), ["49点膜厚測定", "膜厚測定"], "膜厚マップ", "THK"),
    (("シート抵抗", "Rs", "ドーズ", "ビーム"), ["TWシート抵抗測定"], "Rsマップ", "Rs"),
    (("エッチ量", "E/R", "エッチレート"), ["エッチ量測定"], "E/Rマップ", "E/R"),
    (("金属", "汚染"), ["TXRF金属汚染分析"], "TXRFマップ", "TXRF"),
]
# SPC/QC の管理項目名 -> 異常内容にその項目が書かれているとみなすキーワード
SPC_ITEM_WORDS = {
    "CD": ("CD",), "欠陥数": ("欠陥", "スクラッチ", "パーティクル"), "膜厚": ("膜厚", "残膜", "ロールオフ", "過研磨", "成膜レート"),
    "パーティクル": ("パーティクル", "異物"), "シート抵抗": ("シート抵抗", "Rs", "ドーズ"), "重ね合わせ": ("重ね合わせ",),
    "E/R均一性": ("E/R", "エッチレート", "エッチ量"),
}
REWORK = {
    "LIT": ["レジスト剥離→再塗布・再露光（リワーク）", "アッシング・洗浄後に再露光"],
    "CMP": ["追加研磨（タッチアップ）で残膜を規格内へ調整", "再研磨＋後洗浄"],
    "CLN": ["再洗浄（SC1→DHF）を実施", "再洗浄後にパーティクル再検査"],
    "CVD": ["DHFで膜剥離後、再成膜"],
}
REWORK_PARTICLE = {"CMP": ["後洗浄ユニットで再洗浄", "ブラシ洗浄のみ再実施"], "CVD": ["スクラバ洗浄後にパーティクル再検査"]}
NEXT_PROCESS = {"CMP": "配線形成", "CVD": "CMP", "ETC": "洗浄・アッシング", "LIT": "エッチング", "CLN": "成膜", "IMP": "アニール",
                "INS": "次工程", "OHT": "次工程", "ROB": "次工程", "UPW": "次工程", "SCR": "次工程"}
# 規格の一文は、異常内容のキーワードに合うものだけを添える
SPEC_LINE = [
    ("CMP", ("スクラッチ",), ["規格：スクラッチ 15個/wf以下"]), ("CMP", ("残膜", "過研磨", "膜厚"), ["規格：残膜 NU 5%以下"]),
    ("CMP", ("パーティクル",), ["規格：パーティクル(≧0.1µm) 50個/wf以下"]),
    ("CVD", ("膜厚",), ["規格：膜厚NU 3%以下"]), ("CVD", ("パーティクル",), ["規格：パーティクル(≧0.1µm) 30個以下"]),
    ("ETC", ("CD",), ["規格：CD ±3nm"]), ("ETC", ("E/R", "エッチレート", "Vpp", "RF", "反射波"), ["規格：E/R均一性 5%以下"]),
    ("LIT", ("重ね合わせ",), ["規格：重ね合わせ ±8nm"]), ("LIT", ("CD", "フォーカス"), ["規格：CD ±4nm"]),
    ("CLN", ("パーティクル", "ウォーターマーク", "欠陥"), ["規格：パーティクル 20個/wf以下"]),
    ("IMP", ("シート抵抗", "ビーム"), ["規格：Rs均一性 1.5%以下", "規格：ドーズ偏差 ±2%"]),
]

# 発見方法の欄に書く短い表記
DET_SHORT = {
    "装置アラーム": ["装置アラーム", "設備アラーム停止"], "オペレーター": ["作業者発見", "OP発見"], "オペレーター申告": ["OP申告"],
    "FDC監視": ["FDC", "FDC監視"], "後工程からの連絡": ["後工程連絡", "次工程指摘"], "保全巡回": ["保全巡回"],
    "中央監視": ["中央監視アラーム"], "使用先装置からの連絡": ["使用先設備より連絡"],
}


# ---------------------------------------------------------------------------
# 帳票1枚分のデータ
# ---------------------------------------------------------------------------
@dataclass
class LotRow:
    lot: str
    code: str
    size: int
    hold: int
    scrap: int = 0
    rework: int = 0
    tokusai: int = 0
    disp: str = ""
    place: str = ""

    @property
    def ng(self) -> int:
        return self.scrap + self.rework + self.tokusai


@dataclass
class Rec:
    idx: int
    inc: D.Incident
    version: str                 # "A" / "B"
    rev: str                     # "Rev.3" / "Rev.2" / "Rev.1"
    report_id: str
    report_date: date
    due: date
    issuer: D.Person
    issuing_dept: str
    to_section: str              # 宛先の課・係（正式）
    to_dept: str                 # 帳票に書かれた宛先表記
    cc: str
    product_name: str
    product_code: str
    lots: list
    checks: set
    detected_short: str
    discovery_class: str
    finder: str
    occurred_text: str
    symptom: str
    cause_est: str
    treatment: str
    outflow: str
    outflow_detail: str
    remarks: str
    creator: str
    checker: str
    approver: str
    answer_state: str            # 回答済 / 一部回答 / 未回答
    answer_date: date | None
    answerer: str
    answer_staff: str
    answer_checker: str
    answer_approver: str
    answer_dept: str
    ans_cause: str
    ans_why: list
    ans_action: str
    ans_prevention: str
    ans_horiz: str
    ans_result: str
    confirm_date: date | None
    close_judgement: str
    qa_check: str
    qa_comment: str
    map_specs: list              # [(pattern, top_text, bottom_text, caption)]
    flags: list = field(default_factory=list)   # このファイルに入れたゆれ
    qty_comment: bool = False                   # 数量セルに訂正コメントを残す（旧様式のみ）


# ---------------------------------------------------------------------------
# 小物ヘルパー
# ---------------------------------------------------------------------------
def _next_weekday(d: date) -> date:
    while D._is_holiday(d):
        d += timedelta(days=1)
    return d


def _sn(name: str) -> str:
    return name.split(" ")[0]


def _numbered(lines: list[str], style: str) -> str:
    out = []
    for i, s in enumerate(lines, 1):
        if style == "・":
            out.append(f"・{s}")
        elif style == "1)":
            out.append(f"{i}) {s}")
        elif style == "①":
            out.append(f"{chr(0x2460 + i - 1)}{s}")
        else:
            out.append(s)
    return "\n".join(out)


def _strip_cause(c: str) -> str:
    """確定原因の文から補足（推定・メーカー見解・括弧書き）を落とし、発行側の推定原因に使える形にする。"""
    c = re.sub(r"。同一箇所での再発.*$", "", c)
    c = re.sub(r"^(特定できず。|原因不明（再現せず）。)", "", c)
    c = re.sub(r"(（確定は[^（）]*）|（メーカー見解も同様）|。ただし複合要因の可能性あり)", "", c)
    # 末尾の「（設置13年目）」「（前回PMから30日）」のような補足は何段でも落とす
    while re.search(r"（[^（）]*）$", c):
        c = re.sub(r"（[^（）]*）$", "", c)
    c = re.sub(r"(の可能性が高い|の可能性を疑っている|の可能性|と推定|と判断|と思われる)$", "", c)
    return c.strip("。 ")


def _clean_src(s: str) -> str:
    """元トラブルの文に残っている任意記号（「清掃・?交換」の ? など）を、人が書いた形に直す。"""
    if not s:
        return s
    s = re.sub(r"・\?", "・", s)
    s = re.sub(r"(^|\n|[ 　)）.、])\?", r"\1", s)
    return s.replace("?", "")


def _metric_rule(inc: D.Incident):
    """異常内容に出てくる測定項目から (選別検査, マップ名称, マップ表記ラベル) を決める。"""
    sym = unicodedata.normalize("NFKC", inc.symptom)
    for kws, insp, cap, lab in METRIC_RULES:
        if any(k in sym for k in kws):
            return insp, cap, lab
    return None


def _effective_detection(inc: D.Incident) -> str:
    """発見方法。元トラブルの detected_by より、異常内容に書かれた発見経緯（後工程連絡・巡回）を優先する。"""
    sym = inc.symptom
    if "後工程" in sym:
        return "後工程からの連絡"
    if "巡回" in sym and inc.detected_by not in ("保全巡回", "オペレーター"):
        return "オペレーター"
    # 「SPC異常(CD)」のような管理項目名が異常内容と食い違う場合（E/R異常なのにCD等）は項目名を書かない
    m = re.match(r"^(SPC異常|QC)\((.+)\)$", inc.detected_by)
    if m:
        kws = SPC_ITEM_WORDS.get(m.group(2))
        if kws and not any(w in unicodedata.normalize("NFKC", sym) for w in kws):
            return m.group(1)
    return inc.detected_by


def _next_process_name(inc: D.Incident) -> str:
    """異常内容に「後工程(リソ)」とあればその工程名、なければ設備種別の標準的な次工程。"""
    m = re.search(r"後工程[（(]([^）)]+)[）)]", inc.symptom)
    return m.group(1) if m else NEXT_PROCESS.get(D.kind_of(inc.equipment), "次工程")


def _date_text(d: date, r) -> str:
    """年を省略しない日付表記（M/D だけの表記は発行日には使わない）。"""
    return D.fmt_date(d, style=r.choices([0, 1, 2, 3, 4, 6, 8], weights=[30, 12, 8, 14, 12, 4, 6])[0])


_DATE_FORMATS = [
    ("yyyy/mm/dd", lambda d: f"{d.year}/{d.month:02d}/{d.day:02d}"),
    ("yyyy/m/d", lambda d: f"{d.year}/{d.month}/{d.day}"),
    ('yyyy"年"m"月"d"日"', lambda d: f"{d.year}年{d.month}月{d.day}日"),
    ('[$-411]ggge"年"m"月"d"日"', lambda d: f"令和{d.year - 2018}年{d.month}月{d.day}日"),
]


# ---------------------------------------------------------------------------
# 欠陥マップ画像（Pillow）
# ---------------------------------------------------------------------------
def _map_pattern(inc: D.Incident) -> str:
    sub, k = inc.subsystem, D.kind_of(inc.equipment)
    if "エッジ" in inc.symptom or "外周" in inc.symptom:
        return "edge"
    if sub in ("研磨パッド", "コンディショナ", "スラリー供給ライン"):
        return "scratch"
    if sub in ("研磨ヘッド",):
        return "ring"
    if sub in ("ESC(静電チャック)", "ヒーター", "リフトピン"):
        return "edge"
    if sub in ("フォーカス", "ステージ", "アライメント", "光源", "レチクル搬送", "光学系/光源"):
        return "shot"
    if sub in ("ノズル", "乾燥ユニット", "後洗浄ユニット", "スピンチャック", "薬液供給"):
        return "spiral" if sub == "乾燥ユニット" else "center"
    if sub in ("ビームライン", "イオン源", "高圧電源", "MFC/ガス供給", "RF電源", "終点検出"):
        return "gradient"
    if "搬送" in sub or k in ("OHT", "ROB", "AGV", "STK") or sub in ("ハンド/吸着", "グリッパ", "ホイスト", "ロードポート"):
        return "chip"
    if sub in ("チャンバー", "ドライポンプ", "TMP/排気"):
        return "cluster"
    return "random"


def _wafer_map_png(pattern: str, r, top_text: str, bottom_text: str, px: int = 168) -> bytes:
    """ウェーハ外形・ダイグリッド・欠陥点を描いた小さなマップ画像を PNG で返す（2倍で描いて縮小）。"""
    S = 2
    W = px * S
    img = Image.new("RGB", (W, W), (255, 255, 255))
    d = ImageDraw.Draw(img)
    cx = cy = W / 2
    R = W * 0.40
    die = W / 24
    # ダイグリッド（勾配マップはダイごとに色を塗る）
    ga, gb, gc = r.uniform(-1, 1), r.uniform(-1, 1), r.uniform(-1.2, 1.2)
    for ix in range(-12, 12):
        for iy in range(-12, 12):
            x0, y0 = cx + ix * die, cy + iy * die
            corners = [(x0, y0), (x0 + die, y0), (x0, y0 + die), (x0 + die, y0 + die)]
            if not all((x - cx) ** 2 + (y - cy) ** 2 <= R * R for x, y in corners):
                continue
            if pattern == "gradient":
                nx, ny = (x0 + die / 2 - cx) / R, (y0 + die / 2 - cy) / R
                v = 0.5 + 0.28 * (ga * nx + gb * ny) + 0.3 * gc * (nx * nx + ny * ny - 0.5)
                v = min(1.0, max(0.0, v))
                col = (int(255 * min(1, 2 * v)), int(255 * (1 - abs(v - 0.5) * 2) * 0.85 + 30), int(255 * min(1, 2 * (1 - v))))
                d.rectangle([x0, y0, x0 + die, y0 + die], fill=col, outline=(255, 255, 255))
            else:
                d.rectangle([x0, y0, x0 + die, y0 + die], outline=(215, 215, 215))
    d.ellipse([cx - R, cy - R, cx + R, cy + R], outline=(30, 30, 30), width=2 * S)
    d.polygon([(cx - 5 * S, cy + R + 2), (cx + 5 * S, cy + R + 2), (cx, cy + R - 7 * S)], fill=(255, 255, 255), outline=(30, 30, 30))

    pts: list[tuple[float, float]] = []

    def inside(x, y):
        return (x - cx) ** 2 + (y - cy) ** 2 <= (R * 0.98) ** 2

    red = (205, 20, 20)
    if pattern == "scratch":
        for _ in range(r.randint(1, 3)):
            rr = R * r.uniform(0.35, 1.1)
            ox, oy = cx + r.uniform(-R, R) * 0.6, cy + r.uniform(-R, R) * 0.6
            a0 = r.uniform(0, 360)
            ext = r.uniform(25, 70)
            steps = 40
            prev = None
            for s in range(steps + 1):
                a = math.radians(a0 + ext * s / steps)
                x, y = ox + rr * math.cos(a), oy + rr * math.sin(a)
                if inside(x, y) and prev is not None and inside(*prev):
                    d.line([prev, (x, y)], fill=red, width=2 * S)
                prev = (x, y)
                if inside(x, y) and r.random() < 0.25:
                    pts.append((x + r.gauss(0, 3), y + r.gauss(0, 3)))
        pts += [(cx + r.uniform(-R, R), cy + r.uniform(-R, R)) for _ in range(r.randint(3, 10))]
    elif pattern == "edge":
        ac = r.uniform(0, 2 * math.pi)
        for _ in range(r.randint(40, 90)):
            a = ac + r.gauss(0, 0.9)
            rad = R * r.uniform(0.84, 0.97)
            pts.append((cx + rad * math.cos(a), cy + rad * math.sin(a)))
    elif pattern == "ring":
        rr = R * r.uniform(0.45, 0.65)
        for _ in range(r.randint(40, 80)):
            a = r.uniform(0, 2 * math.pi)
            rad = rr + r.gauss(0, R * 0.04)
            pts.append((cx + rad * math.cos(a), cy + rad * math.sin(a)))
    elif pattern == "center":
        for _ in range(r.randint(30, 70)):
            pts.append((cx + r.gauss(0, R * 0.2), cy + r.gauss(0, R * 0.2)))
    elif pattern == "spiral":
        turns = r.uniform(1.5, 2.5)
        for s in range(90):
            t = s / 90
            a = t * turns * 2 * math.pi
            rad = R * 0.9 * t
            pts.append((cx + rad * math.cos(a) + r.gauss(0, 2), cy + rad * math.sin(a) + r.gauss(0, 2)))
    elif pattern == "cluster":
        for _ in range(r.randint(2, 4)):
            a, rad = r.uniform(0, 2 * math.pi), R * r.uniform(0, 0.8)
            ccx, ccy = cx + rad * math.cos(a), cy + rad * math.sin(a)
            pts += [(ccx + r.gauss(0, R * 0.06), ccy + r.gauss(0, R * 0.06)) for _ in range(r.randint(8, 22))]
        pts += [(cx + r.uniform(-R, R), cy + r.uniform(-R, R)) for _ in range(r.randint(5, 15))]
    elif pattern == "shot":
        sw, sh = die * 4, die * 3
        col = r.randint(-3, 2)
        for _ in range(r.randint(2, 5)):
            row = r.randint(-3, 2)
            c2 = col if r.random() < 0.6 else r.randint(-3, 2)
            x0, y0 = cx + c2 * sw, cy + row * sh
            if inside(x0 + sw / 2, y0 + sh / 2):
                d.rectangle([x0, y0, x0 + sw, y0 + sh], fill=(255, 205, 205), outline=(200, 80, 80))
                pts += [(x0 + r.uniform(0, sw), y0 + r.uniform(0, sh)) for _ in range(r.randint(3, 8))]
    elif pattern == "chip":
        a = r.uniform(0, 2 * math.pi)
        ex, ey = cx + R * math.cos(a), cy + R * math.sin(a)
        d.pieslice([ex - 9 * S, ey - 9 * S, ex + 9 * S, ey + 9 * S], 0, 360, fill=(255, 255, 255), outline=red, width=S)
        for _ in range(r.randint(1, 3)):
            b = a + math.pi + r.gauss(0, 0.5)
            ln = R * r.uniform(0.2, 0.6)
            d.line([(ex, ey), (ex + ln * math.cos(b), ey + ln * math.sin(b))], fill=red, width=S)
        pts += [(ex + r.gauss(0, R * 0.15), ey + r.gauss(0, R * 0.15)) for _ in range(r.randint(10, 30))]
    elif pattern == "random":
        pts += [(cx + r.uniform(-R, R), cy + r.uniform(-R, R)) for _ in range(r.randint(15, 60))]
    for x, y in pts:
        if inside(x, y):
            d.ellipse([x - 1.8 * S, y - 1.8 * S, x + 1.8 * S, y + 1.8 * S], fill=red)

    font = ImageFont.load_default(size=10 * S)
    d.text((4 * S, 2 * S), top_text, fill=(0, 0, 0), font=font)
    d.text((4 * S, W - 14 * S), bottom_text, fill=(0, 0, 0), font=font)
    if pattern == "gradient":
        for i in range(40):
            v = i / 39
            col = (int(255 * min(1, 2 * v)), int(255 * (1 - abs(v - 0.5) * 2) * 0.85 + 30), int(255 * min(1, 2 * (1 - v))))
            d.rectangle([W - 10 * S, W - 16 * S - i * S * 2, W - 5 * S, W - 16 * S - (i - 1) * S * 2], fill=col)
    img = img.resize((px, px), Image.LANCZOS)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


# ---------------------------------------------------------------------------
# 対象トラブルの選定
# ---------------------------------------------------------------------------
def _weighted_pick(r, cands: list, weight, k: int) -> list:
    out, pool = [], list(cands)
    for _ in range(min(k, len(pool))):
        ws = [weight(x) for x in pool]
        x = r.choices(pool, weights=ws)[0]
        out.append(x)
        pool.remove(x)
    return out


def select_incidents() -> list[D.Incident]:
    """影響ロットのあるトラブルから30件を選ぶ（期間に分散、品質系を多め、設備の偏りを抑える）。"""
    r = D.rng("f5:select")
    excluded = ("AGV", "STK", "EXH", "PCW", "VAC")
    pool = [i for i in D.standard_incidents() if i.lots_affected and D.kind_of(i.equipment) not in excluded]
    sev_w = {"重大": 1.5, "大": 1.3, "中": 1.0, "小": 0.6}
    used_eq: dict[str, int] = {}
    used_kind: dict[str, int] = {}

    def w(i: D.Incident) -> float:
        base = (3.0 if i.category == "品質" else 1.0) * sev_w[i.severity] * (1.4 if i.scrap_wafers else 1.0)
        if _effective_detection(i) == "後工程からの連絡":
            base *= 2.0          # 流出ありの例を増やす
        return base / (1 + 2.5 * used_eq.get(i.equipment.equipment_id, 0)) / (1 + 0.45 * used_kind.get(D.kind_of(i.equipment), 0))

    def mark(xs):
        for i in xs:
            used_eq[i.equipment.equipment_id] = used_eq.get(i.equipment.equipment_id, 0) + 1
            used_kind[D.kind_of(i.equipment)] = used_kind.get(D.kind_of(i.equipment), 0) + 1

    chosen = []
    # 未回答・一部回答の帳票用に 対応中3件・保留2件、流出ありの例として後工程発見1件、
    # 別紙ロット一覧の例として5ロット以上の案件1件を必ず含める
    recent = datetime(2026, 8, 12)       # 期間末の数週間に発生し、まだ回答期限前のもの
    for status, k, cond in (("対応中", 1, lambda i: i.occurred_at >= recent), ("対応中", 2, lambda i: i.occurred_at < recent),
                            ("保留", 2, lambda i: True), ("完了", 1, lambda i: _effective_detection(i) == "後工程からの連絡"),
                            ("完了", 1, lambda i: len(i.lots_affected) >= 5 and i.reported_at >= datetime(2024, 10, 10))):
        xs = _weighted_pick(r, [i for i in pool if i.status == status and cond(i)], w, k)
        mark(xs)
        chosen += xs
    ids = {i.incident_id for i in chosen}
    start, end = datetime(2023, 4, 1), datetime(2026, 9, 1)
    n_win = N_FILES - len(chosen)
    for wi in range(n_win):
        a = start + (end - start) * wi / n_win
        b = start + (end - start) * (wi + 1) / n_win
        cands = [i for i in pool if a <= i.occurred_at < b and i.status in ("完了", "経過観察") and i.incident_id not in ids]
        xs = _weighted_pick(r, cands, w, 1)
        mark(xs)
        chosen += xs
        ids.update(i.incident_id for i in xs)
    return sorted(chosen, key=lambda i: i.occurred_at)


# ---------------------------------------------------------------------------
# 帳票データの組み立て
# ---------------------------------------------------------------------------
def _person(section: str, role: str | None = None) -> list[D.Person]:
    return [p for p in D.people() if p.section == section and (role is None or p.role == role)]


def _issuer_of(inc: D.Incident) -> D.Person:
    rep = inc.reporter
    if rep.section.startswith("設備保全課") or rep.role == "FE":
        # 保全が見つけた場合でも、連絡票はその時間帯の製造課の班長名で発行される
        rep = _person(f"製造課 {D.crew_of(inc.occurred_at)}班", "班長")[0]
    return rep


def _destination(inc: D.Incident, issuer: D.Person) -> str:
    k = D.kind_of(inc.equipment)
    maint = "設備保全課 施設係" if k in D.UTILITY_KINDS else (
        "設備保全課 保全1係" if D.fab_of(inc.equipment) == "Fab1" else "設備保全課 保全2係")
    if inc.category == "人為":
        crew = f"製造課 {D.crew_of(inc.occurred_at)}班"
        return "生産技術課" if issuer.section.startswith("製造課") else crew
    if inc.category == "品質" and inc.cause_category in ("設定ミス", "調整不良", "前工程起因"):
        return "生産技術課" if issuer.section != "生産技術課" else maint
    return maint


def _lot_rows(inc: D.Incident, r, place: str) -> list[LotRow]:
    rows = []
    for lot in inc.lots_affected:
        lr = D.rng(f"f5:lot:{lot}")
        prefix = lot[:2]
        rev = "A1" if int(lot[2:4]) <= 24 else "A2"
        size = 25 if lot.endswith(".1") else lr.choice([12, 13, 10, 8, 13])
        hold = size if r.random() < 0.75 else max(3, size - r.randint(2, size // 2))
        rows.append(LotRow(lot, f"{PRODUCTS[prefix][1]}-{rev}", size, hold, place=place))
    return rows


def _alloc(rows: list[LotRow], attr: str, n: int) -> None:
    """不良枚数を先頭ロットから保留枚数の範囲で割り付ける。"""
    for row in rows:
        room = row.hold - row.ng
        take = min(room, n)
        setattr(row, attr, getattr(row, attr) + take)
        n -= take
        if n <= 0:
            break


def _detect_sentence(inc: D.Incident, r, nxt: str) -> str:
    det = _effective_detection(inc)
    if det == "装置アラーム":
        return r.choice(["装置アラーム停止により発覚", "アラーム停止時に処理中ウェーハを確認して発見", "装置アラーム（処理中停止）で発覚"])
    if det == "後工程からの連絡" and "後工程" in inc.symptom:
        return r.choice(["後工程からの連絡で発覚", f"{nxt}工程より連絡", "後工程連絡により当工程の処理品を確認"])
    if det.startswith("オペレーター") and "巡回" in inc.symptom:
        return r.choice(["OP巡回時に発見", "巡回中のオペレーターが発見"])
    if det.startswith("オペレーター"):
        return r.choice(["オペレーターが処理後の外観確認で発見", "OP巡回時に発見", "作業者からの申告で発覚"])
    if det == "FDC監視":
        return r.choice(["FDCでパラメータ逸脱を検知", "FDCアラートにより発覚"])
    if det == "SPC異常":
        return r.choice(["SPC管理限界外れで発覚", "SPCチャートのアウトオブコントロールで発覚"])
    if det == "QC":
        return r.choice(["定期QCにてNG", "QCウェーハの測定で規格外"])
    if det.startswith("SPC"):
        return r.choice([f"{det}（管理限界外れ）で発覚", f"SPCチャートのアウトオブコントロールで発覚（{det}）"])
    if det == "後工程からの連絡":
        return r.choice([f"後工程（{nxt}）から異常品の連絡あり", f"{nxt}工程の受入れ検査でNGとの連絡"])
    if det in ("保全巡回",):
        return "保全巡回時に装置異常を発見、処理済みロットを確認"
    if det in ("中央監視", "使用先装置からの連絡"):
        return "ユーティリティ異常（中央監視アラーム）の影響で処理ロットに異常"
    return r.choice([f"{det}で検出", f"{det}にてNG"])


def _estimated_cause(inc: D.Incident, r, issuer: D.Person) -> str:
    sub = inc.subsystem
    base = _strip_cause(inc.cause) if inc.cause else ""
    inv1 = (inc.investigation.split("。")[0] + "。") if inc.investigation else ""
    if inc.category == "人為":
        # 作業起因は「〇〇の不具合？」とは書かない
        opts = ["作業ミスの可能性（当事者に聞き取り中）", "条件設定・操作ミスと思われる", "作業手順の不備が疑われる"]
        if base:
            opts += [f"{base}の可能性", f"{base}と推定"]
        return r.choice(opts)
    if not base:
        opts = [f"不明（{sub}周辺の設備要因を疑う）。調査願います", "調査中。設備側での原因調査をお願いします",
                f"{sub}の不具合が疑われるが詳細不明", "現時点で不明"]
    elif issuer.section.startswith("製造課"):
        opts = [f"設備起因と思われる（{sub}）", f"{sub}の不具合？", f"{base}の可能性（保全一次見解）", "不明。保全にて調査願います",
                f"{sub}まわりの異常と推定"]
    else:
        opts = [f"{base}の可能性", f"{base}と推定（保全一次調査より）", f"設備起因（{sub}）と推定。詳細は宛先部署にて調査願います",
                f"{sub}の異常と思われる。{inv1}", f"保全一次見解：{base}", f"{base}によるものと推定"]
    s = r.choice(opts)
    if r.random() < 0.35:
        s += r.choice(["\n原因調査と対策の回答をお願いします。", "\n→回答願います", "。至急調査願います", "\n（詳細は回答にて）"])
    return s


def _symptom_text(inc: D.Incident, r, lots: list[LotRow], has_map: bool, nxt: str, occ: str) -> str:
    eq = inc.equipment
    k = D.kind_of(eq)
    sym = _clean_src(inc.symptom).rstrip("。")
    det = _detect_sentence(inc, r, nxt)
    total = sum(x.hold for x in lots)
    lots_s = lots[0].lot + (f" 他{len(lots) - 1}ロット" if len(lots) > 1 else "")
    pats = [
        lambda: f"{occ}頃、{eq.equipment_id}（{eq.process}）にて{sym}。\n{det}。",
        lambda: f"【現象】{sym}\n【発見】{det}（{occ}）\n【対象】{lots_s}",
        lambda: f"・現象：{sym}\n・発見経緯：{det}\n・影響範囲：{lots_s}（計{total}枚）",
        lambda: f"{det}。\n{sym}。\n対象ロット：{lots_s}",
        lambda: f"{eq.name}で処理した{lots_s}にて異常発生。\n{sym}。",
        lambda: f"{sym}。\n（{occ}、{det}）",
    ]
    t = r.choice(pats)()
    nsym = unicodedata.normalize("NFKC", sym)
    specs = [s for kk, kws, lines in SPEC_LINE if kk == k and any(w in nsym for w in kws) for s in lines]
    if specs and "規格" not in nsym and "管理値" not in nsym and r.random() < 0.4:
        t += "\n" + r.choice(specs)
    if has_map:
        t += r.choice(["\n分布は右図参照。", "\n※マップ添付", "", "\nワーストウェーハのマップを添付"])
    else:
        t += r.choice(["\n（検査前に保留したためマップ無し）", "\n※アラーム停止のため検査未実施、マップ添付なし"])
    return t


def _treatment_text(rec_rows: list[LotRow], checks: set, r, inspects: list[str], qa: str, place: str, yy: int, rework_txt: str) -> str:
    nl = len(rec_rows)
    tot = sum(x.hold for x in rec_rows)
    scrap = sum(x.scrap for x in rec_rows)
    rw = sum(x.rework for x in rec_rows)
    tok = sum(x.tokusai for x in rec_rows)
    lines = [r.choice([
        f"対象{nl}ロット（計{tot}枚）を{place}にて保留",
        f"該当ロット全数HOLD（{nl}ロット／{tot}枚）",
        f"{rec_rows[0].lot}{'ほか' if nl > 1 else ''}をMES上でHOLDし{place}へ移動",
        f"処理済み{nl}ロットを着工停止・保留（{tot}枚）",
    ])]
    if "選別" in checks:
        insp = r.choice(inspects)
        lines.append(r.choice([f"{insp}にて全数選別を実施", f"保留ウェーハは{insp}で全数確認し良否判定", f"{insp}（全数）で選別、NG品を抜取り"]))
    if "再加工" in checks:
        lines.append(f"{rework_txt}（{rw}枚）")
    if "廃棄" in checks:
        lines.append(r.choice([f"NG{scrap}枚は廃棄（MESスクラップ登録）", f"規格外{scrap}枚をスクラップ処理", f"{scrap}枚廃棄、残りは流動可否を判定"]))
    if "特採" in checks:
        lines.append(r.choice([f"軽微な規格外れ{tok}枚は特採申請（特採No.TS-{yy:02d}-{r.randint(10, 199):03d}）",
                               f"{tok}枚は品証{qa}さん判断で特採（顧客影響なし）"]))
    if "出荷停止" in checks:
        lines.append(r.choice([f"同期間に処理した完成品{r.randint(1, 4)}ロットを出荷停止・隔離", "出荷前在庫を出荷停止とし、倉庫で識別保管",
                               f"{rec_rows[0].lot}を含む出荷予定分を出荷停止（営業へ連絡済）"]))
    if r.random() < 0.35:
        lines.append(r.choice(["保留解除は品証判定後とする", "判定結果は別途連絡します", "設備は保全復旧・QC確認まで着工禁止"]))
    return _numbered(lines, r.choice(["・", "1)", "", "①"]))


def _qa_comment(inc: D.Incident, r, qa_person: D.Person, ans: dict, checks: set, lots: list, inspects: list,
                answer_date: date) -> str:
    """回答済みの連絡票に品証が書くコメント。

    回答内容の結果で書く内容を変える:
      - effective（有効）: 効果確認日が入っている・再発でない・水平展開が完了 → クローズの判断とその根拠
      - partial（一部）  : 効果確認が未了、再発案件、水平展開が予定・検討中 → 残っている確認事項の指示
      - monitoring（監視）: 元トラブルが経過観察 → 監視期間・再発時の依頼
    書き方は記入者で変える（主任は「〜のこと」「クローズ可」の体言止め、担当は「〜します」「〜願います」）。
    """
    if r.random() < 0.25:
        return ""                                # コメント欄は空欄のことも多い
    staff = qa_person.role == "担当"
    md = lambda d: f"{d.month}/{d.day}"          # noqa: E731
    insp = r.choice(inspects)
    n_lot = r.randint(3, 12)
    first_lot = lots[0].lot + ("ほか" if len(lots) > 1 else "")
    scrap = sum(x.scrap for x in lots)
    horiz_open = any(w in (ans["horiz"] or "") for w in ("予定", "検討"))
    rel = inc.related_incident_id
    cands: list[tuple[str, str]] = []            # (主任の書き方, 担当の書き方)
    if inc.status == "経過観察":
        until = answer_date + timedelta(days=r.choice([14, 21, 30, 45]))
        cands += [(f"経過観察期間（〜{md(until)}）終了後に再判定", f"経過観察（{md(until)}まで）の結果を見て判定します"),
                  ("原因推定のため、再発時はログ・サンプル確保のうえ連絡のこと", "再発した場合はログとサンプルを確保して連絡願います"),
                  (f"{insp}のトレンドを週次で品証へ共有のこと", f"{insp}のデータを週1回共有願います"),
                  ("監視結果を見てクローズ判断。それまで保留", "監視結果が出るまで保留とします")]
    elif ans["confirm"] and not inc.recurrence and not horiz_open:
        cd = ans["confirm"]
        cands += [(f"{md(cd)} 効果確認OK。クローズ可", f"{md(cd)}に効果確認しました。クローズとします"),
                  (f"対策後{n_lot}ロット流動、{insp}で異常なし。クローズ", f"対策後{n_lot}ロット流動し、{insp}で異常がないことを確認しました。クローズします"),
                  (f"保留ロット（{first_lot}）判定完了。本件クローズ", f"保留ロット（{first_lot}）の判定が完了したのでクローズします"),
                  ("対策内容確認。効果確認済みにつきクローズ", "対策内容・効果確認とも問題ありません。クローズします"),
                  ("了解。クローズ", "内容確認しました。クローズします")]
        if "特採" in checks:
            cands.append(("特採分の後工程影響なしを確認。クローズ", "特採分は後工程で問題ないことを確認済みです。クローズします"))
        if "出荷停止" in checks:
            cands.append(("出荷停止ロットは全数良品判定、出荷再開済み。クローズ", "出荷停止していたロットは全数良品判定となり、出荷を再開しています。クローズします"))
        if scrap:
            cands.append((f"廃棄{scrap}枚のスクラップ処理を確認。クローズ", f"廃棄{scrap}枚はスクラップ処理済みを確認しました。クローズします"))
    else:
        if horiz_open:
            cands.append(("水平展開の実施結果を次回報告のこと", "水平展開の実施結果が出たら連絡願います"))
        if inc.recurrence:
            cands.append(("再発案件のため、次に再発した場合は8Dで再提出のこと", "再発時は8Dレポートで再提出願います"))
            if rel:
                m = r.choice([2, 3])
                cands.append((f"前回（{rel}）と同一不具合。恒久対策の効果確認を{m}ヶ月継続のこと",
                              f"前回（{rel}）と同じ不具合のため、効果確認を{m}ヶ月継続願います"))
        if not ans["confirm"]:
            due = answer_date + timedelta(days=r.choice([14, 21, 28]))
            cands += [(f"効果確認は品証（{_sn(qa_person.name)}）で実施。確認後クローズ", f"効果確認は{_sn(qa_person.name)}で実施します。確認後にクローズします"),
                      (f"効果確認未了。{md(due)}頃に再確認", f"効果確認がまだのため、{md(due)}頃に再確認します"),
                      (f"対策後の{insp}結果（{n_lot}ロット分）を提出のこと", f"対策後の{insp}結果を{n_lot}ロット分提出願います")]
        if "特採" in checks:
            cands.append(("特採分の後工程での評価結果を確認してからクローズ", "特採分の後工程での評価結果を確認してからクローズします"))
    chief_txt, staff_txt = r.choice(cands)
    if staff:
        return D._noise(staff_txt, D._writer_habit(qa_person.employee_id), r)
    return chief_txt


def _build_record(idx: int, inc: D.Incident, seq_no: int, force_unanswered: str | None, no_map: bool = False,
                  two_maps: bool = False) -> Rec:
    r = D.rng(f"f5:rec:{inc.incident_id}")
    eq = inc.equipment
    k = D.kind_of(eq)
    yy = inc.occurred_at.year % 100
    issuer = _issuer_of(inc)
    habit = D._writer_habit(issuer.employee_id)

    # --- 日付・版 ---
    rd = inc.reported_at.date() + timedelta(days=r.choice([0, 0, 1, 1, 1, 2, 3]))
    if not issuer.section.startswith("製造課"):
        rd = _next_weekday(rd)
    version = "A" if rd >= REV3_START else "B"
    rev = "Rev.3" if version == "A" else ("Rev.2" if rd >= REV2_START else "Rev.1")
    due = _next_weekday(rd + timedelta(days=7 if version == "A" else 10))
    report_id = f"PA-{rd.year}-{seq_no:04d}" if version == "A" else f"工異{rd.year % 100}-{seq_no:03d}"

    # --- 部署 ---
    issuing_dept = issuer.section if issuer.section != "品質保証課" else "品質保証部 品質保証課"
    if issuer.section.startswith("製造課") and r.random() < 0.4:
        issuing_dept = "製造部 " + issuer.section
    to_section = _destination(inc, issuer)
    to_dept = r.choice([to_section, to_section.replace("設備保全課 ", "設備保全課（") + "）" if "設備保全課 " in to_section else to_section,
                        "製造部 " + to_section if not to_section.startswith("品質") else to_section])
    cc_list = [x for x in ("品質保証課", "生産技術課", "製造課") if x not in (issuer.section.split(" ")[0], to_section.split(" ")[0])]
    cc = "、".join(cc_list[: r.choice([1, 2])])

    # --- 品名・ロット ---
    place = r.choice(["工程内保留棚", "ストッカ保留エリア", f"{eq.area.split(' ')[-1]}保留棚", "品証保留棚"])
    lots = _lot_rows(inc, r, place)
    prefixes = {x.lot[:2] for x in lots}
    if len(prefixes) == 1:
        pname, pcode = PRODUCTS[lots[0].lot[:2]][0], lots[0].code
    else:
        pname = r.choice(["複数（ロット一覧参照）", "複数品種", " / ".join(PRODUCTS[p][0] for p in sorted(prefixes))])
        pcode = r.choice(["―", "複数", " / ".join(sorted({x.code for x in lots}))])

    # --- 処置区分と数量 ---
    det = _effective_detection(inc)
    outflow_yes = det == "後工程からの連絡"
    nxt = _next_process_name(inc)
    metric = _metric_rule(inc)
    inspects = metric[0] if metric else INSPECT.get(k, ["外観検査", "欠陥検査"])
    checks: set = set()
    if inc.scrap_wafers > 0:
        checks.add("廃棄")
    if k in REWORK and (k == "LIT" or r.random() < 0.45):
        checks.add("再加工")
    if det not in ("装置アラーム",) or len(lots) >= 2 or r.random() < 0.4:
        checks.add("選別")
    if inc.severity in ("小", "中") and inc.scrap_wafers <= 3 and "再加工" not in checks and r.random() < 0.3:
        checks.add("特採")
    if outflow_yes or inc.severity == "重大":
        checks.add("出荷停止")
    if not checks:
        checks.add("選別")
    total_hold = sum(x.hold for x in lots)
    _alloc(lots, "scrap", min(inc.scrap_wafers, total_hold))
    if "再加工" in checks:
        _alloc(lots, "rework", r.randint(1, max(1, min(25, total_hold - sum(x.ng for x in lots)))))
    if "特採" in checks:
        _alloc(lots, "tokusai", r.randint(1, max(1, min(8, total_hold - sum(x.ng for x in lots)))))
    if sum(x.ng for x in lots) == 0:
        extra = 0 if k == "INS" else r.choice([0, 1, 2, 3, 5])   # 測定機の異常ではウェーハ自体は壊れない
        if extra:
            _alloc(lots, "scrap", extra)
            checks.add("廃棄")
    # 保留枚数を廃棄で使い切って再加工・特採に回す枚数が残らなかった場合は、その区分を■にしない
    for opt, attr in (("廃棄", "scrap"), ("再加工", "rework"), ("特採", "tokusai")):
        if opt in checks and sum(getattr(x, attr) for x in lots) == 0:
            checks.discard(opt)
    if not checks:
        checks.add("選別")
    pending = inc.status == "対応中"
    for row in lots:
        parts = []
        if row.scrap:
            parts.append(r.choice([f"廃棄{row.scrap}", f"廃棄{row.scrap}枚"]))
        if row.rework:
            parts.append(r.choice([f"RW{row.rework}", f"再加工{row.rework}枚"]))
        if row.tokusai:
            parts.append(f"特採{row.tokusai}")
        if "選別" in checks and not parts:
            parts.append("全数選別")
        if "出荷停止" in checks and r.random() < 0.5:
            parts.append("出荷停止")
        row.disp = ("判定待ち" if pending and r.random() < 0.6 else "・".join(parts)) or "HOLD"

    qa = _sn(r.choice(_person("品質保証課")).name)
    rework_txt = ""
    if "再加工" in checks:
        particle = any(w in unicodedata.normalize("NFKC", inc.symptom) for w in ("パーティクル", "異物"))
        rework_txt = r.choice(REWORK_PARTICLE[k] if particle and k in REWORK_PARTICLE else REWORK[k])
    treatment = _treatment_text(lots, checks, r, inspects, qa, place, yy, rework_txt)
    if outflow_yes:
        outflow = "有"
        outflow_detail = r.choice([f"後工程（{nxt}）まで流出、同工程で停止済み", f"{nxt}工程で検出。客先流出なし",
                                   f"後工程へ{len(lots)}ロット流出（出荷前に停止）"])
    else:
        outflow = "無"
        outflow_detail = r.choice(["工程内で停止", "", "当該工程内で全数保留済み", "流出なし（保留済み）", ""])

    # --- 欠陥マップ ---
    pattern = _map_pattern(inc)
    map_specs = []
    map_insp = r.choice(inspects)          # 複数枚貼る場合も同じ検査の結果をそろえる
    if not no_map:
        n_maps = 2 if (two_maps and version == "B" and len(lots) >= 2) else 1
        for mi in range(n_maps):
            lot = lots[mi]
            slot = D.rng(f"f5:slot:{lot.lot}").randint(1, lot.size)
            m = re.search(r"(\d+)個", unicodedata.normalize("NFKC", inc.symptom))
            if metric and metric[2] in ("CD", "OVL"):
                # 測定系の異常はNG点（規格外れ点）を赤で示したマップ
                pattern = "edge" if "エッジ" in inc.symptom else ("shot" if k == "LIT" else pattern)
                bottom = f"{metric[2]} NG={r.randint(3, 17)}pts"
                cap = metric[1]
            elif pattern == "gradient" or (metric and metric[2] in ("THK", "Rs", "E/R")):
                label = metric[2] if metric and metric[2] in ("THK", "Rs", "E/R") else {"CVD": "THK", "IMP": "Rs", "ETC": "E/R"}.get(k, "THK")
                if pattern not in ("gradient", "edge", "ring"):
                    pattern = "gradient"
                # 異常内容に均一性の数値（例「7.9%のばらつき」「NU 24%」）があれば1枚目のマップはその値にそろえる
                pm = re.search(r"(\d+(?:\.\d+)?)%", unicodedata.normalize("NFKC", inc.symptom))
                nu = float(pm.group(1)) if pm and mi == 0 and 1.0 <= float(pm.group(1)) <= 40 else r.uniform(3.2, 9.5)
                bottom = f"{label} NU {nu:.1f}%"
                cap = {"THK": "膜厚マップ", "Rs": "Rsマップ", "E/R": "E/Rマップ"}[label]
            else:
                n_def = int(m.group(1)) if m and mi == 0 else r.randint(12, 180)
                if metric and metric[2] == "PC":
                    bottom, cap = f"PC={n_def}", "パーティクルマップ"
                elif metric:
                    bottom, cap = f"DEF={n_def}", metric[1]
                else:
                    # 異常内容が測定値ではない（アラーム・漏液等）場合は、保留品を検査した結果のマップ
                    bottom, cap = f"DEF={n_def}", f"保留品{map_insp}結果"
            map_specs.append((pattern, f"{lot.lot} #{slot:02d}", bottom, f"{cap}（{lot.lot} #{slot:02d}）"))

    occ_dt = inc.occurred_at
    occ_short = f"{occ_dt.month}/{occ_dt.day} {occ_dt.hour}:{occ_dt.minute:02d}"
    symptom = _symptom_text(inc, r, lots, bool(map_specs), nxt, occ_short)
    cause_est = _estimated_cause(inc, r, issuer)
    symptom = D._noise(symptom, habit, r)
    cause_est = D._noise(cause_est, habit, r)
    treatment = D._noise(treatment, habit, r)

    # --- 発行側の押印 ---
    creator = _sn(issuer.name)
    if issuer.section.startswith("製造課"):
        chief = _person(issuer.section, "班長")[0]
        checker = "" if chief == issuer else _sn(chief.name)
        approver = _sn(MANAGERS["製造課"])
    elif issuer.section == "生産技術課":
        checker = "" if issuer.role == "主任技師" else "原田"
        approver = _sn(MANAGERS["生産技術課"])
    else:
        checker = "" if issuer.role == "主任" else "小川"
        approver = _sn(MANAGERS["品質保証課"])
    if r.random() < 0.08:
        approver = ""                 # 承認印もらい忘れ

    # --- 回答欄 ---
    if force_unanswered or inc.status == "対応中":
        state = "未回答"
    elif inc.status == "保留":
        state = "一部回答"
    else:
        state = "回答済"
    answer_date = None
    if state != "未回答":
        base = inc.completed_at.date() if inc.completed_at else rd
        answer_date = _next_weekday(max(rd + timedelta(days=1), base) + timedelta(days=r.randint(0, 6)))
        if answer_date > SNAPSHOT_DATE:
            state, answer_date = "未回答", None

    ppl_sect = to_section
    if ppl_sect.startswith("設備保全課"):
        cands = [p for p in inc.assignees if p.section == ppl_sect and p.role in ("主任", "担当")] or _person(ppl_sect)
        staff = cands[0]
        chief = (_person(ppl_sect, "係長") or _person(ppl_sect, "主任"))[0]
        mgr = MANAGERS["設備保全課"]
        dept_disp = ppl_sect
    elif ppl_sect == "生産技術課":
        cands = [p for p in inc.assignees if p.section == "生産技術課" and p.role == "技師"] or _person("生産技術課", "技師")
        staff = r.choice(cands)
        chief = _person("生産技術課", "主任技師")[0]
        mgr = MANAGERS["生産技術課"]
        dept_disp = r.choice(["生産技術課", "製造部 生産技術課", "生技"])
    else:
        staff = r.choice([p for p in _person(ppl_sect) if p.role == "オペレーター"])
        chief = _person(ppl_sect, "班長")[0]
        mgr = MANAGERS["製造課"]
        dept_disp = ppl_sect

    ans = dict(answerer="", staff="", checker="", approver="", dept="", cause="", why=[], action="", prevention="",
               horiz="", result="", confirm=None, close="", qa="", qa_comment="")
    if state != "未回答":
        # 係長・主任クラスが回答者として名前を書き、担当欄には実作業者の印
        writer = staff if r.random() < 0.6 else chief
        ans["answerer"] = r.choice([writer.name, _sn(writer.name), f"{_sn(writer.name)}（{writer.role}）"])
        ans["staff"] = _sn(staff.name)
        ans["dept"] = dept_disp
        ans["cause"] = _clean_src(inc.cause) or "調査中"
        ans["action"] = _clean_src(inc.action)
        ans["horiz"] = _clean_src(inc.horizontal_deployment) or r.choice(["なし", "該当なし", "不要（本機固有）"])
        ans["result"] = _clean_src(inc.result)
    if state == "回答済":
        ans["checker"] = _sn(chief.name) if chief != staff else ""
        ans["approver"] = _sn(mgr) if r.random() < 0.9 else ""
        ans["why"] = [_clean_src(s) for s in inc.why_why]
        prev = _clean_src(inc.prevention) or r.choice(["特になし（突発事象のため監視継続）", "現行の管理で対応"])
        if r.random() < 0.12:
            prev = f"詳細は設備トラブル報告書 {inc.incident_id} を参照"
        ans["prevention"] = prev
        if inc.status == "経過観察":
            ans["result"] = (inc.result or "暫定処置で復帰") + r.choice(["（経過観察中）", "。引き続き監視", ""])
            ans["close"] = r.choice(["継続監視", "保留（経過観察）"])
        else:
            cd = _next_weekday(answer_date + timedelta(days=r.randint(7, 30)))
            if cd <= SNAPSHOT_DATE and r.random() < 0.8:
                ans["confirm"] = cd
                ans["close"] = r.choice(["クローズ", "可", "完了"])
                ans["qa"] = qa
        qa_person = next(p for p in _person("品質保証課") if _sn(p.name) == qa)
        ans["qa_comment"] = _qa_comment(inc, r, qa_person, ans, checks, lots, inspects, answer_date)
    elif state == "一部回答":
        ans["why"] = []
        ans["prevention"] = _clean_src(inc.prevention) or r.choice(["恒久対策は部品入荷後に別途回答します", "（検討中・後日回答）"])
        ans["horiz"] = r.choice(["恒久対策確定後に検討", "後日回答"])

    remarks_opts = ["", "", "", f"関連：{inc.incident_id}（設備トラブル報告）", "写真は共有フォルダ参照", f"品証{qa}さんへ連絡済み",
                    "前回同様の不具合あり" if inc.recurrence else "", "顧客影響なし（社内工程内）"]
    remarks = r.choice(remarks_opts)
    flags: list[str] = []
    if state == "未回答" and due + timedelta(days=3) <= SNAPSHOT_DATE:
        remarks = (remarks + "\n" if remarks else "") + f"※回答期限超過（{(due + timedelta(days=3)).month}/{(due + timedelta(days=3)).day} 督促済）"
        flags.append("回答期限超過の督促メモ")
    elif state != "未回答" and answer_date and answer_date > due:
        flags.append("回答日が回答期限超過")

    discovery_class = "後工程" if outflow_yes else ("品証検査" if det in ("マクロ検査",) else "工程内")
    finder = "FDC（自動）" if det == "FDC監視" else _sn(inc.reporter.name)
    detected_short = r.choice(DET_SHORT.get(det, [det]))

    return Rec(
        idx=idx, inc=inc, version=version, rev=rev, report_id=report_id, report_date=rd, due=due, issuer=issuer,
        issuing_dept=issuing_dept, to_section=to_section, to_dept=to_dept, cc=cc, product_name=pname, product_code=pcode,
        lots=lots, checks=checks, detected_short=detected_short, discovery_class=discovery_class, finder=finder,
        occurred_text="", symptom=symptom, cause_est=cause_est, treatment=treatment, outflow=outflow,
        outflow_detail=outflow_detail, remarks=remarks, creator=creator, checker=checker, approver=approver,
        answer_state=state, answer_date=answer_date, answerer=ans["answerer"], answer_staff=ans["staff"],
        answer_checker=ans["checker"], answer_approver=ans["approver"], answer_dept=ans["dept"], ans_cause=ans["cause"],
        ans_why=ans["why"], ans_action=ans["action"], ans_prevention=ans["prevention"], ans_horiz=ans["horiz"],
        ans_result=ans["result"], confirm_date=ans["confirm"], close_judgement=ans["close"], qa_check=ans["qa"],
        qa_comment=ans["qa_comment"], map_specs=map_specs, flags=flags,
    )


# ---------------------------------------------------------------------------
# Excel 描画ヘルパー
# ---------------------------------------------------------------------------
class Sheet:
    """merge・罫線・行高の自動調整をまとめたワークシートの薄いラッパー。"""

    def __init__(self, ws, font_name: str = BASE_FONT):
        self.ws = ws
        self.font_name = font_name
        self.fits: list[tuple] = []

    def box(self, r1, c1, r2, c2, value=None, *, size=10, bold=False, fill=None, h="left", v="center", wrap=True,
            font=None, color="000000", border=True, rotate=0, fmt=None, fit=True):
        ws = self.ws
        cell = ws.cell(r1, c1)
        if value is not None and value != "":
            cell.value = value
        cell.font = Font(name=font or self.font_name, size=size, bold=bold, color=color)
        cell.alignment = Alignment(horizontal=h, vertical=v, wrap_text=wrap, text_rotation=rotate)
        if fill is not None:
            cell.fill = fill
        if fmt:
            cell.number_format = fmt
        if border:
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        if r2 > r1 or c2 > c1:
            ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
        if fit and isinstance(value, str) and value and not rotate:
            self.fits.append((r1, c1, r2, c2, value, size))
        return cell

    def frame(self, r1, c1, r2, c2, side=MED):
        """範囲の外周だけを太線にする（内側の罫線は残す）。"""
        ws = self.ws
        for rr in range(r1, r2 + 1):
            for cc in range(c1, c2 + 1):
                if rr not in (r1, r2) and cc not in (c1, c2):
                    continue
                b = ws.cell(rr, cc).border
                ws.cell(rr, cc).border = Border(
                    left=side if cc == c1 else b.left, right=side if cc == c2 else b.right,
                    top=side if rr == r1 else b.top, bottom=side if rr == r2 else b.bottom)

    def units(self, c1, c2) -> float:
        return sum((self.ws.column_dimensions[get_column_letter(c)].width or 8.43) + 5 / 7 for c in range(c1, c2 + 1))

    def apply_fits(self, default_h: float):
        """文字数と結合幅から必要な行数を見積もり、足りなければ行高を広げる（単一行の範囲から先に処理）。"""
        ws = self.ws
        for r1, c1, r2, c2, text, size in sorted(self.fits, key=lambda x: x[2] - x[0]):
            avail = max(1.0, self.units(c1, c2) - 1.0)
            n_lines = 0
            for line in text.split("\n"):
                u = sum(1.9 if unicodedata.east_asian_width(ch) in "WFA" else 1.0 for ch in line) * size / 10 * 1.04
                n_lines += max(1, math.ceil(u / avail))
            need = n_lines * size * 1.38 + 5
            cur = sum(ws.row_dimensions[rr].height or default_h for rr in range(r1, r2 + 1))
            if need > cur:
                add = (need - cur) / (r2 - r1 + 1)
                for rr in range(r1, r2 + 1):
                    ws.row_dimensions[rr].height = round((ws.row_dimensions[rr].height or default_h) + add, 1)

    def base_font(self, max_row: int, max_col: int):
        """書式未設定のセルも帳票フォントにそろえる（空セルに入力したときの見た目用）。"""
        for rr in range(1, max_row + 1):
            for cc in range(1, max_col + 1):
                c = self.ws.cell(rr, cc)
                if c.font is None or c.font.name == "Calibri":
                    c.font = Font(name=self.font_name, size=10)


def _place_image(ws, png: bytes, col: int, row: int, px: int, off_x: int = 6, off_y: int = 4):
    """画像を貼る。Excel はセル幅を超える列オフセットを切り詰めるため、オフセットを列単位に繰り上げてから指定する。"""
    img = XLImage(io.BytesIO(png))
    img.width = img.height = px
    while True:
        col_px = int((ws.column_dimensions[get_column_letter(col)].width or 8.43) * 7 + 5)
        if off_x < col_px:
            break
        off_x -= col_px
        col += 1
    marker = AnchorMarker(col=col - 1, colOff=pixels_to_EMU(off_x), row=row - 1, rowOff=pixels_to_EMU(off_y))
    img.anchor = OneCellAnchor(_from=marker, ext=XDRPositiveSize2D(pixels_to_EMU(px), pixels_to_EMU(px)))
    ws.add_image(img)


def _put_date(sh: Sheet, r1, c1, r2, c2, d: date | None, r, *, as_cell: bool, answer=False, h="center") -> str:
    """日付を日付セル（表示形式付き）または文字列で書き、人が読む表記を返す。"""
    kw = dict(font=ANSWER_FONT, color=ANSWER_COLOR) if answer else {}
    if d is None:
        sh.box(r1, c1, r2, c2, None, h=h, **kw)
        return ""
    if as_cell:
        fmt, disp = r.choice(_DATE_FORMATS)
        sh.box(r1, c1, r2, c2, datetime(d.year, d.month, d.day), fmt=fmt, h=h, fit=False, **kw)
        return disp(d)
    s = _date_text(d, r)
    sh.box(r1, c1, r2, c2, s, h=h, **kw)
    return s


def _stamp(sh: Sheet, r1, c1, r2, c2, name: str, d: date | None):
    text = f"{name}\n{d.month}/{d.day}" if name and d else name
    sh.box(r1, c1, r2, c2, text or None, size=10, h="center", color=STAMP_COLOR, fit=False)


def _print_setup(ws, *, a3: bool, area: str, rev_text: str):
    ws.page_setup.paperSize = ws.PAPERSIZE_A3 if a3 else ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if a3 else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = area
    ws.page_margins = PageMargins(left=0.35, right=0.35, top=0.55, bottom=0.45, header=0.25, footer=0.2)
    ws.print_options.horizontalCentered = True
    ws.oddHeader.left.text = "社外秘"
    ws.oddHeader.left.size = 9
    ws.oddFooter.left.text = rev_text
    ws.oddFooter.left.size = 8
    ws.oddFooter.right.text = "保存期間：5年"
    ws.oddFooter.right.size = 8
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85 if a3 else 100


def _qty(v: int, as_text: bool, zen: bool) -> str | int:
    if not as_text:
        return v
    s = f"{v}枚"
    return D.to_zenkaku(s, digits_only=True) if zen else s


def _why_text(steps: list[str], r) -> str:
    if not steps:
        return ""
    style = r.choice(["なぜ", "→", "num"])
    if style == "なぜ":
        return "\n".join(f"なぜ{i}：{s}" for i, s in enumerate(steps, 1))
    if style == "→":
        return "\n→".join(steps)
    return "\n".join(f"{i}. {s}" for i, s in enumerate(steps, 1))


# ---------------------------------------------------------------------------
# レイアウトA: A4縦 Rev.3（方眼紙・回答欄は下段）
# ---------------------------------------------------------------------------
LABELS_A = dict(
    report_id="連絡No.", report_date="発行日", answer_due="回答期限", issuing_dept="発行部署", to_dept="宛先", reporter="発行者",
    occurred_at="発生日時", product_name="品名", product_code="品番", process="工程", equipment_id="設備", equipment_name="設備",
    line="ライン", detected_by="発見方法", lots="ロットNo.", qty_input="投入数", qty_defect="不良数", qty_hold="保留数",
    symptom="異常内容", cause_estimated="発生原因（推定）", disposition="処置区分", disposition_detail="処置内容",
    outflow="流出有無", outflow_detail="流出先・範囲", remarks="備考", answer_date="回答日", answer_dept="回答部署",
    answerer="回答者", cause="原因（確定）", why_why="なぜなぜ", action="暫定対策（処置）", prevention="恒久対策",
    horizontal_deployment="水平展開", result="効果確認", approver="承認", checker="確認", creator="作成",
    answer_approver="承認", answer_checker="確認", answer_staff="担当", qa_check="確認印", close_judgement="クローズ判定",
)


def _render_a(rec: Rec, r) -> tuple[Workbook, dict, dict, int]:
    inc, eq = rec.inc, rec.inc.equipment
    wb = Workbook()
    ws = wb.active
    ws.title = r.choices(["連絡票", "工程異常連絡票", "Sheet1"], weights=[55, 35, 10])[0]
    if ws.title == "Sheet1":
        rec.flags.append("シート名が既定のSheet1のまま")
    sh = Sheet(ws)
    NC, W, DH = 36, 2.45, 18.0
    for c in range(1, NC + 1):
        ws.column_dimensions[get_column_letter(c)].width = W
    for rr in range(1, 41):
        ws.row_dimensions[rr].height = DH
    V: dict = {}
    L: dict = dict(LABELS_A)
    lab = dict(fill=FILL_A, size=9, h="center")
    ans_kw = dict(font=ANSWER_FONT, color=ANSWER_COLOR)
    date_cells = r.random() < 0.5

    # タイトル・発行側押印
    sh.box(1, 1, 2, 22, "工程異常連絡票", size=18, bold=True, h="center", border=False, fit=False)
    sh.box(3, 1, 3, 22, "発行部署 → 宛先部署　／　宛先部署は回答期限までに下段回答欄へ記入し返却のこと", size=8, border=False, fit=False)
    sh.box(1, 23, 3, 24, "発\n行", **lab, fit=False)
    for c0, name, val in ((25, "承認", rec.approver), (29, "確認", rec.checker), (33, "作成", rec.creator)):
        sh.box(1, c0, 1, c0 + 3, name, **lab)
        _stamp(sh, 2, c0, 3, c0 + 3, val, rec.report_date if name != "承認" else _next_weekday(rec.report_date + timedelta(days=r.choice([0, 1]))))
    sh.frame(1, 23, 3, 36)
    V.update(approver=rec.approver, checker=rec.checker, creator=rec.creator)
    ws.row_dimensions[4].height = 6

    # ヘッダ項目
    sh.box(5, 1, 5, 5, L["report_id"], **lab)
    sh.box(5, 6, 5, 14, rec.report_id, h="center")
    sh.box(5, 15, 5, 18, L["report_date"], **lab)
    V["report_date"] = _put_date(sh, 5, 19, 5, 25, rec.report_date, r, as_cell=date_cells)
    sh.box(5, 26, 5, 29, L["answer_due"], **lab)
    V["answer_due"] = _put_date(sh, 5, 30, 5, 36, rec.due, r, as_cell=date_cells)
    V["report_id"] = rec.report_id

    occ = inc.occurred_at
    occ_text = r.choice(D.fmt_datetime_variants(occ)[:5] + [f"{occ.year}/{occ.month}/{occ.day} {occ.hour}:{occ.minute:02d}"])
    rows = [
        (6, "issuing_dept", rec.issuing_dept, "to_dept", rec.to_dept),
        (7, "reporter", rec.issuer.name, "occurred_at", occ_text),
        (8, "product_name", rec.product_name, "product_code", rec.product_code),
        (9, "process", eq.process, "equipment_id", f"{eq.equipment_id} {eq.name}"),
        (10, "line", eq.line, "detected_by", rec.detected_short),
    ]
    for rr, k1, v1, k2, v2 in rows:
        sh.box(rr, 1, rr, 5, L[k1], **lab)
        sh.box(rr, 6, rr, 18, v1)
        sh.box(rr, 19, rr, 22, L[k2], **lab)
        sh.box(rr, 23, rr, 36, v2)
    V.update(issuing_dept=rec.issuing_dept, to_dept=rec.to_dept, reporter=rec.issuer.name, occurred_at=occ_text,
             product_name=rec.product_name, product_code=rec.product_code, process=eq.process, equipment_id=eq.equipment_id,
             equipment_name=eq.name, line=eq.line, detected_by=rec.detected_short)

    # ロット表（4行まで。5ロット以上は別紙）
    sh.box(11, 1, 11, 36, "■ 対象ロット・数量", size=10, bold=True, fill=FILL_BAR_A, color="FFFFFF", fit=False)
    heads = [(1, 2, "No."), (3, 10, L["lots"]), (11, 17, "品番"), (18, 21, L["qty_input"]), (22, 25, L["qty_defect"]),
             (26, 29, L["qty_hold"]), (30, 36, "処置・備考")]
    for c1, c2, t in heads:
        sh.box(12, c1, 12, c2, t, **lab)
    overflow = len(rec.lots) > 4
    use_formula = not overflow and r.random() < 0.4
    for i in range(4):
        rr = 13 + i
        lot = rec.lots[i] if i < len(rec.lots) else None
        sh.box(rr, 1, rr, 2, i + 1, h="center", size=9)
        sh.box(rr, 3, rr, 10, lot.lot if lot else None, size=9)
        sh.box(rr, 11, rr, 17, lot.code if lot else None, size=9)
        sh.box(rr, 18, rr, 21, lot.size if lot else None, h="right", size=9)
        sh.box(rr, 22, rr, 25, lot.ng if lot else None, h="right", size=9)
        sh.box(rr, 26, rr, 29, lot.hold if lot else None, h="right", size=9)
        sh.box(rr, 30, rr, 36, lot.disp if lot else None, size=8)
    tot_in, tot_ng, tot_hold = (sum(x.size for x in rec.lots), sum(x.ng for x in rec.lots), sum(x.hold for x in rec.lots))
    sh.box(17, 1, 17, 17, "合計" + ("（別紙含む）" if overflow else ""), **lab)
    for (c1, c2), val, col in (((18, 21), tot_in, "R"), ((22, 25), tot_ng, "V"), ((26, 29), tot_hold, "Z")):
        cell = sh.box(17, c1, 17, c2, f"=SUM({col}13:{col}16)" if use_formula else val, h="right", size=9, bold=True, fit=False)
    sh.box(17, 30, 17, 36, None)
    if use_formula:
        rec.flags.append("合計欄がSUM式（キャッシュ値なし）")
    V.update(lots=[x.lot for x in rec.lots], qty_input=str(tot_in), qty_defect=str(tot_ng), qty_hold=str(tot_hold))
    sh.frame(11, 1, 17, 36)

    # 異常内容＋欠陥マップ
    for rr in range(18, 23):
        ws.row_dimensions[rr].height = 26
    sh.box(18, 1, 23, 5, L["symptom"], **lab, v="center", fit=False)
    sh.box(18, 6, 23, 24, rec.symptom, size=9, v="top")
    n_img = 0
    if rec.map_specs:
        pat, top, bottom, cap = rec.map_specs[0]
        sh.box(18, 25, 22, 36, None)
        sh.box(23, 25, 23, 36, cap, size=8, h="center")
        png = _wafer_map_png(pat, D.rng(f"f5:map:{rec.inc.incident_id}:0"), top, bottom, px=150)
        _place_image(ws, png, 25, 18, 150, off_x=22, off_y=4)
        n_img = 1
        V["defect_map_caption"] = cap
    else:
        sh.box(18, 25, 23, 36, "マップ\nなし", size=9, h="center", color="808080")
        rec.flags.append("欠陥マップ画像なし")
    V["symptom"] = rec.symptom

    # 推定原因・処置
    sh.box(24, 1, 24, 5, L["cause_estimated"], **lab)
    sh.box(24, 6, 24, 36, rec.cause_est, size=9)
    sep = r.choice(["　", "　　", "  "])
    check_text = sep.join(("■" if o in rec.checks else "□") + o for o in OPTIONS)
    sh.box(25, 1, 25, 5, L["disposition"], **lab)
    sh.box(25, 6, 25, 36, check_text, size=10)
    sh.box(26, 1, 26, 5, L["disposition_detail"], **lab)
    sh.box(26, 6, 26, 36, rec.treatment, size=9, v="top")
    sh.box(27, 1, 27, 5, L["outflow"], **lab)
    sh.box(27, 6, 27, 13, ("■有　□無" if rec.outflow == "有" else "□有　■無"), h="center")
    sh.box(27, 14, 27, 18, L["outflow_detail"], **lab)
    sh.box(27, 19, 27, 36, rec.outflow_detail or None, size=9)
    remarks = rec.remarks
    if overflow:
        remarks = ("※対象ロット5件目以降は別紙「ロット一覧」参照" + ("\n" + remarks if remarks else ""))
        rec.flags.append("ロット5件以上を別シートに記載")
    sh.box(28, 1, 28, 5, L["remarks"], **lab)
    sh.box(28, 6, 28, 36, remarks or None, size=9)
    V.update(cause_estimated=rec.cause_est, disposition=[o for o in OPTIONS if o in rec.checks], disposition_detail=rec.treatment,
             outflow=rec.outflow, outflow_detail=rec.outflow_detail, remarks=remarks)
    sh.frame(5, 1, 28, 36)
    ws.row_dimensions[29].height = 8

    # 回答欄（宛先部署記入：明朝・紺色）
    sh.box(30, 1, 30, 36, "▼ 回答欄（宛先部署にて記入し、回答期限までに発行部署へ返却）", size=10, bold=True, fill=FILL_ANS_A, fit=False)
    answered = rec.answer_state != "未回答"
    inline = answered and r.random() < 0.25
    if inline:
        ad = _date_text(rec.answer_date, r)
        txt = f"回答日：{ad}　回答部署：{rec.answer_dept}　回答者：{rec.answerer}"
        sh.box(31, 1, 31, 36, txt, **ans_kw)
        V.update(answer_date=ad, answer_dept=rec.answer_dept, answerer=rec.answerer)
        rec.flags.append("回答日・回答者を1セルに「ラベル：値」で記入")
    else:
        sh.box(31, 1, 31, 5, L["answer_date"], **lab)
        V["answer_date"] = _put_date(sh, 31, 6, 31, 12, rec.answer_date, r, as_cell=date_cells, answer=True)
        sh.box(31, 13, 31, 17, L["answer_dept"], **lab)
        sh.box(31, 18, 31, 26, rec.answer_dept or None, **ans_kw)
        sh.box(31, 27, 31, 29, L["answerer"], **lab)
        sh.box(31, 30, 31, 36, rec.answerer or None, **ans_kw)
        V.update(answer_dept=rec.answer_dept, answerer=rec.answerer)
    why = _why_text(rec.ans_why, r)
    for rr, key, val in ((32, "cause", rec.ans_cause), (33, "why_why", why), (34, "action", rec.ans_action),
                         (35, "prevention", rec.ans_prevention), (36, "horizontal_deployment", rec.ans_horiz), (37, "result", rec.ans_result)):
        sh.box(rr, 1, rr, 5, L[key], **lab)
        sh.box(rr, 6, rr, 36, val or None, size=9, v="top", **ans_kw)
        ws.row_dimensions[rr].height = 30 if not val else DH
    V.update(cause=rec.ans_cause, why_why=rec.ans_why, action=rec.ans_action, prevention=rec.ans_prevention,
             horizontal_deployment=rec.ans_horiz, result=rec.ans_result)

    # 回答部署・品証の押印
    sh.box(38, 1, 40, 3, "回答\n部署", **lab, fit=False)
    for c0, name, val in ((4, "承認", rec.answer_approver), (9, "確認", rec.answer_checker), (14, "担当", rec.answer_staff)):
        sh.box(38, c0, 38, c0 + 4, name, **lab)
        _stamp(sh, 39, c0, 40, c0 + 4, val, rec.answer_date)
    sh.box(38, 19, 40, 21, "品証\n確認", **lab, fit=False)
    sh.box(38, 22, 38, 27, L["qa_check"], **lab)
    _stamp(sh, 39, 22, 40, 27, rec.qa_check, rec.confirm_date)
    sh.box(38, 28, 38, 36, L["close_judgement"], **lab)
    sh.box(39, 28, 40, 36, rec.close_judgement or None, h="center", bold=True, **ans_kw)
    V.update(answer_approver=rec.answer_approver, answer_checker=rec.answer_checker, answer_staff=rec.answer_staff,
             qa_check=rec.qa_check, close_judgement=rec.close_judgement)
    sh.frame(30, 1, 40, 36)

    sh.apply_fits(DH)
    sh.base_font(40, NC)
    _print_setup(ws, a3=False, area="A1:AJ40", rev_text="様式QA-F05 Rev.3（2024.10改訂）")

    if overflow:
        ws2 = wb.create_sheet("ロット一覧")
        s2 = Sheet(ws2)
        for c, wdt in zip("ABCDEFG", (5, 16, 14, 8, 8, 8, 22)):
            ws2.column_dimensions[c].width = wdt
        s2.box(1, 1, 1, 7, f"別紙　対象ロット一覧（{rec.report_id}）", size=12, bold=True, border=False, fit=False)
        for c, t in enumerate(["No.", "ロットNo.", "品番", "投入数", "不良数", "保留数", "処置・備考"], 1):
            s2.box(3, c, 3, c, t, fill=FILL_A, h="center", size=9)
        for i, lot in enumerate(rec.lots, 1):
            for c, v in enumerate([i, lot.lot, lot.code, lot.size, lot.ng, lot.hold, lot.disp], 1):
                s2.box(3 + i, c, 3 + i, c, v, size=9, h="right" if isinstance(v, int) else "left")
        s2.base_font(4 + len(rec.lots), 7)
        _print_setup(ws2, a3=False, area=f"A1:G{3 + len(rec.lots)}", rev_text="様式QA-F05 別紙")
    return wb, V, L, n_img


# ---------------------------------------------------------------------------
# レイアウトB: A3横 Rev.1/Rev.2（左：発行部署記入、右：回答欄）
# ---------------------------------------------------------------------------
LABELS_B2 = dict(
    report_id="管理番号", report_date="発行年月日", answer_due="回答希望日", discovery_class="発見区分", issuing_dept="発信部署",
    reporter="発信者", to_dept="受信部署", cc="写し配布", line="ライン", product_name="製品名", product_code="製品コード",
    process="工程", equipment_id="使用設備", equipment_name="使用設備", occurred_at="発生日時", detected_by="発見方法",
    finder="発見者", lots="ロットNo.", qty_input="投入数", qty_defect="不良数", qty_hold="保留数", symptom="異常の内容",
    cause_estimated="発生原因（推定）", disposition="処置区分", disposition_detail="処置の内容", outflow="流出有無",
    outflow_detail="流出先", remarks="備考", answer_date="回答日", answerer="回答者", answer_dept="回答部署",
    cause="発生原因（確定）", why_why="なぜなぜ分析", action="暫定対策", prevention="恒久対策", horizontal_deployment="水平展開",
    result="効果確認", confirm_date="効果確認日", close_judgement="クローズ判定", related_incident_id="設備トラブル報告No.",
    qa_comment="品証コメント", approver="承認", checker="確認", creator="起票", answer_approver="承認", answer_checker="確認",
    answer_staff="担当", qa_check="確認",
)
LABELS_B1 = {**LABELS_B2, **dict(
    report_id="No.", symptom="不具合内容", cause_estimated="推定原因", disposition="処置", disposition_detail="処置方法",
    outflow="流出", answer_due="回答期日", cause="真因", action="応急処置", prevention="再発防止策", issuing_dept="発行元",
    to_dept="発行先",
)}
B_WIDTHS = {1: 1.5, 2: 14, 3: 11, 4: 11, 5: 11, 6: 9, 7: 9, 8: 9.5, 9: 11, 10: 11, 11: 10, 12: 11, 13: 2,
            14: 14, 15: 7.5, 16: 11, 17: 11, 18: 11, 19: 10, 20: 11, 21: 10, 22: 10, 23: 9, 24: 10, 25: 1.5}


def _render_b(rec: Rec, r) -> tuple[Workbook, dict, dict, int]:
    inc, eq = rec.inc, rec.inc.equipment
    L: dict = dict(LABELS_B1 if rec.rev == "Rev.1" else LABELS_B2)
    wb = Workbook()
    ws = wb.active
    ws.title = r.choices(["工程異常連絡票", "連絡票", "Sheet1"], weights=[65, 25, 10])[0]
    if ws.title == "Sheet1":
        rec.flags.append("シート名が既定のSheet1のまま")
    sh = Sheet(ws)
    DH = 20.0
    for c, wdt in B_WIDTHS.items():
        ws.column_dimensions[get_column_letter(c)].width = wdt
    for rr in range(1, 35):
        ws.row_dimensions[rr].height = DH
    V: dict = {}
    lab = dict(fill=FILL_B, size=10, h="center")
    ans_kw = dict(font=ANSWER_FONT, color=ANSWER_COLOR)
    date_cells = r.random() < 0.6
    qty_text = r.random() < 0.2
    qty_zen = qty_text and r.random() < 0.5

    # ---- 左：タイトル・発信部署押印 ----
    sh.box(1, 2, 2, 8, "工 程 異 常 連 絡 票", size=20, bold=True, h="center", border=False, fit=False)
    sh.box(3, 2, 3, 8, f"様式QA-05 {rec.rev}　※発信部署で起票し、受信部署は右欄に回答すること", size=8, border=False, fit=False)
    sh.box(1, 9, 3, 9, "発信\n部署", **lab, fit=False)
    for c0, name, val in ((10, "承認", rec.approver), (11, "確認", rec.checker), (12, L["creator"], rec.creator)):
        sh.box(1, c0, 1, c0, name, **lab)
        _stamp(sh, 2, c0, 3, c0, val, rec.report_date)
    sh.frame(1, 9, 3, 12)
    V.update(approver=rec.approver, checker=rec.checker, creator=rec.creator)
    ws.row_dimensions[4].height = 8

    sh.box(5, 2, 5, 2, L["report_id"], **lab)
    sh.box(5, 3, 5, 4, rec.report_id, h="center", bold=True)
    sh.box(5, 5, 5, 5, L["report_date"], **lab)
    V["report_date"] = _put_date(sh, 5, 6, 5, 7, rec.report_date, r, as_cell=date_cells)
    sh.box(5, 8, 5, 8, L["answer_due"], **lab)
    V["answer_due"] = _put_date(sh, 5, 9, 5, 10, rec.due, r, as_cell=date_cells)
    sh.box(5, 11, 5, 11, L["discovery_class"], **lab)
    sh.box(5, 12, 5, 12, rec.discovery_class, h="center")
    V.update(report_id=rec.report_id, discovery_class=rec.discovery_class)

    sh.box(6, 2, 6, 2, L["issuing_dept"], **lab)
    sh.box(6, 3, 6, 4, rec.issuing_dept)
    sh.box(6, 5, 6, 5, L["reporter"], **lab)
    sh.box(6, 6, 6, 7, rec.issuer.name)
    sh.box(6, 8, 6, 8, L["to_dept"], **lab)
    sh.box(6, 9, 6, 12, rec.to_dept)
    sh.box(7, 2, 7, 2, L["cc"], **lab)
    sh.box(7, 3, 7, 7, rec.cc or None)
    sh.box(7, 8, 7, 8, L["line"], **lab)
    sh.box(7, 9, 7, 12, eq.line)
    sh.box(8, 2, 8, 2, L["product_name"], **lab)
    sh.box(8, 3, 8, 7, rec.product_name)
    sh.box(8, 8, 8, 8, L["product_code"], **lab)
    sh.box(8, 9, 8, 12, rec.product_code)
    equip_text = r.choice([f"{eq.equipment_id}（{eq.name}）", f"{eq.equipment_id} {eq.name}", f"{eq.name}　{eq.equipment_id}"])
    sh.box(9, 2, 9, 2, L["process"], **lab)
    sh.box(9, 3, 9, 5, eq.process)
    sh.box(9, 6, 9, 7, L["equipment_id"], **lab)
    sh.box(9, 8, 9, 12, equip_text)
    sh.box(10, 2, 10, 2, L["occurred_at"], **lab)
    occ = inc.occurred_at
    if date_cells:
        occ_disp = f"{occ.year}/{occ.month}/{occ.day} {occ.hour}:{occ.minute:02d}"
        sh.box(10, 3, 10, 5, occ.replace(second=0), fmt="yyyy/m/d h:mm", h="left", fit=False)
    else:
        occ_disp = r.choice(D.fmt_datetime_variants(occ))
        sh.box(10, 3, 10, 5, occ_disp)
    sh.box(10, 6, 10, 7, L["detected_by"], **lab)
    sh.box(10, 8, 10, 10, rec.detected_short)
    sh.box(10, 11, 10, 11, L["finder"], **lab)
    sh.box(10, 12, 10, 12, rec.finder, h="center")
    V.update(issuing_dept=rec.issuing_dept, reporter=rec.issuer.name, to_dept=rec.to_dept, cc=rec.cc, line=eq.line,
             product_name=rec.product_name, product_code=rec.product_code, process=eq.process, equipment_id=eq.equipment_id,
             equipment_name=eq.name, occurred_at=occ_disp, detected_by=rec.detected_short, finder=rec.finder)

    # ロット表（6行）
    sh.box(11, 2, 11, 12, "対象ロット", size=10, bold=True, fill=FILL_B, fit=False)
    for c1, c2, t in ((2, 3, L["lots"]), (4, 5, "製品コード"), (6, 6, L["qty_input"]), (7, 7, L["qty_defect"]),
                      (8, 8, L["qty_hold"]), (9, 10, "処置"), (11, 12, "保管場所")):
        sh.box(12, c1, 12, c2, t, **lab)
    for i in range(6):
        rr = 13 + i
        lot = rec.lots[i] if i < len(rec.lots) else None
        sh.box(rr, 2, rr, 3, lot.lot if lot else None, size=10)
        sh.box(rr, 4, rr, 5, lot.code if lot else None, size=10)
        for c, attr in ((6, "size"), (7, "ng"), (8, "hold")):
            sh.box(rr, c, rr, c, _qty(getattr(lot, attr), qty_text, qty_zen) if lot else None, h="right")
        sh.box(rr, 9, rr, 10, lot.disp if lot else None, size=9)
        sh.box(rr, 11, rr, 12, (lot.place if i == 0 or r.random() < 0.3 else "同上") if lot else None, size=9)
    tot = (sum(x.size for x in rec.lots), sum(x.ng for x in rec.lots), sum(x.hold for x in rec.lots))
    sh.box(19, 2, 19, 5, "合計", **lab)
    for c, v in zip((6, 7, 8), tot):
        sh.box(19, c, 19, c, _qty(v, qty_text, qty_zen), h="right", bold=True)
    sh.box(19, 9, 19, 12, None)
    if qty_text:
        rec.flags.append("数量を「25枚」等の文字列で記入" + ("（全角数字）" if qty_zen else ""))
    qty_disp = [(D.to_zenkaku(f"{v}枚", digits_only=True) if qty_zen else f"{v}枚") if qty_text else str(v) for v in tot]
    V.update(lots=[x.lot for x in rec.lots], qty_input=qty_disp[0], qty_defect=qty_disp[1], qty_hold=qty_disp[2])
    roll = r.random() < 0.12
    if (roll or rec.qty_comment) and rec.lots[0].ng:
        ws.cell(13, 7).comment = Comment(f"再検査の結果で{max(0, rec.lots[0].ng - 1)}→{rec.lots[0].ng}枚に訂正（{_sn(rec.issuer.name)}）", _sn(rec.issuer.name))
        rec.flags.append("数量セルに訂正コメント")
    sh.frame(11, 2, 19, 12)

    # 異常の内容（縦書きラベル）＋マップ
    for rr in range(20, 26):
        ws.row_dimensions[rr].height = 30
    vertical = r.random() < 0.5
    sh.box(20, 2, 25, 2, L["symptom"] if not vertical else L["symptom"], **lab, rotate=255 if vertical else 0, fit=False)
    sh.box(20, 3, 25, 8, rec.symptom, size=10, v="top")
    n_img = 0
    if rec.map_specs:
        caps = "／".join(s[3] for s in rec.map_specs)
        sh.box(20, 9, 20, 12, caps, size=8, h="center")
        sh.box(21, 9, 25, 12, None)
        for mi, (pat, top, bottom, cap) in enumerate(rec.map_specs):
            px = 150 if len(rec.map_specs) == 2 else 176
            png = _wafer_map_png(pat, D.rng(f"f5:map:{rec.inc.incident_id}:{mi}"), top, bottom, px=px)
            _place_image(ws, png, 9, 21, px, off_x=8 + mi * 160 if len(rec.map_specs) == 2 else 70, off_y=4)
            n_img += 1
        V["defect_map_caption"] = caps
        if n_img == 2:
            rec.flags.append("欠陥マップ画像2枚")
    else:
        sh.box(20, 9, 25, 12, "（マップなし）", size=9, h="center", color="808080")
        rec.flags.append("欠陥マップ画像なし")
    V["symptom"] = rec.symptom

    sh.box(26, 2, 27, 2, L["cause_estimated"], **lab)
    sh.box(26, 3, 27, 12, rec.cause_est, size=10, v="top")
    sh.box(28, 2, 28, 2, L["disposition"], **lab)
    for j, o in enumerate(OPTIONS):
        sh.box(28, 3 + 2 * j, 28, 3 + 2 * j, "■" if o in rec.checks else "□", h="right", size=12, fit=False)
        sh.box(28, 4 + 2 * j, 28, 4 + 2 * j, o, h="left")
    sh.box(29, 2, 31, 2, L["disposition_detail"], **lab)
    sh.box(29, 3, 31, 12, rec.treatment, size=10, v="top")
    sh.box(32, 2, 32, 2, L["outflow"], **lab)
    sh.box(32, 3, 32, 3, "■ 有" if rec.outflow == "有" else "□ 有", h="center")
    sh.box(32, 4, 32, 4, "■ 無" if rec.outflow == "無" else "□ 無", h="center")
    sh.box(32, 5, 32, 5, L["outflow_detail"], **lab)
    sh.box(32, 6, 32, 12, rec.outflow_detail or None)
    sh.box(33, 2, 34, 2, L["remarks"], **lab)
    sh.box(33, 3, 34, 12, rec.remarks or None, size=9, v="top")
    V.update(cause_estimated=rec.cause_est, disposition=[o for o in OPTIONS if o in rec.checks], disposition_detail=rec.treatment,
             outflow=rec.outflow, outflow_detail=rec.outflow_detail, remarks=rec.remarks)
    sh.frame(5, 2, 34, 12)

    # ---- 右：回答欄 ----
    sh.box(1, 14, 2, 18, "【回答欄】", size=16, bold=True, h="left", border=False, fit=False)
    sh.box(3, 14, 3, 18, "（受信部署記入　回答希望日までに発信部署へ返送）", size=8, border=False, fit=False)
    sh.box(1, 19, 3, 19, "受信\n部署", **lab, fit=False)
    for c0, key, val in ((20, "answer_approver", rec.answer_approver), (21, "answer_checker", rec.answer_checker), (22, "answer_staff", rec.answer_staff)):
        sh.box(1, c0, 1, c0, L[key], **lab)
        _stamp(sh, 2, c0, 3, c0, val, rec.answer_date)
    sh.box(1, 23, 3, 23, "品証", **lab, fit=False)
    sh.box(1, 24, 1, 24, L["qa_check"], **lab)
    _stamp(sh, 2, 24, 3, 24, rec.qa_check, rec.confirm_date)
    sh.frame(1, 19, 3, 24)
    V.update(answer_approver=rec.answer_approver, answer_checker=rec.answer_checker, answer_staff=rec.answer_staff, qa_check=rec.qa_check)

    sh.box(5, 14, 5, 14, L["answer_date"], **lab)
    V["answer_date"] = _put_date(sh, 5, 15, 5, 16, rec.answer_date, r, as_cell=date_cells, answer=True)
    sh.box(5, 17, 5, 17, L["answerer"], **lab)
    sh.box(5, 18, 5, 19, rec.answerer or None, **ans_kw)
    sh.box(5, 20, 5, 20, L["answer_dept"], **lab)
    sh.box(5, 21, 5, 24, rec.answer_dept or None, **ans_kw)
    V.update(answerer=rec.answerer, answer_dept=rec.answer_dept)

    sh.box(6, 14, 8, 14, L["cause"], **lab)
    sh.box(6, 15, 8, 24, rec.ans_cause or None, v="top", **ans_kw)
    sh.box(9, 14, 13, 14, L["why_why"], **lab)
    for i in range(5):
        step = rec.ans_why[i] if i < len(rec.ans_why) else None
        sh.box(9 + i, 15, 9 + i, 15, f"なぜ{i + 1}", size=9, h="center", fill=FILL_ANS_B)
        sh.box(9 + i, 16, 9 + i, 24, step, size=9, **ans_kw)
    sh.box(14, 14, 17, 14, L["action"], **lab)
    sh.box(14, 15, 17, 24, rec.ans_action or None, size=9, v="top", **ans_kw)
    sh.box(18, 14, 21, 14, L["prevention"], **lab)
    sh.box(18, 15, 21, 24, rec.ans_prevention or None, size=9, v="top", **ans_kw)
    sh.box(22, 14, 23, 14, L["horizontal_deployment"], **lab)
    sh.box(22, 15, 23, 24, rec.ans_horiz or None, size=9, v="top", **ans_kw)
    sh.box(24, 14, 25, 14, L["result"], **lab)
    sh.box(24, 15, 25, 24, rec.ans_result or None, size=9, v="top", **ans_kw)
    sh.box(26, 14, 26, 14, L["confirm_date"], **lab)
    V["confirm_date"] = _put_date(sh, 26, 15, 26, 16, rec.confirm_date, r, as_cell=date_cells, answer=True)
    sh.box(26, 17, 26, 17, L["close_judgement"], **lab)
    sh.box(26, 18, 26, 19, rec.close_judgement or None, h="center", bold=True, **ans_kw)
    related = rec.inc.incident_id if rec.answer_state == "回答済" and r.random() < 0.7 else ""
    sh.box(26, 20, 26, 21, L["related_incident_id"], **{**lab, "size": 9})
    sh.box(26, 22, 26, 24, related or None, h="center", **ans_kw)
    sh.box(27, 14, 34, 14, L["qa_comment"], **lab)
    sh.box(27, 15, 34, 24, rec.qa_comment or None, v="top", size=10, **ans_kw)
    V.update(cause=rec.ans_cause, why_why=rec.ans_why, action=rec.ans_action, prevention=rec.ans_prevention,
             horizontal_deployment=rec.ans_horiz, result=rec.ans_result, close_judgement=rec.close_judgement,
             related_incident_id=related, qa_comment=rec.qa_comment)
    sh.frame(5, 14, 34, 24)

    # 8文字以上の項目ラベルはセル幅に収まるよう 9pt に下げてある（テンプレート作成時の手調整）
    for row in ws.iter_rows(min_row=1, max_row=34, max_col=25):
        for c in row:
            if (isinstance(c.value, str) and len(c.value) >= 8 and "\n" not in c.value and c.font.sz and c.font.sz >= 10
                    and c.fill is not None and str(c.fill.fgColor.rgb).endswith("DDEBF7")):
                c.font = Font(name=c.font.name, size=9, bold=c.font.b, color=c.font.color)
    sh.apply_fits(DH)
    sh.base_font(34, 25)
    _print_setup(ws, a3=True, area="A1:Y34", rev_text=f"様式QA-05 {rec.rev}")

    # 記入要領シート（テンプレートに残っている説明シート）
    if r.random() < 0.3:
        guide = wb.create_sheet("記入要領", index=0 if r.random() < 0.3 else None)
        gs = Sheet(guide)
        guide.column_dimensions["A"].width = 4
        guide.column_dimensions["B"].width = 90
        gs.box(1, 1, 1, 2, "工程異常連絡票　記入要領", size=14, bold=True, border=False, fit=False)
        tips = ["発信部署は異常発見後、当日中に起票し受信部署・品質保証課へ配布する。", "対象ロットはMESのHOLD登録と一致させること。",
                "処置区分は該当するものをすべて■にする（選別／再加工／廃棄／特採／出荷停止）。", "特採の場合は品質保証課の承認を得ること。",
                "受信部署は回答希望日までに右欄（回答欄）を記入し返送する。", "なぜなぜ分析は真因に到達するまで記入する（最大5段）。",
                "品質保証課は効果確認後にクローズ判定を行う。"]
        for i, t in enumerate(tips, 3):
            gs.box(i, 1, i, 1, f"{i - 2}.", border=False, fit=False)
            gs.box(i, 2, i, 2, t, border=False, fit=False)
        gs.base_font(3 + len(tips), 2)
        rec.flags.append("記入要領シートあり" + ("（先頭シート）" if wb.sheetnames[0] == "記入要領" else ""))
    return wb, V, L, n_img


# ---------------------------------------------------------------------------
# 決定的な保存
# ---------------------------------------------------------------------------
def _save_deterministic(wb: Workbook, path: Path, created: datetime, creator: str, modifier: str) -> None:
    """openpyxl の保存時刻（文書プロパティ・zip エントリ時刻）を固定し、再実行でバイト一致させる。"""
    wb.properties.creator = creator
    wb.properties.lastModifiedBy = modifier or creator
    wb.properties.title = "工程異常連絡票"
    buf = io.BytesIO()
    wb.save(buf)
    stamp = created.strftime("%Y-%m-%dT%H:%M:%SZ")
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:)", lambda m: m.group(1) + stamp.encode() + m.group(2), data)
            zi = zipfile.ZipInfo(info.filename, date_time=ZIP_TIME)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            dst.writestr(zi, data)
    path.write_bytes(out.getvalue())


def _file_name(rec: Rec, r) -> str:
    eq = rec.inc.equipment.equipment_id
    rd = rec.report_date
    if rec.version == "A":
        pats = [f"工程異常連絡票_{rec.report_id}_{eq}", f"{rec.report_id}_工程異常連絡票（{eq}）", f"工程異常連絡票_{rd:%Y%m%d}_{eq}"]
    else:
        pats = [f"{rec.report_id}_工程異常連絡票_{eq}", f"工異連絡票_{rd:%y%m%d}_{eq}", f"【工程異常】{rec.report_id} {eq}"]
    name = r.choice(pats)
    if rec.answer_state == "回答済" and r.random() < 0.35:
        name += "_回答済"
    return name + ".xlsx"


# ---------------------------------------------------------------------------
# README
# ---------------------------------------------------------------------------
def _readme(rows: list[dict]) -> str:
    from collections import Counter

    ver = Counter(x["layout_version"] for x in rows)
    fol = Counter(x["file"].split("/")[0] for x in rows)
    ans = Counter(x["answer_status"] for x in rows)
    flags = Counter(f for x in rows for f in x["irregularities"])
    lines = [
        "# F5 工程異常連絡票（品質系）サンプル",
        "",
        "`python -m scripts.samples.f5_process_abnormality` で生成（固定シード・再実行でバイト一致）。",
        "元データは `scripts/samples/domain.py` の `standard_incidents()` のうち影響ロット（`lots_affected`）があるトラブル30件。",
        "",
        "## フォルダの構成",
        "様式の版ごとにフォルダを分けてある。**1つのフォルダの中は同じ版（同じ様式・同じシート構成）の .xlsx だけ**なので、"
        "アプリの「帳票取り込み」にフォルダまるごと投入し、帳票の種類とシートを1回選ぶだけで、そのフォルダの全ファイルをまとめて読み取れる。",
        "",
        "```",
        "F5_工程異常連絡票/",
        "├ _README.md          ← このファイル",
        "├ _expected.jsonl     ← 正解値（1行1ファイル）",
    ]
    folders = [n for n in VERSION_FOLDERS.values() if fol.get(n)]
    lines += [f"{'└' if i == len(folders) - 1 else '├'} {name}/   … {fol[name]}ファイル（.xlsx のみ）"
              for i, name in enumerate(folders)]
    lines += [
        "```",
        "",
        f"`_expected.jsonl` の `file` は、この帳票フォルダからの相対パス（例 `{rows[-1]['file']}`、区切りは `/`）。",
        "",
        "## 帳票の運用想定",
        "- 異常を見つけた部署（製造課の各班・生産技術課・品質保証課）が起票し、原因部署（主に設備保全課の各係、設定起因なら生産技術課、人為なら製造課）へ発行する。",
        "- 上段／左側が発行部署の記入欄、下段／右側が宛先（受信）部署の回答欄。回答欄は **ＭＳ 明朝・紺色** で記入され、発行側（ＭＳ Ｐゴシック・黒）と書式が異なる。",
        "- 回答後、品質保証課が効果確認・クローズ判定を行う。押印欄は赤字の「姓＋日付(M/D)」。",
        "",
        "## レイアウト（版）",
        "版ごとに下の「フォルダ」へ入れてある。フォルダ名は版の名前＋使用開始時期。",
        "",
        "| layout_version | フォルダ | 用紙 | 使用期間 | 構成 |",
        "|---|---|---|---|---|",
        "| A3横_Rev.1 | `A3横_Rev1_2023以前/` | A3横 | 〜2023/12 | 左：発行部署欄、右：【回答欄】。ラベルが旧名称（No.／発行元／発行先／不具合内容／推定原因／処置／処置方法／流出／回答期日／真因／応急処置／再発防止策） |",
        "| A3横_Rev.2 | `A3横_Rev2_2024改訂/` | A3横 | 2024/01〜2024/09 | Rev.1 と同配置でラベルを改称（管理番号／発信部署／受信部署／異常の内容／発生原因（推定）／処置区分／処置の内容／流出有無／回答希望日／発生原因（確定）／暫定対策／恒久対策）。なぜなぜ分析5段・品証コメント・設備トラブル報告No.欄は Rev.1 から共通。Rev.3 改訂後も一部の班が使い続けている |",
        "| A4縦_Rev.3 | `A4縦_Rev3_2024改訂/` | A4縦 | 2024/10〜 | 方眼紙（36列）のコンパクト版。回答欄は下段。連絡No.は `PA-YYYY-NNNN`（旧版は `工異YY-NNN`）。ロット表は4行、なぜなぜは1セル |",
        "",
        f"件数（＝各フォルダのファイル数）: " + "、".join(f"{k} {v}件" for k, v in sorted(ver.items())) + "（Rev.3 改訂後も旧様式で起票された例を含む）",
        "",
        "回答状況: " + "、".join(f"{k} {v}件" for k, v in sorted(ans.items())),
        "",
        "## 主な項目",
        "連絡No.／発行日／発行部署→宛先部署／品名・品番／工程／設備／ライン／発生日時／発見方法／対象ロット（ロットNo.・投入数・不良数・保留数・処置）／"
        "異常内容（欠陥マップ画像付き）／発生原因（推定）／処置区分（■/□：選別・再加工・廃棄・特採・出荷停止）／処置内容／流出有無／回答期限／"
        "回答欄（回答日・回答者・原因（確定）・なぜなぜ・暫定対策・恒久対策・水平展開・効果確認）／承認・確認・作成（起票）印／品証確認・クローズ判定",
        "",
        "欠陥マップは Pillow で描いたウェーハ外形＋ダイグリッド＋欠陥点の PNG。分布はサブシステムに合わせている"
        "（研磨パッド＝スクラッチ弧、研磨ヘッド＝リング、ESC/ヒーター＝エッジ、露光＝ショット単位、洗浄ノズル＝中心集中、"
        "乾燥＝渦巻き、注入/成膜/エッチの面内分布＝カラーマップ、搬送＝エッジ欠け、チャンバー＝クラスタ）。画像内の文字は英数字のみ（ロットNo. #スロット、DEF=個数 など）。",
        "マップの名称（キャプション）と選別検査は異常内容の測定項目に合わせる（スクラッチ→欠陥マップ、パーティクル→パーティクルマップ、CD/重ね合わせ→NG点マップ、"
        "膜厚/Rs/E/R→面内分布マップ）。アラーム・漏液など測定値でない異常は「保留品〇〇検査結果」として保留ウェーハの検査マップを貼る。",
        "発見方法は、異常内容に「後工程(リソ)から…の連絡」「巡回時」とある場合はそちらを優先し（元トラブルの detected_by とずれることがある）、後工程連絡なら流出有無＝有。",
        "",
        "## 意図的なゆれ・イレギュラー",
        "- 版によってラベル名・配置・連絡No.体系が異なる（A4縦／A3横、Rev.1→Rev.2 のラベル改称）。`_expected.jsonl` の `labels_used` に各ファイルの実ラベルを記載。",
        "- Rev.1 は処置区分のラベルが「処置」で、ロット表の列見出し「処置」と同じ文字列（ラベルの取り違えが起きやすい）。承認・確認は発行側と回答側で同じラベル名。",
        "- 設備欄は「CMP-103 W-CMP 3号機」「CMP-103（W-CMP 3号機）」のように設備IDと設備名が1セルに同居。",
        "- 日付が「日付セル＋表示形式（yyyy/m/d、yyyy\"年\"m\"月\"d\"日\"、和暦 ggge 等）」のファイルと「文字列（R6.8.25、2024年8月25日、24/8/25 等）」のファイルが混在。",
        "- 回答欄が空（未回答）、一部のみ（保留案件：なぜなぜ空・恒久対策は後日回答）、期限超過の督促メモ付き。",
        "- 回答欄だけ別フォント・別色（受信部署が記入）。A4縦版の一部は「回答日：…　回答部署：…　回答者：…」を1セルに記入（セル内ラベル）。",
        "- 複数品種ロットでは品名が「複数（ロット一覧参照）」「複数品種」、品番が「―」「複数」など。",
        "- 5ロット以上は A4縦版のロット表（4行）に入りきらず、別シート「ロット一覧」に記載（合計は別紙含む）。",
        "- A4縦版の合計欄の一部は `=SUM()` 式（openpyxl 保存のためキャッシュ値なし。data_only で読むと None）。期待値は人が計算した合計。",
        "- A3横版の一部は数量を「25枚」「２５枚」の文字列で記入。訂正履歴をセルコメントで残した例あり。",
        "- 異常内容ラベルが縦書き（text_rotation=255）のファイルあり。押印欄は「姓\\n日付」。承認印の押し忘れ（空欄）あり。",
        "- 欠陥マップ画像なし（アラーム停止で検査前に保留）、2枚貼付のファイルあり。",
        "- シート名が `Sheet1` のまま、テンプレートの「記入要領」シートが残っている（先頭にある）ファイルあり。",
        "- ファイル名の付け方がばらばら（`_回答済` 付き、日付始まり、【工程異常】始まり 等）。",
        "- 記入者の書き癖（全角数字・半角カナ・誤変換・箇条書き記号の違い）は domain の writer habit を流用。",
        "",
        "イレギュラーの内訳（ファイル数）:",
        "",
    ]
    lines += [f"- {k}: {v}" for k, v in flags.most_common()]
    lines += [
        "",
        "## _expected.jsonl",
        "1行1ファイル: `{\"file\", \"layout_version\", \"source_incident_id\", \"answer_status\", \"images\", \"irregularities\", \"values\", \"labels_used\"}`。",
        f"- `file` は版フォルダを含む相対パス（`{rows[0]['file']}`）。区切りは `/` 固定で、Windows でもそのまま `フォルダ / file` で開ける。",
        "- `values` は人が帳票を見て読み取る値（表示どおりの文字列）。ロット・処置区分・なぜなぜはリスト。未記入欄は空文字。",
        "- 押印欄（approver / checker / creator / answer_* / qa_check）は姓のみ（押印日は含めない）。",
        "- キー名は他帳票（F1〜F3）とそろえ、発行者（発信者）は `reporter`、設備トラブル報告No. は `related_incident_id`、影響ロットは `lots`。"
        "F5 固有の項目は `issuing_dept` / `to_dept` / `cc` / `product_name` / `product_code` / `qty_input` / `qty_defect` / `qty_hold` / "
        "`cause_estimated` / `disposition`（■の付いた区分のリスト）/ `disposition_detail` / `outflow` / `outflow_detail` / `answer_*` / `close_judgement` など。",
        "- `cause` / `action` / `prevention` などは回答欄（宛先部署記入）の値。発行側の推定原因は `cause_estimated`。",
        "- `equipment_id` と `equipment_name` は同じ「設備」「使用設備」セルから読み分ける。",
        "- 承認者のうち製造課長・生産技術課長・品質保証課長は domain の人員マスタに無いため本帳票で補った架空名（佐々木／山本／吉田）。",
        "",
        "## ファイル一覧",
        "| file | layout_version | 元トラブル | 設備 | 回答状況 |",
        "|---|---|---|---|---|",
    ]
    lines += [f"| {x['file']} | {x['layout_version']} | {x['source_incident_id']} | {x['values'].get('equipment_id', '')} | {x['answer_status']} |" for x in rows]
    return "\n".join(lines) + "\n"


# ---------------------------------------------------------------------------
# エントリポイント
# ---------------------------------------------------------------------------
def generate(output_root: Path | None = None) -> list[Path]:
    root = Path(output_root) if output_root is not None else D.OUTPUT_ROOT
    out_dir = root / "forms" / FORM_FOLDER
    out_dir.mkdir(parents=True, exist_ok=True)
    # 前回の生成物を消す。版フォルダの中の .xlsx も消し、空になった版フォルダ自体も消す
    # （版の名前を変えたときに古いフォルダが残らないようにする）。_README.md と _expected.jsonl は下で作り直す。
    for old in list(out_dir.rglob("*.xlsx")) + [out_dir / "_expected.jsonl", out_dir / "_README.md"]:
        if old.exists():
            old.unlink()
    for sub in sorted((p for p in out_dir.rglob("*") if p.is_dir()), key=lambda p: -len(p.parts)):
        if not any(sub.iterdir()):
            sub.rmdir()

    incs = select_incidents()
    # 連絡No.の連番: 影響ロットのあるトラブルのうち約半数が連絡票になった想定で、年内の通し番号を振る
    all_lot = [i for i in D.standard_incidents() if i.lots_affected]
    year_rank: dict[str, int] = {}
    cnt: dict[int, int] = {}
    for i in all_lot:
        y = i.reported_at.year
        cnt[y] = cnt.get(y, 0) + 1
        year_rank[i.incident_id] = cnt[y]
    # 期間末の直近2件（完了）と、中盤の1件は未回答にする（期限前・督促中）
    done_idx = [n for n, i in enumerate(incs) if i.status == "完了"]
    force = {done_idx[-1]: "期限前", done_idx[-2]: "期限前", done_idx[len(done_idx) // 2]: "督促中"}
    # 欠陥マップなし（アラーム停止で検査前に保留）2件、マップ2枚貼り（旧様式・複数ロット）2件
    rq = D.rng("f5:quota")
    cand = [n for n, i in enumerate(incs) if _effective_detection(i) == "装置アラーム" and i.category != "品質" and not _metric_rule(i)]
    no_map_idx = set(rq.sample(cand, k=min(2, len(cand))))
    cand = [n for n, i in enumerate(incs) if len(i.lots_affected) >= 2 and i.reported_at.date() < REV3_START - timedelta(days=5)
            and n not in no_map_idx]
    two_map_idx = set(rq.sample(cand, k=min(2, len(cand))))
    # 数量の訂正コメント（旧様式・廃棄ありの案件から2件）
    cand = [n for n, i in enumerate(incs) if i.reported_at.date() < REV3_START - timedelta(days=5) and i.scrap_wafers > 0 and n not in two_map_idx]
    comment_idx = set(rq.sample(cand, k=min(2, len(cand))))

    paths: list[Path] = []
    rows: list[dict] = []
    used_names: set = set()
    used_ids: set = set()
    for n, inc in enumerate(incs):
        seq = int(year_rank[inc.incident_id] * 0.48) + 1
        rec = _build_record(n, inc, seq, force.get(n), n in no_map_idx, n in two_map_idx)
        rec.qty_comment = n in comment_idx
        while rec.report_id in used_ids:
            seq += 1
            rec = _build_record(n, inc, seq, force.get(n), n in no_map_idx, n in two_map_idx)
            rec.qty_comment = n in comment_idx
        used_ids.add(rec.report_id)
        r = D.rng(f"f5:render:{inc.incident_id}")
        # Rev.3 改訂後も、製造課の一部の班は旧様式（A3横 Rev.2）を使い続けている
        if rec.version == "A" and rec.issuer.section in ("製造課 C班",) and r.random() < 0.7:
            rec.version, rec.rev = "B", "Rev.2"
            rec.report_id = f"工異{rec.report_date.year % 100}-{seq:03d}"
            rec.due = _next_weekday(rec.report_date + timedelta(days=10))
            rec.flags.append("Rev.3改訂後に旧様式(A3横Rev.2)で起票")
        wb, V, L, n_img = (_render_a if rec.version == "A" else _render_b)(rec, r)
        layout = "A4縦_Rev.3" if rec.version == "A" else f"A3横_{rec.rev}"
        folder = VERSION_FOLDERS[layout]
        name = _file_name(rec, r)
        while name in used_names:
            name = name.replace(".xlsx", "_2.xlsx")
        used_names.add(name)
        (out_dir / folder).mkdir(parents=True, exist_ok=True)
        path = out_dir / folder / name
        created = datetime.combine(rec.report_date, datetime.min.time()) + timedelta(hours=9, minutes=r.randint(0, 480))
        modifier = rec.answerer if rec.answer_state != "未回答" else ""
        if modifier and " " not in modifier:
            modifier = next((p.name for p in D.people() if p.name.startswith(modifier.split("（")[0])), modifier)
        _save_deterministic(wb, path, created, rec.issuer.name, modifier.split("（")[0])
        paths.append(path)
        if rec.answer_state == "一部回答":
            rec.flags.append("回答欄が一部のみ記入（保留案件）")
        elif rec.answer_state == "未回答":
            rec.flags.append("回答欄が空（未回答）")
        if len({x.lot[:2] for x in rec.lots}) > 1:
            rec.flags.append("複数品種ロットで品名・品番が「複数」表記")
        rows.append(dict(file=f"{folder}/{name}", layout_version=layout, source_incident_id=inc.incident_id, answer_status=rec.answer_state,
                         images=n_img, irregularities=list(dict.fromkeys(rec.flags)),
                         values=V, labels_used={k: L[k] for k in V if k in L}))

    exp = out_dir / "_expected.jsonl"
    exp.write_text("".join(json.dumps(x, ensure_ascii=False) + "\n" for x in rows), encoding="utf-8", newline="\n")
    readme = out_dir / "_README.md"
    readme.write_text(_readme(rows), encoding="utf-8", newline="\n")
    return paths + [exp, readme]


if __name__ == "__main__":
    import time

    t0 = time.time()
    written = generate()
    print(f"{len(written)} files written to {written[-1].parent}  ({time.time() - t0:.1f}s)")
