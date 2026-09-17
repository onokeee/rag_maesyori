"""T4 部品交換・保全作業記録（1シートに複数の表が並ぶ“現場のExcel”）を生成する。

    python -m scripts.samples.t4_maintenance_blocks

出力:
    samples/tables/T4_保全作業記録_部品交換.xlsx
    samples/tables/T4_保全作業記録_部品交換_README.md

- 直近18ヶ月（2025年3月〜2026年8月）を月別シートにし、各シートにライン別のブロック
  （■ L1ライン …）を空行区切りで縦に並べる。ブロックごとに見出し行・データ行・小計行を持つ。
- 行の元データは domain.standard_incidents() の parts_used（突発の部品交換）と、
  ここで決定的に生成する定期交換（PM）計画。トラブルIDや設備ID・人名は帳票系サンプルと共通。
- 見出し文言のゆれ、2段見出しの結合セル、非表示列（内部コード）、セル内改行、
  全角数字・"2個" のような文字列数値、〃（同上）、左右2表のブロックなどを意図的に混ぜる。
- 乱数は domain.rng() からのみ取り、xlsx 内のタイムスタンプも固定するので再実行で同一バイト列になる。
"""
from __future__ import annotations

import io
import re
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import domain as D

STEM = "T4_保全作業記録_部品交換"
N_MONTHS = 18
FIXED_TS = datetime(2026, 9, 1, 8, 30, 0)        # xlsx に書く作成・更新日時（固定）
FORMAT_REVISION = (2025, 10)                     # この月から様式改訂（品番列追加・内部コード列の位置変更）

# ---------------------------------------------------------------------------
# 列定義・見出し文言のゆれ
# ---------------------------------------------------------------------------
KEYS_V1 = ["date", "eq", "eqname", "kind", "ref", "part", "qty", "price", "amount", "content", "workers", "hours",
           "remarks", "code"]
KEYS_V2 = ["date", "code", "eq", "eqname", "kind", "ref", "part", "partno", "qty", "price", "amount", "content",
           "workers", "hours", "remarks"]

LABELS = {
    "date": ["日付", "作業日", "実施日"],
    "code": ["内部コード", "内部コード", "内部ｺｰﾄﾞ"],
    "eq": ["設備ID", "設備No", "装置No"],
    "eqname": ["設備名", "装置名", "設備名称"],
    "kind": ["区分", "作業区分", "保全区分"],
    "ref": ["管理No", "報告書No", "トラブルNo/PM No"],
    "part": ["部品名", "交換部品", "交換部品名"],
    "partno": ["品番", "部品番号", "品番"],
    "qty": ["数量", "個数", "数量(個)"],
    "price": ["単価(円)", "単価", "単価（円）"],
    "amount": ["金額(円)", "金額", "部品代(円)"],
    "content": ["作業内容", "作業内容・処置", "作業内容（詳細）"],
    "workers": ["作業者", "担当", "実施者"],
    "hours": ["時間(h)", "作業時間(h)", "工数(h)"],
    "remarks": ["備考", "備考・特記", "備考"],
}
WIDTH = {"date": 10, "code": 22, "eq": 10, "eqname": 20, "kind": 9, "ref": 16, "part": 28, "partno": 12, "qty": 7,
         "price": 11, "amount": 12, "content": 48, "workers": 18, "hours": 8, "remarks": 26}
# 2段見出しでまとめる列グループ（上段ラベル候補, 対象列）
HEADER_GROUPS = [(["交換部品", "部品", "使用部品"], ("part", "partno", "qty", "price", "amount")),
                 (["作業", "作業実績", "作業"], ("content", "workers", "hours"))]

# ---------------------------------------------------------------------------
# 書式
# ---------------------------------------------------------------------------
THIN = Side(style="thin", color="808080")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILLS = [PatternFill("solid", fgColor=c) for c in ("DDEBF7", "E2EFDA", "FFF2CC", "EDEDED")]
SUB_FILL = PatternFill("solid", fgColor="FCE4D6")
TOTAL_FILL = PatternFill("solid", fgColor="D9D9D9")
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
TOP = Alignment(vertical="top")
BOLD = Font(bold=True)
TITLE_FONT = Font(bold=True, size=14)
BLOCK_FONT = Font(bold=True, size=12, color="1F3864")
NOTE_FONT = Font(size=9, color="595959")

# ---------------------------------------------------------------------------
# ブロック（ライン）定義
# ---------------------------------------------------------------------------
BLOCK_ORDER = ["L1", "L2", "L3", "L4", "L5", "L6", "T1", "T2", "UT"]
BLOCK_TITLES = {
    **{f"L{i}": [f"■ L{i}ライン", f"■L{i}ライン", f"■ L{i}ライン（{'Fab1' if i <= 3 else 'Fab2'}）", f"■ L{i} ライン"]
       for i in range(1, 7)},
    "T1": ["■ 搬送系（Fab1 OHT/AGV）", "■ 搬送系 Fab1", "■搬送 Fab1（OHT・AGV）"],
    "T2": ["■ 搬送系（Fab2 OHT/AGV/ストッカ）", "■ 搬送系 Fab2", "■搬送 Fab2（OHT・AGV・STK）"],
    "T": ["■ 搬送系（Fab1/Fab2）", "■ 搬送系", "■搬送系 全体"],
    "UT": ["■ ユーティリティ（施設係）", "■ 共通設備・ユーティリティ", "■ユーティリティ"],
    "OT": ["■ その他（ライン外・予備品整備）", "■ その他", "■その他（ライン外）"],     # 様式にだけある枠。記入はほぼ無い
}


def _block_key(eq: D.Equipment) -> str:
    if eq.line in ("L1", "L2", "L3", "L4", "L5", "L6"):
        return eq.line
    if eq.line == "L1-L3":
        return "T1"
    if eq.line == "L4-L6":
        return "T2"
    return "UT"


# ---------------------------------------------------------------------------
# 行データ
# ---------------------------------------------------------------------------
@dataclass
class WorkRow:
    """シート上の1行（部品1点 または 部品なし作業1件）。表示ゆれは書き出し時に付ける。"""
    work_date: date
    hour: int
    eq: D.Equipment
    kind: str                 # 突発 / 定期 / 予防
    ref: str                  # トラブルID or PM指示No（空もあり）
    part: D.Part | None
    qty: int | None
    price: int | None         # None = 契約内などで空欄
    content: str
    workers: list
    hours: float | None
    remarks: str
    code: str
    source: str               # "incident" / "pm"
    cont: bool = False        # 同一作業の2行目以降（〃 の対象）
    no_part_label: str = ""   # 部品なし作業の表示（"－（調整のみ）" など）
    severity: str = ""
    contract: bool = False    # 保守契約内で単価・金額を空欄にした行
    order: int = 0            # 同一作業内の並び順

    @property
    def amount(self) -> int | None:
        if self.part is None or self.qty is None or self.price is None:
            return None
        return self.price * self.qty


def _months() -> list[tuple[int, int]]:
    y, m = D.PERIOD_END.year, D.PERIOD_END.month
    out = []
    for _ in range(N_MONTHS):
        out.append((y, m))
        y, m = (y, m - 1) if m > 1 else (y - 1, 12)
    return out[::-1]


def _unit_word(part: D.Part | None) -> str:
    """数量に付ける助数詞（"2個" "4本" のような文字列数値の生成用）。"""
    if part is None:
        return "個"
    n = part.name
    for words, u in ((("樹脂", "キット", "(1式)"), "式"), (("フィルタ", "ランプ", "ホース", "ケーブル", "ベルト", "Oリング", "ボンベ",
                                                         "熱電対", "ワイヤ", "チューブ", "ロール"), "本"),
                     (("パッド", "窓", "ディスク"), "枚"), (("テープ",), "巻"), (("セット",), "組"), (("冷媒",), "缶")):
        if any(w in n for w in words):
            return u
    return "個"


def _core_name(name: str) -> str:
    """部品名の本体（"研磨パッド IC-type 30inch" -> "研磨パッド"）。"""
    return re.split(r"[ (（]", name)[0]


# ---------------------------------------------------------------------------
# 突発（トラブル）由来の行
# ---------------------------------------------------------------------------
_NUM_PREFIX = re.compile(r"^\s*(?:[\d０-９]+[.)．）]\s*|[①-⑳]\s*|[(（][\d０-９]+[)）]\s*|・\s*)")
_ADMIN_WORDS = ("DOWN解除", "復旧連絡", "生産へ引渡し", "員数確認", "保護具着用", "LOTO実施", "（継続対応中）")


def _action_lines(action: str) -> list[str]:
    out = [_NUM_PREFIX.sub("", ln).strip() for ln in action.splitlines()]
    return [ln for ln in out if ln]


def _incident_content(r, inc: D.Incident, part: D.Part | None, first: bool, ditto_p: float) -> str:
    """作業内容セル。トラブル報告の処置欄から部品に関係する手順を抜き出し、セル内改行でつなぐ。"""
    lines = _action_lines(inc.action)
    body = [ln for ln in lines if not any(w in ln for w in _ADMIN_WORDS)] or lines
    core = _core_name(part.name) if part else ""
    hit = [ln for ln in body if core and (core in ln or (len(core) >= 6 and core[2:] in ln))]
    if not first:
        x = r.random()
        if x < ditto_p:
            return r.choice(["〃", "〃", "同上"])
        if hit and r.random() < 0.6:
            return hit[0]
        return r.choice([f"{core}同時交換", f"{core}も交換", f"{core}交換", f"同時に{core}交換"])
    k = r.choice([2, 2, 3, 3, 4]) if len(body) > 2 else len(body)
    sel = body[:k]
    if hit and hit[0] not in sel:
        sel = sel[: max(1, k - 1)] + [hit[0]]
    if inc.status == "対応中" and "（継続対応中）" in inc.action:
        sel.append("（継続対応中）")
    head = ""
    if r.random() < 0.28:
        s1 = re.split(r"[。]", inc.symptom)[0]
        if 6 <= len(s1) <= 42:
            head = r.choice(["【現象】", "現象：", "", "◆"]) + s1
    elif inc.alarm is not None and r.random() < 0.15:
        head = f"{inc.alarm.code} {inc.alarm.message}"
    out = ([head] if head else []) + sel
    return "\n".join(out)


def _incident_remarks(r, inc: D.Incident, rows_parts: list) -> str:
    notes = []
    if inc.status == "対応中":
        notes.append(r.choice(["対応中", "対応中（部品先行交換）", "継続対応中", "未完了"]))
    elif inc.status == "経過観察":
        notes.append(r.choice(["経過観察中", "経過観察", "様子見（再発時連絡）"]))
    if inc.recurrence and inc.related_incident_id and r.random() < 0.5:
        notes.append(r.choice([f"再発（前回 {inc.related_incident_id}）", f"再発 {inc.related_incident_id}参照", "再発"]))
    if inc.scrap_wafers and r.random() < 0.35:
        notes.append(f"スクラップ{inc.scrap_wafers}枚")
    if any(p is not None and p.unit_price_yen >= 1_000_000 for p, _ in rows_parts) and r.random() < 0.6:
        notes.append(r.choice(["予備品払出し、補充発注済み", "高額部品のため課長承認済み", "リビルド品を使用", "予備品から払出し"]))
    if inc.remarks and any(w in inc.remarks for w in ("在庫", "発注", "転用", "見積", "保守契約")) and r.random() < 0.5:
        notes.append(inc.remarks)
    if not notes and r.random() < 0.12:
        notes.append(r.choice(["重要度" + inc.equipment.criticality, f"{inc.severity}トラブル", "報告書提出済み", "報告書未提出"]))
    if len(notes) > 2:
        notes = notes[:2]
    return r.choice(["／", "\n", "、"]).join(notes)


def _incident_rows(months: set) -> list[WorkRow]:
    r = D.rng("t4:incident-rows")
    rows: list[WorkRow] = []
    for inc in D.standard_incidents():
        wd = inc.response_started_at
        if (wd.year, wd.month) not in months:
            continue
        if inc.parts_used:
            parts = list(inc.parts_used)
        elif inc.severity in ("大", "重大") and inc.status in ("完了", "経過観察"):
            parts = [(None, None)]        # 部品交換なしの大規模作業も保全作業として記録する
        else:
            continue
        eid = inc.equipment.equipment_id
        code = f"K531-{eid.replace('-', '')}-{inc.incident_id[3:7]}{inc.incident_id[-5:]}"
        remarks = _incident_remarks(r, inc, parts)
        fe = [p for p in inc.assignees if p.role == "FE"]
        for i, (part, qty) in enumerate(parts):
            first = i == 0
            contract = bool(part is not None and fe and part.maker == inc.equipment.maker and r.random() < 0.3)
            rows.append(WorkRow(
                work_date=wd.date(), hour=wd.hour, eq=inc.equipment,
                kind="突発", ref=inc.incident_id, part=part, qty=qty,
                price=None if (part is None or contract) else part.unit_price_yen,
                content=_incident_content(r, inc, part, first, 0.3),
                workers=list(inc.assignees), hours=inc.work_hours if first else None,
                remarks=remarks if first else ("" if r.random() < 0.9 else "同上"),
                code=code, source="incident", cont=not first, severity=inc.severity, contract=contract, order=i,
                no_part_label=r.choice(["－（調整のみ）", "なし（調整・清掃のみ）", "部品交換なし", "－"]) if part is None else "",
            ))
    return rows


# ---------------------------------------------------------------------------
# 定期交換（PM）計画
# ---------------------------------------------------------------------------
@dataclass(frozen=True)
class PMItem:
    key: str
    part: str | None          # カタログ部品名。None は部品なし作業
    work: str                 # 部品なし作業の名称
    interval: int             # 交換周期（月）
    qty: tuple                # 数量の範囲
    hours: float              # 標準工数（1人あたり）
    acts: tuple = ()          # 作業内容1行目の候補
    conds: tuple = ()         # 旧品状態・測定値の候補
    fe: bool = False          # メーカーFE立会い
    fleet: tuple | None = None   # 搬送車両系: 月あたり対象台数の範囲（interval は使わない）
    only: str = ""            # 対象設備の絞り込み（設備名/型式に含まれる語、"!語" は除外）


def _pm(key, part, interval, qty, hours, acts=(), conds=(), fe=False, fleet=None, only="", work=""):
    return PMItem(key, part, work, interval, qty, hours, tuple(acts), tuple(conds), fe, fleet, only)


PM_ITEMS: dict[str, list[PMItem]] = {
    "CMP": [
        _pm("pad", "研磨パッド IC-type 30inch", 1, (1, 1), 2.0,
            ["研磨パッド交換（使用{pad}枚）", "パッド貼替\nブレークイン ダミー{n}0枚", "定期パッド交換（ライフ{pad}枚到達）", "Platen{pl} パッド交換"],
            ["溝深さ残 {groove}mm（交換基準0.30mm）", "旧パッド表面にグレージングあり", "パッド中央部の摩耗大、エッジ部は問題なし", "貼付け後の気泡なし"]),
        _pm("dresser", "ドレッサディスク (ダイヤ電着)", 3, (1, 1), 1.0,
            ["ドレッサディスク交換", "ドレッサ交換（使用{dress}h）", "コンディショナ ディスク定期交換"],
            ["旧品 ダイヤ脱落なし（ルーペ確認）", "砥粒の摩耗大", "ダウンフォース校正 {load}N", "揺動範囲ティーチング確認"]),
        _pm("ring", "リテーナリング", 3, (1, 1), 2.5,
            ["リテーナリング交換", "ヘッド分解、リテーナリング交換", "RR交換（ヘッド{hd}）"],
            ["リング厚み残 {ring}mm（基準1.5mm）", "溝部の摩耗・欠けなし", "旧品に微小クラック1箇所"]),
        _pm("membrane", "ヘッドメンブレン", 2, (1, 1), 2.0,
            ["メンブレン交換", "ヘッドメンブレン定期交換（ヘッド{hd}）", "メンブレン交換、ゾーン圧校正"],
            ["ゾーン加圧リークテストOK", "旧品 微小亀裂なし", "エッジゾーン部に擦れ跡"]),
        _pm("pou", "POUフィルタ 0.5μm", 1, (2, 4), 0.5,
            ["POUフィルタ交換（{qty}本）", "スラリーPOUフィルタ定期交換", "POUフィルタ交換後DIWフラッシング{n}0分"],
            ["差圧 {dp}kPa→{dp2}kPa", "旧フィルタ内に凝集物少量", "流量 SV150ml/min 安定"]),
        _pm("brush", "PVAブラシロール", 2, (2, 2), 1.5,
            ["後洗浄PVAブラシ交換（上下）", "ブラシロール交換、ブレークイン実施", "PVAブラシ交換"],
            ["旧ブラシ表面の目詰まりあり", "ブラシ回転トルク正常", "ブラシギャップ再調整"]),
        _pm("oring", "ヘッドOリングセット", 6, (1, 1), 1.0, ["ヘッドOリング一式交換", "Oリングセット交換（6M PM）"]),
        _pm("subpad", "サブパッド", 6, (1, 1), 1.0, ["サブパッド交換（パッド交換と同時）", "サブパッド貼替"]),
        _pm("grease", None, 3, (0, 0), 1.0, ["揺動軸・プラテン軸受 グリスアップ", "3ヶ月点検：グリスアップ、ベルト張り確認"],
            ["異音なし", "ベルト張力 規定内"], work="グリスアップ・点検"),
    ],
    "CVD": [
        _pm("oring", "チャンバーOリング (FKM)", 3, (2, 6), 3.0,
            ["チャンバーウェットクリーニング\nOリング交換（{qty}本）", "CH-{chx} 大気開放PM、Oリング交換", "定期PM（ウェットクリーニング）"],
            ["リークレート {lr}mTorr/min", "旧Oリングに圧縮永久ひずみあり", "シャワーヘッド穴詰まりなし"]),
        _pm("ceramic", "セラミックリング", 6, (1, 1), 2.0, ["セラミックリング交換", "CH-{chx} セラミックリング定期交換"],
            ["旧品 表面に堆積物、欠けなし"]),
        _pm("gasket", "VCRガスケット 1/4 (Ni)", 6, (4, 12), 1.5, ["ガスライン継手 定期増締め・ガスケット交換（{qty}箇所）", "VCRガスケット交換"],
            ["Heリークチェック 検出下限以下", "継手部に変色なし"]),
        _pm("tc", "熱電対 K型 (シース)", 12, (1, 2), 1.0, ["熱電対 定期交換", "ヒーター熱電対交換、温度校正"], ["校正後の指示差 {tcd}℃"]),
        _pm("fan", "RF電源 冷却ファン", 12, (1, 2), 0.75, ["RF電源冷却ファン交換", "RF電源ファン定期交換（年次）"], ["旧品 軸受異音あり", "回転数正常"]),
        _pm("pin", "リフトピン (サファイア)", 6, (3, 3), 2.0, ["リフトピン交換（3本）", "リフトピン交換、高さ調整"], ["ピン先端に摩耗", "昇降速度確認OK"]),
        _pm("bellows", "DP排気配管 ベローズ", 12, (1, 1), 2.0, ["DP排気配管ベローズ交換", "排気ベローズ年次交換"], ["内部に副生成物堆積（写真保存）"]),
        _pm("dpcheck", None, 3, (0, 0), 1.0, ["ドライポンプ定期点検（電流・温度・振動）", "DP点検：電流{amp}A、ケーシング温度{dpt}℃"],
            ["振動値 {vib}mm/s 正常"], work="DP点検"),
    ],
    "ETC": [
        _pm("fr", "フォーカスリング (Si)", 2, (1, 1), 3.0,
            ["フォーカスリング交換（RF {rfh}h）", "FR交換\nシーズニング{n}0枚", "{ch} フォーカスリング定期交換"],
            ["旧品厚み {fr}mm", "エロージョン進行、段差{step}mm", "リング内周にデポ付着"]),
        _pm("electrode", "上部電極 (Si)", 6, (1, 1), 5.0, ["上部電極交換", "{ch} 上部電極交換（RF {rfh}h）"],
            ["ガス穴の拡大あり（写真保存）", "電極裏面の冷却シート同時確認"], fe=True),
        _pm("shield", "デポシールド", 3, (1, 1), 3.0, ["デポシールド交換（ウェットクリーニング）", "シールド交換、チャンバー清掃"],
            ["堆積膜の剥離なし", "パーティクル {pc}個"]),
        _pm("hefilter", "He供給ライン フィルタ", 6, (1, 1), 0.5, ["He供給ラインフィルタ交換", "裏面He フィルタ定期交換"], ["He流量 {he}sccm 安定"]),
        _pm("oring", "チャンバーOリング (パーフロ)", 3, (2, 4), 1.5, ["Oリング交換（パーフロ {qty}本）", "チャンバーOリング定期交換"],
            ["ROR {lr}mTorr/min"]),
        _pm("oes", "OESビューポート窓", 6, (1, 1), 1.0, ["OES窓交換", "EPD用ビューポート窓 交換"], ["窓の曇りによる光量低下 {oes}%"]),
        _pm("pin", "リフトピン (サファイア)", 6, (3, 4), 1.5, ["リフトピン交換", "リフトピン交換、ピン高さ調整"]),
        _pm("coolant", "チラー 冷媒 (フッ素系) 10L", 12, (1, 2), 1.0, ["チラー冷媒 入替", "チラー冷媒補充・入替（年次）"], ["比抵抗・液色 正常"]),
    ],
    "LIT": [
        _pm("hepa", "温調ユニット HEPAフィルタ", 12, (1, 2), 3.0, ["温調チャンバーHEPA交換", "HEPAフィルタ年次交換"],
            ["差圧 {dp}Pa→{dp2}Pa", "チャンバー内パーティクル {pc}個/cf"]),
        _pm("lmfilter", "リニアモータ冷却水フィルタ", 6, (1, 1), 1.0, ["リニアモータ冷却水フィルタ交換", "ステージ冷却水フィルタ定期交換"],
            ["流量 {flow}L/min"]),
        _pm("gas", "レーザガス (F2/Ar/Ne混合) ボンベ", 3, (1, 1), 1.0, ["レーザガス交換（ボンベ）", "光源ガス定期交換、ガスリフレッシュ"],
            ["交換後パルスエネルギー安定", "ショット数 {shot}億"], fe=True, only="KrF|ArF"),
        _pm("lamp", "アライメント照明ランプ", 6, (1, 1), 1.5, ["アライメント照明ランプ交換", "ALG照明ランプ交換、光量調整"], ["光量 交換前{oes}%"]),
        _pm("stone", "チャッククリーニングストーン", 3, (1, 1), 0.5, ["チャッククリーニングストーン交換", "ストーン交換、チャック清掃"]),
        _pm("scale", "ステージ エンコーダスケール清掃キット", 6, (1, 1), 2.0, ["エンコーダスケール清掃（定期）", "スケール清掃、ステージキャリブレーション"],
            ["信号強度 {sig}%→98%"], fe=True, only="KrF|ArF"),
        _pm("gripper", "レチクルグリッパパッド", 6, (2, 4), 1.0, ["レチクルグリッパパッド交換", "RH グリッパパッド交換"]),
    ],
    "CLN": [
        _pm("filter", "薬液フィルタ 0.05μm", 2, (1, 3), 1.0, ["薬液フィルタ交換（{chem}ライン）", "{chem}循環フィルタ定期交換"],
            ["差圧 {dp}kPa→{dp2}kPa", "交換後パーティクル {pc}個"]),
        _pm("pin", "スピンチャックピン", 3, (3, 6), 1.5, ["スピンチャックピン交換", "チャックピン交換（{qty}本）"], ["ピン先端摩耗あり"], only="枚葉"),
        _pm("n2", "N2フィルタ (乾燥部)", 6, (1, 1), 0.5, ["乾燥部N2フィルタ交換"], only="枚葉"),
        _pm("nozzle", "薬液ノズル (PFA)", 6, (1, 2), 1.0, ["薬液ノズル交換", "ノズル交換、吐出位置ティーチング"], ["ノズル先端に結晶付着"], only="枚葉"),
        _pm("joint", "PFA配管継手 1/2", 12, (2, 6), 2.0, ["PFA継手 年次交換・増締め", "配管継手交換（{qty}箇所）、加圧試験"], ["加圧試験 {n}時間保持OK"]),
        _pm("claw", "ロボットチャック爪 (PEEK)", 6, (2, 4), 1.0, ["ロボットチャック爪交換", "搬送ロボ爪交換、ティーチング確認"]),
        _pm("heater", None, 6, (0, 0), 1.5, ["槽ヒーター・過昇温センサ点検", "薬液槽 温度校正"], ["温度指示差 {tcd}℃"], work="槽点検", only="RCA|HF"),
    ],
    "IMP": [
        _pm("filament", "フィラメント (タングステン)", 1, (1, 1), 2.0, ["フィラメント交換（ソース{hrs}h）", "イオン源 フィラメント交換"],
            ["旧品 断面細り", "アーク電流安定"]),
        _pm("arc", "アークチャンバー", 4, (1, 1), 4.0, ["アークチャンバー交換（ソースPM）", "ソースPM：アークチャンバー交換、碍子清掃"], ["デポ付着多い"]),
        _pm("liner", "グラファイトライナー", 6, (1, 2), 4.0, ["グラファイトライナー交換", "ビームラインPM：ライナー交換"], ["スパッタ痕あり"]),
        _pm("electrode", "引出し電極", 6, (1, 1), 3.0, ["引出し電極交換", "引出し電極交換、ギャップ調整"], ["電極ギャップ {gap}mm"]),
        _pm("gasket", "ガスシリンダ ガスケット", 2, (1, 2), 1.0, ["ガスシリンダ交換時ガスケット交換", "ガスBOX シリンダ交換（{gas}）"],
            ["リークチェックOK（ガス検知器指示なし）"]),
        _pm("adsorber", "クライオ コンプレッサ吸着器", 12, (1, 1), 2.0, ["クライオ コンプレッサ吸着器 年次交換", "吸着器交換、クライオ再生"]),
        _pm("platen", "プラテン エラストマパッド", 12, (1, 1), 3.0, ["プラテン エラストマパッド交換"], ["冷却性能確認OK"], fe=True),
    ],
    "INS": [
        _pm("xe", "検査光源ランプ (Xe)", 6, (1, 1), 2.0, ["検査光源ランプ交換（点灯{hrs}h）", "Xeランプ交換、光軸調整"],
            ["光量 交換前{oes}%"], only="MI-"),
        _pm("pad", "プリアライナ 吸着パッド", 6, (1, 2), 0.5, ["プリアライナ吸着パッド交換"]),
        _pm("std", "標準ウェーハ (PSL)", 12, (1, 1), 1.0, ["標準ウェーハ更新、感度校正", "PSL標準ウェーハ年次更新"], only="MI-"),
        _pm("hepa", "温調ユニット HEPAフィルタ", 12, (1, 1), 2.0, ["装置HEPA交換"], only="HS-"),
        _pm("calib", None, 3, (0, 0), 1.0, ["定期校正（リファレンスウェーハ測定）", "3ヶ月校正"], ["前回値との差 {dv}%"], work="定期校正"),
    ],
    "OHT": [
        _pm("wheel", "走行ホイール (ウレタン)", 0, (4, 4), 1.0, ["{car}号車 走行ホイール交換（4輪）", "{car}号車 ホイール交換\n走行テスト OK"],
            ["ホイール径 {whl}mm（基準{whl_min}mm）", "旧品 ウレタン剥離あり"], fleet=(3, 7)),
        _pm("belt", "ホイストベルト", 0, (2, 2), 1.5, ["{car}号車 ホイストベルト交換", "{car}号車 ベルト交換、原点調整"],
            ["ベルト伸び {bel}mm", "昇降{n}0回テストOK"], fleet=(1, 3)),
        _pm("gripper", "グリッパユニット", 0, (1, 1), 2.0, ["{car}号車 グリッパユニット交換"], ["把持確認 {n}0回OK"], fleet=(0, 1)),
        _pm("rail", None, 1, (0, 0), 3.0, ["レール清掃・点検（Bay{bay}〜Bay{bay2}）", "月例 走行レール点検、継目段差確認"],
            ["異常なし", "継目段差 {step}mm（基準0.5mm）"], work="レール点検"),
    ],
    "AGV": [
        _pm("wheel", "走行ホイール (ウレタン)", 6, (2, 4), 1.0, ["駆動輪交換", "走行ホイール交換（{qty}個）"]),
        _pm("conv", "移載コンベアベルト", 6, (1, 2), 1.0, ["移載コンベアベルト交換"]),
        _pm("tape", "磁気誘導テープ 10m", 3, (1, 3), 1.5, ["磁気誘導テープ 剥がれ部補修（{qty}本）", "誘導テープ貼替"], ["剥がれ箇所 {n}箇所"],
            only="OT-AGV"),
    ],
    "ROB": [
        _pm("pad", "ロボットハンド吸着パッド", 3, (2, 4), 1.0, ["ハンド吸着パッド交換", "吸着パッド交換、真空圧確認"], ["真空圧 {vac}kPa"]),
        _pm("ejector", "真空エジェクタ", 12, (1, 1), 1.0, ["真空エジェクタ年次交換"]),
        _pm("belt", "ロボット駆動ベルト", 12, (1, 2), 2.0, ["駆動ベルト交換、ティーチング確認", "アーム駆動ベルト交換"]),
    ],
    "STK": [
        _pm("wire", "クレーン走行ワイヤ", 12, (2, 2), 4.0, ["クレーン走行ワイヤ交換", "年次PM：走行ワイヤ交換、張力調整"], ["素線切れ {n}箇所"]),
        _pm("conv", "移載コンベアベルト", 6, (2, 4), 1.5, ["入出庫ポート コンベアベルト交換"]),
        _pm("mfc", "N2パージ用MFC", 24, (1, 2), 1.0, ["N2パージMFC交換（棚{shelf}）"]),
        _pm("crane", None, 3, (0, 0), 3.0, ["クレーン定期点検（走行・昇降・ブレーキ）", "3ヶ月点検：クレーン給脂、ワイヤ張力確認"],
            ["異常なし", "ワイヤ張力 規定内"], work="クレーン点検"),
    ],
    "UPW": [
        _pm("uv", "UVランプ (185nm)", 12, (8, 12), 3.0, ["UV酸化装置 ランプ一斉交換（{qty}本）", "UVランプ年次交換"], ["TOC {toc}ppb"]),
        _pm("ro", "RO膜エレメント", 24, (6, 12), 8.0, ["RO膜エレメント交換（{qty}本）", "RO膜更新"], ["透過水導電率 改善"], fe=True),
        _pm("bearing", "送水ポンプ用ベアリング", 12, (2, 2), 3.0, ["送水ポンプ ベアリング交換", "予備機切替の上 ベアリング交換"], ["振動値 {vib}mm/s"]),
        _pm("resin", "ポリッシャ用イオン交換樹脂 (1式)", 12, (1, 1), 6.0, ["ポリッシャ樹脂 入替", "カートリッジポリッシャ 樹脂交換"],
            ["比抵抗 {mo}MΩ・cm"], fe=True),
        _pm("water", None, 1, (0, 0), 2.0, ["月例 水質分析（比抵抗・TOC・シリカ・微粒子）", "計器指示値の現場照合、記録"],
            ["比抵抗 {mo}MΩ・cm、TOC {toc}ppb で規格内", "全項目規格内"], work="水質点検"),
        _pm("pump", None, 3, (0, 0), 2.0, ["送水ポンプ 振動・電流・軸温度測定", "予備機切替運転（{n}時間）"],
            ["振動値 {vib}mm/s、電流 {amp}A", "異常なし"], work="ポンプ点検"),
    ],
    "EXH": [
        _pm("vbelt", "排気ファン Vベルト (セット)", 6, (1, 1), 2.0, ["Vベルト交換、張り調整", "ファンVベルト定期交換"], ["振動値 {vib}mm/s"]),
        _pm("bearing", "排気ファン軸受", 24, (2, 2), 6.0, ["ファン軸受交換（停止調整済み）"]),
        _pm("vib", None, 1, (0, 0), 1.0, ["ファン振動測定・ベルト点検", "月例点検：振動・軸受温度・静圧"], ["振動値 {vib}mm/s", "排気静圧 異常なし"],
            work="月例点検"),
    ],
    "SCR": [
        _pm("nozzle", "スクラバー スプレーノズル", 3, (4, 8), 2.0, ["スプレーノズル交換（{qty}個）", "ノズル交換・循環槽清掃"], ["pH {ph}"]),
        _pm("ph", "pH計電極", 6, (1, 1), 0.5, ["pH計電極交換、2点校正"]),
        _pm("igniter", "バーナーイグナイタ", 6, (1, 1), 1.0, ["イグナイタ交換", "燃焼部PM：イグナイタ交換"]),
        _pm("flame", "フレームセンサ (UV)", 12, (1, 1), 1.0, ["フレームセンサ交換"]),
        _pm("tank", None, 1, (0, 0), 3.0, ["循環槽清掃・スラッジ除去", "月例PM：循環水入替、ノズル詰まり確認"], ["pH {ph}", "詰まりノズル {n}本 清掃"],
            work="循環槽清掃"),
    ],
    "PCW": [
        _pm("cond", "導電率計センサ", 12, (1, 1), 1.0, ["導電率計センサ交換、校正"]),
        _pm("chem", None, 1, (0, 0), 1.5, ["冷却塔 薬注・水質点検", "月例 冷却水水質分析"], ["導電率 {cond}mS/m"], work="水質点検"),
    ],
    "VAC": [
        _pm("strainer", "封水ストレーナ", 3, (1, 2), 1.0, ["封水ストレーナ交換", "ストレーナ交換・封水量確認"], ["真空圧 {vac}kPa"]),
        _pm("check", None, 1, (0, 0), 1.0, ["真空ポンプ 電流・封水温度点検", "月例点検（真空圧・異音・漏水）"], ["真空圧 {vac}kPa", "異常なし"],
            work="月例点検"),
    ],
    "COMMON": [
        _pm("ups", "UPSバッテリー", 48, (2, 4), 2.0, ["装置UPSバッテリー交換", "UPSバッテリー更新（{qty}個）"], ["放電試験OK"]),
    ],
}
_PM_EXTRA = ["取外し品 外観異常なし", "旧品は廃却", "旧品は{maker}へ返却（リビルド）", "作業前後の写真を共有フォルダへ保存",
             "トルク管理 {torque}N·m", "次回交換予定 {nm}月", "交換履歴台帳へ記入済み", "作業後リークチェックOK", "LOTO実施",
             "予備品残 {stock}", "チェックシート記入済み"]
_PM_HEAD = ["【{iv}M PM】", "定期PM（{iv}ヶ月）", "月例PM", "◇定期交換", "PM"]


def _pm_params(r, eq: D.Equipment, item: PMItem, qty: int) -> D._Params:
    k = D.kind_of(eq)
    fleet_n = 48 if eq.equipment_id == "OHT-801" else 40
    p = D._Params(
        qty=qty, n=r.randint(1, 5), pad=r.randint(780, 1450), groove=f"{r.uniform(0.22, 0.42):.2f}", dress=r.randint(60, 140),
        load=r.randint(20, 45), ring=f"{r.uniform(1.4, 2.3):.1f}", dp=r.randint(18, 45), dp2=r.randint(4, 9), lr=f"{r.uniform(0.3, 1.8):.1f}",
        chx=r.choice("ABC"), rfh=r.randint(280, 520), fr=f"{r.uniform(2.1, 3.4):.2f}", step=f"{r.uniform(0.1, 0.6):.2f}",
        pc=r.randint(3, 28), he=f"{r.uniform(8, 16):.1f}", oes=r.randint(58, 85), flow=f"{r.uniform(3.5, 6.0):.1f}", shot=f"{r.uniform(1.2, 9.8):.1f}",
        sig=r.randint(55, 80), chem=r.choice(["SC1", "SC2", "DHF", "SPM", "DIW"] if "RCA" in eq.process else ["DHF", "DIW", "FPM", "APM"]),
        hrs=r.randint(160, 900), gap=f"{r.uniform(3.8, 5.2):.1f}", gas=r.choice(["PH3", "BF3", "AsH3", "Ar"]), dv=f"{r.uniform(0.01, 0.3):.2f}",
        car=r.randint(1, fleet_n), whl=f"{r.uniform(118.2, 119.6):.1f}", whl_min="118.0", bel=f"{r.uniform(1.2, 3.8):.1f}",
        bay=r.randint(1, 12), vac=-r.randint(70, 88), shelf=f"{r.choice('ABCD')}-{r.randint(1, 40):02d}", toc=f"{r.uniform(0.3, 0.9):.1f}",
        vib=f"{r.uniform(0.8, 2.6):.1f}", mo=f"{r.uniform(18.0, 18.2):.1f}", ph=f"{r.uniform(6.8, 7.6):.1f}", cond=r.randint(35, 80),
        tcd=f"{r.uniform(0.2, 1.5):.1f}", amp=f"{r.uniform(18, 32):.1f}", dpt=r.randint(62, 85), torque=r.choice([8, 12, 15, 20, 25]),
        nm=0, stock=f"{r.randint(0, 6)}", maker=D.MAKER_SHORT.get(eq.maker, eq.maker), iv=item.interval or 1,
        pl=r.randint(1, 2), hd=r.randint(1, 4), ch=r.choice(["PM1", "PM2", "PM3"] if eq.maker == "瑞穂エンジニアリング" else ["CH-A", "CH-B"]),
    )
    p["bay2"] = p["bay"] + r.randint(2, 6)
    p.update({kk: vv for kk, vv in D._qc_params(k, r).items() if kk not in ("n", "pc")} if k in D._QC else {})
    return p


def _item_applies(eq: D.Equipment, item: PMItem) -> bool:
    if item.part is not None and eq.category not in D._part_by_name()[item.part].applicable_categories:
        return False
    if item.only:
        text = f"{eq.name} {eq.model} {eq.process}"
        return any(w in text for w in item.only.split("|"))
    return True


def _workday(r, y: int, m: int, lo: int = 1, hi: int | None = None) -> date:
    """PM は平日に入れることが多い（たまに休日作業）。"""
    last = (date(y + (m == 12), m % 12 + 1, 1) - timedelta(days=1)).day
    hi = min(hi or last, last)
    for _ in range(12):
        d = date(y, m, r.randint(lo, hi))
        if not D._is_holiday(d) or r.random() < 0.08:
            return d
    return d


def _pm_rows(months: list[tuple[int, int]]) -> list[WorkRow]:
    """定期交換計画から PM 行を作る。同一設備・同月の周期到来項目は1つのPM指示（同日・同作業者）にまとめる。"""
    ppl = D.people()
    mset = set(months)
    abs0 = months[0][0] * 12 + months[0][1] - 1
    abs1 = months[-1][0] * 12 + months[-1][1] - 1
    # --- 周期到来の判定（期間の前後1ヶ月も見て、遅れ・前倒しで範囲に入るものを拾う） ---
    tasks: dict = defaultdict(list)       # (eid, y, m) -> [(item, flag)]
    fleet_tasks: list = []
    for eq in D.equipment_master():
        k = D.kind_of(eq)
        items = list(PM_ITEMS.get(k, []))
        if k in D.PROCESS_KINDS and eq.criticality == "A":
            items += PM_ITEMS["COMMON"]
        for item in items:
            if not _item_applies(eq, item):
                continue
            ri = D.rng(f"t4:pm-phase:{eq.equipment_id}:{item.key}")
            if item.fleet:
                for (y, m) in months:
                    rf = D.rng(f"t4:pm-fleet:{eq.equipment_id}:{item.key}:{y}-{m}")
                    lo, hi = item.fleet
                    if eq.equipment_id == "OHT-801":
                        hi += 1          # 古いFab1のOHTは消耗が早い
                    for _ in range(rf.randint(lo, hi)):
                        fleet_tasks.append((eq, item, y, m))
                continue
            phase = ri.randrange(item.interval)
            for a in range(abs0 - 1, abs1 + 2):
                if (a + phase) % item.interval:
                    continue
                rm = D.rng(f"t4:pm-due:{eq.equipment_id}:{item.key}:{a}")
                x = rm.random()
                shift, flag = 0, ""
                if x < 0.07 and item.interval >= 2:
                    shift, flag = 1, "delay"          # 生産都合で翌月へ延期
                elif x < 0.12 and item.interval >= 3:
                    shift, flag = -1, "early"         # 摩耗が早く前倒し（予防交換）
                elif x < 0.15:
                    continue                          # 未実施（記入なし）
                aa = a + shift
                y, m = aa // 12, aa % 12 + 1
                if (y, m) not in mset:
                    continue
                if date(y, m, 28) < eq.installed_date + timedelta(days=45):
                    continue                          # 立上げ直後はメーカー保証期間中のため記録対象外
                tasks[(eq.equipment_id, y, m)].append((item, flag))
    rows: list[WorkRow] = []
    maint = {sect: [p for p in ppl if p.section == sect and p.role in ("主任", "担当")] for sect in
             ("設備保全課 保全1係", "設備保全課 保全2係", "設備保全課 施設係")}

    def crew(r, eq, fe: bool) -> list:
        k = D.kind_of(eq)
        sect = "設備保全課 施設係" if k in D.UTILITY_KINDS else ("設備保全課 保全1係" if D.fab_of(eq) == "Fab1" else "設備保全課 保全2係")
        c = r.sample(maint[sect], k=r.choice([1, 1, 2, 2, 3]))
        if fe:
            c += [p for p in ppl if p.role == "FE" and p.section == eq.maker][:1]
        return c

    def one_row(r, eq, item, d, workers, flag, ref, order, code_no=None) -> WorkRow:
        qty = r.randint(*item.qty) if item.part else None
        P = _pm_params(r, eq, item, qty or 0)
        P["nm"] = (d.month + (item.interval or 1) - 1) % 12 + 1
        writer = workers[0]
        hb = D._writer_habit(writer.employee_id)
        lines = []
        if r.random() < 0.22 and item.interval:
            lines.append(D._fill(r.choice(_PM_HEAD), P))
        lines.append(D._fill(r.choice(item.acts) if item.acts else f"{item.part}交換", P))
        if item.conds and r.random() < 0.6:
            lines.append(D._fill(r.choice(item.conds), P))
        if r.random() < 0.3:
            # 部品なしの点検作業には「旧品」「交換予定」などの部品交換向けの文を付けない
            extra = _PM_EXTRA if item.part else [x for x in _PM_EXTRA if not any(w in x for w in ("交換", "旧品", "取外し", "予備品", "リーク"))]
            if not item.interval:
                extra = [x for x in extra if "次回" not in x]
            lines.append(D._fill(r.choice(extra), P))
        qc = D._QC.get(D.kind_of(eq))
        if qc and item.part and item.hours >= 1.5 and r.random() < 0.45:
            lines.append(D._fill(r.choice(qc), P))
        if hb["numbering"] == "・" and len(lines) >= 3 and r.random() < 0.5:
            lines = ["・" + x for x in lines]
        content = D._noise("\n".join(lines), hb, r)
        std = item.hours * max(1, len([w for w in workers if w.role != "FE"])) * r.uniform(0.8, 1.35)
        hours = max(0.25, round(std * 4) / 4)
        notes = []
        if flag == "delay":
            notes.append(r.choice(["前月予定分（生産都合で延期）", "前月分 実施遅れ", "先月計画→今月実施"]))
        elif flag == "early":
            notes.append(r.choice(["摩耗早く前倒し交換", "予兆あり前倒し", "計画前倒し（次月予定分）"]))
        if item.fe and r.random() < 0.6:
            notes.append(r.choice(["FE立会い", "保守契約内作業", "メーカー作業"]))
        if not notes and r.random() < 0.18:
            pool = ["チェックシート添付", "記録は点検簿参照"]
            if item.part:
                pool += ["予備品在庫 残{stock}", "旧品リビルド依頼中"]
            if item.interval:
                pool += ["次回 {nm}月予定"]
            notes.append(D._fill(r.choice(pool), P))
        kind = "予防" if flag == "early" else "定期"
        yymm = f"{d.year % 100:02d}{d.month:02d}"
        return WorkRow(
            work_date=d, hour=r.choice([9, 10, 13, 14, 15]), eq=eq, kind=kind, ref=ref, part=D._part_by_name()[item.part] if item.part else None,
            qty=qty, price=D._part_by_name()[item.part].unit_price_yen if item.part else None, content=content, workers=workers,
            hours=hours, remarks="、".join(notes),
            code=f"K532-{eq.equipment_id.replace('-', '')}-{yymm}{order if code_no is None else code_no:02d}",
            source="pm", cont=order > 0, no_part_label=item.work if item.part is None else "", order=order,
        )

    # --- 通常設備: 同一設備・同月の項目を1指示にまとめる ---
    for (eid, y, m), its in sorted(tasks.items(), key=lambda kv: (kv[0][1], kv[0][2], kv[0][0])):
        eq = D.equipment_by_id(eid)
        r = D.rng(f"t4:pm-order:{eid}:{y}-{m}")
        lo = eq.installed_date.day if (eq.installed_date.year, eq.installed_date.month) == (y, m) else 1
        d = _workday(r, y, m, lo)
        fe = any(it.fe for it, _ in its)
        workers = crew(r, eq, fe)
        # 周期の長い項目（大物）ほど先に書く
        its = sorted(its, key=lambda t: (-t[0].interval, t[0].key))
        split_date = len(its) >= 4 and r.random() < 0.3          # 項目が多い月は2日に分けて実施
        d2 = _workday(r, y, m, min(d.day + 1, 27)) if split_date else d
        for j, (item, flag) in enumerate(its):
            rows.append(one_row(r, eq, item, d2 if (split_date and j >= len(its) // 2) else d, workers, flag, "", j))
    # --- 搬送車両: 車両ごとに別日の作業 ---
    fleet_no: Counter = Counter()
    for i, (eq, item, y, m) in enumerate(fleet_tasks):
        r = D.rng(f"t4:pm-fleet-row:{eq.equipment_id}:{item.key}:{y}-{m}:{i}")
        d = _workday(r, y, m)
        fleet_no[(eq.equipment_id, y, m)] += 1
        rows.append(one_row(r, eq, item, d, crew(r, eq, False), "", "", 0, code_no=50 + fleet_no[(eq.equipment_id, y, m)]))
    # --- PM指示No（月内で日付順に採番。1割は記入なし） ---
    rows.sort(key=lambda w: (w.work_date, w.eq.equipment_id, w.order, w.content))
    seq: Counter = Counter()
    last_key, last_no = None, ""
    rn = D.rng("t4:pm-order-no")
    for w in rows:
        key = (w.eq.equipment_id, w.work_date.year, w.work_date.month, w.code[:-2]) if w.cont else None
        if w.cont and key == last_key:
            w.ref = last_no
            continue
        ym = (w.work_date.year, w.work_date.month)
        seq[ym] += 1
        w.ref = "" if rn.random() < 0.1 else f"PM{ym[0] % 100:02d}{ym[1]:02d}-{seq[ym]:03d}"
        last_key, last_no = (w.eq.equipment_id, ym[0], ym[1], w.code[:-2]), w.ref
    return rows


# ---------------------------------------------------------------------------
# 表示ゆれ（書き出し時）
# ---------------------------------------------------------------------------
def _zen(s: str) -> str:
    return D.to_zenkaku(s, digits_only=True).replace(",", "，").replace(".", "．")


@dataclass
class SheetStyle:
    """月（シート）ごとの記入者の癖。"""
    ditto_p: float
    text_qty_p: float
    zen_p: float
    text_date_p: float
    text_price_p: float
    date_fmt: str
    worker_sep: str
    worker_full: bool
    subtotal_label: str
    blank_between: tuple
    two_row_p: float
    labels: dict
    creator: D.Person
    checker: D.Person
    kind_words: dict = field(default_factory=dict)    # 区分列の書き方（突発/定期/予防 -> 表記）。ブロックごとに差し替えることがある
    stats: Counter = field(default_factory=Counter)


def _render_date(r, st: SheetStyle, w: WorkRow, prev: WorkRow | None):
    if prev is not None and w.cont and w.work_date == prev.work_date and r.random() < st.ditto_p:
        return "〃"
    if r.random() < st.text_date_p:
        st.stats["text_date"] += 1
        d = w.work_date
        night = "夜" if (w.hour >= 20 or w.hour < 8) and w.source == "incident" and r.random() < 0.6 else ""
        return r.choice([f"{d.month}/{d.day}{night}", f"{d.month}/{d.day}{night}", D.fmt_date(d, r) + night, f"{d.month}月{d.day}日{night}"])
    return w.work_date


def _render_qty(r, st: SheetStyle, w: WorkRow):
    if w.qty is None:
        return None
    unit = _unit_word(w.part)
    if unit == "式" and r.random() < 0.7:
        st.stats["text_number"] += 1
        return f"{w.qty}式"
    x = r.random()
    if x < st.text_qty_p:
        st.stats["text_number"] += 1
        return f"{w.qty}{unit}"
    if x < st.text_qty_p + st.zen_p:
        st.stats["zenkaku_number"] += 1
        return _zen(str(w.qty)) + (unit if r.random() < 0.5 else "")
    return w.qty


def _render_price(r, st: SheetStyle, w: WorkRow):
    if w.part is None:
        return None
    if w.contract:
        st.stats["contract_blank"] += 1
        return r.choice(["契約内", "保守契約", "－"])
    x = r.random()
    if x < st.text_price_p:
        st.stats["text_number"] += 1
        return r.choice([f"{w.price:,}", f"¥{w.price:,}", f"{w.price:,}円"])
    if x < st.text_price_p + st.zen_p:
        st.stats["zenkaku_number"] += 1
        return _zen(f"{w.price:,}")
    return w.price


def _render_hours(r, st: SheetStyle, w: WorkRow):
    if w.hours is None:
        return None
    x = r.random()
    if x < st.text_qty_p * 0.6:
        st.stats["text_number"] += 1
        h = w.hours
        return f"{int(h * 60)}分" if h < 1 else r.choice([f"{h:g}h", f"{h:g}H", f"{h:g}時間"])
    if x < st.text_qty_p * 0.6 + st.zen_p:
        st.stats["zenkaku_number"] += 1
        return _zen(f"{w.hours:g}")
    return w.hours


def _render_workers(r, st: SheetStyle, w: WorkRow) -> str:
    names = []
    for p in w.workers:
        if p.role == "FE":
            # メーカー略称と姓を空白でつなぐと「瑞穂 野口」のように名＋姓に読めるので、必ず括弧か「FE」で区切る
            short = D.MAKER_SHORT.get(p.section, p.section)
            names.append(r.choice([f"{D.surname(p)}({short}FE)", f"{short}FE", f"{D.surname(p)}（{short}）"]))
        else:
            names.append(p.name if st.worker_full else D.surname(p))
    return st.worker_sep.join(names)


def _render_part(r, w: WorkRow) -> str:
    if w.part is None:
        return w.no_part_label
    n = w.part.name
    x = r.random()
    if x < 0.14:
        return _core_name(n)
    if x < 0.24:
        return n.replace(" (", "（").replace("(", "（").replace(")", "）")
    if x < 0.27:
        return D.to_hankaku_kana(n)
    return n


def _render_eq(r, eq: D.Equipment) -> str:
    x = r.random()
    if x < 0.006:
        return D.to_zenkaku(eq.equipment_id)
    if x < 0.012:
        return eq.equipment_id.replace("-", "")
    return eq.equipment_id


# 区分の書き方（記入者ごとの癖）。行ごとに変えることはなく、シート（ブロック）単位で1つに決めて全行で使う
_KIND_WORDS = {"突発": ["突発", "突発", "突発", "故障", "突発(重要度)", "事後"], "定期": ["定期", "定期", "定期交換", "PM"],
               "予防": ["予防", "予防交換", "前倒し"]}


def _kind_style(seed: str) -> dict:
    """区分列の書き方を 突発／定期／予防 ごとに1つ選ぶ。"""
    rk = D.rng(seed)
    return {kind: rk.choice(words) for kind, words in _KIND_WORDS.items()}


def _render_kind(st: SheetStyle, w: WorkRow) -> str:
    word = st.kind_words[w.kind]
    if word == "突発(重要度)":
        return f"突発({w.severity})" if w.severity else "突発"
    return word


def _render_ref(r, w: WorkRow) -> str:
    if w.source == "incident" and r.random() < 0.01:
        return w.ref.replace("TR-", "TR")          # ハイフン抜けの手入力
    return w.ref


def _row_values(r, st: SheetStyle, w: WorkRow, prev: WorkRow | None) -> dict:
    same = prev is not None and w.cont and prev.ref == w.ref and prev.eq is w.eq
    ditto = same and r.random() < st.ditto_p
    v = {
        "date": _render_date(r, st, w, prev),
        "code": w.code,
        "eq": "〃" if ditto else _render_eq(r, w.eq),
        "eqname": "〃" if ditto else w.eq.name,
        "kind": "〃" if ditto and r.random() < 0.5 else _render_kind(st, w),
        "ref": "〃" if ditto and w.ref else _render_ref(r, w),
        "part": _render_part(r, w),
        "partno": w.part.part_no if w.part else "－",
        "qty": _render_qty(r, st, w),
        "price": _render_price(r, st, w),
        "amount": None if w.amount is None else w.amount,
        "content": w.content,
        "workers": ("〃" if ditto else _render_workers(r, st, w)) if not (same and r.random() < 0.3) else "",
        "hours": _render_hours(r, st, w),
        "remarks": w.remarks,
    }
    if ditto:
        st.stats["ditto_rows"] += 1
    if w.part is not None and w.contract:
        v["amount"] = None
    if "\n" in (v["content"] or ""):
        st.stats["multiline_content"] += 1
    return v


# ---------------------------------------------------------------------------
# シート書き出し
# ---------------------------------------------------------------------------
def _write_header(ws, row: int, col0: int, keys: list, labels: dict, two_row: bool, r, fill) -> tuple[int, int]:
    """見出し行を書く。two_row なら上段にグループ名を結合セルで置き、単独列は縦結合。戻り値は (次の行, 結合数)。"""
    merges = 0
    if not two_row:
        for j, k in enumerate(keys):
            c = ws.cell(row=row, column=col0 + j, value=labels[k])
            c.font, c.fill, c.alignment, c.border = BOLD, fill, CENTER, BORDER
        return row + 1, 0
    grouped = {}
    for gl, gkeys in HEADER_GROUPS:
        idx = [j for j, k in enumerate(keys) if k in gkeys]
        if idx and idx == list(range(idx[0], idx[-1] + 1)):
            label = r.choice(gl)
            for j in idx:
                grouped[j] = (label, idx[0], idx[-1])
    for j, k in enumerate(keys):
        top = ws.cell(row=row, column=col0 + j)
        bot = ws.cell(row=row + 1, column=col0 + j)
        for c in (top, bot):
            c.font, c.fill, c.alignment, c.border = BOLD, fill, CENTER, BORDER
        if j in grouped:
            label, a, b = grouped[j]
            if j == a:
                top.value = label
                ws.merge_cells(start_row=row, start_column=col0 + a, end_row=row, end_column=col0 + b)
                merges += 1
            bot.value = labels[k]
        else:
            top.value = labels[k]
            ws.merge_cells(start_row=row, start_column=col0 + j, end_row=row + 1, end_column=col0 + j)
            merges += 1
    return row + 2, merges


def _write_table(ws, r, st: SheetStyle, row: int, col0: int, keys: list, rows: list, labels: dict, two_row: bool,
                 block_label: str, fill) -> int:
    """見出し＋データ＋小計の1表を書き、次に使える行を返す。"""
    row, merges = _write_header(ws, row, col0, keys, labels, two_row, r, fill)
    st.stats["merged_header"] += merges
    st.stats["header_rows"] += 1
    pos = {k: col0 + j for j, k in enumerate(keys)}
    if not rows:
        ws.cell(row=row, column=pos["part"], value=r.choice(["該当なし", "（該当なし）", "なし"])).border = BORDER
        st.stats["empty_tables"] += 1
        return row + 1
    prev = None
    for w in rows:
        vals = _row_values(r, st, w, prev)
        for k in keys:
            v = vals[k]
            c = ws.cell(row=row, column=pos[k], value=v if v != "" else None)
            c.border = BORDER
            if k == "date" and isinstance(v, date):
                c.number_format = st.date_fmt
                c.alignment = TOP
            elif k in ("content", "remarks", "workers", "part"):
                c.alignment = WRAP_TOP
            elif k in ("price", "amount") and isinstance(v, int):
                c.number_format = "#,##0"
                c.alignment = TOP
            else:
                c.alignment = TOP
        st.stats[f"rows_{w.source}"] += 1
        prev = w
        row += 1
    # 小計（金額・時間は元データの数値から計算。契約内の空欄は含まない）
    total_amt = sum(w.amount or 0 for w in rows if not w.contract)
    total_h = sum(w.hours or 0 for w in rows)
    n_work = len({(w.ref, w.eq.equipment_id, w.work_date, w.code[:-2] if w.source == "pm" else w.code) for w in rows})
    label = st.subtotal_label.format(blk=block_label, n=n_work)
    label_col = pos["part"] if r.random() < 0.6 else pos["date"]
    c = ws.cell(row=row, column=label_col, value=label)
    c.font = BOLD
    ca = ws.cell(row=row, column=pos["amount"], value=total_amt)
    ca.number_format, ca.font = "#,##0", BOLD
    ch = ws.cell(row=row, column=pos["hours"], value=round(total_h, 2))
    ch.font = BOLD
    for k in keys:
        cell = ws.cell(row=row, column=pos[k])
        cell.fill, cell.border = TOTAL_FILL, BORDER
    st.stats["subtotal_rows"] += 1
    return row + 1


def _sort_rows(rows: list[WorkRow]) -> list[WorkRow]:
    return sorted(rows, key=lambda w: (w.work_date, w.eq.equipment_id, w.ref or w.code[:-2], w.order))


def _block_note(r, key: str, rows: list[WorkRow]) -> str | None:
    """ブロック末尾にたまに付く手書き風の注記。"""
    if not rows or r.random() > 0.14:
        return None
    w = r.choice(rows)
    cands = [f"※{w.eq.equipment_id} の部品費は{(w.work_date.month % 12) + 1}月に計上予定",
             f"※{w.eq.equipment_id} {w.work_date.month}/{w.work_date.day} 作業分は作業時間に待機時間を含む",
             "※単価は購買システム登録単価（税抜）", "※FE作業費は別途（保守契約）"]
    if w.source == "incident":
        cands.append(f"※{w.ref} の詳細は修理報告書参照")
    return r.choice(cands)


def _write_month_sheet(wb: Workbook, y: int, m: int, rows: list[WorkRow]) -> dict:
    r = D.rng(f"t4:sheet:{y}-{m}")
    ws = wb.create_sheet(f"{y}年{m}月")
    v2 = (y, m) >= FORMAT_REVISION
    keys = KEYS_V2 if v2 else KEYS_V1
    ppl = D.people()
    # 月ごとの書き癖
    base_idx = {k: r.choices([0, 1, 2], weights=[6, 3, 1])[0] for k in LABELS}
    labels = {k: LABELS[k][base_idx[k]] for k in LABELS}
    st = SheetStyle(
        ditto_p=r.choice([0.0, 0.0, 0.35, 0.7]), text_qty_p=r.uniform(0.02, 0.08), zen_p=r.uniform(0.005, 0.03),
        text_date_p=r.uniform(0.01, 0.08), text_price_p=r.uniform(0.01, 0.05), date_fmt=r.choice(["m/d", "m/d", "yyyy/m/d", "mm/dd"]),
        worker_sep=r.choice(["・", "・", "/", "、", "\n"]), worker_full=r.random() < 0.25,
        subtotal_label=r.choice(["小計", "小計（{n}件）", "{blk} 小計", "計", "ブロック計（{n}件）"]),
        blank_between=r.choice([(1, 1), (1, 2), (2, 3), (1, 3)]), two_row_p=r.choice([0.0, 0.2, 0.5]) + (0.15 if v2 else 0.0),
        labels=labels, creator=r.choice([p for p in ppl if p.section.startswith("設備保全課 保全") and p.role in ("主任", "担当")]),
        checker=r.choice([p for p in ppl if p.role == "係長"]),
        kind_words=_kind_style(f"t4:kind-words:{y}-{m}"),
    )
    sheet_kind_words = dict(st.kind_words)
    width = len(keys)
    # --- シート見出し ---
    ws.cell(row=1, column=1, value=f"{y}年{m}月　部品交換・保全作業記録").font = TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)
    upd = date(y + (m == 12), m % 12 + 1, r.randint(2, 9))
    ws.cell(row=2, column=1, value="製造部 設備保全課")
    ws.cell(row=2, column=4, value=f"作成：{D.surname(st.creator)}　確認：{D.surname(st.checker)}")
    ws.cell(row=2, column=8, value=f"最終更新 {D.fmt_date(upd, r)}")
    note = r.choice(["※金額は税抜。FE作業費は含まない", "※突発＝トラブル報告書No、定期＝PM指示Noを管理No欄に記入", "", "※単価は購買単価（税抜）"])
    row = 3
    if note:
        ws.cell(row=3, column=1, value=note).font = NOTE_FONT
        row = 4
    row += 1
    # --- ブロック構成 ---
    by_block: dict = defaultdict(list)
    for w in rows:
        by_block[_block_key(w.eq)].append(w)
    order = list(BLOCK_ORDER)
    if not v2 and r.random() < 0.5:
        # 改訂前の様式では搬送系を1ブロックにまとめていた月がある
        by_block["T"] = by_block.pop("T1", []) + by_block.pop("T2", [])
        order = [b for b in order if b not in ("T1", "T2")]
        order.insert(6, "T")
    if r.random() < 0.25:
        order.remove("UT")
        order.insert(6, "UT")
    if v2 and r.random() < 0.6:
        order.append("OT")          # 改訂後の様式で追加された「その他」枠（記入なし＝該当なし）
    # 左右2表にするブロック（月に1つ）: 突発と定期・予防が両方そこそこあるもの
    cands = [b for b in order if sum(w.kind == "突発" for w in by_block[b]) >= 3 and sum(w.kind != "突発" for w in by_block[b]) >= 3]
    side_block = r.choice(cands) if cands else None
    code_cols = [keys.index("code") + 1]
    info = {"sheet": ws.title, "side_block": side_block, "blocks": 0, "rows": len(rows)}
    fill = r.choice(HEAD_FILLS)
    for b in order:
        brow = _sort_rows(by_block.get(b, []))
        title = r.choice(BLOCK_TITLES[b])
        tc = ws.cell(row=row, column=1, value=title)
        tc.font = BLOCK_FONT
        if r.random() < 0.3:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
            st.stats["merged_title"] += 1
        row += 1
        blk_label = re.sub(r"[■\s]|（.*", "", title)
        blk_labels = dict(labels)
        if r.random() < 0.3:
            for k in r.sample(keys, k=2):
                blk_labels[k] = r.choice(LABELS[k])        # ブロックごとに見出し文言が微妙に違う
        two_row = r.random() < st.two_row_p
        # ライン担当が別の人だと区分の書き方も違う（ブロック内の全行で同じ表記）
        rk = D.rng(f"t4:kind-words-block:{y}-{m}:{b}")
        st.kind_words = _kind_style(f"t4:kind-words:{y}-{m}:{b}") if rk.random() < 0.2 else sheet_kind_words
        info["blocks"] += 1
        if b == side_block:
            left = [w for w in brow if w.kind == "突発"]
            right = [w for w in brow if w.kind != "突発"]
            col_r = width + 2
            code_cols.append(col_r + keys.index("code"))
            for cc, txt in ((1, r.choice(["【突発・故障対応】", "＜突発＞", "突発修理分"])),
                            (col_r, r.choice(["【定期交換・予防保全】", "＜定期・PM＞", "定期交換分"]))):
                x = ws.cell(row=row, column=cc, value=txt)
                x.font, x.fill = BOLD, SUB_FILL
            row += 1
            end_l = _write_table(ws, r, st, row, 1, keys, left, blk_labels, two_row, blk_label, fill)
            end_r = _write_table(ws, r, st, row, col_r, keys, right, blk_labels, False, blk_label, fill)
            row = max(end_l, end_r)
            st.stats["side_by_side_tables"] += 2
        else:
            row = _write_table(ws, r, st, row, 1, keys, brow, blk_labels, two_row, blk_label, fill)
        if two_row:
            st.stats["two_row_header_blocks"] += 1
        nt = _block_note(r, b, brow)
        if nt:
            ws.cell(row=row, column=1, value=nt).font = NOTE_FONT
            st.stats["note_rows"] += 1
            row += 1
        row += r.randint(*st.blank_between)
    # --- 当月合計 ---
    amt = sum(w.amount or 0 for w in rows if not w.contract)
    amt_inc = sum(w.amount or 0 for w in rows if not w.contract and w.kind == "突発")
    pos_amount = keys.index("amount") + 1
    ws.cell(row=row, column=1, value="■ 当月合計").font = BLOCK_FONT
    c = ws.cell(row=row, column=pos_amount, value=amt)
    c.number_format, c.font, c.fill = "#,##0", BOLD, TOTAL_FILL
    ws.cell(row=row, column=keys.index("hours") + 1, value=round(sum(w.hours or 0 for w in rows), 2)).font = BOLD
    ws.cell(row=row, column=keys.index("content") + 1,
            value=f"突発 {amt_inc:,}円 / 定期・予防 {amt - amt_inc:,}円\n件数 {len(rows)}行").alignment = WRAP_TOP
    # --- 列幅・非表示列 ---
    for j, k in enumerate(keys):
        ws.column_dimensions[get_column_letter(j + 1)].width = WIDTH[k]
        if side_block:
            ws.column_dimensions[get_column_letter(width + 2 + j)].width = WIDTH[k]
    ws.column_dimensions[get_column_letter(width + 1)].width = 3
    for cc in code_cols:
        ws.column_dimensions[get_column_letter(cc)].hidden = True
    ws.sheet_view.zoomScale = r.choice([70, 80, 85, 100])
    info.update(dict(st.stats))
    info["hidden_columns"] = [get_column_letter(cc) for cc in code_cols]
    info["version"] = "改訂後" if v2 else "改訂前"
    return info


def _write_guide_sheet(wb: Workbook) -> None:
    """先頭の「記入要領」シート（データではない説明シート）。"""
    ws = wb.active
    ws.title = "記入要領"
    lines = [
        ("部品交換・保全作業記録　記入要領", TITLE_FONT),
        ("製造部 設備保全課（2025年10月 様式改訂）", None),
        ("", None),
        ("1. 月ごとにシートを分け、ライン別（L1〜L6・搬送系・ユーティリティ）のブロックに記入する。", None),
        ("2. 突発修理はトラブル報告書No（TR-yyyy-nnnnn）、定期交換はPM指示No（PMyymm-nnn）を管理No欄に記入。", None),
        ("3. 同一作業で複数部品を交換した場合は部品ごとに行を分ける。日付・設備等は「〃」可。", None),
        ("4. 作業時間は作業者合計の実作業時間（h）。待機時間を含む場合は備考に記入。", None),
        ("5. 金額＝単価×数量（税抜）。保守契約内の部品は単価欄に「契約内」と記入し金額は空欄。", None),
        ("6. ブロック末尾に小計、シート末尾に当月合計を入れる。", None),
        ("7. 2025年10月以降は品番列を追加。内部コード列（経理連携用）は非表示のまま触らないこと。", None),
        ("", None),
        ("区分：突発（故障・トラブル対応）／定期（計画交換・PM）／予防（予兆による前倒し交換）", None),
    ]
    for i, (t, f) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=t)
        if f:
            c.font = f
    ws.column_dimensions["A"].width = 100


def _normalize_xlsx(path: Path) -> None:
    """openpyxl は保存時刻を docProps/core.xml と zip エントリに書くため、固定値へ置き換えて再パックする。"""
    ts = FIXED_TS.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    buf = io.BytesIO()
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:(created|modified)[^>]*>)[^<]*(</dcterms:\2>)",
                              lambda mm: mm.group(1) + ts + mm.group(3), data)
            zi = zipfile.ZipInfo(info.filename, date_time=FIXED_TS.timetuple()[:6])
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = info.external_attr
            zout.writestr(zi, data)
    path.write_bytes(buf.getvalue())


# ---------------------------------------------------------------------------
# README
# ---------------------------------------------------------------------------
def _readme(xlsx: Path, infos: list[dict], inc_rows: list[WorkRow], pm_rows: list[WorkRow]) -> str:
    tot = Counter()
    for i in infos:
        for k, v in i.items():
            if isinstance(v, int):
                tot[k] += v
    months = [i["sheet"] for i in infos]
    n_inc_ids = len({w.ref for w in inc_rows})
    lines = [
        f"# {xlsx.name}",
        "",
        "製造部 設備保全課の「部品交換・保全作業記録」を想定したサンプル。1シートに複数の表（ライン別ブロック）が縦に並び、",
        "見出し文言のゆれ・結合セル・非表示列・小計行・左右2表などを含む“表として読みにくいExcel”のテスト用データ。",
        "",
        "- 生成: `python -m scripts.samples.t4_maintenance_blocks`（固定シード。再実行で同一バイト列）",
        "- 形式: .xlsx（Office Open XML、内部XMLは UTF-8）。この README は UTF-8（BOMなし、LF）",
        f"- 期間: {months[0]}〜{months[-1]}（{len(months)}ヶ月）",
        "- 画像: なし",
        "",
        "## シート構成",
        "",
        "| シート | 内容 |",
        "|---|---|",
        "| 記入要領 | 様式の説明文のみ（データなし） |",
        f"| {months[0]} 〜 {months[-1]} | 月別の作業記録（{len(months)}シート） |",
        "",
        "各月シートのレイアウト:",
        "",
        "1. 1行目 シート表題（A1:F1 結合）、2行目 作成者・確認者・最終更新日、3行目 注記（ない月もある）",
        "2. 空行の後、ブロックが縦に並ぶ。ブロック = `■ L1ライン` などのタイトル行 → 見出し行（1段 or 2段）→ データ行 → 小計行 → （注記行）→ 空行1〜3行",
        "3. ブロック順は L1〜L6 → 搬送系 → ユーティリティ（月によりユーティリティが搬送系の前、改訂前の一部の月は搬送系が1ブロック）。"
        "改訂後の一部の月は末尾に記入のない「■ その他」枠（該当なし）がある",
        "4. シート末尾に `■ 当月合計` 行（金額・時間の合計と、作業内容列に突発/定期の内訳）",
        "",
        "### 列（様式改訂で変わる）",
        "",
        f"- 改訂前（〜2025年9月、A〜N列）: {' / '.join(LABELS[k][0] for k in KEYS_V1)}（内部コード＝**N列・非表示**）",
        f"- 改訂後（2025年10月〜、A〜O列）: {' / '.join(LABELS[k][0] for k in KEYS_V2)}（内部コード＝**B列・非表示**、品番列を追加）",
        "- 左右2表のブロックでは、右側の表が1列空けた位置（改訂前 P列〜、改訂後 Q列〜）から同じ列順で始まる",
        "",
        "見出し文言の候補（月ごとに基本文言が決まり、ブロックごとに一部が別の言い方になる）:",
        "",
        "| 列 | 表記ゆれ |",
        "|---|---|",
    ]
    for k in KEYS_V2:
        lines.append(f"| {k} | {' / '.join(dict.fromkeys(LABELS[k]))} |")
    lines += [
        "",
        "## 行数",
        "",
        f"- データ行 合計 **{len(inc_rows) + len(pm_rows):,} 行**（突発由来 {len(inc_rows):,} 行 = トラブル {n_inc_ids:,} 件分、定期・予防 {len(pm_rows):,} 行）",
        f"- 見出し行ブロック {tot['header_rows']:,}（うち2段見出し {tot['two_row_header_blocks']}）、小計行 {tot['subtotal_rows']:,}、注記行 {tot['note_rows']}、該当なしの表 {tot['empty_tables']}",
        "",
        "| シート | 様式 | データ行 | 突発 | 定期・予防 | ブロック | 左右2表のブロック | 非表示列 |",
        "|---|---|---:|---:|---:|---:|---|---|",
    ]
    for i in infos:
        lines.append(f"| {i['sheet']} | {i['version']} | {i['rows']} | {i.get('rows_incident', 0)} | {i.get('rows_pm', 0)} | {i['blocks']} | "
                     f"{i['side_block'] or '－'} | {', '.join(i['hidden_columns'])} |")
    lines += [
        "",
        "## 元データとの対応（他サンプルとのクロス整合）",
        "",
        "- 突発行: `domain.standard_incidents()` のうち、対応着手日（response_started_at）が当月で `parts_used` があるもの（部品1点=1行）。"
        "部品なしでも重大度 大/重大 で完了・経過観察のものは「－（調整のみ）」行として載せる。",
        "  - 管理No = incident_id、設備ID/設備名 = equipment、作業者 = assignees（FE は「金子(扶桑FE)」等）、時間(h) = work_hours（同一トラブルの1行目のみ）",
        "  - 作業内容 = トラブル報告の処置（action）から事務的な手順を除いて抜粋し、セル内改行で連結。2行目以降は部品に関係する手順か「〃/同上」",
        "  - 数量・単価 = parts_used の数量と部品カタログ単価。金額 = 単価×数量",
        "- 定期・予防行: 設備種別ごとの定期交換計画（交換周期・数量・標準工数）から決定的に生成。同一設備・同月の項目は1つのPM指示にまとめる。"
        "搬送車両（OHT）は号車ごとの行。周期到来の一部は翌月へ延期（備考「前月予定分」）・前倒し（区分「予防」）・未実施。",
        "- 部品名・品番・単価は `domain.parts_catalog()`、人は `domain.people()` と同一。",
        "",
        "## 意図的な“汚さ”（irregularities）",
        "",
        "1. **1シート複数表**: ライン別ブロックが空行（1〜3行、月により異なる）で区切られて縦に並ぶ。タイトル行 `■ L1ライン` は表記ゆれあり（`■L1ライン`、`■ L1ライン（Fab1）` 等）、一部はA〜D列結合",
        "2. **見出しの繰り返しと文言ゆれ**: ブロックごとに見出し行を繰り返す。月・ブロックで `部品名/交換部品`、`数量/個数`、`単価(円)/単価`、`時間(h)/作業時間(h)/工数(h)` などが変わる",
        f"3. **2段見出しの結合セル**: 一部ブロックは上段「交換部品」「作業」を横結合、単独列は2行縦結合（結合 {tot['merged_header']:,} 箇所）",
        "4. **非表示列**: 内部コード列（`K531-…` 突発 / `K532-…` 定期）。改訂前はN列、改訂後はB列。左右2表の右側表の内部コード列も非表示",
        "5. **様式改訂**: 2025年10月から列構成が変わる（品番列の追加、内部コード列の位置移動）",
        f"6. **左右2表**: 各月1ブロックだけ、左に突発・右に定期/予防の2表を横並び（間に空列1列）。行位置は揃っておらず、小計行の位置も左右で異なる（計 {tot['side_by_side_tables']} 表）",
        f"7. **小計行**: 各表の末尾に `小計` / `小計（n件）` / `L1ライン 小計` / `計` / `ブロック計（n件）`（月で異なる）。ラベルは部品名列または日付列。金額・時間はデータ行の数値合計（「契約内」行は0扱い）（{tot['subtotal_rows']:,} 行）",
        f"8. **セル内改行（Alt+Enter）**: 作業内容の多くが複数行（{tot['multiline_content']:,} セル）。備考・作業者（区切りが改行の月）にも改行あり",
        f"9. **文字列の数値**: `2個`・`1式`・`4本`、単価 `18,500` / `¥18,500` / `18,500円`、時間 `1.5h` / `30分` 等（{tot['text_number']:,} セル）",
        f"10. **全角数字**: 数量・単価・時間の一部が `２`、`１８，５００` など（{tot['zenkaku_number']:,} セル）。作業内容にも記入者の癖で全角数字・半角カナが混ざる",
        f"11. **〃（同上）**: 同一作業の2行目以降で日付・設備ID・設備名・管理No・作業者に `〃`、作業内容に `〃/同上`（{tot['ditto_rows']:,} 行。月により使わない）",
        f"12. **日付の型混在**: 通常は日付セル（表示形式 `m/d`・`yyyy/m/d`・`mm/dd` が月で異なり年が見えない月もある）。一部は文字列 `3/12`、`3/12夜`、`R7.3.12`、`2025年3月12日` 等（{tot['text_date']:,} セル）",
        f"13. **契約内**: 保守契約内の部品は単価に `契約内`/`保守契約`/`－`、金額空欄（{tot['contract_blank']:,} 行）",
        "14. **部品名のゆれ**: カタログ名のほか、本体名のみ（`研磨パッド`）、全角括弧、半角カナ（`ｾﾝｻ` 等）。改訂後は品番列で正確に突合できる",
        "15. **ID の手入力ゆれ（ごく少数）**: 設備IDの全角化（`ＣＭＰ－１０３`）・ハイフン抜け（`CMP103`）、管理Noのハイフン抜け（`TR2025-00123`）",
        "16. **区分の表記ゆれ**: `突発`/`故障`/`事後`/`突発(大)`、`定期`/`定期交換`/`PM`、`予防`/`予防交換`/`前倒し`。"
        "書き方はシート（記入者）ごとに決まっており、2割ほどのブロックは別の担当者の書き方になる。同じブロック内で表記が変わることはない",
        f"17. **注記行・該当なし**: ブロック末尾に `※…` の注記行（{tot['note_rows']}）、行のない表は `該当なし` 1行のみ（{tot['empty_tables']}）",
        "18. **説明シート**: 先頭シート「記入要領」はデータではない文章のみ",
        "19. PM指示Noの約1割は空欄。作業時間は同一トラブルの1行目のみ記入（2行目以降は空欄）",
        "",
        "## 読み取り時の注意（期待される前処理）",
        "",
        "- `■` で始まる行をブロック開始、`小計`/`計` を含む行をブロック終端として扱う。見出し行はブロックごとに再判定する",
        "- 2段見出しは上段と下段を結合して列名にする（例: `交換部品/数量`）",
        "- 左右2表は空列で分割してから個別の表として扱う",
        "- `〃`/`同上` は直前行の値で補完する",
        "- 数量・単価・時間の文字列（全角・単位付き・カンマ付き）は数値へ正規化する",
    ]
    return "\n".join(lines) + "\n"


# ---------------------------------------------------------------------------
def generate(output_root: Path | None = None) -> list[Path]:
    root = Path(output_root) if output_root is not None else D.OUTPUT_ROOT
    out_dir = root / "tables"
    out_dir.mkdir(parents=True, exist_ok=True)
    months = _months()
    inc_rows = _incident_rows(set(months))
    pm_rows = _pm_rows(months)
    by_month: dict = defaultdict(list)
    for w in inc_rows + pm_rows:
        by_month[(w.work_date.year, w.work_date.month)].append(w)

    wb = Workbook()
    wb.properties.creator = "設備保全課"
    wb.properties.lastModifiedBy = "設備保全課"
    wb.properties.created = FIXED_TS
    wb.properties.title = "部品交換・保全作業記録"
    _write_guide_sheet(wb)
    infos = [_write_month_sheet(wb, y, m, by_month[(y, m)]) for (y, m) in months]
    xlsx = out_dir / f"{STEM}.xlsx"
    wb.save(xlsx)
    _normalize_xlsx(xlsx)

    readme = out_dir / f"{STEM}_README.md"
    readme.write_text(_readme(xlsx, infos, inc_rows, pm_rows), encoding="utf-8", newline="\n")
    return [xlsx, readme]


if __name__ == "__main__":
    import time

    t0 = time.time()
    for p in generate():
        print(f"{p}  ({p.stat().st_size / 1024:.0f} KB)")
    print(f"done in {time.time() - t0:.1f}s")
