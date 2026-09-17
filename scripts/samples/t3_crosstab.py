"""T3 月別・設備別 停止時間集計（クロス集計表）のサンプルを生成する。

    python -m scripts.samples.t3_crosstab      # samples/tables/ に xlsx と README を出力

保全課の担当者が毎年度 Excel で作っている「月別・設備別 停止時間集計表」を想定する。
値はすべて domain.standard_incidents()（トラブル報告書 TR）と domain.minor_stops()（チョコ停記録 MS）
から集計するので、帳票系サンプルの設備番号・TR番号・人名と突き合わせができる。

シート構成
- FY2023〜FY2026 : 年度別（4月始まり）の停止時間（分）。小計・合計・総合計は数式（キャッシュ値も書き込む）
- 全期間         : 2023/4〜2026/8 の月次推移。値貼り付け（数式なし）
- 件数           : トラブル件数・チョコ停件数の 2 ブロックを縦に並べた同構造の表

openpyxl は数式のキャッシュ値を書けないため、保存後に xlsx(zip) 内のシート XML を書き換えて
<v> に計算済みの値を入れる（load_workbook(data_only=True) で値が読めるようにするため）。
zip のタイムスタンプ・文書プロパティも固定し、何度実行してもバイト単位で同じファイルになる。
"""
from __future__ import annotations

import io
import itertools
import re
import time
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Callable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import domain as dm

STEM = "T3_月別停止時間集計"
FISCAL_YEARS = (2023, 2024, 2025, 2026)
DATA_START = (2023, 4)
DATA_END = (2026, 8)                       # 集計対象の最終月（これより後は空欄）
FIXED_TS = datetime(2026, 9, 7, 17, 42, 0)  # 文書プロパティ・zip エントリの固定時刻

# ライン（設備マスタの line 値）の並び順。搬送系・ユーティリティはライン按分せず設備単位で計上する
LINE_ORDER = ("L1", "L2", "L3", "L4", "L5", "L6", "L1-L3", "L4-L6", "Fab1共通", "Fab2共通", "全Fab共通")

# 着色しきい値（手作業で塗った体の静的な塗りつぶし。条件付き書式ではない）
MIN_RED, MIN_YELLOW = 5000, 3000           # 停止時間（分/月）
TR_RED, TR_YELLOW = 10, 7                  # トラブル件数（件/月）
MS_RED, MS_YELLOW = 30, 20                 # チョコ停件数（件/月）

# ---------------------------------------------------------------------------
# スタイル
# ---------------------------------------------------------------------------
FONT_NAME = "ＭＳ Ｐゴシック"
F_BASE = Font(name=FONT_NAME, size=10)
F_BOLD = Font(name=FONT_NAME, size=10, bold=True)
F_TITLE = Font(name=FONT_NAME, size=14, bold=True)
F_NOTE = Font(name=FONT_NAME, size=9)
F_RED = Font(name=FONT_NAME, size=10, color="9C0006")
FILL_HDR = PatternFill("solid", fgColor="DDEBF7")
FILL_SUB = PatternFill("solid", fgColor="F2F2F2")
FILL_TOTAL = PatternFill("solid", fgColor="D9E1F2")
FILL_RED = PatternFill("solid", fgColor="FFC7CE")
FILL_YELLOW = PatternFill("solid", fgColor="FFEB9C")
_THIN = Side(style="thin", color="808080")
_MED = Side(style="medium", color="404040")
B_ALL = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
B_TOTAL = Border(left=_THIN, right=_THIN, top=_MED, bottom=_MED)
A_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
A_LEFT = Alignment(horizontal="left", vertical="center")
A_RIGHT = Alignment(horizontal="right", vertical="center")


# ---------------------------------------------------------------------------
# 集計
# ---------------------------------------------------------------------------
YM = tuple  # (年, 月)


def fy_of(dt: date | datetime) -> int:
    """年度（4月始まり）。"""
    return dt.year if dt.month >= 4 else dt.year - 1


def fy_months(fy: int) -> list[YM]:
    return [(fy, m) for m in range(4, 13)] + [(fy + 1, m) for m in (1, 2, 3)]


def all_months() -> list[YM]:
    out, (y, m) = [], DATA_START
    while (y, m) <= DATA_END:
        out.append((y, m))
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return out


@dataclass
class Agg:
    """設備×年月 のキーで集計した値一式。キーは (設備番号, 年, 月)。"""
    tr_min: Counter = field(default_factory=Counter)   # トラブル停止時間（分）
    ms_min: Counter = field(default_factory=Counter)   # チョコ停停止時間（分）
    tr_cnt: Counter = field(default_factory=Counter)   # トラブル件数
    ms_cnt: Counter = field(default_factory=Counter)   # チョコ停件数
    total_min: Counter = field(default_factory=Counter)
    incidents: list = field(default_factory=list)
    minor: list = field(default_factory=list)


def aggregate() -> Agg:
    agg = Agg(incidents=dm.standard_incidents(), minor=dm.minor_stops())
    for i in agg.incidents:
        k = (i.equipment.equipment_id, i.occurred_at.year, i.occurred_at.month)
        agg.tr_min[k] += i.downtime_min      # 月をまたぐ停止も発生月に全量計上
        agg.tr_cnt[k] += 1
    for s in agg.minor:
        k = (s.equipment.equipment_id, s.occurred_at.year, s.occurred_at.month)
        agg.ms_min[k] += s.duration_min
        agg.ms_cnt[k] += 1
    agg.total_min.update(agg.tr_min)
    agg.total_min.update(agg.ms_min)
    return agg


def line_groups(eqs: list[dm.Equipment]) -> list[tuple[str, list[dm.Equipment]]]:
    """ライン順に設備をまとめる（空のラインは出さない）。"""
    out = []
    for line in LINE_ORDER:
        members = sorted((e for e in eqs if e.line == line), key=lambda e: e.equipment_id)
        if members:
            out.append((line, members))
    return out


def sum_months(counter: Counter, eqs, months) -> int:
    return sum(counter.get((e.equipment_id, y, m), 0) for e in eqs for (y, m) in months)


def _installed_ym(eq: dm.Equipment) -> YM:
    return (eq.installed_date.year, eq.installed_date.month)


def cell_fn_for(counter: Counter, dash: str = "－", zero_blank: bool = True) -> Callable:
    """設備×月セルに書く値を返す関数を作る。導入前は dash、実績なしは空欄（zero_blank=False なら 0）。"""
    def fn(eq: dm.Equipment, ym: YM):
        if ym > DATA_END:
            return None                      # 未集計月
        if ym < _installed_ym(eq):
            return dash                      # 設備導入前
        v = counter.get((eq.equipment_id, ym[0], ym[1]), 0)
        return v if (v or not zero_blank) else None
    return fn


# ---------------------------------------------------------------------------
# クロス集計ブロックの書き出し
# ---------------------------------------------------------------------------
@dataclass
class Trail:
    """月列の右側に付く列（合計・年度計・前年比など）。"""
    label: str
    kind: str = "sum"                       # "sum": 月列の合計 / "static": 関数で値を与える
    group: str | None = None                # 2段ヘッダの上段（同じ group が続くと横結合）
    span: tuple[int, int] | None = None     # sum 対象の月インデックス [i0, i1)。None なら全月
    eq_fn: Callable | None = None           # static: 設備1台 -> 値
    rows_fn: Callable | None = None         # static: 設備リスト -> 値（小計・総合計行）
    fmt: str = "#,##0"
    width: float = 9.0


@dataclass
class BlockResult:
    header_rows: tuple[int, int]
    first_data_row: int
    last_row: int
    last_col: int
    eq_rows: dict                           # 設備番号 -> 行
    subtotal_rows: dict                     # ライン -> 行
    total_row: int
    month_col0: int
    trail_col0: int


def _numeric(v) -> int | float:
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else 0


def write_block(ws, cache: dict, *, top: int, lead: list, months: list[YM], groups, cell_fn: Callable,
                trailing: list[Trail], formulas: bool, subtotal_func: str = "sum",
                subtotal_label: str = "{line} 小計", total_label: str = "総合計",
                thresholds: tuple[int, int] | None = None, num_fmt: str = "#,##0") -> BlockResult:
    """2段ヘッダ＋ライン別小計＋総合計のクロス集計表を1ブロック書く。

    lead: [(見出し, 列幅, 設備->値)] の先頭は必ずライン列（縦結合）。
    cache: {セル番地: 値}。数式セルのキャッシュ値を記録し、保存後に XML へ書き込む。
    """
    n_lead = len(lead)
    c_m0 = n_lead + 1
    c_t0 = c_m0 + len(months)
    last_col = c_t0 + len(trailing) - 1
    h1, h2 = top, top + 1
    merges: list[tuple[int, int, int, int]] = []
    L = get_column_letter

    def style(c, font=F_BASE, fill=None, border=B_ALL, align=None, fmt=None):
        c.font = font
        c.border = border
        if fill:
            c.fill = fill
        if align:
            c.alignment = align
        if fmt:
            c.number_format = fmt

    def put(row, col, value, formula_cached=None, **kw):
        """値または数式を書く。formula_cached が与えられたらキャッシュ値として記録する。"""
        c = ws.cell(row, col)
        c.value = value
        if formula_cached is not None:
            cache[c.coordinate] = formula_cached
        style(c, **kw)
        return c

    # --- ヘッダ（2段） ---
    for r in (h1, h2):
        for col in range(1, last_col + 1):
            style(ws.cell(r, col), font=F_BOLD, fill=FILL_HDR, align=A_CENTER)
    for j, (label, _w, _fn) in enumerate(lead, start=1):
        ws.cell(h1, j).value = label
        merges.append((h1, j, h2, j))
    for year, grp in itertools.groupby(enumerate(months), key=lambda x: x[1][0]):
        grp = list(grp)
        c0, c1 = c_m0 + grp[0][0], c_m0 + grp[-1][0]
        ws.cell(h1, c0).value = f"{year}年"
        if c1 > c0:
            merges.append((h1, c0, h1, c1))
        for i, (_y, m) in grp:
            ws.cell(h2, c_m0 + i).value = f"{m}月"
    j = 0
    while j < len(trailing):
        t = trailing[j]
        if t.group:
            k = j
            while k + 1 < len(trailing) and trailing[k + 1].group == t.group:
                k += 1
            ws.cell(h1, c_t0 + j).value = t.group
            if k > j:
                merges.append((h1, c_t0 + j, h1, c_t0 + k))
            for q in range(j, k + 1):
                ws.cell(h2, c_t0 + q).value = trailing[q].label
            j = k + 1
        else:
            ws.cell(h1, c_t0 + j).value = t.label
            merges.append((h1, c_t0 + j, h2, c_t0 + j))
            j += 1

    # --- 右側列の書き出し（設備行・小計行・総合計行で共通） ---
    def write_trailing(row, nums, eqs_for_static, eq=None, font=F_BASE, fill=None, border=B_ALL):
        for q, t in enumerate(trailing):
            col = c_t0 + q
            if t.kind == "sum":
                i0, i1 = t.span or (0, len(months))
                val = sum(nums[i0:i1])
                if formulas:
                    f = f"=SUM({L(c_m0 + i0)}{row}:{L(c_m0 + i1 - 1)}{row})"
                    put(row, col, f, formula_cached=val, font=font, fill=fill, border=border, fmt=t.fmt, align=A_RIGHT)
                else:
                    put(row, col, val, font=font, fill=fill, border=border, fmt=t.fmt, align=A_RIGHT)
            else:
                val = t.eq_fn(eq) if eq is not None else t.rows_fn(eqs_for_static)
                align = A_RIGHT if isinstance(val, (int, float)) else A_CENTER
                put(row, col, val, font=font, fill=fill, border=border, fmt=t.fmt, align=align)

    # --- データ行 ---
    row = h2 + 1
    first_data_row = row
    row_nums: dict[int, list] = {}
    eq_rows, sub_rows = {}, {}
    all_eqs = []
    for line, eqs in groups:
        g0 = row
        for eq in eqs:
            all_eqs.append(eq)
            eq_rows[eq.equipment_id] = row
            put(row, 1, line, font=F_BOLD, align=A_CENTER)
            for jj, (_label, _w, fn) in enumerate(lead[1:], start=2):
                put(row, jj, fn(eq), align=A_LEFT if jj == 3 else A_CENTER)
            nums = []
            for i, ym in enumerate(months):
                v = cell_fn(eq, ym)
                n = _numeric(v)
                nums.append(n)
                fill, font = None, F_BASE
                if thresholds and n >= thresholds[0]:
                    fill, font = FILL_RED, F_RED
                elif thresholds and n >= thresholds[1]:
                    fill = FILL_YELLOW
                put(row, c_m0 + i, v, font=font, fill=fill, fmt=num_fmt,
                    align=A_RIGHT if isinstance(v, (int, float)) else A_CENTER)
            row_nums[row] = nums
            write_trailing(row, nums, None, eq=eq)
            row += 1
        g1 = row - 1
        if g1 > g0:
            merges.append((g0, 1, g1, 1))
        # 小計行
        put(row, 1, subtotal_label.format(line=line), font=F_BOLD, fill=FILL_SUB, align=A_CENTER)
        for jj in range(2, n_lead + 1):
            style(ws.cell(row, jj), fill=FILL_SUB)
        merges.append((row, 1, row, n_lead))
        sub_nums = [sum(row_nums[r][i] for r in range(g0, g1 + 1)) for i in range(len(months))]
        for i in range(len(months)):
            col = L(c_m0 + i)
            if formulas:
                f = (f"=SUBTOTAL(9,{col}{g0}:{col}{g1})" if subtotal_func == "subtotal"
                     else f"=SUM({col}{g0}:{col}{g1})")
                put(row, c_m0 + i, f, formula_cached=sub_nums[i], font=F_BOLD, fill=FILL_SUB, fmt=num_fmt, align=A_RIGHT)
            else:
                put(row, c_m0 + i, sub_nums[i], font=F_BOLD, fill=FILL_SUB, fmt=num_fmt, align=A_RIGHT)
        row_nums[row] = sub_nums
        write_trailing(row, sub_nums, eqs, font=F_BOLD, fill=FILL_SUB)
        sub_rows[line] = row
        row += 1

    # 総合計行
    put(row, 1, total_label, font=F_BOLD, fill=FILL_TOTAL, border=B_TOTAL, align=A_CENTER)
    for jj in range(2, n_lead + 1):
        style(ws.cell(row, jj), fill=FILL_TOTAL, border=B_TOTAL)
    merges.append((row, 1, row, n_lead))
    tot_nums = [sum(row_nums[r][i] for r in sub_rows.values()) for i in range(len(months))]
    last_sub = max(sub_rows.values())
    for i in range(len(months)):
        col = L(c_m0 + i)
        if formulas:
            if subtotal_func == "subtotal":
                f = f"=SUBTOTAL(9,{col}{first_data_row}:{col}{last_sub})"
            else:
                f = "=" + "+".join(f"{col}{r}" for r in sub_rows.values())
            put(row, c_m0 + i, f, formula_cached=tot_nums[i], font=F_BOLD, fill=FILL_TOTAL, border=B_TOTAL,
                fmt=num_fmt, align=A_RIGHT)
        else:
            put(row, c_m0 + i, tot_nums[i], font=F_BOLD, fill=FILL_TOTAL, border=B_TOTAL, fmt=num_fmt, align=A_RIGHT)
    write_trailing(row, tot_nums, all_eqs, font=F_BOLD, fill=FILL_TOTAL, border=B_TOTAL)
    total_row = row

    for r0, c0, r1, c1 in merges:
        ws.merge_cells(start_row=r0, start_column=c0, end_row=r1, end_column=c1)
    return BlockResult((h1, h2), first_data_row, total_row, last_col, eq_rows, sub_rows, total_row, c_m0, c_t0)


def set_widths(ws, lead: list, n_months: int, trailing: list[Trail], month_width: float = 7.6):
    for j, (_label, w, _fn) in enumerate(lead, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for i in range(n_months):
        ws.column_dimensions[get_column_letter(len(lead) + 1 + i)].width = month_width
    for q, t in enumerate(trailing):
        ws.column_dimensions[get_column_letter(len(lead) + 1 + n_months + q)].width = t.width


def write_notes(ws, start_row: int, lines: list[str]) -> int:
    for k, text in enumerate(lines):
        c = ws.cell(start_row + k, 1, text)
        c.font = F_NOTE
    return start_row + len(lines)


def write_stamp_box(ws, col0: int, names: list[tuple[str, str | None]]):
    """右上の押印欄（承認/確認/作成）。names: [(欄名, 姓 or None)]"""
    for q, (label, name) in enumerate(names):
        col = col0 + q
        c = ws.cell(1, col, label)
        c.font, c.fill, c.border, c.alignment = F_NOTE, FILL_HDR, B_ALL, A_CENTER
        for r in (2, 3):
            cc = ws.cell(r, col)
            cc.border, cc.alignment, cc.font = B_ALL, A_CENTER, F_BOLD
        ws.cell(2, col).value = name
        ws.merge_cells(start_row=2, start_column=col, end_row=3, end_column=col)


def person(eid: str) -> dm.Person:
    return {p.employee_id: p for p in dm.people()}[eid]


def short_cause(i: dm.Incident, limit: int = 34) -> str:
    """注記用に原因（なければ現象）を短く切る。"""
    text = i.cause or ("原因調査中：" + i.symptom)
    text = re.split(r"[。\n]", text)[0]
    return text if len(text) <= limit else text[:limit] + "…"


# ---------------------------------------------------------------------------
# 年度シート
# ---------------------------------------------------------------------------
@dataclass(frozen=True)
class FyStyle:
    """年度ごとにファイルを作った担当者・テンプレートの違い（意図的な表記ゆれ）。"""
    fy: int
    created: object                 # datetime ならセル書式付き日付、str ならそのまま
    author_id: str
    checker_id: str
    approver_id: str | None         # None = 未承認
    with_criticality: bool          # FY2025 からテンプレートに「重要度」列が追加された
    subtotal_label: str
    subtotal_func: str              # "subtotal"（SUBTOTAL関数）/ "sum"
    sum_label: str
    ratio_label: str
    ratio_fmt: str
    dash: str                       # 導入前の表記（全角「－」/ 半角「-」）


FY_STYLES = (
    FyStyle(2023, datetime(2024, 4, 12), "M12410", "M10544", "M10231", False, "{line} 計", "subtotal",
            "合計", "前年比", "0%", "－"),
    FyStyle(2024, "2025年4月9日", "M12410", "M10544", "M10231", False, "{line} 小計", "sum",
            "合計", "前年比", "0%", "-"),
    FyStyle(2025, "R8.4.8", "M13702", "M10544", "M10231", True, "{line}小計", "sum",
            "合計", "前年比", "0.0%", "－"),
    FyStyle(2026, "2026/09/07 更新", "M13702", "M10688", None, True, "{line}小計", "sum",
            "合計", "前年同期比", "0.0%", "－"),
)


def yoy_ratio(agg: Agg, eqs: list[dm.Equipment], fy: int):
    """前年比（FY2026 は前年同期=4〜8月比）。比較できないときは文字列。"""
    if fy == FISCAL_YEARS[0]:
        return "－"
    cur_m = [ym for ym in fy_months(fy) if ym <= DATA_END]
    prev_m = [(y - 1, m) for (y, m) in cur_m]
    if all(e.installed_date > date(fy - 1, 4, 1) for e in eqs):
        return "新設"
    prev = sum_months(agg.total_min, eqs, prev_m)
    if prev == 0:
        return "－"
    return round(sum_months(agg.total_min, eqs, cur_m) / prev, 3)


def build_fy_sheet(wb: Workbook, cache_all: dict, agg: Agg, st: FyStyle) -> dict:
    fy = st.fy
    ws = wb.create_sheet(f"FY{fy}")
    cache = cache_all.setdefault(ws.title, {})
    fy_end = date(fy + 1, 3, 31)
    eqs = [e for e in dm.equipment_master() if e.installed_date <= fy_end]   # FY2023 には L6 がまだ無い
    months = fy_months(fy)
    lead = [("ライン", 8.5, None), ("設備番号", 10.0, lambda e: e.equipment_id), ("設備名", 24.0, lambda e: e.name)]
    if st.with_criticality:
        lead.append(("重要度", 6.5, lambda e: e.criticality))
    n_valid = sum(1 for ym in months if ym <= DATA_END)
    trailing = [
        Trail(st.sum_label, "sum", span=(0, len(months)), width=10.5),
        Trail(st.ratio_label, "static", eq_fn=lambda e: yoy_ratio(agg, [e], fy),
              rows_fn=lambda es: yoy_ratio(agg, es, fy), fmt=st.ratio_fmt, width=9.5),
    ]
    last_col = len(lead) + len(months) + len(trailing)

    # タイトル・作成情報
    partial = n_valid < 12
    title = f"FY{fy} 月別・設備別 停止時間集計表" + ("（途中集計：8月末まで）" if partial else "")
    ws.cell(1, 1, title).font = F_TITLE
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col - 4)
    ws.cell(2, 1, "作成日").font = F_BOLD
    c = ws.cell(2, 2, st.created)
    c.font, c.alignment = F_BASE, A_LEFT
    if isinstance(st.created, datetime):
        c.number_format = "yyyy/m/d"
    ws.merge_cells("B2:C2")
    au = person(st.author_id)
    ws.cell(3, 1, "作成者").font = F_BOLD
    ws.cell(3, 2, f"{au.section} {au.name}").font = F_BASE
    ws.merge_cells("B3:C3")
    period_end = "2026/8/31" if partial else f"{fy + 1}/3/31"
    ws.cell(4, 1, f"集計期間：{fy}/4/1～{period_end}　単位：分（トラブル停止＋チョコ停）").font = F_BASE
    write_stamp_box(ws, last_col - 2, [
        ("承認", dm.surname(person(st.approver_id)) if st.approver_id else None),
        ("確認", dm.surname(person(st.checker_id))),
        ("作成", dm.surname(au)),
    ])

    res = write_block(ws, cache, top=6, lead=lead, months=months, groups=line_groups(eqs),
                      cell_fn=cell_fn_for(agg.total_min, dash=st.dash), trailing=trailing, formulas=True,
                      subtotal_func=st.subtotal_func, subtotal_label=st.subtotal_label,
                      thresholds=(MIN_RED, MIN_YELLOW))
    set_widths(ws, lead, len(months), trailing)
    ws.freeze_panes = ws.cell(res.first_data_row, res.month_col0)
    ws.sheet_view.zoomScale = 85
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    notes = fy_notes(agg, st, eqs)
    end = write_notes(ws, res.last_row + 2, notes)
    return {"sheet": ws.title, "block": res, "notes": notes, "eqs": eqs, "notes_end": end,
            "grand_total": sum_months(agg.total_min, eqs, months)}


def fy_notes(agg: Agg, st: FyStyle, eqs: list[dm.Equipment]) -> list[str]:
    """年度シート下の注記。数値・TR番号はすべて集計元データから算出する。"""
    fy = st.fy
    months = [ym for ym in fy_months(fy) if ym <= DATA_END]
    incs = [i for i in agg.incidents if fy_of(i.occurred_at) == fy]
    mss = [s for s in agg.minor if fy_of(s.occurred_at) == fy]
    tr_min = sum(i.downtime_min for i in incs)
    ms_min = sum(s.duration_min for s in mss)
    open_n = sum(1 for i in incs if i.status in ("対応中", "保留"))
    top = sorted(incs, key=lambda i: (-i.downtime_min, i.incident_id))[:3]
    by_line = {line: sum_months(agg.total_min, es, months) for line, es in line_groups(eqs)}
    worst_line = max(by_line, key=lambda k: (by_line[k], k))
    by_eq = sorted(((sum_months(agg.total_min, [e], months), e.equipment_id) for e in eqs), reverse=True)[:3]
    top_txt = "、".join(f"{i.incident_id}（{i.equipment.equipment_id} {i.subsystem}：{short_cause(i)}　{i.downtime_min:,}分）"
                       for i in top)
    eq_txt = " / ".join(f"{eid} {v:,}分" for v, eid in by_eq)

    if fy == 2023:
        return [
            "※注記",
            f"※1 停止時間＝トラブル報告書（TR）の停止時間＋チョコ停記録（MS）の停止時間。計画停止（PM・定期点検）は含まない。",
            f"※2 内訳：トラブル {len(incs):,}件 {tr_min:,}分、チョコ停 {len(mss):,}件 {ms_min:,}分。月をまたいだトラブルは発生月に全量計上。",
            "※3 前年比はFY2022以前が旧保全システムのため算出していない（「－」）。",
            f"※4 月3,000分以上は黄、5,000分以上は赤で色付け（手作業）。「－」は設備導入前、空欄は停止なし。",
            f"※5 停止時間上位：{top_txt}",
            "※6 L6（Fab2 3F 新ライン）はFY2024から立上げのため本表には含まない。",
        ]
    ratio_all = yoy_ratio(agg, eqs, fy)
    ratio_txt = f"{ratio_all:.0%}" if isinstance(ratio_all, float) else str(ratio_all)
    # 前年度のファイルをコピーして書き足している体なので、共通の注記も年度ごとに少しずつ言い回しが違う
    head = {
        2024: [f"※1 前年比＝当年度合計÷前年度合計。前年度途中から稼働した設備は「新設」とした。",
               f"※4 着色：月間{MIN_YELLOW:,}分以上＝黄、{MIN_RED:,}分以上＝赤（目視確認用。条件付き書式ではないので値修正時は塗り直すこと）",
               f"※5 「{st.dash}」は設備導入前。搬送系（L1-L3/L4-L6）・ユーティリティ（Fab共通）はライン按分せず設備単位で計上。"],
        2025: [f"※1 前年比は年度合計どうしの比。FY2024途中に立上げたL6設備は比較対象外（「新設」）。",
               f"※4 色付けは前年度と同じ基準（黄：{MIN_YELLOW:,}分～、赤：{MIN_RED:,}分～）。手で塗っているので数値修正時は要注意。",
               f"※5 「{st.dash}」＝導入前。重要度（A/B/C）列を今年度から追加。搬送・ユーティリティは設備単位。"],
        2026: [f"※1 2026年9月以降は未集計（途中集計）。前年同期比＝当年4〜8月計÷前年4〜8月計。",
               f"※4 色付け基準はFY2025と同じ（黄 {MIN_YELLOW:,}分以上／赤 {MIN_RED:,}分以上）。",
               f"※5 「{st.dash}」＝導入前（本年度は該当なし）。9月以降の列は空欄のまま、小計の数式は残している。"],
    }[fy]
    notes = ["※注記", head[0]]
    notes += [
        f"※2 集計元：保全管理システム（TR報告書 {len(incs):,}件／チョコ停 {len(mss):,}件）。停止時間の内訳はトラブル {tr_min:,}分・チョコ停 {ms_min:,}分。",
        f"※3 対応中・保留の案件（{open_n}件）の停止時間は集計時点の暫定値。完了後に修正する場合あり。",
        head[1],
        head[2],
        f"※6 停止時間ワースト設備：{eq_txt}。ライン別では {worst_line} が最多（{by_line[worst_line]:,}分）。全体の{st.ratio_label} {ratio_txt}。",
        f"※7 単発の長時間停止：{top_txt}",
    ]
    # 持病号機の対策効果（domain の対策完了日に合わせたコメント。前年比は実データから計算）
    remarks = {
        2025: [("CVD-205", "ドライポンプ過負荷対策（2025年4月）"), ("CLN-504", "ノズル液だれ対策（2024年12月）")],
        2026: [("LIT-402", "ステージサーボ対策（2025年9月）")],
    }.get(fy, [])
    k = 8
    for eid, what in remarks:
        r = yoy_ratio(agg, [dm.equipment_by_id(eid)], fy)
        if not isinstance(r, float):
            continue
        effect = f"停止時間が減少（{st.ratio_label} {r:.0%}）" if r < 1 else f"改善効果見られず（{st.ratio_label} {r:.0%}）、継続監視"
        notes.append(f"※{k} {eid}：{what}後、{effect}。")
        k += 1
    if fy == 2024:
        notes.append(f"※{k} L6は2024年9月から順次立上げ（初期故障対応含む）。導入前の月は「{st.dash}」。")
    return notes


# ---------------------------------------------------------------------------
# 全期間シート（値貼り付け）
# ---------------------------------------------------------------------------
def build_all_period_sheet(wb: Workbook, cache_all: dict, agg: Agg) -> dict:
    ws = wb.create_sheet("全期間")
    cache = cache_all.setdefault(ws.title, {})
    eqs = dm.equipment_master()
    months = all_months()
    lead = [("ライン", 8.5, None), ("設備番号", 10.0, lambda e: e.equipment_id), ("設備名", 24.0, lambda e: e.name)]
    trailing = []
    for fy in FISCAL_YEARS:
        idx = [i for i, ym in enumerate(months) if fy_of(date(ym[0], ym[1], 1)) == fy]
        label = f"FY{fy}" + ("(4-8月)" if fy == 2026 else "")
        trailing.append(Trail(label, "sum", group="年度計", span=(idx[0], idx[-1] + 1), width=10.5))
    trailing += [
        Trail("総計(分)", "sum", width=11.0),
        Trail("総計(時間)", "static", eq_fn=lambda e: round(sum_months(agg.total_min, [e], months) / 60, 1),
              rows_fn=lambda es: round(sum_months(agg.total_min, es, months) / 60, 1), fmt="#,##0.0", width=10.0),
        Trail("うちトラブル", "static", group="内訳(分)", eq_fn=lambda e: sum_months(agg.tr_min, [e], months),
              rows_fn=lambda es: sum_months(agg.tr_min, es, months), width=11.0),
        Trail("うちチョコ停", "static", group="内訳(分)", eq_fn=lambda e: sum_months(agg.ms_min, [e], months),
              rows_fn=lambda es: sum_months(agg.ms_min, es, months), width=11.0),
    ]
    au = person("M13520")   # 保全2係 阿部 真央
    ws.cell(1, 1, "設備別 停止時間 月次推移（全期間：2023年4月～2026年8月）").font = F_TITLE
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=20)
    ws.cell(2, 1, "作成日").font = F_BOLD
    c = ws.cell(2, 2, FIXED_TS.replace(hour=0, minute=0))
    c.number_format, c.font, c.alignment = 'yyyy"年"m"月"d"日"', F_BASE, A_LEFT
    ws.merge_cells("B2:C2")
    ws.cell(3, 1, "作成者").font = F_BOLD
    ws.cell(3, 2, f"{au.section.replace('設備保全課 ', '')} {dm.surname(au)}（{au.employee_id}）").font = F_BASE
    ws.merge_cells("B3:C3")
    ws.cell(4, 1, "※各年度シートの集計値を値貼り付け（数式なし）。単位：分").font = F_NOTE

    res = write_block(ws, cache, top=5, lead=lead, months=months, groups=line_groups(eqs),
                      cell_fn=cell_fn_for(agg.total_min), trailing=trailing, formulas=False,
                      subtotal_label="{line} 小計", thresholds=(MIN_RED, MIN_YELLOW))
    set_widths(ws, lead, len(months), trailing, month_width=7.0)
    ws.freeze_panes = ws.cell(res.first_data_row, res.month_col0)
    ws.sheet_view.zoomScale = 80

    incs, mss = agg.incidents, agg.minor
    tr_min, ms_min = sum(i.downtime_min for i in incs), sum(s.duration_min for s in mss)
    by_eq = sorted(((sum_months(agg.total_min, [e], months), e.equipment_id) for e in eqs), reverse=True)[:5]
    by_year = {fy: sum_months(agg.total_min, eqs, [ym for ym in fy_months(fy) if ym <= DATA_END]) for fy in FISCAL_YEARS}
    notes = [
        "※注記",
        f"※1 対象：トラブル報告書 {len(incs):,}件（{incs[0].incident_id}～{incs[-1].incident_id}）＋チョコ停記録 {len(mss):,}件。",
        f"※2 停止時間 総計 {tr_min + ms_min:,}分（約{(tr_min + ms_min) / 60:,.0f}時間）＝トラブル {tr_min:,}分＋チョコ停 {ms_min:,}分。",
        "※3 年度計：" + "、".join(f"FY{fy} {v:,}分" for fy, v in by_year.items()) + "（FY2026は8月末まで）。",
        "※4 停止時間の多い設備（全期間）：" + "、".join(f"{eid}（{v / 60:,.0f}h）" for v, eid in by_eq),
        "※5 空欄＝停止なし、「－」＝設備導入前。色は年度シートと同じ基準（3,000分以上 黄／5,000分以上 赤）。",
        "※6 総計(時間)は小数第1位で四捨五入しているため、小計・総合計と各行の和が一致しない場合がある。",
    ]
    write_notes(ws, res.last_row + 2, notes)
    return {"sheet": ws.title, "block": res, "notes": notes, "grand_total": tr_min + ms_min}


# ---------------------------------------------------------------------------
# 件数シート（トラブル件数・チョコ停件数の2ブロック）
# ---------------------------------------------------------------------------
def build_count_sheet(wb: Workbook, cache_all: dict, agg: Agg) -> dict:
    ws = wb.create_sheet("件数")
    cache = cache_all.setdefault(ws.title, {})
    eqs = dm.equipment_master()
    months = all_months()
    lead = [("ライン", 8.5, None), ("設備番号", 10.0, lambda e: e.equipment_id), ("設備名", 24.0, lambda e: e.name)]

    def trailing_for():
        tr = []
        for fy in FISCAL_YEARS:
            idx = [i for i, ym in enumerate(months) if fy_of(date(ym[0], ym[1], 1)) == fy]
            tr.append(Trail(f"FY{fy}", "sum", group="年度計", span=(idx[0], idx[-1] + 1), fmt="#,##0", width=8.5))
        tr.append(Trail("総計", "sum", fmt="#,##0", width=9.0))
        return tr

    au = person("M13520")
    ws.cell(1, 1, "設備別 月別件数（トラブル／チョコ停）").font = F_TITLE
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=20)
    ws.cell(2, 1, "作成日").font = F_BOLD
    ws.cell(2, 2, "2026-09-07").font = F_BASE
    ws.merge_cells("B2:C2")
    ws.cell(3, 1, "作成者").font = F_BOLD
    ws.cell(3, 2, f"保全2係 {au.name}").font = F_BASE
    ws.merge_cells("B3:C3")

    ws.cell(5, 1, "■ トラブル件数（TR：トラブル報告書ベース、単位：件）").font = F_BOLD
    tr_trailing = trailing_for()
    res1 = write_block(ws, cache, top=6, lead=lead, months=months, groups=line_groups(eqs),
                       cell_fn=cell_fn_for(agg.tr_cnt), trailing=tr_trailing, formulas=True,
                       subtotal_label="{line} 小計", thresholds=(TR_RED, TR_YELLOW), num_fmt="0")
    top2 = res1.last_row + 3
    ws.cell(top2 - 1, 1, "■ チョコ停件数（MS：チョコ停記録ベース、単位：件）　※0件も「0」で表示").font = F_BOLD
    ms_trailing = trailing_for()
    res2 = write_block(ws, cache, top=top2, lead=lead, months=months, groups=line_groups(eqs),
                       cell_fn=cell_fn_for(agg.ms_cnt, zero_blank=False), trailing=ms_trailing, formulas=True,
                       subtotal_label="{line}計", total_label="合計", thresholds=(MS_RED, MS_YELLOW), num_fmt="0")
    set_widths(ws, lead, len(months), tr_trailing, month_width=5.5)
    ws.freeze_panes = "D1"
    ws.sheet_view.zoomScale = 80

    incs, mss = agg.incidents, agg.minor
    sev = Counter(i.severity for i in incs)
    no_alarm = sum(1 for s in mss if s.alarm is None)
    auto = sum(1 for s in mss if s.recovered_by == "自動復帰")
    notes = [
        "※注記",
        f"※1 トラブル件数は状態（完了/経過観察/対応中/保留）を問わず全件：計 {len(incs):,}件"
        f"（重大 {sev['重大']}・大 {sev['大']:,}・中 {sev['中']:,}・小 {sev['小']:,}）。",
        f"※2 瞬低（電源瞬時電圧低下）時は複数設備で同時に起票されるため、同じ日に件数が集中する月がある。",
        f"※3 チョコ停件数：計 {len(mss):,}件（うち自動復帰 {auto:,}件、アラームコードなし {no_alarm:,}件）。",
        f"※4 着色：トラブル {TR_YELLOW}件以上 黄／{TR_RED}件以上 赤、チョコ停 {MS_YELLOW}件以上 黄／{MS_RED}件以上 赤。",
        "※5 上段は0件を空欄、下段は0件を「0」と表示（作成者が異なるため表記が揃っていない）。「－」は設備導入前。",
    ]
    write_notes(ws, res2.last_row + 2, notes)
    return {"sheet": ws.title, "blocks": (res1, res2), "notes": notes}


# ---------------------------------------------------------------------------
# 保存（数式キャッシュ値の書き込み・決定的な zip）
# ---------------------------------------------------------------------------
_CELL_F_RE = re.compile(r'<c r="([A-Z]+[0-9]+)"([^>]*)><f>([^<]*)</f>(?:<v\s*/>|<v>[^<]*</v>)?</c>')


def _attrs(tag: str) -> dict:
    return dict(re.findall(r'([\w:]+)="([^"]*)"', tag))


def finalize_xlsx(raw: bytes, cache_all: dict) -> bytes:
    """openpyxl の出力に数式キャッシュ値を差し込み、zip のタイムスタンプを固定して返す。"""
    zin = zipfile.ZipFile(io.BytesIO(raw))
    wb_xml = zin.read("xl/workbook.xml").decode("utf-8")
    rels_xml = zin.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    rid_target = {}
    for tag in re.findall(r"<Relationship\b[^>]*/>", rels_xml):
        a = _attrs(tag)
        rid_target[a["Id"]] = "xl/" + a["Target"].lstrip("/").removeprefix("xl/")
    sheet_path = {}
    for tag in re.findall(r"<sheet\b[^>]*/>", wb_xml):
        a = _attrs(tag)
        sheet_path[rid_target[a["r:id"]]] = a["name"]

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            title = sheet_path.get(info.filename)
            if title is not None and cache_all.get(title):
                cache = cache_all[title]
                hit = set()

                def repl(m):
                    ref = m.group(1)
                    if ref not in cache:
                        return m.group(0)
                    hit.add(ref)
                    v = cache[ref]
                    vs = str(int(v)) if float(v).is_integer() else repr(float(v))
                    return f'<c r="{ref}"{m.group(2)}><f>{m.group(3)}</f><v>{vs}</v></c>'

                data = _CELL_F_RE.sub(repl, data.decode("utf-8")).encode("utf-8")
                missing = set(cache) - hit
                if missing:
                    raise RuntimeError(f"{title}: キャッシュ値を書き込めなかったセル {sorted(missing)[:5]} ...")
            elif info.filename == "docProps/core.xml":
                # openpyxl は保存時刻で modified を上書きするので固定値に戻す（決定的出力のため）
                data = re.sub(r"(<dcterms:modified[^>]*>)[^<]*", r"\g<1>" + FIXED_TS.strftime("%Y-%m-%dT%H:%M:%SZ"),
                              data.decode("utf-8")).encode("utf-8")
            zi = zipfile.ZipInfo(info.filename, date_time=FIXED_TS.timetuple()[:6])
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            zout.writestr(zi, data)
    return out.getvalue()


# ---------------------------------------------------------------------------
# README
# ---------------------------------------------------------------------------
def build_readme(xlsx_path: Path, infos: list[dict], cache_all: dict, agg: Agg, wb: Workbook) -> str:
    L = get_column_letter
    incs, mss = agg.incidents, agg.minor
    # 実際に書いた数式の例（README の記述とファイルの中身を一致させるため、ブックから取り出す）
    fx = {}
    for info in infos[:len(FY_STYLES)]:
        b = info["block"]
        ws = wb[info["sheet"]]
        f_total = ws.cell(b.total_row, b.month_col0).value
        terms = f_total.split("+")
        fx[info["sheet"]] = {
            "sub": ws.cell(next(iter(b.subtotal_rows.values())), b.month_col0).value,
            "total": "+".join(terms[:3]) + ("+…" if len(terms) > 3 else ""),
            "sum": ws.cell(b.first_data_row, b.trail_col0).value,
        }
    tr_min, ms_min = sum(i.downtime_min for i in incs), sum(s.duration_min for s in mss)
    lines = [
        f"# {STEM}.xlsx（T3 月別・設備別 停止時間集計／クロス集計表）",
        "",
        "製造部 設備保全課が年度ごとに作成している「月別・設備別 停止時間集計表」を模したサンプル。",
        "RAG 前処理で *2段結合ヘッダ・縦結合セル・小計行・数式・注記* を含むクロス集計表を扱うためのテストデータ。",
        "",
        "- 生成: `python -m scripts.samples.t3_crosstab`（固定シード。再実行してもバイト単位で同一）",
        f"- 形式: Office Open XML（.xlsx）。セル文字列は XML 内で UTF-8。本 README は UTF-8（BOMなし・LF）",
        "- 読み取り: `openpyxl.load_workbook(path, data_only=True)` で数式セルもキャッシュ値（数値）が得られる",
        f"- 集計元: `domain.standard_incidents()` {len(incs):,}件（停止時間 {tr_min:,}分）＋ `domain.minor_stops()` {len(mss):,}件（{ms_min:,}分）",
        f"- 期間: 2023-04-01 ～ 2026-08-31（年度は4月始まり、FY2026 は 8月末までの途中集計）",
        "",
        "## 値の定義（再計算方法）",
        "",
        "- 停止時間セル = その設備・その年月（発生日時 `occurred_at` の年月）の `Incident.downtime_min` の和 ＋ `MinorStop.duration_min` の和（分）",
        "- 月をまたぐトラブルも発生月に全量計上。対応中・保留のトラブルも `downtime_min` をそのまま計上",
        "- 件数シート: 上段 = Incident 件数、下段 = MinorStop 件数（同じ年月・設備キー）",
        "- 小計 = ライン内の設備行の和、総合計 = 小計の和、合計 = 月列の和（すべて整数で厳密に一致）",
        "- 前年比 = 当年度合計 ÷ 前年度合計（小数3桁に丸めた静的値）。FY2026 は 4〜8月同士の比較",
        "",
        "## シート構成",
        "",
        "| シート | 見出し行 | データ行 | 設備行 | 小計行 | 列数 | 数式セル | 結合範囲 | 総合計(分/件) |",
        "|---|---|---|---|---|---|---|---|---|",
    ]
    for info in infos:
        ws = wb[info["sheet"]]
        blocks = info.get("blocks") or (info["block"],)
        for bi, b in enumerate(blocks):
            name = info["sheet"] + (f"（{'上段 トラブル' if bi == 0 else '下段 チョコ停'}）" if len(blocks) > 1 else "")
            total_cell = ws.cell(b.total_row, b.trail_col0 + (4 if len(blocks) > 1 or info['sheet'] == '全期間' else 0))
            total_val = cache_all[ws.title].get(total_cell.coordinate, total_cell.value)
            n_formula = sum(1 for k in cache_all[ws.title]
                            if b.header_rows[0] <= int(re.sub(r"[A-Z]+", "", k)) <= b.last_row)
            lines.append(
                f"| {name} | {b.header_rows[0]}-{b.header_rows[1]} | {b.first_data_row}-{b.last_row} | {len(b.eq_rows)} | "
                f"{len(b.subtotal_rows)} | {b.last_col}（A-{L(b.last_col)}） | {n_formula} | "
                f"{len(ws.merged_cells.ranges) if bi == 0 else '（同上）'} | {total_val:,} |")
    lines += [
        "",
        "各シートの共通レイアウト:",
        "",
        "- 1〜4行目: タイトル（結合）、作成日、作成者、集計期間・単位。年度シートは右上に押印欄（承認/確認/作成、2〜3行目を縦結合）",
        "- 見出し2段: 上段「2024年」「2025年」などの年を月列の上に横結合、下段「4月」…「3月」。ライン・設備番号・設備名・合計・前年比は2段を縦結合",
        "- A列「ライン」は同じラインの設備行を縦結合（小計行は含まない）。小計行・総合計行は A〜設備名列を横結合したラベル",
        "- ライン順: " + " → ".join(LINE_ORDER) + "（L1-L3/L4-L6 は OHT・AGV など搬送系、Fab共通はユーティリティ）",
        "- 表の下に「※注記」行（A列に長文、結合なし）。上位トラブルの TR番号・設備番号・停止時間を記載",
        "",
        "## 意図的なイレギュラー（前処理のテスト観点）",
        "",
        "1. **数式とキャッシュ値**: 年度シートと件数シートの 合計列・小計行・総合計行は数式。キャッシュ値（計算済みの数値）も書き込み済みで、"
        "`calcPr fullCalcOnLoad=1` のため Excel で開くと再計算される。全期間シートは値貼り付けで数式なし。数式の例:",
        f"    - FY2023: 小計 `{fx['FY2023']['sub']}`、総合計 `{fx['FY2023']['total']}`（SUBTOTAL は小計を二重計上しない）、合計 `{fx['FY2023']['sum']}`",
        f"    - FY2024: 小計 `{fx['FY2024']['sub']}`、総合計 `{fx['FY2024']['total']}`（小計セルの足し算）、合計 `{fx['FY2024']['sum']}`",
        f"    - FY2025/FY2026: 重要度列の分だけ右にずれ、小計 `{fx['FY2025']['sub']}`、合計 `{fx['FY2025']['sum']}`",
        "2. **シートごとに行位置・列位置が違う**: FY2023 は L6 の設備（2024年9月以降導入）が無く行数が少ない。FY2025・FY2026 は「重要度」列が追加され月列が1列右にずれる",
        "3. **表記ゆれ**: 小計ラベルが「L1 計」（FY2023）/「L1 小計」（FY2024・全期間・件数上段）/「L1小計」（FY2025・FY2026）/「L1計」（件数下段）。総合計ラベルは件数下段のみ「合計」。FY2026 の比較列は「前年同期比」",
        "4. **導入前の表記**: 導入前の月は文字列「－」（全角）だが FY2024 シートだけ半角「-」。数値列に文字列が混在する",
        "5. **0 の表記**: 停止時間・トラブル件数は 0 を空欄、チョコ停件数（件数シート下段）は 0 を数値 0 で表示。FY2026 の 9月〜3月は未集計のため設備行は空欄だが、小計・総合計の数式は 0 を返す",
        "6. **前年比の型混在**: 数値（書式 `0%` / `0.0%`）と文字列「－」「新設」が同じ列に混在。FY2023 は全行「－」",
        "7. **作成日の書式違い**: FY2023 は日付セル（`yyyy/m/d`）、FY2024「2025年4月9日」、FY2025「R8.4.8」（和暦文字列）、FY2026「2026/09/07 更新」、全期間は日付セル（`yyyy\"年\"m\"月\"d\"日\"`）、件数「2026-09-07」",
        "8. **押印欄**: 年度シート右上に小さな表（承認/確認/作成）。FY2026 は承認欄が空（未承認）",
        "9. **静的な色付け**: 月間の停止時間 3,000分以上を黄、5,000分以上を赤で塗りつぶし（条件付き書式ではなくセルの塗り）。件数シートはトラブル 7/10件、チョコ停 20/30件がしきい値",
        "10. **1シートに2つの表**: 件数シートは上段（トラブル件数）と下段（チョコ停件数）が同じ見出し構造で縦に並び、間に「■」見出し行がある。ウィンドウ枠固定は列のみ",
        "11. **単位の混在**: 全期間シートは「総計(分)」と「総計(時間)」（小数1位丸め）が並ぶ。時間列は丸めのため行の和と小計が一致しないことがある（注記 ※6）",
        "12. **全角・半角**: タイトル・注記の範囲表記に全角「～」、集計期間は半角数字。設備番号・TR番号は常に半角",
        "",
        "## 他サンプルとの対応",
        "",
        "- 設備番号・設備名・ラインは `domain.equipment_master()` と同一",
        "- 注記に出る TR番号（例: " + ", ".join(i.incident_id for i in sorted(incs, key=lambda i: -i.downtime_min)[:3])
        + "）はトラブル報告書系サンプルの報告番号と一致",
        "- 作成者・押印欄の人名は `domain.people()`（保全課 松本 拓也 / 清水 彩花 / 阿部 真央、係長 中村 浩二 / 斎藤 健太郎、課長 高橋 誠）",
        "",
        "## 検証用の値",
        "",
    ]
    for fy in FISCAL_YEARS:
        ms_ = [ym for ym in fy_months(fy) if ym <= DATA_END]
        eqs = dm.equipment_master()
        lines.append(f"- FY{fy}: 停止時間 {sum_months(agg.total_min, eqs, ms_):,}分"
                     f"（トラブル {sum_months(agg.tr_min, eqs, ms_):,} ＋ チョコ停 {sum_months(agg.ms_min, eqs, ms_):,}）、"
                     f"トラブル {sum_months(agg.tr_cnt, eqs, ms_):,}件、チョコ停 {sum_months(agg.ms_cnt, eqs, ms_):,}件")
    lines.append(f"- 全期間: {tr_min + ms_min:,}分、トラブル {len(incs):,}件、チョコ停 {len(mss):,}件")
    lines.append("")
    return "\n".join(lines)


# ---------------------------------------------------------------------------
def generate(output_root: Path | None = None) -> list[Path]:
    root = Path(output_root) if output_root else dm.OUTPUT_ROOT
    out_dir = root / "tables"
    out_dir.mkdir(parents=True, exist_ok=True)
    agg = aggregate()

    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = person("M12410").name
    wb.properties.lastModifiedBy = person("M13702").name
    wb.properties.title = "月別・設備別 停止時間集計"
    wb.properties.created = datetime(2024, 4, 12, 9, 15, 0)
    wb.properties.modified = FIXED_TS

    cache_all: dict[str, dict] = {}
    infos = [build_fy_sheet(wb, cache_all, agg, st) for st in FY_STYLES]
    infos.append(build_all_period_sheet(wb, cache_all, agg))
    infos.append(build_count_sheet(wb, cache_all, agg))
    wb.active = len(FY_STYLES) - 1          # 最新年度シートを開いた状態で保存

    buf = io.BytesIO()
    wb.save(buf)
    data = finalize_xlsx(buf.getvalue(), cache_all)
    xlsx_path = out_dir / f"{STEM}.xlsx"
    xlsx_path.write_bytes(data)

    readme_path = out_dir / f"{STEM}_README.md"
    readme_path.write_bytes(build_readme(xlsx_path, infos, cache_all, agg, wb).encode("utf-8"))
    return [xlsx_path, readme_path]


if __name__ == "__main__":
    t0 = time.time()
    for p in generate():
        print(f"{p}  ({p.stat().st_size / 1024:.1f} KB)")
    print(f"done in {time.time() - t0:.1f}s")
