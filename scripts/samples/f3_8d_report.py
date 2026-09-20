"""F3 是正処置報告書（8D形式）のサンプルExcelを作る。

    python -m scripts.samples.f3_8d_report      # samples/forms/F3_8D是正処置報告書/ に30ファイル出力

正準トラブル履歴（domain.standard_incidents）から、重大・再発・品質影響のあった案件を30件選び、
設備保全課／品質保証課が書く 8D レポート（D1〜D8）に仕立てる。

- 30件のうち 1/3 は顧客クレーム起点（顧客名・クレームNo・不良数量つき）の様式
- 様式は「2シート（8D報告(1)/(2)）」と「1シート縦長」、方眼紙レイアウトと列レイアウトが混在
- ファイルは様式の版ごとのサブフォルダ（VERSION_DIRS）に分けて出力する。1フォルダ＝同じ様式・同じシート構成なので、
  フォルダごと帳票取り込みに入れれば帳票の種類とシートを1回選ぶだけで全部読み取れる
- D1〜D8 の見出しは日本語名あり／なし（"D1" のみ）など、ファイルごとに表記がゆれる
- D6 の対策前後データは、正準データ上の同一設備・同一サブシステムの発生件数やチョコ停件数から実際に数える
- 抽出精度の検証用に、各ファイルの「人が読んだときの値」と「使われたラベル文字列」を _expected.jsonl に書く
- 同じ入力から毎回同じバイト列のファイルになるよう、zip のタイムスタンプと文書プロパティを固定して保存する
"""
from __future__ import annotations

import bisect
import io
import json
import math
import re
import unicodedata
import zipfile
from contextlib import suppress
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from functools import lru_cache
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.pagebreak import Break, RowBreak
from PIL import Image as PILImage, ImageDraw, ImageFont

from . import domain as D

FOLDER = "F3_8D是正処置報告書"
N_FILES = 30
N_CLAIMS = 10
TODAY = date(2026, 9, 14)             # 報告書の日付はこの日より未来にしない
PERIOD_END_DT = datetime(D.PERIOD_END.year, D.PERIOD_END.month, D.PERIOD_END.day, 23, 59)
EFFECT_DAYS = 90                      # 効果確認の比較期間（対策前・対策後それぞれ）


# ---------------------------------------------------------------------------
# 小物ヘルパー
# ---------------------------------------------------------------------------
_ZEN_DIGITS_RE = re.compile(r"(?<![A-Za-z0-9\-_.,/:])\d+(?:\.\d+)?(?![A-Za-z0-9\-_.,/:])")


def _zen(s: str) -> str:
    """IDやコードに含まれない数字（小数点を含む）だけを全角にする（記入者の癖の再現）。"""
    return _ZEN_DIGITS_RE.sub(lambda m: D.to_zenkaku(m.group(0), digits_only=True).replace(".", "．"), s)


def _tok(s: str) -> set:
    """原因となぜなぜの整合チェック用の語集合（漢字バイグラム・カタカナトライグラム・英数語）。"""
    s = re.sub(r"[（(][^）)]*[）)]", " ", s)
    out = set()
    for m in re.findall(r"[ァ-ヴー]{3,}|[A-Za-z0-9]{2,}|[一-龥]{2,}", s):
        if "一" <= m[0] <= "龥":
            out.update(m[i:i + 2] for i in range(len(m) - 1))
        elif "ァ" <= m[0] <= "ヴ" or m[0] == "ー":
            out.update(m[i:i + 3] for i in range(len(m) - 2))
        else:
            out.add(m)
    return out - {"発生", "確認", "異常", "不良"}


def _core_cause(cause: str) -> str:
    """原因文から「と推定」「（設置12年目）」などの付記を落とした本体。"""
    s = re.split(r"。", cause)[0]
    s = re.sub(r"（[^）]*）$", "", s)
    s = re.sub(r"(と推定|と判断|と思われる)$", "", s)
    return s.strip()


def _first_clause(s: str, limit: int = 48) -> str:
    s = re.split(r"。", s)[0]
    return s if len(s) <= limit else s[:limit].rstrip("、") + "…"


def _split_items(s: str) -> list[str]:
    """再発防止策の連結文（→ ／ 。区切り）を項目に分ける。"""
    items = [x.strip(" 、") for x in re.split(r"→|／|。", s or "")]
    return [x for x in items if len(x) >= 4]


def _sibs(eq: D.Equipment) -> list[str]:
    """同種設備のID。同じメーカー・同じ型式系列の号機を先に並べる。"""
    k = D.kind_of(eq)
    fam = re.split(r"[\s（]", eq.model)[0]
    cands = [e for e in D.equipment_master() if D.kind_of(e) == k and e is not eq]
    cands.sort(key=lambda e: (e.maker != eq.maker, not e.model.startswith(fam[:6]), e.equipment_id))
    return [e.equipment_id for e in cands]


def _sib_word(eq: D.Equipment, sibs: list[str]) -> str:
    """先頭の比較対象が同じメーカーなら「同型機」、違えば「同種設備」と書く。"""
    if not sibs or sibs[0] == "他号機":
        return "同型機"
    return "同型機" if D.equipment_by_id(sibs[0]).maker == eq.maker else "同種設備"


def _found_phrase(det: str) -> str:
    """detected_by（検知手段）を「誰が」欄の言い回しにする。"""
    if det.startswith("オペレーター"):
        return "作業中に発見"
    if det == "技術者の気付き":
        return "データ確認中に気付き"
    if det.endswith("からの連絡"):
        return f"{det}で発覚"
    return f"{det}で発見"


def _how_phrase(det: str) -> str:
    """detected_by を「どのように」欄の書き出しにする。"""
    if det.startswith("オペレーター"):
        return "オペレーターが異常に気付き保全へ連絡。"
    if det.endswith("からの連絡"):
        return f"{det}により発覚。"
    return f"{det}により検知。"


# 初動処置の文から、封じ込め（暫定処置）として表に載せる断片を選ぶためのキーワードと、その効果確認の書き方
_CONTAIN_WORDS = [
    (("HOLD", "リストアップ"), ["不良の拡大を防止", "対象ロットを特定し封じ込め済み"]),
    (("退避",), ["処理中ウェーハへの影響なし", "退避ウェーハは品証判定で流動可"]),
    (("振替",), ["生産への影響を最小化", "仕掛りの滞留なし"]),
    (("インヒビット", "着工停止", "DOWN登録", "停止"), ["不良品の追加発生なし", "誤着工なし"]),
    (("立入", "LOTO", "隔離", "避難"), ["二次災害なし", "安全を確保"]),
    (("切替", "閉止", "バイパス"), ["供給を維持", "影響範囲の拡大なし"]),
]
_CONTAIN_BAD = ("写真", "再発", "リセット", "したが", "するも")


def _containment_fragments(first_response: str) -> list[tuple[str, list[str]]]:
    out = []
    for x in re.split(r"→|。|、", first_response):
        x = x.strip()
        if len(x) < 5 or any(b in x for b in _CONTAIN_BAD):
            continue
        for words, effects in _CONTAIN_WORDS:
            if any(w in x for w in words):
                out.append((x, effects))
                break
    return out


def _wd(ch: str) -> int:
    return 2 if unicodedata.east_asian_width(ch) in "WFA" else 1


def _text_lines(text: str, width_units: float, fsize: float) -> int:
    """結合セル幅（列幅単位の合計）に対して、折り返し後の行数を見積もる。"""
    cap = max(4.0, width_units * (11.0 / fsize) * 0.9 - 1.5)
    n = 0
    for para in str(text).split("\n"):
        w = sum(_wd(c) for c in para)
        n += max(1, math.ceil(w / cap))
    return n


def _yen(n: int) -> str:
    return f"約{n / 10000:,.0f}万円" if n >= 10000 else f"{n:,}円"


# ---------------------------------------------------------------------------
# 表示値（日付・数値）: セルに入れる値と、人が読んだときの表示文字列を対で持つ
# ---------------------------------------------------------------------------
@dataclass
class Dt:
    """日付（with_time=True なら日時）。描画時にファイルの書式設定に従ってセル値／文字列にする。"""
    v: datetime
    with_time: bool = False


@dataclass
class Num:
    v: int
    unit: str = "分"


_DATE_FMTS = [  # (Excelの表示形式, 表示文字列)
    ("yyyy/mm/dd", lambda d: f"{d.year}/{d.month:02d}/{d.day:02d}"),
    ("yyyy/m/d", lambda d: f"{d.year}/{d.month}/{d.day}"),
    ('yyyy"年"m"月"d"日"', lambda d: f"{d.year}年{d.month}月{d.day}日"),
    ('[$-ja-JP]ggge"年"m"月"d"日"', lambda d: f"令和{d.year - 2018}年{d.month}月{d.day}日"),
]
_DATETIME_FMTS = [
    ("yyyy/mm/dd hh:mm", lambda d: f"{d.year}/{d.month:02d}/{d.day:02d} {d.hour:02d}:{d.minute:02d}"),
    ("yyyy/m/d h:mm", lambda d: f"{d.year}/{d.month}/{d.day} {d.hour}:{d.minute:02d}"),
    ('yyyy"年"m"月"d"日" h:mm', lambda d: f"{d.year}年{d.month}月{d.day}日 {d.hour}:{d.minute:02d}"),
]


# ---------------------------------------------------------------------------
# ラベル表記ゆれ
# ---------------------------------------------------------------------------
D_LABELS = {
    "jp": ["D1 チーム編成", "D2 問題の記述", "D3 暫定処置", "D4 根本原因", "D5 恒久対策の選定",
           "D6 恒久対策の実施と効果確認", "D7 再発防止（標準化）", "D8 チームの称賛・完了承認"],
    "jp_colon": ["D1：チームの結成", "D2：問題の明確化", "D3：暫定処置（封じ込め）", "D4：根本原因の追究", "D5：恒久対策の選定",
                 "D6：恒久対策の実施・効果確認", "D7：再発防止（標準化・水平展開）", "D8：チームの称賛と完了承認"],
    "code": [f"D{i}" for i in range(1, 9)],
    "en_jp": ["D1 Team（チーム編成）", "D2 Problem（問題の記述）", "D3 Containment（暫定処置）", "D4 Root Cause（根本原因）",
              "D5 Corrective Action（恒久対策の選定）", "D6 Implement & Verify（実施・効果確認）",
              "D7 Prevent Recurrence（再発防止）", "D8 Congratulate（称賛・完了承認）"],
}

LBL = {
    "report_id": ["管理No.", "8D No.", "報告書No.", "管理番号"],
    "report_date": ["発行日", "作成日", "起票日", "報告日"],
    "revision": ["版数", "Rev.", "改訂"],
    "related_incident_id": ["関連トラブルNo.", "トラブル報告No.", "設備故障報告No.", "関連報告書No."],
    "department": ["起票部署", "発行部署", "担当部署"],
    "reporter": ["作成者", "起票者", "報告者"],
    "equipment_id": ["設備No.", "設備番号", "装置No.", "号機"],
    "equipment_name": ["設備名", "装置名", "設備名称"],
    "line": ["ライン", "生産ライン"],
    "process": ["工程", "工程名"],
    "maker": ["メーカー", "設備メーカー"],
    "severity": ["重要度", "ランク", "重大度"],
    "occurred_at": ["発生日時", "不具合発生日時", "発生日"],
    "recovered_at": ["復旧日時", "生産復帰日時", "復旧完了"],
    "downtime_min": ["停止時間", "ダウンタイム", "設備停止時間"],
    "status": ["ステータス", "状況", "進捗"],
    "customer_name": ["顧客名", "お客様名", "得意先"],
    "claim_no": ["クレームNo.", "顧客クレームNo.", "苦情受付No."],
    "claim_date": ["受付日", "クレーム受付日"],
    "answer_due": ["回答期限", "回答期日"],
    "product_name": ["品名/品番", "製品名", "品番"],
    "defect_qty": ["不良数量", "不良数/受入数", "不具合数"],
    "symptom": ["問題の概要", "現象", "不具合内容", "問題の記述"],
    "alarm": ["アラーム", "発報アラーム", "アラームコード"],
    "lots": ["対象ロット", "影響ロット", "該当ロット"],
    "containment_actions": ["暫定処置", "封じ込め処置", "暫定対策"],
    "action": ["復旧処置", "応急処置内容", "処置内容（復旧）"],
    "parts": ["使用部品", "交換部品"],
    "containment_verification": ["効果確認", "暫定処置の効果確認", "効果の確認"],
    "investigation": ["調査内容", "調査・検証結果", "原因調査"],
    "cause": ["発生原因", "発生原因（なぜ起きたか）", "発生要因"],
    "why_why": ["なぜなぜ分析", "なぜなぜ分析（発生）", "なぜなぜ"],
    "escape_cause": ["流出原因", "流出原因（なぜ流出したか）", "流出要因", "検出できなかった理由"],
    "fishbone": ["特性要因（4M+2）", "要因分析（6M）", "特性要因図"],
    "countermeasure_candidates": ["対策案の比較", "恒久対策案", "対策案の評価"],
    "selection_reason": ["選定理由", "採用理由", "選定の考え方"],
    "permanent_actions": ["実施内容", "恒久対策の実施内容", "実施事項"],
    "effect_data": ["効果確認データ", "効果確認（対策前後比較）", "対策前後の比較"],
    "verification_period": ["確認期間", "効果確認期間"],
    "result": ["効果判定", "効果の確認結果", "結果"],
    "prevention": ["再発防止策", "再発防止", "歯止め"],
    "standard_docs": ["標準化（改訂文書）", "改訂文書", "標準類への反映"],
    "horizontal_deployment": ["水平展開", "横展開", "水平展開先"],
    "recognition": ["チームの称賛", "チームへの称賛", "称賛・謝辞"],
    "approver": ["承認者", "完了承認者"],
    "approval_date": ["承認日", "完了承認日", "クローズ日"],
    "approval_comment": ["承認コメント", "所見", "承認者コメント"],
    "customer_answer_date": ["顧客回答日", "最終回答日"],
    "photo_captions": ["写真", "対策前後の写真"],
}

W5_LABELS = {
    "en_jp": ["What（何が）", "When（いつ）", "Where（どこで）", "Who（誰が）", "Why（なぜ問題か）", "How（どのように）", "How many（どれだけ）"],
    "jp": ["何が", "いつ", "どこで", "誰が", "なぜ（問題点）", "どのように", "どれくらい"],
    "en": ["What", "When", "Where", "Who", "Why", "How", "How much/many"],
    "num": ["①何が(What)", "②いつ(When)", "③どこで(Where)", "④誰が(Who)", "⑤なぜ(Why)", "⑥どのように(How)", "⑦どれだけ(How many)"],
}
W5_KEYS = ["problem_what", "problem_when", "problem_where", "problem_who", "problem_why", "problem_how", "problem_how_many"]

FISH_KEYS = ["fishbone_man", "fishbone_machine", "fishbone_material", "fishbone_method", "fishbone_measurement", "fishbone_environment"]
FISH_LABELS = {
    "jp": ["人", "機械", "材料", "方法", "測定", "環境"],
    "jp_en": ["人（Man）", "機械（Machine）", "材料（Material）", "方法（Method）", "測定（Measurement）", "環境（Environment）"],
    "en_jp": ["Man（人）", "Machine（設備）", "Material（材料）", "Method（方法）", "Measurement（測定）", "Environment（環境）"],
}
FISH_LEGEND = ["◎：主要因　○：影響あり　×：検証の結果 要因でない", "（◎主要因／○寄与要因／×否定）", "凡例　◎=真因　○=要因の可能性あり　×=否定"]


# ---------------------------------------------------------------------------
# 顧客・製品（クレーム様式用。社名はすべて架空）
# ---------------------------------------------------------------------------
_PRODUCTS = {
    "MK": "MK-7720A（ロジックIC）", "SR": "SR-310C（CMOSイメージセンサ）", "PX": "PX-45N（パワーMOSFET）",
    "TQ": "TQ-9100（車載MCU）", "NB": "NB-2204（電源IC）", "HV": "HV-800D（高耐圧ドライバIC）",
}
_CUSTOMERS = {   # 製品系列 -> (顧客名, 受入拠点, クレームNo書式)
    "TQ": ("株式会社北斗オートモーティブ電子", "岐阜工場 受入検査課", "HAE-QC-{yy}{mm}-{n3}"),
    "HV": ("株式会社北斗オートモーティブ電子", "岐阜工場 受入検査課", "HAE-QC-{yy}{mm}-{n3}"),
    "SR": ("株式会社ソラリスデバイス", "長野事業所 品質管理部", "SD{yyyy}{mm}{n2}"),
    "PX": ("ミナミ電装株式会社", "本社工場 部品品質G", "{yy}-CL-{n4}"),
    "NB": ("ミナミ電装株式会社", "本社工場 部品品質G", "{yy}-CL-{n4}"),
    "MK": ("東和インダストリアル株式会社", "品質保証センター", "TWI-Q-{n5}"),
}
_INTERNAL_CUSTOMER = ("社内：組立テスト部（第2工場）", "FT課", "AT-{yy}{mm}-{n3}")

_CLAIM_DEFECT = {
    "CMP": [("配線間ショート不良（研磨スクラッチ起因）", "受入時の電気特性検査", "チップ"), ("外観不良（研磨キズ・マイクロスクラッチ）", "受入外観検査", "枚")],
    "CVD": [("パッシベーション膜厚不足による耐湿性不良", "信頼性試験（HAST）", "個"), ("パーティクル起因のショート不良", "最終テスト", "チップ")],
    "ETC": [("コンタクトオープン不良（エッチング残り）", "FT（ファイナルテスト）", "チップ"), ("ゲートリーク電流不良", "受入特性検査", "チップ")],
    "LIT": [("重ね合わせずれによるVthばらつき", "受入特性検査", "チップ"), ("パターン欠損による機能不良", "FT（ファイナルテスト）", "チップ")],
    "CLN": [("異物付着による外観不良", "受入外観検査", "枚"), ("ウォーターマーク起因のボンディング不良", "組立工程", "個")],
    "IMP": [("Vth規格外れ（ドーズずれ）", "FT（ファイナルテスト）", "チップ"), ("接合リーク電流不良", "受入特性検査", "チップ")],
    "INS": [("欠陥見逃しによる外観不良品の流出", "受入外観検査", "枚")],
}

# HOLDロットの全数確認で測る項目（設備種別ごと）
_LOT_CHECK = {"CMP": "残膜厚・欠陥数", "CVD": "膜厚・パーティクル", "ETC": "CD・欠陥数", "LIT": "重ね合わせ・CD", "CLN": "パーティクル",
              "IMP": "シート抵抗", "INS": "再検査による欠陥数"}

# 材料要因の候補（設備種別ごと）
_MATERIAL = {
    "CMP": ["スラリーのロット変動", "研磨パッドのロット差"], "CVD": ["原料ガス（SiH4/TEOS）の純度・ロット", "交換部品（Oリング）の材質"],
    "ETC": ["エッチングガスの純度", "チャンバーパーツ（フォーカスリング）の材料ロット"], "LIT": ["レジストのロット変動", "レチクル・ペリクルの汚れ"],
    "CLN": ["薬液（SC1/DHF）のロット", "PVAブラシのロット差"], "IMP": ["ドーパントガスの純度・残量", "フィラメント材のロット"],
    "INS": ["標準試料（校正ウェーハ）の劣化"], "OHT": ["FOUPの変形・個体差", "ベルト・ホイール材の品質ロット"],
    "AGV": ["バッテリーセルの個体差"], "ROB": ["FOUP・カセットの個体差"], "STK": ["FOUPの変形・個体差"],
    "UPW": ["原水水質の季節変動", "交換用樹脂・フィルタのロット"], "EXH": ["Vベルトの材質ロット"], "SCR": ["燃料・循環水の水質"],
    "PCW": ["補給水の水質変動", "水処理薬剤の濃度"], "VAC": ["封水の水質"],
}

# サブシステム別の効果確認指標: (項目名, 単位, 目標値の表示, 判定方向 "le"/"ge", 対策前範囲, 対策後範囲, 小数桁)
_METRICS = {
    "スラリー供給ライン": ("スラリー流量変動幅", "mL/min", 10, "le", (18, 45), (3, 8), 0),
    "研磨パッド": ("研磨レート面内均一性", "%", 5.0, "le", (5.8, 9.5), (2.1, 4.2), 1),
    "コンディショナ": ("ドレッサトルク変動", "N·m", 0.5, "le", (0.9, 2.4), (0.15, 0.40), 2),
    "研磨ヘッド": ("ヘッド吸着圧（絶対値）", "kPa", 65, "ge", (12, 40), (68, 80), 0),
    "ウェーハ搬送": ("搬送エラー発生率", "ppm", 50, "le", (180, 900), (0, 40), 0),
    "後洗浄ユニット": ("洗浄後欠陥数（>0.12μm）", "個/枚", 50, "le", (80, 380), (8, 35), 0),
    "終点検出": ("EPD検出失敗率", "%", 0.5, "le", (1.5, 6.0), (0.0, 0.3), 1),
    "MFC/ガス供給": ("MFC流量偏差", "%", 2.0, "le", (3.5, 12.0), (0.3, 1.5), 1),
    "RF電源": ("反射波電力", "W", 20, "le", (60, 250), (3, 15), 0),
    "ヒーター": ("ヒーター温度偏差", "℃", 3.0, "le", (8.0, 45.0), (0.5, 2.5), 1),
    "チャンバー": ("パーティクル（>0.09μm）", "個/枚", 30, "le", (45, 900), (3, 25), 0),
    "ドライポンプ": ("DP駆動電流", "A", 8.0, "le", (9.0, 18.0), (5.5, 7.5), 1),
    "APC/圧力制御": ("チャンバー圧力変動幅", "mTorr", 5, "le", (12, 40), (1, 4), 0),
    "真空搬送": ("搬送位置ずれ", "mm", 0.3, "le", (0.6, 2.0), (0.05, 0.20), 2),
    "ESC(静電チャック)": ("裏面Heリーク流量", "sccm", 1.0, "le", (1.2, 6.0), (0.2, 0.8), 1),
    "TMP/排気": ("TMP軸振動", "mm/s", 0.5, "le", (0.8, 2.2), (0.1, 0.4), 2),
    "チラー": ("チラー循環流量", "L/min", 8.0, "ge", (3.0, 6.5), (8.2, 10.0), 1),
    "リフトピン": ("リフトピン動作時間", "s", 1.5, "le", (2.2, 5.0), (0.8, 1.3), 1),
    "ウェハステージ": ("ステージ位置偏差", "nm", 10, "le", (15, 170), (2, 8), 0),
    "アライメント": ("アライメント計測再現性(3σ)", "nm", 3.0, "le", (4.5, 12.0), (0.8, 2.5), 1),
    "フォーカス": ("フォーカス誤差(3σ)", "nm", 25, "le", (35, 90), (8, 20), 0),
    "レチクル搬送": ("レチクル搬送エラー", "件/月", 1, "le", (3, 9), (0, 1), 0),
    "光源": ("パルスエネルギー安定性(3σ)", "%", 1.0, "le", (1.4, 3.5), (0.3, 0.8), 1),
    "温調チャンバー": ("チャンバー内温度変動", "℃", 0.02, "le", (0.05, 0.20), (0.005, 0.015), 3),
    "スピンチャック": ("チャック回転数偏差", "rpm", 10, "le", (25, 120), (1, 6), 0),
    "ノズル": ("洗浄後パーティクル（>0.1μm）", "個/枚", 20, "le", (35, 300), (2, 15), 0),
    "薬液供給": ("薬液濃度 管理幅外れ", "回/月", 0, "le", (3, 9), (0, 0), 0),
    "薬液槽/ヒーター": ("槽温度偏差", "℃", 1.0, "le", (2.0, 8.0), (0.1, 0.6), 1),
    "乾燥ユニット": ("ウォーターマーク欠陥数", "個/枚", 5, "le", (12, 80), (0, 3), 0),
    "配管/漏液": ("漏液センサ検知", "回/月", 0, "le", (2, 6), (0, 0), 0),
    "排気": ("槽排気風速", "m/s", 0.5, "ge", (0.2, 0.4), (0.6, 0.9), 1),
    "イオン源": ("フィラメント寿命", "h", 300, "ge", (80, 220), (320, 480), 0),
    "高圧電源": ("アーク発生回数", "回/日", 2, "le", (8, 40), (0, 1), 0),
    "ビームライン": ("ドーズ均一性(1σ)", "%", 1.0, "le", (1.3, 3.0), (0.3, 0.8), 1),
    "真空系(クライオ)": ("クライオ到達温度", "K", 15, "le", (18, 35), (9, 13), 0),
    "エンドステーション": ("プラテン搬送エラー", "件/月", 1, "le", (3, 8), (0, 1), 0),
    "光学系/光源": ("光量（初期比）", "%", 85, "ge", (50, 78), (90, 99), 0),
    "校正": ("標準試料の測定値ずれ", "%", 1.0, "le", (1.5, 4.0), (0.1, 0.6), 1),
    "プリアライナ/搬送": ("ノッチ検出失敗率", "%", 0.1, "le", (0.5, 3.0), (0.0, 0.08), 2),
    "画像処理PC": ("検査処理時間", "min/枚", 4.0, "le", (6.0, 15.0), (2.5, 3.8), 1),
    "走行部": ("走行ホイール振動値", "mm/s", 2.0, "le", (3.5, 9.0), (0.6, 1.6), 1),
    "ホイスト": ("ホイストベルト伸び", "mm", 1.0, "le", (1.8, 4.5), (0.0, 0.4), 1),
    "グリッパ": ("グリッパ把持エラー", "件/月", 2, "le", (5, 14), (0, 1), 0),
    "給電": ("給電電圧変動", "V", 5, "le", (9, 25), (1, 4), 0),
    "通信": ("通信断発生", "回/日", 1, "le", (4, 15), (0, 1), 0),
    "バッテリー": ("満充電時電圧", "V", 25.0, "ge", (22.5, 24.6), (25.3, 26.1), 1),
    "移載部": ("移載エラー", "件/月", 1, "le", (3, 10), (0, 1), 0),
    "安全センサ": ("誤停止回数", "回/日", 1, "le", (5, 20), (0, 1), 0),
    "ハンド/吸着": ("ハンド吸着圧（絶対値）", "kPa", 60, "ge", (30, 52), (65, 78), 0),
    "ロードポート": ("マッピングエラー", "件/月", 1, "le", (4, 12), (0, 1), 0),
    "アーム": ("アーム位置ずれ", "mm", 0.2, "le", (0.5, 1.5), (0.02, 0.15), 2),
    "クレーン": ("クレーン位置決めエラー", "件/月", 1, "le", (3, 10), (0, 1), 0),
    "N2パージ": ("FOUP内酸素濃度", "ppm", 1000, "le", (3000, 9000), (200, 700), 0),
    "RO膜": ("RO透過水導電率", "μS/cm", 5.0, "le", (8.0, 20.0), (1.5, 4.0), 1),
    "UV酸化": ("供給水TOC", "ppb", 1.0, "le", (1.5, 6.0), (0.3, 0.8), 1),
    "イオン交換/ポリッシャ": ("供給水比抵抗", "MΩ·cm", 18.0, "ge", (16.5, 17.8), (18.1, 18.2), 1),
    "送水ポンプ": ("送水圧力変動", "MPa", 0.02, "le", (0.04, 0.12), (0.005, 0.015), 3),
    "排気ファン": ("ファン振動値", "mm/s", 4.5, "le", (7.0, 15.0), (1.5, 3.5), 1),
    "循環水": ("循環水pH 管理値外れ", "回/月", 0, "le", (2, 7), (0, 0), 0),
    "燃焼部": ("着火失敗", "回/月", 0, "le", (2, 8), (0, 0), 0),
    "冷却塔/熱交換器": ("PCW供給温度", "℃", 20.0, "le", (21.5, 25.0), (17.5, 19.5), 1),
    "循環ポンプ": ("ポンプ軸受振動", "mm/s", 4.5, "le", (6.0, 12.0), (1.2, 3.0), 1),
    "水質管理": ("冷却水導電率", "mS/m", 50, "le", (60, 120), (25, 45), 0),
    "真空ポンプ": ("到達圧力（絶対圧）", "kPa", 20, "le", (25, 45), (8, 15), 0),
    "封水/冷却": ("封水温度", "℃", 30, "le", (33, 45), (20, 28), 0),
    "ガスBOX": ("ガスBOX内リーク検知", "回/月", 0, "le", (1, 4), (0, 0), 0),
    "操作": ("操作ミスによる停止", "件/月", 0, "le", (2, 5), (0, 0), 0),
    "ホスト通信": ("ホスト通信タイムアウト", "件/月", 1, "le", (4, 15), (0, 1), 0),
    "制御PC": ("制御PCハングアップ", "件/月", 0, "le", (2, 6), (0, 0), 0),
    "ナビゲーション": ("位置ロスト", "回/日", 1, "le", (3, 12), (0, 1), 0),
}


def _metric_row(sub: str, r, force_ng: bool = False) -> tuple:
    """効果確認の指標1行。force_ng なら対策後の値を「改善したが目標未達」の範囲に置く。"""
    name, unit, target, direc, br, ar, nd = _METRICS.get(sub, (f"{sub}関連アラーム", "件/月", 1, "le", (3, 9), (0, 1), 0))
    before = r.uniform(*br)
    if force_ng:
        lo, hi = (target, before) if direc == "le" else (before, target)
        after = lo + (hi - lo) * r.uniform(0.35, 0.8)
    else:
        after = r.uniform(*ar)
    # 判定は表示桁に丸めた値で行う（表示と判定が食い違わないように）
    b_v, a_v = round(before, nd), round(after, nd)
    fmt = (lambda x: f"{x:.{nd}f}") if nd else (lambda x: f"{round(x):,}")
    tgt = ("≦" if direc == "le" else "≧") + (f"{target:.{nd}f}" if nd else f"{target:,}")
    ok = (a_v <= target) if direc == "le" else (a_v >= target)
    if force_ng and ok:     # 丸めで目標内に入ってしまった場合は1桁ぶん外に出す
        step = 10 ** -nd
        a_v = round(target + step if direc == "le" else target - step, nd)
        ok = False
    return name, unit, fmt(b_v), fmt(a_v), tgt, ok


# ---------------------------------------------------------------------------
# 正準データの索引（効果確認データの集計用）
# ---------------------------------------------------------------------------
@dataclass
class _Index:
    mode: dict       # (設備ID, サブシステム) -> [(発生日時, トラブルID)]
    eq: dict         # 設備ID -> [発生日時]
    ms: dict         # (設備ID, サブシステム) -> [チョコ停日時]

    @staticmethod
    def _range(lst, a, b, key=lambda x: x):
        keys = [key(x) for x in lst]
        return lst[bisect.bisect_left(keys, a):bisect.bisect_left(keys, b)]

    def mode_ids(self, eid, sub, a, b) -> list[str]:
        return [i for _, i in self._range(self.mode.get((eid, sub), []), a, b, key=lambda x: x[0])]

    def eq_count(self, eid, a, b) -> int:
        return len(self._range(self.eq.get(eid, []), a, b))

    def ms_count(self, eid, sub, a, b) -> int:
        return len(self._range(self.ms.get((eid, sub), []), a, b))


@lru_cache(maxsize=None)
def _index() -> _Index:
    mode, eq, ms = {}, {}, {}
    for i in D.standard_incidents():
        mode.setdefault((i.equipment.equipment_id, i.subsystem), []).append((i.occurred_at, i.incident_id))
        eq.setdefault(i.equipment.equipment_id, []).append(i.occurred_at)
    for m in D.minor_stops():
        if m.alarm is not None:
            ms.setdefault((m.equipment.equipment_id, m.alarm.subsystem), []).append(m.occurred_at)
    return _Index(mode, eq, ms)


# ---------------------------------------------------------------------------
# 対象案件の選定
# ---------------------------------------------------------------------------
@dataclass
class Plan:
    inc: D.Incident
    claim: bool
    impl_last: date                      # 恒久対策の最終実施日
    claim_date: date | None = None       # 顧客クレーム受付日
    processed_at: datetime | None = None  # 流出ロットの当社処理日時
    effect: dict = field(default_factory=dict)


def _effect(inc: D.Incident, impl_last: date) -> dict:
    """同一設備・同一サブシステムの発生件数などを、対策前後の同じ長さの期間で数える。"""
    ix = _index()
    eid, sub = inc.equipment.equipment_id, inc.subsystem
    b1 = inc.occurred_at + timedelta(seconds=1)
    b0 = inc.occurred_at - timedelta(days=EFFECT_DAYS)
    a0 = datetime.combine(impl_last, time(0, 0))
    a1 = min(a0 + timedelta(days=EFFECT_DAYS), PERIOD_END_DT)
    return dict(
        b0=b0, b1=b1, a0=a0, a1=a1, after_days=max(0, (a1 - a0).days),
        before_ids=ix.mode_ids(eid, sub, b0, b1), after_ids=ix.mode_ids(eid, sub, a0, a1),
        between_ids=[x for x in ix.mode_ids(eid, sub, b1, a0) if x != inc.incident_id],
        ms_before=ix.ms_count(eid, sub, b0, b1), ms_after=ix.ms_count(eid, sub, a0, a1),
        eq_before=ix.eq_count(eid, b0 - timedelta(days=EFFECT_DAYS), b1), eq_after=ix.eq_count(eid, a0, a0 + timedelta(days=2 * EFFECT_DAYS)),
        eq_after_days=max(0, (min(a0 + timedelta(days=2 * EFFECT_DAYS), PERIOD_END_DT) - a0).days),
    )


def _effect_level(eff: dict) -> str:
    """対策後の同一不具合件数から効果を3段階で判定する（good: 0件か1/3以下 / partial: 減少 / ng: 減らず）。"""
    nb, na = len(eff["before_ids"]), len(eff["after_ids"])
    if eff["after_days"] < 30 or na == 0 or (na <= nb / 3 and na <= 2):
        return "good"
    return "partial" if na < nb else "ng"


def _base_ok(i: D.Incident) -> bool:
    if not (i.status == "完了" or (i.status == "経過観察" and i.cause_category != "不明")):
        return False
    if not i.cause or len(i.why_why) < 3 or i.category == "外部要因" or i.completed_at is None:
        return False
    if not (date(2023, 6, 1) <= i.occurred_at.date() <= date(2026, 8, 10)):
        return False
    return len(_tok(i.cause) & _tok(" ".join(i.why_why[1:]))) >= 2


N_PENDING = 2          # 対策後の確認期間が足りず「効果確認中」（承認欄が空）の社内報告書の件数


@lru_cache(maxsize=None)
def _plans() -> tuple:
    r = D.rng("f3_8d:select")
    pool = [i for i in D.standard_incidents() if _base_ok(i)]
    used_eq: dict = {}
    used_kind: dict = {}
    chosen: list[Plan] = []
    ng_count = partial_count = 0

    def pick(cands, lo: date, hi: date, claim: bool, kind_cap: int, pending: bool = False) -> None:
        nonlocal ng_count, partial_count
        cs = [i for i in cands if lo <= i.occurred_at.date() < hi and used_eq.get(i.equipment.equipment_id, 0) < 2
              and used_kind.get(D.kind_of(i.equipment), 0) < kind_cap and all(p.inc is not i for p in chosen)]
        r.shuffle(cs)
        # 重い案件・品質起因（クレーム）を優先
        if claim:
            cs.sort(key=lambda i: (i.category != "品質" and not i.detected_by.startswith(("SPC", "インライン", "後工程", "QC", "パーティクル"))))
        else:
            cs.sort(key=lambda i: {"重大": 0, "大": 1, "中": 2, "小": 3}[i.severity] + (0 if i.recurrence else 0.5) + r.random() * 1.2)
        for i in cs:
            pr = D.rng(f"f3_8d:plan:{i.incident_id}")
            comp = i.completed_at.date()
            if claim:
                processed = i.occurred_at - timedelta(hours=pr.randint(5, 60))
                claim_date = processed.date() + timedelta(days=pr.randint(16, 40))
                impl_last = max(comp + timedelta(days=pr.randint(12, 40)), claim_date + timedelta(days=pr.randint(14, 35)))
            else:
                processed, claim_date = None, None
                impl_last = comp + timedelta(days=pr.randint(8, 45))
            if impl_last > TODAY - timedelta(days=10):
                continue
            eff = _effect(i, impl_last)
            # 効果確認中（対策後60日未満）の案件は pending 枠でだけ採る
            if (eff["after_days"] < 60) != pending:
                continue
            level = _effect_level(eff)
            if level == "ng" and ng_count >= 3:
                continue
            if level == "partial" and partial_count >= 3:
                continue
            ng_count += level == "ng"
            partial_count += level == "partial"
            chosen.append(Plan(i, claim, impl_last, claim_date, processed, eff))
            used_eq[i.equipment.equipment_id] = used_eq.get(i.equipment.equipment_id, 0) + 1
            used_kind[D.kind_of(i.equipment)] = used_kind.get(D.kind_of(i.equipment), 0) + 1
            return

    def buckets(lo: date, hi: date, n: int) -> list[tuple[date, date]]:
        step = (hi - lo).days / n
        return [(lo + timedelta(days=int(step * k)), lo + timedelta(days=int(step * (k + 1)))) for k in range(n)]

    claim_pool = [i for i in pool if D.kind_of(i.equipment) in D.PROCESS_KINDS and i.lots_affected]
    for lo, hi in buckets(date(2023, 6, 1), date(2026, 5, 20), N_CLAIMS):
        pick(claim_pool, lo, hi, True, kind_cap=3)
    internal_pool = [i for i in pool if i.severity in ("重大", "大") or (i.severity == "中" and (i.recurrence or i.lots_affected))]
    for lo, hi in buckets(date(2023, 6, 1), date(2026, 6, 1), N_FILES - N_CLAIMS - N_PENDING):
        pick(internal_pool, lo, hi, False, kind_cap=4)
    # 直近の案件は恒久対策を打ったばかりで効果確認中
    for _ in range(N_PENDING):
        pick(internal_pool, date(2026, 5, 20), date(2026, 8, 10), False, kind_cap=6, pending=True)
    # バケットで埋まらなかった分は期間全体から補う
    for _ in range(20):
        if sum(p.claim for p in chosen) >= N_CLAIMS:
            break
        pick(claim_pool, date(2023, 6, 1), date(2026, 5, 20), True, kind_cap=4)
    for _ in range(40):
        if len(chosen) >= N_FILES:
            break
        pick(internal_pool, date(2023, 6, 1), date(2026, 6, 1), False, kind_cap=6)
    n_pending = sum(p.effect["after_days"] < 60 for p in chosen)
    if len(chosen) != N_FILES or sum(p.claim for p in chosen) != N_CLAIMS or n_pending != N_PENDING:
        raise RuntimeError(f"8D対象案件の選定に失敗: {len(chosen)}件（クレーム{sum(p.claim for p in chosen)}件、効果確認中{n_pending}件）")
    chosen.sort(key=lambda p: p.inc.occurred_at)
    return tuple(chosen)


# ---------------------------------------------------------------------------
# 報告書の中身（レイアウトに依存しない）
# ---------------------------------------------------------------------------
_DOC_TYPES = {
    "pm": ("PM点検基準書（{kind}）", "MS-{kind}-PM-{n:03d}"),
    "wi": ("作業標準書（{sub} 保全手順）", "WI-{kind}-{n:03d}"),
    "fdc": ("FDC監視項目・管理基準（{kind}）", "PE-FDC-{kind}-{n:02d}"),
    "parts": ("定期交換部品基準表", "MS-PRT-{kind}-{n:02d}"),
    "check": ("日常点検チェックシート（{eq}）", "MF-CHK-{eq}"),
    "fmea": ("設備FMEA（{kind}）", "PE-FMEA-{kind}-{n:02d}"),
    "qcp": ("QC工程表（{line}）", "QA-QCP-{line}-{n:02d}"),
    "ship": ("出荷検査基準書", "QA-INS-{n:03d}"),
    "edu": ("保全教育資料（{sub}）", "MS-EDU-{n:03d}"),
}


class _DocRegistry:
    """改訂文書の版数と報告書の年別連番を、報告書の時系列順に積み上げて管理する。

    同じ文書Noは毎回同じ番号で、使われるたびに版が進む。報告書Noは年ごと・様式ごとに連番（間に他部署の8Dが入るので飛び番あり）。
    """

    def __init__(self):
        self.rev: dict = {}
        self.seq: dict = {}

    def next_no(self, prefix: str, year: int, r) -> int:
        key = (prefix, year)
        self.seq[key] = self.seq.get(key, r.randint(2, 9)) + r.randint(1, 6)
        return self.seq[key]

    def use(self, dtype: str, eq: D.Equipment, sub: str) -> tuple[str, str, int]:
        kind = D.kind_of(eq)
        title_t, no_t = _DOC_TYPES[dtype]
        rr = D.rng(f"f3_8d:doc:{dtype}:{kind}:{sub if dtype in ('wi', 'edu') else ''}")
        line = eq.line if eq.line.startswith("L") and len(eq.line) == 2 else "COM"
        no = no_t.format(kind=kind, n=rr.randint(1, 60), eq=eq.equipment_id, line=line)
        title = title_t.format(kind=kind, sub=sub, eq=eq.equipment_id, line=line)
        cur = self.rev.get(no, rr.randint(1, 6))
        self.rev[no] = cur + 1
        return title, no, cur


_ADOPT_BY_CC = {
    "劣化": ["{part}の交換周期を{a}ヶ月→{b}ヶ月に短縮し、定期交換部品に登録", "{metric}をFDCでトレンド監視し、管理値の{p}%到達でワーニング発報",
             "劣化診断（{metric}測定）をPM点検項目に追加"],
    "摩耗": ["{part}の交換基準を暦日管理から使用時間（枚数）管理に変更", "{sub}の摩耗量測定をPM点検項目に追加（基準値明記）",
             "{metric}をFDCで監視し、上昇傾向で予防交換"],
    "汚れ": ["{sub}の清掃周期を{a}週→{b}週に変更", "清掃手順（写真付き）を作業標準書に追加", "{metric}の日常点検を追加し汚れの兆候を早期検知"],
    "異物": ["{sub}周辺の清掃範囲を拡大し、清掃後の異物確認をチェックシート化", "パーティクルQCの頻度を週1→毎日に変更", "部品交換時の異物混入防止手順（養生・拭き取り）を標準化"],
    "締結緩み": ["締付トルク管理（規定トルク明記・合いマーク）を導入", "振動部の締結部をPM時に増し締め点検する項目を追加", "緩み止め（ダブルナット／ねじロック剤）を採用"],
    "断線": ["ケーブル屈曲部に保護チューブを追加し、固定方法を変更", "可動部ケーブルを高屈曲品に変更", "ケーブル外観点検をPM項目に追加"],
    "設定ミス": ["パラメータ変更時のダブルチェック（チェックシート）を導入", "装置パラメータのバックアップと差分確認をPM後の必須作業に追加", "設定変更権限をMESで管理"],
    "作業ミス": ["作業手順書に注意点・禁止事項を写真付きで追記", "該当作業を認定制とし、未認定者の単独作業を禁止", "作業後の相互確認（指差呼称）を標準化"],
    "ソフト不具合": ["メーカー修正版ソフト（Ver.{ver}）へ更新", "ソフト更新時の検証手順（ダミー運転）を標準化", "ログ監視で同事象の予兆をワーニング化"],
    "設計起因": ["メーカー対策品（改良型）へ変更", "{sub}の構造変更（メーカー改造キット適用）", "同型機の改造計画を作成し順次実施"],
    "調整不良": ["調整手順に測定値の記録欄と基準値を追加", "調整後の確認項目（{metric}）を明確化", "調整作業のOJT教育を実施"],
    "施工不良": ["施工後の検査項目（リーク・トルク）を追加", "施工業者への要求仕様を明文化", "施工時の立会い確認を必須化"],
    "能力不足": ["{sub}の能力（容量）アップを実施", "負荷分散の運用ルールを設定", "余裕度を定期的に評価する仕組みを追加"],
    "前工程起因": ["前工程との品質情報連携（異常時の連絡ルール）を明確化", "受入時の確認項目を追加", "前工程の管理値を見直し"],
    "外部要因": ["UPS・瞬低対策の強化", "復電時の自動復帰手順を整備", "非常時連絡網を見直し"],
    "不明": ["{metric}の監視を強化しデータを蓄積", "再発時のログ取得手順を整備", "メーカーと共同で解析を継続"],
}
_CLAIM_ESCAPE_ACTIONS = ["{metric}異常の発生時は、該当期間のロットの{check}を全数確認してから出荷判定する", "インライン欠陥検査のサンプリング率を{a}%→{b}%に変更",
                         "SPC管理限界を見直し（±3σ→±2σで警告）", "装置異常時のロットHOLD遡及範囲を{h}時間→{h2}時間に拡大"]


# ---------------------------------------------------------------------------
# 自由記述の文面（案件の条件に合わせて選ぶ。同じ定型文が多数のファイルに並ばないように）
# ---------------------------------------------------------------------------
def _part_short(name: str) -> str:
    """部品マスタの品名を文中で使う呼び方にする（「研磨パッド IC-type 30inch」→「研磨パッド」、「MFC (SiH4 500sccm)」→「MFC」）。"""
    s = re.sub(r"\s*[（(][^）)]*[）)]", "", name)
    toks = s.split()
    if not toks:
        return name
    # 2語目以降は日本語を含む語だけ残す（型式・サイズの英数字は落とす）
    return " ".join([toks[0]] + [t for t in toks[1:] if re.search(r"[^\x00-\x7F]", t)])


def _claim_why(r, *, internal: bool, cname: str, insp: str, defect: str, unit: str, bad: int, total: int, product: str,
               automotive: bool, n_hold: int) -> str:
    """5W2H の「なぜ問題か」（クレーム様式）。顧客・製品・不良の中身から候補を作り、1〜2文を選ぶ。"""
    rate = bad / total * 100
    if internal:
        whys = ["後工程（組立テスト）で発見された流出不良であり、ウェーハ工程内で止められなかった",
                "後工程での発見だが、同時期のロットが出荷待ちで出荷判定に影響した",
                f"FT歩留まりが{r.uniform(1.5, 6.0):.1f}ポイント低下し、組立テスト部の生産計画に影響した",
                f"同時期に払い出した{r.randint(3, 9)}ロットも再テスト対象となり、後工程の工数が増えた",
                f"{defect}は組立後に見つかると{unit}単位の廃棄になり、損失が大きい"]
    else:
        whys = ["顧客ラインで発見された流出不良であり、顧客の生産計画に影響した",
                f"不良率{rate:.2f}%で、{cname}との品質目標（{r.choice([10, 50, 100])}ppm以下）を大きく超えた",
                f"{cname}で同一ロットの受入品が全数保留となり、顧客側で選別工数が発生した",
                "当社の出荷検査をすり抜けて流出しており、工程内の検出力不足を示している",
                f"{cname}への同種不良の流出は{r.choice(['2', '3', '約4'])}年ぶりで、サプライヤ評価（品質スコア）に影響する",
                f"{product}は{cname}向けの主力品種で、代替品の緊急手配が必要になった"]
        if bad >= 100:
            whys.append(f"不良数が{bad:,}{unit}と多く、顧客から選別費用の負担を求められる可能性がある")
        if "信頼性" in insp:
            whys.append(f"{defect}は信頼性に関わる不良で、市場に流出した場合の影響が大きい")
        if automotive:
            whys += ["車載向け製品であり、顧客から8Dでの回答を要求された",
                     f"車載向け（{product}）のため、顧客の変化点管理・PPAPの見直し対象となった"]
    if n_hold >= 2:
        whys.append(f"社内でも{n_hold}ロットがHOLDとなり、出荷計画の見直しが必要になった")
    picked = r.sample(whys, k=2) if r.random() < 0.3 else [r.choice(whys)]
    return "。".join(picked)


# 作業者（人）の要因のうち、発生原因でないと検証したもの（× 要因 → 検証結果）の候補
def _fish_man_candidates(inc: D.Incident, k: str, sub: str, part_s: str | None, cc: str, metric: str) -> list[str]:
    inv = inc.investigation
    maker = D.MAKER_SHORT.get(inc.equipment.maker, inc.equipment.maker)
    out = []
    m = re.search(r"前回PMは(\d+)日前に実施", inv)
    if m:
        out.append(f"× 前回PM（{m.group(1)}日前）での組付け・締付けミス → 作業記録に特記事項なし")
    if "班長への聞き取り" in inv:
        out.append("× 発生前の誤操作・段取り替え → 班長への聞き取りで該当する作業なし")
    if any(p.role == "FE" for p in inc.assignees):
        out.append(f"× {maker}FEの前回サービス作業の不備 → サービスレポートを確認し、作業内容に問題なし")
    if inc.detected_by.startswith("オペレーター"):
        out.append("× オペレーターの誤操作 → 操作ログ上、発生前に異常な操作なし")
    if part_s:
        out.append(f"× {part_s}交換時の取付け不良 → 前回交換時の記録写真で取付状態に問題なし")
    if cc in ("劣化", "摩耗", "断線"):
        out.append(f"× 日常点検での見落とし → 直近の点検記録では{metric}は基準内")
    if k in D.UTILITY_KINDS:
        out.append("× 中央監視の警報見落とし → 警報発報から対応開始まで手順どおり")
    elif k in D.TRANSPORT_KINDS:
        out.append("× 手動搬送・マニュアル操作による干渉 → MCSログで手動操作なし")
    out.append(f"× 作業者の力量不足 → 対応者は{inc.equipment.category}の保全認定取得者で問題なし")
    return out


# 方法の要因（管理・手順の穴）: 原因区分ごとの候補。{part}{sub}{usage}{metric} を埋めて使う
_METHOD_GAPS = {
    "劣化": ["{part}の交換周期が暦日管理で、{usage}を考慮していない", "{sub}の劣化診断（{metric}の測定）が点検項目にない",
             "{part}の交換実績（寿命データ）を蓄積・評価していない"],
    "摩耗": ["{part}の摩耗量の交換限度値が点検基準書に明記されていない", "交換周期が暦日管理で、{usage}を考慮していない",
             "{sub}の摩耗確認が目視判定のみで、数値で記録していない"],
    "断線": ["可動部ケーブルの点検が目視のみで、屈曲部の内部断線は確認できない", "{sub}の配線経路・固定方法が標準化されていない",
             "ケーブル類の交換周期が設定されていない"],
    "汚れ": ["清掃手順に{sub}の清掃が含まれていない", "{sub}の清掃周期が汚れの付き方（{usage}）に合っていない",
             "清掃後の確認方法（拭き取り・目視の基準）が決まっていない"],
    "異物": ["清掃手順に{sub}周辺の異物除去が含まれていない", "部品交換時の養生・異物混入防止の手順がない",
             "異物の持込み経路（部品・工具・ウェーハ）の管理ルールがない"],
    "締結緩み": ["締付トルクの規定がなく作業者任せ", "合いマーク等による緩み確認の方法が決まっていない", "振動部の増し締め点検がPM項目にない"],
    "作業ミス": ["手順書に注意点・確認項目の記載がなかった", "作業後の確認を1名で完結できる手順になっていた", "作業チェックシートの確認欄が形骸化していた"],
    "設定ミス": ["パラメータ変更時のダブルチェックがルール化されていない", "装置パラメータの変更履歴を残す仕組みがない", "手順書に設定値の確認手順がない"],
    "調整不良": ["調整手順に測定値の記録欄・判定基準がない", "調整後の確認項目（{metric}）が決まっていない", "調整作業のやり方が担当者によって違う"],
    "設計起因": ["同一箇所の再発時にメーカー対策を要求する判断基準がない", "メーカー対策品の情報を入手・評価する仕組みがない",
                 "{sub}の弱点が設備FMEAに反映されていない"],
    "能力不足": ["負荷増加（増産）時に{sub}の能力余裕を確認する手順がない", "設備仕様の余裕度を定期的に評価していない"],
    "その他": ["変更管理・確認手順が明確でなかった", "異常時の判断基準（停止・継続）が決まっていない",
               "{sub}の点検項目・基準が設備導入時のまま見直されていない"],
}
# 設備種別ごとの「使用量」の言い方（交換周期を暦日でなく何で管理すべきか）
_USAGE = {"CMP": "処理枚数", "CLN": "処理枚数", "CVD": "RF積算時間", "ETC": "RF積算時間", "LIT": "ショット数", "IMP": "ビーム照射時間",
          "INS": "稼働時間", "OHT": "走行距離", "AGV": "走行時間", "ROB": "搬送回数", "STK": "搬送回数"}


def _approval_comment(r, cls: str, *, habit: dict, ad: date | None, claim: bool, internal: bool, cname: str, metric: str,
                      sub: str, nb: int, na: int, ids_txt: str, ng_rows: list, hz: str, sib_txt: str, doc_title: str | None,
                      maker: str | None, related: str | None, shift: str, eff_end: date) -> str:
    """D8 承認コメント。効果の確認結果（cls）で評価の一文を選び、案件の状況に応じた指示を0〜2件付ける。

    cls: effective（効果あり・再発ゼロ）/ monitoring（効果ありだが発生・未達指標が残る）/ partial（減少したが再発あり）/
         conditional（効果なし・条件付き承認）/ pending（効果確認中・未承認）
    承認者（課長）の書き癖（番号の振り方・誤変換）を反映する。
    """
    f = lambda d: D.fmt_date(d, style=0)         # noqa: E731
    if cls == "pending":
        text = r.choice([f"効果確認期間終了（{f(eff_end)}）後に完了承認予定。",
                         f"中間確認。{f(eff_end)}までの監視結果がそろった時点で再提出のこと。",
                         f"内容は確認した。{sub}の再発有無を{f(eff_end)}まで見てから完了判断とする。",
                         "効果確認中のため承認保留。監視データは月例会議で報告すること。"])
        return D._noise(text, habit, r)
    head = {
        "effective": ["対策内容・効果確認とも妥当。本8Dをクローズとする。", "効果確認データを確認した。完了承認とする。",
                      f"対策後の再発ゼロ、{metric}も目標内を確認。クローズ可。", "真因に対する対策になっており妥当。承認する。",
                      "対策前後のデータで効果が確認できている。承認。", "なぜなぜから対策まで筋が通っている。クローズとする。",
                      "内容確認済み。完了とする。"],
        # 監視継続の理由（発生が残る / 目標未達の指標が残る）に合う書き方だけを使う
        "monitoring": ([f"効果は認められるが、発生がゼロではない。{r.choice([3, 6])}ヶ月の監視継続を条件に承認。",
                        f"対策後{na}件の発生は減少の範囲内と判断。監視継続を条件に承認する。",
                        "おおむね効果あり。残った発生分の要因を確認のうえ、監視継続で承認する。"] if na else [])
                      + (["目標未達の指標が残っている。監視結果の報告をもって正式クローズとする。",
                          f"{'・'.join(ng_rows[:2])}が目標未達。次回定例で推移を報告のうえクローズとする。"] if ng_rows else [])
                      + [f"承認。ただし{metric}のトレンドは{(ad + timedelta(days=40)).month}月の定例で報告すること。",
                         f"内容は妥当。念のため{r.choice([2, 3])}ヶ月は{sub}の監視を継続し、異常があれば再度報告すること。"],
        "partial": [f"同一不具合は{nb}件→{na}件と減ったが、再発が残っている。監視継続を条件に承認。",
                    f"効果は限定的。再発分（{ids_txt}）の個別要因を分析し、{f(ad + timedelta(days=30))}までに報告のこと。",
                    f"件数は減っており一定の効果は認める。再発分の追加点検結果を次回報告すること。",
                    "対策の方向性は妥当だが効果が十分ではない。追加対策の要否を係内で検討し報告すること。"],
        "conditional": [f"再発が続いているため条件付きで承認。追加対策の計画を{f(ad + timedelta(days=30))}までに提出すること。",
                        f"対策の効果が確認できていない。D4に戻って真因を再検討し、{f(ad + timedelta(days=30))}までに再提出のこと。",
                        f"同一不具合が止まっていない。{sub}の構造見直しも含めた追加対策を{f(ad + timedelta(days=30))}までに計画すること。"]
                       + ([f"条件付き承認。{maker}を交えて追加対策を協議し、計画を{f(ad + timedelta(days=30))}までに提出のこと。"] if maker else []),
    }[cls]
    orders = []
    if any(w in hz for w in ("予定", "検討中")) and sib_txt:
        orders.append(f"{sib_txt}への水平展開の結果は月例会議で報告すること。")
    if doc_title:
        orders.append(r.choice([f"{doc_title}の改訂内容を係内に周知し、周知記録を残すこと。", "改訂文書の周知状況を次回PMで確認すること。"]))
    if claim and not internal:
        short = cname.replace("株式会社", "")
        orders.append(r.choice([f"{short}様への最終回答は品証と内容をすり合わせてから提出のこと。", "顧客回答書の写しを本報告書に添付すること。"]))
    elif claim:
        orders.append("組立テスト部へは対策完了の連絡を忘れずに。")
    if maker and cls in ("effective", "monitoring"):
        orders.append(f"{maker}へは対策内容を共有し、同型機への恒久対策の提案を求めること。")
    if related:
        orders.append(f"前回（{related}）の対策で止められなかった理由も教育資料に残すこと。")
    if shift in ("夜勤", "休日"):
        orders.append(f"{shift}帯の一次対応手順もあわせて見直すこと。")
    if cls == "effective":
        orders.append(r.choice(["良い事例なので課内の改善発表会で共有すること。", "水平展開の進捗は月例会議で報告すること。"]))
    n = r.choices([0, 1, 2], weights=[3, 5, 2])[0] if cls != "conditional" else r.choice([0, 1])
    picked = r.sample(orders, k=min(n, len(orders)))
    text = r.choice(head)
    if len(picked) == 2:
        text += "\n" + D._numbered(picked, habit["numbering"])
    elif picked:
        text += picked[0]
    return D._noise(text, habit, r)


def _pick_team(inc: D.Incident, claim: bool, r) -> dict:
    ppl = D.people()
    eq = inc.equipment
    k = D.kind_of(eq)
    champion = next(p for p in ppl if p.role == "課長")
    if k in D.UTILITY_KINDS:
        sect = "設備保全課 施設係"
    else:
        sect = "設備保全課 保全1係" if D.fab_of(eq) == "Fab1" else "設備保全課 保全2係"
    chief = next((p for p in ppl if p.section == sect and p.role in ("係長",)), None) or next(p for p in ppl if p.section == sect and p.role == "主任")
    qa_lead = next(p for p in ppl if p.section == "品質保証課" and p.role == "主任")
    qa_staff = next(p for p in ppl if p.section == "品質保証課" and p.role == "担当")
    internal = [p for p in inc.assignees if p.role not in ("FE", "課長") and p is not chief]
    fe = [p for p in inc.assignees if p.role == "FE"]
    pe = [p for p in inc.assignees if p.section == "生産技術課"] or [r.choice([p for p in ppl if p.section == "生産技術課"])]
    members = []   # (役割, 人, 担当内容)
    if claim:
        leader = qa_lead
        members.append(("サブリーダー", chief, "設備調査・恒久対策の取りまとめ"))
        author = qa_staff
    else:
        leader = chief
        author = next((p for p in internal if p.section == sect), None) or next(p for p in ppl if p.section == sect and p.role == "主任")
    duties = ["現象調査・部品交換", "復旧作業・動作確認", "点検基準の改訂", "水平展開の実施", "データ収集・記録"]
    for p in internal:
        if p in (leader, author) or any(p is m[1] for m in members) or p.section == "生産技術課":
            continue
        members.append(("メンバー", p, duties[len(members) % len(duties)]))
        if len(members) >= 3:
            break
    if author is not leader and not any(author is m[1] for m in members):
        members.insert(0 if not claim else 1, ("メンバー（事務局）" if r.random() < 0.5 else "メンバー", author,
                                               "報告書作成・進捗管理" if not claim else "顧客窓口・報告書作成"))
    members.append(("メンバー", pe[0], r.choice(["FDC/SPCデータ解析", "プロセス影響評価", "条件検証・データ解析"])))
    if claim or inc.lots_affected:
        if not claim:
            members.append(("メンバー", qa_staff, "ロット影響判定・出荷可否判断"))
    rep = inc.reporter
    if rep.section.startswith("製造課") and len(members) < 7:
        members.append(("メンバー", rep, "発見者・発生状況の聞き取り"))
    for f in fe[:1]:
        members.append(("協力（メーカー）", f, f"{D.MAKER_SHORT.get(f.section, f.section)}側の解析・部品手配"))
    return dict(champion=champion, leader=leader, members=members, author=author, chief=chief)


def _stamp_name(p: D.Person) -> str:
    return D.surname(p)


def build_content(plan: Plan, idx: int, docs: _DocRegistry) -> dict:
    """1件分の 8D 報告書の中身を作る（文字列・表・日付オブジェクト）。"""
    inc, claim = plan.inc, plan.claim
    r = D.rng(f"f3_8d:content:{inc.incident_id}")
    eq = inc.equipment
    k = D.kind_of(eq)
    sub = inc.subsystem
    sibs = _sibs(eq) or ["他号機"]
    zen = r.random() < 0.2
    Z = _zen if zen else (lambda s: s)
    team = _pick_team(inc, claim, r)
    leader, author, champion = team["leader"], team["author"], team["champion"]
    sn = D.surname
    metric = _METRICS.get(sub, (f"{sub}関連アラーム",))[0]
    parts = inc.parts_used
    main_part = parts[0][0].name if parts else f"{sub}構成部品"
    C: dict = dict(claim=claim, zen=zen)
    comp = inc.completed_at
    occ = inc.occurred_at
    eff = plan.effect

    # --- ヘッダ ---
    year = (plan.claim_date or comp.date()).year
    if claim:
        C["report_id"] = f"QA-CL-{year}-{docs.next_no('QA-CL', year, r):03d}"
    else:
        C["report_id"] = f"8D-{year}-{docs.next_no('8D', year, r):03d}"
    C["related_incident_id"] = inc.incident_id
    C["department"] = "品質保証部 品質保証課" if claim else f"製造部 {author.section.split(' ')[0]}"
    C["reporter"] = author.name
    C["equipment_id"] = eq.equipment_id
    C["equipment_name"] = eq.name
    C["line"] = eq.line
    C["process"] = eq.process
    C["maker"] = f"{D.MAKER_SHORT.get(eq.maker, eq.maker)} {eq.model}"
    C["severity"] = inc.severity if r.random() < 0.6 else {"重大": "S（重大）", "大": "A（大）", "中": "B（中）", "小": "C（小）"}[inc.severity]
    C["occurred_at"] = Dt(occ, True)
    C["recovered_at"] = Dt(comp, True)
    C["downtime_min"] = Num(inc.downtime_min)
    C["revision"] = r.choice(["Rev.1", "Rev.2", "第1版", "第2版", "Rev.0", "第3版"])

    # 効果確認の期間が取れているか（取れていなければ効果確認中）と、効果の程度（good / partial / ng）
    complete = eff["after_days"] >= 60
    level = _effect_level(eff)
    ng = level == "ng"

    # --- クレーム情報 ---
    lots = list(inc.lots_affected)
    internal = False
    CW, SHIP = "顧客", "出荷"      # クレーム相手の呼び方と「出荷」の言い方（社内後工程からの苦情なら 組立テスト部／払出）
    if claim:
        prefix = lots[0][:2]
        cust = _INTERNAL_CUSTOMER if r.random() < 0.15 else _CUSTOMERS[prefix]
        internal = cust is _INTERNAL_CUSTOMER
        if internal:
            CW, SHIP = "組立テスト部", "払出"
        cd = plan.claim_date
        pd = plan.processed_at
        n = r.randint(1, 999)
        C["customer_name"] = cust[0]
        C["customer_site"] = cust[1]
        C["claim_no"] = cust[2].format(yy=f"{cd.year % 100:02d}", yyyy=cd.year, mm=f"{cd.month:02d}", n2=f"{n % 100:02d}", n3=f"{n:03d}",
                                       n4=f"{n * 7 % 10000:04d}", n5=f"{n * 13 + 10000:05d}")
        C["claim_date"] = Dt(datetime.combine(cd, time()))
        C["answer_due"] = Dt(datetime.combine(cd + timedelta(days=r.choice([14, 21, 30])), time()))
        C["product_name"] = _PRODUCTS[prefix]
        defect, insp, unit = r.choice(_CLAIM_DEFECT[k])
        wk = pd.isocalendar()[1]
        esc_lot = f"{prefix}{pd.year % 100:02d}{wk:02d}{r.randint(1, 399):03d}.{r.randint(1, 3)}"
        C["escaped_lot"] = esc_lot
        if unit == "枚":
            total, bad = 25, r.randint(2, 9)
        elif unit == "個":
            total, bad = r.choice([2000, 3000, 5000, 10000]), r.randint(4, 80)
        else:
            total, bad = r.choice([12000, 18500, 24000, 30000]), r.randint(12, 480)
        C["defect_qty"] = f"{bad:,}{unit} / {total:,}{unit}（{bad / total * 100:.2f}%）"
        C["defect"] = defect
        found = cd - timedelta(days=r.randint(1, 3))
        C["processed_at"] = pd
        lots = [esc_lot] + lots[:1]
    C["lots"] = "、".join(lots) if lots else ""

    # --- D1 ---
    C["team"] = team

    # --- D2 問題の記述 ---
    short = _first_clause(inc.symptom)
    if claim:
        cname = C["customer_name"].replace("社内：", "")
        C["symptom"] = Z(f"{cname}の{insp}にて、{'前工程からの払出品' if internal else '当社出荷品'} {C['product_name']}（ロット {C['escaped_lot']}）に"
                         f"{C['defect']}が発生（{C['defect_qty']}）。") + "\n" + \
            Z(f"当社調査の結果、同ロットを処理した {eq.equipment_id}（{eq.name}）の{sub}不具合（{inc.incident_id}）に起因すると判断した。")
        what = f"{C['product_name']}の{C['defect']}"
        when = (f"当社処理：{D.fmt_date(pd, style=0)} {pd.hour}:{pd.minute:02d}（{eq.equipment_id}）／{'後工程' if internal else '顧客'}発見：{D.fmt_date(found, style=0)}"
                f"／受付：{D.fmt_date(cd, style=0)}")
        where = f"{cname} {C['customer_site']}（発見）／流出元：{'' if internal else '当社 '}{D.fab_of(eq)} {eq.line} {eq.process}工程"
        who = f"{cname} {C['customer_site']}が{insp}で発見"
        why = _claim_why(r, internal=internal, cname=cname, insp=insp, defect=defect, unit=unit, bad=bad, total=total,
                         product=C["product_name"], automotive=prefix in ("TQ", "HV"), n_hold=len(inc.lots_affected))
        how = f"{insp}で不良を検出。" + f"原因は当社 {sub} の不具合（{_core_cause(inc.cause)}）"
        how_many = f"不良 {C['defect_qty']}、対象ロット {C['escaped_lot']}（{SHIP}済み）" + (f"、社内HOLD {len(inc.lots_affected)}ロット" if inc.lots_affected else "")
    else:
        C["symptom"] = inc.symptom
        what = f"{eq.equipment_id}（{eq.name}）{sub}の不具合：{short}"
        when = f"{D.fmt_date(occ, style=0)} {occ.hour}:{occ.minute:02d} 発生（{inc.shift}）"
        if k in D.UTILITY_KINDS:
            where = f"{eq.area}　{eq.process}系統（{eq.line}）"
        else:
            area = eq.area.split(" ", 1)[1] if " " in eq.area else eq.area
            where = f"{D.fab_of(eq)} {area}　{eq.line} {eq.process}工程"
        who = f"{inc.reporter.section} {sn(inc.reporter)}（{_found_phrase(inc.detected_by)}）"
        whys = []
        if inc.downtime_min >= 600 and k in D.UTILITY_KINDS:
            whys.append(f"重要度{eq.criticality}のユーティリティ設備が約{inc.downtime_min / 60:.0f}時間停止し、{eq.line}の使用先装置に影響")
        elif inc.downtime_min >= 600:
            whys.append(f"重要度{eq.criticality}設備が約{inc.downtime_min / 60:.0f}時間停止し、{eq.line}の生産計画に遅れが発生")
        elif k in D.UTILITY_KINDS:
            whys.append(f"{eq.process}の停止（{inc.downtime_min:,}分）で使用先装置が待機")
        else:
            whys.append(f"{eq.line} {eq.process}工程の設備停止（{inc.downtime_min:,}分）")
        if inc.recurrence and inc.related_incident_id:
            whys.append(f"同一不具合の再発（前回 {inc.related_incident_id}）")
        if inc.lots_affected:
            whys.append(f"製品品質への影響があり、{len(inc.lots_affected)}ロットをHOLD")
        why = "。".join(whys)
        how = _how_phrase(inc.detected_by)
        if inc.alarm and inc.alarm.code not in short:
            how += f"{inc.alarm.code}（{inc.alarm.message}）発報。"
        how += r.choice(["{s}", "{s}の状態で停止", "{s}。{cons}"]).format(
            s=_first_clause(inc.symptom, 60).rstrip("。"),
            cons=r.choice(["処理を中断", "装置DOWN登録", "仕掛りを他号機へ振替", "搬送を一時停止", "生産を停止"]))
        hm = [f"停止時間 {inc.downtime_min:,}分（約{inc.downtime_min / 60:.1f}h）"]
        if inc.lots_affected:
            hm.append(f"影響ロット {len(inc.lots_affected)}（{C['lots']}）")
        if inc.scrap_wafers:
            hm.append(f"廃棄ウェーハ {inc.scrap_wafers}枚")
        hm.append(f"損失額 {_yen(inc.cost_yen)}")
        how_many = "、".join(hm)
    C["w5"] = [Z(x) for x in (what, when, where, who, why, how, how_many)]
    C["alarm"] = f"{inc.alarm.code}　{inc.alarm.message}" if inc.alarm else r.choice(["なし", "－", "アラーム発報なし"])

    # --- D3 暫定処置 ---
    rows = []
    d_occ = occ.date()
    tdate = lambda d: d  # noqa: E731  表の日付はレイアウト側で書式を決める
    # 担当者の割り振り: 設備の作業は保全、出荷・検査は品証、監視・データは生産技術
    maint = [m[1] for m in team["members"] if m[1].section.startswith("設備保全課")] or [team["chief"]]
    qa_people = [m[1] for m in team["members"] if m[1].section == "品質保証課"] or [team["leader"] if claim else author]
    pe_people = [m[1] for m in team["members"] if m[1].section == "生産技術課"] or maint
    sib_w = _sib_word(eq, sibs)
    if claim:
        cd = plan.claim_date
        unit_bad = "枚" if "枚" in C["defect_qty"] else "個"
        rows.append([f"{CW}在庫品（{C['escaped_lot']}）の{'選別' if internal else '返却・選別'}", tdate(cd + timedelta(days=r.randint(1, 3))), sn(author),
                     f"選別の結果、不良{r.randint(1, 5)}{unit_bad}を追加で除去。良品は{CW}承認のうえ使用継続"])
        rows.append([f"社内在庫・仕掛品の全数検査（{r.randint(3, 12)}ロット）", tdate(cd + timedelta(days=r.randint(1, 4))), sn(author),
                     r.choice(["追加の不良なし", f"1ロットで同不良を検出し{SHIP}停止", "全ロット規格内"])])
        rows.append([f"{SHIP}前検査を抜取から全数に暫定切替（{D.fmt_date(cd + timedelta(days=1), style=4)}〜）", tdate(cd + timedelta(days=1)), sn(team["leader"]),
                     f"切替後の{SHIP}品で{CW}からの不良連絡なし"])
        if r.random() < 0.5:
            rows.append([f"代替{'ロットの緊急払出' if internal else '品の緊急出荷'}", tdate(cd + timedelta(days=r.randint(2, 6))), sn(author),
                         "後工程の着工計画に影響なし" if internal else "顧客の生産計画に影響なし（顧客確認済み）"])
    else:
        frags = _containment_fragments(inc.first_response)
        if not frags:
            frags = [(r.choice([f"{eq.equipment_id}を着工停止（DOWN登録）", f"{eq.equipment_id}をインヒビット設定し着工停止"]),
                      ["不良品の追加発生なし", "誤着工なし"])]
        for x, effects in frags[:2]:
            rows.append([x, tdate(d_occ), sn(inc.reporter), r.choice(effects)])
        if inc.lots_affected:
            rows.append([f"対象ロット（{C['lots']}）をHOLDし、{_LOT_CHECK.get(k, '外観')}を全数確認",
                         tdate(d_occ + timedelta(days=r.randint(0, 1))), sn(qa_people[0]),
                         f"廃棄 {inc.scrap_wafers}枚、残りは品証判定で流動可" if inc.scrap_wafers else "全数規格内、品証判定で流動可"])
    if sibs and sibs[0] != "他号機":
        s2 = "、".join(sibs[:r.randint(1, min(3, len(sibs)))])
        rows.append([f"{sib_w}（{s2}）の{sub}緊急点検", tdate(comp.date() + timedelta(days=r.randint(0, 3))), sn(r.choice(maint)),
                     r.choice(["異常なし", "異常なし（写真記録あり）", f"{s2.split('、')[0]}で軽微な劣化を確認し予防交換"])])
    rows.append([f"{sub}の点検を暫定で{r.choice(['1回/日', '1回/シフト', '週2回'])}に強化（恒久対策完了まで）", tdate(comp.date() + timedelta(days=r.randint(0, 2))),
                 sn(author if not claim else r.choice(maint)), r.choice(["期間中の再発なし", "兆候なし、継続中", "期間中の異常なし"])])
    rows.sort(key=lambda x: x[1])
    C["containment_rows"] = [[str(i + 1), Z(a), b, c, Z(d)] for i, (a, b, c, d) in enumerate(rows)]
    C["action"] = inc.action if not claim else f"（{D.fmt_date(comp, style=0)} {inc.incident_id}にて復旧済み）\n{inc.action}"
    C["parts"] = "、".join(f"{p.name}（{p.part_no}）×{q}" for p, q in parts) if parts else r.choice(["なし", "－", "部品交換なし"])
    between = eff["between_ids"]
    ver = []
    if claim:
        ver.append(f"暫定処置後、{CW}からの追加不良連絡なし。初回報告を{D.fmt_date(plan.claim_date + timedelta(days=r.randint(1, 3)), style=0)}に{CW}へ提出")
    else:
        ver.append(r.choice(["復旧後、{qc}", "暫定処置後の確認：{qc}"]).format(qc=inc.result.rstrip("。") if inc.result else "動作確認OK"))
    if between:
        more = "ほか" if len(between) > 3 else ""
        ver.append(f"ただし恒久対策完了までに同一不具合が{len(between)}件発生（{'、'.join(between[:3])}{more}）し、都度暫定処置で対応")
    else:
        ver.append(r.choice(["恒久対策完了まで同一不具合の再発なし", "暫定処置期間中の再発なし", "監視強化期間中、兆候なし"]))
    C["containment_verification"] = Z("。".join(ver) + "。")

    # --- D4 根本原因 ---
    C["investigation"] = inc.investigation
    C["cause"] = inc.cause
    C["why_why"] = list(inc.why_why)
    core = _core_cause(inc.cause)
    esc = []
    if claim:
        esc.append(f"{eq.equipment_id}の異常発生から{inc.detected_by}で検出されるまでの間に処理されたロット（{C['escaped_lot']}）が、既に後工程へ流動・{SHIP}されていた")
        esc.append(r.choice([f"{SHIP}検査は抜取（{r.choice([2, 3, 5])}枚/ロット）であり、局所的に発生した不良を検出できなかった",
                             f"インライン欠陥検査のサンプリングが{r.choice([3, 4, 5])}ロットに1ロットで、対象ロットは検査対象外だった",
                             "SPCは管理限界内で推移しており、異常として判定されなかった（管理限界の設定が広い）"]))
        esc.append(f"ロットHOLDの遡及範囲を発生時刻から{r.choice([4, 6, 8])}時間としたが、異常の兆候はそれ以前から出ていた")
    elif inc.lots_affected:
        esc.append(f"{inc.detected_by}で検出されるまでに{len(inc.lots_affected)}ロットが処理されていた")
        esc.append(r.choice([f"処理中の{metric}をリアルタイムで監視しておらず、処理後の検査まで異常に気付けなかった",
                             f"FDCの{metric}の閾値が管理限界より緩く、ワーニングが出なかった",
                             "QC測定の頻度が1回/日で、発生から検出までに時間差があった"]))
        esc.append("後工程への流出はなし（HOLDで封じ込め済み）")
    else:
        if k in D.UTILITY_KINDS:
            esc.append(r.choice(["使用先装置はインターロックで安全側に停止しており、製品の流出はなし",
                                 "停止中に処理されていた使用先装置のロットは品証判定で影響なしを確認し、流出はなし"]))
        elif k in D.TRANSPORT_KINDS:
            esc.append(r.choice(["搬送が停止しただけでウェーハ外観に異常はなく、製品の流出はなし",
                                 "停止したFOUPは手動で回収し、ウェーハの外観検査で異常なしを確認（流出なし）"]))
        else:
            esc.append(r.choice(["装置インターロックにより処理中に停止したため、不良品の後工程流出はなし",
                                 "装置停止で検知しており、製品の流出はなし",
                                 f"{inc.detected_by}で検知した時点で処理中ウェーハを退避しており、製品への影響・流出なし",
                                 "流出はなし（停止時の仕掛りは品証確認のうえ流動）"]))
        pm_item = metric if sub in metric else f"{sub}の{metric}"
        opts = [f"ただし{sub}の劣化兆候を事前に検知する監視（FDC・点検項目）がなく、突発停止に至った",
                f"停止前に{metric}の変化はあったが、監視項目に含まれていなかった",
                f"PM点検に{pm_item}の確認がなく、劣化の進行を把握できていなかった"]
        if eff["ms_before"] >= 2:
            # 正準データ上、発生前90日に同サブシステムのチョコ停が実際に出ていた場合はその件数を書く
            opts = [f"発生前90日間に{sub}関連のチョコ停が{eff['ms_before']}件出ていたが、都度リセットで復帰しており予兆として保全に共有されなかった",
                    f"{sub}関連のチョコ停（直近90日で{eff['ms_before']}件）を製造側で処置しており、保全への情報共有ルールがなかった"]
        esc.append(r.choice(opts))
    C["escape_cause"] = Z("。".join(esc) + "。")

    # 特性要因（6M）: 発生原因（◎）は原因区分に応じたカテゴリに置き、なぜなぜの最終段（仕組みの原因）は方法か人に置く
    cc = inc.cause_category
    age = int((occ.date() - eq.installed_date).days / 365.25) + 1
    primary = {"作業ミス": "man", "設定ミス": "man", "調整不良": "meth", "施工不良": "meth",
               "異物": "mat", "前工程起因": "mat"}.get(cc, "mach")
    why_last = inc.why_why[-1]
    sys_cat = "man" if any(w in why_last for w in ("作業者", "人員", "教育", "経験", "力量", "新人", "担当者", "要員")) else "meth"
    if why_last == core or primary == sys_cat:
        sys_mark = "○"
    else:
        sys_mark = "◎"
    part_s = _part_short(parts[0][0].name) if parts else None
    man = []
    if primary == "man":
        man.append(f"◎ {core}")
    else:
        # 否定した要因は、調査内容（前回PM記録・聞き取り・FE作業など）で実際に確認したことから選ぶ
        man.append(r.choice(_fish_man_candidates(inc, k, sub, part_s, cc, metric)))
    if sys_cat == "man" and why_last != core:
        man.append(f"{sys_mark} {why_last}")
    if inc.shift == "夜勤":
        man.append(r.choice(["○ 夜勤帯は巡回頻度が少なく、兆候の発見が遅れた", "○ 夜勤は保全が少人数で、初動の切り分けに時間を要した",
                             f"○ 夜間は{sub}の異音・異常表示に気付ける人員が現場にいなかった"]))
    elif inc.shift == "休日":
        man.append(r.choice(["○ 休日帯は巡回頻度が少なく、兆候の発見が遅れた", "○ 休日で保全の常駐がなく、呼出しから到着まで時間がかかった",
                             "○ 休日の一次対応を製造側だけで判断できず、停止が長引いた"]))
    elif r.random() < 0.5:
        man.append(r.choice(["× 引継ぎ漏れ → 引継ぎ簿に記載あり", "× シフト間の連絡不足 → 前直からの申し送りに異常の記載なし",
                             f"× 他作業との干渉 → 発生前後に{eq.equipment_id}周辺での作業なし"]))
    mach = []
    if primary == "mach":
        mach.append(f"◎ {core}")
    else:
        mach.append(r.choice([f"× {sub}本体の故障 → 分解点検の結果、部品に異常なし", f"× {sub}の制御異常 → ログ確認の結果、指令値どおり動作"]))
    if age >= 8:
        mach.append(f"○ 設置{age}年目で{sub}の経年劣化が進行")
    if primary in ("man", "meth"):
        mach.append("○ 誤作業を検知・防止するインターロックがない")
    else:
        mach.append(r.choice([f"○ {eq.model}は{sub}の点検性が悪く、目視確認が難しい構造",
                              f"× {_sib_word(eq, sibs)}（{sibs[0]}）は同条件で異常なし" if sibs[0] != "他号機" else "× 装置改造履歴 → 該当なし"]))
    mat_c = _MATERIAL.get(k, ["交換部品の品質ロット"])
    mat = [f"◎ {core}"] if primary == "mat" else []
    mat.append(f"× {r.choice(mat_c)} → 受入記録を確認、規格内")
    if parts:
        mark = "○" if cc in ("異物", "前工程起因") else "×"
        mat.append(f"{mark} 交換部品（{parts[0][0].name}）のロット不良 → " + ("同ロット品に同傾向あり、メーカーへ調査依頼" if mark == "○" else "同ロット在庫品を確認し異常なし"))
    meth = [f"◎ {core}"] if primary == "meth" else []
    if sys_cat == "meth" and why_last != core:
        meth.append(f"{sys_mark} {why_last}")
    # 管理・手順の穴は原因区分ごとの候補から、部品名・使用量の単位を入れて選ぶ（なぜなぜの最終段と同じ趣旨のものは避ける）
    gap_fill = dict(part=part_s or f"{sub}の消耗部品", sub=sub, usage=_USAGE.get(k, "運転時間"), metric=metric)
    gaps = [g.format(**gap_fill) for g in _METHOD_GAPS.get(cc, _METHOD_GAPS["その他"])]
    gaps = [g for g in gaps if len(_tok(g) & _tok(why_last)) < 2] or gaps
    meth.append(f"○ {r.choice(gaps)}")
    # 測定: 調査内容に兆候・計測の記述があればそれを要因に書く
    munit = _METRICS.get(sub, ("", "件/月"))[1]
    m_fdc = re.search(r"FDCトレンドを遡ると(\d+)日前から兆候", inc.investigation)
    if m_fdc:
        meas = [f"○ {metric}の変化が{m_fdc.group(1)}日前からFDCに出ていたが、ワーニング設定がなかった"]
    elif "目視点検では外観上の異常は見当たらず" in inc.investigation:
        meas = [f"○ 外観では異常が分からず、{metric}を計測しないと劣化を判別できない"]
    elif inc.detected_by == "FDC監視":
        meas = [f"○ FDCの{metric}の閾値が管理限界より緩く、ワーニングが遅れた"]
    else:
        meas = [r.choice([f"○ FDCで{metric}を監視していなかった", f"○ {metric}の管理限界が広く、兆候を捉えられなかった",
                          f"○ {metric}の測定がPM時のみで、測定間隔が長い"])]
    cal = occ.date() - timedelta(days=r.randint(40, 320))
    meas_neg = ["× 計測器の校正 → 期限内、問題なし", f"× {metric}の測定器の校正 → 校正期限内（{cal.year % 100}年{cal.month}月校正）",
                "× 測定者による測定値のばらつき → 2名で再測定し同値"]
    if munit not in ("件/月", "回/日", "回/月", "ppm", "%", "h", "件/3ヶ月"):
        meas_neg.append(f"× センサの指示誤差 → 基準器との比較で誤差{r.choice(['0.5', '1', '2'])}%以内")
    meas.append(r.choice(meas_neg))
    env = []
    if occ.month in (6, 7, 8, 9) and sub in ("チラー", "冷却塔/熱交換器", "温調チャンバー", "封水/冷却", "RF電源", "ドライポンプ"):
        env.append(f"○ 夏季の冷却水温度上昇（{r.uniform(24.5, 29.0):.1f}℃）")
    if k in D.UTILITY_KINDS:
        t_lo, t_hi = (24.0, 34.0) if occ.month in (6, 7, 8, 9) else ((1.0, 11.0) if occ.month in (12, 1, 2) else (9.0, 22.0))
        env.append(f"× 外気温・湿度 → 発生時{r.uniform(t_lo, t_hi):.1f}℃／{r.randint(35, 80)}%で設計条件内")
    else:
        env.append(f"× クリーンルーム温湿度 → {r.uniform(22.6, 23.4):.1f}℃／{r.randint(40, 50)}%で管理範囲内")
    if cc in ("汚れ", "異物"):
        env.append(f"○ {sub}周辺に{'スラリー・研磨くず' if k == 'CMP' else '反応生成物・パーティクル'}が堆積しやすい環境")
    elif r.random() < 0.4:
        env.append(r.choice(["× 近隣工事・振動の影響 → 工事記録なし", "× 隣接装置の搬入・改造工事による振動 → 期間中の工事・搬入なし",
                             "× 瞬低・電源品質 → 受電設備の記録に瞬低なし", "× 静電気の影響 → イオナイザの動作・除電状態に異常なし"]))
    C["fishbone"] = [Z("\n".join(x)) for x in (man, mach, mat, meth, meas, env)]

    # --- D5 恒久対策の選定 ---
    fill = dict(part=main_part, a=r.choice([6, 12, 18, 24]), b=r.choice([3, 4, 6]), p=r.choice([70, 80]), metric=metric, sub=sub,
                ver=f"{r.randint(2, 6)}.{r.randint(0, 9)}.{r.randint(0, 20)}", h=r.choice([4, 6, 8]), h2=r.choice([24, 48]))
    if fill["b"] >= fill["a"]:
        fill["a"], fill["b"] = fill["b"] * 2, fill["b"]
    adopted = [x for x in _split_items(inc.prevention) if not x.startswith(("特になし", "現状の点検"))][:2]
    bank = [x.format(**fill) for x in _ADOPT_BY_CC.get(cc, _ADOPT_BY_CC["不明"])]
    r.shuffle(bank)
    for x in bank:
        if len(adopted) >= r.choice([2, 3]):
            break
        if x not in adopted:
            adopted.append(x)
    if claim:
        adopted.append(r.choice(_CLAIM_ESCAPE_ACTIONS).format(metric=metric, check=_LOT_CHECK.get(k, "外観"), a=r.choice([10, 20, 25]), b=r.choice([50, 100]), h=fill["h"], h2=fill["h2"])
                       .replace("出荷", SHIP))
    rejected = [(f"{sub}ユニット一式の更新", "◎", f"大（{r.randint(300, 2400)}万円）", f"{r.randint(3, 8)}ヶ月", "△", "不採用（費用対効果）"),
                (f"{sub}の点検頻度を毎日に増やすのみ", "△", "小（工数増）", "即日", "△", "不採用（定着性に課題）")]
    if claim:
        rejected.append((f"{SHIP}検査を恒久的に全数検査とする", "○", "大（検査工数 +{0}h/日）".format(r.randint(3, 8)), "1週", "×", "不採用（工程内で対策）"))
    cand = []
    for x in adopted:
        cost = r.choice(["小（{0}万円）".format(r.randint(1, 30)), "小", "中（{0}万円）".format(r.randint(40, 150)), "なし（工数のみ）"])
        cand.append([Z(x), r.choice(["◎", "○"]), cost, r.choice(["1週", "2週", "1ヶ月", "即日"]), r.choice(["◎", "○"]), "採用"])
    rj = r.sample(rejected, k=r.choice([1, 2]) if not claim else 2)
    for row in rj:
        cand.append([row[0], row[1], row[2], row[3], row[4], row[5]])
    order = list(range(len(cand)))
    r.shuffle(order)
    cand = [cand[i] for i in order]
    C["candidates"] = [[str(i + 1)] + row for i, row in enumerate(cand)]
    # 選定理由: 共通の考え方に加え、採用案の中身（監視・部品・使用量管理など）に合う理由だけを候補にする
    watch_words = ("FDC", "監視", "トレンド", "SPC", "ワーニング")
    has_watch = any(w in x for x in adopted for w in watch_words)
    has_prevent = any(not any(w in x for w in watch_words) for x in adopted)
    reasons = [
        "発生原因（{core}）に直接効き、費用・工期とも小さい案を採用。ユニット更新は効果は高いが費用対効果が低く、次期更新計画で検討する。",
        "真因への効果と定着性を重視して選定。点検頻度の増加のみでは工数が増えるだけで歯止めにならないため不採用とした。",
        "係内レビュー（保全・生技・品証）で{n_cand}案を比較し、効果が数値で確認できる案を採用。点検頻度の増加だけの案は歯止めにならないため除外。",
    ]
    if has_watch:
        reasons.append("効果・コスト・実現性の3観点で評価し、◎○の案を採用。{sub}の監視強化は生産技術課と合意済み。")
    if has_watch and has_prevent:
        reasons.append("{sub}の劣化・異常を「起こさない」対策と「早く見つける」対策を1つずつ以上含めることを基準に選定した。")
    if parts and sibs[0] != "他号機":
        reasons.append("予算内（保全費）で即時実施でき、同型機にも展開可能な案を優先。{part}の仕様変更はメーカー回答待ちのため今回は見送り。")
        reasons.append("{part}の在庫・納期を確認し、1ヶ月以内に同型機も含めて実施できる案を選定した。")
    if inc.related_incident_id:
        reasons.append("過去の類似トラブル（{past}）の対策履歴を確認し、効果が実証されている管理方法を採用した。")
    if cc in ("劣化", "摩耗", "汚れ", "断線") and k in _USAGE:
        # 使用量に比例して進む不具合だけ「使用量で管理」の理由を候補にする
        reasons.append("{sub}の不具合は{usage}に比例して進むため、暦日ではなく使用量で管理する案を優先した。ユニット更新は次年度予算で再検討。")
    reason = r.choice(reasons).format(core=core, sub=sub, part=_part_short(main_part), usage=_USAGE.get(k, "運転時間"),
                                      n_cand=len(cand), past=inc.related_incident_id)
    if claim:
        reason += f"流出対策は{SHIP}検査の全数化ではなく、工程内での検出力向上で対応する。"
    C["selection_reason"] = Z(reason)

    # --- D6 実施と効果確認 ---
    impl = plan.impl_last
    n_ad = len(adopted)
    pa = []
    floor = (plan.claim_date + timedelta(days=3)) if claim else comp.date() + timedelta(days=1)
    turn = {"qa": 0, "pe": 0, "mt": 0}
    # 完了日は最後の項目が恒久対策の最終実施日になるよう、項目順に昇順で並べる
    done_dates = sorted(max(floor, impl - timedelta(days=(n_ad - 1 - j) * r.randint(2, 12))) for j in range(n_ad))
    done_dates[-1] = max(done_dates[-1], impl)
    for j, x in enumerate(adopted):
        dd = done_dates[j]
        st_txt = "完了" if complete or j < n_ad - 1 else r.choice(["完了", "完了（効果確認中）"])
        # 実施内容に応じて担当部署を決める（出荷・検査→品証、監視・データ→生産技術、その他→保全）
        if any(w in x for w in ("出荷", "検査", "サンプリング", "顧客", "QC工程", "HOLD")):
            grp, people_ = "qa", qa_people
        elif any(w in x for w in ("FDC", "SPC", "監視", "トレンド", "ログ", "データ")):
            grp, people_ = "pe", pe_people
        else:
            grp, people_ = "mt", maint
        who_ = people_[turn[grp] % len(people_)]
        turn[grp] += 1
        pa.append([str(j + 1), Z(x), sn(who_), dd, st_txt])
    C["permanent_rows"] = pa
    er = []
    nb, na = len(eff["before_ids"]), len(eff["after_ids"])
    ok_mark, ng_mark = r.choice([("○", "△"), ("OK", "NG"), ("良", "要監視")])
    # 同一不具合件数の目標: 0件（多くの様式）か、対策前が多い場合は 1/3以下
    tgt_third = nb >= 3 and r.random() < 0.35
    cnt_ok = (na <= nb / 3) if tgt_third else (na == 0)
    er.append([f"同一不具合の発生件数（{sub}）", "件/3ヶ月", f"{nb}", f"{na}" if eff["after_days"] >= 30 else "－",
               "対策前の1/3以下" if tgt_third else r.choice(["0", "0件"]),
               (ok_mark if cnt_ok else ng_mark) if eff["after_days"] >= 30 else "確認中"])
    force_ng = ng and r.random() < 0.5
    mname, munit, mb, ma, mt, mok = _metric_row(sub, r, force_ng=force_ng)
    er.append([mname, munit, mb, ma if complete else "－", mt, (ok_mark if mok else ng_mark) if complete else "確認中"])
    if eff["ms_before"] + eff["ms_after"] > 0:
        mb_ = eff["ms_before"] / 3
        ma_ = eff["ms_after"] / max(1, eff["after_days"] / 30)
        er.append([f"{sub}関連チョコ停", "件/月", f"{mb_:.1f}", f"{ma_:.1f}" if complete else "－", "対策前の50%以下",
                   (ok_mark if ma_ <= mb_ * 0.5 or ma_ == 0 else ng_mark) if complete else "確認中"])
    if r.random() < 0.55 and eff["eq_before"] > 0:
        mtbf_b = 2 * EFFECT_DAYS * 24 / eff["eq_before"]
        mtbf_a = eff["eq_after_days"] * 24 / eff["eq_after"] if eff["eq_after"] else None
        er.append([f"MTBF（{eq.equipment_id} 全故障）", "h", f"{mtbf_b:,.0f}",
                   ("－" if not complete else (f"{mtbf_a:,.0f}" if mtbf_a else f"{eff['eq_after_days'] * 24:,}以上")), "対策前以上",
                   (ok_mark if (mtbf_a is None or mtbf_a >= mtbf_b) else ng_mark) if complete else "確認中"])
    C["effect_rows"] = er
    C["effect_headers_before"], C["effect_headers_after"] = r.choice(
        [("対策前", "対策後"), ("対策前（実績）", "対策後（実績）"), ("Before", "After"), ("対策前", "対策後（実績）")])
    C["verification_period"] = (f"対策前 {D.fmt_date(eff['b0'], style=0)}〜{D.fmt_date(eff['b1'], style=0)}／"
                                f"対策後 {D.fmt_date(eff['a0'], style=0)}〜{D.fmt_date(eff['a1'], style=0)}")
    ids_txt = "、".join(eff["after_ids"][:3]) + ("ほか" if na > 3 else "")
    if not complete:
        res = f"効果確認中（{D.fmt_date(eff['a0'] + timedelta(days=EFFECT_DAYS), style=0)}まで監視継続）。現時点で同一不具合" + (f"{na}件" if na else "の再発なし") + "。"
        status = "効果確認中"
    elif ng:
        head = (f"{mname}は目標値内となったが" if mok else f"{mname}は{mb}→{ma}{munit}と改善傾向だが目標未達で")
        res = (f"{head}、対策後も同一不具合が{na}件発生（{ids_txt}）。"
               "効果は限定的と判断し、追加対策（D5再検討）を実施する。")
        status = "条件付き完了"
    elif level == "partial":
        res = (f"同一不具合は対策前{nb}件→対策後{na}件に減少（{ids_txt}）。{mname}は目標値内となったが、"
               f"再発がゼロではないため{r.choice(['3ヶ月間の監視を継続する', '次回PMで追加点検を行う', 'メーカーと追加対策を協議中'])}。")
        status = r.choice(["完了（監視継続）", "完了"])
    elif na > 0:
        # 大幅に減ったがゼロではない（good 判定の範囲）
        res = r.choice([
            "対策前{nb}件→対策後{na}件（{ids}）と大幅に減少。{m}も目標値内で推移しており効果ありと判断するが、発生が残るため監視を継続する。",
            "{m}が{mb}→{ma}{u}に改善し、同一不具合も{nb}件→{na}件に減少。残った{na}件（{ids}）は個別に要因を確認し、本対策の効果はありと判定。",
        ]).format(nb=nb, na=na, ids=ids_txt, m=mname, mb=mb, ma=ma, u=munit)
        status = "完了" if cnt_ok else r.choice(["完了（監視継続）", "完了"])
    else:
        res = r.choice([
            "対策後{days}日間、同一不具合の再発なし。{m}も目標値内で安定しており、効果ありと判断。",
            "対策前{nb}件→対策後0件。{m}は目標を満足し、恒久対策の効果を確認した。",
            "{m}が{mb}→{ma}{u}に改善。対策後の再発0件で、効果ありと判定。",
        ]).format(days=eff["after_days"], m=mname, nb=nb, mb=mb, ma=ma, u=munit)
        status = "完了"
    ng_rows = []
    if complete and not ng:
        ng_rows = [row[0] for row in er[1:] if row[5] == ng_mark]
        if ng_rows:
            res += f"ただし{'・'.join(ng_rows[:2])}は目標未達のため、継続して監視する。"
    C["result"] = Z(res)
    C["status"] = status
    C["photos"] = []
    if inc.photos:
        after_cap = "対策後：" + (adopted[0][:24] + ("…" if len(adopted[0]) > 24 else ""))
        C["photos"] = [f"対策前：{inc.photos[0]}", after_cap]

    # --- D7 再発防止（標準化） ---
    dtypes = []
    joined = " ".join(adopted)
    if any(w in joined for w in ("周期", "交換基準", "定期交換")):
        dtypes.append("parts")
    if any(w in joined for w in ("FDC", "監視", "トレンド", "SPC")):
        dtypes.append("fdc")
    if any(w in joined for w in ("手順", "作業", "標準")):
        dtypes.append("wi")
    if any(w in joined for w in ("点検", "チェック")):
        dtypes.append(r.choice(["pm", "check"]))
    if claim:
        dtypes.append("qcp" if internal else r.choice(["ship", "qcp"]))
    if len(dtypes) < 2:
        dtypes += [x for x in ("pm", "fmea", "edu") if x not in dtypes][: 2 - len(dtypes)]
    dtypes = list(dict.fromkeys(dtypes))[:4]
    docs_rows = []
    rev_style = r.choice(["Rev", "版"])
    for j, dt_ in enumerate(dtypes):
        title, no, cur = docs.use(dt_, eq, sub)
        content = {"parts": f"{main_part}の交換周期・基準を改訂", "fdc": f"{metric}を監視項目に追加、管理値設定",
                   "wi": f"{sub}の点検・交換手順に注意点を追記（写真付き）", "pm": f"{sub}の点検項目を追加",
                   "check": f"{sub}の日常点検項目を追加", "fmea": f"故障モード「{_first_clause(core, 30)}」を追加しRPN再評価",
                   "edu": "本事例を教育資料に追加", "ship": f"{C.get('defect', '不良')}の確認項目を追加", "qcp": f"{eq.process}工程の管理項目・頻度を見直し"}[dt_]
        rev_txt = f"Rev.{cur}→Rev.{cur + 1}" if rev_style == "Rev" else f"第{cur}版→第{cur + 1}版"
        docs_rows.append([title, no, Z(content), rev_txt, impl + timedelta(days=r.randint(0, 10))])
    C["docs_rows"] = docs_rows
    prev = inc.prevention if inc.prevention and not inc.prevention.startswith(("特になし", "現状の点検")) else "／".join(adopted[:2])
    C["prevention"] = prev + r.choice(["", "。上記を標準類に反映し歯止めとした", "。改訂内容は保全課内で周知済み", "。周知記録は保全課共有フォルダに保管"])
    # 水平展開: トラブル報告時点の記載を起点に、8D完了時点の状況（実施済み・判断結果）を書き足す
    hz = inc.horizontal_deployment
    has_sib = sibs[0] != "他号機"
    sib_txt = "、".join(sibs[:3])
    hz_done = D.fmt_date(min(impl + timedelta(days=r.randint(5, 25)), TODAY - timedelta(days=3)), style=r.choice([0, 7]))
    if hz and "展開不要" in hz and any(w in joined for w in ("水平展開", "全車両", "同型機", "全号機")):
        hz = ""
    if not hz:
        hz = (f"{sib_w}（{sib_txt}）に{D.fmt_date(impl + timedelta(days=30), style=7)}までに展開予定" if has_sib
              else "同型機なし。類似構造の設備で点検実施")
    elif complete and "検討中" in hz:
        hz = (f"{hz} → 検討の結果、{sib_w}（{sib_txt}）へ展開。{hz_done} 完了" if has_sib and r.random() < 0.7
              else f"{hz} → 検討の結果、構造が異なるため展開不要と判断")
    elif complete and "予定" in hz:
        # 元の記載に期限（○日まで・次回PM）があるので、日付は書かずに結果だけ追記する
        hz = f"{hz} → 実施済み（{r.choice(['異常なし', '1台で軽微な摩耗あり、予防交換', '異常なし、記録は保全課共有フォルダ'])}）"
    elif complete and "定例会で共有" in hz and has_sib:
        hz = f"{hz}。{sib_w}（{sib_txt}）の点検は{hz_done}に完了"
    elif "他Fab" in hz and has_sib:
        other = [s for s in sibs if D.fab_of(D.equipment_by_id(s)) != D.fab_of(eq)]
        if other:
            hz = hz.replace("他Fabの同型機", f"他Fabの{sib_w}（{'、'.join(other[:3])}）")
    C["horizontal_deployment"] = hz

    # --- D8 称賛・完了承認 ---
    # 夜間・休日の復旧に当たったのは保全のメンバー
    mem_names = [sn(p) for p in maint if p is not team["chief"]] or [sn(m[1]) for m in team["members"] if m[0].startswith("メンバー") and m[1].role != "FE"]
    rec = []
    if inc.shift == "夜勤":
        rec.append(f"夜間の突発対応にもかかわらず、{('・'.join(mem_names[:2]) or sn(author))}が迅速に復旧し、生産への影響を最小限に抑えた。")
    elif inc.shift == "休日":
        rec.append(f"休日の呼出しに対し、{('・'.join(mem_names[:2]) or sn(author))}が速やかに出勤・対応した。")
    if inc.reporter.section.startswith("製造課"):
        rec.append(f"{inc.reporter.section.replace('製造課 ', '製造課')} {sn(inc.reporter)}さんの早期発見により、不良の拡大を防止できた。")
    fe = [m[1] for m in team["members"] if m[1].role == "FE"]
    if fe:
        rec.append(f"{D.MAKER_SHORT.get(fe[0].section, fe[0].section)} {sn(fe[0])}FEの協力により、原因部位を早期に特定できた。")
    pe = [m[1] for m in team["members"] if m[1].section == "生産技術課"]
    if pe and r.random() < 0.6:
        rec.append(f"生産技術課 {sn(pe[0])}さんのデータ解析が、監視項目の追加につながった。")
    if claim:
        if internal:
            rec.append(f"受付から{r.randint(1, 3)}日で組立テスト部へ初回報告を行い、後工程の滞留を最小限に抑えた。")
        else:
            rec.append(f"受付から{r.randint(1, 3)}日で顧客へ初回報告を行い、{C['customer_name'].replace('株式会社', '')}様から迅速な対応について評価をいただいた。")
    rec.append(r.choice(["本件の対策は保全改善事例として課内発表会で共有する。", "チーム各位の協力に感謝する。", "関係各位の協力に感謝します。",
                         f"本事例は{(impl + timedelta(days=40)).month}月の保全月例会で水平展開事例として紹介予定。"]))
    C["recognition"] = Z("".join(rec[:4]))
    # 承認コメントは効果確認の結果で書き分ける（effective / monitoring / partial / conditional / pending）
    if status == "効果確認中":
        cls = "pending"
    elif ng:
        cls = "conditional"
    elif level == "partial":
        cls = "partial"
    elif na > 0 or ng_rows or status != "完了":
        cls = "monitoring"
    else:
        cls = "effective"
    fe_maker = D.MAKER_SHORT.get(fe[0].section, fe[0].section) if fe else None
    comment_args = dict(habit=D._writer_habit(champion.employee_id), claim=claim, internal=internal,
                        cname=C.get("customer_name", "").replace("社内：", ""), metric=metric, sub=sub, nb=nb, na=na, ids_txt=ids_txt,
                        ng_rows=ng_rows,
                        hz=hz, sib_txt=sib_txt if has_sib else "", doc_title=docs_rows[0][0] if docs_rows else None, maker=fe_maker,
                        related=inc.related_incident_id if inc.recurrence else None, shift=inc.shift,
                        eff_end=(eff["a0"] + timedelta(days=EFFECT_DAYS)).date())
    if cls == "pending":
        C["approver"], C["approval_date"] = "", None
        C["approval_comment"] = _approval_comment(r, cls, ad=None, **comment_args)
    else:
        ad = min(TODAY, max(impl, eff["a1"].date()) + timedelta(days=r.randint(2, 12)))
        C["approver"] = champion.name
        C["approval_date"] = Dt(datetime.combine(ad, time()))
        C["approval_comment"] = _approval_comment(r, cls, ad=ad, **comment_args)
    if claim:
        C["customer_answer_date"] = Dt(datetime.combine(min(TODAY, (C["approval_date"].v.date() if C["approval_date"] else impl + timedelta(days=5)) + timedelta(days=r.randint(0, 3))), time()))
    # 発行日: 承認依頼の直前（効果確認が終わっていればその後）。効果確認中なら最新の更新日
    if C["approval_date"]:
        rd = C["approval_date"].v.date() - timedelta(days=r.randint(0, 2))
        rd = max(rd, eff["a1"].date() if complete else impl)
        rd = min(rd, C["approval_date"].v.date())
    else:
        rd = max(impl, TODAY - timedelta(days=r.randint(3, 12)))
    C["report_date"] = Dt(datetime.combine(min(TODAY, rd), time()))
    C["stamps"] = dict(approve=_stamp_name(champion) if C["approver"] else "", check=_stamp_name(leader), create=_stamp_name(author))
    return C


# ---------------------------------------------------------------------------
# レイアウト設定（ファイルごと）
# ---------------------------------------------------------------------------
@dataclass
class Style:
    layout_version: str
    form_no: str
    two_sheet: bool
    claim: bool
    grid: str            # "hougan"（方眼紙）/ "cols"（列レイアウト）
    section: str         # "bar"（見出し帯）/ "side"（左端に縦結合の見出し）
    dlabels: list
    w5: str              # "rows" / "grid" / "horizontal"
    w5_labels: list
    fish: str            # "2x3" / "3x2"
    fish_labels: list
    why_style: str       # "rows" / "arrow"
    date_mode: str       # "cell" / "text"
    date_fmt: int
    dt_fmt: int
    text_date_style: int
    tbl_date_style: int
    font: str
    fsize: float
    stamp: str           # "image" / "text"
    photos: bool
    list_sheet: bool
    history_sheet: bool
    labels: dict
    label_fill: str
    bar_fill: str
    head_fill: str
    base_h: float
    gap: int
    fixed_rows: bool     # 表に空行（様式の固定行数）を残す
    adopt_word: tuple
    title: str
    sheet_names: tuple
    page_view: bool


_FONTS = [("ＭＳ Ｐゴシック", 10), ("游ゴシック", 10), ("Meiryo UI", 9.5), ("ＭＳ ゴシック", 10), ("BIZ UDPゴシック", 10)]

# 様式の版ごとの出力先サブフォルダ。「版の呼び名（Windowsで使えない "." は外す）＋その版が使われ始めた年」。
# 年は改訂履歴シート（_history_sheet）の改訂日に合わせている。
# 1つのフォルダの中は同じ様式・同じシート構成のファイルだけなので、フォルダごと帳票取り込みに入れて一括で読み取れる。
VERSION_DIRS = {
    "QA-F-021_Rev.3_2sheet_hougan": "QA-F-021_Rev3_2sheet_hougan_2023改訂",
    "QA-F-021_Rev.3_1sheet_hougan": "QA-F-021_Rev3_1sheet_hougan_2023改訂",
    "QA-F-021_Rev.2_1sheet_cols": "QA-F-021_Rev2_1sheet_cols_2021改訂",
    "QA-F-022_Rev.1_2sheet_hougan": "QA-F-022_Rev1_2sheet_hougan_2020制定",
    "QA-F-022_Rev.2_1sheet_cols": "QA-F-022_Rev2_1sheet_cols_2024改訂",
}


def _make_style(idx: int, claim: bool, plan_rank: int) -> Style:
    r = D.rng(f"f3_8d:style:{idx}")
    if claim:
        two = plan_rank % 5 < 3          # クレーム10件: 2シート6 / 1シート4
        form = "QA-F-022 Rev.1" if two else "QA-F-022 Rev.2"
        grid = "hougan" if two else "cols"
        dstyle = r.choice(["jp_colon", "jp", "en_jp"]) if two else r.choice(["code", "jp_colon"])
        title = r.choice(["顧客クレーム 是正処置報告書（8D）", "8D Report（顧客クレーム）", "是正処置報告書（8D）［顧客クレーム］"])
    else:
        v = plan_rank % 20
        if v < 8:
            form, two, grid = "QA-F-021 Rev.3", True, "hougan"
        elif v < 13:
            form, two, grid = "QA-F-021 Rev.3", False, "hougan"
        else:
            form, two, grid = "QA-F-021 Rev.2", False, "cols"
        dstyle = r.choice(["jp", "en_jp", "jp"]) if form.endswith("3") else r.choice(["code", "code", "jp_colon"])
        title = r.choice(["是正処置報告書（8D）", "8D 是正処置報告書", "8Dレポート（是正処置報告書）", "是正処置報告書"])
    section = "bar" if grid == "hougan" else "side"
    labels = {k: r.choice(v) for k, v in LBL.items()}
    font, fsize = r.choice(_FONTS)
    lv = f"{form.replace(' ', '_')}_{'2sheet' if two else '1sheet'}_{grid}"
    if two:
        names = ("8D報告(1)", "8D報告(2)")
    else:
        names = (r.choice(["8D報告", "8D報告書", "是正処置報告", "8D"]),)
    return Style(
        layout_version=lv, form_no=form, two_sheet=two, claim=claim, grid=grid, section=section, dlabels=D_LABELS[dstyle],
        w5=r.choice(["rows", "grid", "horizontal"] if grid == "hougan" else ["rows", "rows", "grid"]),
        w5_labels=W5_LABELS[r.choice(list(W5_LABELS))], fish=r.choice(["2x3", "3x2"]),
        fish_labels=FISH_LABELS[r.choice(list(FISH_LABELS))], why_style=r.choice(["rows", "arrow"]),
        date_mode=r.choice(["cell", "cell", "text"]), date_fmt=r.choice([0, 0, 1, 2, 3]), dt_fmt=r.choice([0, 1, 2]),
        text_date_style=r.choice([0, 3, 4, 2, 8]), tbl_date_style=r.choice([0, 4, 7, 5]), font=font, fsize=fsize,
        stamp=r.choice(["image", "image", "text"]), photos=r.random() < 0.45, list_sheet=form.endswith("Rev.3") and r.random() < 0.7,
        history_sheet=r.random() < 0.35, labels=labels,
        label_fill=r.choice(["DDEBF7", "E2EFDA", "F2F2F2", "FFF2CC"]), bar_fill=r.choice(["1F4E78", "305496", "595959", "375623"]),
        head_fill=r.choice(["D9E1F2", "EDEDED", "DDEBF7"]), base_h=16.5 if grid == "hougan" else 18.0, gap=r.choice([0, 1]),
        fixed_rows=r.random() < 0.6, adopt_word=r.choice([("採用", "不採用"), ("○採用", "×不採用"), ("採", "否")]), title=title,
        sheet_names=names, page_view=r.random() < 0.3,
    )


# ---------------------------------------------------------------------------
# 描画（ワークシート上の配置）
# ---------------------------------------------------------------------------
THIN = Side(style="thin", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


class Recorder:
    """_expected.jsonl 用に、キーごとの表示値と使われたラベル文字列を記録する。"""

    def __init__(self):
        self.values: dict = {}
        self.labels: dict = {}

    def put(self, key: str, label: str | None, display):
        if key in self.values:
            return
        if display in ("", None, [], {}):
            return
        self.values[key] = display
        if label:
            self.labels[key] = label


def _alloc(c1: int, c2: int, weights: list[float]) -> list[tuple[int, int]]:
    """列 c1..c2 を重みに比例して整数列に割り付ける（各1列以上）。"""
    n = c2 - c1 + 1
    tot = sum(weights)
    raw = [w / tot * n for w in weights]
    cnt = [max(1, int(x)) for x in raw]
    while sum(cnt) < n:
        j = max(range(len(raw)), key=lambda i: raw[i] - cnt[i])
        cnt[j] += 1
    while sum(cnt) > n:
        j = max(range(len(raw)), key=lambda i: (cnt[i] - raw[i]) if cnt[i] > 1 else -99)
        cnt[j] -= 1
    out, c = [], c1
    for k in cnt:
        out.append((c, c + k - 1))
        c += k
    return out


_STAMP_CACHE: dict = {}


def _font_path() -> str | None:
    for p in ("C:/Windows/Fonts/msgothic.ttc", "C:/Windows/Fonts/meiryo.ttc", "C:/Windows/Fonts/YuGothM.ttc"):
        if Path(p).exists():
            return p
    return None


def _pil_font(size: int):
    p = _font_path()
    try:
        return ImageFont.truetype(p, size) if p else ImageFont.load_default()
    except OSError:
        return ImageFont.load_default()


def _stamp_png(name: str) -> bytes:
    """朱色の丸印（認印風）。"""
    if name in _STAMP_CACHE:
        return _STAMP_CACHE[name]
    sz = 96
    im = PILImage.new("RGBA", (sz, sz), (255, 255, 255, 0))
    dr = ImageDraw.Draw(im)
    red = (214, 40, 40, 235)
    dr.ellipse([4, 4, sz - 5, sz - 5], outline=red, width=5)
    chars = list(name)[:3]
    fs = 34 if len(chars) <= 2 else 25
    f = _pil_font(fs)
    total = fs * len(chars)
    y = (sz - total) / 2 - 2
    for ch in chars:
        w = dr.textlength(ch, font=f)
        dr.text(((sz - w) / 2, y), ch, fill=red, font=f)
        y += fs
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    _STAMP_CACHE[name] = buf.getvalue()
    return _STAMP_CACHE[name]


def _photo_png(caption: str, when: date, seed: str, after: bool) -> bytes:
    """現場写真の代わりの簡易イラスト（部品の外観と注記・撮影日付入り）。"""
    r = D.rng(f"f3_8d:photo:{seed}")
    w, h = 320, 220
    base = (r.randint(150, 190), r.randint(150, 190), r.randint(150, 190))
    im = PILImage.new("RGB", (w, h), base)
    dr = ImageDraw.Draw(im)
    for i in range(0, h, 4):   # 奥行き感のあるグラデーション
        c = tuple(max(0, v - i // 6) for v in base)
        dr.line([(0, i), (w, i)], fill=c)
    x0, y0 = r.randint(40, 90), r.randint(40, 70)
    dr.rounded_rectangle([x0, y0, x0 + r.randint(140, 190), y0 + r.randint(80, 110)], radius=10, fill=(205, 205, 210), outline=(90, 90, 90), width=3)
    for _ in range(r.randint(2, 5)):
        cx, cy = r.randint(x0 + 15, x0 + 130), r.randint(y0 + 15, y0 + 70)
        dr.ellipse([cx - 9, cy - 9, cx + 9, cy + 9], fill=(120, 120, 125))
    if after:
        dr.line([(w - 70, 40), (w - 55, 58), (w - 25, 20)], fill=(30, 160, 60), width=7)
    else:
        ex, ey = r.randint(x0 + 30, x0 + 120), r.randint(y0 + 20, y0 + 70)
        dr.ellipse([ex - 32, ey - 22, ex + 32, ey + 22], outline=(230, 30, 30), width=4)
        dr.line([(ex + 30, ey - 18), (ex + 70, ey - 50)], fill=(230, 30, 30), width=3)
    f = _pil_font(14)
    dr.rectangle([0, h - 26, w, h], fill=(0, 0, 0))
    dr.text((6, h - 22), caption[:22], fill=(255, 255, 255), font=f)
    dr.text((w - 92, 6), f"{when.year}/{when.month:02d}/{when.day:02d}", fill=(255, 140, 0), font=_pil_font(13))
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    return buf.getvalue()


class Canvas:
    """1ワークシート分の描画ヘルパー。行の高さは文字量から見積もって最後にまとめて設定する。"""

    def __init__(self, ws, st: Style, rec: Recorder):
        self.ws, self.st, self.rec = ws, st, rec
        if st.grid == "hougan":
            self.widths = [2.63] * 36
        else:
            # 旧版の列レイアウト: 左端がセクション見出し列、残り20列はやや広めの等幅列
            self.widths = [5.2] + [4.45] * 20
        for i, w in enumerate(self.widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        self.n = len(self.widths)
        self.cx0 = 2 if st.section == "side" else 1
        self.row = 1
        self.heights: dict = {}
        self.section_starts: list[int] = []

    # --- 基本 ---
    def units(self, c1: int, c2: int) -> float:
        return sum(self.widths[c1 - 1:c2])

    def fit(self, r1: int, r2: int, width: float, text: str, size: float, pad: float = 5.0):
        lines = _text_lines(text, width, size)
        need = lines * size * 1.45 + pad
        rows = list(range(r1, r2 + 1))
        cur = sum(self.heights.get(x, self.st.base_h) for x in rows)
        if need > cur:
            add = (need - cur) / len(rows)
            for x in rows:
                self.heights[x] = self.heights.get(x, self.st.base_h) + add

    def set_h(self, r: int, h: float):
        self.heights[r] = max(self.heights.get(r, 0), h)

    def put(self, r1, c1, r2, c2, value, *, role="value", fmt=None, h=None, v=None, size=None, bold=False,
            rotate=None, border=True, wrap=True, color=None):
        st = self.st
        cell = self.ws.cell(r1, c1)
        cell.value = value
        if (role in ("label", "head") and size is None and r1 == r2 and isinstance(value, str)
                and _text_lines(value, self.units(c1, c2), st.fsize) >= 2):
            size = max(7.5, st.fsize - 1.5)     # 狭い欄の長いラベルは文字を小さくして収める（よくある手直し）
        fill = {"label": st.label_fill, "head": st.head_fill, "bar": st.bar_fill}.get(role)
        fc = "FFFFFF" if role == "bar" else (color or "000000")
        cell.font = Font(name=st.font, size=size or st.fsize, bold=bold or role in ("label", "head", "bar", "title"), color=fc)
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)
        if h is None:
            h = "center" if role in ("label", "head") else "left"
        if v is None:
            v = "center" if role in ("label", "head", "bar") else "top"
        cell.alignment = Alignment(horizontal=h, vertical=v, wrap_text=wrap, text_rotation=rotate or 0)
        if fmt:
            cell.number_format = fmt
        if border:
            cell.border = BOX
        if (r1, c1) != (r2, c2):
            self.ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
        if isinstance(value, str) and wrap and not rotate:
            self.fit(r1, r2, self.units(c1, c2), value, size or st.fsize)
        return cell

    # --- 値の変換 ---
    def conv(self, val):
        """(セルに入れる値, 表示形式, 人が読む表示文字列) を返す。"""
        st = self.st
        if isinstance(val, Dt):
            if st.date_mode == "cell":
                if val.with_time:
                    f, disp = _DATETIME_FMTS[st.dt_fmt]
                else:
                    f, disp = _DATE_FMTS[st.date_fmt]
                return val.v, f, disp(val.v)
            if val.with_time:
                s = D.fmt_datetime_variants(val.v)[[0, 3, 4, 1, 6][st.text_date_style % 5]]
            else:
                s = D.fmt_date_variants(val.v)[st.text_date_style]
            return s, None, s
        if isinstance(val, Num):
            if st.date_mode == "cell":
                return val.v, '#,##0"分"', f"{val.v:,}分"
            s = f"{val.v:,}分（{val.v / 60:.1f}h）"
            return s, None, s
        if isinstance(val, date):
            s = D.fmt_date_variants(val)[st.tbl_date_style]
            return s, None, s
        if val is None:
            return "", None, ""
        return val, None, val

    # --- ブロック ---
    def kv_row(self, items: list, weights: list | None = None, height: float | None = None):
        """[(key, value)] を ラベル|値 の組で1行に並べる。"""
        if weights is None:
            weights = [2, 4] * len(items)
        spans = _alloc(self.cx0, self.n, weights)
        r = self.row
        for j, (key, val) in enumerate(items):
            label = self.st.labels.get(key, key)
            (a, b), (c, d) = spans[2 * j], spans[2 * j + 1]
            self.put(r, a, r, b, label, role="label")
            cv, f, disp = self.conv(val)
            self.put(r, c, r, d, cv, fmt=f, v="center")
            self.rec.put(key, label, disp)
        if height:
            self.set_h(r, height)
        self.row += 1

    def text_block(self, key: str, text, *, lw: float = 4, min_rows: int = 1, label: str | None = None, top: bool = False):
        label = label or self.st.labels.get(key, key)
        cv, f, disp = self.conv(text)
        r = self.row
        if top:
            self.put(r, self.cx0, r, self.n, label, role="label", h="left")
            r += 1
            self.put(r, self.cx0, r + min_rows - 1, self.n, cv, fmt=f)
            self.row = r + min_rows
        else:
            (a, b), (c, d) = _alloc(self.cx0, self.n, [lw, 24 - lw])
            self.put(r, a, r + min_rows - 1, b, label, role="label")
            self.put(r, c, r + min_rows - 1, d, cv, fmt=f)
            self.row = r + min_rows
        self.rec.put(key, label, disp)

    def caption(self, text: str):
        self.put(self.row, self.cx0, self.row, self.n, text, role="label", h="left")
        self.row += 1

    def table(self, key: str, label: str | None, headers: list[str], rows: list[list], weights: list[float],
              *, fixed: int = 0, center_cols: tuple = (), with_caption: bool = True):
        if label and with_caption:
            self.caption(label)
        spans = _alloc(self.cx0, self.n, weights)
        r = self.row
        for (a, b), hd in zip(spans, headers):
            self.put(r, a, r, b, hd, role="head")
        r += 1
        out = []
        for row in rows:
            rec_row = {}
            for j, ((a, b), val) in enumerate(zip(spans, row)):
                cv, f, disp = self.conv(val)
                self.put(r, a, r, b, cv, fmt=f, h="center" if j in center_cols else "left", v="center" if j in center_cols else "top")
                rec_row[headers[j]] = disp
            out.append(rec_row)
            r += 1
        for _ in range(max(0, fixed - len(rows))):
            for a, b in spans:
                self.put(r, a, r, b, None)
            r += 1
        self.row = r
        self.rec.put(key, label, out)

    def section(self, idx: int, body):
        st = self.st
        label = st.dlabels[idx]
        self.section_starts.append(self.row)
        if st.section == "bar":
            self.put(self.row, 1, self.row, self.n, label, role="bar", h="left", size=st.fsize + 1)
            self.set_h(self.row, 20)
            self.row += 1
            body(label)
        else:
            r0 = self.row
            body(label)
            r1 = self.row - 1
            vertical = len(label) > 3
            self.put(r0, 1, r1, 1, label, role="bar", h="center", v="center", rotate=255 if vertical else None, size=st.fsize + (0 if vertical else 1))
            if vertical:
                need = len(label) * (st.fsize + 2) + 10
                cur = sum(self.heights.get(x, st.base_h) for x in range(r0, r1 + 1))
                if need > cur:
                    for x in range(r0, r1 + 1):
                        self.heights[x] = self.heights.get(x, st.base_h) + (need - cur) / (r1 - r0 + 1)
        if st.gap:
            self.set_h(self.row, 8)
            self.row += 1
        return label

    def image(self, png: bytes, row: int, col: int, w: int, hgt: int, off_x: int = 2, off_y: int = 2):
        img = XLImage(io.BytesIO(png))
        img.width, img.height = w, hgt
        marker = AnchorMarker(col=col - 1, colOff=pixels_to_EMU(off_x), row=row - 1, rowOff=pixels_to_EMU(off_y))
        img.anchor = OneCellAnchor(_from=marker, ext=XDRPositiveSize2D(pixels_to_EMU(w), pixels_to_EMU(hgt)))
        self.ws.add_image(img)

    # --- 仕上げ（行高さ・印刷設定） ---
    def finish(self, footer_left: str):
        ws, st = self.ws, self.st
        last = self.row - 1
        for x in range(1, last + 1):
            ws.row_dimensions[x].height = round(self.heights.get(x, st.base_h), 1)
        ws.sheet_view.showGridLines = False
        if st.page_view:
            ws.sheet_view.view = "pageBreakPreview"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.orientation = "portrait"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_options.horizontalCentered = True
        ws.page_margins = PageMargins(left=0.47, right=0.39, top=0.59, bottom=0.59, header=0.3, footer=0.3)
        ws.print_area = f"A1:{get_column_letter(self.n)}{last}"
        ws.oddFooter.left.text = footer_left
        ws.oddFooter.left.size = 8
        ws.oddFooter.center.text = "&P / &N"
        ws.oddFooter.right.text = "保存期間：10年"
        ws.oddFooter.right.size = 8
        # 改ページ: 横幅をA4縦の印字幅に合わせて縮小したときの1ページ分の高さ（シート上のpt）を超えるところで、
        # 直前のセクション先頭で区切る。印字幅 = 595pt - 左右余白、印字高さ = 842pt - 上下余白（ヘッダ/フッタ分を少し引く）
        width_pt = sum(int(w * 7 + 5) for w in self.widths) * 0.75
        scale = min(1.0, (595 - (0.47 + 0.39) * 72) / width_pt)
        page_h = (842 - (0.59 + 0.59) * 72 - 12) / scale
        acc, page_start = 0.0, 1
        tops = set(self.section_starts)
        cum = {}
        for x in range(1, last + 1):
            cum[x] = acc
            acc += self.heights.get(x, st.base_h)
        breaks = []
        base = 0.0
        prev_top = None
        for x in sorted(tops):
            if cum[x] - base > page_h and prev_top and prev_top > page_start:
                breaks.append(prev_top - 1)
                base = cum[prev_top]
                page_start = prev_top
            prev_top = x
        if acc - base > page_h and prev_top and prev_top > page_start:
            breaks.append(prev_top - 1)
        if breaks:
            ws.row_breaks = RowBreak()
            for b in breaks:
                ws.row_breaks.append(Break(id=b))


# ---------------------------------------------------------------------------
# 各セクションの描画
# ---------------------------------------------------------------------------
def _header(cv: Canvas, C: dict, page: int):
    st = cv.st
    n = cv.n
    r = cv.row
    # 様式番号（右上の小さな文字）
    cv.put(r, 1, r, n, f"様式 {st.form_no}", size=8, h="right", border=False, v="center")
    cv.set_h(r, 13)
    r += 1
    title_span, *boxes = _alloc(1, n, [64, 12, 12, 12])
    ttl = st.title if page == 1 else f"{st.title}　（{page}/2）"
    cv.put(r, title_span[0], r + 2, title_span[1], ttl, role="title", size=16 if st.grid == "hougan" else 15, h="center", v="center", border=False)
    if page == 1:
        roles = ["承認", "確認", "作成"] if "022" not in st.form_no else ["承認", "審査", "作成"]
        keys = ["approve", "check", "create"]
        for (a, b), role, k in zip(boxes, roles, keys):
            cv.put(r, a, r, b, role, role="label", size=st.fsize - 1)
            name = C["stamps"][k]
            if st.stamp == "text" or not name:
                cv.put(r + 1, a, r + 2, b, name, h="center", v="center", size=st.fsize + 1)
                if name:
                    cv.rec.put(f"stamp_{k}", role, name)
            else:
                cv.put(r + 1, a, r + 2, b, None)
                wpx = int(cv.units(a, b) * 7 + 5)
                cv.image(_stamp_png(name), r + 1, a, 38, 38, off_x=max(2, (wpx - 38) // 2), off_y=3)
        cv.set_h(r + 1, 22)
        cv.set_h(r + 2, 22)
    cv.row = r + 3
    cv.set_h(cv.row, 6)
    cv.row += 1
    save_cx0 = cv.cx0
    cv.cx0 = 1       # ヘッダ部は全幅
    if page == 1:
        cv.kv_row([("report_id", C["report_id"]), ("report_date", C["report_date"]), ("revision", C["revision"])])
        if C["claim"]:
            cv.kv_row([("customer_name", C["customer_name"]), ("claim_no", C["claim_no"])], weights=[2, 10, 2, 4])
            cv.kv_row([("claim_date", C["claim_date"]), ("answer_due", C["answer_due"]), ("defect_qty", C["defect_qty"])], weights=[2, 3.4, 2, 3.4, 2, 7.2])
            cv.kv_row([("product_name", C["product_name"]), ("related_incident_id", C["related_incident_id"])], weights=[2, 10, 2, 4])
        else:
            cv.kv_row([("related_incident_id", C["related_incident_id"]), ("department", C["department"]), ("reporter", C["reporter"])])
        cv.kv_row([("equipment_id", C["equipment_id"]), ("equipment_name", C["equipment_name"]), ("line", C["line"])])
        cv.kv_row([("process", C["process"]), ("maker", C["maker"]), ("severity", C["severity"])])
        cv.kv_row([("occurred_at", C["occurred_at"]), ("recovered_at", C["recovered_at"]), ("downtime_min", C["downtime_min"])])
        if C["claim"]:
            cv.kv_row([("department", C["department"]), ("reporter", C["reporter"]), ("status", C["status"])])
        else:
            cv.kv_row([("status", C["status"]), ("lots", C["lots"] or "－")], weights=[2, 4, 2, 10])
    else:
        # 2枚目: 管理No・設備Noを再掲
        items = [("report_id", C["report_id"]), ("equipment_id", C["equipment_id"])]
        spans = _alloc(1, n, [2, 4, 2, 4, 6])
        rr = cv.row
        for j, (key, val) in enumerate(items):
            (a, b), (c, d) = spans[2 * j], spans[2 * j + 1]
            cv.put(rr, a, rr, b, st.labels[key], role="label")
            cv.put(rr, c, rr, d, val, v="center")
        cv.row += 1
    cv.cx0 = save_cx0
    cv.set_h(cv.row, 8)
    cv.row += 1


def _d1(cv: Canvas, C: dict, label: str):
    st = cv.st
    t = C["team"]
    champ_word = random_choice(C, "champ", ["チャンピオン", "推進責任者", "責任者"])
    rows = [[champ_word, t["champion"].name, t["champion"].section, t["champion"].role, "リソース確保・完了承認"],
            ["リーダー", t["leader"].name, t["leader"].section, t["leader"].role, "全体統括・対策立案"]]
    for role, p, duty in t["members"]:
        sect = f"{p.section}（FE）" if p.role == "FE" else p.section
        rows.append([role, p.name, sect, p.role, duty])
    headers = ["役割", "氏名", "所属", "役職", "担当内容"]
    # D1 は表の見出しを付けず、セクション見出しをそのままラベルとする
    cv.table("team_members", label, headers, rows, [2.2, 2.6, 4.6, 1.8, 5.2], fixed=8 if st.fixed_rows else 0, center_cols=(0, 3), with_caption=False)
    cv.rec.put("team_leader", "リーダー", t["leader"].name)
    cv.rec.put("team_champion", champ_word, t["champion"].name)


def random_choice(C: dict, key: str, opts: list):
    """ファイル内で一貫した語の揺れ（同じファイルでは同じ語を使う）。"""
    k = f"_choice_{key}"
    if k not in C:
        C[k] = D.rng(f"f3_8d:choice:{C['report_id']}:{key}").choice(opts)
    return C[k]


def _d2(cv: Canvas, C: dict, label: str):
    st = cv.st
    cv.text_block("symptom", C["symptom"], lw=3.5, min_rows=2)
    if C["claim"]:
        cv.kv_row([("alarm", C["alarm"]), ("lots", C["lots"])], weights=[3.5, 8.5, 3, 9])
    else:
        cv.kv_row([("alarm", C["alarm"])], weights=[3.5, 20.5])
    labs = st.w5_labels
    vals = C["w5"]
    if st.w5 == "rows":
        for key, lab, val in zip(W5_KEYS, labs, vals):
            (a, b), (c, d) = _alloc(cv.cx0, cv.n, [4.5, 19.5])
            cv.put(cv.row, a, cv.row, b, lab, role="label")
            cv.put(cv.row, c, cv.row, d, val)
            cv.rec.put(key, lab, val)
            cv.row += 1
    elif st.w5 == "grid":
        pairs = list(zip(W5_KEYS, labs, vals))
        for i in range(0, 6, 2):
            spans = _alloc(cv.cx0, cv.n, [3, 9, 3, 9])
            for j, (key, lab, val) in enumerate(pairs[i:i + 2]):
                (a, b), (c, d) = spans[2 * j], spans[2 * j + 1]
                cv.put(cv.row, a, cv.row, b, lab, role="label")
                cv.put(cv.row, c, cv.row, d, val)
                cv.rec.put(key, lab, val)
            cv.row += 1
        key, lab, val = pairs[6]
        (a, b), (c, d) = _alloc(cv.cx0, cv.n, [3, 21])
        cv.put(cv.row, a, cv.row, b, lab, role="label")
        cv.put(cv.row, c, cv.row, d, val)
        cv.rec.put(key, lab, val)
        cv.row += 1
    else:   # horizontal: ラベルを横一列、その下に値
        spans = _alloc(cv.cx0, cv.n, [5, 5, 5, 5, 5, 5, 6])
        for (a, b), lab in zip(spans, labs):
            cv.put(cv.row, a, cv.row, b, lab, role="head", size=st.fsize - 1)
        cv.set_h(cv.row, 28)
        cv.row += 1
        for (a, b), key, lab, val in zip(spans, W5_KEYS, labs, vals):
            cv.put(cv.row, a, cv.row, b, val, size=st.fsize - 1)
            cv.rec.put(key, lab, val)
        cv.row += 1


def _d3(cv: Canvas, C: dict, label: str):
    st = cv.st
    cap = st.labels["containment_actions"]
    headers = ["No", "処置内容", "実施日", "担当", "効果確認"]
    cv.table("containment_actions", cap, headers, C["containment_rows"], [1, 9, 3, 2.4, 7], fixed=5 if st.fixed_rows else 0, center_cols=(0, 2, 3))
    cv.text_block("action", C["action"], lw=4, min_rows=1)
    cv.text_block("parts", C["parts"], lw=4)
    cv.text_block("containment_verification", C["containment_verification"], lw=4, min_rows=1)


def _d4(cv: Canvas, C: dict, label: str):
    st = cv.st
    cv.text_block("investigation", C["investigation"], lw=4, min_rows=2)
    cv.text_block("cause", C["cause"], lw=4)
    why = C["why_why"]
    wl = st.labels["why_why"]
    if st.why_style == "rows":
        r0 = cv.row
        (a, b), (c, d), (e, f) = _alloc(cv.cx0, cv.n, [4, 2.5, 17.5])
        for j, w in enumerate(why, 1):
            cv.put(cv.row, c, cv.row, d, f"なぜ{j}", role="head", size=st.fsize - 1)
            cv.put(cv.row, e, cv.row, f, w, v="center")
            cv.row += 1
        cv.put(r0, a, cv.row - 1, b, wl, role="label")
    else:
        text = "\n".join(f"{'　' * (j - 1)}{'→ ' if j > 1 else ''}なぜ{j}：{w}" for j, w in enumerate(why, 1))
        (a, b), (c, d) = _alloc(cv.cx0, cv.n, [4, 20])
        cv.put(cv.row, a, cv.row, b, wl, role="label")
        cv.put(cv.row, c, cv.row, d, text)
        cv.row += 1
    cv.rec.put("why_why", wl, list(why))
    cv.text_block("escape_cause", C["escape_cause"], lw=4)
    # 特性要因（6M）
    fl = st.labels["fishbone"]
    legend = random_choice(C, "legend", FISH_LEGEND)
    (a, b), (c, d) = _alloc(cv.cx0, cv.n, [9, 15])
    cv.put(cv.row, a, cv.row, b, fl, role="label", h="left")
    cv.put(cv.row, c, cv.row, d, legend, size=st.fsize - 1.5, h="right", v="center")
    cv.row += 1
    cv.rec.labels["fishbone"] = fl
    ncol = 3 if st.fish == "2x3" else 2
    spans = _alloc(cv.cx0, cv.n, [1] * ncol)
    items = list(zip(FISH_KEYS, st.fish_labels, C["fishbone"]))
    for i in range(0, 6, ncol):
        chunk = items[i:i + ncol]
        for (a, b), (key, lab, val) in zip(spans, chunk):
            cv.put(cv.row, a, cv.row, b, lab, role="head")
        cv.row += 1
        for (a, b), (key, lab, val) in zip(spans, chunk):
            cv.put(cv.row, a, cv.row + 1, b, val, size=st.fsize - 0.5)
            cv.rec.put(key, lab, val)
        cv.row += 2


def _d5(cv: Canvas, C: dict, label: str):
    st = cv.st
    yes, no = st.adopt_word
    rows = []
    for row in C["candidates"]:
        rr = list(row)
        rr[6] = yes if rr[6] == "採用" else rr[6].replace("不採用", no)
        rows.append(rr)
    headers = ["No", "対策案", "効果", "コスト", "工期", "評価", "採否"]
    cv.table("countermeasure_candidates", st.labels["countermeasure_candidates"], headers, rows, [1, 10, 1.6, 3.4, 1.8, 1.6, 3],
             fixed=5 if st.fixed_rows else 0, center_cols=(0, 2, 4, 5, 6))
    if st.list_sheet:
        dv = DataValidation(type="list", formula1="=リスト!$A$1:$A$4", allow_blank=True)
    else:
        dv = DataValidation(type="list", formula1='"◎,○,△,×"', allow_blank=True)
    cv.ws.add_data_validation(dv)
    spans = _alloc(cv.cx0, cv.n, [1, 10, 1.6, 3.4, 1.8, 1.6, 3])
    first = cv.row - max(len(rows), 5 if st.fixed_rows else 0)
    for col in (spans[2][0], spans[5][0]):
        dv.add(f"{get_column_letter(col)}{first}:{get_column_letter(col)}{cv.row - 1}")
    cv.text_block("selection_reason", C["selection_reason"], lw=4)


def _d6(cv: Canvas, C: dict, label: str):
    st = cv.st
    cv.table("permanent_actions", st.labels["permanent_actions"], ["No", "実施内容", "担当", "完了日", "状況"], C["permanent_rows"],
             [1, 12, 2.6, 3, 3.4], fixed=4 if st.fixed_rows else 0, center_cols=(0, 2, 3, 4))
    headers = ["項目", "単位", C["effect_headers_before"], C["effect_headers_after"], "目標値", "判定"]
    cv.table("effect_data", st.labels["effect_data"], headers, C["effect_rows"], [8, 3, 3.2, 3.2, 4.2, 2.4], center_cols=(1, 2, 3, 4, 5))
    cv.text_block("verification_period", C["verification_period"], lw=4)
    cv.text_block("result", C["result"], lw=4)
    if st.photos and C["photos"]:
        pl = st.labels["photo_captions"]
        cv.caption(pl)
        (a, b), (c, d) = _alloc(cv.cx0, cv.n, [1, 1])
        top = cv.row
        nrows = 9 if st.grid == "hougan" else 8
        cv.put(top, a, top + nrows - 1, b, None)
        cv.put(top, c, top + nrows - 1, d, None)
        per = 172 / nrows
        for x in range(top, top + nrows):
            cv.set_h(x, per)
        when = C["permanent_rows"][0][3] if C["permanent_rows"] else C["recovered_at"].v.date()
        seed = C["report_id"]
        cv.image(_photo_png(C["photos"][0], C["recovered_at"].v.date(), seed + "b", False), top, a, 240, 165, off_x=10, off_y=6)
        cv.image(_photo_png(C["photos"][1], when, seed + "a", True), top, c, 240, 165, off_x=10, off_y=6)
        cv.row = top + nrows
        cv.put(cv.row, a, cv.row, b, C["photos"][0], h="center", v="center", size=st.fsize - 1)
        cv.put(cv.row, c, cv.row, d, C["photos"][1], h="center", v="center", size=st.fsize - 1)
        cv.row += 1
        cv.rec.put("photo_captions", pl, list(C["photos"]))


def _d7(cv: Canvas, C: dict, label: str):
    st = cv.st
    cv.text_block("prevention", C["prevention"], lw=4)
    cv.table("standard_docs", st.labels["standard_docs"], ["文書名", "文書No.", "改訂内容", "版", "改訂日"], C["docs_rows"],
             [5.5, 4, 8.5, 3.2, 2.8], fixed=4 if st.fixed_rows else 0, center_cols=(3, 4))
    cv.text_block("horizontal_deployment", C["horizontal_deployment"], lw=4)


def _d8(cv: Canvas, C: dict, label: str):
    st = cv.st
    cv.text_block("recognition", C["recognition"], lw=4, min_rows=2)
    items = [("approver", C["approver"] or ""), ("approval_date", C["approval_date"])]
    if C["claim"]:
        items.append(("customer_answer_date", C["customer_answer_date"]))
        cv.kv_row(items, weights=[2.5, 5.5, 2.5, 5.5, 2.5, 5.5])
    else:
        cv.kv_row(items, weights=[3, 9, 3, 9])
    cv.text_block("approval_comment", C["approval_comment"], lw=4)


def _history_sheet(wb: Workbook, st: Style):
    ws = wb.create_sheet("改訂履歴")
    rows = [("版", "改訂日", "改訂内容", "承認"), ("Rev.0", "2016/04/01", "制定（ISO9001対応）", "品証部長"),
            ("Rev.1", "2018/10/01", "D4に流出原因欄を追加", "品証部長"), ("Rev.2", "2021/04/01", "特性要因(6M)欄を追加、D7に改訂文書No欄", "品証部長")]
    if st.form_no.endswith("Rev.3"):
        rows.append(("Rev.3", "2023/04/01", "2シート構成に変更、効果確認データ表を追加", "品証部長"))
    if "022" in st.form_no:
        rows = [("版", "改訂日", "改訂内容", "承認"), ("Rev.1", "2020/07/01", "顧客クレーム用として QA-F-021 から派生", "品証部長")] + \
               ([("Rev.2", "2024/01/15", "1シート構成（顧客提出用）を追加", "品証部長")] if st.form_no.endswith("Rev.2") else [])
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v)
            c.border = BOX
            c.font = Font(name=st.font, size=10, bold=i == 1)
    for col, w in zip("ABCD", (8, 12, 48, 12)):
        ws.column_dimensions[col].width = w


def _list_sheet(wb: Workbook, st: Style):
    ws = wb.create_sheet("リスト")
    for i, v in enumerate(["◎", "○", "△", "×"], 1):
        ws.cell(i, 1, v)
    for i, v in enumerate(["完了", "効果確認中", "条件付き完了", "継続対応"], 1):
        ws.cell(i, 2, v)
    ws.sheet_state = "hidden"


def render(C: dict, st: Style) -> tuple[Workbook, Recorder]:
    wb = Workbook()
    rec = Recorder()
    ws1 = wb.active
    ws1.title = st.sheet_names[0]
    cv = Canvas(ws1, st, rec)
    footer = f"{st.form_no}"
    _header(cv, C, 1)
    sections = [_d1, _d2, _d3, _d4, _d5, _d6, _d7, _d8]
    for i, fn in enumerate(sections):
        if st.two_sheet and i == 4:
            cv.finish(footer)
            ws2 = wb.create_sheet(st.sheet_names[1])
            cv = Canvas(ws2, st, rec)
            _header(cv, C, 2)
        lab = cv.section(i, lambda label, fn=fn: fn(cv, C, label))
        rec.labels.setdefault(f"section_d{i + 1}", lab)
    cv.finish(footer)
    if st.list_sheet:
        _list_sheet(wb, st)
    if st.history_sheet:
        _history_sheet(wb, st)
    # 文書プロパティ（作成者・更新者）
    wb.properties.creator = C["reporter"]
    wb.properties.lastModifiedBy = C["approver"] or C["reporter"]
    wb.properties.title = st.title
    return wb, rec


# ---------------------------------------------------------------------------
# 決定的な保存（zip内タイムスタンプ・文書プロパティの日時を固定）
# ---------------------------------------------------------------------------
def _save_deterministic(wb: Workbook, path: Path, stamp: datetime):
    wb.properties.created = stamp
    buf = io.BytesIO()
    wb.save(buf)
    iso = stamp.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)", rb"\g<1>" + iso + rb"\g<2>", data)
                data = re.sub(rb"(<dcterms:created[^>]*>)[^<]*(</dcterms:created>)", rb"\g<1>" + iso + rb"\g<2>", data)
            zi = zipfile.ZipInfo(info.filename, date_time=(stamp.year, stamp.month, stamp.day, stamp.hour, stamp.minute, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            dst.writestr(zi, data)
    path.write_bytes(out.getvalue())


def _file_name(C: dict, idx: int, r) -> str:
    eid = C["equipment_id"]
    rd = C["report_date"].v
    if C["claim"]:
        cust = re.sub(r"株式会社|（株）|社内：|（.*?）", "", C["customer_name"]).strip()
        opts = [f"【8D】{C['report_id']}_{cust}.xlsx", f"8D報告書_{C['claim_no']}_{eid}.xlsx", f"顧客クレーム8D_{C['report_id']}.xlsx"]
    else:
        opts = [f"8D報告書_{C['report_id']}_{eid}.xlsx", f"是正処置報告書_{eid}_{rd:%Y%m%d}.xlsx", f"{C['report_id']}_{eid}_8D.xlsx"]
    return r.choice(opts)


# ---------------------------------------------------------------------------
# README
# ---------------------------------------------------------------------------
def _readme(entries: list[dict], styles: list[Style]) -> str:
    from collections import Counter
    lv = Counter(s.layout_version for s in styles)
    lines = [
        "# F3 是正処置報告書（8D形式）サンプル",
        "",
        "`python -m scripts.samples.f3_8d_report` で生成（固定シード。再実行してもバイト単位で同一のファイルになる）。",
        "内容は正準トラブル履歴（`scripts/samples/domain.py` の `standard_incidents()`）から選んだ30件に基づく。",
        "`関連トラブルNo.` から元のトラブル記録（TR-YYYY-NNNNN）をたどれる。",
        "",
        "## 帳票の構成",
        "",
        "| 区分 | 内容 |",
        "|---|---|",
        "| ヘッダ | 様式番号、表題、承認/確認（審査）/作成の押印欄、管理No・発行日・版数、関連トラブルNo、設備No・設備名・ライン・工程・メーカー、発生/復旧日時、停止時間、ステータス |",
        "| クレーム様式のみ | 顧客名、クレームNo、受付日、回答期限、不良数量、品名/品番、流出ロット |",
        "| D1 | チーム編成表（役割・氏名・所属・役職・担当内容）。チャンピオン（課長）、リーダー、メンバー、メーカーFE |",
        "| D2 | 問題の概要、アラーム、5W2H（What/When/Where/Who/Why/How/How many） |",
        "| D3 | 暫定処置表（No・処置内容・実施日・担当・効果確認）、復旧処置、使用部品、暫定処置の効果確認 |",
        "| D4 | 調査内容、発生原因、なぜなぜ分析、流出原因（発生原因と別欄）、特性要因6セル（人/機械/材料/方法/測定/環境、◎○×の評価つき） |",
        "| D5 | 対策案の比較表（効果・コスト・工期・評価・採否）、選定理由 |",
        "| D6 | 恒久対策の実施表、対策前後の効果確認データ表、確認期間、効果判定、（一部）対策前後の写真 |",
        "| D7 | 再発防止策、標準化（改訂文書名・文書No・改訂内容・版・改訂日）、水平展開 |",
        "| D8 | チームの称賛、承認者・承認日（クレーム様式は顧客回答日も）、承認コメント |",
        "",
        "## 様式バージョン（layout_version）",
        "",
        "様式の版ごとにサブフォルダを分けてある（サブフォルダの中は .xlsx だけ）。",
        "1つのフォルダの中は同じ様式・同じシート構成のファイルなので、フォルダごと `帳票取り込み` に入れれば、",
        "帳票の種類と読み取るシートを1回選ぶだけで、そのフォルダのファイルをまとめて読み取れる。",
        "",
        "| フォルダ | layout_version | 件数 | 特徴 |",
        "|---|---|---|---|",
    ]
    desc = {
        "QA-F-021_Rev.3_2sheet_hougan": "社内用の現行版。方眼紙（36列）、見出し帯。「8D報告(1)」にヘッダ〜D4、「8D報告(2)」にD5〜D8",
        "QA-F-021_Rev.3_1sheet_hougan": "社内用の現行版を1シートにまとめたもの（縦長、セクション単位で改ページ）",
        "QA-F-021_Rev.2_1sheet_cols": "社内用の旧版。列レイアウト（左端の見出し列＋20列）、左端列に縦結合のセクション見出し（\"D1\" のみの表記が多い）",
        "QA-F-022_Rev.1_2sheet_hougan": "顧客クレーム用。方眼紙、2シート。顧客名・クレームNo等のブロックがヘッダに入る",
        "QA-F-022_Rev.2_1sheet_cols": "顧客クレーム用の顧客提出版。列レイアウト・1シート",
    }
    for k, v in sorted(lv.items()):
        lines.append(f"| `{VERSION_DIRS[k]}/` | `{k}` | {v} | {desc.get(k, '')} |")
    lines += [
        "",
        f"顧客クレーム様式は {sum(s.claim for s in styles)} 件（全体の1/3）。顧客名・製品名はすべて架空。",
        "",
        "## 意図的なゆらぎ（抽出テスト用）",
        "",
        "- **シート構成**: 2シート（`8D報告(1)` / `8D報告(2)`）と1シート（`8D報告` `8D報告書` `是正処置報告` `8D`）。2枚目のヘッダに管理No・設備Noを再掲",
        "- **非帳票シート**: 一部に非表示の `リスト` シート（入力規則の選択肢）や `改訂履歴` シート（様式の改訂履歴）がある",
        "- **セクション見出し**: `D1 チーム編成` / `D1：チームの結成` / `D1`（日本語名なし）/ `D1 Team（チーム編成）`。列レイアウトでは左端の縦結合セルに縦書き",
        "- **項目ラベルの表記ゆれ**: 例 `管理No.`/`8D No.`/`報告書No.`/`管理番号`、`発行日`/`作成日`/`起票日`/`報告日`、`設備No.`/`設備番号`/`装置No.`/`号機`。実際に使ったラベルは `_expected.jsonl` の `labels_used`",
        "- **5W2H の配置**: 縦並び（ラベル左・値右）/ 2列グリッド / 横並び（ラベル行の下に値行）。ラベルも `What（何が）`/`何が`/`What`/`①何が(What)` とゆれる",
        "- **なぜなぜ分析**: `なぜ1〜なぜN` の行形式と、1セル内に矢印で連ねた形式",
        "- **特性要因**: 2行×3列 と 3行×2列。各セルは「◎/○/× 要因 → 検証結果」の複数行テキスト。凡例セルは項目ではない",
        "- **日付**: 実日付セル（表示形式 `yyyy/mm/dd`・`yyyy/m/d`・`yyyy\"年\"m\"月\"d\"日\"`・和暦 `ggge年m月d日`）と文字列（`2025/03/14`・`2025年3月14日`・`R7.3.14` 等）。表の中の日付は `3/14` のように年なしの場合もある",
        "- **停止時間**: 数値セル（表示形式 `#,##0\"分\"`）または文字列 `1,032分（17.2h）`",
        "- **押印欄**: 丸印の画像（セルは空）または氏名の文字列。効果確認中の報告書は承認欄が空",
        "- **写真**: 一部のファイルは D6 に対策前後の写真（画像2枚＋キャプションセル）",
        "- **空欄の表現**: `なし` / `－` / 空セル。様式の固定行数ぶんの空行が表の下に残るファイルがある",
        "- **全角数字**: 一部の報告書は本文中の数字が全角（IDやコードは半角のまま）。元トラブル記録の記入者の癖（半角カナ・誤変換）もそのまま残る",
        "- **採否の表記**: `採用/不採用`、`○採用/×不採用`、`採/否`。判定は `○/△`、`OK/NG`、`良/要監視`",
        "- **表示モード**: 一部シートは改ページプレビューで保存",
        "- **文字サイズ**: 狭い欄に入る長いラベル（`顧客クレームNo.` など）は本文より小さいフォントで折り返している",
        "- **効果確認表の見出し**: `対策前/対策後`、`対策前（実績）/対策後（実績）`、`Before/After`、`対策前/対策後（実績）` の組み合わせ",
        "",
        "## 中身の整合性",
        "",
        "- D6 の「同一不具合の発生件数」「関連チョコ停」「MTBF」は、正準データの同一設備・同一サブシステムを対策前後90日（MTBFは180日）で実際に数えた値。対策後の期間が足りない報告書は `効果確認中`",
        "- 対策後も再発が減らない案件（最大3件）は判定 `△`/`NG`/`要監視` になり、ステータスは `条件付き完了`。減ったがゼロでない案件は `完了（監視継続）` のことがある",
        f"- 直近の{N_PENDING}件は恒久対策の実施から日が浅く `効果確認中`。指標の対策後の値は `－`、判定は `確認中`（同一不具合件数は対策後30日以上経っていれば途中経過を記入）。承認者・承認日・承認印は空欄",
        "- 同一不具合件数の目標は `0`/`0件` または `対策前の1/3以下`。判定記号は目標と対策後の値から決めており、効果判定の文章と食い違わない",
        "- D3 の担当は作業内容に合わせる（初動は発見者、ロット判定は品証、同型機点検は保全）。D6 の担当も出荷・検査→品証、監視・データ→生産技術、設備→保全",
        "- 特性要因の ◎ は原因区分に応じたカテゴリ（作業ミス・設定ミス→人、異物・前工程起因→材料、調整不良・施工不良→方法、その他→機械）に置き、なぜなぜの最終段は方法（人に関する内容なら人）に置く",
        "- D7 の文書Noは設備種別ごとに固定で、報告書の時系列順に版数が上がる（同じ文書が複数の報告書に出てくると版が進む）",
        "- クレーム様式は、流出ロットの当社処理 → 顧客発見 → 受付 → 暫定処置 → 恒久対策 → 承認 → 顧客回答 の順に日付が並ぶ",
        "",
        "## _expected.jsonl",
        "",
        "1行1ファイル: `{\"file\", \"layout_version\", \"values\", \"labels_used\"}`。`_expected.jsonl` はこのフォルダ直下に置く。",
        "",
        "- `file` はこのフォルダからの相対パス（`版フォルダ/ファイル名.xlsx`。区切りは `/`）",
        "- `values` は人が画面で読む表示どおりの文字列（日付セルは表示形式を適用した文字列、停止時間は `1,032分` など）",
        "- 表の項目（`team_members` `containment_actions` `countermeasure_candidates` `permanent_actions` `effect_data` `standard_docs`）は、列見出しをキーにした行dictのリスト。様式の空行は含まない",
        "- `why_why` と `photo_captions` は文字列のリスト",
        "- 押印欄は文字で書かれているときだけ `stamp_approve` / `stamp_check` / `stamp_create` を入れる（画像の印影は値に含めない）",
        "- 空欄の項目はキー自体を出さない（例: 効果確認中の `approver` / `approval_date`）",
        "- `labels_used` はそのファイルでその値に対応するラベルセルの文字列。表はキャプションまたはセクション見出し",
        "",
        "## ファイル一覧",
        "",
        "| ファイル（版フォルダ/ファイル名） | layout_version | 設備 | 関連トラブル | ステータス |",
        "|---|---|---|---|---|",
    ]
    for e in entries:
        v = e["values"]
        lines.append(f"| {e['file']} | `{e['layout_version']}` | {v.get('equipment_id', '')} | {v.get('related_incident_id', '')} | {v.get('status', '')} |")
    return "\n".join(lines) + "\n"


# ---------------------------------------------------------------------------
# エントリポイント
# ---------------------------------------------------------------------------
def generate(output_root: Path | None = None) -> list[Path]:
    root = Path(output_root) if output_root is not None else D.OUTPUT_ROOT
    out_dir = root / "forms" / FOLDER
    out_dir.mkdir(parents=True, exist_ok=True)
    # 古いファイルの掃除（ファイル名や版フォルダを変えても前回の出力が残らないように）。
    # 版フォルダの中まで消し、いまは作らない版のフォルダは空にしたうえで削除する。_README.md と _expected.jsonl は残す。
    keep_dirs = set(VERSION_DIRS.values())
    for old in out_dir.glob("*.xlsx"):
        old.unlink()
    for sub in sorted(p for p in out_dir.iterdir() if p.is_dir()):
        for old in sub.glob("*.xlsx"):
            old.unlink()
        if sub.name not in keep_dirs:
            with suppress(OSError):     # 中身が空になったときだけ消える
                sub.rmdir()
    plans = _plans()
    docs = _DocRegistry()
    entries, styles, paths = [], [], []
    names_used = set()
    claim_rank = internal_rank = 0
    for idx, plan in enumerate(plans):
        C = build_content(plan, idx, docs)
        if plan.claim:
            st = _make_style(idx, True, claim_rank)
            claim_rank += 1
        else:
            st = _make_style(idx, False, [3, 11, 17, 5, 14, 0, 8, 19, 2, 12, 6, 16, 9, 1, 15, 4, 13, 7, 18, 10][internal_rank % 20])
            internal_rank += 1
        wb, rec = render(C, st)
        name = _file_name(C, idx, D.rng(f"f3_8d:name:{idx}"))
        if name in names_used:
            name = name.replace(".xlsx", f"_{idx:02d}.xlsx")
        names_used.add(name)
        sub = VERSION_DIRS[st.layout_version]          # 様式の版ごとにサブフォルダを分ける
        path = out_dir / sub / name
        path.parent.mkdir(parents=True, exist_ok=True)
        rd = C["report_date"].v
        _save_deterministic(wb, path, datetime(rd.year, rd.month, rd.day, 17, 30))
        labels = {k: v for k, v in rec.labels.items() if not k.startswith("section_")}
        for i in range(8):
            labels[f"section_d{i + 1}"] = rec.labels[f"section_d{i + 1}"]
        # file はこのフォルダからの相対パス（区切りは "/"。評価スクリプトが folder / e["file"] で開く）
        entries.append(dict(file=f"{sub}/{name}", layout_version=st.layout_version, values=rec.values, labels_used=labels))
        styles.append(st)
        paths.append(path)
    exp = out_dir / "_expected.jsonl"
    exp.write_text("".join(json.dumps(e, ensure_ascii=False) + "\n" for e in entries), encoding="utf-8")
    readme = out_dir / "_README.md"
    readme.write_text(_readme(entries, styles), encoding="utf-8")
    return paths + [exp, readme]


if __name__ == "__main__":
    import sys
    import time as _time

    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except (AttributeError, ValueError):
        pass
    t0 = _time.time()
    out = generate(Path(sys.argv[1]) if len(sys.argv) > 1 else None)
    for p in out:
        print(p)
    print(f"{len(out)} files ({_time.time() - t0:.1f}s)")
