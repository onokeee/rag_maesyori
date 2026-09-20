"""T6 装置トラブルカルテ（5,000行×20列）。表は C3 起点（上に2行・左に2列の余白）。

現場の台帳らしく、1行＝1件のカルテ。対応内容は複数人が日付つきで書き足したログ。
"""
from __future__ import annotations

from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from scripts.samples import domain
from scripts.samples.domain import rng

ROWS = 5000
START_ROW = 3          # 見出しの行
START_COL = 3          # 見出しの左端（C列）

HEADERS = [
    "カルテNo", "発生日時", "ライン", "工程", "設備番号", "設備名", "機種",
    "現象区分", "現象", "停止時間(分)", "生産影響(枚)", "原因区分", "原因",
    "処置区分", "対応内容", "使用部品", "部品費(円)", "担当者", "完了日時", "状態",
]

SYMPTOM_KINDS = ("停止", "異音・異臭", "品質異常", "アラーム", "漏れ", "温度異常", "通信異常", "その他")
CAUSE_KINDS = ("部品劣化", "調整ずれ", "汚れ・付着", "設定ミス", "ソフト不具合", "ユーティリティ", "原因不明", "作業ミス")
ACTION_KINDS = ("部品交換", "清掃", "調整", "再起動", "ソフト更新", "応急処置", "経過観察")
STATES = ("完了", "完了", "完了", "完了", "対応中", "保留")

SYMPTOMS = {
    "停止": ["搬送アームが原点復帰せず停止", "チャンバ扉のインターロックで停止", "ロードポートでウェーハ検知できず停止",
           "レシピ実行中に非常停止", "ステージ移動中に位置ずれエラーで停止"],
    "異音・異臭": ["搬送部から金属的な異音", "真空ポンプから異音、振動大", "駆動部から焦げたにおい",
              "ファンフィルタから風切り音", "コンプレッサ室から断続的な異音"],
    "品質異常": ["膜厚が規格上限を超過", "パターン欠陥が連続発生", "面内均一性が悪化",
             "エッチング残りを検出", "レジスト塗布ムラが発生"],
    "アラーム": ["ALM-2031 原点復帰エラー", "ALM-1450 チャンバ圧力異常", "ALM-3302 温度偏差大",
             "ALM-5012 通信タイムアウト", "ALM-0088 排気流量低下"],
    "漏れ": ["純水配管の継手からにじみ", "薬液ドレンから滴下を確認", "N2配管の接続部でリーク",
          "冷却水ホースから漏れ", "真空リーク（到達圧に届かず）"],
    "温度異常": ["ヒータ温度が設定値に届かない", "チラー出口温度が上昇", "ボート温度の偏差が拡大",
             "冷却水温が高め推移", "サセプタ温度のばらつき大"],
    "通信異常": ["ホストとの通信断が断続", "PLCとの応答遅延", "レシピ転送に失敗",
             "MESへの実績送信エラー", "センサ信号のノイズで誤検知"],
    "その他": ["定期点検で摩耗を確認", "オペレータからの申告で調査", "前工程の影響で待機時間が延長",
            "表示灯の球切れ", "作業手順の問い合わせ"],
}

CAUSES = {
    "部品劣化": ["ベアリングの摩耗", "Oリングの劣化（硬化）", "ベルトの伸び", "電極の消耗", "ポンプオイルの劣化"],
    "調整ずれ": ["ティーチング位置のずれ", "センサ感度の設定ずれ", "リミットの位置ずれ", "圧力設定のずれ"],
    "汚れ・付着": ["パーティクルの堆積", "薬液残渣の付着", "フィルタの目詰まり", "センサ受光部の汚れ"],
    "設定ミス": ["レシピのパラメータ誤設定", "アラーム閾値の設定誤り", "前回作業の戻し忘れ"],
    "ソフト不具合": ["制御ソフトの既知バグ", "通信リトライ処理の不備", "ログ肥大による処理遅延"],
    "ユーティリティ": ["N2供給圧の低下", "冷却水流量の低下", "電源電圧の瞬低", "排気圧の変動"],
    "原因不明": ["再現せず、監視継続", "調査中（メーカー問い合わせ中）"],
    "作業ミス": ["部品の取り付け向き誤り", "手順の抜け", "工具の置き忘れによる干渉"],
}



def _wb_font(ws, cell, *, bold=False, size=11):
    ws[cell].font = Font(bold=bold, size=size)


FIRST_BODIES = (
    "オペレータより連絡を受け現場へ。ラインは停止のまま、後工程に待ちが出ている旨を連絡。",
    "現場確認。状況を再現し、操作パネルの表示とワークの位置を写真で記録した。",
    "オペレータから聞き取り。直前のレシピ変更・部品交換はなし。前直でも一度同じ症状が出ていたとのこと。",
    "アラーム履歴を確認。同一アラームが過去3か月で2回発生しており、いずれも再起動で復旧していた。",
    "受付。第一報を保全班に展開し、代替機での流動可否を製造に確認中。",
)
MID_BODIES = (
    "分解清掃を実施。摺動部にパーティクルの堆積あり、清掃後に手動で動作確認。",
    "パラメータを前回の良品条件に戻して様子見。30分連続運転で再現なし。",
    "メーカーのサービスへ問い合わせ。ログ送付済みで回答待ち。",
    "仮復旧させ生産再開。翌日の定期停止で本対応を行う方針とした。",
    "念のため隣号機も同じ箇所を点検。異常なし。",
    "電源を落としてセンサの受光部を清掃。感度を規定値に調整。",
    "ティーチング位置を再設定。5回の往復で位置ずれがないことを確認。",
    "配管の継手を増し締め。リークチェックで漏れなし。",
    "エラーログを解析。通信タイムアウトの直前にノイズと思われる異常値を確認。",
    "保全ミーティングで共有。同型機への水平展開を検討することとした。",
    "予備品在庫を確認。在庫なしのため発注をかけた。",
    "清掃後も同じ症状が再発。別の原因を疑い、駆動部の電流値を測定して記録。",
    "測定結果を前回データと比較。基準値内だが上昇傾向が見られる。",
    "製造と調整し、次の段取り替えのタイミングで停止させてもらうことにした。",
    "一時的に手動運転に切り替えて対応。オペレータへ操作手順を説明した。",
)
PART_BODIES = (
    "{part}を手配（納期1週間）。それまでは応急処置で運転する。",
    "{part}が入荷。受け入れ確認後、交換作業の段取りを組んだ。",
    "{part}を交換。交換前後の数値を記録し、点検表に添付した。",
    "{part}の交換後、慣らし運転を1時間実施。異音なし。",
)
CLOSE_BODIES = (
    "動作確認OK。生産復帰。",
    "試運転3回実施、異常なし。クローズとする。",
    "1週間経過し再発なしを確認。完了とする。",
    "点検表と履歴に記録して完了。次回点検時に同じ箇所を重点確認する。",
    "製造へ復帰を連絡。稼働率への影響は当日分のみ。",
)
FOLLOW_BODIES = (
    "（追記）同様の事象が隣のラインでも発生したとの情報あり。情報を共有。",
    "（追記）交換した部品をメーカーへ返却し、解析を依頼した。",
    "（追記）標準作業書に点検項目を追加する方向で検討中。",
    "（追記）月次の保全会議で報告予定。",
    "（追記）その後の稼働データを確認。異常なし。",
)


def _make_log(r, occurred: datetime, people_names: list[str], part_name: str | None, state: str) -> str:
    """1つのセルに複数人が日付つきで書き足した対応ログ（下ほど新しい）。"""
    part = part_name or "該当部品"
    count = r.choice([6, 8, 8, 9, 10, 10, 12, 14])
    writers = [r.choice(people_names) for _ in range(count)]
    # 同じ人が続けて書くことも多い
    for i in range(1, count):
        if r.random() < 0.35:
            writers[i] = writers[i - 1]

    bodies: list[str] = [r.choice(FIRST_BODIES)]
    mid = list(MID_BODIES)
    r.shuffle(mid)
    parts_used = 0
    for i in range(1, count - 1):
        if part_name and parts_used < 2 and r.random() < 0.3:
            bodies.append(PART_BODIES[parts_used].format(part=part))
            parts_used += 1
        elif mid:
            bodies.append(mid.pop())
        else:
            bodies.append(r.choice(MID_BODIES))
    if state == "完了":
        bodies.append(r.choice(CLOSE_BODIES))
        if r.random() < 0.4:
            bodies.append(r.choice(FOLLOW_BODIES))
    elif state == "対応中":
        bodies.append(r.choice(["現在も調査中。", "部品の入荷待ち。", "メーカー回答待ちで経過観察中。"]))
    else:
        bodies.append(r.choice(["保留。次回の定期停止で対応する。", "優先度低のため保留とした。"]))

    lines, t = [], occurred
    for i, body in enumerate(bodies):
        t = t + timedelta(hours=r.choice([0, 1, 2, 4, 6, 20, 24, 48, 72, 120, 168]),
                          minutes=r.choice([5, 10, 20, 35, 50]))
        stamp = r.choice([
            f"{t.month}/{t.day}",
            f"{t:%Y/%m/%d}",
            f"{t.month}/{t.day} {t.hour}:{t.minute:02d}",
            f"{t.month}月{t.day}日",
        ])
        # 記入者を書く人と書かない人がいる（姓だけのことも）
        who = writers[i] if i < len(writers) else writers[-1]
        writer = who.split()[0] if r.random() < 0.5 else who
        if r.random() < 0.65:
            body = f"{writer}{r.choice(['：', ': ', ' ', '　'])}{body}"
        lines.append(f"【{stamp}】{body}")
    return "\n".join(lines)


def build(path: Path) -> Path:
    r = rng("t6-karte")
    equipment = domain.equipment_master()
    persons = domain.people()
    parts = domain.parts_catalog()
    names = [p.name for p in persons]

    wb = Workbook()
    ws = wb.active
    ws.title = "トラブルカルテ"

    # 表の外の余白（現場の台帳らしく、上にタイトルを置く）
    ws.cell(row=1, column=START_COL, value="装置トラブルカルテ").font = Font(bold=True, size=14)
    ws.cell(row=2, column=START_COL, value="製造部 設備保全課").font = Font(size=9)
    ws.cell(row=2, column=START_COL + 6, value=f"出力日: {date(2026, 4, 1):%Y/%m/%d}").font = Font(size=9)

    head_fill = PatternFill("solid", fgColor="DCE6F1")
    thin = Side(style="thin", color="AAAAAA")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for i, name in enumerate(HEADERS):
        cell = ws.cell(row=START_ROW, column=START_COL + i, value=name)
        cell.font = Font(bold=True)
        cell.fill = head_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    start = datetime(2023, 4, 1, 8, 0)
    row_no = START_ROW + 1
    for n in range(1, ROWS + 1):
        eq = r.choice(equipment)
        occurred = start + timedelta(days=r.randint(0, 1090), hours=r.randint(0, 15), minutes=r.choice([0, 5, 10, 20, 35, 50]))
        kind = r.choice(SYMPTOM_KINDS)
        cause_kind = r.choice(CAUSE_KINDS)
        action_kind = r.choice(ACTION_KINDS)
        state = r.choice(STATES)
        usable = [p for p in parts if eq.category in p.applicable_categories] or parts
        part = r.choice(usable) if action_kind == "部品交換" else (r.choice(usable) if r.random() < 0.15 else None)
        stop_min = r.choice([5, 10, 15, 20, 30, 45, 60, 90, 120, 180, 240, 480])
        if kind == "その他":
            stop_min = r.choice([0, 0, 5, 10])
        worker = r.choice(persons)
        done = occurred + timedelta(hours=r.randint(1, 96)) if state == "完了" else None

        values = [
            f"KT-{occurred:%y}-{n:05d}",
            occurred,
            eq.line,
            eq.process,
            eq.equipment_id,
            eq.name,
            eq.model,
            kind,
            r.choice(SYMPTOMS[kind]),
            stop_min,
            stop_min * r.choice([0, 0, 2, 5, 8, 12]),
            cause_kind,
            r.choice(CAUSES[cause_kind]),
            action_kind,
            _make_log(r, occurred, names, part.name if part else None, state),
            f"{part.name}（{part.part_no}）×{r.choice([1, 1, 1, 2, 4])}" if part else "",
            part.unit_price_yen if part else "",
            worker.name,
            done,
            state,
        ]
        # 現場のばらつき: たまに空欄・記号・同上
        if r.random() < 0.04:
            values[12] = r.choice(["", "－", "調査中"])
        if r.random() < 0.02:
            values[8] = "同上"
        if r.random() < 0.03:
            values[17] = ""

        for i, v in enumerate(values):
            cell = ws.cell(row=row_no, column=START_COL + i, value=v)
            cell.border = border
            if isinstance(v, datetime):
                cell.number_format = "yyyy/mm/dd hh:mm"
            elif i in (9, 10, 16) and isinstance(v, int):
                cell.number_format = "#,##0"
            if i == 14:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        row_no += 1

    widths = [14, 17, 8, 8, 12, 22, 14, 11, 34, 11, 12, 11, 26, 10, 60, 26, 11, 10, 17, 8]
    for i, w in enumerate(widths):
        ws.column_dimensions[get_column_letter(START_COL + i)].width = w
    ws.freeze_panes = ws.cell(row=START_ROW + 1, column=START_COL)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


if __name__ == "__main__":
    out = build(Path("samples/tables/T6_装置トラブルカルテ.xlsx"))
    print(out, out.stat().st_size)
