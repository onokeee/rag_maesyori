"""一覧表（Excel/CSV、1行＝1件）の取り込み。

CSV/Excel の読み取り（source）、表の範囲判定（detect）、列の対応づけ（mapping・spec）、正規化（normalize）、
チェック（checks）、記録の Markdown（markdown）、zip（outputs）、DB（store）、ジョブ（pipeline）を
この順に1ファイルにまとめてある。
"""
from __future__ import annotations

import codecs
import copy
import csv
import difflib
import functools
import gzip
import hashlib
import io
import json
import math
import os
import re
import shutil
import sqlite3
import threading
import time as _time
import unicodedata
import warnings
import zipfile
import zlib
from collections import Counter, defaultdict
from dataclasses import dataclass, field, asdict, fields
from datetime import date, datetime, time, timedelta
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Iterator, Protocol, Callable
from xml.parsers import expat

from flask import current_app
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.comments.comment_sheet import CommentSheet
from openpyxl.packaging.relationship import RelationshipList, get_dependents, get_rels_path
from openpyxl.styles.numbers import BUILTIN_FORMATS, BUILTIN_FORMATS_MAX_SIZE
from openpyxl.utils import column_index_from_string, range_boundaries
from openpyxl.utils.cell import coordinate_to_tuple
from openpyxl.utils.datetime import CALENDAR_MAC_1904, MAC_EPOCH, WINDOWS_EPOCH, from_excel
from openpyxl.worksheet._reader import INLINE_STRING, WorksheetReader, WorkSheetParser
from openpyxl.worksheet.table import Table
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.xml.constants import COMMENTS_NS
from openpyxl.xml.functions import fromstring

from app import core
from app.core import (
    UploadError,
    nfkc_keep_enclosed,
    escape_md_line,
    estimate_tokens,
    join_blocks,
    md_bullet,
    md_filename,
    upload_path,
    JobCancelled,
    JobError,
    start_job,
)



# ====================================================================================================
# 元 tables/source.py
# 表ソースの共通モデル。ExcelもCSVも読み込み後はこの形にそろえる。
#
# 見出しの検出・列の対応づけ・正規化は、すべてこのインターフェースだけを使う。
# ====================================================================================================

EXCEL_EXTENSIONS = {".xlsx", ".xlsm"}
CSV_EXTENSIONS = {".csv", ".tsv", ".txt"}

# 制御文字と見えない文字（ゼロ幅スペース・BOM・向き指定・ソフトハイフン。excel.text._INVISIBLE_RE と同じ集合）
_CONTROL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f\u200b-\u200d\u2060\ufeff\u202a-\u202e\u2066-\u2069\u00ad]")
_EXCEL_ERRORS = {"#N/A", "#DIV/0!", "#REF!", "#VALUE!", "#NAME?", "#NUM!", "#NULL!", "#GETTING_DATA"}


@dataclass
class CellInfo:
    value: object
    text: str
    number_format: str | None = None
    bold: bool = False
    strike: bool = False
    fill: bool = False
    merged_anchor: tuple[int, int] | None = None  # 結合範囲の左上 (行, 列)。1始まり。左上セル自身も自分を指す


@dataclass
class SourceRow:
    index: int  # 1始まり。CSVはレコード番号（Excelで開いたときの行番号と同じ）
    cells: list[CellInfo]
    hidden: bool | None = None  # CSVは不明なので None

    def text(self, col: int) -> str:
        """0始まりの列位置の文字列。範囲外は空。"""
        return self.cells[col].text if 0 <= col < len(self.cells) else ""

    def cell(self, col: int) -> CellInfo | None:
        return self.cells[col] if 0 <= col < len(self.cells) else None

    @property
    def is_blank(self) -> bool:
        return not any(c.text for c in self.cells)


@dataclass
class SheetInfo:
    name: str
    hidden: bool
    max_row: int
    max_col: int
    table_ranges: list[str]
    date1904: bool = False
    hidden_columns: list[int] = field(default_factory=list)  # 1始まりの列番号
    freeze_panes: str | None = None
    auto_filter: str | None = None


class TableSource(Protocol):
    kind: str  # "csv" | "excel"
    file_name: str

    def sheets(self) -> list[SheetInfo]: ...

    def rows(self, sheet: str, start: int = 1, limit: int | None = None) -> Iterator[SourceRow]: ...


def open_source(path: Path, file_name: str, options: dict | None = None) -> TableSource:
    """拡張子で CSV / Excel を振り分けて開く。"""
    ext = Path(file_name or str(path)).suffix.lower()
    if ext in EXCEL_EXTENSIONS:

        return ExcelSource(path, file_name, options)
    if ext in CSV_EXTENSIONS:

        return CsvSource(path, file_name, options)
    raise UploadError(f"対応していない形式です（{ext or '拡張子なし'}）。.xlsx / .xlsm / .csv / .tsv / .txt を選んでください")


def clean_text(text: str) -> str:
    """改行をLFにそろえ、制御文字を除いて前後の空白を削る。"""
    if not text:
        return ""
    s = text
    if "_x000D_" in s:
        s = s.replace("_x000D_", "")
    if "\r" in s:
        s = s.replace("\r\n", "\n").replace("\r", "\n")
    if _CONTROL_RE.search(s):  # 多くのセルは置き換えるものがないので、先に確かめてから置き換える
        s = _CONTROL_RE.sub("", s)
    return s.strip()


def cell_text(value) -> str:
    """セル値を表示用の文字列にする。"""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        if value.time() == time(0):
            return value.date().isoformat()
        if value.second:
            return value.strftime("%Y-%m-%d %H:%M:%S")
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M:%S" if value.second else "%H:%M")
    if isinstance(value, timedelta):
        minutes = int(value.total_seconds() // 60)
        return f"{minutes // 60}:{minutes % 60:02d}"
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return repr(round(value, 10))
    return clean_text(str(value))


# ---- 値の種類の判定（見出しらしさ・型推定で共用） ----

NA_TOKENS = {"-", "－", "―", "‐", "ー", "N/A", "n/a", "NA", "#N/A"}
_NUM_RE = re.compile(r"^[△▲+\-]?[¥]?(\d{1,3}(,\d{3})+|\d+)(\.\d+)?[%]?-?$")
_DATE_RE = re.compile(r"^(\d{4})[/\-.年](\d{1,2})[/\-.月](\d{1,2})日?(\s*\(?[月火水木金土日]?\)?)?"
                      r"((?:\s+|T)\d{1,2}:\d{2}(:\d{2}(\.\d{1,7})?)?(\s*(Z|[+\-]\d{2}:?\d{2}))?)?$")
_ERA_RE = re.compile(r"^(R|H|S|令和|平成|昭和)\s*(\d{1,2}|元)[./年](\d{1,2})[./月](\d{1,2})日?(\s+\d{1,2}:\d{2}(:\d{2})?)?$")
_MD_RE = re.compile(r"^(\d{1,2})/(\d{1,2})$")
_YMD8_RE = re.compile(r"^(19|20)\d{2}(0[1-9]|1[0-2])(0[1-9]|[12]\d|3[01])$")
_TIME_RE = re.compile(r"^\d{1,2}:\d{2}(:\d{2})?$")
_CODE_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9\-_/.#]*$")


def value_kind(value, text: str) -> str:
    """セル値の種類: blank/number/date/datetime/time/text/code/string/bool/error"""
    if not text:
        return "blank"
    if isinstance(value, bool):
        return "bool"
    if isinstance(value, datetime):
        return "date" if value.time() == time(0) else "datetime"
    if isinstance(value, date):
        return "date"
    if isinstance(value, time):
        return "time"
    if isinstance(value, (int, float, timedelta)):
        return "number"
    s = unicodedata.normalize("NFKC", text).strip()
    if s in _EXCEL_ERRORS:
        return "error"
    if "\n" in s or len(s) > 40:
        return "text"
    if _YMD8_RE.match(s):
        return "date"
    m = _DATE_RE.match(s) or _ERA_RE.match(s)
    if m:
        return "datetime" if ":" in s else "date"
    if _MD_RE.match(s):
        return "date"
    if _TIME_RE.match(s):
        return "time"
    if _NUM_RE.match(s) and not (len(s) > 1 and s[0] == "0" and s[1].isdigit()):
        return "number"
    if _CODE_RE.match(s) and any(ch.isdigit() for ch in s):
        return "code"
    return "string"


# ====================================================================================================
# 元 tables/csv_source.py
# CSV/TSV の読み込み。文字コード・区切り文字・前置き行を判定し、1レコードずつ返す。
#
# - 文字コード: BOM → UTF-8(厳密) → CP932 → shift_jis_2004。全件をデコードして確かめる
# - 区切り文字: 「最も多い列数に一致する行の割合」が高い候補を採用する
# - 引用符内の改行に対応するため csv.reader(newline='') に読ませ、行単位の分割はしない
# ====================================================================================================

SNIFF_BYTES = 1024 * 1024
SNIFF_RECORDS = 200
DELIMITERS = [",", "\t", ";", "|"]
REPLACEMENT_CHAR = "〓"
_CHUNK = 1024 * 1024
_EXCEL_WRAPPED = re.compile(r'^="(.*)"$', re.S)
_TRAILER_WORDS = re.compile(r"(合計|総計|件数|END|EOF|以上)", re.I)
# 1つの値がこれより多くの物理行にまたがるなら、" の閉じ忘れとみなす（セルの中の改行としては多すぎる）
MAX_RECORD_LINES = 500
# 列数の上限。見出しごとに辞書を引くので、列が桁違いに多いと画面の処理が何秒も止まる（一覧表は多くても数百列）
MAX_COLUMNS = 2000
_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_ZIP_MAGIC = b"PK\x03\x04"

csv.field_size_limit(16 * 1024 * 1024)

# 読めない文字を〓に置き換えるエラーハンドラ。置き換えた回数をスレッドごとに数える
_geta_state = threading.local()


def _geta_handler(err: UnicodeDecodeError):
    _geta_state.count = getattr(_geta_state, "count", 0) + 1
    return REPLACEMENT_CHAR, err.end


codecs.register_error("tables_geta", _geta_handler)


@dataclass
class CsvSniff:
    encoding: str
    bom: bool
    delimiter: str
    preamble_rows: int
    header_row: int
    trailer_rows: int
    confidence: float
    warnings: list[str] = field(default_factory=list)
    decode_error_line: int | None = None  # 全件デコードで失敗した物理行（おおよそ）
    max_columns: int = 0  # 先頭のレコードで最も多い列数（末尾の空列を除く）


def sniff_csv(path) -> CsvSniff:
    path = Path(path)
    with path.open("rb") as f:
        head = f.read(SNIFF_BYTES)
    warnings: list[str] = []
    if not head:
        return CsvSniff("utf-8", False, ",", 0, 1, 0, 0.0, ["ファイルが空です"])

    encoding, bom, error_line = _detect_encoding(path, head)
    if error_line is not None:
        warnings.append(
            f"{error_line}行目付近に、文字コード {encoding} として読めない文字があります。"
            f"文字コードを選び直すか、「読めない文字を{REPLACEMENT_CHAR}に置き換える」を選んでください"
        )

    if encoding in ("cp932", "shift_jis_2004"):
        utf8_lines = _utf8_lines(path)
        if utf8_lines:
            listed = "・".join(str(n) for n in utf8_lines[:5]) + ("…" if len(utf8_lines) > 5 else "")
            warnings.append(
                f"UTF-8 の行が混ざっています（{listed}行目）。{encoding.upper()} として読むと、その行の字が化けます。"
                "文字コードをそろえてから取り込み直してください"
            )

    truncated = path.stat().st_size > len(head)
    text = _decode_head(head, encoding)
    delimiter, ratio, records = _detect_delimiter(text, truncated, path.suffix.lower() == ".tsv")
    if ratio == 0.0:
        warnings.append("区切り文字を判定できませんでした（列が1つだけです）。固定長テキストは対象外です")

    header_row = _guess_header_record(records)
    trailer_rows = _count_trailer_rows(path, encoding, delimiter, _modal_width(records))
    confidence = round(ratio * (1.0 if error_line is None else 0.6), 3)
    max_columns = max((_width(r) for r in records), default=0)
    if not records and truncated:
        # 先頭のレコードが読み取り範囲より長い（何十万列もの見出し行など）。区切り文字の数で列数の見当をつける
        first = text.split("\n", 1)[0]
        max_columns = max(first.count(d) for d in (delimiter, ",", "\t", ";", "|")) + 1
    return CsvSniff(
        encoding=encoding,
        bom=bom,
        delimiter=delimiter,
        preamble_rows=max(0, header_row - 1),
        header_row=header_row,
        trailer_rows=trailer_rows,
        confidence=confidence,
        warnings=warnings,
        decode_error_line=error_line,
        max_columns=max_columns,
    )


# ---- 文字コード ----

def _detect_encoding(path: Path, head: bytes) -> tuple[str, bool, int | None]:
    if head.startswith(codecs.BOM_UTF8):
        return "utf-8-sig", True, _full_decode_error_line(path, "utf-8-sig")
    if head.startswith(codecs.BOM_UTF16_LE) or head.startswith(codecs.BOM_UTF16_BE):
        return "utf-16", True, _full_decode_error_line(path, "utf-16")
    sample = head[:4096]
    if len(sample) >= 4:
        even_nul = sample[0::2].count(0) / max(1, len(sample[0::2]))
        odd_nul = sample[1::2].count(0) / max(1, len(sample[1::2]))
        if odd_nul > 0.3 and even_nul < 0.05:
            return "utf-16-le", False, _full_decode_error_line(path, "utf-16-le")
        if even_nul > 0.3 and odd_nul < 0.05:
            return "utf-16-be", False, _full_decode_error_line(path, "utf-16-be")

    best: tuple[int, str, int] | None = None  # (失敗位置, 文字コード, 行)
    cp932_line: int | None = None
    for enc in ("utf-8", "cp932", "shift_jis_2004"):
        err = _full_decode_error(path, enc)
        if err is None:
            if enc == "shift_jis_2004" and cp932_line is not None and _has_cp932_only_chars(path):
                # CP932 の拡張文字（髙・﨑 など）を含むファイルに読めないバイトが混じっただけ。
                # Shift_JIS-2004 で読むと 髙→郄 のように別の字に化けるので、CP932 のまま読めない行を知らせる
                return "cp932", False, cp932_line
            return enc, False, None
        pos, line = err
        if enc == "cp932":
            cp932_line = line
        if best is None or pos > best[0]:
            best = (pos, enc, line)
    assert best is not None
    return best[1], False, best[2]


def _has_cp932_only_chars(path: Path) -> bool:
    """CP932 として正しく読めて、Shift_JIS-2004 だと別の字になる行があるか（NEC・IBM 拡張文字など）。"""
    with path.open("rb") as f:
        for raw in f:
            try:
                text = raw.decode("cp932")
            except UnicodeDecodeError:
                continue
            if text != raw.decode("shift_jis_2004", errors="replace"):
                return True
    return False


def _full_decode_error(path: Path, encoding: str) -> tuple[int, int] | None:
    """全件を分割デコードし、失敗した (バイト位置, 物理行) を返す。成功なら None。"""
    decoder = codecs.getincrementaldecoder(encoding)(errors="strict")
    consumed = 0
    newlines = 0
    with path.open("rb") as f:
        while True:
            chunk = f.read(_CHUNK)
            final = not chunk
            try:
                decoder.decode(chunk, final=final)
            except UnicodeDecodeError as e:
                # チャンク内の位置はデコーダのバッファ分ずれるが、行の目安には十分
                offset = max(0, min(len(chunk), e.start))
                return consumed + offset, newlines + chunk[:offset].count(b"\n") + 1
            if final:
                return None
            consumed += len(chunk)
            newlines += chunk.count(b"\n")


def _utf8_lines(path: Path, limit: int = 6) -> list[int]:
    """Shift_JIS 系と判定したファイルで、UTF-8 として正しく読める日本語の行（2つのファイルをつないだ等）の行番号。

    UTF-8 の日本語のバイト列は CP932 としてもほぼ読めてしまい、エラーにならずに字が化けるため、行ごとに確かめる。
    """
    found: list[int] = []
    with path.open("rb") as f:
        for n, line in enumerate(f, 1):
            if line.isascii():
                continue
            try:
                line.decode("utf-8")
            except UnicodeDecodeError:
                continue
            found.append(n)
            if len(found) >= limit:
                break
    return found


def _full_decode_error_line(path: Path, encoding: str) -> int | None:
    err = _full_decode_error(path, encoding)
    return None if err is None else err[1]


def _decode_head(head: bytes, encoding: str) -> str:
    decoder = codecs.getincrementaldecoder(encoding)(errors="tables_geta")
    return decoder.decode(head, final=False).replace("\x00", "")


# ---- 区切り文字と見出し ----

def _parse_records(text: str, delimiter: str, truncated: bool) -> list[list[str]]:
    records: list[list[str]] = []
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=False)
    try:
        for rec in reader:
            records.append(rec)
            if len(records) > SNIFF_RECORDS:
                break
    except csv.Error:
        pass
    if truncated and len(records) <= SNIFF_RECORDS and records:
        records.pop()  # 途中で切れた最後のレコードは使わない
    return records[:SNIFF_RECORDS]


def _width(rec: list[str]) -> int:
    """末尾の空列を除いた列数。"""
    n = len(rec)
    while n and not rec[n - 1].strip():
        n -= 1
    return n


def _modal_width(records: list[list[str]]) -> int:
    widths = [len(r) for r in records if len(r) >= 2]
    if not widths:
        return 1
    return Counter(widths).most_common(1)[0][0]


def _detect_delimiter(text: str, truncated: bool, prefer_tab: bool) -> tuple[str, float, list[list[str]]]:
    best: tuple[tuple, str, list[list[str]]] | None = None
    for delim in DELIMITERS:
        records = _parse_records(text, delim, truncated)
        nonblank = [r for r in records if _width(r) > 0]
        if not nonblank:
            continue
        mode = _modal_width(nonblank)
        if mode < 2:
            ratio = 0.0
        else:
            ratio = sum(1 for r in nonblank if len(r) == mode) / len(nonblank)
        header_like = _header_likeness(nonblank, mode)
        rank = (round(ratio, 2), header_like, 1 if (prefer_tab and delim == "\t") else 0, mode)
        if best is None or rank > best[0]:
            best = (rank, delim, records)
    if best is None:
        return ("\t" if prefer_tab else ","), 0.0, []
    return best[1], best[0][0], best[2]


def _header_likeness(records: list[list[str]], mode: int) -> float:
    for rec in records[:30]:
        if len(rec) == mode:
            cells = [c.strip() for c in rec if c.strip()]
            if not cells:
                continue
            strings = sum(1 for c in cells if value_kind(c, c) in ("string", "code"))
            return round(strings / len(cells), 2)
    return 0.0


def _guess_header_record(records: list[list[str]]) -> int:
    """最頻の列数に一致し、次の行も同じ列数になる最初のレコード（1始まり）。"""
    mode = _modal_width(records)
    if mode < 2:
        return 1
    for i, rec in enumerate(records):
        if len(rec) != mode:
            continue
        cells = [c.strip() for c in rec if c.strip()]
        if len(cells) < max(2, mode * 0.5):
            continue
        strings = sum(1 for c in cells if value_kind(c, c) in ("string", "code"))
        if strings / len(cells) < 0.7:
            continue
        following = [r for r in records[i + 1:i + 4] if _width(r) > 0]
        if following and all(len(r) != mode for r in following):
            continue
        return i + 1
    return 1


def _count_trailer_rows(path: Path, encoding: str, delimiter: str, mode: int) -> int:
    """末尾の件数行・合計行（列数が合わない短い行）を数える。"""
    size = path.stat().st_size
    if size == 0 or mode < 2:
        return 0
    with path.open("rb") as f:
        f.seek(max(0, size - 65536))
        tail = f.read()
    text = tail.decode(encoding.replace("-sig", ""), errors="tables_geta").replace("\r\n", "\n").replace("\r", "\n")
    lines = text.split("\n")
    if size > 65536:
        lines = lines[1:]  # 先頭は途中から始まるので捨てる
    count = 0
    for line in reversed(lines):
        if not line.strip():
            continue
        if '"' in line or count >= 5:
            break
        cells = next(csv.reader([line], delimiter=delimiter))
        if len(cells) < mode * 0.5 and (_width(cells) <= 3 or _TRAILER_WORDS.search(line)):
            count += 1
            continue
        break
    return count


# ---- 本体 ----

class CsvSource:
    kind = "csv"

    def __init__(self, path, file_name: str | None = None, options: dict | None = None):
        self.path = Path(path)
        self.file_name = file_name or self.path.name
        options = dict(options or {})
        self.sniff: CsvSniff | None = None
        if not options.get("encoding") or not options.get("delimiter"):
            self.sniff = sniff_csv(self.path)
        self.encoding: str = options.get("encoding") or self.sniff.encoding
        self.delimiter: str = options.get("delimiter") or self.sniff.delimiter
        if self.sniff is not None and self.sniff.max_columns > MAX_COLUMNS:
            raise UploadError(f"列数が上限（{MAX_COLUMNS:,}列）を超えています（{self.sniff.max_columns:,}列）。"
                              "不要な列を削除して保存し直してください")
        if not isinstance(self.delimiter, str) or len(self.delimiter) != 1:
            raise UploadError("区切り文字は1文字で選んでください")
        _reject_binary(self.path)
        errors = str(options.get("errors") or "strict")
        self.replace_errors = errors.startswith("replace")
        self.replaced_rows: list[int] = []  # 〓に置き換えたレコード番号（直近の rows() 走査分）
        # " が閉じていないため、以降の行を1つの値として読んだレコード番号（直近の rows() 走査分）
        self.unclosed_quote_row: int | None = None
        # " は閉じているが、1つの値が MAX_RECORD_LINES 行より多くにまたがるレコード番号（閉じ忘れの疑い）
        self.long_record_row: int | None = None
        self._stats: tuple[int, int] | None = None

    @property
    def warnings(self) -> list[str]:
        return list(self.sniff.warnings) if self.sniff else []

    def close(self) -> None:
        pass

    def sheets(self) -> list[SheetInfo]:
        if self._stats is None:
            max_row = max_col = 0
            for row in self.rows(self.file_name):
                max_row = row.index
                max_col = max(max_col, len(row.cells))
            self._stats = (max_row, max_col)
        return [SheetInfo(name=self.file_name, hidden=False, max_row=self._stats[0], max_col=self._stats[1], table_ranges=[])]

    def rows(self, sheet: str | None = None, start: int = 1, limit: int | None = None) -> Iterator[SourceRow]:
        errors = "tables_geta" if self.replace_errors else "strict"
        self.replaced_rows = []
        self.unclosed_quote_row = None
        self.long_record_row = None
        index = 0
        emitted = 0
        state = {"eof": False}

        def physical_lines(f):
            for line in f:
                yield line.replace("\x00", "")
            state["eof"] = True  # レコードの途中でファイルが終わった（" が閉じていない）ときだけ、ここまで読まれる

        try:
            with self.path.open("r", encoding=self.encoding, errors=errors, newline="") as f:
                reader = csv.reader(physical_lines(f), delimiter=self.delimiter, strict=False)
                line_before = 0
                for record in reader:
                    index += 1
                    if self.replace_errors and any(REPLACEMENT_CHAR in v for v in record):
                        self.replaced_rows.append(index)
                    if self.unclosed_quote_row is None and state["eof"]:
                        self.unclosed_quote_row = index
                    elif self.long_record_row is None and reader.line_num - line_before > MAX_RECORD_LINES:
                        self.long_record_row = index
                    line_before = reader.line_num
                    if len(record) > MAX_COLUMNS and _width(record) > MAX_COLUMNS:
                        # 先読みで分からなかった列数の上限もここで止める（アップロード時に全行を読むので、そこで断る）
                        raise UploadError(f"{index}行目の列数が上限（{MAX_COLUMNS:,}列）を超えています"
                                          f"（{_width(record):,}列）。不要な列を削除して保存し直してください")
                    if index < start:
                        continue
                    if limit is not None and emitted >= limit:
                        return
                    emitted += 1
                    yield SourceRow(index=index, cells=[_csv_cell(v) for v in record], hidden=None)
        except LookupError as e:
            raise UploadError(f"文字コード {self.encoding} は使えません。文字コードを選び直してください") from e
        except csv.Error as e:
            # 「field larger than field limit」など。" の閉じ忘れで残りのファイル全体が1つの値になったときに起きる
            # 英語の例外文は画面に出さず、原因の見当を日本語で示す（元の例外は from e で残す）
            hint = ("1つの値が大きすぎます。\" の閉じ忘れの可能性があります" if "field larger" in str(e)
                    else "区切り文字や \" の置き方が崩れています")
            raise UploadError(
                f"{index + 1}行目付近から、\" が閉じていないなどの理由で CSV として読めません（{hint}）。ファイルを確認してください"
            ) from e
        except UnicodeError as e:
            # UTF-16 の「BOM が無い」などは UnicodeDecodeError ではなく UnicodeError で上がる
            raise UploadError(
                f"{index + 1}行目付近で、文字コード {self.encoding} として読めない文字がありました。"
                f"文字コードを選び直すか、「読めない文字を{REPLACEMENT_CHAR}に置き換える」を選んでください"
            ) from e


def _reject_binary(path: Path) -> None:
    """拡張子が .csv / .txt / .tsv でも、中身が Excel やバイナリなら読む前に止める（文字コードの問題に見せない）。"""
    try:
        with path.open("rb") as f:
            head = f.read(4096)
    except OSError:
        return
    if head.startswith(_ZIP_MAGIC) or head.startswith(_OLE_MAGIC):
        raise UploadError("中身はExcelファイルです。拡張子を .xlsx にして選び直してください")
    if head.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return
    if len(head) >= 4:
        even_nul = head[0::2].count(0) / max(1, len(head[0::2]))
        odd_nul = head[1::2].count(0) / max(1, len(head[1::2]))
        if (odd_nul > 0.3 and even_nul < 0.05) or (even_nul > 0.3 and odd_nul < 0.05):
            return  # BOMなしの UTF-16
    control = sum(1 for b in head if b < 0x20 and b not in (0x09, 0x0A, 0x0D))
    if head and control > 0.1 * len(head):
        raise UploadError("テキストのCSVではありません（中身がバイナリです）。CSV・TSV・テキストのファイルを選んでください")


def _csv_cell(raw: str) -> CellInfo:
    if not raw:
        return CellInfo(value=None, text="")
    text = clean_text(raw)
    m = _EXCEL_WRAPPED.match(text) if text.startswith("=") else None
    if m:
        text = m.group(1).strip()
    return CellInfo(value=text if text else None, text=text)


# ====================================================================================================
# 元 tables/excel_source.py
# Excel(.xlsx/.xlsm) の読み込み。シートの XML を先頭から順に読み、ブック全体をメモリに載せない。
#
# 通常モード（load_workbook(data_only=True)）で開いたときと同じ CellInfo を返す:
# - セルの値の変換は openpyxl の WorkSheetParser をそのまま使う（日付・時刻・共有文字列・インライン文字列）
# - 結合セル・列の非表示・ウィンドウ枠・オートフィルタ・テーブル・ハイパーリンク・コメントは、
#   sheetData を除いたシートの XML を openpyxl の WorksheetReader で読む（結合範囲の値の消去やリンク先の値もそのまま）
# - 値のある最終行・最終列とセル数は、開くときに expat で1回数える（セル数の上限もここで確かめる）
# 行が昇順に並んでいないシート（まれ）は、そのシートだけ全行を読んでから並べ直す。
# ====================================================================================================

DEFAULT_MAX_CELLS = 500_000
# これより列の多いシートは、各行をその行の最後のセルまでにする（XFD1 のような遠くの値1つで
# 行数×16,384列のセルを作って読み込みが止まらないように）。256 は旧形式（.xls）の列数の上限
PAD_MAX_COLUMNS = 256
_DIMENSION_RE = re.compile(rb'<dimension ref="([A-Z]+\d+(?::[A-Z]+\d+)?)"')
_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_SHEETDATA_START = re.compile(rb"<((?:[A-Za-z_][\w.\-]*:)?)sheetData\b[^>]*?(/?)>")
_SHEETDATA_END = re.compile(rb"</(?:[A-Za-z_][\w.\-]*:)?sheetData\s*>")
_STAT_KEYS = ("cell_count", "max_row", "max_col", "ordered", "uncached_formulas", "value_cols")
# Excel の上限。これを超える位置にセルがあるファイルは Excel が書いたものではない（壊れている・手で作った）
EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_COLUMNS = 16_384
_OPEN_ERROR = "Excelファイルとして開けませんでした。.xlsx 形式で保存し直してください"


def _too_many_cells(count: int, limit: int) -> UploadError:
    return UploadError(
        f"セル数が上限（{limit:,}セル）を超えています（約{count:,}セル）。"
        "CSVで保存してから取り込んでください（書式による判定は使えなくなります）"
    )


def estimate_cell_count(path) -> int:
    """シートXMLの <dimension> から、読み込み前にセル数の目安を出す。"""
    total = 0
    try:
        with zipfile.ZipFile(path) as zf:
            for name in zf.namelist():
                if not (name.startswith("xl/worksheets/sheet") and name.endswith(".xml")):
                    continue
                with zf.open(name) as f:
                    head = f.read(4096)
                m = _DIMENSION_RE.search(head)
                if not m:
                    continue
                ref = m.group(1).decode()
                if ":" not in ref:
                    total += 1
                    continue
                min_col, min_row, max_col, max_row = range_boundaries(ref)
                # <dimension> は書式だけのセル（遠くの XFD1048576 など）まで広がるので、シートの XML の大きさで抑える。
                # <c> 1個は少なくとも約8バイトなので、これより多いセルはありえない（実際の件数は開いたあと数える）
                total += min((max_row - min_row + 1) * (max_col - min_col + 1), zf.getinfo(name).file_size // 8)
    except (zipfile.BadZipFile, OSError, ValueError):
        return 0   # zip でない・<dimension> が読めない: 目安なしで進める（開くときに案内する）
    except (zlib.error, EOFError) as exc:
        # 圧縮データが壊れている（zipfile は zlib.error / EOFError を包まずに投げる）
        raise UploadError("Excelファイルとして読み込めません（ファイルが壊れている可能性があります）") from exc
    return total


# ---- シートの XML: sheetData の外（結合・列・リンクなど） ----

def _split_sheet_xml(f) -> bytes:
    """sheetData の中身を除いたシートの XML（行が大きくても前後だけを持つ）。"""
    buf = b""
    while True:
        chunk = f.read(_CHUNK)
        buf += chunk
        m = _SHEETDATA_START.search(buf)
        if m:
            break
        if not chunk:
            return buf  # sheetData がない
    before, prefix = buf[:m.start()], m.group(1)
    empty = b"<" + prefix + b"sheetData/>"
    if m.group(2):  # <sheetData/>
        return before + empty + buf[m.end():] + f.read()
    rest = buf[m.end():]
    while True:
        e = _SHEETDATA_END.search(rest)
        if e:
            return before + empty + rest[e.end():] + f.read()
        chunk = f.read(_CHUNK)
        if not chunk:
            return before + empty  # 閉じタグがない（openpyxl でも読めない形）
        rest = rest[-64:] + chunk


def _sheet_extras(wb, zf: zipfile.ZipFile, ws_ro, shared_strings, table_names: set[str]) -> dict:
    """sheetData 以外を通常モードと同じ手順で読む。_cells には結合範囲・リンク・コメントで作られたセルだけが入る。"""
    member = ws_ro._worksheet_path
    with zf.open(member) as f:
        stripped = _split_sheet_xml(f)
    ws = Worksheet(wb, ws_ro.title)
    rels_path = get_rels_path(member)
    ws._rels = get_dependents(zf, rels_path) if rels_path in zf.namelist() else RelationshipList()
    reader = WorksheetReader(ws, io.BytesIO(stripped), shared_strings, True, False)
    reader.bind_all()
    for r in ws._rels.find(COMMENTS_NS):
        comment_sheet = CommentSheet.from_tree(fromstring(zf.read(r.target)))
        for ref, comment in comment_sheet.comments:
            try:
                ws[ref].comment = comment
            except AttributeError:
                continue
    tables = []
    for target in reader.tables:
        table = Table.from_tree(fromstring(zf.read(target)))
        name = table.name.lower()
        if name in table_names or name in wb.defined_names:
            raise ValueError(f"Table with name {table.name} already exists")
        table_names.add(name)
        tables.append(table)
    anchors: dict[tuple[int, int], tuple[int, int]] = {}
    for rng in ws.merged_cells.ranges:
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                anchors[(r, c)] = (rng.min_row, rng.min_col)
    hidden_columns: list[int] = []
    for key, dim in ws.column_dimensions.items():
        if dim.hidden:
            lo = dim.min or column_index_from_string(key)
            # Excel の列は 16,384 列まで。手で作った max="20000000" で全部の列番号を並べない
            hi = min(dim.max or lo, EXCEL_MAX_COLUMNS)
            hidden_columns.extend(range(lo, hi + 1))
    by_name = {}
    for t in tables:
        by_name[t.name] = t  # TableList と同じく名前で上書き
    return {
        "extra_cells": dict(ws._cells), "anchors": anchors, "hidden_columns": sorted(set(hidden_columns)),
        "freeze_panes": ws.freeze_panes, "auto_filter": ws.auto_filter.ref or None,
        "table_ranges": [t.ref for t in by_name.values()],
    }


# ---- シートの XML: 値のある範囲とセル数（expat で1回なめる） ----

def _scan_cells(f, extra: dict, shared_strings) -> dict:
    """WorkSheetParser と同じ規則で、値のある最終行・最終列・セル数・行の並びを調べる。"""
    ROW, V, IS, F = f"{_MAIN_NS} row", f"{_MAIN_NS} v", f"{_MAIN_NS} is", f"{_MAIN_NS} f"
    depth = 0
    row_depth = cell_depth = is_depth = r_depth = 0
    row_no = col_no = 0
    last_row = 0
    ordered = True
    count = max_row = max_col = 0
    seen_extra: set = set()
    value_cols: set[int] = set()  # 値のある列（列数の上限の確認用。遠くのメモ1つは1列と数える）
    # セル1つぶんの状態
    cell_row = cell_col = 0
    cell_type = "n"
    v_state = 0  # 0: 未読 1: 読み中 2: 読んだ（findtext と同じく最初の v だけ）
    v_text = None
    is_found = False
    plain = None
    rich: list = []
    text: list[str] = []
    capture = 0  # 1: v 2: is/t 3: is/r/t
    has_formula = False
    # Excel で計算されていない（値が保存されていない）数式のセルの数。列番号（1始まり）→ 数
    uncached: dict[int, int] = {}

    def local(name: str) -> str:
        return name[name.rfind(" ") + 1:]

    def start(name, attrs):
        nonlocal depth, row_depth, cell_depth, is_depth, r_depth, row_no, col_no, last_row, ordered
        nonlocal cell_row, cell_col, cell_type, v_state, v_text, is_found, plain, rich, capture, has_formula
        depth += 1
        if cell_depth:
            if depth == cell_depth + 1:
                if name == F:
                    has_formula = True
                elif name == V and v_state == 0:
                    v_state, capture = 1, 1
                    text.clear()
                elif name == IS and not is_found:
                    is_found, is_depth = True, depth
            elif is_depth:
                if depth == is_depth + 1:
                    tag = local(name)
                    if tag == "t":
                        capture = 2
                        text.clear()
                    elif tag == "r":
                        r_depth = depth
                        rich.append(None)
                elif r_depth and depth == r_depth + 1 and local(name) == "t":
                    capture = 3
                    text.clear()
            return
        if row_depth:
            if depth == row_depth + 1:
                cell_depth = depth
                ref = attrs.get("r")
                if ref:
                    cell_row, cell_col = coordinate_to_tuple(ref)
                    col_no = cell_col
                    if cell_row != row_no:
                        ordered = False
                else:
                    col_no += 1
                    cell_row, cell_col = row_no, col_no
                cell_type = attrs.get("t", "n")
                v_state, v_text, is_found, plain, capture = 0, None, False, None, 0
                has_formula = False
                rich = []
            return
        if name == ROW:
            row_depth = depth
            ref = attrs.get("r")
            if ref:
                try:
                    row_no = int(ref)
                except ValueError:
                    row_no = int(float(ref))
            else:
                row_no += 1
            col_no = 0
            if row_no <= last_row:
                ordered = False
            last_row = row_no

    def end(name):
        nonlocal depth, row_depth, cell_depth, is_depth, r_depth, count, max_row, max_col
        nonlocal v_state, v_text, plain, capture
        if capture and depth == (cell_depth + 1 if capture == 1 else is_depth + 1 if capture == 2 else r_depth + 1):
            value = "".join(text) if text else None  # ElementTree の .text と同じく、文字がなければ None
            if capture == 1:
                v_state, v_text = 2, value
            elif capture == 2:
                plain = value
            else:
                rich[-1] = value
            capture = 0
        elif r_depth and depth == r_depth:
            r_depth = 0
        elif is_depth and depth == is_depth:
            is_depth = 0
        elif cell_depth and depth == cell_depth:
            count += 1
            # 値のない数式（プログラムが書いて Excel で保存していないファイル）。
            # 空文字列を返す数式は Excel が t="str" と空の v で保存するので数えない
            if has_formula and not (v_state == 2 and (v_text or cell_type == "str")):
                uncached[cell_col] = uncached.get(cell_col, 0) + 1
            if cell_type == "inlineStr":
                value = (plain or "") + "".join(t for t in rich if t) if is_found else None
            else:
                value = (v_text or None) if v_state == 2 else None
                if value is not None and cell_type == "s":
                    value = shared_strings[int(value)]
            key = (cell_row, cell_col)
            m = extra.get(key)
            if m is not None:
                seen_extra.add(key)
                if isinstance(m, MergedCell):
                    value = None
                elif value is None:
                    value = m._value
            if value is not None and value != "":
                value_cols.add(cell_col)
                if cell_row > max_row:
                    max_row = cell_row
                if cell_col > max_col:
                    max_col = cell_col
            cell_depth = 0
        elif row_depth and depth == row_depth:
            row_depth = 0
        depth -= 1

    def chars(data):
        if capture:
            text.append(data)

    parser = expat.ParserCreate(namespace_separator=" ")
    parser.buffer_text = True
    parser.StartElementHandler = start
    parser.EndElementHandler = end
    parser.CharacterDataHandler = chars
    parser.ParseFile(f)
    for key, m in extra.items():
        if key in seen_extra or isinstance(m, MergedCell) or m._value is None or m._value == "":
            continue
        max_row, max_col = max(max_row, key[0]), max(max_col, key[1])
        value_cols.add(key[1])
    return {"cell_count": count + len(extra) - len(seen_extra), "max_row": max_row, "max_col": max_col, "ordered": ordered,
            "uncached_formulas": {str(c): n for c, n in sorted(uncached.items())}, "value_cols": len(value_cols)}


# ---- シートの XML: セルの値（WorkSheetParser） ----

def _inline_text(element) -> str:
    """Text.from_tree(element).content と同じ文字列（直下の t は最後のもの＋各 r の t。ふりがな rPh は除く）。"""
    plain = None
    runs = []
    for el in element:
        tag = el.tag.rpartition("}")[2]
        if tag == "t":
            plain = el.text
        elif tag == "r":
            text = None
            for sub in el:
                if sub.tag.rpartition("}")[2] == "t":
                    text = sub.text
            if text is not None:
                runs.append(text)
    return (plain or "") + "".join(runs)


class _SheetParser(WorkSheetParser):
    """インライン文字列だけを速く読む WorkSheetParser（値は同じ。Serialisable を作らない）。"""

    def parse_cell(self, element):
        if element.get("t") != "inlineStr" or not self.data_only:
            return super().parse_cell(element)
        coordinate = element.get("r")
        style_id = element.get("s", 0)
        if style_id:
            style_id = int(style_id)
        if coordinate:
            row, column = coordinate_to_tuple(coordinate)
            self.col_counter = column
        else:
            self.col_counter += 1
            row, column = self.row_counter, self.col_counter
        value, data_type = None, "inlineStr"
        child = element.find(INLINE_STRING)
        if child is not None:
            value, data_type = _inline_text(child), "s"
        return {"row": row, "column": column, "value": value, "data_type": data_type, "style_id": style_id}


def _row_hidden(attrs: dict | None) -> bool:
    """RowDimension(hidden=...) と同じ変換（'false' 'f' '0' と空は False）。"""
    value = (attrs or {}).get("hidden")
    return bool(value) and value not in ("false", "f", "0")


class ExcelSource:
    kind = "excel"

    def __init__(self, path, file_name: str | None = None, options: dict | None = None):
        self.path = Path(path)
        self.file_name = file_name or self.path.name
        options = dict(options or {})
        self.max_cells = int(options.get("max_cells") or DEFAULT_MAX_CELLS)
        # 同じファイルを前に開いたときの sheet_stats()（取り込みの控え）。あれば値のある範囲を数え直さない
        known_stats = options.get("sheet_stats") or {}
        self._shared_strings = []

        # 範囲情報（dimension）は書式だけの行で大きく出ることがあるので、目安の4倍を超えるときだけ先に止める
        estimated = estimate_cell_count(self.path)
        if estimated > self.max_cells * 4:
            raise _too_many_cells(estimated, self.max_cells)
        wb = None
        try:
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                wb = load_workbook(self.path, read_only=True, data_only=True)
                self._sheets = self._read_sheets(wb, known_stats)
        except UploadError:
            raise
        except Exception as e:  # openpyxl・XML は形式ごとに様々な例外を出す
            raise UploadError(_OPEN_ERROR) from e
        finally:
            if wb is not None:
                wb.close()

        actual = sum(s["cell_count"] for s in self._sheets.values())
        if actual > self.max_cells:
            raise _too_many_cells(actual, self.max_cells)
        self.date1904 = wb.epoch == CALENDAR_MAC_1904
        self._epoch = wb.epoch
        self._date_formats = wb._date_formats
        self._timedelta_formats = wb._timedelta_formats
        self._styles: dict[int, tuple] = {}
        self._wb_styles = (wb._cell_styles, wb._fonts, wb._fills, wb._number_formats)
        self._default_style = self._style_tuple(None)

    def _read_sheets(self, wb, known_stats: dict) -> dict[str, dict]:
        sheets: dict[str, dict] = {}
        table_names: set[str] = set()
        zf = wb._archive
        use_known = (bool(known_stats) and set(known_stats) == set(wb.sheetnames)
                     and all(isinstance(v, dict) and set(_STAT_KEYS) <= set(v) for v in known_stats.values()))
        for ws in wb.worksheets:
            self._shared_strings = ws._shared_strings
            extras = _sheet_extras(wb, zf, ws, ws._shared_strings, table_names)
            if use_known:
                scan = {k: known_stats[ws.title][k] for k in _STAT_KEYS}
            else:
                with zf.open(ws._worksheet_path) as f:
                    scan = _scan_cells(f, extras["extra_cells"], ws._shared_strings)
            if int(scan["max_row"]) > EXCEL_MAX_ROWS or int(scan["max_col"]) > EXCEL_MAX_COLUMNS:
                # 行番号だけを大きくしたセル（D5000000 など）。rows() が空の行を何百万も作って止まったようになる
                raise UploadError(f"シート「{ws.title}」に Excel の上限（{EXCEL_MAX_ROWS:,}行・{EXCEL_MAX_COLUMNS:,}列）を"
                                  "超える位置のセルがあります。ファイルが壊れている可能性があります。"
                                  "Excel で開いて保存し直してから取り込んでください")
            sheets[ws.title] = {**extras, **scan, "member": ws._worksheet_path, "state": ws.sheet_state}
        return sheets

    def sheet_stats(self) -> dict[str, dict]:
        """シートごとの値のある範囲・セル数・行の並び（JSON にできる形。次に開くときの options["sheet_stats"]）。"""
        return {name: {k: meta[k] for k in _STAT_KEYS} for name, meta in self._sheets.items()}

    def uncached_formulas(self, sheet: str) -> dict[int, int]:
        """Excel で計算されていない数式のセルの数（列番号1始まり → 数）。"""
        return {int(c): n for c, n in (self._meta(sheet).get("uncached_formulas") or {}).items()}

    def close(self) -> None:
        pass

    def _style_tuple(self, style_array) -> tuple:
        """(表示形式, 太字, 取り消し線, 塗り)。style_array が None なら既定（全部0）のスタイル。"""
        _cell_styles, fonts, fills, number_formats = self._wb_styles
        num_id = style_array.numFmtId if style_array is not None else 0
        font = fonts[style_array.fontId if style_array is not None else 0]
        fill = fills[style_array.fillId if style_array is not None else 0]
        if num_id < BUILTIN_FORMATS_MAX_SIZE:
            number_format = BUILTIN_FORMATS.get(num_id, "General")
        else:
            number_format = number_formats[num_id - BUILTIN_FORMATS_MAX_SIZE]
        return (number_format, bool(font is not None and font.b), bool(font is not None and font.strike),
                bool(fill is not None and fill.fill_type not in (None, "none")))

    def _style(self, style_id) -> tuple:
        found = self._styles.get(style_id)
        if found is None:
            found = self._styles[style_id] = self._style_tuple(self._wb_styles[0][style_id])
        return found

    def sheets(self) -> list[SheetInfo]:
        return [SheetInfo(
            name=name,
            hidden=meta["state"] != "visible",
            max_row=meta["max_row"],
            max_col=meta["max_col"],
            table_ranges=list(meta["table_ranges"]),
            date1904=self.date1904,
            hidden_columns=list(meta["hidden_columns"]),
            freeze_panes=meta["freeze_panes"],
            auto_filter=meta["auto_filter"],
        ) for name, meta in self._sheets.items()]

    def _meta(self, sheet: str) -> dict:
        if sheet not in self._sheets:
            raise UploadError(f"シート「{sheet}」が見つかりません")
        return self._sheets[sheet]

    def _parsed_rows(self, zf: zipfile.ZipFile, meta: dict, end: int) -> Iterator[tuple[int, list[dict], dict | None]]:
        """(行番号, セルの dict, 行の属性) を行番号の昇順に返す。end より後は読まない。"""
        with zf.open(meta["member"]) as f:
            parser = _SheetParser(f, self._shared_strings, data_only=True, epoch=self._epoch,
                                     date_formats=self._date_formats, timedelta_formats=self._timedelta_formats)
            if meta["ordered"]:
                for idx, cells in parser.parse():
                    if idx > end:
                        return
                    yield idx, cells, parser.row_dimensions.pop(str(idx), None)
                return
            # 行が昇順でないシート: 全行を読んでから並べる（通常モードと同じく後のセルで上書き）
            rows: dict[int, dict[int, dict]] = {}
            for _idx, cells in parser.parse():
                for d in cells:
                    rows.setdefault(d["row"], {})[d["column"]] = d
            for idx in sorted(rows):
                if idx > end:
                    return
                yield idx, list(rows[idx].values()), parser.row_dimensions.get(str(idx))

    def rows(self, sheet: str, start: int = 1, limit: int | None = None) -> Iterator[SourceRow]:
        meta = self._meta(sheet)
        if int(meta.get("value_cols") or 0) > MAX_COLUMNS:
            # CSV と同じ列数の上限。見出しごとに判定・辞書引きをするので、列が桁違いに多いと画面が何十秒も止まる
            raise UploadError(f"シート「{sheet}」の列数が上限（{MAX_COLUMNS:,}列）を超えています"
                              f"（値のある列が{int(meta['value_cols']):,}列）。不要な列を削除して保存し直してください")
        anchors = meta["anchors"]
        extra = meta["extra_cells"]
        extra_rows = {r for r, _c in extra}
        anchor_rows = {r for r, _c in anchors}
        max_col = meta["max_col"]
        begin = max(1, start)
        end = meta["max_row"] if limit is None else min(meta["max_row"], start + limit - 1)
        if end < begin:
            return
        default = self._default_style
        style_of = self._style
        # 列が多いシート（遠くの列に値が1つあるだけのことが多い）は、行を最後の列まで埋めない
        # （CSV と同じく行ごとに長さが違う。読む側は範囲外を空として扱う）
        # 遠くの行に値が1つあるだけのシート（行×列がセル数の上限を超える＝ほとんどが空）も同じく埋めない
        trim = max_col > PAD_MAX_COLUMNS or meta["max_row"] * max_col > self.max_cells
        anchor_last: dict[int, int] = {}
        extra_last: dict[int, int] = {}
        if trim:
            for (ar, ac) in anchors:
                if ac <= max_col and ac > anchor_last.get(ar, 0):
                    anchor_last[ar] = ac
            for (er, ec) in extra:
                if ec <= max_col and ec > extra_last.get(er, 0):
                    extra_last[er] = ec
        done = -1  # 読み終えた行（-1: まだ行を読み始めていない）
        try:
            with zipfile.ZipFile(self.path) as zf:
                done = 0
                parsed = self._parsed_rows(zf, meta, end)
                pending = next(parsed, None)
                for r in range(1, end + 1):
                    while pending is not None and pending[0] < r:
                        pending = next(parsed, None)
                    if pending is not None and pending[0] == r:
                        _idx, xml_cells, attrs = pending
                        pending = next(parsed, None)
                    else:
                        xml_cells, attrs = (), None
                    if r < begin:
                        continue
                    by_col = {}
                    for d in xml_cells:
                        if d["column"] <= max_col:
                            by_col[d["column"]] = d
                    has_extra = r in extra_rows
                    has_anchor = r in anchor_rows
                    width = max_col
                    if trim:
                        # 遠くの列に値が1つだけある表でも、行ごとに全列を作らない（その行の最後のセルまで）
                        width = max([d["column"] for d in by_col.values()] + [0])
                        if has_anchor:
                            width = max(width, anchor_last[r])
                        if has_extra:
                            width = max(width, extra_last[r])
                    row_cells = []
                    for c in range(1, width + 1):
                        anchor = anchors.get((r, c)) if has_anchor else None
                        m = extra.get((r, c)) if has_extra else None
                        d = by_col.get(c)
                        if m is not None and (isinstance(m, MergedCell) or d is None):
                            # 結合範囲の左上以外（値は消える）や、リンク・コメントで作られたセルは既定のスタイル
                            value = None if isinstance(m, MergedCell) else m._value
                            number_format, bold, strike, fill = default
                        elif d is not None:
                            value = d["value"]
                            if value is None and m is not None:
                                value = m._value
                            number_format, bold, strike, fill = style_of(d["style_id"])
                        else:
                            row_cells.append(CellInfo(value=None, text="", merged_anchor=anchor))
                            continue
                        row_cells.append(CellInfo(value=value, text=cell_text(value), number_format=number_format,
                                                  bold=bold, strike=strike, fill=fill, merged_anchor=anchor))
                    done = r
                    yield SourceRow(index=r, cells=row_cells, hidden=_row_hidden(attrs))
        except (UploadError, GeneratorExit):
            raise
        except Exception as e:
            if done >= 0:
                # 途中の行の値（<v>NaN</v> など Excel 以外のソフトが書いた値）が読めない。場所を示す
                raise UploadError(f"{max(done, begin - 1) + 1}行目付近の値を読めません（NaN など）。"
                                  "Excelで開いて値を直し、保存し直してください") from e
            raise UploadError(_OPEN_ERROR) from e


# ====================================================================================================
# 元 tables/detect.py
# 見出し帯・データ範囲・行の分類・表の種類の判定。
#
# - 見出し行: アンカー語 → 自動の点数付け（文字列の割合・辞書の語・下の行の型がそろっているか など）
# - 2段見出し: 横に結合された上段を右へ埋めて「上_下」でつなぐ（CSVは上段の空欄を右へ埋める）
# - データの終わり: 空行3行、合計行（含む）、注記行・表題行・見出しの再出現（含まない）
# - MVPでは1シート1表。下や右に別の表があれば警告する
# ====================================================================================================

HEAD_ROWS = 80  # 見出し探索のために読む行数
HEADER_SEARCH_ROWS = 40  # 見出し行の候補にする先頭行数
PREVIEW_ROWS = 60  # row_classes に全行を載せる先頭行数
MAX_ROW_CLASSES = 500
BLANK_ROWS_END = 3
LOOKAHEAD_ROWS = 300
KEY_SAMPLE_ROWS = 50

_BLOCK_TITLE_RE = re.compile(r"^\s*[■□◆◇●○▼▽★☆]")
# 空行の後に1つだけ書かれた「以上」「作成者：…」などの書き添え（表の終わりの印）
_FOOTER_RE = re.compile(r"^(?:以上[。.]?|(?:作成者|作成日|作成|記入者|承認者?|確認者|出典|備考)\s*[:：\s].*|(?:作成者|作成日|承認|出典|備考)[:：]?)$")
_NOTE_RE = re.compile(r"^\s*(※|\*\d|（注|\(注|注[記意]?\s*[\d０-９]*\s*[:：)）.．、\s])")
# 集計の語はラベル全体（前置き＋語＋短い後置き）として見る。「合計カウンタ不良」のような文の一部は合計行にしない
_TOTAL_RE = re.compile(r"^.{0,15}?(総合計|合計|総計|累計)(件数|数|額|金額|時間|値)?$")
# 「計」だけの語は「設計」「会計」と区別するため、単独か、区切り・数字・英字・月/年/度/期/週/日の後ろだけ
_SUBTOTAL_RE = re.compile(r"(小計|平均)$|(^|[\s・_:：/／\-－0-9A-Za-z月年度期週日])計$")
_PAREN_TAIL_RE = re.compile(r"\s*[（(][^()（）]*[)）]\s*$")
_UNIT_RE = re.compile(r"^(.*?)\s*[（(\[［]\s*([^()（）\[\]［］]{1,12})\s*[)）\]］]\s*$")
_KNOWN_UNITS = {
    "分", "秒", "時間", "時", "日", "日数", "h", "hr", "hrs", "H", "min", "sec", "s", "円", "千円", "万円", "百万円",
    "個", "件", "回", "本", "枚", "台", "式", "人", "名", "箇所", "ヶ所", "%", "％", "mm", "cm", "m", "km", "μm", "um",
    "nm", "g", "kg", "t", "l", "L", "ml", "mL", "kPa", "Pa", "MPa", "V", "A", "W", "kW", "kWh", "℃", "°C", "pcs",
    "ppm", "rpm", "L/min", "sccm", "slm", "Torr", "件数", "人時", "工数", "h/人",
}
_MONTH_RES = [
    re.compile(r"^(\d{1,2})月(分)?$"),
    re.compile(r"^(\d{4})[/\-.年](\d{1,2})月?$"),
    re.compile(r"^(R|H|令和|平成)(\d{1,2}|元)[./年](\d{1,2})月?$"),
    re.compile(r"^(\d{2})/(\d{1,2})$"),
    re.compile(r"^(\d{4})-(\d{2})-01$"),  # 月初日の日付セルを見出しにした表
    re.compile(r"^(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?$", re.I),
]


@dataclass
class RowClass:
    index: int
    kind: str  # title/header/data/continuation/subtotal/note/blank/excluded
    reason: str = ""


@dataclass
class LayoutGuess:
    sheet: str
    table_kind: str  # list/crosstab/form_like/unknown
    header_rows: list[int]
    data_start: int
    data_end: int
    headers: list[str]  # 2段は "上_下" で結合、結合セルは右へ埋める
    row_classes: list[RowClass]  # 先頭60行＋データ範囲の要約
    confidence: float
    warnings: list[str]
    header_levels: list[list[str]] = field(default_factory=list)  # 見出しの段ごとの元の文字列（表示用）
    key_columns: list[int] = field(default_factory=list)  # 継続行の判定に使う列（0始まり）
    counts: dict[str, int] = field(default_factory=dict)  # データ範囲内の行の種類ごとの件数


# ---- 公開関数 ----

def split_header_unit(header: str) -> tuple[str, str]:
    """「停止時間(分)」→ ("停止時間", "分")。単位らしくない括弧は分けない。"""
    text = str(header or "").strip()
    m = _UNIT_RE.match(text)
    if not m or not m.group(1).strip():
        return text, ""
    unit = unicodedata.normalize("NFKC", m.group(2)).strip()
    if unit in _KNOWN_UNITS or unit.lower() in {u.lower() for u in _KNOWN_UNITS if u.isascii()}:
        return m.group(1).strip(), unit
    return text, ""


def is_month_label(text: str) -> bool:
    """見出しが月・年月として読めるか（「4月」「2026/04」「R8.4」「2025年_4月」など）。"""
    s = unicodedata.normalize("NFKC", str(text or "")).strip().replace(" ", "")
    if not s:
        return False
    last = s.split("_")[-1]
    return any(r.match(last) for r in _MONTH_RES) or any(r.match(s) for r in _MONTH_RES)


def guess_layout(source, sheet, anchors: list[str] | None = None, header_row: int | None = None,
                 data_end: int | None = None, header_rows: list[int] | None = None,
                 max_scan_rows: int | None = None, scan_memo: dict | None = None) -> LayoutGuess:
    """表の見出し帯・データ範囲・種類を推定する。

    header_rows を渡すとそのまま使う。header_row だけなら2段見出しかを自動で確かめる。
    max_scan_rows を渡すと、データの終わりの探索をその行数で打ち切る（簡易判定用）。
    scan_memo（dict）を渡すと、全行をなめる部分（キー列・行の分類・データの終わり）の結果をそこに控え、
    同じ見出し帯・同じ条件ならなめ直さない（自動判定のあとに同じ見出し行を指定し直したときなど）。値は JSON にできる形。
    """
    head = list(source.rows(sheet, 1, HEAD_ROWS))
    warnings: list[str] = []
    by_index = {r.index: r for r in head}
    wanted = sorted(header_rows) if header_rows else ([header_row] if header_row else [])
    if wanted and wanted[-1] > HEAD_ROWS:
        # 手で指定した見出し行が先頭の読み取り範囲（HEAD_ROWS 行）より下にあるときは、
        # その行と前後（2段見出しの確認・表の幅の確認に使う下の行）も読む
        start = max(1, wanted[0] - 1)
        for r in source.rows(sheet, start, wanted[-1] - start + 33):
            by_index.setdefault(r.index, r)
    if not any(not r.is_blank for r in by_index.values()):
        return LayoutGuess(sheet, "unknown", [], 1, 0, [], [RowClass(r.index, "blank") for r in head[:PREVIEW_ROWS]],
                           0.0, ["表が見つかりません（先頭に値のある行がありません）"])
    is_csv = getattr(source, "kind", "") == "csv"

    score = 1.0
    if header_rows:
        rows_h = sorted(header_rows)
    else:
        if header_row:
            best = header_row
        else:
            best, score = _choose_header_row(head, anchors)
            if best is None:
                return _no_table(sheet, head, "見出し行が見つかりません（2列以上に文字が並ぶ行がありません）")
        rows_h = _header_band(by_index, best, is_csv)

    levels, headers = _build_headers(by_index, rows_h, is_csv)
    width, split_warning = _table_width(by_index, rows_h, levels)
    if split_warning:
        warnings.append(split_warning)
    header_is_data = _header_looks_like_data(by_index.get(rows_h[-1]), width)
    headers = _dedupe((headers + [""] * width)[:width])
    levels = [(lv + [""] * width)[:width] for lv in levels]
    data_start = rows_h[-1] + 1

    header_norms = {_norm(h) for h in headers if h and not h.startswith("列")}
    ctx = _Ctx(width=width, header_norms=header_norms, key_cols=[], is_csv=is_csv, header_rows=rows_h,
               left=_left_edge(headers))
    memo_key = json.dumps([sheet, rows_h, headers, width, is_csv, data_end, max_scan_rows], ensure_ascii=False)
    found = scan_memo.get(memo_key) if scan_memo is not None else None
    if found is not None:
        ctx.key_cols = list(found["key_cols"])
        classes = [RowClass(i, k, r) for i, k, r in found["classes"]]
        counts, end, scan_warnings = dict(found["counts"]), found["end"], list(found["warnings"])
    else:
        ctx.key_cols = _auto_key_columns(source, sheet, data_start, ctx)
        classes, counts, end, scan_warnings = _scan(
            source, sheet, head, ctx, data_start, data_end, max_scan_rows)
        if scan_memo is not None:
            scan_memo[memo_key] = {"key_cols": list(ctx.key_cols), "classes": [[c.index, c.kind, c.reason] for c in classes],
                                   "counts": dict(counts), "end": end, "warnings": list(scan_warnings)}
    warnings.extend(scan_warnings)

    months = sum(1 for h in headers if is_month_label(h))
    data_rows = counts.get("data", 0)
    if months >= 3 and months >= 0.3 * max(1, len([h for h in headers if not h.startswith("列")])):
        table_kind = "crosstab"
    elif _labels_down_first_column(head, rows_h, data_start, end):
        table_kind = "form_like"
    elif data_rows >= 3 and score >= 0.4:
        table_kind = "list"
    elif _form_like(head):
        table_kind = "form_like"
    else:
        table_kind = "unknown"
    if header_is_data and table_kind != "crosstab":
        # 見出し行のないCSVでは1行目の値が見出し（＝取り込み設定に残る列名）になり、その記録も md に出ない
        warnings.append("見出し行がデータのように見えます（日付・数値や長い文章が並んでいます）。"
                        "見出し行のない表には対応していません。1行目に列名を入れてから取り込み直してください")
    if table_kind in ("form_like", "unknown"):
        warnings.append("一覧表の形に見えません（見出しの下に同じ形の行が続いていません）")

    hidden_cols = _hidden_columns(source, sheet, width)
    if hidden_cols:
        warnings.append(f"非表示の列があります（{', '.join(hidden_cols)}列）")

    confidence = round(max(0.0, min(1.0, score)) * min(1.0, data_rows / 5), 3)
    return LayoutGuess(
        sheet=sheet, table_kind=table_kind, header_rows=rows_h, data_start=data_start, data_end=end,
        headers=headers, row_classes=classes, confidence=confidence, warnings=warnings,
        header_levels=levels, key_columns=ctx.key_cols, counts=counts,
    )


def classify_rows(source, sheet, layout: LayoutGuess, key_columns: list[int] | None = None
                  ) -> Iterator[tuple[SourceRow, RowClass]]:
    """データ範囲（data_start〜data_end）の行を分類しながら返す。"""
    ctx = _ctx_from_layout(layout)
    ctx.is_csv = getattr(source, "kind", "") == "csv"
    if key_columns is not None:
        ctx.key_cols = list(key_columns)
    if layout.data_end < layout.data_start:
        return
    limit = layout.data_end - layout.data_start + 1
    after_blank = False
    for row in source.rows(sheet, layout.data_start, limit):
        rc = _classify(row, ctx, after_blank)
        after_blank = rc.kind == "blank"
        yield row, rc


def sample_data_rows(source, sheet, layout: LayoutGuess, n: int = 200) -> list[SourceRow]:
    """列の対応づけ候補に使う、先頭のデータ行（継続行・小計などを除く）。"""
    out: list[SourceRow] = []
    for row, rc in classify_rows(source, sheet, layout):
        if rc.kind == "data":
            out.append(row)
            if len(out) >= n:
                break
    return out


def kind_from_layout(layout) -> str:
    """表の形の見立てから「行が並ぶ表か」を決める（控えを使う ImportSource.list_kind と共通）。"""
    if layout.table_kind in ("list", "crosstab") and layout.counts.get("data", 0) >= 10:
        return layout.table_kind
    return ""


def list_kind(source, sheet) -> str:
    """行が並ぶ表か（見出し行の下に同じ形の行が10行以上）。"list" / "crosstab" / ""（どちらでもない）。"""
    try:
        layout = guess_layout(source, sheet, max_scan_rows=200)
    except Exception:
        return ""
    return kind_from_layout(layout)


# ---- 内部: 行の事実 ----

@dataclass
class _Ctx:
    width: int
    header_norms: set[str]
    key_cols: list[int]
    is_csv: bool
    header_rows: list[int]
    left: int = 0  # 表の左端の列（0始まり）。注記・表題はこの列から始まる行だけにする


def _left_edge(headers: list[str]) -> int:
    """見出しのある最初の列（空の見出しは「列N」になっている）。"""
    return next((i for i, h in enumerate(headers) if h and not re.fullmatch(r"列\d+", h)), 0)


def _ctx_from_layout(layout: LayoutGuess) -> _Ctx:
    width = len(layout.headers)
    norms = {_norm(h) for h in layout.headers if h and not h.startswith("列")}
    return _Ctx(width=width, header_norms=norms, key_cols=list(layout.key_columns),
                is_csv=False, header_rows=list(layout.header_rows), left=_left_edge(layout.headers))


def _norm(text: str) -> str:
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", str(text or ""))).lower()


def _is_header_cell(cell) -> bool:
    if not cell.text:
        return False
    return isinstance(cell.value, str) and value_kind(cell.value, cell.text) in ("string", "code") \
        and len(cell.text) <= 40 and "\n" not in cell.text


def _headerish(row: SourceRow) -> bool:
    cells = [c for c in row.cells if c.text]
    if len(cells) < 2:
        return False
    return sum(1 for c in cells if _is_header_cell(c)) / len(cells) >= 0.8


def _last_col(row: SourceRow) -> int:
    for i in range(len(row.cells) - 1, -1, -1):
        c = row.cells[i]
        if c.text or (c.merged_anchor and c.merged_anchor[0] == row.index and c.merged_anchor[1] <= i + 1
                      and row.cells[c.merged_anchor[1] - 1].text):
            return i + 1
    return 0


def _covered_positions(row: SourceRow) -> int:
    """値のあるセル＋同じ行で横に結合された範囲の数。"""
    n = 0
    for i, c in enumerate(row.cells):
        if c.text:
            n += 1
        elif c.merged_anchor and c.merged_anchor[0] == row.index and c.merged_anchor[1] - 1 < len(row.cells) \
                and row.cells[c.merged_anchor[1] - 1].text:
            n += 1
    return n


# ---- 内部: 見出し行の選択 ----

def _choose_header_row(head: list[SourceRow], anchors: list[str] | None) -> tuple[int | None, float]:

    scores: dict[int, float] = {}
    candidates = head[:HEADER_SEARCH_ROWS]
    for pos, row in enumerate(candidates):
        cells = [c for c in row.cells if c.text]
        if len(cells) < 2:
            continue
        texts = [c.text for c in cells]
        hdr_ratio = sum(1 for c in cells if _is_header_cell(c)) / len(cells)
        unique = len(set(texts)) / len(texts)
        below = [r for r in head[pos + 1:pos + 13] if not r.is_blank]
        width_below = max([_last_col(r) for r in below] or [0])
        coverage = min(1.0, _covered_positions(row) / max(1, _last_col(row), width_below))
        if below:
            consist = sum(1 for r in below if sum(1 for c in r.cells if c.text) >= max(2, 0.3 * len(cells))) / len(below)
            contrast = sum(1 for r in below if not _headerish(r)) / len(below)
        else:
            consist = contrast = 0.0
        dict_hits = 0
        for t in texts:
            if len(t) <= 20:
                found = lookup_header(t)
                if found and found[1] == "dictionary":
                    dict_hits += 1
        bold = sum(1 for c in cells if c.bold) / len(cells)
        score = (0.3 * hdr_ratio + 0.1 * unique + 0.15 * coverage + 0.15 * consist + 0.15 * contrast
                 + 0.1 * min(1.0, dict_hits / 3) + 0.05 * bold)
        if _BLOCK_TITLE_RE.match(texts[0]) or _NOTE_RE.match(texts[0]):
            score -= 0.3
        if len(cells) == 2 and len(texts[0]) <= 10 and not _is_header_cell(cells[1]):
            score -= 0.2  # 「出力日時, 2026/09/14」のような前置き行
        # 上にも同じ形の行が並んでいるなら、見出しではなくデータの途中
        above = [r for r in head[max(0, pos - 10):pos] if not r.is_blank]
        if above:
            same_shape = sum(1 for r in above if sum(1 for c in r.cells if c.text) >= 0.6 * len(cells))
            score -= 0.3 * same_shape / max(len(above), 3)
        scores[row.index] = score

    if not scores:
        return None, 0.0
    if anchors:
        anchor_norms = [_norm(a) for a in anchors if a]
        best_hits = 0
        best_row = None
        for row in candidates:
            norms = [_norm(c.text) for c in row.cells if c.text]
            hits = sum(1 for a in anchor_norms if any(a and a in n for n in norms))
            if hits > best_hits or (hits == best_hits and hits and best_row is not None
                                    and scores.get(row.index, 0) > scores.get(best_row, 0)):
                best_hits, best_row = hits, row.index
        if best_row is not None:
            return best_row, max(scores.get(best_row, 0.5), 0.5 + 0.5 * best_hits / max(1, len(anchor_norms)))
    best = max(scores, key=lambda i: (round(scores[i], 3), -i))
    return best, scores[best]


def _header_band(by_index: dict[int, SourceRow], best: int, is_csv: bool) -> list[int]:
    """選んだ行と上下の行から、2段見出しかどうかを決める。"""
    row = by_index.get(best)
    above = by_index.get(best - 1)
    below = by_index.get(best + 1)
    after_below = by_index.get(best + 2)
    after = by_index.get(best + 1)
    if row is None:
        return [best]
    if above is not None and not above.is_blank and _pair_is_header(above, row, after, is_csv):
        return [best - 1, best]
    if below is not None and not below.is_blank and _pair_is_header(row, below, after_below, is_csv):
        return [best, best + 1]
    return [best]


def _pair_is_header(top: SourceRow, bottom: SourceRow, after: SourceRow | None, is_csv: bool) -> bool:
    if after is not None and not after.is_blank and _headerish(after):
        return False
    top_cells = [c for c in top.cells if c.text]
    bottom_cells = [c for c in bottom.cells if c.text]
    if len(top_cells) < 2 or len(bottom_cells) < 2:
        return False  # 上段が1セルだけならタイトル行とみなす
    if _BLOCK_TITLE_RE.match(top_cells[0].text) or _NOTE_RE.match(top_cells[0].text):
        return False
    if not all(_is_header_cell(c) for c in top_cells) or not _headerish(bottom):
        return False
    horizontal = any(
        c.merged_anchor and c.merged_anchor[0] == top.index and c.merged_anchor[1] != i + 1
        for i, c in enumerate(top.cells)
    )
    vertical = any(c.merged_anchor and c.merged_anchor[0] == top.index for c in bottom.cells)
    if horizontal or vertical:
        return True
    if is_csv:
        # 結合のないCSV: 上段がまばらで、上段の値の真下にも下段の値がある
        if len(top_cells) >= 2 and len(top_cells) <= 0.6 * len(bottom_cells):
            positions = [i for i, c in enumerate(top.cells) if c.text]
            if all(i < len(bottom.cells) and bottom.cells[i].text for i in positions):
                return True
        # 結合した2段見出しを書き出したCSV: 上段の見出しは横に広がる列の左端にだけあり、右隣の空欄の下に下段がある
        return _csv_spanned_pair(top, bottom)
    return False


def _csv_spanned_pair(top: SourceRow, bottom: SourceRow) -> bool:
    """「管理No,発生,,設備,,停止時間」の下に「,日付,時刻,番号,名称,」がある形か。"""
    def text(row: SourceRow, i: int) -> str:
        return row.cells[i].text if i < len(row.cells) else ""

    under_blank = 0
    for i, c in enumerate(bottom.cells):
        if not c.text:
            continue
        if not _is_header_cell(c):
            return False
        if not text(top, i):
            under_blank += 1
        elif text(top, i + 1) or not text(bottom, i + 1):
            return False  # 上段の見出しの右隣が空いていない（横に広がっていない）
    return under_blank >= 1


def _build_headers(by_index: dict[int, SourceRow], rows_h: list[int], is_csv: bool
                   ) -> tuple[list[list[str]], list[str]]:
    rows = [by_index.get(r) for r in rows_h]
    ncols = max((len(r.cells) for r in rows if r is not None), default=0)
    levels: list[list[str]] = []
    for level, row in enumerate(rows):
        values = [""] * ncols
        if row is None:
            levels.append(values)
            continue
        last_text = ""
        last_pos = -1
        for i in range(ncols):
            cell = row.cells[i] if i < len(row.cells) else None
            text = _clean_header(cell.text) if cell else ""
            anchor = cell.merged_anchor if cell else None
            if text:
                values[i] = text
                last_text, last_pos = text, i
                continue
            if anchor and anchor[0] == row.index and anchor[1] - 1 != i:
                src = row.cells[anchor[1] - 1] if anchor[1] - 1 < len(row.cells) else None
                values[i] = _clean_header(src.text) if src else ""  # 横の結合は右へ埋める
            elif anchor and anchor[0] != row.index:
                values[i] = ""  # 上の段からの縦の結合は重ねない
            elif is_csv and level < len(rows) - 1 and last_pos >= 0:
                lower = rows[-1]
                if lower is not None and i < len(lower.cells) and lower.cells[i].text:
                    values[i] = last_text  # CSVの2段見出しは上段の空欄を右へ埋める
        levels.append(values)
    headers = []
    for i in range(ncols):
        parts: list[str] = []
        for lv in levels:
            if lv[i] and (not parts or parts[-1] != lv[i]):
                parts.append(lv[i])
        headers.append("_".join(parts))
    return levels, headers


def _clean_header(text: str) -> str:
    return re.sub(r"[ \t]+", " ", (text or "").replace("\n", "")).strip()


def _dedupe(headers: list[str]) -> list[str]:
    seen: Counter[str] = Counter()
    out = []
    for i, h in enumerate(headers):
        name = h or f"列{i + 1}"
        seen[name] += 1
        # 「備考_2」だと2段見出しの「上_下」と区別できず、下段「2」が表示名になるので括弧で番号を付ける
        out.append(name if seen[name] == 1 else f"{name}({seen[name]})")
    return out


# 見出しの右にこれ以上の空の列を挟んで見出しが1つだけあれば、表とは別の書き込み（メモなど）とみなす
_FAR_HEADER_GAP = 20


def _table_width(by_index: dict[int, SourceRow], rows_h: list[int], levels: list[list[str]]) -> tuple[int, str]:
    """見出しの右端。空の見出し列を挟んで右に別の表があれば、左の表だけにする。"""
    ncols = len(levels[0]) if levels else 0
    filled = [any(lv[i] for lv in levels) for i in range(ncols)]
    if not any(filled):
        return 0, ""
    last = max(i for i, f in enumerate(filled) if f) + 1
    first = min(i for i, f in enumerate(filled) if f)
    data_rows = [by_index[r] for r in sorted(by_index) if r > rows_h[-1]][:30]
    data_rows = [r for r in data_rows if not r.is_blank]
    for gap in range(first + 1, last):
        if filled[gap]:
            continue
        right = sum(1 for i in range(gap + 1, last) if filled[i])
        if right == 1:
            # 大きく離れた列に見出しが1つだけ（XFD1 のメモなど）。下に値が無ければ表に含めない
            # （含めると16,384列の表になり、列の対応づけも読み込みも止まったようになる）
            far = last - 1
            used = sum(1 for r in data_rows if far < len(r.cells) and r.cells[far].text)
            if far - gap >= _FAR_HEADER_GAP and not (data_rows and used > 0.1 * len(data_rows)):
                from openpyxl.utils import get_column_letter

                # 列の範囲を指定する画面は無いので、Excel 側で直す手順を案内する
                return gap, (f"見出しから大きく離れた {get_column_letter(far + 1)}列 の値は表に含めません。"
                             "表の一部なら、Excel でその列を表の右隣に移してから取り込んでください")
        if right < 2:
            continue
        used = sum(1 for r in data_rows if gap < len(r.cells) and r.cells[gap].text)
        if data_rows and used > 0.1 * len(data_rows):
            continue
        from openpyxl.utils import get_column_letter

        # 列の範囲を指定する画面は無いので、Excel 側で分ける手順を案内する
        return gap, (f"見出しの右側（{get_column_letter(gap + 2)}列以降）に別の表があるようです。"
                     "最初の表だけを読み取ります。右側の表は、Excel で別のシートか別のファイルに分けてから取り込んでください")
    # 右端の見出しが空欄でも、その列に値が続いていれば表の列に含める（黙って捨てない。見出しは「列N」になる）
    while data_rows and sum(1 for r in data_rows if last < len(r.cells) and r.cells[last].text) > 0.1 * len(data_rows):
        last += 1
    return last, ""


def _header_looks_like_data(row: SourceRow | None, width: int) -> bool:
    """見出し行の値の半分以上が日付・数値か、40字を超える文章があるか。"""
    if row is None:
        return False
    filled = [c for c in (row.cells[:width] if width else row.cells) if c.text]
    if len(filled) < 2:
        return False
    if any(len(c.text) > 40 for c in filled):
        return True
    data_like = sum(1 for c in filled if value_kind(c.value, c.text) in ("number", "date", "datetime"))
    return data_like * 2 >= len(filled)


# ---- 内部: 行の分類 ----

def _aggregate_label(row: SourceRow, width: int, key_cols: list[int] | None = None) -> str | None:
    """小計・合計行なら "subtotal" / "total"。

    「ブロック計」「部署計」のように語の後ろに「計」だけが付くラベルは、「設計」「会計」「流量計」と
    見分けられないので、ほかのキー列（日付・番号。集計値の数値は除く）がすべて空の行のときだけ小計にする。
    """
    cells = row.cells[:width] if width else row.cells
    filled = [(i, c) for i, c in enumerate(cells) if c.text]
    nonempty = [c for _i, c in filled]
    found: list[tuple[int, str]] = []
    label_col: dict[int, int] = {}
    for pos, (col, c) in enumerate(filled[:3]):
        if not isinstance(c.value, str) or len(c.text) > 30 or "\n" in c.text:
            continue
        if "計" not in c.text and "平均" not in c.text:
            continue
        raw = _PAREN_TAIL_RE.sub("", unicodedata.normalize("NFKC", c.text)).strip()
        label = raw.replace(" ", "")
        if label and len(label) <= 25:
            found.append((pos, raw))
            label_col[pos] = col
    if not found:
        return None
    strings = sum(1 for c in nonempty if c.text not in NA_TOKENS and value_kind(c.value, c.text) != "number")
    for pos, label in found:
        # 「P-004, 2026/04/04, 稼働時間累計, …」「pH計, 2026/08/01, 7.1」のように、番号・日付を持つ行の
        # 項目名（計器名など）は合計・小計行ではない。先頭列のラベルは同じ行の右側の番号・日付も見る
        others = filled[:pos] if pos > 0 else filled[1:]
        if _row_has_own_key(cells, others, key_cols or [], label_col[pos]):
            continue
        if _TOTAL_RE.match(label.replace(" ", "")) and strings <= 3:
            return "total"
        if pos == 0 and strings <= 2:
            if _SUBTOTAL_RE.search(label):
                return "subtotal"
            others = [k for k in key_cols or [] if k != label_col[pos]]
            if label.endswith("計") and others and all(
                    k >= len(cells) or not cells[k].text or value_kind(cells[k].value, cells[k].text) == "number"
                    for k in others):
                return "subtotal"
    return None


def _row_has_own_key(cells: list, before: list, key_cols: list[int], label_col: int) -> bool:
    """ラベルより左に番号・日付があるか、ほかのキー列に数値でない値があるか（＝データ行の印）。"""
    if any(value_kind(c.value, c.text) in ("code", "date", "datetime") for _i, c in before):
        return True
    return any(k != label_col and k < len(cells) and cells[k].text and cells[k].text not in NA_TOKENS
               and value_kind(cells[k].value, cells[k].text) != "number" for k in key_cols)


def _classify(row: SourceRow, ctx: _Ctx, after_blank: bool = False) -> RowClass:
    """行の種類。after_blank: 直前が空行か（空行の後の書き添え・継続行の扱いに使う）。"""
    width = ctx.width or len(row.cells)
    cells = row.cells[:width]
    nonempty = [c for c in cells if c.text]
    if not nonempty:
        if any(c.text for c in row.cells[width:]):
            return RowClass(row.index, "excluded", "見出しの範囲外にだけ値がある行")
        return RowClass(row.index, "blank")
    if row.hidden:
        return RowClass(row.index, "excluded", "非表示の行")
    strikes = sum(1 for c in nonempty if c.strike)
    if strikes and (strikes >= 0.5 * len(nonempty) or any(k < len(cells) and cells[k].strike for k in ctx.key_cols)):
        return RowClass(row.index, "excluded", "取り消し線の行")
    first = nonempty[0].text
    if ctx.header_norms and len(first) <= 40 and _norm(first) in ctx.header_norms:
        norms = {_norm(c.text) for c in nonempty}
        hits = len(norms & ctx.header_norms)
        if hits >= max(2, 0.6 * len(ctx.header_norms)):
            return RowClass(row.index, "header", "見出しの再出現")
    # 注記・表題は表の左端の列から書かれた行だけ。文章の列だけに「●再発防止…」「※後日交換」とある行は
    # 前の記録の続き（継続行）なので、ここでは決めずに下の判定へ回す
    at_left = next(i for i, c in enumerate(cells) if c.text) <= ctx.left
    if len(nonempty) <= 2 and at_left and _NOTE_RE.match(first):
        return RowClass(row.index, "note", "注記")
    if len(nonempty) <= 2 and at_left and _BLOCK_TITLE_RE.match(first):
        return RowClass(row.index, "title", "表題")
    if after_blank and len(nonempty) == 1 and len(first) <= 20             and _FOOTER_RE.match(unicodedata.normalize("NFKC", first).strip()):
        return RowClass(row.index, "note", "表の下の書き添え")
    agg = _aggregate_label(row, width, ctx.key_cols)
    if agg:
        return RowClass(row.index, "subtotal", "合計" if agg == "total" else "小計")
    if ctx.is_csv and width >= 4 and len(row.cells) < 0.5 * width and len(nonempty) <= 2:
        return RowClass(row.index, "excluded", "列数が見出しと合わない")
    if ctx.key_cols and all(not _filled(row, k) for k in ctx.key_cols):
        if _has_own_values(row, ctx) or after_blank:
            # キー列（日付など）を書き省いただけの行。数値・日付を持つので別の記録。
            # 空行を挟んだ行も前の記録の続きとは限らないので、黙って混ぜずに1件として出す
            return RowClass(row.index, "data")
        return RowClass(row.index, "continuation", "キー列が空で文章だけの行")
    if ctx.key_cols and _keys_merged_from_above(row, ctx):
        return RowClass(row.index, "continuation", "キー列が上の行から縦に結合された行")
    return RowClass(row.index, "data")


def _filled(row: SourceRow, col: int) -> bool:
    cell = row.cell(col)
    if cell is None:
        return False
    return bool(cell.text) or (cell.merged_anchor is not None and cell.merged_anchor != (row.index, col + 1))


def _keys_merged_from_above(row: SourceRow, ctx: _Ctx) -> bool:
    """キー列がすべて上のデータ行からの縦結合（または空）で、この行だけの値が文字だけか。

    管理No・日付を縦に結合して1件の対応内容を複数行に書いた表を、1件にまとめるため。
    数値・日付がこの行にあるなら別の記録（同じ日の部品交換の2件目など）として扱う。
    """
    top = ctx.header_rows[-1] if ctx.header_rows else 0
    merged = 0
    for k in ctx.key_cols:
        cell = row.cell(k)
        if cell is None or cell.text:
            if cell is not None and cell.text:
                return False
            continue
        anchor = cell.merged_anchor
        if anchor is not None and top < anchor[0] < row.index:
            merged += 1
    if not merged:
        return False
    return not _has_own_values(row, ctx)


def _has_own_values(row: SourceRow, ctx: _Ctx) -> bool:
    """キー列以外に数値・日付・時刻の値があるか（＝前の記録の続きではなく、この行だけの記録の印）。"""
    keys = set(ctx.key_cols)
    return any(i not in keys and c.text and value_kind(c.value, c.text) in ("number", "date", "datetime", "time")
               for i, c in enumerate(row.cells[:ctx.width or len(row.cells)]))


def _auto_key_columns(source, sheet, data_start: int, ctx: _Ctx) -> list[int]:
    """継続行の判定に使うキー列。値がほぼ埋まった短い列（日付・コードを優先）を最大3列。"""
    rows = []
    blanks = 0
    for row in source.rows(sheet, data_start, KEY_SAMPLE_ROWS * 3):
        rc = _classify(row, ctx)
        if rc.kind == "blank":
            blanks += 1
            if blanks >= BLANK_ROWS_END:
                break
            continue
        blanks = 0
        if rc.kind in ("title", "note", "header"):
            break
        if rc.kind == "data":
            rows.append(row)
            if len(rows) >= KEY_SAMPLE_ROWS:
                break
    if len(rows) < 3:
        return []
    candidates = []
    for col in range(ctx.width):
        filled = sum(1 for r in rows if _filled(r, col))
        if filled < 0.9 * len(rows):
            continue
        kinds = Counter(value_kind(r.cells[col].value, r.cells[col].text) for r in rows if r.cell(col) and r.cells[col].text)
        if kinds.get("text", 0) > 0.2 * len(rows):
            continue
        priority = 0 if (kinds.get("date", 0) + kinds.get("datetime", 0) + kinds.get("code", 0)) >= 0.5 * len(rows) else 1
        candidates.append((priority, col))
    return [col for _, col in sorted(candidates)[:3]]


def _scan(source, sheet, head: list[SourceRow], ctx: _Ctx, data_start: int, data_end: int | None,
          max_scan_rows: int | None):
    """先頭からデータの終わりまで分類する。戻り値: (row_classes, counts, data_end, warnings)"""
    classes: list[RowClass] = []
    warnings: list[str] = []
    counts: Counter[str] = Counter()
    extra = 0

    def record(rc: RowClass):
        nonlocal extra
        if rc.index <= PREVIEW_ROWS:
            classes.append(rc)
        elif rc.kind not in ("data", "blank") and extra < MAX_ROW_CLASSES:
            classes.append(rc)
            extra += 1

    for row in head:
        if row.index >= data_start:
            break
        if row.index in ctx.header_rows:
            kind = "header"
        elif row.is_blank:
            kind = "blank"
        elif _NOTE_RE.match(next(c.text for c in row.cells if c.text)):
            kind = "note"
        else:
            kind = "title"
        record(RowClass(row.index, kind, "表の上" if kind in ("title", "note") else ""))

    last_nonblank = data_start - 1
    blank_run = 0
    blank_total = 0
    end: int | None = None
    stop_reason = ""
    scanned = 0
    after_rows: list[SourceRow] = []
    prev_blank = False
    for row in source.rows(sheet, data_start):
        if end is not None:
            after_rows.append(row)
            if len(after_rows) >= LOOKAHEAD_ROWS:
                break
            continue
        scanned += 1
        rc = _classify(row, ctx, prev_blank)
        prev_blank = rc.kind == "blank"
        if data_end is not None:
            if row.index > data_end:
                end = data_end
                after_rows.append(row)
                continue
            counts[rc.kind] += 1
            record(rc)
            continue
        if rc.kind == "blank":
            blank_run += 1
            record(rc)
            if blank_run >= BLANK_ROWS_END:
                end = last_nonblank
                stop_reason = "blank"
            continue
        if rc.kind in ("note", "title") or (rc.kind == "header" and rc.reason == "見出しの再出現"):
            end = last_nonblank
            after_rows.append(row)
            stop_reason = rc.kind
            continue
        blank_total += blank_run  # 途中の空行はデータ範囲に数える
        blank_run = 0
        counts[rc.kind] += 1
        record(rc)
        last_nonblank = row.index
        if rc.kind == "subtotal" and rc.reason == "合計":
            end = row.index
            stop_reason = "total"
            continue
        if max_scan_rows is not None and scanned >= max_scan_rows:
            end = row.index
            stop_reason = "limit"
            break
    if end is None:
        end = data_end if data_end is not None else last_nonblank

    if data_end is None:
        counts["blank"] = blank_total
    if not counts.get("blank"):
        counts.pop("blank", None)
    # 終わりより後の行（先頭60行の表示用と、別の表があるかの確認）
    other_table_row = None
    prev_blank = True
    for row in after_rows:
        rc = _classify(row, ctx, prev_blank)
        prev_blank = rc.kind == "blank"
        if rc.kind == "blank":
            kind, reason = "blank", ""
        elif rc.kind in ("note", "title"):
            kind, reason = rc.kind, "表の下"
        else:
            kind, reason = "excluded", f"データの範囲外（{end}行目で終了）"
        if row.index <= PREVIEW_ROWS or (kind != "blank" and extra < 20):
            if row.index > PREVIEW_ROWS:
                extra += 1
            classes.append(RowClass(row.index, kind, reason))
        if other_table_row is None and (
            (rc.kind == "title" and row.index > end)
            or (rc.kind not in ("blank", "note") and sum(1 for c in row.cells if c.text) >= 3)
        ):
            other_table_row = row.index
    if stop_reason != "limit" and other_table_row is not None:
        warnings.append(
            f"データの終わり（{end}行目）より下にも表らしい行があります（{other_table_row}行目）。"
            "最初の表だけを読み取ります。残りは範囲を指定して1つずつ取り込んでください"
        )
    if stop_reason == "header":
        warnings.append(f"{end + 1}行目で見出しが再び出たので、そこで読み取りを止めました")
    if counts.get("excluded"):
        warnings.append(f"除外した行が{counts['excluded']}行あります（非表示・取り消し線など）")
    classes.sort(key=lambda rc: rc.index)
    return classes, dict(counts), end, warnings


# ---- 内部: その他 ----

def _no_table(sheet: str, head: list[SourceRow], message: str) -> LayoutGuess:
    classes = [RowClass(r.index, "blank" if r.is_blank else "title") for r in head[:PREVIEW_ROWS]]
    kind = "form_like" if _form_like(head) else "unknown"
    return LayoutGuess(sheet, kind, [], 1, 0, [], classes, 0.0, [message])


def _labels_down_first_column(head: list[SourceRow], rows_h: list[int], data_start: int, end: int) -> bool:
    """見出し行が2列以下で、下の行の先頭列に項目名（辞書の語）が並ぶなら帳票（ラベル: 値）。"""

    header = next((r for r in head if r.index == rows_h[-1]), None)
    if header is None or sum(1 for c in header.cells if c.text) > 3:
        return False
    firsts = [next(c for c in r.cells if c.text) for r in head if data_start <= r.index <= end and not r.is_blank]
    labels = sum(1 for c in firsts if _is_header_cell(c) and (lookup_header(c.text) or ("", ""))[1] == "dictionary")
    return len(firsts) >= 3 and labels >= max(3, 0.5 * len(firsts))


def _form_like(head: list[SourceRow]) -> bool:
    """「ラベル: 値」が並ぶ帳票らしい形か。"""
    pairs = 0
    for row in head:
        cells = [c for c in row.cells if c.text]
        if 2 <= len(cells) <= 6 and _is_header_cell(cells[0]) and len(cells[0].text) <= 12:
            pairs += 1
        elif len(cells) == 1 and re.search(r"[:：]\s*\S", cells[0].text) and len(cells[0].text) <= 40:
            pairs += 1
    return pairs >= 5


def _hidden_columns(source, sheet, width: int) -> list[str]:
    if getattr(source, "kind", "") != "excel":
        return []
    from openpyxl.utils import get_column_letter

    for info in source.sheets():
        if info.name == sheet:
            return [get_column_letter(c) for c in info.hidden_columns if c <= width]
    return []


# ====================================================================================================
# 元 tables/dictionary.py
# 一覧表用の標準キー辞書。
#
# 帳票の pattern/dictionary.py は、保存済み帳票の停止ラベルにも使われているので変更しない。
# キー名は意味が同じものだけ帳票とそろえる（equipment_id, equipment_name, downtime, symptom, cause, alarm など）。
# ====================================================================================================

@dataclass(frozen=True)
class StdColumn:
    key: str
    display: str
    type: str  # code/string/text/date/datetime/time/number/enum/status
    role: str  # key/date/entity/entity_label/category/measure/text/log/person/attribute
    synonyms: tuple[str, ...] = field(default_factory=tuple)
    unit: str = ""
    md: str = "attribute"  # body/attribute/omit
    log_candidate: bool = False  # 役割「経過の記録」（log）になりやすい列


def _c(key, display, type_, role, synonyms, unit="", md="attribute", log_candidate=False) -> StdColumn:
    return StdColumn(key, display, type_, role, tuple(synonyms), unit, md, log_candidate)


STANDARD_COLUMNS: list[StdColumn] = [
    # 識別・日時
    _c("record_no", "管理No", "code", "key",
       ["管理No", "管理番号", "管理NO", "故障No", "トラブルNo", "トラブル番号", "台帳No", "台帳番号", "報告書No",
        "報告番号", "記録No", "記録番号", "整理番号", "受付番号", "伝票No", "カルテNo", "カルテ番号", "受付No", "案件No", "案件番号", "問合せNo", "No", "番号"]),
    _c("occurred_at", "発生日時", "datetime", "date",
       ["発生日時", "発生日", "故障発生日", "故障発生日時", "発生年月日", "日付", "年月日", "作業日", "実施日",
        "起票日", "発見日", "発見日時"]),
    _c("occurred_time", "発生時刻", "time", "attribute", ["発生時刻", "時刻", "発生時間"]),
    _c("reported_at", "報告日時", "datetime", "attribute", ["報告日時", "報告日", "連絡日時"]),
    _c("started_at", "対応開始日時", "datetime", "attribute", ["対応開始日時", "対応開始日", "着手日", "作業開始日時"]),
    _c("completed_at", "完了日", "datetime", "attribute", ["完了日", "完了日時", "復旧日時", "復旧日", "終了日", "終了日時"]),
    _c("due_date", "期限", "date", "attribute", ["期限", "対策期限", "完了予定日", "予定日"]),
    # 設備・場所
    _c("equipment_id", "設備番号", "code", "entity",
       ["設備番号", "設備No", "設備NO", "設備コード", "設備ID", "装置番号", "装置No", "装置ID", "装置コード",
        "機番", "号機", "対象設備", "設備"]),
    _c("equipment_name", "設備名", "string", "entity_label", ["設備名", "装置名", "設備名称", "装置名称", "機器名"]),
    _c("equipment_class", "設備分類", "enum", "attribute", ["設備分類", "設備分類コード", "設備区分", "装置区分", "設備種別"]),
    _c("line", "ライン", "string", "attribute", ["ライン", "ラインコード", "ライン名", "製造ライン"]),
    _c("process", "工程", "string", "attribute", ["工程", "工程名", "工程コード"]),
    _c("location", "設置場所", "string", "attribute", ["設置場所", "場所", "エリア"]),
    _c("part_location", "部位", "string", "attribute", ["部位", "部位コード", "故障部位", "ユニット"]),
    # 区分
    _c("record_type", "記録区分", "enum", "category", ["記録区分", "記録種別"]),
    _c("failure_category", "故障区分", "enum", "category",
       ["故障区分", "故障区分コード", "故障分類", "区分", "不具合区分", "作業区分", "保全区分", "トラブル区分"]),
    _c("severity", "重要度", "enum", "attribute", ["重要度", "重要度コード", "重大度", "影響度", "ランク"]),
    _c("shift", "シフト", "enum", "attribute", ["シフト", "シフトコード", "勤務帯", "直"]),
    _c("status", "状態", "status", "category", ["状態", "状態コード", "ステータス", "進捗状況", "対応状況"]),
    _c("recurrence", "再発", "enum", "attribute", ["再発", "再発フラグ", "再発有無"]),
    _c("judgement", "判定", "enum", "attribute", ["判定", "効果判定", "評価"]),
    # 文章
    _c("alarm", "アラーム", "code", "attribute", ["アラーム", "アラームコード", "エラーコード", "アラーム番号", "警報"]),
    _c("symptom", "現象", "text", "text",
       ["現象", "故障内容", "不具合内容", "不具合現象", "症状", "トラブル内容", "事象", "故障現象"], md="body"),
    _c("initial_action", "初期対応", "text", "text", ["初期対応", "応急処置", "暫定対策", "暫定処置", "初動"], md="body"),
    _c("investigation", "調査内容", "text", "text", ["調査内容", "調査結果", "調査"], md="body"),
    _c("cause", "原因", "text", "text", ["原因", "推定原因", "故障原因", "真因", "原因内容", "真因(なぜなぜ要約)"], md="body"),
    _c("cause_category", "原因区分", "enum", "category", ["原因区分", "原因区分コード", "原因分類"]),
    _c("why_analysis", "なぜなぜ分析", "text", "text", ["なぜなぜ分析", "なぜなぜ"], md="body"),
    _c("action", "処置", "text", "text",
       ["処置", "処置内容", "対策", "対策内容", "作業内容", "修理内容", "処置・対策", "作業内容・処置"],
       md="body", log_candidate=True),
    _c("response_log", "対応内容", "text", "log",
       ["対応内容", "対応履歴", "対応経過", "経過", "経緯", "対応記録", "経過記録", "進捗", "対応メモ"],
       md="body", log_candidate=True),
    _c("permanent_action", "恒久対策", "text", "text", ["恒久対策", "再発防止策", "再発防止対策", "恒久処置"], md="body"),
    _c("horizontal_deployment", "水平展開", "text", "text", ["水平展開"], md="body"),
    _c("result", "結果", "text", "text", ["結果", "効果確認", "処置結果"], md="body"),
    _c("remarks", "備考", "text", "attribute", ["備考", "特記事項", "備考・特記", "メモ", "コメント"]),
    # 部品・数量・金額・時間
    _c("parts_used", "使用部品", "text", "attribute", ["使用部品", "使用部品一覧"]),
    _c("part_name", "部品名", "string", "attribute", ["部品名", "交換部品", "交換部品名", "部品"]),
    _c("part_no", "品番", "code", "attribute", ["品番", "部品番号", "部品コード", "型番"]),
    _c("quantity", "数量", "number", "measure", ["数量", "個数", "使用数", "交換数"]),
    _c("unit_price", "単価", "number", "measure", ["単価"], unit="円"),
    _c("cost", "費用", "number", "measure", ["費用", "金額", "部品代", "修理費", "コスト"], unit="円"),
    _c("downtime", "停止時間", "number", "measure", ["停止時間", "ダウンタイム", "設備停止時間", "ライン停止時間"], unit="分"),
    _c("work_hours", "作業工数", "number", "measure", ["作業工数", "工数", "作業時間", "時間"], unit="h"),
    _c("scrap_qty", "廃棄枚数", "number", "measure", ["廃棄枚数", "廃棄数", "不良数"]),
    _c("attachments", "添付数", "number", "attribute", ["添付数", "添付"]),
    # 人・組織
    _c("worker", "担当者", "string", "person",
       ["担当者", "担当", "作業者", "実施者", "対応者", "担当者社員番号", "保全担当"]),
    _c("reporter", "報告者", "string", "person", ["報告者", "起票者", "報告者社員番号", "連絡者"]),
    _c("approver", "承認者", "string", "person", ["承認者", "確認者"]),
    _c("department", "部署", "string", "attribute", ["部署", "担当部署", "起票部署", "部門", "部門コード", "担当部門コード", "報告者部門コード"]),
    # 関連・管理用
    _c("related_no", "関連番号", "code", "attribute", ["関連管理番号", "関連トラブルNo", "関連No", "関連番号"]),
    _c("lot", "対象ロット", "string", "attribute", ["対象ロット", "ロット", "ロットNo"]),
    _c("internal_code", "内部コード", "code", "attribute", ["内部コード"]),
    _c("registered_at", "登録日時", "datetime", "attribute", ["登録日時", "作成日時"]),
    _c("updated_at", "更新日時", "datetime", "attribute", ["更新日時", "最終更新日時"]),
    _c("registered_by", "登録者", "string", "attribute", ["登録者ID", "登録者", "更新者ID", "更新者"]),
]

_STRIP_RE = re.compile(r"[\s・()（）\[\]【】「」『』_\-‐－/／:：.。、,，#＃*＊]")


def norm_header(text) -> str:
    """見出し比較用の正規化（NFKC・小文字・空白と区切り記号の除去）。"""
    s = unicodedata.normalize("NFKC", str(text or "")).lower()
    return _STRIP_RE.sub("", s)


_SYNONYMS: dict[str, StdColumn] = {}
for _col in STANDARD_COLUMNS:
    for _syn in (_col.display, *_col.synonyms):
        _SYNONYMS.setdefault(norm_header(_syn), _col)
_SIMILAR_KEYS = sorted(
    (k for k in _SYNONYMS if len(k) >= 2 and not (k.isascii() and len(k) < 4)),
    key=len, reverse=True,
)


def lookup_header(header: str) -> tuple[StdColumn, str] | None:
    """見出しから標準キーを探す。戻り値: (標準列, "dictionary" | "similar")"""

    name, _unit = split_header_unit(header)
    candidates = [header, name]
    if "_" in str(header):  # 2段見出し「交換部品_品番」は下段でも照合する
        lower = str(header).rsplit("_", 1)[-1]
        candidates += [lower, split_header_unit(lower)[0]]
    for candidate in candidates:
        if norm_header(candidate) in _SYNONYMS:
            return _SYNONYMS[norm_header(candidate)], "dictionary"
    n = norm_header(name)
    if not n:
        return None
    for key in _SIMILAR_KEYS:  # 長い同義語から順に、見出しに含まれるものを探す
        if key in n:
            return _SYNONYMS[key], "similar"
    close = difflib.get_close_matches(n, list(_SYNONYMS), n=1, cutoff=0.8)
    if close:
        return _SYNONYMS[close[0]], "similar"
    return None


# ====================================================================================================
# 元 tables/source_cache.py
# 取り込みごとの読み取り結果の控え（シート一覧・シートの大きさ・先頭の行・表の形の推定・列の見本の行）。
#
# 画面を開くたびに表ファイル全体を開き直して読み直さないように、取り込みのフォルダに小さな JSON で持つ。
# ファイル（ハッシュ）と読み込み設定（文字コード・区切り・読めない文字の扱い・セル数の上限）を鍵にし、
# どれかが変わったら中身を捨てて作り直す（読み取り・判定のコードを直したときも作り直す）。
# 中身は元のファイルから同じ計算で作れるものだけ（消してもよい）。
#
# ImportSource は TableSource として振る舞う: 控えにある先頭の行は控えから返し、それ以外を読むときだけ元のファイルを開く。
# ====================================================================================================

CACHE_FILE = "source_cache.json"
CACHE_VERSION = 1
MAX_LAYOUTS = 24  # 見出し行を何度も指定し直したときに増えすぎないように
MAX_SAMPLES = 4


class _NotCacheable(Exception):
    """JSON にできない値（控えに入れずに毎回読む）。"""


# ---- 値・行・表の形の JSON 化 ------------------------------------------------------------

def _enc_value(v):
    if v is None or isinstance(v, (bool, int, float, str)):
        return v
    if isinstance(v, datetime):
        return {"dt": v.isoformat()}
    if isinstance(v, date):
        return {"d": v.isoformat()}
    if isinstance(v, time):
        return {"t": v.isoformat()}
    if isinstance(v, timedelta):
        return {"td": [v.days, v.seconds, v.microseconds]}
    raise _NotCacheable(type(v).__name__)


def _dec_value(v):
    if not isinstance(v, dict):
        return v
    if "dt" in v:
        return datetime.fromisoformat(v["dt"])
    if "d" in v:
        return date.fromisoformat(v["d"])
    if "t" in v:
        return time.fromisoformat(v["t"])
    days, seconds, micro = v["td"]
    return timedelta(days=days, seconds=seconds, microseconds=micro)


def _enc_row(row: SourceRow) -> list:
    cells = []
    for c in row.cells:
        if c.value is None and not c.text and c.number_format is None and not (c.bold or c.strike or c.fill) \
                and c.merged_anchor is None:
            cells.append(0)  # 何もない空欄
            continue
        flags = (1 if c.bold else 0) | (2 if c.strike else 0) | (4 if c.fill else 0)
        cells.append([_enc_value(c.value), c.text, c.number_format, flags,
                      list(c.merged_anchor) if c.merged_anchor else None])
    return [row.index, row.hidden, cells]


def _dec_row(data: list) -> SourceRow:
    cells = []
    for c in data[2]:
        if c == 0:
            cells.append(CellInfo(value=None, text=""))
            continue
        value, text, number_format, flags, anchor = c
        cells.append(CellInfo(value=_dec_value(value), text=text, number_format=number_format, bold=bool(flags & 1),
                              strike=bool(flags & 2), fill=bool(flags & 4),
                              merged_anchor=tuple(anchor) if anchor else None))
    return SourceRow(index=data[0], cells=cells, hidden=data[1])


def _dec_layout(data: dict) -> LayoutGuess:
    d = dict(data)
    d["row_classes"] = [RowClass(**rc) for rc in d["row_classes"]]
    return LayoutGuess(**d)


def layout_key(sheet: str, anchors=None, header_row=None, data_end=None, header_rows=None, max_scan_rows=None) -> str:
    """guess_layout の結果を左右する引数だけで鍵を作る（見出し行を指定したときはアンカー語・見出し候補は使われない）。"""
    header_rows = sorted(int(r) for r in header_rows) if header_rows else None
    header_row = int(header_row) if header_row and not header_rows else None
    anchors = list(anchors) if anchors and not header_rows and not header_row else None
    return json.dumps([sheet, anchors, header_row, header_rows, int(data_end) if data_end else None,
                       int(max_scan_rows) if max_scan_rows else None], ensure_ascii=False)


# ---- 控えのファイル ----------------------------------------------------------------------------

_CODE_MODULES = ("tables.py",)   # 読み取り・判定のコードはこのファイルにまとめてある


@functools.lru_cache(maxsize=1)
def _code_version() -> str:
    """読み取り・判定のコードの中身のハッシュ。コードを直したら古い控えを使わない。"""
    digest = hashlib.sha256()
    base = Path(__file__).resolve().parent
    for name in _CODE_MODULES:
        try:
            digest.update((base / name).read_bytes())
        except OSError:
            digest.update(name.encode("utf-8"))
    return digest.hexdigest()[:16]


class ImportSource:
    """元の表ソースの代わりに使う、控え付きの表ソース（kind・file_name・sheets・rows は元と同じ）。

    open_real(sheet_stats) で元のファイルを開く（Excel はシートの大きさの控えを渡すと数え直さない）。
    real に開いた元のソースを渡せば、それを使う（ジョブで読み込みと同じソースを使うとき）。
    """

    def __init__(self, directory: Path, key: str, kind: str, file_name: str,
                 open_real: Callable[[dict | None], object], real=None):
        self.directory = Path(directory)
        self.path = self.directory / CACHE_FILE
        try:
            # 置き場所はここで用意する（消したあとに書き戻さないよう _save では作らない）
            self.directory.mkdir(parents=True, exist_ok=True)
        except OSError:
            pass
        self.key = key
        self.kind = kind
        self.file_name = file_name
        self._open_real = open_real
        self._real = real
        self._data = self._load()
        if real is not None:
            self._remember_stats(real)

    # ---- 控えの読み書き ----

    def _load(self) -> dict:
        try:
            data = json.loads(self.path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            data = None
        version = [CACHE_VERSION, _code_version()]
        if not isinstance(data, dict) or data.get("version") != version or data.get("key") != self.key:
            data = {"version": version, "key": self.key}
        return data

    def _save(self) -> None:
        """一時ファイルに書いて置き換える。書けなくても処理は続ける（次に作り直すだけ）。

        ここではフォルダを作らない: 取り込みを消した（purge_table_import が imports/<id>/ を消した）あとに
        別のタブの画面処理が控えを書くと、元の表の先頭行・見本の行が入った控えごとフォルダが
        復活してしまう（design.md 3.3「データを残さない」）。フォルダを作るのは __init__ だけで、
        消えたあとに書こうとした分は捨てる。
        """
        if not self.directory.is_dir():
            return
        tmp = self.path.with_name(f".{self.path.name}.{os.getpid()}.{threading.get_ident()}.tmp")
        try:
            tmp.write_text(json.dumps(self._data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
            os.replace(tmp, self.path)
        except OSError:
            try:
                tmp.unlink(missing_ok=True)
            except OSError:
                pass

    def _put(self, section: str, key: str, value, limit: int | None = None) -> None:
        # 別のリクエストやジョブが足した分を消さないよう、書く直前に読み直して足す
        latest = self._load()
        entries = dict(latest.get(section) or {})
        entries.pop(key, None)
        entries[key] = value
        if limit is not None:
            while len(entries) > limit:
                entries.pop(next(iter(entries)))
        latest[section] = entries
        self._data = latest
        self._save()

    # ---- 元のソース ----

    @property
    def real(self):
        """元のファイルを開いたソース（初めて使うときに開く）。"""
        if self._real is None:
            stats = self._data.get("sheet_stats") if self.kind == "excel" else None
            self._real = self._open_real(stats)
            self._remember_stats(self._real)
        return self._real

    def _remember_stats(self, real) -> None:
        if self.kind == "excel" and not self._data.get("sheet_stats") and hasattr(real, "sheet_stats"):
            self._data = self._load()
            self._data["sheet_stats"] = real.sheet_stats()
            self._save()

    def __getattr__(self, name):
        # date1904・replaced_rows・sniff など、控えにない属性は元のソースから
        if name.startswith("_"):
            raise AttributeError(name)
        return getattr(self.real, name)

    # ---- TableSource ----

    def sheets(self) -> list[SheetInfo]:
        cached = self._data.get("sheets")
        if cached is None:
            infos = self.real.sheets()
            self._data = self._load()
            self._data["sheets"] = [asdict(s) for s in infos]
            self._save()
            return infos
        return [SheetInfo(**s) for s in cached]

    def rows(self, sheet: str, start: int = 1, limit: int | None = None) -> Iterator[SourceRow]:
        if limit is None or start < 1 or start + limit - 1 > HEAD_ROWS:
            return self.real.rows(sheet, start, limit)
        return iter([r for r in self._head(sheet) if start <= r.index < start + limit])

    def _head(self, sheet: str) -> list[SourceRow]:
        """先頭 HEAD_ROWS 行（見出しの探索・範囲の画面の表示・年度の手がかりに使う範囲）。"""
        cached = (self._data.get("head") or {}).get(sheet)
        if cached is not None:
            return [_dec_row(r) for r in cached]
        rows = list(self.real.rows(sheet, 1, HEAD_ROWS))
        try:
            self._put("head", sheet, [_enc_row(r) for r in rows])
        except _NotCacheable:
            return rows
        return rows

    # ---- 計算結果の控え ----

    def layout(self, sheet: str, anchors=None, header_row=None, data_end=None, header_rows=None,
               max_scan_rows=None) -> LayoutGuess:
        """guess_layout と同じ結果（同じ引数で前に計算していれば控えから）。"""
        key = layout_key(sheet, anchors, header_row, data_end, header_rows, max_scan_rows)
        cached = (self._data.get("layouts") or {}).get(key)
        if cached is not None:
            return _dec_layout(cached)
        memo = dict(self._data.get("scans") or {})
        known = set(memo)
        layout = guess_layout(self, sheet, anchors=anchors, header_row=header_row, data_end=data_end,
                              header_rows=header_rows, max_scan_rows=max_scan_rows, scan_memo=memo)
        for scan_key in set(memo) - known:
            self._put("scans", scan_key, memo[scan_key], MAX_LAYOUTS)
        self._put("layouts", key, asdict(layout), MAX_LAYOUTS)
        return layout

    def list_kind(self, sheet: str) -> str:
        """detect.list_kind と同じ判定（簡易判定の表の形を控えから使う）。"list" / "crosstab" / ""。"""
        try:
            layout = self.layout(sheet, max_scan_rows=200)
        except Exception:
            return ""
        return kind_from_layout(layout)

    def sample_rows(self, sheet: str, layout: LayoutGuess, n: int = 200) -> list[SourceRow]:
        """sample_data_rows と同じ行（同じ表の形なら控えから）。"""
        key = hashlib.sha256(json.dumps([sheet, asdict(layout), n], ensure_ascii=False, sort_keys=True)
                             .encode("utf-8")).hexdigest()
        cached = (self._data.get("samples") or {}).get(key)
        if cached is not None:
            return [_dec_row(r) for r in cached]
        rows = sample_data_rows(self, sheet, layout, n)
        try:
            self._put("samples", key, [_enc_row(r) for r in rows], MAX_SAMPLES)
        except _NotCacheable:
            pass
        return rows


# ====================================================================================================
# 元 tables/spec.py
# 一覧表の取り込み設定（TableSpec）。JSON で保存する（dataclass ⇔ dict、検証、spec_hash）。
#
# - 列の定義（ColumnSpec）、「経過の記録」の列の段（LogStageSpec。画面の役割名は「経過の記録」）、AI の custom 段を持つ。
# - 取り込み時の見出しとの照合（resolve_columns）と、画面の候補からの設定作成（spec_from_suggestions）もここに置く。
# ====================================================================================================

COLUMN_TYPES = ("code", "string", "text", "date", "datetime", "time", "number", "enum", "status")
COLUMN_ROLES = ("key", "date", "entity", "entity_label", "category", "measure", "text", "log", "person", "attribute")
MD_MODES = ("body", "attribute", "omit")
GROUP_BY = ("month", "entity_month")
ROW_POLICIES = ("exclude_with_warning", "include")
CONTINUATION_POLICIES = ("merge_into_previous", "keep")
CUSTOM_OUTPUT_TYPES = ("text", "choice")

DEFAULT_NA_TOKENS = ["-", "－", "―", "‐", "N/A", "n/a", "NA", "#N/A", "該当なし"]
_KEY_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_]{0,63}$")


def _default_header() -> dict:
    return {"anchors": [], "rows": 1}


def _default_exclude() -> dict:
    return {"hidden_rows": "exclude_with_warning", "strike_rows": "exclude_with_warning"}


# 判定に使っていない項目・前の版にあって今は無い設定。古い JSON に残っていても読み捨てる
RETIRED_KEYS = {"header": ("search_rows",), "exclude": ("aggregate_keywords",),
                # lightrag_hint: ファイル名のヒントは付けない。dedupe_timeline/omit_person: 中身は削らない。
                # max_records_per_file: 記録ファイルは月ごとで、件数では分けない。
                # dataset_card/summaries: 出すのは RAG に入れる記録ファイルだけ（集計・説明は作らない。6.3）。
                "markdown": ("lightrag_hint", "dedupe_timeline", "omit_person", "max_records_per_file",
                             "dataset_card", "summaries")}


def _default_record() -> dict:
    return {"key": [], "fallback_key": ["occurred_at", "equipment_id", "symptom:20"]}


def _default_period() -> dict:
    # 期間の置き換えはしない（取り込みごとにその内容だけで md を作る）。日付の列だけを持つ
    return {"date_column": "occurred_at"}


def _default_checks() -> dict:
    return {"type_error_rate": {"warn": 0.02, "block": 0.10}, "reconcile_tolerance": 0}


def _default_markdown() -> dict:
    return {"file_prefix": "", "group_by": "month", "records": True, "title_columns": []}


@dataclass
class ColumnSpec:
    key: str
    display: str
    headers: list[str] = field(default_factory=list)
    type: str = "string"  # code/string/text/date/datetime/time/number/enum/status
    role: str = "attribute"  # key/date/entity/entity_label/category/measure/text/log/person/attribute
    unit: str = ""
    unit_conversions: dict = field(default_factory=dict)  # {"h": 60} = 1h を列の単位で60
    required: bool = False
    md: str = "attribute"  # body/attribute/omit
    fill_down_blank: bool = False
    normalize: list[str] = field(default_factory=lambda: ["nfkc"])  # nfkc / upper
    allowed: list[str] = field(default_factory=list)
    value_map: dict = field(default_factory=dict)
    description: str = ""


@dataclass
class LogStageSpec:
    column: str
    enabled_ai: bool = False
    context_columns: list[str] = field(default_factory=list)
    people: list[dict] = field(default_factory=list)  # {name, aliases, org}。アプリ内だけで使う
    groups: list[str] = field(default_factory=list)
    glossary: dict = field(default_factory=dict)
    entry_types: list[str] = field(default_factory=list)
    instruction: str = ""
    incident: bool = True
    run_if: dict = field(default_factory=dict)
    limits: dict = field(default_factory=dict)
    splitter: dict = field(default_factory=dict)  # ai.SplitOptions.from_dict の形
    mask: list[str] = field(default_factory=lambda: ["phone", "email"])


@dataclass
class CustomStageSpec:
    id: str
    inputs: list[str] = field(default_factory=list)
    prompt: str = ""
    output_type: str = "text"  # text/choice
    choices: list[str] = field(default_factory=list)
    max_chars: int = 80
    fallback: str = "不明"
    target_key: str = ""
    quote_required: bool = True


@dataclass
class TableSpec:
    name: str
    description: str = ""
    file_types: list[str] = field(default_factory=lambda: ["xlsx", "xlsm", "csv"])
    name_patterns: list[str] = field(default_factory=list)
    header: dict = field(default_factory=_default_header)
    exclude: dict = field(default_factory=_default_exclude)
    continuation_rows: str = "merge_into_previous"
    na_tokens: list[str] = field(default_factory=lambda: list(DEFAULT_NA_TOKENS))
    fiscal_year_start_month: int = 4
    columns: list[ColumnSpec] = field(default_factory=list)
    record: dict = field(default_factory=_default_record)
    period: dict = field(default_factory=_default_period)
    log_stage: LogStageSpec | None = None
    custom_stages: list[CustomStageSpec] = field(default_factory=list)
    markdown: dict = field(default_factory=_default_markdown)
    checks: dict = field(default_factory=_default_checks)

    # ---- 参照の補助 ----
    def column(self, key: str) -> ColumnSpec | None:
        for col in self.columns:
            if col.key == key:
                return col
        return None

    def columns_with_role(self, role: str) -> list[ColumnSpec]:
        return [c for c in self.columns if c.role == role]

    def first_role(self, *roles: str) -> ColumnSpec | None:
        for role in roles:
            for col in self.columns:
                if col.role == role:
                    return col
        return None

    @property
    def date_key(self) -> str:
        key = (self.period or {}).get("date_column") or ""
        if key and self.column(key):
            return key
        col = self.first_role("date")
        return col.key if col else key

    @property
    def file_prefix(self) -> str:
        return str((self.markdown or {}).get("file_prefix") or self.name)

# ---- 経過の記録の基準日 ----------------------------------------------------------------

_BASE_DATE_RE = re.compile(r"(\d{4})[-/](\d{1,2})[-/](\d{1,2})")


def _sget(obj, name: str, default=None):
    """dataclass でも dict でも同じように読む（AI整形は spec を dict のまま持つことがある）。"""
    if obj is None:
        return default
    value = obj.get(name, None) if isinstance(obj, dict) else getattr(obj, name, None)
    return default if value is None else value


def base_date_from(values: dict, spec) -> date | None:
    """経過の記録の相対日付（「翌週」など）を解くときの基準日。無ければ None（＝年不明）。

    探す順: 期間の日付列 → role='date' の列すべて → occurred_at。'-' でも '/' でも読む。
    AI整形（ai.runner）と Markdown（tables.markdown）で同じ日付にするため、ここ1か所に置く。
    """
    keys = [_sget(_sget(spec, "period", {}) or {}, "date_column", None)]
    for col in _sget(spec, "columns", []) or []:
        if str(_sget(col, "role", "")) == "date":
            keys.append(str(_sget(col, "key", "")))
    keys.append("occurred_at")
    for key in keys:
        if not key:
            continue
        m = _BASE_DATE_RE.search(str((values or {}).get(key) or ""))
        if m:
            try:
                return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
            except ValueError:
                continue
    return None


# ---- dict ⇔ dataclass ----------------------------------------------------------------

def _pick(cls, d: dict) -> dict:
    names = {f.name for f in fields(cls)}
    return {k: copy.deepcopy(v) for k, v in (d or {}).items() if k in names}


def _merged(default: dict, value, drop: tuple = ()) -> dict:
    out = copy.deepcopy(default)
    if isinstance(value, dict):
        out.update(copy.deepcopy(value))
    for name in drop:
        out.pop(name, None)
    return out


def _as_list(value) -> list:
    if value is None:
        return []
    if isinstance(value, (list, tuple)):
        return list(value)
    return [value]


def _str_field(obj, name: str) -> None:
    value = getattr(obj, name)
    if value is not None and not isinstance(value, str):
        setattr(obj, name, str(value))


def _custom_stage_from(d: dict) -> CustomStageSpec:
    stage = CustomStageSpec(**_pick(CustomStageSpec, d))
    for name in ("id", "output_type", "target_key"):
        _str_field(stage, name)
    stage.inputs = [str(v) for v in _as_list(stage.inputs)]
    stage.choices = [str(v) for v in _as_list(stage.choices)]
    return stage


def _column_from(d: dict) -> ColumnSpec:
    if isinstance(d, ColumnSpec):
        return d
    col = ColumnSpec(**_pick(ColumnSpec, d))
    # 手で書いた JSON の「"key": 1」なども文字列にして検証に回す（検証の途中で例外にしない）
    for name in ("key", "display", "type", "role", "md"):
        _str_field(col, name)
    col.headers = [str(h) for h in _as_list(col.headers) if str(h).strip()]
    col.normalize = [str(n) for n in _as_list(col.normalize)]
    col.allowed = [str(a) for a in _as_list(col.allowed)]
    col.value_map = dict(col.value_map or {})
    col.unit_conversions = dict(col.unit_conversions or {})
    col.required = bool(col.required)
    col.fill_down_blank = bool(col.fill_down_blank)
    col.unit = str(col.unit or "")
    col.description = str(col.description or "")
    return col


def spec_from_dict(d: dict) -> TableSpec:
    """dict（JSON）から TableSpec を作る。足りない項目は既定値。"""
    if not isinstance(d, dict):
        raise ValueError("取り込み設定は JSON のオブジェクトで指定してください")
    data = copy.deepcopy(d)
    spec = TableSpec(**_pick(TableSpec, {k: v for k, v in data.items()
                                         if k not in ("columns", "log_stage", "custom_stages", "markdown", "header",
                                                      "data_end", "exclude", "record", "period", "checks")}))
    spec.name = str(spec.name or "")
    spec.columns = [_column_from(c) for c in _as_list(data.get("columns"))]
    spec.header = _merged(_default_header(), data.get("header"), RETIRED_KEYS["header"])
    spec.exclude = _merged(_default_exclude(), data.get("exclude"), RETIRED_KEYS["exclude"])
    spec.record = _merged(_default_record(), data.get("record"))
    spec.period = _merged(_default_period(), data.get("period"))
    spec.checks = _merged(_default_checks(), data.get("checks"))
    spec.markdown = _merged(_default_markdown(), data.get("markdown"), RETIRED_KEYS["markdown"])
    log = data.get("log_stage")
    spec.log_stage = LogStageSpec(**_pick(LogStageSpec, log)) if isinstance(log, dict) and log.get("column") else None
    if spec.log_stage is not None:
        st = spec.log_stage
        # 手で書いた JSON の「"mask": "email"」「"phone,email"」も規則の並びとして読む（1文字ずつにしない）
        if isinstance(st.mask, str):
            st.mask = [m for m in re.split(r"[,、，\s]+", st.mask) if m]
        st.mask = [str(m) for m in _as_list(st.mask)]
        for name in ("context_columns", "groups", "entry_types"):
            setattr(st, name, [str(v) for v in _as_list(getattr(st, name))])
    spec.custom_stages = [_custom_stage_from(c) for c in _as_list(data.get("custom_stages")) if isinstance(c, dict)]
    spec.na_tokens = [str(t) for t in _as_list(spec.na_tokens)]
    spec.name_patterns = [str(p) for p in _as_list(spec.name_patterns)]
    spec.file_types = [str(t) for t in _as_list(spec.file_types)]
    spec.fiscal_year_start_month = int(spec.fiscal_year_start_month or 4)
    return spec


def spec_to_dict(spec: TableSpec) -> dict:
    return asdict(spec)


def spec_json(spec: TableSpec) -> str:
    """保存用の JSON（キー順固定）。"""
    return json.dumps(spec_to_dict(spec), ensure_ascii=False, sort_keys=True)


def spec_hash(spec: TableSpec) -> str:
    canonical = json.dumps(spec_to_dict(spec), ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()


# ---- 検証 ---------------------------------------------------------------------------

def validate_spec(spec: TableSpec) -> list[str]:
    """設定の問題点（日本語）。空なら取り込みに使える。"""
    errors: list[str] = []
    if not str(spec.name or "").strip():
        errors.append("表の名前を入力してください")
    if not spec.columns:
        errors.append("列を1つ以上設定してください")
    keys: set[str] = set()
    for col in spec.columns:
        label = col.display or col.key
        if not _KEY_RE.match(col.key or ""):
            errors.append(f"列「{label}」のキー「{col.key}」は半角英数字と _ で指定してください")
        elif col.key in keys:
            errors.append(f"キー「{col.key}」が重複しています")
        keys.add(col.key)
        if not str(col.display or "").strip():
            errors.append(f"列「{col.key}」の表示名を入力してください")
        if col.type not in COLUMN_TYPES:
            errors.append(f"列「{label}」の型「{col.type}」は使えません")
        if col.role not in COLUMN_ROLES:
            errors.append(f"列「{label}」の役割「{col.role}」は使えません")
        if col.md not in MD_MODES:
            errors.append(f"列「{label}」のmdでの扱い「{col.md}」は使えません")
        for unit, factor in (col.unit_conversions or {}).items():
            if not isinstance(factor, (int, float)) or isinstance(factor, bool) or factor <= 0:
                errors.append(f"列「{label}」の単位換算「{unit}」の倍率が正しくありません")
    if len([c for c in spec.columns if c.role == "entity"]) > 1:
        errors.append("役割「対象（設備・製品・顧客など）」の列は1つだけにしてください")

    all_keys = set(keys)
    record = spec.record or {}

    def _check_key_parts(parts: list, label: str, check_missing: bool = True) -> None:
        # 記録キーの部品は「列」または「列:文字数」（normalize._key_part が切り詰める）
        for part in parts:
            base, _, length = str(part).partition(":")
            if length and not length.isdigit():
                errors.append(f"{label}の「{part}」の文字数は「列:20」のように数字で指定してください")
            elif check_missing and base not in all_keys:
                errors.append(f"{label}の列「{base}」がありません")

    _check_key_parts(_as_list(record.get("key")), "記録キー")
    # fallback_key の既定はよくある列名なので、既定のままのときはその表に無くても問題にしない
    fallback = _as_list(record.get("fallback_key"))
    _check_key_parts(fallback, "記録キーの代わり",
                     check_missing=[str(p) for p in fallback] != _default_record()["fallback_key"])
    date_column = (spec.period or {}).get("date_column")
    col = spec.column(date_column) if date_column else None
    if col is not None and col.type not in ("date", "datetime"):
        errors.append(f"日付の列「{col.display}」の型を日付にしてください")
    if not 1 <= int(spec.fiscal_year_start_month or 0) <= 12:
        errors.append("年度の開始月は1〜12で指定してください")
    if spec.continuation_rows not in CONTINUATION_POLICIES:
        errors.append("継続行の扱いが正しくありません")
    header = spec.header if isinstance(spec.header, dict) else {}
    value = header.get("rows")
    if isinstance(value, bool) or not isinstance(value, int) or not 1 <= value <= 1000:
        errors.append("見出しの行数（header.rows）は1以上の整数で指定してください")
    for name in ("hidden_rows", "strike_rows"):
        if (spec.exclude or {}).get(name) not in ROW_POLICIES:
            errors.append("非表示行・取り消し線の行の扱いが正しくありません")
            break

    md = spec.markdown or {}
    if md.get("group_by") not in GROUP_BY:
        errors.append("記録ファイルのまとめ方は month / entity_month から選んでください")
    elif md.get("group_by") == "entity_month" and not spec.first_role("entity"):
        errors.append("対象×月でまとめるには、役割「対象（設備・製品・顧客など）」の列が必要です")
    for key in _as_list(md.get("title_columns")):
        base = str(key).split(":")[0]
        if base not in all_keys:
            errors.append(f"見出しに使う列「{base}」がありません")

    if spec.log_stage is not None:
        col = spec.column(spec.log_stage.column)
        if col is None:
            errors.append(f"経過の記録の列「{spec.log_stage.column}」がありません")
        for key in spec.log_stage.context_columns:
            if key not in keys:
                errors.append(f"AI整形に添える列「{key}」がありません")
        errors.extend(_log_stage_errors(spec.log_stage))
    try:
        tolerance = float((spec.checks or {}).get("reconcile_tolerance") or 0)
        if not 0 <= tolerance < float("inf"):
            raise ValueError
    except (TypeError, ValueError):
        errors.append("突き合わせの許容差（checks.reconcile_tolerance）は0以上の数値で指定してください")
    stage_ids: set[str] = set()
    for stage in spec.custom_stages:
        if not _KEY_RE.match(stage.id or "") or stage.id in stage_ids:
            errors.append(f"AIの追加処理のID「{stage.id}」が正しくないか重複しています")
        stage_ids.add(stage.id)
        if stage.output_type not in CUSTOM_OUTPUT_TYPES:
            errors.append(f"AIの追加処理「{stage.id}」の出力の種類が正しくありません")
        if stage.output_type == "choice" and not stage.choices:
            errors.append(f"AIの追加処理「{stage.id}」の選択肢を入力してください")
        for key in stage.inputs:
            if key not in keys:
                errors.append(f"AIの追加処理「{stage.id}」の入力列「{key}」がありません")
    rates = (spec.checks or {}).get("type_error_rate") or {}
    if not isinstance(rates, dict):
        errors.append("型エラーの割合の上限（checks.type_error_rate）は warn と block を持つ形で指定してください")
        rates = {}
    try:
        if not 0 <= float(rates.get("warn", 0.02)) <= float(rates.get("block", 0.10)) <= 1:
            errors.append("型エラーの割合の上限は 0〜1 で、警告 ≦ 確定を止める にしてください")
    except (TypeError, ValueError):
        errors.append("型エラーの割合の上限は数値で指定してください")
    return errors


# 区切りの正規表現（JSON で取り込んだ設定だけが持つ）の上限。画面からは設定しない
_MAX_SPLIT_PATTERNS = 20
_MAX_PATTERN_CHARS = 200
# 量指定子を含むグループにさらに量指定子が付く形（(.+)+ など）と、選択（|）を含むグループに量指定子が付く形
# （(?:\d|\d)* など。選択肢が重なると同じく極端に遅くなる）。どちらも受け付けない
_NESTED_QUANTIFIER_RE = re.compile(r"\([^)]*[+*|][^)]*\)[+*{]")


def _log_stage_errors(stage: LogStageSpec) -> list[str]:
    """AI整形の設定（JSON で取り込んだときだけ画面に出ない項目）の問題点。"""
    from app.ai import normalize_rules

    errors: list[str] = []
    bad = [m for m in stage.mask if not normalize_rules([m])]
    if bad:
        errors.append(f"AI整形の伏せ字の規則「{'、'.join(bad)}」は使えません（phone / email / person / amount か"
                      "電話番号 / メール / 人名 / 金額）")
    if not isinstance(stage.people, list) or not all(
            isinstance(p, dict) and isinstance(p.get("name"), str) and p["name"].strip()
            and isinstance(p.get("aliases", []), list) for p in stage.people):
        errors.append("AI整形の人名一覧（people）は、name（名前）を持つ項目の並びで指定してください")
    for name, label in (("glossary", "用語集"), ("splitter", "区切り"), ("limits", "上限"), ("run_if", "実行条件")):
        if not isinstance(getattr(stage, name), dict):
            errors.append(f"AI整形の{label}（{name}）の書き方が正しくありません")
    splitter = stage.splitter if isinstance(stage.splitter, dict) else {}
    for key in ("extra_anchors", "not_date_patterns"):
        patterns = splitter.get(key)
        if patterns is None:
            continue
        if not isinstance(patterns, list) or len(patterns) > _MAX_SPLIT_PATTERNS:
            errors.append(f"AI整形の区切りの正規表現（{key}）は{_MAX_SPLIT_PATTERNS}個までの並びで指定してください")
            continue
        for pat in patterns:
            text = str(pat)[:40]
            if not isinstance(pat, str) or len(pat) > _MAX_PATTERN_CHARS:
                errors.append(f"AI整形の区切りの正規表現「{text}」は{_MAX_PATTERN_CHARS}文字以内の文字列にしてください")
                continue
            try:
                re.compile(pat)
            except re.error:
                errors.append(f"AI整形の区切りの正規表現「{text}」が正しくありません")
                continue
            if _NESTED_QUANTIFIER_RE.search(pat):
                errors.append(f"AI整形の区切りの正規表現「{text}」は処理が極端に遅くなる形（(…+)+ など）です")
    for key in ("sentence_split_min_chars", "order_tolerance_days"):
        if splitter.get(key) is not None:
            try:
                int(splitter[key])
            except (TypeError, ValueError):
                errors.append(f"AI整形の区切りの {key} は整数で指定してください")
    # 上限・実行条件の数値（aiproc/runner.py が int() で読む。数値でないと分割プレビュー・AI整形が止まる）
    limits = stage.limits if isinstance(stage.limits, dict) else {}
    for key in ("max_segments", "max_input_tokens"):
        if key in limits and not _is_int_at_least(limits[key], 1):
            errors.append(f"AI整形の上限（limits）の {key} は1以上の整数で指定してください")
    if isinstance(stage.run_if, dict) and not _run_if_ok(stage.run_if, 0):
        errors.append("AI整形の実行条件（run_if）の min_chars・min_segments は0以上の整数で、"
                      "any・all は条件の並びで指定してください")
    return errors


def _is_int_at_least(value, low: int) -> bool:
    if isinstance(value, bool):
        return False
    try:
        n = int(value)
    except (TypeError, ValueError):
        return False
    return n >= low and (not isinstance(value, float) or value == n)


def _run_if_ok(rule, depth: int) -> bool:
    """実行条件（ai.runner._run_if の形）の数値と入れ子を確かめる。"""
    if not isinstance(rule, dict):
        return True   # 条件でないものは runner が「条件なし」として扱う
    if depth > 10:
        return False
    for key in ("any", "all"):
        if key in rule:
            items = rule[key]
            if items is not None and not isinstance(items, list):
                return False
            if not all(_run_if_ok(r, depth + 1) for r in items or []):
                return False
    return all(_is_int_at_least(rule[k], 0) for k in ("min_chars", "min_segments") if k in rule)


# ---- 見出しとの照合 --------------------------------------------------------------------

@dataclass
class ColumnResolution:
    positions: dict[str, int]  # 列キー → 見出しの位置（0始まり）
    missing_required: list[str]  # 見つからない必須列の表示名
    missing: list[str]  # 見つからない列の表示名（必須以外も含む）
    unused_headers: list[str]  # どの列にも対応しない見出し


def resolve_columns(spec: TableSpec, headers: list[str]) -> ColumnResolution:
    """設定の列を、読み取った見出しの位置に対応づける。"""

    def norms_of(header: str) -> set[str]:
        name = split_header_unit(header)[0]
        out = {norm_header(header), norm_header(name)}
        if "_" in header:
            lower = header.rsplit("_", 1)[-1]
            out |= {norm_header(lower), norm_header(split_header_unit(lower)[0])}
        return {n for n in out if n}

    header_norms = [norms_of(h) for h in headers]
    exact = [norm_header(h) for h in headers]
    positions: dict[str, int] = {}
    used: set[int] = set()

    # 完全一致を先に、その後に単位・2段見出しの下段で一致したものを割り当てる
    for strict in (True, False):
        for col in spec.columns:
            if col.key in positions:
                continue
            candidates = [c for c in (list(col.headers) + [col.display]) if c]
            cand_exact = [norm_header(c) for c in candidates]
            cand_loose = {norm_header(split_header_unit(c)[0]) for c in candidates} | set(cand_exact)
            cand_loose.discard("")
            for pos in range(len(headers)):
                if pos in used:
                    continue
                hit = exact[pos] in cand_exact if strict else bool(header_norms[pos] & cand_loose)
                if hit:
                    positions[col.key] = pos
                    used.add(pos)
                    break
    missing = [c.display for c in spec.columns if c.key not in positions]
    missing_required = [c.display for c in spec.columns if c.required and c.key not in positions]
    unused = [h for pos, h in enumerate(headers)
              if pos not in used and not (h.startswith("列") and h[1:].isdigit())]
    return ColumnResolution(positions, missing_required, missing, unused)


# ---- 候補からの設定作成 -----------------------------------------------------------------

def _get(obj, name, default=None):
    if isinstance(obj, dict):
        return obj.get(name, default)
    return getattr(obj, name, default)


def spec_from_suggestions(name: str, layout, suggestions, options: dict | None = None) -> TableSpec:
    """見出しの判定結果（LayoutGuess）と列の候補（ColumnSuggestion）から取り込み設定を作る。

    options: description, name_patterns, file_types, group_by, fiscal_year_start_month
    """
    options = dict(options or {})
    spec = TableSpec(name=name)  # ファイル名の先頭は空（＝表の名前。TableSpec.file_prefix）
    spec.description = str(options.get("description") or "")
    if options.get("name_patterns"):
        spec.name_patterns = [str(p) for p in options["name_patterns"]]
    if options.get("file_types"):
        spec.file_types = [str(t) for t in options["file_types"]]
    if options.get("fiscal_year_start_month"):
        spec.fiscal_year_start_month = int(options["fiscal_year_start_month"])
    header_rows = list(_get(layout, "header_rows", []) or [])
    spec.header["rows"] = max(1, len(header_rows))

    columns: list[ColumnSpec] = []
    keys: set[str] = set()
    for s in suggestions:
        header = _get(s, "header", "")
        key = _get(s, "key") or f"col{int(_get(s, 'index', len(columns))) + 1}"
        base, n = key, 2
        while key in keys:
            key = f"{base}_{n}"
            n += 1
        keys.add(key)
        type_ = _get(s, "type", "string") or "string"
        if type_ not in COLUMN_TYPES:
            type_ = "string"
        role = _get(s, "role", "attribute") or "attribute"
        if role not in COLUMN_ROLES:
            role = "attribute"
        md = _get(s, "md", "attribute") or "attribute"
        columns.append(ColumnSpec(
            key=key, display=_get(s, "display", "") or header, headers=[header] if header else [],
            type=type_, role=role, unit=_get(s, "unit", "") or "", md=md if md in MD_MODES else "attribute",
            fill_down_blank=bool(_get(s, "fill_down_blank", False)),
            normalize=["nfkc", "upper"] if (role == "entity" and type_ == "code") else ["nfkc"],
        ))
    # 役割の重複を直す（entity・entity_label・key は1列だけ）
    for role in ("entity", "entity_label", "key"):
        for extra in [c for c in columns if c.role == role][1:]:
            extra.role = "attribute"
    spec.columns = columns

    entity = spec.first_role("entity")
    date_col = spec.first_role("date")
    key_col = spec.first_role("key")
    text_col = spec.first_role("text")
    if key_col is not None:
        key_col.required = True
    if date_col is not None:
        date_col.required = True
    spec.record = {
        "key": [key_col.key] if key_col else [],
        "fallback_key": [c for c in (
            date_col.key if date_col else None,
            entity.key if entity else None,
            f"{text_col.key}:20" if text_col else None,
        ) if c],
    }
    spec.period = {"date_column": date_col.key if date_col else ""}
    log_col = next((c for c in columns if c.role == "log"), None)
    if log_col is not None:
        spec.log_stage = LogStageSpec(column=log_col.key, context_columns=[
            c.key for c in (spec.first_role("entity_label"), entity, spec.columns_with_role("text")[0]
                            if spec.columns_with_role("text") else None) if c is not None])
    if options.get("group_by") in GROUP_BY:
        spec.markdown["group_by"] = options["group_by"]
    return spec


# ====================================================================================================
# 元 tables/records.py
# 記録の値の見方（設備の列、日付の月、数値の書き方）。
#
# `tables/summaries.py`（月次集計・設備別年度集計・データセット説明）から、記録ファイルの作成に必要な
# 共通の処理だけを残したもの。集計は 2026-09-20 に外した（docs/design.md 6.3）。
# 書き方（桁区切り・単位）を1か所にまとめて、同じデータから同じ md が出るようにする。
# ====================================================================================================

# ---- 書式 -----------------------------------------------------------------------------

def round_half_up(value: float, digits: int = 0) -> float | int:
    if not math.isfinite(value):
        return value  # 桁あふれした値（inf）は丸めない（Decimal の quantize が落ちるため）
    q = Decimal(1).scaleb(-digits)
    d = Decimal(str(value)).quantize(q, rounding=ROUND_HALF_UP)
    return int(d) if digits == 0 else float(d)


def fmt_number(value, max_decimals: int = 2) -> str:
    """桁区切り付き。小数は最大2桁（末尾の0は削る）。"""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(int(value))
    if isinstance(value, int):
        return f"{value:,}"
    try:
        f = float(value)
    except (TypeError, ValueError):
        return str(value)
    if not math.isfinite(f):
        return str(value)  # 桁あふれした値などで、丸め（Decimal）を落とさない
    if f.is_integer():
        return f"{int(f):,}"
    r = round_half_up(f, max_decimals)
    text = f"{r:,.{max_decimals}f}".rstrip("0").rstrip(".")
    return text


def unit_label(unit: str) -> str:
    u = (unit or "").strip()
    return {"h": "時間", "H": "時間", "hr": "時間", "min": "分"}.get(u, u)


# ---- 月 -------------------------------------------------------------------------------

def month_label(month: str) -> str:
    return f"{int(month[:4])}年{int(month[5:7])}月"


def month_first_day(month: str) -> str:
    return f"{month}-01"


def month_last_day(month: str) -> str:
    # 翌月を経由しない（9999-12 の翌月は日付にできない。9999/12/31 は「期限なし」によく使われる）
    import calendar
    from datetime import date

    y, m = int(month[:4]), int(month[5:7])
    return date(y, m, calendar.monthrange(y, m)[1]).isoformat()


def is_month(text) -> bool:
    s = str(text or "")
    return len(s) >= 7 and s[4] == "-" and s[:4].isdigit() and s[5:7].isdigit() and 1 <= int(s[5:7]) <= 12


# ---- 設備の列 --------------------------------------------------------------------------

def entity_columns(spec: TableSpec):
    entity = spec.first_role("entity")
    label = spec.first_role("entity_label")
    return entity, label


_CODE_NAME_PAREN = re.compile(r"^([0-9A-Za-z][0-9A-Za-z\-_/.]{0,19})[ 　]*[（(][ 　]*(.+?)[ 　]*[)）]$")
_CODE_NAME_SPACE = re.compile(r"^([0-9A-Za-z][0-9A-Za-z\-_/.]{0,19})[ 　]+(\S.*)$")
# 「名前（番号）」の並び（「Oxideエッチャ 2号機（ETC-302）」）。番号が後ろにあっても同じ設備として扱う。
_NAME_CODE_PAREN = re.compile(r"^(.+?)[ 　]*[（(][ 　]*([0-9A-Za-z][0-9A-Za-z\-_/.]{0,19})[ 　]*[)）]$")
_HAS_DIGIT = re.compile(r"\d")
_HAS_ALPHA = re.compile(r"[A-Za-z]")
_CODE_ONLY = re.compile(r"^[0-9A-Za-z][0-9A-Za-z\-_/.]*$")
# 番号の後ろの括弧が名前ではなく但し書きのとき（「IMP-602（推定）」）。番号だけを残す
_QUALIFIERS = {"推定", "仮", "予定", "調査中", "不明", "未定", "確認中", "暫定", "候補", "要確認", "代替", "予備", "旧", "新"}
# 設備の列の値がこれだけなら、設備が決まっていない（設備ごとのファイル分けに入れない）
_PLACEHOLDER_ENTITIES = {"推定", "仮", "予定", "調査中", "不明", "未定", "確認中", "暫定", "要確認",
                         "〃", "′′", "同上", "々", "仝"}  # 補えなかった「上と同じ」の記号も設備ではない


def split_entity_code(text) -> tuple[str, str]:
    """「ETC-302(OXIDEエッチャ 2号機)」「CVD-203 W-CVD 3号機」「Oxideエッチャ 2号機（ETC-302）」→ (設備番号, 名前)。

    「IMP-602（推定）」のように括弧が但し書きなら (番号, "")。分けられなければ (原文, "")。設備名の列がない台帳で、同じ設備が「番号だけ」「番号＋名前」「名前＋番号」と
    揺れるとファイル分けが割れるため、番号にそろえる。
    """
    s = " ".join(str(text or "").split())
    # 後ろの括弧の番号を先に見る: 「Fab1 OHTシステム（OHT-801）」の番号は Fab1 ではなく OHT-801
    for pattern in (_NAME_CODE_PAREN, _CODE_NAME_PAREN, _CODE_NAME_SPACE):
        m = pattern.match(s)
        if not m:
            continue
        if pattern is _NAME_CODE_PAREN:
            code, name = m.group(2), m.group(1).strip()
        else:
            code, name = m.group(1), m.group(2).strip()
        if not name or not _HAS_DIGIT.search(code) or not _HAS_ALPHA.search(code):
            continue
        if pattern is _CODE_NAME_PAREN and name in _QUALIFIERS:
            return code, ""  # 番号は残し、但し書きは設備名にしない
        if all(_CODE_ONLY.match(p) for p in re.split(r"[\s,、/]+", name) if p):
            continue  # 名前の側も番号だけ（「CMP-101 / CMP-102」）なら分けない
        return code, name
    return s, ""


def entity_value(values: dict, spec: TableSpec) -> tuple[str, str]:
    """(設備番号, 設備名)。設備名の列がない code 列では「番号(名前)」を分ける（ファイル分けの単位をそろえる）。"""
    entity, label = entity_columns(spec)
    if entity is None:
        return "", ""
    raw = " ".join(str(values.get(entity.key) or "").split())
    if unicodedata.normalize("NFKC", raw).strip("（）() ") in _PLACEHOLDER_ENTITIES:
        return "", ""  # 「調査中」だけの値は設備ではない（記録の本文には原文のまま出る）
    if label is not None:
        return raw, " ".join(str(values.get(label.key) or "").split())
    if entity.type == "code" and raw:
        code, name = split_entity_code(raw)
        if name or code != raw:
            return code, name
    return raw, ""


def entity_display(values: dict, spec: TableSpec) -> tuple[str, str, str]:
    """(設備番号, 設備名, 表示「設備名（設備番号）」)。"""
    eid, name = entity_value(values, spec)
    if eid and name and name != eid:
        return eid, name, f"{name}（{eid}）"
    return eid, name, eid or name


# ====================================================================================================
# 元 tables/mapping.py
# 列の対応づけの候補（見出しと先頭のデータから、キー・型・単位・役割・md での扱いを推す）。
#
# 照合の順番: 表用の標準キー辞書（tables.dictionary）→ 似た文字列（候補として出すだけ）。
# 取り込み設定は保存しないので（利用者の指示 2026-09-20）、保存済みの設定との照合は行わない。
# ====================================================================================================

_LOG_LINE_RE = re.compile(
    r"^\s*[【\[(（<]?\s*(\d{4}[/\-.年])?\d{1,2}[/\-.月]\d{1,2}日?|^\s*(R|令和)\d{1,2}\.\d{1,2}\.\d{1,2}",
    re.M,
)
_MERGED = object()  # 結合範囲の左上以外のセル（空欄ではない扱い）
_NUMBER_TYPES = {"number"}
_DATE_TYPES = {"date", "datetime"}

@dataclass
class ColumnSuggestion:
    index: int
    header: str
    unit: str
    key: str | None
    display: str
    type: str
    role: str
    examples: list[str]
    type_error_rate: float
    blank_rate: float
    md: str  # body/attribute/omit
    fill_down_blank: bool = False
    log: bool = False
    matched_by: str = ""  # dictionary/similar/none
    inferred_type: str = ""  # 値から推定した型（type と違えば画面で知らせる）
    # md が "omit" になった理由（画面の説明用）: blank（空欄だけ）/template（取り込み設定のとおり）
    omit_reason: str = ""


def suggest_columns(headers: list[str], sample_rows) -> list[ColumnSuggestion]:
    """見出しと先頭のデータ行から、列ごとの標準キー・型・役割などの候補を作る。

    sample_rows: SourceRow のリスト、または値のリストのリスト。
    見出しも値も無い列（表がA列から始まらないときの左の空列など）は、列の一覧に出さない。
    """
    grid = [_row_values(r) for r in sample_rows]
    out: list[ColumnSuggestion] = []
    for i, header in enumerate(headers):
        unit = split_header_unit(header)[1]
        values = [row[i] if i < len(row) else (None, "") for row in grid]
        if _is_empty_column(header, values):
            continue
        stats = _column_stats(values)

        found = lookup_header(header)
        if found and found[1] == "similar" and found[0].role == "person" and _looks_like_prose(values, stats):
            # 「対応」「作業」が似た語の「対応者」「作業者」に当たっても、値が文章なら人名の列ではない
            found = None
        if found:
            std, how = found
            s = ColumnSuggestion(
                index=i, header=header, unit=unit or std.unit, key=std.key,
                display=_display_name(header), type=std.type, role=std.role,
                examples=stats["examples"], type_error_rate=0.0, blank_rate=stats["blank_rate"],
                md=std.md, matched_by=how, inferred_type=stats["type"],
            )
        else:
            role = "text" if stats["type"] == "text" else "attribute"
            s = ColumnSuggestion(
                index=i, header=header, unit=unit, key=None, display=_display_name(header), type=stats["type"], role=role,
                examples=stats["examples"], type_error_rate=0.0, blank_rate=stats["blank_rate"],
                md="body" if role == "text" else "attribute", matched_by="none", inferred_type=stats["type"],
            )
        if s.md == "omit":
            s.omit_reason = "dictionary"
        s.type_error_rate = _type_error_rate(values, s.type)
        if stats["log_like"] and s.role in ("text", "log", "attribute") and s.type == "text":
            s.role, s.log = "log", True
        if s.role == "log":
            s.log = True
        if s.blank_rate >= 1.0:
            s.md, s.omit_reason = "omit", "blank"
        if s.role in ("entity", "entity_label") and _looks_filled_down(values):
            s.fill_down_blank = True
        out.append(s)
    _keep_one_log(out)
    _dedupe_keys(out)
    return out


def _keep_one_log(out: list[ColumnSuggestion]) -> None:
    """役割「経過の記録」（log）は1列だけ。辞書で当たった列、なければ最初の列を残し、
    ほかは長文にする（何も変えずに保存しただけで「経過の記録の列は1つだけに」と断られないように）。"""
    logs = [s for s in out if s.role == "log"]
    if len(logs) <= 1:
        return
    rank = {"dictionary": 0, "similar": 1}
    best = min(logs, key=lambda s: (rank.get(s.matched_by, 3), s.index))
    for s in logs:
        if s is not best:
            s.role, s.log = "text", False
            if s.md != "omit":
                s.md = "body"


# ---- 内部 ----

def _display_name(header: str) -> str:
    """表示名の候補。2段見出しは下段、単位は除く（「交換部品_単価(円)」→「単価」）。"""

    if is_month_label(header):
        return header  # クロス集計の年月見出しは年を残す
    lower = header.rsplit("_", 1)[-1] if "_" in header else header
    name = split_header_unit(lower)[0]
    return name or header


def _row_values(row) -> list[tuple[object, str]]:
    cells = getattr(row, "cells", None)
    if cells is not None:
        index = getattr(row, "index", 0)
        return [
            (_MERGED, "") if (not c.text and c.merged_anchor and c.merged_anchor != (index, i + 1)) else (c.value, c.text)
            for i, c in enumerate(cells)
        ]
    return [(v, cell_text(v)) for v in row]


def _is_empty_column(header: str, values: list[tuple[object, str]]) -> bool:
    """見出しが空欄で、読み取った行のどこにも値が無い列か。

    見出しの空欄は、この時点では「列3」のような仮の名前になっている（tables.detect._dedupe）。
    値を見る深さは、列の一覧を作るために読んだデータ行（画面からは先頭200行）そのもの。
    追加で読み直さずに済み、例（examples）や空欄率を出すのと同じ範囲なので、
    画面に出る中身と食い違わない。空欄でも下の方に値が出てくる列は、この範囲で見つかれば残る。
    見出しが空でも値がある列（別のところで問題として知らせる）は残す。
    """
    h = str(header or "").strip()
    if h and not (h.startswith("列") and h[1:].isdigit()):
        return False
    return all(_is_blank(value, text) for value, text in values)


def _is_blank(value, text: str) -> bool:
    if value is _MERGED:
        return False
    return not text or text.strip() in NA_TOKENS


def _looks_like_prose(values: list[tuple[object, str]], stats: dict) -> bool:
    """値が人名ではなく文章か（長文、または平均で人名より明らかに長い）。"""
    if stats["type"] == "text":
        return True
    texts = [t for v, t in values if not _is_blank(v, t) and v is not _MERGED]
    return bool(texts) and sum(len(t) for t in texts) / len(texts) > 15


def _column_stats(values: list[tuple[object, str]]) -> dict:
    total = len(values)
    filled = [(v, t) for v, t in values if not _is_blank(v, t)]
    nonblank = [(v, t) for v, t in filled if v is not _MERGED]
    kinds = Counter(value_kind(v, t) for v, t in nonblank)
    examples: list[str] = []
    for _v, t in nonblank:
        short = t if len(t) <= 60 else t[:60] + "…"
        if short not in examples:
            examples.append(short)
        if len(examples) >= 3:
            break
    log_hits = sum(1 for _v, t in nonblank if len(_LOG_LINE_RE.findall(t)) >= 2)
    return {
        "type": _infer_type(kinds, nonblank),
        "blank_rate": round(1 - len(filled) / total, 3) if total else 1.0,
        "examples": examples,
        "log_like": bool(nonblank) and log_hits >= max(2, 0.3 * len(nonblank)),
    }


def _infer_type(kinds: Counter, nonblank: list[tuple[object, str]]) -> str:
    n = sum(kinds.values())
    if not n:
        return "string"
    if kinds.get("text", 0) >= 0.2 * n:
        return "text"
    top, count = kinds.most_common(1)[0]
    if top in ("date", "datetime"):
        return "datetime" if kinds.get("datetime", 0) >= 0.2 * n else "date"
    if top == "time":
        return "time"
    if top == "number":
        # 先頭ゼロ付きの数字（コード）が混ざるならコード
        return "code" if kinds.get("code", 0) >= 0.2 * n else "number"
    if top == "code":
        return "code"
    distinct = len({unicodedata.normalize("NFKC", t) for _v, t in nonblank})
    if n >= 20 and distinct <= 12 and distinct <= 0.2 * n:
        return "enum"
    return "string"


def _type_error_rate(values: list[tuple[object, str]], type_: str) -> float:
    nonblank = [(v, t) for v, t in values if v is not _MERGED and not _is_blank(v, t)]
    if not nonblank:
        return 0.0
    if type_ == "number":
        ok_kinds = _NUMBER_TYPES
    elif type_ in ("date", "datetime"):
        ok_kinds = _DATE_TYPES
    elif type_ == "time":
        ok_kinds = {"time", "datetime"}
    elif type_ == "code":
        ok_kinds = {"code", "number", "string"}
    else:
        return 0.0
    errors = sum(1 for v, t in nonblank if value_kind(v, t) not in ok_kinds)
    return round(errors / len(nonblank), 3)


def _looks_filled_down(values: list[tuple[object, str]]) -> bool:
    """空欄が「上と同じ」の意味で使われていそうか（値の後に空欄が続く）。"""
    texts = [t for v, t in values if v is not _MERGED]
    blanks = sum(1 for t in texts if not t)
    if not texts or blanks < 0.2 * len(texts) or blanks > 0.9 * len(texts):
        return False
    return bool(texts[0])


def _dedupe_keys(items: list[ColumnSuggestion]) -> None:
    """同じ標準キーが複数の列に付いたら、辞書で完全一致した列を優先し、残りは番号を付ける。"""
    by_key: dict[str, list[ColumnSuggestion]] = {}
    for s in items:
        if s.key:
            by_key.setdefault(s.key, []).append(s)
    rank = {"dictionary": 0, "similar": 1}
    for key, group in by_key.items():
        if len(group) < 2:
            continue
        group.sort(key=lambda s: (rank.get(s.matched_by, 3), s.index))
        for n, s in enumerate(group[1:], start=2):
            s.key = f"{key}_{n}"


# ====================================================================================================
# 元 tables/checks.py
# 取り込み結果のチェック（確定を止めるエラーと警告）。
#
# - エラー: 必須の見出しがない、型エラーの割合が上限超え
# - 警告: 型エラー、許可値にない区分、年を補った日付、除外した行、日付が空の行、
#         読めない文字の置き換え、使っていない列、値が空の必須列、合計の照合
# 行ごとの型エラーなどは normalize.read_records が Issue にし、ここでは全体を見たチェックを足す。
# ====================================================================================================

LEVEL_LABELS = {"error": "エラー", "warning": "警告"}


@dataclass
class Issue:
    level: str  # error/warning
    code: str
    message: str
    row: int | None = None
    column: str | None = None

    def to_dict(self) -> dict:
        return asdict(self)


def has_blocking(issues) -> bool:
    return any(_level(i) == "error" for i in issues)


def count_levels(issues) -> dict[str, int]:
    c = Counter(_level(i) for i in issues)
    return {"error": c.get("error", 0), "warning": c.get("warning", 0)}


def _level(issue) -> str:
    return issue.level if isinstance(issue, Issue) else str((issue or {}).get("level"))




def _values(record) -> dict:
    return _get(record, "values", {}) or {}


def run_checks(records, spec, stats) -> list[Issue]:
    """全体を見たチェック。stats は ImportStats か、その dict。"""
    st = stats if isinstance(stats, dict) else stats.to_dict()
    issues: list[Issue] = []

    for display in st.get("missing_required") or []:
        issues.append(Issue("error", "required_missing", f"必須の列「{display}」の見出しが見つかりません"))

    rates = (spec.checks or {}).get("type_error_rate") or {}
    warn, block = float(rates.get("warn", 0.02)), float(rates.get("block", 0.10))
    checked = st.get("type_checked") or {}
    for key, errors in sorted((st.get("type_errors") or {}).items()):
        total = checked.get(key) or 0
        if not total or not errors:
            continue
        rate = errors / total
        col = spec.column(key)
        label = col.display if col else key
        if rate > block:
            issues.append(Issue("error", "type_error_rate",
                                f"列「{label}」で値を変換できない行が多すぎます（{errors}/{total}件、{rate:.0%}）。"
                                "列の型か見出しの位置を確認してください", column=label))
        elif rate >= warn:
            issues.append(Issue("warning", "type_error_rate",
                                f"列「{label}」で値を変換できない行があります（{errors}/{total}件、{rate:.1%}）", column=label))

    dups = st.get("duplicate_keys") or {}
    if dups:
        sample = "、".join(f"{k}（{n}件）" for k, n in list(sorted(dups.items()))[:5])
        issues.append(Issue("warning", "duplicate_key",
                            f"記録キーが同じ行が{len(dups)}種類あります: {sample}。出現順の番号で区別します"))

    date_key = spec.date_key
    if date_key:
        undated = sum(1 for rec in records if not _values(rec).get(date_key))
        if undated:
            issues.append(Issue("warning", "undated", f"日付が空の行が{undated}件あります（「日付なし」のファイルに入れます）"))
        issues += _date_outliers(records, date_key)

    # 許可値
    for col in spec.columns:
        if col.allowed:
            allowed = set(col.allowed)
            bad = Counter(str(_values(r).get(col.key)) for r in records
                          if _values(r).get(col.key) not in (None, "") and str(_values(r).get(col.key)) not in allowed)
            if bad:
                detail = "、".join(f"{v}（{n}件）" for v, n in bad.most_common(8))
                issues.append(Issue("warning", "not_allowed",
                                    f"列「{col.display}」に選択肢にない値があります: {detail}", column=col.display))
        if col.required and col.key in (st.get("positions") or {}):
            empty = sum(1 for r in records if _values(r).get(col.key) in (None, ""))
            if empty:
                issues.append(Issue("warning", "required_empty",
                                    f"必須の列「{col.display}」が空の行が{empty}件あります", column=col.display))
    if st.get("year_inferred"):
        issues.append(Issue("warning", "year_inferred",
                            f"年のない日付に、年を補った行が{st['year_inferred']}件あります（{st.get('year_context_label') or '年度'}から）"))
    if st.get("year_missing"):
        issues.append(Issue("warning", "year_missing",
                            f"年のない日付で、年を補えなかった行が{st['year_missing']}件あります。"
                            "タイトル・シート名・ファイル名に年度がないためです"))
    for reason, n in sorted((st.get("excluded") or {}).items()):
        if reason in ("非表示の行", "取り消し線の行"):
            issues.append(Issue("warning", "excluded_rows", f"{reason}を{n}行除外しました"))
    if st.get("included_hidden"):
        issues.append(Issue("warning", "included_rows", f"非表示・取り消し線の行を{st['included_hidden']}行取り込みました"))
    if st.get("error_values"):
        issues.append(Issue("warning", "excel_error", f"Excelのエラー値（#N/A など）を空欄として扱ったセルが{st['error_values']}個あります"))
    if st.get("uncached_formulas"):
        by_col = st["uncached_formulas"]
        names = "、".join(f"「{k}」" for k in list(by_col)[:5])
        issues.append(Issue("warning", "uncached_formula",
                            f"Excelで計算されていない数式のセルが{sum(by_col.values())}個あります（列{names}）。"
                            "値が空欄として読まれています。Excelで開いて保存し直してから取り込んでください"))
    if st.get("unclosed_quote_row"):
        row = st["unclosed_quote_row"]
        issues.append(Issue("error", "unclosed_quote",
                            f"{row}行目の \" が閉じていないため、以降の行が1つの値になっています。"
                            "元のファイルの \" を直してから、もう一度取り込んでください", row=row))
    if st.get("long_record_row"):
        row = st["long_record_row"]
        issues.append(Issue("warning", "long_record",
                            f"{row}行目の値が{MAX_RECORD_LINES}行以上あります。\" の閉じ忘れでないか確認してください", row=row))
    if st.get("replaced_rows"):
        rows = st["replaced_rows"]
        issues.append(Issue("warning", "replaced_chars",
                            f"読めない文字を〓に置き換えた行が{len(rows)}行あります（{', '.join(str(r) for r in rows[:10])}行目など）"))
    if st.get("unused_headers"):
        names = "、".join(st["unused_headers"][:10])
        issues.append(Issue("warning", "unused_headers", f"設定にない列があります（取り込みません）: {names}"))
    for missing in st.get("missing_optional") or []:
        issues.append(Issue("warning", "column_missing", f"列「{missing}」の見出しが見つかりません（空欄として扱います）"))
    tolerance = float((spec.checks or {}).get("reconcile_tolerance") or 0)
    for rc in st.get("reconcile") or []:
        diff = abs(float(rc.get("expected") or 0) - float(rc.get("actual") or 0))
        if diff > tolerance + 1e-6:
            issues.append(Issue("warning", "reconcile",
                                f"{rc.get('label') or '合計'}が合いません（列「{rc.get('column')}」: 表の値 {_num(rc.get('expected'))}、"
                                f"行の合計 {_num(rc.get('actual'))}）", row=rc.get("row"), column=rc.get("column")))
    if not records and not any(i.level == "error" for i in issues):
        # 記録が1件も無いまま確定すると、中身の無い zip を渡してしまう。確定を止めて理由も出す
        excluded = sum((st.get("excluded") or {}).values())
        extra = f"（取り消し線・非表示などで{excluded}行を除外しました）" if excluded else ""
        issues.append(Issue("error", "no_records",
                            f"取り込める行がありません{extra}。表の範囲か元のファイルを見直してください"))
    return issues


DATE_OUTLIER_YEARS = 5  # 記録の年の中央値からこれより離れた日付は、打ち間違いの疑い


def _date_outliers(records, date_key: str) -> list[Issue]:
    """ほかの記録から何年も離れた日付（2025年のデータに 2052年 など）。記録ファイルは年月ごとに分かれるので、
    離れた月に1件だけの記録ファイルができる原因になる。"""
    dated = []
    for rec in records:
        s = str(_values(rec).get(date_key) or "")
        if len(s) >= 4 and s[:4].isdigit():
            dated.append((int(s[:4]), s[:10], (_get(rec, "source") or {}).get("row")))
    if len(dated) < 3:
        return []
    years = sorted(y for y, _d, _r in dated)
    median = years[len(years) // 2]
    far = [(d, r) for y, d, r in dated if abs(y - median) > DATE_OUTLIER_YEARS]
    if not far:
        return []
    sample = "、".join(f"{r}行目（{d}）" if r else d for d, r in far[:5]) + ("など" if len(far) > 5 else "")
    return [Issue("warning", "date_outlier",
                  f"ほかの記録から{DATE_OUTLIER_YEARS}年より離れた日付が{len(far)}件あります: {sample}。"
                  "年の打ち間違いでないか確認してください", row=far[0][1])]


def _num(value) -> str:
    try:
        f = float(value)
    except (TypeError, ValueError):
        return str(value)
    return f"{int(f):,}" if f.is_integer() else f"{f:,.2f}"


# ====================================================================================================
# 元 tables/normalize.py
# 行の読み込みと正規化（ステップ5）。
#
# 1行ずつ: 行の分類 → 継続行の連結 → 結合範囲と「空欄＝上と同じ」列の補完 →
# 型の変換（和暦・年のない日付・8桁日付・6桁時刻・△▲/末尾マイナス・桁区切り・%・timedelta→分・単位換算）→
# 記録キー（出現順の番号付き）。値は NFKC だけで揃える（名寄せ辞書は使わない）。
# ただし md に出す文章の値は丸数字などの囲み文字（①Ⓐ㋐㊤…）を原文どおり残す（①→1 だと番号と本文の区切りが消えるため）。
# 比較・照合に使う正規化（日付・数値の解析、NA 判定、コードの突き合わせ）は従来どおり NFKC だけをかける。
# 数値・日付以外の文章列は元の値を残さない（NFKC と空白の畳み込みだけなので）。
# ====================================================================================================

MAX_ROW_ISSUES = 2000
PROGRESS_EVERY = 500

_SPACES = re.compile(r"[^\S\n]+")
_ALL_SPACES = re.compile(r"\s+")
_WEEKDAY = re.compile(r"\s*[(（]\s*[月火水木金土日](?:曜日?)?\s*[)）]\s*")
# 秒の小数（SQL Server などの「.000」）と時差（「Z」「+09:00」）は読み捨てる（秒までで足り、時刻は書かれたまま）
_TIME_PART = (r"(?:\s*T?\s*(\d{1,2})\s*[:時]\s*(\d{1,2})\s*(?:[:分]\s*(\d{1,2})(?:\.\d{1,7})?\s*秒?)?分?"
              r"(?:\s*(?:Z|[+\-]\d{2}:?\d{2}))?)?")
_YMD = re.compile(r"^(\d{4})\s*[/\-.年]\s*(\d{1,2})\s*[/\-.月]\s*(\d{1,2})\s*日?" + _TIME_PART + r"$")
_ERA_YMD = re.compile(r"^(令和|平成|昭和|R|H|S)\s*(\d{1,2}|元)\s*[./年\-]\s*(\d{1,2})\s*[./月\-]\s*(\d{1,2})\s*日?"
                      + _TIME_PART + r"$", re.I)
_MD = re.compile(r"^(\d{1,2})\s*[/月]\s*(\d{1,2})\s*日?" + _TIME_PART + r"$")
_YMD8 = re.compile(r"^((?:19|20)\d{2})(0[1-9]|1[0-2])(0[1-9]|[12]\d|3[01])(?:\s*(\d{2})(\d{2})(\d{2})?|\s+(\d{1,2}):(\d{2})(?::(\d{2}))?)?$")
_TIME_TEXT = re.compile(r"^(\d{1,2})\s*[:時]\s*(\d{1,2})\s*(?:[:分]\s*(\d{1,2})\s*秒?)?分?$")
_TIME_DIGITS = re.compile(r"^(\d{2})(\d{2})(\d{2})?$")
_NUMBER = re.compile(
    r"^(?P<sign>[△▲\-−+])?\s*[¥￥]?\s*(?P<int>\d{1,3}(?:,\d{3})+|\d+)?(?P<dec>\.\d+)?(?:[eE](?P<exp>[+\-]?\d+))?"
    r"\s*(?P<tail>-)?\s*(?P<unit>[^\d\s,.\-+]{1,6})?$")
_ERA_BASE = {"令和": 2018, "R": 2018, "平成": 1988, "H": 1988, "昭和": 1925, "S": 1925}

# 単位の換算表（列の単位ごと。値は「元の単位1つが列の単位でいくつか」）
_UNIT_ALIASES = {"h": "時間", "hr": "時間", "hrs": "時間", "hour": "時間", "hours": "時間", "H": "時間",
                 "min": "分", "mins": "分", "sec": "秒", "yen": "円", "¥": "円", "￥": "円",
                 "s": "秒", "時": "時間", "％": "%"}
_BUILTIN_CONVERSIONS = {
    "分": {"分": 1, "時間": 60, "秒": 1 / 60, "日": 1440},
    "時間": {"時間": 1, "分": 1 / 60, "秒": 1 / 3600, "日": 24},
    "秒": {"秒": 1, "分": 60, "時間": 3600},
    "円": {"円": 1, "千円": 1000, "万円": 10000, "百万円": 1000000},
    "千円": {"千円": 1, "円": 0.001, "万円": 10},
    "%": {"%": 1},
}

_TYPE_LABELS = {"date": "日付", "datetime": "日時", "time": "時刻", "number": "数値"}


# ---- データ構造 ---------------------------------------------------------------------

@dataclass
class RecordRow:
    key: str
    values: dict[str, object]
    originals: dict[str, str]
    source: dict  # {"file", "sheet", "row"}
    warnings: list[str] = field(default_factory=list)

    def to_dict(self) -> dict:
        return {"key": self.key, "values": self.values, "originals": self.originals,
                "source": self.source, "warnings": self.warnings}


# 「上と同じ」の記号。セル全体がこれだけなら直前のデータ行の値で補う（NFKC で ″ は ′′ になる）
DITTO_MARKS = {"〃", "″", "′′", "同上", "々", "仝"}
# 記録番号（key）と担当者（person）も補う。key が「〃」のままだと見出し・出典・記録キーが「〃」になる
DITTO_ROLES = ("key", "date", "entity", "entity_label", "category", "attribute", "person")
_ISO_DATE_RE = re.compile(r"\d{4}-\d{2}-\d{2}")


def is_ditto(text) -> bool:
    """セル全体が「〃」「同上」などの「上と同じ」の記号か。"""
    t = str(text or "").strip()
    return bool(t) and (t in DITTO_MARKS or unicodedata.normalize("NFKC", t).strip() in DITTO_MARKS)


@dataclass
class ImportStats:
    file: str = ""
    sheet: str = ""
    rows_scanned: int = 0
    data_rows: int = 0
    records: int = 0
    blank_rows: int = 0
    excluded: dict = field(default_factory=dict)  # 理由 → 行数
    continuation_merged: int = 0
    subtotal_rows: int = 0
    included_hidden: int = 0
    type_errors: dict = field(default_factory=dict)
    type_checked: dict = field(default_factory=dict)
    year_inferred: int = 0
    year_missing: int = 0
    year_context: dict = field(default_factory=dict)
    year_context_label: str = ""
    error_values: int = 0
    filled_down: dict = field(default_factory=dict)
    ditto_filled: dict = field(default_factory=dict)  # 列キー → 「〃」「同上」を直前の行の値で補った数
    unused_headers: list = field(default_factory=list)
    missing_required: list = field(default_factory=list)
    missing_optional: list = field(default_factory=list)
    positions: dict = field(default_factory=dict)
    duplicate_keys: dict = field(default_factory=dict)
    reconcile: list = field(default_factory=list)
    replaced_rows: list = field(default_factory=list)
    uncached_formulas: dict = field(default_factory=dict)  # 列の表示名 → Excel で計算されていない数式のセルの数
    unclosed_quote_row: int | None = None  # CSV の " が閉じていないため、以降を1つの値として読んだ行
    long_record_row: int | None = None  # CSV の1つの値がとても多くの行にまたがる行（閉じ忘れの疑い）
    date_min: str | None = None
    date_max: str | None = None
    months: dict = field(default_factory=dict)
    layout_warnings: list = field(default_factory=list)
    row_issues_truncated: int = 0
    key_columns: list = field(default_factory=list)
    elapsed_sec: float = 0.0

    def to_dict(self) -> dict:
        return asdict(self)


# ---- 年・年度の手がかり -------------------------------------------------------------------

_FY_PATTERNS = [
    (re.compile(r"(?:FY|ＦＹ)\s*'?(\d{4}|\d{2})(?!\d)", re.I), lambda m: _year4(m.group(1))),
    (re.compile(r"(?<!\d)(\d{4})\s*年度"), lambda m: int(m.group(1))),
    (re.compile(r"(令和|平成|R|H)\s*(\d{1,2}|元)\s*年度"), lambda m: _era_year(m.group(1), m.group(2))),
]
_YM_PATTERNS = [
    re.compile(r"(?<!\d)(\d{4})\s*年\s*(\d{1,2})\s*月"),
    re.compile(r"(?<!\d)(\d{4})[/\-._](\d{1,2})(?![\d/\-.])"),
    re.compile(r"(?<!\d)(\d{4})(0[1-9]|1[0-2])(?!\d)"),
]
_Y_PATTERN = re.compile(r"(?<!\d)(\d{4})\s*年(?!度)")


def _year4(text: str) -> int:
    n = int(text)
    return n + 2000 if n < 100 else n


def _era_year(era: str, num: str) -> int:
    base = _ERA_BASE.get(era.upper() if len(era) == 1 else era, 2018)
    return base + (1 if num == "元" else int(num))


def find_year_context(sources: list[tuple[str, str]]) -> dict:
    """タイトル行・シート名・ファイル名から年度・年月を探す。sources: [(出どころの名前, 文字列)]（優先順）。

    戻り値: {"fiscal_year", "year", "month", "source"}（見つからなければ値は None）
    """
    for label, text in sources:
        s = unicodedata.normalize("NFKC", str(text or ""))
        if not s.strip():
            continue
        for pattern, conv in _FY_PATTERNS:
            m = pattern.search(s)
            if m:
                fy = conv(m)
                if 1950 <= fy <= 2100:
                    return {"fiscal_year": fy, "year": None, "month": None, "source": label}
        for pattern in _YM_PATTERNS:
            m = pattern.search(s)
            if m and 1950 <= int(m.group(1)) <= 2100 and 1 <= int(m.group(2)) <= 12:
                return {"fiscal_year": None, "year": int(m.group(1)), "month": int(m.group(2)), "source": label}
        m = _Y_PATTERN.search(s)
        if m and 1950 <= int(m.group(1)) <= 2100:
            return {"fiscal_year": None, "year": int(m.group(1)), "month": None, "source": label}
    return {"fiscal_year": None, "year": None, "month": None, "source": ""}


def year_context_for(source, sheet: str, layout, file_name: str) -> dict:
    """表の上のタイトル行 → シート名 → ファイル名の順で年度・年月を探す。"""
    sources: list[tuple[str, str]] = []
    top = (layout.header_rows[0] - 1) if getattr(layout, "header_rows", None) else 0
    if top > 0:
        texts = []
        for row in source.rows(sheet, 1, min(top, 20)):
            texts.append(" ".join(c.text for c in row.cells if c.text))
        sources.append(("タイトル行", "\n".join(texts)))
    if getattr(source, "kind", "") == "excel":
        sources.append(("シート名", sheet))
    sources.append(("ファイル名", Path(file_name or "").stem))
    return find_year_context(sources)


# ---- 値の変換 ---------------------------------------------------------------------------

@dataclass
class ConvertContext:
    fiscal_year: int | None = None
    year: int | None = None
    month: int | None = None
    fiscal_start: int = 4
    date1904: bool = False
    na_tokens: frozenset = frozenset()


def _nfkc(text) -> str:
    """比較・照合用の NFKC（囲み文字も ①→1 にそろえる）。日付・数値の解析、NA 判定、キーの照合に使う。"""
    return unicodedata.normalize("NFKC", str(text or ""))


def nfkc_text(text, keep_enclosed: bool = True) -> str:
    """NFKC＋行内の空白の畳み込み（改行は残す）。

    md に出す値は囲み文字（①Ⓐ㋐㊤…）を原文どおり残す（core.mdtext.nfkc_keep_enclosed）。
    コードのように比較・突き合わせに使う値は keep_enclosed=False で従来どおり NFKC だけをかける。
    """
    if text is None:
        return ""
    base = nfkc_keep_enclosed(str(text or "")) if keep_enclosed else _nfkc(text)
    s = base.replace("_x000D_", "").replace("\r\n", "\n").replace("\r", "\n")
    if "\n" not in s:  # 1行なら行に分けない（結果は同じ）
        return _SPACES.sub(" ", s).strip()
    lines = [_SPACES.sub(" ", line).strip() for line in s.split("\n")]
    return "\n".join(lines).strip("\n")


def _fmt_dt(y: int, mo: int, d: int, hh=None, mm=None, ss=None) -> tuple[str | None, str | None]:
    try:
        day = date(int(y), int(mo), int(d))
    except (TypeError, ValueError):
        return None, None
    if hh is None or hh == "":
        return day.isoformat(), None
    try:
        t = time(int(hh), int(mm or 0), int(ss or 0))
    except ValueError:
        return None, None
    return day.isoformat(), _fmt_time(t)


def _fmt_time(t: time) -> str:
    return t.strftime("%H:%M:%S") if t.second else t.strftime("%H:%M")


def _join_dt(d: str | None, t: str | None, type_: str) -> str | None:
    if d is None:
        return None
    if type_ == "datetime" and t and t not in ("00:00",):
        return f"{d} {t}"
    return d


def parse_date_text(text: str, cctx: ConvertContext) -> tuple[str | None, str | None, str | None]:
    """文字列の日付・日時。戻り値: (日付, 時刻, 印)。印は year_inferred / year_missing / None。"""
    s = _WEEKDAY.sub(" ", _nfkc(text)).strip()
    s = re.sub(r"\s+", " ", s)
    m = _YMD.match(s)
    if m:
        d, t = _fmt_dt(m.group(1), m.group(2), m.group(3), m.group(4), m.group(5), m.group(6))
        return d, t, None
    m = _YMD8.match(s)
    if m:
        hh, mm, ss = (m.group(4), m.group(5), m.group(6)) if m.group(4) else (m.group(7), m.group(8), m.group(9))
        d, t = _fmt_dt(m.group(1), m.group(2), m.group(3), hh, mm, ss)
        return d, t, None
    m = _ERA_YMD.match(s)
    if m:
        era = m.group(1)
        y = _era_year(era.upper() if len(era) == 1 else era, m.group(2))
        d, t = _fmt_dt(y, m.group(3), m.group(4), m.group(5), m.group(6), m.group(7))
        return d, t, None
    m = _MD.match(s)
    if m:
        mo = int(m.group(1))
        if cctx.fiscal_year:
            y = cctx.fiscal_year if mo >= cctx.fiscal_start else cctx.fiscal_year + 1
        elif cctx.year:
            y = cctx.year
        else:
            return None, None, "year_missing"
        d, t = _fmt_dt(y, mo, m.group(2), m.group(3), m.group(4), m.group(5))
        return d, t, ("year_inferred" if d else None)
    return None, None, None


def _excel_serial_date(value: float, date1904: bool) -> datetime | None:
    low = 3000
    if not (low < value < 80000):
        return None
    try:
        dt = from_excel(value, epoch=MAC_EPOCH if date1904 else WINDOWS_EPOCH)
    except (ValueError, OverflowError):
        return None
    return dt if isinstance(dt, datetime) else None


def _parse_time_text(text: str) -> str | None:
    s = _nfkc(text).strip()
    m = _TIME_TEXT.match(s) or _TIME_DIGITS.match(s)
    if not m:
        return None
    try:
        t = time(int(m.group(1)), int(m.group(2)), int(m.group(3) or 0))
    except ValueError:
        return None
    return _fmt_time(t)


def _unit_norm(unit: str) -> str:
    u = _nfkc(unit).strip()
    return _UNIT_ALIASES.get(u, _UNIT_ALIASES.get(u.lower(), u))


def unit_factor(src_unit: str, target_unit: str, conversions: dict | None = None) -> float | None:
    """src_unit の値を target_unit に直す倍率。換算できなければ None。"""
    if not src_unit:
        return 1.0
    conversions = conversions or {}
    for name in (src_unit, _unit_norm(src_unit)):
        if name in conversions:
            return float(conversions[name])
    src, dst = _unit_norm(src_unit), _unit_norm(target_unit)
    if not dst or src == dst:
        return 1.0
    return _BUILTIN_CONVERSIONS.get(dst, {}).get(src)


def _clean_number(v: float) -> int | float:
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, int):
        return v
    v = round(float(v), 9)
    if v.is_integer() and abs(v) < 1e15:
        return int(v)
    return v


def parse_number_text(text: str) -> tuple[float | None, str]:
    """数値の文字列。戻り値: (値, 末尾の単位)。△▲・末尾マイナス・桁区切り・全角数字に対応。"""
    s = _nfkc(text).strip().replace(" ", "")
    if not s:
        return None, ""
    m = _NUMBER.match(s)
    if not m or (m.group("int") is None and m.group("dec") is None):
        return None, ""
    # 指数は文字列のまま float に渡す（10 ** 309 などで OverflowError にしない）。桁あふれ・inf は数値にしない
    number = float((m.group("int") or "0").replace(",", "") + (m.group("dec") or "")
                   + (f"e{m.group('exp')}" if m.group("exp") else ""))
    if not math.isfinite(number):
        return None, ""
    if (m.group("sign") and m.group("sign") in "△▲-−") or m.group("tail"):
        number = -number
    return number, m.group("unit") or ""


def _zero_pad(value: float, number_format: str | None) -> str:
    n = int(value)
    fmt = (number_format or "").strip()
    if fmt and set(fmt) == {"0"} and len(fmt) > 1:
        return str(n).zfill(len(fmt))
    return str(n)


_NON_ASCII = re.compile(r"[^\x00-\x7f]")


def _upper_code(text: str) -> str:
    """code 列の upper。同じセルに「番号＋設備名」が入っている台帳があるので、番号の部分だけ大文字にする。

    「Oxideエッチャ 2号機（ETC-302）」→「Oxideエッチャ 2号機（ETC-302）」（設備名はそのまま）。
    英数字だけの値は今までどおり全体を大文字にする（「etc-302」→「ETC-302」）。
    """
    if not _NON_ASCII.search(text):
        return text.upper()

    code, name = split_entity_code(text)
    if not name or code.upper() == code:
        return text
    low, target = text.lower(), code.lower()
    at = low.find(target) if low.startswith(target) else low.rfind(target)
    if at < 0:
        return text
    return text[:at] + code.upper() + text[at + len(code):]


def convert_cell(value, text: str, col, cctx: ConvertContext, number_format: str | None = None,
                 header_unit: str = "") -> tuple[object, str | None, str | None]:
    """セル1つを列の型に変換する。戻り値: (値, エラー文, 印)。印: year_inferred / year_missing / error_value"""
    if value is None and not text:
        return None, None, None
    stripped = _nfkc(text).strip()
    if isinstance(value, str) or value is None:
        if not stripped:
            return None, None, None
        if stripped in _EXCEL_ERRORS:
            return None, None, "error_value"
        if stripped in cctx.na_tokens:
            return None, None, None
    t = col.type
    if t in ("date", "datetime"):
        return _convert_date(value, text, t, cctx)
    if t == "time":
        return _convert_time(value, text)
    if t == "number":
        return _convert_number(value, text, col, number_format, header_unit)
    if t == "code":
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            out = _zero_pad(value, number_format) if float(value).is_integer() else str(_clean_number(value))
        else:
            out = _SPACES.sub(" ", nfkc_text(text, keep_enclosed=False)).strip()  # コードは突き合わせに使うので NFKC のみ
        if "upper" in col.normalize:
            out = _upper_code(out)
        return out, None, None
    if t == "text":
        out = nfkc_text(text)
        return (out or None), None, None
    # string / enum / status
    out = nfkc_text(text) if "nfkc" in col.normalize or not col.normalize else str(text).strip()
    if "upper" in col.normalize:
        out = out.upper()
    if col.value_map:
        mapped = col.value_map.get(out)
        if mapped is None:
            mapped = {nfkc_text(k): v for k, v in col.value_map.items()}.get(out)
        if mapped is not None:
            out = str(mapped)
    return (out or None), None, None


def _convert_date(value, text, type_, cctx):
    label = _TYPE_LABELS[type_]
    if isinstance(value, date) and value.year < 1900:
        # 日付の表示形式のセルに 0 や負の数（空の計算結果など）。1899年の日付にせず、変換できない値として知らせる
        return text, f"「{text}」を{label}に変換できません", None
    if isinstance(value, datetime):
        d, tm = value.date().isoformat(), (_fmt_time(value.time()) if value.time() != time(0) else None)
        return _join_dt(d, tm, type_), None, None
    if isinstance(value, date):
        return value.isoformat(), None, None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if isinstance(value, int) or float(value).is_integer():
            d, tm, _ = parse_date_text(str(int(value)), cctx)
            if d:
                return _join_dt(d, tm, type_), None, None
        dt = _excel_serial_date(float(value), cctx.date1904)
        if dt is not None:
            tm = _fmt_time(dt.time()) if dt.time() != time(0) else None
            return _join_dt(dt.date().isoformat(), tm, type_), None, None
        return str(_clean_number(value)), f"「{text}」を{label}に変換できません", None
    if isinstance(value, (time, timedelta)):
        return text, f"「{text}」を{label}に変換できません", None
    d, tm, flag = parse_date_text(text, cctx)
    if d:
        return _join_dt(d, tm, type_), None, flag
    if flag == "year_missing":
        return nfkc_text(text), f"「{text}」は年がないため{label}にできません", flag
    return nfkc_text(text), f"「{_short(text)}」を{label}に変換できません", None


def _convert_time(value, text):
    if isinstance(value, datetime):
        return _fmt_time(value.time()), None, None
    if isinstance(value, time):
        return _fmt_time(value), None, None
    if isinstance(value, timedelta):
        secs = int(round(value.total_seconds()))
        if 0 <= secs < 86400:
            return _fmt_time(time(secs // 3600, secs % 3600 // 60, secs % 60)), None, None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        v = float(value)
        if 0 <= v < 1:
            secs = int(round(v * 86400))
            secs = min(secs, 86399)
            return _fmt_time(time(secs // 3600, secs % 3600 // 60, secs % 60)), None, None
        parsed = _parse_time_text(str(int(v))) if v.is_integer() else None
        if parsed:
            return parsed, None, None
        return str(_clean_number(v)), f"「{text}」を時刻に変換できません", None
    parsed = _parse_time_text(text)
    if parsed:
        return parsed, None, None
    return nfkc_text(text), f"「{_short(text)}」を時刻に変換できません", None


def _convert_number(value, text, col, number_format, header_unit):
    target = col.unit or ""
    if isinstance(value, bool):
        return None, f"「{text}」を数値に変換できません", None
    if isinstance(value, timedelta) or isinstance(value, time):
        minutes = value.total_seconds() / 60 if isinstance(value, timedelta) else value.hour * 60 + value.minute + value.second / 60
        factor = unit_factor("分", target or "分", col.unit_conversions)
        if factor is None:
            return None, f"「{text}」（時間）を{target}に換算できません", None
        return _clean_number(minutes * factor), None, None
    if isinstance(value, (int, float)):
        v = float(value)
        if number_format and "%" in number_format:
            v *= 100
        factor = unit_factor(header_unit, target, col.unit_conversions) if header_unit else 1.0
        if factor is None:
            # 文字列のときと同じく、換算できない単位は黙って元の数のまま使わない
            return _clean_number(v), f"「{text}」の単位「{header_unit}」を{target}に換算できません", None
        if not math.isfinite(v * factor):
            return None, f"「{_short(text)}」を数値に変換できません", None
        return _clean_number(v * factor), None, None
    minutes = _duration_minutes(text)
    if minutes is not None and (target or header_unit):
        # CSV に書き出された [h]:mm（「1:30」）や「1時間30分」。Excel の時間のセルと同じく分として換算する
        factor = unit_factor("分", target or header_unit, col.unit_conversions)
        if factor is not None:
            return _clean_number(minutes * factor), None, None
    number, unit = parse_number_text(text)
    if number is None:
        return nfkc_text(text), f"「{_short(text)}」を数値に変換できません", None
    src_unit = unit or header_unit
    factor = unit_factor(src_unit, target, col.unit_conversions)
    if factor is None:
        if unit and not target:
            factor = 1.0  # 列に単位がなければ数値部分だけを使う
        else:
            return nfkc_text(text), f"「{_short(text)}」の単位「{unit or header_unit}」を{target}に換算できません", None
    if not math.isfinite(number * factor):
        return nfkc_text(text), f"「{_short(text)}」を数値に変換できません", None
    return _clean_number(number * factor), None, None


_DURATION_HMS = re.compile(r"^(\d{1,4}):([0-5]\d)(?::([0-5]\d))?$")
_DURATION_JA = re.compile(r"^(\d{1,4})時間(\d{1,2})分$")


def _duration_minutes(text) -> float | None:
    """「1:30」「1:30:00」「1時間30分」を分にする（それ以外は None）。"""
    s = "".join(unicodedata.normalize("NFKC", str(text or "")).split())
    m = _DURATION_HMS.match(s)
    if m:
        return int(m.group(1)) * 60 + int(m.group(2)) + int(m.group(3) or 0) / 60
    m = _DURATION_JA.match(s)
    if m:
        return int(m.group(1)) * 60 + int(m.group(2))
    return None


def _short(text, n: int = 30) -> str:
    s = str(text or "").replace("\n", " ")
    return s if len(s) <= n else s[:n] + "…"


# ---- 読み込み本体 ---------------------------------------------------------------------------

def _is_text_column(col) -> bool:
    return col.type == "text" or col.role in ("text", "log")


def read_records(source, source_opts: dict | None, layout, spec: TableSpec, on_progress=None) -> tuple[list[RecordRow], list[Issue], ImportStats]:
    """表ソースの全行を読み、正規化した記録・問題・件数を返す。

    source_opts: {"sheet", "file_name", "fiscal_year"（年度を手で指定するとき）, "key_columns"（継続行の判定列、0始まり）}
    on_progress(done, total): 一定行ごとに呼ぶ（ジョブの進捗・中止確認に使う）
    """
    started = _time.monotonic()
    opts = dict(source_opts or {})
    sheet = opts.get("sheet") or layout.sheet
    file_name = opts.get("file_name") or getattr(source, "file_name", "") or ""
    is_csv = getattr(source, "kind", "") == "csv"
    stats = ImportStats(file=file_name, sheet="" if is_csv else sheet,
                        layout_warnings=list(getattr(layout, "warnings", []) or []))
    issues: list[Issue] = []

    res = resolve_columns(spec, list(layout.headers))
    stats.positions = dict(res.positions)
    stats.unused_headers = list(res.unused_headers)
    stats.missing_required = list(res.missing_required)
    stats.missing_optional = [d for d in res.missing if d not in res.missing_required]

    ctx_year = year_context_for(source, sheet, layout, file_name)
    if opts.get("fiscal_year"):
        ctx_year = {"fiscal_year": int(opts["fiscal_year"]), "year": None, "month": None, "source": "指定"}
    stats.year_context = ctx_year
    if ctx_year.get("fiscal_year"):
        stats.year_context_label = f"{ctx_year['source']}の{ctx_year['fiscal_year']}年度"
    elif ctx_year.get("year"):
        stats.year_context_label = f"{ctx_year['source']}の{ctx_year['year']}年"
    na = frozenset(_nfkc(t).strip() for t in (spec.na_tokens or []))
    cctx = ConvertContext(fiscal_year=ctx_year.get("fiscal_year"), year=ctx_year.get("year"), month=ctx_year.get("month"),
                          fiscal_start=int(spec.fiscal_year_start_month or 4),
                          date1904=bool(getattr(source, "date1904", False)), na_tokens=na)

    columns = [(c, res.positions[c.key]) for c in spec.columns if c.key in res.positions]
    header_units = {c.key: split_header_unit(layout.headers[pos])[1] for c, pos in columns}
    measure_cols = [(c, pos) for c, pos in columns if c.type == "number"]
    policies = spec.exclude or {}

    anchors: dict[tuple[int, int], tuple[object, str, str | None]] = {}
    prev_raw: dict[str, tuple[object, str, str | None]] = {}
    last_row_raw: dict[str, tuple[object, str, str | None]] = {}  # 直前のデータ行の値（「〃」の補完に使う）
    pending: list[RecordRow] = []
    last_record: RecordRow | None = None
    running: dict[int, float] = {}
    grand: dict[int, float] = {}
    total_rows = max(1, layout.data_end - layout.data_start + 1)
    row_issue_count = 0

    def add_row_issue(issue: Issue):
        nonlocal row_issue_count
        row_issue_count += 1
        if row_issue_count <= MAX_ROW_ISSUES:
            issues.append(issue)
        else:
            stats.row_issues_truncated += 1

    def cell_raw(row, pos: int) -> tuple[object, str, str | None]:
        cell = row.cell(pos)
        if cell is None:
            return None, "", None
        if not cell.text and cell.merged_anchor and cell.merged_anchor != (row.index, pos + 1):
            found = anchors.get(cell.merged_anchor)
            if found is not None:
                return found
        return cell.value, cell.text, cell.number_format

    def convert(col, raw, row_index: int, record_warnings: list[str], originals: dict, header_unit: str = ""):
        value, text, number_format = raw
        out, error, flag = convert_cell(value, text, col, cctx, number_format, header_unit)
        if col.type in ("date", "datetime", "time", "number"):
            if value is not None or text:
                stats.type_checked[col.key] = stats.type_checked.get(col.key, 0) + (0 if out is None and not error else 1)
        if flag == "error_value":
            stats.error_values += 1
        elif flag == "year_inferred":
            stats.year_inferred += 1
            record_warnings.append(f"{col.display}の年を{stats.year_context_label or '年度'}から補った")
        elif flag == "year_missing":
            stats.year_missing += 1
        if error:
            stats.type_errors[col.key] = stats.type_errors.get(col.key, 0) + 1
            record_warnings.append(error)
            add_row_issue(Issue("warning", "type_error", error, row=row_index, column=col.display))
        if out is not None and not _is_text_column(col) and text and str(out) != text:
            originals[col.key] = text
        return out

    def reconcile_row(row, label: str, level: str):
        acc = running if level == "小計" else grand
        positions = [(c.display, pos, c, header_units.get(c.key, "")) for c, pos in measure_cols]
        for display, pos, col, unit in positions:
            value, text, fmt = cell_raw(row, pos)
            number, error, _flag = convert_cell(value, text, col, cctx, fmt, unit)
            if error or not isinstance(number, (int, float)) or isinstance(number, bool):
                continue
            stats.reconcile.append({"row": row.index, "label": f"{label}の行" if label else level,
                                    "column": display, "expected": _clean_number(number),
                                    "actual": _clean_number(acc.get(pos, 0.0))})
        if level == "小計":
            running.clear()

    def accumulate(pos: int, number):
        if isinstance(number, (int, float)):
            running[pos] = running.get(pos, 0.0) + number
            grand[pos] = grand.get(pos, 0.0) + number

    # 継続行の判定に使う列: 指定 → 設定のキー・日付の列 → 自動判定
    key_columns = list(opts.get("key_columns") or [])
    if not key_columns:
        key_columns = sorted(pos for c, pos in columns if c.role in ("key", "date"))
    if not key_columns:
        key_columns = list(layout.key_columns or [])
    if not opts.get("key_columns"):
        # 「空欄は上の値」の列は空欄が当たり前なので、継続行の判定には使わない
        fill_down = {pos for c, pos in columns if c.fill_down_blank}
        key_columns = [k for k in key_columns if k not in fill_down]
    stats.key_columns = key_columns
    for n, (row, rc) in enumerate(classify_rows(source, sheet, layout, key_columns=key_columns)):
        stats.rows_scanned += 1
        if on_progress is not None and n % PROGRESS_EVERY == 0:
            on_progress(n, total_rows)
        for i, cell in enumerate(row.cells):
            if cell.merged_anchor is not None and cell.text and cell.merged_anchor == (row.index, i + 1):
                anchors[cell.merged_anchor] = (cell.value, cell.text, cell.number_format)
        kind, reason = rc.kind, rc.reason
        if kind == "blank":
            stats.blank_rows += 1
            continue
        if kind == "excluded" and reason in ("非表示の行", "取り消し線の行"):
            policy = policies.get("hidden_rows" if reason == "非表示の行" else "strike_rows", "exclude_with_warning")
            if policy == "include":
                kind = "data"
                stats.included_hidden += 1
        if kind == "subtotal":
            stats.subtotal_rows += 1
            stats.excluded[reason or "小計"] = stats.excluded.get(reason or "小計", 0) + 1
            label = next((c.text for c in row.cells if c.text), "")
            reconcile_row(row, _nfkc(label).strip()[:20], "合計" if reason == "合計" else "小計")
            continue
        if kind in ("note", "title", "header", "excluded"):
            key = reason or {"note": "注記", "title": "表題", "header": "見出し"}.get(kind, "除外")
            stats.excluded[key] = stats.excluded.get(key, 0) + 1
            continue
        if kind == "continuation":
            if spec.continuation_rows == "merge_into_previous" and last_record is not None:
                merged = False
                for col, pos in columns:
                    if not _is_text_column(col):
                        continue
                    text = row.text(pos)
                    if not text:
                        continue
                    addition = nfkc_text(text)
                    current = last_record.values.get(col.key)
                    last_record.values[col.key] = f"{current}\n{addition}" if current else addition
                    merged = True
                if merged:
                    stats.continuation_merged += 1
                    continue
            kind = "data"

        stats.data_rows += 1
        record_warnings: list[str] = []
        originals: dict[str, str] = {}
        values: dict[str, object] = {}
        row_raw: dict[str, tuple[object, str, str | None]] = {}
        for col, pos in columns:
            raw = cell_raw(row, pos)
            if col.role in DITTO_ROLES and is_ditto(raw[1]):
                # 「〃」「同上」は「上と同じ」の書き方。空欄は上の値の設定に関わらず、直前のデータ行の値で補う
                above = last_row_raw.get(col.key)
                if above is not None:
                    raw = above
                    stats.ditto_filled[col.key] = stats.ditto_filled.get(col.key, 0) + 1
                else:
                    add_row_issue(Issue("warning", "ditto_unfilled",
                                        f"{col.display}の「{raw[1]}」を補えませんでした（上の行に値がありません）",
                                        row=row.index, column=col.display))
            row_raw[col.key] = raw
            if col.fill_down_blank and not raw[1] and raw[0] is None:
                if col.key in prev_raw:
                    raw = prev_raw[col.key]
                    row_raw[col.key] = raw
                    stats.filled_down[col.key] = stats.filled_down.get(col.key, 0) + 1
            elif raw[1] or raw[0] is not None:
                prev_raw[col.key] = raw
            if raw[0] is None and not raw[1]:
                continue  # 空欄は変換しても値・問題・件数が増えない
            out = convert(col, raw, row.index, record_warnings, originals, header_units.get(col.key, ""))
            if out is not None:
                values[col.key] = out
        src = {"file": file_name, "sheet": stats.sheet, "row": row.index}

        for col, pos in measure_cols:
            accumulate(pos, values.get(col.key))
        last_row_raw = {k: v for k, v in row_raw.items() if v[1] or v[0] is not None}
        record = RecordRow("", values, originals, src, record_warnings)
        pending.append(record)
        last_record = record

    uncached = getattr(source, "uncached_formulas", None)
    if callable(uncached):
        by_col = uncached(sheet)
        for col, pos in columns:
            n = by_col.get(pos + 1)
            if n:
                stats.uncached_formulas[col.display] = stats.uncached_formulas.get(col.display, 0) + n
    replaced = getattr(source, "replaced_rows", None)
    if replaced:
        stats.replaced_rows = list(replaced)
    unclosed = getattr(source, "unclosed_quote_row", None)
    if isinstance(unclosed, int):
        stats.unclosed_quote_row = unclosed
    long_record = getattr(source, "long_record_row", None)
    if isinstance(long_record, int):
        stats.long_record_row = long_record
    _assign_keys(pending, spec, stats)
    date_key = spec.date_key
    months: Counter[str] = Counter()
    dates = []
    for rec in pending:
        v = rec.values.get(date_key)
        if v:
            s = str(v)
            if _ISO_DATE_RE.fullmatch(s[:10]):
                dates.append(s[:10])  # 変換できず原文のまま残った「不明」などは日付の範囲に入れない
            if len(s) >= 7 and s[4] == "-":
                months[s[:7]] += 1
    stats.months = dict(sorted(months.items()))
    if stats.ditto_filled:
        names = {c.key: c.display for c in spec.columns}
        detail = "、".join(f"{names.get(k, k)} {n}件" for k, n in stats.ditto_filled.items())
        issues.append(Issue("warning", "ditto_filled", f"「〃」「同上」を直前の行の値で補いました（{detail}）"))
    stats.date_min = min(dates) if dates else None
    stats.date_max = max(dates) if dates else None
    stats.records = len(pending)
    if stats.row_issues_truncated:
        issues.append(Issue("warning", "issues_truncated",
                            f"行ごとの問題が多いため、{MAX_ROW_ISSUES}件を超えた分（{stats.row_issues_truncated}件）は省略しました"))
    if on_progress is not None:
        on_progress(total_rows, total_rows)
    stats.elapsed_sec = round(_time.monotonic() - started, 2)
    return pending, issues, stats


def _key_part(values: dict, part: str) -> str:
    name, _, length = str(part).partition(":")
    value = values.get(name)
    if value in (None, ""):
        return ""
    s = str(value)
    if length.isdigit():
        s = _ALL_SPACES.sub("", s)[: int(length)]
    return s


def _assign_keys(records: list[RecordRow], spec: TableSpec, stats: ImportStats) -> None:
    """記録キー。キー列がそろわない行は代わりのキー。同じキーは出現順に #2, #3 を付ける。"""
    record = spec.record or {}
    key_parts = [str(p) for p in (record.get("key") or [])]
    fallback = [str(p) for p in (record.get("fallback_key") or [])]
    bases: list[tuple[str, bool]] = []
    for rec in records:
        base, primary = "", False
        if key_parts:
            parts = [_key_part(rec.values, p) for p in key_parts]
            if all(parts):
                base, primary = "|".join(parts), True
        if not base and fallback:
            parts = [_key_part(rec.values, p) for p in fallback]
            if any(parts):
                base = "|".join(parts)
        if not base:
            base = f"行{rec.source.get('row')}"
        bases.append((base, primary))
    counts = Counter(b for b, _ in bases)
    seen: Counter[str] = Counter()
    for rec, (base, primary) in zip(records, bases):
        seen[base] += 1
        rec.key = base if seen[base] == 1 else f"{base}#{seen[base]}"
        if counts[base] > 1 and primary:
            stats.duplicate_keys[base] = counts[base]


# ====================================================================================================
# 元 tables/markdown.py
# 一覧表の Markdown 生成（記録ファイル）。
#
# 決まり（docs/design.md 6章）: 同じ入力から同じバイト列。生成日時・取込ID・行番号の一覧を本文に書かない。
# レコード内に空行を入れない。パイプ表を使わない。出す列の中身は1文字も削らない（人名・コードの列も出す）。
# 「経過の記録」の列（role=log）は logproc のルール出力（対応の時系列）＋ AI 照合に通った要点だけを出す。
# 大きい記録は「（続きn/m）」に分けて、どの部分にも管理No・設備・日付を書く（切られても身元が分かるように）。
# ====================================================================================================

# 記録1件の上限（推定トークン）。これを超えたら「（続きn/m）」に分ける。
# 取り込み側の設定は見えないので、狭い固定窓（600/overlap 50）でも記録が途中で切られない大きさにする。
# オフライン評価（T1・T2・T5 × F1200/100・F600/50・P2000）で 300/400/600 を比べて 400 を選んだ（docs/design.md 6.5）。
RECORD_TOKEN_BUDGET = 400
TITLE_TEXT_CHARS = 40


@dataclass
class MdFile:
    name: str
    text: str
    kind: str = "records"  # 作るのは記録ファイルだけ（集計・データセット説明は外した）

    @property
    def data(self) -> bytes:
        return self.text.encode("utf-8")


# ---- 経過の記録（role=log。logproc） -------------------------------------------------------

def people_index_for(spec: TableSpec, records: list[dict]):
    """記入者の判定に使う人物の索引（人物一覧＋担当列の値）。"""
    from app.ai import PeopleIndex

    stage = spec.log_stage
    person_keys = [c.key for c in spec.columns if c.role == "person"]
    names = sorted({str(r.get("values", {}).get(k)) for r in records for k in person_keys
                    if r.get("values", {}).get(k)})
    return PeopleIndex(stage.people if stage else None, column_names=names,
                       groups=(stage.groups or None) if stage else None)


def _base_date(values: dict, spec: TableSpec) -> date | None:
    # AI整形（ai.runner）と同じ探し方にする。片方だけ日付が出て md が「年不明」になるのを防ぐ
    return base_date_from(values, spec)


def parse_log_cell(spec: TableSpec, values: dict, people=None):
    """ログ列のセルをルールで分割する（マスク → parse_log）。AI 整形も同じ関数で分割して ID をそろえる。"""
    from app.ai import PeopleIndex, SplitOptions, mask_text, parse_log

    stage = spec.log_stage
    if stage is None:
        return None
    text = values.get(stage.column)
    if text in (None, ""):
        text = ""
    people = people or PeopleIndex(stage.people, groups=stage.groups or None)
    rules = list(stage.mask or [])
    if rules:
        names = people.names() if any(r in ("person", "人名") for r in rules) else ()
        text, _spans = mask_text(str(text), rules, names=names)
    options = SplitOptions.from_dict(stage.splitter or {})
    return parse_log(str(text), base_date=_base_date(values, spec), people=people, options=options)


def _ai_item(ai_results: dict | None, key: str) -> tuple[str | None, dict]:
    item = (ai_results or {}).get(key)
    if not isinstance(item, dict):
        return None, {}
    status = str(item.get("status") or "ok")
    result = item.get("result") if isinstance(item.get("result"), dict) else item
    return status, result


def _glossary(text: str, glossary: dict | None) -> str:
    if not glossary or not text:
        return text
    from app.ai import apply_glossary

    return apply_glossary(text, glossary)


def ai_point_lines(result: dict, glossary: dict | None = None) -> list[str]:
    """AI の incident から「対応の要点」の行を作る（照合に通った項目だけが result に残っている前提）。"""
    inc = result.get("incident") if isinstance(result.get("incident"), dict) else {}
    lines: list[str] = []

    def g(text) -> str:
        return _glossary(" ".join(str(text or "").split()), glossary)

    rc = inc.get("root_cause") if isinstance(inc.get("root_cause"), dict) else {}
    if rc.get("v"):
        lines.append(f"原因: {g(rc['v'])}" + (f"（{rc['certainty']}）" if rc.get("certainty") else ""))
    for field_name, label in (("temporary_actions", "暫定処置"), ("permanent_actions", "恒久処置")):
        items = [g(a.get("v")) for a in inc.get(field_name) or [] if isinstance(a, dict) and a.get("v")]
        if items:
            lines.append(f"{label}: {'、'.join(items)}")
    parts = []
    for p in inc.get("parts") or []:
        if not isinstance(p, dict):
            continue
        text = " ".join(x for x in (g(p.get("name")), str(p.get("model") or "").strip(),
                                     str(p.get("qty_q") or "").strip()) if x)
        if text:
            parts.append(text)
    if parts:
        lines.append(f"使用部品: {'、'.join(parts)}")
    rec = inc.get("recurrence") if isinstance(inc.get("recurrence"), dict) else {}
    if rec.get("v"):
        lines.append(f"再発: {g(rec['v'])}" + (f"（{rec['count_q']}）" if rec.get("count_q") else ""))
    fs = inc.get("final_state") if isinstance(inc.get("final_state"), dict) else {}
    if fs.get("v"):
        lines.append(f"最終状態: {g(fs['v'])}")
    return [escape_md_line(line) for line in lines]


def _segment_types(result: dict) -> dict[str, list[str]]:
    types: dict[str, list[str]] = {}
    for entry in result.get("entries") or []:
        if not isinstance(entry, dict):
            continue
        t = [str(x) for x in entry.get("t") or [] if x]
        for seg in entry.get("segs") or []:
            types.setdefault(str(seg), [])
            for x in t:
                if x not in types[str(seg)]:
                    types[str(seg)].append(x)
    return types


# ---- 値の書き方 ------------------------------------------------------------------------

def _one_line(text) -> str:
    return " ".join(str(text or "").split())


def _time_column(spec: TableSpec):
    date_key = spec.date_key
    for key in ("occurred_time", f"{date_key.removesuffix('_at').removesuffix('_date')}_time"):
        col = spec.column(key)
        if col is not None and col.type == "time":
            return col
    return None


def format_value(col, value) -> str:
    if value is None or value == "":
        return ""
    if col.type == "number" and isinstance(value, (int, float)) and not isinstance(value, bool):
        return f"{fmt_number(value)}{unit_label(col.unit)}"
    return str(value)


def _date_text(value, time_value=None) -> str:
    s = str(value)
    if time_value and len(s) == 10:
        s = f"{s} {time_value}"
    if is_month(s):
        s += f"（{month_label(s[:7])}）"
    return s


def _entity_label_name(spec: TableSpec) -> str:
    entity, _label = entity_columns(spec)
    if entity is None:
        return "設備"
    return "設備" if entity.key.startswith("equipment") else entity.display


def _is_hidden(col, spec: TableSpec) -> bool:
    """出さない列か。画面で「出さない」にした列だけ（人名・コードの列も既定では出す）。"""
    return col.md == "omit"


_BRACKETS = {"（": "）", "(": ")", "「": "」", "『": "』", "【": "】", "［": "］", "[": "]", "〔": "〕", "《": "》",
             "〈": "〉", "｛": "｝", "{": "}", "“": "”"}
_TITLE_TAG = re.compile(r"^[【\[][^】\]]{0,12}[】\]][ 　]*|^■?[ 　]*発生(?:日時)?[ 　]*[:：][ 　]*")
# 行頭の「R05.04.01 11:45(休日)、」のような日付・時刻（見出しの末尾に日付が入るので、現象の字数を使わない）
_TITLE_LEAD_DATE = re.compile(
    r"^(?:(?:[RHSrhs][ 　]?\d{1,2}|\d{2,4})[./年\-][ 　]?\d{1,2}[./月\-][ 　]?\d{1,2}日?"
    r"|\d{1,2}[:：]\d{2}(?:[:：]\d{2})?"
    r"|[（(][^）)]{0,10}[)）]"
    r"|[\s頃、,。.・~〜\-]+)+")
# 日付・時刻のすぐ後にこれが続くときは、時刻が文の一部（「0:28以降、…」）なので外さない
_TITLE_LEAD_KEEP = ("以降", "以後", "以前", "から", "まで", "より", "前後", "過ぎ", "すぎ")


def _strip_lead_date(text: str) -> str:
    m = _TITLE_LEAD_DATE.match(text)
    if m is None or text[m.end():].startswith(_TITLE_LEAD_KEEP):
        return text
    return text[m.end():]


def _clip_title_text(text: str, limit: int) -> str:
    """見出し用に切り詰める。閉じない括弧の前で切り、切ったことが分かるように「…」を付ける。"""
    if limit <= 0 or len(text) <= limit:
        return text
    cut = text[:limit]
    opened: list[int] = []
    for i, ch in enumerate(cut):
        if ch in _BRACKETS:
            opened.append(i)
        elif opened and ch == _BRACKETS[cut[opened[-1]]]:
            opened.pop()
    if opened:
        cut = cut[:opened[0]]
    cut = cut.rstrip().rstrip("、,。.")
    return (cut or text[:limit].rstrip()) + "…"


# 2行目以降の行頭の「現象:」「設備：」のような項目名（数字で始まるもの＝時刻は含めない）
_TITLE_LABEL = re.compile(r"^(?![\d０-９])[^\s:：、。]{1,8}[ 　]*[:：][ 　]*")


def _title_text_line(raw) -> str:
    """見出しに使う文章。1行目の「【発生】」「発生日時:」などの札と、それに続く日付・時刻を外した残り。

    「【発生】R05.04.01 11:45(休日)」のように日付だけの1行目は使わず（見出しの日付と同じものを繰り返さない）、
    次の行（「設備:」の行は飛ばす）の札・項目名・日付を外した残りを使う。
    """
    lines = str(raw).split("\n")
    first = _one_line(lines[0])
    body = _strip_lead_date(_TITLE_TAG.sub("", first).strip()).strip()
    if body or not first:
        return body
    # 1行目が札と日付だけ（「発生:2024-04-28 14:50」）なら、次の行から探す。「設備:」の行は見出しの設備と重なるので飛ばす
    for line in lines[1:]:
        line = _one_line(line)
        tag = _TITLE_TAG.match(line)
        text = _strip_lead_date(_TITLE_TAG.sub("", line).strip()).strip()
        label = _TITLE_LABEL.match(text)
        if (tag and "設備" in tag.group(0)) or (label and "設備" in label.group(0)):
            continue
        text = _strip_lead_date(_TITLE_LABEL.sub("", text).strip()).strip()
        if text:
            return text
    return ""


def record_title(values: dict, spec: TableSpec, source: dict | None = None) -> str:
    """見出し: 【管理No】設備名（設備番号）現象の先頭40字｜日付"""
    md = spec.markdown or {}
    pieces: list[str] = []
    title_columns = list(md.get("title_columns") or [])
    entity, label = entity_columns(spec)
    if title_columns:
        for item in title_columns:
            key, _, length = str(item).partition(":")
            col = spec.column(key)
            if entity is not None and key in (entity.key, label.key if label else None):
                text = entity_display(values, spec)[2]
            else:
                raw = values.get(key)
                text = _title_text_line(raw) if raw not in (None, "") else ""
                limit = int(length) if length.isdigit() else TITLE_TEXT_CHARS
                if col is not None and col.type in ("text", "string") and len(text) > limit:
                    text = _clip_title_text(text, limit)
                if col is not None and col.role == "key" and text:
                    text = f"【{text}】"
            if text and text not in pieces:
                pieces.append(text)
    else:
        key_col = spec.first_role("key")
        if key_col is not None and values.get(key_col.key):
            pieces.append(f"【{_one_line(values[key_col.key])}】")
        display = entity_display(values, spec)[2]
        if display:
            pieces.append(display)
        text_col = spec.column("symptom") or spec.first_role("text")
        if text_col is not None and values.get(text_col.key):
            first = _title_text_line(values[text_col.key])
            if first:
                pieces.append(_clip_title_text(first, TITLE_TEXT_CHARS))
    title = ""
    for piece in pieces:
        if title and not title.endswith(("】", "）")):
            title += " "
        title += piece
    date_value = values.get(spec.date_key)
    if date_value:
        title += f"｜{str(date_value)[:10]}"
    title = _one_line(title)
    if title:
        return title
    # 見出しの材料が何も無い表（識別番号・設備・長文・日付のどれも無い）。全部の記録が同じ見出しに
    # なると RAG のチャンクを見分けられないので、先頭のほうの列の値をつないで見出しにする
    parts: list[str] = []
    for col in spec.columns:
        if _is_hidden(col, spec) or col.type == "text":
            continue
        text = _one_line(str(values.get(col.key) or ""))
        if text:
            parts.append(_clip_title_text(text, 20))
        if len(parts) >= 3:
            break
    if parts:
        return " ".join(parts)
    row = (source or {}).get("row")
    return f"{row}行目の記録" if row else "（見出しなし）"


def record_blocks(record: dict, spec: TableSpec, ai_results: dict | None = None, people=None) -> list[list[str]]:
    """1件分のブロック。推定トークンが RECORD_TOKEN_BUDGET を超えたら「（続きn/m）」に分ける。

    分けても文字は1つも消さない。2つ目以降には管理No・設備・日付を書き直す（その部分だけで身元が分かるように）。
    """
    lines, repeat = _record_lines(record, spec, ai_results, people)
    joined = "\n".join(lines)
    # 推定は1文字あたり最大1.1トークン。それでも上限以下なら数えずに済む（大半の記録。判定の結果は同じ）
    if len(joined) * 11 <= RECORD_TOKEN_BUDGET * 10 or estimate_tokens(joined) <= RECORD_TOKEN_BUDGET:
        return [lines]
    return _split_record(lines[0], lines[1:], repeat)


def record_block(record: dict, spec: TableSpec, ai_results: dict | None = None, people=None) -> list[str]:
    """1件分の行（画面の下書き表示用）。分かれる記録は続きの見出しも含めて続けて返す。"""
    return [line for block in record_blocks(record, spec, ai_results, people) for line in block]


def _record_lines(record: dict, spec: TableSpec, ai_results: dict | None, people) -> tuple[list[str], list[str]]:
    """1件分の行と、分けたときに書き直す行（管理No・設備・日付）。"""
    values = record.get("values", {}) or {}
    key = record.get("key", "")
    lines = [f"## {record_title(values, spec, record.get('source'))}"]
    entity, label = entity_columns(spec)
    key_col = spec.first_role("key")
    time_col = _time_column(spec)
    date_key = spec.date_key
    log_key = spec.log_stage.column if spec.log_stage else None
    entity_done = False
    repeat: list[str] = []
    status, result = _ai_item(ai_results, key)

    def add(col, bullet: list[str]) -> None:
        """行を足す。管理No・設備・日付の行は、記録を分けたときに書き直すので控えておく。"""
        lines.extend(bullet)
        if col is not None and (col.key == date_key or (key_col is not None and col.key == key_col.key)
                                or (entity is not None and col.key in (entity.key, label.key if label else ""))):
            repeat.extend(bullet)

    for col in spec.columns:
        if _is_hidden(col, spec):
            continue
        if time_col is not None and col.key == time_col.key and values.get(date_key):
            continue
        if entity is not None and label is not None and col.key in (entity.key, label.key):
            if entity_done:
                continue
            entity_done = True
            eid, name, display = entity_display(values, spec)
            if eid and name and not _is_hidden(entity, spec) and not _is_hidden(label, spec):
                add(entity, md_bullet(_entity_label_name(spec), display))
                continue
            for c in (entity, label):
                if not _is_hidden(c, spec):
                    add(c, md_bullet(c.display, format_value(c, values.get(c.key))))
            continue
        if entity is not None and label is None and col.key == entity.key:
            # 設備名の列がない台帳。見出し・集計と同じ「設備名（設備番号）」の書き方にそろえる（列の名前はそのまま）
            _eid, name, display = entity_display(values, spec)
            if name and display:
                add(col, md_bullet(col.display, display))
                continue
        value = values.get(col.key)
        if col.key == log_key:
            lines += _log_lines(col, values, spec, status, result, people)
            continue
        if value in (None, ""):
            continue
        if col.key == date_key:
            time_value = values.get(time_col.key) if time_col is not None else None
            add(col, md_bullet(col.display, _date_text(value, time_value)))
        else:
            add(col, md_bullet(col.display, format_value(col, value)))
    if status == "ok":
        for stage in spec.custom_stages:
            item = (result.get("custom") or {}).get(stage.id) if isinstance(result.get("custom"), dict) else None
            v = item.get("v") if isinstance(item, dict) else item
            if v:
                target = spec.column(stage.target_key) if stage.target_key else None
                lines += md_bullet(f"{target.display if target else stage.id}（AI分類）", _one_line(v))
    source = record.get("source", {}) or {}
    lines += md_bullet("出典", _source_text(values, spec, source))
    return lines, repeat


# ---- 大きい記録を「（続きn/m）」に分ける（文字は1つも消さない） ------------------------------

_MIN_PART_TOKENS = 80          # 見出し・書き直す行を引いても、これだけは中身に使う
_BULLET_LABEL = re.compile(r"^- ([^:]{1,40}): ")


def _tokens(lines: list[str]) -> int:
    return estimate_tokens("\n".join(lines)) if lines else 0


def _bullet_groups(body: list[str]) -> list[list[str]]:
    """箇条書きのまとまり（`- 項目:` の行と、それに続く2字下げの行）に分ける。"""
    groups: list[list[str]] = []
    for line in body:
        if line.startswith("- ") or not groups:
            groups.append([line])
        else:
            groups[-1].append(line)
    return groups


def _split_group(group: list[str], budget: int) -> list[list[str]]:
    """1つの箇条書きが budget に収まらないときに小分けにする（行も文も消さない）。"""
    if _tokens(group) <= budget:
        return [group]
    head = group[0]
    if len(group) > 1:
        # 複数行の値（対応の時系列など）。見出しの行を「（続き）」で繰り返して2字下げの行を分ける
        cont = f"{head[:-1]}（続き）:" if head.endswith(":") else head
        out, cur = [], [head]
        for line in group[1:]:
            if len(cur) > 1 and _tokens(cur) + estimate_tokens(line) > budget:
                out.append(cur)
                cur = [cont]
            cur.append(line)
        return out + [cur]
    m = _BULLET_LABEL.match(head)
    if m is None:
        return [group]  # 項目名が読めない1行。切らずにそのまま出す
    # 1行の長い値。「。」の後ろで分け、項目名を「（続き）」で繰り返す
    label, text = m.group(1), head[m.end():]
    pieces = [p for p in re.split(r"(?<=。)", text) if p]
    out, cur = [], ""
    for piece in pieces:
        if cur and estimate_tokens(f"- {label}（続き）: {cur}{piece}") > budget:
            out.append([f"- {label}: {cur}" if not out else f"- {label}（続き）: {cur}"])
            cur = ""
        cur += piece
    if cur:
        out.append([f"- {label}: {cur}" if not out else f"- {label}（続き）: {cur}"])
    return out or [group]


def _split_record(title: str, body: list[str], repeat: list[str]) -> list[list[str]]:
    """見出し・本文を、1つあたり RECORD_TOKEN_BUDGET に収まるブロックの並びにする。"""
    # 2つ目以降は「見出し（続きn/m）」＋管理No・設備・日付を書き直すので、その分を引いた残りが中身に使える
    overhead = estimate_tokens(f"{title}（続き00/00）") + _tokens(repeat)
    budget = max(_MIN_PART_TOKENS, RECORD_TOKEN_BUDGET - overhead)
    groups = [g for group in _bullet_groups(body) for g in _split_group(group, budget)]
    parts: list[list[str]] = []
    cur: list[str] = []
    for g in groups:
        if cur and _tokens(cur) + _tokens(g) > budget:
            parts.append(cur)
            cur = []
        cur += g
    if cur or not parts:
        parts.append(cur)
    if len(parts) == 1:
        return [[title] + parts[0]]
    total = len(parts)
    out = []
    for i, part in enumerate(parts, start=1):
        head = f"{title}（{i}/{total}）" if i == 1 else f"{title}（続き{i}/{total}）"
        again = [] if i == 1 else [ln for ln in repeat if ln not in part]
        out.append([head] + again + part)
    return out


# ---- 経過の記録の列の行 ---------------------------------------------------------------


def _log_lines(col, values: dict, spec: TableSpec, status, result: dict, people) -> list[str]:
    """ログ列の行（対応の要点＋対応の時系列）。書かれた文は1つも省かない。"""
    stage = spec.log_stage
    parse = parse_log_cell(spec, values, people)
    if parse is None or parse.kind == "empty":
        return []
    out: list[str] = []
    types: dict[str, list[str]] = {}
    if status == "ok" and result:
        points = ai_point_lines(result, stage.glossary)
        if points:
            out.append("- 対応の要点（AI抽出）:")
            out += [f"  {p}" for p in points]
        types = _segment_types(result)
    _eid, entity_name, entity_text = entity_display(values, spec)
    entity_label = entity_name or entity_text
    from app.ai import render_timeline

    timeline = render_timeline(parse, entity_label, types=types or None, glossary=stage.glossary or None)
    if parse.kind == "header_cell":
        out += md_bullet(f"{col.display}（見出しごと）", timeline)
        return out
    if not timeline:
        return out
    out.append("- 対応の時系列:")
    out += [f"  {line}" for line in timeline]
    return out


def _source_text(values: dict, spec: TableSpec, source: dict) -> str:
    file_name = str(source.get("file") or "")
    key_col = spec.first_role("key")
    if key_col is not None and values.get(key_col.key):
        return f"{file_name}（{key_col.display} {_one_line(values[key_col.key])}）"
    row = source.get("row")
    sheet = source.get("sheet")
    if row:
        where = f"シート「{sheet}」{row}行目" if sheet else f"{row}行目"
        return f"{file_name}（{where}）"
    return file_name


# ---- ファイルの組み立て ---------------------------------------------------------------------

def _sort_key(record: dict, spec: TableSpec, time_col) -> tuple:
    values = record.get("values", {}) or {}
    d = str(values.get(spec.date_key) or "")
    t = str(values.get(time_col.key) or "") if time_col is not None else ""
    # 日付が同じ（または日付の列が無い）ときは元の表の順に並べる。記録キーの文字くらべだと
    # 「行10」「行100」「行11」の順になり、元のExcelと突き合わせられなくなる
    row = (record.get("source") or {}).get("row")
    return (0 if d else 1, d, t, int(row) if isinstance(row, int) else 0, str(record.get("key", "")))


class _Names:
    """ファイル名の重複を避ける（安全化で同じ名前になった場合だけ _2 を付ける）。

    Windows のフォルダは大文字・小文字を区別しないので、比較は casefold して行う
    （「ETC-302号機」と「Etc-302号機」が同じファイルに上書きされて記録が消えるのを防ぐ）。
    """

    def __init__(self):
        self.used: set[str] = set()

    def make(self, parts: list[str]) -> str:
        name = md_filename(parts)
        n = 2
        while name.casefold() in self.used:
            name = md_filename(parts + [str(n)])
            n += 1
        self.used.add(name.casefold())
        return name


def render_all(spec: TableSpec, records: list[dict], ai_results: dict | None) -> list[MdFile]:
    """全 md ファイルを作る（決定的）。records は確定済み全行（RecordRow.to_dict の形）。

    作るのは RAG に入れる記録ファイルだけ（集計・データセット説明は 2026-09-20 に外した。docs/design.md 6.3）。
    """
    md = spec.markdown or {}
    names = _Names()
    time_col = _time_column(spec)
    ordered = sorted(records, key=lambda r: _sort_key(r, spec, time_col))
    if not md.get("records", True):
        return []
    return _record_files(spec, ordered, ai_results or {}, names)


def _record_files(spec: TableSpec, ordered: list[dict], ai_results: dict, names: _Names) -> list[MdFile]:
    """記録ファイル（月ごと、または対象×月ごとに1ファイル。件数では分けない）。"""
    md = spec.markdown or {}
    prefix = spec.file_prefix
    by_entity = md.get("group_by") == "entity_month"
    entity, _label = entity_columns(spec)
    groups: dict[tuple, list[dict]] = defaultdict(list)
    for rec in ordered:
        values = rec.get("values", {}) or {}
        d = str(values.get(spec.date_key) or "")
        month = d[:7] if is_month(d) else ""
        eid = entity_value(values, spec)[0] if (by_entity and entity is not None) else ""
        groups[(eid, month) if by_entity else ("", month)].append(rec)
    people = people_index_for(spec, ordered) if spec.log_stage else None
    files: list[MdFile] = []
    for (eid, month) in sorted(groups, key=lambda g: (g[0], g[1] == "", g[1])):
        recs = groups[(eid, month)]
        parts = [prefix]
        if by_entity:
            parts.append(eid or "設備不明")
        parts.append(month or "日付なし")
        blocks: list[list[str]] = [_record_file_header(spec, recs, month, eid)]
        for r in recs:
            blocks += record_blocks(r, spec, ai_results, people)
        files.append(MdFile(names.make(parts), join_file(blocks), "records"))
    return files


def _record_file_header(spec: TableSpec, group: list[dict], month: str, eid: str) -> list[str]:
    name = spec.name
    scope_month = month_label(month) if month else "日付なし"
    entity_disp = ""
    if eid:
        entity_disp = entity_display(group[0].get("values", {}), spec)[2] or eid
    title = f"# {name}"
    if entity_disp:
        title += f" {entity_disp}"
    title += f" {scope_month}の記録" if month else " 日付なしの記録"
    body = [f"- データ種別: {name}（1行＝1件）の記録"]
    if entity_disp:
        body += md_bullet(_entity_label_name(spec), entity_disp)
    if month:
        body.append(f"- 対象期間: {month_first_day(month)}〜{month_last_day(month)}")
    scope = f"{eid}の{scope_month}" if eid else scope_month
    body.append(f"- このファイルの記録: {len(group):,}件（{scope}の全件）")
    return _Block(title, body)


class _Block(list):
    """ファイル先頭の「# 見出し / 空行 / 本文」。join_blocks はブロック内の空行を落とすので別ブロックに分けて出す。"""

    def __init__(self, title: str, body: list[str]):
        super().__init__([title] + body)
        self.title = title
        self.body = body


def join_file(blocks: list[list[str]]) -> str:
    expanded: list[list[str]] = []
    for block in blocks:
        if isinstance(block, _Block):
            expanded.append([block.title])
            expanded.append(block.body)
        else:
            expanded.append(block)
    return join_blocks(expanded)


# ====================================================================================================
# 元 tables/outputs.py
# 画面で見る問題一覧の CSV と、ダウンロード用 zip（RAG に入れる md だけ）。
#
# 投入済みとの差分管理はしない（取り込みごとに、その取り込みの全 md を zip にまとめて渡す）。
# ====================================================================================================

_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
_FORMULA_PREFIX = ("=", "+", "-", "@", "\t", "\r")


# ---- 問題一覧の CSV（確認画面から見るためのもの。zip には入れない） ---------------------------------------

def guard_formula(value) -> str:
    """Excel で開いたときに式として実行されないよう、= + - @（全角も）で始まる文字列の先頭に ' を付ける。"""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)):
        return str(value)
    s = str(value)
    # 日本語の Excel は全角の ＝ ＋ － ＠ で始まる値も式として読むので、先頭の1文字は NFKC で比べる
    if s.startswith(_FORMULA_PREFIX) or unicodedata.normalize("NFKC", s[:1]).startswith(_FORMULA_PREFIX):
        return "'" + s
    return s


def csv_bytes(header: list[str], rows) -> bytes:
    """UTF-8（BOM付き）・CRLF の CSV。全セルに式の対策をする。"""
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\r\n")
    writer.writerow([guard_formula(h) for h in header])
    for row in rows:
        writer.writerow([guard_formula(v) for v in row])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def issues_csv(issues) -> bytes:

    rows = []
    for issue in issues:
        d = issue.to_dict() if hasattr(issue, "to_dict") else dict(issue)
        rows.append([LEVEL_LABELS.get(d.get("level"), d.get("level")), d.get("code", ""), d.get("row") or "",
                     d.get("column") or "", d.get("message", "")])
    return csv_bytes(["種類", "コード", "行", "列", "内容"], rows)


# ---- zip --------------------------------------------------------------------------------

def _zip_write(zf: zipfile.ZipFile, name: str, data: bytes) -> None:
    info = zipfile.ZipInfo(name, date_time=_ZIP_TIME)  # 時刻を固定して、同じ内容なら同じ zip にする
    info.compress_type = zipfile.ZIP_DEFLATED
    info.external_attr = 0o644 << 16
    zf.writestr(info, data)


def build_zip(md_files: list[tuple[str, bytes]]) -> bytes:
    """RAG に入れる md だけをフォルダ分けせずに入れる（zip を開いて、そのまま LightRAG にドラッグできるように）。"""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in sorted(md_files):
            _zip_write(zf, name, data)
    return buf.getvalue()


# ====================================================================================================
# 元 tables/store.py
# 一覧表の DB アクセス（取り込み1件と、その取り込みが使う取り込み設定）。
#
# 取り込み設定は保存しない（利用者の指示 2026-09-20:「表の方には、取り込み設定を保持しておく機能はいらない」）。
# 取り込みごとに列の対応づけを決め、その設定（TableSpec）を取り込みの行（spec_json）に持つ。
# 接続は models.core.get_db()（リクエスト中もジョブの app_context 中も使える）。conn を渡せばそれを使う。
# JSON の列は読み出し時に dict/list に直した値を別名（source, stats）で付け、spec_json は TableSpec にする。
# ====================================================================================================

IMPORT_JSON_COLUMNS = {"source_json": "source", "stats_json": "stats"}


def _db(conn=None) -> sqlite3.Connection:
    return conn if conn is not None else core.get_db()


def _loads(text, default):
    try:
        value = json.loads(text) if text else default
    except (TypeError, ValueError):
        return default
    return value if isinstance(value, type(default)) else default


def _dumps(value) -> str:
    return value if isinstance(value, str) else json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)


def _spec_or_none(spec_text):
    try:
        return spec_from_dict(_loads(spec_text, {})) if spec_text else None
    except ValueError:
        return None


# ---- 取り込み ------------------------------------------------------------------------------

def _decode_import(row) -> dict | None:
    if row is None:
        return None
    d = dict(row)
    for column, alias in IMPORT_JSON_COLUMNS.items():
        d[alias] = _loads(d.get(column), {})
    d["spec"] = _spec_or_none(d.get("spec_json"))
    return d


def create_import(file_name: str, file_hash: str, stored_path: str, source: dict | None = None, conn=None,
                  session_id: str | None = None) -> int:
    """取り込みを1件作る。session_id は置いたブラウザ（views.current_session_id）。

    template_id / template_version_id は取り込み自身の番号にそろえる（設定はもう無いが、AI整形の控え
    （ai_items）がこの番号で取り込みを束ねている。design.md 3.2）。
    """
    db = _db(conn)
    ts = core.now()
    cur = db.execute("""INSERT INTO table_imports (file_name, file_hash, stored_path, source_json, session_id,
                        status, created_at, updated_at)
                        VALUES (?, ?, ?, ?, ?, 'uploaded', ?, ?)""",
                     (file_name, file_hash, stored_path, _dumps(source or {}), session_id, ts, ts))
    import_id = cur.lastrowid
    db.execute("UPDATE table_imports SET template_id = id, template_version_id = id WHERE id = ?", (import_id,))
    db.commit()
    return import_id


def get_import(import_id: int, conn=None) -> dict | None:
    return _decode_import(_db(conn).execute("SELECT * FROM table_imports WHERE id = ?", (import_id,)).fetchone())


def save_spec(import_id: int, spec: TableSpec, conn=None) -> None:
    """この取り込みが使う取り込み設定を保存する（列の対応づけを保存するたびに上書きする）。"""
    update_import(import_id, conn=conn, spec_json=spec_json(spec), spec_hash=spec_hash(spec))


def update_import(import_id: int, conn=None, commit: bool = True, **columns) -> None:
    """列を更新する。source/stats は dict のまま渡してよい（*_json に保存）。"""
    if not columns:
        return
    sets, args = [], []
    for name, value in columns.items():
        column = f"{name}_json" if name in ("source", "stats") else name
        if column.endswith("_json") and not isinstance(value, str):
            value = _dumps(value)
        sets.append(f"{column} = ?")
        args.append(value)
    sets.append("updated_at = ?")
    args.append(core.now())
    db = _db(conn)
    db.execute(f"UPDATE table_imports SET {', '.join(sets)} WHERE id = ?", (*args, import_id))
    if commit:
        db.commit()


# 取り込み1件を消すのは core/purge.py の purge_table_import（ファイルと DB の行をまとめて消す。design.md 3.3）


def list_imports(status: str | list | None = None, limit: int = 100, conn=None,
                 session_id: str | None = None) -> list[dict]:
    """取り込みの一覧。session_id を渡すとそのブラウザの分だけ（持ち主の分からない古い行は含む）。"""
    where, args = [], []
    if session_id:
        where.append("(session_id IS NULL OR session_id = ?)")
        args.append(session_id)
    if status:
        statuses = [status] if isinstance(status, str) else list(status)
        where.append(f"status IN ({','.join('?' * len(statuses))})")
        args += statuses
    sql = ("SELECT * FROM table_imports" + (f" WHERE {' AND '.join(where)}" if where else "")
           + " ORDER BY id DESC LIMIT ?")
    return [_decode_import(r) for r in _db(conn).execute(sql, (*args, limit)).fetchall()]


# ====================================================================================================
# 元 tables/pipeline.py
# 一覧表の取り込みの実行関数（ジョブ本体）と、画面から呼ぶ補助。
#
# - run_read: 保存した範囲と設定で全行を読む → 正規化 → チェック → rows.jsonl.gz / issues.json / issues.csv
# - run_render: その取り込みの記録（＋照合に通った AI 整形の結果）から全 md を作り、取り込みのフォルダに保存
# - build_download: RAG に入れる md だけをまとめた zip
# 期間の置き換え・投入済みとの差分・取り消しはしない（取り込みごとに、その取り込みの内容だけで md を作る）。
# 置き場所: TABLES_DIR/imports/<import_id>/（rows.jsonl.gz, issues.*, md/, preview_md/, source_cache.json）
# 画面とジョブは import_source() の控え（tables.source_cache）を通して表を読み、同じ計算を繰り返さない。
# ====================================================================================================

ROWS_FILE = "rows.jsonl.gz"
ISSUES_CSV = "issues.csv"
ISSUES_JSON = "issues.json"
MD_DIR = "md"
PREVIEW_DIR = "preview_md"


class PipelineError(JobError):
    """利用者に見せる日本語メッセージの処理エラー（そのまま画面に出る）。"""


# ---- ファイル -----------------------------------------------------------------------------

def import_dir(import_id: int) -> Path:
    return Path(current_app.config["TABLES_DIR"]) / "imports" / str(int(import_id))


def import_files(import_id: int) -> dict[str, Path]:
    base = import_dir(import_id)
    return {"dir": base, "rows": base / ROWS_FILE, "issues_csv": base / ISSUES_CSV, "issues_json": base / ISSUES_JSON,
            "md": base / MD_DIR, "preview": base / PREVIEW_DIR}


# 取り込み1件を消すのは core/purge.py の purge_table_import（アップロードしたファイル・このフォルダ・
# DB の行・AI整形の控えをまとめて消す。design.md 3.3）


def write_atomic(path: Path, data: bytes) -> None:
    """一時ファイルに書いてから置き換える（途中で止まっても前のファイルが残る）。

    置き場所（imports/<id>/）は作らない。消された取り込みのフォルダを書き戻さないため（design.md 3.3）。
    """
    path = Path(path)
    tmp = path.with_name(f".{path.name}.tmp")
    with open(tmp, "wb") as f:
        f.write(data)
    os.replace(tmp, path)


def write_rows(path: Path, records) -> None:
    """記録を gzip の JSON Lines で保存（mtime=0・キー順固定で決定的）。一時ファイルに直接書いてから置き換える。

    置き場所は作らない（write_atomic と同じ）。
    """
    path = Path(path)
    tmp = path.with_name(f".{path.name}.tmp")
    with open(tmp, "wb") as raw, gzip.GzipFile(filename="", mode="wb", fileobj=raw, mtime=0, compresslevel=6) as gz:
        buf = io.BufferedWriter(gz, 1024 * 1024)
        for rec in records:
            d = rec.to_dict() if hasattr(rec, "to_dict") else dict(rec)
            buf.write(json.dumps(d, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8") + b"\n")
        buf.flush()
        buf.detach()
    os.replace(tmp, path)


def load_rows(import_id: int, offset: int = 0, limit: int | None = None) -> list[dict]:
    path = import_files(import_id)["rows"]
    if not path.exists():
        return []
    with gzip.open(path, "rt", encoding="utf-8") as f:
        rows = [json.loads(line) for line in f if line.strip()]
    return rows[offset: offset + limit] if limit is not None else rows[offset:]


def load_rows_page(import_id: int, offset: int, limit: int) -> tuple[list[dict], int]:
    """offset から limit 件の記録と全件数。範囲外の行は JSON を読まずに数えるだけ。"""
    path = import_files(import_id)["rows"]
    if not path.exists():
        return [], 0
    rows: list[dict] = []
    total = 0
    with gzip.open(path, "rb") as f:
        for line in f:
            if not line.strip():
                continue
            if offset <= total < offset + limit:
                rows.append(json.loads(line))
            total += 1
    return rows, total


def load_issues(import_id: int) -> list[dict]:
    path = import_files(import_id)["issues_json"]
    if not path.exists():
        return []
    return json.loads(path.read_text(encoding="utf-8"))


def md_text(directory: Path, name: str) -> str | None:
    """フォルダ内の md の中身（名前はそのフォルダ直下の .md だけ許す）。"""
    path = directory / name
    if path.parent != directory or path.suffix != ".md" or not path.exists():
        return None
    return path.read_text(encoding="utf-8")


def md_paths(import_id: int) -> list[Path]:
    directory = import_files(import_id)["md"]
    return sorted(directory.glob("*.md"), key=lambda p: p.name) if directory.exists() else []


def _write_md_dir(target: Path, files: list[MdFile]) -> None:
    from app.core import path_limit

    tmp = target.with_name(target.name + ".new")
    limit = path_limit()
    if limit is not None and files:
        longest = max(files, key=lambda f: len(f.name)).name
        if len(os.path.abspath(tmp)) + 1 + len(longest) > limit:
            # Windows のパスの長さの上限（260文字）を超えると、書けずに分かりにくいエラーで止まる
            raise PipelineError(
                f"Markdownのファイル名が長すぎて、サーバーのデータの置き場所に書けません（最長 {len(longest)}文字）。"
                "「表の名前」を短くするか、アプリを浅いフォルダに置いてください")
    shutil.rmtree(tmp, ignore_errors=True)
    # 親（imports/<id>/）は作り直さない。消された取り込みのフォルダを復活させないため（design.md 3.3）
    tmp.mkdir()
    for f in files:
        (tmp / f.name).write_bytes(f.data)
    shutil.rmtree(target, ignore_errors=True)
    tmp.rename(target)


# ---- 設定と表の範囲 ---------------------------------------------------------------------------

def spec_for_import(imp: dict):
    """その取り込みが使う取り込み設定（取り込みの行が持つ。tables.store）。列の対応づけ前は None。"""
    return (imp or {}).get("spec")


def _source_options(imp: dict) -> dict:
    src = imp.get("source") or {}
    options = {key: src[key] for key in ("encoding", "delimiter", "errors") if src.get(key)}
    options["max_cells"] = current_app.config.get("EXCEL_MAX_CELLS", 500000)
    return options


def open_import_source(imp: dict, sheet_stats: dict | None = None):
    options = _source_options(imp)
    if sheet_stats:
        options["sheet_stats"] = sheet_stats
    return open_source(upload_path(imp["stored_path"]), imp["file_name"], options)


def import_source(imp: dict, real=None) -> ImportSource:
    """控え付きの表ソース（元のファイルは必要になったときだけ開く）。ファイルや読み込み設定が変われば控えは作り直す。"""
    key = json.dumps([imp.get("file_hash") or "", imp["stored_path"], imp["file_name"], _source_options(imp)],
                     ensure_ascii=False, sort_keys=True)
    kind = "excel" if Path(imp["file_name"]).suffix.lower() in EXCEL_EXTENSIONS else "csv"
    return ImportSource(import_dir(imp["id"]), key, kind, imp["file_name"],
                        lambda stats: open_import_source(imp, stats), real=real)


def layout_for_import(source, imp: dict, spec):
    """保存した範囲（見出し行・データ終了行）で表の形を決め直す。未指定なら自動判定。同じ条件の結果は控えから。"""
    cached = source if isinstance(source, ImportSource) else import_source(imp, real=source)
    src = imp.get("source") or {}
    sheet = src.get("sheet")
    if not sheet:
        sheets = [s for s in cached.sheets() if not s.hidden] or cached.sheets()
        sheet = sheets[0].name
    header_rows = [int(r) for r in (src.get("header_rows") or []) if r] or None
    anchors = list((spec.header or {}).get("anchors") or []) if spec is not None else None
    return cached.layout(sheet, anchors=anchors or None, header_row=src.get("header_row") or None,
                         data_end=src.get("data_end_row") or None, header_rows=header_rows)


# ---- 読み込み（ジョブ） ----------------------------------------------------------------------------

def run_read(ctx, import_id: int) -> dict:
    """ジョブ: 保存した範囲と設定で全行を読み、正規化・チェックして行データと問題一覧を書く。"""
    imp = get_import(import_id)
    spec = spec_for_import(imp) if imp else None
    if spec is None:
        update_import(import_id, status="failed", stats={"error": "列の対応づけが決まっていません"})
        raise PipelineError("列の対応づけが決まっていません")
    try:
        ctx.progress(phase="読み込み", done=0, total=0)
        cached = import_source(imp)
        source = cached.real
        layout = layout_for_import(cached, imp, spec)
        ctx.check_cancel()

        def on_progress(done, total):
            ctx.progress(phase="読み込み", done=done, total=total)
            ctx.check_cancel()

        records, row_issues, stats = read_records(
            source, {"sheet": layout.sheet, "file_name": imp["file_name"]}, layout, spec, on_progress=on_progress)
        issues = row_issues + run_checks(records, spec, stats)
        files = import_files(import_id)
        ctx.check_cancel()
        ctx.progress(phase="保存", done=len(records), total=len(records))
        if get_import(import_id) is None:
            # 読み込みの間に渡し終えて（または削除されて）消えた。行データ・問題一覧を書き戻さない（design.md 3.3）
            raise PipelineError("取り込みが削除されたため、読み込んだ内容は保存しませんでした")
        write_rows(files["rows"], records)
        write_atomic(files["issues_json"], json.dumps([i.to_dict() for i in issues], ensure_ascii=False).encode("utf-8"))
        write_atomic(files["issues_csv"], issues_csv(issues))
        for directory in (files["md"], files["preview"]):
            shutil.rmtree(directory, ignore_errors=True)
        st = stats.to_dict()
        st.update({
            "spec_hash": imp.get("spec_hash") or "",
            "layout": {"sheet": layout.sheet, "table_kind": layout.table_kind, "header_rows": layout.header_rows,
                       "data_start": layout.data_start, "data_end": layout.data_end, "headers": layout.headers},
            "issue_counts": count_levels(issues),
        })
        update_import(import_id, status="preview", stats=st, rows_path=str(files["rows"]),
                            issues_path=str(files["issues_csv"]))
        return {"records": len(records), **st["issue_counts"]}
    except JobCancelled:
        update_import(import_id, status="uploaded")
        raise
    except Exception as exc:
        # 例外の型と本文は core.jobs のログに残す。画面には Python の例外名を出さない
        message = str(exc) if isinstance(exc, (PipelineError, UploadError)) else (
            "表の読み込み中に予期しないエラーが起きました。もう一度読み込んでも直らない場合は、"
            "表の範囲・見出し行や列の対応づけを見直してください")
        update_import(import_id, status="failed", stats={**(imp.get("stats") or {}), "error": message})
        raise


def start_read_job(import_id: int) -> int:
    return _start_import_job(import_id, "table_read", "reading", run_read)


def _start_import_job(import_id: int, kind: str, status: str, fn) -> int:
    """取り込みの状態を先に書いてからジョブを始める。

    ジョブを先に始めると、ジョブがすぐ終わって書いた「preview」「confirmed」を、あとの「reading」「confirming」で
    上書きしてしまい、待ち画面が「途中で止まりました」と誤って出す。job_id だけはジョブを作ってから書く。
    """
    before = (get_import(import_id) or {}).get("status")
    update_import(import_id, status=status)
    try:
        job_id = start_job(kind, "table_import", import_id, lambda ctx: fn(ctx, import_id), {"import_id": import_id})
    except Exception:
        if before is not None:
            update_import(import_id, status=before)
        raise
    update_import(import_id, job_id=job_id)
    return job_id


# ---- Markdown の作成 ---------------------------------------------------------------------------

def usable_ai_results(import_id: int, imp: dict, spec) -> dict:
    """照合に通った AI の結果のうち、今の行の内容と合うものだけ（md 用の形 {key: {"status","result"}}）。"""
    if spec.log_stage is None or not imp.get("template_id"):
        return {}
    from app import ai

    accepted = ai.results_for_render(imp["template_id"], "log", import_id=import_id)
    if not accepted:
        return {}
    data = ai.load_rows_for_ai(import_id)
    by_key = ai.items_by_key(imp["template_id"], "log", import_id=import_id)
    out = {}
    for w in ai.prepare_works(data, ["log"], list(accepted)):
        item = by_key.get(w.row_key)
        if item is None or ai.is_outdated(item, item.get("template_version_id"), w.source_hash, w.context_hash,
                                                w.segments_hash):
            continue
        out[w.row_key] = {"status": "ok", "result": accepted[w.row_key]}
    return out


def render_files(import_id: int, imp: dict, spec, records: list[dict] | None = None) -> list[MdFile]:
    if records is None:
        records = load_rows(import_id)
    return render_all(spec, records, usable_ai_results(import_id, imp, spec))


def _preview_signature(import_id: int, imp: dict, spec) -> str:
    """プレビューの md を作った入力（設定・行データ・AI の結果）の目印。"""

    rows_path = import_files(import_id)["rows"]
    row = core.get_db().execute(
        # この取り込みの AI の結果だけ（usable_ai_results が読む範囲と同じ）。同じ設定の別の取り込みでは変わらない
        "SELECT COUNT(*), COALESCE(MAX(updated_at), '') FROM ai_items WHERE template_id = ? AND import_id = ?",
        (imp.get("template_id") or 0, import_id)).fetchone()
    return json.dumps([spec_hash(spec), rows_path.stat().st_mtime_ns if rows_path.exists() else 0, list(row)])


def _preview_index(import_id: int, signature: str) -> dict | None:
    index_path = import_files(import_id)["preview"] / "index.json"
    try:
        index = json.loads(index_path.read_text(encoding="utf-8"))
    except (ValueError, OSError):
        return None
    return index if isinstance(index, dict) and index.get("signature") == signature else None


def _md_dir_from_preview(import_id: int, signature: str) -> int | None:
    """同じ入力で作ったプレビューの md を md/ に置く（作り直さない）。置いたファイル数。そろっていなければ None。

    作成は決定的なので、作り直しても同じバイト列になる。中身は読まずにハードリンク（できなければコピー）で置く
    （Windows では書いたばかりの多数の小さなファイルを読み直すと遅いため）。ファイルは書き換えずに作り直すので共有してよい。
    """
    index = _preview_index(import_id, signature)
    if index is None:
        return None
    files = import_files(import_id)
    preview, target = files["preview"], files["md"]
    names = [item["name"] for item in index["files"]]
    try:
        if any((preview / item["name"]).stat().st_size != item["size"] for item in index["files"]):
            return None
    except (OSError, KeyError, TypeError):
        return None
    tmp = target.with_name(target.name + ".new")
    shutil.rmtree(tmp, ignore_errors=True)
    try:
        tmp.mkdir()  # 親は作り直さない（消された取り込みのフォルダを復活させない）
    except OSError:
        return None
    try:
        for name in names:
            try:
                os.link(preview / name, tmp / name)
            except OSError:
                shutil.copyfile(preview / name, tmp / name)
    except OSError:
        shutil.rmtree(tmp, ignore_errors=True)
        return None
    shutil.rmtree(target, ignore_errors=True)
    tmp.rename(target)
    return len(names)


def preview_signature(import_id: int, imp: dict, spec) -> str:
    """プレビューの md を作った入力の目印（画面から、作り直しが要るかを見るために使う）。"""
    return _preview_signature(import_id, imp, spec)


def _record_count(f) -> int | None:
    """記録ファイル1つに入っている記録の件数（画面の「記録」の列）。

    大きい記録は「（続きn/m）」の見出しに分かれるが、それは1件の記録の続きなので数えない
    （数えると、上に出る「記録件数 60件」と食い違う。tables.markdown._split_record）。
    """
    if f.kind != "records":
        return None
    return sum(1 for line in f.text.split("\n") if line.startswith("## ") and "（続き" not in line)


def ready_preview_files(import_id: int, imp: dict, spec) -> list[dict] | None:
    """すでに作ってあるプレビューの一覧。作っていなければ None（作るのはジョブ run_preview）。"""
    index = _preview_index(import_id, _preview_signature(import_id, imp, spec))
    return index["files"] if index is not None else None


def preview_files(import_id: int, imp: dict, spec, ctx=None) -> list[dict]:
    """プレビュー用に全 md を作る（行・設定・AI結果が変わらなければ前回の結果を使う）。"""
    base = import_files(import_id)
    signature = _preview_signature(import_id, imp, spec)
    index_path = base["preview"] / "index.json"
    index = _preview_index(import_id, signature)
    if index is not None:
        return index["files"]
    if ctx is not None:
        ctx.progress(phase="Markdownの作成", done=1, total=3)
    files = render_files(import_id, imp, spec)
    if ctx is not None:
        ctx.check_cancel()
        ctx.progress(phase="保存", done=2, total=3)
    if get_import(import_id) is None:
        # 作っている間に取り込みが削除された。消したフォルダに記録の md を作り直さない（design.md 3.3）
        raise PipelineError("取り込みが削除されたため、Markdownの下書きは作りませんでした")
    _write_md_dir(base["preview"], files)
    listing = [{"name": f.name, "kind": f.kind, "size": len(f.data), "records": _record_count(f)}
               for f in files]
    index_path.write_text(json.dumps({"signature": signature, "files": listing}, ensure_ascii=False), encoding="utf-8")
    return listing


def run_preview(ctx, import_id: int) -> dict:
    """ジョブ: 「内容とファイルの確認」に出す md をまとめて作る（確定はしない）。

    件数が多いと十数秒かかるので、画面（GET）の中では作らず、待ち画面で進み具合を出せるようにする。
    """
    imp = get_import(import_id)
    spec = spec_for_import(imp)
    if spec is None:
        raise PipelineError("取り込み設定が見つかりません")
    ctx.progress(phase="記録の読み込み", done=0, total=3)
    ctx.check_cancel()
    files = preview_files(import_id, imp, spec, ctx=ctx)
    ctx.progress(phase="完了", done=3, total=3)
    return {"files": len(files)}


def start_preview_job(import_id: int, signature: str) -> int:
    return start_job("table_preview", "table_import", import_id, lambda ctx: run_preview(ctx, import_id),
                          {"import_id": import_id, "signature": signature})


def run_render(ctx, import_id: int) -> dict:
    """ジョブ: この取り込みの記録から全 Markdown を作り、取り込みのフォルダに保存する。"""
    imp = get_import(import_id)
    try:
        spec = spec_for_import(imp)
        if spec is None:
            raise PipelineError("取り込み設定が見つかりません")
        ctx.progress(phase="記録の読み込み", done=0, total=3)
        records = load_rows(import_id)
        ctx.check_cancel()
        ctx.progress(phase="Markdownの作成", done=1, total=3)
        if get_import(import_id) is None:
            # 読み込みの間に渡し終えて（または削除されて）消えた。md を作り直さない（design.md 3.3）
            raise PipelineError("取り込みが削除されたため、Markdownは作りませんでした")
        # 確認画面で同じ入力から作ったプレビューがあれば、それを置く
        file_count = _md_dir_from_preview(import_id, _preview_signature(import_id, imp, spec))
        if file_count is None:
            files = render_files(import_id, imp, spec, records)
            ctx.check_cancel()
            ctx.progress(phase="保存", done=2, total=3)
            if get_import(import_id) is None:
                raise PipelineError("取り込みが削除されたため、Markdownは作りませんでした")
            _write_md_dir(import_files(import_id)["md"], files)
            file_count = len(files)
        stats = imp.get("stats") or {}
        stats["output"] = {"files": file_count, "records": len(records)}
        update_import(import_id, status="confirmed", confirmed_at=core.now(), stats=stats)
        ctx.progress(phase="完了", done=3, total=3)
        return {"files": file_count, "records": len(records)}
    except Exception:
        update_import(import_id, status="preview")
        raise


def start_render_job(import_id: int) -> int:
    return _start_import_job(import_id, "table_render", "confirming", run_render)


# ---- ダウンロード ----------------------------------------------------------------------------

def build_download(import_id: int) -> bytes:
    """RAG に入れる md だけの zip（フォルダ分けなし。管理用の CSV は作らない）。"""
    return build_zip([(p.name, p.read_bytes()) for p in md_paths(import_id)])
