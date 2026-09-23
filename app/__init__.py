"""Flask アプリ本体。設定・土台（DBを含む）・画面を1ファイルにまとめてある。起動は flask --app app serve。

- 設定（旧 config.py）: env ファイルの読み込み、待ち受け先（HOST・PORT）、受け付ける宛先の名前。
- 土台（旧 core.py / models/database.py）: SQLite のスキーマ・移行・データアクセス、安全なファイル名、
  Markdown テキスト処理、アップロードの保存と事前チェック、帳票登録の Excel の一時的な覚え、
  ジョブ実行、取り込んだデータの削除。
- 画面（旧 views.py）: 3つの Blueprint（帳票取り込み・表の取り込み・帳票登録）と段の断片描画。
- create_app とエラー画面、起動コマンド serve。

読み取りは app/extract.py（帳票と一覧表）、AI は app/ai.py（起動時には読み込まない）。
同じ名前で中身の違うものは分けてある（JOB_KIND_LABELS / ROW_KIND_LABELS、_dumps_value /
_dumps_extraction、save_ai_connection_row と画面側の save_ai_connection ルート）。
"""
from __future__ import annotations

import hashlib
import io
import json
import logging
import os
import posixpath
import queue
import re
import secrets
import shutil
import socket
import sqlite3
import sys
import threading
import time
import unicodedata
import zipfile
import zlib
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from functools import lru_cache
from pathlib import Path
from types import SimpleNamespace
from typing import Callable, TYPE_CHECKING
from urllib.parse import quote, urlsplit
from uuid import uuid4

from dotenv import load_dotenv
from flask import abort, Blueprint, current_app, flash, Flask, g, has_app_context, jsonify, redirect, render_template, request, send_file, session, url_for


####################################################################################################
# 設定（旧 config.py）
####################################################################################################

# ====================================================================================================
# 設定（元 config.py）
# ====================================================================================================

BASE_DIR = Path(__file__).resolve().parent.parent   # プロジェクトの根（env・instance/・data/・uploads/ の置き場所）

# 接続設定は aiagent_minimal_rag_tougou と同じく、ドット無しの "env" ファイル（export KEY="..." 形式も可）から読む
load_dotenv(BASE_DIR / "env")


def _split(value: str) -> list[str]:
    return [v.strip() for v in value.replace(",", ";").split(";") if v.strip()]


def _optional_float(name: str) -> float | None:
    raw = os.environ.get(name, "").strip()
    return float(raw) if raw else None


class Config:
    # セッション署名鍵は create_app が FLASK_SECRET_KEY または .flask_secret ファイルから設定する
    SECRET_KEY = None
    DATABASE = Path(os.environ.get("DATABASE", BASE_DIR / "instance" / "app.db"))
    UPLOAD_DIR = Path(os.environ.get("UPLOAD_DIR", BASE_DIR / "uploads"))
    # 画面から保存する設定（model_settings.yaml）の置き場所
    DATA_DIR = Path(os.environ.get("DATA_DIR", BASE_DIR / "data"))
    # 一覧表の行データ・状態ファイルの置き場所（DATA_DIR/tables。create_app で DATA_DIR に合わせて決め直す）
    TABLES_DIR = Path(os.environ.get("TABLES_DIR", DATA_DIR / "tables"))
    # 帳票（1ファイル＝1件）
    ALLOWED_EXTENSIONS = {".xlsx", ".xlsm"}
    # 一覧表（Excel/CSV・1行＝1件）
    TABLE_ALLOWED_EXTENSIONS = {".xlsx", ".xlsm", ".csv", ".tsv", ".txt"}
    TABLE_MAX_UPLOAD_BYTES = int(os.environ.get("TABLE_MAX_UPLOAD_BYTES", str(200 * 1024 * 1024)))
    # 帳票1ファイルの上限（画面を開くたびにブックを開き直すので、一覧表とは別に持つ）
    FORM_MAX_UPLOAD_BYTES = int(os.environ.get("FORM_MAX_UPLOAD_BYTES", str(200 * 1024 * 1024)))
    # 1回のリクエストの上限（一覧表の大きいCSV/Excelを想定）。受け口はここ。
    # 200MB 固定にしていたころは、TABLE_MAX_UPLOAD_BYTES を広げても Flask が先に 413 で断り、
    # 設定が効かなかった（2026-09-23 のレビューで実測）
    MAX_CONTENT_LENGTH = max(200 * 1024 * 1024, TABLE_MAX_UPLOAD_BYTES)
    # Excel のセル数の上限（これを超えるシートは読み込まない）
    EXCEL_MAX_CELLS = int(os.environ.get("EXCEL_MAX_CELLS", "500000"))
    # 読み込み・下書き・AI整形を同時に動かす本数（列ごと。core/jobs.py）。
    # 社内LANのサーバーで数人が同時に使うので、1本だと誰かの長い読み込みでほかの人が待たされる。
    # 1〜8 に丸められる。列ごとの本数はプロセス内でその列を最初に使ったときに決まり、あとから減らない
    JOB_WORKERS = int(os.environ.get("JOB_WORKERS", "3"))

    # ---- LLM（OpenAI互換API） ----
    OPENAI_BASE_URL = os.environ.get("OPENAI_BASE_URL", "https://api.openai.com/v1").strip().rstrip("/")
    OPENAI_API_KEY = os.environ.get("OPENAI_API_KEY", "").strip()
    OPENAI_MODEL = os.environ.get("OPENAI_MODEL", "gpt-5.6-sol").strip()
    OPENAI_MODELS = _split(os.environ.get("OPENAI_MODELS", ""))
    OPENAI_TEMPERATURE = float(os.environ.get("OPENAI_TEMPERATURE", "0"))
    OPENAI_TOP_P = _optional_float("OPENAI_TOP_P")
    # OPENAI_MAX_TOKENS・LLM_RATE_LIMIT_RETRIES・LLM_RATE_LIMIT_MAX_WAIT は持たない。
    # 読み込むだけでどこからも使っておらず、書かれた既定値（3回・20秒）も実際の値と違っていた
    # （再試行は ai.py の MAX_RETRIES=5・MAX_WAIT=120.0。Retry-After の優先や、429 が3行続いたら
    #  一時停止する決まりと噛み合っているので、外から2つだけ変えられる形にはしない。2026-09-23）


# ====================================================================================================
# アプリ（元 app.py）
# ====================================================================================================

_SECRET_FILE = BASE_DIR / ".flask_secret"

# --- 「/」は帳票取り込みへ ------------------------------------------------------------
# ホーム画面は無い。画面は「帳票取り込み」「表の取り込み」「帳票登録」の3つだけで、
# 作業中の一覧もダウンロード待ちの一覧も持たない（利用者の指示 2026-09-20）。
# ここに残るのは「/」を開いたときの転送だけ（古いブックマーク・開いたままの画面の戻り先も拾う）。
home_bp = Blueprint("home", __name__)


@home_bp.get("/", endpoint="index")
def to_forms():
    return redirect(url_for("forms.new"))


@home_bp.get("/guide", endpoint="guide")
def guide():
    """解説（帳票の Markdown がどう作られるか）。上部タブの4つ目。読むだけの画面で、データには触れない
    （利用者の求め 2026-09-22）。中身は templates/base.html の screen == "guide" の節。"""
    from flask import render_template

    return render_template("base.html", screen="guide")


def _secret_key() -> str:
    """セッション署名鍵（画面のメッセージ表示に使う）。再起動しても変わらないようファイルに保持する。"""
    env = os.getenv("FLASK_SECRET_KEY", "").strip()
    if env:
        return env
    if _SECRET_FILE.exists():
        return _SECRET_FILE.read_text(encoding="utf-8").strip()
    key = secrets.token_urlsafe(48)
    _SECRET_FILE.write_text(key, encoding="utf-8")
    try:
        os.chmod(_SECRET_FILE, 0o600)
    except OSError:
        pass
    return key


# --- デバッガは開かせない -----------------------------------------------------------
# デバッガが開いていると、例外が出たときにブラウザからこのサーバの Python を実行できる。
# PIN はユーザー名・MACアドレス・マシンIDから作る固定値で、守りにならない。

_TRUE = ("1", "true", "yes", "on", "t", "y")
_DEBUG_ARGS = ("--debugger", "--debug")
_NO_DEBUG_MSG = (
    "\n[app] デバッグモードでは起動しません。\n"
    "  理由: デバッガが開くと、例外が出たときにブラウザからこのサーバの Python を実行できます。\n"
    "  断ったもの: flask run の --debug / --debugger と、環境変数 FLASK_DEBUG=1。\n"
    "  通常の起動: flask --app app serve\n"
    "  詳しいエラーを見たいとき: app/__init__.py の DEBUG を True にして flask --app app serve\n"
)


def _refuse_debugger() -> None:
    if str(os.environ.get("FLASK_DEBUG") or "").strip().lower() in _TRUE:
        raise SystemExit(_NO_DEBUG_MSG)
    if any(a in _DEBUG_ARGS for a in sys.argv[1:]):
        raise SystemExit(_NO_DEBUG_MSG)


# --- 待ち受け先と、受け付ける宛先の名前 -------------------------------------------------
# サーバ（JupyterLab のターミナルなど）で起動し、社内LANの他のPCから数人で開いて使う（design.md 0）。
# 既定はこのサーバの中からだけ開ける 127.0.0.1。LAN に出すときは起動時に環境変数で渡す:
#   HOST=0.0.0.0 PORT=5000 flask --app app serve
_ALL_ADDRESSES = ("", "0.0.0.0", "::")


def _env_port(default: int = 5000) -> int:
    raw = (os.environ.get("PORT") or "").strip()
    if not raw:
        return default
    try:
        port = int(raw)
    except ValueError:
        raise SystemExit(f"\n[app] PORT が数字ではありません: {raw!r}\n  例: PORT=5000 flask --app app serve\n") from None
    if not 1 <= port <= 65535:
        raise SystemExit(f"\n[app] PORT が範囲外です: {port}（1〜65535）\n")
    return port


HOST = (os.environ.get("HOST") or "127.0.0.1").strip() or "127.0.0.1"
PORT = _env_port()


def _is_loopback(host) -> bool:
    # 空文字は 0.0.0.0 と同じで全てのネットワークから届くので loopback に含めない
    return str(host or "").strip().lower() in ("127.0.0.1", "localhost", "::1")


def _hostname_only(value) -> str:
    """"mypc:5000" や "[::1]:5000" から、宛先の名前だけを取り出す（小文字・末尾のドットは落とす）。"""
    try:
        name = urlsplit(f"//{str(value or '').strip()}").hostname or ""
    except ValueError:
        return ""
    return name.rstrip(".").lower()


def _lan_address() -> str:
    """LAN の他のPCから届くこのサーバのアドレス。取れなければ空（通信はしない）。"""
    sock = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    try:
        sock.connect(("192.0.2.1", 9))   # TEST-NET-1。UDP なので何も送らず、どの口から出るかだけ決めさせる
        return sock.getsockname()[0]
    except OSError:
        return ""
    finally:
        sock.close()


@lru_cache(maxsize=1)
def _machine_names() -> tuple[str, ...]:
    """このサーバ自身の名前とアドレス（LAN から届く宛先）。起動中は変わらない前提で1回だけ調べる。"""
    names: set[str] = set()
    try:
        hostname = socket.gethostname()
    except OSError:
        hostname = ""
    if hostname:
        names.update({hostname.lower(), hostname.split(".")[0].lower()})
        try:
            canonical, aliases, addresses = socket.gethostbyname_ex(hostname)
        except OSError:
            pass
        else:
            names.update(n.lower() for n in (canonical, *aliases) if n)
            names.update(addresses)
    lan = _lan_address()
    if lan:
        names.add(lan)
    # 日本語のPC名のときブラウザは punycode（xn--…）で送ってくるので、その形も受け付ける
    for name in list(names):
        try:
            names.add(name.encode("idna").decode("ascii").lower())
        except (UnicodeError, ValueError):
            pass
    return tuple(sorted(n for n in names if n))


def allowed_hosts(host: str | None = None) -> set[str]:
    """受け付ける宛先の名前（Host ヘッダー）。

    ログインが無いので、宛先の名前まで確かめる。攻撃者のドメインがこのサーバのアドレスを指していると、
    そのページとこのアプリが同一オリジンになり、画面の中身を読み取られてしまう（DNSリバインディング）。
    受け付けるのは、ループバックと、LAN に出しているときはこのサーバ自身の名前・アドレス。
    ロードバランサや別名で開くときは ALLOWED_HOSTS（`;` か `,` 区切り）で足す。`*` ですべて受け付ける。
    """
    bound = (HOST if host is None else host).strip().lower()
    names = {"127.0.0.1", "localhost", "::1"}
    for entry in (os.environ.get("ALLOWED_HOSTS") or "").replace(",", ";").split(";"):
        entry = entry.strip()
        if entry == "*":
            return {"*"}
        if entry:
            names.add(_hostname_only(entry) or entry.lower())
    if bound not in _ALL_ADDRESSES:
        names.add(_hostname_only(bound) or bound)
    if bound in _ALL_ADDRESSES or not _is_loopback(bound):
        names.update(_machine_names())   # LAN に出しているときだけ、PC名・LANのアドレスでも開ける
    return names


def _home_host() -> str:
    """画面や起動時の案内に出す「開けるアドレス」の宛先。"""
    bound = HOST.strip().lower()
    if _is_loopback(bound):
        return "127.0.0.1"
    if bound in _ALL_ADDRESSES:
        return _lan_address() or (_machine_names()[0] if _machine_names() else "127.0.0.1")
    return bound


def startup_notice(port: int | None = None) -> str:
    """起動したときに出す案内。LAN に出しているときは、ログインが無いことを必ず知らせる。"""
    url = f"http://{_home_host()}:{port or PORT}/"
    if _is_loopback(HOST):
        return f"[app] {url} で起動しました（このサーバの中からだけ開けます）"
    line = "=" * 68
    return (
        f"\n{line}\n"
        f"[app] 社内LANに公開して起動しました。ほかのPCからは次のアドレスを開いてください。\n"
        f"\n      {url}\n\n"
        f"  ・ログインはありません。このアドレスを知っている人は誰でも使えます。\n"
        f"    いま取り込んでいる帳票・一覧表の中身も、開かれれば見えます。\n"
        f"    信頼できる社内LANの中だけで使ってください。\n"
        f"  ・ダウンロードするとサーバーからデータが消えます。ダウンロードしていないものも\n"
        f"    しばらくすると捨てます。要るものはその場でダウンロードしてください。\n"
        f"  ・止めるとき: Ctrl+C（nohup で動かしているときは kill <PID>）\n"
        f"{line}\n"
    )


####################################################################################################
# 土台（旧 core.py・models/database.py）
####################################################################################################

# ====================================================================================================
# 元 models/database.py
# SQLite のスキーマ・マイグレーションとデータアクセス（帳票）。
#
# 一覧表・ジョブ・AI のテーブルもここで作る。一覧表のデータアクセスは tables/store.py、
# ジョブは core/jobs.py が持つ。
# ====================================================================================================

if TYPE_CHECKING:
    from app.extract import FieldDef, PatternDef, SheetDef

BUSY_TIMEOUT_MS = 5000

# 帳票の状態（導出値）: unread=読み取り前 / reviewing=確認中 / confirmed=確定済み / modified=修正中
# 画面に出す語は templates/base.html の STATE_LABELS が持つ
_STATE_SQL = """CASE
    WHEN d.data_json IS NULL THEN 'unread'
    WHEN d.confirmed_json IS NULL THEN 'reviewing'
    WHEN d.confirmed_json != d.data_json THEN 'modified'
    ELSE 'confirmed' END"""

# ---- スキーマ -----------------------------------------------------------------

# 第1版（作り直し前）のスキーマ。古いDBにも新規DBにも同じ形で適用できるよう IF NOT EXISTS
_SCHEMA_V1 = """
CREATE TABLE IF NOT EXISTS patterns (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name TEXT NOT NULL,
    version TEXT NOT NULL DEFAULT 'v1',
    description TEXT NOT NULL DEFAULT '',
    status TEXT NOT NULL DEFAULT 'draft',          -- draft / active / inactive
    image_processing TEXT NOT NULL DEFAULT 'none', -- none / vision
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS pattern_sheets (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    pattern_id INTEGER NOT NULL REFERENCES patterns(id) ON DELETE CASCADE,
    sheet_name TEXT NOT NULL
);

-- 古いDBにある required 列（項目の「必須」・シートの「必須」）はもう使わない。画面に必須の設定が
-- 無いので新しいDBでは作らず、読み書きもしない（古いDBには既定値のまま残る。あっても無くても同じ）
CREATE TABLE IF NOT EXISTS pattern_fields (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    pattern_id INTEGER NOT NULL REFERENCES patterns(id) ON DELETE CASCADE,
    sort_order INTEGER NOT NULL DEFAULT 0,
    field_name TEXT NOT NULL,
    display_name TEXT NOT NULL,
    candidates TEXT NOT NULL DEFAULT '[]',         -- JSON配列
    data_type TEXT NOT NULL DEFAULT 'string',
    extraction_rule TEXT NOT NULL DEFAULT '{}'     -- JSON {"direction", "columns", "section"}
);

CREATE TABLE IF NOT EXISTS pattern_samples (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    pattern_id INTEGER NOT NULL REFERENCES patterns(id) ON DELETE CASCADE,
    file_name TEXT NOT NULL,
    file_hash TEXT NOT NULL,
    stored_path TEXT NOT NULL,
    created_at TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS documents (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    file_name TEXT NOT NULL,
    file_hash TEXT NOT NULL,
    stored_path TEXT NOT NULL,
    pattern_id INTEGER REFERENCES patterns(id) ON DELETE SET NULL,
    status TEXT NOT NULL DEFAULT 'uploaded',       -- 旧: uploaded / extracted / registered（使わない）
    data_json TEXT,
    markdown TEXT,                                 -- 使わない（互換のため残す）
    created_at TEXT NOT NULL,
    registered_at TEXT                             -- 使わない（confirmed_at へ移行）
);
"""

_SCHEMA_TABLES = """
CREATE TABLE IF NOT EXISTS table_templates (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name TEXT NOT NULL UNIQUE,
    description TEXT NOT NULL DEFAULT '',
    current_version_id INTEGER,                    -- table_template_versions.id（循環するので外部キーにしない）
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS table_template_versions (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    template_id INTEGER NOT NULL REFERENCES table_templates(id) ON DELETE CASCADE,
    version INTEGER NOT NULL,
    spec_json TEXT NOT NULL,
    spec_hash TEXT NOT NULL,
    used INTEGER NOT NULL DEFAULT 0,
    note TEXT NOT NULL DEFAULT '',
    created_at TEXT NOT NULL,
    UNIQUE (template_id, version)
);

CREATE TABLE IF NOT EXISTS table_imports (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    template_id INTEGER REFERENCES table_templates(id) ON DELETE SET NULL,
    template_version_id INTEGER REFERENCES table_template_versions(id) ON DELETE SET NULL,
    file_name TEXT NOT NULL,
    file_hash TEXT NOT NULL,
    stored_path TEXT NOT NULL,
    source_json TEXT NOT NULL DEFAULT '{}',        -- {"kind","encoding","delimiter","preamble_rows","sheet","header_row","header_rows","data_end_row"}
    status TEXT NOT NULL DEFAULT 'uploaded',       -- uploaded / reading / preview / confirming / confirmed / discarded / failed
    stats_json TEXT NOT NULL DEFAULT '{}',
    issues_path TEXT,
    rows_path TEXT,
    job_id INTEGER,
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL,
    confirmed_at TEXT
);
CREATE INDEX IF NOT EXISTS idx_table_imports_template ON table_imports(template_id);

CREATE TABLE IF NOT EXISTS jobs (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    kind TEXT NOT NULL,
    ref_type TEXT NOT NULL DEFAULT '',
    ref_id INTEGER,
    status TEXT NOT NULL DEFAULT 'queued',         -- queued / running / paused / done / failed / cancelled / interrupted
    params_json TEXT NOT NULL DEFAULT '{}',
    progress_json TEXT NOT NULL DEFAULT '{}',
    message TEXT NOT NULL DEFAULT '',
    cancel_requested INTEGER NOT NULL DEFAULT 0,
    pause_requested INTEGER NOT NULL DEFAULT 0,
    heartbeat_at TEXT,
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_jobs_status ON jobs(status);
CREATE INDEX IF NOT EXISTS idx_jobs_ref ON jobs(ref_type, ref_id);

CREATE TABLE IF NOT EXISTS llm_calls (
    cache_key TEXT PRIMARY KEY,
    raw_text TEXT,
    parsed_json TEXT,
    model TEXT,
    params_json TEXT,
    structured_mode TEXT,
    finish_reason TEXT,
    tokens_in INTEGER,
    tokens_out INTEGER,
    latency_ms INTEGER,
    created_at TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS ai_items (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    template_id INTEGER NOT NULL,
    stage_id TEXT NOT NULL,
    row_key TEXT NOT NULL,
    template_version_id INTEGER,
    source_hash TEXT,
    context_hash TEXT,
    segments_hash TEXT,
    cache_key TEXT,
    status TEXT NOT NULL DEFAULT 'pending',        -- pending / ok / flagged / rule_only / error / skipped / outdated / excluded
    result_json TEXT,
    checks_json TEXT,
    attempts INTEGER NOT NULL DEFAULT 0,
    error TEXT,
    job_id INTEGER,
    updated_at TEXT NOT NULL,
    UNIQUE (template_id, stage_id, row_key)
);
CREATE INDEX IF NOT EXISTS idx_ai_items_status ON ai_items(template_id, stage_id, status);
"""


def _column_names(conn: sqlite3.Connection, table: str) -> set[str]:
    return {r[1] for r in conn.execute(f"PRAGMA table_info({table})").fetchall()}


def _add_column(conn: sqlite3.Connection, table: str, name: str, definition: str) -> None:
    """列が無ければ追加する（古いDBに途中まで手で足した場合も通るように）。"""
    if name not in _column_names(conn, table):
        conn.execute(f"ALTER TABLE {table} ADD COLUMN {name} {definition}")


def _run_script(conn: sqlite3.Connection, script: str) -> None:
    # executescript は途中で COMMIT するので、トランザクション内では1文ずつ実行する
    for statement in script.split(";"):
        lines = [ln for ln in statement.splitlines() if ln.strip()]
        if lines:
            conn.execute("\n".join(lines))


def _m1_base(conn: sqlite3.Connection) -> None:
    _run_script(conn, _SCHEMA_V1)


def _m2_forms(conn: sqlite3.Connection) -> None:
    _add_column(conn, "patterns", "title_fields", "TEXT NOT NULL DEFAULT '[]'")
    _add_column(conn, "patterns", "md_options", "TEXT NOT NULL DEFAULT '{}'")
    _add_column(conn, "patterns", "version_no", "INTEGER NOT NULL DEFAULT 1")
    _add_column(conn, "pattern_fields", "unit", "TEXT NOT NULL DEFAULT ''")
    _add_column(conn, "pattern_fields", "rag_output", "TEXT NOT NULL DEFAULT 'show'")
    _add_column(conn, "documents", "confirmed_json", "TEXT")
    _add_column(conn, "documents", "confirmed_at", "TEXT")
    _add_column(conn, "documents", "title", "TEXT NOT NULL DEFAULT ''")
    # 旧「登録済み」は確定済みとして引き継ぐ
    conn.execute("""UPDATE documents SET confirmed_json = data_json,
                        confirmed_at = COALESCE(registered_at, created_at)
                    WHERE status = 'registered' AND data_json IS NOT NULL AND confirmed_json IS NULL""")
    conn.execute("UPDATE documents SET confirmed_at = registered_at "
                 "WHERE confirmed_at IS NULL AND registered_at IS NOT NULL AND confirmed_json IS NOT NULL")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_documents_hash ON documents(file_hash)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_documents_pattern ON documents(pattern_id)")


def _m3_tables(conn: sqlite3.Connection) -> None:
    _run_script(conn, _SCHEMA_TABLES)


def _m4_form_batches(conn: sqlite3.Connection) -> None:
    """帳票のまとめ取り込み（1回の選択で複数ファイル）。同じ batch_id の帳票を順に確認して zip で渡す。"""
    _add_column(conn, "documents", "batch_id", "TEXT NOT NULL DEFAULT ''")
    _add_column(conn, "documents", "batch_order", "INTEGER NOT NULL DEFAULT 0")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_documents_batch ON documents(batch_id)")


# 使っていない表（8.1 で外した機能のもの）。取り込みを指す列が無く purge の探索に乗らないので、
# 行が入れば永久に残ってしまう。作らない機能の表は置かない（design.md 3.3）
_UNUSED_TABLES = ("table_outputs", "table_downloads", "table_template_samples", "alias_entries")


def _m5_purge_scope(conn: sqlite3.Connection) -> None:
    """AI整形の控えを「取り込み単位」で消せるようにし、使っていない表を落とす。

    ai_items は (設定, 段, 行) で一意なので、設定単位で消すと同じ設定で作業中の別の取り込みの結果まで
    消えていた（課金済みの応答キャッシュも道連れ）。import_id は次の _m6 で表を作り直すときに持たせる。
    """
    for table in _UNUSED_TABLES:
        conn.execute(f"DROP TABLE IF EXISTS {table}")


_AI_ITEMS_COLUMNS = ("id, template_id, stage_id, row_key, template_version_id, source_hash, context_hash, "
                     "segments_hash, cache_key, status, result_json, checks_json, attempts, error, job_id, "
                     "updated_at, import_id")


def _m6_ai_items_per_import(conn: sqlite3.Connection) -> None:
    """ai_items を「取り込み × 段 × 行」で一意にする（表を作り直す。SQLite は UNIQUE を外せないため）。

    (設定, 段, 行) で一意だと、同じ設定・同じ行の2つ目の取り込みは1つ目の結果を「処理済み」とみなして
    自分の行を持たず、1つ目をダウンロード（削除）すると2つ目の AI 結果まで消えていた（design.md 3.3）。
    """
    _add_column(conn, "ai_items", "import_id", "INTEGER")
    conn.execute("""CREATE TABLE ai_items_new (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        template_id INTEGER NOT NULL,
        stage_id TEXT NOT NULL,
        row_key TEXT NOT NULL,
        template_version_id INTEGER,
        source_hash TEXT,
        context_hash TEXT,
        segments_hash TEXT,
        cache_key TEXT,
        status TEXT NOT NULL DEFAULT 'pending',
        result_json TEXT,
        checks_json TEXT,
        attempts INTEGER NOT NULL DEFAULT 0,
        error TEXT,
        job_id INTEGER,
        updated_at TEXT NOT NULL,
        import_id INTEGER,
        UNIQUE (import_id, template_id, stage_id, row_key)
    )""")
    conn.execute(f"INSERT INTO ai_items_new ({_AI_ITEMS_COLUMNS}) SELECT {_AI_ITEMS_COLUMNS} FROM ai_items")
    conn.execute("DROP TABLE ai_items")
    conn.execute("ALTER TABLE ai_items_new RENAME TO ai_items")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_ai_items_status ON ai_items(template_id, stage_id, status)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_ai_items_import ON ai_items(import_id)")


def _m7_llm_calls_owner(conn: sqlite3.Connection) -> None:
    """AI の生の応答に「どの取り込みが払ったか」を持たせる（design.md 3.3）。

    ai_items が覚えているのは最後に使った応答のキーだけなので、再依頼で直した行の1回目の応答や、
    一時停止の直前に受け取った応答は、どの ai_items からも参照されない。持ち主が分からないと、
    その取り込みを消しても（ほかの取り込みが作業中の間は）生の応答が残ってしまう。
    古いDBの行は NULL（持ち主が分からない）のまま、これまでどおり扱う。
    """
    _add_column(conn, "llm_calls", "import_id", "INTEGER")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_llm_calls_import ON llm_calls(import_id)")


def _m8_session_scope(conn: sqlite3.Connection) -> None:
    """取り込んだものに「どのブラウザのものか」を持たせる（社内LANで数人が同時に使うため）。

    ログインは無いので利用者は分からないが、セッションのクッキー（views.current_session_id）で
    ブラウザごとの作業場所は分けられる。ほかのブラウザの帳票・取り込みは見えない（404）。
    古いDBの行は NULL（持ち主が分からない）のままで、これまでどおり扱う（起動時の片付けで消える）。
    帳票の種類・取り込み設定は「設定」なので分けない（みんなで使う）。
    """
    _add_column(conn, "documents", "session_id", "TEXT")
    _add_column(conn, "table_imports", "session_id", "TEXT")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_documents_session ON documents(session_id)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_table_imports_session ON table_imports(session_id)")


_TABLE_IMPORTS_COLUMNS = ("id, template_id, template_version_id, file_name, file_hash, stored_path, source_json, "
                          "status, stats_json, issues_path, rows_path, job_id, created_at, updated_at, "
                          "confirmed_at, session_id, spec_json, spec_hash")

def _m9_import_spec(conn: sqlite3.Connection) -> None:
    """取り込み設定の保存をやめ、その取り込みが使う設定（spec）を取り込みの行に持たせる。

    利用者の指示（2026-09-20）:「表の方には、取り込み設定を保持しておく機能はいらない」。
    取り込みごとに列の対応づけを決めるので、設定を名前で残して選び直す仕組み（table_templates と
    その版）は要らない。表を作り直すのは table_templates への外部キーを外すため（参照先の表を
    落とすと、外部キーを有効にした接続では table_imports への書き込みがすべて落ちる）。
    template_id / template_version_id は取り込み自身の id にそろえて残す（aiproc・core.purge が
    AI整形の控えを束ねる鍵に使っている）。
    """
    _add_column(conn, "table_imports", "spec_json", "TEXT NOT NULL DEFAULT ''")
    _add_column(conn, "table_imports", "spec_hash", "TEXT NOT NULL DEFAULT ''")
    if conn.execute("SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = 'table_template_versions'").fetchone():
        # まだダウンロードしていない取り込みが使っている版を、取り込みの行に写す
        conn.execute("""UPDATE table_imports SET
            spec_json = COALESCE((SELECT v.spec_json FROM table_template_versions v
                                  WHERE v.id = table_imports.template_version_id), ''),
            spec_hash = COALESCE((SELECT v.spec_hash FROM table_template_versions v
                                  WHERE v.id = table_imports.template_version_id), '')""")
    conn.execute("""CREATE TABLE table_imports_new (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        template_id INTEGER,                           -- = id（AI整形の控えを束ねる鍵。設定はもう無い）
        template_version_id INTEGER,                   -- = id（AI整形の控えの「作り直し判定」に使う）
        file_name TEXT NOT NULL,
        file_hash TEXT NOT NULL,
        stored_path TEXT NOT NULL,
        source_json TEXT NOT NULL DEFAULT '{}',        -- {"kind","encoding","delimiter","preamble_rows","sheet","header_rows","data_end_row"}
        status TEXT NOT NULL DEFAULT 'uploaded',       -- uploaded / reading / preview / confirming / confirmed / failed
        stats_json TEXT NOT NULL DEFAULT '{}',
        issues_path TEXT,
        rows_path TEXT,
        job_id INTEGER,
        created_at TEXT NOT NULL,
        updated_at TEXT NOT NULL,
        confirmed_at TEXT,
        session_id TEXT,
        spec_json TEXT NOT NULL DEFAULT '',            -- この取り込みの取り込み設定（tables.spec.TableSpec の JSON）
        spec_hash TEXT NOT NULL DEFAULT ''
    )""")
    conn.execute(f"INSERT INTO table_imports_new ({_TABLE_IMPORTS_COLUMNS}) "
                 f"SELECT {_TABLE_IMPORTS_COLUMNS} FROM table_imports")
    conn.execute("DROP TABLE table_imports")
    conn.execute("ALTER TABLE table_imports_new RENAME TO table_imports")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_table_imports_session ON table_imports(session_id)")
    conn.execute("DROP TABLE IF EXISTS table_template_versions")
    conn.execute("DROP TABLE IF EXISTS table_templates")
    conn.execute("UPDATE table_imports SET template_id = id, template_version_id = id")
    # AI整形の控えは「設定」ではなく取り込みで束ねる。持ち主の分からない古い行（import_id なし）は消す
    conn.execute("DELETE FROM ai_items WHERE import_id IS NULL")
    conn.execute("UPDATE ai_items SET template_id = import_id, template_version_id = import_id")


def _m10_drop_unused(conn: sqlite3.Connection) -> None:
    """使わなくなったものを落とす。

    - 互換ビュー table_template_versions: _m9 が一時的に置いたもの。aiproc.runner.load_spec が
      table_imports.spec_json を直接読むようになったので要らない。
    - table_imports.period_json: 書く側がもう無く、常に '{}' のまま（期間は取り込み設定が持つ）。
    """
    conn.execute("DROP VIEW IF EXISTS table_template_versions")
    if "period_json" in _column_names(conn, "table_imports"):
        conn.execute("ALTER TABLE table_imports DROP COLUMN period_json")


def _m11_document_touch(conn: sqlite3.Connection) -> None:
    """帳票に「最後にさわった時刻」を持たせる（見回りが作業中の帳票を消さないように）。

    一覧表（table_imports）は updated_at を持っていて、design.md 3.3 の「2時間さわられて
    いないものを捨てる」どおりに動く。帳票だけ取り込んだ時刻で切られていた。
    """
    _add_column(conn, "documents", "updated_at", "TEXT")
    conn.execute("UPDATE documents SET updated_at = COALESCE(confirmed_at, created_at) WHERE updated_at IS NULL")


def _m12_drop_pattern_samples(conn: sqlite3.Connection) -> None:
    """見本の Excel の控え（pattern_samples）を落とす。

    利用者の指示（2026-09-21）:「帳票登録で、見本のExcelは置かずに、設定だけ保持するように
    してほしい」。置いた Excel はその場で読み取るだけで、保存も控えもしない（views/form_types.py）。
    控えが残っていると、もう無いファイルを指し続ける行になるので表ごと落とす。
    uploads/samples に残ったファイルは、起動時に core.files.remove_sample_dir がフォルダごと片付ける。
    """
    conn.execute("DROP TABLE IF EXISTS pattern_samples")


def _m13_ai_connections(conn: sqlite3.Connection) -> None:
    """AI接続（APIキー・接続先・モデル）を「ブラウザごと」に持つ（利用者の指示 2026-09-21）。

    利用者の指示:「AI接続の設定は、もともとの位置ヘッダーの画面右上『AI接続』に移動させる。全部空欄にしておいて
    cookieでユーザー毎に登録内容をずっと保持させるようにしてほしい」。
    これまでは data/model_settings.yaml 1つを全員で使っていたので、社内LANで数人が使うと互いのキーを
    上書きし合い、1つのキーを共有していた。持ち主は帳票・取り込みと同じ session_id（views.current_session_id）。
    この表は「設定」なので、取り込みを捨てる片付け（core.purge_session / sweep_stale / purge_all_pending）の
    対象にしない（core.SETTINGS_TABLES）。最後の接続確認の結果もここに持ち、ヘッダーはそれを表示するだけ
    （画面を開くたびに確認しに行かない）。
    """
    conn.execute("""CREATE TABLE IF NOT EXISTS ai_connections (
        session_id TEXT PRIMARY KEY,                   -- ブラウザの作業場所（views.current_session_id）
        api_key TEXT NOT NULL DEFAULT '',              -- 平文（暗号化はしていない。画面にそう書く）
        chat_url TEXT NOT NULL DEFAULT '',             -- …/chat/completions までのフルパス（空欄＝サーバー共通の値）
        models_url TEXT NOT NULL DEFAULT '',           -- …/models までのフルパス（空欄＝サーバー共通の値）
        model TEXT NOT NULL DEFAULT '',                -- 使うモデル（空欄＝サーバー共通の既定）
        models_json TEXT NOT NULL DEFAULT '[]',        -- APIから取得した候補（datalist に出すだけ）
        last_check_ok INTEGER,                         -- NULL=未確認 / 1=つながった / 0=つながらなかった
        last_check_at TEXT,
        last_check_detail TEXT NOT NULL DEFAULT '',    -- 最後の確認の結果（つながらなかった理由）
        created_at TEXT NOT NULL,
        updated_at TEXT NOT NULL
    )""")


def _m14_pattern_books(conn: sqlite3.Connection) -> None:
    """帳票の種類に、登録したときのシートの中身（覚えたシート）を持たせる（利用者の指示 2026-09-22）。

    利用者の指示:「再度シートを置かなくても、登録したときにシートのセル番地と文字情報を記憶しておけばだせるはず」。
    Excel のファイルは今までどおり保存しない（2026-09-21「見本のExcelは置かずに、設定だけ保持する」）。
    覚えるのはシートの中身だけ: シート名・セルの番地と文字（型のある値）・結合・太字・塗りの有る無し。
    セルの色は覚えない（同日の指示「セル色の情報は不要」。塗りの有る無しと「同じ塗りか」の番号だけ。
    forms.fill_group）。書き方は forms.WorkbookInfo.to_json。種類1つにつき1つ（最後に置いた Excel の分）で、
    種類を消せば一緒に消える（ON DELETE CASCADE）。この表は「設定」なので、取り込みを捨てる片付けの
    対象にしない（core.SETTINGS_TABLES）。
    """
    conn.execute("""CREATE TABLE IF NOT EXISTS pattern_books (
        pattern_id INTEGER PRIMARY KEY REFERENCES patterns(id) ON DELETE CASCADE,
        file_name TEXT NOT NULL,                       -- 置いた Excel のファイル名（表示用。ファイルそのものは無い）
        file_hash TEXT NOT NULL,                       -- その Excel の sha256（ブラウザが送る book_hash と同じもの）
        book_json TEXT NOT NULL,                       -- 覚えたシートの中身（forms.WorkbookInfo.to_json）
        saved_at TEXT NOT NULL                         -- 置いた日時
    )""")


# PRAGMA user_version = 適用済みの件数。追加は末尾にだけ行う
MIGRATIONS: list[Callable[[sqlite3.Connection], None]] = [_m1_base, _m2_forms, _m3_tables, _m4_form_batches,
                                                          _m5_purge_scope, _m6_ai_items_per_import,
                                                          _m7_llm_calls_owner, _m8_session_scope, _m9_import_spec,
                                                          _m10_drop_unused, _m11_document_touch,
                                                          _m12_drop_pattern_samples, _m13_ai_connections,
                                                          _m14_pattern_books]


def migrate(conn: sqlite3.Connection) -> int:
    """未適用のマイグレーションを1つずつトランザクションで適用し、適用後の版を返す。"""
    current = conn.execute("PRAGMA user_version").fetchone()[0]
    for number, migration in enumerate(MIGRATIONS[current:], start=current + 1):
        conn.execute("BEGIN IMMEDIATE")
        try:
            migration(conn)
            conn.execute(f"PRAGMA user_version = {number}")
            conn.commit()
        except Exception:
            conn.rollback()
            raise
    return max(current, len(MIGRATIONS))


# ---- 接続 -----------------------------------------------------------------------

def connect(path: str | Path | None = None) -> sqlite3.Connection:
    """g を使わない接続（バックグラウンドスレッド用）。呼び出し側で close する。

    path 省略時はアプリ設定の DATABASE（app_context が必要）。
    """
    if path is None:
        path = current_app.config["DATABASE"]
    conn = sqlite3.connect(path, timeout=BUSY_TIMEOUT_MS / 1000, check_same_thread=False)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    conn.execute(f"PRAGMA busy_timeout = {BUSY_TIMEOUT_MS}")
    conn.execute("PRAGMA journal_mode = WAL")
    # 消した行の中身をその場でゼロ埋めする（design.md 3.3）。これが無いと、行を消しても解放された
    # ページに帳票の値・対象の名前・人名が平文で残り、DBファイル（と app.db-wal）から読めてしまう。
    # VACUUM は他の接続があると失敗することがあるので、消し方そのものを安全側にしておく。
    conn.execute("PRAGMA secure_delete = ON")
    return conn


def get_db() -> sqlite3.Connection:
    if "db" not in g:
        g.db = connect()
    return g.db


def close_db(_exc=None) -> None:
    db = g.pop("db", None)
    if db is not None:
        db.close()


def init_app(app) -> None:
    Path(app.config["DATABASE"]).parent.mkdir(parents=True, exist_ok=True)
    conn = connect(app.config["DATABASE"])
    try:
        migrate(conn)
    finally:
        conn.close()
    app.teardown_appcontext(close_db)


def now() -> str:
    return datetime.now().isoformat(timespec="seconds")


def _all(sql: str, args=()) -> list[dict]:
    return [dict(r) for r in get_db().execute(sql, args).fetchall()]


def _one(sql: str, args=()) -> dict | None:
    row = get_db().execute(sql, args).fetchone()
    return dict(row) if row else None


def _exec(sql: str, args=()) -> int:
    db = get_db()
    cur = db.execute(sql, args)
    db.commit()
    return cur.lastrowid


def _dumps_value(value) -> str:
    return value if isinstance(value, str) else json.dumps(value, ensure_ascii=False)


# ---- patterns ---------------------------------------------------------------

def list_patterns() -> list[dict]:
    return _all("""
        SELECT p.*,
               (SELECT COUNT(*) FROM pattern_fields f WHERE f.pattern_id = p.id) AS field_count,
               (SELECT COUNT(*) FROM documents d WHERE d.pattern_id = p.id AND d.confirmed_json IS NOT NULL) AS document_count,
               (SELECT COUNT(*) FROM documents d WHERE d.pattern_id = p.id AND d.confirmed_json IS NULL) AS working_count
        FROM patterns p ORDER BY p.name, p.version
    """)


def count_active_patterns() -> int:
    return get_db().execute("SELECT COUNT(*) FROM patterns WHERE status = 'active'").fetchone()[0]


def create_pattern(name: str) -> int:
    ts = now()
    return _exec(
        "INSERT INTO patterns (name, version, description, created_at, updated_at) VALUES (?, 'v1', '', ?, ?)",
        (name, ts, ts),
    )


def _field_def(row: dict) -> FieldDef:
    from app.extract import FieldDef, PatternDef, SheetDef   # extract が core を読み込むので、ここで
    fd = FieldDef(
        field_name=row["field_name"],
        display_name=row["display_name"],
        candidates=json.loads(row["candidates"]),
        data_type=row["data_type"],
        direction=json.loads(row["extraction_rule"]).get("direction", "auto"),
    )
    # unit / rag_output は WP-forms が FieldDef に追加する。未追加でも属性として持たせる
    fd.unit = row.get("unit") or ""
    fd.rag_output = row.get("rag_output") or "show"
    rule = json.loads(row["extraction_rule"])
    fd.table_columns = list(rule.get("columns") or [])
    # 探す区画（発行側・回答側など、同じラベルが並ぶ帳票でどちらを読むか）
    fd.section = rule.get("section") or ""
    # クリックで作った項目の控え（見本でのシート名とセル番地）
    fd.sheet_name = rule.get("sheet_name") or ""
    fd.label_cell = rule.get("label_cell") or ""
    fd.cell = rule.get("cell") or ""
    fd.renamed = bool(rule.get("renamed"))
    return fd


def _extraction_rule(f: FieldDef) -> dict:
    """pattern_fields.extraction_rule の JSON。明細表は列見出しも持つ。"""
    rule = {"direction": f.direction}
    columns = getattr(f, "table_columns", None)
    if f.data_type == "table" and columns:
        rule["columns"] = list(columns)
    section = getattr(f, "section", "") or ""
    if section:
        rule["section"] = section
    for key in ("sheet_name", "label_cell", "cell"):
        value = getattr(f, key, "") or ""
        if value:
            rule[key] = value
    if getattr(f, "renamed", False):
        rule["renamed"] = 1   # 見出しを手で直した項目
    return rule


def load_pattern(pattern_id: int | None) -> PatternDef | None:
    from app.extract import FieldDef, PatternDef, SheetDef   # extract が core を読み込むので、ここで
    row = _one("SELECT * FROM patterns WHERE id = ?", (pattern_id,))
    if row is None:
        return None
    sheets = _all("SELECT * FROM pattern_sheets WHERE pattern_id = ? ORDER BY id", (pattern_id,))
    fields = _all("SELECT * FROM pattern_fields WHERE pattern_id = ? ORDER BY sort_order, id", (pattern_id,))
    pattern = PatternDef(
        id=row["id"],
        name=row["name"],
        version=row["version"],
        description=row["description"],
        status=row["status"],
        image_processing=row["image_processing"],
        sheets=[SheetDef(s["sheet_name"]) for s in sheets],
        fields=[_field_def(f) for f in fields],
    )
    pattern.title_fields = json.loads(row.get("title_fields") or "[]")
    pattern.md_options = json.loads(row.get("md_options") or "{}")
    pattern.version_no = row.get("version_no") or 1
    return pattern


def load_active_patterns() -> list[PatternDef]:
    ids = [r["id"] for r in _all("SELECT id FROM patterns WHERE status = 'active' ORDER BY name, version")]
    return [load_pattern(i) for i in ids]


def save_pattern(pattern: PatternDef, status: str) -> None:
    """種類のメタ情報を更新し、シート・項目定義を置き換える。保存ごとに version_no を+1。"""
    db = get_db()
    with db:
        db.execute(
            """UPDATE patterns SET name=?, version=?, description=?, image_processing=?, status=?,
                      title_fields=?, md_options=?, version_no=version_no + 1, updated_at=? WHERE id=?""",
            (pattern.name, pattern.version, pattern.description, pattern.image_processing, status,
             json.dumps(list(getattr(pattern, "title_fields", None) or []), ensure_ascii=False),
             json.dumps(dict(getattr(pattern, "md_options", None) or {}), ensure_ascii=False),
             now(), pattern.id),
        )
        db.execute("DELETE FROM pattern_sheets WHERE pattern_id = ?", (pattern.id,))
        db.execute("DELETE FROM pattern_fields WHERE pattern_id = ?", (pattern.id,))
        db.executemany(
            "INSERT INTO pattern_sheets (pattern_id, sheet_name) VALUES (?, ?)",
            [(pattern.id, s.sheet_name) for s in pattern.sheets],
        )
        db.executemany(
            """INSERT INTO pattern_fields
               (pattern_id, sort_order, field_name, display_name, candidates, data_type, extraction_rule,
                unit, rag_output)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            [
                (pattern.id, i, f.field_name, f.display_name, json.dumps(f.candidates, ensure_ascii=False),
                 f.data_type, json.dumps(_extraction_rule(f), ensure_ascii=False),
                 getattr(f, "unit", "") or "", getattr(f, "rag_output", "show") or "show")
                for i, f in enumerate(pattern.fields)
            ],
        )
    row = db.execute("SELECT version_no FROM patterns WHERE id = ?", (pattern.id,)).fetchone()
    if row is not None:
        try:
            pattern.version_no = row[0]
        except AttributeError:
            pass


def set_pattern_status(pattern_id: int, status: str) -> None:
    _exec("UPDATE patterns SET status = ?, updated_at = ? WHERE id = ?", (status, now(), pattern_id))


def delete_pattern(pattern_id: int) -> None:
    _exec("DELETE FROM patterns WHERE id = ?", (pattern_id,))


# 見本の Excel は保存しないので、その控え（add_sample / list_samples …）は持たない。
# 帳票の種類が持つのは設定（シート名・見出しのセル・値のセル・読み取る向き・項目名）と、
# 登録したときのシートの中身（pattern_books。下）だけ。


# ---- pattern_books（覚えたシート）-------------------------------------------------------
# 帳票の種類1つにつき、最後に置いた Excel のシートの中身（forms.WorkbookInfo.to_json）を1つ持つ。
# Excel のファイルではない（セルの番地と文字・結合・太字・塗りの有る無しだけ。色は無い）。

def save_pattern_book(pattern_id: int, file_name: str, file_hash: str, book_json: str) -> None:
    """覚えたシートを入れる（前のものがあれば入れ替える）。"""
    _exec(
        """INSERT INTO pattern_books (pattern_id, file_name, file_hash, book_json, saved_at)
           VALUES (?, ?, ?, ?, ?)
           ON CONFLICT(pattern_id) DO UPDATE SET file_name = excluded.file_name, file_hash = excluded.file_hash,
                                                book_json = excluded.book_json, saved_at = excluded.saved_at""",
        (pattern_id, file_name, file_hash, book_json, now()),
    )


def load_pattern_book(pattern_id: int | None) -> dict | None:
    """覚えたシート（book_json を含む行）。無ければ None（覚える前に登録した種類）。"""
    return _one("SELECT * FROM pattern_books WHERE pattern_id = ?", (pattern_id,))


def pattern_book_meta(pattern_id: int | None) -> dict | None:
    """覚えたシートの見出し（ファイル名・sha256・置いた日時・JSON の大きさ）。中身は読まない。"""
    return _one("SELECT pattern_id, file_name, file_hash, saved_at, LENGTH(book_json) AS size "
                "FROM pattern_books WHERE pattern_id = ?", (pattern_id,))


def delete_pattern_book(pattern_id: int) -> None:
    """覚えたシートだけを捨てる（種類は残る。テスト・覚える前の状態に戻すとき用）。"""
    _exec("DELETE FROM pattern_books WHERE pattern_id = ?", (pattern_id,))


# ---- documents --------------------------------------------------------------
# 状態は導出する: data_json なし→unread、confirmed_json なし→reviewing、
# confirmed_json != data_json→modified、それ以外→confirmed

def create_document(file_name: str, file_hash: str, stored_path: str, pattern_id: int | None = None,
                    batch_id: str = "", batch_order: int = 0, session_id: str | None = None) -> int:
    """帳票を1件作る。session_id は取り込んだブラウザ（views.current_session_id）。"""
    return _exec(
        "INSERT INTO documents (file_name, file_hash, stored_path, pattern_id, batch_id, batch_order, session_id, "
        "created_at, updated_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
        (file_name, file_hash, stored_path, pattern_id, batch_id, batch_order, session_id, now(), now()),
    )


# 持ち主で絞る条件。持ち主の分からない行（この仕組みより前のDBの行）は、これまでどおり誰からでも扱える
_OWNED = "(d.session_id IS NULL OR d.session_id = ?)"


def _owner_where(session_id: str | None) -> tuple[str, list]:
    return (_OWNED, [session_id]) if session_id else ("", [])


def get_document(doc_id: int) -> dict | None:
    return _one(f"""
        SELECT d.*, {_STATE_SQL} AS state,
               p.name AS pattern_name, p.version AS pattern_version
        FROM documents d LEFT JOIN patterns p ON p.id = d.pattern_id
        WHERE d.id = ?
    """, (doc_id,))


_LIST_COLUMNS = f"""
        SELECT d.id, d.file_name, d.file_hash, d.stored_path, d.pattern_id, d.title, d.batch_id, d.batch_order,
               d.session_id, d.created_at, d.confirmed_at, {_STATE_SQL} AS state,
               p.name AS pattern_name, p.version AS pattern_version,
               -- 帳票の種類が削除されても、読み取ったときの種類名を表示に使う
               CASE WHEN json_valid(d.data_json) THEN json_extract(d.data_json, '$.pattern.name') END
                   AS extraction_pattern_name
        FROM documents d LEFT JOIN patterns p ON p.id = d.pattern_id"""


def list_batch_documents(batch_id: str, session_id: str | None = None) -> list[dict]:
    """まとめ取り込み（1回の選択で複数ファイル）の帳票を、選んだ順に返す。

    session_id を渡すと、そのブラウザの帳票だけを返す（ほかの人のまとまりは見えない）。
    """
    if not batch_id:
        return []
    owner, args = _owner_where(session_id)
    where = "WHERE d.batch_id = ?" + (f" AND {owner}" if owner else "")
    return _all(f"{_LIST_COLUMNS} {where} ORDER BY d.batch_order, d.id", (batch_id, *args))


def list_confirmed_documents(ids: list[int] | None = None, session_id: str | None = None) -> list[dict]:
    """確定済みの版を持つ帳票（修正中も確定済みの版を持つ）。ids・session_id 指定で絞り込み。"""
    sql = f"""SELECT d.*, {_STATE_SQL} AS state, p.name AS pattern_name, p.version AS pattern_version
              FROM documents d LEFT JOIN patterns p ON p.id = d.pattern_id
              WHERE d.confirmed_json IS NOT NULL"""
    owner, args = _owner_where(session_id)
    if owner:
        sql += f" AND {owner}"
    if ids is not None:
        if not ids:
            return []
        sql += f" AND d.id IN ({', '.join('?' for _ in ids)})"
        return _all(sql + " ORDER BY d.id", (*args, *ids))
    return _all(sql + " ORDER BY d.id", tuple(args))


def save_draft(doc_id: int, data_json, title: str | None = None) -> None:
    """作業中の値を保存する（dict も可）。title は検索用（タイトル項目の値）。"""
    if title is None:
        _exec("UPDATE documents SET data_json = ?, updated_at = ? WHERE id = ?",
              (_dumps_value(data_json), now(), doc_id))
    else:
        _exec("UPDATE documents SET data_json = ?, title = ?, updated_at = ? WHERE id = ?",
              (_dumps_value(data_json), title, now(), doc_id))


def confirm_document(doc_id: int, title: str | None = None) -> bool:
    """作業中の値を確定済みの版にする。読み取り前（data_json なし）なら False。"""
    sets, args = ["confirmed_json = data_json", "confirmed_at = ?", "updated_at = ?"], [now(), now()]
    if title is not None:
        sets.append("title = ?")
        args.append(title)
    db = get_db()
    cur = db.execute(f"UPDATE documents SET {', '.join(sets)} WHERE id = ? AND data_json IS NOT NULL",
                     (*args, doc_id))
    db.commit()
    return cur.rowcount > 0


def reset_document(doc_id: int, pattern_id: int | None, data_json) -> None:
    """選び直して再読み取り：作業中の値を置き換える（確定済みの版は残る→修正中になる）。"""
    _exec("UPDATE documents SET pattern_id = ?, data_json = ?, updated_at = ? WHERE id = ?",
          (pattern_id, None if data_json is None else _dumps_value(data_json), now(), doc_id))


def find_confirmed_by_hash(file_hash: str, exclude_id: int | None = None,
                           session_id: str | None = None) -> dict | None:
    """同じ中身の確定済みの帳票（取り込み直しの注意に出す）。ほかの人の帳票は知らせない。"""
    owner, args = _owner_where(session_id)
    sql = ("SELECT d.id, d.file_name, d.title FROM documents d WHERE d.file_hash = ? "
           "AND d.confirmed_json IS NOT NULL AND d.id != ?")
    if owner:
        sql += f" AND {owner}"
    return _one(sql + " ORDER BY d.id LIMIT 1",
                (file_hash, exclude_id if exclude_id is not None else -1, *args))


def update_document(doc_id: int, **columns) -> None:
    assignments = ", ".join(f"{name} = ?" for name in columns)
    _exec(f"UPDATE documents SET {assignments} WHERE id = ?", (*columns.values(), doc_id))


# 帳票を消すのは core/purge.py（元のファイルと DB の行をまとめて消す）


# ---- ai_connections（ブラウザごとの AI接続） ---------------------------------------------------
# 持ち主は session_id（views.current_session_id）。行は「保存」か「接続の確認」で初めてできる（読むだけでは作らない）。
# 取り込みの片付け（core.purge_*）では消えない「設定」。長く使われていない行だけ core.sweep_stale_ai_connections が消す。

AI_CONNECTION_FIELDS = ("api_key", "chat_url", "models_url", "model")


def get_ai_connection(session_id: str | None) -> dict | None:
    """そのブラウザの AI接続の行（無ければ None）。models は list にして返す。"""
    if not session_id:
        return None
    row = _one("SELECT * FROM ai_connections WHERE session_id = ?", (session_id,))
    if row is None:
        return None
    try:
        row["models"] = [str(m) for m in json.loads(row.get("models_json") or "[]") if str(m).strip()]
    except ValueError:
        row["models"] = []
    return row


def save_ai_connection_row(session_id: str, *, api_key: str | None = None, chat_url: str | None = None,
                       models_url: str | None = None, model: str | None = None,
                       models: list[str] | None = None) -> dict:
    """そのブラウザの AI接続を保存する（None の項目は変えない）。接続先・キーを変えたら確認の結果は「未確認」に戻す。"""
    ts = now()
    db = get_db()
    with db:
        row = db.execute("SELECT * FROM ai_connections WHERE session_id = ?", (session_id,)).fetchone()
        if row is None:
            db.execute("INSERT INTO ai_connections (session_id, created_at, updated_at) VALUES (?, ?, ?)",
                       (session_id, ts, ts))
        sets, args = ["updated_at = ?"], [ts]
        for name, value in (("api_key", api_key), ("chat_url", chat_url), ("models_url", models_url), ("model", model)):
            if value is not None:
                sets.append(f"{name} = ?")
                args.append(str(value))
        if models is not None:
            sets.append("models_json = ?")
            args.append(json.dumps([str(m) for m in models], ensure_ascii=False))
        # つながるかどうかは接続先・キー・モデルで決まる。どれかが変わったら前の結果は当てにならない
        changed = row is None or any(value is not None and str(value) != (row[name] or "")
                                     for name, value in (("api_key", api_key), ("chat_url", chat_url),
                                                         ("models_url", models_url), ("model", model)))
        if changed:
            sets.append("last_check_ok = NULL")
            sets.append("last_check_at = NULL")
            sets.append("last_check_detail = ''")
        db.execute(f"UPDATE ai_connections SET {', '.join(sets)} WHERE session_id = ?", (*args, session_id))
    return get_ai_connection(session_id) or {}


def set_ai_connection_check(session_id: str, ok: bool, detail: str = "") -> None:
    """接続の確認の結果（ヘッダーの「接続中／つながりません」と「最終確認」）を覚える。行が無ければ作る。"""
    ts = now()
    db = get_db()
    with db:
        db.execute("INSERT OR IGNORE INTO ai_connections (session_id, created_at, updated_at) VALUES (?, ?, ?)",
                   (session_id, ts, ts))
        db.execute("UPDATE ai_connections SET last_check_ok = ?, last_check_at = ?, last_check_detail = ?, "
                   "updated_at = ? WHERE session_id = ?",
                   (1 if ok else 0, ts, str(detail or "")[:500], ts, session_id))


def clear_ai_connection_key(session_id: str) -> bool:
    """［キーを消す］（共有PC）。接続先・モデルは残し、キーと確認の結果だけ消す。戻り値: 行があったか。"""
    db = get_db()
    with db:
        cur = db.execute("UPDATE ai_connections SET api_key = '', last_check_ok = NULL, last_check_at = NULL, "
                         "last_check_detail = '', updated_at = ? WHERE session_id = ?", (now(), session_id))
    return cur.rowcount > 0






# ====================================================================================================
# 元 core/naming.py
# 出力ファイル名（Windows と LightRAG の両方で安全な名前）。
#
# LightRAG のファイル名ヒント（`.[legacy-R(...)]`）は付けない。ヒントはサーバー側の取り込み設定より
# 優先されてしまう上に、そのサーバーが知らない書き方だと取り込みが HTTP 400 で断られる。
# 名前は「安定・意味が分かる・重複しない」だけを満たし、チャンクへの耐性は本文の作り方（記録の分割）で確保する。
# ====================================================================================================

_UNSAFE = re.compile(r'[\\/:*?"<>|\[\]\s\x00-\x1f\x7f]')
_WINDOWS_RESERVED = {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))}


def safe_filename_part(text, max_len: int = 60) -> str:
    """ファイル名の1部品。NFKC、禁止文字・空白・角かっこを _ に、'.[' を除去、前後の . _ を除去。

    ここは本文と違って NFKC をそのままかける（囲み文字は残さない）。ファイル名は同じ文書に対して
    安定・一意であればよく、①などを残すと OS や LightRAG 側での扱いが揺れるため。
    """
    if text is None:
        return ""
    s = unicodedata.normalize("NFKC", str(text))
    s = s.replace(".[", "")  # LightRAG のヒント記法と誤認されないように
    s = _UNSAFE.sub("_", s)
    s = re.sub(r"_+", "_", s).strip("._")
    s = s[:max_len].strip("._")
    if s.split(".")[0].upper() in _WINDOWS_RESERVED:
        s = f"{s}_"
    return s


def md_filename(parts: list[str]) -> str:
    """部品を _ で連結した .md ファイル名。空の部品は飛ばす。"""
    safe = [p for p in (safe_filename_part(x) for x in parts) if p]
    return f"{'_'.join(safe) or '無題'}.md"


# ====================================================================================================
# 元 core/mdtext.py
# Markdown 出力のテキスト処理（帳票・一覧表で共通）。
#
# 決まり（docs/design.md 6章）: UTF-8・LF、レコード内に空行を入れない、パイプ表を使わない（`- 項目: 値`）、
# 複数行の値は2文字下げの連続行。
# ====================================================================================================

_SPACES = re.compile(r"[^\S\n]+")          # 改行以外の空白の連続
_HEADING = re.compile(r"^(#{1,6})(\s|$)")
_LIST = re.compile(r"^([-*+])(\s|$)")
_ORDERED = re.compile(r"^(\d{1,9})([.)])(\s|$)")
_RULE = re.compile(r"^(?:-[ \t]*){3,}$|^(?:=[ \t]*){3,}$|^(?:\*[ \t]*){3,}$|^(?:_[ \t]*){3,}$")
_FENCE = re.compile(r"^(```|~~~)")


def _enclosed_marks() -> str:
    """NFKC で囲みが外れてしまう囲み文字（丸数字 ①、丸英字 Ⓐ、丸カナ ㋐、丸漢字 ㊤ など）を集める。"""
    marks = []
    for start, end in ((0x2460, 0x24FF), (0x3240, 0x32FF), (0x1F100, 0x1F1FF)):
        for cp in range(start, end + 1):
            ch = chr(cp)
            if "CIRCLED" in unicodedata.name(ch, "") and unicodedata.normalize("NFKC", ch) != ch:
                marks.append(ch)
    return "".join(marks)


# ①→1 のように囲みが外れると「①破損…」が「1破損…」になり、番号と本文の区切りが消えて読めなくなる。
# 丸数字は帳票の手順・項目番号でごく普通に使われるので、囲み文字だけは NFKC をかけずに残す。
# （㈱→(株)、⑴→(1) のように区切りが残る表記はそのまま NFKC で正規化する）
_ENCLOSED = re.compile(f"([{re.escape(_enclosed_marks())}])")


def nfkc_keep_enclosed(text: str) -> str:
    """NFKC 正規化。ただし囲み文字（①Ⓐ㋐㊤…）はそのまま残す。

    囲み文字は前後の文字と結合しないので、そこで区切って正規化しても結果は変わらない。
    """
    if not _ENCLOSED.search(text):
        return unicodedata.normalize("NFKC", text)
    return "".join(part if _ENCLOSED.fullmatch(part) else unicodedata.normalize("NFKC", part)
                   for part in _ENCLOSED.split(text))


def nfkc_value(text) -> str:
    """値の正規化。NFKC（囲み文字は残す）＋空白の畳み込み（改行は保持、行末の空白と前後の空行は除く）。"""
    if text is None:
        return ""
    s = nfkc_keep_enclosed(str(text)).replace("\r\n", "\n").replace("\r", "\n")
    lines = [_SPACES.sub(" ", line).strip() for line in s.split("\n")]
    return "\n".join(lines).strip("\n")


def escape_md_line(line: str) -> str:
    """行頭の見出し・箇条書き・引用・番号付きリスト、区切り線、コードフェンスとして解釈されないようにする。"""
    body = line.lstrip(" \t")
    indent = line[: len(line) - len(body)]
    if not body:
        return line
    if _FENCE.match(body):
        body = "\\" + body[0] + "\\" + body[1] + "\\" + body[2:]
    elif _RULE.match(body):
        body = "\\" + body
    elif _HEADING.match(body) or _LIST.match(body) or body.startswith(">"):
        body = "\\" + body
    else:
        m = _ORDERED.match(body)
        if m:
            body = m.group(1) + "\\" + body[len(m.group(1)):]
    return indent + body


def _value_lines(value) -> list[str]:
    if value is None:
        return []
    items = value if isinstance(value, (list, tuple)) else [value]
    lines: list[str] = []
    for item in items:
        if item is None:
            continue
        text = str(item).replace("\r\n", "\n").replace("\r", "\n")
        lines += [ln.strip() for ln in text.split("\n") if ln.strip()]
    return lines


def md_bullet(label: str, value) -> list[str]:
    """`- label: value`。複数行は `- label:` の後に2文字下げで続ける。値が空なら出さない（[]）。"""
    lines = _value_lines(value)
    if not lines:
        return []
    label = " ".join(str(label).split())
    if len(lines) == 1:
        return [f"- {label}: {lines[0]}"]
    return [f"- {label}:"] + [f"  {escape_md_line(ln)}" for ln in lines]


# estimate_tokens 用：UTF-8 のバイトを種類の印に置き換える表（数字 → "0"、ASCII の記号 → "."、それ以外 → "x"）
_TOKEN_CLASS = bytes(
    ord("0") if 0x30 <= b <= 0x39 else ord(".") if (0x21 <= b <= 0x2F or 0x3A <= b <= 0x40 or 0x5B <= b <= 0x60
                                                     or 0x7B <= b <= 0x7E) else ord("x")
    for b in range(256))


def estimate_tokens(text: str) -> int:
    """推定トークン数（実トークン以上になる見積もり）。

    非ASCII 1文字=1.1、ASCII の記号 1文字=1、数字は「連続する数字のまとまり1つ=1 ＋ 3桁ごとに1」、
    それ以外の ASCII（英字・空白・改行）2文字=1。
    LightRAG の o200k_base は数字を3桁ずつに区切り、記号（- : / . = など）もほぼ1文字ずつ別のトークンにする。
    日時・品番・計測値の多い記録（「2023-09-01 09:44」は実10トークン）を英字と同じ2文字=1で数えると、
    実トークンより3割ほど少なく見積もり、記録の上限（チャンク 1,500）を超えることがあった。
    実出力 142,952 ブロックとの実測で、この式は実/推定の最大 0.99（旧式は 1.32）、全体では 23% 多めに見積もる。
    例外：まれな漢字（髙・﨑 など）は1文字が2〜3トークンになる。記録全体では他の文字の余裕に吸収される。
    """
    if not text:
        return 0
    ascii_count = len(text.encode("ascii", "ignore"))  # ASCII の文字数（1文字ずつ数えるより速い）
    marks = text.encode("utf-8", "surrogatepass").translate(_TOKEN_CLASS)
    digits = marks.count(b"0")
    punct = marks.count(b".")
    digit_runs = marks.count(b"x0") + marks.count(b".0") + (marks[:1] == b"0")   # 数字のまとまりの数
    other = ascii_count - digits - punct
    # 30分の1トークン単位の整数で数えて ceil する（浮動小数の誤差を避ける）
    total = 33 * (len(text) - ascii_count) + 15 * other + 30 * punct + 10 * digits + 30 * digit_runs
    return -(-total // 30)


def join_blocks(blocks: list[list[str]]) -> str:
    """ブロック間に空行1つ、ブロック内の空行は除く、末尾改行1つ、LF。"""
    parts: list[str] = []
    for block in blocks:
        lines: list[str] = []
        for line in block or []:
            for ln in str(line).replace("\r\n", "\n").replace("\r", "\n").split("\n"):
                if ln.strip():
                    lines.append(ln.rstrip())
        if lines:
            parts.append("\n".join(lines))
    return "\n\n".join(parts) + "\n" if parts else ""


# ====================================================================================================
# 元 core/files.py
# アップロードファイルの保存・事前チェック・削除。
#
# 保存は分割して読みながら sha256 を計算する（ブック全体を解析してからハッシュを取らない）。
# Excel の事前チェックは openpyxl で開く前に、先頭バイトと zip の目次だけで行う。
# 帳票登録の見本の Excel は保存しない（read_upload でメモリに読むだけ）ので、事前チェックは
# パスのほかに「中身そのもの（bytes）」も受け取れる。
# ====================================================================================================

logger = logging.getLogger(__name__)

CHUNK_SIZE = 1024 * 1024

OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
ZIP_MAGIC = b"PK\x03\x04"
STRICT_NS = "http://purl.oclc.org/ooxml/spreadsheetml/main"

# zip 展開の上限（値は仮置き）
ZIP_MAX_TOTAL = 500 * 1024 * 1024
ZIP_MAX_PART = 200 * 1024 * 1024
ZIP_MAX_RATIO = 100
# 小さいパーツは圧縮率が高くても害がないので、この大きさ以上だけ圧縮率を見る
ZIP_RATIO_MIN_BYTES = 16 * 1024 * 1024
# 結合セルの面積（ブック全体の合計）の上限。openpyxl は開くときに結合範囲のセルを1つずつ作るので、
# シート全体の結合（A1:XFD1048576 など）が1つあるだけで、上のサイズ制限より前に固まる（数秒〜終わらない）。
# 行全体の結合（A1:XFD1 = 16,384 セル）や列全体の結合（A:A = 約105万セル）は通す。
MAX_MERGED_CELLS = 2_000_000
# 帳票（通常モードで開く）の結合セルの面積の上限。openpyxl は結合範囲ごとに中のセルを1つずつ片付け、
# 左上のセルに罫線があれば範囲の辺のセルを1つずつ作るので、面積に比例して遅くなる（行全体の結合 120 個
# ＝約200万セルで約16秒、罫線付きの列全体の結合1つで1分以上）。帳票は画面を開くたびにブックを開き直すので、
# 1回あたり2秒程度に収まるこの値で断る（行全体の結合なら12個まで通る。ふつうの帳票は数千セル）。
# 一覧表は読み取り専用モードで開き結合を展開しないので、上の MAX_MERGED_CELLS のまま。
FORM_MAX_MERGED_CELLS = 200_000
# ハイパーリンク・コメントの範囲の面積（ブック全体の合計）の上限。openpyxl はリンク（<hyperlink ref>）とコメント
# （<comment ref>）の範囲のセルを1つずつ作る（帳票の通常モードでも、一覧表の tables.excel_source でも）。
# シート全体を指す範囲1つで開くのが終わらなくなる（A1:XFD50 の約82万セルで一覧表は約28秒）。
# ふつうのリンク・コメントは1セルか数セルなので、1回あたり2秒以内に収まるこの値で断る。
MAX_LINKED_CELLS = 50_000
# セル数の上限（ブック全体の <c> の数）。openpyxl は開くときにセルを1つずつ作るので、
# 圧縮すると小さいが展開すると大量のセルがあるブック（64KB で 100万セルなど）は、上のサイズ制限を通っても固まる。
# 一覧表は tables.excel_source が別に上限（EXCEL_MAX_CELLS）を持ち「CSVで保存」と案内するので、ここは最後の砦の値。
MAX_CELLS = 1_000_000
# 図形（描画）の上限。openpyxl は開くときに、シートが参照する描画部品（xl/drawings/*.xml）をシートごとに読み直し、
# 図形（アンカー）を1つずつ作る。1つの大きな描画を多数のシートから参照させると、圧縮後は小さくても
# 開くたびに（書類の画面を開くたびにも）数十秒〜終わらない。そこで「描画部品の図形数 × 参照するシート数」と
# 「描画部品の展開後の大きさ × 参照するシート数」の合計を数えて上限を設ける。
# 普通の帳票は1シートに数十個程度の図形なので、十分に余裕のある値にしている。
MAX_DRAWING_ANCHORS = 10_000
MAX_DRAWING_BYTES = 20 * 1024 * 1024
# 図形（アンカー）として数える要素の名前（DrawingML の spreadsheetDrawing）
_ANCHOR_NAMES = ("absoluteAnchor", "oneCellAnchor", "twoCellAnchor")
_DRAWING_REL_SUFFIX = "/drawing"
# openpyxl がシートの数だけ読み直す部品の関係の種類。シートの部品（worksheet）は <sheet> 1つにつき1回、
# コメントの部品（comments）はそれを参照するシートを読むたびに読み直される。1つの部品を複数のシートから
# 参照させると（Excel は作らないが、手で作れば開ける）、結合セル・セル・行・コメントの範囲もその回数だけ
# 作られるので、部品1つ分だけ数えていた上限をいくらでもすり抜けられる。参照される回数を掛けて数える。
_REREAD_REL_SUFFIXES = ("/worksheet", "/comments")
_RELS_SUFFIX = ".rels"
_DRAWING_ERROR = "図形や画像の数が多すぎるため読み込めません。不要な図形・画像を削除して保存し直してください"
# セルとして数える要素の名前空間（SpreadsheetML 本体の <c> だけ。グラフの <c:chart> などは名前空間が違う）
_SHEET_NAMESPACES = ("http://schemas.openxmlformats.org/spreadsheetml/2006/main", STRICT_NS)
# workbook.xml の名前空間判定で読む先頭バイト数
_WORKBOOK_HEAD_BYTES = 64 * 1024
# 掴まれているファイルを消すときの再試行（ウイルス対策のスキャンなどは短時間で終わる）
REMOVE_RETRIES = 3
REMOVE_RETRY_WAIT = 0.05
# アップロードの保存先（UPLOAD_DIR 直下）と save_upload が付けるファイル名の形。
# 帳票登録の見本（samples）はもう保存しないので、ここには入れない（前の版の残骸は remove_sample_dir が消す）
UPLOAD_SUBDIRS = ("documents", "tables")
_STORED_NAME = re.compile(r"[0-9a-f]{32}\.[A-Za-z0-9]+")
# 前の版が見本の Excel を置いていたフォルダ（UPLOAD_DIR 直下）
SAMPLE_SUBDIR = "samples"


class UploadError(Exception):
    """利用者に見せる日本語メッセージを持つ例外。"""


@dataclass
class StoredFile:
    stored_path: str   # UPLOAD_DIR からの相対パス（/ 区切り）
    file_name: str     # 元のファイル名（表示・ダウンロード用）
    file_hash: str     # sha256
    size: int


@dataclass
class MemoryFile:
    """保存せずにメモリへ読んだアップロード（帳票登録の見本の Excel）。"""

    data: bytes        # ブックの中身そのもの
    file_name: str     # 元のファイル名（表示用）
    file_hash: str     # sha256

    @property
    def size(self) -> int:
        return len(self.data)


def original_name(storage) -> str:
    # secure_filename は日本語を消してしまうので、表示用には元のファイル名（パス部分除去）を使う
    return Path((getattr(storage, "filename", None) or "").replace("\\", "/")).name


def _normalize_ext(ext: str) -> str:
    ext = ext.strip().lower()
    return ext if ext.startswith(".") else f".{ext}"


def _format_size(size: int) -> str:
    return f"{size / 1024 / 1024:.0f}MB" if size >= 1024 * 1024 else f"{size / 1024:.0f}KB"


def upload_path(stored_path: str) -> Path:
    """保存パス → 実ファイルのパス。UPLOAD_DIR の外を指すパスは拒否する。"""
    base = Path(current_app.config["UPLOAD_DIR"]).resolve()
    path = (base / stored_path).resolve()
    if path != base and base not in path.parents:
        raise UploadError("保存先のパスが不正です")
    return path


def _checked_name(storage, allowed: set[str]) -> tuple[str, str]:
    """元のファイル名と拡張子。名前が無い・扱えない拡張子なら UploadError。"""
    name = original_name(storage)
    if not name:
        raise UploadError("ファイルを選択してください")
    ext = Path(name).suffix.lower()
    allowed_exts = {_normalize_ext(e) for e in allowed}
    if ext not in allowed_exts:
        kinds = " / ".join(sorted(allowed_exts))
        hint = "（.xls はExcelで .xlsx に保存し直してください）" if ext == ".xls" else ""
        raise UploadError(f"{kinds} のファイルを選んでください{hint}")
    return name, ext


def read_upload(storage, allowed: set[str], max_bytes: int) -> MemoryFile:
    """アップロードをディスクに置かずにメモリへ読む。拡張子・サイズ・空ファイルを確認する。

    帳票登録の見本の Excel に使う。置いた Excel はサーバーに残さない（利用者の指示 2026-09-21
    「見本のExcelは置かずに、設定だけ保持するようにしてほしい」）ので、保存先を作らずに
    中身と sha256 だけを返す。呼び出し側は読み取った結果（WorkbookInfo）だけを短い間覚えておく。
    """
    name, _ext = _checked_name(storage, allowed)
    digest = hashlib.sha256()
    chunks: list[bytes] = []
    size = 0
    stream = getattr(storage, "stream", storage)
    while True:
        chunk = stream.read(CHUNK_SIZE)
        if not chunk:
            break
        size += len(chunk)
        if size > max_bytes:
            raise UploadError(f"ファイルが大きすぎます（上限 {_format_size(max_bytes)}）")
        digest.update(chunk)
        chunks.append(chunk)
    if size == 0:
        raise UploadError("ファイルが空です")
    return MemoryFile(data=b"".join(chunks), file_name=name, file_hash=digest.hexdigest())


def save_upload(storage, subdir: str, allowed: set[str], max_bytes: int) -> StoredFile:
    """アップロードを UPLOAD_DIR/subdir に保存する。拡張子・サイズ・空ファイルを確認し、失敗時は消して UploadError。"""
    name, ext = _checked_name(storage, allowed)
    stored = f"{subdir.strip('/')}/{uuid4().hex}{ext}"
    dest = upload_path(stored)
    dest.parent.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256()
    size = 0
    stream = getattr(storage, "stream", storage)
    try:
        with open(dest, "wb") as out:
            while True:
                chunk = stream.read(CHUNK_SIZE)
                if not chunk:
                    break
                size += len(chunk)
                if size > max_bytes:
                    raise UploadError(f"ファイルが大きすぎます（上限 {_format_size(max_bytes)}）")
                digest.update(chunk)
                out.write(chunk)
        if size == 0:
            raise UploadError("ファイルが空です")
    except BaseException:
        dest.unlink(missing_ok=True)
        raise
    return StoredFile(stored_path=stored, file_name=name, file_hash=digest.hexdigest(), size=size)


def precheck_excel(source, max_cells: int | None = None, max_merged: int | None = None,
                   max_rows: int | None = None) -> None:
    """Excel（.xlsx/.xlsm）として開いてよいかを、中身を展開せずに確かめる。問題があれば UploadError。

    source: ファイルのパス、またはブックの中身そのもの（bytes）。帳票登録の見本の Excel は
      保存しないので bytes で渡す（zipfile はどちらも同じように読める）。
    max_cells: セル数の上限（省略時は MAX_CELLS）。
    max_merged: 結合セルの面積の上限（省略時は MAX_MERGED_CELLS。帳票は FORM_MAX_MERGED_CELLS を渡す）。
    max_rows: 行（<row>）の数の上限（省略時は max_merged と同じ値）。通常モードの openpyxl は、セルの無い
      <row ht="20" customHeight="1"/> のような行にも行の書式を1つずつ作るので、結合セルと同じくらい開くのが遅くなる
      （100万行で約9秒）。帳票は画面を開くたびに開き直すので、結合セルと同じ目安（1回あたり2秒程度）で断る。
      一覧表は読み取り専用モードで開き行の書式を作らないので、MAX_MERGED_CELLS のままで困らない。
    """
    if isinstance(source, (bytes, bytearray)):
        head, book = bytes(source[:8]), io.BytesIO(source)
    else:
        book = Path(source)
        with open(book, "rb") as f:
            head = f.read(8)
    if head.startswith(OLE_MAGIC[:4]):
        raise UploadError("パスワード付きのブック、または .xls 形式です。パスワードを外して .xlsx 形式で保存し直してください")
    if not head.startswith(ZIP_MAGIC):
        raise UploadError("Excelファイル（.xlsx / .xlsm）として読み込めません。Excelで開いて .xlsx 形式で保存し直してください")
    try:
        with zipfile.ZipFile(book) as zf:
            infos = zf.infolist()
            names = {i.filename for i in infos}
            if "xl/workbook.bin" in names:
                raise UploadError(".xlsb（バイナリブック）形式は対象外です。Excelで .xlsx 形式で保存し直してください")
            _check_zip_limits(infos)
            workbook = "xl/workbook.xml" if "xl/workbook.xml" in names else next(
                (n for n in sorted(names) if n.lower().endswith("workbook.xml")), None)
            if workbook is None:
                raise UploadError("Excelファイル（.xlsx / .xlsm）ではありません（ブックの情報が見つかりません）")
            with zf.open(workbook) as wb:
                text = wb.read(_WORKBOOK_HEAD_BYTES).decode("utf-8", errors="ignore")
            # シートの置き場所と名前（拡張子）は workbook.xml.rels で自由に決められる（sheet1.dat でも開ける）ので、
            # 名前で選ばずに部品をすべて見る（XML でない部品は、読み始めてすぐ読めなくなって終わる）
            max_merged = max_merged or MAX_MERGED_CELLS
            _check_sheet_parts(zf, sorted(names), max_cells or MAX_CELLS, max_merged, max_rows or max_merged)
    # zlib.error / EOFError: 圧縮データが壊れている（zipfile はこれらを包まずにそのまま投げる）
    except (zipfile.BadZipFile, zipfile.LargeZipFile, OSError, RuntimeError, NotImplementedError,
            zlib.error, EOFError) as exc:
        # 例外の種類は画面に出さず（利用者には意味がない）、ログにだけ残す
        logger.warning("Excelファイルの事前チェックで読み込めませんでした: %s", exc.__class__.__name__)
        raise UploadError("Excelファイルとして読み込めません（ファイルが壊れている可能性があります）") from exc
    root = re.search(r"<(?:\w+:)?workbook\b[^>]*>", text)
    if STRICT_NS in (root.group(0) if root else text[:2048]):
        raise UploadError("Strict Open XML 形式のブックです。Excelの「名前を付けて保存」で通常の「Excel ブック (.xlsx)」を選んで保存し直してください")
    # シートが1つも無いブックは断る（読み取り先が無い）。先頭だけ読んでいるので、
    # シートの一覧の終わりが読んだ範囲に無いときは判断しない
    sheets_end = re.search(r"<(?:\w+:)?sheets\s*/>|</(?:\w+:)?sheets>", text)
    if sheets_end and not re.search(r"<(?:\w+:)?sheet\b", text[:sheets_end.end()]):
        raise UploadError("シートがないブックです。シートのあるブックを選んでください")


def _merged_area(ref: str | None) -> int:
    from openpyxl.utils.cell import range_boundaries

    try:
        min_col, min_row, max_col, max_row = range_boundaries(str(ref or "").strip().upper())
    except (ValueError, TypeError):
        return 0   # 読めない範囲は openpyxl 側で扱いが決まる（ここでは数えない）
    if None in (min_col, min_row, max_col, max_row):
        return 0
    return (abs(max_col - min_col) + 1) * (abs(max_row - min_row) + 1)


class _PartCounter:
    """XML の部品1つを分割して読み、結合範囲の面積・セル数・行数・リンクとコメントの範囲の面積を数える（xml.parsers.expat）。

    文字列を正規表現で探すのではなく XML として読むので、ref = "..."（= の前後の空白）、文字参照（&#88;）、
    どんな名前空間の接頭辞（<x.y:c> など）でも、openpyxl（xml.etree＝同じ expat）と同じ解釈で数える。
    """

    def __init__(self, merged: int, cells: int, max_cells: int, max_merged: int | None = None, *,
                 rows: int = 0, linked: int = 0, max_rows: int | None = None, factor: int = 1):
        from xml.parsers import expat

        self.merged, self.cells, self.max_cells = merged, cells, max_cells
        self.max_merged = max_merged or MAX_MERGED_CELLS
        self.rows, self.linked = rows, linked
        self.max_rows = max_rows or self.max_merged
        self.factor = max(int(factor), 1)   # この部品が openpyxl に読み直される回数（その回数だけ数える）
        self.anchors = 0              # この部品の図形（アンカー）の数
        self.drawing_targets = []     # この部品（.rels）が参照する描画部品の Target
        self.reread_targets = []      # この部品（.rels）が参照する、シートごとに読み直される部品の Target
        self.parser = expat.ParserCreate(namespace_separator=" ")
        self.parser.StartElementHandler = self._start
        # DTD で実体を定義して大量に展開させる細工は、正しいブックには無いので断る
        self.parser.EntityDeclHandler = self._entity

    def _start(self, name: str, attrs: dict) -> None:
        namespace, _, local = name.rpartition(" ")
        if local == "mergeCell":
            self.merged += _merged_area(attrs.get("ref")) * self.factor
            if self.merged > self.max_merged:
                raise UploadError("結合セルの範囲が大きすぎます（シート全体・列全体の結合など）。"
                                  "不要な結合を解除して保存し直してください")
        elif local == "c" and namespace in _SHEET_NAMESPACES:
            self.cells += self.factor
            if self.cells > self.max_cells:
                raise UploadError(f"セル数が上限（{self.max_cells:,} セル）を超えています。"
                                  "不要なシート・範囲を削除して保存し直してください")
        elif local == "row" and namespace in _SHEET_NAMESPACES:
            self.rows += self.factor
            if self.rows > self.max_rows:
                raise UploadError(f"行数が上限（{self.max_rows:,} 行）を超えています（高さなどの書式だけの行も数えます）。"
                                  "不要な行を削除して保存し直してください")
        elif local in ("hyperlink", "comment") and namespace in _SHEET_NAMESPACES:
            self.linked += _merged_area(attrs.get("ref")) * self.factor
            if self.linked > MAX_LINKED_CELLS:
                raise UploadError("ハイパーリンクまたはコメントの範囲が大きすぎます（シート全体・列全体を指すリンクなど）。"
                                  "不要なリンク・コメントを削除して保存し直してください")
        elif local in _ANCHOR_NAMES:
            self.anchors += 1
        elif local == "Relationship" and attrs.get("TargetMode") != "External":
            rel_type = str(attrs.get("Type", ""))
            if rel_type.endswith(_DRAWING_REL_SUFFIX):
                self.drawing_targets.append(str(attrs.get("Target", "")))
            else:
                for suffix in _REREAD_REL_SUFFIXES:
                    if rel_type.endswith(suffix):
                        self.reread_targets.append((suffix, str(attrs.get("Target", ""))))
                        break

    def _entity(self, *_args) -> None:
        raise UploadError("Excelファイルとして読み込めません（不正なファイルの可能性があります）")


def _check_sheet_parts(zf: zipfile.ZipFile, parts: list[str], max_cells: int, max_merged: int | None = None,
                       max_rows: int | None = None) -> None:
    """部品の XML を分割して読み、結合範囲の面積・セル数・行数・リンクとコメントの範囲の面積の合計を数え、上限を超えたら UploadError。

    openpyxl で開く前に確かめる（開いた時点で結合範囲のセルと、すべてのセルが作られてしまうため）。
    XML として読めなくなった部品は、そこまでの分だけ数える（画像などの XML でない部品はすぐ終わる）。
    openpyxl も同じ expat で読むので、読めない部品のその先の結合・セルは作られない（開くこと自体が失敗する）。
    """
    merged = 0
    cells = 0
    rows = 0
    linked = 0
    anchors: dict[str, int] = {}       # 部品 → 図形の数
    references: dict[str, int] = {}    # 描画部品 → 参照される回数（シートごとに読み直されるため）
    rereads = _reread_counts(zf, parts)
    for part in parts:
        counter = _PartCounter(merged, cells, max_cells, max_merged, rows=rows, linked=linked, max_rows=max_rows,
                               factor=rereads.get(part, 1))
        _parse_part(zf, part, counter)
        merged, cells, rows, linked = counter.merged, counter.cells, counter.rows, counter.linked
        anchors[part] = counter.anchors
        for target in counter.drawing_targets:
            drawing = _resolve_rel_target(part, target)
            # 参照元のシートの部品を複数の <sheet> が指していれば、その描画もその回数だけ読み直される
            references[drawing] = references.get(drawing, 0) + rereads.get(_rels_owner(part), 1)
    _check_drawings(zf, anchors, references)


def _parse_part(zf: zipfile.ZipFile, part: str, counter: "_PartCounter") -> None:
    """部品1つを分割して読み、counter に数えさせる（XML として読めなくなったら、そこまでの分で終わる）。"""
    from xml.parsers import expat

    with zf.open(part) as f:
        try:
            while True:
                chunk = f.read(CHUNK_SIZE)
                counter.parser.Parse(chunk, not chunk)
                if not chunk:
                    break
        except expat.ExpatError:
            pass


def _reread_counts(zf: zipfile.ZipFile, parts: list[str]) -> dict[str, int]:
    """部品 → openpyxl がそれを読む回数（シートの部品・コメントの部品だけ。ふつうのブックはどれも1回）。

    シートの部品は <sheet>（workbook.xml.rels の worksheet の関係）1つにつき1回、コメントの部品は
    それを参照するシートを読むたびに1回読まれる。関係を書いた .rels は小さいので、本体を数える前に先に読む。
    """
    targets: list[tuple[str, str, str]] = []    # (関係の種類, .rels の部品名, 参照先の部品名)
    for part in parts:
        if not part.endswith(_RELS_SUFFIX):
            continue
        counter = _PartCounter(0, 0, MAX_CELLS)
        _parse_part(zf, part, counter)
        targets += [(kind, part, _resolve_rel_target(part, t)) for kind, t in counter.reread_targets]
    counts: dict[str, int] = {}
    for kind, _rels_part, target in targets:
        if kind == "/worksheet":
            counts[target] = counts.get(target, 0) + 1
    # コメントの部品は、それを参照するシートを読むたびに読み直される（そのシートの部品を複数の <sheet> が
    # 指していれば、その回数だけ増える）ので、持ち主のシートの回数を足す
    for kind, rels_part, target in targets:
        if kind == "/comments":
            counts[target] = counts.get(target, 0) + max(counts.get(_rels_owner(rels_part), 0), 1)
    return counts


def _rels_owner(rels_part: str) -> str:
    """.rels の持ち主の部品（xl/worksheets/_rels/sheet1.xml.rels → xl/worksheets/sheet1.xml）。"""
    folder = posixpath.dirname(posixpath.dirname(rels_part))
    return posixpath.normpath(posixpath.join(folder, posixpath.basename(rels_part)[:-len(_RELS_SUFFIX)]))


def _resolve_rel_target(rels_part: str, target: str) -> str:
    """.rels の Target を zip 内のパスにする（xl/worksheets/_rels/sheet1.xml.rels の ../drawings/d.xml → xl/drawings/d.xml）。"""
    if target.startswith("/"):
        return target.lstrip("/")
    folder = posixpath.dirname(posixpath.dirname(rels_part))   # _rels の1つ上＝参照元の部品があるフォルダ
    return posixpath.normpath(posixpath.join(folder, target))


def _check_drawings(zf: zipfile.ZipFile, anchors: dict[str, int], references: dict[str, int]) -> None:
    """描画部品を読み直す回数（参照するシートの数）を掛けて、図形の数と大きさの合計が上限を超えたら UploadError。"""
    total_anchors = 0
    total_bytes = 0
    for drawing, count in references.items():
        if drawing not in anchors:
            continue   # 存在しない部品は openpyxl も読まない
        total_anchors += anchors[drawing] * count
        total_bytes += zf.getinfo(drawing).file_size * count
        if total_anchors > MAX_DRAWING_ANCHORS or total_bytes > MAX_DRAWING_BYTES:
            raise UploadError(_DRAWING_ERROR)


def _check_zip_limits(infos: list[zipfile.ZipInfo]) -> None:
    total = 0
    compressed = 0
    for info in infos:
        total += info.file_size
        compressed += info.compress_size
        if info.file_size > ZIP_MAX_PART:
            raise UploadError(f"ブックの中身が大きすぎます（1つの部品が展開後 {_format_size(info.file_size)}、上限 {_format_size(ZIP_MAX_PART)}）")
        if info.file_size >= ZIP_RATIO_MIN_BYTES and info.file_size > ZIP_MAX_RATIO * max(info.compress_size, 1):
            raise UploadError("ブックの圧縮率が異常に高いため読み込みを中止しました（壊れているか、不正なファイルの可能性があります）")
    if total > ZIP_MAX_TOTAL:
        raise UploadError(f"ブックの中身が大きすぎます（展開後の合計 {_format_size(total)}、上限 {_format_size(ZIP_MAX_TOTAL)}）")
    # 圧縮率はブック全体でも見る。部品ごとの確認は小さい部品を見ないので、上の大きさに足りない部品を
    # 並べるだけで、小さいファイルをいくらでも大きく展開させられる（700KB が 465MB ＝ 事前チェックだけで30秒以上）。
    # ふつうのブックは5〜7倍程度なので、部品ごとと同じ100倍で断る（小さいブックは見ない点も同じ）。
    if total >= ZIP_RATIO_MIN_BYTES and total > ZIP_MAX_RATIO * max(compressed, 1):
        raise UploadError("ブックの圧縮率が異常に高いため読み込みを中止しました（壊れているか、不正なファイルの可能性があります）")


def remove_upload(stored_path: str | None) -> bool:
    """アップロードしたファイルを消す。消せたら True。

    Windows では他のプロセス（ウイルス対策のスキャン、Excel で開いたまま、同期ソフト）や、解析に失敗した
    ブックを掴んだままの openpyxl のせいで消せないことがある。「ダウンロードしたら消えます」（design.md 3.3）と
    言い切っている以上、消せないときも中身だけは 0 バイトに切り詰めて、元のデータが残らないようにする。
    例外は投げない（ここで落とすと、日本語のエラー案内の代わりに 500 になる）。
    残った空ファイルは次の起動時に remove_orphan_uploads が片付ける。
    """
    if not stored_path:
        return True
    try:
        path = upload_path(stored_path)
    except UploadError:
        return False
    for attempt in range(REMOVE_RETRIES):
        try:
            path.unlink(missing_ok=True)
            return True
        except OSError:
            time.sleep(REMOVE_RETRY_WAIT * (attempt + 1))
    try:
        with open(path, "r+b") as f:
            f.truncate(0)
    except OSError:
        pass
    return False


def remove_sample_dir(base) -> int:
    """前の版が置いた見本の Excel（uploads/samples）をフォルダごと消す。消したファイルの件数を返す。

    帳票登録は見本の Excel を受け取っても置かなくなった（利用者の指示 2026-09-21）。
    前の版で置いたままのファイルが残ることがあるので、起動時に片付ける。
    """
    folder = Path(base) / SAMPLE_SUBDIR
    if not folder.is_dir():
        return 0
    count = sum(1 for p in folder.rglob("*") if p.is_file())
    shutil.rmtree(folder, ignore_errors=True)
    return count


def remove_orphan_import_dirs(tables_dir, known_ids) -> int:
    """DB に無い取り込みの imports/<id>/ フォルダを消す。消した件数を返す。

    取り込みを消し損ねたとき（削除の途中で落ちた・Windows でファイルを掴まれていた）に、読み込んだ行や
    作った Markdown が残り続けないように起動時に片付ける。作業中の取り込みは DB に行があるので消さない。
    """
    root = Path(tables_dir) / "imports"
    if not root.is_dir():
        return 0
    known = {int(i) for i in known_ids}
    removed = 0
    for path in sorted(root.iterdir()):
        if not path.is_dir() or not path.name.isdigit() or int(path.name) in known:
            continue
        shutil.rmtree(path, ignore_errors=True)
        if not path.exists():
            removed += 1
    return removed


# 起動時の片付けで、これより新しいファイルは消さない（保存してから DB の行を作るまでの間を守る）。
# core.jobs.STALE_AFTER（2分）と同じ長さ。
ORPHAN_GRACE_SECONDS = 120


def remove_orphan_uploads(base, known_paths) -> int:
    """DB のどこからも参照されていないアップロード済みファイルを消す。消した件数を返す。

    取り込みの途中で失敗し、消し損ねたファイルが残ることがある（画面からは消せない）ので起動時に片付ける。
    save_upload が作った名前（uuid + 拡張子）のファイルだけを対象にする。
    """
    base = Path(base)
    known = {str(p).replace("\\", "/").strip("/") for p in known_paths if p}
    removed = 0
    for subdir in UPLOAD_SUBDIRS:
        for path in sorted((base / subdir).glob("*")):
            if not path.is_file() or not _STORED_NAME.fullmatch(path.name):
                continue
            if f"{subdir}/{path.name}" in known:
                continue
            try:
                if time.time() - path.stat().st_mtime < ORPHAN_GRACE_SECONDS:
                    continue   # できたばかり: 別に起動しているアプリが取り込みの途中（DB の行を作る前）かもしれない
            except OSError:
                continue
            try:
                path.unlink()
            except OSError:
                continue
            removed += 1
    return removed


# ---- Windows のパスの長さ ------------------------------------------------------------
# 長いパスが有効でない Windows では、パス全体が 260 文字を超えるファイルを書けない。
# 置き場所（DATA_DIR 配下）に md を書く前に確かめるために使う（tables/pipeline.py）。

MAX_PATH_CHARS = 259


def _long_paths_enabled() -> bool:
    import os
    if os.name != "nt":
        return True
    try:
        import winreg
        with winreg.OpenKey(winreg.HKEY_LOCAL_MACHINE, r"SYSTEM\CurrentControlSet\Control\FileSystem") as key:
            value, _kind = winreg.QueryValueEx(key, "LongPathsEnabled")
        return bool(value)
    except OSError:
        return False


def path_limit() -> int | None:
    """書けるファイルのパスの長さの上限（上限が無ければ None）。"""
    return None if _long_paths_enabled() else MAX_PATH_CHARS


# ====================================================================================================
# 元 core/workbook_cache.py
# 置かれた Excel を「読み取った形」だけ、しばらくメモリに覚えておく置き場。
#
# 帳票登録は見本の Excel をサーバーに置かない（利用者の指示 2026-09-21「見本のExcelは置かずに、
# 設定だけ保持するようにしてほしい」）。ブラウザはファイルを選んだまま持っているので、セルを
# クリックするたびに同じ Excel を送り直してくる。そのたびにブックを開き直すと（結合セル・図形を
# 1つずつ作るので）数秒かかることがあるため、開いた結果（WorkbookInfo）だけをここに置く。
# 登録したときのシートの中身（セルの番地と文字）は帳票の種類と一緒に DB が覚える
# （pattern_books。2026-09-22）。ここはそれを開いた結果も同じように短い間だけ置く。
#
# 決まりごと:
#   - ディスクには何も書かない。Excel の中身（bytes）も持たない（開いた結果だけ）。
#   - アプリを終えれば消える（残るのは読み取りの設定と覚えたシートだけ、という約束を崩さない）。
#   - 鍵は「ブラウザの作業場所（session_id）＋ファイルの sha256」。ほかの人が置いたブックは見えない。
#   - 古いもの・多すぎるもの・大きすぎるものは、置いた順に落とす。
#   - waitress は1つのプロセスを複数のスレッドで回すので、出し入れは錠（Lock）の中で行う。
# ====================================================================================================

# 覚えておく時間。帳票登録は「置く → セルをクリックする → 使用開始」を続けて行う作業なので、
# 手が止まっても1回分の作業（30分）で足りる。切れてもブラウザが送り直すので読み直すだけ。
TTL_SECONDS = 30 * 60
# 覚えておく数と、その元になった Excel の大きさの合計。数人が同時に使っても RAM を食い尽くさない値。
MAX_ENTRIES = 8
MAX_BYTES = 64 * 1024 * 1024


@dataclass
class Book:
    """ブラウザが置いた Excel を読み取った結果（中身そのものは持たない）。

    帳票の種類が覚えているシート（pattern_books。views._stored_book が戻す）も同じ形で持つ。
    そのときの size は覚えた JSON の大きさ、saved_at は置いた日時。
    """

    file_name: str
    file_hash: str
    size: int                      # 元の Excel の大きさ（置ける量を数えるため）
    info: object                   # excel.workbook.WorkbookInfo
    used_at: float = field(default_factory=time.monotonic)
    saved_at: str = ""             # 帳票の種類と一緒に覚えたシートなら、その Excel を置いた日時


_books: dict[tuple[str, str], Book] = {}
_lock = threading.Lock()


def _expired(book: Book, now: float) -> bool:
    return now - book.used_at > TTL_SECONDS


def _evict(now: float) -> None:
    """古いもの → 入れた順、の順に落とす（錠の中から呼ぶ）。"""
    for key in [k for k, b in _books.items() if _expired(b, now)]:
        del _books[key]
    total = sum(b.size for b in _books.values())
    while _books and (len(_books) > MAX_ENTRIES or total > MAX_BYTES):
        key, dropped = next(iter(_books.items()))   # dict は入れた順に並ぶ＝いちばん古いもの
        del _books[key]
        total -= dropped.size


def put(session_id: str, book: Book) -> Book:
    """読み取った結果を覚える（同じブックを置き直したら入れ替える）。"""
    key = (session_id or "", book.file_hash)
    with _lock:
        book.used_at = time.monotonic()
        _books.pop(key, None)       # 入れ直して「いちばん新しい」にする
        _books[key] = book
        _evict(book.used_at)
    return book


def get(session_id: str, file_hash: str) -> Book | None:
    """覚えているブックを返す（無ければ None。呼び出し側はブラウザに置き直してもらう）。"""
    key = (session_id or "", file_hash or "")
    now = time.monotonic()
    with _lock:
        book = _books.get(key)
        if book is None:
            return None
        if _expired(book, now):
            del _books[key]
            return None
        book.used_at = now
        _books.pop(key)
        _books[key] = book          # 使ったものを「いちばん新しい」にする
        return book


def clear() -> None:
    """全部忘れる（テスト・片付け用）。"""
    with _lock:
        _books.clear()


def count() -> int:
    with _lock:
        return len(_books)


# ====================================================================================================
# 元 core/jobs.py
# バックグラウンドジョブ（ワーカースレッド＋キュー。AI整形だけ別の列で動かす）。
#
# - 状態・進捗・一時停止/中止の要求は jobs テーブルに持つ（画面はポーリングで読む）。
# - fn(ctx) は start_job を呼んだアプリの app_context 内で実行する。fn 内で DB を使うときは connect()。
# - 実行中は一定間隔で heartbeat_at を更新する。起動時に recover_interrupted() で止まったジョブを「中断」にする。
#   動いている間も、このプロセスが持っていないジョブの heartbeat が古ければ、参照したときに「中断」にする
#   （アプリを閉じてすぐ起動し直したとき、前のプロセスのジョブが「処理中」のまま残らないように）。
# - AI整形（ai_format）は何時間もかかったり一時停止したりするので、読み込み・プレビュー・Markdown作成とは
#   別の列（lane）で動かす（止めている間に、ほかの取り込みの作業まで止まらないように）。
#   同じ取り込みのジョブは、列が違っても同時には動かさない（前のジョブが終わるまで待たせる）。
# - 社内LANのサーバーで数人が同時に使う（2026-09-20 の利用者の指示）ので、列ごとに数本のワーカーを持つ
#   （JOB_WORKERS。既定 3）。誰かが3万行の表を読んでいる間、ほかの人の読み込みが待たされないようにする。
#   取り込みごとの直列（同じ取り込みを2つのジョブが同時に触らない）は、ワーカーが増えても守る。
# ====================================================================================================

HEARTBEAT_INTERVAL = 30          # 秒。fn が呼ばなくても実行中はこの間隔で更新する
STALE_AFTER = timedelta(minutes=2)
PAUSE_POLL = 0.3                 # 一時停止中に再開・中止を確かめる間隔（秒）
# 進捗・メッセージ・heartbeat の書き込みが「database is locked」になったときの再試行（回数と待ち秒の単位）。
# 取れなければその1回の更新だけ見送る（進捗の表示のためにジョブ全体を失敗にしない）
SOFT_UPDATE_RETRIES = 2
SOFT_UPDATE_BACKOFF = 0.2

ACTIVE_STATUSES = ("queued", "running", "paused")
TERMINAL_STATUSES = ("done", "failed", "cancelled", "interrupted")

STATUS_LABELS = {
    "queued": "待機中",
    "running": "処理中",
    "paused": "一時停止中",
    "done": "完了",
    "failed": "エラー",
    "cancelled": "中止",
    "interrupted": "中断",
}

# ジョブの種類 → 動かす列（書いていない種類は "main"）
LANE_OF_KIND = {"ai_format": "ai"}

# 列ごとのワーカー本数（app.config["JOB_WORKERS"]）。数人が同時に使うので、1本だと
# 誰かの長い読み込みでほかの人が待たされる。増やしすぎても DB とディスクの取り合いになるだけ
DEFAULT_WORKERS = 3
MAX_WORKERS = 8
IDLE_POLL = 1.0                  # 待つものが何も無いときに、キューを見直す間隔（秒）

_queues: dict[str, queue.Queue] = {}
# キューから取り出したが、まだ始めていないジョブ（列ごと。同じ列のワーカーで分け合う）
_pending: dict[str, list] = {}
# (列, 何本目) → ワーカーのスレッド
_workers: dict[tuple[str, int], threading.Thread] = {}
_worker_lock = threading.Lock()
# 各ワーカーがいま実行している取り込み（ref_type, ref_id）。同じ取り込みのジョブを2つ同時に動かさないため
_executing: dict[tuple[str, int], tuple] = {}
# 各ワーカーがいま実行しているジョブ（DB のパス, job_id）。待機中のジョブに「何を待っているか」を出すため
_executing_job: dict[tuple[str, int], tuple[str, int]] = {}
# このプロセスが登録したジョブ（DB のパス, job_id）。これ以外で heartbeat が古いものは持ち主がいない
_owned: set[tuple[str, int]] = set()


class JobCancelled(Exception):
    """中止の要求を受けて fn の処理を打ち切るときに投げる。"""


class JobError(Exception):
    """利用者にそのまま見せてよい日本語メッセージを持つエラー。

    これを継承した例外（aiproc.runner.AIJobError、tables.pipeline.PipelineError）は、そのメッセージを
    画面にそのまま出す。それ以外の例外は Python の例外名が画面に出ないようにし、内容はログにだけ残す。
    """


def error_message(exc: Exception) -> str:
    """ジョブの失敗を画面に出す日本語の文にする。詳しい内容（例外名）はログに残してあるので出さない。"""
    text = " ".join(str(exc).split())
    if isinstance(exc, JobError) and text:
        return text
    return "処理中にエラーが発生しました。もう一度実行してください（詳しい内容はアプリのログに記録しました）"


def _now() -> str:
    return now()


class JobContext:
    """fn に渡す操作口。スレッドをまたいで呼んでもよい（内部でロック）。"""

    def __init__(self, job_id: int, conn=None):
        self.job_id = job_id
        self._conn = conn or connect()
        self._db = _db_key()  # 待機中のジョブの heartbeat を更新するとき、このプロセスのものを選ぶため
        self._lock = threading.Lock()
        row = self._row()
        self._progress: dict = json.loads(row["progress_json"] or "{}") if row else {}
        self.params: dict = json.loads(row["params_json"] or "{}") if row else {}

    # ---- 内部 ----
    def _row(self):
        with self._lock:
            return self._conn.execute("SELECT * FROM jobs WHERE id = ?", (self.job_id,)).fetchone()

    def _update(self, sql_sets: str, args=()) -> None:
        with self._lock:
            self._conn.execute(f"UPDATE jobs SET {sql_sets} WHERE id = ?", (*args, self.job_id))
            self._conn.commit()

    def _soft_update(self, sql_sets: str, args=()) -> None:
        """進捗など、1回飛ばしても次の更新で追いつく書き込み。ロック中なら少し待って再試行し、だめなら見送る。"""
        for attempt in range(SOFT_UPDATE_RETRIES + 1):
            try:
                self._update(sql_sets, args)
                return
            except sqlite3.OperationalError as exc:
                text = str(exc).lower()
                if "locked" not in text and "busy" not in text:
                    raise
                with self._lock:
                    try:
                        self._conn.rollback()
                    except sqlite3.Error:
                        pass
                if attempt < SOFT_UPDATE_RETRIES:
                    time.sleep(SOFT_UPDATE_BACKOFF * (attempt + 1))
        logging.getLogger(__name__).warning("ジョブ %s の進捗の更新を見送りました（データベースが使用中）", self.job_id)

    def _flags(self) -> tuple[bool, bool]:
        row = self._row()
        if row is None:
            return True, False  # ジョブが消された → 中止扱い
        return bool(row["cancel_requested"]), bool(row["pause_requested"])

    def close(self) -> None:
        with self._lock:
            self._conn.close()

    # ---- 公開 ----
    def progress(self, **kw) -> None:
        """進捗を上書きマージして保存する（例: done=10, total=100, phase="読み込み"）。heartbeat も更新。"""
        self._progress.update(kw)
        ts = _now()
        self._soft_update("progress_json = ?, heartbeat_at = ?, updated_at = ?",
                          (json.dumps(self._progress, ensure_ascii=False, default=str), ts, ts))

    def message(self, text: str) -> None:
        self._soft_update("message = ?, updated_at = ?", (text, _now()))

    def heartbeat(self) -> None:
        self._soft_update("heartbeat_at = ?", (_now(),))

    def touch_queued(self) -> None:
        """このプロセスのキューで待っているジョブの heartbeat も更新する。

        待機中のジョブは自分では heartbeat を打たない。誤って2つ目を起動したとき、その起動時の
        recover_interrupted（別プロセス）が、1つ目で2分以上待っているジョブを「中断」にしないように。
        """
        with _worker_lock:
            owned = {job_id for db, job_id in _owned if db == self._db}
        if not owned:
            return
        with self._lock:
            ids = [r[0] for r in self._conn.execute("SELECT id FROM jobs WHERE status = 'queued'") if r[0] in owned]
            ts = _now()
            for i in range(0, len(ids), 500):  # SQLite の変数の上限を超えないよう分ける
                part = ids[i:i + 500]
                self._conn.execute(f"UPDATE jobs SET heartbeat_at = ? WHERE status = 'queued' AND id IN "
                                   f"({', '.join('?' for _ in part)})", (ts, *part))
            self._conn.commit()

    def is_cancelled(self) -> bool:
        return self._flags()[0]

    def should_stop(self) -> bool:
        """中止または一時停止が要求されていれば True（新しい処理を出さないための判定）。"""
        cancel, pause = self._flags()
        return cancel or pause

    def wait_if_paused(self) -> bool:
        """一時停止の要求があれば再開まで待つ。続けてよければ True、中止なら False。"""
        cancel, pause = self._flags()
        if cancel:
            return False
        if not pause:
            return True
        self._update("status = 'paused', heartbeat_at = ?, updated_at = ?", (_now(), _now()))
        last_beat = time.monotonic()
        while True:
            time.sleep(PAUSE_POLL)
            cancel, pause = self._flags()
            if cancel:
                return False
            if not pause:
                self._update("status = 'running', heartbeat_at = ?, updated_at = ?", (_now(), _now()))
                return True
            if time.monotonic() - last_beat >= HEARTBEAT_INTERVAL:
                self.heartbeat()
                last_beat = time.monotonic()

    def check_cancel(self) -> None:
        """一時停止中なら待ち、中止が要求されていれば JobCancelled を投げる。"""
        if not self.wait_if_paused():
            raise JobCancelled()


# ---- 実行 -----------------------------------------------------------------------

def start_job(kind: str, ref_type: str, ref_id: int | None, fn: Callable[[JobContext], dict | None],
              params: dict | None = None) -> int:
    """ジョブを登録してキューに入れ、job_id を返す。app_context 内で呼ぶ。"""
    app = current_app._get_current_object()
    ts = _now()
    conn = connect()
    try:
        cur = conn.execute(
            """INSERT INTO jobs (kind, ref_type, ref_id, status, params_json, progress_json, created_at, updated_at)
               VALUES (?, ?, ?, 'queued', ?, '{}', ?, ?)""",
            (kind, ref_type or "", ref_id, json.dumps(params or {}, ensure_ascii=False, default=str), ts, ts),
        )
        conn.commit()
        job_id = cur.lastrowid
    finally:
        conn.close()
    lane = LANE_OF_KIND.get(kind, "main")
    ref = (ref_type or "", ref_id) if ref_id is not None else None
    with _worker_lock:
        _owned.add((_db_key(), job_id))
        q = _queues.setdefault(lane, queue.Queue())
    q.put((app, job_id, fn, ref, _db_key()))
    _ensure_workers(lane)
    return job_id


def _db_key() -> str:
    return str(current_app.config["DATABASE"])


def worker_count() -> int:
    """1つの列で同時に動かすジョブの本数（app.config["JOB_WORKERS"]、既定 3）。"""
    try:
        value = int(current_app.config.get("JOB_WORKERS") or DEFAULT_WORKERS)
    except (KeyError, RuntimeError, TypeError, ValueError):
        value = DEFAULT_WORKERS
    return max(1, min(MAX_WORKERS, value))


def _ensure_workers(lane: str) -> None:
    """その列のワーカーを必要な本数まで立てる（落ちていたら立て直す）。"""
    count = worker_count()
    with _worker_lock:
        for index in range(count):
            key = (lane, index)
            worker = _workers.get(key)
            if worker is None or not worker.is_alive():
                worker = threading.Thread(target=_worker_loop, args=(lane, index), daemon=True,
                                          name=f"job-worker-{lane}-{index}")
                _workers[key] = worker
                worker.start()


def _next_item(lane: str, key: tuple[str, int]):
    """次に実行するものを取り出す。同じ取り込みを別のワーカーが実行中なら、その取り込みの分は後回しにする。

    先頭から見て最初に動けるものを選ぶので、同じ取り込みの後ろのジョブが前のジョブを追い越すことはない。
    待ち行列（_pending）は同じ列のワーカーで分け合うので、どのワーカーも必ず時間を区切って見直す
    （別のワーカーが取り出して置いた分に、いつまでも誰も手を付けない、が起きないように）。
    """
    q = _queues[lane]
    while True:
        with _worker_lock:
            pending = _pending.setdefault(lane, [])
            try:
                while True:
                    pending.append(q.get_nowait())
            except queue.Empty:
                pass
            busy = {ref for other, ref in _executing.items() if other != key and ref is not None}
            for i, item in enumerate(pending):
                if item[3] is None or item[3] not in busy:
                    _executing[key] = item[3]
                    _executing_job[key] = (item[4], item[1])
                    return pending.pop(i)
            waiting = bool(pending)
        try:
            item = q.get(timeout=PAUSE_POLL if waiting else IDLE_POLL)
        except queue.Empty:
            continue
        with _worker_lock:
            _pending.setdefault(lane, []).append(item)


def _worker_loop(lane: str = "main", index: int = 0) -> None:
    key = (lane, index)
    while True:
        app, job_id, fn, _ref, _db = _next_item(lane, key)
        try:
            _run(app, job_id, fn)
        except Exception:  # ワーカーは止めない
            try:
                app.logger.exception("ジョブ %s の実行管理でエラー", job_id)
            except Exception:
                pass
            _fail_left_open(app, job_id)
        finally:
            with _worker_lock:
                _executing.pop(key, None)
                _executing_job.pop(key, None)


def _fail_left_open(app, job_id: int) -> None:
    """実行管理の途中で落ちた（最後の状態の書き込みが DB のロックで失敗した など）ジョブを「エラー」で閉じる。

    閉じないと「処理中」のまま残り、中止も再実行もできなくなる。ここも失敗したら諦める（ログは残してある）。
    """
    try:
        with app.app_context():
            conn = connect()
            try:
                marks = ", ".join("?" for _ in ACTIVE_STATUSES)
                conn.execute(f"UPDATE jobs SET status = 'failed', pause_requested = 0, message = ?, updated_at = ? "
                             f"WHERE id = ? AND status IN ({marks})",
                             (error_message(RuntimeError()), _now(), job_id, *ACTIVE_STATUSES))
                conn.commit()
            finally:
                conn.close()
    except Exception:
        pass


class _Ticker:
    """fn が長い呼び出しで止まっていても heartbeat を更新し続ける（このプロセスの待機中のジョブの分も）。"""

    def __init__(self, ctx: JobContext):
        self._ctx = ctx
        self._stop = threading.Event()
        self._thread = threading.Thread(target=self._loop, name=f"job-{ctx.job_id}-heartbeat", daemon=True)

    def _loop(self) -> None:
        while not self._stop.wait(HEARTBEAT_INTERVAL):
            try:
                self._ctx.heartbeat()
                # 待機中のジョブは、このプロセスでいずれかのジョブが動いている（か一時停止している）間だけ
                # 待たされる。そのジョブの Ticker が代わりに更新する
                self._ctx.touch_queued()
            except Exception:
                return

    def __enter__(self):
        self._thread.start()
        return self

    def __exit__(self, *exc):
        self._stop.set()
        self._thread.join(timeout=5)


def _run(app, job_id: int, fn: Callable[[JobContext], dict | None]) -> None:
    with app.app_context():
        ctx = JobContext(job_id)
        try:
            row = ctx._row()
            if row is None or row["status"] != "queued":
                return  # 待機中に中止・中断された
            if row["cancel_requested"]:
                ctx._update("status = 'cancelled', message = ?, updated_at = ?", ("中止しました", _now()))
                return
            ts = _now()
            with ctx._lock:  # 確認と開始の間に中止されていたら実行しない
                cur = ctx._conn.execute("UPDATE jobs SET status = 'running', heartbeat_at = ?, updated_at = ? "
                                        "WHERE id = ? AND status = 'queued'", (ts, ts, job_id))
                ctx._conn.commit()
            if cur.rowcount == 0:
                return
            try:
                with _Ticker(ctx):
                    result = fn(ctx)
            except JobCancelled:
                ctx._update("status = 'cancelled', pause_requested = 0, message = ?, updated_at = ?",
                            ("中止しました", _now()))
                return
            except Exception as exc:
                app.logger.exception("ジョブ %s（%s）でエラー", job_id, row["kind"])
                ctx._update("status = 'failed', pause_requested = 0, message = ?, updated_at = ?",
                            (error_message(exc), _now()))
                return
            if result is not None:
                ctx.progress(result=result)
            if ctx.is_cancelled():
                ctx._update("status = 'cancelled', pause_requested = 0, message = CASE WHEN message = '' "
                            "THEN '中止しました' ELSE message END, updated_at = ?", (_now(),))
            else:
                ctx._update("status = 'done', pause_requested = 0, updated_at = ?", (_now(),))
        finally:
            ctx.close()


# ---- 参照・操作（画面から） --------------------------------------------------------

def _decode(row) -> dict | None:
    if row is None:
        return None
    job = dict(row)
    job["params"] = json.loads(job.get("params_json") or "{}")
    job["progress"] = json.loads(job.get("progress_json") or "{}")
    job["result"] = job["progress"].get("result")
    job["status_label"] = STATUS_LABELS.get(job["status"], job["status"])
    job["finished"] = job["status"] in TERMINAL_STATUSES
    return job


def _with_conn(fn):
    conn = connect()
    try:
        return fn(conn)
    finally:
        conn.close()


INTERRUPTED_MESSAGE = "アプリの終了などで処理が中断されました。もう一度実行してください"


def _reap_orphan(conn, row):
    """持ち主のいないジョブ（このプロセスが登録しておらず、heartbeat が2分以上古い）を「中断」にして返す。

    起動時の recover_interrupted は2分以上古いものしか直さない（誤って2つ目を起動したときに、動いている
    1つ目のジョブを中断にしないため）。閉じてすぐ起動し直すと前のジョブが「処理中」のまま残るので、
    参照したときにもう一度確かめる。
    """
    if row is None or row["status"] not in ACTIVE_STATUSES:
        return row
    with _worker_lock:
        if (_db_key(), row["id"]) in _owned:
            return row
    cutoff = (datetime.now() - STALE_AFTER).isoformat(timespec="seconds")
    if (row["heartbeat_at"] or row["updated_at"] or row["created_at"] or "") >= cutoff:
        return row
    marks = ", ".join("?" for _ in ACTIVE_STATUSES)
    conn.execute(f"UPDATE jobs SET status = 'interrupted', pause_requested = 0, message = ?, updated_at = ? "
                 f"WHERE id = ? AND status IN ({marks})", (INTERRUPTED_MESSAGE, _now(), row["id"], *ACTIVE_STATUSES))
    conn.commit()
    return conn.execute("SELECT * FROM jobs WHERE id = ?", (row["id"],)).fetchone()


JOB_KIND_LABELS = {"ai_format": "AI整形"}


def _waiting_note(conn, job: dict | None) -> dict | None:
    """待機中のジョブに、何が終わるのを待っているかを付ける（job["waiting_for"] と、空なら job["message"]）。

    AI整形は一時停止している間も列と取り込みを使ったままなので、その後ろのジョブは「待機中」のまま動かない。
    理由が画面に出ないと、利用者は止まっているのか分からない（再開・中止すれば動き出すことも伝える）。
    """
    if job is None or job["status"] != "queued":
        return job
    lane = LANE_OF_KIND.get(job["kind"], "main")
    ref = (job.get("ref_type") or "", job.get("ref_id")) if job.get("ref_id") is not None else None
    db_key = _db_key()
    with _worker_lock:
        running = dict(_executing)
        running_jobs = dict(_executing_job)
        lane_size = sum(1 for k, w in _workers.items() if k[0] == lane and w.is_alive())
    mine = {k: v for k, v in running_jobs.items() if v[0] == db_key and v[1] != job["id"]}
    blocker, same_ref = None, False
    for key, owner in mine.items():        # 同じ取り込みを別のワーカーが実行中
        if ref is not None and running.get(key) == ref:
            blocker, same_ref = owner[1], True
            break
    if blocker is None:
        # 同じ列のワーカーが全部ふさがっている（空くまで動けない）
        busy_here = [owner[1] for key, owner in mine.items() if key[0] == lane]
        if busy_here and len(busy_here) >= max(1, lane_size):
            blocker = busy_here[0]
    if blocker is None:
        return job
    row = conn.execute("SELECT id, kind, ref_type, ref_id, status FROM jobs WHERE id = ?", (blocker,)).fetchone()
    if row is None or row["status"] not in ACTIVE_STATUSES:
        return job
    job["waiting_for"] = {"job_id": row["id"], "kind": row["kind"], "ref_type": row["ref_type"],
                          "ref_id": row["ref_id"], "status": row["status"], "same_ref": same_ref}
    if not job.get("message"):
        label = JOB_KIND_LABELS.get(row["kind"], "前の処理")
        whose = "この取り込みの" if same_ref else ("別の取り込みの" if row["kind"] in JOB_KIND_LABELS else "")
        if row["status"] == "paused":
            job["message"] = (f"{whose}{label}が一時停止中のため待っています。"
                              f"{label}の画面で再開するか中止すると始まります")
        else:
            job["message"] = f"{whose}{label}が終わるのを待っています"
    return job


def get_job(job_id: int) -> dict | None:
    """ジョブ1件（params / progress / result / status_label / finished を付けて返す）。"""
    return _with_conn(lambda c: _waiting_note(c, _decode(_reap_orphan(
        c, c.execute("SELECT * FROM jobs WHERE id = ?", (job_id,)).fetchone()))))


def latest_job(ref_type: str, ref_id: int, kind: str | None = None) -> dict | None:
    sql, args = "SELECT * FROM jobs WHERE ref_type = ? AND ref_id = ?", [ref_type, ref_id]
    if kind:
        sql += " AND kind = ?"
        args.append(kind)
    return _with_conn(lambda c: _waiting_note(c, _decode(_reap_orphan(
        c, c.execute(sql + " ORDER BY id DESC LIMIT 1", args).fetchone()))))


def _request(job_id: int, sets: str, statuses=ACTIVE_STATUSES) -> bool:
    def run(conn):
        marks = ", ".join("?" for _ in statuses)
        cur = conn.execute(f"UPDATE jobs SET {sets}, updated_at = ? WHERE id = ? AND status IN ({marks})",
                           (_now(), job_id, *statuses))
        conn.commit()
        return cur.rowcount > 0
    return _with_conn(run)


def request_pause(job_id: int) -> bool:
    return _request(job_id, "pause_requested = 1", ("queued", "running"))


def request_resume(job_id: int) -> bool:
    return _request(job_id, "pause_requested = 0")


def request_cancel(job_id: int) -> bool:
    """中止を要求する。待機中のジョブはその場で中止にする。"""
    ok = _request(job_id, "cancel_requested = 1")
    _request(job_id, "status = 'cancelled', message = '中止しました'", ("queued",))
    return ok


def recover_interrupted() -> int:
    """起動時：実行中・待機中・一時停止中のまま heartbeat が2分以上古いジョブを「中断」にする。件数を返す。"""
    cutoff = (datetime.now() - STALE_AFTER).isoformat(timespec="seconds")

    def run(conn):
        marks = ", ".join("?" for _ in ACTIVE_STATUSES)
        cur = conn.execute(
            f"""UPDATE jobs SET status = 'interrupted', pause_requested = 0, message = ?, updated_at = ?
                WHERE status IN ({marks}) AND COALESCE(heartbeat_at, updated_at, created_at) < ?""",
            (INTERRUPTED_MESSAGE, _now(), *ACTIVE_STATUSES, cutoff),
        )
        conn.commit()
        return cur.rowcount
    return _with_conn(run)


def reap_orphan_jobs() -> int:
    """持ち主のいないジョブ（このプロセスが登録しておらず、生存印が2分以上古い）を「中断」にする。

    起動をすり抜けたジョブ（止めた直後に起動し直すと、生存印がまだ新しくて recover_interrupted に
    引っかからない）は、これまでブラウザがそのジョブを見に来たときしか閉じられなかった。
    見に来る人がいないと「実行中」のまま残り、その取り込みが片付けの対象から外れ続けて、
    元ファイル（最大200MB）と作った md が次の再起動まで消えなかった（2026-09-23 のレビュー）。
    このプロセスが動かしているジョブ（_owned）は、生存印が古くても閉じない。
    """
    cutoff = (datetime.now() - STALE_AFTER).isoformat(timespec="seconds")

    def run(conn):
        marks = ", ".join("?" for _ in ACTIVE_STATUSES)
        rows = conn.execute(
            f"""SELECT id FROM jobs WHERE status IN ({marks})
                AND COALESCE(heartbeat_at, updated_at, created_at) < ?""",
            (*ACTIVE_STATUSES, cutoff)).fetchall()
        with _worker_lock:
            ids = [r[0] for r in rows if (_db_key(), r[0]) not in _owned]
        if not ids:
            return 0
        conn.execute(
            f"""UPDATE jobs SET status = 'interrupted', pause_requested = 0, message = ?, updated_at = ?
                WHERE id IN ({", ".join("?" for _ in ids)})""",
            (INTERRUPTED_MESSAGE, _now(), *ids))
        conn.commit()
        return len(ids)
    return _with_conn(run)


def wait_job(job_id: int, timeout: float = 30.0, statuses=TERMINAL_STATUSES) -> dict | None:
    """指定の状態になるまで待つ（テスト・同期実行用）。時間切れなら最後の状態を返す。"""
    deadline = time.monotonic() + timeout
    while True:
        job = get_job(job_id)
        if job is None or job["status"] in statuses or time.monotonic() >= deadline:
            return job
        time.sleep(0.05)


# ====================================================================================================
# 元 core/purge.py
# 取り込んだデータを消す（design.md 3.3「データを残さない」）。
#
# このアプリはダウンロードが終わった時点で、その取り込みに属するものをすべて消す。
# 消すもの: アップロードした元のファイル、imports/<id>/（読み込んだ行・控え・作った md）、
#           DB の行（documents / table_imports / jobs / ai_items / llm_calls とそれらの子）。
# 残すもの: 帳票の種類（patterns とその子）・ブラウザごとの AI接続（ai_connections）。これらは「設定」で
#           データではない（SETTINGS_TABLES。取り込みを指す列が後から足されても、ここからは消さない）。
#
# DB の行は表の名前を決め打ちせず、その取り込みを指す列（document_id / import_id）を持つ表を
# sqlite_master から探して消す（後から表が増えても消し残さないため）。
# ====================================================================================================

log = logging.getLogger(__name__)

# 取り込み1件を指す外部キーの列名（この列を持つ表は、その取り込みと一緒に消す）
DOCUMENT_REF_COLUMNS = ("document_id",)
IMPORT_REF_COLUMNS = ("import_id", "table_import_id")
# 取り込みを指す列があっても、まとめては消さない表。llm_calls は AI の生の応答で、同じ文面の行なら
# 別の取り込みの ai_items が同じ応答を使っていることがある（消すと再実行で再課金になる）。
# 持ち主が消えた分は _delete_orphan_llm_calls が「参照されていなければ消す」で扱う。
SHARED_TABLES = ("llm_calls",)
# 「設定」の表。取り込みを捨てる片付け（purge_* / sweep_stale / purge_session）では決して消さない。
# ai_connections はブラウザごとの AI接続（APIキー・接続先。利用者の指示 2026-09-21「ずっと保持」）で、
# 帳票・一覧表と同じ session_id を持ち主に持つが、その人の取り込みを捨てても残す。
# pattern_books は帳票の種類が覚えているシートの中身（セルの番地と文字。Excel のファイルではない。
# 利用者の指示 2026-09-22）で、種類と一緒にしか消えない。
SETTINGS_TABLES = ("patterns", "pattern_sheets", "pattern_fields", "pattern_books", "ai_connections")


def _tables_with_column(db, column: str) -> list[str]:
    names = []
    for (table,) in db.execute("SELECT name FROM sqlite_master WHERE type = 'table'").fetchall():
        if table.startswith("sqlite_") or table in SHARED_TABLES or table in SETTINGS_TABLES:
            continue
        if column in {row[1] for row in db.execute(f'PRAGMA table_info("{table}")')}:
            names.append(table)
    return names


def _delete_by_columns(db, columns, value) -> int:
    removed = 0
    for column in columns:
        for table in _tables_with_column(db, column):
            removed += db.execute(f'DELETE FROM "{table}" WHERE "{column}" = ?', (value,)).rowcount
    return removed


def _delete_jobs(db, ref_type: str, ref_id: int) -> int:
    return db.execute("DELETE FROM jobs WHERE ref_type = ? AND ref_id = ?", (ref_type, ref_id)).rowcount


_UNREFERENCED_LLM_CALLS = ("cache_key NOT IN (SELECT cache_key FROM ai_items WHERE cache_key IS NOT NULL)")
# 持ち主の取り込みがもう無い応答（持ち主の分からない古いDBの分＝NULL も含む）
NO_LIVE_OWNER = "(import_id IS NULL OR import_id NOT IN (SELECT id FROM table_imports))"


def _ai_work_alive(db) -> bool:
    """まだある取り込みに、AI整形の結果かAI整形のジョブ（一時停止・中断を含む）が残っているか。

    ai_items が覚えているのは最後に使った応答のキーだけで、再依頼で直した行の1回目の応答や、
    一時停止の直前に受け取った応答（ai_items の行がまだ無い）は、どこからも参照されない。
    それでも再実行・再開ではキャッシュとして引く（引けないと再課金になる）ので、
    AI整形が残っている取り込みがある間は、持ち主の分からない（import_id が NULL の古いDBの）
    参照されない応答を消さない。持ち主の分かる応答は、その取り込みを消すときに一緒に消える。
    """
    return db.execute(
        "SELECT 1 FROM ai_items WHERE import_id IN (SELECT id FROM table_imports) "
        "UNION ALL SELECT 1 FROM jobs WHERE kind = 'ai_format' AND ref_type = 'table_import' "
        "AND ref_id IN (SELECT id FROM table_imports) LIMIT 1").fetchone() is not None


def _delete_orphan_llm_calls(db, keys=()) -> int:
    """どの ai_items からも参照されていない AI の生の応答を消す（試し実行の分もここで消える）。

    keys: 消した取り込みが使っていた応答のキー。ほかの行が使っておらず、まだある取り込みが払った分でも
    なければ、いつでもすぐ消す。消す取り込みは、別の取り込みが払った応答をキャッシュとして引いている
    ことがある（同じ文面なら同じキーになる）ので、払った取り込みがまだあるうちは消さない
    （消すと、その取り込みの再開・再実行で同じ応答をもう一度買うことになる）。その応答は、払った
    取り込みを消すときに下の「持ち主の取り込みがもう無い分」で消えるので、残り続けることはない。
    もう無い取り込みが払った応答（llm_calls.import_id）も、ほかから使われていなければすぐ消す。
    ai_items は最後に使った応答のキーしか覚えていないので、再依頼で直した行の1回目の応答や、
    一時停止の直前に受け取った応答は keys に入らない。持ち主の列があれば、それも一緒に消せる。
    持ち主の分からない応答（古いDBの分）は、AI整形が残っている取り込みが1件も無いときだけ消す（_ai_work_alive）。
    """
    removed = 0
    keys = sorted({k for k in keys if k})
    for start in range(0, len(keys), 500):
        chunk = keys[start:start + 500]
        marks = ", ".join("?" for _ in chunk)
        removed += db.execute(f"DELETE FROM llm_calls WHERE cache_key IN ({marks}) "
                              f"AND {NO_LIVE_OWNER} AND {_UNREFERENCED_LLM_CALLS}", chunk).rowcount
    # 持ち主の取り込みがもう無い分（design.md 3.3「取り込んだデータはダウンロードが終わった時点で消す」）
    removed += db.execute(f"DELETE FROM llm_calls WHERE import_id IS NOT NULL "
                          f"AND import_id NOT IN (SELECT id FROM table_imports) "
                          f"AND {_UNREFERENCED_LLM_CALLS}").rowcount
    # まだ使われていて残した分は、持ち主が居なくなったので「持ち主なし」に戻す
    # （使っている取り込みを消すときに、その keys で消える）
    db.execute("UPDATE llm_calls SET import_id = NULL WHERE import_id IS NOT NULL "
               "AND import_id NOT IN (SELECT id FROM table_imports)")
    if not _ai_work_alive(db):
        removed += db.execute(f"DELETE FROM llm_calls WHERE {_UNREFERENCED_LLM_CALLS}").rowcount
    return removed


def delete_orphan_ai_items(db) -> int:
    """もう無い取り込みを指す AI整形の結果を消す（取り込みを消したあとに、動いていた AI の書き込みが残った分など）。

    import_id の無い行（古いDBの分。持ち主の取り込みが分からない）は、その取り込み設定の取り込みが
    1件も残っていなければ消す（残っていれば、その取り込みを消すときに purge_table_import が消す）。
    """
    removed = db.execute("DELETE FROM ai_items WHERE import_id IS NOT NULL "
                         "AND import_id NOT IN (SELECT id FROM table_imports)").rowcount
    removed += db.execute("DELETE FROM ai_items WHERE import_id IS NULL AND template_id NOT IN "
                          "(SELECT template_id FROM table_imports WHERE template_id IS NOT NULL)").rowcount
    return removed


def sweep_orphan_ai(db) -> int:
    """持ち主の無い AI整形の結果と、どこからも使われない AI の生の応答を消して縮める（戻り値: 消した件数）。

    取り込み設定を消したとき（ai_items は消えるが llm_calls が残る）や起動時の片付けで使う。
    """
    removed = delete_orphan_ai_items(db) + _delete_orphan_llm_calls(db)
    db.commit()
    if removed:
        _shrink(db)
    return removed


def _mark_incomplete() -> None:
    """この要求の中で、消し切れなかったもの（掴まれていて消せなかったファイル）があったと印を付ける。"""
    if has_app_context():
        g.purge_incomplete = True


def purge_incomplete() -> bool:
    """この要求の中で消したときに、消し切れなかったファイルがあったか（保存先フォルダの結果の画面で知らせるため）。

    消せなかったファイルも中身は 0 バイトに切り詰めてあり、次の起動時に片付く。例外にはしない
    （消すこと自体は DB の行まで終わっていて、画面を 500 にするより「消し切れなかった」と知らせる方がよい）。
    """
    return has_app_context() and bool(g.get("purge_incomplete"))


def _truncate_leftovers(folder: Path) -> None:
    """消し切れなかったフォルダに残ったファイルを 0 バイトに切り詰める（切り詰められないものは諦める）。"""
    for path in folder.rglob("*"):
        try:
            if path.is_file():
                with open(path, "r+b") as f:
                    f.truncate(0)
        except OSError:
            pass


def purge_after_send(response, fn, *args):
    """本文を最後まで送り終えてから消す（design.md 3.3）。

    送る前に消すと、ブラウザを閉じた・通信が切れた・保存先の空きが足りない、のどれでも
    「手元にファイルが無いのにサーバー側は消えている」ことになり、取り戻せない（再ダウンロードもできない）。
    本文を最後まで渡しきったときだけ消し、途中で切れたときは残す（もう一度ダウンロードできる）。

    「渡しきった」はサーバー（waitress）の送信の溜めに入れ終えたことで、相手が受け取ったことではない。
    溜めの上限は app.OUTBUF_HIGH_WATERMARK（16KB）で、OS の溜めと合わせて最後の数十KBは受け取られる前に
    消すことになる（数十KBより小さい md は、途中で切れても消える。design.md 3.3）。

    send_file の応答は direct_passthrough が立っていて、そのままだと WSGI が close のときの
    呼び出し（call_on_close）を拾わないので、ここで解除して本文を包み直す。
    """
    if response.status_code != 200:
        # 206（Range で一部だけ）・304 などは本文の全部を渡していない。ダウンロードマネージャーや
        # 途中からの再開、ウイルス対策のプロキシが送ってくることがある。消すと残りを取り戻せないので消さない。
        return response
    app = current_app._get_current_object()
    body = response.response
    sent: list[bool] = []

    def stream():
        try:
            for chunk in body:
                yield chunk
            sent.append(True)   # 最後まで渡した
        finally:
            close = getattr(body, "close", None)
            if close is not None:
                close()

    response.direct_passthrough = False
    response.response = stream()

    @response.call_on_close
    def _purge() -> None:
        if not sent:
            return   # 送信が途中で終わった: 消さずに残す
        try:
            with app.app_context():
                fn(*args)
        except Exception:  # 応答は送り終えているので、ここで落としても利用者に伝える先が無い
            log.exception("ダウンロードしたデータの削除に失敗しました")

    return response


# ---- 帳票 --------------------------------------------------------------------------

# 帳票を消したときに呼ぶ関数（画面側がメモリに持っている帳票ごとの目印を一緒に捨てるため）。
# 引数は消した帳票の id のリスト。
_DOCUMENT_PURGE_HOOKS: list = []


def on_documents_purged(fn):
    """帳票を消したあとに fn(doc_ids) を呼ぶよう登録する（同じ関数は1回だけ）。デコレーターとしても使える。"""
    if fn not in _DOCUMENT_PURGE_HOOKS:
        _DOCUMENT_PURGE_HOOKS.append(fn)
    return fn


def purge_documents(doc_ids) -> int:
    """帳票を消す（元のファイルと DB の行）。戻り値: 消した帳票の件数。"""
    ids = [int(i) for i in doc_ids]
    if not ids:
        return 0
    db = get_db()
    marks = ", ".join("?" for _ in ids)
    stored = [row[0] for row in db.execute(f"SELECT stored_path FROM documents WHERE id IN ({marks})", ids)]
    # DB の行を先に消す。ファイルを先に消すと、DB の削除が失敗したとき（ロック・強制終了）に
    # 「ファイルの無い帳票」が作業中として残ってしまう。逆なら残るのは行の無いファイルで、起動時に片付く。
    removed = 0
    for doc_id in ids:
        _delete_by_columns(db, DOCUMENT_REF_COLUMNS, doc_id)
        _delete_jobs(db, "document", doc_id)
        removed += db.execute("DELETE FROM documents WHERE id = ?", (doc_id,)).rowcount
    db.commit()
    for path in stored:
        if not remove_upload(path):
            _mark_incomplete()
    for hook in _DOCUMENT_PURGE_HOOKS:
        try:
            hook(ids)
        except Exception:  # 目印の片付けに失敗しても、消すこと自体は終わっている
            log.exception("帳票を消したあとの片付けに失敗しました")
    _shrink(db)
    return removed


def purge_batch(batch_id: str) -> int:
    """まとめ取り込み1回分（同じ batch_id の帳票）をすべて消す。"""
    if not batch_id:
        return 0
    ids = [row[0] for row in get_db().execute("SELECT id FROM documents WHERE batch_id = ?", (batch_id,))]
    return purge_documents(ids)


# ---- 一覧表 ------------------------------------------------------------------------

def import_dir(import_id: int) -> Path:
    return Path(current_app.config["TABLES_DIR"]) / "imports" / str(int(import_id))


def purge_table_import(import_id: int) -> int:
    """一覧表の取り込み1件を消す（元のファイル・imports/<id>/・DB の行・AI整形の控え）。取り込み設定は残す。"""
    import_id = int(import_id)
    db = get_db()
    row = db.execute("SELECT stored_path, template_id FROM table_imports WHERE id = ?", (import_id,)).fetchone()
    if row is None:
        # DB の行が無くてもフォルダが残っていることがあるので、そこだけ片付ける
        shutil.rmtree(import_dir(import_id), ignore_errors=True)
        return 0
    # DB の行を先に消し、ファイルはそのあと（purge_documents と同じ理由。残ったフォルダ・ファイルは起動時に片付く）
    # この取り込みが使っていた AI の応答のキー（行を消す前に控える。消したあと、ほかで使われていなければ消す）
    keys = [r[0] for r in db.execute("SELECT cache_key FROM ai_items WHERE import_id = ? AND cache_key IS NOT NULL",
                                     (import_id,))]
    if row["template_id"] is not None:
        keys += [r[0] for r in db.execute("SELECT cache_key FROM ai_items WHERE template_id = ? AND import_id IS NULL "
                                          "AND cache_key IS NOT NULL", (row["template_id"],))]
    _delete_by_columns(db, IMPORT_REF_COLUMNS, import_id)
    _delete_jobs(db, "table_import", import_id)
    removed = db.execute("DELETE FROM table_imports WHERE id = ?", (import_id,)).rowcount
    # AI整形の結果の控え（行ごとの整形結果）は取り込んだデータそのものなので一緒に消す。この取り込みの分だけ
    # （ai_items.import_id）を消すので、同じ設定で作業中の別の取り込みの結果と応答キャッシュは残る。
    # import_id が無い行は古いDBの分（持ち主が分からない）なので、その設定の分をまとめて消す。
    if row["template_id"] is not None:
        db.execute("DELETE FROM ai_items WHERE template_id = ? AND import_id IS NULL", (row["template_id"],))
    delete_orphan_ai_items(db)
    _delete_orphan_llm_calls(db, keys)
    db.commit()
    if not remove_upload(row["stored_path"]):
        _mark_incomplete()
    folder = import_dir(import_id)
    shutil.rmtree(folder, ignore_errors=True)
    if folder.exists():
        # Windows で別のスレッド・プロセスがファイルを掴んでいると消し残る。黙って成功扱いにせず記録し、
        # remove_upload と同じく残ったファイルの中身を 0 バイトに切り詰める（名前は残っても読み込んだ行・md は残さない）。
        # DB の行はもう無いので、残ったフォルダは次の起動時に remove_orphan_import_dirs が片付ける。
        _truncate_leftovers(folder)
        _mark_incomplete()
        log.warning("取り込みのフォルダを消し切れませんでした（中身は空にし、次の起動時に片付けます）: imports/%s",
                    import_id)
    _shrink(db)
    return removed


def _shrink(db) -> None:
    """消したあとの後始末（design.md 3.3）。

    - wal_checkpoint(TRUNCATE): 消す前後の内容が残る app.db-wal を切り詰める（強制終了時に拾われないように）
    - VACUUM: 解放したページを手放してファイルの大きさも戻す（中身は PRAGMA secure_delete で消えている）
    どちらも他の接続が使っていると実行できない。そのときは次に消したときに縮む。

    VACUUM は DB 全体を書き直すあいだ書き込みを止める。大きな app.db だと busy_timeout より長くなり、
    動いている AI整形などのジョブの書き込みが「database is locked」で失敗する。なので、待機中・実行中・
    一時停止中のジョブがあるときは VACUUM をしない（消した中身は secure_delete で上書き済み。
    ファイルの大きさは、ジョブが無いときの次の削除か起動時の片付けで戻る）。

    さらに、空きページが少ないうちは VACUUM をしない。帳票の取り込みはジョブを作らないので、
    .md を1件ダウンロードするだけでも上の判定を素通りして毎回 DB 全体を書き直していた
    （118MB で1秒、その間ほかの人の自動保存が最長0.75秒待たされた。2026-09-23 のレビューで実測）。
    """
    forget_id_counters(db)
    # VACUUM → checkpoint の順。逆にすると VACUUM の結果が app.db-wal に残り、
    # instance/app.db 自体は縮まない（実測 121,260KB のまま。2026-09-23 のレビュー）
    vacuum = not _jobs_active(db) and _worth_vacuum(db)
    statements = (["VACUUM"] if vacuum else []) + ["PRAGMA wal_checkpoint(TRUNCATE)"]
    for statement in statements:
        try:
            db.execute(statement)
        except sqlite3.Error:
            pass


# 空きページがこの枚数とこの割合の両方を超えたときだけ VACUUM する（小さい削除で毎回書き直さない）
VACUUM_MIN_FREE_PAGES = 512          # 4KiB ページで約 2MB
VACUUM_MIN_FREE_RATIO = 0.25


def _worth_vacuum(db) -> bool:
    """縮める価値があるか（空きページが十分たまっているか）。"""
    try:
        free = db.execute("PRAGMA freelist_count").fetchone()[0]
        total = db.execute("PRAGMA page_count").fetchone()[0]
    except (sqlite3.Error, TypeError, IndexError):
        return True    # 数えられないときは今までどおり縮める
    return bool(free >= VACUUM_MIN_FREE_PAGES and total and free >= VACUUM_MIN_FREE_RATIO * total)


def _jobs_active(db) -> bool:
    try:
        row = db.execute("SELECT 1 FROM jobs WHERE status IN ('queued', 'running', 'paused') LIMIT 1").fetchone()
        return row is not None
    except sqlite3.Error:
        return True   # 確かめられないときは、動いているかもしれないジョブを止めない側に倒す


# ---- 番号の続きを忘れる ------------------------------------------------------------
# 取り込むたびに行を作り、ダウンロードで消す表。AUTOINCREMENT の表は、消したあとも sqlite_sequence に
# 「これまでに使った一番大きい番号」が残り、何件取り込んだかの記録になってしまう（design.md 3.3「履歴は持たない」）。
WORK_TABLES = ("documents", "table_imports", "jobs", "ai_items")
# 表が空になったら、続きの番号をこの範囲の乱数にする。最初の番号（1〜）とも前回の番号とも重ならないので、
# 開いたままの古い画面やブラウザに残った画面が、消した番号で別の取り込みを開いたり書き換えたりしない
# （重なる確率は 1回あたり 取り込み件数 / 約4.5×10^15。JavaScript で正確に扱える 2^53 未満に収める）。
_ID_BASE_MIN = 2 ** 31
_ID_BASE_MAX = 2 ** 52


def forget_id_counters(db) -> int:
    """空になった取り込みの表の、番号の続き（sqlite_sequence）を乱数に置き換える。戻り値: 置き換えた表の数。

    AUTOINCREMENT は外さない（外すと、消した番号がすぐ使い回され、古い画面が別の取り込みを指す）。
    行が残っている表は置き換えない（作業中の取り込みの番号が見えている間は、続きの番号を隠しても意味が無い）。
    """
    changed = 0
    try:
        for table in WORK_TABLES:
            base = _ID_BASE_MIN + secrets.randbelow(_ID_BASE_MAX - _ID_BASE_MIN)
            # 空かどうかの確認と置き換えを1文で行う（あいだに別のスレッドが行を足しても番号が重ならない）
            changed += db.execute(f'UPDATE sqlite_sequence SET seq = ? WHERE name = ? '
                                  f'AND NOT EXISTS (SELECT 1 FROM "{table}")', (base, table)).rowcount
        db.commit()
    except sqlite3.Error:
        # sqlite_sequence がまだ無い（AUTOINCREMENT の表に1行も入れていない新しいDB）か、ロック中。
        # 次に消したとき・次の起動時にもう一度行う
        db.rollback()
        return 0
    return changed


# ---- まとめて捨てる（作業中の表示を持たない） ----------------------------------------------
# この アプリは「作業中の一覧」も「ダウンロード待ちの一覧」も持たない（利用者の指示 2026-09-20）。
# 「その場でダウンロードしない限り、その場ですぐ捨てる」（利用者の指示 2026-09-20）ので、捨てる機会は4つ:
#   - 画面を離れたとき: その人が触っていた分を捨てる（purge_session / discard_documents / discard_table_imports）
#   - 新しいファイルを置いたとき: 同じ人の前の分を捨てる（同上）
#   - 動いている間: IDLE_HOURS さわられていないものを捨てる（sweep_stale）
#   - 起動時: ダウンロードしていない帳票・一覧表をすべて捨てる（purge_all_pending）
# あとの2つ（時間切れ・起動時）は、動いているジョブが付いているものには手を出さない（_busy_ids）。
# 残すのは「設定」（帳票の種類・ブラウザごとの AI接続 = SETTINGS_TABLES）だけで、これは purge の対象ではない。
# AI接続だけは、クッキーの寿命（約1年。views.SESSION_LIFETIME）より長くさわられていない行を
# sweep_stale_ai_connections が消す（もう戻って来ないブラウザの APIキーを DB に残さない）。

# 放っておかれた取り込みを捨てるまでの時間。数人が同時に使う社内LANの置き方（2026-09-20）では、
# 画面を閉じた合図（sendBeacon）が届かないこと（ブラウザの強制終了・スリープ・LANの切断）があるので、
# 短めにして「残らないこと」を優先する。長くしたいときはこの数字だけを変える。
IDLE_HOURS = 2
STALE_HOURS = IDLE_HOURS   # 旧名（app.py が参照している）


def _import_ids(db, where: str = "", args=()) -> list[int]:
    return [row[0] for row in db.execute(f"SELECT id FROM table_imports {where}", args)]


def _document_ids(db, where: str = "", args=()) -> list[int]:
    return [row[0] for row in db.execute(f"SELECT id FROM documents {where}", args)]


def purge_all_pending() -> tuple[int, int]:
    """ダウンロードしていない帳票・一覧表をすべて捨てる。戻り値: (帳票の件数, 一覧表の件数)。

    ダウンロードが終わったものはその時点で消えている（purge_after_send）ので、DB に残っている
    帳票・取り込みは「途中のもの」か「確定したがダウンロードしていないもの」しかない。
    続きを開く入口（作業中の一覧）を持たないので、起動時にまとめて捨てる。

    動いているジョブが付いているものには手を出さない。起動直後は _recover_jobs が動いていた
    ジョブを「中断」に直したあとなので、ふつうは1件も残らない。それでも、別のプロセスが同じ DB を
    見ているとき（JupyterLab のターミナルで二重に起動したとき）に、他方が処理中の取り込みを
    消してしまわないようにする。
    """
    db = get_db()
    busy_docs = _busy_ids(db, "document")
    busy_imports = _busy_ids(db, "table_import")
    forms = purge_documents([i for i in _document_ids(db) if i not in busy_docs])
    tables = 0
    for import_id in _import_ids(db):
        if import_id not in busy_imports:
            tables += purge_table_import(import_id)
    return forms, tables


def purge_old_sample_files() -> int:
    """前の版が置いた見本の Excel（uploads/samples）を片付ける。戻り値: 消した件数。

    帳票登録は見本の Excel を置かなくなった（利用者の指示 2026-09-21「見本のExcelは置かずに、
    設定だけ保持するようにしてほしい」）。置いた Excel はその場で読み取るだけで保存しないので、
    ふつうは1件も無い。前の版で置いたままのファイルだけを、起動時にフォルダごと捨てる。
    """
    from flask import current_app


    return remove_sample_dir(current_app.config["UPLOAD_DIR"])


def _stale_before(hours: float) -> str:
    from datetime import datetime, timedelta

    return (datetime.now() - timedelta(hours=hours)).isoformat(timespec="seconds")


# 動いているジョブ（待機中・実行中・一時停止中）が指している取り込みは捨てない。
# 大きな一覧表の読み込みや AI整形は24時間を超えることがあり、途中で消すとジョブが
# 「もう無い行」を書きに行って失敗する（利用者から見れば、放っておいたら処理が消えた、になる）。
# ジョブが終われば updated_at がその時刻になるので、次の回以降に改めて対象になる。
_BUSY = ("SELECT ref_id FROM jobs WHERE ref_type = ? AND ref_id IS NOT NULL "
         "AND status IN ('queued', 'running', 'paused')")


def _busy_ids(db, ref_type: str) -> set[int]:
    try:
        return {row[0] for row in db.execute(_BUSY, (ref_type,))}
    except sqlite3.Error:
        return set()   # 確かめられないときは何も捨てない側に倒せないので、呼び出し元で空集合＝全部対象


def _sweep_orphan_uploads() -> int:
    """行の無いアップロードファイル（置いたあと DB への登録が失敗したもの）を片付ける。

    以前は起動時にしか見ていなかったので、動かしている間にできた孤児（最大200MB）が
    次の再起動まで残っていた（2026-09-23 のレビュー）。
    """
    try:
        db = get_db()
        known: set[str] = set()
        for (table,) in db.execute("SELECT name FROM sqlite_master WHERE type = 'table'").fetchall():
            columns = {row[1] for row in db.execute(f'PRAGMA table_info("{table}")')}
            if "stored_path" in columns:
                known |= {row[0] for row in db.execute(f'SELECT stored_path FROM "{table}"') if row[0]}
        return remove_orphan_uploads(current_app.config["UPLOAD_DIR"], known)
    except (sqlite3.Error, OSError, RuntimeError, KeyError):
        return 0     # 片付けは best effort。ここで見回り全体を止めない


def sweep_stale(hours: float = STALE_HOURS) -> tuple[int, int]:
    """しばらくさわられていない帳票・一覧表を捨てる。戻り値: (帳票の件数, 一覧表の件数)。

    帳票も一覧表も「最後にさわった日時」（途中保存・読み取り → 確定 → 取り込み の順に見る）で切る。
    まとめ取り込みは、そのまとまりのどれか1件でも新しければ、まとまりごと残す（50件を上から順に
    見ていくと、まだ手が届いていない帳票だけが画面から消えてしまうため）。
    動いているジョブが付いているものは、そのジョブが終わるまで残す。
    """
    try:
        reap_orphan_jobs()   # 書き込みなので database is locked で落ちうる。片付け全体は止めない
    except sqlite3.Error:
        pass
    _sweep_orphan_uploads()
    db = get_db()
    limit = _stale_before(hours)
    busy_docs = _busy_ids(db, "document")
    busy_imports = _busy_ids(db, "table_import")
    # 見るのは updated_at が先。confirmed_at は「最後に確定した時刻」で、そのあと直し続けても
    # 進まない。先に見ていたころは、10:00 に確定して 12:30 まで直していた帳票が 12:05 の見回りで
    # 消えていた（画面の説明とも逆。2026-09-23 のレビューで実測）
    forms = purge_documents([i for i in _document_ids(
        db,
        "WHERE COALESCE(updated_at, confirmed_at, created_at) < ? "
        "AND (batch_id = '' OR NOT EXISTS (SELECT 1 FROM documents s WHERE s.batch_id = documents.batch_id "
        "AND COALESCE(s.updated_at, s.confirmed_at, s.created_at) >= ?))",
        (limit, limit)) if i not in busy_docs])
    tables = 0
    for import_id in _import_ids(
            db, "WHERE COALESCE(updated_at, confirmed_at, created_at) < ?", (limit,)):
        if import_id in busy_imports:
            continue
        tables += purge_table_import(import_id)
    sweep_stale_ai_connections()
    return forms, tables


# ブラウザごとの AI接続を消すまでの日数。クッキーの寿命（views.SESSION_LIFETIME = 365日）を過ぎたブラウザは
# 同じ id で戻って来られないので、その行（APIキー）を持ち続ける理由が無い。保存・接続の確認・AI整形の
# 開始のたびに updated_at が進むので、使っている人の分は消えない。
AI_CONNECTION_KEEP_DAYS = 400


def sweep_stale_ai_connections(days: float = AI_CONNECTION_KEEP_DAYS) -> int:
    """クッキーの寿命より長くさわられていないブラウザの AI接続（APIキー）を消す。戻り値: 消した件数。"""
    db = get_db()
    limit = _stale_before(days * 24)
    try:
        removed = db.execute("DELETE FROM ai_connections WHERE updated_at < ?", (limit,)).rowcount
        db.commit()
    except sqlite3.Error:
        db.rollback()
        return 0
    return removed


# ---- 使っている人ごとに捨てる ---------------------------------------------------------
# 社内LANに置いて数人が同時に使う（利用者の指示 2026-09-20）ので、「作業中のもの」は
# 全員分がひとつの DB に混ざっている。画面を離れた人の分だけを捨てられるように、
# documents / table_imports は持ち主（session_id。views.current_session_id() がブラウザごとに配る）を持つ。
#
# 持ち主の列がまだ無い古い DB でも動くようにしてある:
#   - 番号を指して捨てる（discard_documents / discard_table_imports）… 持ち主を確かめずに捨てる
#     （1人で使っていた頃と同じ動き。指された番号は、その画面が自分で取り込んだものしかない）
#   - まとめて捨てる（purge_session）… 誰のものか分からないので何も捨てない（他人の分を巻き込まない）

SESSION_COLUMN = "session_id"


def _has_session_column(db, table: str) -> bool:
    try:
        return SESSION_COLUMN in {row[1] for row in db.execute(f'PRAGMA table_info("{table}")')}
    except sqlite3.Error:
        return False


def _owned_ids(db, table: str, ids, session_id) -> list[int]:
    """ids のうち、いま DB にあって、その人のもの（持ち主の列が無ければ確かめない）。

    もう無い番号・他人の番号はここで落ちるので、呼び出しを何度繰り返しても2回目からは何もしない
    （画面を閉じる合図は同じものが2回届くことがある）。
    """
    wanted = sorted({int(i) for i in ids})
    if not wanted:
        return []
    marks = ", ".join("?" for _ in wanted)
    sql = f'SELECT id FROM "{table}" WHERE id IN ({marks})'
    args = list(wanted)
    if session_id and _has_session_column(db, table):
        sql += f' AND "{SESSION_COLUMN}" = ?'
        args.append(session_id)
    try:
        return [row[0] for row in db.execute(sql, args)]
    except sqlite3.Error:
        return []


def _session_ids(db, table: str, session_id) -> list[int]:
    if not session_id or not _has_session_column(db, table):
        return []
    try:
        return [row[0] for row in db.execute(f'SELECT id FROM "{table}" WHERE "{SESSION_COLUMN}" = ?', (session_id,))]
    except sqlite3.Error:
        return []


def discard_documents(doc_ids, session_id=None) -> int:
    """指された帳票のうち、その人のもので、処理中でないものを捨てる。戻り値: 捨てた件数。

    画面を離れたとき・新しいファイルを置いたときに呼ぶ。もう無いもの・他人のもの・処理中のものは
    黙って飛ばす（何度呼んでも安全で、無駄な読み書きもしない）。
    """
    db = get_db()
    ids = _owned_ids(db, "documents", doc_ids, session_id)
    if not ids:
        return 0
    busy = _busy_ids(db, "document")
    return purge_documents([i for i in ids if i not in busy])


def discard_table_imports(import_ids, session_id=None) -> int:
    """指された一覧表の取り込みのうち、その人のもので、処理中でないものを捨てる。戻り値: 捨てた件数。"""
    db = get_db()
    ids = _owned_ids(db, "table_imports", import_ids, session_id)
    if not ids:
        return 0
    busy = _busy_ids(db, "table_import")
    removed = 0
    for import_id in ids:
        if import_id not in busy:
            removed += purge_table_import(import_id)
    return removed


def purge_session(session_id, *, include_busy: bool = False,
                  documents: bool = True, tables: bool = True) -> tuple[int, int]:
    """その人の、まだダウンロードしていない帳票・一覧表をすべて捨てる。戻り値: (帳票の件数, 一覧表の件数)。

    捨てるのは、元のファイル・imports/<id>/（読み込んだ行・控え・作った md）・DB の行・
    AI整形の控えと、ほかから使われなくなった AI の応答（purge_documents / purge_table_import と同じ）。
    残すのは設定（帳票の種類・一覧表の取り込み設定・AI接続）だけ。

    documents / tables で片方だけにできる。帳票取り込みの画面を閉じた合図で
    表の取り込みの画面（同じブラウザの別のタブ）の作業まで巻き込まないために使う。

    動いているジョブ（読み込み・下書き・AI整形）が付いているものは、そのジョブが壊れるので捨てない。
    捨て損ねた分は、ジョブが終わったあと sweep_stale が IDLE_HOURS で片付ける。
    どうしても今すぐ捨てるときだけ include_busy=True（ジョブは次の書き込みで失敗して終わる）。
    """
    if not session_id:
        return 0, 0
    db = get_db()
    doc_ids = _session_ids(db, "documents", session_id) if documents else []
    import_ids = _session_ids(db, "table_imports", session_id) if tables else []
    if not include_busy:
        busy_docs = _busy_ids(db, "document")
        busy_imports = _busy_ids(db, "table_import")
        doc_ids = [i for i in doc_ids if i not in busy_docs]
        import_ids = [i for i in import_ids if i not in busy_imports]
    forms = purge_documents(doc_ids)
    tables = 0
    for import_id in import_ids:
        tables += purge_table_import(import_id)
    return forms, tables


####################################################################################################
# 画面（旧 views.py）
####################################################################################################

from app.extract import (
    apply_manual_values,
    extract_document,
    is_blank_value,
    refresh_summary,
    clean_table_value,
    is_table_value,
    parse_table_text,
    table_text_lines,
    EXCEL_ERROR_WARNING,
    WorkbookInfo,
    load_workbook_info,
    build_markdown,
    markdown_filename,
    rank_patterns,
    rank_batch,
    table_like_sheets,
    suggest_title_fields,
    click_field,
    merge_labels,
    merge_target,
    same_sheet_field,
    separate_names,
    split_rows,
    table_cells,
    pattern_to_meta,
    pattern_to_rows,
    rows_to_pattern,
    match_pattern,
    PatternDef,
)
from app.extract import (
    count_levels,
    has_blocking,
    suggest_columns,
    ai_point_lines,
    people_index_for,
    record_block,
    open_source,
    spec_from_suggestions,
    validate_spec,
    build_download,
    create_import,
    get_import,
    import_files,
    import_source,
    issues_csv as build_issues_csv,
    load_issues,
    load_rows,
    load_rows_page,
    md_paths,
    md_text,
    preview_signature,
    ready_preview_files,
    save_spec,
    spec_for_import,
    start_preview_job,
    start_read_job,
    start_render_job,
    update_import,
)


# ====================================================================================================
# 元 views/__init__.py
# ====================================================================================================

# ---- 使っている人ごとの作業場所 ---------------------------------------------------------
# 社内LANのサーバーで動かし、数人が同時に別々のPCから使う（2026-09-20 の利用者の指示）。
# ログインは無いので「誰か」は分からないが、「どのブラウザか」はセッションのクッキーで分かる。
# 取り込んだ帳票・一覧表はそのブラウザのものとして持ち主を記録し、ほかのブラウザからは
# 見えない・触れないようにする（views/forms.py・views/tables.py の 404）。
# 「設定」のうち帳票の種類はみんなで使うので分けない。AI接続（APIキー・接続先）だけは、同じ id で
# ブラウザごとに持つ（利用者の指示 2026-09-21「cookieでユーザー毎に登録内容をずっと保持」。ai.py・core.ai_connections）。
# そのためクッキーは約1年もたせる（SESSION_LIFETIME。create_app が PERMANENT_SESSION_LIFETIME に入れる）。
# 長くしても分かれ方は変わらない: id は uuid4 で、署名鍵（SECRET_KEY）付きの HttpOnly・SameSite=Lax のクッキー
# にしか無く、ほかのブラウザからは推測も持ち出しもできない。

SESSION_ID_KEY = "sid"
SESSION_LIFETIME = timedelta(days=365)


def current_session_id() -> str:
    """このブラウザの作業場所の id（クッキーに無ければ作る）。クッキーは約1年もたせる。"""
    sid = session.get(SESSION_ID_KEY)
    if not isinstance(sid, str) or len(sid) != 32:
        sid = uuid4().hex
        session[SESSION_ID_KEY] = sid
    session.permanent = True   # 有効期限（PERMANENT_SESSION_LIFETIME）付きのクッキーにする。要求のたびに延びる
    return sid


def owns(row) -> bool:
    """帳票・取り込みの行がこのブラウザのものか。

    session_id が空の行は持ち主が分からない（この仕組みを入れる前のDBの行・テストで直接作った行）ので、
    これまでどおり誰からでも扱えるものとして扱う。起動時の片付けで消えるので、普段は残らない。
    """
    owner = (row or {}).get("session_id") if hasattr(row, "get") else None
    return not owner or owner == current_session_id()


# ---- ダウンロードのファイル名 ---------------------------------------------------------
# Content-Disposition には UTF-8 の名前（filename*）と、それを読めない相手向けの ASCII の名前
# （filename）が並ぶ。日本語だけの名前だと ASCII 側が「_2026-09-19.zip」のように意味を失い、
# ログや古い環境では何のファイルか分からなくなるので、英字の既定名にそろえる。

_ASCII_UNSAFE = re.compile(r"[^A-Za-z0-9._-]+")


def ascii_download_name(name: str, default_stem: str) -> str:
    """name から作る ASCII のファイル名。英字が残らなければ default_stem を使う（拡張子は保つ）。"""
    stem = _ASCII_UNSAFE.sub("_", Path(name).stem).strip("._-")
    if not re.search(r"[A-Za-z]", stem):
        stem = f"{default_stem}_{stem}".strip("._-") if stem else default_stem
    return f"{stem}{Path(name).suffix}"


def set_download_name(response, name: str, default_stem: str):
    """添付ファイル名を「意味のある ASCII 名 ＋ 日本語の名前」にそろえる。"""
    fallback = ascii_download_name(name, default_stem)
    encoded = quote(name, safe="")
    response.headers["Content-Disposition"] = (
        f"attachment; filename=\"{fallback}\"; filename*=UTF-8''{encoded}")
    return response


# ---- 書き込みの発火元を確かめる（他サイトからの操作を断る） ----------------------------
# 127.0.0.1 だけで待ち受けても、「利用者が開いた別のサイトのページが、そのブラウザから
# このアプリへ POST する」ことは防げない（AI接続先の書き換え・帳票の削除などができてしまう）。
# ブラウザは POST に必ず Origin を付け、最近のブラウザは Sec-Fetch-Site も付けるので、
# 他サイト発と分かる書き込みだけを断る（curl などヘッダの無い要求はアプリを直接たたく操作として許す）。

SAFE_METHODS = ("GET", "HEAD", "OPTIONS")
SAME_SITE_FETCH = ("same-origin", "none")

# GET でもデータを消すルート（Markdown をダウンロードすると、その帳票・取り込みを消す。design.md 3.3）。
# 他サイトのページの <img>・リンク・window.open からでも GET は出せるので、書き込みと同じく発火元を確かめる。
# アドレス欄に打った URL（Sec-Fetch-Site: none）とアプリ内のクリック（same-origin）は通す。
PURGING_ENDPOINTS = frozenset({"forms.download_md", "forms.download_batch", "tables.download_zip"})


def _origin_of(url: str) -> str:
    parts = urlsplit(url)
    return f"{parts.scheme}://{parts.netloc}" if parts.scheme and parts.netloc else ""


def is_cross_site_write(req=None) -> bool:
    """他サイトのページから出された書き込み要求（データを消すダウンロードを含む）なら True。"""
    req = req if req is not None else request
    purging = req.endpoint in PURGING_ENDPOINTS
    if req.method in SAFE_METHODS and not purging:
        return False
    site = (req.headers.get("Sec-Fetch-Site") or "").strip().lower()
    if site and site not in SAME_SITE_FETCH:
        return True
    host = req.host_url.rstrip("/")
    origin = (req.headers.get("Origin") or "").strip()
    if origin and origin.rstrip("/") != host:
        return True
    if purging and not site:
        # Sec-Fetch-Site を付けない古いブラウザ向け。Referer が別サイトなら断る（無ければ手入力として通す）
        referer = (req.headers.get("Referer") or "").strip()
        return bool(referer) and _origin_of(referer) != host
    return False


def render_part(template: str, part: str, **ctx) -> str:
    """画面のテンプレート（forms.html など）の中のマクロ part_* を1つだけ描いて、HTML の断片を返す。

    段の中身は fetch でそのつど取りに来るので、ページ全体ではなく断片だけを描く。マクロは render_template と同じ
    文脈（url_for・request・ctx の変数）を見る。マクロに引数があれば ctx の同じ名前の値を渡す。
    """
    current_app.update_template_context(ctx)
    module = current_app.jinja_env.get_template(template).make_module(ctx)
    fn = getattr(module, part)
    return str(fn(**{k: ctx[k] for k in fn.arguments if k in ctx}))


# ====================================================================================================
# 元 views/forms.py
# 帳票取り込み（1ファイル＝1件）。画面は /forms の1枚だけ。
#
# 上から順に「ファイルを置く」「帳票の種類とシート」「読み取り結果」「確定してダウンロード」の
# 欄が現れる。画面の移動はなく、どの操作も fetch でこのファイルのルートを呼び、返ってきた
# HTML の断片をその場に入れ替える（views/forms.py のルートは JSON か HTML の断片を返す）。
#
# Markdown は常にデータから作る（読み取り結果のプレビュー＝作業中の値、ダウンロード＝確定済みの値）。
# 複数ファイルをまとめて置くと「取り込みのまとまり（batch）」になり、同じフォームの帳票として
# まとめて1回だけ種類とシートを決め、読み取り結果は全部の帳票を縦に並べて見ていく（タブで1件ずつ
# 切り替えない。利用者の指示 2026-09-20）。最後に zip でまとめて渡す。
# ダウンロードしたデータはその場で消す（design.md 3.3）。消すのは Markdown を作り終え、本文を送り終えたあとだけ
# （途中で切れたときは消さない: core.purge.purge_after_send）。
# ====================================================================================================

forms_bp = Blueprint("forms", __name__, url_prefix="/forms")

# 左のシートプレビューに出す範囲の上限（大きいシートで画面が重くならないように）
GRID_MAX_ROWS = 300
GRID_MAX_COLS = 60
CONFIRMED_STATES = ("confirmed", "modified")
MAX_BATCH_FILES = 50
# ダウンロードでデータが消えることの案内（画面の文言・確認ダイアログで使う）
FORMS_DELETE_ON_DOWNLOAD_NOTE = ("ダウンロードすると、この帳票の元のファイルと読み取り結果はサーバーから消えます。"
                           "同じものをもう一度ダウンロードすることはできません。")
FORMS_DELETE_ON_DOWNLOAD_CONFIRM = "ダウンロードすると、この帳票のデータはサーバーから消えます。もう一度ダウンロードすることはできません。"
# 確定したあとに直した（修正中の）帳票は、直した値が .md に入るよう確定し直してから渡す
MODIFIED_DOWNLOAD_CONFIRM = "確定し直してから、直した値で Markdown を作ります。" + FORMS_DELETE_ON_DOWNLOAD_CONFIRM
BATCH_DELETE_CONFIRM = ("ダウンロードすると、このまとまりの帳票のデータはサーバーからすべて消えます。"
                        "もう一度ダウンロードすることはできません。")
LOST_WORK_MESSAGE = "読み取り直すと、手で修正した値は失われます"
# 別のタブ・古い画面から保存・確定されたときの案内
STALE_MESSAGE = "別の画面で内容が変わりました。画面を読み込み直してください"


# ---- 共通 -------------------------------------------------------------------------

def _get_document(doc_id: int) -> dict:
    """この画面（このブラウザ）の帳票を返す。ほかの人の帳票は「無い」として扱う（404）。

    403 にすると「その番号の帳票はある」ことが分かってしまうので、404 にそろえる。
    """
    doc = get_document(doc_id)
    if doc is None or not owns(doc):
        abort(404)
    return doc


def _load_info(doc: dict) -> WorkbookInfo | None:
    """元のファイルを読む。消えている・壊れている場合は None。"""
    try:
        return load_workbook_info(upload_path(doc["stored_path"]))
    except Exception:
        return None


def _data(doc: dict, column: str = "data_json") -> dict | None:
    raw = doc.get(column)
    if not raw:
        return None
    try:
        return json.loads(raw)
    except ValueError:
        return None


def _dumps_extraction(extraction: dict) -> str:
    # 状態（修正中かどうか）は文字列の比較で決まるので、保存は常にこの形にそろえる
    return json.dumps(extraction, ensure_ascii=False)


def _search_title(doc: dict, extraction: dict) -> str:
    """一覧に出す見出し（Markdown の1行目）。"""
    first = build_markdown(doc, extraction).split("\n", 1)[0]
    return first[2:].strip() if first.startswith("# ") else first.strip()


def _display_text(value) -> str:
    if is_table_value(value):
        return table_json(value)
    return "" if value is None else str(value)


@forms_bp.app_template_filter("table_json")
def table_json(value) -> str:
    """明細表の値を画面の入力欄（hidden）に入れる JSON。"""
    return json.dumps(value, ensure_ascii=False) if is_table_value(value) else ""


@forms_bp.app_template_filter("field_text")
def field_text(value) -> str:
    """一覧・読み取りテスト用の表示。明細表は1行1明細の「品番: X／品名: Y」。"""
    if is_table_value(value):
        return "\n".join(table_text_lines(value))
    return "" if value is None else str(value)


def _same_text(a: str, b: str) -> bool:
    norm = lambda s: s.replace("\r\n", "\n").replace("\r", "\n").strip()  # noqa: E731
    return norm(a) == norm(b)


def _apply_values(extraction: dict, values: dict, confirmed: dict | None = None) -> bool:
    """画面の入力値を反映する。表示中の値と同じ文字列の項目は触らない（勝手に「手で修正」にしない）。

    confirmed（確定済みの版）を渡すと、確定済みと同じ値に戻した項目は確定済みの項目そのものに戻す。
    戻り値: 変わった項目があれば True。
    """
    before = _dumps_extraction(extraction)
    fields = {f["field_name"]: f for f in extraction["fields"]}
    changed = {}
    for name, text in values.items():
        f = fields.get(name)
        if f is None or text is None:
            continue
        text = text if isinstance(text, str) else json.dumps(text, ensure_ascii=False)
        if f["data_type"] == "table":
            parsed, warning = parse_table_text(text)
            # 行の無い表（列見出しだけ）と空の値は同じもの
            if warning is None and (parsed == clean_table_value(f["value"])
                                    or (is_blank_value(parsed) and is_blank_value(f["value"]))):
                continue
        elif _same_text(text, _display_text(f["value"])):
            continue
        changed[f"value-{name}"] = text
    if changed:
        apply_manual_values(extraction, changed)
        _restore_confirmed_fields(extraction, confirmed)
    else:
        refresh_summary(extraction)
    return _dumps_extraction(extraction) != before


# 手で直すと変わる項目のキー。これ以外（表示名・Markdownへの出し方など）が違う項目は戻さない
_EDIT_KEYS = {"value", "unit", "warning", "edited"}
_LOCATION_KEYS = {"sheet", "value_cell"}


def _restore_confirmed_fields(extraction: dict, confirmed: dict | None) -> None:
    """確定済みと同じ値（数値は単位も）になった項目を、確定済みの項目の中身に戻す。"""
    new_pattern, old_pattern = extraction.get("pattern", {}), (confirmed or {}).get("pattern", {})
    if (not confirmed or old_pattern.get("id") != new_pattern.get("id")
            or old_pattern.get("version_no") != new_pattern.get("version_no")
            or confirmed.get("sheets") != extraction.get("sheets")):
        return  # 別の種類・版・シートで読み直した版は比べない
    before = {f["field_name"]: f for f in confirmed.get("fields", [])}
    for i, f in enumerate(extraction["fields"]):
        old = before.get(f["field_name"])
        same_value = old is not None and (old.get("value") == f.get("value") or (
            f.get("data_type") == "table" and is_blank_value(old.get("value")) and is_blank_value(f.get("value"))))
        if (old is not None and old != f and same_value
                and (old.get("unit") or "") == (f.get("unit") or "")
                and {k: v for k, v in old.items() if k not in _EDIT_KEYS | _LOCATION_KEYS}
                == {k: v for k, v in f.items() if k not in _EDIT_KEYS | _LOCATION_KEYS}):
            extraction["fields"][i] = json.loads(json.dumps(old))
    refresh_summary(extraction)


def _is_blank(value) -> bool:
    return is_blank_value(value)


_UNIT_WARNING_PREFIXES = ("この項目の単位は", "単位が書かれていません")


def _field_status(f: dict) -> dict:
    """項目の状態タグ: 値の出どころ（auto/manual/blank）と要確認かどうか。"""
    blank = _is_blank(f["value"])
    if f.get("edited"):
        source = "manual"
    elif blank:
        source = "blank"
    else:
        source = "auto"
    # 単位の食い違い・単位なしの警告は、手で直した値でも Markdown に誤った単位で出るので要確認のまま
    unit_issue = f["data_type"] == "number" and str(f.get("warning") or "").startswith(_UNIT_WARNING_PREFIXES)
    date_issue = f["data_type"] == "date" and bool(f.get("warning"))
    value = f["value"]
    number_issue = f["data_type"] == "number" and (
        "数値の部分だけ" in str(f.get("warning") or "")
        or not isinstance(value, (int, float)) or isinstance(value, bool))
    error_value = source == "blank" and str(f.get("warning") or "").startswith(EXCEL_ERROR_WARNING)
    issue = error_value or (
        not blank and bool(f.get("warning")) and (source == "auto" or unit_issue or date_issue or number_issue))
    return {"source": source, "issue": bool(issue), "blank": blank}


def _summary(doc: dict, extraction: dict) -> dict:
    """読み取り結果の欄の数値（チップ・Markdownプレビュー・項目ごとの状態）。"""
    statuses = {f["field_name"]: _field_status(f) for f in extraction["fields"]}
    return {
        "markdown": build_markdown(doc, extraction),
        "file_name": markdown_filename(doc, extraction),
        "state": doc["state"],
        "counts": {
            "issue": sum(1 for s in statuses.values() if s["issue"]),
            "manual": sum(1 for s in statuses.values() if s["source"] == "manual"),
            "blank": sum(1 for s in statuses.values() if s["blank"]),
        },
        "fields": {name: {**s, "warning": f.get("warning") or ""}
                   for name, s, f in ((f["field_name"], statuses[f["field_name"]], f) for f in extraction["fields"])},
    }


def _sheet_grids(info: WorkbookInfo | None, sheet_names: list[str]) -> list[dict]:
    """元のシートを HTML の表にするための行データ（結合セルは rowspan/colspan）。"""
    if info is None:
        return []
    from openpyxl.utils import get_column_letter

    sheets = []
    for name in sheet_names:
        grid = info.grids.get(name)
        if grid is None:
            continue
        max_row, max_col = min(grid.max_row, GRID_MAX_ROWS), min(grid.max_col, GRID_MAX_COLS)
        rows = []
        for r in range(1, max_row + 1):
            cells = []
            for c in range(1, max_col + 1):
                top, left, bottom, right = grid.bounds(r, c)
                if (top, left) != (r, c):
                    continue  # 結合範囲の左上以外は描かない
                cell = grid.cells.get((r, c))
                cells.append({
                    "coord": f"{get_column_letter(c)}{r}",
                    "text": cell.text if cell else "",
                    "rowspan": min(bottom, max_row) - r + 1,
                    "colspan": min(right, max_col) - c + 1,
                    "label": bool(cell and (cell.bold or cell.filled)),
                })
            rows.append({"index": r, "cells": cells})
        sheets.append({
            "name": name,
            "rows": rows,
            "letters": [get_column_letter(c) for c in range(1, max_col + 1)],
            "truncated": grid.max_row > GRID_MAX_ROWS or grid.max_col > GRID_MAX_COLS,
            "images": len(info.images_in([name])),
        })
    return sheets


def _version(doc: dict) -> str:
    """読み取り結果を出したときの作業データの版（古い画面からの保存・確定を見分ける）。"""
    return hashlib.sha1(str(doc.get("data_json") or "").encode("utf-8")).hexdigest()[:16]


def _forms_payload() -> dict:
    payload = request.get_json(force=True, silent=True)
    return payload if isinstance(payload, dict) else {}


def _is_stale(doc: dict) -> bool:
    """画面が送ってきた版が今の作業データと違うか。版を送らない呼び出しは比べない。"""
    sent = _forms_payload().get("version") or request.form.get("version")
    return bool(sent) and str(sent) != _version(doc)


# ---- 1枚の画面 ----------------------------------------------------------------------

@forms_bp.get("/", endpoint="index")
@forms_bp.get("/new", endpoint="new")
def forms_page():
    """帳票取り込みの1枚の画面。ここから先はすべて fetch で欄が増えていく。"""
    current_session_id()   # 画面を開いた時点で作業場所（クッキー）を決めておく（同時に置かれても取り違えない）
    return render_template("base.html", screen="forms",
                           active_pattern_count=count_active_patterns(),
                           pattern_count=len(list_patterns()),
                           max_files=MAX_BATCH_FILES,
                           max_mb=(current_app.config.get("MAX_CONTENT_LENGTH") or 0) // (1024 * 1024))


# ---- ファイルを置く ------------------------------------------------------------------

def upload_error_text(storage, exc: Exception, position: int | None = None) -> str:
    """取り込めなかったファイルのメッセージ。まとめて置いたときは「3件目のファイル」と位置で示す。"""
    message = str(exc)
    name = original_name(storage)
    if name and message.startswith(f"{name}: "):
        message = message[len(name) + 2:]
    return f"{position}件目のファイル: {message}" if position else message


def _store_document(storage, batch_id: str = "", order: int = 0) -> int:
    """1ファイルを保存して帳票を作る。読めないファイルは保存先から消して UploadError。"""
    cfg = current_app.config
    # 帳票1ファイルの上限はここ。一覧表むけの設定（TABLE_MAX_UPLOAD_BYTES）を広げても
    # 帳票の上限は動かさない（別の話なので連動させない。2026-09-23 のレビュー）
    stored = save_upload(storage, "documents", cfg["ALLOWED_EXTENSIONS"], cfg["FORM_MAX_UPLOAD_BYTES"])
    try:
        path = upload_path(stored.stored_path)
        precheck_excel(path, cfg.get("EXCEL_MAX_CELLS"), max_merged=FORM_MAX_MERGED_CELLS)
        try:
            info = load_workbook_info(path)
        except Exception as exc:
            current_app.logger.warning("帳票を読み込めませんでした: %s", exc.__class__.__name__)
            raise UploadError("Excelファイルとして読み込めませんでした") from exc
        if not info.grids:
            raise UploadError("シートがないブックです。シートのあるブックを選んでください")
    except Exception:
        remove_upload(stored.stored_path)   # 思わぬエラーでもアップロードしたファイルを残さない（design.md 3.3）
        raise
    return create_document(stored.file_name, stored.file_hash, stored.stored_path,
                              batch_id=batch_id, batch_order=order, session_id=current_session_id())


@forms_bp.post("/upload", endpoint="upload")
def forms_upload():
    """置かれたファイルを取り込む（fetch）。複数なら1つのまとまり（batch）にして、同じ画面で1件ずつ読む。"""
    storages = [s for s in request.files.getlist("file") if s is not None and s.filename]
    if not storages:
        return jsonify(error="ファイルを置いてください"), 400
    if len(storages) > MAX_BATCH_FILES:
        return jsonify(error=f"一度に置けるのは{MAX_BATCH_FILES}ファイルまでです（置かれたのは{len(storages)}ファイル）"), 400
    batch_id = uuid4().hex if len(storages) > 1 else ""
    docs, errors = [], []
    for order, storage in enumerate(storages):
        try:
            doc_id = _store_document(storage, batch_id=batch_id, order=order)
        except UploadError as exc:
            errors.append(upload_error_text(storage, exc, order + 1 if len(storages) > 1 else None))
            continue
        doc = get_document(doc_id)
        docs.append({"id": doc_id, "file_name": doc["file_name"] if doc else ""})
    if not docs:
        return jsonify(error=errors[0] if errors else "取り込めるファイルがありませんでした", errors=errors), 400
    return jsonify(docs=docs, batch_id=batch_id, errors=errors)


# ---- 画面を離れたので捨てる ------------------------------------------------------------
# 「その場でダウンロードしない限り、その場ですぐ捨てる」（利用者の指示 2026-09-20）。
# 画面を閉じた・隠したときに static/app.js の ragDiscard がここへ「捨てて」と送ってくる。
# 送り主は navigator.sendBeacon なので、
#   - 中身の型は text/plain（get_json(force=True) で読む）
#   - 応答は読めず、やり直しもできない（いつでも 204 を返し、4xx にしない）
# もう無い番号・ほかの人の番号・処理中のものは core.purge 側で黙って外れる。

@forms_bp.post("/discard", endpoint="discard")
def forms_discard():
    """この画面（このブラウザ）の、まだダウンロードしていない帳票を捨てる。"""
    payload = request.get_json(force=True, silent=True)
    # sendBeacon は中身の形を選べない。辞書以外（"x" や [1,2]）が届いても 204 を返す
    payload = payload if isinstance(payload, dict) else {}
    ids = payload.get("doc_ids") or []
    sid = current_session_id()
    try:
        if ids:
            discard_documents(ids, sid)
        else:
            purge_session(sid, tables=False)   # 表の取り込み（別のタブ）は巻き込まない
    except Exception as exc:   # 捨て損ねてもブラウザには伝えられない。時間切れの片付けに任せる
        current_app.logger.warning("帳票の片付けに失敗しました: %s", exc.__class__.__name__)
    return "", 204


# ---- 2 帳票の種類とシート ---------------------------------------------------------------

def _id_list(raw: str) -> list[int]:
    # 桁の大きい数は SQLite に渡すと OverflowError（500）になるので、ここで落とす
    return [n for n in (int(part) for part in raw.split(",") if part.strip().isdigit())
            if 0 < n < 2 ** 63]


@forms_bp.get("/type", endpoint="type_all")
def type_all_fragment():
    """帳票の種類とシートの欄（まとめて置いた分すべてに同じ設定を使う）。?ids=1,2,3"""
    docs = _docs_of(_id_list(request.args.get("ids", "")))
    if not docs:
        return jsonify(error="取り込んだ帳票がありません。ファイルを置き直してください"), 404
    return _type_response(docs)


@forms_bp.get("/<int:doc_id>/type")
def type_fragment(doc_id: int):
    """帳票1件の「帳票の種類とシート」（まとまりでも同じ欄を使う）。"""
    return _type_response([_get_document(doc_id)])


def _batch_matches(ranked: list[dict]) -> list[SimpleNamespace]:
    """置かれたファイル全部をまとめた帳票の種類の候補（合う順。並べ方は forms.rank_batch）。

    同じフォームの帳票をまとめて置く前提なので、種類は1つだけ選ぶ。件数で違う分は
    「3項目中2〜3項目」のように幅で出し、ファイルごとの数は下のファイルの行に出す。
    先頭（先に選んでおく種類）は、最も多くのファイルで最も合った種類。「見つかった割合」だけで
    並べると 4項目中4項目の小さな種類が 33項目中32項目の本物に勝つので、割合ではなく
    forms.match_score（見つかった数・割合・シート名の一致）の点で並べる。
    """
    out = []
    for b in rank_batch(ranked):
        lo, hi, total = b.found_min, b.found_max, b.total_fields
        out.append(SimpleNamespace(
            pattern=b.pattern, total_fields=total, found_fields=lo, sheet_names=b.sheet_names,
            votes=b.votes, score=b.score, files=len(b.matches),
            found_label=(f"{total}項目中{lo}項目が見つかりました" if lo == hi
                         else f"{total}項目中{lo}〜{hi}項目が見つかりました")))
    return out


def _type_response(docs: list[dict]):
    patterns = load_active_patterns()
    ranked: list[dict] = []
    sheets: list[dict] = []
    by_name: dict[str, dict] = {}
    table_sheets: list[str] = []
    # ブックは1件ずつ開いて、必要な数だけ取り出したらすぐ手放す。50件を同時にメモリへ広げると
    # サーバーが落ちるので、置かれたファイル全部の WorkbookInfo を持ち続けない
    for doc in docs:
        info = _load_info(doc)
        if info is None:
            return jsonify(error="元のファイルを読み込めませんでした。ファイルを置き直してください"), 409
        ranked.append({m.pattern.id: m for m in rank_patterns(info, patterns)})
        for name, grid in info.grids.items():
            sheet = by_name.get(name)
            if sheet is None:
                sheet = {"name": name, "cells": len(grid.cells), "images": len(info.images_in([name])),
                         "hidden": grid.hidden, "files": 0}
                by_name[name] = sheet
                sheets.append(sheet)
            sheet["files"] += 1
        table_sheets.extend(n for n in table_like_sheets(info) if n not in table_sheets)
        info = grid = None   # 次の1件を開く前にブックを手放す

    matches = _batch_matches(ranked)
    suggested = {str(m.pattern.id): m.sheet_names for m in matches}
    first = docs[0]
    selected_id = first["pattern_id"] if any(m.pattern.id == first["pattern_id"] for m in matches) else None
    selected_id = selected_id or (matches[0].pattern.id if matches else None)
    # 読み取り済みの帳票が読んだシートを全部合わせる。1件分の "sheets" には、そのファイルに
    # 在ったシートしか残らないので、先頭ファイルだけを見ると選んだチェックが消えてしまう
    selected_sheets: list[str] = []
    for d in docs:
        data = _data(d) if d["pattern_id"] == selected_id else None
        for name in (data or {}).get("sheets") or []:
            if name not in selected_sheets:
                selected_sheets.append(name)
    if not selected_sheets:
        selected_sheets = suggested.get(str(selected_id)) or []

    files = [{"id": d["id"], "file_name": d["file_name"],
              "found": ranked[i].get(selected_id).found_fields if selected_id in ranked[i] else 0}
             for i, d in enumerate(docs)]
    # 種類を選び直したときにファイルごとの項目数を出し直すための表（画面の JS が使う）
    file_counts = {str(pid): {str(d["id"]): ranked[i][pid].found_fields for i, d in enumerate(docs)}
                   for pid in (ranked[0] if ranked else {})}
    html = render_part("base.html", "part_type", doc=first,
        docs=docs,
        files=files,
        file_counts=file_counts,
        matches=matches,
        sheets=sheets,
        suggested=suggested,
        selected_id=selected_id,
        selected_sheets=selected_sheets,
        table_sheets=table_sheets,
        duplicate=any(find_confirmed_by_hash(d["file_hash"], exclude_id=d["id"],
                                                session_id=current_session_id()) for d in docs),
        lost_work=any(d["state"] in CONFIRMED_STATES for d in docs),
        form_types_url=url_for("form_types.index"),
    )
    return jsonify(html=html, file_name=first["file_name"], has_types=bool(matches),
                   docs=[{"id": d["id"], "file_name": d["file_name"]} for d in docs])


@forms_bp.post("/read", endpoint="read_all")
def read_all():
    """置かれた帳票を1つの種類・シートでまとめて読み取り、全部の読み取り結果を返す。"""
    return _read_documents(_docs_of(_id_list(request.form.get("ids", ""))))


@forms_bp.get("/review", endpoint="review_all")
def review_all():
    """いま保存されている読み取り結果を描き直す（読み取りはやり直さない）。

    まとめ置きから1件だけ外したときに使う。以前は残りの帳票の③まで画面から消え、
    ［読み取る］で取り直すしかなく、手で直した値が警告なしに失われていた
    （2026-09-23 のレビューで実測）。ここは保存済みの内容をそのまま描くので何も失われない。
    """
    ids = [d["id"] for d in _docs_of(_id_list(request.args.get("ids", ""))) if d.get("data_json")]
    if not ids:
        return jsonify(error="読み取り結果がありません。［読み取る］を押してください"), 404
    return _review_response(ids)


@forms_bp.post("/<int:doc_id>/read")
def read(doc_id: int):
    """帳票1件を読み取る（まとまりでも同じ道すじを通る）。"""
    return _read_documents([_get_document(doc_id)])


def _read_documents(docs: list[dict]):
    if not docs:
        return jsonify(error="取り込んだ帳票がありません。ファイルを置き直してください"), 404
    pattern = load_pattern(request.form.get("pattern_id", type=int))
    sheets = request.form.getlist("sheets")
    if pattern is None or not sheets:
        return jsonify(error="帳票の種類と読み取るシートを選んでください"), 400
    if any(d["state"] in CONFIRMED_STATES for d in docs) and request.form.get("acknowledge") != "on":
        return jsonify(error=f"{LOST_WORK_MESSAGE}。確認のチェックを入れてから読み取り直してください"), 400

    read_ids, errors = [], []

    def failed(doc: dict, reason: str) -> None:
        """読み取れなかった帳票。③に並ばないので、前の読み取り結果も残さない
        （残すと、画面に出ていない帳票を④が数えて zip の件数と合わなくなる）。"""
        errors.append(f"{doc['file_name']}: {reason}")
        if doc["state"] not in CONFIRMED_STATES and doc.get("data_json"):
            reset_document(doc["id"], doc["pattern_id"], None)

    for doc in docs:
        info = _load_info(doc)
        if info is None:
            failed(doc, "元のファイルを読み込めませんでした")
            continue
        use = [s for s in sheets if s in info.grids]
        if not use:
            failed(doc, "選んだシートがファイルにありません")
            continue
        extraction = extract_document(info, pattern, use)
        reset_document(doc["id"], pattern.id, _dumps_extraction(extraction))
        if doc["state"] not in CONFIRMED_STATES:
            update_document(doc["id"], title=_search_title(doc, extraction))
        read_ids.append(doc["id"])
    if not read_ids:
        return jsonify(error=errors[0] if errors else "読み取れる帳票がありませんでした", errors=errors), 400
    return _review_response(read_ids, errors=errors)


def _state_label(doc: dict, summary: dict) -> str:
    """読み取り結果の見出しに出す、その帳票の今の状態。"""
    if doc["state"] == "confirmed":
        return "確定済み"
    if doc["state"] == "modified":
        return "修正中"
    issue = summary["counts"]["issue"]
    return f"要確認 {issue}件" if issue else "未確定"


STATE_COLORS = {"確定済み": "green", "修正中": "violet", "未確定": "gray"}


def _review_response(doc_ids, errors=None):
    """読み取り結果の欄（帳票を縦に並べた HTML の断片）。

    重くならないよう、元のシートの表を作るのは先頭の帳票だけにする。残りは画面に入った時点で
    GET /forms/<id>/grid を呼んで読み込む（static/review.js）。
    """
    ids = [doc_ids] if isinstance(doc_ids, int) else list(doc_ids)
    items, docs = [], []
    for no, doc_id in enumerate(ids, start=1):
        doc = _get_document(doc_id)
        extraction = _data(doc)
        if extraction is None:
            if len(ids) == 1:
                return jsonify(error="まだ読み取りをしていません"), 409
            continue
        summary = _summary(doc, extraction)
        info = _load_info(doc) if no == 1 else None
        label = _state_label(doc, summary)
        items.append({
            "no": no, "doc": doc, "extraction": extraction, "summary": summary,
            "grids": _sheet_grids(info, extraction["sheets"]) if no == 1 else None,
            "info_missing": no == 1 and info is None,
            "version": _version(doc), "state_label": label,
            "state_color": STATE_COLORS.get(label, "amber"),
        })
        docs.append({"id": doc["id"], "file_name": doc["file_name"], "state": doc["state"],
                     "version": _version(doc), "summary": summary, "state_label": label})
    if not items:
        return jsonify(error="まだ読み取りをしていません"), 409
    html = render_part("base.html", "part_review", items=items,
        confirmed=sum(1 for d in docs if d["state"] in CONFIRMED_STATES),
    )
    first = items[0]
    return jsonify(html=html, version=first["version"], summary=first["summary"], doc_id=first["doc"]["id"],
                   docs=docs, errors=errors or [])


@forms_bp.get("/<int:doc_id>/grid")
def grid_fragment(doc_id: int):
    """1件の帳票の「元のシート」（読み取り結果の欄が画面に入ったときに読み込む）。"""
    doc = _get_document(doc_id)
    extraction = _data(doc)
    if extraction is None:
        return jsonify(error="まだ読み取りをしていません"), 409
    info = _load_info(doc)
    html = render_part("base.html", "part_grid", grids=_sheet_grids(info, extraction["sheets"]),
                           info_missing=info is None)
    return jsonify(html=html, doc_id=doc_id)


# ---- 3 読み取り結果（その場で直す・途中保存） -----------------------------------------------

# 途中保存をした画面の目印（doc_id → (保存後の版, 画面の目印)）。
_DRAFT_TOKENS: dict[int, tuple[str, str]] = {}


@on_documents_purged
def _forget_draft_tokens(doc_ids) -> None:
    for doc_id in doc_ids:
        _DRAFT_TOKENS.pop(int(doc_id), None)


def _page_token() -> str:
    token = _forms_payload().get("page_token")
    return str(token)[:64] if token else ""


def _saved_by_same_page(doc_id: int, doc: dict) -> bool:
    token = _page_token()
    saved = _DRAFT_TOKENS.get(doc_id)
    return bool(token) and saved is not None and saved == (_version(doc), token)


def _json_values() -> dict:
    payload = _forms_payload()
    if payload:
        values = payload.get("values", payload)
        return {str(k): v for k, v in values.items()} if isinstance(values, dict) else {}
    return {key[len("value-"):]: request.form[key] for key in request.form.keys() if key.startswith("value-")}


@forms_bp.post("/<int:doc_id>/draft")
def draft(doc_id: int):
    """入力内容の途中保存（fetch）。変わった項目だけ保存し、204 を返す。"""
    doc = _get_document(doc_id)
    extraction = _data(doc)
    if extraction is None:
        return jsonify(error="まだ読み取りをしていません"), 409
    if _is_stale(doc) and not _saved_by_same_page(doc_id, doc):
        return jsonify(error=STALE_MESSAGE), 409
    version = _version(doc)
    if _apply_values(extraction, _json_values(), _data(doc, "confirmed_json")):
        title = None if doc["state"] in CONFIRMED_STATES else _search_title(doc, extraction)
        new_json = _dumps_extraction(extraction)
        save_draft(doc_id, new_json, title=title)
        version = _version({"data_json": new_json})
        token = _page_token()
        if token:
            _DRAFT_TOKENS[doc_id] = (version, token)
    return "", 204, {"X-Doc-Version": version}


@forms_bp.post("/<int:doc_id>/preview")
def preview(doc_id: int):
    """入力中の値で作った Markdown と状態（保存はしない）。"""
    doc = _get_document(doc_id)
    extraction = _data(doc)
    if extraction is None:
        return jsonify(error="まだ読み取りをしていません"), 409
    if _is_stale(doc):
        return jsonify(error=STALE_MESSAGE), 409
    _apply_values(extraction, _json_values(), _data(doc, "confirmed_json"))
    return jsonify(_summary(doc, extraction))


# ---- 4 確定してダウンロード ---------------------------------------------------------------

@forms_bp.post("/<int:doc_id>/confirm", endpoint="confirm")
def forms_confirm(doc_id: int):
    """読み取り結果を確定する（fetch）。値は途中保存で入っているので、ここでは版だけ確かめる。"""
    doc = _get_document(doc_id)
    extraction = _data(doc)
    if extraction is None:
        return jsonify(error="まだ読み取りをしていません"), 409
    if _is_stale(doc):
        return jsonify(error=STALE_MESSAGE), 409
    values = _json_values()
    if values and _apply_values(extraction, values, _data(doc, "confirmed_json")):
        save_draft(doc_id, _dumps_extraction(extraction))
    confirm_document(doc_id, title=_search_title(doc, extraction))
    return jsonify(ok=True, doc_id=doc_id)


def _docs_of(ids: list[int]) -> list[dict]:
    """この画面の帳票だけ（ほかのブラウザの帳票は、番号を送られても無いものとして外す）。"""
    return [d for d in (get_document(i) for i in ids) if d is not None and owns(d)]


@forms_bp.get("/finish", endpoint="finish")
def finish_fragment():
    """確定してダウンロードの欄（HTML の断片）。?ids=1,2,3 は同じ画面で扱っている帳票。"""
    docs = _docs_of(_id_list(request.args.get("ids", "")))
    if not docs:
        return jsonify(html="", ready=False, confirmed=0, total=0)
    # 読み取れていない帳票は確定できない＝zip に入らないので、件数には数えない
    # （数えると「残り12件も確定して…」と書いてあるのに .md が4件しか入らない zip になる）
    ready = [d for d in docs if d.get("data_json")]
    unread = [d for d in docs if not d.get("data_json")]
    current_id = request.args.get("current", type=int)
    current = next((d for d in ready if d["id"] == current_id), None) or (ready[0] if ready else docs[0])
    confirmed = [d for d in ready if d["state"] in CONFIRMED_STATES]
    pending = [d for d in ready if d["state"] not in CONFIRMED_STATES]
    batch_id = current.get("batch_id") or ""
    working = _data(current)
    read_yet = working is not None or any(_data(d) is not None for d in docs)
    extraction = _data(current, "confirmed_json") or working
    html = render_part("base.html", "part_finish", docs=docs,
        ready=ready,
        unread=unread,
        current=current,
        confirmed=confirmed,
        pending=pending,
        batch_id=batch_id if len(docs) > 1 else "",
        read_yet=read_yet,
        to_confirm=_to_confirm(ready),
        confirmed_states=CONFIRMED_STATES,
        file_name=markdown_filename(current, extraction) if extraction else "",
        markdown=build_markdown(current, _data(current, "confirmed_json")) if current["state"] in CONFIRMED_STATES else "",
        delete_note=FORMS_DELETE_ON_DOWNLOAD_NOTE,
        delete_confirm=(MODIFIED_DOWNLOAD_CONFIRM if current["state"] == "modified"
                        else FORMS_DELETE_ON_DOWNLOAD_CONFIRM),
        batch_confirm=_batch_zip_confirm(ready, unread),
    )
    return jsonify(html=html, ready=bool(confirmed), confirmed=len(confirmed), total=len(ready),
                   read_yet=read_yet,
                   next_id=(pending[0]["id"] if pending else None))


def _to_confirm(docs: list[dict]) -> list[dict]:
    """まとめてのダウンロードの前に確定し直す帳票（未確定と、直したまま確定していない修正中）。"""
    return [d for d in docs if d["state"] != "confirmed"]


def _batch_zip_confirm(docs: list[dict], unread: list[dict] | None = None) -> str:
    """まとまりの zip ダウンロードの確認文（確定していない分はこのボタンで確定してから渡す）。

    docs は zip に入る帳票（読み取り済み）だけ。読み取れていない帳票はサーバーに残るので、
    「すべて消えます」とは言わない。
    """
    rest = _to_confirm(docs)
    head = (f"まだ確定していない{len(rest)}件も確定してから、{len(docs)}件をまとめて zip でダウンロードします。"
            if rest else f"{len(docs)}件をまとめて zip でダウンロードします。")
    if unread:
        return (head + "ダウンロードすると、渡したこの"
                f"{len(docs)}件のデータはサーバーから消えます。"
                f"まだ読み取れていない{len(unread)}件はサーバーに残ります（②に戻ってもう一度読み取ってください）。")
    if rest:
        return head + "ダウンロードすると、このまとまりの帳票のデータはサーバーからすべて消えます。"
    return BATCH_DELETE_CONFIRM


# ---- ダウンロード -------------------------------------------------------------------

@forms_bp.get("/<int:doc_id>/download.md")
def download_md(doc_id: int):
    """Markdown を渡し、渡し終えた帳票のデータを消す（design.md 3.3）。"""
    name, body = _single_markdown(doc_id)
    response = send_file(io.BytesIO(body), mimetype="text/markdown", as_attachment=True, download_name=name,
                         conditional=False)   # Range でも全体を返す
    set_download_name(response, name, "form")
    return purge_after_send(response, purge_documents, [doc_id])


def _single_markdown(doc_id: int) -> tuple[str, bytes]:
    doc = _get_document(doc_id)
    if doc["confirmed_json"] is None:
        abort(404)   # まだ確定していない帳票は渡さない
    # 直したまま確定していない（修正中）帳票は、直した値で作る（design.md 3.3）。画面の JS を
    # 通さずにリンクを開いたとき（中クリック・「名前を付けてリンク先を保存」）でも同じにする
    extraction = _data(doc) if doc["state"] == "modified" else _data(doc, "confirmed_json")
    if extraction is None:
        abort(404)
    return markdown_filename(doc, extraction), build_markdown(doc, extraction).encode("utf-8")


def _unique_name(name: str, used: set[str]) -> str:
    """まとまりの中で重ならない名前（大文字・小文字だけの違いも重なりとみなす）。"""
    stem, suffix = Path(name).stem, Path(name).suffix or ".md"
    candidate, n = name, 2
    while candidate.casefold() in used:
        candidate = f"{stem}_{n}{suffix}"
        n += 1
    used.add(candidate.casefold())
    return candidate


def _batch_markdown_files(confirmed_ids: list[int]) -> list[tuple[str, bytes]]:
    files, used = [], set()
    for doc in list_confirmed_documents(confirmed_ids, session_id=current_session_id()):
        # 修正中（確定したあとに直した）帳票は、直した値で作る（1件の .md と同じ）
        column = "data_json" if doc["state"] == "modified" else "confirmed_json"
        try:
            extraction = json.loads(doc[column])
        except (TypeError, ValueError):
            continue
        files.append((_unique_name(markdown_filename(doc, extraction), used),
                      build_markdown(doc, extraction).encode("utf-8")))
    return files


@forms_bp.get("/batches/<batch_id>/download.zip")
def download_batch(batch_id: str):
    """まとめ取り込み1回分の Markdown を zip で渡し、渡し終えた分のデータを消す。

    確定済みの帳票だけを zip にして、その分だけ消す（未確定の帳票はサーバーに残る）。
    """
    docs = list_batch_documents(batch_id, session_id=current_session_id())
    if not docs:
        abort(404)   # ほかのブラウザのまとまりも「無い」として扱う
    pending = [d for d in docs if d["state"] not in CONFIRMED_STATES]
    confirmed_ids = [d["id"] for d in docs if d["state"] in CONFIRMED_STATES]
    if not confirmed_ids:
        abort(404)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, body in _batch_markdown_files(confirmed_ids):
            zf.writestr(name, body)
    buffer.seek(0)
    zip_name = f"帳票Markdown_{datetime.now():%Y%m%d_%H%M%S}.zip"
    response = send_file(buffer, mimetype="application/zip", as_attachment=True, download_name=zip_name,
                         conditional=False)   # Range でも全体を返す
    set_download_name(response, zip_name, "forms_markdown")
    if pending:
        return purge_after_send(response, purge_documents, confirmed_ids)
    return purge_after_send(response, purge_batch, batch_id)


@forms_bp.get("/<int:doc_id>/original")
def original(doc_id: int):
    doc = _get_document(doc_id)
    try:
        path = upload_path(doc["stored_path"])
    except UploadError:
        abort(404)
    if not path.exists():
        abort(404)
    return send_file(path, as_attachment=True, download_name=doc["file_name"])


@forms_bp.post("/<int:doc_id>/delete", endpoint="delete")
def forms_delete(doc_id: int):
    """この帳票の取り込みをやめる（元のファイルと読み取り結果を消す）。"""
    _get_document(doc_id)
    purge_documents([doc_id])
    incomplete = purge_incomplete()
    return jsonify(ok=True, incomplete=incomplete)


# 画面を作り直す前の URL（お気に入り・古いリンク）は1枚の画面へ送る
@forms_bp.get("/<int:doc_id>", endpoint="legacy")
@forms_bp.get("/<int:doc_id>/done", endpoint="legacy")
def forms_legacy(doc_id: int):
    return redirect(url_for(".index"))


# ====================================================================================================
# 元 views/form_types.py
# 帳票登録。画面は /form-types の1枚だけ。
#
# 上から「登録済みの帳票の種類」、その下に「新しく登録する」。Excel を1つ置くと、同じ画面に
# 名前（ファイル名から入れる）・置いた Excel のシート・読み取る項目・読み取りテストの結果・
# ［使用開始］が現れる。画面の移動はなく、どの操作も fetch でこのファイルのルートを呼び、
# HTML の断片を入れ替える。
#
# 項目は「見出しのセル → 値のセル」をクリックするだけで作る。キー名・型・単位・探す見出しは
# pattern.clicks が見本の値から決めるので、画面には出さない。
# 保存しても使用中にはしない。使用中になるのは［使用開始］を押したときだけ。
#
# 置いた Excel のファイルはサーバーに残さない（利用者の指示 2026-09-21「見本のExcelは置かずに、設定だけ
# 保持するようにしてほしい」）。受け取った要求の中で読み取り、中身（bytes）はそのまま捨てる。
# その代わり、読み取ったシートの中身（セルの番地と文字・結合・太字・塗りの有る無し）を帳票の種類と一緒に
# DB が覚える（core.pattern_books。利用者の指示 2026-09-22「再度シートを置かなくても、登録したときに
# シートのセル番地と文字情報を記憶しておけばだせるはず」）。開き直したときはそれでシートを出し、同じ
# クリックの操作で直せる。セルの色は覚えない（同日の指示「セル色の情報は不要」）。
# 置いた直後の操作（セルのクリック・項目の作り直し・読み取りテスト）ではブラウザが同じ Excel を送り直して
# くることがある。開き直すのが遅いので、読み取った結果だけを core.workbook_cache が短い間メモリに覚えておく。
# 残すのは設定（シート名・見出しのセル・値のセル・読み取る向き・項目名）と覚えたシートだけ。
# ====================================================================================================

form_types_bp = Blueprint("form_types", __name__, url_prefix="/form-types")

# excel/extractor.number_unit の単位不明の警告の書き出し
_NO_UNIT_WARNING = "単位が書かれていません"
TEST_NO_UNIT_WARNING = ("帳票に単位が書かれていません。取り込んだあとの「読み取り結果」で、"
                        "値に単位（分・時間など）を付けて入力できます")


def _get_pattern(pattern_id: int) -> PatternDef:
    pattern = load_pattern(pattern_id)
    if pattern is None:
        abort(404)
    return pattern


def _confirmed_count(pattern_id: int) -> int:
    row = get_db().execute(
        "SELECT COUNT(*) FROM documents WHERE pattern_id = ? AND confirmed_json IS NOT NULL", (pattern_id,)
    ).fetchone()
    return row[0] if row else 0


# ---- ブラウザが置いた Excel（ファイルは保存しない。シートの中身だけ覚える） -------------------
# 置かれた Excel は、この要求の中で読み取って中身を捨てる。読み取ったシートの中身は帳票の種類と
# 一緒に DB が覚える（_remember_book）。次の操作のときはブラウザが同じファイルを送り直してくることが
# あるので、2回目からは読み取った結果（core.workbook_cache）を使い回す。送ってこなければ覚えたシートで出す。

BOOK_FIELD = "book"            # ブラウザが送ってくる Excel（<input type=file name=book>）
BOOK_HASH_FIELD = "book_hash"  # 送り直さずに、さっき読んだブックを指すとき（sha256）
# 覚えたシートが無い種類（覚える前に登録したもの）で、Excel も送られてこなかったとき
NO_BOOK_ERROR = ("この帳票のExcelをもう一度置いてください（この種類は登録したときのシートを覚えていません。"
                 "置くとシートを覚えて、次からは置かずに開けます）")
# ここで受け取る Excel の大きさの上限。保存せずにメモリで読む（数人が同時に置く）ので、
# 取り込みの上限（MAX_CONTENT_LENGTH＝まとめ置きの合計）より小さくしておく。
# 帳票は1枚の紙なので、写真付きでもこの大きさに収まる（見本のいちばん大きいもので約0.1MB）。
BOOK_MAX_BYTES = 50 * 1024 * 1024
# 覚えるシートの中身（JSON）の上限。帳票は1枚の紙なので、ふつうは数KB〜数十KB（見本の報告書で約4KB）。
# 何万セルもある一覧表のようなブックは帳票ではないので、覚えずに断る（画面に減らし方を書く）。
SNAPSHOT_MAX_BYTES = 4 * 1024 * 1024
# 覚えない、と画面に書くもの。読み取りに要らないものは覚えない（利用者の指示 2026-09-22「セル色の情報は不要」）
SNAPSHOT_NOTE = ("Excel のファイルは保存しません。登録したときのシートの中身（セルの番地と文字）だけを覚えておき、"
                 "あとから開いたときにその画面を出します。セルの色は覚えません"
                 "（塗りの有る無しだけを、見出しの判定のために覚えます）。")


def _snapshot_json(book: Book) -> str:
    """覚えるシートの中身（JSON）。大きすぎるブックは UploadError（帳票ではない大きさ）。"""
    text = book.info.to_json()
    size = len(text.encode("utf-8"))
    if size > SNAPSHOT_MAX_BYTES:
        raise UploadError(
            f"このExcelはシートの中身が大きすぎて覚えられません（{size / (1024 * 1024):.1f}MB。"
            f"上限 {SNAPSHOT_MAX_BYTES // (1024 * 1024)}MB）。帳票の様式だけの Excel（記入例が1件のもの）にするか、"
            "使わないシート・値の入った余分な行や列を消して、置き直してください")
    return text


def _remember_book(pattern_id: int, book: Book, book_json: str | None = None) -> None:
    """置かれた Excel のシートの中身を、この種類の「覚えたシート」にする（前のものと入れ替える）。"""
    save_pattern_book(pattern_id, book.file_name, book.file_hash, book_json or _snapshot_json(book))
    book.saved_at = now()


def _stored_book(pattern_id: int) -> Book | None:
    """この種類が覚えているシート。開いた結果は core.workbook_cache にも置く（クリックのたびに JSON を読み直さない）。"""
    row = load_pattern_book(pattern_id)
    if row is None:
        return None
    known = get(current_session_id(), row["file_hash"])
    if known is not None and known.file_name == row["file_name"]:
        known.saved_at = row["saved_at"]
        return known
    try:
        info = WorkbookInfo.from_json(row["book_json"])
    except (ValueError, TypeError, KeyError, AttributeError) as exc:
        current_app.logger.warning("覚えたシートを読めませんでした（pattern %s）: %s", pattern_id, exc.__class__.__name__)
        return None
    if not info.grids:
        return None
    return put(current_session_id(), Book(file_name=row["file_name"], file_hash=row["file_hash"],
                                          size=len(row["book_json"]), info=info, saved_at=row["saved_at"]))


def _read_book(storage) -> Book:
    """置かれた Excel をメモリで読み、読み取った結果だけを覚える。読めなければ UploadError。

    受け取ったファイルは werkzeug が 500KB まではメモリに、それを超える分だけ OS の一時ファイルに
    置く（要求が終わると消える、名前の無いファイル）。こちらからディスクに書くことはしない。
    """
    cfg = current_app.config
    limit = min(cfg["MAX_CONTENT_LENGTH"] or BOOK_MAX_BYTES, BOOK_MAX_BYTES)
    memory = read_upload(storage, cfg["ALLOWED_EXTENSIONS"], limit)
    known = get(current_session_id(), memory.file_hash)
    if known is not None:
        known.file_name = memory.file_name   # 同じ中身を別の名前で置き直したとき
        return known
    precheck_excel(memory.data, cfg.get("EXCEL_MAX_CELLS"), max_merged=FORM_MAX_MERGED_CELLS)
    try:
        # BytesIO で渡すので、どこにもファイルを作らずに読める
        info = load_workbook_info(io.BytesIO(memory.data))
    except Exception as exc:
        current_app.logger.warning("置かれたExcelを読み込めませんでした: %s", exc.__class__.__name__)
        raise UploadError("Excelファイルとして読み込めませんでした") from exc
    if not info.grids:
        raise UploadError("シートがないブックです。シートのあるブックを選んでください")
    return put(current_session_id(),
                              Book(file_name=memory.file_name, file_hash=memory.file_hash,
                                   size=memory.size, info=info))


def _request_book(pattern_id: int | None = None) -> tuple[Book | None, str]:
    """この操作で見ている Excel。戻り値: (ブック, エラー文)。

    順に探す: ブラウザが送ってきたファイル → さっき読んだブック（sha256。core.workbook_cache）→
    その種類が覚えているシート（core.pattern_books）。
    どれも無ければ (None, "")＝シートの無い画面（覚える前に登録した種類。項目の一覧と見出しの手直しはできる）。
    """
    storage = request.files.get(BOOK_FIELD)
    if storage is not None and storage.filename:
        try:
            return _read_book(storage), ""
        except UploadError as exc:
            return None, upload_error_text(storage, exc)
    file_hash = (request.form.get(BOOK_HASH_FIELD) or request.args.get(BOOK_HASH_FIELD) or "").strip()
    if file_hash:
        known = get(current_session_id(), file_hash)
        if known is not None:
            return known, ""
    if pattern_id is not None:
        return _stored_book(pattern_id), ""
    return None, ""


# ---- 1枚の画面 --------------------------------------------------------------------

@form_types_bp.get("/", endpoint="index")
def form_types_page():
    # 帳票の種類は「設定」なのでみんなで使う（ブラウザごとに分けない）。
    # 作業場所のクッキーだけは、ここで開いたときにも決めておく（ほかの画面での取り違えを防ぐ）
    current_session_id()
    return render_template("base.html", screen="form_types", list_html=_list_html(),
                           max_mb=BOOK_MAX_BYTES // (1024 * 1024), snapshot_note=SNAPSHOT_NOTE)


def _list_html() -> str:
    return render_part("base.html", "part_list", patterns=list_patterns())


# 画面を作り直す前の URL（お気に入り・古いリンク）は1枚の画面へ送る
@form_types_bp.get("/new", endpoint="legacy")
@form_types_bp.get("/<int:pattern_id>/build", endpoint="legacy")
@form_types_bp.get("/<int:pattern_id>/edit", endpoint="legacy")
@form_types_bp.get("/<int:pattern_id>/review", endpoint="legacy")
@form_types_bp.get("/<int:pattern_id>/test", endpoint="legacy")
def form_types_legacy(pattern_id: int | None = None):
    return redirect(url_for(".index"))


# ---- 新しく登録する ---------------------------------------------------------------

@form_types_bp.post("/new")
def create():
    """Excel を1つ置いて帳票の種類を作る（fetch）。名前は画面でファイル名から入れてある。

    作るのは設定と、覚えたシート（置かれた Excel のシートの中身）。Excel のファイルはサーバーに残さない。
    覚えられない大きさのブックなら種類も作らない（半端な種類を残さない）。
    """
    storage = request.files.get(BOOK_FIELD)
    if storage is None or not storage.filename:
        return jsonify(error="帳票のExcelファイルを置いてください"), 400
    name = request.form.get("name", "").strip() or Path(storage.filename).stem.strip()
    if not name:
        return jsonify(error="帳票の種類の名前を入れてください"), 400
    try:
        book = _read_book(storage)
        book_json = _snapshot_json(book)
    except UploadError as exc:
        return jsonify(error=upload_error_text(storage, exc)), 400
    pattern_id = create_pattern(name)
    _remember_book(pattern_id, book, book_json)
    return jsonify(pattern_id=pattern_id, html=_build_html(pattern_id, book), list_html=_list_html(),
                   message=f"「{name}」を作りました。読み取りたい欄の見出しと値をクリックしてください"
                           "（Excel のファイルは保存せず、シートの中身だけを覚えました）")


# ---- 読み取る欄をクリックして決める ＋ 読み取りテスト ------------------------------------

def _build_html(pattern_id: int, book: Book | None = None, notes: list[str] | None = None) -> str:
    """登録中の帳票の種類の欄（HTML の断片）。book が無ければシートの無い（設定だけの）画面。

    remembered は覚えたシートの見出し（ファイル名・置いた日時）。いま出しているブックがそれなら画面に書く。
    """
    pattern = _get_pattern(pattern_id)
    info = book.info if book is not None else None
    grids = _sheet_grids(info, list(info.grids)) if info is not None else []
    for g in grids:
        g["click_cells"] = table_cells(info.grids[g["name"]])
    remembered = pattern_book_meta(pattern_id)
    if remembered is not None and (book is None or book.file_hash != remembered["file_hash"]):
        remembered = None   # 覚えたシートとは別の Excel を見ている（置き替えの途中）
    return render_part("base.html", "part_build", pattern=pattern,
        book=book,
        grids=grids,
        remembered=remembered,
        rows=_field_view_rows(pattern, info),
        test=_test_result(pattern, book),
        confirmed_count=_confirmed_count(pattern_id),
        notes=notes or [],
        snapshot_note=SNAPSHOT_NOTE,
    )


@form_types_bp.get("/<int:pattern_id>/panel")
@form_types_bp.post("/<int:pattern_id>/panel")
def build_fragment(pattern_id: int):
    """項目の一覧とシートを出す（GET）。Excel を一緒に置くと（POST）、そのシートに置き替えて覚え直す。

    保存済みの種類を開き直したときは、登録したときに覚えたシートで同じ画面を出す（Excel は要らない）。
    別の Excel（書き方の違う同じ帳票）を置くと、シートが替わり、覚えたシートもそれに入れ替わる（種類は増えない）。
    覚える前に登録した種類（覚えたシートが無い）は、Excel を置くまでシートが出ない。
    """
    _get_pattern(pattern_id)
    book, error = _request_book(pattern_id)
    if error:
        return jsonify(error=error), 400
    message = ""
    placed = request.files.get(BOOK_FIELD)
    # 覚えたシートを入れ替えるのは、ファイルを実際に置いたときだけ。名前の無い空の部品（ファイルを選ばずに送った form）と
    # 合図（book_hash）だけの要求では、いま見ているブックが別のものでも覚えは変えない
    if request.method == "POST" and book is not None and placed is not None and placed.filename:
        remembered = pattern_book_meta(pattern_id)
        if remembered is None or remembered["file_hash"] != book.file_hash or remembered["file_name"] != book.file_name:
            try:
                _remember_book(pattern_id, book)
            except UploadError as exc:
                return jsonify(error=str(exc)), 400
        message = (f"「{book.file_name}」を読み込みました。読み取りたい欄の見出しと値をクリックしてください"
                   "（Excel のファイルは保存せず、シートの中身だけを覚えました。次からは置かずに開けます）")
    return jsonify(html=_build_html(pattern_id, book), list_html=_list_html(), message=message)


def _soften_unit_warning(f: dict) -> None:
    """この欄では値を直せないので、単位なしの警告はどこで直せるかを書く。"""
    if f.get("data_type") == "number" and str(f.get("warning") or "").startswith(_NO_UNIT_WARNING):
        f["warning"] = TEST_NO_UNIT_WARNING


def _field_view_rows(pattern: PatternDef, info) -> list[dict]:
    """項目の一覧（見出し・置いた Excel で見つかった値・セル）。値はいま登録されている設定で読み直す。

    Excel を置いていないとき（info が None）は値の欄を空にする。見出しの手直しはそれでもできる。
    """
    found: dict[str, dict] = {}
    if info is not None and pattern.fields:
        sheets = [s.sheet_name for s in pattern.sheets if s.sheet_name in info.grids] or list(info.grids)[:1]
        found = {f["field_name"]: f for f in extract_document(info, pattern, sheets)["fields"]}
    for f in found.values():
        _soften_unit_warning(f)
    rows = []
    for fd in pattern.fields:
        f = found.get(fd.field_name) or {}
        rows.append({
            "field": fd,
            "label": (fd.candidates[0] if fd.candidates else "") or "（見出しなし）",
            "value": f.get("value"),
            "warning": f.get("warning") or "",
            "sheet": f.get("sheet") or fd.sheet_name,
            "label_cell": f.get("label_cell") or fd.label_cell,
            "value_cell": f.get("value_cell") or fd.cell,
        })
    return rows


def _test_result(pattern: PatternDef, book: Book | None) -> dict | None:
    """いま置いている Excel を、この設定で読み取った結果（Markdown・見つかった件数）。"""
    if book is None or not pattern.fields:
        return None
    info = book.info
    match = match_pattern(info, pattern)
    sheets = match.sheet_names or info.sheet_names[:1]
    extraction = extract_document(info, pattern, sheets)
    doc = {"id": 0, "file_name": book.file_name, "file_hash": book.file_hash}
    return {
        "sheets": sheets,
        "found": sum(1 for f in extraction["fields"] if f["value"] not in (None, "")),
        "total": len(extraction["fields"]),
        "markdown": build_markdown(doc, extraction),
        "file_name": markdown_filename(doc, extraction),
    }


@form_types_bp.post("/<int:pattern_id>/fields")
def add_field(pattern_id: int):
    """クリックした見出しセル（と値セル）から項目を1つ作る（fetch）。

    どのセルを指しているかは、ブラウザが一緒に送ってくる Excel か、その種類が覚えているシートで確かめる
    （Excel のファイルはサーバーに置いていないため）。
    """
    pattern = _get_pattern(pattern_id)
    book, error = _request_book(pattern_id)
    if error:
        return jsonify(error=error), 400
    if book is None:
        return jsonify(error=NO_BOOK_ERROR), 400
    info = book.info
    sheet = request.form.get("sheet", "")
    grid = info.grids.get(sheet)
    if grid is None:
        return jsonify(error="シートが見つかりません。同じ帳票のExcelを置き直してください"), 400

    label_cell = request.form.get("label_cell", "")
    value_cell = request.form.get("value_cell", "")
    row, error = click_field(grid, label_cell, value_cell, {f.field_name for f in pattern.fields})
    if row is None:
        return jsonify(error=error), 400
    if any(f.sheet_name == sheet and f.label_cell == row["label_cell"] and f.cell == row["cell"]
           for f in pattern.fields):
        return jsonify(html=_build_html(pattern_id, book), list_html=_list_html(),
                       message="そのセルはもう項目になっています")

    sheet_rows, field_rows = pattern_to_rows(pattern)
    # 番号と名前を1つのセルにまとめた「使用設備」欄は、設備番号・設備名の2項目になる
    added, separated, merged = [], [], []
    for part in split_rows(row, {f.field_name for f in pattern.fields}):
        # 別の見本で書き方の違う同じ欄（「設備No」と「設備番号」）をクリックしたときは、新しい項目にせず
        # その項目の探す見出しに足す
        same = merge_target(field_rows, part, grid)
        if same is not None:
            merge_labels(same, part)
            merged.append(same["display_name"])
            continue
        # 同じ見本の別のセル（「担当者」と「報告者」）なら、辞書の名前が同じでも別の項目にする
        twin = same_sheet_field(field_rows, part, grid)
        if twin is not None:
            separate_names(twin, part)
            separated.append((part["display_name"], twin["display_name"]))
        else:
            added.append(part["display_name"])
        field_rows.append(part)
    message = "。".join(_add_messages(row, added, separated, merged))
    if not any(r["sheet_name"] == sheet for r in sheet_rows):
        sheet_rows.append({"use": True, "sheet_name": sheet})
    _save_rows(pattern, sheet_rows, field_rows)
    return jsonify(html=_build_html(pattern_id, book), list_html=_list_html(), message=message)


def _add_messages(row: dict, added: list[str], separated: list[tuple[str, str]], merged: list[str]) -> list[str]:
    """クリックの結果の知らせ（項目にした／別の項目にした／見出しに足した、のどれをしたか）。"""
    out = []
    if added:
        out.append(f"「{'」「'.join(added)}」を項目にしました")
    for name, twin in separated:
        out.append(f"「{name}」を「{twin}」とは別の項目にしました（同じ帳票の別のセルなので、両方を読み取ります）")
    if merged:
        label = (row["candidates"].splitlines() or [""])[0]
        out.append(f"「{'」「'.join(merged)}」の見出しに「{label}」を足しました"
                   "（書き方の違う同じ欄なので、1つの項目として読み取ります）")
    return out


@form_types_bp.post("/<int:pattern_id>/fields/<field_name>/delete")
def delete_field(pattern_id: int, field_name: str):
    pattern = _get_pattern(pattern_id)
    sheet_rows, field_rows = pattern_to_rows(pattern)
    rest = [r for r in field_rows if r["field_name"] != field_name]
    if len(rest) == len(field_rows):
        abort(404)
    kept = {r["sheet_name"] for r in rest if r["sheet_name"]}
    sheet_rows = [s for s in sheet_rows if not kept or s["sheet_name"] in kept]
    # 読み取る項目が無くなった使用中の種類は、使用を停止する。そのままだと帳票取り込みの候補に出て、
    # 中身の無い Markdown ができてしまう（［使用開始］も同じ決まりで断っている）
    stopped = not rest and pattern.status == "active"
    _save_rows(pattern, sheet_rows, rest, status="inactive" if stopped else None)
    message = "項目を削除しました"
    if stopped:
        message += "。読み取る項目が無くなったので、この種類の使用を停止しました（帳票取り込みの候補に出なくなります）"
    # 覚えたシート（またはブラウザが一緒に送ってきた Excel）で、シートを出したままにする
    return jsonify(html=_build_html(pattern_id, _request_book(pattern_id)[0]), list_html=_list_html(), message=message)


@form_types_bp.post("/<int:pattern_id>/fields/<field_name>/label")
def rename_field(pattern_id: int, field_name: str):
    """読み取る項目の見出しを手で直す（fetch）。

    直すのは Markdown に書き出す名前だけ。探す見出し（クリックしたときの見出しの言葉）と
    読み取るセルはそのままにする。書き出す名前は項目どうしで重ならないようにする。
    """
    pattern = _get_pattern(pattern_id)
    payload = request.get_json(force=True, silent=True) or {}
    name = " ".join(str(payload.get("name") or request.form.get("name", "")).split())
    if not name:
        return jsonify(error="見出しを入れてください"), 400
    sheet_rows, field_rows = pattern_to_rows(pattern)
    target = next((r for r in field_rows if r["field_name"] == field_name), None)
    if target is None:
        abort(404)
    if any(r["display_name"] == name for r in field_rows if r is not target):
        return jsonify(error=f"「{name}」はほかの項目が使っています。別の見出しにしてください"), 400
    # 探す見出しを持たない項目（見出しのない表など）は、書き出す名前をそのまま探していた。
    # 書き替えで探す先が変わらないよう、いまの見出しを探す見出しとして控えてから名前を変える
    if not target["candidates"] and not target["cell"]:
        target["candidates"] = target["display_name"]
    target["display_name"] = name
    target["renamed"] = True   # このあと別の欄をクリックしても、手で付けた見出しに戻さない
    _save_rows(pattern, sheet_rows, field_rows)
    return jsonify(html=_build_html(pattern_id, _request_book(pattern_id)[0]), list_html=_list_html(),
                   message=f"見出しを「{name}」にしました")


@form_types_bp.post("/<int:pattern_id>/name")
def rename(pattern_id: int):
    pattern = _get_pattern(pattern_id)
    payload = request.get_json(force=True, silent=True) or {}
    name = str(payload.get("name") or request.form.get("name", "")).strip()
    if not name:
        return jsonify(error="帳票の種類の名前を入れてください"), 400
    sheet_rows, field_rows = pattern_to_rows(pattern)
    _save_rows(pattern, sheet_rows, field_rows, name=name)
    return jsonify(ok=True, list_html=_list_html(), message="名前を変えました")


def _save_rows(pattern: PatternDef, sheet_rows: list[dict], field_rows: list[dict], name: str | None = None,
               status: str | None = None) -> None:
    """状態は変えずに保存する（使用開始は［使用開始］を押したときだけ）。タイトル項目は自動で決める。

    status を渡したときだけ状態も変える（読み取る項目が無くなったら使用を停止する）。
    """
    meta = pattern_to_meta(pattern)
    if name:
        meta["name"] = name
    meta["title_fields"] = suggest_title_fields(field_rows)
    save_pattern(rows_to_pattern(pattern.id, meta, sheet_rows, field_rows), status or pattern.status)


# ---- 使用開始・停止・削除 --------------------------------------------------------------

@form_types_bp.post("/<int:pattern_id>/status")
def change_status(pattern_id: int):
    pattern = _get_pattern(pattern_id)
    payload = request.get_json(force=True, silent=True) or {}
    status = payload.get("status") or request.form.get("status")
    # 画面から出来るのは「使用開始」だけ（「使用を停止」のボタンは無くした。利用者の指示 2026-09-22）。
    # 使わなくなった種類は削除する。項目を全部消したときだけ、delete_field が使用中を解く
    if status != "active":
        abort(400)
    if not pattern.fields:
        return jsonify(error="読み取る項目がありません。シートで見出しのセルと値のセルをクリックしてください"), 400
    set_pattern_status(pattern_id, status)
    # 残すのは設定と覚えたシートだけ（Excel のファイルはもともと置いていない）。画面はそのまま続けて使える
    message = f"「{pattern.name}」の使用を開始しました。帳票取り込みの候補に出ます"
    return jsonify(ok=True, status=status, html=_build_html(pattern_id, _request_book(pattern_id)[0]),
                   list_html=_list_html(), message=message)


@form_types_bp.post("/<int:pattern_id>/delete", endpoint="delete")
def form_types_delete(pattern_id: int):
    pattern = _get_pattern(pattern_id)
    delete_pattern(pattern_id)
    return jsonify(ok=True, list_html=_list_html(), message=f"帳票の種類「{pattern.name}」を削除しました")


# ====================================================================================================
# 元 views/tables.py
# 表の取り込み（design.md 2.4）。1画面で全部できる（利用者の指示 2026-09-20）。
#
# 画面は /tables の1枚だけ。ファイルを置く → 読み取り方 → 表の範囲 → 列の対応づけ → AI整形（任意）
# → 内容の確認 → 確定してダウンロード（zip）を、同じ画面の「段」として順に開く。
# 段の中身はこの blueprint が HTML の断片（panel）として返し、保存・実行は JSON でやりとりする。
# 画面の移動（①→②→③）は無く、URL は変わらない（static/tables.js）。
# 範囲の決まり（利用者の判断）: 期間の置き換え・投入済みとの差分・取り消しはしない。取り込みごとに
# その取り込みの記録だけから全 Markdown を作り、全ファイルを zip で渡す。クロス集計と名寄せ辞書は扱わない。
# ダウンロードした取り込みのデータは、zip を送り終えたあとに消す（design.md 3.3・core.purge.purge_after_send）。
# 読み込み・Markdown 作成・AI整形は core.jobs のジョブ（tables.pipeline / ai.runner）で動かし、
# 同じ画面の中に進み具合を出す。表の範囲・列の段は取り込みの控え（tables.source_cache）を通して読む。
# ====================================================================================================

tables_bp = Blueprint("tables", __name__, url_prefix="/tables")

EXCEL_EXT = {".xlsx", ".xlsm"}
GRID_ROWS = 60
GRID_COLS = 30
GRID_CELL_CHARS = 40
DATA_PAGE = 100
ISSUES_SHOWN = 200

# 画面で選べる役割（列の対応づけ）。ここに無い役割（人名・数値・区分など）は候補のまま使い、画面では「その他」に見せる。
# 名前はどんな表にも当てはまる言い方にする（利用者の指示 2026-09-21。entity は「設備」、log は
# 「AI整形の対象（追記ログ）」と呼んでいたが、設備の記録以外の表には当てはまらなかった）。
# 中で使う名前（key/date/entity/log/attribute）は変えない（前に保存した取り込み設定がそのまま読める）。
SCREEN_ROLES = [("key", "識別番号"), ("date", "日付"), ("entity", "対象（顧客・案件・製品など）"),
                ("log", "経過の記録（1つのセルに日付ごとに書き足した列）"), ("attribute", "その他")]
SCREEN_ROLE_KEYS = {role for role, _label in SCREEN_ROLES}
ROW_KIND_LABELS = {"header": "見出し", "data": "データ", "subtotal": "小計・合計", "note": "注記", "continuation": "継続行",
               "excluded": "除外", "title": "表題", "blank": "空行"}
TABLE_KIND_LABELS = {"list": "一覧表", "crosstab": "クロス集計", "form_like": "帳票らしい", "unknown": "不明"}
SCOPE_LABELS = {"pending": "まだ整形していない行と、内容が変わった行", "errors": "エラーになった行だけ",
                "flagged": "要確認の行だけ", "all": "すべての行をやり直す"}
# CSV の文字コード・区切り文字は画面の選択肢だけを受け付ける（他の値は読み込みで落ちて画面が開けなくなるため）
# ダウンロードでデータが消えることの案内（design.md 3.3）
TABLES_DELETE_ON_DOWNLOAD_NOTE = ("zip をダウンロードすると、この取り込みのデータはサーバーから消えます"
                           "（もう一度ダウンロードすることはできません）。")
TABLES_DELETE_ON_DOWNLOAD_CONFIRM = ("ダウンロードすると、この取り込みのデータはサーバーから消えます。"
                              "もう一度ダウンロードすることはできません。")
ENCODING_CHOICES = [("utf-8-sig", "UTF-8（BOM付き）"), ("utf-8", "UTF-8"), ("cp932", "CP932（Shift_JIS）"),
                    ("euc_jp", "EUC-JP"), ("iso2022_jp", "ISO-2022-JP（JISコード）"),
                    ("shift_jis_2004", "Shift_JIS 2004"), ("utf-16", "UTF-16"),
                    ("utf-16-le", "UTF-16LE（BOMなし）"), ("utf-16-be", "UTF-16BE（BOMなし）")]
BUSY_MESSAGE = "処理中は変更できません。終わるか中止してから変更してください"
DELIMITER_CHOICES = [(",", "カンマ"), ("\t", "タブ"), (";", "セミコロン"), ("|", "縦棒")]


# ---- 共通 ---------------------------------------------------------------------------------

def _load_import(import_id: int) -> dict:
    """この画面（このブラウザ）の取り込みを返す。ほかの人の取り込みは「無い」として扱う（404）。

    社内LANで数人が同時に使うので、番号を打ち替えただけでほかの人の表を読めてはいけない。
    403 にすると「その番号の取り込みはある」ことが分かってしまうので、404 にそろえる。
    """
    imp = get_import(import_id)
    if imp is None or not owns(imp):
        abort(404)
    return imp


def _spec_for(imp: dict):
    return spec_for_import(imp)


def _is_excel(imp: dict) -> bool:
    return Path(imp["file_name"]).suffix.lower() in EXCEL_EXT


def _open(imp: dict):
    """控え付きの表ソース。元のファイルは控えにない範囲を読むときだけ開く。"""
    return import_source(imp)


def _sheet(imp: dict, source) -> str:
    if not _is_excel(imp):
        return source.file_name  # CSV はファイル名が1つのシート（全行を数えずに済ませる）
    src = imp.get("source") or {}
    names = [s.name for s in source.sheets()]
    if src.get("sheet") in names:
        return src["sheet"]
    visible = [s.name for s in source.sheets() if not s.hidden]
    if not (visible or names):
        raise UploadError("シートがないブックです。シートのあるブックを選んでください")
    return (visible or names)[0]


def _layout(imp: dict, source, sheet: str, header_rows=None, data_end=None, use_saved=True):
    src = imp.get("source") or {}
    if use_saved:
        header_rows = header_rows or [int(r) for r in src.get("header_rows") or [] if r] or None
        data_end = data_end or src.get("data_end_row") or None
    return source.layout(sheet, header_rows=header_rows or None, data_end=data_end or None)


def _after_read_panel(spec) -> str:
    """読み込みのあとに開く段。経過の記録の列があれば AI整形、なければ内容の確認。"""
    return "ai" if spec is not None and spec.log_stage is not None else "preview"


def _json_error(message: str, status: int = 400, **extra):
    return jsonify({"error": message, **extra}), status


def _tables_payload() -> dict:
    data = request.get_json(silent=True)
    return data if isinstance(data, dict) else {}


def _int(value, default=None):
    try:
        return int(str(value).strip())
    except (TypeError, ValueError):
        return default


# 行番号の範囲「3-4」を広げる上限（見出し行の帯なので広くはならない）。static/tables.js の parseRows と同じ規則
_MAX_ROW_SPAN = 10


def _row_no(value) -> int | None:
    """行番号の入力（全角数字も可）。数字だけのときだけ読む（static/tables.js の parseEnd と同じ規則）。"""
    text = unicodedata.normalize("NFKC", str(value if value is not None else "")).strip()
    return int(text) if text.isascii() and text.isdigit() and int(text) > 0 else None


def _int_list(value) -> list[int]:
    """見出し行の入力（「1,2」「1、2」「1 2」「１，２」「1-2」）。static/tables.js の parseRows と同じ規則で読む。"""
    if isinstance(value, (list, tuple)):
        items = [str(x) for x in value]
    else:
        items = re.split(r"[,、\s]+", unicodedata.normalize("NFKC", str(value or "")))
    out: set[int] = set()
    for item in items:
        m = re.fullmatch(r"(\d+)-(\d+)", item.strip())
        if m and item.isascii():
            a, b = int(m.group(1)), int(m.group(2))
            if 0 < a <= b and b - a < _MAX_ROW_SPAN:
                out.update(range(a, b + 1))
            continue
        n = _row_no(item)
        if n:
            out.add(n)
    return sorted(out)


def _ai_running(import_id: int) -> bool:
    job = latest_job("table_import", import_id, kind="ai_format")
    return bool(job and not job.get("finished"))


# AIの試し実行（同期の要求でジョブが無い）をしている取り込み → 実行中の数。実行中に取り込みを消すと、
# 終わった試し実行が AI の結果・応答を書き戻して残してしまう（design.md 3.3）
_TRIALS: Counter[int] = Counter()
_TRIALS_LOCK = threading.Lock()
TRIAL_BUSY_MESSAGE = "AIの試し実行中は削除・ダウンロード・保存できません。試し実行が終わってからもう一度押してください"


def _trial_running(import_id: int) -> bool:
    with _TRIALS_LOCK:
        return _TRIALS[import_id] > 0


def _processing(imp: dict) -> bool:
    """読み込み中・Markdown作成中・AI整形の実行中か。途中で設定やデータを変えると、古い設定の結果が残ったり、
    消したはずの AI の結果が書き戻されたりする（design.md 3.3）。"""
    return imp.get("status") in ("reading", "confirming") or _ai_running(imp["id"])


# ---- 1画面のやりとり（段の URL と、段の中身） ------------------------------------------------------

def _urls(import_id: int) -> dict:
    """画面（static/tables.js）が使う URL。panel は末尾の NAME を段の名前に置き換えて使う。

    列の対応づけ・AI整形・プレビュー・ダウンロードの URL は、それぞれの段の HTML が
    data-* 属性や <a href> で持っているので、ここには入れない。
    """
    def u(endpoint: str, **kw) -> str:
        return url_for(endpoint, import_id=import_id, **kw)

    return {
        "panel": url_for("tables.panel", import_id=import_id, name="NAME"),
        "source": u("tables.save_source"), "detect": u("tables.layout_detect"), "layout": u("tables.save_layout"),
        "read": u("tables.reread"), "cancel": u("tables.cancel_job"), "delete": u("tables.delete_import"),
        "preview_start": u("tables.start_preview"), "confirm": u("tables.confirm"),
    }


def _job_info(job: dict | None, import_id: int | None = None) -> dict | None:
    if job is None:
        return None
    return {"id": job["id"], "kind": job.get("kind"), "status": job.get("status"),
            "finished": bool(job.get("finished")), "message": job.get("message") or "",
            "progress": job.get("progress") or {},
            "url": url_for("tables.api_job", job_id=job["id"]),
            "cancel_url": url_for("tables.cancel_job", import_id=import_id) if import_id else None}


def _running_job(imp: dict) -> dict | None:
    """読み込み中・確定処理中なら、その進み具合。終わっているのに状態が残っていれば直す。"""
    if imp.get("status") not in ("reading", "confirming"):
        return None
    job = get_job(imp["job_id"]) if imp.get("job_id") else None
    if job is None or job.get("finished"):
        fresh = get_import(imp["id"])
        if fresh and fresh.get("status") in ("reading", "confirming"):
            if fresh["status"] == "reading":
                stats = {**(fresh.get("stats") or {}), "error": "読み込みが途中で止まりました。もう一度読み込んでください"}
                update_import(imp["id"], status="failed", stats=stats)
            else:
                update_import(imp["id"], status="preview")
        return None
    return job


def _panel(html: str, **extra):
    return jsonify({"html": html, **extra})


def _locked(reason: str, **extra):
    """まだ使えない段（灰色で1行だけ理由を出す）。"""
    return jsonify({"html": "", "locked": reason, **extra})


@tables_bp.get("/")
@tables_bp.get("/new")
def new():
    """表の取り込みの画面（1枚）。段の中身はここでは出さず、ファイルを置いたあとに取りに来る。"""
    current_session_id()   # 画面を開いた時点で作業場所（クッキー）を決めておく（同時に置かれても取り違えない）
    return render_template("base.html", screen="tables")


@tables_bp.get("/imports/<int:import_id>/panel/<name>")
def panel(import_id: int, name: str):
    imp = _load_import(import_id)
    handler = _PANELS.get(name)
    if handler is None:
        abort(404)
    return handler(imp)


# ---- ファイルを置く ------------------------------------------------------------------------

@tables_bp.post("/upload", endpoint="upload")
def tables_upload():
    cfg = current_app.config
    try:
        stored = save_upload(request.files.get("file"), "tables", set(cfg["TABLE_ALLOWED_EXTENSIONS"]),
                             int(cfg["TABLE_MAX_UPLOAD_BYTES"]))
    except UploadError as exc:
        return _json_error(str(exc))
    path = upload_path(stored.stored_path)
    source_info: dict = {}
    try:
        if Path(stored.file_name).suffix.lower() in EXCEL_EXT:
            precheck_excel(path)
            source = open_source(path, stored.file_name, {"max_cells": cfg.get("EXCEL_MAX_CELLS", 500000)})
            sheets = [s for s in source.sheets() if not s.hidden] or source.sheets()
            if not sheets:
                raise UploadError("シートがないブックです。シートのあるブックを選んでください")
            # 開けても値を読めないブック（<v>NaN</v> など）は、取り込みを作る前にここで断る。
            # 先頭だけでなく全行を読む（下の方の値が読めないと、取り込みが読み込みの段から先へ進めない）
            for _row in source.rows(sheets[0].name):
                pass
            source_info = {"kind": "excel", "sheet": sheets[0].name}
        else:
            source = open_source(path, stored.file_name)
            sniff = source.sniff
            source_info = {"kind": "csv", "encoding": source.encoding, "delimiter": source.delimiter,
                           "errors": "strict"}
            if sniff is not None:
                source_info.update(bom=sniff.bom, preamble_rows=sniff.preamble_rows, sniff_warnings=sniff.warnings,
                                   confidence=sniff.confidence)
    except UploadError as exc:
        remove_upload(stored.stored_path)
        return _json_error(str(exc))
    except Exception:
        remove_upload(stored.stored_path)   # 思わぬエラーでもアップロードしたファイルを残さない（design.md 3.3）
        raise
    try:
        # DB への登録も同じ try に入れる。ここで落ちると（別の人の VACUUM 中の database is locked、
        # ディスク満杯など）、置いたファイル（最大200MB）だけが行の無い孤児として残り、
        # 再起動するまで消えなかった（2026-09-23 のレビュー）
        import_id = create_import(stored.file_name, stored.file_hash, stored.stored_path, source=source_info,
                                  session_id=current_session_id())
        if source_info.get("kind") == "excel":
            import_source(get_import(import_id), real=source).sheets()  # シート一覧を控えに入れる
    except Exception:
        remove_upload(stored.stored_path)
        raise
    return jsonify({"import_id": import_id, "file_name": stored.file_name, "urls": _urls(import_id)})


# ---- 画面を離れたので捨てる ------------------------------------------------------------------
# 「その場でダウンロードしない限り、その場ですぐ捨てる」（利用者の指示 2026-09-20）。
# 画面を閉じた・隠したときに static/app.js の ragDiscard がここへ「捨てて」と送ってくる。
# navigator.sendBeacon で届くので中身の型は text/plain（get_json(force=True) で読む）。
# 応答は読めず、やり直しもできないので、いつでも 204 を返す（もう無い番号・ほかの人の番号・
# 処理中のものは core.purge 側で黙って外れる）。

@tables_bp.post("/discard", endpoint="discard")
def tables_discard():
    """この画面（このブラウザ）の、まだダウンロードしていない取り込みを捨てる。"""
    payload = request.get_json(force=True, silent=True) or {}
    ids = payload.get("import_ids") or []
    sid = current_session_id()
    try:
        if ids:
            discard_table_imports(ids, sid)
        else:
            purge_session(sid, documents=False)   # 帳票取り込み（別のタブ）は巻き込まない
    except Exception as exc:   # 捨て損ねてもブラウザには伝えられない。時間切れの片付けに任せる
        current_app.logger.warning("取り込みの片付けに失敗しました: %s", exc.__class__.__name__)
    return "", 204


# ---- 読み取り方（文字コード・区切り・シート） ---------------------------------------------------------

def _source_note(imp: dict, src: dict) -> str:
    """段の見出しに出す1行（どう読むか）。"""
    if _is_excel(imp):
        return f"シート: {src.get('sheet') or ''}"
    delimiters = dict(DELIMITER_CHOICES)
    return f"{src.get('encoding') or ''}／{delimiters.get(src.get('delimiter'), src.get('delimiter') or '')}"


def _panel_source(imp: dict):
    import_id = imp["id"]
    src = imp.get("source") or {}
    error = None
    sheets, sheet, list_like = [], None, {}
    try:
        source_obj = _open(imp)
        sheets = source_obj.sheets() if _is_excel(imp) else []
        sheet = _sheet(imp, source_obj)
        if _is_excel(imp) and len([s for s in sheets if not s.hidden]) <= 8:
            # 「一覧表らしい」「クロス集計らしい（対応していません）」を出し分ける（表の範囲の判定と合わせる）
            list_like = {s.name: source_obj.list_kind(s.name) for s in sheets if not s.hidden}
    except UploadError as exc:
        error = str(exc)
    html = render_part("base.html", "part_source", imp=imp, src=src, error=error, sheets=sheets, sheet=sheet, list_like=list_like,
        encodings=ENCODING_CHOICES, delimiters=DELIMITER_CHOICES)
    note = _source_note(imp, {**src, "sheet": sheet or src.get("sheet")})
    return _panel(html, note=note, error=error, import_id=import_id)


@tables_bp.post("/imports/<int:import_id>/source")
def save_source(import_id: int):
    """読み取り方（シート・文字コード・区切り）。変えるたびに保存する。"""
    imp = _load_import(import_id)
    if _processing(imp):
        return _json_error(BUSY_MESSAGE, 409)
    form = _tables_payload() or request.form
    src = dict(imp.get("source") or {})
    before = (src.get("sheet"), src.get("encoding"), src.get("delimiter"), src.get("errors"))
    if _is_excel(imp):
        if form.get("sheet"):
            src["sheet"] = form["sheet"]
    else:
        allowed = {"encoding": {v for v, _label in ENCODING_CHOICES},
                   "delimiter": {v for v, _label in DELIMITER_CHOICES}}
        for name in ("encoding", "delimiter"):
            value = form.get(name)
            if value and value not in allowed[name]:
                return _json_error("この画面にない文字コード・区切り文字は選べません")
            if value:
                src[name] = value
        src["errors"] = "replace" if form.get("replace_errors") in ("on", True, "true", "1") else "strict"
    columns: dict = {"source": src, "status": "uploaded"}   # 読み取り方を保存したら読み込みからやり直す
    if (src.get("sheet"), src.get("encoding"), src.get("delimiter"), src.get("errors")) != before:
        src.pop("header_rows", None)
        src.pop("data_end_row", None)
        # 置いたときの判定（文字コード・区切りの見立て）で出した注意書きと「前置き行」は、
        # 読み方を変えたら合わなくなる。捨てないと、文字コードを直したあとも
        # 「UTF-8 の行が混ざっています」が出続けて消せなかった（2026-09-23 のレビュー）
        # 捨てるのは置いたときの見立ての注意書きだけ。前置き行の数（preamble_rows）は
        # ③「表の範囲」が使う値なので残す（捨てると 0行になって実際と食い違う）
        for stale in ("sniff_warnings", "confidence"):
            src.pop(stale, None)
        # 見出しが変わるので、この取り込みの列の対応づけは捨てて決め直す
        columns.update(spec_json="", spec_hash="")
    update_import(import_id, **columns)
    return jsonify({"ok": True, "note": _source_note(imp, src), "next": "layout",
                    "reset": ["layout", "columns", "ai", "preview", "done"]})


# ---- 表の範囲（見出し行・データの終わり） ----------------------------------------------------------

def _grid(source, sheet: str, layout) -> tuple[list[dict], int]:
    classes = {rc.index: rc for rc in layout.row_classes}
    rows, width = [], 0
    for row in source.rows(sheet, 1, GRID_ROWS):
        cells = [(c.text or "")[:GRID_CELL_CHARS] for c in row.cells[:GRID_COLS]]
        while cells and not cells[-1]:
            cells.pop()
        width = max(width, len(cells))
        rc = classes.get(row.index)
        rows.append({"index": row.index, "cells": cells, "kind": rc.kind if rc else ("blank" if row.is_blank else "title"),
                     "reason": rc.reason if rc else "", "hidden": bool(row.hidden),
                     "strike": any(getattr(c, "strike", False) for c in row.cells if c.text)})
    return rows, max(width, 1)


def _layout_json(layout) -> dict:
    return {
        "sheet": layout.sheet, "table_kind": layout.table_kind,
        "table_kind_label": TABLE_KIND_LABELS.get(layout.table_kind, layout.table_kind),
        "header_rows": layout.header_rows, "data_start": layout.data_start, "data_end": layout.data_end,
        "headers": layout.headers, "warnings": layout.warnings, "errors": list(getattr(layout, "errors", [])),
        "counts": {ROW_KIND_LABELS.get(k, k): v for k, v in (layout.counts or {}).items()},
        "rows": {str(rc.index): {"kind": rc.kind, "reason": rc.reason} for rc in layout.row_classes},
    }


def _panel_layout(imp: dict):
    job = _running_job(imp)
    if job is not None and job.get("kind") == "table_read":
        return _panel("", job=_job_info(job, imp["id"]), reading=True)
    try:
        source_obj = _open(imp)
        sheet = _sheet(imp, source_obj)
        guess = _layout(imp, source_obj, sheet)
        rows, width = _grid(source_obj, sheet, guess)
    except UploadError as exc:
        return _locked(str(exc))
    src = imp.get("source") or {}
    html = render_part("base.html", "part_layout", imp=imp, layout=guess, info=_layout_json(guess), rows=rows,
                           width=width, sheet=sheet, data_end_saved=src.get("data_end_row") or "")
    return _panel(html, note=f"{guess.data_start}〜{guess.data_end}行目" if guess.data_end >= guess.data_start else "")


@tables_bp.post("/imports/<int:import_id>/layout/detect")
def layout_detect(import_id: int):
    imp = _load_import(import_id)
    data = _tables_payload()
    try:
        source_obj = _open(imp)
        sheet = _sheet(imp, source_obj)
        guess = _layout(imp, source_obj, sheet, header_rows=_int_list(data.get("header_rows")) or None,
                        data_end=_row_no(data.get("data_end")), use_saved=False)
    except UploadError as exc:
        return _json_error(str(exc))
    return jsonify(_layout_json(guess))


@tables_bp.post("/imports/<int:import_id>/layout")
def save_layout(import_id: int):
    """この範囲で読み込む。設定が無い・見出しが変わったときは列の対応づけへ、そうでなければ読み込みを始める。"""
    imp = _load_import(import_id)
    if _processing(imp):
        return _json_error(BUSY_MESSAGE, 409)
    data = _tables_payload() or request.form
    header_rows = _int_list(data.get("header_rows"))
    data_end = _row_no(data.get("data_end_row"))
    try:
        source_obj = _open(imp)
        sheet = _sheet(imp, source_obj)
        guess = _layout(imp, source_obj, sheet, header_rows=header_rows or None, data_end=data_end, use_saved=False)
    except UploadError as exc:
        return _json_error(str(exc))
    if guess.table_kind == "crosstab":
        return _json_error("月別集計のようなクロス集計の表には対応していません。1行＝1件の一覧表を選んでください")
    if guess.errors:
        return _json_error(guess.errors[0])
    if not guess.headers:
        return _json_error("見出し行とデータの範囲が見つかりません。見出し行の番号を指定してください")
    if guess.data_end < guess.data_start:
        # 見出しは見つかっている。直すのは見出し行ではなくデータの範囲なので、そう伝える
        last_header = guess.header_rows[-1] if guess.header_rows else 0
        if data_end:
            return _json_error(f"データの終わりに{data_end}行目を指定すると、データの行が1行もありません"
                               f"（見出しは{last_header}行目）。空欄にするか、見出し行より下の行を指定してください")
        return _json_error(f"見出し行（{last_header}行目）より下にデータの行が1行もありません。"
                           "見出し行の番号を確かめるか、データのある表を選んでください")
    src = dict(imp.get("source") or {})
    src.update(sheet=sheet, header_rows=guess.header_rows, header_row=None, data_end_row=data_end or None)
    update_import(import_id, source=src)
    # 取り込み設定は保存しないので、範囲を決めたら必ず「列の対応づけ」へ進む（利用者の指示 2026-09-20）
    return jsonify({"ok": True, "next": "columns", "reset": ["columns", "ai", "preview", "done"]})


# ---- 列の対応づけ ------------------------------------------------------------------------------
# 画面で決めるのは「使う・役割」だけ。キー・型・単位・md での扱い・空欄＝上と同じは、見出しと値から
# 候補づくり（tables.mapping.suggest_columns・tables.dictionary）が決める（利用者の指示 2026-09-20）。


def _screen_role(role: str) -> str:
    """候補の役割を画面の選択肢に寄せる（画面に出さない役割＝人名・数値・区分などは「その他」）。"""
    return role if role in SCREEN_ROLE_KEYS else "attribute"


def _suggest(imp: dict):
    """保存した範囲で見出しと先頭のデータを読み、列ごとの候補を作る。戻り値: (layout, suggestions)"""
    source_obj = _open(imp)
    sheet = _sheet(imp, source_obj)
    guess = _layout(imp, source_obj, sheet)
    samples = source_obj.sample_rows(sheet, guess, 200)
    return guess, suggest_columns(guess.headers, samples)


def _unused_note(sugg) -> str:
    """はじめから「使わない」にしてある列の、その理由（見出しの下に小さく出す。ふつうの列は空）。

    理由が読めないと、チェックの外れている列を入れ直してよいのか分からない。
    「知らせ」の列は 2026-09-21 にやめた（ほとんどの行で空で、表の上の行と同じことを書いていた）。
    型エラー・空欄の割合は表の上（_columns_todo）に出すので、ここには書かない。
    「値は「日付」らしい」も出さない（日付の役割は日付として読める列にしか出ないので、読んでも直せない）。
    """
    if sugg.md != "omit":
        return ""
    return ("空欄だけなので、はじめから使わない設定にしています" if sugg.omit_reason == "blank"
            else "記録に不要な管理用の列らしいので、はじめから使わない設定にしています")


def _column_row(sugg) -> dict:
    """画面の1行（使う・見出し・役割・値の例と、使わない設定にしている理由）。"""
    return {"index": sugg.index, "header": sugg.header, "use": sugg.md != "omit",
            "role": _screen_role(sugg.role), "type": sugg.type, "examples": list(sugg.examples or [])[:3],
            "note": _unused_note(sugg)}


DATE_TYPES = ("date", "datetime")


def _roles_for(row: dict) -> list[tuple[str, str]]:
    """その列で選べる役割。

    「日付」は値が日付として読める列にだけ出す（日付として読めない列を日付にすると保存が通らず、
    直しようのない行き止まりになる）。年月だけの列（2025-03・202503）は日付として読めるので出る
    ＝ その月のファイルに分かれる。
    """
    return [(role, label) for role, label in SCREEN_ROLES
            if role != "date" or row["type"] in DATE_TYPES or row["role"] == "date"]


def _has_date_column(pairs) -> bool:
    """日付として読める列があるか（1つも無ければ「日付の列を選べ」とは言わない）。"""
    return any(sugg.type in DATE_TYPES for _row, sugg in pairs)


# 「列の対応づけは決まっている」の決まり（利用者の問い 2026-09-20「列の対応付けを行う意味は？」）。
# この段で決められるのは「出す／出さない」と四つの役割だけで、ふつうの一覧表ならそのどちらも
# 候補づくり（tables.mapping）が見出しと値から決めている。決めることが無いのに22行の表を出すのは
# 意味がないので、次の条件をすべて満たすときは表をたたんで要約1行だけ出す。
#   1. 識別番号の列がちょうど1つで、見出しが辞書と完全一致している（matched_by == "dictionary"）
#   2. 日付の列がちょうど1つで、同じく完全一致している
#   3. 対象の列は0か1つ。1つなら完全一致している（0なら要約に「対象の列はありません」と書く）
#   4. 経過の記録の列は0か1つ。1つなら完全一致している（0なら要約にそう書く）
#   5. 出す列のどれにも、出すかどうかを決め直す理由が無い
#      ＝ 読めない値がある（type_error_rate > 0）／ほとんど空欄（blank_rate >= UNSURE_BLANK_RATE）
# 似た語で当たっただけ（matched_by == "similar"）や、値の並びから当てた（"none"）列が四つの役割に
# 付いていると 1〜4 で外れる。役割が本当に合っているかは人にしか決められないので、表を開く。
# 半分くらい空欄なのは決め直す理由にしない（出しても困らないため。ここに出さなければ画面のどこにも出ない）。
UNSURE_BLANK_RATE = 0.9

# 四つの役割（画面で決められるもの）と、要約・上の行に出す短い名前（プルダウンの但し書きまでは書かない）。
# 識別番号と日付は無いと決まらない
_DECIDED_ROLES = [("key", "識別番号", True), ("date", "日付", True),
                  ("entity", "対象", False), ("log", "経過の記録", False)]


def _role_columns(pairs, role: str) -> list[tuple[dict, object]]:
    return [pair for pair in pairs if pair[0]["use"] and pair[0]["role"] == role]


def _columns_todo(pairs) -> list[str]:
    """表を開いて決めてもらうことを並べる（空なら要約だけでよい）。pairs: [(画面の1行, 候補)]"""
    todo: list[str] = []
    for role, label, required in _DECIDED_ROLES:
        found = _role_columns(pairs, role)
        if len(found) > 1:
            todo.append(f"{label}の列が{len(found)}つあります。1つにしてください")
        elif not found:
            # 日付として読める列が1つも無い表では、選びようがないので求めない
            if required and (role != "date" or _has_date_column(pairs)):
                todo.append(f"{label}の列が決まっていません。1つ選んでください")
        elif found[0][1].matched_by != "dictionary":
            todo.append(f"「{found[0][0]['header']}」を{label}として読み取ります。これでよいか確かめてください")
    for row, sugg in pairs:
        if not row["use"]:
            continue
        if sugg.type_error_rate:
            todo.append(f"列「{row['header']}」に読み取れない値があります"
                        f"（{round(sugg.type_error_rate * 100, 1)}%）。出すかどうか決めてください")
        elif sugg.blank_rate >= UNSURE_BLANK_RATE:
            todo.append(f"列「{row['header']}」はほとんど空欄です"
                        f"（{int(round(sugg.blank_rate * 100))}%）。出すかどうか決めてください")
    return todo


def _columns_summary(pairs) -> str:
    """決まっているときに出す1行。見つけた役割と、出す列・出さない列の数を正直に書く。"""
    named, missing = [], []
    for role, label, _required in _DECIDED_ROLES:
        found = _role_columns(pairs, role)
        if found:
            named.append(f"{found[0][0]['header']}＝{label}")
        elif role == "date":
            missing.append("日付の列はありません。記録は「日付なし」の1ファイルにまとめます。")
        else:
            missing.append(f"{label}の列はありません。")
    used = sum(1 for row, _s in pairs if row["use"])
    left = len(pairs) - used
    return ("、".join(named) + "として読み取ります。" + "".join(missing)
            + f"{len(pairs)}列のうち{used}列を Markdown に出します"
            + (f"（残る{left}列は出しません）。" if left else "（出さない列はありません）。"))


def _default_table_name(imp: dict) -> str:
    """表の名前の初期値。ファイル名から作る（設定は残らないので、名前もこの取り込みだけのもの）。"""
    return Path(imp["file_name"]).stem


def _build_spec(payload: dict, guess, suggestions):
    """画面の選択（使う・役割）と列の候補から取り込み設定を作る。戻り値: (spec, errors)"""
    choices: dict[int, dict] = {}
    for row in payload.get("columns") or []:
        if isinstance(row, dict):
            choices[_int(row.get("index"), -1)] = row
    used: list[dict] = []
    date_errors: list[str] = []
    for sugg in suggestions:
        choice = choices.get(sugg.index)
        if not (bool(choice.get("use")) if choice else sugg.md != "omit"):
            continue
        role = str((choice or {}).get("role") or "")
        if role not in SCREEN_ROLE_KEYS:
            role = _screen_role(sugg.role)
        # 画面でそのままなら候補の役割（人名・数値・区分など）を活かす。選び直したときだけその役割にする
        role = sugg.role if role == _screen_role(sugg.role) else role
        if role == "date" and sugg.type not in DATE_TYPES:
            # 画面には出さない選び方だが、古い画面から送られたときに「型を日付にしてください」
            # （利用者には直せない指示）ではなく、何が起きているかを返す
            date_errors.append(f"列「{sugg.header}」の値は日付として読めないので、日付の列にはできません")
        type_ = "text" if role == "log" else sugg.type
        # 「使う」列は必ず md に出す（候補が「出さない」でも、チェックを入れたのだから出す）
        md = sugg.md if sugg.md != "omit" else ("body" if type_ == "text" else "attribute")
        used.append({"index": sugg.index, "header": sugg.header, "key": sugg.key, "display": sugg.display,
                     "type": type_, "role": role, "unit": sugg.unit, "md": md,
                     "fill_down_blank": bool(sugg.fill_down_blank)})
    name = str(payload.get("name") or "").strip()
    header_rows = list(getattr(guess, "header_rows", None) or [1])
    spec = spec_from_suggestions(name, {"table_kind": "list", "header_rows": header_rows}, used)
    errors = validate_spec(spec)
    if date_errors:
        # 「型を日付にしてください」は画面に直す場所が無いので、こちらの言い方に置き換える
        errors = date_errors + [e for e in errors if not e.startswith("日付の列「")]
    if sum(1 for u in used if u["role"] == "log") > 1:
        # 黙って最初の列だけを経過の記録にしない（2列目は日付ごとに分けられず1行につながれて出てしまう）
        errors.insert(0, "経過の記録の列は1つだけにしてください")
    return spec, errors


def _panel_columns(imp: dict):
    import_id = imp["id"]
    job = _running_job(imp)
    if job is not None and job.get("kind") == "table_read":
        return _panel("", job=_job_info(job, import_id), reading=True)
    try:
        guess, suggestions = _suggest(imp)
    except UploadError as exc:
        return _locked(str(exc))
    if not guess.headers:
        return _locked("見出しが見つかりません。上の「表の範囲」で見出し行を指定してください")
    rows = [_column_row(s) for s in suggestions]
    spec = _spec_for(imp)
    if spec is not None:
        # この取り込みで一度保存していれば、そのときの「使う・役割」を残す（開き直しても選び直さずに済む）
        by_header = {}
        for col in spec.columns:
            for header in (col.headers or [col.display]):
                by_header[header] = col
        for row in rows:
            col = by_header.get(row["header"])
            row["use"] = col is not None
            if col is not None:
                row["role"] = _screen_role(col.role)
    for row in rows:
        row["roles"] = _roles_for(row)
    # 決めることが無ければ表をたたんで要約1行にする（表は隠すだけで残すので、保存で送る中身は同じ）
    pairs = list(zip(rows, suggestions))
    todo = _columns_todo(pairs)
    html = render_part("base.html", "part_columns", rows=rows, todo=todo,
        summary="" if todo else _columns_summary(pairs),
        name=spec.name if spec is not None else _default_table_name(imp),
        save_url=url_for("tables.save_columns", import_id=import_id))
    return _panel(html, note=f"{sum(1 for r in rows if r['use'])}／{len(rows)}列")


@tables_bp.post("/imports/<int:import_id>/columns")
def save_columns(import_id: int):
    imp = _load_import(import_id)
    if _processing(imp):
        return _json_error(BUSY_MESSAGE, 409)
    try:
        guess, suggestions = _suggest(imp)
    except UploadError as exc:
        return _json_error(str(exc))
    spec, errors = _build_spec(_tables_payload(), guess, suggestions)
    if errors:
        return _json_error(errors[0], errors=errors)
    save_spec(import_id, spec)
    job_id = start_read_job(import_id)
    return jsonify({"ok": True, "next": _after_read_panel(spec), "reset": ["ai", "preview", "done"],
                    "job": _job_info(get_job(job_id), import_id), "reading": True})


@tables_bp.post("/imports/<int:import_id>/read")
def reread(import_id: int):
    """読み込み直し（失敗したとき・もう一度読むとき）。"""
    imp = _load_import(import_id)
    spec = _spec_for(imp)
    if spec is None:
        return _json_error("先に列の対応づけを保存してください")
    if _processing(imp):
        # AI整形の実行中・一時停止中に読み込み直すと、読み込みが AI整形の後ろで待ち続ける
        return _json_error(BUSY_MESSAGE, 409)
    job_id = start_read_job(import_id)
    return jsonify({"ok": True, "next": _after_read_panel(spec), "reset": ["ai", "preview", "done"],
                    "job": _job_info(get_job(job_id), import_id), "reading": True})


# 中止できるジョブ（取り込みの job_id に入るもの）と、中止したときに戻す状態
CANCELLABLE_JOBS = {"table_read": "uploaded", "table_render": "preview"}


@tables_bp.post("/imports/<int:import_id>/cancel")
def cancel_job(import_id: int):
    """［中止］: 読み込み・Markdown作成のジョブを止める。"""
    imp = _load_import(import_id)
    job = get_job(imp["job_id"]) if imp.get("job_id") else None
    if job is None or job["kind"] not in CANCELLABLE_JOBS or job.get("finished"):
        return _json_error("中止できる処理がありません（すでに終わっている可能性があります）")
    request_cancel(job["id"])
    after = get_job(job["id"])
    # 取り込みの状態は読み直す（この間にジョブが成功して preview/confirmed を書いていたら、巻き戻してはいけない）
    fresh = get_import(import_id)
    if (after and after["status"] == "cancelled" and fresh is not None
            and fresh["status"] in ("reading", "confirming")):
        # 待機中のまま中止されたときは本体が動かないので、ここで取り込みの状態を戻す
        update_import(import_id, status=CANCELLABLE_JOBS[job["kind"]])
    return jsonify({"ok": True, "message": "処理を中止しました"})


@tables_bp.post("/imports/<int:import_id>/delete")
def delete_import(import_id: int):
    """間違えて取り込んだファイルを消す（読み込んだ内容・作った Markdown も消える）。"""
    imp = _load_import(import_id)
    if imp["status"] in ("reading", "confirming"):
        return _json_error("処理中の取り込みは削除できません。終わるか中止してから削除してください")
    if _ai_running(import_id):
        return _json_error("AI整形の実行中は削除できません。AI整形を中止してから削除してください")
    if _trial_running(import_id):
        return _json_error(TRIAL_BUSY_MESSAGE)
    preview_job = latest_job("table_import", import_id, kind="table_preview")
    if preview_job is not None and not preview_job.get("finished"):
        # 下書きの作成中に消すと、開いているファイルが消え残ったり、消したあとに記録の md が書き戻されたりする
        return _json_error("Markdownの下書きを作っている間は削除できません。終わってから削除してください")
    purge_table_import(import_id)
    # ファイル名は出さない: 取引先名や「社外秘」を含みうる名前を画面に残さない（design.md 3.3）
    message = "取り込みを削除しました（元のファイル・読み込んだ内容・作った Markdown を消しました）"
    if purge_incomplete():
        # 使用中などで消し切れなかったファイルは中身を 0 バイトにしてある（次の起動時に片付く。design.md 3.3）
        message = "取り込みを削除しました（一部のファイルは使用中で消し切れず、中身を空にしました。次の起動時に片付きます）"
    return jsonify({"ok": True, "message": message, "reload": True})


def _not_ready_reason(imp: dict, spec) -> str:
    """読み込みが終わっていない・失敗したときに、段に出す1行。"""
    if spec is None:
        return "先に「列の対応づけ」を保存してください"
    if imp["status"] == "failed":
        return (imp.get("stats") or {}).get("error") or "表を読み込めませんでした"
    if imp["status"] == "uploaded":
        return "先に「表の範囲」で読み込んでください"
    return ""


# ---- AI整形（任意。経過の記録の列があるときだけ） --------------------------------------------------

def _ai_job(import_id: int) -> dict | None:
    return latest_job("table_import", import_id, kind="ai_format")


def _ai_connection_ctx() -> dict:
    """AI整形の段に要る AI接続の状態（設定そのものはヘッダーの「AI接続」。ai_header_ctx）。"""
    from app import ai

    ready = ai.is_configured()
    external, model = False, ""
    if ready:
        try:
            settings = ai.job_client_settings()
            external = not ai.is_local_endpoint(settings.get("chat_url") or "")
            model = settings.get("model") or ""
        except Exception:
            ready = False
    return {"ai_ready": ready, "external": external, "model": model, "ai_state": ai.connection_status()}


def _panel_ai(imp: dict):
    import_id = imp["id"]
    job = _running_job(imp)
    if job is not None:
        return _panel("", job=_job_info(job, import_id), reading=True)
    spec = _spec_for(imp)
    if spec is not None and spec.log_stage is None:
        return _locked("経過の記録の列がないので、この取り込みでは使いません")
    reason = _not_ready_reason(imp, spec)
    if reason:
        return _locked(reason)
    from app import ai

    log_key = spec.log_stage.column
    col = spec.column(log_key)
    row_choices = []
    for rec in load_rows(import_id):
        text = str((rec.get("values") or {}).get(log_key) or "").strip()
        if text:
            label = rec["key"] + "｜" + " ".join(text.split())[:40]
            row_choices.append((rec["key"], label))
            if len(row_choices) >= 200:
                break
    ai_job = _ai_job(import_id)
    html = render_part("base.html", "part_ai", imp=imp, log_display=col.display if col else log_key, row_choices=row_choices,
        trial_keys=[k for k, _ in row_choices[:10]], job=ai_job, job_active=bool(ai_job and not ai_job.get("finished")),
        job_url=url_for("tables.api_job", job_id=ai_job["id"]) if ai_job else None,
        counts=ai.counts(imp["template_id"], "log", import_id=import_id),
        scopes=SCOPE_LABELS, item_labels=ai.STATUS_LABELS, **_ai_connection_ctx())
    return _panel(html, note=f"経過の記録の列: {col.display if col else log_key}",
                  job=_job_info(ai_job, import_id) if ai_job and not ai_job.get("finished") else None)


@tables_bp.post("/imports/<int:import_id>/ai/split-preview")
def ai_split_preview(import_id: int):
    from app import ai
    from app.ai import format_author, format_when, review_notes
    from app.ai import render_timeline
    from app import ai

    _load_import(import_id)
    row_key = str(_tables_payload().get("row_key") or "")
    try:
        data = ai.load_rows_for_ai(import_id)
        works = ai.prepare_works(data, ["log"], [row_key])
    except ai.AIJobError as exc:
        return _json_error(str(exc))
    if not works:
        return _json_error("行が見つかりません")
    w = works[0]
    parse = w.parse
    # id（s1, s2…）は AI との受け渡し用で画面には出さない。画面は no（順番1、順番2…）と、原文（text）の中の
    # 位置 start/end で「原文と並べて」見せる（利用者の指示 2026-09-21）
    segments = [{
        "id": s.id, "no": n, "body": s.body, "start": int(s.start), "end": int(s.end),
        "when": format_when(s.when), "when_estimated": bool(s.when and s.when.estimated),
        "author": format_author(s.author), "author_estimated": bool(s.author and s.author.estimated),
        "marks": list(s.marks or []), "identifiers": list(s.identifiers or []), "quantities": list(s.quantities or []),
        "plans": list(s.plans or []),
    } for n, s in enumerate(parse.segments, start=1)]
    return jsonify({
        "row_key": row_key, "kind": parse.kind, "order": parse.order, "text": parse.text, "segments": segments,
        "route": w.route, "reason": w.reason, "ai_ready": ai.is_configured(), "notes": review_notes(parse),
        "timeline": render_timeline(parse, w.entity_label, glossary=(w.stage.glossary or None) if w.stage else None),
        "sent_text": ai.messages_text(w.messages) if w.messages else "",
    })


@tables_bp.post("/imports/<int:import_id>/ai/trial")
def ai_trial(import_id: int):
    from app import ai

    imp = _load_import(import_id)
    spec = _spec_for(imp)
    if spec is None or spec.log_stage is None:
        return _json_error("経過の記録の列がありません")
    if imp["status"] == "confirmed":
        # 試し実行の結果は下書きに入るが、zip は確定したときの md を渡す。食い違わないよう断る
        return _json_error("確定後は試し実行できません（結果が確定した Markdown に入らないため）。"
                           "試すときは列の対応づけからもう一度読み込んでください")
    if _ai_running(import_id):
        # 全件実行と試し実行が同じ行に書き戻すので、別の設定の結果が混ざる（2026-09-23 のレビュー）
        return _json_error("AI整形が動いています（一時停止中を含む）。"
                           "終わるか中止してから試し実行してください")
    payload = _tables_payload()
    row_key = str(payload.get("row_key") or "")
    try:
        settings = ai.job_client_settings()
        if not ai.is_local_endpoint(settings.get("chat_url") or "") and not payload.get("confirm_external"):
            return _json_error("対応内容が外部のAIサービスに送信されます。確認のチェックを入れてください")
        with _TRIALS_LOCK:
            _TRIALS[import_id] += 1
        try:
            data = ai.load_rows_for_ai(import_id)
            result = ai.trial_row(import_id, row_key, stage_ids=["log"], data=data)
        finally:
            with _TRIALS_LOCK:
                _TRIALS[import_id] -= 1
                if _TRIALS[import_id] <= 0:
                    del _TRIALS[import_id]
    except ai.LLMNotConfigured:
        return _json_error("AI接続が設定されていません。画面右上の「AI接続」で設定してください")
    except ai.LLMCallError as exc:
        if exc.kind == "fatal":
            # 接続先に届かない・キーが拒否された。ヘッダーの表示もそれに合わせる（確認しに行く手間はかけない）
            ai.record_check(current_session_id(), False, [{"name": "試し実行", "ok": False, "detail": str(exc)}])
        return _json_error(str(exc))
    except ai.AIJobError as exc:
        return _json_error(str(exc))
    stage = next((s for s in result["stages"] if s.get("stage_id") == "log"), None) or {}
    accepted = stage.get("result") if stage.get("status") in ("ok", "flagged") else None
    record = next((r for r in data.rows if r["key"] == row_key), None)
    md = ""
    if record is not None:
        ai_results = {row_key: {"status": "ok", "result": accepted}} if accepted else {}
        md = "\n".join(record_block(record, spec, ai_results, people_index_for(spec, data.rows)))
    checks = stage.get("checks") or {}
    issues = [i.get("message", "") for i in (checks.get("issues") or []) if isinstance(i, dict)]
    return jsonify({
        "row_key": row_key, "status": stage.get("status"), "route": stage.get("route"), "reason": stage.get("reason"),
        "error": stage.get("error") or "", "points": ai_point_lines(accepted, spec.log_stage.glossary) if accepted else [],
        "timeline": stage.get("timeline") or [], "issues": issues, "markdown": md, "cached": stage.get("cached"),
        "stats": ai.trial_stats([result]),
    })


@tables_bp.post("/imports/<int:import_id>/ai/estimate")
def ai_estimate(import_id: int):
    from app import ai

    _load_import(import_id)
    data = _tables_payload()
    scope = data.get("scope") if data.get("scope") in SCOPE_LABELS else "pending"
    try:
        settings = ai.job_client_settings()
        result = ai.estimate(import_id, [s for s in data.get("trials") or [] if isinstance(s, dict)],
                                          scope=scope, concurrency=_int(data.get("concurrency")) or None,
                                          settings=settings, stage_ids=["log"])
    except ai.LLMNotConfigured:
        return _json_error("AI接続が設定されていません")
    except ai.AIJobError as exc:
        return _json_error(str(exc))
    return jsonify(result)


@tables_bp.post("/imports/<int:import_id>/ai/run")
def ai_run(import_id: int):
    from app import ai

    imp = _load_import(import_id)
    data = _tables_payload()
    if imp["status"] != "preview":
        return _json_error("表の読み込みが終わってから実行してください")
    job = _ai_job(import_id)
    if job and not job.get("finished"):
        return _json_error("AI整形はすでに実行中です")
    scope = data.get("scope") if data.get("scope") in SCOPE_LABELS else "pending"
    concurrency = min(16, max(1, _int(data.get("concurrency"), 1) or 1))
    try:
        settings = ai.job_client_settings()
        if not ai.is_local_endpoint(settings.get("chat_url") or "") and not data.get("confirm_external"):
            return _json_error("対応内容が外部のAIサービスに送信されます。確認のチェックを入れてください")
        # 始める前に接続を確かめる（利用者の指示: AI整形を始めるときに確認）。ヘッダーの表示もここで更新される。
        # つながらなければジョブを作らずに理由を返す（何百行も失敗させてから気づかせない）
        ok, steps = ai.check_connection()
        ai.record_check(current_session_id(), ok, steps)
        if not ok:
            failed = next((s for s in steps if not s.get("ok")), steps[-1])
            return _json_error(f"AIにつながりません（{failed['name']}: {failed['detail']}）。"
                               "画面右上の「AI接続」を確認してください", ai_status=ai.connection_status())
        job_id = ai.start_ai_job(import_id, scope=scope, concurrency=concurrency, stage_ids=["log"])
    except ai.LLMNotConfigured:
        return _json_error("AI接続が設定されていません。画面右上の「AI接続」で設定してください")
    return jsonify({"ok": True, "job_id": job_id, "job_url": url_for("tables.api_job", job_id=job_id),
                    "ai_status": ai.connection_status()})


@tables_bp.post("/imports/<int:import_id>/ai/<action>")
def ai_control(import_id: int, action: str):
    _load_import(import_id)
    handlers = {"pause": request_pause, "resume": request_resume, "cancel": request_cancel}
    if action not in handlers:
        abort(404)
    job = _ai_job(import_id)
    ok = bool(job) and handlers[action](job["id"])
    if not ok:
        return _json_error("AI整形のジョブを操作できませんでした（すでに終わっている可能性があります）")
    return jsonify({"ok": True})


# ---- AI接続（ヘッダーの右上。どの画面でも同じパネルが開く。設定画面は無い） -----------------------------
# 利用者の指示（2026-09-21）:「AI接続の設定は、もともとの位置ヘッダーの画面右上『AI接続』に移動させる。
# 全部空欄にしておいてcookieでユーザー毎に登録内容をずっと保持させるようにしてほしい」
# 「AI接続はヘッダー上で、接続中 か 未接続 一目で分かるように」
# 設定はブラウザごと（current_session_id ごと。ai.py・core.ai_connections）。URL は昔のまま /tables/ の下に
# 置いてある（画面には出ない。base.html の data-* 属性で渡す）。

@tables_bp.app_context_processor
def ai_header_ctx() -> dict:
    """どの画面のヘッダーにも「AI接続」の状態を出す。確認しには行かず、覚えている結果を出すだけ。"""
    from app import ai

    try:
        status = ai.connection_status()
    except Exception as exc:   # DB が読めないときでも画面は出す
        status = {"state": "off", "state_label": "未接続", "ready": False, "checked_text": "", "check_detail": str(exc),
                  "sub_text": "クリックして設定", "panel_sub_text": "",
                  "chat_url": "", "models_url": "", "model": "", "models": [], "api_key_saved": False,
                  "api_key_source": "", "effective": {}, "fallback": {},
                  "placeholders": {"chat_url": "https://api.openai.com/v1/chat/completions",
                                   "models_url": "https://api.openai.com/v1/models", "model": ""}}
    return {"ai_header": status,
            "ai_urls": {"save": url_for("tables.save_ai_connection"), "test": url_for("tables.test_ai_connection"),
                        "models": url_for("tables.refresh_ai_models"), "clear": url_for("tables.clear_ai_key")}}


@tables_bp.post("/ai-connection")
def save_ai_connection():
    """ヘッダーの「AI接続」の［保存する］。保存したあと、その設定でつながるかを確かめる（結果はヘッダーに出る）。"""
    from app import ai

    data = _tables_payload()
    sid = current_session_id()
    try:
        ai.save_browser(sid, data)
    except ValueError as exc:
        return _json_error(str(exc))
    except sqlite3.Error as exc:
        return _json_error(f"保存できませんでした（{exc.__class__.__name__}）。少し待ってから、もう一度保存してください")
    ok, steps = ai.check_connection()
    ai.record_check(sid, ok, steps)
    status = ai.connection_status()
    if not status["ready"]:
        message = "保存しました（APIキーと接続先がそろうと接続を確かめます）"
    else:
        message = "保存しました。" + ("AIにつながりました" if ok else "AIにつながりません（下の結果を確認してください）")
    return jsonify({"ok": True, "message": message, "connected": ok, "steps": steps, "status": status})


@tables_bp.post("/ai-connection/clear")
def clear_ai_key():
    """［キーを消す］（共有PCで使い終わったとき）。このブラウザの APIキーだけ消す。接続先・モデルは残す。"""
    from app import ai

    status = ai.clear_browser_key(current_session_id())
    return jsonify({"ok": True, "message": "このブラウザのAPIキーを消しました", "status": status})


@tables_bp.post("/ai-connection/models")
def refresh_ai_models():
    """APIからモデル一覧を取得し、このブラウザの候補（入力欄の候補）として覚える。"""
    from app import ai

    try:
        catalog = ai.model_catalog(refresh=True)
    except Exception as exc:
        return _json_error(f"モデル一覧を取得できませんでした: {ai.friendly_error(exc)}")
    if catalog:
        save_ai_connection_row(current_session_id(), models=catalog)   # 入力欄の候補（datalist）として覚える
        ai.forget_browser()
    return jsonify({"ok": True, "models": catalog, "message": f"APIからモデル一覧を取得しました（{len(catalog)}件）",
                    "status": ai.connection_status()})


@tables_bp.post("/ai-connection/test")
def test_ai_connection():
    """接続の確認: モデル一覧の取得と、1回の短いチャット。パネルを開いたとき・［接続を確かめる］で呼ぶ。"""
    from app import ai

    ok, steps = ai.check_connection()
    ai.record_check(current_session_id(), ok, steps)
    return jsonify({"ok": ok, "steps": steps, "status": ai.connection_status()})


# ---- 内容の確認 ---------------------------------------------------------------------------

def _preview_job(import_id: int, imp: dict, spec, retry: bool = False) -> dict:
    """確認の段の md を作るジョブ。同じ入力のジョブがあればそれを見せ、なければ始める。

    retry: 同じ入力で失敗したジョブを作り直す（［もう一度作る］。ファイルの一時的なロックなどで失敗したとき）。
    """
    signature = preview_signature(import_id, imp, spec)
    job = latest_job("table_import", import_id, kind="table_preview")
    if job is not None and (job.get("params") or {}).get("signature") == signature:
        if not job.get("finished") or (job["status"] == "failed" and not retry):
            return job  # 動いている途中、または同じ入力で失敗したまま（開き直すたびに作り直さない）
    return get_job(start_preview_job(import_id, signature))


def _display_with_unit(col) -> str:
    """データ表の見出し。Markdown は「停止時間: 5分」と単位を付けて書くので、見出しにも単位を添える。"""
    unit = (col.unit or "").strip()
    if unit and unit not in col.display:
        return f"{col.display}（{unit}）"
    return col.display


def _panel_preview(imp: dict, page: int = 1):
    import_id = imp["id"]
    job = _running_job(imp)
    if job is not None:
        return _panel("", job=_job_info(job, import_id), reading=True)
    spec = _spec_for(imp)
    reason = _not_ready_reason(imp, spec)
    if reason:
        return _locked(reason, failed=imp["status"] == "failed",
                       retry_url=url_for("tables.reread", import_id=import_id) if spec is not None else None)
    if _ai_running(import_id):
        return _locked("AI整形が動いています（一時停止中を含む）。再開して終わらせるか、中止してから確認してください")
    files = ready_preview_files(import_id, imp, spec)
    if files is None:
        # 件数分の md を作るのに時間がかかるので、ジョブにして同じ画面に進み具合を出す
        draft = _preview_job(import_id, imp, spec)
        if draft.get("finished") and draft["status"] != "done":
            return _panel(render_part("base.html", "part_preview_failed", error=draft.get("message") or "Markdownの下書きを作れませんでした"),
                          failed=True)
        return _panel("", job=_job_info(draft, import_id), building=True)
    stats = imp.get("stats") or {}
    issues = load_issues(import_id)
    page = max(1, page)
    data_rows, total = load_rows_page(import_id, (page - 1) * DATA_PAGE, DATA_PAGE)
    total_pages = max(1, (total + DATA_PAGE - 1) // DATA_PAGE)
    if page > total_pages:
        page = total_pages
        data_rows, _total = load_rows_page(import_id, (page - 1) * DATA_PAGE, DATA_PAGE)
    blocking = has_blocking(issues)
    # 確定を止めているエラーを先に出す。行番号順のまま先頭200件を切っていたころは、
    # エラーが201件目以降だと「エラーが残っているため確定できません」と出るのに
    # 問題一覧にはエラーが1件も無く、何を直せばよいか分からなかった（2026-09-23 のレビューで実測）
    shown = sorted(issues, key=lambda i: 0 if i.get("level") == "error" else 1)
    html = render_part("base.html", "part_preview", imp=imp, spec=spec, stats=stats, issues=shown[:ISSUES_SHOWN],
        issue_total=len(issues), counts=count_levels(issues), blocking=blocking, files=files, data_rows=data_rows,
        page=page, total_pages=total_pages, columns=[(c.key, _display_with_unit(c)) for c in spec.columns],
        confirmed=imp["status"] == "confirmed", delete_note=TABLES_DELETE_ON_DOWNLOAD_NOTE)
    return _panel(html, note=f"{stats.get('records') or 0}件・{len(files)}ファイル", blocking=blocking,
                  confirmed=imp["status"] == "confirmed")


@tables_bp.post("/imports/<int:import_id>/preview")
def start_preview(import_id: int):
    """確認の段の Markdown の下書きを作り直す（失敗したときの［もう一度作る］）。"""
    imp = _load_import(import_id)
    spec = _spec_for(imp)
    if spec is None or imp["status"] not in ("preview", "confirmed"):
        return _json_error("表の読み込みが終わってから確認してください")
    job = _preview_job(import_id, imp, spec, retry=True)
    return jsonify({"ok": True, "job": _job_info(job, import_id), "building": True})


@tables_bp.get("/imports/<int:import_id>/preview/file")
def preview_file(import_id: int):
    _load_import(import_id)
    text = md_text(import_files(import_id)["preview"], request.args.get("name", ""))
    if text is None:
        return _json_error("ファイルが見つかりません", 404)
    return jsonify({"name": request.args.get("name"), "text": text})


@tables_bp.get("/imports/<int:import_id>/issues.csv")
def issues_csv(import_id: int):
    imp = _load_import(import_id)
    data = build_issues_csv(load_issues(import_id))
    name = f"{Path(imp['file_name']).stem}_問題一覧.csv"
    return set_download_name(
        send_file(io.BytesIO(data), mimetype="text/csv", as_attachment=True, download_name=name), name, "issues")


@tables_bp.post("/imports/<int:import_id>/confirm", endpoint="confirm")
def tables_confirm(import_id: int):
    """確定して Markdown を作る（できたら同じ画面にダウンロードのボタンを出す）。"""
    imp = _load_import(import_id)
    spec = _spec_for(imp)
    if spec is None or imp["status"] not in ("preview", "confirmed"):
        return _json_error("この取り込みはまだ確定できる状態ではありません")
    if _ai_running(import_id):
        return _json_error("AI整形の実行中は確定できません。終わるか中止してから確定してください")
    if has_blocking(load_issues(import_id)):
        return _json_error("エラーが残っているため確定できません。問題一覧を確認して、範囲や列の対応づけを直してください")
    job_id = start_render_job(import_id)
    return jsonify({"ok": True, "next": "done", "job": _job_info(get_job(job_id), import_id), "building": True})


# ---- 確定してダウンロード（zip） ------------------------------------------------------------------

def _download_base(imp: dict, spec) -> str:
    stem = safe_filename_part(spec.file_prefix if spec is not None else Path(imp["file_name"]).stem)
    return f"{stem}_{datetime.now().strftime('%Y-%m-%d')}"


def _panel_done(imp: dict):
    import_id = imp["id"]
    job = _running_job(imp)
    if job is not None:
        return _panel("", job=_job_info(job, import_id), building=True)
    if imp["status"] != "confirmed":
        return _locked("上の「内容の確認」で［確定してMarkdownを作成］を押してください")
    spec = _spec_for(imp)
    files = [{"name": p.name, "size": p.stat().st_size} for p in md_paths(import_id)]
    if not files:
        # 記録0件のまま確定された取り込み（渡せるものが無いので、ダウンロードのボタンは出さない）
        return _locked("この取り込みから作られた Markdown はありません（取り込める行がありませんでした）。"
                       "表の範囲か元のファイルを見直してください")
    html = render_part("base.html", "part_done", imp=imp, files=files, stats=imp.get("stats") or {},
                           delete_note=TABLES_DELETE_ON_DOWNLOAD_NOTE, delete_confirm=TABLES_DELETE_ON_DOWNLOAD_CONFIRM)
    return _panel(html, note=f"{len(files)}ファイル")


@tables_bp.get("/imports/<int:import_id>/download.zip")
def download_zip(import_id: int):
    """zip を渡し、渡し終えた取り込みのデータを消す（design.md 3.3）。

    zip は全体をメモリに作ってから消すので、消す処理で中身が欠けることはない。作れなかったときは何も消さない。
    """
    imp = _load_import(import_id)
    spec = _spec_for(imp)
    # 断る理由はこの画面に直接書く。以前は知らせ（flash）を積んで転送していたが、別の段の描画が
    # 知らせを先に食べることがあり、理由の出ない空の画面に飛ばされていた（2026-09-23 のレビュー）
    reason = None
    if imp["status"] != "confirmed" or spec is None:
        reason = "先に [確定してMarkdownを作成] を押してください"
    elif not md_paths(import_id):
        # 確定はしたが記録が0件だった（押したばかりのボタンをもう一度押せ、とは言わない）
        reason = ("作成された Markdown がありません（取り込める行がありませんでした）。"
                  "表の範囲か元のファイルを見直してください")
    elif _ai_running(import_id):
        reason = "AI整形の実行中はダウンロードできません。終わるか中止してからダウンロードしてください"
    elif _trial_running(import_id):
        reason = TRIAL_BUSY_MESSAGE
    if reason:
        if request.accept_mimetypes.best == "application/json" or request.is_json:
            return jsonify(error=reason), 409
        return render_template("base.html", code=409, reason=reason), 409
    try:
        data = build_download(import_id)
    except FileNotFoundError:
        # 同時に押した別のダウンロードが渡し終えてデータを消した（ダブルクリックなど）。500 ではなく「ダウンロード済み」。
        return _handout_lost(import_id, "ダウンロード")
    name = f"{_download_base(imp, spec)}.zip"
    response = send_file(io.BytesIO(data), mimetype="application/zip", as_attachment=True, download_name=name,
                         conditional=False)   # Range でも全体を返す（一部だけ渡して消すことが無いように）
    set_download_name(response, name, "records")
    return purge_after_send(response, purge_table_import, import_id)


def _handout_lost(import_id: int, what: str):
    """渡す途中で md が消えたとき。取り込みが残っていて確定済みでなければ、別のタブで読み込み直しが始まった
    （「ダウンロード済み」の 404 は事実と違うので、画面に戻して知らせる）。消えていれば渡し終えた・削除した。"""
    now = get_import(import_id)
    if now is None or now["status"] == "confirmed":
        abort(404)
    flash(f"読み込み直しが始まったため、{what}を取りやめました。確定し直してからもう一度押してください", "error")
    return redirect(url_for("tables.new"))


# ---- 段の割り当て ---------------------------------------------------------------------------

def _panel_rows(imp: dict):
    """データ（100行ずつ）のページ送りだけ作り直す。"""
    return _panel_preview(imp, page=max(1, _int(request.args.get("page"), 1) or 1))


_PANELS = {"source": _panel_source, "layout": _panel_layout, "columns": _panel_columns, "ai": _panel_ai,
           "preview": _panel_rows, "done": _panel_done}


# ---- ジョブの進捗（JSON） ------------------------------------------------------------------------------

def _job_visible(job: dict) -> bool:
    """このブラウザの取り込みのジョブか（ほかの人のジョブは進み具合も見せない）。

    取り込みが消えている（ダウンロード済み・削除済み）ジョブは、持ち主が分からないので見せない。
    """
    if (job.get("ref_type") or "") != "table_import" or job.get("ref_id") is None:
        return False
    imp = get_import(job["ref_id"])
    return imp is not None and owns(imp)


def api_job(job_id: int):
    job = get_job(job_id)
    if job is None or not _job_visible(job):
        return _json_error("ジョブが見つかりません", 404)
    return jsonify({k: job.get(k) for k in ("id", "kind", "status", "status_label", "progress", "message", "result",
                                            "finished")})


@tables_bp.record_once
def _register_api(setup_state) -> None:
    # /api/jobs は /tables の外に置く（design.md 2.4）
    setup_state.app.add_url_rule("/api/jobs/<int:job_id>", "tables.api_job", api_job)


####################################################################################################
# create_app・エラー画面・起動コマンド
####################################################################################################

def create_app(overrides: dict | None = None) -> Flask:
    _refuse_debugger()
    app = Flask(__name__)
    app.config.from_object(Config)
    # クッキーは約1年もたせる（views.SESSION_LIFETIME。ヘッダーの「AI接続」の設定をブラウザごとに
    # 「ずっと保持」するため。利用者の指示 2026-09-21）。views.current_session_id が session.permanent を立てる
    app.config.update(SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Lax",
                      PERMANENT_SESSION_LIFETIME=SESSION_LIFETIME)
    if overrides:
        app.config.update(overrides)
    if not app.config.get("SECRET_KEY"):
        app.config["SECRET_KEY"] = _secret_key()

    # DATA_DIR だけ差し替えたとき（テスト等）は一覧表の置き場所もそれに合わせる
    if not (overrides and "TABLES_DIR" in overrides) and "TABLES_DIR" not in os.environ:
        app.config["TABLES_DIR"] = Path(app.config["DATA_DIR"]) / "tables"
    for key in ("UPLOAD_DIR", "DATA_DIR", "TABLES_DIR"):
        app.config[key] = Path(app.config[key])
        app.config[key].mkdir(parents=True, exist_ok=True)
    init_app(app)
    if not app.config.get("ALLOWED_HOSTS"):
        app.config["ALLOWED_HOSTS"] = allowed_hosts()
    else:   # 設定から渡されたものも、Host と同じ形（ポート無し・小文字）に揃え、ループバックは必ず足す
        given = {h if h == "*" else (_hostname_only(h) or str(h).lower()) for h in app.config["ALLOWED_HOSTS"]}
        app.config["ALLOWED_HOSTS"] = {"*"} if "*" in given else given | {"127.0.0.1", "localhost", "::1"}

    app.register_blueprint(home_bp)
    for bp in (forms_bp, form_types_bp, tables_bp):
        app.register_blueprint(bp)

    @app.before_request
    def _refuse_other_host():
        # 宛先の名前（Host）が、運用者が意図した名前かを確かめる（allowed_hosts のとおり）。
        # 待ち受けているアドレスを攻撃者のドメインが指していると、そのページと同一オリジンになり、
        # 画面の中身を読み取られてしまう（DNSリバインディング）。
        allowed = app.config["ALLOWED_HOSTS"]
        try:
            parts = urlsplit(f"//{request.host}")
            host, port = parts.hostname, parts.port   # ポート番号・[] を外したホスト名
        except ValueError:
            host, port = None, None
        host = (host or "").rstrip(".").lower()
        if "*" not in allowed and host not in allowed:
            # 汎用の 400 画面（ホームへ = 同じアドレスの / ）では、また断られるだけで開き方が分からない。
            # 開けるアドレスを示す（ポートは送られてきた Host のもの。読めなければ起動時の PORT）
            home_url = f"http://{_home_host()}:{port or PORT}/"
            text = f"このアドレスでは開けません。このアプリは {home_url} で開いてください"
            if request.accept_mimetypes.best == "application/json" or request.is_json:
                return jsonify(error=text), 400
            return render_template("base.html", code=400, host_refused=True, home_url=home_url), 400

    @app.before_request
    def _refuse_cross_site_write():
        # 他のサイトのページから、利用者のブラウザ経由で書き込ませない（views.is_cross_site_write のとおり）。
        if is_cross_site_write(request):
            abort(403)

    @app.after_request
    def _no_store(response):
        # 画面・JSON には取り込んだ値や Markdown が載る。ダウンロードでサーバー側から消しても（design.md 3.3）、
        # ブラウザのキャッシュ（戻る・再表示）に残らないよう、保存させない。静的ファイル（CSS・JS）は除く
        if request.endpoint != "static":
            response.headers["Cache-Control"] = "no-store"
        # ほかのサイトのページの中（iframe）に表示させない。枠の中でのクリックは同じサイトからの操作になり、
        # is_cross_site_write では防げない（削除ボタンを押させる、など）。静的ファイルも含めてすべての応答に付ける
        response.headers["X-Frame-Options"] = "DENY"
        response.headers["Content-Security-Policy"] = "frame-ancestors 'none'"
        return response

    @app.errorhandler(403)
    def _forbidden(_exc):
        # 画面からの JSON 送信（static/app.js の postJson）には JSON で返す。トーストに理由が出る。
        if request.accept_mimetypes.best == "application/json" or request.is_json:
            return jsonify(error="ほかのサイトのページから送られてきた操作に見えたため受け付けませんでした。"
                                 "取り込みの画面からやり直してください"), 403
        return render_template("base.html", code=403), 403

    @app.errorhandler(404)
    def _not_found(_exc):
        # 403 と同じく、画面からの JSON 送信（static/app.js の postJson）には JSON で返す。
        # ダウンロード済みでデータが消えたあとに、開いたままの画面が途中保存・プレビューを送ると
        # ここに来る。HTML を返すとトーストに「通信に失敗しました」としか出ず、理由が伝わらない。
        if request.accept_mimetypes.best == "application/json" or request.is_json:
            return jsonify(error="このデータはサーバーに残っていません（ダウンロード済みか、削除されています）。"
                                 "取り込みの画面からやり直してください"), 404
        return render_template("base.html", code=404), 404

    @app.errorhandler(413)
    def _too_large(_exc):
        # MAX_CONTENT_LENGTH を超えた送信。Werkzeug の英語の画面を出さない（design.md 2: エラー画面は日本語）
        limit = app.config.get("MAX_CONTENT_LENGTH") or 0
        text = f"{limit / 1024 / 1024:.0f}MB" if limit >= 1024 * 1024 else f"{limit / 1024:.0f}KB"
        if request.accept_mimetypes.best == "application/json" or request.is_json:
            return jsonify(error=f"送った内容が大きすぎます（1回に合計 {text} まで）。分けて送ってください"), 413
        return render_template("base.html", code=413, limit=text), 413

    @app.errorhandler(400)
    @app.errorhandler(405)
    def _bad_request(exc):
        # 送信専用の URL をアドレス欄から開いた（405）など。Werkzeug の英語の画面を出さない（design.md 2）
        if request.accept_mimetypes.best == "application/json" or request.is_json:
            return jsonify(error="この操作は受け付けられませんでした。取り込みの画面から開き直してください"), exc.code
        return render_template("base.html", code=400), exc.code

    @app.errorhandler(500)
    def _server_error(_exc):
        # 画面の中からの送信に HTML を返すと、トーストに「通信に失敗しました」としか出ず
        # 理由が伝わらない（ほかのエラーと同じ扱いにする。2026-09-23 のレビュー）
        if request.accept_mimetypes.best == "application/json" or request.is_json:
            return jsonify(error="サーバー側でエラーが発生しました。画面を開き直してもう一度お試しください"
                                 "（詳しい内容はアプリのログに記録しました）"), 500
        return render_template("base.html", code=500), 500

    @app.cli.command("serve")
    def _serve() -> None:
        """本番用サーバー（waitress）で待ち受ける。待ち受け先は環境変数 HOST・PORT。"""
        # flask --app app run は Flask の開発サーバーを使うので、通常はこちらを使う。開発サーバーだと
        # WAITRESS_OPTIONS の outbuf_high_watermark が効かず、途中で切れたダウンロードに気づけない
        # （送り終えたように見えた時点でデータを消してしまう。design.md 3.3「欠けないダウンロード」）。
        if DEBUG and not _is_loopback(HOST):
            # DEBUG=True の Flask はデバッガを開く。LAN に出す起動では絶対に開かせない（_NO_DEBUG_MSG と同じ理由）
            raise SystemExit("\n[app] DEBUG = True のまま LAN のアドレス（HOST=%s）では起動しません。\n"
                             "  理由: デバッガが開くと、例外が出たときにブラウザからこのサーバの Python を実行できます。\n"
                             "  詳しいエラーを見たいときは HOST を外して（127.0.0.1 で）起動してください。\n" % HOST)
        print(startup_notice())
        if DEBUG:
            app.run(host=HOST, port=PORT, debug=True, use_reloader=False)
            return
        try:
            from waitress import serve
        except ImportError:
            print("[app] waitress が無いため Flask の開発サーバで起動します（pip install waitress を推奨）")
            app.run(host=HOST, port=PORT, debug=False)
        else:
            serve(app, host=HOST, port=PORT, **WAITRESS_OPTIONS)

    _recover_jobs(app)
    _purge_pending(app)
    _cleanup_leftovers(app)
    _start_sweeper(app)
    # ヘッダーは「帳票取り込み／表の取り込み／帳票登録」の3つだけ。使うAIモデルの選択は
    # 「表の取り込み」画面の AI整形の段の中（AI接続）に移したので、共通の値は渡さない。
    return app


# 途中で放り出されたものを捨てる間隔（design.md 3.3）。起動時に全部捨て、動いている間は
# core.purge.STALE_HOURS より古いものをこの間隔で捨てる（数人で使うサーバーに置きっぱなしになるため）。
# 捨てるまでの時間（STALE_HOURS）を短くしたときは、見回りもそれに合わせて短くする。
# 見回りの間隔だけ延びると「2時間で捨てます」と言いながら3時間残ることになるので、上限は10分にする。
SWEEP_INTERVAL_SECONDS = 10 * 60


def _sweep_interval(stale_hours: float) -> int:
    return max(60, min(SWEEP_INTERVAL_SECONDS, int(stale_hours * 3600 / 2) or SWEEP_INTERVAL_SECONDS))


def _purge_pending(app: Flask) -> None:
    """起動時：ダウンロードしていない帳票・一覧表をすべて捨てる（作業中の一覧を持たないため）。"""
    if app.config.get("TESTING"):
        return

    try:
        with app.app_context():
            forms_removed, tables_removed = purge_all_pending()
            samples_removed = purge_old_sample_files()
        if forms_removed or tables_removed:
            print(f"[app] 途中だった取り込みを捨てました（帳票 {forms_removed} 件・一覧表 {tables_removed} 件）")
        if samples_removed:
            print(f"[app] 前の版が置いた見本のExcelを捨てました（{samples_removed} 件・見本はもう置きません）")
    except Exception as exc:  # 起動は止めない
        print(f"[app] 途中だった取り込みの片付けに失敗しました（{exc.__class__.__name__}: {exc}）")


def _start_sweeper(app: Flask) -> None:
    """動いている間：しばらくさわられていない帳票・一覧表を捨て続ける（daemon スレッド）。"""
    if app.config.get("TESTING"):
        return

    interval = _sweep_interval(STALE_HOURS)

    def loop() -> None:
        while True:
            _stop.wait(interval)
            try:
                with app.app_context():
                    forms_removed, tables_removed = sweep_stale(STALE_HOURS)
                if forms_removed or tables_removed:
                    print(f"[app] {STALE_HOURS}時間さわられていない取り込みを捨てました"
                          f"（帳票 {forms_removed} 件・一覧表 {tables_removed} 件）")
            except Exception as exc:   # 次の回でやり直す
                print(f"[app] 古い取り込みの片付けに失敗しました（{exc.__class__.__name__}: {exc}）")

    _stop = threading.Event()
    threading.Thread(target=loop, name="purge-sweeper", daemon=True).start()


def _cleanup_leftovers(app: Flask) -> None:
    """起動時：DB から参照されていない残骸（アップロードファイル・取り込みのフォルダ）を消す。

    データを残さない方針（design.md 3.3）でも、削除の途中で落ちたときやファイルを掴まれていたときに
    残ることがあるので、ここで片付ける。作業中のものは DB に行があるので消さない。
    """
    try:
        with app.app_context():
            db = get_db()
            known: set[str] = set()
            for (table,) in db.execute("SELECT name FROM sqlite_master WHERE type = 'table'").fetchall():
                columns = {row[1] for row in db.execute(f'PRAGMA table_info("{table}")')}
                if "stored_path" in columns:
                    known |= {row[0] for row in db.execute(f'SELECT stored_path FROM "{table}"') if row[0]}
            removed = remove_orphan_uploads(app.config["UPLOAD_DIR"], known)
            import_ids = {row[0] for row in db.execute("SELECT id FROM table_imports")}
            removed_dirs = remove_orphan_import_dirs(app.config["TABLES_DIR"], import_ids)
            # もう無い取り込みの AI整形の結果と、どこからも使われない AI の応答も消す（データを残さない）。
            # ai_items が先に消えていて（取り込み設定の削除など）応答だけ残っていることもあるので、毎回見る。
            sweep_orphan_ai(db)
            # 前の版で空になった表に残っている番号の続き（何件取り込んだか）も忘れる（design.md 3.3「履歴は持たない」）
            if forget_id_counters(db):
                db.execute("PRAGMA wal_checkpoint(TRUNCATE)")   # 置き換える前の値を app.db-wal に残さない
        if removed:
            print(f"[app] 参照されていないアップロードファイルを {removed} 件片付けました")
        if removed_dirs:
            print(f"[app] 参照されていない取り込みのフォルダを {removed_dirs} 件片付けました")
    except Exception as exc:  # 起動は止めない
        print(f"[app] 残ったファイルの整理に失敗しました（{exc.__class__.__name__}: {exc}）")


def _recover_jobs(app: Flask) -> None:
    """前回の終了時に動いていたジョブを「中断」にする。"""
    try:
        with app.app_context():
            recover_interrupted()
    except Exception as exc:  # 起動は止めない
        print(f"[app] 中断したジョブの整理に失敗しました（{exc.__class__.__name__}: {exc}）")


# --- 起動の設定 -------------------------------------------------------------
# 待ち受け先（HOST・PORT）はファイルの上のほうで環境変数から読む。
THREADS = 8
# waitress が応答の本文を先読みして溜める上限（既定は 16MB）。
# 溜められる分はアプリ側では「送り終えた」ように見えるため、既定のままだと 16MB 未満の zip は
# 途中で通信が切れても消えてしまう（core/purge.purge_after_send・design.md 3.3「欠けないダウンロード」）。
# 小さくすると、送れた分だけ読み進めるので途中で切れたことに気づける。ループバックでは速度はほぼ変わらない
# （30MB の本文で 16MB: 約130ms、16KB: 約100ms）。それでも最後の数十KB（この上限＋OS の送受信の溜め。
# 2026-09-19 の測定で約48KB）は相手が受け取る前に送り終えたことになるので、それより小さい md・zip は
# 途中で切れても消える（design.md 3.3・8.0）。
OUTBUF_HIGH_WATERMARK = 16 * 1024
WAITRESS_OPTIONS = {"threads": THREADS, "outbuf_high_watermark": OUTBUF_HIGH_WATERMARK}
# エラー画面に詳細を出すか。通常は False のまま。
DEBUG = False
