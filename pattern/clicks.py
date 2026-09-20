"""クリックしたセルから読み取る項目を作る。

画面では「見出しのセル → 値のセル」の順にクリックするだけ。項目のキー名・型・単位・探す見出しは、
クリックした2つのセルとその値からここで決める（画面には出さない）。
"""
from __future__ import annotations

import re

from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string

from excel.tables import SAME_COLUMNS_RATIO, detect_tables, section_of, table_text_lines, table_title
from excel.text import normalize_label, split_code_name
from excel.workbook import SheetGrid
from pattern.builder import _SECTION_PREFIX, _guess_type, _lookup, _make_suggestion, _next_field_name
from pattern.dictionary import (BY_FIELD_NAME, COMBINED_EQUIPMENT_LABELS, COMBINED_EQUIPMENT_NORMS,
                                COMBINED_EQUIPMENT_PARTS)

BASE_ROW = {
    "use": True, "required": False, "rag_output": "show", "table_columns": "", "section": "",
    "unit": "", "direction": "auto", "sheet_name": "", "label_cell": "", "cell": "", "examples": [],
    # 辞書で分かった項目名（「設備番号」など）。別の見本で書き方の違う同じ欄をクリックしたときに、
    # 新しい項目にせず、その項目の探す見出しに足すために使う（保存はしない）
    "base_name": "",
}


def cell_key(grid: SheetGrid, coord) -> tuple[int, int] | None:
    """セル番地（"C5" / "C5:D6"）を、結合を考えた左上セルの (行, 列) にする。読めなければ None。"""
    text = str(coord or "").split(":")[0].strip()
    if not text:
        return None
    try:
        letters, row = coordinate_from_string(text.upper())
        col = column_index_from_string(letters)
    except Exception:
        return None
    if row < 1 or col < 1:
        return None
    top, left, _, _ = grid.bounds(row, col)
    return (top, left)


def coord_of(key: tuple[int, int]) -> str:
    return f"{get_column_letter(key[1])}{key[0]}"


def click_field(grid: SheetGrid, label_coord, value_coord="", used_names=()) -> tuple[dict | None, str]:
    """クリックした2つのセルから項目行を作る。戻り値: (項目行, エラーメッセージ)。

    - 見出しのセルだけ: 明細表の見出し（列見出しの1つ目）なら明細表の項目、そうでなければ値が空のままの項目
    - 同じセルを2回: 「設備番号：EQ-001」ならセル内の見出しと値、そうでなければ見出しのない「値だけ」の項目
    - 別のセル: 見出しと値（右・下・それ以外）
    """
    used = set(used_names)
    label_key = cell_key(grid, label_coord)
    if label_key is None:
        return None, "セルを選び直してください"
    label = grid.cells.get(label_key)
    value_key = cell_key(grid, value_coord) if value_coord else None

    if value_key is None:
        table = _table_row(grid, label_key, used)
        if table is not None:
            return table, ""
        if label is None or not label.text.strip():
            return None, "文字のないセルは見出しに選べません"
        return _label_row(grid, label, None, "auto", used), ""

    if value_key == label_key:
        if label is not None and label.inline:
            # 「設備番号：EQ-001」のように1つのセルに見出しと値が入っている
            return _label_row(grid, label, label_key, "same_cell", used), ""
        if label is None or not label.text.strip():
            return None, "空のセルは項目にできません"
        return _value_only_row(grid, label_key, used), ""

    if label is None or not label.text.strip():
        return None, "文字のないセルは見出しに選べません"
    return _label_row(grid, label, value_key, _direction(grid, label_key, value_key), used), ""


def _direction(grid: SheetGrid, label_key: tuple[int, int], value_key: tuple[int, int]) -> str:
    top, left, bottom, right = grid.bounds(*label_key)
    vtop, vleft, vbottom, vright = grid.bounds(*value_key)
    if vtop <= bottom and vbottom >= top and vleft > right:
        return "right"
    if vleft <= right and vright >= left and vtop > bottom:
        return "below"
    return "auto"


def _label_row(grid: SheetGrid, label, value_key, direction: str, used: set[str]) -> dict:
    value_cell = grid.cells.get(value_key) if value_key else None
    if direction == "same_cell":
        label_text = re.split(r"[:：]", label.text, maxsplit=1)[0].strip()
        sug = _make_suggestion(label.inline[0], label_text, label.inline[1], label)
        value_text = label.inline[1]
    else:
        label_text = label.text.rstrip(":：").strip()
        value = value_cell.value if value_cell is not None else ""
        sug = _make_suggestion(label.norm, label_text, value, label)
        value_text = value_cell.text if value_cell is not None else ""
    return {
        **BASE_ROW,
        "field_name": _field_name(sug.field_name, used),
        "base_name": sug.field_name or "",
        "display_name": sug.display_name or label_text,
        # クリックした見出しと、辞書で分かる同じ意味の見出し（「設備番号」に対する「設備No」など）。
        # 書き方の違う帳票でも同じ欄を読めるようにする（値は書き替えない）
        "candidates": "\n".join(dict.fromkeys([label_text, *sug.synonyms])),
        "data_type": sug.data_type,
        "unit": sug.unit,
        "direction": direction,
        # クリックした見出しが入っている区画（「▼ 回答欄」）。発行側と回答側に同じ欄がある帳票で、
        # 人が指したほうの欄を読むための目印（その区画のある帳票でだけ効く）
        "section": section_of(grid, label),
        "sheet_name": grid.name,
        "label_cell": coord_of((label.row, label.col)),
        "cell": coord_of(value_key) if value_key else "",
        "examples": [value_text] if value_text else [],
    }


def split_rows(row: dict, used_names=()) -> list[dict]:
    """クリックで作った項目行を、必要なら複数の項目行にする。

    「使用設備：ROB-821（ウェーハソーター 1号機）」のように番号と名前を1つのセルにまとめた欄は、
    設備番号・設備名の2項目にする。文字は1つも消さず、同じセルの読む場所を分けるだけ
    （帳票ごとに「対象設備」「使用設備」「設備」と書き方が違うので、探す見出しは全部入れる）。
    """
    used = set(used_names)
    if row["data_type"] == "table":
        return [row]
    labels = [l for l in row["candidates"].splitlines() if l.strip()]
    if not any(normalize_label(l) in COMBINED_EQUIPMENT_NORMS for l in labels):
        return [row]
    pieces = split_code_name((row.get("examples") or [""])[0])
    if pieces is None:
        return [row]
    out = []
    for field_name, index in COMBINED_EQUIPMENT_PARTS.items():
        entry = BY_FIELD_NAME[field_name]
        out.append({**row,
                    "field_name": _field_name(field_name, used),
                    "base_name": field_name,
                    "display_name": entry.display_name,
                    "candidates": "\n".join(dict.fromkeys([*labels, *COMBINED_EQUIPMENT_LABELS, *entry.synonyms])),
                    "data_type": entry.data_type,
                    "unit": "",
                    "examples": [pieces[index]]})
    return out


def _value_only_row(grid: SheetGrid, key: tuple[int, int], used: set[str]) -> dict:
    """見出しのない「値だけ」の項目（クリックしたセルの番地で読む）。"""
    cell = grid.cells[key]
    coord = coord_of(key)
    return {
        **BASE_ROW,
        "field_name": _field_name(None, used),
        "display_name": f"値（{coord}）",
        "candidates": "",
        "data_type": _guess_type(cell.value),
        "sheet_name": grid.name,
        "cell": coord,
        "examples": [cell.text],
    }


def _table_row(grid: SheetGrid, key: tuple[int, int], used: set[str]) -> dict | None:
    """クリックしたセルが明細表の見出し（または列見出しの1つ目）なら、その表を1つの項目にする。"""
    table = _table_at(grid, key)
    if table is None:
        return None
    anchor = table.anchor
    head = anchor if anchor is not None else table.header[0]
    label_text = " ".join(head.text.split())
    entry = _lookup(head)
    display = (_SECTION_PREFIX.sub("", table_title(head)).strip() or table_title(head)) if anchor else label_text
    value = table.to_value()
    lines = table_text_lines(value) if value else []
    return {
        **BASE_ROW,
        "field_name": _field_name(entry.field_name if entry else None, used),
        "base_name": entry.field_name if entry else "",
        "display_name": display,
        "candidates": "\n".join(dict.fromkeys([label_text, *(entry.synonyms if entry else ())])),
        "data_type": "table",
        "section": section_of(grid, head),
        "table_columns": "\n".join(" ".join(h.text.split()) for h in table.header),
        "sheet_name": grid.name,
        "label_cell": coord_of((head.row, head.col)),
        "examples": lines[:1],
    }


def _table_at(grid: SheetGrid, key: tuple[int, int]):
    """クリックしたセルを見出し（アンカー）か列見出しに持つ明細表。

    見本での記入が1行だけの表（「使用部品」に部品1つ）も、列見出しのすぐ上にある見出しを指したときは
    明細表にする（人が「この表」と指しているので、見本の行数では決めない）。列見出しを指したときは
    今までどおり、行が並ぶ形に見える表だけ（押印欄の「承認｜確認｜作成」を表にしないため）。
    """
    for table in detect_tables(grid):
        if _table_heading(table) == key:
            return table
        if table.has_room and any((h.row, h.col) == key for h in table.header):
            return table
    return None


def _table_heading(table) -> tuple[int, int] | None:
    """列見出しのすぐ上にある、その表の見出しセル。離れていれば None（帳票のタイトルを表の見出しにしない）。"""
    anchor = table.anchor
    if anchor is None or not table.header or anchor.max_row + 1 != min(h.row for h in table.header):
        return None
    return (anchor.row, anchor.col)


def _field_name(name: str | None, used: set[str]) -> str:
    if not name or name in used:
        name = _next_field_name(used)
    used.add(name)
    return name


def table_cells(grid: SheetGrid) -> set[str]:
    """1回のクリックで明細表の項目になるセル（見出しと列見出し）の番地。`_table_at` と同じ判断。"""
    out: set[str] = set()
    for table in detect_tables(grid):
        heading = _table_heading(table)
        if heading is not None:
            out.add(coord_of(heading))
        if table.has_room:
            out.update(coord_of((h.row, h.col)) for h in table.header)
    return out


def merge_target(field_rows: list[dict], row: dict) -> dict | None:
    """クリックで作った項目行が、登録済みのどの項目と同じ欄か（別の見本で書き方が違うだけの欄）。

    辞書で同じ項目と分かるもの（「設備番号」と「設備No」）と、列見出しがほとんど同じ明細表。
    見つかれば、新しい項目にせずその項目の探す見出しに足す。
    """
    base = row.get("base_name") or ""
    if base:
        same = next((r for r in field_rows
                     if r["field_name"] == base and r["data_type"] == row["data_type"]), None)
        if same is not None:
            return same
    if row["data_type"] != "table":
        return None
    columns = _column_norms(row)
    if not columns:
        return None
    for other in field_rows:
        if other["data_type"] != "table":
            continue
        have = _column_norms(other)
        if have and len(columns & have) / len(columns | have) >= SAME_COLUMNS_RATIO:
            return other
    return None


def merge_labels(target: dict, row: dict) -> None:
    """同じ欄と分かった項目に、クリックした見出しを足す（値や型は変えない）。"""
    labels = target["candidates"].splitlines() + row["candidates"].splitlines()
    target["candidates"] = "\n".join(dict.fromkeys(l for l in labels if l.strip()))
    if row["data_type"] == "table":
        columns = target["table_columns"].splitlines() + row["table_columns"].splitlines()
        target["table_columns"] = "\n".join(dict.fromkeys(c for c in columns if c.strip()))


def _column_norms(row: dict) -> set[str]:
    return {normalize_label(c) for c in (row.get("table_columns") or "").splitlines()} - {""}
